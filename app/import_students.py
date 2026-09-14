from __future__ import annotations

import argparse
import csv
import math
import re
import shutil
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import func, select
from sqlalchemy.orm import selectinload

from app.database import DATABASE_PATH, SessionLocal
from app.models import (
    Classroom,
    ImportBatch,
    School,
    SchoolYear,
    Student,
    StudentEnrollment,
    User,
)


PROJECT_DIR = Path(__file__).resolve().parent.parent
EXCEL_PATH = PROJECT_DIR / "uploads" / "hoc_sinh_hoa_sen_2025_2026.xlsx"
SHEET_NAME = "Sheet1"
SCHOOL_YEAR_CODE = "2025-2026"
BACKUP_DIR = PROJECT_DIR / "data" / "backups"
EXPORT_DIR = PROJECT_DIR / "exports"


HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "school_name": ("Trường",),
    "class_group": ("Nhóm/Lớp",),
    "class_name": ("Lớp",),
    "student_code": ("Mã định danh Bộ GD&ĐT", "Mã học sinh"),
    "full_name": ("Họ tên", "Họ và tên"),
    "date_of_birth": ("Ngày sinh",),
    "gender": ("Giới tính",),
    "ethnic_group": ("Dân tộc",),
    "status": ("Trạng thái",),
    "province": ("Tỉnh/Thành phố",),
    "commune_name": ("Xã/Phường",),
    "birth_place": ("Nơi sinh",),
    "current_address": ("Chỗ ở hiện nay",),
    "hamlet": ("Thôn/Xóm",),
    "contact_phone": ("SĐT liên hệ", "Số điện thoại"),
    "area": ("Khu vực",),
    "disability_type": ("Loại khuyết tật",),
    "policy_object": ("Đối tượng chính sách",),
    "tuition_exemption": ("Miễn học phí",),
    "tuition_reduction": ("Giảm học phí",),
    "study_support": ("Hỗ trợ chi phí học tập",),
    "lunch_support": ("Hỗ trợ ăn trưa",),
    "citizen_id": ("Số CCCD",),
    "personal_id": ("Số định danh cá nhân",),
    "father_name": ("Tên cha",),
    "father_job": ("Nghề nghiệp cha",),
    "father_birth_year": ("Năm sinh cha",),
    "mother_name": ("Tên mẹ",),
    "mother_job": ("Nghề nghiệp mẹ",),
    "mother_birth_year": ("Năm sinh mẹ",),
    "representative_name": ("Tên người Đ.Đầu",),
    "guardian_name": ("Tên người đỡ đầu",),
}

REQUIRED_FIELDS = {
    "school_name",
    "class_name",
    "student_code",
    "full_name",
    "date_of_birth",
}

STATUS_MAP = {
    "dang hoc": "DANG_HOC",
    "chuyen den ky 1": "CHUYEN_DEN_KY_1",
    "chuyen den ky 2": "CHUYEN_DEN_KY_2",
    "chuyen den": "CHUYEN_DEN",
    "chuyen di": "CHUYEN_DI",
    "thoi hoc": "THOI_HOC",
}


try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, OSError):
    pass


def bo_dau(text: str) -> str:
    normalized = unicodedata.normalize("NFD", text)
    return "".join(
        character
        for character in normalized
        if unicodedata.category(character) != "Mn"
    )


def khoa_so_sanh(value: Any) -> str:
    if value is None:
        return ""

    text = bo_dau(str(value)).lower().strip()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def chuan_hoa_tieu_de(value: Any) -> str:
    if value is None:
        return ""

    return " ".join(str(value).strip().split()).lower()


def la_gia_tri_36(value: Any) -> bool:
    if value is None or isinstance(value, bool):
        return False

    if isinstance(value, (int, float)):
        return math.isfinite(float(value)) and float(value) == 36.0

    return str(value).strip() == "36"


def chuan_hoa_van_ban(
    value: Any,
    *,
    coi_36_la_trong: bool = False,
) -> str | None:
    if value is None:
        return None

    if coi_36_la_trong and la_gia_tri_36(value):
        return None

    text = " ".join(str(value).strip().split())
    return text or None


def chuan_hoa_ma(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if not math.isfinite(value):
            return ""

        if value.is_integer():
            return str(int(value))

        return format(value, ".15g")

    text = str(value).strip()

    if re.fullmatch(r"[+-]?\d+\.0+", text):
        return text.split(".")[0]

    return re.sub(r"\s+", "", text)


def chuan_hoa_so_dien_thoai(value: Any) -> str | None:
    text = chuan_hoa_van_ban(
        value,
        coi_36_la_trong=True,
    )

    if text is None:
        return None

    compact = re.sub(r"[\s.()\-]", "", text)

    if compact.startswith("+84"):
        compact = "0" + compact[3:]

    return compact or None


def chuan_hoa_nam_sinh(value: Any) -> int | None:
    if value is None or la_gia_tri_36(value):
        return None

    try:
        year = int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None

    current_year = datetime.now().year

    if 1900 <= year <= current_year:
        return year

    return None


def chuan_hoa_ngay_sinh(
    value: Any,
    workbook_epoch,
) -> date | None:
    if value is None or la_gia_tri_36(value):
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if isinstance(value, (int, float)):
        try:
            converted = from_excel(
                value,
                epoch=workbook_epoch,
            )

            if isinstance(converted, datetime):
                return converted.date()

            if isinstance(converted, date):
                return converted
        except (TypeError, ValueError, OverflowError):
            return None

    text = str(value).strip()

    for date_format in (
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d.%m.%Y",
    ):
        try:
            return datetime.strptime(
                text,
                date_format,
            ).date()
        except ValueError:
            continue

    return None


def chuan_hoa_gioi_tinh(value: Any) -> str | None:
    text = chuan_hoa_van_ban(
        value,
        coi_36_la_trong=True,
    )

    if text is None:
        return None

    key = khoa_so_sanh(text)

    if key == "nam":
        return "Nam"

    if key in {"nu", "gai"}:
        return "Nữ"

    return text


def chuan_hoa_trang_thai(value: Any) -> str:
    text = chuan_hoa_van_ban(
        value,
        coi_36_la_trong=True,
    )

    if text is None:
        return "DANG_HOC"

    key = khoa_so_sanh(text)

    if key in STATUS_MAP:
        return STATUS_MAP[key]

    code = re.sub(r"[^A-Z0-9]+", "_", bo_dau(text).upper())
    return code.strip("_") or "DANG_HOC"


def ghep_dia_chi(*parts: str | None) -> str | None:
    values: list[str] = []

    for part in parts:
        if part and part not in values:
            values.append(part)

    return ", ".join(values) if values else None


def tao_ghi_chu(record: dict[str, Any]) -> str | None:
    note_fields = (
        ("Khu vực", "area"),
        ("Đối tượng chính sách", "policy_object"),
        ("Miễn học phí", "tuition_exemption"),
        ("Giảm học phí", "tuition_reduction"),
        ("Hỗ trợ chi phí học tập", "study_support"),
        ("Hỗ trợ ăn trưa", "lunch_support"),
        ("Tên người đại diện", "representative_name"),
        ("Tên người đỡ đầu", "guardian_name"),
    )

    notes: list[str] = []

    for label, key in note_fields:
        value = record.get(key)

        if value:
            notes.append(f"{label}: {value}")

    return "; ".join(notes) if notes else None


def tim_dong_tieu_de(worksheet) -> int | None:
    required_header_keys = {
        chuan_hoa_tieu_de("Mã định danh Bộ GD&ĐT"),
        chuan_hoa_tieu_de("Họ tên"),
    }

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=min(20, worksheet.max_row),
            values_only=True,
        ),
        start=1,
    ):
        row_keys = {
            chuan_hoa_tieu_de(value)
            for value in row
            if value is not None
        }

        if required_header_keys.issubset(row_keys):
            return row_number

    return None


def tao_ban_do_cot(header_values: tuple[Any, ...]) -> dict[str, int]:
    normalized_headers = {
        chuan_hoa_tieu_de(value): index
        for index, value in enumerate(header_values)
        if value is not None
    }

    column_map: dict[str, int] = {}

    for field_name, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            alias_key = chuan_hoa_tieu_de(alias)

            if alias_key in normalized_headers:
                column_map[field_name] = normalized_headers[alias_key]
                break

    return column_map


def lay_gia_tri(
    row: tuple[Any, ...],
    column_map: dict[str, int],
    field_name: str,
) -> Any:
    index = column_map.get(field_name)

    if index is None or index >= len(row):
        return None

    return row[index]


def tao_ban_ghi(
    row: tuple[Any, ...],
    row_number: int,
    column_map: dict[str, int],
    workbook_epoch,
) -> dict[str, Any]:
    raw: dict[str, Any] = {
        field_name: lay_gia_tri(
            row,
            column_map,
            field_name,
        )
        for field_name in HEADER_ALIASES
    }

    province = chuan_hoa_van_ban(
        raw["province"],
        coi_36_la_trong=True,
    )
    commune_name = chuan_hoa_van_ban(
        raw["commune_name"],
        coi_36_la_trong=True,
    )
    hamlet = chuan_hoa_van_ban(
        raw["hamlet"],
        coi_36_la_trong=True,
    )

    personal_id = chuan_hoa_van_ban(
        raw["personal_id"],
        coi_36_la_trong=True,
    )

    if personal_id is None:
        personal_id = chuan_hoa_van_ban(
            raw["citizen_id"],
            coi_36_la_trong=True,
        )

    record: dict[str, Any] = {
        "row_number": row_number,
        "school_name": chuan_hoa_van_ban(raw["school_name"]),
        "class_group": chuan_hoa_van_ban(
            raw["class_group"],
            coi_36_la_trong=True,
        ),
        "class_name": chuan_hoa_van_ban(raw["class_name"]),
        "student_code": chuan_hoa_ma(raw["student_code"]),
        "full_name": chuan_hoa_van_ban(raw["full_name"]),
        "date_of_birth": chuan_hoa_ngay_sinh(
            raw["date_of_birth"],
            workbook_epoch,
        ),
        "gender": chuan_hoa_gioi_tinh(raw["gender"]),
        "ethnic_group": chuan_hoa_van_ban(
            raw["ethnic_group"],
            coi_36_la_trong=True,
        ),
        "status": chuan_hoa_trang_thai(raw["status"]),
        "birth_place": chuan_hoa_van_ban(
            raw["birth_place"],
            coi_36_la_trong=True,
        ),
        "permanent_address": ghep_dia_chi(
            province,
            commune_name,
            hamlet,
        ),
        "current_address": chuan_hoa_van_ban(
            raw["current_address"],
            coi_36_la_trong=True,
        ),
        "contact_phone": chuan_hoa_so_dien_thoai(
            raw["contact_phone"]
        ),
        "disability_type": chuan_hoa_van_ban(
            raw["disability_type"],
            coi_36_la_trong=True,
        ),
        "personal_id": personal_id,
        "father_name": chuan_hoa_van_ban(
            raw["father_name"],
            coi_36_la_trong=True,
        ),
        "father_job": chuan_hoa_van_ban(
            raw["father_job"],
            coi_36_la_trong=True,
        ),
        "father_birth_year": chuan_hoa_nam_sinh(
            raw["father_birth_year"]
        ),
        "mother_name": chuan_hoa_van_ban(
            raw["mother_name"],
            coi_36_la_trong=True,
        ),
        "mother_job": chuan_hoa_van_ban(
            raw["mother_job"],
            coi_36_la_trong=True,
        ),
        "mother_birth_year": chuan_hoa_nam_sinh(
            raw["mother_birth_year"]
        ),
        "area": chuan_hoa_van_ban(
            raw["area"],
            coi_36_la_trong=True,
        ),
        "policy_object": chuan_hoa_van_ban(
            raw["policy_object"],
            coi_36_la_trong=True,
        ),
        "tuition_exemption": chuan_hoa_van_ban(
            raw["tuition_exemption"],
            coi_36_la_trong=True,
        ),
        "tuition_reduction": chuan_hoa_van_ban(
            raw["tuition_reduction"],
            coi_36_la_trong=True,
        ),
        "study_support": chuan_hoa_van_ban(
            raw["study_support"],
            coi_36_la_trong=True,
        ),
        "lunch_support": chuan_hoa_van_ban(
            raw["lunch_support"],
            coi_36_la_trong=True,
        ),
        "representative_name": chuan_hoa_van_ban(
            raw["representative_name"],
            coi_36_la_trong=True,
        ),
        "guardian_name": chuan_hoa_van_ban(
            raw["guardian_name"],
            coi_36_la_trong=True,
        ),
    }

    record["notes"] = tao_ghi_chu(record)
    return record


def them_van_de(
    issues: list[dict[str, Any]],
    level: str,
    row_number: int,
    student_code: str,
    full_name: str | None,
    message: str,
) -> None:
    issues.append(
        {
            "level": level,
            "row_number": row_number,
            "student_code": student_code,
            "full_name": full_name or "",
            "message": message,
        }
    )


def doc_va_kiem_tra_excel() -> dict[str, Any]:
    result: dict[str, Any] = {
        "records": [],
        "errors": [],
        "warnings": [],
        "header_row": None,
        "total_rows": 0,
        "school_count": 0,
        "class_count": 0,
        "existing_student_count": 0,
    }

    if not EXCEL_PATH.exists():
        them_van_de(
            result["errors"],
            "LỖI",
            0,
            "",
            "",
            f"Không tìm thấy file: {EXCEL_PATH}",
        )
        return result

    workbook = None

    try:
        workbook = load_workbook(
            EXCEL_PATH,
            read_only=True,
            data_only=True,
        )

        if SHEET_NAME not in workbook.sheetnames:
            them_van_de(
                result["errors"],
                "LỖI",
                0,
                "",
                "",
                f"Không tìm thấy trang tính '{SHEET_NAME}'.",
            )
            return result

        worksheet = workbook[SHEET_NAME]
        header_row = tim_dong_tieu_de(worksheet)
        result["header_row"] = header_row

        if header_row is None:
            them_van_de(
                result["errors"],
                "LỖI",
                0,
                "",
                "",
                "Không tìm thấy dòng tiêu đề học sinh.",
            )
            return result

        header_values = next(
            worksheet.iter_rows(
                min_row=header_row,
                max_row=header_row,
                values_only=True,
            )
        )
        column_map = tao_ban_do_cot(header_values)

        missing_fields = REQUIRED_FIELDS - set(column_map)

        if missing_fields:
            them_van_de(
                result["errors"],
                "LỖI",
                header_row,
                "",
                "",
                "Thiếu các cột bắt buộc: "
                + ", ".join(sorted(missing_fields)),
            )
            return result

        seen_codes: dict[str, int] = {}

        for row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=header_row + 1,
                values_only=True,
            ),
            start=header_row + 1,
        ):
            if all(value is None for value in row):
                continue

            record = tao_ban_ghi(
                row,
                row_number,
                column_map,
                workbook.epoch,
            )

            if not any(
                (
                    record["student_code"],
                    record["full_name"],
                    record["school_name"],
                    record["class_name"],
                )
            ):
                continue

            result["total_rows"] += 1
            code = record["student_code"]
            name = record["full_name"]

            row_has_error = False

            if not code:
                them_van_de(
                    result["errors"],
                    "LỖI",
                    row_number,
                    code,
                    name,
                    "Thiếu mã học sinh.",
                )
                row_has_error = True

            if not name:
                them_van_de(
                    result["errors"],
                    "LỖI",
                    row_number,
                    code,
                    name,
                    "Thiếu họ tên học sinh.",
                )
                row_has_error = True

            if record["date_of_birth"] is None:
                them_van_de(
                    result["errors"],
                    "LỖI",
                    row_number,
                    code,
                    name,
                    "Ngày sinh không hợp lệ hoặc bị trống.",
                )
                row_has_error = True

            if not record["school_name"]:
                them_van_de(
                    result["errors"],
                    "LỖI",
                    row_number,
                    code,
                    name,
                    "Thiếu tên trường.",
                )
                row_has_error = True

            if not record["class_name"]:
                them_van_de(
                    result["errors"],
                    "LỖI",
                    row_number,
                    code,
                    name,
                    "Thiếu tên lớp.",
                )
                row_has_error = True

            if code in seen_codes:
                them_van_de(
                    result["errors"],
                    "LỖI",
                    row_number,
                    code,
                    name,
                    f"Mã học sinh trùng với dòng {seen_codes[code]}.",
                )
                row_has_error = True
            elif code:
                seen_codes[code] = row_number

            dob = record["date_of_birth"]

            if dob is not None and dob.year < 2020:
                them_van_de(
                    result["warnings"],
                    "CẢNH BÁO",
                    row_number,
                    code,
                    name,
                    f"Năm sinh {dob.year} cần nhà trường kiểm tra.",
                )

            phone = record["contact_phone"]

            if phone and not re.fullmatch(r"\d{9,12}", phone):
                them_van_de(
                    result["warnings"],
                    "CẢNH BÁO",
                    row_number,
                    code,
                    name,
                    f"Số điện thoại '{phone}' có định dạng chưa chuẩn.",
                )

            if not row_has_error:
                result["records"].append(record)

    except Exception as error:
        them_van_de(
            result["errors"],
            "LỖI",
            0,
            "",
            "",
            f"Không đọc được file Excel: {type(error).__name__}: {error}",
        )

    finally:
        if workbook is not None:
            workbook.close()

    return result


def lay_ma_xa_tu_ten_file() -> str | None:
    match = re.match(r"^(\d+)_", EXCEL_PATH.stem)
    return match.group(1) if match else None


def doi_chieu_co_so_du_lieu(data: dict[str, Any]) -> None:
    if data["errors"]:
        return

    with SessionLocal() as db:
        school_year = db.scalar(
            select(SchoolYear).where(
                SchoolYear.code == SCHOOL_YEAR_CODE
            )
        )

        if school_year is None:
            them_van_de(
                data["errors"],
                "LỖI",
                0,
                "",
                "",
                f"Chưa có năm học {SCHOOL_YEAR_CODE} trong cơ sở dữ liệu.",
            )
            return

        schools = db.scalars(
            select(School).options(
                selectinload(School.commune)
            )
        ).all()

        schools_by_name: dict[str, list[School]] = defaultdict(list)

        for school in schools:
            schools_by_name[khoa_so_sanh(school.name)].append(school)

        file_commune_code = lay_ma_xa_tu_ten_file()
        resolved_school_ids: set[int] = set()
        class_keys: set[tuple[int, str]] = set()

        for record in data["records"]:
            matches = schools_by_name.get(
                khoa_so_sanh(record["school_name"]),
                [],
            )

            if len(matches) > 1 and file_commune_code:
                matches = [
                    school
                    for school in matches
                    if school.commune.code == file_commune_code
                ]

            if len(matches) == 0:
                them_van_de(
                    data["errors"],
                    "LỖI",
                    record["row_number"],
                    record["student_code"],
                    record["full_name"],
                    "Không tìm thấy trường trong danh mục: "
                    f"'{record['school_name']}'.",
                )
                continue

            if len(matches) > 1:
                them_van_de(
                    data["errors"],
                    "LỖI",
                    record["row_number"],
                    record["student_code"],
                    record["full_name"],
                    "Tên trường trùng ở nhiều xã; cần đặt mã xã ở đầu tên file.",
                )
                continue

            school = matches[0]
            record["school_id"] = school.id
            record["school_code"] = school.code
            record["commune_code"] = school.commune.code
            record["commune_id"] = school.commune_id

            resolved_school_ids.add(school.id)
            class_keys.add(
                (
                    school.id,
                    khoa_so_sanh(record["class_name"]),
                )
            )

        codes = [
            record["student_code"]
            for record in data["records"]
            if record.get("school_id") is not None
        ]

        existing_count = 0

        if codes:
            existing_count = db.scalar(
                select(func.count(Student.id)).where(
                    Student.code.in_(codes)
                )
            ) or 0

        data["school_count"] = len(resolved_school_ids)
        data["class_count"] = len(class_keys)
        data["existing_student_count"] = existing_count


def xuat_bao_cao(data: dict[str, Any]) -> Path | None:
    issues = data["errors"] + data["warnings"]

    if not issues:
        return None

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = EXPORT_DIR / f"bao_cao_kiem_tra_hoc_sinh_{timestamp}.csv"

    with path.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        writer = csv.writer(file)
        writer.writerow(
            [
                "Muc_do",
                "Dong_Excel",
                "Ma_hoc_sinh",
                "Ho_ten",
                "Noi_dung",
            ]
        )

        for issue in issues:
            writer.writerow(
                [
                    issue["level"],
                    issue["row_number"],
                    issue["student_code"],
                    issue["full_name"],
                    issue["message"],
                ]
            )

    return path


def in_ket_qua_kiem_tra(data: dict[str, Any]) -> None:
    print()
    print("=" * 72)
    print("KẾT QUẢ KIỂM TRA FILE HỌC SINH")
    print("=" * 72)
    print(f"File: {EXCEL_PATH}")
    print(f"Trang tính: {SHEET_NAME}")
    print(f"Dòng tiêu đề: {data['header_row']}")
    print(f"Số dòng dữ liệu đã đọc: {data['total_rows']}")
    print(f"Số dòng hợp lệ: {len(data['records'])}")
    print(f"Số trường được nhận diện: {data['school_count']}")
    print(f"Số lớp được nhận diện: {data['class_count']}")
    print(
        "Học sinh đã tồn tại trong cơ sở dữ liệu: "
        f"{data['existing_student_count']}"
    )
    print(f"Số lỗi: {len(data['errors'])}")
    print(f"Số cảnh báo: {len(data['warnings'])}")

    if data["errors"]:
        print()
        print("CÁC LỖI ĐẦU TIÊN:")

        for issue in data["errors"][:20]:
            print(
                f"- Dòng {issue['row_number']}: "
                f"{issue['message']}"
            )

    if data["warnings"]:
        print()
        print("CÁC CẢNH BÁO ĐẦU TIÊN:")

        for issue in data["warnings"][:20]:
            print(
                f"- Dòng {issue['row_number']}: "
                f"{issue['message']}"
            )

    report_path = xuat_bao_cao(data)

    if report_path is not None:
        print()
        print("Báo cáo kiểm tra đã được lưu tại:")
        print(report_path)

    print("=" * 72)


def sao_luu_co_so_du_lieu() -> Path:
    if not DATABASE_PATH.exists():
        raise FileNotFoundError(
            f"Không tìm thấy cơ sở dữ liệu: {DATABASE_PATH}"
        )

    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = (
        BACKUP_DIR
        / f"phocap_before_student_import_{timestamp}.db"
    )
    shutil.copy2(DATABASE_PATH, backup_path)
    return backup_path


def tao_student(record: dict[str, Any]) -> Student:
    return Student(
        code=record["student_code"],
        full_name=record["full_name"],
        date_of_birth=record["date_of_birth"],
        gender=record["gender"],
        ethnic_group=record["ethnic_group"],
        birth_place=record["birth_place"],
        permanent_address=record["permanent_address"],
        current_address=record["current_address"],
        father_name=record["father_name"],
        father_birth_year=record["father_birth_year"],
        father_job=record["father_job"],
        mother_name=record["mother_name"],
        mother_birth_year=record["mother_birth_year"],
        mother_job=record["mother_job"],
        contact_phone=record["contact_phone"],
        personal_id=record["personal_id"],
        disability_type=record["disability_type"],
        notes=record["notes"],
        is_active=True,
    )


def cap_nhat_import_batch(
    batch_id: int,
    *,
    status: str,
    inserted_rows: int = 0,
    updated_rows: int = 0,
    skipped_rows: int = 0,
    error_rows: int = 0,
    error_report_path: str | None = None,
) -> None:
    with SessionLocal() as db:
        batch = db.get(ImportBatch, batch_id)

        if batch is None:
            return

        batch.status = status
        batch.inserted_rows = inserted_rows
        batch.updated_rows = updated_rows
        batch.skipped_rows = skipped_rows
        batch.error_rows = error_rows
        batch.error_report_path = error_report_path
        batch.completed_at = datetime.now()
        db.commit()


def nhap_du_lieu(data: dict[str, Any]) -> bool:
    if data["errors"]:
        print("Không thể nhập vì file vẫn còn lỗi.")
        return False

    backup_path = sao_luu_co_so_du_lieu()
    report_path = xuat_bao_cao(data)

    with SessionLocal() as db:
        creator = db.scalar(
            select(User).where(
                User.username == "admin.sogddt"
            )
        )

        batch = ImportBatch(
            file_name=EXCEL_PATH.name,
            template_version="1.0",
            operation_type="NHAP_HOC_SINH_BAN_DAU",
            status="DANG_XU_LY",
            total_rows=data["total_rows"],
            valid_rows=len(data["records"]),
            error_rows=len(data["errors"]),
            error_report_path=(
                str(report_path)
                if report_path is not None
                else None
            ),
            created_by_user_id=(
                creator.id if creator is not None else None
            ),
        )
        db.add(batch)
        db.commit()
        db.refresh(batch)
        batch_id = batch.id

    inserted_students = 0
    inserted_enrollments = 0
    skipped_rows = 0
    runtime_warnings: list[dict[str, Any]] = []

    try:
        with SessionLocal() as db:
            school_year = db.scalar(
                select(SchoolYear).where(
                    SchoolYear.code == SCHOOL_YEAR_CODE
                )
            )

            if school_year is None:
                raise RuntimeError(
                    f"Không tìm thấy năm học {SCHOOL_YEAR_CODE}."
                )

            school_ids = {
                record["school_id"]
                for record in data["records"]
            }

            classrooms = db.scalars(
                select(Classroom).where(
                    Classroom.school_year_id == school_year.id,
                    Classroom.school_id.in_(school_ids),
                )
            ).all()

            classroom_map: dict[tuple[int, str], Classroom] = {
                (
                    classroom.school_id,
                    khoa_so_sanh(classroom.name),
                ): classroom
                for classroom in classrooms
            }

            for record in data["records"]:
                key = (
                    record["school_id"],
                    khoa_so_sanh(record["class_name"]),
                )

                if key not in classroom_map:
                    classroom = Classroom(
                        school_id=record["school_id"],
                        school_year_id=school_year.id,
                        code=record["class_group"],
                        name=record["class_name"],
                        is_active=True,
                    )
                    db.add(classroom)
                    classroom_map[key] = classroom

            db.flush()

            codes = [
                record["student_code"]
                for record in data["records"]
            ]

            existing_students = db.scalars(
                select(Student).where(
                    Student.code.in_(codes)
                )
            ).all()

            student_map = {
                student.code: student
                for student in existing_students
            }

            for record in data["records"]:
                code = record["student_code"]

                if code not in student_map:
                    student = tao_student(record)
                    db.add(student)
                    student_map[code] = student
                    inserted_students += 1

            db.flush()

            student_ids = [
                student.id
                for student in student_map.values()
            ]

            existing_enrollments = db.scalars(
                select(StudentEnrollment).where(
                    StudentEnrollment.school_year_id
                    == school_year.id,
                    StudentEnrollment.student_id.in_(student_ids),
                )
            ).all()

            enrollment_by_student: dict[
                int,
                list[StudentEnrollment],
            ] = defaultdict(list)

            for enrollment in existing_enrollments:
                enrollment_by_student[
                    enrollment.student_id
                ].append(enrollment)

            for record in data["records"]:
                student = student_map[record["student_code"]]
                classroom = classroom_map[
                    (
                        record["school_id"],
                        khoa_so_sanh(record["class_name"]),
                    )
                ]

                student_enrollments = enrollment_by_student.get(
                    student.id,
                    [],
                )

                same_enrollment = next(
                    (
                        enrollment
                        for enrollment in student_enrollments
                        if enrollment.school_id == record["school_id"]
                        and enrollment.class_id == classroom.id
                    ),
                    None,
                )

                if same_enrollment is not None:
                    skipped_rows += 1
                    continue

                other_current_enrollment = next(
                    (
                        enrollment
                        for enrollment in student_enrollments
                        if enrollment.is_current
                    ),
                    None,
                )

                if other_current_enrollment is not None:
                    skipped_rows += 1
                    them_van_de(
                        runtime_warnings,
                        "CẢNH BÁO",
                        record["row_number"],
                        record["student_code"],
                        record["full_name"],
                        "Học sinh đã có lớp khác trong cùng năm học; "
                        "không tự động thay đổi.",
                    )
                    continue

                enrollment = StudentEnrollment(
                    student_id=student.id,
                    school_year_id=school_year.id,
                    school_id=record["school_id"],
                    class_id=classroom.id,
                    status=record["status"],
                    enrollment_date=None,
                    transfer_date=None,
                    is_current=True,
                )
                db.add(enrollment)
                inserted_enrollments += 1

            db.commit()

        if runtime_warnings:
            data["warnings"].extend(runtime_warnings)
            report_path = xuat_bao_cao(data)

        cap_nhat_import_batch(
            batch_id,
            status="HOAN_THANH",
            inserted_rows=inserted_students,
            updated_rows=0,
            skipped_rows=skipped_rows,
            error_rows=0,
            error_report_path=(
                str(report_path)
                if report_path is not None
                else None
            ),
        )

    except Exception as error:
        cap_nhat_import_batch(
            batch_id,
            status="THAT_BAI",
            inserted_rows=inserted_students,
            updated_rows=0,
            skipped_rows=skipped_rows,
            error_rows=1,
            error_report_path=(
                str(report_path)
                if report_path is not None
                else None
            ),
        )

        print()
        print("=" * 72)
        print("NHẬP HỌC SINH KHÔNG THÀNH CÔNG")
        print("=" * 72)
        print(f"{type(error).__name__}: {error}")
        print("Bản sao lưu cơ sở dữ liệu:")
        print(backup_path)
        return False

    with SessionLocal() as db:
        total_students = db.scalar(
            select(func.count(Student.id))
        ) or 0
        total_classes = db.scalar(
            select(func.count(Classroom.id)).where(
                Classroom.school_year_id
                == select(SchoolYear.id)
                .where(SchoolYear.code == SCHOOL_YEAR_CODE)
                .scalar_subquery()
            )
        ) or 0
        total_enrollments = db.scalar(
            select(func.count(StudentEnrollment.id)).where(
                StudentEnrollment.school_year_id
                == select(SchoolYear.id)
                .where(SchoolYear.code == SCHOOL_YEAR_CODE)
                .scalar_subquery()
            )
        ) or 0

    print()
    print("=" * 72)
    print("NHẬP HỌC SINH THÀNH CÔNG")
    print("=" * 72)
    print(f"Học sinh mới: {inserted_students}")
    print(f"Hồ sơ lớp học mới: {inserted_enrollments}")
    print(f"Dòng đã tồn tại và bỏ qua: {skipped_rows}")
    print(f"Tổng học sinh trong hệ thống: {total_students}")
    print(f"Tổng lớp của năm học {SCHOOL_YEAR_CODE}: {total_classes}")
    print(
        f"Tổng hồ sơ học sinh năm học {SCHOOL_YEAR_CODE}: "
        f"{total_enrollments}"
    )
    print("Bản sao lưu cơ sở dữ liệu:")
    print(backup_path)

    if report_path is not None:
        print("Báo cáo cảnh báo:")
        print(report_path)

    print("=" * 72)
    return True


def tao_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Kiểm tra hoặc nhập học sinh từ file Excel "
            "vào phần mềm phổ cập giáo dục."
        )
    )
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument(
        "--check",
        action="store_true",
        help="Chỉ kiểm tra file, không thay đổi cơ sở dữ liệu.",
    )
    mode.add_argument(
        "--import",
        dest="do_import",
        action="store_true",
        help="Sao lưu và nhập dữ liệu học sinh.",
    )
    return parser


def main() -> None:
    args = tao_parser().parse_args()
    data = doc_va_kiem_tra_excel()
    doi_chieu_co_so_du_lieu(data)
    in_ket_qua_kiem_tra(data)

    if args.check:
        if data["errors"]:
            print("KẾT LUẬN: File chưa đạt yêu cầu để nhập.")
            raise SystemExit(1)

        print("KẾT LUẬN: File hợp lệ, có thể tiến hành nhập.")
        return

    if args.do_import:
        if data["errors"]:
            print("Dừng nhập vì file vẫn còn lỗi.")
            raise SystemExit(1)

        if not nhap_du_lieu(data):
            raise SystemExit(1)


if __name__ == "__main__":
    main()
