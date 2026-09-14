from __future__ import annotations

import argparse
import csv
import math
import re
import shutil
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete, func, or_, select
from sqlalchemy.exc import SQLAlchemyError

from app.database import DATABASE_PATH, SessionLocal
from app.models import Commune, School, User


# =========================================================
# THIẾT LẬP ĐƯỜNG DẪN
# =========================================================

PROJECT_DIR = Path(__file__).resolve().parent.parent

EXCEL_PATH = (
    PROJECT_DIR
    / "uploads"
    / "PHO_CAP_GIAO_DUC.xlsx"
)

SHEET_NAME = "Danhmuc"

BACKUP_DIR = PROJECT_DIR / "data" / "backups"

EXPORT_DIR = PROJECT_DIR / "exports"


# Bốn cột bắt buộc trong trang tính Danhmuc
EXPECTED_HEADERS = (
    "ID_Xa",
    "Ten_Xa",
    "ID_Truong",
    "Ten_Truong",
)


# Thiết lập hiển thị tiếng Việt trong Terminal Windows
try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, OSError):
    pass


# =========================================================
# CÁC HÀM CHUẨN HÓA DỮ LIỆU
# =========================================================

def la_o_trong(value: Any) -> bool:
    """
    Kiểm tra ô Excel có đang trống hay không.
    """

    if value is None:
        return True

    return str(value).strip() == ""


def chuan_hoa_ma(value: Any) -> str:
    """
    Chuyển mã xã hoặc mã trường thành chuỗi văn bản.

    Hàm này:
    - Giữ được mã có chữ như 40412L85.
    - Giữ được mã có khoảng trắng như MN Happy Home.
    - Chuyển số 16681.0 thành 16681.
    - Không ép mã thành số nguyên khi mã chứa chữ.
    """

    if value is None:
        return ""

    if isinstance(value, bool):
        return ""

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if not math.isfinite(value):
            return ""

        if value.is_integer():
            return str(int(value))

        return format(value, ".15g")

    text = str(value).strip()

    # Thu gọn nhiều khoảng trắng liên tiếp thành một khoảng trắng
    text = re.sub(r"\s+", " ", text)

    # Trường hợp Excel trả về dạng văn bản 16681.0
    if re.fullmatch(r"[+-]?\d+\.0+", text):
        return text.split(".")[0]

    return text


def chuan_hoa_ten(value: Any) -> str:
    """
    Chuẩn hóa tên xã và tên trường:
    - Xóa khoảng trắng đầu và cuối.
    - Thu gọn nhiều khoảng trắng liên tiếp.
    """

    if value is None:
        return ""

    return " ".join(str(value).strip().split())


# =========================================================
# ĐỌC VÀ KIỂM TRA FILE EXCEL
# =========================================================

def doc_du_lieu_excel() -> dict[str, Any]:
    """
    Đọc trang tính Danhmuc và kiểm tra toàn bộ dữ liệu.

    Kết quả trả về gồm:
    - Danh sách xã.
    - Danh sách trường.
    - Danh sách lỗi.
    - Các thông tin thống kê.
    """

    ket_qua: dict[str, Any] = {
        "communes": [],
        "schools": [],
        "errors": [],
        "duplicate_school_names": {},
        "non_numeric_school_codes": 0,
        "total_data_rows": 0,
    }

    if not EXCEL_PATH.exists():
        ket_qua["errors"].append(
            {
                "row": 0,
                "message": (
                    "Không tìm thấy file Excel tại: "
                    f"{EXCEL_PATH}"
                ),
                "values": [],
            }
        )

        return ket_qua

    workbook = None

    try:
        workbook = load_workbook(
            filename=EXCEL_PATH,
            read_only=True,
            data_only=True,
        )

        if SHEET_NAME not in workbook.sheetnames:
            ket_qua["errors"].append(
                {
                    "row": 0,
                    "message": (
                        f"Không tìm thấy trang tính "
                        f"'{SHEET_NAME}'."
                    ),
                    "values": [],
                }
            )

            return ket_qua

        worksheet = workbook[SHEET_NAME]

        header_row = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                min_col=1,
                max_col=4,
                values_only=True,
            )
        )

        actual_headers = tuple(
            str(value).strip()
            if value is not None
            else ""
            for value in header_row
        )

        if actual_headers != EXPECTED_HEADERS:
            ket_qua["errors"].append(
                {
                    "row": 1,
                    "message": (
                        "Tên cột không đúng. "
                        f"Yêu cầu: {EXPECTED_HEADERS}. "
                        f"Thực tế: {actual_headers}."
                    ),
                    "values": list(header_row),
                }
            )

            return ket_qua

        commune_by_code: dict[str, str] = {}
        commune_code_by_name: dict[str, str] = {}

        school_by_code: dict[
            str,
            tuple[str, str],
        ] = {}

        commune_order: list[str] = []

        school_name_to_codes: dict[
            str,
            set[str],
        ] = defaultdict(set)

        for excel_row_number, row in enumerate(
            worksheet.iter_rows(
                min_row=2,
                min_col=1,
                max_col=4,
                values_only=True,
            ),
            start=2,
        ):
            # Bỏ qua dòng hoàn toàn trống
            if all(la_o_trong(value) for value in row):
                continue

            ket_qua["total_data_rows"] += 1

            raw_commune_code = row[0]
            raw_commune_name = row[1]
            raw_school_code = row[2]
            raw_school_name = row[3]

            commune_code = chuan_hoa_ma(
                raw_commune_code
            )

            commune_name = chuan_hoa_ten(
                raw_commune_name
            )

            school_code = chuan_hoa_ma(
                raw_school_code
            )

            school_name = chuan_hoa_ten(
                raw_school_name
            )

            row_errors: list[str] = []

            if not commune_code:
                row_errors.append("Thiếu mã xã.")

            if not commune_name:
                row_errors.append("Thiếu tên xã.")

            if not school_code:
                row_errors.append("Thiếu mã trường.")

            if not school_name:
                row_errors.append("Thiếu tên trường.")

            if len(commune_code) > 30:
                row_errors.append(
                    "Mã xã dài quá 30 ký tự."
                )

            if len(commune_name) > 200:
                row_errors.append(
                    "Tên xã dài quá 200 ký tự."
                )

            if len(school_code) > 30:
                row_errors.append(
                    "Mã trường dài quá 30 ký tự."
                )

            if len(school_name) > 300:
                row_errors.append(
                    "Tên trường dài quá 300 ký tự."
                )

            # Một mã xã không được đi với hai tên xã khác nhau
            existing_commune_name = commune_by_code.get(
                commune_code
            )

            if (
                commune_code
                and existing_commune_name is not None
                and existing_commune_name != commune_name
            ):
                row_errors.append(
                    "Mã xã đã xuất hiện với một tên xã khác: "
                    f"'{existing_commune_name}'."
                )

            # Một tên xã không được đi với hai mã khác nhau
            existing_commune_code = (
                commune_code_by_name.get(commune_name)
            )

            if (
                commune_name
                and existing_commune_code is not None
                and existing_commune_code != commune_code
            ):
                row_errors.append(
                    "Tên xã đã xuất hiện với một mã xã khác: "
                    f"'{existing_commune_code}'."
                )

            # Mã trường phải duy nhất trong toàn bộ file
            existing_school = school_by_code.get(
                school_code
            )

            if (
                school_code
                and existing_school is not None
            ):
                old_commune_code, old_school_name = (
                    existing_school
                )

                row_errors.append(
                    "Mã trường bị trùng. "
                    f"Đã xuất hiện tại xã {old_commune_code}, "
                    f"tên trường '{old_school_name}'."
                )

            if row_errors:
                ket_qua["errors"].append(
                    {
                        "row": excel_row_number,
                        "message": " ".join(row_errors),
                        "values": [
                            raw_commune_code,
                            raw_commune_name,
                            raw_school_code,
                            raw_school_name,
                        ],
                    }
                )

                continue

            if commune_code not in commune_by_code:
                commune_by_code[commune_code] = (
                    commune_name
                )

                commune_code_by_name[commune_name] = (
                    commune_code
                )

                commune_order.append(commune_code)

            school_by_code[school_code] = (
                commune_code,
                school_name,
            )

            school_name_to_codes[school_name].add(
                school_code
            )

            if not school_code.isdigit():
                ket_qua[
                    "non_numeric_school_codes"
                ] += 1

        ket_qua["communes"] = [
            {
                "code": commune_code,
                "name": commune_by_code[commune_code],
            }
            for commune_code in commune_order
        ]

        ket_qua["schools"] = [
            {
                "code": school_code,
                "name": school_name,
                "commune_code": commune_code,
            }
            for school_code, (
                commune_code,
                school_name,
            ) in school_by_code.items()
        ]

        ket_qua["duplicate_school_names"] = {
            school_name: sorted(codes)
            for school_name, codes
            in school_name_to_codes.items()
            if len(codes) > 1
        }

        return ket_qua

    except Exception as error:
        ket_qua["errors"].append(
            {
                "row": 0,
                "message": (
                    "Không thể đọc file Excel: "
                    f"{type(error).__name__}: {error}"
                ),
                "values": [],
            }
        )

        return ket_qua

    finally:
        if workbook is not None:
            workbook.close()


# =========================================================
# XUẤT BÁO CÁO LỖI
# =========================================================

def xuat_bao_cao_loi(
    errors: list[dict[str, Any]],
) -> Path | None:
    """
    Xuất các dòng lỗi ra file CSV trong thư mục exports.
    """

    if not errors:
        return None

    EXPORT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    report_path = (
        EXPORT_DIR
        / f"loi_import_danhmuc_{timestamp}.csv"
    )

    with report_path.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        writer = csv.writer(file)

        writer.writerow(
            [
                "Dong_Excel",
                "Loi",
                "ID_Xa",
                "Ten_Xa",
                "ID_Truong",
                "Ten_Truong",
            ]
        )

        for error in errors:
            values = list(error.get("values", []))

            while len(values) < 4:
                values.append("")

            writer.writerow(
                [
                    error.get("row", ""),
                    error.get("message", ""),
                    values[0],
                    values[1],
                    values[2],
                    values[3],
                ]
            )

    return report_path


# =========================================================
# HIỂN THỊ KẾT QUẢ KIỂM TRA
# =========================================================

def in_ket_qua_kiem_tra(
    data: dict[str, Any],
) -> None:
    """
    In kết quả kiểm tra file Excel ra Terminal.
    """

    print()
    print("=" * 70)
    print("KẾT QUẢ KIỂM TRA FILE DANH MỤC")
    print("=" * 70)

    print(f"File: {EXCEL_PATH}")
    print(f"Trang tính: {SHEET_NAME}")
    print(
        "Số dòng dữ liệu đã đọc: "
        f"{data['total_data_rows']}"
    )
    print(
        "Số xã hợp lệ: "
        f"{len(data['communes'])}"
    )
    print(
        "Số trường hợp lệ: "
        f"{len(data['schools'])}"
    )
    print(
        "Mã trường có chứa chữ hoặc ký tự khác số: "
        f"{data['non_numeric_school_codes']}"
    )
    print(
        "Số tên trường xuất hiện nhiều lần: "
        f"{len(data['duplicate_school_names'])}"
    )
    print(
        "Số lỗi phát hiện: "
        f"{len(data['errors'])}"
    )

    if data["duplicate_school_names"]:
        print()
        print(
            "Các tên trường bị trùng tên "
            "(được phép vì mã trường khác nhau):"
        )

        for school_name, codes in (
            data["duplicate_school_names"].items()
        ):
            print(
                f"- {school_name}: "
                f"{', '.join(codes)}"
            )

    if data["errors"]:
        print()
        print("CÁC LỖI ĐẦU TIÊN:")

        for error in data["errors"][:20]:
            print(
                f"- Dòng {error['row']}: "
                f"{error['message']}"
            )

        if len(data["errors"]) > 20:
            print(
                f"... và còn "
                f"{len(data['errors']) - 20} lỗi khác."
            )

        report_path = xuat_bao_cao_loi(
            data["errors"]
        )

        if report_path is not None:
            print()
            print(
                "Đã xuất báo cáo lỗi tại:"
            )
            print(report_path)

    print("=" * 70)


# =========================================================
# SAO LƯU CƠ SỞ DỮ LIỆU
# =========================================================

def sao_luu_co_so_du_lieu() -> Path:
    """
    Sao lưu phocap.db trước khi thay thế danh mục.
    """

    if not DATABASE_PATH.exists():
        raise FileNotFoundError(
            f"Không tìm thấy cơ sở dữ liệu: "
            f"{DATABASE_PATH}"
        )

    BACKUP_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    backup_path = (
        BACKUP_DIR
        / f"phocap_before_import_{timestamp}.db"
    )

    shutil.copy2(
        DATABASE_PATH,
        backup_path,
    )

    return backup_path


# =========================================================
# NHẬP DỮ LIỆU VÀO SQLITE
# =========================================================

def thay_the_danh_muc(
    data: dict[str, Any],
) -> bool:
    """
    Xóa toàn bộ xã và trường thử nghiệm,
    sau đó nhập danh mục chính thức từ Excel.

    Tài khoản Sở không bị xóa.

    Chương trình sẽ dừng nếu đang có tài khoản
    cấp xã, trường hoặc giáo viên gắn với đơn vị.
    """

    if data["errors"]:
        print(
            "Không thể nhập dữ liệu vì file Excel "
            "vẫn còn lỗi."
        )
        return False

    backup_path = sao_luu_co_so_du_lieu()

    print()
    print(
        "Đã sao lưu cơ sở dữ liệu tại:"
    )
    print(backup_path)

    with SessionLocal() as db:
        try:
            assigned_user_count = db.scalar(
                select(func.count(User.id)).where(
                    or_(
                        User.commune_id.is_not(None),
                        User.school_id.is_not(None),
                    )
                )
            )

            assigned_user_count = (
                assigned_user_count or 0
            )

            if assigned_user_count > 0:
                print()
                print("=" * 70)
                print("KHÔNG THỂ THAY THẾ DANH MỤC")
                print("=" * 70)
                print(
                    "Hiện có "
                    f"{assigned_user_count} tài khoản "
                    "đã gắn với xã hoặc trường."
                )
                print(
                    "Cần xử lý các tài khoản đó "
                    "trước khi thay thế danh mục."
                )
                return False

            # Xóa trường trước vì trường đang tham chiếu xã
            db.execute(delete(School))

            # Sau đó mới xóa xã
            db.execute(delete(Commune))

            db.flush()

            commune_models: list[Commune] = []

            for item in data["communes"]:
                commune = Commune(
                    code=item["code"],
                    name=item["name"],
                    is_active=True,
                )

                commune_models.append(commune)

            db.add_all(commune_models)

            # Cần flush để SQLAlchemy tạo ID cho từng xã
            db.flush()

            commune_id_by_code = {
                commune.code: commune.id
                for commune in commune_models
            }

            school_models: list[School] = []

            for item in data["schools"]:
                commune_id = commune_id_by_code.get(
                    item["commune_code"]
                )

                if commune_id is None:
                    raise ValueError(
                        "Không xác định được xã cho trường "
                        f"{item['code']} - {item['name']}."
                    )

                school = School(
                    commune_id=commune_id,
                    code=item["code"],
                    name=item["name"],
                    address=None,
                    is_active=True,
                )

                school_models.append(school)

            db.add_all(school_models)

            db.commit()

        except (
            SQLAlchemyError,
            ValueError,
            KeyError,
        ) as error:
            db.rollback()

            print()
            print("=" * 70)
            print("NHẬP DỮ LIỆU KHÔNG THÀNH CÔNG")
            print("=" * 70)
            print(
                f"{type(error).__name__}: {error}"
            )
            print(
                "Cơ sở dữ liệu đã được hoàn tác."
            )
            print(
                "Bản sao lưu vẫn còn tại:"
            )
            print(backup_path)

            return False

    # Kiểm tra lại sau khi giao dịch đã hoàn thành
    with SessionLocal() as db:
        commune_count = db.scalar(
            select(func.count(Commune.id))
        )

        school_count = db.scalar(
            select(func.count(School.id))
        )

        commune_count = commune_count or 0
        school_count = school_count or 0

    expected_communes = len(data["communes"])
    expected_schools = len(data["schools"])

    if (
        commune_count != expected_communes
        or school_count != expected_schools
    ):
        print()
        print("=" * 70)
        print("CẢNH BÁO: SỐ LƯỢNG SAU NHẬP CHƯA KHỚP")
        print("=" * 70)
        print(
            f"Yêu cầu: {expected_communes} xã, "
            f"{expected_schools} trường."
        )
        print(
            f"Thực tế: {commune_count} xã, "
            f"{school_count} trường."
        )

        return False

    print()
    print("=" * 70)
    print("NHẬP DANH MỤC THÀNH CÔNG")
    print("=" * 70)
    print(
        f"Đã nhập: {commune_count} xã."
    )
    print(
        f"Đã nhập: {school_count} trường/cơ sở."
    )
    print(
        "Tài khoản cấp Sở được giữ nguyên."
    )
    print(
        "Tất cả xã và trường vừa nhập "
        "được đặt trạng thái Đang hoạt động."
    )
    print(
        "Bản sao lưu cơ sở dữ liệu:"
    )
    print(backup_path)
    print("=" * 70)

    return True


# =========================================================
# XỬ LÝ LỆNH
# =========================================================

def tao_bo_phan_tich_lenh() -> argparse.ArgumentParser:
    """
    Tạo hai chế độ chạy:

    --check:
        Chỉ kiểm tra, không thay đổi cơ sở dữ liệu.

    --replace:
        Sao lưu DB, xóa danh mục cũ
        và nhập toàn bộ danh mục mới.
    """

    parser = argparse.ArgumentParser(
        description=(
            "Kiểm tra hoặc nhập danh mục xã, trường "
            "từ file Excel vào phần mềm phổ cập."
        )
    )

    mode_group = parser.add_mutually_exclusive_group(
        required=True
    )

    mode_group.add_argument(
        "--check",
        action="store_true",
        help=(
            "Chỉ kiểm tra file Excel, "
            "không thay đổi cơ sở dữ liệu."
        ),
    )

    mode_group.add_argument(
        "--replace",
        action="store_true",
        help=(
            "Sao lưu DB, xóa danh mục cũ "
            "và nhập danh mục mới."
        ),
    )

    return parser


def main() -> None:
    parser = tao_bo_phan_tich_lenh()
    args = parser.parse_args()

    data = doc_du_lieu_excel()

    in_ket_qua_kiem_tra(data)

    if args.check:
        if data["errors"]:
            print(
                "KẾT LUẬN: File chưa đạt yêu cầu."
            )
            sys.exit(1)

        print(
            "KẾT LUẬN: File hợp lệ, "
            "có thể tiến hành nhập dữ liệu."
        )
        return

    if args.replace:
        if data["errors"]:
            print(
                "Dừng nhập vì file Excel có lỗi."
            )
            sys.exit(1)

        success = thay_the_danh_muc(data)

        if not success:
            sys.exit(1)


if __name__ == "__main__":
    main()