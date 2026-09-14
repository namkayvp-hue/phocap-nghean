from __future__ import annotations

import json
import re
from copy import copy
import shutil
import sqlite3
import unicodedata
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.page import PageMargins
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload, with_loader_criteria

from app.database import DATABASE_PATH
from app.models import Classroom, School, SchoolYear, User
from app.survey_models import (
    Household,
    SurveyFileExchangeLog,
    SurveyForm,
    SurveyFormInvestigator,
    SurveyPerson,
    SurveyPersonYearRecord,
)


APP_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = APP_DIR / "report_templates" / "Bieu_mau_PCGDMN_2025.xlsx"
STAGING_DIR = Path(DATABASE_PATH).parent / "import_staging"
TRACKING_SHEET_NAME = "Sổ theo dõi PCGDMN"
META_SHEET_NAME = "_PCGDMN_META"
EXPORT_VERSION = "PCGDMN_HO_V1"
MULTI_EXPORT_VERSION = "PCGDMN_NHIEU_HO_V1"
MULTI_LAYOUT_VERSION = 1
DATA_START_ROW = 11
DATA_END_ROW = 110
META_TABLE_ROW = 20
MAX_UPLOAD_BYTES = 15 * 1024 * 1024
MULTI_BLOCK_MIN_ROWS = 32
MULTI_EXTRA_MEMBER_ROWS = 5
MULTI_MARKER_COLUMN = 21  # U - ẩn, không nằm trong vùng in
MULTI_JSON_COLUMN = 22    # V - ẩn, không nằm trong vùng in
MULTI_BLOCK_START = "__PCGDMN_BLOCK_START__"
MULTI_PERSON_ROW = "__PCGDMN_PERSON_ROW__"
MULTI_BLOCK_END = "__PCGDMN_BLOCK_END__"


templates = Jinja2Templates(directory=str(APP_DIR / "templates"))


VISIBLE_HEADERS = [
    "STT",
    "Số phiếu",
    "Mã hộ",
    "Thôn/xóm/khối/bản",
    "Chủ hộ",
    "Địa chỉ hộ",
    "Điện thoại chủ hộ",
    "Trạng thái phiếu",
    "Mã đối tượng",
    "Họ và tên",
    "Ngày sinh",
    "Giới tính",
    "Dân tộc",
    "Quan hệ với chủ hộ",
    "Số định danh cá nhân",
    "Mã Bộ GD&ĐT",
    "Trạng thái cư trú",
    "Địa chỉ thường trú",
    "Chỗ ở hiện nay",
    "Tình trạng học tập",
    "Tên trường",
    "Tên lớp",
    "Tình trạng khuyết tật",
    "Dạng khuyết tật",
    "Mức độ khuyết tật",
    "Có giấy xác nhận khuyết tật",
    "Học hòa nhập",
    "Được hỗ trợ khuyết tật",
    "Nội dung hỗ trợ khuyết tật",
    "Hoàn thành CT GDMN 5 tuổi",
    "Đi học đủ ngày theo quy định",
    "Đi học chuyên cần",
    "Được chuẩn bị tiếng Việt",
    "Theo dõi biểu đồ cân nặng",
    "Suy dinh dưỡng nhẹ cân",
    "Theo dõi biểu đồ chiều cao",
    "Suy dinh dưỡng thấp còi",
    "Hoàn cảnh đặc biệt",
    "Ghi chú đối tượng",
    "Ghi chú năm học",
]

TECHNICAL_HEADERS = [
    "__survey_form_id",
    "__household_id",
    "__survey_person_id",
    "__school_year_id",
    "__school_id",
    "__class_id",
    "__is_active",
    "__person_updated_at",
    "__record_updated_at",
    "__parent_source",
]

META_ROW_HEADERS = [
    "visible_row",
    "person_id",
    "person_code",
    "person_updated_at",
    "record_id",
    "record_updated_at",
    "parent_source",
    "original_name",
    "original_birth_date",
    "original_gender",
    "original_ethnic_group",
    "original_relationship_to_head",
    "original_parent_name",
    "original_hamlet",
    "original_program_class",
    "original_school_name",
    "original_notes",
]


class HouseholdExcelError(RuntimeError):
    pass


def _normalize(value: Any) -> str:
    text = " ".join(str(value or "").strip().lower().split())
    text = unicodedata.normalize("NFD", text)
    normalized = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return normalized.replace("đ", "d")


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _int_value(value: Any) -> int | None:
    if value in (None, "") or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    raw = str(value).strip()
    return int(raw) if raw.isdigit() else None


def _date_value(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def _datetime_value(value: Any) -> datetime | None:
    raw = _text(value)
    if not raw:
        return None
    try:
        return datetime.fromisoformat(raw)
    except ValueError:
        return None


def _split_name(value: str | None) -> tuple[str, str]:
    parts = [part for part in str(value or "").strip().split() if part]
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return " ".join(parts[:-1]), parts[-1]


def _join_name(left: Any, right: Any) -> str:
    return " ".join(part for part in (_text(left), _text(right)) if part).strip()


def _find_tracking_sheet(workbook: Any) -> Any:
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip() == TRACKING_SHEET_NAME:
            return workbook[sheet_name]
    raise HouseholdExcelError(
        "Không tìm thấy trang Sổ theo dõi PCGDMN trong tệp Excel."
    )


def _school_year_start(code: str) -> int:
    try:
        return int(str(code).split("-", 1)[0])
    except (TypeError, ValueError):
        return datetime.now().year


def _age_in_school_year(person: SurveyPerson, school_year_code: str) -> int | None:
    if person.date_of_birth is None:
        return None
    start_year = _school_year_start(school_year_code)
    reference = date(start_year + 1, 7, 15)
    return reference.year - person.date_of_birth.year - (
        (reference.month, reference.day)
        < (person.date_of_birth.month, person.date_of_birth.day)
    )


def _program_class_text(
    person: SurveyPerson,
    record: SurveyPersonYearRecord | None,
    school_year_code: str,
) -> str:
    if record is None:
        return ""
    if str(record.learning_status or "").upper() not in {"DANG_HOC", "CHUYEN_DEN"}:
        return ""
    age = _age_in_school_year(person, school_year_code)
    program = "CTNT" if isinstance(age, int) and age <= 2 else "CTMG"
    class_name = ""
    if record.classroom is not None:
        class_name = str(record.classroom.name or "").strip()
    if not class_name:
        class_name = str(record.class_name_reported or "").strip()
    return f"{program}\n{class_name}" if class_name else program


def _school_name(record: SurveyPersonYearRecord | None) -> str:
    if record is None:
        return ""
    if record.school is not None:
        return str(record.school.name or "").strip()
    return str(record.school_name_reported or "").strip()


def _current_parent(person: SurveyPerson) -> tuple[str, str]:
    if str(person.mother_name or "").strip():
        return str(person.mother_name).strip(), "MOTHER"
    if str(person.father_name or "").strip():
        return str(person.father_name).strip(), "FATHER"
    return "", "NONE"


def _notes_for_row(person: SurveyPerson, record: SurveyPersonYearRecord | None) -> str:
    notes: list[str] = []
    if record is not None and str(record.disability_status or "").upper() == "CO_KHUYET_TAT":
        notes.append(str(record.disability_type or "Trẻ khuyết tật"))

    learning = str(record.learning_status or "").upper() if record is not None else ""
    residency = str(person.residency_status or "").upper()
    movement = learning if learning in {"CHUYEN_DEN", "CHUYEN_DI"} else (
        residency if residency in {"CHUYEN_DEN", "CHUYEN_DI"} else ""
    )
    if movement == "CHUYEN_DEN":
        notes.append("Biến động: Chuyển đến")
    elif movement == "CHUYEN_DI":
        notes.append("Biến động: Chuyển đi")
    elif residency in {"TAM_TRU", "TAM_VANG"}:
        notes.append({"TAM_TRU": "Tạm trú", "TAM_VANG": "Tạm vắng"}[residency])

    if record is not None:
        labels = {
            "BO_HOC": "Bỏ học",
            "THOI_HOC": "Thôi học",
            "CHUA_DI_HOC": "Chưa đi học",
            "TAM_NGHI": "Tạm nghỉ",
        }
        if learning in labels:
            notes.append(labels[learning])
        if record.notes:
            notes.extend(
                item.strip()
                for item in re.split(r"[;\n]+", str(record.notes))
                if item.strip()
            )
    if person.notes:
        notes.extend(
            item.strip()
            for item in re.split(r"[;\n]+", str(person.notes))
            if item.strip()
        )
    if not person.personal_id:
        notes.append("Thiếu số định danh")

    result: list[str] = []
    seen: set[str] = set()
    for item in notes:
        cleaned = " ".join(str(item).split())
        key = _normalize(cleaned)
        if cleaned and key not in seen:
            seen.add(key)
            result.append(cleaned)
    return "; ".join(result)


def _meta_values(sheet: Any) -> dict[str, str]:
    values: dict[str, str] = {}
    for row in range(1, META_TABLE_ROW):
        key = _text(sheet.cell(row, 1).value)
        if key:
            values[key] = _text(sheet.cell(row, 2).value)
    return values


def _meta_row_map(sheet: Any) -> dict[int, dict[str, Any]]:
    header_map = {
        _text(sheet.cell(META_TABLE_ROW, col).value): col
        for col in range(1, sheet.max_column + 1)
        if _text(sheet.cell(META_TABLE_ROW, col).value)
    }
    result: dict[int, dict[str, Any]] = {}
    if "visible_row" not in header_map:
        return result
    for row in range(META_TABLE_ROW + 1, sheet.max_row + 1):
        visible_row = _int_value(sheet.cell(row, header_map["visible_row"]).value)
        if visible_row is None:
            continue
        result[visible_row] = {
            key: sheet.cell(row, col).value
            for key, col in header_map.items()
        }
    return result


def _record_map_for_people(
    db: Session,
    person_ids: list[int],
) -> tuple[dict[tuple[int, int], SurveyPersonYearRecord], dict[int, SchoolYear]]:
    if not person_ids:
        return {}, {}
    records = list(
        db.scalars(
            select(SurveyPersonYearRecord)
            .options(
                selectinload(SurveyPersonYearRecord.school),
                selectinload(SurveyPersonYearRecord.classroom),
                selectinload(SurveyPersonYearRecord.school_year),
            )
            .where(SurveyPersonYearRecord.survey_person_id.in_(person_ids))
            .order_by(
                SurveyPersonYearRecord.updated_at.asc(),
                SurveyPersonYearRecord.id.asc(),
            )
        ).all()
    )
    record_map: dict[tuple[int, int], SurveyPersonYearRecord] = {}
    years: dict[int, SchoolYear] = {}
    for item in records:
        record_map[(item.survey_person_id, item.school_year_id)] = item
        years[item.school_year_id] = item.school_year
    return record_map, years


def _write_metadata(
    workbook: Any,
    *,
    batch: Any,
    survey_form: SurveyForm,
    rows: list[dict[str, Any]],
    exported_at: datetime,
    current_year_column: int,
) -> None:
    meta = workbook.create_sheet(META_SHEET_NAME)
    meta.sheet_state = "veryHidden"
    values = [
        ("VERSION", EXPORT_VERSION),
        ("LAYOUT_VERSION", 3),
        ("BATCH_ID", batch.id),
        ("BATCH_CODE", batch.code),
        ("FORM_ID", survey_form.id),
        ("FORM_NUMBER", survey_form.form_number),
        ("HOUSEHOLD_ID", survey_form.household_id),
        ("HOUSEHOLD_CODE", survey_form.household.code),
        ("HOUSEHOLD_HEAD_NAME", survey_form.household.head_name or ""),
        ("COMMUNE_ID", batch.commune_id),
        ("COMMUNE_NAME", batch.commune.name),
        ("SCHOOL_YEAR_ID", batch.school_year_id),
        ("SCHOOL_YEAR_CODE", batch.school_year.code),
        ("EXPORTED_AT", exported_at.isoformat()),
        ("FORM_UPDATED_AT", survey_form.updated_at.isoformat()),
        ("DATA_START_ROW", DATA_START_ROW),
        ("DATA_END_ROW", DATA_END_ROW),
        ("CURRENT_YEAR_COLUMN", current_year_column),
    ]
    for row_number, (key, value) in enumerate(values, start=1):
        meta.cell(row_number, 1).value = key
        meta.cell(row_number, 2).value = value

    for col, header in enumerate(META_ROW_HEADERS, start=1):
        meta.cell(META_TABLE_ROW, col).value = header
    for index, row_data in enumerate(rows, start=META_TABLE_ROW + 1):
        for col, header in enumerate(META_ROW_HEADERS, start=1):
            meta.cell(index, col).value = row_data.get(header)


def _gender_for_excel(value: Any) -> str:
    normalized = _normalize(value)
    if normalized in {"nu", "female", "f"}:
        return "Nữ"
    if normalized in {"nam", "male", "m"}:
        return "Nam"
    return _text(value)


def _gender_from_excel(
    value: Any,
    *,
    explicit: bool,
) -> tuple[str | None, str | None]:
    normalized = _normalize(value)
    if normalized in {"nu", "female", "f", "x"}:
        return "Nữ", None
    if normalized in {"nam", "male", "m"}:
        return "Nam", None
    if not normalized:
        if explicit:
            return None, "Giới tính đang để trống."
        # Phiếu cũ chỉ đánh dấu x ở cột Nữ; ô trống được hiểu là Nam.
        return "Nam", None
    return None, "Giới tính phải ghi Nam hoặc Nữ."


def _prepare_compact_person_headers(tracking: Any) -> None:
    """
    Dùng lại đúng hai cột J:K của mẫu đã chốt.

    Không chèn cột và không dịch bất kỳ cột năm học nào:
    - J: Quan hệ với chủ hộ.
    - K: Họ tên cha/mẹ hoặc người đỡ đầu.
    """
    original_style = copy(tracking["J6"]._style)
    original_font = copy(tracking["J6"].font)
    original_fill = copy(tracking["J6"].fill)
    original_border = copy(tracking["J6"].border)
    original_alignment = copy(tracking["J6"].alignment)
    original_protection = copy(tracking["J6"].protection)

    if "J6:K9" in {str(item) for item in tracking.merged_cells.ranges}:
        tracking.unmerge_cells("J6:K9")

    for cell_ref in ("J6", "K6"):
        cell = tracking[cell_ref]
        cell._style = copy(original_style)
        cell.font = copy(original_font)
        cell.fill = copy(original_fill)
        cell.border = copy(original_border)
        cell.alignment = copy(original_alignment)
        cell.protection = copy(original_protection)

    tracking.merge_cells("J6:J9")
    tracking.merge_cells("K6:K9")
    tracking["F6"] = "Giới\ntính"
    tracking["J6"] = "Quan hệ với\nchủ hộ"
    tracking["K6"] = "Họ tên cha/mẹ hoặc\nngười đỡ đầu"

    for coordinate in ("F6", "J6", "K6"):
        tracking[coordinate].alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    # Chỉ đánh lại số thứ tự cột; vị trí vật lý A:T giữ nguyên.
    tracking["J10"] = 7
    tracking["K10"] = 8
    for column, number in zip(range(12, 19), range(9, 16)):
        tracking.cell(10, column).value = number
    tracking["S10"] = 16
    tracking["T10"] = 17


def _build_household_workbook(
    *,
    db: Session,
    batch: Any,
    survey_form: SurveyForm,
) -> bytes:
    if not TEMPLATE_PATH.is_file():
        raise HouseholdExcelError(f"Không tìm thấy tệp mẫu: {TEMPLATE_PATH}")
    workbook = load_workbook(TEMPLATE_PATH)
    tracking = _find_tracking_sheet(workbook)
    for sheet in list(workbook.worksheets):
        if sheet is not tracking:
            workbook.remove(sheet)
    tracking.title = TRACKING_SHEET_NAME

    if tracking.max_row > DATA_END_ROW:
        tracking.delete_rows(DATA_END_ROW + 1, tracking.max_row - DATA_END_ROW)
    for row in tracking.iter_rows(
        min_row=DATA_START_ROW,
        max_row=DATA_END_ROW,
        min_col=1,
        max_col=20,
    ):
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None

    # Chỉ thay khu vực tiêu đề ở ba dòng đầu của phiếu. Toàn bộ bảng
    # Sổ theo dõi từ dòng 4 trở xuống được giữ nguyên như tệp mẫu đã duyệt.
    for merged_range in list(tracking.merged_cells.ranges):
        if merged_range.min_row <= 3:
            tracking.unmerge_cells(str(merged_range))
    for row_number in range(1, 4):
        for column_number in range(1, 21):
            tracking.cell(row_number, column_number).value = None

    tracking.merge_cells("A1:P1")
    tracking.merge_cells("Q1:T1")
    tracking.merge_cells("A2:T2")
    tracking.merge_cells("A3:J3")
    tracking.merge_cells("K3:T3")

    head_name = " ".join(str(survey_form.household.head_name or "").split())
    hamlet_name = " ".join(
        str(
            survey_form.household.hamlet_name
            or survey_form.hamlet_name_snapshot
            or ""
        ).split()
    )
    household_address = " ".join(
        str(survey_form.household.address or "").split()
    )
    commune_name = " ".join(str(batch.commune.name or "").split())

    tracking["A1"] = "PHIẾU ĐIỀU TRA PHỔ CẬP GIÁO DỤC"
    tracking["Q1"] = f"Số phiếu: {survey_form.form_number}"
    tracking["A2"] = (
        f"Xóm/khối: {hamlet_name or 'Chưa cập nhật'}; "
        f"Xã/phường: {commune_name}; Tỉnh: Nghệ An; "
        f"Năm học: {batch.school_year.code}"
    )
    tracking["A3"] = f"Họ và tên chủ hộ: {head_name or 'Chưa cập nhật'}"
    tracking["K3"] = f"Địa chỉ: {household_address or 'Chưa cập nhật'}"

    base_font_name = tracking["E4"].font.name or "Times New Roman"
    tracking["A1"].font = Font(
        name=base_font_name,
        size=17,
        bold=True,
    )
    tracking["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=False,
    )
    tracking["Q1"].font = Font(
        name=base_font_name,
        size=11,
        bold=False,
    )
    tracking["Q1"].alignment = Alignment(
        horizontal="right",
        vertical="center",
        wrap_text=True,
    )
    for coordinate in ("A2", "A3", "K3"):
        tracking[coordinate].font = Font(
            name=base_font_name,
            size=11,
            bold=True,
        )
        tracking[coordinate].alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )

    tracking.row_dimensions[1].height = 26
    tracking.row_dimensions[2].height = 22
    tracking.row_dimensions[3].height = 22

    # Chỉ chia lại vùng J:K đã có sẵn; không chèn hoặc dịch cột.
    _prepare_compact_person_headers(tracking)

    start_year = _school_year_start(batch.school_year.code)
    first_start = start_year - 6
    year_column_by_start: dict[int, int] = {}
    for offset, column in enumerate(range(12, 19)):
        year_start = first_start + offset
        year_column_by_start[year_start] = column
        tracking.cell(7, column).value = year_start
        tracking.cell(8, column).value = year_start + 1
    current_year_column = year_column_by_start[start_year]

    people = sorted(
        [item for item in survey_form.household.people if item.is_active],
        key=lambda item: (
            item.date_of_birth is None,
            item.date_of_birth or date.max,
            item.full_name,
            item.id,
        ),
    )
    if len(people) > DATA_END_ROW - DATA_START_ROW + 1:
        raise HouseholdExcelError(
            "Hộ có quá nhiều thành viên so với số dòng của phiếu Excel."
        )

    record_map, years = _record_map_for_people(db, [item.id for item in people])
    metadata_rows: list[dict[str, Any]] = []

    for offset in range(DATA_END_ROW - DATA_START_ROW + 1):
        excel_row = DATA_START_ROW + offset
        person = people[offset] if offset < len(people) else None
        if person is None:
            metadata_rows.append({"visible_row": excel_row})
            continue

        parent_name, parent_source = _current_parent(person)
        name_left, name_right = _split_name(person.full_name)
        tracking.cell(excel_row, 1).value = offset + 1
        tracking.cell(excel_row, 2).value = survey_form.form_number
        tracking.cell(excel_row, 3).value = name_left
        tracking.cell(excel_row, 4).value = name_right
        tracking.cell(excel_row, 5).value = person.date_of_birth
        tracking.cell(excel_row, 5).number_format = "dd/mm/yyyy"
        tracking.cell(excel_row, 6).value = _gender_for_excel(person.gender)
        tracking.cell(excel_row, 7).value = survey_form.household.hamlet_name or survey_form.hamlet_name_snapshot or ""
        tracking.cell(excel_row, 8).value = batch.commune.name
        tracking.cell(excel_row, 9).value = person.ethnic_group or ""
        tracking.cell(excel_row, 10).value = person.relationship_to_head or ""
        tracking.cell(excel_row, 11).value = parent_name

        current_record = record_map.get((person.id, batch.school_year_id))
        for (person_id, year_id), record in record_map.items():
            if person_id != person.id:
                continue
            year = years.get(year_id)
            if year is None:
                continue
            year_start = _school_year_start(year.code)
            column = year_column_by_start.get(year_start)
            if column is not None:
                tracking.cell(excel_row, column).value = _program_class_text(person, record, year.code)

        current_program = _program_class_text(person, current_record, batch.school_year.code)
        current_school = _school_name(current_record)
        notes = _notes_for_row(person, current_record)
        tracking.cell(excel_row, 19).value = current_school
        tracking.cell(excel_row, 20).value = notes

        for column in range(1, 21):
            tracking.cell(excel_row, column).alignment = Alignment(
                horizontal=(
                    "left" if column in {3, 4, 7, 8, 9, 10, 11, 19, 20} else "center"
                ),
                vertical="center",
                wrap_text=True,
            )

        metadata_rows.append({
            "visible_row": excel_row,
            "person_id": person.id,
            "person_code": person.code,
            "person_updated_at": person.updated_at.isoformat(),
            "record_id": current_record.id if current_record else None,
            "record_updated_at": current_record.updated_at.isoformat() if current_record else None,
            "parent_source": parent_source,
            "original_name": person.full_name,
            "original_birth_date": person.date_of_birth.isoformat() if person.date_of_birth else "",
            "original_gender": person.gender or "",
            "original_ethnic_group": person.ethnic_group or "",
            "original_relationship_to_head": person.relationship_to_head or "",
            "original_parent_name": parent_name,
            "original_hamlet": survey_form.household.hamlet_name or "",
            "original_program_class": current_program,
            "original_school_name": current_school,
            "original_notes": notes,
        })

    tracking.print_area = f"A1:T{max(DATA_START_ROW + max(len(people), 1) + 4, 32)}"
    tracking.page_setup.orientation = tracking.ORIENTATION_LANDSCAPE
    tracking.page_setup.paperSize = tracking.PAPERSIZE_A4
    tracking.page_setup.fitToWidth = 1
    tracking.page_setup.fitToHeight = 1
    tracking.sheet_properties.pageSetUpPr.fitToPage = True
    tracking.sheet_properties.pageSetUpPr.autoPageBreaks = False
    tracking.page_margins = PageMargins(
        left=0.15,
        right=0.15,
        top=0.20,
        bottom=0.20,
        header=0.0,
        footer=0.0,
    )
    tracking.print_options.horizontalCentered = True
    tracking.print_options.verticalCentered = False
    tracking.freeze_panes = None
    _write_metadata(
        workbook,
        batch=batch,
        survey_form=survey_form,
        rows=metadata_rows,
        exported_at=datetime.now(),
        current_year_column=current_year_column,
    )
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_.-]+", "_", str(value or "").strip())
    return cleaned.strip("._") or "phieu"


def _excel_scope_filters(request: Request) -> list[Any]:
    """Giới hạn phiếu Excel đúng phạm vi từng cấp, không đổi giao diện.

    ADMIN/SO xem toàn tỉnh; Xã/phường được xử lý các phiếu của đợt thuộc
    địa bàn mình; Trường chỉ xử lý phiếu đã giao cho nhân sự thuộc trường;
    Giáo viên chỉ xử lý phiếu được phân công trực tiếp.
    """

    from app.permissions import (
        COMMUNE_ROLE_CODE,
        SCHOOL_ROLE_CODE,
        TEACHER_ROLE_CODE,
        is_province_reader,
        normalize_role_code,
    )

    user = dict(request.scope.get("auth_user") or {})
    role_code = normalize_role_code(user.get("role_code"))

    if is_province_reader(role_code) or role_code == COMMUNE_ROLE_CODE:
        return []

    if role_code == SCHOOL_ROLE_CODE:
        school_id = user.get("school_id")
        if school_id is None:
            return [SurveyForm.id == -1]
        return [
            SurveyForm.investigators.any(
                SurveyFormInvestigator.user.has(
                    User.school_id == int(school_id)
                )
            )
        ]

    if role_code == TEACHER_ROLE_CODE:
        user_id = user.get("id")
        if user_id is None:
            return [SurveyForm.id == -1]
        return [
            SurveyForm.investigators.any(
                SurveyFormInvestigator.user_id == int(user_id)
            )
        ]

    return [SurveyForm.id == -1]


def _log_exchange(
    *,
    db: Session,
    request: Request,
    batch: Any,
    action: str,
    file_name: str,
    form_total: int,
    person_total: int,
    summary: dict[str, int] | None = None,
) -> None:
    user = dict(request.scope.get("auth_user") or {})
    summary = summary or {}
    db.add(SurveyFileExchangeLog(
        school_year_id=batch.school_year_id,
        survey_batch_id=batch.id,
        commune_id=batch.commune_id,
        action=action,
        status="THANH_CONG",
        file_name=file_name,
        actor_user_id=user.get("id"),
        actor_name_snapshot=user.get("full_name"),
        actor_role_snapshot=user.get("role_name") or user.get("role_code"),
        form_total=form_total,
        person_total=person_total,
        updated_household_total=summary.get("updated_households", 0),
        updated_person_total=summary.get("updated_people", 0),
        created_person_total=summary.get("created_people", 0),
        updated_year_record_total=summary.get("updated_year_records", 0),
        created_year_record_total=summary.get("created_year_records", 0),
        notes=(
            "Xuất/nhập cùng một file Phiếu điều tra A4; dữ liệu kỹ thuật giữ ở trang ẩn."
        ),
    ))



def _assignment_sort_key(survey_form: SurveyForm, request: Request) -> tuple[Any, ...]:
    """Xếp phiếu theo đúng thứ tự phân công trong phạm vi tài khoản."""
    from app.permissions import (
        SCHOOL_ROLE_CODE,
        TEACHER_ROLE_CODE,
        normalize_role_code,
    )

    user = dict(request.scope.get("auth_user") or {})
    role_code = normalize_role_code(user.get("role_code"))
    relevant: list[SurveyFormInvestigator] = []
    for assignment in survey_form.investigators:
        if role_code == TEACHER_ROLE_CODE:
            if assignment.user_id == user.get("id"):
                relevant.append(assignment)
        elif role_code == SCHOOL_ROLE_CODE:
            if assignment.user is not None and assignment.user.school_id == user.get("school_id"):
                relevant.append(assignment)
        else:
            relevant.append(assignment)

    if relevant:
        first = min(
            relevant,
            key=lambda item: (
                item.created_at or datetime.max,
                item.id or 0,
            ),
        )
        assigned_at = first.created_at or datetime.max
        assignment_id = first.id or 0
    else:
        assigned_at = survey_form.created_at or datetime.max
        assignment_id = survey_form.id or 0

    return (
        assigned_at,
        assignment_id,
        str(survey_form.form_number or ""),
        survey_form.id,
    )


def _copy_visible_block(
    *,
    source: Any,
    target: Any,
    source_end_row: int,
    target_start_row: int,
) -> None:
    """Sao chép nguyên mẫu A:T của một hộ xuống vị trí mới."""
    row_offset = target_start_row - 1
    for source_row in range(1, source_end_row + 1):
        target_row = source_row + row_offset
        source_dimension = source.row_dimensions[source_row]
        target_dimension = target.row_dimensions[target_row]
        target_dimension.height = source_dimension.height
        target_dimension.hidden = source_dimension.hidden
        target_dimension.outlineLevel = source_dimension.outlineLevel

        for column in range(1, 21):
            source_cell = source.cell(source_row, column)
            if isinstance(source_cell, MergedCell):
                continue
            target_cell = target.cell(target_row, column)
            target_cell.value = source_cell.value
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = copy(source_cell.alignment)
            target_cell.protection = copy(source_cell.protection)
            if source_cell.hyperlink:
                target_cell._hyperlink = copy(source_cell.hyperlink)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)

    for merged_range in source.merged_cells.ranges:
        if merged_range.max_row > source_end_row or merged_range.max_col > 20:
            continue
        shifted = copy(merged_range)
        shifted.shift(row_shift=row_offset, col_shift=0)
        target.merge_cells(str(shifted))


def _copy_block_data_validations(
    *,
    source: Any,
    target: Any,
    source_end_row: int,
    target_start_row: int,
    validation_items: list[Any] | None = None,
) -> None:
    row_offset = target_start_row - 1
    items = validation_items if validation_items is not None else source.data_validations.dataValidation
    for validation in items:
        shifted_ranges: list[str] = []
        for cell_range in validation.ranges.ranges:
            if cell_range.min_col > 20:
                continue
            min_row = max(1, cell_range.min_row)
            max_row = min(source_end_row, cell_range.max_row)
            if min_row > max_row:
                continue
            min_col = max(1, cell_range.min_col)
            max_col = min(20, cell_range.max_col)
            shifted_ranges.append(
                f"{get_column_letter(min_col)}{min_row + row_offset}:"
                f"{get_column_letter(max_col)}{max_row + row_offset}"
            )
        if not shifted_ranges:
            continue
        new_validation = copy(validation)
        new_validation.sqref = " ".join(shifted_ranges)
        target.add_data_validation(new_validation)


def _write_multi_metadata(
    workbook: Any,
    *,
    batch: Any,
    form_total: int,
    exported_at: datetime,
) -> None:
    meta = workbook.create_sheet(META_SHEET_NAME)
    meta.sheet_state = "veryHidden"
    values = [
        ("VERSION", MULTI_EXPORT_VERSION),
        ("LAYOUT_VERSION", MULTI_LAYOUT_VERSION),
        ("BATCH_ID", batch.id),
        ("BATCH_CODE", batch.code),
        ("COMMUNE_ID", batch.commune_id),
        ("COMMUNE_NAME", batch.commune.name),
        ("SCHOOL_YEAR_ID", batch.school_year_id),
        ("SCHOOL_YEAR_CODE", batch.school_year.code),
        ("FORM_TOTAL", form_total),
        ("EXPORTED_AT", exported_at.isoformat()),
        ("ORDER", "ASSIGNMENT_CREATED_AT_ASSIGNMENT_ID_FORM_NUMBER"),
        ("VISIBLE_SHEET", TRACKING_SHEET_NAME),
        ("TECHNICAL_COLUMNS", "U:V"),
    ]
    for row_number, (key, value) in enumerate(values, start=1):
        meta.cell(row_number, 1).value = key
        meta.cell(row_number, 2).value = value


def _build_multi_household_workbook(
    *,
    db: Session,
    batch: Any,
    forms: list[SurveyForm],
) -> bytes:
    """Tạo một tệp Excel, các hộ nối tiếp nhau và mỗi hộ ngắt một trang in."""
    if not forms:
        raise HouseholdExcelError("Không có phiếu để xuất Excel.")

    master_workbook: Any | None = None
    master_tracking: Any | None = None
    next_start_row = 1
    last_visible_row = 0

    for block_index, survey_form in enumerate(forms, start=1):
        individual_content = _build_household_workbook(
            db=db,
            batch=batch,
            survey_form=survey_form,
        )
        individual_workbook = load_workbook(BytesIO(individual_content))
        source_tracking = _find_tracking_sheet(individual_workbook)
        source_meta = individual_workbook[META_SHEET_NAME]
        source_meta_values = _meta_values(source_meta)
        source_row_meta = _meta_row_map(source_meta)

        person_total = sum(1 for item in survey_form.household.people if item.is_active)
        block_rows = max(
            MULTI_BLOCK_MIN_ROWS,
            DATA_START_ROW + person_total + MULTI_EXTRA_MEMBER_ROWS - 1,
        )
        block_rows = min(block_rows, DATA_END_ROW)

        if master_workbook is None:
            first_validation_items = [
                copy(item) for item in source_tracking.data_validations.dataValidation
            ]
            master_workbook = individual_workbook
            master_tracking = source_tracking
            master_workbook.remove(master_workbook[META_SHEET_NAME])
            if master_tracking.max_row > block_rows:
                master_tracking.delete_rows(
                    block_rows + 1,
                    master_tracking.max_row - block_rows,
                )
            master_tracking.data_validations.dataValidation = []
            target_start_row = 1
        else:
            assert master_tracking is not None
            target_start_row = next_start_row
            _copy_visible_block(
                source=source_tracking,
                target=master_tracking,
                source_end_row=block_rows,
                target_start_row=target_start_row,
            )

        assert master_tracking is not None
        if block_index == 1:
            # Khối đầu đã có sẵn từ tệp mẫu; chỉ cần giới hạn validation đúng vùng dữ liệu.
            _copy_block_data_validations(
                source=source_tracking,
                target=master_tracking,
                source_end_row=block_rows,
                target_start_row=target_start_row,
                validation_items=first_validation_items,
            )
        else:
            _copy_block_data_validations(
                source=source_tracking,
                target=master_tracking,
                source_end_row=block_rows,
                target_start_row=target_start_row,
            )

        block_start_payload = {
            "block_index": block_index,
            "form_id": survey_form.id,
            "household_id": survey_form.household_id,
            "form_number": survey_form.form_number,
            "household_code": survey_form.household.code,
            "form_updated_at": survey_form.updated_at.isoformat(),
            "layout_version": int(source_meta_values.get("LAYOUT_VERSION") or 3),
            "current_year_column": int(source_meta_values.get("CURRENT_YEAR_COLUMN") or 18),
        }
        master_tracking.cell(target_start_row, MULTI_MARKER_COLUMN).value = MULTI_BLOCK_START
        master_tracking.cell(target_start_row, MULTI_JSON_COLUMN).value = json.dumps(
            block_start_payload,
            ensure_ascii=False,
        )

        for relative_row in range(DATA_START_ROW, block_rows + 1):
            absolute_row = target_start_row + relative_row - 1
            row_payload = source_row_meta.get(relative_row)
            if row_payload and _int_value(row_payload.get("person_id")) is not None:
                master_tracking.cell(absolute_row, MULTI_MARKER_COLUMN).value = MULTI_PERSON_ROW
                master_tracking.cell(absolute_row, MULTI_JSON_COLUMN).value = json.dumps(
                    row_payload,
                    ensure_ascii=False,
                    default=_text,
                )

        last_visible_row = target_start_row + block_rows - 1
        end_marker_row = last_visible_row + 1
        master_tracking.cell(end_marker_row, MULTI_MARKER_COLUMN).value = MULTI_BLOCK_END
        master_tracking.cell(end_marker_row, MULTI_JSON_COLUMN).value = json.dumps(
            {
                "block_index": block_index,
                "form_id": survey_form.id,
                "household_id": survey_form.household_id,
            },
            ensure_ascii=False,
        )
        master_tracking.row_dimensions[end_marker_row].hidden = True
        master_tracking.row_dimensions[end_marker_row].height = 2
        master_tracking.row_breaks.append(Break(id=last_visible_row))
        next_start_row = end_marker_row + 1

    assert master_workbook is not None and master_tracking is not None
    for column in range(1, 21):
        letter = get_column_letter(column)
        # Kích thước cột được giữ nguyên từ khối đầu tiên.
        if master_tracking.column_dimensions[letter].width is None:
            master_tracking.column_dimensions[letter].width = 10
    for column in (MULTI_MARKER_COLUMN, MULTI_JSON_COLUMN):
        letter = get_column_letter(column)
        master_tracking.column_dimensions[letter].hidden = True
        master_tracking.column_dimensions[letter].width = 2

    master_tracking.print_area = f"A1:T{last_visible_row}"
    master_tracking.page_setup.orientation = master_tracking.ORIENTATION_LANDSCAPE
    master_tracking.page_setup.paperSize = master_tracking.PAPERSIZE_A4
    master_tracking.page_setup.fitToWidth = 1
    master_tracking.page_setup.fitToHeight = 0
    master_tracking.sheet_properties.pageSetUpPr.fitToPage = True
    master_tracking.sheet_properties.pageSetUpPr.autoPageBreaks = False
    master_tracking.page_margins = PageMargins(
        left=0.15,
        right=0.15,
        top=0.20,
        bottom=0.20,
        header=0.0,
        footer=0.0,
    )
    master_tracking.print_options.horizontalCentered = True
    master_tracking.print_options.verticalCentered = False
    master_tracking.freeze_panes = None
    _write_multi_metadata(
        master_workbook,
        batch=batch,
        form_total=len(forms),
        exported_at=datetime.now(),
    )
    master_workbook.calculation.fullCalcOnLoad = True
    master_workbook.calculation.forceFullCalc = True
    output = BytesIO()
    master_workbook.save(output)
    return output.getvalue()


# === BAI_13B_12_V2_4_1_3_INVESTIGATOR_NAMES_START ===
def _field_excel_investigator_names(
    db: Session,
    forms: list[SurveyForm],
) -> dict[int, list[str]]:
    """Lấy tối đa 3 người điều tra thực tế của từng phiếu, bỏ tài khoản đơn vị trường."""
    result: dict[int, list[str]] = {}

    for survey_form in forms:
        rows = db.execute(
            select(
                User.full_name,
                User.username,
                SurveyFormInvestigator.order_number,
                SurveyFormInvestigator.notes,
                SurveyFormInvestigator.id,
            )
            .join(
                SurveyFormInvestigator,
                SurveyFormInvestigator.user_id == User.id,
            )
            .where(
                SurveyFormInvestigator.survey_form_id == int(survey_form.id)
            )
            .order_by(
                SurveyFormInvestigator.order_number,
                SurveyFormInvestigator.id,
            )
        ).all()

        candidates: list[tuple[str, str, str]] = []
        for full_name, username, _order_number, notes, _assignment_id in rows:
            name = " ".join(str(full_name or "").split())
            login_name = str(username or "").strip().lower()
            note_text = str(notes or "")

            if not name:
                continue

            # Tài khoản đơn vị trường không phải cá nhân ký phiếu.
            if login_name.startswith("truong_"):
                continue

            candidates.append((name, login_name, note_text))

        # Ưu tiên đúng thành viên của "Tổ điều tra 3 cấp" nếu lịch sử còn bản ghi khác.
        team_rows = [
            item
            for item in candidates
            if "tổ điều tra 3 cấp" in item[2].casefold()
        ]
        chosen = team_rows if len(team_rows) >= 3 else candidates

        names: list[str] = []
        seen: set[str] = set()
        for name, _login_name, _note_text in chosen:
            key = name.casefold()
            if key in seen:
                continue
            seen.add(key)
            names.append(name)
            if len(names) >= 3:
                break

        result[int(survey_form.id)] = names

    return result
# === BAI_13B_12_V2_4_1_3_INVESTIGATOR_NAMES_END ===


def export_household_excel(
    *,
    batch_id: int,
    request: Request,
    db: Session,
):
    from app.routers.surveys import lay_dot_dieu_tra

    batch = lay_dot_dieu_tra(db, batch_id)
    if batch is None:
        return RedirectResponse(url="/dieu-tra", status_code=303)

    forms = list(db.scalars(
        select(SurveyForm)
        .join(Household, Household.id == SurveyForm.household_id)
        .options(
            selectinload(SurveyForm.household).selectinload(Household.people),
            selectinload(SurveyForm.investigators).selectinload(
                SurveyFormInvestigator.user
            ),
            with_loader_criteria(
                SurveyPerson,
                SurveyPerson.is_active.is_(True),
                include_aliases=True,
            ),
        )
        .where(
            SurveyForm.survey_batch_id == batch.id,
            *_excel_scope_filters(request),
        )
    ).unique().all())
    if not forms:
        return RedirectResponse(
            url=f"/dieu-tra/{batch.id}/ho-dan?status=no_export_data",
            status_code=303,
        )

    forms.sort(key=lambda item: _assignment_sort_key(item, request))
    person_total = sum(
        sum(1 for person in item.household.people if person.is_active)
        for item in forms
    )

    if len(forms) == 1:
        survey_form = forms[0]
        content = _build_household_workbook(
            db=db,
            batch=batch,
            survey_form=survey_form,
        )
        file_name = (
            f"{_safe_filename(survey_form.form_number)}_"
            f"{_safe_filename(survey_form.household.code)}.xlsx"
        )
        action = "XUAT_DON"
    else:
        content = _build_multi_household_workbook(
            db=db,
            batch=batch,
            forms=forms,
        )
        file_name = (
            f"phieu_dieu_tra_{_safe_filename(batch.commune.code)}_"
            f"{_safe_filename(batch.school_year.code.replace('-', ''))}_"
            f"{len(forms)}_ho.xlsx"
        )
        action = "XUAT_HANG_LOAT_XLSX"

    # === BAI_13B_12_V2_4_1_FIELD_EXCEL_EXPORT ===
    try:
        from app.household_field_excel_v240 import (
            add_field_form_view_to_bytes,
        )
        investigator_names_by_form_id = _field_excel_investigator_names(
            db,
            forms,
        )
        content = add_field_form_view_to_bytes(
            content,
            investigator_names_by_form_id=investigator_names_by_form_id,
        )
    except Exception as exc:
        raise HouseholdExcelError(
            "Không dựng được trang Phiếu điều tra A4 trong file Excel."
        ) from exc

    _log_exchange(
        db=db,
        request=request,
        batch=batch,
        action=action,
        file_name=file_name,
        form_total=len(forms),
        person_total=person_total,
    )
    db.commit()
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )


def _parse_program_class(value: Any) -> tuple[str, str]:
    raw = _text(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    if not raw:
        return "", ""
    lines = [line.strip() for line in raw.split("\n") if line.strip()]
    if not lines:
        return "", ""
    first = _normalize(lines[0]).upper()
    if first in {"CTNT", "CTMG"}:
        return first, " ".join(lines[1:]).strip()
    match = re.match(r"^(CTNT|CTMG)\s+(.+)$", lines[0], flags=re.IGNORECASE)
    if match:
        return match.group(1).upper(), match.group(2).strip()
    return "", " ".join(lines).strip()


def _learning_from_notes(notes: str, has_school_data: bool, existing: str | None) -> str:
    key = _normalize(notes)
    if "chuyen di" in key:
        return "CHUYEN_DI"
    if "chuyen den" in key:
        return "CHUYEN_DEN"
    if "bo hoc" in key:
        return "BO_HOC"
    if "thoi hoc" in key:
        return "THOI_HOC"
    if "chua di hoc" in key:
        return "CHUA_DI_HOC"
    if "da chet" in key or key == "chet" or "; chet" in key:
        return "DA_CHET"
    if has_school_data:
        return "DANG_HOC"
    return existing or "CHUA_XAC_DINH"


def _resolve_school_class(
    db: Session,
    *,
    school_name: str,
    class_name: str,
    commune_id: int,
    school_year_id: int,
) -> tuple[int | None, int | None]:
    school_id: int | None = None
    class_id: int | None = None
    if school_name:
        schools = list(db.scalars(select(School).where(School.is_active.is_(True))).all())
        exact = [item for item in schools if _normalize(item.name) == _normalize(school_name)]
        same_commune = [item for item in exact if item.commune_id == commune_id]
        selected = same_commune[0] if len(same_commune) == 1 else (exact[0] if len(exact) == 1 else None)
        if selected is not None:
            school_id = selected.id
    if school_id is not None and class_name:
        classes = list(db.scalars(
            select(Classroom).where(
                Classroom.school_id == school_id,
                Classroom.school_year_id == school_year_id,
                Classroom.is_active.is_(True),
            )
        ).all())
        exact_classes = [item for item in classes if _normalize(item.name) == _normalize(class_name)]
        if len(exact_classes) == 1:
            class_id = exact_classes[0].id
    return school_id, class_id


def _load_accessible_form(
    *,
    db: Session,
    request: Request,
    batch: Any,
    form_id: int,
) -> SurveyForm | None:
    return db.scalar(
        select(SurveyForm)
        .join(Household, Household.id == SurveyForm.household_id)
        .options(
            selectinload(SurveyForm.household).selectinload(Household.people),
            selectinload(SurveyForm.person_year_records),
        )
        .where(
            SurveyForm.id == form_id,
            SurveyForm.survey_batch_id == batch.id,
            *_excel_scope_filters(request),
        )
    )



def _copy_multi_block_to_single_workbook(
    *,
    workbook: Any,
    tracking: Any,
    global_meta: dict[str, str],
    block_start_row: int,
    block_end_row: int,
    block_payload: dict[str, Any],
) -> bytes:
    """Tách một khối hộ thành tệp một hộ để dùng lại bộ kiểm tra đã ổn định."""
    visible_end_row = block_end_row - 1
    relative_end_row = visible_end_row - block_start_row + 1
    if relative_end_row < DATA_START_ROW:
        raise HouseholdExcelError("Khối phiếu Excel không đủ số dòng theo mẫu.")
    if relative_end_row > DATA_END_ROW:
        raise HouseholdExcelError(
            f"Phiếu {block_payload.get('form_number') or '—'} vượt quá "
            f"{DATA_END_ROW - DATA_START_ROW + 1} dòng đối tượng."
        )

    single = load_workbook(TEMPLATE_PATH)
    single_tracking = _find_tracking_sheet(single)
    for sheet in list(single.worksheets):
        if sheet is not single_tracking:
            single.remove(sheet)
    single_tracking.title = TRACKING_SHEET_NAME

    for merged_range in list(single_tracking.merged_cells.ranges):
        single_tracking.unmerge_cells(str(merged_range))
    if single_tracking.max_row > relative_end_row:
        single_tracking.delete_rows(
            relative_end_row + 1,
            single_tracking.max_row - relative_end_row,
        )
    for row in single_tracking.iter_rows(
        min_row=1,
        max_row=relative_end_row,
        min_col=1,
        max_col=20,
    ):
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None

    for absolute_row in range(block_start_row, visible_end_row + 1):
        relative_row = absolute_row - block_start_row + 1
        source_dimension = tracking.row_dimensions[absolute_row]
        target_dimension = single_tracking.row_dimensions[relative_row]
        target_dimension.height = source_dimension.height
        target_dimension.hidden = source_dimension.hidden
        for column in range(1, 21):
            source_cell = tracking.cell(absolute_row, column)
            if isinstance(source_cell, MergedCell):
                continue
            target_cell = single_tracking.cell(relative_row, column)
            target_cell.value = source_cell.value
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = copy(source_cell.alignment)
            target_cell.protection = copy(source_cell.protection)
            target_cell.number_format = source_cell.number_format

    for merged_range in tracking.merged_cells.ranges:
        if (
            merged_range.min_row >= block_start_row
            and merged_range.max_row <= visible_end_row
            and merged_range.max_col <= 20
        ):
            shifted = copy(merged_range)
            shifted.shift(row_shift=1 - block_start_row, col_shift=0)
            single_tracking.merge_cells(str(shifted))

    meta = single.create_sheet(META_SHEET_NAME)
    meta.sheet_state = "veryHidden"
    values = [
        ("VERSION", EXPORT_VERSION),
        ("LAYOUT_VERSION", block_payload.get("layout_version") or 3),
        ("BATCH_ID", global_meta.get("BATCH_ID")),
        ("BATCH_CODE", global_meta.get("BATCH_CODE")),
        ("FORM_ID", block_payload.get("form_id")),
        ("FORM_NUMBER", block_payload.get("form_number")),
        ("HOUSEHOLD_ID", block_payload.get("household_id")),
        ("HOUSEHOLD_CODE", block_payload.get("household_code") or ""),
        ("COMMUNE_ID", global_meta.get("COMMUNE_ID")),
        ("COMMUNE_NAME", global_meta.get("COMMUNE_NAME")),
        ("SCHOOL_YEAR_ID", global_meta.get("SCHOOL_YEAR_ID")),
        ("SCHOOL_YEAR_CODE", global_meta.get("SCHOOL_YEAR_CODE")),
        ("EXPORTED_AT", global_meta.get("EXPORTED_AT")),
        ("FORM_UPDATED_AT", block_payload.get("form_updated_at") or ""),
        ("DATA_START_ROW", DATA_START_ROW),
        ("DATA_END_ROW", DATA_END_ROW),
        ("CURRENT_YEAR_COLUMN", block_payload.get("current_year_column") or 18),
    ]
    for row_number, (key, value) in enumerate(values, start=1):
        meta.cell(row_number, 1).value = key
        meta.cell(row_number, 2).value = value
    for column, header in enumerate(META_ROW_HEADERS, start=1):
        meta.cell(META_TABLE_ROW, column).value = header

    metadata_index = META_TABLE_ROW + 1
    for absolute_row in range(block_start_row + DATA_START_ROW - 1, block_end_row):
        marker = _text(tracking.cell(absolute_row, MULTI_MARKER_COLUMN).value)
        if marker != MULTI_PERSON_ROW:
            continue
        raw_payload = _text(tracking.cell(absolute_row, MULTI_JSON_COLUMN).value)
        try:
            row_payload = json.loads(raw_payload) if raw_payload else {}
        except json.JSONDecodeError as exc:
            raise HouseholdExcelError(
                f"Thông tin kỹ thuật ở dòng {absolute_row} đã bị hỏng."
            ) from exc
        row_payload["visible_row"] = absolute_row - block_start_row + 1
        for column, header in enumerate(META_ROW_HEADERS, start=1):
            meta.cell(metadata_index, column).value = row_payload.get(header)
        metadata_index += 1

    single_tracking.print_area = f"A1:T{relative_end_row}"
    single_tracking.page_setup.orientation = single_tracking.ORIENTATION_LANDSCAPE
    single_tracking.page_setup.paperSize = single_tracking.PAPERSIZE_A4
    single_tracking.page_setup.fitToWidth = 1
    single_tracking.page_setup.fitToHeight = 1
    single_tracking.sheet_properties.pageSetUpPr.fitToPage = True
    output = BytesIO()
    single.save(output)
    return output.getvalue()


def _parse_multi_household_workbook_loaded(
    *,
    workbook: Any,
    tracking: Any,
    meta: dict[str, str],
    db: Session,
    request: Request,
    batch: Any,
) -> dict[str, Any]:
    if _int_value(meta.get("BATCH_ID")) != batch.id or meta.get("BATCH_CODE") != batch.code:
        raise HouseholdExcelError("File không thuộc đúng đợt điều tra đang mở.")

    blocks: list[tuple[int, int, dict[str, Any]]] = []
    active_start: int | None = None
    active_payload: dict[str, Any] | None = None
    for row_number in range(1, tracking.max_row + 1):
        marker = _text(tracking.cell(row_number, MULTI_MARKER_COLUMN).value)
        if marker == MULTI_BLOCK_START:
            if active_start is not None:
                raise HouseholdExcelError("File có khối hộ bị lồng hoặc thiếu dòng kết thúc.")
            raw_payload = _text(tracking.cell(row_number, MULTI_JSON_COLUMN).value)
            try:
                active_payload = json.loads(raw_payload)
            except json.JSONDecodeError as exc:
                raise HouseholdExcelError(
                    f"Thông tin kỹ thuật của hộ tại dòng {row_number} đã bị hỏng."
                ) from exc
            active_start = row_number
        elif marker == MULTI_BLOCK_END:
            if active_start is None or active_payload is None:
                raise HouseholdExcelError("File có dòng kết thúc hộ nhưng thiếu dòng bắt đầu.")
            blocks.append((active_start, row_number, active_payload))
            active_start = None
            active_payload = None
    if active_start is not None:
        raise HouseholdExcelError("Hộ cuối cùng trong file bị thiếu dòng kết thúc kỹ thuật.")
    if not blocks:
        raise HouseholdExcelError("Không tìm thấy hộ nào trong file Excel nhiều hộ.")

    declared_total = _int_value(meta.get("FORM_TOTAL"))
    if declared_total is not None and declared_total != len(blocks):
        raise HouseholdExcelError("Số hộ trong file không khớp thông tin kỹ thuật ban đầu.")

    previews: list[dict[str, Any]] = []
    combined_rows: list[dict[str, Any]] = []
    combined_errors: list[str] = []
    combined_conflicts: list[str] = []
    combined_warnings: list[str] = []
    summary = {
        "updated_people": 0,
        "created_people": 0,
        "unchanged_people": 0,
        "error_total": 0,
        "conflict_total": 0,
        "warning_total": 0,
    }

    for block_start, block_end, block_payload in blocks:
        single_content = _copy_multi_block_to_single_workbook(
            workbook=workbook,
            tracking=tracking,
            global_meta=meta,
            block_start_row=block_start,
            block_end_row=block_end,
            block_payload=block_payload,
        )
        single_preview = parse_household_workbook(
            content=single_content,
            db=db,
            request=request,
            batch=batch,
        )
        form_number = single_preview["form"].form_number
        household_code = single_preview["household"].code
        for row in single_preview["rows"]:
            row = dict(row)
            row["excel_row"] = f"{form_number} / {row['excel_row']}"
            row["details"] = f"Hộ {household_code}: {row['details']}"
            combined_rows.append(row)
        combined_errors.extend(
            f"Phiếu {form_number}: {message}" for message in single_preview["errors"]
        )
        combined_conflicts.extend(
            f"Phiếu {form_number}: {message}" for message in single_preview["conflicts"]
        )
        combined_warnings.extend(
            f"Phiếu {form_number}: {message}" for message in single_preview["warnings"]
        )
        for key in summary:
            summary[key] += int(single_preview["summary"].get(key, 0) or 0)
        previews.append(single_preview)

    first = previews[0]
    return {
        "form": first["form"],
        "household": first["household"],
        "form_previews": previews,
        "preview_title": f"{len(previews)} phiếu – {len(previews)} hộ gia đình",
        "meta": meta,
        "actions": [action for item in previews for action in item["actions"]],
        "rows": combined_rows,
        "errors": combined_errors,
        "conflicts": combined_conflicts,
        "warnings": combined_warnings,
        "summary": summary,
        "can_confirm": not combined_errors and not combined_conflicts,
    }

def parse_household_workbook(
    *,
    content: bytes,
    db: Session,
    request: Request,
    batch: Any,
) -> dict[str, Any]:
    try:
        workbook = load_workbook(BytesIO(content), data_only=False)
    except Exception as exc:
        raise HouseholdExcelError("Không đọc được file Excel hoặc file đã bị hỏng.") from exc

    # === BAI_13B_12_V2_4_1_FIELD_EXCEL_IMPORT ===
    try:
        from app.household_field_excel_v240 import (
            sync_field_form_to_tracking,
        )
        sync_field_form_to_tracking(workbook)
    except Exception as exc:
        raise HouseholdExcelError(
            "Trang Phiếu điều tra trong file Excel bị sai cấu trúc hoặc đã bị đổi tên."
        ) from exc

    tracking = _find_tracking_sheet(workbook)
    if META_SHEET_NAME not in workbook.sheetnames:
        raise HouseholdExcelError(
            "File không có thông tin kỹ thuật ẩn. Chỉ nhập lại file do phần mềm xuất."
        )
    meta_sheet = workbook[META_SHEET_NAME]
    meta = _meta_values(meta_sheet)
    version = meta.get("VERSION")
    if version == MULTI_EXPORT_VERSION:
        return _parse_multi_household_workbook_loaded(
            workbook=workbook,
            tracking=tracking,
            meta=meta,
            db=db,
            request=request,
            batch=batch,
        )
    if version != EXPORT_VERSION:
        raise HouseholdExcelError("Phiên bản file điều tra không đúng hoặc đã bị thay đổi.")
    layout_version = _int_value(meta.get("LAYOUT_VERSION")) or 1
    explicit_gender = layout_version >= 2
    compact_relationship_layout = layout_version >= 3
    inserted_relationship_layout = layout_version == 2
    if _int_value(meta.get("BATCH_ID")) != batch.id or meta.get("BATCH_CODE") != batch.code:
        raise HouseholdExcelError("File không thuộc đúng đợt điều tra đang mở.")
    form_id = _int_value(meta.get("FORM_ID"))
    if form_id is None:
        raise HouseholdExcelError("File thiếu mã phiếu kỹ thuật.")
    survey_form = _load_accessible_form(
        db=db, request=request, batch=batch, form_id=form_id
    )
    if survey_form is None:
        raise HouseholdExcelError("Tài khoản không có quyền cập nhật phiếu trong file này.")
    if _int_value(meta.get("HOUSEHOLD_ID")) != survey_form.household_id:
        raise HouseholdExcelError("Mã hộ trong file không khớp dữ liệu hệ thống.")

    row_meta = _meta_row_map(meta_sheet)
    default_current_year_column = 19 if inserted_relationship_layout else 18
    current_year_column = (
        _int_value(meta.get("CURRENT_YEAR_COLUMN"))
        or default_current_year_column
    )
    minimum_year_column = 13 if inserted_relationship_layout else 12
    maximum_year_column = 19 if inserted_relationship_layout else 18
    if (
        current_year_column < minimum_year_column
        or current_year_column > maximum_year_column
    ):
        raise HouseholdExcelError("Cột năm học hiện tại trong file không hợp lệ.")

    if inserted_relationship_layout:
        relationship_column = 10
        parent_left_column = 11
        parent_right_column = 12
        school_column = 20
        notes_column = 21
    elif compact_relationship_layout:
        relationship_column = 10
        parent_left_column = 11
        parent_right_column = None
        school_column = 19
        notes_column = 20
    else:
        relationship_column = None
        parent_left_column = 10
        parent_right_column = 11
        school_column = 19
        notes_column = 20

    existing_people = {item.id: item for item in survey_form.household.people}
    current_records = {
        item.survey_person_id: item
        for item in survey_form.person_year_records
        if item.school_year_id == batch.school_year_id
    }
    actions: list[dict[str, Any]] = []
    rows: list[dict[str, Any]] = []
    errors: list[str] = []
    conflicts: list[str] = []
    warnings: list[str] = []
    seen_name_birth: set[tuple[str, date]] = set()
    hamlets: set[str] = set()

    for excel_row in range(DATA_START_ROW, DATA_END_ROW + 1):
        meta_row = row_meta.get(excel_row, {})
        person_id = _int_value(meta_row.get("person_id"))
        full_name = _join_name(
            tracking.cell(excel_row, 3).value,
            tracking.cell(excel_row, 4).value,
        )
        birth_date = _date_value(tracking.cell(excel_row, 5).value)
        gender, gender_error = _gender_from_excel(
            tracking.cell(excel_row, 6).value,
            explicit=explicit_gender,
        )
        hamlet = " ".join(_text(tracking.cell(excel_row, 7).value).split())
        commune_name = " ".join(_text(tracking.cell(excel_row, 8).value).split())
        ethnic_group = " ".join(_text(tracking.cell(excel_row, 9).value).split())
        relationship_to_head = (
            " ".join(
                _text(tracking.cell(excel_row, relationship_column).value).split()
            )
            if relationship_column is not None
            else ""
        )
        if parent_right_column is None:
            parent_name = " ".join(
                _text(tracking.cell(excel_row, parent_left_column).value).split()
            )
        else:
            parent_name = _join_name(
                tracking.cell(excel_row, parent_left_column).value,
                tracking.cell(excel_row, parent_right_column).value,
            )
        program_class = _text(tracking.cell(excel_row, current_year_column).value)
        program, class_name = _parse_program_class(program_class)
        school_name = " ".join(
            _text(tracking.cell(excel_row, school_column).value).split()
        )
        notes = " ".join(
            _text(tracking.cell(excel_row, notes_column).value).split()
        )

        if not full_name:
            if person_id is not None:
                warnings.append(
                    f"Dòng {excel_row}: người cũ đang bị để trống; hệ thống sẽ không tự xóa."
                )
                rows.append({
                    "excel_row": excel_row,
                    "full_name": meta_row.get("original_name") or "—",
                    "status": "BỎ QUA AN TOÀN",
                    "details": "Dòng cũ bị để trống, không xóa dữ liệu.",
                })
            continue

        row_errors: list[str] = []
        if birth_date is None:
            row_errors.append("Ngày sinh đang trống hoặc sai định dạng.")
        if gender_error:
            row_errors.append(gender_error)
        if not ethnic_group:
            row_errors.append("Dân tộc đang để trống.")
        if relationship_column is not None and not relationship_to_head:
            row_errors.append("Quan hệ với chủ hộ đang để trống.")
        if commune_name and _normalize(commune_name) != _normalize(batch.commune.name):
            row_errors.append("Tên xã/phường không đúng với đợt điều tra.")
        if hamlet:
            hamlets.add(_normalize(hamlet))
        if birth_date is not None:
            duplicate_key = (_normalize(full_name), birth_date)
            if duplicate_key in seen_name_birth:
                row_errors.append("Họ tên và ngày sinh bị trùng với dòng khác trong file.")
            seen_name_birth.add(duplicate_key)

        person = existing_people.get(person_id) if person_id is not None else None
        if relationship_column is None:
            if person is not None:
                relationship_to_head = person.relationship_to_head or ""
            else:
                row_errors.append(
                    "File bố cục cũ không có cột Quan hệ với chủ hộ; "
                    "không thể thêm người mới bằng file này."
                )
        if person_id is not None and person is None:
            row_errors.append("Mã người kỹ thuật không còn thuộc hộ này.")
        if person is None and birth_date is not None:
            duplicates = [
                item for item in existing_people.values()
                if item.is_active
                and _normalize(item.full_name) == _normalize(full_name)
                and item.date_of_birth == birth_date
            ]
            if duplicates:
                row_errors.append(
                    "Người cùng họ tên và ngày sinh đã tồn tại trong hộ; không tạo mới để tránh trùng."
                )

        current_record = current_records.get(person.id) if person is not None else None
        existing_learning = current_record.learning_status if current_record else None
        has_school_data = bool(program or class_name or school_name)
        learning_status = _learning_from_notes(notes, has_school_data, existing_learning)
        school_id, class_id = _resolve_school_class(
            db,
            school_name=school_name,
            class_name=class_name,
            commune_id=batch.commune_id,
            school_year_id=batch.school_year_id,
        )

        if row_errors:
            errors.extend(f"Dòng {excel_row}: {message}" for message in row_errors)
            rows.append({
                "excel_row": excel_row,
                "full_name": full_name,
                "status": "LỖI",
                "details": " ".join(row_errors),
            })
            continue

        parent_source = _text(meta_row.get("parent_source")) or "NONE"
        changed_fields: list[str] = []
        if person is None:
            status = "THÊM MỚI"
            changed_fields = ["Tạo mới đối tượng", "Tạo hồ sơ năm học"]
        else:
            comparisons = [
                ("Họ tên", person.full_name, full_name),
                ("Ngày sinh", person.date_of_birth, birth_date),
                ("Giới tính", person.gender, gender),
                ("Dân tộc", person.ethnic_group, ethnic_group),
                ("Quan hệ với chủ hộ", person.relationship_to_head, relationship_to_head),
                ("Cha/mẹ/người đỡ đầu", meta_row.get("original_parent_name") or "", parent_name),
                ("Thôn/xóm", meta_row.get("original_hamlet") or "", hamlet),
                ("Chương trình/lớp", meta_row.get("original_program_class") or "", program_class),
                ("Cơ sở GDMN", meta_row.get("original_school_name") or "", school_name),
                ("Ghi chú", meta_row.get("original_notes") or "", notes),
            ]
            for label, old, new in comparisons:
                old_value = old.isoformat() if isinstance(old, date) else _text(old)
                new_value = new.isoformat() if isinstance(new, date) else _text(new)
                if _normalize(old_value) != _normalize(new_value):
                    changed_fields.append(label)
            status = "CẬP NHẬT" if changed_fields else "KHÔNG ĐỔI"

            if changed_fields:
                exported_person_time = _datetime_value(meta_row.get("person_updated_at"))
                if exported_person_time and person.updated_at > exported_person_time + timedelta(seconds=1):
                    conflicts.append(
                        f"Dòng {excel_row}: hồ sơ {person.full_name} đã được sửa trên hệ thống sau khi xuất file."
                    )
                exported_record_time = _datetime_value(meta_row.get("record_updated_at"))
                if current_record and exported_record_time and current_record.updated_at > exported_record_time + timedelta(seconds=1):
                    conflicts.append(
                        f"Dòng {excel_row}: dữ liệu năm học của {person.full_name} đã thay đổi sau khi xuất file."
                    )

        original_program = _text(meta_row.get("original_program_class"))
        original_school = _text(meta_row.get("original_school_name"))
        original_notes = _text(meta_row.get("original_notes"))
        year_changed = (
            person is None
            or _normalize(original_program) != _normalize(program_class)
            or _normalize(original_school) != _normalize(school_name)
            or _normalize(original_notes) != _normalize(notes)
            or (current_record is not None and current_record.learning_status != learning_status)
        )
        notes_changed = _normalize(original_notes) != _normalize(notes)

        action = {
            "excel_row": excel_row,
            "person_id": person.id if person else None,
            "full_name": full_name,
            "birth_date": birth_date,
            "gender": gender,
            "ethnic_group": ethnic_group,
            "relationship_to_head": relationship_to_head,
            "parent_name": parent_name,
            "parent_source": parent_source,
            "hamlet": hamlet,
            "program": program,
            "class_name": class_name,
            "school_name": school_name,
            "school_id": school_id,
            "class_id": class_id,
            "notes": notes,
            "learning_status": learning_status,
            "changed_fields": changed_fields,
            "year_changed": year_changed,
            "notes_changed": notes_changed,
            "status": status,
        }
        actions.append(action)
        rows.append({
            "excel_row": excel_row,
            "full_name": full_name,
            "status": status,
            "details": ", ".join(changed_fields) if changed_fields else "Không có thay đổi.",
        })

    if len(hamlets) > 1:
        errors.append("Các dòng trong cùng một phiếu đang có nhiều tên thôn/xóm khác nhau.")

    summary = {
        "updated_people": sum(item["status"] == "CẬP NHẬT" for item in actions),
        "created_people": sum(item["status"] == "THÊM MỚI" for item in actions),
        "unchanged_people": sum(item["status"] == "KHÔNG ĐỔI" for item in actions),
        "error_total": len(errors),
        "conflict_total": len(conflicts),
        "warning_total": len(warnings),
    }
    return {
        "form": survey_form,
        "household": survey_form.household,
        "meta": meta,
        "actions": actions,
        "rows": rows,
        "errors": errors,
        "conflicts": conflicts,
        "warnings": warnings,
        "summary": summary,
        "can_confirm": not errors and not conflicts,
    }


def _cleanup_staging() -> None:
    STAGING_DIR.mkdir(parents=True, exist_ok=True)
    cutoff = datetime.now() - timedelta(hours=24)
    for path in STAGING_DIR.glob("*"):
        try:
            if datetime.fromtimestamp(path.stat().st_mtime) < cutoff:
                if path.is_dir():
                    shutil.rmtree(path, ignore_errors=True)
                else:
                    path.unlink(missing_ok=True)
        except OSError:
            pass


def _stage_file(
    *,
    content: bytes,
    batch_id: int,
    user_id: int | None,
    file_name: str,
) -> str:
    _cleanup_staging()
    token = uuid4().hex
    folder = STAGING_DIR / token
    folder.mkdir(parents=True, exist_ok=False)
    (folder / "upload.xlsx").write_bytes(content)
    (folder / "meta.json").write_text(json.dumps({
        "batch_id": batch_id,
        "user_id": user_id,
        "file_name": file_name,
        "created_at": datetime.now().isoformat(),
    }, ensure_ascii=False), encoding="utf-8")
    return token


def _load_staged_file(
    *,
    token: str,
    batch_id: int,
    user_id: int | None,
) -> tuple[bytes, str, Path]:
    if not re.fullmatch(r"[0-9a-f]{32}", str(token or "")):
        raise HouseholdExcelError("Mã xác nhận file không hợp lệ.")
    folder = STAGING_DIR / token
    meta_path = folder / "meta.json"
    file_path = folder / "upload.xlsx"
    if not meta_path.is_file() or not file_path.is_file():
        raise HouseholdExcelError("File chờ xác nhận đã hết hạn hoặc không còn tồn tại.")
    metadata = json.loads(meta_path.read_text(encoding="utf-8"))
    if int(metadata.get("batch_id") or 0) != batch_id:
        raise HouseholdExcelError("File chờ xác nhận không thuộc đợt hiện tại.")
    if metadata.get("user_id") != user_id:
        raise HouseholdExcelError("File chờ xác nhận thuộc phiên đăng nhập khác.")
    return file_path.read_bytes(), str(metadata.get("file_name") or ""), folder


def _backup_database() -> Path:
    database_path = Path(DATABASE_PATH)
    backup_dir = database_path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / f"before_household_excel_import_{datetime.now():%Y%m%d_%H%M%S}.db"
    source = sqlite3.connect(str(database_path))
    target = sqlite3.connect(str(backup_path))
    try:
        source.backup(target)
    finally:
        target.close()
        source.close()
    return backup_path


def _next_person_code(db: Session) -> str:
    sequence = int(db.scalar(select(func.count(SurveyPerson.id))) or 0) + 1
    while True:
        candidate = f"DT-{sequence:08d}"
        if db.scalar(select(SurveyPerson.id).where(SurveyPerson.code == candidate)) is None:
            return candidate
        sequence += 1


def _apply_preview(
    *,
    db: Session,
    request: Request,
    batch: Any,
    parsed: dict[str, Any],
    file_name: str,
) -> dict[str, Any]:
    if not parsed["can_confirm"]:
        raise HouseholdExcelError("File còn lỗi hoặc xung đột nên chưa thể cập nhật.")

    previews = parsed.get("form_previews") or [parsed]
    backup_path = _backup_database()
    summary = {
        "updated_households": 0,
        "updated_people": 0,
        "created_people": 0,
        "updated_year_records": 0,
        "created_year_records": 0,
    }

    for form_preview in previews:
        survey_form: SurveyForm = form_preview["form"]
        household: Household = form_preview["household"]
        changed_hamlets = {
            item["hamlet"] for item in form_preview["actions"]
            if item["hamlet"] and item["status"] in {"CẬP NHẬT", "THÊM MỚI"}
        }
        if len(changed_hamlets) == 1:
            new_hamlet = next(iter(changed_hamlets))
            if _normalize(household.hamlet_name) != _normalize(new_hamlet):
                household.hamlet_name = new_hamlet
                survey_form.hamlet_name_snapshot = new_hamlet
                summary["updated_households"] += 1

        for action in form_preview["actions"]:
            if action["status"] == "KHÔNG ĐỔI":
                continue
            person = db.get(SurveyPerson, action["person_id"]) if action["person_id"] else None
            if person is None:
                person = SurveyPerson(
                    code=_next_person_code(db),
                    household_id=household.id,
                    full_name=action["full_name"],
                    date_of_birth=action["birth_date"],
                    gender=action["gender"],
                    ethnic_group=action["ethnic_group"],
                    relationship_to_head=action["relationship_to_head"],
                    residency_status=(
                        action["learning_status"]
                        if action["learning_status"] in {"CHUYEN_DEN", "CHUYEN_DI"}
                        else "THUONG_TRU"
                    ),
                    current_address=household.address,
                    is_active=True,
                )
                if action["parent_name"]:
                    person.mother_name = action["parent_name"]
                db.add(person)
                db.flush()
                summary["created_people"] += 1
            else:
                person.full_name = action["full_name"]
                person.date_of_birth = action["birth_date"]
                person.gender = action["gender"]
                person.ethnic_group = action["ethnic_group"]
                person.relationship_to_head = action["relationship_to_head"]
                if action["learning_status"] in {"CHUYEN_DEN", "CHUYEN_DI"}:
                    person.residency_status = action["learning_status"]
                elif (
                    person.residency_status in {"CHUYEN_DEN", "CHUYEN_DI"}
                    and action["notes_changed"]
                ):
                    person.residency_status = "THUONG_TRU"
                if action["parent_source"] == "FATHER":
                    person.father_name = action["parent_name"] or None
                else:
                    person.mother_name = action["parent_name"] or None
                summary["updated_people"] += 1

            record = db.scalar(select(SurveyPersonYearRecord).where(
                SurveyPersonYearRecord.survey_form_id == survey_form.id,
                SurveyPersonYearRecord.survey_person_id == person.id,
                SurveyPersonYearRecord.school_year_id == batch.school_year_id,
            ))
            if record is None:
                record = SurveyPersonYearRecord(
                    survey_form_id=survey_form.id,
                    survey_person_id=person.id,
                    school_year_id=batch.school_year_id,
                    learning_status=action["learning_status"],
                    school_id=action["school_id"],
                    class_id=action["class_id"],
                    school_name_reported=action["school_name"] or None,
                    class_name_reported=action["class_name"] or None,
                    notes=action["notes"] or None,
                    is_reviewed=True,
                )
                db.add(record)
                summary["created_year_records"] += 1
            elif action["year_changed"]:
                record.learning_status = action["learning_status"]
                record.school_id = action["school_id"]
                record.class_id = action["class_id"]
                record.school_name_reported = action["school_name"] or None
                record.class_name_reported = action["class_name"] or None
                if action["notes_changed"]:
                    record.notes = action["notes"] or None
                record.is_reviewed = True
                summary["updated_year_records"] += 1

    total_actions = sum(len(item["actions"]) for item in previews)
    _log_exchange(
        db=db,
        request=request,
        batch=batch,
        action="NHAP_EXCEL_NHIEU_HO" if len(previews) > 1 else "NHAP_EXCEL",
        file_name=file_name,
        form_total=len(previews),
        person_total=total_actions,
        summary=summary,
    )
    db.commit()
    summary["backup_path"] = str(backup_path)
    return summary


def _render_import(
    *,
    request: Request,
    batch: Any,
    file_name: str = "",
    errors: list[dict[str, Any]] | None = None,
    summary: dict[str, Any] | None = None,
    preview: dict[str, Any] | None = None,
    confirm_token: str = "",
    status_code: int = 200,
) -> HTMLResponse:
    return templates.TemplateResponse(
        request=request,
        name="surveys/excel_import.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "batch": batch,
            "file_name": file_name,
            "errors": errors or [],
            "summary": summary,
            "preview": preview,
            "confirm_token": confirm_token,
        },
        status_code=status_code,
    )


def import_household_excel_page(
    *,
    batch_id: int,
    request: Request,
    db: Session,
):
    from app.routers.surveys import lay_dot_dieu_tra

    batch = lay_dot_dieu_tra(db, batch_id)
    if batch is None:
        return RedirectResponse(url="/dieu-tra", status_code=303)
    return _render_import(request=request, batch=batch)


async def import_household_excel_post(
    *,
    batch_id: int,
    request: Request,
    excel_file: UploadFile | None,
    confirm_token: str,
    db: Session,
):
    from app.routers.surveys import lay_dot_dieu_tra

    batch = lay_dot_dieu_tra(db, batch_id)
    if batch is None:
        return RedirectResponse(url="/dieu-tra", status_code=303)
    if batch.status == "DA_KET_THUC":
        return _render_import(
            request=request,
            batch=batch,
            errors=[{"row": "—", "message": "Đợt điều tra đã kết thúc, không thể nhập Excel."}],
            status_code=400,
        )

    user = dict(request.scope.get("auth_user") or {})
    try:
        if confirm_token:
            content, file_name, folder = _load_staged_file(
                token=confirm_token,
                batch_id=batch.id,
                user_id=user.get("id"),
            )
            parsed = parse_household_workbook(
                content=content, db=db, request=request, batch=batch
            )
            if not parsed["can_confirm"]:
                return _render_import(
                    request=request,
                    batch=batch,
                    file_name=file_name,
                    preview=parsed,
                    confirm_token=confirm_token,
                    errors=[
                        {"row": "—", "message": message}
                        for message in (parsed["errors"] + parsed["conflicts"])
                    ],
                    status_code=400,
                )
            summary = _apply_preview(
                db=db,
                request=request,
                batch=batch,
                parsed=parsed,
                file_name=file_name,
            )
            shutil.rmtree(folder, ignore_errors=True)
            return _render_import(
                request=request,
                batch=batch,
                file_name=file_name,
                summary=summary,
            )

        if excel_file is None:
            raise HouseholdExcelError("Chưa chọn file Excel cần kiểm tra.")
        file_name = excel_file.filename or ""
        if not file_name.lower().endswith(".xlsx"):
            raise HouseholdExcelError("Chỉ chấp nhận file Excel định dạng .xlsx.")
        content = await excel_file.read()
        if not content:
            raise HouseholdExcelError("File tải lên đang trống.")
        if len(content) > MAX_UPLOAD_BYTES:
            raise HouseholdExcelError("File lớn quá 15 MB.")
        parsed = parse_household_workbook(
            content=content, db=db, request=request, batch=batch
        )
        token = _stage_file(
            content=content,
            batch_id=batch.id,
            user_id=user.get("id"),
            file_name=file_name,
        )
        return _render_import(
            request=request,
            batch=batch,
            file_name=file_name,
            preview=parsed,
            confirm_token=token,
            errors=[
                {"row": "—", "message": message}
                for message in (parsed["errors"] + parsed["conflicts"])
            ],
            status_code=400 if not parsed["can_confirm"] else 200,
        )
    except HouseholdExcelError as exc:
        db.rollback()
        return _render_import(
            request=request,
            batch=batch,
            file_name=(excel_file.filename if excel_file else ""),
            errors=[{"row": "—", "message": str(exc)}],
            status_code=400,
        )
    except Exception as exc:
        db.rollback()
        return _render_import(
            request=request,
            batch=batch,
            file_name=(excel_file.filename if excel_file else ""),
            errors=[{
                "row": "—",
                "message": f"Có lỗi khi xử lý file; chưa cập nhật dữ liệu. Chi tiết: {exc}",
            }],
            status_code=500,
        )
