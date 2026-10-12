"""Fill the eight appendix-2 Excel templates, retaining their source layout."""
from copy import copy
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment
from openpyxl.worksheet.merge import MergedCellRange
from openpyxl.worksheet.cell_range import CellRange
from app.xmc_appendices import TEMPLATE_DIR, COMPULSORY_FORMS
from app.structured_report_catalog import STRUCTURED_FORM_BY_CODE, flatten_fields
from app.services.xmc_appendix_reports import put, ratio
from app.services.compulsory_report_data import (
    Population, attr, known, disabled, capable, accessible, completed, studying,
    in_grade, ppc, location, program, upper, tri_all, school_data, total,
)


def setcell(ws, row, col, value, percent=False):
    if isinstance(ws.cell(row,col), MergedCell):
        for merged in ws.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
                row, col = merged.min_row, merged.min_col
                break
    put(ws, row, col, value, percent=percent)


def expand(ws, start, capacity, required):
    extra = max(0, required - capacity)
    if not extra:
        return start + capacity
    at = start + capacity
    merges = list(ws.merged_cells.ranges)
    heights = {r:d.height for r,d in ws.row_dimensions.items()}
    styles = [copy(c._style) for c in ws[start]]
    repeated = [m for m in merges if m.min_row == start and m.max_row == start]
    for m in merges:
        ws.unmerge_cells(str(m))
    ws.insert_rows(at, extra)
    for m in merges:
        ws.merge_cells(start_row=m.min_row + (extra if m.min_row >= at else 0),
                       end_row=m.max_row + (extra if m.max_row >= at else 0),
                       start_column=m.min_col, end_column=m.max_col)
    for r,h in sorted(heights.items(), reverse=True):
        ws.row_dimensions[r + (extra if r >= at else 0)].height = h
    for r in range(at, at+extra):
        ws.row_dimensions[r].height = heights.get(start, 25)
        for c, style in enumerate(styles,1):
            ws.cell(r,c)._style = copy(style)
        for m in repeated:
            # These ranges clone disjoint template cells into newly inserted rows.
            # Avoid MultiCellRange.add's quadratic scan across thousands of rows.
            coord=CellRange(min_row=r,max_row=r,min_col=m.min_col,max_col=m.max_col).coord
            merged=MergedCellRange(ws,coord)
            ws.merged_cells.ranges.add(merged)
            ws._clean_merge_range(merged)
    return at + extra


def population_value(x, indicator):
    return {3: lambda: True, 4: lambda: x['female'], 5: lambda: x['ethnic'],
            6: lambda: disabled(x), 7: lambda: capable(x), 8: lambda: accessible(x),
            9: lambda: ppc(x)}[indicator]()


def match_location(x, bucket):
    loc = location(x)
    return tri_all(None if loc is None else loc == bucket, True if bucket == 'inbound' else ppc(x))


def th_predicate(indicator):
    if indicator <= 9:
        return lambda x: population_value(x, indicator)
    if 10 <= indicator <= 24:
        g = (indicator-10)//3 + 1
        bucket = ('local','other','inbound')[(indicator-10)%3]
        return lambda x: tri_all(in_grade(x,g,g), match_location(x,bucket))
    if 25 <= indicator <= 27:
        bucket = ('local','other','inbound')[indicator-25]
        return lambda x: tri_all(completed(x,'TH'), match_location(x,bucket))
    if indicator == 28:
        return lambda x: tri_all(completed(x,'TH'), None if ppc(x) is None else not ppc(x))
    if 29 <= indicator <= 31:
        bucket = ('local','other','inbound')[indicator-29]
        return lambda x: tri_all(known(attr(x,'is_repeating_grade')), in_grade(x,1,5), match_location(x,bucket))
    if 32 <= indicator <= 34:
        bucket = ('local','other','inbound')[indicator-32]
        return lambda x: tri_all(None if attr(x,'learning_status') in (None,'CHUA_XAC_DINH') else attr(x,'learning_status')=='BO_HOC', match_location(x,bucket))
    return lambda x: tri_all(ppc(x), None if attr(x,'learning_status') in (None,'CHUA_XAC_DINH') else attr(x,'learning_status')=='CHUA_DI_HOC')


def thcs_predicate(indicator):
    if indicator <= 9:
        return lambda x: population_value(x,indicator)
    if 11 <= indicator <= 22:
        g = (indicator-11)//3 + 6
        bucket = ('local','other','inbound')[(indicator-11)%3]
        return lambda x: tri_all(in_grade(x,g,g),match_location(x,bucket))
    if 23 <= indicator <= 25:
        outcome = lambda x: completed(x,'THCS')
        bucket = ('local','other','inbound')[indicator-23]
    elif 26 <= indicator <= 28:
        outcome = upper
        bucket = ('local','other','inbound')[indicator-26]
    else:
        outcome = lambda x: None if attr(x,'learning_status') in (None,'CHUA_XAC_DINH') else attr(x,'learning_status') in {'BO_HOC','THOI_HOC'}
        bucket = ('local','other','inbound')[indicator-29]
    return lambda x: tri_all(outcome(x),match_location(x,bucket))


def fill_age_table(wb, pop, level, year):
    primary = level == 'TH'
    first, second = (wb['Trang_04'],wb['Trang_05']) if primary else (wb['Trang_09'],wb['Trang_10'])
    ages = list(range(6,15)) if primary else list(range(11,19))
    agecols = [6,7,8,9,10,12,13,14,15] if primary else [6,8,10,12,14,17,20,23]
    for age,col in zip(ages,agecols):
        setcell(first,7,col,year-age)
    for ws, labelcol in ((first,5),(second,4 if primary else 8)):
        for row in list(ws):
            value = ws.cell(row[0].row,labelcol).value
            if not isinstance(value,(int,float)) and not str(value or '').isdigit():
                continue
            indicator = int(value)
            if indicator < 3 or (not primary and indicator==10):
                continue
            if indicator > (35 if primary else 31):
                continue
            pred = th_predicate(indicator) if primary else thcs_predicate(indicator)
            cols = [c.column for c in row if c.column > labelcol and not isinstance(c,MergedCell)]
            if primary:
                groups = [(a,a) for a in range(6,11)]+[(6,10)]+[(a,a) for a in range(11,15)]+[(11,14)]
                if len(cols)!=len(groups):
                    raise ValueError(f'Unexpected primary table grid {ws.title}:{row[0].row}')
                for col,(lo,hi) in zip(cols,groups):
                    setcell(ws,row[0].row,col,pop.number(lo,hi,pred))
            elif indicator <= 9:
                for col,age in zip(agecols,ages):
                    setcell(ws,row[0].row,col,pop.number(age,age,pred))
                setcell(ws,row[0].row,26,pop.number(15,18,pred))
            else:
                groups = [(age,prog) for age in ages for prog in (('PT','TX') if age<15 else ('PT','TX','GDNN'))]
                if len(cols)!=len(groups)+1:
                    raise ValueError(f'Unexpected secondary table grid {ws.title}:{row[0].row}: {len(cols)}')
                for col,(age,prog) in zip(cols,groups):
                    def selected(x, prog=prog):
                        p = program(x)
                        return tri_all(pred(x), None if p is None else p==prog)
                    setcell(ws,row[0].row,col,pop.number(age,age,selected))
                setcell(ws,row[0].row,cols[-1],pop.number(15,18,pred))
    metrics = pop.indicators()
    if primary:
        for row,key in [(11,'entry6'),(16,'primary11'),(17,'primary_study11'),(18,'primary6_14'),(19,'primary14'),(20,'disabled6_14')]:
            n,d,rate=metrics[key]
            setcell(second,row,6,n)
            setcell(second,row,7,rate,True)
    else:
        for row,key in [(20,'secondary15_18'),(21,'upper15_18'),(23,'disabled11_18')]:
            n,d,rate=metrics[key]
            setcell(second,row,12,n)
            setcell(second,row,13,rate,True)


def fill_completion(ws, pop, level):
    people = [x for x in pop.items if completed(x,level) is True]
    people.sort(key=lambda x: (x['person'].full_name,x['person'].id))
    expand(ws,9,4,len(people))
    if level=='THCS':
        ws['G8']='Đã hoàn thành chương trình giáo dục trung học cơ sở'
    for index in range(max(4,len(people))):
        values = [None]*9
        if index<len(people):
            x=people[index]; p=x['person']
            # Current school and school-year are not proof of graduation place/date.
            values=[index+1,p.full_name,p.date_of_birth,p.gender,p.ethnic_group,p.household.address,'X',None,None]
        for col,value in enumerate(values,1):
            put(ws,9+index,col,value)
        ws.cell(9+index,3).number_format='dd/mm/yyyy'
    ws.print_title_rows='8:8'


def resource_mapping(code):
    level='THCS' if 'THCS' in code else 'TH'
    kind='CSVC' if 'CSVC' in code else 'GV'
    fields=flatten_fields(STRUCTURED_FORM_BY_CODE[level+'_01_'+kind])
    mapping={f['excel_col']:f['code'] for f in fields if f.get('export',True)}
    if code=='02-TH_GV':
        cols=[1,2,3,4,5,6,7,8,9,10,11,12]+list(range(15,36))
        mapping[14]='teacher_ratio'
        # Old sessions_1_2 does not unambiguously mean two sessions/day.
        mapping[6]='two_sessions_per_day'
        return 12,3,{cols[c-1]:key for c,key in mapping.items()}
    if code=='06-THCS_GV':
        # Source form drops the three rank columns and adds a site-count column.
        cols=[1,3,4,5,7,8,10,11,12,13,15,17,18,20,21]+list(range(22,44))+[45,46,47,48]
        keys=['school_site_count','principal','vice_principal','teacher_total','teacher_permanent','teacher_contract',
              'teacher_female','teacher_ethnic','teacher_ratio','qual_postgrad','qual_university','qual_college','qual_secondary',
              'subject_math','subject_literature','subject_physics','subject_chemistry','subject_biology','subject_history','subject_geography',
              'subject_music','subject_art','subject_pe','subject_civic','subject_technology','subject_it','subject_english','subject_russian',
              'subject_french','subject_other','team_leader','prof_excellent','prof_good','prof_average','prof_poor',
              'employee_library','employee_equipment_lab','employee_office','employee_health']
        assert len(cols[2:])==len(keys)
        return 11,6,dict(zip(cols[2:],keys))
    mapping[10 if level=='TH' else 8]='room_ratio'
    if level=='TH':
        return 11,3,{1+(c-1)*3:key for c,key in mapping.items()}
    return 12,5,mapping


def fill_resources(ws, rows, code):
    start,capacity,mapping=resource_mapping(code)
    end=expand(ws,start,capacity,len(rows))
    namecol=4 if code=='03-TH_CSVC' else 3 if code=='06-THCS_GV' else 2
    for index in range(max(capacity,len(rows))):
        data=rows[index][1] if index<len(rows) else {}
        setcell(ws,start+index,1,index+1 if index<len(rows) else None)
        setcell(ws,start+index,namecol,rows[index][0].name if index<len(rows) else None)
        for col,key in mapping.items():
            setcell(ws,start+index,col,data.get(key))
    # Preserve merged total label in source form 06.
    setcell(ws,end,1 if code=='06-THCS_GV' else namecol,'Cộng')
    for col,key in mapping.items():
        if key in {'room_ratio','teacher_ratio'}:
            denominator=total(d.get('class_total') for _,d in rows)
            keys=['teacher_total'] if key=='teacher_ratio' else ['room_permanent','room_semi_permanent','room_temporary']+(['room_rent_borrow'] if code=='03-TH_CSVC' else [])
            numerator=total(total(d.get(k) for k in keys) for _,d in rows)
            value=round(numerator/denominator,2) if numerator is not None and denominator else None
        else:
            value=total(d.get(key) for _,d in rows)
        setcell(ws,end,col,value)
    ws.print_title_rows=f'7:{start-1}' if code!='07-THCS_CSVC' else '8:11'
    if code in {'02-TH_GV','06-THCS_GV'}:
        shift=end-(start+capacity)
        denom=total(d.get('teacher_total') for _,d in rows)
        values=[total(total(d.get(k) for k in keys) for _,d in rows) for keys in
                [('qual_university','qual_postgrad'),('qual_postgrad',),('prof_excellent','prof_good','prof_average')]]
        for r,n in zip(([20,23,24] if code=='02-TH_GV' else [23,24,25]),values):
            setcell(ws,r+shift,5 if code=='02-TH_GV' else 14,n)
            setcell(ws,r+shift,7 if code=='02-TH_GV' else 16,ratio(n,denom),True)


def build_compulsory(code, db, ctx):
    bundle=code=='bo-pl2'
    wb=load_workbook(TEMPLATE_DIR/('Phu_luc_2_Giao_duc_bat_buoc.xlsx' if bundle else f'Mau_{code}.xlsx'))
    pop=Population(ctx)
    codes=[c for c,_,_ in COMPULSORY_FORMS] if bundle else [code]
    resources={}
    for form in codes:
        pages=next(p for c,_,p in COMPULSORY_FORMS if c==form)
        ws=wb[f'Trang_{pages[0]:02d}']
        level='THCS' if 'THCS' in form else 'TH'
        if form.endswith(('_TN','_TTN')):
            fill_age_table(wb,pop,level,ctx['reference_year'])
        elif form.endswith('_HT'):
            fill_completion(ws,pop,level)
        else:
            if level not in resources:
                resources[level]=school_data(db,ctx,level)
            fill_resources(ws,resources[level],form)
    for ws in wb:
        for row in ws:
            for cell in row:
                s=str(cell.value or '')
                if s.startswith(('Tỉnh/xã:', 'Đơn vị:')):
                    put(ws,cell.row,cell.column,f"Phạm vi: {ctx['label']} · Năm học {ctx['year'].code}")
                elif s.startswith('Thời điểm:'):
                    put(ws,cell.row,cell.column,f"Năm báo cáo: {ctx['reference_year']} · Số liệu hiện có tại thời điểm xuất")
        ws.page_setup.fitToWidth=1
        ws.page_setup.fitToHeight=0
        ws.print_area=ws.dimensions
    audit=wb.create_sheet('Kiem_tra_du_lieu')
    for row in [
        ['Phạm vi',ctx['label']],['Năm học',ctx['year'].code],
        ['Nguồn','Phiếu điều tra trong phạm vi; dữ liệu mạng lưới, đội ngũ, CSVC đúng năm học. Không sửa dữ liệu gốc.'],
        ['Quy tắc','Ô trống: thiếu xác nhận hoặc mẫu số bằng 0. Số 0 chỉ dùng khi tính được. Tổng tỷ lệ tính lại từ tử/mẫu số.'],
        ['Phạm vi số liệu','Trường/giáo viên: dân số từ phiếu được phân công; đội ngũ và CSVC giới hạn trường của tài khoản.'],
        ['Số phải phổ cập','Kế thừa quy ước thường trú và không thuộc diện ngừng theo dõi. Nơi học dùng trường thông tin đã xác nhận trên phiếu.'],
        ['Lớp/chương trình','Không suy lớp từ tuổi. Chương trình PT/TX/GDNN lấy trường dữ liệu hoặc trường THCS/GDTX được khai báo; chưa rõ để trống.'],
        ['Mẫu hoàn thành','Chỉ lấy xác nhận hoàn thành chương trình. Chưa có trường năm/trường hoàn thành chuyên biệt nên để trống, không lấy năm/trường đang học. Mẫu 08 sửa nhãn cột G từ tiểu học thành THCS theo tiêu đề biểu.'],
        ['Đội ngũ/CSVC','Ưu tiên số liệu nhập theo biểu đúng năm; tiếp đến hồ sơ đội ngũ/mạng lưới đúng năm. Hồ sơ chưa đầy đủ cần đối chiếu trước khi ký.'],
        ['Dữ liệu thiếu ngày sinh',sum(x['age'] is None for x in pop.items)],
    ]: audit.append(row)
    audit.append(['Mã hồ sơ','Họ tên','Thông tin cần kiểm tra'])
    for x in pop.items:
        missing=[]
        if x['age'] is None: missing.append('Ngày sinh; chưa thể phân nhóm tuổi')
        if x['record'] is None: missing.append('Hồ sơ năm học')
        elif x['age'] is not None and 6<=x['age']<=18:
            for key in ['completed_primary_program','completed_lower_secondary_program','disability_can_learn','disability_access_education','study_location_scope','current_education_program']:
                if attr(x,key) in (None,'CHUA_XAC_DINH'): missing.append(key)
        if missing: audit.append([x['person'].code,x['person'].full_name,'; '.join(missing)])
    for level,rows in resources.items():
        for school,data in rows:
            audit.append([level,school.name,'; '.join(k for k in ['teacher_total','class_total','room_permanent','room_semi_permanent','room_temporary'] if data.get(k) is None)])
    for col,width in [('A',30),('B',85),('C',90)]: audit.column_dimensions[col].width=width
    for row in audit:
        for cell in row:
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            if isinstance(cell.value,str): cell.data_type='s'
    output=BytesIO(); wb.save(output); output.seek(0)
    return output
