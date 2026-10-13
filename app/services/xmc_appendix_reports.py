"""Aggregate new XMC appendices from scoped survey records; unknown is not zero."""
from collections import defaultdict
from copy import copy
from datetime import date
from io import BytesIO
import re

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.cell.cell import MergedCell
from sqlalchemy import select

from app.xmc_appendices import TEMPLATE_DIR, XMC_FORMS
from app.xmc_models import XmcEvidence
from app.pcgd_xmc_report_builders_v1 import _norm, _scope_meta
from app.services.xmc_rules_353 import personal_levels, commune_level, LEGAL_BASIS, EFFECTIVE_DATE


def flag(value):
    return value if isinstance(value, bool) else None


def collect(people, evidence, year):
    result = []
    for p, r in people:
        age = year - p.date_of_birth.year if p.date_of_birth else None
        if p.residency_status in {"CHUYEN_DI", "DA_CHET", "CHET"}:
            continue
        data = evidence.get(r.id, {}) if r else {}
        g1, g2, conflict = personal_levels(r, data)
        gender = _norm(p.gender)
        female = gender in {"NU", "FEMALE", "F"} if gender in {"NU", "FEMALE", "F", "NAM", "MALE", "M"} else None
        ethnic = None if not p.ethnic_group or _norm(p.ethnic_group) in {"CHUA XAC DINH", "KHONG", "NONE"} else _norm(p.ethnic_group) != "KINH"
        result.append(dict(person=p, record=r, age=age, data=data, g1=g1, g2=g2,
                           female=female, ethnic=ethnic, conflict=conflict))
    return result


def population(items):
    """All, female, minority, female minority; blank if classification incomplete."""
    groups = (lambda x: True, lambda x: x["female"], lambda x: x["ethnic"],
              lambda x: None if x["female"] is None or x["ethnic"] is None else x["female"] and x["ethnic"])
    counts = []
    for pred in groups:
        values = [pred(x) for x in items]
        counts.append(None if None in values else sum(v is True for v in values))
    return counts


def counts(items, predicate):
    values = [predicate(x) for x in items]
    if any(v is None for v in values):
        return [None] * 4
    return population([x for x, value in zip(items, values) if value])


def negative(value):
    return None if value is None else not value


def ratio(n, d):
    return round(100 * n / d, 2) if n is not None and d else None


def put(ws, row, col, value, *, percent=False):
    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        raise ValueError(f"Invalid report mapping: {ws.title}!{cell.coordinate}")
    # Business-rule rates use 0–100; Excel percentage cells store 0–1.
    cell.value = value / 100 if percent and value is not None else value
    # Names and notes must remain text, even if they look like Excel formulas.
    if isinstance(value, str):
        cell.data_type = "s"
    if percent:
        cell.number_format = "0%" if value is not None and float(value).is_integer() else "0.00%"
    elif isinstance(value, float):
        cell.number_format = "0.00"
    if value is None:
        cell.comment = Comment("Chưa đủ dữ liệu xác nhận hoặc mẫu số bằng 0; không quy đổi thành 0.", "Phổ cập")


def columns(ws):
    for row in ws:
        cells = [c for c in row if re.fullmatch(r"\(\d+\)", str(c.value or "").strip())]
        if len(cells) >= 10:
            return row[0].row, [c.column for c in cells]
    raise ValueError(f"Missing numbered columns in {ws.title}")


def expand_rows(ws, start, count):
    """Expand the three sample rows, moving footer merges and preserving styles."""
    extra = max(0, count - 3)
    if not extra:
        return
    at = start + 3
    merges = list(ws.merged_cells.ranges)
    dims = {i: copy(d) for i, d in ws.row_dimensions.items()}
    row_merges = [m for m in merges if m.min_row == start and m.max_row == start]
    for m in merges:
        ws.unmerge_cells(str(m))
    ws.insert_rows(at, extra)
    for m in merges:
        r0 = m.min_row + (extra if m.min_row >= at else 0)
        r1 = m.max_row + (extra if m.max_row >= at else 0)
        ws.merge_cells(start_row=r0, end_row=r1, start_column=m.min_col, end_column=m.max_col)
    for i, dim in sorted(dims.items(), reverse=True):
        dest = i + extra if i >= at else i
        ws.row_dimensions[dest] = copy(dim)
        ws.row_dimensions[dest].index = dest
    for r in range(at, at + extra):
        ws.row_dimensions[r].height = dims[start].height
        for c in range(1, ws.max_column + 1):
            source = ws.cell(start, c)
            dest = ws.cell(r, c)
            dest._style = copy(source._style)
        for m in row_merges:
            ws.merge_cells(start_row=r, end_row=r, start_column=m.min_col, end_column=m.max_col)
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f"1:{start - 1}"
    ws.print_area = ws.dimensions


def list_rows(ws, values, total=None, percent_columns=()):
    header, cols = columns(ws)
    start = header + 1
    expand_rows(ws, start, len(values))
    for index in range(max(3, len(values))):
        row = values[index] if index < len(values) else [None] * len(cols)
        for logical_col, (col, value) in enumerate(zip(cols, row), 1):
            put(ws, start + index, col, value, percent=logical_col in percent_columns)
    if total is not None:
        for logical_col, (col, value) in enumerate(zip(cols, total), 1):
            put(ws, start + max(3, len(values)), col, value, percent=logical_col in percent_columns)


def age_rows(ws, items, code, year, literacy_basis):
    _, cols = columns(ws)
    for row in ws:
        a = str(row[1].value or "").strip()
        label = " ".join(str(c.value or "") for c in row[:2])
        ages = None
        if (code == "01" and re.fullmatch(r"\d+ tuổi", a)) or (code == "07" and a.isdigit()):
            age = int(a.split()[0])
            ages = [age]
            put(ws, row[0].row, cols[2] if code == "01" else cols[0], year - age)
        elif re.search(r"(?:Cộng|TC:)\s*(\d+)\s*-\s*(\d+)", label):
            match = re.search(r"(?:Cộng|TC:)\s*(\d+)\s*-\s*(\d+)", label)
            ages = range(int(match[1]), int(match[2]) + 1)
        if ages is None:
            continue
        group = [x for x in items if x["age"] in ages]
        if code == "01":
            pop = population(group)
            fail1 = counts(group, lambda x: negative(x["g1"]))
            fail2 = counts(group, lambda x: negative(x["g2"]))
            passed = counts(group, lambda x: x.get(literacy_basis)) if literacy_basis else [None] * 4
            ill = counts(group, lambda x: x["data"].get("illiterate"))
            values = pop + fail1 + fail2 + passed + [ratio(passed[0], pop[0])] + ill + [ratio(ill[0], pop[0])]
            for logical_col, (c, value) in enumerate(zip(cols[3:], values), 4):
                put(ws, row[0].row, c, value, percent=logical_col in {20, 25})
        else:
            values = culture_values(group)
            for c, value in zip(cols[2:], values):
                put(ws, row[0].row, c, value)


def culture_values(group):
    from app.routers.household_updates import _level_from_record
    pop = population(group)
    disabled = counts(group, lambda x: None if not x["record"] or x["record"].disability_status == "CHUA_XAC_DINH"
                      else x["record"].disability_status == "CO_KHUYET_TAT")[0]
    result = [pop[0], pop[2], pop[1], disabled]
    known = all(x["record"] and x["record"].learning_status != "CHUA_XAC_DINH" for x in group)
    for attending in (True, False):
        subset = [x for x in group if x["record"] and
                  (x["record"].learning_status == "DANG_HOC" if attending else
                   x["record"].learning_status in {"CHUA_DI_HOC", "BO_HOC", "THOI_HOC"})]
        if not known:
            result.extend([None] * 13)
            continue
        result.append(len(subset))
        if attending:
            result.append(sum(_level_from_record(x["record"], x["record"].school, x["record"].classroom) == "MN" for x in subset))
        else:
            result.append(sum(x["record"].learning_status == "CHUA_DI_HOC" for x in subset))
        grades = []
        for x in subset:
            r = x["record"]
            name = r.class_name_reported or (r.classroom.name if r.classroom else "")
            match = re.fullmatch(r"(?:LOP\s*)?(\d{1,2})(?:\s*[A-Z][A-Z0-9]*)?", _norm(name))
            # Preschool and people never enrolled do not need a grade 1–9.
            exempt = (_level_from_record(r, r.school, r.classroom) == "MN"
                      or r.learning_status == "CHUA_DI_HOC")
            grades.append(int(match[1]) if match else (0 if exempt else None))
        for grade in range(1, 10):
            result.append(None if None in grades else sum(g == grade for g in grades))
        result.extend([counts(subset, lambda x: flag(x["record"].completed_primary_program))[0],
                       counts(subset, lambda x: flag(x["record"].completed_lower_secondary_program))[0]])
    return result


def unit_values(items, code):
    values = []
    for low, high in ((15, 35), (15, 60)):
        group = [x for x in items if x["age"] is not None and low <= x["age"] <= high]
        if code in {"03", "04"}:
            pop = population(group)
            outcomes = counts(group, lambda x: negative(x["g1" if code == "03" else "g2"]))
        elif code in {"05", "06"}:
            level = "1" if code == "05" else "2"
            pop = counts(group, lambda x: x["data"].get("enrolled" + level))
            def recognized(x):
                enrolled = x["data"].get("enrolled" + level)
                if enrolled is False:
                    return False
                return x["data"].get("recognized" + level) if enrolled is True else None
            outcomes = counts(group, recognized)
        else:
            pop = population(group)
            p1 = counts(group, lambda x: x["g1"])[0]
            p2 = counts(group, lambda x: x["g2"])[0]
            values.extend([pop[0], p1, ratio(p1, pop[0]), p2, ratio(p2, pop[0])])
            continue
        values.extend(pop)
        for number, denominator in zip(outcomes, pop):
            values.extend([number, ratio(number, denominator)])
    return values


def conclusion(items, commune, allow):
    if not allow:
        return "Không áp dụng"
    return commune_level(items, commune.is_special_difficulty_area)


def build_report(code, people, evidence, year, year_code, scope_label, communes,
                 school_scope=False, literacy_basis="g1"):
    if code == "bo-xmc":
        wb = load_workbook(TEMPLATE_DIR / "Phu_luc_3_Xoa_mu_chu.xlsx")
        codes = [c for c, _, _ in XMC_FORMS]
    else:
        wb = load_workbook(TEMPLATE_DIR / f"Mau_{code}_XMC.xlsx")
        codes = [code]
    items = collect(people, evidence, year)
    by_commune = defaultdict(list)
    for x in items:
        by_commune[x["person"].household.commune_id].append(x)
    groups = [(c, by_commune[c.id]) for c in communes]
    for form in codes:
        pages = next(p for c, _, p in XMC_FORMS if c == form)
        if form in {"01", "07"}:
            age_rows(wb[f"Trang_{pages[0]:02d}"], items, form, year, literacy_basis)
            continue
        ws = wb[f"Trang_{pages[0]:02d}"]
        if form in {"02", "08"}:
            rows = []
            for x in sorted(items, key=lambda x: (_norm(x["person"].full_name), x["person"].id)):
                p, d = x["person"], x["data"]
                if x["age"] is None or not 15 <= x["age"] <= 60:
                    continue
                if form == "02" and not (x["g1"] is False or x["g2"] is False):
                    continue
                if form == "08" and d.get("relapsed") is not True:
                    continue
                if form == "02":
                    rows.append([len(rows) + 1, p.full_name, p.date_of_birth.strftime("%d/%m/%Y"), p.gender,
                                 p.ethnic_group, d.get("policy"), p.household.address,
                                 "X" if x["g1"] is False else None, "X" if x["g2"] is False else None,
                                 p.personal_id, p.household.head_name])
                else:
                    rows.append([len(rows) + 1, p.full_name, p.personal_id, p.date_of_birth.year,
                                 p.gender, p.ethnic_group, p.household.address] +
                                [d.get(k) for k in ("recognized_year", "support_status", "support_start_year",
                                                   "support_result", "support_duration", "note")])
            list_rows(ws, rows)
        else:
            outcomes = []
            rows = []
            for c, group in groups:
                values = [len(rows) + 1, c.name if not school_scope else scope_label] + unit_values(group, form)
                if form == "09":
                    result = conclusion(group, c, not school_scope)
                    outcomes.append(result)
                    values.append(result)
                rows.append(values)
            total = [None, "Cộng"] + unit_values(items, form)
            if form == "09":
                total.append(None)
            percent_columns = {5, 7, 10, 12} if form == "09" else {8, 10, 12, 14, 20, 22, 24, 26}
            list_rows(ws, rows, total, percent_columns=percent_columns)
            if form == "09":
                for row in ws:
                    for cell in row:
                        if str(cell.value or "").startswith("Số đơn vị cấp xã đạt chuẩn"):
                            target = ws.cell(cell.row, 11)
                            level = 1 if str(cell.value).endswith("1") else 2
                            target.value = (sum(isinstance(x, int) and x >= level for x in outcomes)
                                            if not school_scope and all(isinstance(x, int) or x == "Chưa đạt" for x in outcomes) and outcomes else None)
    for ws in wb:
        for row in ws:
            for cell in row:
                value = str(cell.value or "")
                if value.startswith("Tỉnh/Thành phố:"):
                    province_name = _scope_meta(scope_label="", selected_commune_id=None, selected_school_id=None, year=year)["place"]
                    put(ws, cell.row, cell.column, "Tỉnh/Thành phố: " + province_name)
                elif value.startswith("Xã/Phường/Đặc khu:"):
                    put(ws, cell.row, cell.column, "Phạm vi: " + scope_label)
                elif value.startswith("Năm học:"):
                    put(ws, cell.row, cell.column, "Năm học: " + year_code)
                elif value.startswith("Thời điểm:"):
                    put(ws, cell.row, cell.column, "Thời điểm xuất: " + date.today().strftime("%d/%m/%Y"))
        ws.print_area = ws.dimensions
    audit = wb.create_sheet("Kiem_tra_du_lieu")
    audit.append(["Nội dung", "Giá trị"])
    for row in (("Phạm vi", scope_label), ("Năm học", year_code),
                ("Tuổi", f"Năm {year} trừ năm sinh (tuổi theo năm điều tra)."),
                ("Căn cứ", LEGAL_BASIS),
                ("Trạng thái áp dụng", "Tính thử theo quy định mới trước ngày hiệu lực" if date.today() < EFFECTIVE_DATE else "Theo quy định mới"),
                ("Nguồn mức biết chữ", "Điều 23: giai đoạn 1 XMC hoặc hoàn thành lớp 3; giai đoạn 2 XMC hoặc hoàn thành chương trình tiểu học. Kế thừa chỉ báo mức 1/2 đã xác nhận trên phiếu cũ; tái mù đã xác minh không tính là đạt hiện tại."),
                ("Quy tắc ô trống", "Chưa đủ dữ liệu hoặc mẫu số bằng 0. Không coi hồ sơ chưa xác minh là không đạt."),
                ("Mẫu 01 - Đạt chuẩn biết chữ", {"g1":"Mức độ 1", "g2":"Mức độ 2"}.get(literacy_basis, "Chưa xác nhận quy ước mức độ")),
                ("Mẫu 05/06", "Học viên và công nhận trong năm từ phần dữ liệu bổ sung XMC; không suy từ dân số biết chữ."),
                ("Mẫu 09", "Điều 24–25: ít nhất 90% nhóm 15–60 đạt mức tương ứng; xã đặc biệt khó khăn dùng nhóm 15–35. Chỉ đánh giá chỉ tiêu dân số, không thay quyết định công nhận; không xét cho trường/giáo viên. Số xã cuối biểu là số xã đáp ứng chỉ tiêu tính toán, không phải số quyết định công nhận."),
                ("Cấp tỉnh", "Điều 24–25 yêu cầu 100% xã được công nhận mức tương ứng. Không suy quyết định công nhận cấp tỉnh từ tỷ lệ dân số toàn tỉnh."),
                ("Danh sách", "Mẫu 02/08 chỉ gồm người có xác nhận phù hợp; hồ sơ thiếu nguồn được liệt kê phía dưới."),
                ("Số đối tượng trong phạm vi", len(items))):
        audit.append(row)
    audit.append(["Mã đối tượng", "Họ tên", "Thông tin cần bổ sung"])
    for x in items:
        missing = []
        if x["age"] is None:
            missing.append("Ngày sinh")
        if x["age"] is not None and 15 <= x["age"] <= 60:
            for k in ("g1", "g2"):
                if x[k] is None:
                    missing.append("Mức biết chữ " + k[-1])
            from app.xmc_models import BOOL_FIELDS
            for k in ("illiterate", "enrolled1", "enrolled2", "relapsed"):
                if x["data"].get(k) is None:
                    missing.append(BOOL_FIELDS[k])
            for level in ("1", "2"):
                if x["data"].get("enrolled" + level) is True and x["data"].get("recognized" + level) is None:
                    missing.append(BOOL_FIELDS["recognized" + level])
        if x["age"] is not None and 0 <= x["age"] <= 60:
            r = x["record"]
            if r is None:
                missing.append("Hồ sơ năm học")
            elif r.learning_status == "CHUA_XAC_DINH":
                missing.append("Tình trạng học tập (mẫu 07)")
            if r is None or r.disability_status == "CHUA_XAC_DINH":
                missing.append("Tình trạng khuyết tật (mẫu 07)")
            if x["female"] is None:
                missing.append("Giới tính")
            if x["ethnic"] is None:
                missing.append("Dân tộc")
        if x["conflict"]:
            missing.append("Mâu thuẫn giữa mức biết chữ và tình trạng chưa biết chữ/tái mù chữ")
        if missing:
            audit.append([x["person"].code, x["person"].full_name, "; ".join(missing)])
    for col, width in (("A", 35), ("B", 65), ("C", 90)):
        audit.column_dimensions[col].width = width
    for row in audit:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, str):
                cell.data_type = "s"
    stream = BytesIO()
    from app.services.report_export_date import stamp_workbook
    stamp_workbook(wb)
    wb.save(stream)
    stream.seek(0)
    return stream
