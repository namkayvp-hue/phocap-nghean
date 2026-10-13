"""Join PDF page fragments into continuous Excel forms at export time."""
from copy import copy

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import RowBreak, ColBreak


def column_maps(sheets):
    boundaries = []
    for ws in sheets:
        widths = [ws.column_dimensions[get_column_letter(c)].width for c in range(1, ws.max_column + 1)]
        total = sum(widths)
        points = [0]
        for width in widths:
            points.append(points[-1] + width / total)
        boundaries.append([round(p, 6) for p in points])
    grid = sorted(set(p for points in boundaries for p in points))
    return grid, [[grid.index(p) + 1 for p in points] for points in boundaries]


def join_compulsory_pages(wb):
    for first, second in [('Trang_04', 'Trang_05'), ('Trang_09', 'Trang_10')]:
        if first not in wb.sheetnames or second not in wb.sheetnames:
            continue
        sources = [wb[first], wb[second]]
        grid, maps = column_maps(sources)
        target = wb.create_sheet(first + '_joined', wb.index(sources[0]))
        width = sum(sources[0].column_dimensions[get_column_letter(c)].width for c in range(1, sources[0].max_column + 1))
        for col in range(1, len(grid)):
            target.column_dimensions[get_column_letter(col)].width = (grid[col] - grid[col - 1]) * width
        offset = 0
        for source, mapping in zip(sources, maps):
            merged = {(m.min_row, m.min_col): m for m in source.merged_cells.ranges}
            for row in source:
                for cell in row:
                    if isinstance(cell, MergedCell):
                        continue
                    end_row, end_col = cell.row, cell.column
                    area = merged.get((cell.row, cell.column))
                    if area:
                        end_row, end_col = area.max_row, area.max_col
                    left, right = mapping[cell.column - 1], mapping[end_col] - 1
                    dest = target.cell(offset + cell.row, left, cell.value)
                    dest._style = copy(cell._style)
                    dest.data_type = cell.data_type
                    dest.comment = copy(cell.comment)
                    if right > left or end_row > cell.row:
                        target.merge_cells(start_row=offset + cell.row, end_row=offset + end_row,
                                           start_column=left, end_column=right)
            for row, dim in source.row_dimensions.items():
                target.row_dimensions[offset + row].height = dim.height
            offset += source.max_row
        target.page_setup = copy(sources[0].page_setup)
        target.page_margins = copy(sources[0].page_margins)
        target.sheet_properties = copy(sources[0].sheet_properties)
        target.print_options = copy(sources[0].print_options)
        target.print_title_rows = '7:8'
        target.freeze_panes = 'A9'
        target.sheet_view.showGridLines = False
        for source in sources:
            wb.remove(source)
        target.title = first
    for ws in wb:
        if ws.title in {'Trang_04', 'Trang_06', 'Trang_07', 'Trang_08', 'Trang_09', 'Trang_11', 'Trang_12', 'Trang_13'}:
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.row_breaks = RowBreak()
            ws.col_breaks = ColBreak()
            ws.print_area = ws.dimensions
            if not ws.print_title_rows:
                ws.print_title_rows = {'Trang_06':'7:11', 'Trang_07':'7:10', 'Trang_08':'8:8',
                                       'Trang_11':'7:10', 'Trang_12':'8:11', 'Trang_13':'8:8'}.get(ws.title, '7:8')
    return wb
