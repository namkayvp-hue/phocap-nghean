from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from io import BytesIO
from math import ceil
from typing import Any
from urllib.parse import quote, urlencode

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.database import get_db
from app.models import Commune, Role, School, SchoolYear, User
from app.permissions import ADMIN_ROLE_CODES, SCHOOL_ROLE_CODE, normalize_role_code
from app.routers.report_center import (
    _choose_year_id,
    _load_school_years,
    _load_scope_options,
    _parse_optional_int,
)
from app.routers.surveys import lay_thong_tin_nguoi_dung
from app.staff_models import StaffMember, StaffYearRecord, SchoolStaffYearSummary


router = APIRouter(prefix="/doi-ngu", tags=["Quản lý đội ngũ"])
templates = Jinja2Templates(directory="app/templates")

STATUS_LABELS = {
    "DANG_LAM_VIEC": "Đang làm việc",
    "TAM_NGHI": "Tạm nghỉ",
    "CHUYEN_DI": "Chuyển đi",
    "NGHI_VIEC": "Nghỉ việc",
    "NGHI_HUU": "Nghỉ hưu",
}
POSITION_LABELS = {
    "CBQL": "Cán bộ quản lý",
    "GIAO_VIEN": "Giáo viên",
    "NHAN_VIEN": "Nhân viên",
}
EMPLOYMENT_LABELS = {
    "CHUA_XAC_DINH": "Chưa xác định",
    "BIEN_CHE": "Biên chế",
    "HOP_DONG_LAM_VIEC": "Hợp đồng làm việc",
    "HOP_DONG_LAO_DONG": "Hợp đồng lao động",
}
TEACHING_LEVEL_LABELS = {
    "KHONG_DAY": "Không trực tiếp giảng dạy",
    "NHA_TRE": "Nhà trẻ",
    "MAU_GIAO": "Mẫu giáo",
    "TIEU_HOC": "Tiểu học",
    "THCS": "Trung học cơ sở",
    "LIEN_CAP_TH_THCS": "Tiểu học & THCS",
}
TEACHING_AGE_LABELS = {
    "KHONG_XAC_DINH": "Chưa xác định",
    "NHA_TRE": "Nhà trẻ",
    "TUOI_3_4": "Mẫu giáo 3–4 tuổi",
    "TUOI_5": "Mẫu giáo 5 tuổi",
}
QUALIFICATION_LEVEL_LABELS = {
    "CHUA_XAC_DINH": "Chưa xác định",
    "TRUNG_CAP": "Trung cấp",
    "CAO_DANG": "Cao đẳng",
    "DAI_HOC": "Đại học",
    "THAC_SI": "Thạc sĩ",
    "TIEN_SI": "Tiến sĩ",
    "KHAC": "Khác",
}
QUALIFICATION_STANDARD_LABELS = {
    "CHUA_XAC_DINH": "Chưa xác định",
    "CHUA_DAT": "Chưa đạt chuẩn",
    "DAT_CHUAN": "Đạt chuẩn",
    "TREN_CHUAN": "Trên chuẩn",
}
PROFESSIONAL_STANDARD_LABELS = {
    "CHUA_DANH_GIA": "Chưa đánh giá",
    "CHUA_DAT": "Chưa đạt",
    "DAT": "Đạt",
    "KHA": "Khá",
    "TOT": "Tốt",
}
INSTITUTION_GROUP_LABELS = {
    "TRUONG_MAM_NON": "Trường mầm non",
    "CO_SO_GDMN_DOC_LAP": "Cơ sở GDMN độc lập",
}


def _clean(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _normalize(value: Any) -> str:
    text = _clean(value).lower().replace("đ", "d")
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_nonnegative_int(value: Any) -> int:
    try:
        result = int(str(value or "0").strip())
    except (TypeError, ValueError):
        return 0
    return max(result, 0)


def _can_manage(user: dict[str, Any]) -> bool:
    role = normalize_role_code(user.get("role_code"))
    return role in ADMIN_ROLE_CODES or role == SCHOOL_ROLE_CODE


def _ensure_school_in_scope(user: dict[str, Any], school_id: int) -> bool:
    role = normalize_role_code(user.get("role_code"))
    if role in ADMIN_ROLE_CODES:
        return True
    if role == SCHOOL_ROLE_CODE:
        return user.get("school_id") is not None and int(user["school_id"]) == int(school_id)
    return False


def _next_staff_code(db: Session) -> str:
    max_id = db.scalar(select(func.max(StaffMember.id))) or 0
    return f"NS-{int(max_id) + 1:08d}"


def _load_page_scope(
    db: Session,
    request: Request,
    school_year_id: int | None,
    commune_id: int | None,
    school_id: int | None,
) -> dict[str, Any]:
    user = lay_thong_tin_nguoi_dung(request)
    years = _load_school_years(db)
    selected_year_id = _choose_year_id(years, school_year_id)
    communes, schools, selected_commune_id, selected_school_id, scope_label = _load_scope_options(
        db, user, commune_id, school_id
    )

    # V4.2: trường đã sáp nhập/ngừng hoạt động vẫn phải tra cứu được ở
    # những năm học mà trường còn hồ sơ đội ngũ. Không đưa trường inactive
    # vào dữ liệu năm hiện hành nếu không có hồ sơ của chính năm đó.
    if selected_year_id is not None:
        role = normalize_role_code(user.get("role_code"))
        archived_statement = (
            select(School)
            .join(StaffYearRecord, StaffYearRecord.school_id == School.id)
            .where(
                School.is_active.is_(False),
                StaffYearRecord.school_year_id == selected_year_id,
            )
            .distinct()
        )
        if role == SCHOOL_ROLE_CODE and user.get("school_id") is not None:
            archived_statement = archived_statement.where(
                School.id == int(user["school_id"])
            )
        elif role not in ADMIN_ROLE_CODES and user.get("commune_id") is not None:
            archived_statement = archived_statement.where(
                School.commune_id == int(user["commune_id"])
            )
        elif selected_commune_id is not None:
            archived_statement = archived_statement.where(
                School.commune_id == selected_commune_id
            )

        archived_schools = list(db.scalars(archived_statement).all())
        school_map = {int(item.id): item for item in schools}
        for item in archived_schools:
            school_map.setdefault(int(item.id), item)
        schools = sorted(
            school_map.values(),
            key=lambda item: (_normalize(item.name), int(item.id)),
        )

        if school_id is not None and int(school_id) in school_map:
            selected_school_id = int(school_id)
            selected_school = school_map[int(school_id)]
            suffix = "" if selected_school.is_active else " (đã sáp nhập/ngừng hoạt động)"
            scope_label = f"Trường {selected_school.name}{suffix}"

    return {
        "user": user,
        "school_years": years,
        "selected_year_id": selected_year_id,
        "communes": communes,
        "schools": schools,
        "selected_commune_id": selected_commune_id,
        "selected_school_id": selected_school_id,
        "scope_label": scope_label,
    }


def _scope_conditions(scope: dict[str, Any]) -> list[Any]:
    conditions: list[Any] = [StaffYearRecord.school_year_id == scope["selected_year_id"]]
    if scope["selected_school_id"] is not None:
        conditions.append(StaffYearRecord.school_id == scope["selected_school_id"])
    elif scope["selected_commune_id"] is not None:
        conditions.append(School.commune_id == scope["selected_commune_id"])
    return conditions


def _infer_institution_group(school_name: str) -> str:
    normalized = _normalize(school_name)
    independent_tokens = (
        "nhom tre",
        "lop mam non",
        "co so mn",
        "csmn",
        "doc lap",
        "tgd ",
        "nhom lop",
    )
    return (
        "CO_SO_GDMN_DOC_LAP"
        if any(token in normalized for token in independent_tokens)
        else "TRUONG_MAM_NON"
    )


IMPORT_HEADERS = [
    "STT",
    "Năm học",
    "Mã xã/phường",
    "Tên xã/phường",
    "Mã trường",
    "Tên trường",
    "Tên đăng nhập hiện có",
    "Mã cán bộ",
    "Mã định danh Bộ GD&ĐT",
    "Họ và tên",
    "Ngày sinh",
    "Giới tính",
    "Dân tộc",
    "Số định danh cá nhân",
    "Số điện thoại",
    "Email",
    "Vị trí việc làm",
    "Chức vụ",
    "Hình thức tuyển dụng",
    "Loại hợp đồng",
    "Ngày tuyển dụng",
    "Trình độ đào tạo",
    "Mức chuẩn đào tạo",
    "Kết quả chuẩn nghề nghiệp",
    "Cấp học trực tiếp",
    "Môn dạy",
    "Chuyên ngành chính",
    "Đánh giá viên chức",
    "Dạy nhóm/lớp",
    "Hưởng chế độ, chính sách",
    "Trạng thái công tác",
    "Ghi chú",
]

THIN_BORDER = Border(
    left=Side(style="thin", color="B7C9DC"),
    right=Side(style="thin", color="B7C9DC"),
    top=Side(style="thin", color="B7C9DC"),
    bottom=Side(style="thin", color="B7C9DC"),
)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")


def _account_staff_code(username: str | None) -> str:
    value = _clean(username)
    if "." in value:
        return value.split(".", 1)[1].strip()
    return value


def _eligible_account_statement(
    user: dict[str, Any],
    selected_commune_id: int | None = None,
    selected_school_id: int | None = None,
):
    statement = (
        select(User)
        .join(User.role)
        .join(User.school)
        .options(selectinload(User.role), selectinload(User.school).selectinload(School.commune))
        .where(
            User.is_active.is_(True),
            User.school_id.is_not(None),
            Role.code.in_([SCHOOL_ROLE_CODE, "GIAO_VIEN"]),
        )
    )
    role = normalize_role_code(user.get("role_code"))
    if role == SCHOOL_ROLE_CODE and user.get("school_id") is not None:
        statement = statement.where(User.school_id == int(user["school_id"]))
    elif role not in ADMIN_ROLE_CODES and user.get("commune_id") is not None:
        statement = statement.where(School.commune_id == int(user["commune_id"]))
    if selected_school_id is not None:
        statement = statement.where(User.school_id == selected_school_id)
    elif selected_commune_id is not None:
        statement = statement.where(School.commune_id == selected_commune_id)
    return statement.order_by(School.name, User.full_name)


def _staff_record_conditions(scope: dict[str, Any]) -> list[Any]:
    return _scope_conditions(scope)


def _build_data_audit(db: Session, scope: dict[str, Any]) -> dict[str, Any]:
    accounts = list(
        db.scalars(
            _eligible_account_statement(
                scope["user"],
                scope["selected_commune_id"],
                scope["selected_school_id"],
            )
        ).all()
    )
    conditions = _staff_record_conditions(scope)
    records = list(
        db.scalars(
            select(StaffYearRecord)
            .join(StaffYearRecord.staff_member)
            .join(StaffYearRecord.school)
            .options(selectinload(StaffYearRecord.staff_member), selectinload(StaffYearRecord.school))
            .where(*conditions)
        ).all()
    )
    converted_codes = {
        _clean(record.staff_member.ministry_staff_code)
        for record in records
        if _clean(record.staff_member.ministry_staff_code)
    }
    unconverted_accounts = [
        account
        for account in accounts
        if _account_staff_code(account.username) not in converted_codes
    ]
    account_school_ids = {int(account.school_id) for account in accounts if account.school_id}
    record_school_ids = {int(record.school_id) for record in records}
    incomplete_records = [
        record
        for record in records
        if record.position_group == "GIAO_VIEN"
        and (
            record.qualification_standard == "CHUA_XAC_DINH"
            or record.professional_standard == "CHUA_DANH_GIA"
            or record.teaching_age_group == "KHONG_XAC_DINH"
        )
    ]
    no_code_records = [record for record in records if not _clean(record.staff_member.ministry_staff_code)]
    ministry_counts: dict[str, int] = {}
    personal_counts: dict[str, int] = {}
    for record in records:
        ministry = _clean(record.staff_member.ministry_staff_code)
        personal = _clean(record.staff_member.personal_id)
        if ministry:
            ministry_counts[ministry] = ministry_counts.get(ministry, 0) + 1
        if personal:
            personal_counts[personal] = personal_counts.get(personal, 0) + 1
    duplicate_codes = sum(value > 1 for value in ministry_counts.values())
    duplicate_personal_ids = sum(value > 1 for value in personal_counts.values())
    return {
        "eligible_accounts": len(accounts),
        "staff_records": len(records),
        "converted_accounts": len(accounts) - len(unconverted_accounts),
        "unconverted_accounts": len(unconverted_accounts),
        "incomplete_records": len(incomplete_records),
        "records_without_code": len(no_code_records),
        "schools_with_accounts": len(account_school_ids),
        "schools_with_records": len(record_school_ids),
        "schools_without_records": len(account_school_ids - record_school_ids),
        "duplicate_codes": duplicate_codes,
        "duplicate_personal_ids": duplicate_personal_ids,
    }


def _parse_page(value: Any, default: int, minimum: int, maximum: int) -> int:
    try:
        result = int(str(value or default).strip())
    except (TypeError, ValueError):
        result = default
    return min(max(result, minimum), maximum)


def _excel_response(workbook: Workbook, filename: str) -> StreamingResponse:
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    encoded = quote(filename)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )


def _style_header(ws: Any, row: int, start_col: int, end_col: int) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row, col)
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def _style_data_region(ws: Any, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _set_column_widths(ws: Any, widths: dict[int, float]) -> None:
    from openpyxl.utils import get_column_letter

    for column_index, width in widths.items():
        ws.column_dimensions[get_column_letter(column_index)].width = width


def _write_school_catalog(ws: Any, schools: list[School]) -> None:
    headers = ["STT", "Mã xã/phường", "Tên xã/phường", "Mã trường", "Tên trường"]
    ws.append(headers)
    _style_header(ws, 1, 1, len(headers))
    for index, school in enumerate(schools, start=1):
        ws.append([
            index,
            school.commune.code if school.commune else "",
            school.commune.name if school.commune else "",
            school.code,
            school.name,
        ])
    _style_data_region(ws, 2, max(ws.max_row, 2), 1, len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{max(ws.max_row, 1)}"
    _set_column_widths(ws, {1: 7, 2: 16, 3: 28, 4: 18, 5: 45})
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 2).number_format = "@"
        ws.cell(row, 4).number_format = "@"


def _write_choice_catalog(ws: Any) -> dict[str, tuple[int, int]]:
    catalogs = {
        "GioiTinh": ["Nam", "Nữ", "Khác", "Chưa xác định"],
        "ViTri": list(POSITION_LABELS.values()),
        "TuyenDung": list(EMPLOYMENT_LABELS.values()),
        "TrinhDo": list(QUALIFICATION_LEVEL_LABELS.values()),
        "MucChuan": list(QUALIFICATION_STANDARD_LABELS.values()),
        "ChuanNgheNghiep": list(PROFESSIONAL_STANDARD_LABELS.values()),
        "CapHoc": list(TEACHING_LEVEL_LABELS.values()),
        "DayNhomLop": list(TEACHING_AGE_LABELS.values()),
        "CoKhong": ["Có", "Không"],
        "TrangThai": list(STATUS_LABELS.values()),
    }
    ranges: dict[str, tuple[int, int]] = {}
    for col_index, (name, values) in enumerate(catalogs.items(), start=1):
        ws.cell(1, col_index, name)
        ws.cell(1, col_index).fill = HEADER_FILL
        ws.cell(1, col_index).font = Font(color="FFFFFF", bold=True)
        for row_index, value in enumerate(values, start=2):
            ws.cell(row_index, col_index, value)
        ranges[name] = (col_index, len(values) + 1)
    ws.sheet_state = "hidden"
    return ranges


def _add_import_validations(ws: Any, ranges: dict[str, tuple[int, int]]) -> None:
    from openpyxl.utils import get_column_letter

    header_to_catalog = {
        "Giới tính": "GioiTinh",
        "Vị trí việc làm": "ViTri",
        "Hình thức tuyển dụng": "TuyenDung",
        "Loại hợp đồng": "TuyenDung",
        "Trình độ đào tạo": "TrinhDo",
        "Mức chuẩn đào tạo": "MucChuan",
        "Kết quả chuẩn nghề nghiệp": "ChuanNgheNghiep",
        "Cấp học trực tiếp": "CapHoc",
        "Dạy nhóm/lớp": "DayNhomLop",
        "Hưởng chế độ, chính sách": "CoKhong",
        "Trạng thái công tác": "TrangThai",
    }
    header_index = {value: index + 1 for index, value in enumerate(IMPORT_HEADERS)}
    for header, catalog in header_to_catalog.items():
        col_index = header_index[header]
        catalog_col, last_row = ranges[catalog]
        formula = f"'Danh muc lua chon'!${get_column_letter(catalog_col)}$2:${get_column_letter(catalog_col)}${last_row}"
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(validation)
        validation.add(f"{get_column_letter(col_index)}2:{get_column_letter(col_index)}50000")


def _create_import_template_workbook(
    db: Session,
    scope: dict[str, Any],
) -> Workbook:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Nhap danh sach"
    ws.append(IMPORT_HEADERS)
    _style_header(ws, 1, 1, len(IMPORT_HEADERS))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(IMPORT_HEADERS)).column_letter}1"
    ws.row_dimensions[1].height = 42
    _set_column_widths(
        ws,
        {
            1: 7, 2: 13, 3: 16, 4: 28, 5: 18, 6: 42, 7: 24, 8: 18, 9: 24,
            10: 30, 11: 14, 12: 14, 13: 20, 14: 18, 15: 28, 16: 22, 17: 24,
            18: 24, 19: 24, 20: 16, 21: 22, 22: 22, 23: 26, 24: 22, 25: 25,
            26: 24, 27: 22, 28: 35,
        },
    )
    selected_year_code = next(
        (year.code for year in scope["school_years"] if year.id == scope["selected_year_id"]),
        scope["school_years"][0].code if scope["school_years"] else "",
    )
    for row in range(2, 200):
        ws.cell(row, 1, row - 1)
        ws.cell(row, 2, selected_year_code)
    text_columns = [3, 5, 7, 8, 9, 13, 14]
    for col in text_columns:
        for row in range(2, 50001):
            ws.cell(row, col).number_format = "@"
    choice_ws = workbook.create_sheet("Danh muc lua chon")
    choice_ranges = _write_choice_catalog(choice_ws)
    _add_import_validations(ws, choice_ranges)
    schools = list(
        db.scalars(
            select(School)
            .options(selectinload(School.commune))
            .where(School.is_active.is_(True))
            .order_by(School.name)
        ).all()
    )
    school_ws = workbook.create_sheet("Danh muc truong")
    _write_school_catalog(school_ws, schools)
    accounts = list(
        db.scalars(
            _eligible_account_statement(
                scope["user"],
                scope["selected_commune_id"],
                scope["selected_school_id"],
            )
        ).all()
    )
    codes_in_year = set(
        db.scalars(
            select(StaffMember.ministry_staff_code)
            .join(StaffYearRecord, StaffYearRecord.staff_member_id == StaffMember.id)
            .join(School, School.id == StaffYearRecord.school_id)
            .where(*_scope_conditions(scope), StaffMember.ministry_staff_code.is_not(None))
        ).all()
    )
    suggestion_ws = workbook.create_sheet("Tai khoan goi y")
    suggestion_headers = [
        "STT", "Năm học", "Mã xã/phường", "Tên xã/phường", "Mã trường", "Tên trường",
        "Tên đăng nhập", "Mã định danh gợi ý", "Họ và tên", "Vai trò tài khoản", "Đã chuyển vào đội ngũ",
    ]
    suggestion_ws.append(suggestion_headers)
    _style_header(suggestion_ws, 1, 1, len(suggestion_headers))
    for index, account in enumerate(accounts, start=1):
        code = _account_staff_code(account.username)
        suggestion_ws.append([
            index,
            next((year.code for year in scope["school_years"] if year.id == scope["selected_year_id"]), ""),
            account.school.commune.code if account.school and account.school.commune else "",
            account.school.commune.name if account.school and account.school.commune else "",
            account.school.code if account.school else "",
            account.school.name if account.school else "",
            account.username,
            code,
            account.full_name,
            account.role.name if account.role else "",
            "Có" if code in codes_in_year else "Chưa",
        ])
    _style_data_region(suggestion_ws, 2, max(suggestion_ws.max_row, 2), 1, len(suggestion_headers))
    suggestion_ws.freeze_panes = "A2"
    suggestion_ws.auto_filter.ref = f"A1:K{max(suggestion_ws.max_row, 1)}"
    _set_column_widths(suggestion_ws, {1: 7, 2: 13, 3: 16, 4: 28, 5: 18, 6: 42, 7: 24, 8: 24, 9: 30, 10: 22, 11: 22})
    for row in range(2, suggestion_ws.max_row + 1):
        for col in (3, 5, 7, 8):
            suggestion_ws.cell(row, col).number_format = "@"
    guide_ws = workbook.create_sheet("Huong dan nhap")
    guide_rows = [
        ["HƯỚNG DẪN NHẬP DANH SÁCH ĐỘI NGŨ"],
        ["1", "Nhập dữ liệu tại trang 'Nhap danh sach'; không đổi tên các cột ở dòng 1."],
        ["2", "Có thể sao chép thông tin từ trang 'Tai khoan goi y' để chuyển lại danh sách tài khoản hiện có."],
        ["3", "Bắt buộc có Họ và tên, Mã trường hoặc Tên trường. Nên có Mã định danh Bộ GD&ĐT để chống trùng."],
        ["4", "Thứ tự nhận diện khi nhập: Mã định danh Bộ GD&ĐT → Số định danh cá nhân → Họ tên + ngày sinh."],
        ["5", "Dòng đã có sẽ được cập nhật; dòng không xác định được trường hoặc bị mâu thuẫn định danh sẽ bị bỏ qua."],
        ["6", "Các ô có danh sách chọn phải dùng đúng giá trị trong danh mục để báo cáo MN-01 GV chính xác."],
        ["7", "Không nhập số liệu minh họa. Chỉ nhập dữ liệu thực tế của đơn vị và năm học được chọn."],
    ]
    for row in guide_rows:
        guide_ws.append(row)
    guide_ws.merge_cells("A1:F1")
    guide_ws["A1"].fill = HEADER_FILL
    guide_ws["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    guide_ws["A1"].alignment = Alignment(horizontal="center")
    _set_column_widths(guide_ws, {1: 8, 2: 110})
    for row in range(2, guide_ws.max_row + 1):
        guide_ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    return workbook


def _create_current_data_workbook(db: Session, scope: dict[str, Any]) -> Workbook:
    workbook = Workbook()
    summary_ws = workbook.active
    summary_ws.title = "Thong ke"
    audit = _build_data_audit(db, scope)
    summary_ws.append(["KIỂM TRA DỮ LIỆU ĐỘI NGŨ HIỆN CÓ", "Số lượng"])
    _style_header(summary_ws, 1, 1, 2)
    labels = [
        ("Tài khoản trường/giáo viên đủ điều kiện", audit["eligible_accounts"]),
        ("Hồ sơ đội ngũ trong năm học", audit["staff_records"]),
        ("Tài khoản đã chuyển", audit["converted_accounts"]),
        ("Tài khoản chưa chuyển", audit["unconverted_accounts"]),
        ("Trường có tài khoản", audit["schools_with_accounts"]),
        ("Trường đã có hồ sơ đội ngũ", audit["schools_with_records"]),
        ("Trường chưa có hồ sơ đội ngũ", audit["schools_without_records"]),
        ("Giáo viên thiếu thông tin báo cáo", audit["incomplete_records"]),
        ("Hồ sơ thiếu mã định danh Bộ GD&ĐT", audit["records_without_code"]),
        ("Nhóm mã cán bộ bị trùng", audit["duplicate_codes"]),
        ("Nhóm số định danh cá nhân bị trùng", audit["duplicate_personal_ids"]),
    ]
    for label, value in labels:
        summary_ws.append([label, value])
    _style_data_region(summary_ws, 2, summary_ws.max_row, 1, 2)
    _set_column_widths(summary_ws, {1: 52, 2: 18})
    records_ws = workbook.create_sheet("Ho so doi ngu")
    records_ws.append(IMPORT_HEADERS + ["Trạng thái hoàn thiện báo cáo"])
    _style_header(records_ws, 1, 1, len(IMPORT_HEADERS) + 1)
    records = list(
        db.scalars(
            select(StaffYearRecord)
            .join(StaffYearRecord.staff_member)
            .join(StaffYearRecord.school)
            .options(
                selectinload(StaffYearRecord.staff_member),
                selectinload(StaffYearRecord.school).selectinload(School.commune),
                selectinload(StaffYearRecord.school_year),
            )
            .where(*_scope_conditions(scope))
            .order_by(School.name, StaffMember.full_name)
        ).all()
    )
    for index, record in enumerate(records, start=1):
        member = record.staff_member
        school = record.school
        incomplete = (
            record.position_group == "GIAO_VIEN"
            and (
                record.qualification_standard == "CHUA_XAC_DINH"
                or record.professional_standard == "CHUA_DANH_GIA"
                or record.teaching_age_group == "KHONG_XAC_DINH"
            )
        )
        records_ws.append([
            index,
            record.school_year.code,
            school.commune.code if school.commune else "",
            school.commune.name if school.commune else "",
            school.code,
            school.name,
            "",
            member.code,
            member.ministry_staff_code or "",
            member.full_name,
            member.date_of_birth,
            member.gender or "",
            member.ethnic_group or "",
            member.personal_id or "",
            member.phone or "",
            member.email or "",
            POSITION_LABELS.get(record.position_group, record.position_group),
            record.position_title or "",
            EMPLOYMENT_LABELS.get(record.employment_type, record.employment_type),
            EMPLOYMENT_LABELS.get(record.employment_type, record.employment_type),
            record.recruitment_date,
            QUALIFICATION_LEVEL_LABELS.get(record.qualification_level, record.qualification_level),
            QUALIFICATION_STANDARD_LABELS.get(record.qualification_standard, record.qualification_standard),
            PROFESSIONAL_STANDARD_LABELS.get(record.professional_standard, record.professional_standard),
            TEACHING_LEVEL_LABELS.get(record.teaching_level, record.teaching_level),
            record.teaching_subject or "",
            record.main_specialty or "",
            record.staff_evaluation or "",
            TEACHING_AGE_LABELS.get(record.teaching_age_group, record.teaching_age_group),
            "Có" if record.receives_policy else "Không",
            STATUS_LABELS.get(record.status_code, record.status_code),
            record.notes or "",
            "Thiếu thông tin" if incomplete else "Đủ dữ liệu cơ bản",
        ])
    _style_data_region(records_ws, 2, max(records_ws.max_row, 2), 1, len(IMPORT_HEADERS) + 1)
    records_ws.freeze_panes = "A2"
    records_ws.auto_filter.ref = f"A1:{records_ws.cell(1, len(IMPORT_HEADERS)+1).column_letter}{max(records_ws.max_row,1)}"
    _set_column_widths(records_ws, {1: 7, 2: 13, 3: 16, 4: 28, 5: 18, 6: 42, 7: 24, 8: 18, 9: 24, 10: 30, 11: 14, 12: 14, 13: 20, 14: 18, 15: 28, 16: 22, 17: 24, 18: 24, 19: 24, 20: 16, 21: 22, 22: 22, 23: 26, 24: 22, 25: 25, 26: 24, 27: 22, 28: 35, 29: 24})
    for row in range(2, records_ws.max_row + 1):
        records_ws.cell(row, 11).number_format = "dd/mm/yyyy"
        records_ws.cell(row, 20).number_format = "dd/mm/yyyy"
        for col in (3, 5, 7, 8, 9, 13, 14):
            records_ws.cell(row, col).number_format = "@"
    accounts_ws = workbook.create_sheet("Tai khoan chua chuyen")
    headers = ["STT", "Năm học", "Mã xã/phường", "Tên xã/phường", "Mã trường", "Tên trường", "Tên đăng nhập", "Mã định danh gợi ý", "Họ và tên", "Vai trò"]
    accounts_ws.append(headers)
    _style_header(accounts_ws, 1, 1, len(headers))
    converted_codes = {
        _clean(record.staff_member.ministry_staff_code)
        for record in records
        if _clean(record.staff_member.ministry_staff_code)
    }
    accounts = list(db.scalars(_eligible_account_statement(scope["user"], scope["selected_commune_id"], scope["selected_school_id"])).all())
    year_code = next((year.code for year in scope["school_years"] if year.id == scope["selected_year_id"]), "")
    row_index = 0
    for account in accounts:
        code = _account_staff_code(account.username)
        if code in converted_codes:
            continue
        row_index += 1
        accounts_ws.append([
            row_index,
            year_code,
            account.school.commune.code if account.school and account.school.commune else "",
            account.school.commune.name if account.school and account.school.commune else "",
            account.school.code if account.school else "",
            account.school.name if account.school else "",
            account.username,
            code,
            account.full_name,
            account.role.name if account.role else "",
        ])
    _style_data_region(accounts_ws, 2, max(accounts_ws.max_row, 2), 1, len(headers))
    accounts_ws.freeze_panes = "A2"
    accounts_ws.auto_filter.ref = f"A1:J{max(accounts_ws.max_row,1)}"
    _set_column_widths(accounts_ws, {1: 7, 2: 13, 3: 16, 4: 28, 5: 18, 6: 42, 7: 24, 8: 24, 9: 30, 10: 22})
    schools = list(db.scalars(select(School).options(selectinload(School.commune)).where(School.is_active.is_(True)).order_by(School.name)).all())
    school_ws = workbook.create_sheet("Danh muc truong")
    _write_school_catalog(school_ws, schools)
    return workbook


def _map_label_code(value: Any, labels: dict[str, str], default: str) -> str:
    normalized = _normalize(value)
    if not normalized:
        return default
    for code, label in labels.items():
        if normalized == _normalize(code) or normalized == _normalize(label):
            return code
    for code, label in labels.items():
        label_norm = _normalize(label)
        if normalized in label_norm or label_norm in normalized:
            return code
    return default


def _map_boolean(value: Any) -> bool:
    return _normalize(value) in {"co", "1", "true", "x", "yes"}


def _match_staff_member(
    db: Session,
    ministry_code: str,
    personal_id: str,
    full_name: str,
    birth_date: date | None,
) -> tuple[StaffMember | None, str | None]:
    by_code = None
    by_personal = None
    if ministry_code:
        by_code = db.scalar(select(StaffMember).where(StaffMember.ministry_staff_code == ministry_code))
    if personal_id:
        matches = list(db.scalars(select(StaffMember).where(StaffMember.personal_id == personal_id)).all())
        if len(matches) == 1:
            by_personal = matches[0]
        elif len(matches) > 1:
            return None, "Số định danh cá nhân đang trùng nhiều hồ sơ"
    if by_code is not None and by_personal is not None and by_code.id != by_personal.id:
        return None, "Mã Bộ GD&ĐT và số định danh đang thuộc hai người khác nhau"
    member = by_code or by_personal
    if member is not None:
        return member, None
    if full_name and birth_date is not None:
        matches = list(
            db.scalars(
                select(StaffMember).where(
                    func.lower(StaffMember.full_name) == full_name.lower(),
                    StaffMember.date_of_birth == birth_date,
                )
            ).all()
        )
        if len(matches) == 1:
            return matches[0], None
        if len(matches) > 1:
            return None, "Họ tên và ngày sinh trùng nhiều hồ sơ"
    return None, None


def _record_form_context(
    *,
    scope: dict[str, Any],
    record: StaffYearRecord | None = None,
    status: str | None = None,
) -> dict[str, Any]:
    return {
        "nguoi_dung": scope["user"],
        "school_years": scope["school_years"],
        "schools": scope["schools"],
        "selected_year_id": record.school_year_id if record else scope["selected_year_id"],
        "selected_school_id": record.school_id if record else scope["selected_school_id"],
        "record": record,
        "member": record.staff_member if record else None,
        "status": status,
        "status_labels": STATUS_LABELS,
        "position_labels": POSITION_LABELS,
        "employment_labels": EMPLOYMENT_LABELS,
        "teaching_level_labels": TEACHING_LEVEL_LABELS,
        "teaching_age_labels": TEACHING_AGE_LABELS,
        "qualification_level_labels": QUALIFICATION_LEVEL_LABELS,
        "qualification_standard_labels": QUALIFICATION_STANDARD_LABELS,
        "professional_standard_labels": PROFESSIONAL_STANDARD_LABELS,
    }


@router.get("", response_class=HTMLResponse)
def staff_list(
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
    position_group: str = "",
    q: str = "",
    page: str = "1",
    page_size: str = "50",
    audit: str = "",
    status: str | None = None,
    db: Session = Depends(get_db),
) -> HTMLResponse:
    scope = _load_page_scope(
        db,
        request,
        _parse_optional_int(school_year_id),
        _parse_optional_int(commune_id),
        _parse_optional_int(school_id),
    )
    conditions = _scope_conditions(scope)
    normalized_position = str(position_group or "").strip().upper()
    if normalized_position in POSITION_LABELS:
        conditions.append(StaffYearRecord.position_group == normalized_position)
    keyword = _clean(q)
    if keyword:
        pattern = f"%{keyword}%"
        conditions.append(
            or_(
                StaffMember.full_name.ilike(pattern),
                StaffMember.ministry_staff_code.ilike(pattern),
                StaffMember.personal_id.ilike(pattern),
                School.name.ilike(pattern),
            )
        )
    current_page = _parse_page(page, 1, 1, 100000)
    current_page_size = _parse_page(page_size, 50, 20, 200)
    base_query = (
        select(StaffYearRecord)
        .join(StaffYearRecord.staff_member)
        .join(StaffYearRecord.school)
        .options(
            selectinload(StaffYearRecord.staff_member),
            selectinload(StaffYearRecord.school).selectinload(School.commune),
            selectinload(StaffYearRecord.school_year),
        )
        .where(*conditions)
    )
    total_records = int(
        db.scalar(
            select(func.count(StaffYearRecord.id))
            .join(StaffYearRecord.staff_member)
            .join(StaffYearRecord.school)
            .where(*conditions)
        )
        or 0
    )
    page_count = max(ceil(total_records / current_page_size), 1)
    current_page = min(current_page, page_count)
    records = list(
        db.scalars(
            base_query
            .order_by(School.name, StaffMember.full_name)
            .offset((current_page - 1) * current_page_size)
            .limit(current_page_size)
        ).all()
    )
    active_conditions = [*conditions, StaffYearRecord.is_active.is_(True), StaffYearRecord.status_code == "DANG_LAM_VIEC"]
    position_counts = {
        row[0]: int(row[1])
        for row in db.execute(
            select(StaffYearRecord.position_group, func.count(StaffYearRecord.id))
            .join(StaffYearRecord.staff_member)
            .join(StaffYearRecord.school)
            .where(*active_conditions)
            .group_by(StaffYearRecord.position_group)
        ).all()
    }
    incomplete = int(
        db.scalar(
            select(func.count(StaffYearRecord.id))
            .join(StaffYearRecord.staff_member)
            .join(StaffYearRecord.school)
            .where(
                *active_conditions,
                StaffYearRecord.position_group == "GIAO_VIEN",
                or_(
                    StaffYearRecord.qualification_standard == "CHUA_XAC_DINH",
                    StaffYearRecord.professional_standard == "CHUA_DANH_GIA",
                    StaffYearRecord.teaching_age_group == "KHONG_XAC_DINH",
                ),
            )
        )
        or 0
    )
    summary = {
        "total": sum(position_counts.values()),
        "managers": position_counts.get("CBQL", 0),
        "teachers": position_counts.get("GIAO_VIEN", 0),
        "employees": position_counts.get("NHAN_VIEN", 0),
        "incomplete": incomplete,
    }
    query_params: dict[str, Any] = {
        "school_year_id": scope["selected_year_id"],
        "page_size": current_page_size,
    }
    if scope["selected_commune_id"] is not None:
        query_params["commune_id"] = scope["selected_commune_id"]
    if scope["selected_school_id"] is not None:
        query_params["school_id"] = scope["selected_school_id"]
    if normalized_position:
        query_params["position_group"] = normalized_position
    if keyword:
        query_params["q"] = keyword
    audit_data = _build_data_audit(db, scope) if audit == "1" else None
    return templates.TemplateResponse(
        request=request,
        name="staff/list.html",
        context={
            "nguoi_dung": scope["user"],
            "school_years": scope["school_years"],
            "selected_year_id": scope["selected_year_id"],
            "communes": scope["communes"],
            "schools": scope["schools"],
            "selected_commune_id": scope["selected_commune_id"],
            "selected_school_id": scope["selected_school_id"],
            "scope_label": scope["scope_label"],
            "position_group": normalized_position,
            "q": keyword,
            "records": records,
            "total_records": total_records,
            "summary": summary,
            "can_manage": _can_manage(scope["user"]),
            "status": status,
            "position_labels": POSITION_LABELS,
            "status_labels": STATUS_LABELS,
            "employment_labels": EMPLOYMENT_LABELS,
            "teaching_level_labels": TEACHING_LEVEL_LABELS,
            "qualification_level_labels": QUALIFICATION_LEVEL_LABELS,
            "teaching_age_labels": TEACHING_AGE_LABELS,
            "qualification_standard_labels": QUALIFICATION_STANDARD_LABELS,
            "professional_standard_labels": PROFESSIONAL_STANDARD_LABELS,
            "base_query": urlencode(query_params),
            "scope_query": urlencode({key: value for key, value in query_params.items() if key not in {"position_group", "q", "page_size"}}),
            "page": current_page,
            "page_size": current_page_size,
            "page_count": page_count,
            "audit": audit_data,
        },
    )


@router.get("/kiem-tra-du-lieu")
def staff_audit_redirect(
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
) -> RedirectResponse:
    params: dict[str, Any] = {"audit": 1}
    if _parse_optional_int(school_year_id) is not None:
        params["school_year_id"] = _parse_optional_int(school_year_id)
    if _parse_optional_int(commune_id) is not None:
        params["commune_id"] = _parse_optional_int(commune_id)
    if _parse_optional_int(school_id) is not None:
        params["school_id"] = _parse_optional_int(school_id)
    return RedirectResponse(f"/doi-ngu?{urlencode(params)}", status_code=303)


@router.get("/xuat-mau-excel")
def export_staff_import_template(
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
    db: Session = Depends(get_db),
) -> StreamingResponse:
    scope = _load_page_scope(db, request, _parse_optional_int(school_year_id), _parse_optional_int(commune_id), _parse_optional_int(school_id))
    workbook = _create_import_template_workbook(db, scope)
    year_code = next((year.code for year in scope["school_years"] if year.id == scope["selected_year_id"]), "nam_hoc")
    return _excel_response(workbook, f"mau_nhap_doi_ngu_{year_code}.xlsx")


@router.get("/xuat-danh-sach-excel")
def export_current_staff_data(
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
    db: Session = Depends(get_db),
) -> StreamingResponse:
    scope = _load_page_scope(db, request, _parse_optional_int(school_year_id), _parse_optional_int(commune_id), _parse_optional_int(school_id))
    workbook = _create_current_data_workbook(db, scope)
    year_code = next((year.code for year in scope["school_years"] if year.id == scope["selected_year_id"]), "nam_hoc")
    return _excel_response(workbook, f"kiem_tra_doi_ngu_hien_co_{year_code}.xlsx")


@router.get("/them", response_class=HTMLResponse)
def staff_create_page(
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
    db: Session = Depends(get_db),
) -> HTMLResponse:
    scope = _load_page_scope(
        db,
        request,
        _parse_optional_int(school_year_id),
        _parse_optional_int(commune_id),
        _parse_optional_int(school_id),
    )
    if not _can_manage(scope["user"]):
        return RedirectResponse("/doi-ngu?status=forbidden", status_code=303)
    return templates.TemplateResponse(
        request=request,
        name="staff/form.html",
        context=_record_form_context(scope=scope),
    )


@router.post("/them")
def staff_create(
    request: Request,
    school_year_id: int = Form(...),
    school_id: int = Form(...),
    ministry_staff_code: str = Form(""),
    full_name: str = Form(...),
    date_of_birth: str = Form(""),
    gender: str = Form(""),
    ethnic_group: str = Form(""),
    personal_id: str = Form(""),
    phone: str = Form(""),
    email: str = Form(""),
    status_code: str = Form("DANG_LAM_VIEC"),
    position_group: str = Form("GIAO_VIEN"),
    position_title: str = Form(""),
    employment_type: str = Form("CHUA_XAC_DINH"),
    recruitment_date: str = Form(""),
    teaching_level: str = Form("KHONG_DAY"),
    teaching_subject: str = Form(""),
    main_specialty: str = Form(""),
    staff_evaluation: str = Form(""),
    teaching_age_group: str = Form("KHONG_XAC_DINH"),
    receives_policy: str | None = Form(None),
    qualification_level: str = Form("CHUA_XAC_DINH"),
    qualification_standard: str = Form("CHUA_XAC_DINH"),
    professional_standard: str = Form("CHUA_DANH_GIA"),
    notes: str = Form(""),
    db: Session = Depends(get_db),
) -> RedirectResponse:
    user = lay_thong_tin_nguoi_dung(request)
    if not _can_manage(user) or not _ensure_school_in_scope(user, school_id):
        return RedirectResponse("/doi-ngu?status=forbidden", status_code=303)
    normalized_code = _clean(ministry_staff_code) or None
    member = None
    if normalized_code:
        member = db.scalar(
            select(StaffMember).where(StaffMember.ministry_staff_code == normalized_code)
        )
    if member is None:
        member = StaffMember(
            code=_next_staff_code(db),
            ministry_staff_code=normalized_code,
            full_name=_clean(full_name),
            date_of_birth=_parse_date(date_of_birth),
            gender=_clean(gender) or None,
            ethnic_group=_clean(ethnic_group) or None,
            personal_id=_clean(personal_id) or None,
            phone=_clean(phone) or None,
            email=_clean(email) or None,
            notes=_clean(notes) or None,
        )
        db.add(member)
        db.flush()
    else:
        member.full_name = _clean(full_name)
        member.date_of_birth = _parse_date(date_of_birth)
        member.gender = _clean(gender) or None
        member.ethnic_group = _clean(ethnic_group) or None
        member.personal_id = _clean(personal_id) or None
        member.phone = _clean(phone) or None
        member.email = _clean(email) or None
    existing = db.scalar(
        select(StaffYearRecord).where(
            StaffYearRecord.staff_member_id == member.id,
            StaffYearRecord.school_year_id == school_year_id,
        )
    )
    if existing is not None:
        db.rollback()
        return RedirectResponse(
            f"/doi-ngu?school_year_id={school_year_id}&school_id={school_id}&status=duplicate",
            status_code=303,
        )
    record = StaffYearRecord(
        staff_member_id=member.id,
        school_year_id=school_year_id,
        school_id=school_id,
        status_code=status_code if status_code in STATUS_LABELS else "DANG_LAM_VIEC",
        position_group=position_group if position_group in POSITION_LABELS else "GIAO_VIEN",
        position_title=_clean(position_title) or None,
        employment_type=(
            employment_type if employment_type in EMPLOYMENT_LABELS else "CHUA_XAC_DINH"
        ),
        recruitment_date=_parse_date(recruitment_date),
        teaching_level=(
            teaching_level if teaching_level in TEACHING_LEVEL_LABELS else "KHONG_DAY"
        ),
        teaching_subject=_clean(teaching_subject) or None,
        main_specialty=_clean(main_specialty) or None,
        staff_evaluation=_clean(staff_evaluation) or None,
        teaching_age_group=(
            teaching_age_group
            if teaching_age_group in TEACHING_AGE_LABELS
            else "KHONG_XAC_DINH"
        ),
        receives_policy=receives_policy == "1",
        qualification_level=(
            qualification_level
            if qualification_level in QUALIFICATION_LEVEL_LABELS
            else "CHUA_XAC_DINH"
        ),
        qualification_standard=(
            qualification_standard
            if qualification_standard in QUALIFICATION_STANDARD_LABELS
            else "CHUA_XAC_DINH"
        ),
        professional_standard=(
            professional_standard
            if professional_standard in PROFESSIONAL_STANDARD_LABELS
            else "CHUA_DANH_GIA"
        ),
        notes=_clean(notes) or None,
        created_by_user_id=user.get("id"),
    )
    db.add(record)
    db.commit()
    return RedirectResponse(
        f"/doi-ngu?school_year_id={school_year_id}&school_id={school_id}&status=created",
        status_code=303,
    )


@router.get("/{record_id}/sua", response_class=HTMLResponse)
def staff_edit_page(
    record_id: int,
    request: Request,
    db: Session = Depends(get_db),
) -> HTMLResponse:
    record = db.scalar(
        select(StaffYearRecord)
        .options(
            selectinload(StaffYearRecord.staff_member),
            selectinload(StaffYearRecord.school),
            selectinload(StaffYearRecord.school_year),
        )
        .where(StaffYearRecord.id == record_id)
    )
    if record is None:
        return RedirectResponse("/doi-ngu?status=not_found", status_code=303)
    user = lay_thong_tin_nguoi_dung(request)
    if not _can_manage(user) or not _ensure_school_in_scope(user, record.school_id):
        return RedirectResponse("/doi-ngu?status=forbidden", status_code=303)
    scope = _load_page_scope(db, request, record.school_year_id, record.school.commune_id, record.school_id)
    return templates.TemplateResponse(
        request=request,
        name="staff/form.html",
        context=_record_form_context(scope=scope, record=record),
    )


@router.post("/{record_id}/sua")
def staff_edit(
    record_id: int,
    request: Request,
    school_year_id: int = Form(...),
    school_id: int = Form(...),
    ministry_staff_code: str = Form(""),
    full_name: str = Form(...),
    date_of_birth: str = Form(""),
    gender: str = Form(""),
    ethnic_group: str = Form(""),
    personal_id: str = Form(""),
    phone: str = Form(""),
    email: str = Form(""),
    status_code: str = Form("DANG_LAM_VIEC"),
    position_group: str = Form("GIAO_VIEN"),
    position_title: str = Form(""),
    employment_type: str = Form("CHUA_XAC_DINH"),
    recruitment_date: str = Form(""),
    teaching_level: str = Form("KHONG_DAY"),
    teaching_subject: str = Form(""),
    main_specialty: str = Form(""),
    staff_evaluation: str = Form(""),
    teaching_age_group: str = Form("KHONG_XAC_DINH"),
    receives_policy: str | None = Form(None),
    qualification_level: str = Form("CHUA_XAC_DINH"),
    qualification_standard: str = Form("CHUA_XAC_DINH"),
    professional_standard: str = Form("CHUA_DANH_GIA"),
    notes: str = Form(""),
    db: Session = Depends(get_db),
) -> RedirectResponse:
    record = db.scalar(
        select(StaffYearRecord)
        .options(selectinload(StaffYearRecord.staff_member))
        .where(StaffYearRecord.id == record_id)
    )
    if record is None:
        return RedirectResponse("/doi-ngu?status=not_found", status_code=303)
    user = lay_thong_tin_nguoi_dung(request)
    if not _can_manage(user) or not _ensure_school_in_scope(user, record.school_id) or not _ensure_school_in_scope(user, school_id):
        return RedirectResponse("/doi-ngu?status=forbidden", status_code=303)
    member = record.staff_member
    normalized_code = _clean(ministry_staff_code) or None
    duplicate_member = None
    if normalized_code:
        duplicate_member = db.scalar(
            select(StaffMember).where(
                StaffMember.ministry_staff_code == normalized_code,
                StaffMember.id != member.id,
            )
        )
    if duplicate_member is not None:
        return RedirectResponse(f"/doi-ngu/{record_id}/sua?status=duplicate_code", status_code=303)
    duplicate_record = db.scalar(
        select(StaffYearRecord).where(
            StaffYearRecord.staff_member_id == member.id,
            StaffYearRecord.school_year_id == school_year_id,
            StaffYearRecord.id != record.id,
        )
    )
    if duplicate_record is not None:
        return RedirectResponse(f"/doi-ngu/{record_id}/sua?status=duplicate", status_code=303)
    member.ministry_staff_code = normalized_code
    member.full_name = _clean(full_name)
    member.date_of_birth = _parse_date(date_of_birth)
    member.gender = _clean(gender) or None
    member.ethnic_group = _clean(ethnic_group) or None
    member.personal_id = _clean(personal_id) or None
    member.phone = _clean(phone) or None
    member.email = _clean(email) or None
    member.notes = _clean(notes) or None
    record.school_year_id = school_year_id
    record.school_id = school_id
    record.status_code = status_code if status_code in STATUS_LABELS else "DANG_LAM_VIEC"
    record.position_group = position_group if position_group in POSITION_LABELS else "GIAO_VIEN"
    record.position_title = _clean(position_title) or None
    record.employment_type = employment_type if employment_type in EMPLOYMENT_LABELS else "CHUA_XAC_DINH"
    record.recruitment_date = _parse_date(recruitment_date)
    record.teaching_level = teaching_level if teaching_level in TEACHING_LEVEL_LABELS else "KHONG_DAY"
    record.teaching_subject = _clean(teaching_subject) or None
    record.main_specialty = _clean(main_specialty) or None
    record.staff_evaluation = _clean(staff_evaluation) or None
    record.teaching_age_group = teaching_age_group if teaching_age_group in TEACHING_AGE_LABELS else "KHONG_XAC_DINH"
    record.receives_policy = receives_policy == "1"
    record.qualification_level = qualification_level if qualification_level in QUALIFICATION_LEVEL_LABELS else "CHUA_XAC_DINH"
    record.qualification_standard = qualification_standard if qualification_standard in QUALIFICATION_STANDARD_LABELS else "CHUA_XAC_DINH"
    record.professional_standard = professional_standard if professional_standard in PROFESSIONAL_STANDARD_LABELS else "CHUA_DANH_GIA"
    record.notes = _clean(notes) or None
    db.commit()
    return RedirectResponse(
        f"/doi-ngu?school_year_id={school_year_id}&school_id={school_id}&status=updated",
        status_code=303,
    )


@router.post("/{record_id}/ngung-theo-doi")
def staff_archive(record_id: int, request: Request, db: Session = Depends(get_db)) -> RedirectResponse:
    record = db.get(StaffYearRecord, record_id)
    if record is None:
        return RedirectResponse("/doi-ngu?status=not_found", status_code=303)
    user = lay_thong_tin_nguoi_dung(request)
    if not _can_manage(user) or not _ensure_school_in_scope(user, record.school_id):
        return RedirectResponse("/doi-ngu?status=forbidden", status_code=303)
    record.is_active = False
    db.commit()
    return RedirectResponse(
        f"/doi-ngu?school_year_id={record.school_year_id}&school_id={record.school_id}&status=archived",
        status_code=303,
    )


@router.post("/dong-bo-tai-khoan")
def sync_accounts(
    request: Request,
    school_year_id: int = Form(...),
    school_id: str = Form(""),
    db: Session = Depends(get_db),
) -> RedirectResponse:
    user = lay_thong_tin_nguoi_dung(request)
    if not _can_manage(user):
        return RedirectResponse("/doi-ngu?status=forbidden", status_code=303)
    selected_school_id = _parse_optional_int(school_id)
    statement = (
        select(User)
        .options(selectinload(User.role), selectinload(User.school))
        .where(User.is_active.is_(True), User.school_id.is_not(None))
    )
    if selected_school_id is not None:
        if not _ensure_school_in_scope(user, selected_school_id):
            return RedirectResponse("/doi-ngu?status=forbidden", status_code=303)
        statement = statement.where(User.school_id == selected_school_id)
    elif normalize_role_code(user.get("role_code")) == SCHOOL_ROLE_CODE:
        statement = statement.where(User.school_id == int(user["school_id"]))
    created = 0
    skipped = 0
    for account in db.scalars(statement).all():
        role_code = normalize_role_code(account.role.code)
        if role_code not in {SCHOOL_ROLE_CODE, "GIAO_VIEN"}:
            continue
        ministry_code = account.username.split(".", 1)[-1] if "." in account.username else None
        member = None
        if ministry_code:
            member = db.scalar(
                select(StaffMember).where(StaffMember.ministry_staff_code == ministry_code)
            )
        if member is None:
            member = StaffMember(
                code=_next_staff_code(db),
                ministry_staff_code=ministry_code,
                full_name=account.full_name,
            )
            db.add(member)
            db.flush()
        existing = db.scalar(
            select(StaffYearRecord).where(
                StaffYearRecord.staff_member_id == member.id,
                StaffYearRecord.school_year_id == school_year_id,
            )
        )
        if existing is not None:
            skipped += 1
            continue
        is_manager = role_code == SCHOOL_ROLE_CODE
        db.add(
            StaffYearRecord(
                staff_member_id=member.id,
                school_year_id=school_year_id,
                school_id=int(account.school_id),
                position_group="CBQL" if is_manager else "GIAO_VIEN",
                position_title="Cán bộ quản lý" if is_manager else "Giáo viên",
                teaching_level="KHONG_DAY" if is_manager else "MAU_GIAO",
                teaching_age_group="KHONG_XAC_DINH",
                created_by_user_id=user.get("id"),
            )
        )
        created += 1
    db.commit()
    sid = selected_school_id or user.get("school_id") or ""
    return RedirectResponse(
        f"/doi-ngu?school_year_id={school_year_id}&school_id={sid}&status=synced_{created}_{skipped}",
        status_code=303,
    )


def _find_header_map(ws: Any) -> tuple[int, dict[str, int]]:
    for row_number in range(1, min(ws.max_row, 50) + 1):
        values = [_normalize(ws.cell(row_number, col).value) for col in range(1, ws.max_column + 1)]
        has_name = any(value in {"ho ten", "ho va ten"} for value in values)
        has_school = any(value in {"ten truong", "truong"} or "ten truong" in value for value in values)
        if has_name and has_school:
            return row_number, {value: index + 1 for index, value in enumerate(values) if value}
    raise ValueError("Không tìm thấy dòng tiêu đề có cột Họ tên và Tên trường.")


def _column(header_map: dict[str, int], *names: str) -> int | None:
    for name in names:
        normalized = _normalize(name)
        if normalized in header_map:
            return header_map[normalized]
    return None


def _map_position(value: Any) -> str:
    text = _normalize(value)
    if "can bo quan ly" in text or "hieu truong" in text or "pho hieu truong" in text:
        return "CBQL"
    if "nhan vien" in text:
        return "NHAN_VIEN"
    return "GIAO_VIEN"


def _map_status(value: Any) -> str:
    text = _normalize(value)
    if "nghi huu" in text:
        return "NGHI_HUU"
    if "nghi viec" in text:
        return "NGHI_VIEC"
    if "tam nghi" in text:
        return "TAM_NGHI"
    if "chuyen" in text:
        return "CHUYEN_DI"
    return "DANG_LAM_VIEC"


def _map_teaching(value: Any, concurrent: Any) -> tuple[str, str]:
    text = f"{_normalize(value)} {_normalize(concurrent)}"
    if "nha tre" in text:
        return "NHA_TRE", "NHA_TRE"
    if "mau giao" in text:
        return "MAU_GIAO", "KHONG_XAC_DINH"
    return "KHONG_DAY", "KHONG_XAC_DINH"


@router.post("/nhap-excel")
def import_staff_excel(
    request: Request,
    school_year_id: int = Form(...),
    school_id: str = Form(""),
    excel_file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> RedirectResponse:
    user = lay_thong_tin_nguoi_dung(request)
    if not _can_manage(user):
        return RedirectResponse("/doi-ngu?status=forbidden", status_code=303)
    forced_school_id = _parse_optional_int(school_id)
    if forced_school_id is not None and not _ensure_school_in_scope(user, forced_school_id):
        return RedirectResponse("/doi-ngu?status=forbidden", status_code=303)
    created = updated = skipped = errors = 0
    try:
        content = excel_file.file.read()
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        ws = workbook["Nhap danh sach"] if "Nhap danh sach" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        header_row, header_map = _find_header_map(ws)
        columns = {
            "year": _column(header_map, "Năm học"),
            "commune_code": _column(header_map, "Mã xã/phường"),
            "commune": _column(header_map, "Tên xã/phường", "UBND Xã", "Xã/phường"),
            "school_code": _column(header_map, "Mã trường"),
            "school": _column(header_map, "Tên trường"),
            "username": _column(header_map, "Tên đăng nhập hiện có"),
            "code": _column(header_map, "Mã định danh Bộ GD&ĐT", "Mã cán bộ"),
            "name": _column(header_map, "Họ và tên", "Họ tên"),
            "birth": _column(header_map, "Ngày sinh"),
            "gender": _column(header_map, "Giới tính"),
            "ethnic": _column(header_map, "Dân tộc"),
            "personal": _column(header_map, "Số định danh cá nhân"),
            "phone": _column(header_map, "Số điện thoại", "Điện thoại"),
            "email": _column(header_map, "Email"),
            "position": _column(header_map, "Vị trí việc làm"),
            "title": _column(header_map, "Chức vụ"),
            "employment": _column(header_map, "Hình thức tuyển dụng", "Loại hợp đồng"),
            "recruit": _column(header_map, "Ngày tuyển dụng"),
            "qualification": _column(header_map, "Trình độ đào tạo"),
            "qualification_standard": _column(header_map, "Mức chuẩn đào tạo", "Đạt chuẩn"),
            "professional": _column(header_map, "Kết quả chuẩn nghề nghiệp", "Chuẩn nghề nghiệp"),
            "teaching_level": _column(header_map, "Cấp học trực tiếp"),
            "subject": _column(header_map, "Môn dạy"),
            "specialty": _column(header_map, "Chuyên ngành chính", "Chuyên ngành"),
            "evaluation": _column(header_map, "Đánh giá viên chức"),
            "teaching_age": _column(header_map, "Dạy nhóm/lớp", "Dạy nhóm lớp"),
            "policy": _column(header_map, "Hưởng chế độ, chính sách", "Hưởng chính sách"),
            "status": _column(header_map, "Trạng thái công tác", "Trạng thái"),
            "notes": _column(header_map, "Ghi chú"),
        }
        if columns["name"] is None:
            raise ValueError("Tệp Excel thiếu cột Họ và tên.")
        valid_year = db.get(SchoolYear, school_year_id)
        if valid_year is None:
            raise ValueError("Năm học không tồn tại.")

        # V4.2: import theo đúng cặp xã/phường + trường. Tên trường trùng ở
        # hai địa bàn không còn bị dict ghi đè và gán nhầm sang xã khác.
        # Trường inactive chỉ được dùng khi chính năm đang nhập đã có hồ sơ
        # lịch sử tại trường đó (phục vụ cập nhật dữ liệu năm cũ).
        historical_school_ids = set(
            db.scalars(
                select(StaffYearRecord.school_id)
                .where(StaffYearRecord.school_year_id == school_year_id)
                .distinct()
            ).all()
        )
        schools = list(
            db.scalars(
                select(School)
                .options(selectinload(School.commune))
                .where(
                    or_(
                        School.is_active.is_(True),
                        School.id.in_(historical_school_ids) if historical_school_ids else School.id == -1,
                    )
                )
            ).all()
        )
        school_by_code = {
            _normalize(item.code): item
            for item in schools
            if _normalize(item.code)
        }
        school_by_name: dict[str, list[School]] = {}
        for item in schools:
            school_by_name.setdefault(_normalize(item.name), []).append(item)
        for row_number in range(header_row + 1, ws.max_row + 1):
            full_name = _clean(ws.cell(row_number, columns["name"]).value)
            if not full_name:
                continue
            row_year = _clean(ws.cell(row_number, columns["year"]).value) if columns["year"] else ""
            if row_year and row_year != valid_year.code:
                skipped += 1
                continue
            target_school_id = forced_school_id
            if target_school_id is None:
                school_obj = None
                if columns["school_code"]:
                    school_obj = school_by_code.get(
                        _normalize(ws.cell(row_number, columns["school_code"]).value)
                    )
                if school_obj is None and columns["school"]:
                    school_value = _normalize(
                        ws.cell(row_number, columns["school"]).value
                    )
                    commune_code_value = (
                        _normalize(ws.cell(row_number, columns["commune_code"]).value)
                        if columns["commune_code"]
                        else ""
                    )
                    commune_name_value = (
                        _normalize(ws.cell(row_number, columns["commune"]).value)
                        if columns["commune"]
                        else ""
                    )

                    candidates = list(school_by_name.get(school_value, []))
                    if not candidates and school_value:
                        candidates = [
                            item
                            for item in schools
                            if (
                                school_value == _normalize(item.name)
                                or school_value in _normalize(item.name)
                                or _normalize(item.name) in school_value
                            )
                        ]

                    if commune_code_value:
                        candidates = [
                            item for item in candidates
                            if item.commune is not None
                            and _normalize(item.commune.code) == commune_code_value
                        ]
                    elif commune_name_value:
                        candidates = [
                            item for item in candidates
                            if item.commune is not None
                            and (
                                commune_name_value == _normalize(item.commune.name)
                                or commune_name_value in _normalize(item.commune.name)
                                or _normalize(item.commune.name) in commune_name_value
                            )
                        ]

                    if len(candidates) == 1:
                        school_obj = candidates[0]
                    elif len(candidates) > 1:
                        errors += 1
                        continue
                target_school_id = school_obj.id if school_obj else None
            if target_school_id is None or not _ensure_school_in_scope(user, target_school_id):
                skipped += 1
                continue
            username = _clean(ws.cell(row_number, columns["username"]).value) if columns["username"] else ""
            ministry_code = _clean(ws.cell(row_number, columns["code"]).value) if columns["code"] else ""
            if not ministry_code and username:
                ministry_code = _account_staff_code(username)
            personal_id = _clean(ws.cell(row_number, columns["personal"]).value) if columns["personal"] else ""
            birth_date = _parse_date(ws.cell(row_number, columns["birth"]).value) if columns["birth"] else None
            member, match_error = _match_staff_member(db, ministry_code, personal_id, full_name, birth_date)
            if match_error:
                errors += 1
                continue
            if member is None:
                member = StaffMember(
                    code=_next_staff_code(db),
                    ministry_staff_code=ministry_code or None,
                    full_name=full_name,
                    date_of_birth=birth_date,
                    gender=_clean(ws.cell(row_number, columns["gender"]).value) if columns["gender"] else None,
                    ethnic_group=_clean(ws.cell(row_number, columns["ethnic"]).value) if columns["ethnic"] else None,
                    personal_id=personal_id or None,
                    phone=_clean(ws.cell(row_number, columns["phone"]).value) if columns["phone"] else None,
                    email=_clean(ws.cell(row_number, columns["email"]).value) if columns["email"] else None,
                    notes=_clean(ws.cell(row_number, columns["notes"]).value) if columns["notes"] else None,
                )
                db.add(member)
                db.flush()
            else:
                member.full_name = full_name
                if birth_date is not None:
                    member.date_of_birth = birth_date
                if ministry_code and not _clean(member.ministry_staff_code):
                    member.ministry_staff_code = ministry_code
                if personal_id and not _clean(member.personal_id):
                    member.personal_id = personal_id
                for field_name, column_name in (("gender", "gender"), ("ethnic_group", "ethnic"), ("phone", "phone"), ("email", "email"), ("notes", "notes")):
                    column_index = columns[column_name]
                    if column_index:
                        value = _clean(ws.cell(row_number, column_index).value)
                        if value:
                            setattr(member, field_name, value)
            existing = db.scalar(select(StaffYearRecord).where(StaffYearRecord.staff_member_id == member.id, StaffYearRecord.school_year_id == school_year_id))
            position_value = ws.cell(row_number, columns["position"]).value if columns["position"] else ""
            position_group = _map_position(position_value)
            if columns["position"]:
                position_group = _map_label_code(position_value, POSITION_LABELS, position_group)
            defaults = {
                "school_id": target_school_id,
                "status_code": _map_label_code(ws.cell(row_number, columns["status"]).value if columns["status"] else "", STATUS_LABELS, "DANG_LAM_VIEC"),
                "position_group": position_group,
                "position_title": _clean(ws.cell(row_number, columns["title"]).value) if columns["title"] else (_clean(position_value) or POSITION_LABELS[position_group]),
                "employment_type": _map_label_code(ws.cell(row_number, columns["employment"]).value if columns["employment"] else "", EMPLOYMENT_LABELS, "CHUA_XAC_DINH"),
                "recruitment_date": _parse_date(ws.cell(row_number, columns["recruit"]).value) if columns["recruit"] else None,
                "qualification_level": _map_label_code(ws.cell(row_number, columns["qualification"]).value if columns["qualification"] else "", QUALIFICATION_LEVEL_LABELS, "CHUA_XAC_DINH"),
                "qualification_standard": _map_label_code(ws.cell(row_number, columns["qualification_standard"]).value if columns["qualification_standard"] else "", QUALIFICATION_STANDARD_LABELS, "CHUA_XAC_DINH"),
                "professional_standard": _map_label_code(ws.cell(row_number, columns["professional"]).value if columns["professional"] else "", PROFESSIONAL_STANDARD_LABELS, "CHUA_DANH_GIA"),
                "teaching_level": _map_label_code(ws.cell(row_number, columns["teaching_level"]).value if columns["teaching_level"] else "", TEACHING_LEVEL_LABELS, "KHONG_DAY"),
                "teaching_subject": _clean(ws.cell(row_number, columns["subject"]).value) if columns["subject"] else None,
                "main_specialty": _clean(ws.cell(row_number, columns["specialty"]).value) if columns["specialty"] else None,
                "staff_evaluation": _clean(ws.cell(row_number, columns["evaluation"]).value) if columns["evaluation"] else None,
                "source_status_label": _clean(ws.cell(row_number, columns["status"]).value) if columns["status"] else None,
                "teaching_age_group": _map_label_code(ws.cell(row_number, columns["teaching_age"]).value if columns["teaching_age"] else "", TEACHING_AGE_LABELS, "KHONG_XAC_DINH"),
                "receives_policy": _map_boolean(ws.cell(row_number, columns["policy"]).value) if columns["policy"] else False,
                "notes": _clean(ws.cell(row_number, columns["notes"]).value) if columns["notes"] else None,
            }
            if existing is None:
                db.add(StaffYearRecord(staff_member_id=member.id, school_year_id=school_year_id, created_by_user_id=user.get("id"), **defaults))
                created += 1
            else:
                existing.school_id = target_school_id
                for key, value in defaults.items():
                    column_key = {
                        "status_code": "status", "position_group": "position", "position_title": "title",
                        "employment_type": "employment", "recruitment_date": "recruit",
                        "qualification_level": "qualification", "qualification_standard": "qualification_standard",
                        "professional_standard": "professional", "teaching_level": "teaching_level",
                        "teaching_subject": "subject", "main_specialty": "specialty",
                        "staff_evaluation": "evaluation", "source_status_label": "status",
                        "teaching_age_group": "teaching_age", "receives_policy": "policy", "notes": "notes",
                    }.get(key)
                    if key == "school_id" or column_key is None or columns.get(column_key):
                        if value not in (None, "") or key in {"receives_policy", "school_id"}:
                            setattr(existing, key, value)
                existing.is_active = True
                updated += 1
        db.commit()
        return RedirectResponse(
            f"/doi-ngu?school_year_id={school_year_id}&school_id={forced_school_id or ''}&status=imported_{created}_{updated}_{skipped}_{errors}&audit=1",
            status_code=303,
        )
    except Exception:
        db.rollback()
        return RedirectResponse(
            f"/doi-ngu?school_year_id={school_year_id}&school_id={forced_school_id or ''}&status=import_error",
            status_code=303,
        )


@router.get("/cau-hinh-lop", response_class=HTMLResponse)
def class_config_page(
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
    status: str | None = None,
    db: Session = Depends(get_db),
) -> HTMLResponse:
    scope = _load_page_scope(
        db,
        request,
        _parse_optional_int(school_year_id),
        _parse_optional_int(commune_id),
        _parse_optional_int(school_id),
    )
    schools = scope["schools"]
    if scope["selected_school_id"] is not None:
        schools = [item for item in schools if item.id == scope["selected_school_id"]]
    elif scope["selected_commune_id"] is not None:
        schools = [item for item in schools if item.commune_id == scope["selected_commune_id"]]
    summaries = {
        item.school_id: item
        for item in db.scalars(
            select(SchoolStaffYearSummary).where(
                SchoolStaffYearSummary.school_year_id == scope["selected_year_id"],
                SchoolStaffYearSummary.school_id.in_([school.id for school in schools] or [-1]),
            )
        ).all()
    }
    rows = []
    for school in schools:
        summary = summaries.get(school.id)
        rows.append(
            {
                "school": school,
                "summary": summary,
                "institution_group": summary.institution_group if summary else _infer_institution_group(school.name),
                "total_groups_classes": summary.total_groups_classes if summary else 0,
                "preschool_classes_3_4": summary.preschool_classes_3_4 if summary else 0,
                "preschool_classes_5": summary.preschool_classes_5 if summary else 0,
            }
        )
    return templates.TemplateResponse(
        request=request,
        name="staff/class_config.html",
        context={
            "nguoi_dung": scope["user"],
            "school_years": scope["school_years"],
            "selected_year_id": scope["selected_year_id"],
            "communes": scope["communes"],
            "selected_commune_id": scope["selected_commune_id"],
            "selected_school_id": scope["selected_school_id"],
            "rows": rows,
            "status": status,
            "can_manage": _can_manage(scope["user"]),
            "institution_group_labels": INSTITUTION_GROUP_LABELS,
        },
    )


@router.post("/cau-hinh-lop/{school_id}")
def class_config_save(
    school_id: int,
    request: Request,
    school_year_id: int = Form(...),
    institution_group: str = Form("TRUONG_MAM_NON"),
    total_groups_classes: str = Form("0"),
    preschool_classes_3_4: str = Form("0"),
    preschool_classes_5: str = Form("0"),
    notes: str = Form(""),
    db: Session = Depends(get_db),
) -> RedirectResponse:
    user = lay_thong_tin_nguoi_dung(request)
    if not _can_manage(user) or not _ensure_school_in_scope(user, school_id):
        return RedirectResponse("/doi-ngu?status=forbidden", status_code=303)
    summary = db.scalar(
        select(SchoolStaffYearSummary).where(
            SchoolStaffYearSummary.school_id == school_id,
            SchoolStaffYearSummary.school_year_id == school_year_id,
        )
    )
    if summary is None:
        summary = SchoolStaffYearSummary(
            school_id=school_id,
            school_year_id=school_year_id,
        )
        db.add(summary)
    summary.institution_group = (
        institution_group
        if institution_group in INSTITUTION_GROUP_LABELS
        else "TRUONG_MAM_NON"
    )
    summary.total_groups_classes = _parse_nonnegative_int(total_groups_classes)
    summary.preschool_classes_3_4 = _parse_nonnegative_int(preschool_classes_3_4)
    summary.preschool_classes_5 = _parse_nonnegative_int(preschool_classes_5)
    summary.notes = _clean(notes) or None
    summary.updated_by_user_id = user.get("id")
    db.commit()
    school = db.get(School, school_id)
    commune_id = school.commune_id if school else ""
    return RedirectResponse(
        f"/doi-ngu/cau-hinh-lop?school_year_id={school_year_id}&commune_id={commune_id}&status=saved",
        status_code=303,
    )
