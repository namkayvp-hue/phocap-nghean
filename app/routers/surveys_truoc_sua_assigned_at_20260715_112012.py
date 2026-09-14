from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from math import ceil
from pathlib import Path
from typing import Annotated, Any
from urllib.parse import urlencode
from uuid import uuid4
import sqlite3
import unicodedata

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import (
    HTMLResponse,
    RedirectResponse,
    StreamingResponse,
)
from fastapi.templating import Jinja2Templates
from sqlalchemy import func, or_, select, text
from sqlalchemy.exc import IntegrityError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import (
    Session,
    selectinload,
    with_loader_criteria,
)

from app.database import DATABASE_PATH, get_db
from app.models import (
    Classroom,
    Commune,
    School,
    SchoolYear,
    Student,
)
from app.survey_models import (
    Household,
    SurveyBatch,
    SurveyForm,
    SurveyPerson,
    SurveyPersonYearRecord,
)


APP_DIR = Path(__file__).resolve().parent.parent

templates = Jinja2Templates(
    directory=str(APP_DIR / "templates")
)

router = APIRouter(
    prefix="/dieu-tra",
    tags=["Điều tra phổ cập"],
)


BATCH_STATUS_LABELS = {
    "CHUAN_BI": "Chuẩn bị",
    "DANG_DIEU_TRA": "Đang điều tra",
    "DA_KET_THUC": "Đã kết thúc",
}

FORM_STATUS_LABELS = {
    "CHUA_DIEU_TRA": "Chưa điều tra",
    "DANG_DIEU_TRA": "Đang điều tra",
    "DA_HOAN_THANH": "Đã hoàn thành",
    "CAN_BO_SUNG": "Cần bổ sung",
}

STATUS_MESSAGES = {
    "created": "Đã tạo đợt điều tra mới thành công.",
    "household_created": (
        "Đã thêm hộ gia đình và tạo phiếu điều tra thành công."
    ),
    "household_attached": (
        "Hộ gia đình đã có trong hệ thống; "
        "đã tạo phiếu cho đợt điều tra này."
    ),
    "person_created": (
        "Đã thêm đối tượng điều tra vào hộ gia đình thành công."
    ),
    "person_updated": (
        "Đã cập nhật thông tin đối tượng thành công."
    ),
    "person_deactivated": (
        "Đã ngừng theo dõi đối tượng. Lịch sử năm học vẫn được giữ nguyên."
    ),
    "year_record_saved": (
        "Đã lưu thông tin theo dõi năm học thành công."
    ),
    "form_updated": (
        "Đã cập nhật trạng thái và thông tin phiếu điều tra thành công."
    ),
    "investigators_saved": (
        "Đã lưu phân công người điều tra cho phiếu thành công."
    ),
    "investigators_cleared": (
        "Đã xóa toàn bộ phân công người điều tra của phiếu."
    ),
}

HOUSEHOLD_PAGE_SIZE = 20
PERSON_PAGE_SIZE = 20

RESIDENCY_STATUS_LABELS = {
    "THUONG_TRU": "Thường trú",
    "TAM_TRU": "Tạm trú",
    "TAM_VANG": "Tạm vắng",
    "CHUYEN_DEN": "Chuyển đến",
    "CHUYEN_DI": "Chuyển đi",
    "CHUA_XAC_DINH": "Chưa xác định",
}


LEARNING_STATUS_LABELS = {
    "DANG_HOC": "Đang học",
    "CHUA_DI_HOC": "Chưa đi học",
    "CHUYEN_DEN": "Chuyển đến",
    "CHUYEN_DI": "Chuyển đi",
    "TAM_NGHI": "Tạm nghỉ",
    "THOI_HOC": "Thôi học",
    "BO_HOC": "Bỏ học",
    "DA_TOT_NGHIEP": "Đã tốt nghiệp",
    "KHONG_THUOC_DIEN": "Không thuộc diện theo dõi",
}

BOOLEAN_YEAR_FIELDS = (
    "completed_preschool_5",
    "attends_required_days",
    "attends_regularly",
    "prepared_vietnamese",
    "weight_monitored",
    "underweight",
    "height_monitored",
    "stunted",
)


def chuan_hoa_van_ban(value: str) -> str:
    """Xóa khoảng trắng thừa trong văn bản."""

    return " ".join(value.strip().split())


def chuan_hoa_ma(value: str) -> str:
    """Chuẩn hóa mã định danh nhưng vẫn giữ dạng văn bản."""

    return "".join(value.strip().split()).upper()


def chuan_hoa_so_dien_thoai(value: str) -> str | None:
    """Giữ số điện thoại dưới dạng văn bản."""

    normalized = "".join(value.strip().split())
    return normalized or None


def chuyen_ngay(
    value: str,
    ten_truong: str,
    bat_buoc: bool = False,
) -> tuple[date | None, str | None]:
    value = value.strip()

    if not value:
        if bat_buoc:
            return None, f"{ten_truong} không được để trống."
        return None, None

    try:
        result = date.fromisoformat(value)
    except ValueError:
        return None, f"{ten_truong} không đúng định dạng."

    return result, None


def tao_ma_dot_dieu_tra(
    db: Session,
    commune: Commune,
    school_year: SchoolYear,
) -> str:
    year_code = school_year.code.replace("-", "")

    existing_count = db.scalar(
        select(func.count(SurveyBatch.id)).where(
            SurveyBatch.commune_id == commune.id,
            SurveyBatch.school_year_id == school_year.id,
        )
    )

    sequence = int(existing_count or 0) + 1

    while True:
        candidate = (
            f"DT-{commune.code}-{year_code}-{sequence:03d}"
        )

        duplicate = db.scalar(
            select(SurveyBatch.id).where(
                SurveyBatch.code == candidate
            )
        )

        if duplicate is None:
            return candidate

        sequence += 1


def tao_ma_ho(
    db: Session,
    commune: Commune,
) -> str:
    """Tạo mã hộ duy nhất trong toàn hệ thống."""

    existing_count = db.scalar(
        select(func.count(Household.id)).where(
            Household.commune_id == commune.id
        )
    )

    sequence = int(existing_count or 0) + 1

    while True:
        candidate = (
            f"HO-{commune.code}-{sequence:06d}"
        )

        duplicate = db.scalar(
            select(Household.id).where(
                Household.code == candidate
            )
        )

        if duplicate is None:
            return candidate

        sequence += 1


def tao_so_phieu(
    db: Session,
    batch: SurveyBatch,
) -> str:
    """Tạo số phiếu duy nhất trong một đợt điều tra."""

    year_code = batch.school_year.code.replace("-", "")

    existing_count = db.scalar(
        select(func.count(SurveyForm.id)).where(
            SurveyForm.survey_batch_id == batch.id
        )
    )

    sequence = int(existing_count or 0) + 1

    while True:
        candidate = (
            f"PH-{batch.commune.code}-"
            f"{year_code}-{sequence:06d}"
        )

        duplicate = db.scalar(
            select(SurveyForm.id).where(
                SurveyForm.survey_batch_id == batch.id,
                SurveyForm.form_number == candidate,
            )
        )

        if duplicate is None:
            return candidate

        sequence += 1


def lay_du_lieu_danh_muc(
    db: Session,
) -> tuple[list[SchoolYear], list[Commune]]:
    school_years = db.scalars(
        select(SchoolYear)
        .where(SchoolYear.is_active.is_(True))
        .order_by(SchoolYear.code.desc())
    ).all()

    communes = db.scalars(
        select(Commune)
        .where(Commune.is_active.is_(True))
        .order_by(Commune.name.asc())
    ).all()

    return school_years, communes


def lay_dot_dieu_tra(
    db: Session,
    batch_id: int,
) -> SurveyBatch | None:
    return db.scalar(
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
        )
        .where(SurveyBatch.id == batch_id)
    )


def tao_url_danh_sach_ho(
    batch_id: int,
    page: int,
    q: str,
    hamlet: str,
    form_status: str = "",
) -> str:
    parameters: dict[str, str | int] = {
        "page": page,
    }

    if q:
        parameters["q"] = q

    if hamlet:
        parameters["hamlet"] = hamlet

    if form_status:
        parameters["form_status"] = form_status

    return (
        f"/dieu-tra/{batch_id}/ho-dan?"
        + urlencode(parameters)
    )


def tao_ma_doi_tuong(db: Session) -> str:
    """Tạo mã đối tượng duy nhất, ví dụ DT-00000001."""

    existing_count = db.scalar(
        select(func.count(SurveyPerson.id))
    )

    sequence = int(existing_count or 0) + 1

    while True:
        candidate = f"DT-{sequence:08d}"

        duplicate = db.scalar(
            select(SurveyPerson.id).where(
                SurveyPerson.code == candidate
            )
        )

        if duplicate is None:
            return candidate

        sequence += 1


def lay_phieu_ho(
    db: Session,
    batch_id: int,
    household_id: int,
) -> SurveyForm | None:
    """Lấy đúng phiếu của hộ trong đợt điều tra."""

    return db.scalar(
        select(SurveyForm)
        .options(
            selectinload(SurveyForm.household)
            .selectinload(Household.people),
            selectinload(SurveyForm.survey_batch)
            .selectinload(SurveyBatch.school_year),
            selectinload(SurveyForm.survey_batch)
            .selectinload(SurveyBatch.commune),
        )
        .where(
            SurveyForm.survey_batch_id == batch_id,
            SurveyForm.household_id == household_id,
        )
    )


def lay_danh_sach_nguoi_dieu_tra(
    db: Session,
    commune_id: int,
) -> list[dict[str, Any]]:
    """Lấy cán bộ quản lý và giáo viên đang hoạt động trong xã."""

    rows = db.execute(
        text(
            """
            SELECT
                u.id,
                u.username,
                u.full_name,
                u.commune_id,
                u.school_id,
                r.code AS role_code,
                r.name AS role_name,
                s.code AS school_code,
                s.name AS school_name
            FROM users AS u
            JOIN roles AS r
                ON r.id = u.role_id
            LEFT JOIN schools AS s
                ON s.id = u.school_id
            WHERE u.is_active = 1
              AND u.commune_id = :commune_id
              AND r.code IN ('TRUONG', 'GIAO_VIEN')
            ORDER BY
                CASE WHEN s.name IS NULL THEN 1 ELSE 0 END,
                s.name COLLATE NOCASE,
                CASE WHEN r.code = 'TRUONG' THEN 0 ELSE 1 END,
                u.full_name COLLATE NOCASE,
                u.id
            """
        ),
        {"commune_id": commune_id},
    ).mappings().all()

    return [dict(row) for row in rows]


def lay_phan_cong_hien_tai(
    db: Session,
    survey_form_id: int,
) -> list[dict[str, Any]]:
    """Lấy danh sách người đã được phân công vào một phiếu."""

    rows = db.execute(
        text(
            """
            SELECT
                sfi.id AS assignment_id,
                sfi.user_id,
                sfi.order_number,
                sfi.is_primary,
                sfi.assigned_at,
                sfi.notes,
                u.username,
                u.full_name,
                r.code AS role_code,
                r.name AS role_name,
                s.code AS school_code,
                s.name AS school_name
            FROM survey_form_investigators AS sfi
            JOIN users AS u
                ON u.id = sfi.user_id
            JOIN roles AS r
                ON r.id = u.role_id
            LEFT JOIN schools AS s
                ON s.id = u.school_id
            WHERE sfi.survey_form_id = :survey_form_id
            ORDER BY
                sfi.is_primary DESC,
                sfi.order_number ASC,
                sfi.id ASC
            """
        ),
        {"survey_form_id": survey_form_id},
    ).mappings().all()

    return [dict(row) for row in rows]


def hien_thi_trang_phan_cong(
    *,
    request: Request,
    db: Session,
    survey_form: SurveyForm,
    status: str | None = None,
    thong_bao_loi: str | None = None,
    status_code: int = 200,
):
    """Hiển thị trang phân công người điều tra cho một phiếu."""

    candidates = lay_danh_sach_nguoi_dieu_tra(
        db=db,
        commune_id=survey_form.survey_batch.commune_id,
    )
    assignments = lay_phan_cong_hien_tai(
        db=db,
        survey_form_id=survey_form.id,
    )

    selected_ids = {
        int(item["user_id"])
        for item in assignments
    }
    primary_user_id = next(
        (
            int(item["user_id"])
            for item in assignments
            if bool(item["is_primary"])
        ),
        None,
    )

    school_map: dict[int, dict[str, Any]] = {}
    role_counts = {
        "TRUONG": 0,
        "GIAO_VIEN": 0,
    }

    for candidate in candidates:
        role_code = str(candidate.get("role_code") or "")
        if role_code in role_counts:
            role_counts[role_code] += 1

        school_id = candidate.get("school_id")
        if school_id is None:
            continue

        school_id = int(school_id)
        school_map.setdefault(
            school_id,
            {
                "id": school_id,
                "code": candidate.get("school_code") or "",
                "name": candidate.get("school_name") or "",
            },
        )

    schools = sorted(
        school_map.values(),
        key=lambda item: str(item["name"]).casefold(),
    )

    assignment_notes = next(
        (
            str(item["notes"])
            for item in assignments
            if item.get("notes")
        ),
        "",
    )

    return templates.TemplateResponse(
        request=request,
        name="surveys/investigators.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "batch": survey_form.survey_batch,
            "survey_form": survey_form,
            "household": survey_form.household,
            "candidates": candidates,
            "schools": schools,
            "assignments": assignments,
            "selected_ids": selected_ids,
            "primary_user_id": primary_user_id,
            "assignment_notes": assignment_notes,
            "role_counts": role_counts,
            "thong_bao": STATUS_MESSAGES.get(status),
            "thong_bao_loi": thong_bao_loi,
        },
        status_code=status_code,
    )


def tao_url_danh_sach_doi_tuong(
    batch_id: int,
    household_id: int,
    page: int,
    q: str,
    residency_status: str,
) -> str:
    parameters: dict[str, str | int] = {
        "page": page,
    }

    if q:
        parameters["q"] = q

    if residency_status:
        parameters["residency_status"] = residency_status

    return (
        f"/dieu-tra/{batch_id}/ho-dan/"
        f"{household_id}/doi-tuong?"
        + urlencode(parameters)
    )


def hien_thi_trang_doi_tuong(
    *,
    request: Request,
    db: Session,
    survey_form: SurveyForm,
    q: str = "",
    residency_status: str = "",
    page: int = 1,
    status: str | None = None,
    thong_bao_loi: str | None = None,
    form_data: dict[str, Any] | None = None,
    status_code: int = 200,
):
    """Dùng chung khi mở trang và khi biểu mẫu có lỗi."""

    q = q.strip()[:100]

    if residency_status not in RESIDENCY_STATUS_LABELS:
        residency_status = ""

    filters = [
        SurveyPerson.household_id == survey_form.household_id,
        SurveyPerson.is_active.is_(True),
    ]

    if q:
        search_value = f"%{q}%"
        filters.append(
            or_(
                SurveyPerson.code.ilike(search_value),
                SurveyPerson.full_name.ilike(search_value),
                SurveyPerson.personal_id.ilike(search_value),
                SurveyPerson.ministry_student_code.ilike(search_value),
            )
        )

    if residency_status:
        filters.append(
            SurveyPerson.residency_status == residency_status
        )

    total_records = db.scalar(
        select(func.count(SurveyPerson.id)).where(*filters)
    )
    total_records = int(total_records or 0)

    total_pages = max(
        1,
        ceil(total_records / PERSON_PAGE_SIZE),
    )

    page = max(1, page)
    page = min(page, total_pages)
    offset = (page - 1) * PERSON_PAGE_SIZE

    people = db.scalars(
        select(SurveyPerson)
        .options(selectinload(SurveyPerson.student))
        .where(*filters)
        .order_by(
            SurveyPerson.date_of_birth.asc(),
            SurveyPerson.full_name.asc(),
            SurveyPerson.id.asc(),
        )
        .offset(offset)
        .limit(PERSON_PAGE_SIZE)
    ).all()

    students = db.scalars(
        select(Student)
        .where(Student.is_active.is_(True))
        .order_by(
            Student.full_name.asc(),
            Student.date_of_birth.asc(),
            Student.code.asc(),
        )
    ).all()

    used_student_ids = {
        item
        for item in db.scalars(
            select(SurveyPerson.student_id).where(
                SurveyPerson.student_id.is_not(None),
                SurveyPerson.is_active.is_(True),
            )
        ).all()
        if item is not None
    }

    start_page = max(1, page - 2)
    end_page = min(total_pages, page + 2)

    page_links = [
        {
            "number": page_number,
            "active": page_number == page,
            "url": tao_url_danh_sach_doi_tuong(
                batch_id=survey_form.survey_batch_id,
                household_id=survey_form.household_id,
                page=page_number,
                q=q,
                residency_status=residency_status,
            ),
        }
        for page_number in range(start_page, end_page + 1)
    ]

    previous_url = None
    if page > 1:
        previous_url = tao_url_danh_sach_doi_tuong(
            batch_id=survey_form.survey_batch_id,
            household_id=survey_form.household_id,
            page=page - 1,
            q=q,
            residency_status=residency_status,
        )

    next_url = None
    if page < total_pages:
        next_url = tao_url_danh_sach_doi_tuong(
            batch_id=survey_form.survey_batch_id,
            household_id=survey_form.household_id,
            page=page + 1,
            q=q,
            residency_status=residency_status,
        )

    if form_data is None:
        form_data = {
            "student_id": None,
            "full_name": "",
            "date_of_birth": "",
            "gender": "",
            "ethnic_group": "",
            "relationship_to_head": "",
            "personal_id": "",
            "ministry_student_code": "",
            "permanent_address": survey_form.household.address,
            "current_address": survey_form.household.address,
            "residency_status": "THUONG_TRU",
            "special_circumstances": "",
            "notes": "",
        }

    return templates.TemplateResponse(
        request=request,
        name="surveys/people.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "survey_form": survey_form,
            "batch": survey_form.survey_batch,
            "household": survey_form.household,
            "people": people,
            "students": students,
            "used_student_ids": used_student_ids,
            "q": q,
            "selected_residency_status": residency_status,
            "residency_status_labels": RESIDENCY_STATUS_LABELS,
            "page": page,
            "page_size": PERSON_PAGE_SIZE,
            "total_records": total_records,
            "total_pages": total_pages,
            "page_links": page_links,
            "previous_url": previous_url,
            "next_url": next_url,
            "thong_bao": STATUS_MESSAGES.get(status),
            "thong_bao_loi": thong_bao_loi,
            "form_data": form_data,
        },
        status_code=status_code,
    )


# =========================================================
# QUẢN LÝ ĐỢT ĐIỀU TRA
# =========================================================

@router.get("", response_class=HTMLResponse)
def danh_sach_dot_dieu_tra(
    request: Request,
    school_year_id: int | None = None,
    commune_id: int | None = None,
    survey_status: str = "",
    status: str | None = None,
    db: Session = Depends(get_db),
):
    school_years, communes = lay_du_lieu_danh_muc(db)

    valid_year_ids = {item.id for item in school_years}
    valid_commune_ids = {item.id for item in communes}

    if school_year_id not in valid_year_ids:
        school_year_id = None

    if commune_id not in valid_commune_ids:
        commune_id = None

    if survey_status not in BATCH_STATUS_LABELS:
        survey_status = ""

    filters = []

    if school_year_id is not None:
        filters.append(
            SurveyBatch.school_year_id == school_year_id
        )

    if commune_id is not None:
        filters.append(
            SurveyBatch.commune_id == commune_id
        )

    if survey_status:
        filters.append(
            SurveyBatch.status == survey_status
        )

    statement = (
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
            selectinload(SurveyBatch.survey_forms),
        )
        .where(*filters)
        .order_by(
            SurveyBatch.created_at.desc(),
            SurveyBatch.id.desc(),
        )
    )

    survey_batches = db.scalars(statement).all()

    return templates.TemplateResponse(
        request=request,
        name="surveys/index.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "school_years": school_years,
            "communes": communes,
            "survey_batches": survey_batches,
            "selected_school_year_id": school_year_id,
            "selected_commune_id": commune_id,
            "selected_status": survey_status,
            "status_labels": BATCH_STATUS_LABELS,
            "thong_bao": STATUS_MESSAGES.get(status),
            "today": date.today().isoformat(),
        },
    )


@router.post("/them")
def them_dot_dieu_tra(
    request: Request,
    school_year_id: Annotated[int, Form()],
    commune_id: Annotated[int, Form()],
    name: Annotated[str, Form()],
    start_date: Annotated[str, Form()] = "",
    end_date: Annotated[str, Form()] = "",
    notes: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    errors: list[str] = []

    name = chuan_hoa_van_ban(name)
    notes = notes.strip()

    if not name:
        errors.append(
            "Tên đợt điều tra không được để trống."
        )

    if len(name) > 300:
        errors.append(
            "Tên đợt điều tra dài quá 300 ký tự."
        )

    school_year = db.get(
        SchoolYear,
        school_year_id,
    )

    if school_year is None or not school_year.is_active:
        errors.append("Năm học được chọn không hợp lệ.")

    commune = db.get(
        Commune,
        commune_id,
    )

    if commune is None or not commune.is_active:
        errors.append("Xã được chọn không hợp lệ.")

    parsed_start_date, start_error = chuyen_ngay(
        start_date,
        "Ngày bắt đầu",
        bat_buoc=True,
    )

    if start_error:
        errors.append(start_error)

    parsed_end_date, end_error = chuyen_ngay(
        end_date,
        "Ngày kết thúc",
        bat_buoc=False,
    )

    if end_error:
        errors.append(end_error)

    if (
        parsed_start_date is not None
        and parsed_end_date is not None
        and parsed_end_date < parsed_start_date
    ):
        errors.append(
            "Ngày kết thúc không được trước ngày bắt đầu."
        )

    if errors:
        school_years, communes = lay_du_lieu_danh_muc(db)

        survey_batches = db.scalars(
            select(SurveyBatch)
            .options(
                selectinload(SurveyBatch.school_year),
                selectinload(SurveyBatch.commune),
                selectinload(SurveyBatch.survey_forms),
            )
            .order_by(
                SurveyBatch.created_at.desc(),
                SurveyBatch.id.desc(),
            )
        ).all()

        return templates.TemplateResponse(
            request=request,
            name="surveys/index.html",
            context={
                "nguoi_dung": request.scope.get("auth_user"),
                "school_years": school_years,
                "communes": communes,
                "survey_batches": survey_batches,
                "selected_school_year_id": school_year_id,
                "selected_commune_id": commune_id,
                "selected_status": "",
                "status_labels": BATCH_STATUS_LABELS,
                "thong_bao_loi": " ".join(errors),
                "form_data": {
                    "school_year_id": school_year_id,
                    "commune_id": commune_id,
                    "name": name,
                    "start_date": start_date,
                    "end_date": end_date,
                    "notes": notes,
                },
                "today": date.today().isoformat(),
            },
            status_code=400,
        )

    user = request.scope.get("auth_user") or {}

    batch = SurveyBatch(
        code=tao_ma_dot_dieu_tra(
            db=db,
            commune=commune,
            school_year=school_year,
        ),
        name=name,
        school_year_id=school_year_id,
        commune_id=commune_id,
        status="CHUAN_BI",
        start_date=parsed_start_date,
        end_date=parsed_end_date,
        created_by_user_id=user.get("id"),
        notes=notes or None,
    )

    db.add(batch)

    try:
        db.commit()
    except IntegrityError:
        db.rollback()

        return RedirectResponse(
            url="/dieu-tra",
            status_code=303,
        )

    return RedirectResponse(
        url="/dieu-tra?status=created",
        status_code=303,
    )


# =========================================================
# QUẢN LÝ HỘ GIA ĐÌNH TRONG MỘT ĐỢT ĐIỀU TRA
# =========================================================

@router.get(
    "/{batch_id}/ho-dan",
    response_class=HTMLResponse,
)
def danh_sach_ho_dan(
    batch_id: int,
    request: Request,
    q: str = "",
    hamlet: str = "",
    form_status: str = "",
    page: int = 1,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    batch = lay_dot_dieu_tra(db, batch_id)

    if batch is None:
        return RedirectResponse(
            url="/dieu-tra",
            status_code=303,
        )

    q = q.strip()[:100]
    hamlet = chuan_hoa_van_ban(hamlet)[:200]

    if form_status not in FORM_STATUS_LABELS:
        form_status = ""

    filters = [
        SurveyForm.survey_batch_id == batch.id,
    ]

    if q:
        search_value = f"%{q}%"

        filters.append(
            or_(
                SurveyForm.form_number.ilike(search_value),
                Household.code.ilike(search_value),
                Household.head_name.ilike(search_value),
                Household.address.ilike(search_value),
                Household.phone.ilike(search_value),
                Household.hamlet_name.ilike(search_value),
            )
        )

    if hamlet:
        filters.append(
            Household.hamlet_name == hamlet
        )

    if form_status:
        filters.append(
            SurveyForm.status == form_status
        )

    total_records = db.scalar(
        select(func.count(SurveyForm.id))
        .join(
            Household,
            Household.id == SurveyForm.household_id,
        )
        .where(*filters)
    )

    total_records = int(total_records or 0)

    total_pages = max(
        1,
        ceil(total_records / HOUSEHOLD_PAGE_SIZE),
    )

    page = max(1, page)
    page = min(page, total_pages)

    offset = (
        page - 1
    ) * HOUSEHOLD_PAGE_SIZE

    survey_forms = db.scalars(
        select(SurveyForm)
        .join(
            Household,
            Household.id == SurveyForm.household_id,
        )
        .options(
            selectinload(SurveyForm.household)
            .selectinload(Household.people),
            selectinload(SurveyForm.investigators),
            with_loader_criteria(
                SurveyPerson,
                SurveyPerson.is_active.is_(True),
                include_aliases=True,
            ),
        )
        .where(*filters)
        .order_by(
            Household.hamlet_name.asc(),
            Household.head_name.asc(),
            SurveyForm.id.asc(),
        )
        .offset(offset)
        .limit(HOUSEHOLD_PAGE_SIZE)
    ).all()

    investigator_summaries: dict[int, list[dict[str, Any]]] = {}

    if survey_forms:
        form_ids = [int(item.id) for item in survey_forms]
        placeholders = ", ".join(
            f":form_id_{index}"
            for index in range(len(form_ids))
        )
        parameters = {
            f"form_id_{index}": form_id
            for index, form_id in enumerate(form_ids)
        }

        assignment_rows = db.execute(
            text(
                f"""
                SELECT
                    sfi.survey_form_id,
                    sfi.user_id,
                    sfi.is_primary,
                    sfi.order_number,
                    u.full_name,
                    u.username,
                    r.name AS role_name,
                    s.name AS school_name
                FROM survey_form_investigators AS sfi
                JOIN users AS u
                    ON u.id = sfi.user_id
                JOIN roles AS r
                    ON r.id = u.role_id
                LEFT JOIN schools AS s
                    ON s.id = u.school_id
                WHERE sfi.survey_form_id IN ({placeholders})
                ORDER BY
                    sfi.survey_form_id,
                    sfi.is_primary DESC,
                    sfi.order_number,
                    sfi.id
                """
            ),
            parameters,
        ).mappings().all()

        for row in assignment_rows:
            form_id = int(row["survey_form_id"])
            investigator_summaries.setdefault(
                form_id,
                [],
            ).append(dict(row))

    assigned_form_count = int(
        db.execute(
            text(
                """
                SELECT COUNT(DISTINCT sfi.survey_form_id)
                FROM survey_form_investigators AS sfi
                JOIN survey_forms AS sf
                    ON sf.id = sfi.survey_form_id
                WHERE sf.survey_batch_id = :batch_id
                """
            ),
            {"batch_id": batch.id},
        ).scalar()
        or 0
    )
    hamlet_names = db.scalars(
        select(Household.hamlet_name)
        .join(
            SurveyForm,
            SurveyForm.household_id == Household.id,
        )
        .where(
            SurveyForm.survey_batch_id == batch.id,
            Household.hamlet_name.is_not(None),
            Household.hamlet_name != "",
        )
        .distinct()
        .order_by(Household.hamlet_name.asc())
    ).all()

    status_rows = db.execute(
        select(
            SurveyForm.status,
            func.count(SurveyForm.id),
        )
        .where(SurveyForm.survey_batch_id == batch.id)
        .group_by(SurveyForm.status)
    ).all()

    status_counts = {
        code: 0
        for code in FORM_STATUS_LABELS
    }

    for status_code, count_value in status_rows:
        status_counts[status_code] = int(count_value or 0)

    batch_total_records = sum(status_counts.values())
    unassigned_form_count = max(
        0,
        batch_total_records - assigned_form_count,
    )
    completed_records = status_counts.get("DA_HOAN_THANH", 0)
    progress_percent = (
        round(completed_records * 100 / batch_total_records)
        if batch_total_records
        else 0
    )

    start_page = max(1, page - 2)
    end_page = min(total_pages, page + 2)

    page_links = [
        {
            "number": page_number,
            "active": page_number == page,
            "url": tao_url_danh_sach_ho(
                batch_id=batch.id,
                page=page_number,
                q=q,
                hamlet=hamlet,
                form_status=form_status,
            ),
        }
        for page_number in range(
            start_page,
            end_page + 1,
        )
    ]

    previous_url = None

    if page > 1:
        previous_url = tao_url_danh_sach_ho(
            batch_id=batch.id,
            page=page - 1,
            q=q,
            hamlet=hamlet,
            form_status=form_status,
        )

    next_url = None

    if page < total_pages:
        next_url = tao_url_danh_sach_ho(
            batch_id=batch.id,
            page=page + 1,
            q=q,
            hamlet=hamlet,
            form_status=form_status,
        )

    return templates.TemplateResponse(
        request=request,
        name="surveys/households.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "batch": batch,
            "survey_forms": survey_forms,
            "investigator_summaries": investigator_summaries,
            "assigned_form_count": assigned_form_count,
            "unassigned_form_count": unassigned_form_count,
            "hamlet_names": hamlet_names,
            "q": q,
            "selected_hamlet": hamlet,
            "selected_form_status": form_status,
            "status_counts": status_counts,
            "batch_total_records": batch_total_records,
            "progress_percent": progress_percent,
            "page": page,
            "page_size": HOUSEHOLD_PAGE_SIZE,
            "total_records": total_records,
            "total_pages": total_pages,
            "page_links": page_links,
            "previous_url": previous_url,
            "next_url": next_url,
            "form_status_labels": FORM_STATUS_LABELS,
            "thong_bao": STATUS_MESSAGES.get(status),
        },
    )



# =========================================================
# PHÂN CÔNG NGƯỜI ĐIỀU TRA
# =========================================================

@router.get(
    "/{batch_id}/ho-dan/{household_id}/phan-cong",
    response_class=HTMLResponse,
)
def trang_phan_cong_nguoi_dieu_tra(
    batch_id: int,
    household_id: int,
    request: Request,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )

    if survey_form is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    return hien_thi_trang_phan_cong(
        request=request,
        db=db,
        survey_form=survey_form,
        status=status,
    )


@router.post(
    "/{batch_id}/ho-dan/{household_id}/phan-cong/luu",
)
async def luu_phan_cong_nguoi_dieu_tra(
    batch_id: int,
    household_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )

    if survey_form is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    submitted = await request.form()
    raw_ids = submitted.getlist("investigator_ids")

    selected_ids: list[int] = []
    seen_ids: set[int] = set()

    for raw_value in raw_ids:
        try:
            user_id = int(str(raw_value))
        except (TypeError, ValueError):
            continue

        if user_id <= 0 or user_id in seen_ids:
            continue

        seen_ids.add(user_id)
        selected_ids.append(user_id)

    if not selected_ids:
        return hien_thi_trang_phan_cong(
            request=request,
            db=db,
            survey_form=survey_form,
            thong_bao_loi=(
                "Bạn cần chọn ít nhất một người điều tra. "
                "Nếu muốn bỏ toàn bộ phân công, hãy dùng nút "
                "Xóa toàn bộ phân công."
            ),
            status_code=400,
        )

    primary_value = str(
        submitted.get("primary_user_id") or ""
    ).strip()

    try:
        primary_user_id = int(primary_value)
    except ValueError:
        primary_user_id = selected_ids[0]

    if primary_user_id not in seen_ids:
        primary_user_id = selected_ids[0]

    assignment_notes = chuan_hoa_van_ban(
        str(submitted.get("assignment_notes") or "")
    )[:1000]

    placeholders = ", ".join(
        f":user_id_{index}"
        for index in range(len(selected_ids))
    )
    parameters: dict[str, Any] = {
        "commune_id": survey_form.survey_batch.commune_id,
    }
    parameters.update(
        {
            f"user_id_{index}": user_id
            for index, user_id in enumerate(selected_ids)
        }
    )

    allowed_ids = {
        int(value)
        for value in db.execute(
            text(
                f"""
                SELECT u.id
                FROM users AS u
                JOIN roles AS r
                    ON r.id = u.role_id
                WHERE u.id IN ({placeholders})
                  AND u.is_active = 1
                  AND u.commune_id = :commune_id
                  AND r.code IN ('TRUONG', 'GIAO_VIEN')
                """
            ),
            parameters,
        ).scalars().all()
    }

    if allowed_ids != seen_ids:
        return hien_thi_trang_phan_cong(
            request=request,
            db=db,
            survey_form=survey_form,
            thong_bao_loi=(
                "Danh sách có tài khoản không còn hoạt động, "
                "không thuộc đúng xã/phường hoặc không đúng vai trò. "
                "Hãy tải lại trang và chọn lại."
            ),
            status_code=400,
        )

    now_value = datetime.now().isoformat(
        sep=" ",
        timespec="seconds",
    )

    try:
        db.execute(
            text(
                """
                DELETE FROM survey_form_investigators
                WHERE survey_form_id = :survey_form_id
                """
            ),
            {"survey_form_id": survey_form.id},
        )

        for order_number, user_id in enumerate(
            selected_ids,
            start=1,
        ):
            db.execute(
                text(
                    """
                    INSERT INTO survey_form_investigators (
                        survey_form_id,
                        user_id,
                        order_number,
                        is_primary,
                        assigned_at,
                        notes,
                        created_at
                    ) VALUES (
                        :survey_form_id,
                        :user_id,
                        :order_number,
                        :is_primary,
                        :assigned_at,
                        :notes,
                        :created_at
                    )
                    """
                ),
                {
                    "survey_form_id": survey_form.id,
                    "user_id": user_id,
                    "order_number": order_number,
                    "is_primary": (
                        1 if user_id == primary_user_id else 0
                    ),
                    "assigned_at": now_value,
                    "notes": assignment_notes or None,
                    "created_at": now_value,
                },
            )

        db.commit()
    except Exception:
        db.rollback()

        return hien_thi_trang_phan_cong(
            request=request,
            db=db,
            survey_form=survey_form,
            thong_bao_loi=(
                "Không thể lưu phân công. Dữ liệu cũ vẫn được giữ nguyên."
            ),
            status_code=500,
        )

    return RedirectResponse(
        url=(
            f"/dieu-tra/{batch_id}/ho-dan/"
            f"{household_id}/phan-cong?status=investigators_saved"
        ),
        status_code=303,
    )


@router.post(
    "/{batch_id}/ho-dan/{household_id}/phan-cong/xoa",
)
def xoa_phan_cong_nguoi_dieu_tra(
    batch_id: int,
    household_id: int,
    db: Session = Depends(get_db),
):
    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )

    if survey_form is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    try:
        db.execute(
            text(
                """
                DELETE FROM survey_form_investigators
                WHERE survey_form_id = :survey_form_id
                """
            ),
            {"survey_form_id": survey_form.id},
        )
        db.commit()
    except Exception:
        db.rollback()
        return RedirectResponse(
            url=(
                f"/dieu-tra/{batch_id}/ho-dan/"
                f"{household_id}/phan-cong"
            ),
            status_code=303,
        )

    return RedirectResponse(
        url=(
            f"/dieu-tra/{batch_id}/ho-dan/"
            f"{household_id}/phan-cong?status=investigators_cleared"
        ),
        status_code=303,
    )


# =========================================================
# TRẠNG THÁI PHIẾU VÀ TIẾN ĐỘ ĐIỀU TRA
# =========================================================

def hien_thi_trang_cap_nhat_phieu(
    *,
    request: Request,
    survey_form: SurveyForm,
    active_people_count: int,
    status: str | None = None,
    thong_bao_loi: str | None = None,
    form_data: dict[str, Any] | None = None,
    status_code: int = 200,
):
    """Hiển thị biểu mẫu cập nhật trạng thái một phiếu điều tra."""

    if form_data is None:
        form_data = {
            "form_status": survey_form.status,
            "survey_date": (
                survey_form.survey_date.isoformat()
                if survey_form.survey_date
                else ""
            ),
            "village_head_name": (
                survey_form.village_head_name or ""
            ),
            "household_representative_name": (
                survey_form.household_representative_name or ""
            ),
            "household_confirmed": (
                survey_form.household_confirmed_at is not None
            ),
            "commune_confirmed": (
                survey_form.commune_confirmed_at is not None
            ),
            "notes": survey_form.notes or "",
        }

    return templates.TemplateResponse(
        request=request,
        name="surveys/form_status.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "batch": survey_form.survey_batch,
            "survey_form": survey_form,
            "household": survey_form.household,
            "active_people_count": active_people_count,
            "form_status_labels": FORM_STATUS_LABELS,
            "thong_bao": STATUS_MESSAGES.get(status),
            "thong_bao_loi": thong_bao_loi,
            "form_data": form_data,
            "today": date.today().isoformat(),
        },
        status_code=status_code,
    )


@router.get(
    "/{batch_id}/ho-dan/{household_id}/phieu",
    response_class=HTMLResponse,
)
def trang_cap_nhat_phieu(
    batch_id: int,
    household_id: int,
    request: Request,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )

    if survey_form is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    active_people_count = int(
        db.scalar(
            select(func.count(SurveyPerson.id)).where(
                SurveyPerson.household_id == household_id,
                SurveyPerson.is_active.is_(True),
            )
        )
        or 0
    )

    return hien_thi_trang_cap_nhat_phieu(
        request=request,
        survey_form=survey_form,
        active_people_count=active_people_count,
        status=status,
    )


@router.post(
    "/{batch_id}/ho-dan/{household_id}/phieu/cap-nhat"
)
def cap_nhat_phieu(
    batch_id: int,
    household_id: int,
    request: Request,
    form_status: Annotated[str, Form()],
    survey_date: Annotated[str, Form()] = "",
    village_head_name: Annotated[str, Form()] = "",
    household_representative_name: Annotated[str, Form()] = "",
    household_confirmed: Annotated[str, Form()] = "",
    commune_confirmed: Annotated[str, Form()] = "",
    notes: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )

    if survey_form is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    active_people_count = int(
        db.scalar(
            select(func.count(SurveyPerson.id)).where(
                SurveyPerson.household_id == household_id,
                SurveyPerson.is_active.is_(True),
            )
        )
        or 0
    )

    form_status = form_status.strip().upper()
    village_head_name = chuan_hoa_van_ban(village_head_name)
    household_representative_name = chuan_hoa_van_ban(
        household_representative_name
    )
    notes = notes.strip()

    parsed_survey_date, date_error = chuyen_ngay(
        survey_date,
        "Ngày điều tra",
        bat_buoc=form_status != "CHUA_DIEU_TRA",
    )

    errors: list[str] = []

    if form_status not in FORM_STATUS_LABELS:
        errors.append("Trạng thái phiếu không hợp lệ.")

    if date_error:
        errors.append(date_error)

    if parsed_survey_date and parsed_survey_date > date.today():
        errors.append("Ngày điều tra không được lớn hơn ngày hiện tại.")

    if len(village_head_name) > 200:
        errors.append("Tên trưởng thôn/xóm dài quá 200 ký tự.")

    if len(household_representative_name) > 200:
        errors.append("Tên người đại diện hộ dài quá 200 ký tự.")

    if form_status == "DA_HOAN_THANH" and active_people_count == 0:
        errors.append(
            "Không thể hoàn thành phiếu vì hộ chưa có đối tượng đang hoạt động."
        )

    if form_status == "CAN_BO_SUNG" and not notes:
        errors.append(
            "Phiếu cần bổ sung phải ghi rõ nội dung cần bổ sung."
        )

    form_data = {
        "form_status": form_status,
        "survey_date": survey_date,
        "village_head_name": village_head_name,
        "household_representative_name": (
            household_representative_name
        ),
        "household_confirmed": bool(household_confirmed),
        "commune_confirmed": bool(commune_confirmed),
        "notes": notes,
    }

    if errors:
        return hien_thi_trang_cap_nhat_phieu(
            request=request,
            survey_form=survey_form,
            active_people_count=active_people_count,
            thong_bao_loi=" ".join(errors),
            form_data=form_data,
            status_code=400,
        )

    now = datetime.now()

    survey_form.status = form_status
    survey_form.survey_date = parsed_survey_date
    survey_form.village_head_name = village_head_name or None
    survey_form.household_representative_name = (
        household_representative_name or None
    )
    survey_form.household_confirmed_at = (
        survey_form.household_confirmed_at or now
        if household_confirmed
        else None
    )
    survey_form.commune_confirmed_at = (
        survey_form.commune_confirmed_at or now
        if commune_confirmed
        else None
    )
    survey_form.notes = notes or None

    # Đồng bộ ảnh chụp thông tin hộ tại thời điểm cập nhật phiếu.
    survey_form.head_name_snapshot = survey_form.household.head_name
    survey_form.address_snapshot = survey_form.household.address
    survey_form.hamlet_name_snapshot = survey_form.household.hamlet_name

    db.commit()

    return RedirectResponse(
        url=(
            f"/dieu-tra/{batch_id}/ho-dan/{household_id}/phieu?"
            + urlencode({"status": "form_updated"})
        ),
        status_code=303,
    )


@router.get(
    "/{batch_id}/tien-do",
    response_class=HTMLResponse,
)
def tien_do_dieu_tra(
    batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    batch = lay_dot_dieu_tra(db, batch_id)

    if batch is None:
        return RedirectResponse(url="/dieu-tra", status_code=303)

    status_rows = db.execute(
        select(
            SurveyForm.status,
            func.count(SurveyForm.id),
        )
        .where(SurveyForm.survey_batch_id == batch.id)
        .group_by(SurveyForm.status)
    ).all()

    status_counts = {code: 0 for code in FORM_STATUS_LABELS}
    for status_code, count_value in status_rows:
        status_counts[status_code] = int(count_value or 0)

    total_forms = sum(status_counts.values())
    completed_forms = status_counts.get("DA_HOAN_THANH", 0)
    progress_percent = (
        round(completed_forms * 100 / total_forms)
        if total_forms
        else 0
    )

    total_people = int(
        db.scalar(
            select(func.count(SurveyPerson.id))
            .join(
                SurveyForm,
                SurveyForm.household_id == SurveyPerson.household_id,
            )
            .where(
                SurveyForm.survey_batch_id == batch.id,
                SurveyPerson.is_active.is_(True),
            )
        )
        or 0
    )

    hamlet_rows = db.execute(
        select(
            Household.hamlet_name,
            SurveyForm.status,
            func.count(SurveyForm.id),
        )
        .join(
            Household,
            Household.id == SurveyForm.household_id,
        )
        .where(SurveyForm.survey_batch_id == batch.id)
        .group_by(Household.hamlet_name, SurveyForm.status)
        .order_by(Household.hamlet_name.asc())
    ).all()

    grouped: dict[str, dict[str, int]] = {}
    for hamlet_name, status_code, count_value in hamlet_rows:
        display_name = hamlet_name or "Chưa xác định thôn/xóm"
        grouped.setdefault(
            display_name,
            {code: 0 for code in FORM_STATUS_LABELS},
        )
        grouped[display_name][status_code] = int(count_value or 0)

    hamlet_progress: list[dict[str, Any]] = []
    for hamlet_name, counts in grouped.items():
        hamlet_total = sum(counts.values())
        hamlet_completed = counts.get("DA_HOAN_THANH", 0)
        hamlet_progress.append(
            {
                "hamlet_name": hamlet_name,
                "counts": counts,
                "total": hamlet_total,
                "progress_percent": (
                    round(hamlet_completed * 100 / hamlet_total)
                    if hamlet_total
                    else 0
                ),
            }
        )

    return templates.TemplateResponse(
        request=request,
        name="surveys/progress.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "batch": batch,
            "form_status_labels": FORM_STATUS_LABELS,
            "status_counts": status_counts,
            "total_forms": total_forms,
            "total_people": total_people,
            "progress_percent": progress_percent,
            "hamlet_progress": hamlet_progress,
        },
    )


# =========================================================
# XUẤT EXCEL ĐI ĐIỀU TRA
# =========================================================

def gia_tri_co_khong(value: bool | None) -> str:
    """Đổi dữ liệu Boolean sang chữ dễ đọc trong Excel."""

    if value is None:
        return ""

    return "Có" if value else "Không"


def tao_tieu_de_excel(
    worksheet,
    *,
    title: str,
    last_column: int,
) -> None:
    """Tạo phần tiêu đề chung của trang Excel điều tra."""

    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=last_column,
    )
    title_cell = worksheet.cell(row=1, column=1, value=title)
    title_cell.font = Font(size=16, bold=True)
    title_cell.alignment = Alignment(horizontal="center")


@router.get("/{batch_id}/ho-dan/xuat-excel")
def xuat_excel_di_dieu_tra(
    batch_id: int,
    db: Session = Depends(get_db),
):
    """Xuất toàn bộ hộ và đối tượng của một đợt ra Excel."""

    batch = lay_dot_dieu_tra(db, batch_id)

    if batch is None:
        return RedirectResponse(
            url="/dieu-tra",
            status_code=303,
        )

    survey_forms = db.scalars(
        select(SurveyForm)
        .join(
            Household,
            Household.id == SurveyForm.household_id,
        )
        .options(
            selectinload(SurveyForm.household)
            .selectinload(Household.people),
            with_loader_criteria(
                SurveyPerson,
                SurveyPerson.is_active.is_(True),
                include_aliases=True,
            ),
        )
        .where(SurveyForm.survey_batch_id == batch.id)
        .order_by(
            Household.hamlet_name.asc(),
            Household.head_name.asc(),
            SurveyForm.id.asc(),
        )
    ).all()

    person_ids = [
        person.id
        for survey_form in survey_forms
        for person in survey_form.household.people
        if person.is_active
    ]

    year_records = db.scalars(
        select(SurveyPersonYearRecord).where(
            SurveyPersonYearRecord.survey_person_id.in_(person_ids),
            SurveyPersonYearRecord.school_year_id
            == batch.school_year_id,
        )
    ).all() if person_ids else []

    record_by_person_id = {
        item.survey_person_id: item
        for item in year_records
    }

    school_ids = {
        item.school_id
        for item in year_records
        if item.school_id is not None
    }
    class_ids = {
        item.class_id
        for item in year_records
        if item.class_id is not None
    }

    schools = {
        item.id: item
        for item in db.scalars(
            select(School).where(School.id.in_(school_ids))
        ).all()
    } if school_ids else {}

    classrooms = {
        item.id: item
        for item in db.scalars(
            select(Classroom).where(Classroom.id.in_(class_ids))
        ).all()
    } if class_ids else {}

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "DIEU_TRA"

    visible_headers = [
        "STT",
        "Số phiếu",
        "Mã hộ",
        "Thôn/xóm/khối/bản",
        "Chủ hộ",
        "Địa chỉ hộ",
        "Điện thoại chủ hộ",
        "Trạng thái phiếu",
        "Mã đối tượng",
        "Họ và tên",
        "Ngày sinh",
        "Giới tính",
        "Dân tộc",
        "Quan hệ với chủ hộ",
        "Số định danh cá nhân",
        "Mã Bộ GD&ĐT",
        "Trạng thái cư trú",
        "Địa chỉ thường trú",
        "Chỗ ở hiện nay",
        "Tình trạng học tập",
        "Tên trường",
        "Tên lớp",
        "Hoàn thành CT GDMN 5 tuổi",
        "Đi học đủ ngày theo quy định",
        "Đi học chuyên cần",
        "Được chuẩn bị tiếng Việt",
        "Theo dõi biểu đồ cân nặng",
        "Suy dinh dưỡng nhẹ cân",
        "Theo dõi biểu đồ chiều cao",
        "Suy dinh dưỡng thấp còi",
        "Hoàn cảnh đặc biệt",
        "Ghi chú đối tượng",
        "Ghi chú năm học",
    ]

    technical_headers = [
        "__survey_form_id",
        "__household_id",
        "__survey_person_id",
        "__school_year_id",
        "__school_id",
        "__class_id",
        "__is_active",
    ]

    all_headers = visible_headers + technical_headers
    last_column = len(all_headers)

    tao_tieu_de_excel(
        worksheet,
        title="DANH SÁCH ĐIỀU TRA PHỔ CẬP GIÁO DỤC",
        last_column=last_column,
    )

    worksheet.cell(
        row=2,
        column=1,
        value=f"Đợt điều tra: {batch.code} – {batch.name}",
    )
    worksheet.cell(
        row=3,
        column=1,
        value=(
            f"Năm học: {batch.school_year.code} | "
            f"Xã/phường: {batch.commune.name} "
            f"(Mã {batch.commune.code})"
        ),
    )
    worksheet.cell(
        row=4,
        column=1,
        value=(
            "Chỉ sửa dữ liệu từ hàng 7 trở xuống. "
            "Không đổi tên cột, không xóa cột kỹ thuật và không đổi tên sheet."
        ),
    )

    header_row = 6
    for column_index, header in enumerate(all_headers, start=1):
        cell = worksheet.cell(
            row=header_row,
            column=column_index,
            value=header,
        )
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    thin_side = Side(style="thin", color="D9E2F3")
    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    row_number = header_row + 1
    sequence = 1

    for survey_form in survey_forms:
        household = survey_form.household
        active_people = [
            item
            for item in household.people
            if item.is_active
        ]
        active_people.sort(
            key=lambda item: (
                item.date_of_birth or date.max,
                item.full_name,
                item.id,
            )
        )

        # Hộ chưa có người vẫn được xuất một dòng để cán bộ điều tra
        # có thể bổ sung dữ liệu ngoài thực địa.
        people_for_export: list[SurveyPerson | None] = (
            active_people if active_people else [None]
        )

        for person in people_for_export:
            record = (
                record_by_person_id.get(person.id)
                if person is not None
                else None
            )

            school = (
                schools.get(record.school_id)
                if record is not None
                else None
            )
            classroom = (
                classrooms.get(record.class_id)
                if record is not None
                else None
            )

            school_name = ""
            class_name = ""
            if record is not None:
                school_name = (
                    record.school_name_reported
                    or (school.name if school is not None else "")
                )
                class_name = (
                    record.class_name_reported
                    or (
                        classroom.name
                        if classroom is not None
                        else ""
                    )
                )

            values = [
                sequence,
                survey_form.form_number,
                household.code,
                household.hamlet_name or "",
                household.head_name,
                household.address,
                household.phone or "",
                FORM_STATUS_LABELS.get(
                    survey_form.status,
                    survey_form.status,
                ),
                person.code if person is not None else "",
                person.full_name if person is not None else "",
                (
                    person.date_of_birth
                    if person is not None
                    else ""
                ),
                person.gender or "" if person is not None else "",
                (
                    person.ethnic_group or ""
                    if person is not None
                    else ""
                ),
                (
                    person.relationship_to_head or ""
                    if person is not None
                    else ""
                ),
                (
                    person.personal_id or ""
                    if person is not None
                    else ""
                ),
                (
                    person.ministry_student_code or ""
                    if person is not None
                    else ""
                ),
                (
                    RESIDENCY_STATUS_LABELS.get(
                        person.residency_status,
                        person.residency_status,
                    )
                    if person is not None
                    else ""
                ),
                (
                    person.permanent_address or ""
                    if person is not None
                    else ""
                ),
                (
                    person.current_address or ""
                    if person is not None
                    else ""
                ),
                (
                    LEARNING_STATUS_LABELS.get(
                        record.learning_status,
                        record.learning_status,
                    )
                    if record is not None
                    else ""
                ),
                school_name,
                class_name,
                (
                    gia_tri_co_khong(
                        record.completed_preschool_5
                    )
                    if record is not None
                    else ""
                ),
                (
                    gia_tri_co_khong(
                        record.attends_required_days
                    )
                    if record is not None
                    else ""
                ),
                (
                    gia_tri_co_khong(
                        record.attends_regularly
                    )
                    if record is not None
                    else ""
                ),
                (
                    gia_tri_co_khong(
                        record.prepared_vietnamese
                    )
                    if record is not None
                    else ""
                ),
                (
                    gia_tri_co_khong(record.weight_monitored)
                    if record is not None
                    else ""
                ),
                (
                    gia_tri_co_khong(record.underweight)
                    if record is not None
                    else ""
                ),
                (
                    gia_tri_co_khong(record.height_monitored)
                    if record is not None
                    else ""
                ),
                (
                    gia_tri_co_khong(record.stunted)
                    if record is not None
                    else ""
                ),
                (
                    (
                        record.special_circumstances
                        if record is not None
                        and record.special_circumstances
                        else person.special_circumstances or ""
                    )
                    if person is not None
                    else ""
                ),
                person.notes or "" if person is not None else "",
                record.notes or "" if record is not None else "",
                survey_form.id,
                household.id,
                person.id if person is not None else "",
                batch.school_year_id,
                record.school_id if record is not None else "",
                record.class_id if record is not None else "",
                True if person is not None else "",
            ]

            for column_index, value in enumerate(values, start=1):
                cell = worksheet.cell(
                    row=row_number,
                    column=column_index,
                    value=value,
                )
                cell.border = border
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

            worksheet.cell(
                row=row_number,
                column=11,
            ).number_format = "dd/mm/yyyy"

            row_number += 1
            sequence += 1

    last_data_row = max(header_row + 1, row_number - 1)

    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.auto_filter.ref = (
        f"A{header_row}:"
        f"{get_column_letter(len(visible_headers))}{last_data_row}"
    )
    worksheet.row_dimensions[header_row].height = 42

    widths = {
        1: 7,
        2: 28,
        3: 22,
        4: 22,
        5: 24,
        6: 34,
        7: 19,
        8: 18,
        9: 18,
        10: 26,
        11: 14,
        12: 12,
        13: 14,
        14: 22,
        15: 24,
        16: 20,
        17: 20,
        18: 34,
        19: 34,
        20: 22,
        21: 32,
        22: 24,
        23: 22,
        24: 24,
        25: 20,
        26: 23,
        27: 23,
        28: 24,
        29: 23,
        30: 23,
        31: 34,
        32: 34,
        33: 34,
    }

    for column_index, width in widths.items():
        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = width

    technical_start = len(visible_headers) + 1
    for column_index in range(technical_start, last_column + 1):
        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].hidden = True

    # Danh mục dùng cho các ô lựa chọn trong file điều tra.
    list_sheet = workbook.create_sheet("DANH_MUC")
    list_values = {
        1: ["Giới tính", "Nam", "Nữ", "Khác"],
        2: [
            "Trạng thái cư trú",
            *RESIDENCY_STATUS_LABELS.values(),
        ],
        3: [
            "Tình trạng học tập",
            *LEARNING_STATUS_LABELS.values(),
        ],
        4: ["Có/Không", "Có", "Không"],
    }

    for column_index, items in list_values.items():
        for item_index, value in enumerate(items, start=1):
            list_sheet.cell(
                row=item_index,
                column=column_index,
                value=value,
            )

    list_sheet.sheet_state = "hidden"

    validation_gender = DataValidation(
        type="list",
        formula1="'DANH_MUC'!$A$2:$A$4",
        allow_blank=True,
    )
    validation_residency = DataValidation(
        type="list",
        formula1=(
            f"'DANH_MUC'!$B$2:$B$"
            f"{len(RESIDENCY_STATUS_LABELS) + 1}"
        ),
        allow_blank=True,
    )
    validation_learning = DataValidation(
        type="list",
        formula1=(
            f"'DANH_MUC'!$C$2:$C$"
            f"{len(LEARNING_STATUS_LABELS) + 1}"
        ),
        allow_blank=True,
    )
    validation_boolean = DataValidation(
        type="list",
        formula1="'DANH_MUC'!$D$2:$D$3",
        allow_blank=True,
    )

    for validation in (
        validation_gender,
        validation_residency,
        validation_learning,
        validation_boolean,
    ):
        worksheet.add_data_validation(validation)

    validation_gender.add(f"L7:L5000")
    validation_residency.add(f"Q7:Q5000")
    validation_learning.add(f"T7:T5000")
    for column_letter in (
        "W", "X", "Y", "Z", "AA", "AB", "AC", "AD"
    ):
        validation_boolean.add(
            f"{column_letter}7:{column_letter}5000"
        )

    guide = workbook.create_sheet("HUONG_DAN", 0)
    guide["A1"] = "HƯỚNG DẪN SỬ DỤNG FILE ĐIỀU TRA"
    guide["A1"].font = Font(size=16, bold=True)
    guide["A3"] = "1. Mở sheet DIEU_TRA để rà soát và cập nhật dữ liệu."
    guide["A4"] = (
        "2. Không đổi tên sheet, không đổi tiêu đề cột và không xóa "
        "các cột kỹ thuật đang được ẩn."
    )
    guide["A5"] = (
        "3. Không sửa Số phiếu, Mã hộ và Mã đối tượng đã có."
    )
    guide["A6"] = (
        "4. Để bổ sung người mới vào hộ: sao chép một dòng của đúng hộ, "
        "giữ Số phiếu và Mã hộ, sau đó xóa Mã đối tượng và nhập thông tin mới."
    )
    guide["A7"] = (
        "5. Các cột Có/Không, giới tính, cư trú và tình trạng học tập "
        "đã có danh sách lựa chọn."
    )
    guide["A8"] = (
        "6. Sau khi điều tra xong, lưu nguyên định dạng .xlsx để sử dụng "
        "chức năng Nhập Excel cập nhật ở bài tiếp theo."
    )
    guide["A10"] = f"Mã đợt: {batch.code}"
    guide["A11"] = f"Năm học: {batch.school_year.code}"
    guide["A12"] = f"Xã/phường: {batch.commune.name}"
    guide.column_dimensions["A"].width = 115
    for row in range(1, 13):
        guide.cell(row=row, column=1).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    file_name = (
        f"dieu_tra_{batch.commune.code}_"
        f"{batch.school_year.code.replace('-', '')}.xlsx"
    )

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{file_name}"'
            )
        },
    )


# =========================================================
# NHẬP EXCEL CẬP NHẬT
# =========================================================

EXCEL_IMPORT_REQUIRED_HEADERS = (
    "Số phiếu",
    "Mã hộ",
    "Thôn/xóm/khối/bản",
    "Chủ hộ",
    "Địa chỉ hộ",
    "Điện thoại chủ hộ",
    "Trạng thái phiếu",
    "Mã đối tượng",
    "Họ và tên",
    "Ngày sinh",
    "Giới tính",
    "Dân tộc",
    "Quan hệ với chủ hộ",
    "Số định danh cá nhân",
    "Mã Bộ GD&ĐT",
    "Trạng thái cư trú",
    "Địa chỉ thường trú",
    "Chỗ ở hiện nay",
    "Tình trạng học tập",
    "Tên trường",
    "Tên lớp",
    "Hoàn thành CT GDMN 5 tuổi",
    "Đi học đủ ngày theo quy định",
    "Đi học chuyên cần",
    "Được chuẩn bị tiếng Việt",
    "Theo dõi biểu đồ cân nặng",
    "Suy dinh dưỡng nhẹ cân",
    "Theo dõi biểu đồ chiều cao",
    "Suy dinh dưỡng thấp còi",
    "Hoàn cảnh đặc biệt",
    "Ghi chú đối tượng",
    "Ghi chú năm học",
    "__survey_form_id",
    "__household_id",
    "__survey_person_id",
    "__school_year_id",
    "__school_id",
    "__class_id",
)


def gia_tri_excel_dang_chu(value: Any) -> str:
    """Đọc ô Excel thành văn bản, không tạo đuôi .0 ngoài ý muốn."""

    if value is None:
        return ""
    if isinstance(value, bool):
        return "Có" if value else "Không"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def khoa_so_sanh_excel(value: Any) -> str:
    """Khóa so sánh không phân biệt dấu, hoa thường và khoảng trắng."""

    text = gia_tri_excel_dang_chu(value)
    text = unicodedata.normalize("NFD", text)
    text = "".join(
        char
        for char in text
        if unicodedata.category(char) != "Mn"
    )
    return " ".join(text.lower().strip().split())


def so_nguyen_excel(value: Any) -> int | None:
    """Đọc ID kỹ thuật từ Excel."""

    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value).strip()
    return int(text) if text.isdigit() else None


def ngay_excel(value: Any) -> date | None:
    """Đọc ngày Excel hoặc chuỗi ngày thông dụng."""

    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value).strip()
    for format_code in (
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d.%m.%Y",
    ):
        try:
            return datetime.strptime(text, format_code).date()
        except ValueError:
            continue
    return None


def co_khong_excel(value: Any) -> tuple[bool | None, bool]:
    """Trả về (giá trị, hợp lệ). Ô trống tương ứng None."""

    if value in (None, ""):
        return None, True
    if isinstance(value, bool):
        return value, True
    key = khoa_so_sanh_excel(value)
    if key in {"co", "yes", "true", "1", "x"}:
        return True, True
    if key in {"khong", "no", "false", "0"}:
        return False, True
    return None, False


def tao_ban_sao_truoc_khi_nhap_excel() -> Path:
    """Tạo bản sao SQLite trước khi nhập dữ liệu."""

    database_path = Path(DATABASE_PATH)
    backup_dir = database_path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = backup_dir / f"before_excel_import_{timestamp}.db"

    source = sqlite3.connect(str(database_path))
    target = sqlite3.connect(str(backup_path))
    try:
        source.backup(target)
    finally:
        target.close()
        source.close()
    return backup_path


def hien_thi_nhap_excel(
    *,
    request: Request,
    batch: SurveyBatch,
    errors: list[dict[str, Any]] | None = None,
    summary: dict[str, Any] | None = None,
    file_name: str = "",
    status_code: int = 200,
):
    return templates.TemplateResponse(
        request=request,
        name="surveys/excel_import.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "batch": batch,
            "errors": errors or [],
            "summary": summary,
            "file_name": file_name,
        },
        status_code=status_code,
    )


@router.get(
    "/{batch_id}/ho-dan/nhap-excel",
    response_class=HTMLResponse,
)
def trang_nhap_excel(
    batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    batch = lay_dot_dieu_tra(db, batch_id)
    if batch is None:
        return RedirectResponse(url="/dieu-tra", status_code=303)
    return hien_thi_nhap_excel(request=request, batch=batch)


@router.post(
    "/{batch_id}/ho-dan/nhap-excel",
    response_class=HTMLResponse,
)
async def nhap_excel_cap_nhat(
    batch_id: int,
    request: Request,
    excel_file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """Kiểm tra toàn bộ file rồi mới cập nhật trong một giao dịch."""

    batch = lay_dot_dieu_tra(db, batch_id)
    if batch is None:
        return RedirectResponse(url="/dieu-tra", status_code=303)

    if batch.status == "DA_KET_THUC":
        return hien_thi_nhap_excel(
            request=request,
            batch=batch,
            errors=[{
                "row": "—",
                "message": "Đợt điều tra đã kết thúc, không thể nhập Excel.",
            }],
            file_name=excel_file.filename or "",
            status_code=400,
        )

    file_name = excel_file.filename or ""
    if not file_name.lower().endswith(".xlsx"):
        return hien_thi_nhap_excel(
            request=request,
            batch=batch,
            errors=[{
                "row": "—",
                "message": "Chỉ chấp nhận file Excel định dạng .xlsx.",
            }],
            file_name=file_name,
            status_code=400,
        )

    content = await excel_file.read()
    if not content:
        return hien_thi_nhap_excel(
            request=request,
            batch=batch,
            errors=[{"row": "—", "message": "File tải lên đang trống."}],
            file_name=file_name,
            status_code=400,
        )
    if len(content) > 15 * 1024 * 1024:
        return hien_thi_nhap_excel(
            request=request,
            batch=batch,
            errors=[{
                "row": "—",
                "message": "File lớn quá 15 MB. Hãy kiểm tra lại file điều tra.",
            }],
            file_name=file_name,
            status_code=400,
        )

    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception:
        return hien_thi_nhap_excel(
            request=request,
            batch=batch,
            errors=[{
                "row": "—",
                "message": "Không đọc được file Excel. File có thể bị hỏng hoặc sai định dạng.",
            }],
            file_name=file_name,
            status_code=400,
        )

    if "DIEU_TRA" not in workbook.sheetnames:
        return hien_thi_nhap_excel(
            request=request,
            batch=batch,
            errors=[{
                "row": "—",
                "message": "Không tìm thấy sheet DIEU_TRA. Không đổi tên sheet của file xuất từ phần mềm.",
            }],
            file_name=file_name,
            status_code=400,
        )

    worksheet = workbook["DIEU_TRA"]
    header_row = 6
    header_map: dict[str, int] = {}
    for column_index in range(1, worksheet.max_column + 1):
        header = gia_tri_excel_dang_chu(
            worksheet.cell(row=header_row, column=column_index).value
        )
        if header:
            header_map[header] = column_index

    missing_headers = [
        header
        for header in EXCEL_IMPORT_REQUIRED_HEADERS
        if header not in header_map
    ]
    if missing_headers:
        return hien_thi_nhap_excel(
            request=request,
            batch=batch,
            errors=[{
                "row": "6",
                "message": (
                    "Thiếu hoặc đã đổi tên các cột: "
                    + ", ".join(missing_headers)
                ),
            }],
            file_name=file_name,
            status_code=400,
        )

    def cell_value(row_number: int, header: str) -> Any:
        return worksheet.cell(
            row=row_number,
            column=header_map[header],
        ).value

    residency_reverse = {
        khoa_so_sanh_excel(label): code
        for code, label in RESIDENCY_STATUS_LABELS.items()
    }
    learning_reverse = {
        khoa_so_sanh_excel(label): code
        for code, label in LEARNING_STATUS_LABELS.items()
    }
    form_status_reverse = {
        khoa_so_sanh_excel(label): code
        for code, label in FORM_STATUS_LABELS.items()
    }

    survey_forms = {
        item.id: item
        for item in db.scalars(
            select(SurveyForm)
            .options(selectinload(SurveyForm.household))
            .where(SurveyForm.survey_batch_id == batch.id)
        ).all()
    }
    household_ids = {item.household_id for item in survey_forms.values()}
    households = {
        item.id: item
        for item in db.scalars(
            select(Household).where(Household.id.in_(household_ids))
        ).all()
    } if household_ids else {}

    existing_people = {
        item.id: item
        for item in db.scalars(
            select(SurveyPerson).where(
                SurveyPerson.household_id.in_(household_ids)
            )
        ).all()
    } if household_ids else {}

    errors: list[dict[str, Any]] = []
    actions: list[dict[str, Any]] = []
    household_values: dict[int, tuple[str, ...]] = {}
    seen_existing_person_ids: dict[int, int] = {}
    seen_personal_ids: dict[str, int] = {}
    seen_ministry_codes: dict[str, int] = {}
    seen_name_birth: dict[tuple[int, str, date], int] = {}

    for row_number in range(header_row + 1, worksheet.max_row + 1):
        form_id = so_nguyen_excel(cell_value(row_number, "__survey_form_id"))
        household_id = so_nguyen_excel(cell_value(row_number, "__household_id"))
        school_year_id = so_nguyen_excel(cell_value(row_number, "__school_year_id"))
        visible_person_code = chuan_hoa_ma(
            gia_tri_excel_dang_chu(cell_value(row_number, "Mã đối tượng"))
        )
        technical_person_id = so_nguyen_excel(
            cell_value(row_number, "__survey_person_id")
        )

        visible_values = [
            cell_value(row_number, header)
            for header in EXCEL_IMPORT_REQUIRED_HEADERS[:33]
        ]
        if not any(value not in (None, "") for value in visible_values) and not form_id:
            continue

        row_errors: list[str] = []
        survey_form = survey_forms.get(form_id) if form_id else None
        household = households.get(household_id) if household_id else None

        if survey_form is None:
            row_errors.append("ID phiếu kỹ thuật không thuộc đợt điều tra này.")
        if household is None:
            row_errors.append("ID hộ kỹ thuật không thuộc đợt điều tra này.")
        if survey_form is not None and survey_form.household_id != household_id:
            row_errors.append("ID phiếu và ID hộ không khớp nhau.")
        if school_year_id != batch.school_year_id:
            row_errors.append("ID năm học không đúng với đợt điều tra.")

        if survey_form is not None:
            if chuan_hoa_ma(gia_tri_excel_dang_chu(
                cell_value(row_number, "Số phiếu")
            )) != chuan_hoa_ma(survey_form.form_number):
                row_errors.append("Số phiếu đã bị thay đổi.")
        if household is not None:
            if chuan_hoa_ma(gia_tri_excel_dang_chu(
                cell_value(row_number, "Mã hộ")
            )) != chuan_hoa_ma(household.code):
                row_errors.append("Mã hộ đã bị thay đổi.")

        hamlet_name = chuan_hoa_van_ban(gia_tri_excel_dang_chu(
            cell_value(row_number, "Thôn/xóm/khối/bản")
        ))
        head_name = chuan_hoa_van_ban(gia_tri_excel_dang_chu(
            cell_value(row_number, "Chủ hộ")
        ))
        address = chuan_hoa_van_ban(gia_tri_excel_dang_chu(
            cell_value(row_number, "Địa chỉ hộ")
        ))
        phone = chuan_hoa_so_dien_thoai(gia_tri_excel_dang_chu(
            cell_value(row_number, "Điện thoại chủ hộ")
        )) or ""
        form_status_key = khoa_so_sanh_excel(
            cell_value(row_number, "Trạng thái phiếu")
        )
        form_status = form_status_reverse.get(form_status_key)
        if not head_name:
            row_errors.append("Chủ hộ không được để trống.")
        if not address:
            row_errors.append("Địa chỉ hộ không được để trống.")
        if form_status is None:
            row_errors.append("Trạng thái phiếu không hợp lệ.")

        household_tuple = (
            hamlet_name,
            head_name,
            address,
            phone,
            form_status or "",
        )
        if household_id is not None:
            previous_household_tuple = household_values.get(household_id)
            if previous_household_tuple is None:
                household_values[household_id] = household_tuple
            elif previous_household_tuple != household_tuple:
                row_errors.append(
                    "Thông tin hộ không thống nhất giữa các dòng của cùng một hộ. "
                    "Hãy sao chép cùng thông tin chủ hộ, địa chỉ, điện thoại và trạng thái phiếu cho mọi dòng."
                )

        full_name = chuan_hoa_van_ban(gia_tri_excel_dang_chu(
            cell_value(row_number, "Họ và tên")
        ))
        birth_date = ngay_excel(cell_value(row_number, "Ngày sinh"))
        gender = chuan_hoa_van_ban(gia_tri_excel_dang_chu(
            cell_value(row_number, "Giới tính")
        ))
        ethnic_group = chuan_hoa_van_ban(gia_tri_excel_dang_chu(
            cell_value(row_number, "Dân tộc")
        ))
        relationship = chuan_hoa_van_ban(gia_tri_excel_dang_chu(
            cell_value(row_number, "Quan hệ với chủ hộ")
        ))
        personal_id = chuan_hoa_ma(gia_tri_excel_dang_chu(
            cell_value(row_number, "Số định danh cá nhân")
        ))
        ministry_code = chuan_hoa_ma(gia_tri_excel_dang_chu(
            cell_value(row_number, "Mã Bộ GD&ĐT")
        ))
        residency_key = khoa_so_sanh_excel(
            cell_value(row_number, "Trạng thái cư trú")
        )
        residency_status = residency_reverse.get(residency_key)
        permanent_address = chuan_hoa_van_ban(gia_tri_excel_dang_chu(
            cell_value(row_number, "Địa chỉ thường trú")
        ))
        current_address = chuan_hoa_van_ban(gia_tri_excel_dang_chu(
            cell_value(row_number, "Chỗ ở hiện nay")
        ))
        special_circumstances = gia_tri_excel_dang_chu(
            cell_value(row_number, "Hoàn cảnh đặc biệt")
        ).strip()
        person_notes = gia_tri_excel_dang_chu(
            cell_value(row_number, "Ghi chú đối tượng")
        ).strip()

        has_person_data = any((
            visible_person_code,
            full_name,
            birth_date,
            gender,
            ethnic_group,
            relationship,
            personal_id,
            ministry_code,
        ))

        person: SurveyPerson | None = None
        is_new_person = False
        # Xóa Mã đối tượng trên dòng sao chép nghĩa là thêm người mới,
        # dù ID kỹ thuật ẩn vẫn còn do Excel sao chép cả dòng.
        if visible_person_code:
            if technical_person_id is None:
                row_errors.append("Đối tượng đã có mã nhưng thiếu ID kỹ thuật.")
            else:
                person = existing_people.get(technical_person_id)
                if person is None or person.household_id != household_id:
                    row_errors.append("Đối tượng kỹ thuật không thuộc đúng hộ.")
                elif not person.is_active:
                    row_errors.append("Đối tượng này đã ngừng theo dõi.")
                elif chuan_hoa_ma(person.code) != visible_person_code:
                    row_errors.append("Mã đối tượng không khớp ID kỹ thuật.")
                elif technical_person_id in seen_existing_person_ids:
                    row_errors.append(
                        f"Đối tượng bị lặp lại; đã xuất hiện ở dòng {seen_existing_person_ids[technical_person_id]}."
                    )
                else:
                    seen_existing_person_ids[technical_person_id] = row_number
        elif has_person_data:
            is_new_person = True
            technical_person_id = None

        if person is not None or is_new_person:
            if not full_name:
                row_errors.append("Họ và tên đối tượng không được để trống.")
            if birth_date is None:
                row_errors.append("Ngày sinh không hợp lệ hoặc đang để trống.")
            elif birth_date > date.today():
                row_errors.append("Ngày sinh không được lớn hơn ngày hiện tại.")
            if gender not in {"Nam", "Nữ", "Khác"}:
                row_errors.append("Giới tính phải là Nam, Nữ hoặc Khác.")
            if not ethnic_group:
                row_errors.append("Dân tộc không được để trống.")
            if not relationship:
                row_errors.append("Quan hệ với chủ hộ không được để trống.")
            if residency_status is None:
                row_errors.append("Trạng thái cư trú không hợp lệ.")

            if personal_id:
                previous_row = seen_personal_ids.get(personal_id)
                if previous_row is not None:
                    row_errors.append(
                        f"Số định danh bị trùng với dòng {previous_row}."
                    )
                else:
                    seen_personal_ids[personal_id] = row_number
                duplicate = db.scalar(
                    select(SurveyPerson).where(
                        SurveyPerson.personal_id == personal_id,
                        SurveyPerson.is_active.is_(True),
                        SurveyPerson.id != (person.id if person else -1),
                    )
                )
                if duplicate is not None:
                    row_errors.append("Số định danh đã được dùng cho đối tượng khác.")

            if ministry_code:
                previous_row = seen_ministry_codes.get(ministry_code)
                if previous_row is not None:
                    row_errors.append(
                        f"Mã Bộ GD&ĐT bị trùng với dòng {previous_row}."
                    )
                else:
                    seen_ministry_codes[ministry_code] = row_number
                duplicate = db.scalar(
                    select(SurveyPerson).where(
                        SurveyPerson.ministry_student_code == ministry_code,
                        SurveyPerson.is_active.is_(True),
                        SurveyPerson.id != (person.id if person else -1),
                    )
                )
                if duplicate is not None:
                    row_errors.append("Mã Bộ GD&ĐT đã được dùng cho đối tượng khác.")

            if household_id is not None and full_name and birth_date:
                name_birth_key = (
                    household_id,
                    full_name.lower(),
                    birth_date,
                )
                previous_row = seen_name_birth.get(name_birth_key)
                if previous_row is not None:
                    row_errors.append(
                        f"Họ tên và ngày sinh bị trùng với dòng {previous_row}."
                    )
                else:
                    seen_name_birth[name_birth_key] = row_number
                duplicate = db.scalar(
                    select(SurveyPerson).where(
                        SurveyPerson.household_id == household_id,
                        func.lower(SurveyPerson.full_name) == full_name.lower(),
                        SurveyPerson.date_of_birth == birth_date,
                        SurveyPerson.is_active.is_(True),
                        SurveyPerson.id != (person.id if person else -1),
                    )
                )
                if duplicate is not None:
                    row_errors.append("Người cùng họ tên và ngày sinh đã tồn tại trong hộ.")

        learning_key = khoa_so_sanh_excel(
            cell_value(row_number, "Tình trạng học tập")
        )
        learning_status = learning_reverse.get(learning_key) if learning_key else None
        school_name = chuan_hoa_van_ban(gia_tri_excel_dang_chu(
            cell_value(row_number, "Tên trường")
        ))
        class_name = chuan_hoa_van_ban(gia_tri_excel_dang_chu(
            cell_value(row_number, "Tên lớp")
        ))
        year_notes = gia_tri_excel_dang_chu(
            cell_value(row_number, "Ghi chú năm học")
        ).strip()

        boolean_headers = (
            ("completed_preschool_5", "Hoàn thành CT GDMN 5 tuổi"),
            ("attends_required_days", "Đi học đủ ngày theo quy định"),
            ("attends_regularly", "Đi học chuyên cần"),
            ("prepared_vietnamese", "Được chuẩn bị tiếng Việt"),
            ("weight_monitored", "Theo dõi biểu đồ cân nặng"),
            ("underweight", "Suy dinh dưỡng nhẹ cân"),
            ("height_monitored", "Theo dõi biểu đồ chiều cao"),
            ("stunted", "Suy dinh dưỡng thấp còi"),
        )
        boolean_values: dict[str, bool | None] = {}
        for field_name, header in boolean_headers:
            parsed_value, valid = co_khong_excel(cell_value(row_number, header))
            boolean_values[field_name] = parsed_value
            if not valid:
                row_errors.append(f"Cột {header} chỉ được nhập Có, Không hoặc để trống.")

        has_year_data = any((
            learning_key,
            school_name,
            class_name,
            special_circumstances,
            year_notes,
            *[value is not None for value in boolean_values.values()],
        ))
        if has_year_data and person is None and not is_new_person:
            row_errors.append("Có dữ liệu năm học nhưng dòng chưa có đối tượng.")
        if has_year_data and learning_status is None:
            row_errors.append("Tình trạng học tập không hợp lệ hoặc đang để trống.")
        if learning_status in {"DANG_HOC", "CHUYEN_DEN"} and not school_name:
            row_errors.append("Đối tượng đang học cần có tên trường.")

        if row_errors:
            errors.extend(
                {"row": row_number, "message": message}
                for message in row_errors
            )
            continue

        actions.append({
            "row": row_number,
            "survey_form": survey_form,
            "household": household,
            "household_values": household_tuple,
            "person": person,
            "is_new_person": is_new_person,
            "person_values": {
                "full_name": full_name,
                "date_of_birth": birth_date,
                "gender": gender,
                "ethnic_group": ethnic_group,
                "relationship_to_head": relationship,
                "personal_id": personal_id or None,
                "ministry_student_code": ministry_code or None,
                "permanent_address": permanent_address or None,
                "current_address": current_address or None,
                "residency_status": residency_status,
                "special_circumstances": special_circumstances or None,
                "notes": person_notes or None,
            },
            "has_person_data": person is not None or is_new_person,
            "has_year_data": has_year_data,
            "year_values": {
                "learning_status": learning_status,
                "school_name": school_name,
                "class_name": class_name,
                "special_circumstances": special_circumstances or None,
                "notes": year_notes or None,
                **boolean_values,
            },
            "technical_school_id": so_nguyen_excel(
                cell_value(row_number, "__school_id")
            ),
            "technical_class_id": so_nguyen_excel(
                cell_value(row_number, "__class_id")
            ),
        })

    if errors:
        db.rollback()
        return hien_thi_nhap_excel(
            request=request,
            batch=batch,
            errors=errors[:200],
            file_name=file_name,
            status_code=400,
        )

    if not actions:
        return hien_thi_nhap_excel(
            request=request,
            batch=batch,
            errors=[{
                "row": "—",
                "message": "Không tìm thấy dòng dữ liệu hợp lệ để cập nhật.",
            }],
            file_name=file_name,
            status_code=400,
        )

    backup_path = tao_ban_sao_truoc_khi_nhap_excel()
    updated_households: set[int] = set()
    updated_people = 0
    created_people = 0
    updated_year_records = 0
    created_year_records = 0

    try:
        for action in actions:
            survey_form = action["survey_form"]
            household = action["household"]
            hamlet_name, head_name, address, phone, form_status = action[
                "household_values"
            ]

            if household.id not in updated_households:
                household.hamlet_name = hamlet_name or None
                household.head_name = head_name
                household.address = address
                household.phone = phone or None
                household.is_active = True
                survey_form.head_name_snapshot = head_name
                survey_form.address_snapshot = address
                survey_form.hamlet_name_snapshot = hamlet_name or None
                survey_form.status = form_status
                updated_households.add(household.id)

            person = action["person"]
            if action["is_new_person"]:
                values = action["person_values"]
                student: Student | None = None
                inferred: list[Student] = []
                if values["ministry_student_code"]:
                    item = db.scalar(select(Student).where(
                        Student.code == values["ministry_student_code"]
                    ))
                    if item is not None:
                        inferred.append(item)
                if values["personal_id"]:
                    item = db.scalar(select(Student).where(
                        Student.personal_id == values["personal_id"]
                    ))
                    if item is not None and item not in inferred:
                        inferred.append(item)
                if len(inferred) == 1:
                    candidate = inferred[0]
                    existing_link = db.scalar(select(SurveyPerson).where(
                        SurveyPerson.student_id == candidate.id,
                        SurveyPerson.is_active.is_(True),
                    ))
                    if existing_link is None:
                        student = candidate

                person = SurveyPerson(
                    code=f"TMP-{uuid4().hex[:12]}",
                    household_id=household.id,
                    student_id=student.id if student is not None else None,
                    is_active=True,
                    **values,
                )
                db.add(person)
                db.flush()
                person.code = f"DT-{person.id:08d}"
                created_people += 1
                if survey_form.status == "CHUA_DIEU_TRA":
                    survey_form.status = "DANG_DIEU_TRA"
            elif action["has_person_data"]:
                for field_name, value in action["person_values"].items():
                    setattr(person, field_name, value)
                updated_people += 1

            if action["has_year_data"] and person is not None:
                year_values = action["year_values"]
                school: School | None = None
                classroom: Classroom | None = None
                school_name_key = khoa_so_sanh_excel(year_values["school_name"])
                class_name_key = khoa_so_sanh_excel(year_values["class_name"])

                technical_school_id = action["technical_school_id"]
                if technical_school_id:
                    candidate = db.get(School, technical_school_id)
                    if (
                        candidate is not None
                        and candidate.commune_id == batch.commune_id
                        and khoa_so_sanh_excel(candidate.name) == school_name_key
                    ):
                        school = candidate
                if school is None and school_name_key:
                    candidates = db.scalars(select(School).where(
                        School.commune_id == batch.commune_id,
                        School.is_active.is_(True),
                    )).all()
                    matches = [
                        item for item in candidates
                        if khoa_so_sanh_excel(item.name) == school_name_key
                    ]
                    if len(matches) == 1:
                        school = matches[0]

                technical_class_id = action["technical_class_id"]
                if school is not None and technical_class_id:
                    candidate = db.get(Classroom, technical_class_id)
                    if (
                        candidate is not None
                        and candidate.school_id == school.id
                        and candidate.school_year_id == batch.school_year_id
                        and khoa_so_sanh_excel(candidate.name) == class_name_key
                    ):
                        classroom = candidate
                if school is not None and classroom is None and class_name_key:
                    candidates = db.scalars(select(Classroom).where(
                        Classroom.school_id == school.id,
                        Classroom.school_year_id == batch.school_year_id,
                    )).all()
                    matches = [
                        item for item in candidates
                        if khoa_so_sanh_excel(item.name) == class_name_key
                    ]
                    if len(matches) == 1:
                        classroom = matches[0]

                record = db.scalar(select(SurveyPersonYearRecord).where(
                    SurveyPersonYearRecord.survey_person_id == person.id,
                    SurveyPersonYearRecord.school_year_id == batch.school_year_id,
                ))
                if record is None:
                    record = SurveyPersonYearRecord(
                        survey_form_id=survey_form.id,
                        survey_person_id=person.id,
                        school_year_id=batch.school_year_id,
                        learning_status=year_values["learning_status"],
                    )
                    db.add(record)
                    created_year_records += 1
                else:
                    updated_year_records += 1

                record.survey_form_id = survey_form.id
                record.school_id = school.id if school is not None else None
                record.class_id = classroom.id if classroom is not None else None
                record.school_name_reported = year_values["school_name"] or None
                record.class_name_reported = year_values["class_name"] or None
                record.learning_status = year_values["learning_status"]
                record.special_circumstances = year_values[
                    "special_circumstances"
                ]
                record.notes = year_values["notes"]
                for field_name in BOOLEAN_YEAR_FIELDS:
                    setattr(record, field_name, year_values[field_name])

        db.commit()
    except Exception as exc:
        db.rollback()
        return hien_thi_nhap_excel(
            request=request,
            batch=batch,
            errors=[{
                "row": "—",
                "message": (
                    "Không thể hoàn tất cập nhật. Dữ liệu chưa được thay đổi. "
                    f"Chi tiết kỹ thuật: {type(exc).__name__}."
                ),
            }],
            file_name=file_name,
            status_code=500,
        )

    return hien_thi_nhap_excel(
        request=request,
        batch=batch,
        file_name=file_name,
        summary={
            "updated_households": len(updated_households),
            "updated_people": updated_people,
            "created_people": created_people,
            "updated_year_records": updated_year_records,
            "created_year_records": created_year_records,
            "backup_path": str(backup_path),
        },
    )


@router.post("/{batch_id}/ho-dan/them")
def them_ho_dan(
    batch_id: int,
    request: Request,
    head_name: Annotated[str, Form()],
    address: Annotated[str, Form()],
    hamlet_name: Annotated[str, Form()] = "",
    phone: Annotated[str, Form()] = "",
    notes: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    batch = lay_dot_dieu_tra(db, batch_id)

    if batch is None:
        return RedirectResponse(
            url="/dieu-tra",
            status_code=303,
        )

    head_name = chuan_hoa_van_ban(head_name)
    address = chuan_hoa_van_ban(address)
    hamlet_name = chuan_hoa_van_ban(hamlet_name)
    phone_value = chuan_hoa_so_dien_thoai(phone)
    notes = notes.strip()

    errors: list[str] = []

    if batch.status == "DA_KET_THUC":
        errors.append(
            "Đợt điều tra đã kết thúc, không thể thêm hộ."
        )

    if not head_name:
        errors.append(
            "Họ và tên chủ hộ không được để trống."
        )

    if len(head_name) > 200:
        errors.append(
            "Họ và tên chủ hộ dài quá 200 ký tự."
        )

    if not address:
        errors.append(
            "Địa chỉ hộ gia đình không được để trống."
        )

    if len(hamlet_name) > 200:
        errors.append(
            "Tên thôn/xóm dài quá 200 ký tự."
        )

    if phone_value and len(phone_value) > 30:
        errors.append(
            "Số điện thoại dài quá 30 ký tự."
        )

    form_data: dict[str, Any] = {
        "head_name": head_name,
        "address": address,
        "hamlet_name": hamlet_name,
        "phone": phone,
        "notes": notes,
    }

    existing_household = db.scalar(
        select(Household).where(
            Household.commune_id == batch.commune_id,
            Household.head_name == head_name,
            Household.address == address,
        )
    )

    if existing_household is not None:
        existing_form = db.scalar(
            select(SurveyForm).where(
                SurveyForm.survey_batch_id == batch.id,
                SurveyForm.household_id == existing_household.id,
            )
        )

        if existing_form is not None:
            errors.append(
                "Hộ gia đình này đã có phiếu "
                "trong đợt điều tra hiện tại."
            )

    if errors:
        # Hiển thị lại danh sách trang đầu để người dùng sửa dữ liệu.
        survey_forms = db.scalars(
            select(SurveyForm)
            .options(
                selectinload(SurveyForm.household)
                .selectinload(Household.people),
                selectinload(SurveyForm.investigators),
                with_loader_criteria(
                    SurveyPerson,
                    SurveyPerson.is_active.is_(True),
                    include_aliases=True,
                ),
            )
            .where(
                SurveyForm.survey_batch_id == batch.id
            )
            .order_by(SurveyForm.id.asc())
            .limit(HOUSEHOLD_PAGE_SIZE)
        ).all()

        hamlet_names = db.scalars(
            select(Household.hamlet_name)
            .join(
                SurveyForm,
                SurveyForm.household_id == Household.id,
            )
            .where(
                SurveyForm.survey_batch_id == batch.id,
                Household.hamlet_name.is_not(None),
                Household.hamlet_name != "",
            )
            .distinct()
            .order_by(Household.hamlet_name.asc())
        ).all()

        total_records = db.scalar(
            select(func.count(SurveyForm.id)).where(
                SurveyForm.survey_batch_id == batch.id
            )
        )

        status_rows = db.execute(
            select(
                SurveyForm.status,
                func.count(SurveyForm.id),
            )
            .where(SurveyForm.survey_batch_id == batch.id)
            .group_by(SurveyForm.status)
        ).all()
        status_counts = {code: 0 for code in FORM_STATUS_LABELS}
        for status_code, count_value in status_rows:
            status_counts[status_code] = int(count_value or 0)
        batch_total_records = sum(status_counts.values())
        progress_percent = (
            round(
                status_counts.get("DA_HOAN_THANH", 0)
                * 100
                / batch_total_records
            )
            if batch_total_records
            else 0
        )

        return templates.TemplateResponse(
            request=request,
            name="surveys/households.html",
            context={
                "nguoi_dung": request.scope.get("auth_user"),
                "batch": batch,
                "survey_forms": survey_forms,
                "hamlet_names": hamlet_names,
                "q": "",
                "selected_hamlet": "",
                "selected_form_status": "",
                "status_counts": status_counts,
                "batch_total_records": batch_total_records,
                "progress_percent": progress_percent,
                "page": 1,
                "page_size": HOUSEHOLD_PAGE_SIZE,
                "total_records": int(total_records or 0),
                "total_pages": max(
                    1,
                    ceil(
                        int(total_records or 0)
                        / HOUSEHOLD_PAGE_SIZE
                    ),
                ),
                "page_links": [],
                "previous_url": None,
                "next_url": None,
                "form_status_labels": FORM_STATUS_LABELS,
                "thong_bao_loi": " ".join(errors),
                "form_data": form_data,
            },
            status_code=400,
        )

    message_status = "household_created"

    try:
        if existing_household is None:
            household = Household(
                code=tao_ma_ho(
                    db=db,
                    commune=batch.commune,
                ),
                commune_id=batch.commune_id,
                head_name=head_name,
                hamlet_name=hamlet_name or None,
                address=address,
                phone=phone_value,
                notes=notes or None,
                is_active=True,
            )

            db.add(household)
            db.flush()
        else:
            household = existing_household
            message_status = "household_attached"

            if hamlet_name:
                household.hamlet_name = hamlet_name

            if phone_value:
                household.phone = phone_value

            if notes:
                household.notes = notes

            household.is_active = True

        survey_form = SurveyForm(
            survey_batch_id=batch.id,
            household_id=household.id,
            form_number=tao_so_phieu(
                db=db,
                batch=batch,
            ),
            head_name_snapshot=household.head_name,
            address_snapshot=household.address,
            hamlet_name_snapshot=household.hamlet_name,
            survey_date=None,
            status="CHUA_DIEU_TRA",
            village_head_name=None,
            household_representative_name=None,
            household_confirmed_at=None,
            commune_confirmed_at=None,
            notes=None,
        )

        db.add(survey_form)
        db.commit()

    except IntegrityError:
        db.rollback()

        return RedirectResponse(
            url=f"/dieu-tra/{batch.id}/ho-dan",
            status_code=303,
        )

    return RedirectResponse(
        url=(
            f"/dieu-tra/{batch.id}/ho-dan?"
            + urlencode({"status": message_status})
        ),
        status_code=303,
    )


# =========================================================
# QUẢN LÝ ĐỐI TƯỢNG TRONG HỘ GIA ĐÌNH
# =========================================================

@router.get(
    "/{batch_id}/ho-dan/{household_id}/doi-tuong",
    response_class=HTMLResponse,
)
def danh_sach_doi_tuong(
    batch_id: int,
    household_id: int,
    request: Request,
    q: str = "",
    residency_status: str = "",
    page: int = 1,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )

    if survey_form is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    return hien_thi_trang_doi_tuong(
        request=request,
        db=db,
        survey_form=survey_form,
        q=q,
        residency_status=residency_status,
        page=page,
        status=status,
    )


@router.post(
    "/{batch_id}/ho-dan/{household_id}/doi-tuong/them"
)
def them_doi_tuong(
    batch_id: int,
    household_id: int,
    request: Request,
    full_name: Annotated[str, Form()],
    student_id: Annotated[str, Form()] = "",
    date_of_birth: Annotated[str, Form()] = "",
    gender: Annotated[str, Form()] = "",
    ethnic_group: Annotated[str, Form()] = "",
    relationship_to_head: Annotated[str, Form()] = "",
    personal_id: Annotated[str, Form()] = "",
    ministry_student_code: Annotated[str, Form()] = "",
    permanent_address: Annotated[str, Form()] = "",
    current_address: Annotated[str, Form()] = "",
    residency_status: Annotated[str, Form()] = "THUONG_TRU",
    special_circumstances: Annotated[str, Form()] = "",
    notes: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )

    if survey_form is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    full_name = chuan_hoa_van_ban(full_name)
    gender = chuan_hoa_van_ban(gender)
    ethnic_group = chuan_hoa_van_ban(ethnic_group)
    relationship_to_head = chuan_hoa_van_ban(
        relationship_to_head
    )
    personal_id = chuan_hoa_ma(personal_id)
    ministry_student_code = chuan_hoa_ma(
        ministry_student_code
    )
    permanent_address = chuan_hoa_van_ban(
        permanent_address
    )
    current_address = chuan_hoa_van_ban(
        current_address
    )
    special_circumstances = special_circumstances.strip()
    notes = notes.strip()

    form_data: dict[str, Any] = {
        "student_id": int(student_id) if student_id.isdigit() else None,
        "full_name": full_name,
        "date_of_birth": date_of_birth,
        "gender": gender,
        "ethnic_group": ethnic_group,
        "relationship_to_head": relationship_to_head,
        "personal_id": personal_id,
        "ministry_student_code": ministry_student_code,
        "permanent_address": permanent_address,
        "current_address": current_address,
        "residency_status": residency_status,
        "special_circumstances": special_circumstances,
        "notes": notes,
    }

    errors: list[str] = []

    if survey_form.survey_batch.status == "DA_KET_THUC":
        errors.append(
            "Đợt điều tra đã kết thúc, không thể thêm đối tượng."
        )

    selected_student: Student | None = None

    if student_id.strip():
        if not student_id.isdigit():
            errors.append("Học sinh liên kết không hợp lệ.")
        else:
            selected_student = db.get(Student, int(student_id))

            if (
                selected_student is None
                or not selected_student.is_active
            ):
                errors.append(
                    "Không tìm thấy học sinh được chọn."
                )

    # Có thể tự tìm hồ sơ học sinh theo mã Bộ hoặc số định danh.
    inferred_students: list[Student] = []

    if ministry_student_code:
        item = db.scalar(
            select(Student).where(
                Student.code == ministry_student_code
            )
        )
        if item is not None:
            inferred_students.append(item)

    if personal_id:
        item = db.scalar(
            select(Student).where(
                Student.personal_id == personal_id
            )
        )
        if item is not None and item not in inferred_students:
            inferred_students.append(item)

    if selected_student is None and len(inferred_students) == 1:
        selected_student = inferred_students[0]
        form_data["student_id"] = selected_student.id

    if len(inferred_students) > 1:
        errors.append(
            "Mã Bộ và số định danh đang trỏ đến hai học sinh khác nhau."
        )

    if (
        selected_student is not None
        and any(
            item.id != selected_student.id
            for item in inferred_students
        )
    ):
        errors.append(
            "Học sinh được chọn không khớp với mã Bộ hoặc số định danh đã nhập."
        )

    if selected_student is not None:
        existing_link = db.scalar(
            select(SurveyPerson).where(
                SurveyPerson.student_id == selected_student.id,
                SurveyPerson.is_active.is_(True),
            )
        )

        if existing_link is not None:
            errors.append(
                "Học sinh này đã được liên kết với một đối tượng điều tra khác."
            )

        if not full_name:
            full_name = selected_student.full_name
            form_data["full_name"] = full_name

        if not date_of_birth and selected_student.date_of_birth:
            date_of_birth = selected_student.date_of_birth.isoformat()
            form_data["date_of_birth"] = date_of_birth

        if not gender and selected_student.gender:
            gender = selected_student.gender
            form_data["gender"] = gender

        if not ethnic_group and selected_student.ethnic_group:
            ethnic_group = selected_student.ethnic_group
            form_data["ethnic_group"] = ethnic_group

        if not personal_id and selected_student.personal_id:
            personal_id = selected_student.personal_id
            form_data["personal_id"] = personal_id

        if not ministry_student_code and selected_student.code:
            ministry_student_code = selected_student.code
            form_data["ministry_student_code"] = (
                ministry_student_code
            )

        if not permanent_address and selected_student.permanent_address:
            permanent_address = selected_student.permanent_address
            form_data["permanent_address"] = permanent_address

        if not current_address and selected_student.current_address:
            current_address = selected_student.current_address
            form_data["current_address"] = current_address

    if not full_name:
        errors.append("Họ và tên đối tượng không được để trống.")

    if len(full_name) > 200:
        errors.append("Họ và tên dài quá 200 ký tự.")

    parsed_birth_date, birth_error = chuyen_ngay(
        date_of_birth,
        "Ngày sinh",
        bat_buoc=True,
    )

    if birth_error:
        errors.append(birth_error)

    if parsed_birth_date and parsed_birth_date > date.today():
        errors.append(
            "Ngày sinh không được lớn hơn ngày hiện tại."
        )

    if not gender:
        errors.append("Giới tính không được để trống.")
    elif gender not in {"Nam", "Nữ", "Khác"}:
        errors.append("Giới tính không hợp lệ.")

    if not ethnic_group:
        errors.append("Dân tộc không được để trống.")

    if not relationship_to_head:
        errors.append(
            "Quan hệ với chủ hộ không được để trống."
        )

    if residency_status not in RESIDENCY_STATUS_LABELS:
        errors.append("Trạng thái cư trú không hợp lệ.")

    if len(personal_id) > 50:
        errors.append("Số định danh dài quá 50 ký tự.")

    if len(ministry_student_code) > 50:
        errors.append("Mã Bộ GD&ĐT dài quá 50 ký tự.")

    if personal_id:
        duplicate_personal_id = db.scalar(
            select(SurveyPerson).where(
                SurveyPerson.personal_id == personal_id,
                SurveyPerson.is_active.is_(True),
            )
        )
        if duplicate_personal_id is not None:
            errors.append(
                "Số định danh cá nhân đã được sử dụng cho đối tượng khác."
            )

    if ministry_student_code:
        duplicate_ministry_code = db.scalar(
            select(SurveyPerson).where(
                SurveyPerson.ministry_student_code
                == ministry_student_code,
                SurveyPerson.is_active.is_(True),
            )
        )
        if duplicate_ministry_code is not None:
            errors.append(
                "Mã Bộ GD&ĐT đã được sử dụng cho đối tượng khác."
            )

    duplicate_name = db.scalar(
        select(SurveyPerson).where(
            SurveyPerson.household_id == household_id,
            func.lower(SurveyPerson.full_name) == full_name.lower(),
            SurveyPerson.date_of_birth == parsed_birth_date,
            SurveyPerson.is_active.is_(True),
        )
    )

    if duplicate_name is not None:
        errors.append(
            "Đối tượng có cùng họ tên và ngày sinh đã tồn tại trong hộ."
        )

    if errors:
        return hien_thi_trang_doi_tuong(
            request=request,
            db=db,
            survey_form=survey_form,
            thong_bao_loi=" ".join(errors),
            form_data=form_data,
            status_code=400,
        )

    person = SurveyPerson(
        # Mã tạm duy nhất để lấy ID an toàn trước khi sinh mã chính thức.
        code=f"TMP-{uuid4().hex[:12]}",
        household_id=household_id,
        student_id=(
            selected_student.id
            if selected_student is not None
            else None
        ),
        full_name=full_name,
        date_of_birth=parsed_birth_date,
        gender=gender or None,
        ethnic_group=ethnic_group or None,
        relationship_to_head=relationship_to_head or None,
        personal_id=personal_id or None,
        ministry_student_code=ministry_student_code or None,
        permanent_address=permanent_address or None,
        current_address=current_address or None,
        residency_status=residency_status,
        special_circumstances=special_circumstances or None,
        notes=notes or None,
        is_active=True,
    )

    db.add(person)
    db.flush()

    # Dùng khóa chính đã được CSDL cấp để tránh trùng mã khi nhiều
    # người cùng nhập dữ liệu. Mã đã cấp không bị dùng lại vì bản ghi
    # được xóa mềm bằng is_active=False.
    person.code = f"DT-{person.id:08d}"

    if survey_form.status == "CHUA_DIEU_TRA":
        survey_form.status = "DANG_DIEU_TRA"

    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return hien_thi_trang_doi_tuong(
            request=request,
            db=db,
            survey_form=survey_form,
            thong_bao_loi=(
                "Không thể lưu đối tượng. Mã hoặc thông tin "
                "định danh có thể đã tồn tại."
            ),
            form_data=form_data,
            status_code=400,
        )

    return RedirectResponse(
        url=(
            f"/dieu-tra/{batch_id}/ho-dan/"
            f"{household_id}/doi-tuong?status=person_created"
        ),
        status_code=303,
    )


# =========================================================
# CHỈNH SỬA VÀ NGỪNG THEO DÕI ĐỐI TƯỢNG
# =========================================================

def khoa_so_sanh_khong_dau(value: str | None) -> str:
    """Chuẩn hóa văn bản để so sánh không phân biệt dấu và hoa thường."""

    normalized = unicodedata.normalize(
        "NFD",
        value or "",
    )

    without_marks = "".join(
        character
        for character in normalized
        if unicodedata.category(character) != "Mn"
    )

    return " ".join(
        without_marks.lower().strip().split()
    )


def la_quan_he_chu_ho(value: str | None) -> bool:
    """Nhận diện các cách ghi phổ biến của quan hệ chủ hộ."""

    return khoa_so_sanh_khong_dau(value) in {
        "chu ho",
        "chủ hộ",
    }


def la_doi_tuong_chu_ho(
    person: SurveyPerson,
    household: Household,
) -> bool:
    """Xác định đối tượng đang đại diện vai trò chủ hộ."""

    if la_quan_he_chu_ho(person.relationship_to_head):
        return True

    return (
        khoa_so_sanh_khong_dau(person.full_name)
        == khoa_so_sanh_khong_dau(household.head_name)
    )


def lay_doi_tuong_sua(
    db: Session,
    household_id: int,
    person_id: int,
) -> SurveyPerson | None:
    """Lấy một đối tượng đang hoạt động thuộc đúng hộ để sửa."""

    return db.scalar(
        select(SurveyPerson)
        .options(selectinload(SurveyPerson.student))
        .where(
            SurveyPerson.id == person_id,
            SurveyPerson.household_id == household_id,
            SurveyPerson.is_active.is_(True),
        )
    )


def hien_thi_form_sua_doi_tuong(
    *,
    request: Request,
    db: Session,
    survey_form: SurveyForm,
    person: SurveyPerson,
    thong_bao_loi: str | None = None,
    form_data: dict[str, Any] | None = None,
    status_code: int = 200,
):
    """Hiển thị biểu mẫu sửa đối tượng và giữ lại dữ liệu khi có lỗi."""

    students = db.scalars(
        select(Student)
        .where(Student.is_active.is_(True))
        .order_by(
            Student.full_name.asc(),
            Student.date_of_birth.asc(),
            Student.code.asc(),
        )
    ).all()

    used_student_ids = {
        item
        for item in db.scalars(
            select(SurveyPerson.student_id).where(
                SurveyPerson.student_id.is_not(None),
                SurveyPerson.is_active.is_(True),
                SurveyPerson.id != person.id,
            )
        ).all()
        if item is not None
    }

    if form_data is None:
        form_data = {
            "student_id": person.student_id,
            "full_name": person.full_name,
            "date_of_birth": (
                person.date_of_birth.isoformat()
                if person.date_of_birth
                else ""
            ),
            "gender": person.gender or "",
            "ethnic_group": person.ethnic_group or "",
            "relationship_to_head": (
                person.relationship_to_head or ""
            ),
            "personal_id": person.personal_id or "",
            "ministry_student_code": (
                person.ministry_student_code or ""
            ),
            "permanent_address": person.permanent_address or "",
            "current_address": person.current_address or "",
            "residency_status": (
                person.residency_status or "CHUA_XAC_DINH"
            ),
            "special_circumstances": (
                person.special_circumstances or ""
            ),
            "notes": person.notes or "",
        }

    return templates.TemplateResponse(
        request=request,
        name="surveys/person_edit.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "survey_form": survey_form,
            "batch": survey_form.survey_batch,
            "household": survey_form.household,
            "person": person,
            "students": students,
            "used_student_ids": used_student_ids,
            "residency_status_labels": RESIDENCY_STATUS_LABELS,
            "form_data": form_data,
            "thong_bao_loi": thong_bao_loi,
        },
        status_code=status_code,
    )


@router.get(
    "/{batch_id}/ho-dan/{household_id}/doi-tuong/{person_id}/sua",
    response_class=HTMLResponse,
)
def sua_doi_tuong_form(
    batch_id: int,
    household_id: int,
    person_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )

    if survey_form is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    person = lay_doi_tuong_sua(
        db=db,
        household_id=household_id,
        person_id=person_id,
    )

    if person is None:
        return RedirectResponse(
            url=(
                f"/dieu-tra/{batch_id}/ho-dan/"
                f"{household_id}/doi-tuong"
            ),
            status_code=303,
        )

    return hien_thi_form_sua_doi_tuong(
        request=request,
        db=db,
        survey_form=survey_form,
        person=person,
    )


@router.post(
    "/{batch_id}/ho-dan/{household_id}/doi-tuong/{person_id}/sua"
)
def cap_nhat_doi_tuong(
    batch_id: int,
    household_id: int,
    person_id: int,
    request: Request,
    full_name: Annotated[str, Form()],
    student_id: Annotated[str, Form()] = "",
    date_of_birth: Annotated[str, Form()] = "",
    gender: Annotated[str, Form()] = "",
    ethnic_group: Annotated[str, Form()] = "",
    relationship_to_head: Annotated[str, Form()] = "",
    personal_id: Annotated[str, Form()] = "",
    ministry_student_code: Annotated[str, Form()] = "",
    permanent_address: Annotated[str, Form()] = "",
    current_address: Annotated[str, Form()] = "",
    residency_status: Annotated[str, Form()] = "THUONG_TRU",
    special_circumstances: Annotated[str, Form()] = "",
    notes: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )

    if survey_form is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    person = lay_doi_tuong_sua(
        db=db,
        household_id=household_id,
        person_id=person_id,
    )

    if person is None:
        return RedirectResponse(
            url=(
                f"/dieu-tra/{batch_id}/ho-dan/"
                f"{household_id}/doi-tuong"
            ),
            status_code=303,
        )

    full_name = chuan_hoa_van_ban(full_name)
    gender = chuan_hoa_van_ban(gender)
    ethnic_group = chuan_hoa_van_ban(ethnic_group)
    relationship_to_head = chuan_hoa_van_ban(
        relationship_to_head
    )
    personal_id = chuan_hoa_ma(personal_id)
    ministry_student_code = chuan_hoa_ma(
        ministry_student_code
    )
    permanent_address = chuan_hoa_van_ban(
        permanent_address
    )
    current_address = chuan_hoa_van_ban(
        current_address
    )
    special_circumstances = special_circumstances.strip()
    notes = notes.strip()

    form_data: dict[str, Any] = {
        "student_id": int(student_id) if student_id.isdigit() else None,
        "full_name": full_name,
        "date_of_birth": date_of_birth,
        "gender": gender,
        "ethnic_group": ethnic_group,
        "relationship_to_head": relationship_to_head,
        "personal_id": personal_id,
        "ministry_student_code": ministry_student_code,
        "permanent_address": permanent_address,
        "current_address": current_address,
        "residency_status": residency_status,
        "special_circumstances": special_circumstances,
        "notes": notes,
    }

    errors: list[str] = []

    if survey_form.survey_batch.status == "DA_KET_THUC":
        errors.append(
            "Đợt điều tra đã kết thúc, không thể sửa đối tượng."
        )

    selected_student: Student | None = None

    if student_id.strip():
        if not student_id.isdigit():
            errors.append("Học sinh liên kết không hợp lệ.")
        else:
            selected_student = db.get(Student, int(student_id))

            if (
                selected_student is None
                or not selected_student.is_active
            ):
                errors.append(
                    "Không tìm thấy học sinh được chọn."
                )

    if selected_student is not None:
        existing_link = db.scalar(
            select(SurveyPerson).where(
                SurveyPerson.student_id == selected_student.id,
                SurveyPerson.id != person.id,
                SurveyPerson.is_active.is_(True),
            )
        )

        if existing_link is not None:
            errors.append(
                "Học sinh này đã được liên kết với đối tượng khác."
            )

        if ministry_student_code and (
            selected_student.code != ministry_student_code
        ):
            errors.append(
                "Mã Bộ GD&ĐT không khớp với học sinh đã chọn."
            )

        if (
            personal_id
            and selected_student.personal_id
            and selected_student.personal_id != personal_id
        ):
            errors.append(
                "Số định danh không khớp với học sinh đã chọn."
            )

        if not full_name:
            full_name = selected_student.full_name
            form_data["full_name"] = full_name

        if not date_of_birth and selected_student.date_of_birth:
            date_of_birth = selected_student.date_of_birth.isoformat()
            form_data["date_of_birth"] = date_of_birth

        if not gender and selected_student.gender:
            gender = selected_student.gender
            form_data["gender"] = gender

        if not ethnic_group and selected_student.ethnic_group:
            ethnic_group = selected_student.ethnic_group
            form_data["ethnic_group"] = ethnic_group

        if not personal_id and selected_student.personal_id:
            personal_id = selected_student.personal_id
            form_data["personal_id"] = personal_id

        if not ministry_student_code and selected_student.code:
            ministry_student_code = selected_student.code
            form_data["ministry_student_code"] = ministry_student_code

        if not permanent_address and selected_student.permanent_address:
            permanent_address = selected_student.permanent_address
            form_data["permanent_address"] = permanent_address

        if not current_address and selected_student.current_address:
            current_address = selected_student.current_address
            form_data["current_address"] = current_address

    if not full_name:
        errors.append("Họ và tên đối tượng không được để trống.")

    if len(full_name) > 200:
        errors.append("Họ và tên dài quá 200 ký tự.")

    parsed_birth_date, birth_error = chuyen_ngay(
        date_of_birth,
        "Ngày sinh",
        bat_buoc=True,
    )

    if birth_error:
        errors.append(birth_error)

    if parsed_birth_date and parsed_birth_date > date.today():
        errors.append(
            "Ngày sinh không được lớn hơn ngày hiện tại."
        )

    if not gender:
        errors.append("Giới tính không được để trống.")
    elif gender not in {"Nam", "Nữ", "Khác"}:
        errors.append("Giới tính không hợp lệ.")

    if not ethnic_group:
        errors.append("Dân tộc không được để trống.")

    if not relationship_to_head:
        errors.append(
            "Quan hệ với chủ hộ không được để trống."
        )

    if residency_status not in RESIDENCY_STATUS_LABELS:
        errors.append("Trạng thái cư trú không hợp lệ.")

    if len(personal_id) > 50:
        errors.append("Số định danh dài quá 50 ký tự.")

    if len(ministry_student_code) > 50:
        errors.append("Mã Bộ GD&ĐT dài quá 50 ký tự.")

    if personal_id:
        duplicate_personal_id = db.scalar(
            select(SurveyPerson).where(
                SurveyPerson.personal_id == personal_id,
                SurveyPerson.id != person.id,
                SurveyPerson.is_active.is_(True),
            )
        )

        if duplicate_personal_id is not None:
            errors.append(
                "Số định danh cá nhân đã được dùng cho đối tượng khác."
            )

    if ministry_student_code:
        duplicate_ministry_code = db.scalar(
            select(SurveyPerson).where(
                SurveyPerson.ministry_student_code
                == ministry_student_code,
                SurveyPerson.id != person.id,
                SurveyPerson.is_active.is_(True),
            )
        )

        if duplicate_ministry_code is not None:
            errors.append(
                "Mã Bộ GD&ĐT đã được dùng cho đối tượng khác."
            )

    duplicate_name = db.scalar(
        select(SurveyPerson).where(
            SurveyPerson.household_id == household_id,
            func.lower(SurveyPerson.full_name) == full_name.lower(),
            SurveyPerson.date_of_birth == parsed_birth_date,
            SurveyPerson.id != person.id,
            SurveyPerson.is_active.is_(True),
        )
    )

    if duplicate_name is not None:
        errors.append(
            "Đối tượng có cùng họ tên và ngày sinh đã tồn tại trong hộ."
        )

    other_people = db.scalars(
        select(SurveyPerson).where(
            SurveyPerson.household_id == household_id,
            SurveyPerson.id != person.id,
            SurveyPerson.is_active.is_(True),
        )
    ).all()

    replacement_head = next(
        (
            item
            for item in other_people
            if la_quan_he_chu_ho(item.relationship_to_head)
        ),
        None,
    )

    was_head = la_doi_tuong_chu_ho(
        person,
        survey_form.household,
    )
    will_be_head = la_quan_he_chu_ho(
        relationship_to_head
    )

    if will_be_head and replacement_head is not None:
        errors.append(
            "Trong hộ đã có một đối tượng được ghi là Chủ hộ."
        )

    if was_head and not will_be_head and replacement_head is None:
        errors.append(
            "Đối tượng này đang là chủ hộ. Hãy sửa một đối tượng khác "
            "thành quan hệ Chủ hộ trước."
        )

    if errors:
        return hien_thi_form_sua_doi_tuong(
            request=request,
            db=db,
            survey_form=survey_form,
            person=person,
            thong_bao_loi=" ".join(errors),
            form_data=form_data,
            status_code=400,
        )

    person.student_id = (
        selected_student.id
        if selected_student is not None
        else None
    )
    person.full_name = full_name
    person.date_of_birth = parsed_birth_date
    person.gender = gender or None
    person.ethnic_group = ethnic_group or None
    person.relationship_to_head = relationship_to_head or None
    person.personal_id = personal_id or None
    person.ministry_student_code = ministry_student_code or None
    person.permanent_address = permanent_address or None
    person.current_address = current_address or None
    person.residency_status = residency_status
    person.special_circumstances = special_circumstances or None
    person.notes = notes or None

    if will_be_head:
        survey_form.household.head_name = full_name
        survey_form.head_name_snapshot = full_name
    elif was_head and replacement_head is not None:
        survey_form.household.head_name = replacement_head.full_name
        survey_form.head_name_snapshot = replacement_head.full_name

    try:
        db.commit()
    except IntegrityError:
        db.rollback()

        return hien_thi_form_sua_doi_tuong(
            request=request,
            db=db,
            survey_form=survey_form,
            person=person,
            thong_bao_loi=(
                "Không thể cập nhật đối tượng. Thông tin định danh "
                "hoặc liên kết học sinh có thể đã tồn tại."
            ),
            form_data=form_data,
            status_code=400,
        )

    return RedirectResponse(
        url=(
            f"/dieu-tra/{batch_id}/ho-dan/"
            f"{household_id}/doi-tuong?status=person_updated"
        ),
        status_code=303,
    )


@router.post(
    "/{batch_id}/ho-dan/{household_id}/doi-tuong/{person_id}/ngung-theo-doi"
)
def ngung_theo_doi_doi_tuong(
    batch_id: int,
    household_id: int,
    person_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )

    if survey_form is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    person = lay_doi_tuong_sua(
        db=db,
        household_id=household_id,
        person_id=person_id,
    )

    if person is None:
        return RedirectResponse(
            url=(
                f"/dieu-tra/{batch_id}/ho-dan/"
                f"{household_id}/doi-tuong"
            ),
            status_code=303,
        )

    errors: list[str] = []

    if survey_form.survey_batch.status == "DA_KET_THUC":
        errors.append(
            "Đợt điều tra đã kết thúc, không thể ngừng theo dõi."
        )

    if la_doi_tuong_chu_ho(person, survey_form.household):
        other_people = db.scalars(
            select(SurveyPerson).where(
                SurveyPerson.household_id == household_id,
                SurveyPerson.id != person.id,
                SurveyPerson.is_active.is_(True),
            )
        ).all()

        replacement_head = next(
            (
                item
                for item in other_people
                if la_quan_he_chu_ho(item.relationship_to_head)
            ),
            None,
        )

        if replacement_head is None:
            errors.append(
                "Không thể ngừng theo dõi chủ hộ. Hãy sửa một đối tượng "
                "khác thành quan hệ Chủ hộ trước."
            )
        else:
            survey_form.household.head_name = replacement_head.full_name
            survey_form.head_name_snapshot = replacement_head.full_name

    if errors:
        return hien_thi_trang_doi_tuong(
            request=request,
            db=db,
            survey_form=survey_form,
            thong_bao_loi=" ".join(errors),
            status_code=400,
        )

    person.is_active = False

    try:
        db.commit()
    except IntegrityError:
        db.rollback()

        return hien_thi_trang_doi_tuong(
            request=request,
            db=db,
            survey_form=survey_form,
            thong_bao_loi=(
                "Không thể ngừng theo dõi đối tượng ở thời điểm này."
            ),
            status_code=400,
        )

    return RedirectResponse(
        url=(
            f"/dieu-tra/{batch_id}/ho-dan/"
            f"{household_id}/doi-tuong?status=person_deactivated"
        ),
        status_code=303,
    )


# =========================================================
# THEO DÕI THÔNG TIN TỪNG NĂM HỌC
# =========================================================

def lay_doi_tuong_trong_ho(
    db: Session,
    household_id: int,
    person_id: int,
) -> SurveyPerson | None:
    """Lấy một đối tượng đang hoạt động thuộc đúng hộ."""

    return db.scalar(
        select(SurveyPerson)
        .options(selectinload(SurveyPerson.student))
        .where(
            SurveyPerson.id == person_id,
            SurveyPerson.household_id == household_id,
            SurveyPerson.is_active.is_(True),
        )
    )


def chuyen_id_tuy_chon(value: str) -> int | None:
    """Chuyển ID từ biểu mẫu; giá trị trống trả về None."""

    value = value.strip()
    if not value:
        return None
    if not value.isdigit():
        return None
    return int(value)


def tao_du_lieu_lich_su_nam_hoc(
    db: Session,
    person_id: int,
) -> list[dict[str, Any]]:
    """Ghép bản ghi năm học với tên năm, trường và lớp để hiển thị."""

    records = db.scalars(
        select(SurveyPersonYearRecord)
        .where(
            SurveyPersonYearRecord.survey_person_id == person_id
        )
        .order_by(
            SurveyPersonYearRecord.school_year_id.desc(),
            SurveyPersonYearRecord.id.desc(),
        )
    ).all()

    year_ids = {item.school_year_id for item in records}
    school_ids = {
        item.school_id for item in records if item.school_id is not None
    }
    class_ids = {
        item.class_id for item in records if item.class_id is not None
    }

    years = {
        item.id: item
        for item in db.scalars(
            select(SchoolYear).where(SchoolYear.id.in_(year_ids))
        ).all()
    } if year_ids else {}

    schools = {
        item.id: item
        for item in db.scalars(
            select(School).where(School.id.in_(school_ids))
        ).all()
    } if school_ids else {}

    classrooms = {
        item.id: item
        for item in db.scalars(
            select(Classroom).where(Classroom.id.in_(class_ids))
        ).all()
    } if class_ids else {}

    result: list[dict[str, Any]] = []
    for record in records:
        year = years.get(record.school_year_id)
        school = schools.get(record.school_id)
        classroom = classrooms.get(record.class_id)
        result.append(
            {
                "record": record,
                "school_year_code": (
                    year.code if year is not None else str(record.school_year_id)
                ),
                "school_name": (
                    record.school_name_reported
                    or (school.name if school is not None else "—")
                ),
                "class_name": (
                    record.class_name_reported
                    or (classroom.name if classroom is not None else "—")
                ),
                "learning_status_label": LEARNING_STATUS_LABELS.get(
                    record.learning_status,
                    record.learning_status,
                ),
            }
        )

    return result


@router.get(
    "/{batch_id}/ho-dan/{household_id}/doi-tuong/nam-hoc",
    response_class=HTMLResponse,
)
def tong_quan_theo_doi_nam_hoc(
    request: Request,
    batch_id: int,
    household_id: int,
    db: Session = Depends(get_db),
):
    """Danh sách tất cả đối tượng trong hộ và số hồ sơ năm học."""

    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )
    if survey_form is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    people = db.scalars(
        select(SurveyPerson)
        .where(
            SurveyPerson.household_id == household_id,
            SurveyPerson.is_active.is_(True),
        )
        .order_by(
            SurveyPerson.date_of_birth.asc(),
            SurveyPerson.full_name.asc(),
        )
    ).all()

    person_ids = [item.id for item in people]
    records = db.scalars(
        select(SurveyPersonYearRecord)
        .where(
            SurveyPersonYearRecord.survey_person_id.in_(person_ids)
        )
        .order_by(
            SurveyPersonYearRecord.school_year_id.desc(),
            SurveyPersonYearRecord.id.desc(),
        )
    ).all() if person_ids else []

    year_ids = {item.school_year_id for item in records}
    years = {
        item.id: item
        for item in db.scalars(
            select(SchoolYear).where(SchoolYear.id.in_(year_ids))
        ).all()
    } if year_ids else {}

    grouped: dict[int, list[SurveyPersonYearRecord]] = {}
    for record in records:
        grouped.setdefault(record.survey_person_id, []).append(record)

    rows: list[dict[str, Any]] = []
    for person in people:
        person_records = grouped.get(person.id, [])
        latest = person_records[0] if person_records else None
        latest_year = years.get(latest.school_year_id) if latest else None
        rows.append(
            {
                "person": person,
                "record_count": len(person_records),
                "latest_year_code": latest_year.code if latest_year else "—",
                "latest_status": (
                    LEARNING_STATUS_LABELS.get(
                        latest.learning_status,
                        latest.learning_status,
                    )
                    if latest
                    else "Chưa có dữ liệu"
                ),
            }
        )

    return templates.TemplateResponse(
        request=request,
        name="surveys/year_overview.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "survey_form": survey_form,
            "batch": survey_form.survey_batch,
            "household": survey_form.household,
            "rows": rows,
        },
    )


@router.get(
    "/{batch_id}/ho-dan/{household_id}/doi-tuong/{person_id}/nam-hoc",
    response_class=HTMLResponse,
)
def theo_doi_nam_hoc_doi_tuong(
    request: Request,
    batch_id: int,
    household_id: int,
    person_id: int,
    school_year_id: int | None = None,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )
    person = lay_doi_tuong_trong_ho(
        db=db,
        household_id=household_id,
        person_id=person_id,
    )

    if survey_form is None or person is None:
        return RedirectResponse(
            url=(
                f"/dieu-tra/{batch_id}/ho-dan/"
                f"{household_id}/doi-tuong"
            ),
            status_code=303,
        )

    school_years = db.scalars(
        select(SchoolYear)
        .order_by(SchoolYear.code.desc(), SchoolYear.id.desc())
    ).all()
    valid_year_ids = {item.id for item in school_years}

    if school_year_id not in valid_year_ids:
        school_year_id = survey_form.survey_batch.school_year_id
    if school_year_id not in valid_year_ids and school_years:
        school_year_id = school_years[0].id

    selected_record = None
    if school_year_id is not None:
        selected_record = db.scalar(
            select(SurveyPersonYearRecord).where(
                SurveyPersonYearRecord.survey_person_id == person_id,
                SurveyPersonYearRecord.school_year_id == school_year_id,
            )
        )

    schools = db.scalars(
        select(School)
        .options(selectinload(School.commune))
        .where(School.is_active.is_(True))
        .order_by(School.name.asc())
    ).all()

    selected_school_id = (
        selected_record.school_id if selected_record else None
    )
    selected_class_id = (
        selected_record.class_id if selected_record else None
    )

    history_rows = tao_du_lieu_lich_su_nam_hoc(
        db=db,
        person_id=person_id,
    )

    return templates.TemplateResponse(
        request=request,
        name="surveys/year_records.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "survey_form": survey_form,
            "batch": survey_form.survey_batch,
            "household": survey_form.household,
            "person": person,
            "school_years": school_years,
            "schools": schools,
            "selected_school_year_id": school_year_id,
            "selected_record": selected_record,
            "selected_school_id": selected_school_id,
            "selected_class_id": selected_class_id,
            "history_rows": history_rows,
            "learning_status_labels": LEARNING_STATUS_LABELS,
            "thong_bao": STATUS_MESSAGES.get(status),
            "thong_bao_loi": None,
        },
    )


@router.post(
    "/{batch_id}/ho-dan/{household_id}/doi-tuong/{person_id}/nam-hoc/luu"
)
def luu_theo_doi_nam_hoc(
    request: Request,
    batch_id: int,
    household_id: int,
    person_id: int,
    school_year_id: Annotated[int, Form()],
    learning_status: Annotated[str, Form()] = "DANG_HOC",
    school_id: Annotated[str, Form()] = "",
    class_id: Annotated[str, Form()] = "",
    school_name_reported: Annotated[str, Form()] = "",
    class_name_reported: Annotated[str, Form()] = "",
    completed_preschool_5: Annotated[bool, Form()] = False,
    attends_required_days: Annotated[bool, Form()] = False,
    attends_regularly: Annotated[bool, Form()] = False,
    prepared_vietnamese: Annotated[bool, Form()] = False,
    weight_monitored: Annotated[bool, Form()] = False,
    underweight: Annotated[bool, Form()] = False,
    height_monitored: Annotated[bool, Form()] = False,
    stunted: Annotated[bool, Form()] = False,
    special_circumstances: Annotated[str, Form()] = "",
    notes: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    survey_form = lay_phieu_ho(
        db=db,
        batch_id=batch_id,
        household_id=household_id,
    )
    person = lay_doi_tuong_trong_ho(
        db=db,
        household_id=household_id,
        person_id=person_id,
    )

    if survey_form is None or person is None:
        return RedirectResponse(
            url=(
                f"/dieu-tra/{batch_id}/ho-dan/"
                f"{household_id}/doi-tuong"
            ),
            status_code=303,
        )

    errors: list[str] = []

    if survey_form.survey_batch.status == "DA_KET_THUC":
        errors.append(
            "Đợt điều tra đã kết thúc, không thể cập nhật năm học."
        )

    school_year = db.get(SchoolYear, school_year_id)
    if school_year is None:
        errors.append("Năm học không hợp lệ.")

    if learning_status not in LEARNING_STATUS_LABELS:
        errors.append("Tình trạng học tập không hợp lệ.")

    parsed_school_id = chuyen_id_tuy_chon(school_id)
    parsed_class_id = chuyen_id_tuy_chon(class_id)

    selected_school = (
        db.get(School, parsed_school_id)
        if parsed_school_id is not None
        else None
    )
    selected_class = (
        db.get(Classroom, parsed_class_id)
        if parsed_class_id is not None
        else None
    )

    if school_id.strip() and parsed_school_id is None:
        errors.append("Trường được chọn không hợp lệ.")
    elif parsed_school_id is not None and selected_school is None:
        errors.append("Không tìm thấy trường được chọn.")

    if class_id.strip() and parsed_class_id is None:
        errors.append("Lớp được chọn không hợp lệ.")
    elif parsed_class_id is not None and selected_class is None:
        errors.append("Không tìm thấy lớp được chọn.")

    if selected_class is not None:
        if selected_school is None:
            errors.append("Cần chọn trường trước khi chọn lớp.")
        elif selected_class.school_id != selected_school.id:
            errors.append("Lớp không thuộc trường đã chọn.")
        elif selected_class.school_year_id != school_year_id:
            errors.append("Lớp không thuộc năm học đã chọn.")

    school_name_reported = chuan_hoa_van_ban(
        school_name_reported
    )
    class_name_reported = chuan_hoa_van_ban(
        class_name_reported
    )
    special_circumstances = special_circumstances.strip()
    notes = notes.strip()

    if selected_school is not None:
        school_name_reported = selected_school.name
    if selected_class is not None:
        class_name_reported = selected_class.name

    statuses_requiring_school = {
        "DANG_HOC",
        "CHUYEN_DEN",
        "TAM_NGHI",
    }
    if (
        learning_status in statuses_requiring_school
        and selected_school is None
        and not school_name_reported
    ):
        errors.append(
            "Đối tượng đang học cần chọn trường hoặc nhập tên trường ngoài hệ thống."
        )

    if len(school_name_reported) > 300:
        errors.append("Tên trường dài quá 300 ký tự.")
    if len(class_name_reported) > 200:
        errors.append("Tên lớp dài quá 200 ký tự.")

    existing_record = db.scalar(
        select(SurveyPersonYearRecord).where(
            SurveyPersonYearRecord.survey_person_id == person_id,
            SurveyPersonYearRecord.school_year_id == school_year_id,
        )
    )

    if errors:
        school_years = db.scalars(
            select(SchoolYear)
            .order_by(SchoolYear.code.desc(), SchoolYear.id.desc())
        ).all()
        schools = db.scalars(
            select(School)
            .options(selectinload(School.commune))
            .where(School.is_active.is_(True))
            .order_by(School.name.asc())
        ).all()

        form_record = existing_record or SurveyPersonYearRecord(
            survey_form_id=survey_form.id,
            survey_person_id=person.id,
            school_year_id=school_year_id,
            learning_status=learning_status,
        )
        form_record.school_id = parsed_school_id
        form_record.class_id = parsed_class_id
        form_record.school_name_reported = school_name_reported or None
        form_record.class_name_reported = class_name_reported or None
        form_record.learning_status = learning_status
        form_record.completed_preschool_5 = completed_preschool_5
        form_record.attends_required_days = attends_required_days
        form_record.attends_regularly = attends_regularly
        form_record.prepared_vietnamese = prepared_vietnamese
        form_record.weight_monitored = weight_monitored
        form_record.underweight = underweight
        form_record.height_monitored = height_monitored
        form_record.stunted = stunted
        form_record.special_circumstances = special_circumstances or None
        form_record.notes = notes or None

        return templates.TemplateResponse(
            request=request,
            name="surveys/year_records.html",
            context={
                "nguoi_dung": request.scope.get("auth_user"),
                "survey_form": survey_form,
                "batch": survey_form.survey_batch,
                "household": survey_form.household,
                "person": person,
                "school_years": school_years,
                "schools": schools,
                "selected_school_year_id": school_year_id,
                "selected_record": form_record,
                "selected_school_id": parsed_school_id,
                "selected_class_id": parsed_class_id,
                "history_rows": tao_du_lieu_lich_su_nam_hoc(
                    db=db,
                    person_id=person_id,
                ),
                "learning_status_labels": LEARNING_STATUS_LABELS,
                "thong_bao": None,
                "thong_bao_loi": " ".join(errors),
            },
            status_code=400,
        )

    record = existing_record
    if record is None:
        record = SurveyPersonYearRecord(
            survey_form_id=survey_form.id,
            survey_person_id=person.id,
            school_year_id=school_year_id,
            learning_status=learning_status,
        )
        db.add(record)

    record.survey_form_id = survey_form.id
    record.school_id = parsed_school_id
    record.class_id = parsed_class_id
    record.school_name_reported = school_name_reported or None
    record.class_name_reported = class_name_reported or None
    record.learning_status = learning_status
    record.completed_preschool_5 = completed_preschool_5
    record.special_circumstances = special_circumstances or None
    record.attends_required_days = attends_required_days
    record.attends_regularly = attends_regularly
    record.prepared_vietnamese = prepared_vietnamese
    record.weight_monitored = weight_monitored
    record.underweight = underweight
    record.height_monitored = height_monitored
    record.stunted = stunted
    record.notes = notes or None

    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return RedirectResponse(
            url=(
                f"/dieu-tra/{batch_id}/ho-dan/{household_id}/"
                f"doi-tuong/{person_id}/nam-hoc?school_year_id="
                f"{school_year_id}"
            ),
            status_code=303,
        )

    return RedirectResponse(
        url=(
            f"/dieu-tra/{batch_id}/ho-dan/{household_id}/"
            f"doi-tuong/{person_id}/nam-hoc?school_year_id="
            f"{school_year_id}&status=year_record_saved"
        ),
        status_code=303,
    )

