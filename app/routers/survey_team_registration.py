from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any
import re
import unicodedata

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.database import get_db
from app.permissions import (
    COMMUNE_ROLE_CODE,
    SCHOOL_ROLE_CODE,
    TEACHER_ROLE_CODE,
    normalize_role_code,
)

APP_DIR = Path(__file__).resolve().parent.parent
templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

router = APIRouter(
    prefix="/dieu-tra/phan-cong-to-dieu-tra",
    tags=["Tổ điều tra phổ cập"],
)

LEVEL_LABELS = {
    "MN": "Mầm non",
    "TH": "Tiểu học",
    "THCS": "THCS",
    "KHAC": "Chưa xác định",
}

SUBMISSION_LABELS = {
    "DRAFT": "Bản nháp",
    "SENT": "Đã gửi xã/phường",
}

SCHEMA_SQL = [
    """
    CREATE TABLE IF NOT EXISTS survey_investigation_participants (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        survey_batch_id INTEGER NOT NULL,
        school_id INTEGER NOT NULL,
        user_id INTEGER NOT NULL,
        level_code VARCHAR(20) NOT NULL,
        created_at DATETIME NOT NULL,
        updated_at DATETIME NOT NULL,
        UNIQUE(survey_batch_id, user_id),
        FOREIGN KEY(survey_batch_id) REFERENCES survey_batches(id),
        FOREIGN KEY(school_id) REFERENCES schools(id),
        FOREIGN KEY(user_id) REFERENCES users(id)
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS survey_participant_submissions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        survey_batch_id INTEGER NOT NULL,
        school_id INTEGER NOT NULL,
        status VARCHAR(20) NOT NULL DEFAULT 'DRAFT',
        sent_at DATETIME NULL,
        sent_by_user_id INTEGER NULL,
        updated_at DATETIME NOT NULL,
        UNIQUE(survey_batch_id, school_id),
        FOREIGN KEY(survey_batch_id) REFERENCES survey_batches(id),
        FOREIGN KEY(school_id) REFERENCES schools(id),
        FOREIGN KEY(sent_by_user_id) REFERENCES users(id)
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS survey_team_registration_logs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        survey_batch_id INTEGER NOT NULL,
        school_id INTEGER NOT NULL,
        action VARCHAR(30) NOT NULL,
        actor_user_id INTEGER NULL,
        actor_name_snapshot VARCHAR(200) NULL,
        participant_count INTEGER NOT NULL DEFAULT 0,
        notes TEXT NULL,
        created_at DATETIME NOT NULL
    )
    """,
    """
    CREATE INDEX IF NOT EXISTS
        ix_survey_investigation_participants_batch_school
    ON survey_investigation_participants(survey_batch_id, school_id)
    """,
    """
    CREATE INDEX IF NOT EXISTS
        ix_survey_participant_submissions_batch
    ON survey_participant_submissions(survey_batch_id, status)
    """,
]


def _user(request: Request) -> dict[str, Any]:
    return dict(request.scope.get("auth_user") or {})


def _role(request: Request) -> str:
    return normalize_role_code(_user(request).get("role_code"))


def _forbidden() -> RedirectResponse:
    return RedirectResponse(url="/?status=forbidden", status_code=303)


def _ensure_schema(db: Session) -> None:
    for statement in SCHEMA_SQL:
        db.execute(text(statement))


def _normalize(value: Any) -> str:
    raw = str(value or "").strip().lower()
    raw = unicodedata.normalize("NFD", raw)
    raw = "".join(
        char for char in raw
        if unicodedata.category(char) != "Mn"
    )
    return re.sub(r"\s+", " ", raw)


def _table_exists(db: Session, table_name: str) -> bool:
    value = db.execute(
        text(
            """
            SELECT 1
            FROM sqlite_master
            WHERE type = 'table'
              AND name = :name
            LIMIT 1
            """
        ),
        {"name": table_name},
    ).scalar()
    return value is not None


def _school_level(
    db: Session,
    *,
    school_id: int,
    school_name: str,
    school_year_id: int,
) -> str:
    if _table_exists(db, "school_network_year_data"):
        value = db.execute(
            text(
                """
                SELECT level_code
                FROM school_network_year_data
                WHERE school_id = :school_id
                  AND school_year_id = :school_year_id
                  AND level_code IN ('TH', 'THCS')
                ORDER BY id DESC
                LIMIT 1
                """
            ),
            {
                "school_id": int(school_id),
                "school_year_id": int(school_year_id),
            },
        ).scalar()
        if value in {"TH", "THCS"}:
            return str(value)

    key = _normalize(school_name)
    if (
        "mam non" in key
        or "mau giao" in key
        or re.search(r"(^|\s)mn(\s|$)", key)
    ):
        return "MN"
    if "trung hoc co so" in key or "thcs" in key:
        return "THCS"
    if (
        "tieu hoc" in key
        or (
            re.search(r"(^|\s)th(\s|$)", key)
            and "thpt" not in key
            and "thcs" not in key
        )
    ):
        return "TH"
    return "KHAC"


def _school_row(db: Session, school_id: int) -> dict[str, Any] | None:
    row = db.execute(
        text(
            """
            SELECT
                s.id, s.code, s.name, s.commune_id,
                c.name AS commune_name
            FROM schools AS s
            JOIN communes AS c ON c.id = s.commune_id
            WHERE s.id = :school_id
              AND s.is_active = 1
            LIMIT 1
            """
        ),
        {"school_id": int(school_id)},
    ).mappings().first()
    return dict(row) if row else None


def _batch_row(db: Session, batch_id: int) -> dict[str, Any] | None:
    row = db.execute(
        text(
            """
            SELECT
                sb.id, sb.commune_id, sb.school_year_id,
                sb.name, sb.status,
                sy.code AS school_year_code,
                c.name AS commune_name
            FROM survey_batches AS sb
            JOIN school_years AS sy ON sy.id = sb.school_year_id
            JOIN communes AS c ON c.id = sb.commune_id
            WHERE sb.id = :batch_id
            LIMIT 1
            """
        ),
        {"batch_id": int(batch_id)},
    ).mappings().first()
    return dict(row) if row else None


def _available_batches(
    db: Session,
    *,
    commune_id: int,
) -> list[dict[str, Any]]:
    rows = db.execute(
        text(
            """
            SELECT
                sb.id, sb.name, sb.status,
                sy.code AS school_year_code
            FROM survey_batches AS sb
            JOIN school_years AS sy ON sy.id = sb.school_year_id
            WHERE sb.commune_id = :commune_id
            ORDER BY
                sb.school_year_id DESC,
                sb.created_at DESC,
                sb.id DESC
            """
        ),
        {"commune_id": int(commune_id)},
    ).mappings().all()
    return [dict(row) for row in rows]


def _selected_batch(
    db: Session,
    *,
    commune_id: int,
    batch_id: int | None,
) -> dict[str, Any] | None:
    batches = _available_batches(db, commune_id=commune_id)
    if not batches:
        return None
    if batch_id is not None:
        for item in batches:
            if int(item["id"]) == int(batch_id):
                return _batch_row(db, int(item["id"]))
    return _batch_row(db, int(batches[0]["id"]))


def _teacher_rows(
    db: Session,
    *,
    school_id: int,
) -> list[dict[str, Any]]:
    rows = db.execute(
        text(
            """
            SELECT u.id, u.username, u.full_name
            FROM users AS u
            JOIN roles AS r ON r.id = u.role_id
            WHERE u.school_id = :school_id
              AND u.is_active = 1
              AND UPPER(r.code) = :role_code
            ORDER BY u.full_name COLLATE NOCASE, u.id
            """
        ),
        {
            "school_id": int(school_id),
            "role_code": TEACHER_ROLE_CODE,
        },
    ).mappings().all()
    return [dict(row) for row in rows]



# =========================================================
# BAI_13B_10_V1_3_STAFF_TO_TEACHER_ACCOUNTS_START
# Lấy GV đang làm việc từ hồ sơ Đội ngũ để đối chiếu tài khoản đăng nhập.
# =========================================================

def _staff_teacher_rows_for_batch(
    db: Session,
    *,
    school_id: int,
    school_year_id: int,
) -> list[dict[str, Any]]:
    if (
        not _table_exists(db, "staff_members")
        or not _table_exists(db, "staff_year_records")
        or not _table_exists(db, "school_years")
    ):
        return []

    target_code = db.execute(
        text(
            """
            SELECT code
            FROM school_years
            WHERE id = :school_year_id
            LIMIT 1
            """
        ),
        {"school_year_id": int(school_year_id)},
    ).scalar()

    target_match = re.search(r"(\d{4})", str(target_code or ""))
    target_start = int(target_match.group(1)) if target_match else 9999

    rows = db.execute(
        text(
            """
            SELECT
                syr.id AS record_id,
                sm.id AS staff_member_id,
                sm.ministry_staff_code,
                sm.code AS internal_staff_code,
                sm.full_name,
                sy.code AS source_year_code,
                CAST(SUBSTR(sy.code, 1, 4) AS INTEGER) AS source_year_start
            FROM staff_year_records AS syr
            JOIN staff_members AS sm
              ON sm.id = syr.staff_member_id
            JOIN school_years AS sy
              ON sy.id = syr.school_year_id
            WHERE syr.school_id = :school_id
              AND syr.is_active = 1
              AND sm.is_active = 1
              AND UPPER(syr.position_group) = 'GIAO_VIEN'
              AND UPPER(syr.status_code) = 'DANG_LAM_VIEC'
              AND CAST(SUBSTR(sy.code, 1, 4) AS INTEGER) <= :target_start
            ORDER BY
                sm.id,
                source_year_start DESC,
                syr.id DESC
            """
        ),
        {
            "school_id": int(school_id),
            "target_start": int(target_start),
        },
    ).mappings().all()

    # Mỗi người chỉ lấy bản ghi năm gần nhất, tránh trùng qua các năm học.
    result: list[dict[str, Any]] = []
    seen: set[int] = set()
    for row in rows:
        item = dict(row)
        staff_member_id = int(item["staff_member_id"])
        if staff_member_id in seen:
            continue
        seen.add(staff_member_id)
        result.append(item)
    return result


def _normalized_person_name(value: Any) -> str:
    return _normalize(value)


# === BAI_13B_10_V1_3_1_DUPLICATE_NAME_FIX ===
def _teacher_account_matches(
    db: Session,
    *,
    school_id: int,
    staff_row: dict[str, Any],
) -> tuple[str, dict[str, Any] | None]:
    ministry_code = str(
        staff_row.get("ministry_staff_code")
        or staff_row.get("internal_staff_code")
        or ""
    ).strip()

    # V1.3.1:
    # Khi hồ sơ Đội ngũ đã có MÃ CÁN BỘ/GV thì mã là định danh ưu tiên.
    # Không được ghép theo họ tên nếu username gv.<mã> chưa tồn tại,
    # vì một trường có thể có 2 giáo viên trùng họ tên.
    expected_username = (
        f"gv.{ministry_code}".lower()
        if ministry_code
        else ""
    )

    if expected_username:
        row = db.execute(
            text(
                """
                SELECT
                    u.id,
                    u.username,
                    u.full_name,
                    u.school_id,
                    u.is_active,
                    UPPER(r.code) AS role_code
                FROM users AS u
                JOIN roles AS r
                  ON r.id = u.role_id
                WHERE LOWER(u.username) = :username
                LIMIT 1
                """
            ),
            {"username": expected_username},
        ).mappings().first()

        if row is not None:
            account = dict(row)

            if (
                str(account.get("role_code") or "")
                == TEACHER_ROLE_CODE
                and int(account.get("school_id") or 0)
                == int(school_id)
            ):
                return (
                    "ACTIVE"
                    if int(account.get("is_active") or 0) == 1
                    else "LOCKED",
                    account,
                )

            # Username gv.<mã> đã tồn tại nhưng thuộc sai trường/vai trò:
            # phải dừng ở CONFLICT, tuyệt đối không chiếm dụng.
            return "CONFLICT", account

        # Có mã GV nhưng chưa có đúng username -> đây là tài khoản còn thiếu.
        # KHÔNG fallback theo họ tên.
        return "MISSING", None

    # Chỉ khi hồ sơ thật sự KHÔNG CÓ mã cán bộ/GV mới được phép
    # đối chiếu theo họ tên trong cùng trường và cùng vai trò.
    candidates = db.execute(
        text(
            """
            SELECT
                u.id,
                u.username,
                u.full_name,
                u.school_id,
                u.is_active,
                UPPER(r.code) AS role_code
            FROM users AS u
            JOIN roles AS r
              ON r.id = u.role_id
            WHERE u.school_id = :school_id
              AND UPPER(r.code) = :role_code
            """
        ),
        {
            "school_id": int(school_id),
            "role_code": TEACHER_ROLE_CODE,
        },
    ).mappings().all()

    target_name = _normalized_person_name(
        staff_row.get("full_name")
    )
    name_matches = [
        dict(row)
        for row in candidates
        if _normalized_person_name(row.get("full_name"))
        == target_name
    ]

    if len(name_matches) == 1:
        account = name_matches[0]
        return (
            "ACTIVE"
            if int(account.get("is_active") or 0) == 1
            else "LOCKED",
            account,
        )

    return "MISSING", None

def _safe_teacher_username(
    db: Session,
    *,
    code: str,
) -> str:
    clean_code = re.sub(
        r"[^a-zA-Z0-9_-]+",
        "",
        str(code or "").strip(),
    ).lower()
    if not clean_code:
        clean_code = "khongma"

    base = f"gv.{clean_code}"[:100]
    candidate = base
    index = 2

    while db.execute(
        text(
            """
            SELECT 1
            FROM users
            WHERE LOWER(username) = :username
            LIMIT 1
            """
        ),
        {"username": candidate.lower()},
    ).scalar() is not None:
        suffix = f".{index}"
        candidate = f"{base[:100-len(suffix)]}{suffix}"
        index += 1

    return candidate


def _temporary_teacher_password() -> str:
    import secrets

    upper = "ABCDEFGHJKLMNPQRSTUVWXYZ"
    lower = "abcdefghijkmnopqrstuvwxyz"
    digits = "23456789"
    symbols = "@#$%"
    required = [
        secrets.choice(upper),
        secrets.choice(lower),
        secrets.choice(digits),
        secrets.choice(symbols),
    ]
    alphabet = upper + lower + digits + symbols
    required.extend(secrets.choice(alphabet) for _ in range(8))
    secrets.SystemRandom().shuffle(required)
    return "".join(required)

# BAI_13B_10_V1_3_STAFF_TO_TEACHER_ACCOUNTS_END


def _submission(
    db: Session,
    *,
    batch_id: int,
    school_id: int,
) -> dict[str, Any] | None:
    row = db.execute(
        text(
            """
            SELECT
                sps.id, sps.status, sps.sent_at,
                sps.sent_by_user_id, sps.updated_at,
                u.full_name AS sent_by_name
            FROM survey_participant_submissions AS sps
            LEFT JOIN users AS u ON u.id = sps.sent_by_user_id
            WHERE sps.survey_batch_id = :batch_id
              AND sps.school_id = :school_id
            LIMIT 1
            """
        ),
        {
            "batch_id": int(batch_id),
            "school_id": int(school_id),
        },
    ).mappings().first()
    return dict(row) if row else None


def _participant_ids(
    db: Session,
    *,
    batch_id: int,
    school_id: int,
) -> list[int]:
    rows = db.execute(
        text(
            """
            SELECT user_id
            FROM survey_investigation_participants
            WHERE survey_batch_id = :batch_id
              AND school_id = :school_id
            ORDER BY id
            """
        ),
        {
            "batch_id": int(batch_id),
            "school_id": int(school_id),
        },
    ).scalars().all()
    return [int(value) for value in rows]


def _save_submission(
    db: Session,
    *,
    batch_id: int,
    school_id: int,
    status: str,
    actor_user_id: int | None,
) -> None:
    now_value = datetime.now().isoformat(sep=" ", timespec="seconds")
    existing = db.execute(
        text(
            """
            SELECT id
            FROM survey_participant_submissions
            WHERE survey_batch_id = :batch_id
              AND school_id = :school_id
            LIMIT 1
            """
        ),
        {
            "batch_id": int(batch_id),
            "school_id": int(school_id),
        },
    ).scalar()

    sent_at = now_value if status == "SENT" else None
    sent_by = actor_user_id if status == "SENT" else None

    params = {
        "batch_id": int(batch_id),
        "school_id": int(school_id),
        "status": status,
        "sent_at": sent_at,
        "sent_by_user_id": sent_by,
        "updated_at": now_value,
    }

    if existing is None:
        db.execute(
            text(
                """
                INSERT INTO survey_participant_submissions (
                    survey_batch_id, school_id, status,
                    sent_at, sent_by_user_id, updated_at
                ) VALUES (
                    :batch_id, :school_id, :status,
                    :sent_at, :sent_by_user_id, :updated_at
                )
                """
            ),
            params,
        )
    else:
        db.execute(
            text(
                """
                UPDATE survey_participant_submissions
                SET status = :status,
                    sent_at = :sent_at,
                    sent_by_user_id = :sent_by_user_id,
                    updated_at = :updated_at
                WHERE survey_batch_id = :batch_id
                  AND school_id = :school_id
                """
            ),
            params,
        )


def _log(
    db: Session,
    *,
    batch_id: int,
    school_id: int,
    request: Request,
    action: str,
    count: int,
    notes: str = "",
) -> None:
    user = _user(request)
    db.execute(
        text(
            """
            INSERT INTO survey_team_registration_logs (
                survey_batch_id, school_id, action,
                actor_user_id, actor_name_snapshot,
                participant_count, notes, created_at
            ) VALUES (
                :batch_id, :school_id, :action,
                :actor_user_id, :actor_name,
                :participant_count, :notes, :created_at
            )
            """
        ),
        {
            "batch_id": int(batch_id),
            "school_id": int(school_id),
            "action": action,
            "actor_user_id": user.get("id"),
            "actor_name": str(user.get("full_name") or "")[:200] or None,
            "participant_count": int(count),
            "notes": notes[:1000] or None,
            "created_at": datetime.now().isoformat(
                sep=" ", timespec="seconds"
            ),
        },
    )


def _validate_school_scope(
    db: Session,
    request: Request,
    batch_id: int,
) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    user = _user(request)
    school_id = user.get("school_id")
    if school_id is None:
        return None, None

    school = _school_row(db, int(school_id))
    batch = _batch_row(db, int(batch_id))
    if school is None or batch is None:
        return None, None
    if int(school["commune_id"]) != int(batch["commune_id"]):
        return None, None
    return school, batch


@router.get("/giao-vien-tham-gia", response_class=HTMLResponse)
def school_participants_page(
    request: Request,
    batch_id: int | None = None,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    if _role(request) != SCHOOL_ROLE_CODE:
        return _forbidden()

    _ensure_schema(db)
    user = _user(request)
    school_id = user.get("school_id")
    if school_id is None:
        return _forbidden()

    school = _school_row(db, int(school_id))
    if school is None:
        return _forbidden()

    batches = _available_batches(
        db,
        commune_id=int(school["commune_id"]),
    )
    batch = _selected_batch(
        db,
        commune_id=int(school["commune_id"]),
        batch_id=batch_id,
    )
    teachers = _teacher_rows(db, school_id=int(school_id))
    staff_teachers: list[dict[str, Any]] = []
    if batch is not None:
        staff_teachers = _staff_teacher_rows_for_batch(
            db,
            school_id=int(school_id),
            school_year_id=int(batch["school_year_id"]),
        )

    selected_ids: list[int] = []
    submission = None
    level_code = "KHAC"

    if batch is not None:
        selected_ids = _participant_ids(
            db,
            batch_id=int(batch["id"]),
            school_id=int(school_id),
        )
        submission = _submission(
            db,
            batch_id=int(batch["id"]),
            school_id=int(school_id),
        )
        level_code = _school_level(
            db,
            school_id=int(school_id),
            school_name=str(school["name"]),
            school_year_id=int(batch["school_year_id"]),
        )

    status_messages = {
        "draft_saved": "Đã lưu bản nháp danh sách giáo viên tham gia.",
        "sent": "Đã gửi danh sách giáo viên tham gia về xã/phường.",
        "withdrawn": "Đã thu hồi danh sách để trường tiếp tục chỉnh sửa.",
        "no_staff": "Không tìm thấy hồ sơ giáo viên đang làm việc trong dữ liệu Đội ngũ.",
        "no_teacher_role": "Hệ thống chưa có vai trò GIAO_VIEN để tạo tài khoản.",
    }

    return templates.TemplateResponse(
        request=request,
        name="survey_teams/school_participants.html",
        context={
            "nguoi_dung": user,
            "school": school,
            "batches": batches,
            "batch": batch,
            "teachers": teachers,
            "staff_teacher_count": len(staff_teachers),
            "selected_ids": set(selected_ids),
            "submission": submission,
            "level_code": level_code,
            "level_label": LEVEL_LABELS.get(level_code, level_code),
            "submission_labels": SUBMISSION_LABELS,
            "message": status_messages.get(status),
        },
    )



# =========================================================
# BAI_13B_10_V1_3_SYNC_ROUTE_START
# Trường chủ động tạo tài khoản GV còn thiếu từ hồ sơ Đội ngũ.
# POST trả về Excel mật khẩu tạm; không lưu mật khẩu rõ trong CSDL.
# =========================================================

@router.post("/giao-vien-tham-gia/dong-bo-tai-khoan")
async def sync_teacher_accounts_from_staff(
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != SCHOOL_ROLE_CODE:
        return _forbidden()

    _ensure_schema(db)

    form = await request.form()
    try:
        batch_id = int(form.get("batch_id") or 0)
    except (TypeError, ValueError):
        batch_id = 0

    school, batch = _validate_school_scope(
        db,
        request,
        batch_id,
    )
    if school is None or batch is None:
        return _forbidden()

    # Nếu trường đã gửi danh sách về xã thì không tạo thêm trong bước này,
    # tránh làm người dùng hiểu rằng danh sách đã gửi tự thay đổi.
    submission = _submission(
        db,
        batch_id=batch_id,
        school_id=int(school["id"]),
    )
    if submission is not None and submission.get("status") == "SENT":
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/giao-vien-tham-gia"
                f"?batch_id={batch_id}"
            ),
            status_code=303,
        )

    staff_rows = _staff_teacher_rows_for_batch(
        db,
        school_id=int(school["id"]),
        school_year_id=int(batch["school_year_id"]),
    )
    if not staff_rows:
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/giao-vien-tham-gia"
                f"?batch_id={batch_id}&status=no_staff"
            ),
            status_code=303,
        )

    teacher_role_id = db.execute(
        text(
            """
            SELECT id
            FROM roles
            WHERE UPPER(code) = :role_code
            LIMIT 1
            """
        ),
        {"role_code": TEACHER_ROLE_CODE},
    ).scalar()
    if teacher_role_id is None:
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/giao-vien-tham-gia"
                f"?batch_id={batch_id}&status=no_teacher_role"
            ),
            status_code=303,
        )

    from app.security import ma_hoa_mat_khau

    created: list[dict[str, Any]] = []
    reused: list[dict[str, Any]] = []
    locked: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []

    now_value = datetime.now().isoformat(
        sep=" ",
        timespec="seconds",
    )

    try:
        for staff_row in staff_rows:
            state, account = _teacher_account_matches(
                db,
                school_id=int(school["id"]),
                staff_row=staff_row,
            )

            base_info = {
                "full_name": str(staff_row.get("full_name") or ""),
                "staff_code": str(
                    staff_row.get("ministry_staff_code")
                    or staff_row.get("internal_staff_code")
                    or ""
                ),
                "source_year": str(staff_row.get("source_year_code") or ""),
            }

            if state == "ACTIVE" and account is not None:
                reused.append({
                    **base_info,
                    "username": str(account.get("username") or ""),
                    "note": "Tài khoản đang hoạt động - giữ nguyên mật khẩu.",
                })
                continue

            if state == "LOCKED" and account is not None:
                locked.append({
                    **base_info,
                    "username": str(account.get("username") or ""),
                    "note": (
                        "Tài khoản đang khóa - hệ thống KHÔNG tự mở khóa. "
                        "Trường cần kiểm tra trước khi sử dụng."
                    ),
                })
                continue

            if state == "CONFLICT" and account is not None:
                conflicts.append({
                    **base_info,
                    "username": str(account.get("username") or ""),
                    "note": (
                        "Tên đăng nhập theo mã giáo viên đã thuộc tài khoản "
                        "khác trường/vai trò - KHÔNG tự sửa."
                    ),
                })
                continue

            password = _temporary_teacher_password()
            username = _safe_teacher_username(
                db,
                code=base_info["staff_code"],
            )

            db.execute(
                text(
                    """
                    INSERT INTO users (
                        username,
                        password_hash,
                        full_name,
                        role_id,
                        commune_id,
                        school_id,
                        is_active,
                        created_at
                    ) VALUES (
                        :username,
                        :password_hash,
                        :full_name,
                        :role_id,
                        :commune_id,
                        :school_id,
                        1,
                        :created_at
                    )
                    """
                ),
                {
                    "username": username,
                    "password_hash": ma_hoa_mat_khau(password),
                    "full_name": base_info["full_name"],
                    "role_id": int(teacher_role_id),
                    "commune_id": int(school["commune_id"]),
                    "school_id": int(school["id"]),
                    "created_at": now_value,
                },
            )

            created.append({
                **base_info,
                "username": username,
                "temporary_password": password,
                "note": "Tài khoản mới - đổi mật khẩu sau khi bàn giao.",
            })

        # Chuẩn bị file bàn giao TRƯỚC commit để nếu tạo Excel lỗi
        # thì toàn bộ INSERT còn có thể rollback.
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        from fastapi.responses import StreamingResponse

        workbook = Workbook()
        ws = workbook.active
        ws.title = "Tai_khoan_moi"

        headers = [
            "STT",
            "Họ và tên",
            "Mã cán bộ/GV",
            "Tên đăng nhập",
            "Mật khẩu tạm",
            "Trường",
            "Cấp học",
            "Nguồn đội ngũ",
            "Ghi chú",
        ]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        level_code = _school_level(
            db,
            school_id=int(school["id"]),
            school_name=str(school["name"]),
            school_year_id=int(batch["school_year_id"]),
        )
        level_label = LEVEL_LABELS.get(level_code, level_code)

        for index, item in enumerate(created, start=1):
            ws.append([
                index,
                item["full_name"],
                item["staff_code"],
                item["username"],
                item["temporary_password"],
                school["name"],
                level_label,
                item["source_year"],
                item["note"],
            ])

        if not created:
            ws.append([
                "",
                "Không tạo tài khoản mới.",
                "",
                "",
                "",
                school["name"],
                level_label,
                "",
                "Các tài khoản phù hợp đã tồn tại hoặc cần kiểm tra xung đột/khóa.",
            ])

        def add_status_sheet(title: str, rows: list[dict[str, Any]]) -> None:
            sheet = workbook.create_sheet(title)
            sheet.append([
                "STT",
                "Họ và tên",
                "Mã cán bộ/GV",
                "Tên đăng nhập",
                "Nguồn đội ngũ",
                "Ghi chú",
            ])
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )
            for index, item in enumerate(rows, start=1):
                sheet.append([
                    index,
                    item["full_name"],
                    item["staff_code"],
                    item["username"],
                    item["source_year"],
                    item["note"],
                ])
            if not rows:
                sheet.append(["", "Không có", "", "", "", ""])

            for width, col in zip(
                [8, 28, 20, 26, 18, 55],
                "ABCDEF",
            ):
                sheet.column_dimensions[col].width = width

        add_status_sheet("Da_co_tai_khoan", reused)
        add_status_sheet("Tai_khoan_dang_khoa", locked)
        add_status_sheet("Can_kiem_tra", conflicts)

        for width, col in zip(
            [8, 28, 20, 26, 22, 34, 15, 18, 50],
            "ABCDEFGHI",
        ):
            ws.column_dimensions[col].width = width

        ws.freeze_panes = "A2"
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        db.commit()

    except Exception:
        db.rollback()
        raise

    # Không ghi mật khẩu rõ xuống file hệ thống hoặc database.
    safe_school = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        _normalize(school["name"]).replace(" ", "_"),
    ).strip("_") or f"school_{school['id']}"
    filename = (
        f"tai_khoan_giao_vien_{safe_school}_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )

# BAI_13B_10_V1_3_SYNC_ROUTE_END


@router.post("/giao-vien-tham-gia/luu")
async def save_school_participants(
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != SCHOOL_ROLE_CODE:
        return _forbidden()

    _ensure_schema(db)
    form = await request.form()

    try:
        batch_id = int(form.get("batch_id") or 0)
    except (TypeError, ValueError):
        batch_id = 0

    school, batch = _validate_school_scope(db, request, batch_id)
    if school is None or batch is None:
        return _forbidden()

    submission = _submission(
        db,
        batch_id=batch_id,
        school_id=int(school["id"]),
    )
    if submission and submission.get("status") == "SENT":
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/giao-vien-tham-gia"
                f"?batch_id={batch_id}"
            ),
            status_code=303,
        )

    valid_teacher_ids = {
        int(item["id"])
        for item in _teacher_rows(
            db,
            school_id=int(school["id"]),
        )
    }
    selected_ids: list[int] = []
    for raw in form.getlist("teacher_ids"):
        try:
            value = int(raw)
        except (TypeError, ValueError):
            continue
        if value in valid_teacher_ids:
            selected_ids.append(value)
    selected_ids = sorted(set(selected_ids))

    level_code = _school_level(
        db,
        school_id=int(school["id"]),
        school_name=str(school["name"]),
        school_year_id=int(batch["school_year_id"]),
    )

    db.execute(
        text(
            """
            DELETE FROM survey_investigation_participants
            WHERE survey_batch_id = :batch_id
              AND school_id = :school_id
            """
        ),
        {
            "batch_id": batch_id,
            "school_id": int(school["id"]),
        },
    )

    now_value = datetime.now().isoformat(sep=" ", timespec="seconds")
    for user_id in selected_ids:
        db.execute(
            text(
                """
                INSERT INTO survey_investigation_participants (
                    survey_batch_id, school_id, user_id,
                    level_code, created_at, updated_at
                ) VALUES (
                    :batch_id, :school_id, :user_id,
                    :level_code, :created_at, :updated_at
                )
                """
            ),
            {
                "batch_id": batch_id,
                "school_id": int(school["id"]),
                "user_id": int(user_id),
                "level_code": level_code,
                "created_at": now_value,
                "updated_at": now_value,
            },
        )

    _save_submission(
        db,
        batch_id=batch_id,
        school_id=int(school["id"]),
        status="DRAFT",
        actor_user_id=_user(request).get("id"),
    )
    _log(
        db,
        batch_id=batch_id,
        school_id=int(school["id"]),
        request=request,
        action="SAVE_DRAFT",
        count=len(selected_ids),
    )
    db.commit()

    return RedirectResponse(
        url=(
            "/dieu-tra/phan-cong-to-dieu-tra/giao-vien-tham-gia"
            f"?batch_id={batch_id}&status=draft_saved"
        ),
        status_code=303,
    )


@router.post("/giao-vien-tham-gia/gui-xa")
async def send_school_participants(
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != SCHOOL_ROLE_CODE:
        return _forbidden()

    _ensure_schema(db)
    form = await request.form()
    try:
        batch_id = int(form.get("batch_id") or 0)
    except (TypeError, ValueError):
        batch_id = 0

    school, batch = _validate_school_scope(db, request, batch_id)
    if school is None or batch is None:
        return _forbidden()

    level_code = _school_level(
        db,
        school_id=int(school["id"]),
        school_name=str(school["name"]),
        school_year_id=int(batch["school_year_id"]),
    )
    participant_ids = _participant_ids(
        db,
        batch_id=batch_id,
        school_id=int(school["id"]),
    )

    if not participant_ids or level_code not in {"MN", "TH", "THCS"}:
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/giao-vien-tham-gia"
                f"?batch_id={batch_id}"
            ),
            status_code=303,
        )

    _save_submission(
        db,
        batch_id=batch_id,
        school_id=int(school["id"]),
        status="SENT",
        actor_user_id=_user(request).get("id"),
    )
    _log(
        db,
        batch_id=batch_id,
        school_id=int(school["id"]),
        request=request,
        action="SEND_TO_COMMUNE",
        count=len(participant_ids),
    )
    db.commit()

    return RedirectResponse(
        url=(
            "/dieu-tra/phan-cong-to-dieu-tra/giao-vien-tham-gia"
            f"?batch_id={batch_id}&status=sent"
        ),
        status_code=303,
    )


@router.post("/giao-vien-tham-gia/thu-hoi")
async def withdraw_school_participants(
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != SCHOOL_ROLE_CODE:
        return _forbidden()

    _ensure_schema(db)
    form = await request.form()
    try:
        batch_id = int(form.get("batch_id") or 0)
    except (TypeError, ValueError):
        batch_id = 0

    school, batch = _validate_school_scope(db, request, batch_id)
    if school is None or batch is None:
        return _forbidden()

    # === BAI_13B_10_V2_WITHDRAW_LOCK_START ===
    # Khi xã/phường đã tạo dự thảo tổ, trường không được thu hồi danh sách
    # vì sẽ làm thay đổi cơ cấu 1 MN + 1 TH + 1 THCS.
    try:
        team_draft_exists = db.execute(
            text(
                "SELECT 1 FROM survey_investigation_teams "
                "WHERE survey_batch_id = :batch_id LIMIT 1"
            ),
            {"batch_id": int(batch_id)},
        ).scalar()
    except Exception:
        team_draft_exists = None

    if team_draft_exists is not None:
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/giao-vien-tham-gia"
                f"?batch_id={batch_id}"
            ),
            status_code=303,
        )
    # === BAI_13B_10_V2_WITHDRAW_LOCK_END ===

    participant_ids = _participant_ids(
        db,
        batch_id=batch_id,
        school_id=int(school["id"]),
    )
    _save_submission(
        db,
        batch_id=batch_id,
        school_id=int(school["id"]),
        status="DRAFT",
        actor_user_id=None,
    )
    _log(
        db,
        batch_id=batch_id,
        school_id=int(school["id"]),
        request=request,
        action="WITHDRAW",
        count=len(participant_ids),
        notes=(
            "Thu hồi trước khi xã/phường tạo tổ điều tra. "
            "Bài 13B-10 V2 sẽ khóa thu hồi sau khi đã chốt tổ."
        ),
    )
    db.commit()

    return RedirectResponse(
        url=(
            "/dieu-tra/phan-cong-to-dieu-tra/giao-vien-tham-gia"
            f"?batch_id={batch_id}&status=withdrawn"
        ),
        status_code=303,
    )


@router.get("/danh-sach-truong", response_class=HTMLResponse)
def commune_submissions_page(
    request: Request,
    batch_id: int | None = None,
    db: Session = Depends(get_db),
):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()

    _ensure_schema(db)
    user = _user(request)
    commune_id = user.get("commune_id")
    if commune_id is None:
        return _forbidden()

    batches = _available_batches(db, commune_id=int(commune_id))
    batch = _selected_batch(
        db,
        commune_id=int(commune_id),
        batch_id=batch_id,
    )

    school_rows: list[dict[str, Any]] = []
    summary = {
        "MN": 0,
        "TH": 0,
        "THCS": 0,
        "sent_schools": 0,
        "draft_schools": 0,
        "total_selected": 0,
    }

    if batch is not None:
        schools = db.execute(
            text(
                """
                SELECT s.id, s.code, s.name
                FROM schools AS s
                WHERE s.commune_id = :commune_id
                  AND s.is_active = 1
                ORDER BY s.name COLLATE NOCASE, s.id
                """
            ),
            {"commune_id": int(commune_id)},
        ).mappings().all()

        for raw_school in schools:
            school = dict(raw_school)
            level_code = _school_level(
                db,
                school_id=int(school["id"]),
                school_name=str(school["name"]),
                school_year_id=int(batch["school_year_id"]),
            )
            teacher_count = db.execute(
                text(
                    """
                    SELECT COUNT(*)
                    FROM users AS u
                    JOIN roles AS r ON r.id = u.role_id
                    WHERE u.school_id = :school_id
                      AND u.is_active = 1
                      AND UPPER(r.code) = :role_code
                    """
                ),
                {
                    "school_id": int(school["id"]),
                    "role_code": TEACHER_ROLE_CODE,
                },
            ).scalar()

            selected_count = db.execute(
                text(
                    """
                    SELECT COUNT(*)
                    FROM survey_investigation_participants
                    WHERE survey_batch_id = :batch_id
                      AND school_id = :school_id
                    """
                ),
                {
                    "batch_id": int(batch["id"]),
                    "school_id": int(school["id"]),
                },
            ).scalar()

            submission = _submission(
                db,
                batch_id=int(batch["id"]),
                school_id=int(school["id"]),
            )
            submission_status = (
                str(submission.get("status"))
                if submission
                else "CHUA_DANG_KY"
            )

            if submission_status == "SENT":
                summary["sent_schools"] += 1
                if level_code in {"MN", "TH", "THCS"}:
                    summary[level_code] += int(selected_count or 0)
            elif submission_status == "DRAFT":
                summary["draft_schools"] += 1

            summary["total_selected"] += int(selected_count or 0)

            school_rows.append(
                {
                    **school,
                    "level_code": level_code,
                    "level_label": LEVEL_LABELS.get(level_code, level_code),
                    "teacher_count": int(teacher_count or 0),
                    "selected_count": int(selected_count or 0),
                    "submission": submission,
                    "submission_status": submission_status,
                }
            )

    return templates.TemplateResponse(
        request=request,
        name="survey_teams/commune_submissions.html",
        context={
            "nguoi_dung": user,
            "batches": batches,
            "batch": batch,
            "schools": school_rows,
            "summary": summary,
            "submission_labels": SUBMISSION_LABELS,
        },
    )


# =========================================================
# BAI_13B_10_V2_RANDOM_TEAMS_START
# GHÉP TỔ 3 CẤP + CHIA ĐỀU HỘ Ở TRẠNG THÁI DỰ THẢO
# =========================================================

TEAM_SCHEMA_SQL = [
    """
    CREATE TABLE IF NOT EXISTS survey_investigation_teams (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        survey_batch_id INTEGER NOT NULL,
        commune_id INTEGER NOT NULL,
        team_number INTEGER NOT NULL,
        status VARCHAR(20) NOT NULL DEFAULT 'DRAFT',
        generation_code VARCHAR(60) NOT NULL,
        created_by_user_id INTEGER NULL,
        created_at DATETIME NOT NULL,
        updated_at DATETIME NOT NULL,
        UNIQUE(survey_batch_id, team_number),
        FOREIGN KEY(survey_batch_id) REFERENCES survey_batches(id),
        FOREIGN KEY(commune_id) REFERENCES communes(id),
        FOREIGN KEY(created_by_user_id) REFERENCES users(id)
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS survey_investigation_team_members (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        team_id INTEGER NOT NULL,
        user_id INTEGER NOT NULL,
        school_id INTEGER NOT NULL,
        level_code VARCHAR(20) NOT NULL,
        order_number INTEGER NOT NULL,
        created_at DATETIME NOT NULL,
        UNIQUE(team_id, user_id),
        UNIQUE(team_id, level_code),
        FOREIGN KEY(team_id) REFERENCES survey_investigation_teams(id)
            ON DELETE CASCADE,
        FOREIGN KEY(user_id) REFERENCES users(id),
        FOREIGN KEY(school_id) REFERENCES schools(id)
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS survey_investigation_team_forms (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        team_id INTEGER NOT NULL,
        survey_form_id INTEGER NOT NULL,
        assignment_order INTEGER NOT NULL,
        created_at DATETIME NOT NULL,
        UNIQUE(survey_form_id),
        UNIQUE(team_id, assignment_order),
        FOREIGN KEY(team_id) REFERENCES survey_investigation_teams(id)
            ON DELETE CASCADE,
        FOREIGN KEY(survey_form_id) REFERENCES survey_forms(id)
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS survey_team_generation_logs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        survey_batch_id INTEGER NOT NULL,
        commune_id INTEGER NOT NULL,
        generation_code VARCHAR(60) NOT NULL,
        action VARCHAR(30) NOT NULL,
        team_count INTEGER NOT NULL DEFAULT 0,
        household_count INTEGER NOT NULL DEFAULT 0,
        mn_count INTEGER NOT NULL DEFAULT 0,
        th_count INTEGER NOT NULL DEFAULT 0,
        thcs_count INTEGER NOT NULL DEFAULT 0,
        reserve_count INTEGER NOT NULL DEFAULT 0,
        actor_user_id INTEGER NULL,
        actor_name_snapshot VARCHAR(200) NULL,
        notes TEXT NULL,
        created_at DATETIME NOT NULL
    )
    """,
    """
    CREATE INDEX IF NOT EXISTS ix_survey_investigation_teams_batch
    ON survey_investigation_teams(survey_batch_id, status)
    """,
    """
    CREATE INDEX IF NOT EXISTS ix_survey_investigation_team_members_team
    ON survey_investigation_team_members(team_id, level_code)
    """,
    """
    CREATE INDEX IF NOT EXISTS ix_survey_investigation_team_forms_team
    ON survey_investigation_team_forms(team_id, assignment_order)
    """,
]


def _ensure_team_schema(db: Session) -> None:
    for statement in TEAM_SCHEMA_SQL:
        db.execute(text(statement))


def _commune_batch_scope(
    db: Session,
    request: Request,
    batch_id: int,
) -> dict[str, Any] | None:
    if _role(request) != COMMUNE_ROLE_CODE:
        return None

    user = _user(request)
    commune_id = user.get("commune_id")
    if commune_id is None:
        return None

    batch = _batch_row(db, int(batch_id))
    if batch is None:
        return None

    if int(batch["commune_id"]) != int(commune_id):
        return None

    return batch


def _sent_participants(
    db: Session,
    *,
    batch_id: int,
) -> list[dict[str, Any]]:
    rows = db.execute(
        text(
            """
            SELECT
                sip.user_id,
                sip.school_id,
                sip.level_code,
                u.full_name,
                u.username,
                s.name AS school_name
            FROM survey_investigation_participants AS sip
            JOIN survey_participant_submissions AS sps
              ON sps.survey_batch_id = sip.survey_batch_id
             AND sps.school_id = sip.school_id
             AND sps.status = 'SENT'
            JOIN users AS u
              ON u.id = sip.user_id
             AND u.is_active = 1
            JOIN schools AS s
              ON s.id = sip.school_id
             AND s.is_active = 1
            WHERE sip.survey_batch_id = :batch_id
              AND sip.level_code IN ('MN', 'TH', 'THCS')
            ORDER BY
                sip.level_code,
                s.name COLLATE NOCASE,
                u.full_name COLLATE NOCASE,
                u.id
            """
        ),
        {"batch_id": int(batch_id)},
    ).mappings().all()

    return [dict(row) for row in rows]


def _batch_forms(
    db: Session,
    *,
    batch_id: int,
) -> list[dict[str, Any]]:
    rows = db.execute(
        text(
            """
            SELECT
                sf.id AS survey_form_id,
                sf.form_number,
                h.id AS household_id,
                h.code AS household_code,
                h.head_name,
                h.hamlet_name,
                h.address
            FROM survey_forms AS sf
            JOIN households AS h
              ON h.id = sf.household_id
            WHERE sf.survey_batch_id = :batch_id
            ORDER BY
                COALESCE(h.hamlet_name, '') COLLATE NOCASE,
                h.head_name COLLATE NOCASE,
                sf.id
            """
        ),
        {"batch_id": int(batch_id)},
    ).mappings().all()

    return [dict(row) for row in rows]


def _draft_team_exists(
    db: Session,
    *,
    batch_id: int,
) -> bool:
    value = db.execute(
        text(
            """
            SELECT 1
            FROM survey_investigation_teams
            WHERE survey_batch_id = :batch_id
            LIMIT 1
            """
        ),
        {"batch_id": int(batch_id)},
    ).scalar()
    return value is not None


def _delete_draft_teams(
    db: Session,
    *,
    batch_id: int,
) -> None:
    team_ids = [
        int(value)
        for value in db.execute(
            text(
                """
                SELECT id
                FROM survey_investigation_teams
                WHERE survey_batch_id = :batch_id
                  AND status = 'DRAFT'
                """
            ),
            {"batch_id": int(batch_id)},
        ).scalars().all()
    ]

    if not team_ids:
        return

    placeholders = ", ".join(
        f":team_id_{index}"
        for index in range(len(team_ids))
    )
    params = {
        f"team_id_{index}": team_id
        for index, team_id in enumerate(team_ids)
    }

    db.execute(
        text(
            f"""
            DELETE FROM survey_investigation_team_forms
            WHERE team_id IN ({placeholders})
            """
        ),
        params,
    )
    db.execute(
        text(
            f"""
            DELETE FROM survey_investigation_team_members
            WHERE team_id IN ({placeholders})
            """
        ),
        params,
    )
    db.execute(
        text(
            f"""
            DELETE FROM survey_investigation_teams
            WHERE id IN ({placeholders})
            """
        ),
        params,
    )


def _generation_log(
    db: Session,
    *,
    request: Request,
    batch_id: int,
    commune_id: int,
    generation_code: str,
    action: str,
    team_count: int,
    household_count: int,
    mn_count: int,
    th_count: int,
    thcs_count: int,
    reserve_count: int,
    notes: str = "",
) -> None:
    user = _user(request)
    db.execute(
        text(
            """
            INSERT INTO survey_team_generation_logs (
                survey_batch_id,
                commune_id,
                generation_code,
                action,
                team_count,
                household_count,
                mn_count,
                th_count,
                thcs_count,
                reserve_count,
                actor_user_id,
                actor_name_snapshot,
                notes,
                created_at
            ) VALUES (
                :batch_id,
                :commune_id,
                :generation_code,
                :action,
                :team_count,
                :household_count,
                :mn_count,
                :th_count,
                :thcs_count,
                :reserve_count,
                :actor_user_id,
                :actor_name_snapshot,
                :notes,
                :created_at
            )
            """
        ),
        {
            "batch_id": int(batch_id),
            "commune_id": int(commune_id),
            "generation_code": generation_code,
            "action": action,
            "team_count": int(team_count),
            "household_count": int(household_count),
            "mn_count": int(mn_count),
            "th_count": int(th_count),
            "thcs_count": int(thcs_count),
            "reserve_count": int(reserve_count),
            "actor_user_id": user.get("id"),
            "actor_name_snapshot": str(
                user.get("full_name") or ""
            )[:200] or None,
            "notes": notes[:2000] or None,
            "created_at": datetime.now().isoformat(
                sep=" ",
                timespec="seconds",
            ),
        },
    )


def _team_preview_data(
    db: Session,
    *,
    batch_id: int,
) -> dict[str, Any]:
    participants = _sent_participants(
        db,
        batch_id=batch_id,
    )

    by_level = {
        "MN": [],
        "TH": [],
        "THCS": [],
    }
    for item in participants:
        level_code = str(item.get("level_code") or "")
        if level_code in by_level:
            by_level[level_code].append(item)

    team_rows = db.execute(
        text(
            """
            SELECT
                t.id,
                t.team_number,
                t.status,
                t.generation_code,
                t.created_at,
                COUNT(tf.id) AS household_count
            FROM survey_investigation_teams AS t
            LEFT JOIN survey_investigation_team_forms AS tf
              ON tf.team_id = t.id
            WHERE t.survey_batch_id = :batch_id
            GROUP BY
                t.id,
                t.team_number,
                t.status,
                t.generation_code,
                t.created_at
            ORDER BY t.team_number
            """
        ),
        {"batch_id": int(batch_id)},
    ).mappings().all()

    teams: list[dict[str, Any]] = []
    used_user_ids: set[int] = set()

    for raw_team in team_rows:
        team = dict(raw_team)

        member_rows = db.execute(
            text(
                """
                SELECT
                    tm.user_id,
                    tm.level_code,
                    tm.order_number,
                    u.full_name,
                    u.username,
                    s.name AS school_name
                FROM survey_investigation_team_members AS tm
                JOIN users AS u ON u.id = tm.user_id
                JOIN schools AS s ON s.id = tm.school_id
                WHERE tm.team_id = :team_id
                ORDER BY tm.order_number
                """
            ),
            {"team_id": int(team["id"])},
        ).mappings().all()

        members = [dict(row) for row in member_rows]
        for item in members:
            used_user_ids.add(int(item["user_id"]))

        hamlet_rows = db.execute(
            text(
                """
                SELECT
                    COALESCE(h.hamlet_name, 'Chưa xác định') AS hamlet_name,
                    COUNT(*) AS household_count
                FROM survey_investigation_team_forms AS tf
                JOIN survey_forms AS sf
                  ON sf.id = tf.survey_form_id
                JOIN households AS h
                  ON h.id = sf.household_id
                WHERE tf.team_id = :team_id
                GROUP BY COALESCE(h.hamlet_name, 'Chưa xác định')
                ORDER BY
                    household_count DESC,
                    hamlet_name COLLATE NOCASE
                """
            ),
            {"team_id": int(team["id"])},
        ).mappings().all()

        team["members"] = members
        team["hamlet_summary"] = [
            dict(row) for row in hamlet_rows
        ]
        teams.append(team)

    reserve: list[dict[str, Any]] = []
    for item in participants:
        if int(item["user_id"]) not in used_user_ids:
            reserve.append(item)

    form_total = len(
        _batch_forms(
            db,
            batch_id=batch_id,
        )
    )
    team_count = len(teams)
    counts = [
        int(item.get("household_count") or 0)
        for item in teams
    ]

    return {
        "teams": teams,
        "team_status": str(teams[0].get("status") or "") if teams else "",
        "reserve": reserve,
        "form_total": form_total,
        "team_count": team_count,
        "mn_count": len(by_level["MN"]),
        "th_count": len(by_level["TH"]),
        "thcs_count": len(by_level["THCS"]),
        "possible_team_count": min(
            len(by_level["MN"]),
            len(by_level["TH"]),
            len(by_level["THCS"]),
        ) if all(by_level.values()) else 0,
        "min_households": min(counts) if counts else 0,
        "max_households": max(counts) if counts else 0,
    }


def _team_changes_locked(db: Session, batch_id: int) -> bool:
    from app.survey_models import SurveyBatch
    from app.survey_workflow_models import SurveyCommuneExecutionState
    from sqlalchemy import select

    batch = db.get(SurveyBatch, batch_id)
    if batch is None or batch.is_locked:
        return True
    state = db.scalar(select(SurveyCommuneExecutionState).where(
        SurveyCommuneExecutionState.survey_batch_id == batch_id
    ))
    return bool(state and (state.is_commune_locked or state.is_province_locked))


def _change_team_member(db: Session, request: Request, *, batch_id: int,
                        team_id: int, old_user_id: int, new_user_id: int) -> None:
    """Hoán đổi cùng cấp hoặc thay bằng dự phòng trong một giao dịch."""
    batch = _commune_batch_scope(db, request, batch_id)
    if batch is None:
        raise ValueError("member_scope")
    # Giữ khóa ghi trong suốt kiểm tra và hoán đổi trên SQLite.
    changed = db.execute(text("""UPDATE survey_investigation_teams
        SET updated_at=updated_at WHERE id=:team AND survey_batch_id=:batch
        AND commune_id=:commune AND status='DRAFT'"""),
        {"team": team_id, "batch": batch_id, "commune": batch["commune_id"]})
    if changed.rowcount != 1 or _team_changes_locked(db, batch_id):
        raise ValueError("member_locked")
    if old_user_id == new_user_id:
        raise ValueError("member_invalid")
    teams = {row["id"]: dict(row) for row in db.execute(text(
        "SELECT id,status,commune_id,team_number,generation_code FROM survey_investigation_teams WHERE survey_batch_id=:batch"
    ), {"batch": batch_id}).mappings()}
    members = [dict(row) for row in db.execute(text("""SELECT m.* FROM survey_investigation_team_members m
        JOIN survey_investigation_teams t ON t.id=m.team_id WHERE t.survey_batch_id=:batch"""),
        {"batch": batch_id}).mappings()]
    old_rows = [m for m in members if m["user_id"] == old_user_id]
    new_rows = [m for m in members if m["user_id"] == new_user_id]
    if len(old_rows) != 1 or old_rows[0]["team_id"] != team_id or len(new_rows) > 1:
        raise ValueError("member_stale")
    old = old_rows[0]
    replacement = new_rows[0] if new_rows else None
    participants = {p["user_id"]: p for p in _sent_participants(db, batch_id=batch_id)}
    incoming = participants.get(new_user_id)
    outgoing = participants.get(old_user_id)
    if not incoming or not outgoing or incoming["level_code"] != old["level_code"] or outgoing["level_code"] != old["level_code"]:
        raise ValueError("member_invalid")
    affected = {team_id}
    if replacement:
        if replacement["team_id"] == team_id or replacement["level_code"] != old["level_code"]:
            raise ValueError("member_invalid")
        affected.add(replacement["team_id"])
    for affected_id in affected:
        team = teams[affected_id]
        if team["status"] != "DRAFT" or team["commune_id"] != batch["commune_id"]:
            raise ValueError("member_locked")
        slots = [m for m in members if m["team_id"] == affected_id]
        if len(slots) != 3 or {m["level_code"] for m in slots} != {"MN", "TH", "THCS"}:
            raise ValueError("member_invalid")
    db.execute(text("UPDATE survey_investigation_team_members SET user_id=:user,school_id=:school WHERE id=:id"),
               {"user": new_user_id, "school": incoming["school_id"], "id": old["id"]})
    if replacement:
        db.execute(text("UPDATE survey_investigation_team_members SET user_id=:user,school_id=:school WHERE id=:id"),
                   {"user": old_user_id, "school": outgoing["school_id"], "id": replacement["id"]})
    now = datetime.now()
    for affected_id in affected:
        db.execute(text("UPDATE survey_investigation_teams SET updated_at=:now WHERE id=:id"), {"now": now, "id": affected_id})
    destination = f"tổ {teams[replacement['team_id']]['team_number']}" if replacement else "dự phòng"
    _generation_log(db, request=request, batch_id=batch_id, commune_id=batch["commune_id"],
        generation_code=teams[team_id]["generation_code"], action="CHANGE_MEMBER", team_count=len(affected),
        household_count=0, mn_count=0, th_count=0, thcs_count=0, reserve_count=0,
        notes=f"Tổ {teams[team_id]['team_number']}: {outgoing['full_name']} (user_id={old_user_id}) đổi với {incoming['full_name']} (user_id={new_user_id}) từ {destination}; cấp {old['level_code']}.")


@router.post("/lap-to/dieu-chinh-giao-vien")
async def change_team_member(request: Request, db: Session = Depends(get_db)):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()
    form = await request.form()
    try:
        batch_id, team_id, old_id, new_id = [int(form.get(key) or 0) for key in
            ("batch_id", "team_id", "old_user_id", "new_user_id")]
    except (TypeError, ValueError):
        return _forbidden()
    if _commune_batch_scope(db, request, batch_id) is None:
        return _forbidden()
    from sqlalchemy.exc import IntegrityError
    try:
        _change_team_member(db, request, batch_id=batch_id, team_id=team_id, old_user_id=old_id, new_user_id=new_id)
        db.commit()
        status = "member_changed"
    except ValueError as exc:
        db.rollback()
        status = str(exc) if str(exc) in {"member_scope", "member_locked", "member_invalid", "member_stale"} else "member_invalid"
    except IntegrityError:
        db.rollback()
        status = "member_stale"
    return RedirectResponse(f"/dieu-tra/phan-cong-to-dieu-tra/lap-to?batch_id={batch_id}&status={status}", status_code=303)


@router.get("/lap-to", response_class=HTMLResponse)
def commune_team_builder_page(
    request: Request,
    batch_id: int | None = None,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()

    _ensure_schema(db)
    _ensure_team_schema(db)

    user = _user(request)
    commune_id = user.get("commune_id")
    if commune_id is None:
        return _forbidden()

    batches = _available_batches(
        db,
        commune_id=int(commune_id),
    )
    batch = _selected_batch(
        db,
        commune_id=int(commune_id),
        batch_id=batch_id,
    )

    preview = {
        "teams": [],
        "reserve": [],
        "form_total": 0,
        "team_count": 0,
        "mn_count": 0,
        "th_count": 0,
        "thcs_count": 0,
        "possible_team_count": 0,
        "min_households": 0,
        "max_households": 0,
    }

    if batch is not None:
        preview = _team_preview_data(
            db,
            batch_id=int(batch["id"]),
        )

    member_options = {"MN": [], "TH": [], "THCS": []}
    changes_locked = batch is None or _team_changes_locked(db, int(batch["id"]))
    if batch is not None and not changes_locked:
        locations = {m["user_id"]: team for team in preview["teams"] for m in team["members"]}
        for person in _sent_participants(db, batch_id=int(batch["id"])):
            location = locations.get(person["user_id"])
            if location is None or location["status"] == "DRAFT":
                member_options[person["level_code"]].append({**person,
                    "location": f"Tổ {location['team_number']}" if location else "Dự phòng"})

    status_messages = {
        "member_changed": "Đã điều chỉnh giáo viên. Mỗi tổ vẫn đủ 3 cấp; hộ và địa bàn của tổ được giữ nguyên.",
        "member_scope": "Không thể điều chỉnh tổ ngoài phạm vi xã hoặc đợt điều tra.",
        "member_locked": "Tổ đã chốt/gửi hoặc đợt điều tra đã khóa; không thể điều chỉnh giáo viên.",
        "member_invalid": "Không thể điều chỉnh: cần chọn giáo viên cùng cấp trong danh sách trường đã gửi và giữ đủ 3 cấp trong mỗi tổ.",
        "member_stale": "Thành viên tổ đã thay đổi. Hãy tải lại trang và chọn lại giáo viên.",
        "sent": (
            "Đã chốt và gửi phân công xuống Trường/Giáo viên."
        ),
        "already_sent": (
            "Phân công này đã được chốt trước đó; không ghi lặp."
        ),
        "no_draft": (
            "Chưa có dự thảo tổ để chốt và gửi."
        ),
        "generated": (
            "Đã tạo dự thảo tổ 3 cấp và chia đều hộ dân. "
            "Phân công hiện CHƯA gửi xuống trường/giáo viên."
        ),
        "regenerated": (
            "Đã tạo lại dự thảo tổ và phân bổ lại hộ dân."
        ),
        "cleared": (
            "Đã xóa bản dự thảo tổ. "
            "Danh sách giáo viên các trường vẫn được giữ nguyên."
        ),
        "not_ready": (
            "Chưa đủ giáo viên đã gửi ở cả 3 cấp MN, TH và THCS "
            "để tạo tổ."
        ),
        "no_household": (
            "Đợt điều tra chưa có hộ/phiếu để phân công."
        ),
    }

    return templates.TemplateResponse(
        request=request,
        name="survey_teams/commune_team_builder.html",
        context={
            "nguoi_dung": user,
            "batches": batches,
            "batch": batch,
            "preview": preview,
            "member_options": member_options,
            "changes_locked": changes_locked,
            "message": status_messages.get(status),
        },
    )


@router.post("/lap-to/tao-ngau-nhien")
async def generate_random_teams(
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()

    _ensure_schema(db)
    _ensure_team_schema(db)

    form = await request.form()
    try:
        batch_id = int(form.get("batch_id") or 0)
    except (TypeError, ValueError):
        batch_id = 0

    batch = _commune_batch_scope(
        db,
        request,
        batch_id,
    )
    if batch is None:
        return _forbidden()

    participants = _sent_participants(
        db,
        batch_id=batch_id,
    )
    forms = _batch_forms(
        db,
        batch_id=batch_id,
    )

    by_level = {
        "MN": [],
        "TH": [],
        "THCS": [],
    }
    for item in participants:
        level_code = str(item.get("level_code") or "")
        if level_code in by_level:
            by_level[level_code].append(item)

    if not forms:
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/lap-to"
                f"?batch_id={batch_id}&status=no_household"
            ),
            status_code=303,
        )

    if not all(by_level.values()):
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/lap-to"
                f"?batch_id={batch_id}&status=not_ready"
            ),
            status_code=303,
        )

    # === BAI_13B_10_V2_1_NO_EMPTY_TEAM ===
    # Số tổ không được vượt số hộ/phiếu.
    # Ví dụ 6 GV mỗi cấp nhưng chỉ có 5 hộ thì chỉ lập 5 tổ,
    # 3 GV còn lại (1 MN + 1 TH + 1 THCS) nằm ở Dự phòng.
    team_count = min(
        len(by_level["MN"]),
        len(by_level["TH"]),
        len(by_level["THCS"]),
        len(forms),
    )
    if team_count <= 0:
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/lap-to"
                f"?batch_id={batch_id}&status=not_ready"
            ),
            status_code=303,
        )

    # Không cho ghi đè tổ đã chốt ở V3.
    locked_exists = db.execute(
        text(
            """
            SELECT 1
            FROM survey_investigation_teams
            WHERE survey_batch_id = :batch_id
              AND status != 'DRAFT'
            LIMIT 1
            """
        ),
        {"batch_id": batch_id},
    ).scalar()
    if locked_exists is not None:
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/lap-to"
                f"?batch_id={batch_id}"
            ),
            status_code=303,
        )

    # Nếu đã có dự thảo, xóa dự thảo cũ để tạo lại.
    existed_before = _draft_team_exists(
        db,
        batch_id=batch_id,
    )
    _delete_draft_teams(
        db,
        batch_id=batch_id,
    )

    now = datetime.now()
    generation_code = (
        f"GEN-{batch_id}-"
        f"{now.strftime('%Y%m%d%H%M%S%f')}"
    )

    # Seed được lưu gián tiếp trong generation_code để lần tạo đã lưu
    # không thay đổi khi người dùng mở lại trang.
    import random
    seed_value = (
        int(now.timestamp() * 1_000_000)
        ^ int(_user(request).get("id") or 0)
        ^ int(batch_id)
    )
    rng = random.Random(seed_value)

    for level_code in ("MN", "TH", "THCS"):
        rng.shuffle(by_level[level_code])

    rng.shuffle(forms)

    created_at = now.isoformat(
        sep=" ",
        timespec="seconds",
    )

    team_ids: list[int] = []

    for index in range(team_count):
        cursor = db.execute(
            text(
                """
                INSERT INTO survey_investigation_teams (
                    survey_batch_id,
                    commune_id,
                    team_number,
                    status,
                    generation_code,
                    created_by_user_id,
                    created_at,
                    updated_at
                ) VALUES (
                    :batch_id,
                    :commune_id,
                    :team_number,
                    'DRAFT',
                    :generation_code,
                    :created_by_user_id,
                    :created_at,
                    :updated_at
                )
                """
            ),
            {
                "batch_id": batch_id,
                "commune_id": int(batch["commune_id"]),
                "team_number": index + 1,
                "generation_code": generation_code,
                "created_by_user_id": _user(request).get("id"),
                "created_at": created_at,
                "updated_at": created_at,
            },
        )
        team_id = int(cursor.lastrowid)
        team_ids.append(team_id)

        members = (
            ("MN", by_level["MN"][index], 1),
            ("TH", by_level["TH"][index], 2),
            ("THCS", by_level["THCS"][index], 3),
        )

        for level_code, member, order_number in members:
            db.execute(
                text(
                    """
                    INSERT INTO survey_investigation_team_members (
                        team_id,
                        user_id,
                        school_id,
                        level_code,
                        order_number,
                        created_at
                    ) VALUES (
                        :team_id,
                        :user_id,
                        :school_id,
                        :level_code,
                        :order_number,
                        :created_at
                    )
                    """
                ),
                {
                    "team_id": team_id,
                    "user_id": int(member["user_id"]),
                    "school_id": int(member["school_id"]),
                    "level_code": level_code,
                    "order_number": order_number,
                    "created_at": created_at,
                },
            )

    # Chia vòng tròn: chênh lệch số hộ giữa các tổ tối đa 1.
    team_assignment_order = {
        team_id: 0
        for team_id in team_ids
    }

    for index, form_row in enumerate(forms):
        team_id = team_ids[index % team_count]
        team_assignment_order[team_id] += 1

        db.execute(
            text(
                """
                INSERT INTO survey_investigation_team_forms (
                    team_id,
                    survey_form_id,
                    assignment_order,
                    created_at
                ) VALUES (
                    :team_id,
                    :survey_form_id,
                    :assignment_order,
                    :created_at
                )
                """
            ),
            {
                "team_id": team_id,
                "survey_form_id": int(
                    form_row["survey_form_id"]
                ),
                "assignment_order": int(
                    team_assignment_order[team_id]
                ),
                "created_at": created_at,
            },
        )

    # Kiểm tra toàn vẹn trước commit.
    assigned_count = int(
        db.execute(
            text(
                """
                SELECT COUNT(*)
                FROM survey_investigation_team_forms AS tf
                JOIN survey_investigation_teams AS t
                  ON t.id = tf.team_id
                WHERE t.survey_batch_id = :batch_id
                """
            ),
            {"batch_id": batch_id},
        ).scalar()
        or 0
    )

    member_count = int(
        db.execute(
            text(
                """
                SELECT COUNT(*)
                FROM survey_investigation_team_members AS tm
                JOIN survey_investigation_teams AS t
                  ON t.id = tm.team_id
                WHERE t.survey_batch_id = :batch_id
                """
            ),
            {"batch_id": batch_id},
        ).scalar()
        or 0
    )

    if assigned_count != len(forms):
        db.rollback()
        raise RuntimeError(
            "Không phân đủ toàn bộ phiếu/hộ cho các tổ."
        )

    if member_count != team_count * 3:
        db.rollback()
        raise RuntimeError(
            "Một hoặc nhiều tổ chưa đủ đúng 3 thành viên."
        )

    reserve_count = (
        len(by_level["MN"])
        + len(by_level["TH"])
        + len(by_level["THCS"])
        - team_count * 3
    )

    _generation_log(
        db,
        request=request,
        batch_id=batch_id,
        commune_id=int(batch["commune_id"]),
        generation_code=generation_code,
        action=(
            "REGENERATE_DRAFT"
            if existed_before
            else "GENERATE_DRAFT"
        ),
        team_count=team_count,
        household_count=len(forms),
        mn_count=len(by_level["MN"]),
        th_count=len(by_level["TH"]),
        thcs_count=len(by_level["THCS"]),
        reserve_count=reserve_count,
        notes=(
            f"Seed={seed_value}; "
            "chia vòng tròn sau khi xáo trộn danh sách hộ."
        ),
    )

    db.commit()

    return RedirectResponse(
        url=(
            "/dieu-tra/phan-cong-to-dieu-tra/lap-to"
            f"?batch_id={batch_id}&status="
            + (
                "regenerated"
                if existed_before
                else "generated"
            )
        ),
        status_code=303,
    )


# =========================================================
# BAI_13B_10_V3_2_FINALIZE_SEND_START
# Chốt tổ 3 cấp và ghi chính thức vào survey_form_investigators.
# =========================================================

@router.post("/lap-to/chot-gui")
async def finalize_and_send_team_assignments_v3_2(
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()

    _ensure_schema(db)
    _ensure_team_schema(db)

    form = await request.form()
    try:
        batch_id = int(form.get("batch_id") or 0)
    except (TypeError, ValueError):
        batch_id = 0

    batch = _commune_batch_scope(db, request, batch_id)
    if batch is None:
        return _forbidden()

    teams = [
        dict(row)
        for row in db.execute(
            text(
                """
                SELECT id, team_number, status, generation_code
                FROM survey_investigation_teams
                WHERE survey_batch_id = :batch_id
                ORDER BY team_number
                """
            ),
            {"batch_id": int(batch_id)},
        ).mappings().all()
    ]

    if not teams:
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/lap-to"
                f"?batch_id={batch_id}&status=no_draft"
            ),
            status_code=303,
        )

    if any(str(row.get("status") or "") != "DRAFT" for row in teams):
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/lap-to"
                f"?batch_id={batch_id}&status=already_sent"
            ),
            status_code=303,
        )

    now_value = datetime.now().isoformat(
        sep=" ",
        timespec="seconds",
    )
    total_forms = 0

    try:
        for team in teams:
            team_id = int(team["id"])
            team_number = int(team["team_number"])

            members = [
                dict(row)
                for row in db.execute(
                    text(
                        """
                        SELECT
                            tm.user_id,
                            tm.level_code,
                            tm.order_number,
                            u.is_active
                        FROM survey_investigation_team_members tm
                        JOIN users u ON u.id = tm.user_id
                        WHERE tm.team_id = :team_id
                        ORDER BY tm.order_number
                        """
                    ),
                    {"team_id": team_id},
                ).mappings().all()
            ]

            levels = {
                str(item.get("level_code") or "")
                for item in members
            }
            if (
                len(members) != 3
                or levels != {"MN", "TH", "THCS"}
                or any(int(item.get("is_active") or 0) != 1 for item in members)
            ):
                raise RuntimeError(
                    f"Tổ {team_number} không đủ đúng 3 giáo viên "
                    "MN + TH + THCS đang hoạt động."
                )

            form_ids = [
                int(value)
                for value in db.execute(
                    text(
                        """
                        SELECT survey_form_id
                        FROM survey_investigation_team_forms
                        WHERE team_id = :team_id
                        ORDER BY assignment_order
                        """
                    ),
                    {"team_id": team_id},
                ).scalars().all()
            ]

            # === BAI_13B_12_V2_2_FINALIZE_WITH_AREA_START ===
            if not form_ids:
                quota_total = int(
                    db.execute(
                        text(
                            """
                            SELECT COALESCE(SUM(household_quota), 0)
                            FROM survey_team_area_assignments
                            WHERE team_id = :team_id
                            """
                        ),
                        {"team_id": team_id},
                    ).scalar()
                    or 0
                )

                if quota_total <= 0:
                    raise RuntimeError(
                        f"Tổ {team_number} chưa được giao địa bàn/chỉ tiêu hộ."
                    )

                invalid_area = db.execute(
                    text(
                        """
                        SELECT
                            a.name,
                            a.expected_households,
                            COALESCE(
                                (
                                    SELECT SUM(ta.household_quota)
                                    FROM survey_team_area_assignments ta
                                    JOIN survey_investigation_teams it
                                      ON it.id = ta.team_id
                                    WHERE ta.area_id = a.id
                                      AND it.survey_batch_id = :batch_id
                                ),
                                0
                            ) AS assigned_total
                        FROM survey_commune_areas a
                        JOIN survey_batches sb
                          ON sb.commune_id = a.commune_id
                         AND sb.school_year_id = a.school_year_id
                        WHERE sb.id = :batch_id
                          AND a.is_active = 1
                          AND a.expected_households > 0
                          AND COALESCE(
                                (
                                    SELECT SUM(ta2.household_quota)
                                    FROM survey_team_area_assignments ta2
                                    JOIN survey_investigation_teams it2
                                      ON it2.id = ta2.team_id
                                    WHERE ta2.area_id = a.id
                                      AND it2.survey_batch_id = :batch_id
                                ),
                                0
                              ) != a.expected_households
                        LIMIT 1
                        """
                    ),
                    {"batch_id": int(batch_id)},
                ).mappings().first()

                if invalid_area is not None:
                    raise RuntimeError(
                        "Địa bàn "
                        f"{invalid_area['name']} chưa được giao đủ chỉ tiêu "
                        f"({invalid_area['assigned_total']}/"
                        f"{invalid_area['expected_households']})."
                    )

                # V2.2: chốt/gửi tổ trước; phiếu được tạo khi đi thực địa.
                continue
            # === BAI_13B_12_V2_2_FINALIZE_WITH_AREA_END ===

            for survey_form_id in form_ids:
                correct_batch = db.execute(
                    text(
                        """
                        SELECT 1
                        FROM survey_forms
                        WHERE id = :survey_form_id
                          AND survey_batch_id = :batch_id
                        LIMIT 1
                        """
                    ),
                    {
                        "survey_form_id": survey_form_id,
                        "batch_id": int(batch_id),
                    },
                ).scalar()

                if correct_batch is None:
                    raise RuntimeError(
                        f"Phiếu {survey_form_id} không thuộc đợt #{batch_id}."
                    )

                db.execute(
                    text(
                        """
                        DELETE FROM survey_form_investigators
                        WHERE survey_form_id = :survey_form_id
                        """
                    ),
                    {"survey_form_id": survey_form_id},
                )

                for order_number, member in enumerate(members, start=1):
                    db.execute(
                        text(
                            """
                            INSERT INTO survey_form_investigators (
                                survey_form_id,
                                user_id,
                                order_number,
                                is_primary,
                                signed_at,
                                notes,
                                created_at
                            ) VALUES (
                                :survey_form_id,
                                :user_id,
                                :order_number,
                                :is_primary,
                                :signed_at,
                                :notes,
                                :created_at
                            )
                            """
                        ),
                        {
                            "survey_form_id": survey_form_id,
                            "user_id": int(member["user_id"]),
                            "order_number": order_number,
                            "is_primary": 1 if order_number == 1 else 0,
                            "signed_at": now_value,
                            "notes": (
                                f"Tổ điều tra 3 cấp số {team_number}; "
                                f"{member['level_code']}; chốt {now_value}"
                            ),
                            "created_at": now_value,
                        },
                    )

                total_forms += 1

        bad_form = db.execute(
            text(
                """
                SELECT tf.survey_form_id
                FROM survey_investigation_team_forms tf
                JOIN survey_investigation_teams t
                  ON t.id = tf.team_id
                LEFT JOIN survey_form_investigators sfi
                  ON sfi.survey_form_id = tf.survey_form_id
                WHERE t.survey_batch_id = :batch_id
                GROUP BY tf.survey_form_id
                HAVING COUNT(sfi.id) != 3
                LIMIT 1
                """
            ),
            {"batch_id": int(batch_id)},
        ).scalar()

        if bad_form is not None:
            raise RuntimeError(
                f"Phiếu {bad_form} chưa có đúng 3 người điều tra."
            )

        db.execute(
            text(
                """
                UPDATE survey_investigation_teams
                SET status = 'SENT',
                    updated_at = :updated_at
                WHERE survey_batch_id = :batch_id
                  AND status = 'DRAFT'
                """
            ),
            {
                "batch_id": int(batch_id),
                "updated_at": now_value,
            },
        )

        _generation_log(
            db,
            request=request,
            batch_id=int(batch_id),
            commune_id=int(batch["commune_id"]),
            generation_code=str(
                teams[0].get("generation_code") or ""
            ),
            action="FINALIZE_SEND",
            team_count=len(teams),
            household_count=total_forms,
            mn_count=len(teams),
            th_count=len(teams),
            thcs_count=len(teams),
            reserve_count=0,
            notes=(
                f"Chốt và gửi {total_forms} phiếu; "
                f"{total_forms * 3} lượt phân công."
            ),
        )

        db.commit()

    except Exception:
        db.rollback()
        raise

    return RedirectResponse(
        url=(
            "/dieu-tra/phan-cong-to-dieu-tra/lap-to"
            f"?batch_id={batch_id}&status=sent"
        ),
        status_code=303,
    )

# BAI_13B_10_V3_2_FINALIZE_SEND_END


@router.post("/lap-to/xoa-du-thao")
async def clear_team_draft(
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()

    _ensure_schema(db)
    _ensure_team_schema(db)

    form = await request.form()
    try:
        batch_id = int(form.get("batch_id") or 0)
    except (TypeError, ValueError):
        batch_id = 0

    batch = _commune_batch_scope(
        db,
        request,
        batch_id,
    )
    if batch is None:
        return _forbidden()

    locked_exists = db.execute(
        text(
            """
            SELECT 1
            FROM survey_investigation_teams
            WHERE survey_batch_id = :batch_id
              AND status != 'DRAFT'
            LIMIT 1
            """
        ),
        {"batch_id": batch_id},
    ).scalar()

    if locked_exists is None:
        _delete_draft_teams(
            db,
            batch_id=batch_id,
        )

    _generation_log(
        db,
        request=request,
        batch_id=batch_id,
        commune_id=int(batch["commune_id"]),
        generation_code=(
            f"CLEAR-{batch_id}-"
            f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}"
        ),
        action="CLEAR_DRAFT",
        team_count=0,
        household_count=0,
        mn_count=0,
        th_count=0,
        thcs_count=0,
        reserve_count=0,
        notes="Xóa dự thảo trước khi chốt phân công.",
    )

    db.commit()

    return RedirectResponse(
        url=(
            "/dieu-tra/phan-cong-to-dieu-tra/lap-to"
            f"?batch_id={batch_id}&status=cleared"
        ),
        status_code=303,
    )


@router.get(
    "/lap-to/{team_id}",
    response_class=HTMLResponse,
)
def commune_team_detail(
    team_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()

    _ensure_team_schema(db)

    user = _user(request)
    commune_id = user.get("commune_id")
    if commune_id is None:
        return _forbidden()

    team = db.execute(
        text(
            """
            SELECT
                t.id,
                t.survey_batch_id,
                t.commune_id,
                t.team_number,
                t.status,
                t.generation_code,
                sb.name AS batch_name,
                sy.code AS school_year_code
            FROM survey_investigation_teams AS t
            JOIN survey_batches AS sb
              ON sb.id = t.survey_batch_id
            JOIN school_years AS sy
              ON sy.id = sb.school_year_id
            WHERE t.id = :team_id
              AND t.commune_id = :commune_id
            LIMIT 1
            """
        ),
        {
            "team_id": int(team_id),
            "commune_id": int(commune_id),
        },
    ).mappings().first()

    if team is None:
        return _forbidden()

    members = [
        dict(row)
        for row in db.execute(
            text(
                """
                SELECT
                    tm.level_code,
                    tm.order_number,
                    u.full_name,
                    u.username,
                    s.name AS school_name
                FROM survey_investigation_team_members AS tm
                JOIN users AS u ON u.id = tm.user_id
                JOIN schools AS s ON s.id = tm.school_id
                WHERE tm.team_id = :team_id
                ORDER BY tm.order_number
                """
            ),
            {"team_id": int(team_id)},
        ).mappings().all()
    ]

    households = [
        dict(row)
        for row in db.execute(
            text(
                """
                SELECT
                    tf.assignment_order,
                    sf.form_number,
                    h.code AS household_code,
                    h.head_name,
                    h.hamlet_name,
                    h.address,
                    (
                        SELECT COUNT(*)
                        FROM survey_people AS sp
                        WHERE sp.household_id = h.id
                          AND sp.is_active = 1
                    ) AS people_count
                FROM survey_investigation_team_forms AS tf
                JOIN survey_forms AS sf
                  ON sf.id = tf.survey_form_id
                JOIN households AS h
                  ON h.id = sf.household_id
                WHERE tf.team_id = :team_id
                ORDER BY tf.assignment_order
                """
            ),
            {"team_id": int(team_id)},
        ).mappings().all()
    ]

    return templates.TemplateResponse(
        request=request,
        name="survey_teams/commune_team_detail.html",
        context={
            "nguoi_dung": user,
            "team": dict(team),
            "members": members,
            "households": households,
        },
    )

# BAI_13B_10_V2_RANDOM_TEAMS_END

# ============================================================
# BAI_13B_12_V2_2_AREA_WORKFLOW_START
# Danh mục địa bàn + giao địa bàn/chỉ tiêu cho tổ điều tra.
# ============================================================
from datetime import datetime as _b1322_datetime

_B1322_AREA_TYPES = {
    "THON": "Thôn",
    "XOM": "Xóm",
    "KHOI": "Khối",
    "TO": "Tổ",
    "BAN": "Bản",
}

_B1322_SCHEMA_SQL = (
    """
    CREATE TABLE IF NOT EXISTS survey_commune_areas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        commune_id INTEGER NOT NULL,
        school_year_id INTEGER NOT NULL,
        code VARCHAR(80) NOT NULL,
        name VARCHAR(200) NOT NULL,
        area_type VARCHAR(20) NOT NULL,
        expected_households INTEGER NOT NULL DEFAULT 0,
        notes TEXT NULL,
        is_active BOOLEAN NOT NULL DEFAULT 1,
        created_by_user_id INTEGER NULL,
        created_at DATETIME NOT NULL,
        updated_at DATETIME NOT NULL,
        UNIQUE(commune_id, school_year_id, code),
        FOREIGN KEY(commune_id) REFERENCES communes(id),
        FOREIGN KEY(school_year_id) REFERENCES school_years(id),
        FOREIGN KEY(created_by_user_id) REFERENCES users(id)
    )
    """,
    """
    CREATE UNIQUE INDEX IF NOT EXISTS
    uq_survey_commune_areas_name
    ON survey_commune_areas(
        commune_id,
        school_year_id,
        area_type,
        name COLLATE NOCASE
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS survey_team_area_assignments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        team_id INTEGER NOT NULL,
        area_id INTEGER NOT NULL,
        household_quota INTEGER NOT NULL DEFAULT 0,
        notes TEXT NULL,
        created_by_user_id INTEGER NULL,
        created_at DATETIME NOT NULL,
        updated_at DATETIME NOT NULL,
        UNIQUE(team_id, area_id),
        FOREIGN KEY(team_id)
            REFERENCES survey_investigation_teams(id)
            ON DELETE CASCADE,
        FOREIGN KEY(area_id)
            REFERENCES survey_commune_areas(id),
        FOREIGN KEY(created_by_user_id)
            REFERENCES users(id)
    )
    """,
    """
    CREATE INDEX IF NOT EXISTS
    ix_survey_team_area_assignments_area
    ON survey_team_area_assignments(area_id)
    """,
    """
    CREATE TABLE IF NOT EXISTS survey_form_areas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        survey_form_id INTEGER NOT NULL UNIQUE,
        area_id INTEGER NOT NULL,
        created_at DATETIME NOT NULL,
        FOREIGN KEY(survey_form_id)
            REFERENCES survey_forms(id),
        FOREIGN KEY(area_id)
            REFERENCES survey_commune_areas(id)
    )
    """,
    """
    CREATE INDEX IF NOT EXISTS
    ix_survey_form_areas_area
    ON survey_form_areas(area_id)
    """,
)


def _b1322_ensure_area_schema(db: Session) -> None:
    for statement in _B1322_SCHEMA_SQL:
        db.execute(text(statement))
    db.commit()


def _b1322_int(value, default=0):
    try:
        return int(str(value or "").strip())
    except (TypeError, ValueError):
        return default


def _b1322_redirect(*, school_year_id=None, batch_id=None, status=None):
    params = []
    if school_year_id:
        params.append(f"school_year_id={int(school_year_id)}")
    if batch_id:
        params.append(f"batch_id={int(batch_id)}")
    if status:
        params.append(f"status={status}")
    suffix = ("?" + "&".join(params)) if params else ""
    return RedirectResponse(
        url="/dieu-tra/phan-cong-to-dieu-tra/dia-ban" + suffix,
        status_code=303,
    )


def _b1322_area_messages(status):
    messages = {
        "saved": ("ok", "Đã lưu địa bàn."),
        "duplicate": ("error", "Tên địa bàn đã tồn tại trong cùng loại và năm học."),
        "invalid": ("error", "Dữ liệu địa bàn chưa hợp lệ."),
        "locked": ("error", "Địa bàn đã gắn với nhiệm vụ đã chốt/gửi nên không thể thay đổi."),
        "deactivated": ("ok", "Đã ngừng sử dụng địa bàn."),
        "activated": ("ok", "Đã kích hoạt lại địa bàn."),
        "assign_saved": ("ok", "Đã lưu giao địa bàn/chỉ tiêu cho các tổ."),
        "distributed": ("ok", "Đã chia đều chỉ tiêu theo số hộ dự kiến."),
        "over_expected": ("error", "Tổng chỉ tiêu giao cho một địa bàn vượt số hộ dự kiến."),
        "no_teams": ("error", "Chưa có tổ điều tra để giao địa bàn."),
        "no_areas": ("error", "Chưa có địa bàn đang hoạt động có số hộ dự kiến."),
        "team_locked": ("error", "Tổ đã chốt/gửi; không thể thay đổi chỉ tiêu."),
    }
    return messages.get(str(status or ""), (None, None))


@router.get("/dia-ban", response_class=HTMLResponse)
def b1322_commune_area_page(
    request: Request,
    school_year_id: str | None = None,
    batch_id: str | None = None,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()

    _b1322_ensure_area_schema(db)
    user = _user(request)
    commune_id = _b1322_int(user.get("commune_id"))
    if not commune_id:
        return _forbidden()

    years = [
        dict(row)
        for row in db.execute(
            text(
                """
                SELECT id, code, name, is_active
                FROM school_years
                ORDER BY code DESC, id DESC
                """
            )
        ).mappings().all()
    ]

    selected_year_id = _b1322_int(school_year_id)
    if not selected_year_id and years:
        active = next(
            (item for item in years if bool(item.get("is_active"))),
            None,
        )
        selected_year_id = int(
            (active or years[0])["id"]
        )

    selected_year = next(
        (
            item for item in years
            if int(item["id"]) == int(selected_year_id)
        ),
        None,
    )

    batches = []
    if selected_year is not None:
        batches = [
            dict(row)
            for row in db.execute(
                text(
                    """
                    SELECT id, code, name, status, is_locked
                    FROM survey_batches
                    WHERE commune_id = :commune_id
                      AND school_year_id = :school_year_id
                    ORDER BY created_at DESC, id DESC
                    """
                ),
                {
                    "commune_id": commune_id,
                    "school_year_id": int(selected_year["id"]),
                },
            ).mappings().all()
        ]

    selected_batch_id = _b1322_int(batch_id)
    selected_batch = next(
        (
            item for item in batches
            if int(item["id"]) == int(selected_batch_id)
        ),
        None,
    )

    areas = []
    if selected_year is not None:
        rows = db.execute(
            text(
                """
                SELECT
                    a.id,
                    a.code,
                    a.name,
                    a.area_type,
                    a.expected_households,
                    a.notes,
                    a.is_active,
                    COALESCE(
                        (
                            SELECT COUNT(*)
                            FROM survey_form_areas fa
                            JOIN survey_forms sf
                              ON sf.id = fa.survey_form_id
                            JOIN survey_batches sb
                              ON sb.id = sf.survey_batch_id
                            WHERE fa.area_id = a.id
                              AND sb.commune_id = a.commune_id
                              AND sb.school_year_id = a.school_year_id
                        ),
                        0
                    ) AS actual_households
                FROM survey_commune_areas a
                WHERE a.commune_id = :commune_id
                  AND a.school_year_id = :school_year_id
                ORDER BY
                    a.is_active DESC,
                    a.area_type,
                    a.name COLLATE NOCASE,
                    a.id
                """
            ),
            {
                "commune_id": commune_id,
                "school_year_id": int(selected_year["id"]),
            },
        ).mappings().all()

        for row in rows:
            item = dict(row)
            expected = max(
                0,
                _b1322_int(item.get("expected_households")),
            )
            actual = max(
                0,
                _b1322_int(item.get("actual_households")),
            )
            item["area_type_label"] = _B1322_AREA_TYPES.get(
                str(item.get("area_type") or ""),
                str(item.get("area_type") or ""),
            )
            # === BAI_13B_12_V2_3_AREA_STATS ===
            difference = expected - actual
            item["remaining_households"] = max(0, difference)
            item["excess_households"] = max(0, -difference)
            item["progress_percent"] = (
                round(actual * 100 / expected)
                if expected
                else (100 if actual else 0)
            )
            item["assigned_quota_selected"] = 0
            areas.append(item)

    teams = []
    assignment_values = {}
    teams_locked = False

    if selected_batch is not None:
        team_rows = db.execute(
            text(
                """
                SELECT id, team_number, status
                FROM survey_investigation_teams
                WHERE survey_batch_id = :batch_id
                ORDER BY team_number
                """
            ),
            {"batch_id": int(selected_batch["id"])},
        ).mappings().all()

        for row in team_rows:
            team = dict(row)
            member_rows = db.execute(
                text(
                    """
                    SELECT
                        tm.level_code,
                        u.full_name,
                        COALESCE(s.name, '') AS school_name
                    FROM survey_investigation_team_members tm
                    JOIN users u ON u.id = tm.user_id
                    LEFT JOIN schools s ON s.id = tm.school_id
                    WHERE tm.team_id = :team_id
                    ORDER BY tm.order_number
                    """
                ),
                {"team_id": int(team["id"])},
            ).mappings().all()

            team["members_text"] = " · ".join(
                f"{m['level_code']}: {m['full_name']}"
                for m in member_rows
            )

            team["assigned_quota"] = _b1322_int(
                db.execute(
                    text(
                        """
                        SELECT COALESCE(SUM(household_quota), 0)
                        FROM survey_team_area_assignments
                        WHERE team_id = :team_id
                        """
                    ),
                    {"team_id": int(team["id"])},
                ).scalar()
            )

            team["actual_count"] = _b1322_int(
                db.execute(
                    text(
                        """
                        SELECT COUNT(*)
                        FROM survey_investigation_team_forms tf
                        WHERE tf.team_id = :team_id
                        """
                    ),
                    {"team_id": int(team["id"])},
                ).scalar()
            )

            teams.append(team)

        teams_locked = any(
            str(team.get("status") or "") != "DRAFT"
            for team in teams
        )

        assignment_rows = db.execute(
            text(
                """
                SELECT
                    ta.team_id,
                    ta.area_id,
                    ta.household_quota
                FROM survey_team_area_assignments ta
                JOIN survey_investigation_teams t
                  ON t.id = ta.team_id
                WHERE t.survey_batch_id = :batch_id
                """
            ),
            {"batch_id": int(selected_batch["id"])},
        ).mappings().all()

        area_assigned = {}
        for row in assignment_rows:
            team_id_value = int(row["team_id"])
            area_id_value = int(row["area_id"])
            quota = max(0, _b1322_int(row["household_quota"]))
            assignment_values[
                f"{team_id_value}:{area_id_value}"
            ] = quota
            area_assigned[area_id_value] = (
                area_assigned.get(area_id_value, 0) + quota
            )

        for area in areas:
            area["assigned_quota_selected"] = area_assigned.get(
                int(area["id"]),
                0,
            )

    active_areas = [
        item for item in areas if bool(item.get("is_active"))
    ]

    area_summary = {
        "active_count": len(active_areas),
        "expected_total": sum(
            _b1322_int(a["expected_households"])
            for a in active_areas
        ),
        "actual_total": sum(
            _b1322_int(a["actual_households"])
            for a in active_areas
        ),
    }
    total_difference = (
        area_summary["expected_total"]
        - area_summary["actual_total"]
    )
    area_summary["remaining_total"] = max(0, total_difference)
    area_summary["excess_total"] = max(0, -total_difference)

    kind, message = _b1322_area_messages(status)

    return templates.TemplateResponse(
        request=request,
        name="survey_teams/commune_areas_v2_2.html",
        context={
            "nguoi_dung": user,
            "years": years,
            "selected_year": selected_year,
            "batches": batches,
            "selected_batch": selected_batch,
            "areas": areas,
            "active_areas": active_areas,
            "area_summary": area_summary,
            "teams": teams,
            "teams_locked": teams_locked,
            "assignment_values": assignment_values,
            "message": message,
            "message_kind": kind,
        },
    )


@router.post("/dia-ban/luu")
async def b1322_save_area(
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()
    _b1322_ensure_area_schema(db)

    form = await request.form()
    user = _user(request)
    commune_id = _b1322_int(user.get("commune_id"))
    school_year_id = _b1322_int(form.get("school_year_id"))
    batch_id = _b1322_int(form.get("batch_id"))
    area_id = _b1322_int(form.get("area_id"))
    area_type = str(form.get("area_type") or "").strip().upper()
    name = " ".join(
        str(form.get("name") or "").strip().split()
    )
    expected = _b1322_int(form.get("expected_households"), -1)
    notes = str(form.get("notes") or "").strip()[:500]

    if (
        not commune_id
        or not school_year_id
        or area_type not in _B1322_AREA_TYPES
        or not name
        or len(name) > 200
        or expected < 0
    ):
        return _b1322_redirect(
            school_year_id=school_year_id,
            batch_id=batch_id,
            status="invalid",
        )

    actual_count = 0
    if area_id:
        existing = db.execute(
            text(
                """
                SELECT id
                FROM survey_commune_areas
                WHERE id = :area_id
                  AND commune_id = :commune_id
                  AND school_year_id = :school_year_id
                LIMIT 1
                """
            ),
            {
                "area_id": area_id,
                "commune_id": commune_id,
                "school_year_id": school_year_id,
            },
        ).scalar()
        if existing is None:
            return _b1322_redirect(
                school_year_id=school_year_id,
                batch_id=batch_id,
                status="invalid",
            )

        locked_ref = db.execute(
            text(
                """
                SELECT 1
                FROM survey_team_area_assignments ta
                JOIN survey_investigation_teams t
                  ON t.id = ta.team_id
                WHERE ta.area_id = :area_id
                  AND t.status != 'DRAFT'
                LIMIT 1
                """
            ),
            {"area_id": area_id},
        ).scalar()
        if locked_ref is not None:
            return _b1322_redirect(
                school_year_id=school_year_id,
                batch_id=batch_id,
                status="locked",
            )

        # === BAI_13B_12_V2_3_EXPECTED_IS_PLAN_START ===
        # Số hộ dự kiến là kế hoạch quản lý, có thể thấp hơn số hộ thực tế.
        # Không chặn sửa danh mục chỉ vì thực tế đã vượt dự kiến.
        # === BAI_13B_12_V2_3_EXPECTED_IS_PLAN_END ===

    try:
        now = _b1322_datetime.now()
        if area_id:
            db.execute(
                text(
                    """
                    UPDATE survey_commune_areas
                    SET name = :name,
                        area_type = :area_type,
                        expected_households = :expected,
                        notes = :notes,
                        updated_at = :updated_at
                    WHERE id = :area_id
                    """
                ),
                {
                    "name": name,
                    "area_type": area_type,
                    "expected": expected,
                    "notes": notes or None,
                    "updated_at": now,
                    "area_id": area_id,
                },
            )
        else:
            code = (
                f"{area_type}-{commune_id}-{school_year_id}-"
                f"{now.strftime('%Y%m%d%H%M%S%f')}"
            )
            db.execute(
                text(
                    """
                    INSERT INTO survey_commune_areas (
                        commune_id,
                        school_year_id,
                        code,
                        name,
                        area_type,
                        expected_households,
                        notes,
                        is_active,
                        created_by_user_id,
                        created_at,
                        updated_at
                    ) VALUES (
                        :commune_id,
                        :school_year_id,
                        :code,
                        :name,
                        :area_type,
                        :expected,
                        :notes,
                        1,
                        :created_by_user_id,
                        :created_at,
                        :updated_at
                    )
                    """
                ),
                {
                    "commune_id": commune_id,
                    "school_year_id": school_year_id,
                    "code": code,
                    "name": name,
                    "area_type": area_type,
                    "expected": expected,
                    "notes": notes or None,
                    "created_by_user_id": user.get("id"),
                    "created_at": now,
                    "updated_at": now,
                },
            )
        db.commit()
    except Exception:
        db.rollback()
        return _b1322_redirect(
            school_year_id=school_year_id,
            batch_id=batch_id,
            status="duplicate",
        )

    return _b1322_redirect(
        school_year_id=school_year_id,
        batch_id=batch_id,
        status="saved",
    )


@router.post("/dia-ban/ngung-dung")
async def b1322_deactivate_area(
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()
    _b1322_ensure_area_schema(db)
    form = await request.form()
    user = _user(request)
    commune_id = _b1322_int(user.get("commune_id"))
    area_id = _b1322_int(form.get("area_id"))
    school_year_id = _b1322_int(form.get("school_year_id"))
    batch_id = _b1322_int(form.get("batch_id"))

    locked_ref = db.execute(
        text(
            """
            SELECT 1
            FROM survey_team_area_assignments ta
            JOIN survey_investigation_teams t
              ON t.id = ta.team_id
            WHERE ta.area_id = :area_id
              AND t.status != 'DRAFT'
            LIMIT 1
            """
        ),
        {"area_id": area_id},
    ).scalar()
    form_ref = db.execute(
        text(
            """
            SELECT 1 FROM survey_form_areas
            WHERE area_id = :area_id
            LIMIT 1
            """
        ),
        {"area_id": area_id},
    ).scalar()

    if locked_ref is not None or form_ref is not None:
        return _b1322_redirect(
            school_year_id=school_year_id,
            batch_id=batch_id,
            status="locked",
        )

    db.execute(
        text(
            """
            DELETE FROM survey_team_area_assignments
            WHERE area_id = :area_id
            """
        ),
        {"area_id": area_id},
    )
    db.execute(
        text(
            """
            UPDATE survey_commune_areas
            SET is_active = 0,
                updated_at = :updated_at
            WHERE id = :area_id
              AND commune_id = :commune_id
            """
        ),
        {
            "updated_at": _b1322_datetime.now(),
            "area_id": area_id,
            "commune_id": commune_id,
        },
    )
    db.commit()
    return _b1322_redirect(
        school_year_id=school_year_id,
        batch_id=batch_id,
        status="deactivated",
    )


@router.post("/dia-ban/kich-hoat")
async def b1322_activate_area(
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()
    _b1322_ensure_area_schema(db)
    form = await request.form()
    user = _user(request)
    area_id = _b1322_int(form.get("area_id"))
    school_year_id = _b1322_int(form.get("school_year_id"))
    batch_id = _b1322_int(form.get("batch_id"))

    db.execute(
        text(
            """
            UPDATE survey_commune_areas
            SET is_active = 1,
                updated_at = :updated_at
            WHERE id = :area_id
              AND commune_id = :commune_id
            """
        ),
        {
            "updated_at": _b1322_datetime.now(),
            "area_id": area_id,
            "commune_id": _b1322_int(user.get("commune_id")),
        },
    )
    db.commit()
    return _b1322_redirect(
        school_year_id=school_year_id,
        batch_id=batch_id,
        status="activated",
    )


def _b1322_batch_teams_and_areas(db, request, batch_id):
    batch = _commune_batch_scope(db, request, int(batch_id))
    if batch is None:
        return None, [], []

    teams = [
        dict(row)
        for row in db.execute(
            text(
                """
                SELECT id, team_number, status
                FROM survey_investigation_teams
                WHERE survey_batch_id = :batch_id
                ORDER BY team_number
                """
            ),
            {"batch_id": int(batch_id)},
        ).mappings().all()
    ]

    areas = [
        dict(row)
        for row in db.execute(
            text(
                """
                SELECT id, name, area_type, expected_households
                FROM survey_commune_areas
                WHERE commune_id = :commune_id
                  AND school_year_id = :school_year_id
                  AND is_active = 1
                  AND expected_households > 0
                ORDER BY name COLLATE NOCASE, id
                """
            ),
            {
                "commune_id": int(batch["commune_id"]),
                "school_year_id": int(batch["school_year_id"]),
            },
        ).mappings().all()
    ]
    return batch, teams, areas


@router.post("/dia-ban/phan-cong-luu")
async def b1322_save_area_assignments(
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()
    _b1322_ensure_area_schema(db)

    form = await request.form()
    batch_id = _b1322_int(form.get("batch_id"))
    school_year_id = _b1322_int(form.get("school_year_id"))
    batch, teams, areas = _b1322_batch_teams_and_areas(
        db, request, batch_id
    )

    if batch is None or not teams:
        return _b1322_redirect(
            school_year_id=school_year_id,
            batch_id=batch_id,
            status="no_teams",
        )

    if any(str(t["status"]) != "DRAFT" for t in teams):
        return _b1322_redirect(
            school_year_id=school_year_id,
            batch_id=batch_id,
            status="team_locked",
        )

    team_ids = {int(t["id"]) for t in teams}
    area_ids = {int(a["id"]) for a in areas}
    values = {}
    area_totals = {area_id: 0 for area_id in area_ids}

    for key, raw in form.multi_items():
        if not str(key).startswith("quota__"):
            continue
        parts = str(key).split("__")
        if len(parts) != 3:
            continue
        team_id = _b1322_int(parts[1])
        area_id = _b1322_int(parts[2])
        quota = _b1322_int(raw, -1)
        if (
            team_id not in team_ids
            or area_id not in area_ids
            or quota < 0
        ):
            return _b1322_redirect(
                school_year_id=school_year_id,
                batch_id=batch_id,
                status="invalid",
            )
        if quota:
            values[(team_id, area_id)] = quota
            area_totals[area_id] += quota

    expected_map = {
        int(a["id"]): int(a["expected_households"] or 0)
        for a in areas
    }
    if any(
        area_totals.get(area_id, 0) > expected
        for area_id, expected in expected_map.items()
    ):
        return _b1322_redirect(
            school_year_id=school_year_id,
            batch_id=batch_id,
            status="over_expected",
        )

    now = _b1322_datetime.now()
    try:
        db.execute(
            text(
                """
                DELETE FROM survey_team_area_assignments
                WHERE team_id IN (
                    SELECT id
                    FROM survey_investigation_teams
                    WHERE survey_batch_id = :batch_id
                )
                """
            ),
            {"batch_id": batch_id},
        )
        user_id = _user(request).get("id")
        for (team_id, area_id), quota in values.items():
            db.execute(
                text(
                    """
                    INSERT INTO survey_team_area_assignments (
                        team_id,
                        area_id,
                        household_quota,
                        created_by_user_id,
                        created_at,
                        updated_at
                    ) VALUES (
                        :team_id,
                        :area_id,
                        :household_quota,
                        :created_by_user_id,
                        :created_at,
                        :updated_at
                    )
                    """
                ),
                {
                    "team_id": team_id,
                    "area_id": area_id,
                    "household_quota": quota,
                    "created_by_user_id": user_id,
                    "created_at": now,
                    "updated_at": now,
                },
            )
        db.commit()
    except Exception:
        db.rollback()
        return _b1322_redirect(
            school_year_id=school_year_id,
            batch_id=batch_id,
            status="invalid",
        )

    return _b1322_redirect(
        school_year_id=school_year_id,
        batch_id=batch_id,
        status="assign_saved",
    )


@router.post("/dia-ban/chia-deu")
async def b1322_auto_distribute_areas(
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != COMMUNE_ROLE_CODE:
        return _forbidden()
    _b1322_ensure_area_schema(db)

    form = await request.form()
    batch_id = _b1322_int(form.get("batch_id"))
    school_year_id = _b1322_int(form.get("school_year_id"))
    batch, teams, areas = _b1322_batch_teams_and_areas(
        db, request, batch_id
    )

    if batch is None or not teams:
        return _b1322_redirect(
            school_year_id=school_year_id,
            batch_id=batch_id,
            status="no_teams",
        )
    if not areas:
        return _b1322_redirect(
            school_year_id=school_year_id,
            batch_id=batch_id,
            status="no_areas",
        )
    if any(str(t["status"]) != "DRAFT" for t in teams):
        return _b1322_redirect(
            school_year_id=school_year_id,
            batch_id=batch_id,
            status="team_locked",
        )

    total = sum(
        int(a["expected_households"] or 0) for a in areas
    )
    team_count = len(teams)
    base = total // team_count
    extra = total % team_count
    capacities = {
        int(team["id"]): base + (
            1 if index < extra else 0
        )
        for index, team in enumerate(teams)
    }

    allocations = {}
    team_index = 0
    ordered_team_ids = [int(t["id"]) for t in teams]

    for area in areas:
        remaining = int(area["expected_households"] or 0)
        area_id = int(area["id"])
        while remaining > 0 and team_index < len(ordered_team_ids):
            team_id = ordered_team_ids[team_index]
            capacity = capacities.get(team_id, 0)
            if capacity <= 0:
                team_index += 1
                continue
            amount = min(remaining, capacity)
            allocations[(team_id, area_id)] = (
                allocations.get((team_id, area_id), 0)
                + amount
            )
            capacities[team_id] -= amount
            remaining -= amount
            if capacities[team_id] <= 0:
                team_index += 1

    now = _b1322_datetime.now()
    try:
        db.execute(
            text(
                """
                DELETE FROM survey_team_area_assignments
                WHERE team_id IN (
                    SELECT id
                    FROM survey_investigation_teams
                    WHERE survey_batch_id = :batch_id
                )
                """
            ),
            {"batch_id": batch_id},
        )
        user_id = _user(request).get("id")
        for (team_id, area_id), quota in allocations.items():
            if quota <= 0:
                continue
            db.execute(
                text(
                    """
                    INSERT INTO survey_team_area_assignments (
                        team_id, area_id, household_quota,
                        created_by_user_id, created_at, updated_at
                    ) VALUES (
                        :team_id, :area_id, :quota,
                        :user_id, :created_at, :updated_at
                    )
                    """
                ),
                {
                    "team_id": team_id,
                    "area_id": area_id,
                    "quota": quota,
                    "user_id": user_id,
                    "created_at": now,
                    "updated_at": now,
                },
            )
        db.commit()
    except Exception:
        db.rollback()
        return _b1322_redirect(
            school_year_id=school_year_id,
            batch_id=batch_id,
            status="invalid",
        )

    return _b1322_redirect(
        school_year_id=school_year_id,
        batch_id=batch_id,
        status="distributed",
    )
# ============================================================
# BAI_13B_12_V2_2_AREA_WORKFLOW_END
# ============================================================

# ============================================================
# BAI_13B_12_V2_2_TEAM_NEW_HOUSEHOLD_START
# Tổ đã chốt/gửi trực tiếp tạo hộ thực tế tại địa bàn được giao.
# URL nằm trong router tổ điều tra để không xung đột route household_id.
# ============================================================
from uuid import uuid4 as _b1322_uuid4
from fastapi.responses import JSONResponse as _b1322_JSONResponse


def _b1322_team_context_for_teacher(
    db: Session,
    *,
    batch_id: int,
    user_id: int,
):
    batch = db.execute(
        text(
            """
            SELECT
                sb.id,
                sb.code,
                sb.name,
                sb.status,
                sb.is_locked,
                sb.commune_id,
                sb.school_year_id,
                sy.code AS school_year_code
            FROM survey_batches sb
            JOIN school_years sy
              ON sy.id = sb.school_year_id
            WHERE sb.id = :batch_id
            LIMIT 1
            """
        ),
        {"batch_id": int(batch_id)},
    ).mappings().first()

    if batch is None:
        return None, None, [], []

    team = db.execute(
        text(
            """
            SELECT
                t.id,
                t.team_number,
                t.status,
                t.commune_id,
                t.survey_batch_id
            FROM survey_investigation_teams t
            JOIN survey_investigation_team_members tm
              ON tm.team_id = t.id
            WHERE t.survey_batch_id = :batch_id
              AND tm.user_id = :user_id
              AND t.status = 'SENT'
            ORDER BY t.team_number
            LIMIT 1
            """
        ),
        {
            "batch_id": int(batch_id),
            "user_id": int(user_id),
        },
    ).mappings().first()

    if team is None:
        return dict(batch), None, [], []

    members = [
        dict(row)
        for row in db.execute(
            text(
                """
                SELECT
                    tm.user_id,
                    tm.level_code,
                    tm.order_number,
                    u.full_name,
                    COALESCE(s.name, '') AS school_name
                FROM survey_investigation_team_members tm
                JOIN users u ON u.id = tm.user_id
                LEFT JOIN schools s ON s.id = tm.school_id
                WHERE tm.team_id = :team_id
                  AND u.is_active = 1
                ORDER BY tm.order_number
                """
            ),
            {"team_id": int(team["id"])},
        ).mappings().all()
    ]

    labels = {
        "THON": "Thôn",
        "XOM": "Xóm",
        "KHOI": "Khối",
        "TO": "Tổ",
        "BAN": "Bản",
    }

    assignment_rows = db.execute(
        text(
            """
            SELECT
                ta.area_id,
                ta.household_quota,
                a.name AS area_name,
                a.area_type,
                COALESCE(
                    (
                        SELECT COUNT(*)
                        FROM survey_investigation_team_forms tf
                        JOIN survey_form_areas fa
                          ON fa.survey_form_id = tf.survey_form_id
                        WHERE tf.team_id = ta.team_id
                          AND fa.area_id = ta.area_id
                    ),
                    0
                ) AS actual_count
            FROM survey_team_area_assignments ta
            JOIN survey_commune_areas a
              ON a.id = ta.area_id
            WHERE ta.team_id = :team_id
              AND a.is_active = 1
              AND ta.household_quota > 0
            ORDER BY a.name COLLATE NOCASE, a.id
            """
        ),
        {"team_id": int(team["id"])},
    ).mappings().all()

    # === BAI_13B_12_V2_3_TEACHER_AREA_SCOPE_START ===
    assignments = []
    for row in assignment_rows:
        item = dict(row)
        item["area_type_label"] = labels.get(
            str(item.get("area_type") or ""),
            str(item.get("area_type") or ""),
        )
        difference = (
            int(item["household_quota"] or 0)
            - int(item["actual_count"] or 0)
        )
        item["remaining_count"] = max(0, difference)
        item["excess_count"] = max(0, -difference)
        # Số hộ dự kiến/chỉ tiêu là kế hoạch, KHÔNG phải trần cấm nhập.
        assignments.append(item)

    # Tương thích luồng TEST/legacy: trước V2.2 có thể đã chia phiếu cho
    # tổ rồi mới khai báo địa bàn. Nếu chưa có bảng giao địa bàn cho tổ,
    # suy ra phạm vi của tổ từ các phiếu hiện có đã nối survey_form_areas.
    if not assignments:
        legacy_rows = db.execute(
            text(
                """
                SELECT
                    a.id AS area_id,
                    a.name AS area_name,
                    a.area_type,
                    COUNT(tf.survey_form_id) AS actual_count
                FROM survey_investigation_team_forms tf
                JOIN survey_form_areas fa
                  ON fa.survey_form_id = tf.survey_form_id
                JOIN survey_commune_areas a
                  ON a.id = fa.area_id
                WHERE tf.team_id = :team_id
                  AND a.is_active = 1
                GROUP BY a.id, a.name, a.area_type
                ORDER BY a.name COLLATE NOCASE, a.id
                """
            ),
            {"team_id": int(team["id"])},
        ).mappings().all()

        for row in legacy_rows:
            item = dict(row)
            actual_count = int(item.get("actual_count") or 0)
            item["household_quota"] = actual_count
            item["remaining_count"] = 0
            item["excess_count"] = 0
            item["area_type_label"] = labels.get(
                str(item.get("area_type") or ""),
                str(item.get("area_type") or ""),
            )
            assignments.append(item)

    return dict(batch), dict(team), members, assignments
    # === BAI_13B_12_V2_3_TEACHER_AREA_SCOPE_END ===


@router.get("/nhap-ho-moi/{batch_id}/quyen")
def b1322_teacher_new_house_permission(
    batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != TEACHER_ROLE_CODE:
        return _b1322_JSONResponse({"can_create": False})

    _b1322_ensure_area_schema(db)
    user = _user(request)
    user_id = _b1322_int(user.get("id"))
    if not user_id:
        return _b1322_JSONResponse({"can_create": False})

    batch, team, members, assignments = (
        _b1322_team_context_for_teacher(
            db,
            batch_id=int(batch_id),
            user_id=user_id,
        )
    )

    # === BAI_13B_12_V2_3_4_PERMISSION_BUTTON_START ===
    # Hiển thị nút khi giáo viên thuộc tổ 3 người đã SENT và đợt còn hoạt động.
    # Không dùng "đã có giao địa bàn" làm điều kiện ẨN nút.
    can_create = bool(
        batch
        and not bool(batch.get("is_locked"))
        and str(batch.get("status") or "") != "DA_KET_THUC"
        and team
        and len(members) == 3
    )
    # === BAI_13B_12_V2_3_4_PERMISSION_BUTTON_END ===

    return _b1322_JSONResponse(
        {
            "can_create": can_create,
            "team_number": (
                int(team["team_number"]) if team else None
            ),
            "remaining_quota": sum(
                int(item.get("remaining_count") or 0)
                for item in assignments
            ),
            "excess_households": sum(
                int(item.get("excess_count") or 0)
                for item in assignments
            ),
        }
    )


@router.get(
    "/nhap-ho-moi/{batch_id}",
    response_class=HTMLResponse,
)
def b1322_teacher_new_house_page(
    batch_id: int,
    request: Request,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    if _role(request) != TEACHER_ROLE_CODE:
        return _forbidden()

    _b1322_ensure_area_schema(db)
    user = _user(request)
    user_id = _b1322_int(user.get("id"))
    batch, team, members, assignments = (
        _b1322_team_context_for_teacher(
            db,
            batch_id=int(batch_id),
            user_id=user_id,
        )
    )

    if (
        not batch
        or bool(batch.get("is_locked"))
        or str(batch.get("status") or "") == "DA_KET_THUC"
        or team is None
        or len(members) != 3
    ):
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    # Dùng dict lồng nhau để template vẫn truy cập batch.school_year.code.
    batch_view = dict(batch)
    batch_view["school_year"] = {
        "code": batch_view.get("school_year_code")
    }

    messages = {
        "invalid": "Thông tin hộ hoặc địa bàn chưa hợp lệ.",
        "quota_full": "Chỉ tiêu dự kiến đã đạt; V2.3 vẫn cho phép nhập đủ hộ thực tế.",
        "invalid_area": "Địa bàn này chưa thuộc phạm vi được giao cho tổ.",
        "save_error": "Không lưu được hộ mới. Dữ liệu đã được rollback.",
    }

    return templates.TemplateResponse(
        request=request,
        name="surveys/new_household_team_v2_2.html",
        context={
            "nguoi_dung": user,
            "batch": batch_view,
            "team": team,
            "members": members,
            "assignments": assignments,
            "message": messages.get(str(status or "")),
        },
    )


@router.post("/nhap-ho-moi/{batch_id}/them")
async def b1322_teacher_create_new_house(
    batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) != TEACHER_ROLE_CODE:
        return _forbidden()

    _b1322_ensure_area_schema(db)
    user = _user(request)
    user_id = _b1322_int(user.get("id"))

    batch, team, members, assignments = (
        _b1322_team_context_for_teacher(
            db,
            batch_id=int(batch_id),
            user_id=user_id,
        )
    )

    if (
        not batch
        or bool(batch.get("is_locked"))
        or str(batch.get("status") or "") == "DA_KET_THUC"
        or team is None
        or len(members) != 3
    ):
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/ho-dan",
            status_code=303,
        )

    form = await request.form()
    area_id = _b1322_int(form.get("area_id"))
    head_name = " ".join(
        str(form.get("head_name") or "").strip().split()
    )
    address = " ".join(
        str(form.get("address") or "").strip().split()
    )
    phone = " ".join(
        str(form.get("phone") or "").strip().split()
    )
    notes = str(form.get("notes") or "").strip()

    if (
        not area_id
        or not head_name
        or len(head_name) > 200
        or not address
        or len(phone) > 30
    ):
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/"
                f"nhap-ho-moi/{batch_id}?status=invalid"
            ),
            status_code=303,
        )

    assignment = next(
        (
            item for item in assignments
            if int(item["area_id"]) == int(area_id)
        ),
        None,
    )
    if assignment is None:
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/"
                f"nhap-ho-moi/{batch_id}?status=invalid_area"
            ),
            status_code=303,
        )

    # === BAI_13B_12_V2_3_NO_HOUSEHOLD_CAP_START ===
    # Không chặn khi đạt/vượt số hộ dự kiến. Điều tra PCGD/XMC phải thu đủ
    # hộ thực tế phát hiện ngoài địa bàn.
    # === BAI_13B_12_V2_3_NO_HOUSEHOLD_CAP_END ===

    now = _b1322_datetime.now()
    temp_house = (
        "TMP-HO-" + _b1322_uuid4().hex[:20].upper()
    )
    temp_form = (
        "TMP-PH-" + _b1322_uuid4().hex[:20].upper()
    )

    try:
        cursor = db.execute(
            text(
                """
                INSERT INTO households (
                    code,
                    commune_id,
                    head_name,
                    hamlet_name,
                    address,
                    phone,
                    notes,
                    is_active,
                    created_at,
                    updated_at
                ) VALUES (
                    :code,
                    :commune_id,
                    :head_name,
                    :hamlet_name,
                    :address,
                    :phone,
                    :notes,
                    1,
                    :created_at,
                    :updated_at
                )
                """
            ),
            {
                "code": temp_house,
                "commune_id": int(batch["commune_id"]),
                "head_name": head_name,
                "hamlet_name": str(assignment["area_name"]),
                "address": address,
                "phone": phone or None,
                "notes": notes or None,
                "created_at": now,
                "updated_at": now,
            },
        )
        household_id = int(cursor.lastrowid)

        household_code = (
            f"HO-{int(batch['commune_id']):05d}-"
            f"{household_id:08d}"
        )
        db.execute(
            text(
                """
                UPDATE households
                SET code = :code
                WHERE id = :household_id
                """
            ),
            {
                "code": household_code,
                "household_id": household_id,
            },
        )

        cursor = db.execute(
            text(
                """
                INSERT INTO survey_forms (
                    survey_batch_id,
                    household_id,
                    form_number,
                    head_name_snapshot,
                    address_snapshot,
                    hamlet_name_snapshot,
                    survey_date,
                    status,
                    village_head_name,
                    household_representative_name,
                    household_confirmed_at,
                    commune_confirmed_at,
                    notes,
                    created_at,
                    updated_at
                ) VALUES (
                    :survey_batch_id,
                    :household_id,
                    :form_number,
                    :head_name,
                    :address,
                    :hamlet_name,
                    NULL,
                    'CHUA_DIEU_TRA',
                    NULL,
                    NULL,
                    NULL,
                    NULL,
                    :notes,
                    :created_at,
                    :updated_at
                )
                """
            ),
            {
                "survey_batch_id": int(batch_id),
                "household_id": household_id,
                "form_number": temp_form,
                "head_name": head_name,
                "address": address,
                "hamlet_name": str(assignment["area_name"]),
                "notes": (
                    "[V2.2] Hộ được tổ điều tra tạo mới "
                    "trực tiếp tại thực địa."
                ),
                "created_at": now,
                "updated_at": now,
            },
        )
        survey_form_id = int(cursor.lastrowid)

        form_number = (
            f"PH-{int(batch_id):06d}-"
            f"{survey_form_id:08d}"
        )
        db.execute(
            text(
                """
                UPDATE survey_forms
                SET form_number = :form_number
                WHERE id = :survey_form_id
                """
            ),
            {
                "form_number": form_number,
                "survey_form_id": survey_form_id,
            },
        )

        db.execute(
            text(
                """
                INSERT INTO survey_form_areas (
                    survey_form_id,
                    area_id,
                    created_at
                ) VALUES (
                    :survey_form_id,
                    :area_id,
                    :created_at
                )
                """
            ),
            {
                "survey_form_id": survey_form_id,
                "area_id": int(area_id),
                "created_at": now,
            },
        )

        next_order = int(
            db.execute(
                text(
                    """
                    SELECT COALESCE(MAX(assignment_order), 0) + 1
                    FROM survey_investigation_team_forms
                    WHERE team_id = :team_id
                    """
                ),
                {"team_id": int(team["id"])},
            ).scalar()
            or 1
        )

        db.execute(
            text(
                """
                INSERT INTO survey_investigation_team_forms (
                    team_id,
                    survey_form_id,
                    assignment_order,
                    created_at
                ) VALUES (
                    :team_id,
                    :survey_form_id,
                    :assignment_order,
                    :created_at
                )
                """
            ),
            {
                "team_id": int(team["id"]),
                "survey_form_id": survey_form_id,
                "assignment_order": next_order,
                "created_at": now,
            },
        )

        for member in members:
            order_number = int(member["order_number"])
            db.execute(
                text(
                    """
                    INSERT INTO survey_form_investigators (
                        survey_form_id,
                        user_id,
                        order_number,
                        is_primary,
                        signed_at,
                        notes,
                        created_at
                    ) VALUES (
                        :survey_form_id,
                        :user_id,
                        :order_number,
                        :is_primary,
                        :signed_at,
                        :notes,
                        :created_at
                    )
                    """
                ),
                {
                    "survey_form_id": survey_form_id,
                    "user_id": int(member["user_id"]),
                    "order_number": order_number,
                    "is_primary": 1 if order_number == 1 else 0,
                    "signed_at": now,
                    "notes": (
                        f"Tổ điều tra 3 cấp số "
                        f"{int(team['team_number'])}; "
                        f"{member['level_code']}; "
                        "hộ tạo mới tại thực địa."
                    ),
                    "created_at": now,
                },
            )

        # === BAI_13B_12_V2_3_ALLOW_OVER_EXPECTED_START ===
        # Không rollback khi số hộ thực tế vượt kế hoạch.
        # === BAI_13B_12_V2_3_ALLOW_OVER_EXPECTED_END ===

        db.commit()

    except Exception:
        db.rollback()
        return RedirectResponse(
            url=(
                "/dieu-tra/phan-cong-to-dieu-tra/"
                f"nhap-ho-moi/{batch_id}?status=save_error"
            ),
            status_code=303,
        )

    return RedirectResponse(
        url=(
            f"/dieu-tra/{batch_id}/ho-dan/"
            f"{household_id}/nhap-nhanh"
            "?status=household_created"
        ),
        status_code=303,
    )
# ============================================================
# BAI_13B_12_V2_2_TEAM_NEW_HOUSEHOLD_END
# ============================================================
