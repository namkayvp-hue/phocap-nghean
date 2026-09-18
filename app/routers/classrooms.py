from __future__ import annotations

from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
from typing import Annotated, Any
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, File, Form, Query, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from fastapi.templating import Jinja2Templates
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from app.database import get_db
from app.models import Classroom, Commune, School, SchoolYear, StudentEnrollment
from app.permissions import (
    ADMIN_ROLE_CODES,
    COMMUNE_ROLE_CODE,
    DEPARTMENT_ROLE_CODE,
    SCHOOL_ROLE_CODE,
    normalize_role_code,
)
from app.report_input_models import SchoolNetworkYearData
from app.survey_models import SurveyBatch, SurveyPersonYearRecord
from app.survey_workflow_models import (
    SurveyCommuneExecutionState,
    SurveySchoolAssignment,
)
from app.thpt_reference_models import THPTSchoolReference


APP_DIR = Path(__file__).resolve().parent.parent
templates = Jinja2Templates(directory=str(APP_DIR / "templates"))
router = APIRouter(prefix="/lop-hoc", tags=["Danh mục lớp học"])


LEVEL_LABELS = {
    "MN": "Mầm non",
    "TH": "Tiểu học",
    "THCS": "THCS",
    "THPT": "THPT",
}

STATUS_MESSAGES = {
    "created": "Đã thêm lớp vào danh mục.",
    "updated": "Đã cập nhật lớp.",
    "locked": "Đã khóa lớp. Dữ liệu cũ vẫn được giữ nguyên.",
    "unlocked": "Đã mở lại lớp.",
    "copied": "Đã sao chép danh mục lớp từ năm học trước.",
    "duplicate": "Tên lớp đã tồn tại trong đúng trường và năm học.",
    "invalid": "Thông tin lớp chưa hợp lệ.",
    "invalid_scope": "Trường không thuộc phạm vi quản lý của tài khoản.",
    "invalid_level": "Trường không thuộc cấp học đang chọn.",
    "survey_locked": "Dữ liệu của năm học/đơn vị đã bị khóa; không thể thay đổi danh mục lớp.",
    "forbidden": "Tài khoản không có quyền quản lý danh mục lớp.",
    "no_previous": "Không có danh mục lớp năm trước để sao chép.",
}


READ_ROLES = frozenset({
    *ADMIN_ROLE_CODES,
    DEPARTMENT_ROLE_CODE,
    COMMUNE_ROLE_CODE,
    SCHOOL_ROLE_CODE,
})

WRITE_ROLES = frozenset({
    *ADMIN_ROLE_CODES,
    COMMUNE_ROLE_CODE,
    SCHOOL_ROLE_CODE,
})


def actor_of(request: Request) -> dict[str, Any]:
    return dict(request.scope.get("auth_user") or {})


def role_of(request: Request) -> str:
    return normalize_role_code(actor_of(request).get("role_code"))


def normalize_level(value: str | None) -> str | None:
    code = str(value or "").strip().upper()
    return code if code in LEVEL_LABELS else None


def normalize_name(value: str) -> str:
    return " ".join(str(value or "").strip().split())


def normalize_code(value: str) -> str | None:
    code = "".join(str(value or "").strip().split()).upper()
    return code or None


def allowed_school(request: Request, school: School | None) -> bool:
    if school is None:
        return False

    actor = actor_of(request)
    role = role_of(request)

    if role in ADMIN_ROLE_CODES or role == DEPARTMENT_ROLE_CODE:
        return True

    if role == COMMUNE_ROLE_CODE:
        commune_id = actor.get("commune_id")
        return commune_id is not None and int(commune_id) == int(school.commune_id)

    if role == SCHOOL_ROLE_CODE:
        school_id = actor.get("school_id")
        return school_id is not None and int(school_id) == int(school.id)

    return False


def can_write_school(request: Request, school: School | None) -> bool:
    return role_of(request) in WRITE_ROLES and allowed_school(request, school)


def level_school_ids(
    db: Session,
    level_code: str | None,
    school_year_id: int | None,
) -> set[int] | None:
    """Trả về tập trường thuộc cấp; None nghĩa là không lọc cấp."""

    if level_code is None:
        return None

    if level_code == "THPT":
        return {
            int(item)
            for item in db.scalars(
                select(THPTSchoolReference.official_school_id)
                .where(THPTSchoolReference.official_school_id.is_not(None))
                .distinct()
            ).all()
            if item is not None
        }

    return {
        school_id
        for school_id, levels in network_levels_by_school(db, school_year_id).items()
        if level_code in levels
    }


def network_levels_by_school(
    db: Session, school_year_id: int | None,
) -> dict[int, set[str]]:
    """D?ng n?m g?n nh?t c?a t?ng tr??ng, kh?ng l?y ph?n lo?i t? n?m t??ng lai."""
    statement = (
        select(
            SchoolNetworkYearData.school_id,
            SchoolNetworkYearData.level_code,
            SchoolYear.code,
        )
        .join(SchoolYear, SchoolYear.id == SchoolNetworkYearData.school_year_id)
        .order_by(SchoolYear.code.desc())
    )
    year = db.get(SchoolYear, school_year_id) if school_year_id else None
    if year is not None:
        statement = statement.where(SchoolYear.code <= year.code)
    latest_year: dict[int, str] = {}
    result: dict[int, set[str]] = {}
    for school_id, level, year_code in db.execute(statement):
        latest_year.setdefault(school_id, year_code)
        if year_code == latest_year[school_id]:
            result.setdefault(school_id, set()).add(str(level).strip().upper())
    return result


def school_levels(
    db: Session, school_id: int, school_year_id: int | None = None,
) -> list[str]:
    levels = network_levels_by_school(db, school_year_id).get(school_id, set())
    is_thpt = db.scalar(
        select(THPTSchoolReference.id).where(
            THPTSchoolReference.official_school_id == school_id
        ).limit(1)
    )
    if is_thpt is not None:
        levels.add("THPT")
    return [code for code in LEVEL_LABELS if code in levels]


def load_years(db: Session) -> list[SchoolYear]:
    return list(
        db.scalars(
            select(SchoolYear)
            .where(SchoolYear.is_active.is_(True))
            .order_by(SchoolYear.code.desc(), SchoolYear.id.desc())
        ).all()
    )


def load_communes(request: Request, db: Session) -> list[Commune]:
    actor = actor_of(request)
    role = role_of(request)

    statement = select(Commune).where(Commune.is_active.is_(True))
    if role == COMMUNE_ROLE_CODE and actor.get("commune_id") is not None:
        statement = statement.where(Commune.id == int(actor["commune_id"]))
    elif role == SCHOOL_ROLE_CODE and actor.get("school_id") is not None:
        school = db.get(School, int(actor["school_id"]))
        if school is None:
            return []
        statement = statement.where(Commune.id == int(school.commune_id))
    elif role not in {*ADMIN_ROLE_CODES, DEPARTMENT_ROLE_CODE}:
        return []

    return list(db.scalars(statement.order_by(Commune.name.asc())).all())


def load_schools(
    request: Request,
    db: Session,
    *,
    level_code: str | None,
    school_year_id: int | None,
    commune_id: int | None,
) -> list[School]:
    actor = actor_of(request)
    role = role_of(request)

    statement = (
        select(School)
        .options(selectinload(School.commune))
        .where(School.is_active.is_(True))
    )

    if role == COMMUNE_ROLE_CODE:
        actor_commune = actor.get("commune_id")
        if actor_commune is None:
            return []
        statement = statement.where(School.commune_id == int(actor_commune))
    elif role == SCHOOL_ROLE_CODE:
        actor_school = actor.get("school_id")
        if actor_school is None:
            return []
        statement = statement.where(School.id == int(actor_school))
    elif role not in {*ADMIN_ROLE_CODES, DEPARTMENT_ROLE_CODE}:
        return []

    if commune_id is not None:
        statement = statement.where(School.commune_id == commune_id)

    ids = level_school_ids(db, level_code, school_year_id)
    if ids is not None:
        if not ids:
            return []
        statement = statement.where(School.id.in_(ids))

    return list(db.scalars(statement.order_by(School.name.asc())).all())


def class_catalog_lock_message(
    db: Session,
    *,
    school_year_id: int,
    school: School,
) -> str | None:
    """Khóa danh mục lớp cùng với khóa dữ liệu điều tra của năm học."""

    batches = list(
        db.scalars(
            select(SurveyBatch)
            .where(
                SurveyBatch.school_year_id == school_year_id,
                SurveyBatch.commune_id == school.commune_id,
            )
            .order_by(SurveyBatch.id.desc())
        ).all()
    )

    for batch in batches:
        state = db.scalar(
            select(SurveyCommuneExecutionState).where(
                SurveyCommuneExecutionState.survey_batch_id == batch.id
            )
        )

        if state is not None and bool(state.is_province_locked):
            return "Cấp Sở đã khóa năm học/đợt điều tra."

        if bool(batch.is_locked) or (
            state is not None and bool(state.is_commune_locked)
        ):
            return "Xã/phường đã khóa dữ liệu điều tra."

        assignment = db.scalar(
            select(SurveySchoolAssignment).where(
                SurveySchoolAssignment.survey_batch_id == batch.id,
                SurveySchoolAssignment.school_id == school.id,
            )
        )
        if assignment is not None and bool(assignment.is_locked):
            return "Trường đã được khóa trong đợt điều tra."

    return None


def redirect_list(
    *,
    school_year_id: int | None,
    level_code: str | None,
    commune_id: int | None,
    school_id: int | None,
    status: str,
    extra: dict[str, int | str] | None = None,
) -> RedirectResponse:
    params: dict[str, int | str] = {"status": status}
    if school_year_id is not None:
        params["school_year_id"] = school_year_id
    if level_code:
        params["cap"] = level_code
    if commune_id is not None:
        params["commune_id"] = commune_id
    if school_id is not None:
        params["school_id"] = school_id
    if extra:
        params.update(extra)
    return RedirectResponse(f"/lop-hoc?{urlencode(params)}", status_code=303)


def validate_target(
    request: Request,
    db: Session,
    *,
    school_year_id: int,
    school_id: int,
    level_code: str | None,
) -> tuple[SchoolYear | None, School | None, str | None]:
    year = db.get(SchoolYear, school_year_id)
    school = db.get(School, school_id)

    if year is None or not year.is_active or school is None or not school.is_active:
        return year, school, "invalid"

    if not can_write_school(request, school):
        return year, school, "invalid_scope"

    ids = level_school_ids(db, level_code, school_year_id)
    if ids is not None and school.id not in ids:
        return year, school, "invalid_level"

    if class_catalog_lock_message(
        db,
        school_year_id=school_year_id,
        school=school,
    ):
        return year, school, "survey_locked"

    return year, school, None


@router.get("", response_class=HTMLResponse)
def list_classes(
    request: Request,
    school_year_id: int | None = Query(default=None),
    cap: str | None = Query(default=None),
    commune_id: int | None = Query(default=None),
    school_id: int | None = Query(default=None),
    q: str = Query(default="", max_length=100),
    status: str | None = Query(default=None),
    created: int = Query(default=0, ge=0),
    reactivated: int = Query(default=0, ge=0),
    skipped: int = Query(default=0, ge=0),
    copied: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
):
    role = role_of(request)
    actor = actor_of(request)
    if role not in READ_ROLES:
        return RedirectResponse("/?status=forbidden", 303)

    years = load_years(db)
    valid_year_ids = {int(item.id) for item in years}
    if school_year_id not in valid_year_ids:
        school_year_id = int(years[0].id) if years else None

    level_code = normalize_level(cap)

    if role == SCHOOL_ROLE_CODE and actor.get("school_id") is not None:
        school_id = int(actor["school_id"])
        own_levels = school_levels(db, school_id, school_year_id)
        if own_levels and level_code not in own_levels:
            level_code = own_levels[0]
    elif level_code is None:
        level_code = "MN"

    communes = load_communes(request, db)
    valid_commune_ids = {int(item.id) for item in communes}

    if role == COMMUNE_ROLE_CODE and actor.get("commune_id") is not None:
        commune_id = int(actor["commune_id"])
    elif role == SCHOOL_ROLE_CODE and actor.get("school_id") is not None:
        own_school = db.get(School, int(actor["school_id"]))
        commune_id = int(own_school.commune_id) if own_school else None
    elif commune_id not in valid_commune_ids:
        commune_id = None

    schools = load_schools(
        request,
        db,
        level_code=level_code,
        school_year_id=school_year_id,
        commune_id=commune_id,
    )
    valid_school_ids = {int(item.id) for item in schools}
    if school_id not in valid_school_ids:
        if role == SCHOOL_ROLE_CODE and schools:
            school_id = int(schools[0].id)
        else:
            school_id = None

    selected_school = db.get(School, school_id) if school_id else None
    selected_year = db.get(SchoolYear, school_year_id) if school_year_id else None

    keyword = normalize_name(q)[:100]
    classes: list[Classroom] = []
    usage_by_class: dict[int, dict[str, int]] = {}

    if selected_school is not None and selected_year is not None:
        statement = select(Classroom).where(
            Classroom.school_id == selected_school.id,
            Classroom.school_year_id == selected_year.id,
        )
        classes = list(
            db.scalars(
                statement.order_by(
                    Classroom.is_active.desc(),
                    Classroom.name.asc(),
                    Classroom.id.asc(),
                )
            ).all()
        )
        if keyword:
            keyword_key = keyword.casefold()
            classes = [
                item for item in classes
                if keyword_key in normalize_name(item.name).casefold()
            ]

        for classroom in classes:
            enrollment_count = int(
                db.scalar(
                    select(func.count(StudentEnrollment.id)).where(
                        StudentEnrollment.class_id == classroom.id
                    )
                )
                or 0
            )
            survey_count = int(
                db.scalar(
                    select(func.count(SurveyPersonYearRecord.id)).where(
                        SurveyPersonYearRecord.class_id == classroom.id
                    )
                )
                or 0
            )
            usage_by_class[classroom.id] = {
                "students": enrollment_count,
                "surveys": survey_count,
                "total": enrollment_count + survey_count,
            }

    lock_message = None
    if selected_school is not None and selected_year is not None:
        lock_message = class_catalog_lock_message(
            db,
            school_year_id=selected_year.id,
            school=selected_school,
        )

    can_manage = (
        selected_school is not None
        and role in WRITE_ROLES
        and can_write_school(request, selected_school)
        and lock_message is None
    )

    thpt_unlinked = 0
    if level_code == "THPT":
        thpt_unlinked = int(
            db.scalar(
                select(func.count(THPTSchoolReference.id)).where(
                    THPTSchoolReference.official_school_id.is_(None)
                )
            )
            or 0
        )

    detail_parts: list[str] = []
    if created:
        detail_parts.append(f"Thêm mới {created} lớp")
    if reactivated:
        detail_parts.append(f"mở lại {reactivated} lớp")
    if skipped:
        detail_parts.append(f"bỏ qua {skipped} lớp đã tồn tại")
    if copied:
        detail_parts.append(f"sao chép {copied} lớp")

    return templates.TemplateResponse(
        request=request,
        name="classrooms/list.html",
        context={
            "nguoi_dung": actor,
            "levels": (
                {code: LEVEL_LABELS[code] for code in own_levels}
                if role == SCHOOL_ROLE_CODE and actor.get("school_id") is not None and own_levels
                else LEVEL_LABELS
            ),
            "selected_level": level_code,
            "selected_level_label": LEVEL_LABELS.get(level_code, "Tất cả cấp"),
            "school_years": years,
            "selected_school_year_id": school_year_id,
            "selected_year": selected_year,
            "communes": communes,
            "selected_commune_id": commune_id,
            "schools": schools,
            "selected_school_id": school_id,
            "selected_school": selected_school,
            "classes": classes,
            "usage_by_class": usage_by_class,
            "q": keyword,
            "status": status,
            "message": STATUS_MESSAGES.get(status),
            "detail_message": "; ".join(detail_parts),
            "can_manage": can_manage,
            "lock_message": lock_message,
            "active_count": sum(1 for item in classes if item.is_active),
            "inactive_count": sum(1 for item in classes if not item.is_active),
            "thpt_unlinked": thpt_unlinked,
        },
    )


def read_class_excel(content: bytes) -> list[str]:
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("File Excel không được vượt quá 5 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError("File Excel có dữ liệu giải nén quá lớn.")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Không đọc được file. Hãy dùng file .xlsx theo mẫu.") from exc
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(max_row=1002, max_col=20)
        header = next(rows, ())
        columns = [i for i, cell in enumerate(header)
                   if normalize_name(str(cell.value or "")).casefold() == "tên lớp"]
        if len(columns) != 1:
            raise ValueError("Dòng đầu phải có đúng một cột Tên lớp. Hãy tải file mẫu.")
        column = columns[0]
        names: list[str] = []
        seen: set[str] = set()
        for number, row in enumerate(rows, start=2):
            cell = row[column]
            if cell.value is None or not str(cell.value).strip():
                continue
            if number > 1001:
                raise ValueError("File chỉ được chứa tối đa 1.000 dòng dữ liệu.")
            name = normalize_name(str(cell.value))
            if cell.data_type in {"f", "e"} or len(name) > 200 or ";" in name:
                raise ValueError(f"Dòng {number}: tên lớp không hợp lệ; không dùng công thức hoặc dấu chấm phẩy.")
            if name.casefold() not in seen:
                names.append(name)
                seen.add(name.casefold())
        if sheet.max_row and sheet.max_row > 1002:
            raise ValueError("File chỉ được chứa tối đa 1.000 dòng dữ liệu.")
        if not names or len(names) > 200:
            raise ValueError("File phải có từ 1 đến 200 tên lớp khác nhau.")
        return names
    finally:
        workbook.close()


@router.get("/mau-excel")
def class_excel_template(request: Request, cap: str = "TH"):
    if role_of(request) not in READ_ROLES:
        return RedirectResponse("/?status=forbidden", 303)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Danh sách lớp"
    sheet.append(["STT", "Tên lớp"])
    examples = {"MN": ["3 tuổi A", "4 tuổi A"], "THCS": ["6A", "6B"], "THPT": ["10A1", "10A2"]}
    for index, name in enumerate(examples.get(cap, ["1A", "1B"]), 1):
        sheet.append([index, name])
    sheet.column_dimensions["A"].width = 10
    sheet.column_dimensions["B"].width = 35
    sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="mau_danh_sach_lop.xlsx"'})


@router.post("/nhap-excel", response_class=HTMLResponse)
def preview_class_excel(
    request: Request,
    school_year_id: Annotated[int, Form()],
    school_id: Annotated[int, Form()],
    file: Annotated[UploadFile, File()],
    cap: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    level = normalize_level(cap)
    year, school, error = validate_target(request, db, school_year_id=school_year_id,
                                           school_id=school_id, level_code=level)
    if error:
        file.file.close()
        return redirect_list(school_year_id=school_year_id, school_id=school_id,
                             commune_id=school.commune_id if school else None,
                             level_code=level, status=error)
    names = []
    message = ""
    try:
        if not str(file.filename or "").lower().endswith(".xlsx"):
            raise ValueError("Vui lòng chọn file Excel .xlsx.")
        names = read_class_excel(file.file.read(5 * 1024 * 1024 + 1))
    except ValueError as exc:
        message = str(exc)
    finally:
        file.file.close()
    existing = {normalize_name(item.name).casefold(): item for item in db.scalars(
        select(Classroom).where(Classroom.school_id == school_id, Classroom.school_year_id == school_year_id)
    )}
    rows = []
    for name in names:
        item = existing.get(name.casefold())
        action = "Thêm mới" if item is None else ("Bỏ qua: đã tồn tại" if item.is_active else "Mở lại lớp đã khóa")
        rows.append({"name": name, "action": action})
    return templates.TemplateResponse(request=request, name="classrooms/import_preview.html",
        context={"nguoi_dung": actor_of(request), "school": school, "year": year,
                 "cap": level or "", "rows": rows, "names": "\n".join(names), "error": message},
        status_code=400 if message else 200)


@router.post("/them")
def create_classes(
    request: Request,
    school_year_id: Annotated[int, Form()],
    school_id: Annotated[int, Form()],
    ten_lop: Annotated[str, Form()],
    cap: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    level_code = normalize_level(cap)
    year, school, error = validate_target(
        request,
        db,
        school_year_id=school_year_id,
        school_id=school_id,
        level_code=level_code,
    )
    commune_id = int(school.commune_id) if school else None
    if error:
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status=error,
        )

    names: list[str] = []
    seen: set[str] = set()
    for raw_line in str(ten_lop or "").replace(";", "\n").splitlines():
        name = normalize_name(raw_line)
        key = name.casefold()
        if not name or key in seen:
            continue
        if len(name) > 200:
            return redirect_list(
                school_year_id=school_year_id,
                level_code=level_code,
                commune_id=commune_id,
                school_id=school_id,
                status="invalid",
            )
        names.append(name)
        seen.add(key)

    if not names or len(names) > 200:
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status="invalid",
        )

    created_count = 0
    reactivated_count = 0
    skipped_count = 0

    existing_rows = list(
        db.scalars(
            select(Classroom).where(
                Classroom.school_id == school_id,
                Classroom.school_year_id == school_year_id,
            )
        ).all()
    )
    existing_by_name = {
        normalize_name(item.name).casefold(): item
        for item in existing_rows
    }

    try:
        for name in names:
            key = name.casefold()
            existing = existing_by_name.get(key)
            if existing is None:
                classroom = Classroom(
                    school_id=school_id,
                    school_year_id=school_year_id,
                    code=None,
                    name=name,
                    is_active=True,
                )
                db.add(classroom)
                existing_by_name[key] = classroom
                created_count += 1
            elif not existing.is_active:
                existing.is_active = True
                reactivated_count += 1
            else:
                skipped_count += 1
        db.commit()
    except IntegrityError:
        db.rollback()
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status="duplicate",
        )

    return redirect_list(
        school_year_id=school_year_id,
        level_code=level_code,
        commune_id=commune_id,
        school_id=school_id,
        status="created",
        extra={
            "created": created_count,
            "reactivated": reactivated_count,
            "skipped": skipped_count,
        },
    )


@router.post("/{class_id}/cap-nhat")
def update_class(
    class_id: int,
    request: Request,
    ten_lop: Annotated[str, Form()],
    school_year_id: Annotated[int, Form()],
    school_id: Annotated[int, Form()],
    cap: Annotated[str, Form()] = "",
    ma_lop: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    level_code = normalize_level(cap)
    year, school, error = validate_target(
        request,
        db,
        school_year_id=school_year_id,
        school_id=school_id,
        level_code=level_code,
    )
    commune_id = int(school.commune_id) if school else None
    if error:
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status=error,
        )

    classroom = db.get(Classroom, class_id)
    name = normalize_name(ten_lop)
    code = normalize_code(ma_lop)

    if (
        classroom is None
        or classroom.school_id != school_id
        or classroom.school_year_id != school_year_id
        or not name
        or len(name) > 200
        or (code is not None and len(code) > 50)
    ):
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status="invalid",
        )

    duplicate = next(
        (
            item
            for item in db.scalars(
                select(Classroom).where(
                    Classroom.school_id == school_id,
                    Classroom.school_year_id == school_year_id,
                    Classroom.id != classroom.id,
                )
            ).all()
            if normalize_name(item.name).casefold() == name.casefold()
        ),
        None,
    )
    if duplicate is not None:
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status="duplicate",
        )

    classroom.name = name
    classroom.code = code
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status="duplicate",
        )

    return redirect_list(
        school_year_id=school_year_id,
        level_code=level_code,
        commune_id=commune_id,
        school_id=school_id,
        status="updated",
    )


@router.post("/{class_id}/khoa-mo")
def toggle_class(
    class_id: int,
    request: Request,
    school_year_id: Annotated[int, Form()],
    school_id: Annotated[int, Form()],
    cap: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    level_code = normalize_level(cap)
    year, school, error = validate_target(
        request,
        db,
        school_year_id=school_year_id,
        school_id=school_id,
        level_code=level_code,
    )
    commune_id = int(school.commune_id) if school else None
    if error:
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status=error,
        )

    classroom = db.get(Classroom, class_id)
    if (
        classroom is None
        or classroom.school_id != school_id
        or classroom.school_year_id != school_year_id
    ):
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status="invalid",
        )

    classroom.is_active = not bool(classroom.is_active)
    db.commit()

    return redirect_list(
        school_year_id=school_year_id,
        level_code=level_code,
        commune_id=commune_id,
        school_id=school_id,
        status="unlocked" if classroom.is_active else "locked",
    )


@router.post("/sao-chep-nam-truoc")
def copy_previous_year(
    request: Request,
    school_year_id: Annotated[int, Form()],
    school_id: Annotated[int, Form()],
    cap: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    level_code = normalize_level(cap)
    year, school, error = validate_target(
        request,
        db,
        school_year_id=school_year_id,
        school_id=school_id,
        level_code=level_code,
    )
    commune_id = int(school.commune_id) if school else None
    if error:
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status=error,
        )

    previous_year = db.scalar(
        select(SchoolYear)
        .where(SchoolYear.code < year.code)
        .order_by(SchoolYear.code.desc(), SchoolYear.id.desc())
        .limit(1)
    )
    if previous_year is None:
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status="no_previous",
        )

    source_classes = list(
        db.scalars(
            select(Classroom)
            .where(
                Classroom.school_id == school_id,
                Classroom.school_year_id == previous_year.id,
                Classroom.is_active.is_(True),
            )
            .order_by(Classroom.name.asc())
        ).all()
    )
    if not source_classes:
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status="no_previous",
        )

    target_rows = list(
        db.scalars(
            select(Classroom).where(
                Classroom.school_id == school_id,
                Classroom.school_year_id == school_year_id,
            )
        ).all()
    )
    target_by_name = {
        normalize_name(item.name).casefold(): item
        for item in target_rows
    }

    copied_count = 0
    for source in source_classes:
        key = normalize_name(source.name).casefold()
        existing = target_by_name.get(key)
        if existing is None:
            classroom = Classroom(
                school_id=school_id,
                school_year_id=school_year_id,
                code=source.code,
                name=source.name,
                is_active=True,
            )
            db.add(classroom)
            target_by_name[key] = classroom
            copied_count += 1
        elif not existing.is_active:
            existing.is_active = True
            copied_count += 1

    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return redirect_list(
            school_year_id=school_year_id,
            level_code=level_code,
            commune_id=commune_id,
            school_id=school_id,
            status="duplicate",
        )

    return redirect_list(
        school_year_id=school_year_id,
        level_code=level_code,
        commune_id=commune_id,
        school_id=school_id,
        status="copied",
        extra={"copied": copied_count},
    )
