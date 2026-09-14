from __future__ import annotations

import json
import re
import unicodedata
from collections import defaultdict
from io import BytesIO
from math import ceil
from pathlib import Path
from typing import Annotated, Any, Iterable
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, Form, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, or_, select, text
from sqlalchemy.orm import Session, selectinload

from app.comparison_models import (
    StudentSurveyComparisonSignoff,
    StudentSurveyResolutionLog,
)
from app.database import get_db
from app.models import (
    Classroom,
    School,
    Student,
    StudentEnrollment,
    User,
)
from app.permissions import (
    SCHOOL_ROLE_CODE,
    TEACHER_ROLE_CODE,
    normalize_role_code,
    is_admin_role,
)
from app.routers.students import STATUS_LABELS as STUDENT_STATUS_LABELS
from app.routers.surveys import (
    LEARNING_STATUS_LABELS,
    lay_thong_tin_nguoi_dung,
    tao_bo_loc_dot_theo_nguoi_dung,
    tao_bo_loc_phieu_theo_nguoi_dung,
)
from app.survey_models import (
    Household,
    SurveyBatch,
    SurveyForm,
    SurveyPerson,
    SurveyPersonYearRecord,
)


APP_DIR = Path(__file__).resolve().parent.parent

templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

router = APIRouter(
    prefix="/dieu-tra",
    tags=["Đối chiếu điều tra với học sinh"],
)

PAGE_SIZE_OPTIONS = (20, 50, 100, 200)

RESULT_LABELS = {
    "KHOP_HOAN_TOAN": "Khớp hoàn toàn",
    "DIEU_TRA_CHUA_CO_HOC_SINH": (
        "Có trong điều tra, chưa có trong dữ liệu học sinh"
    ),
    "HOC_SINH_CHUA_CO_DIEU_TRA": (
        "Có trong dữ liệu học sinh, chưa có trong điều tra"
    ),
    "SAI_KHAC_TRUONG_LOP": "Sai khác trường hoặc lớp",
    "SAI_KHAC_TRANG_THAI": "Sai khác trạng thái học tập",
    "NGHI_TRUNG": "Nghi trùng hoặc xung đột nhận diện",
    "THIEU_THONG_TIN": "Thiếu thông tin để đối chiếu",
}

RESULT_CLASSES = {
    "KHOP_HOAN_TOAN": "good",
    "DIEU_TRA_CHUA_CO_HOC_SINH": "danger",
    "HOC_SINH_CHUA_CO_DIEU_TRA": "danger",
    "SAI_KHAC_TRUONG_LOP": "warning",
    "SAI_KHAC_TRANG_THAI": "warning",
    "NGHI_TRUNG": "danger",
    "THIEU_THONG_TIN": "warning",
}

MATCH_BASIS_LABELS = {
    "LIEN_KET": "Liên kết hồ sơ sẵn có",
    "SO_DINH_DANH": "Số định danh cá nhân",
    "MA_HOC_SINH": "Mã học sinh",
    "HO_TEN_NGAY_SINH_GIOI_TINH": (
        "Họ tên + ngày sinh + giới tính"
    ),
    "HO_TEN_NGAY_SINH": "Họ tên + ngày sinh",
    "CHUA_KHOP": "Chưa tìm được hồ sơ khớp",
    "XUNG_DOT": "Nhiều hồ sơ cùng khớp",
}

ACTIVE_SURVEY_STATUSES = {"DANG_HOC", "CHUYEN_DEN"}
ACTIVE_STUDENT_STATUSES = {
    "DANG_HOC",
    "CHUYEN_DEN_KY_1",
    "CHUYEN_DEN_KY_2",
}

RESOLUTION_ACTION_LABELS = {
    "LIEN_KET": "Đã liên kết hồ sơ học sinh",
    "BO_LIEN_KET": "Đã bỏ liên kết học sinh",
    "DA_RA_SOAT": "Đã rà soát",
    "KHONG_TIM_THAY": "Không tìm thấy hồ sơ phù hợp",
    "XAC_NHAN_SAI_KHAC": "Đã xác nhận sai khác",
    "CHO_XAC_MINH": "Chờ xác minh thêm",
}

RESOLUTION_DONE_ACTIONS = {
    "LIEN_KET",
    "DA_RA_SOAT",
    "KHONG_TIM_THAY",
    "XAC_NHAN_SAI_KHAC",
}

RESOLUTION_STATE_LABELS = {
    "": "Tất cả trạng thái xử lý",
    "PENDING": "Còn chờ xử lý",
    "RECORDED": "Đã ghi nhận xử lý",
}

SIGNOFF_ACTION_LABELS = {
    "CHOT": "Đã chốt kết quả đối chiếu",
    "MO_LAI": "Đã mở lại để xử lý",
}

SIGNOFF_STATUS_MESSAGES = {
    "signed": (
        "success",
        "Đã chốt kết quả đối chiếu và lưu ảnh chụp số liệu vào biên bản.",
    ),
    "reopened": (
        "success",
        "Đã mở lại kết quả đối chiếu. Các chức năng xử lý thủ công hoạt động trở lại.",
    ),
    "pending": (
        "warning",
        "Chưa thể chốt vì vẫn còn dòng cần xử lý chưa được ghi nhận.",
    ),
    "already_signed": (
        "warning",
        "Kết quả đối chiếu đang ở trạng thái đã chốt.",
    ),
    "not_signed": (
        "warning",
        "Kết quả đối chiếu chưa được chốt nên không cần mở lại.",
    ),
    "invalid_note": (
        "warning",
        "Hãy nhập lý do hoặc căn cứ tối thiểu 5 ký tự.",
    ),
    "database_error": (
        "warning",
        "Không thể lưu trạng thái chốt. Hệ thống đã hoàn tác giao dịch.",
    ),
    "forbidden": (
        "warning",
        "Chỉ tài khoản quản trị cấp Sở mới được chốt hoặc mở lại kết quả đối chiếu.",
    ),
}


# =========================================================
# HÀM CHUẨN HÓA
# =========================================================


def _clean_text(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _normalize_code(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).upper()


def _normalize_text(value: Any) -> str:
    text = _clean_text(value).lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(
        character
        for character in text
        if unicodedata.category(character) != "Mn"
    )
    return re.sub(r"\s+", " ", text).strip()


def _normalize_gender(value: Any) -> str:
    text = _normalize_text(value)
    if text in {"nam", "male", "m"}:
        return "NAM"
    if text in {"nu", "nữ", "female", "f"}:
        return "NU"
    return text.upper()



def _parse_optional_int(value: Any) -> int | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return int(text)
    except (TypeError, ValueError):
        return None

def _school_name(record: SurveyPersonYearRecord | None) -> str:
    if record is None:
        return "—"
    if record.school is not None:
        return record.school.name
    return _clean_text(record.school_name_reported) or "Chưa xác định"


def _class_name(record: SurveyPersonYearRecord | None) -> str:
    if record is None:
        return "—"
    if record.classroom is not None:
        return record.classroom.name
    return _clean_text(record.class_name_reported) or "Chưa xác định"


def _student_status_label(enrollment: StudentEnrollment | None) -> str:
    if enrollment is None:
        return "Không có xếp lớp năm học này"
    code = str(enrollment.status or "")
    return STUDENT_STATUS_LABELS.get(code, code or "Chưa xác định")


def _survey_status_label(record: SurveyPersonYearRecord | None) -> str:
    if record is None:
        return "Chưa có hồ sơ năm học"
    code = str(record.learning_status or "CHUA_XAC_DINH")
    return LEARNING_STATUS_LABELS.get(code, code)


def _scope_label(request: Request) -> str:
    user = lay_thong_tin_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))

    if role_code in {SCHOOL_ROLE_CODE, TEACHER_ROLE_CODE}:
        return "PHẠM VI TOÀN TRƯỜNG"
    if role_code == "XA":
        return "PHẠM VI TOÀN XÃ/PHƯỜNG"
    return "PHẠM VI TOÀN TỈNH"


def _chunked(values: Iterable[Any], size: int = 800) -> Iterable[list[Any]]:
    chunk: list[Any] = []
    for value in values:
        chunk.append(value)
        if len(chunk) >= size:
            yield chunk
            chunk = []
    if chunk:
        yield chunk




# === V12_BASELINE_SOURCE_HELPERS_START ===
RECONCILIATION_PLAN_FILE = APP_DIR.parent / "data" / "school_merger_approved_plans.json"


def _baseline_state_for_batch(
    db: Session,
    batch: SurveyBatch,
) -> dict[str, Any] | None:
    try:
        row = db.execute(
            text(
                "SELECT s.*, sy.code AS source_year_code "
                "FROM student_reconciliation_source_states s "
                "JOIN school_years sy ON sy.id=s.source_school_year_id "
                "WHERE s.target_school_year_id=:target_year_id "
                "AND s.commune_id=:commune_id "
                "AND s.status='LOCKED' LIMIT 1"
            ),
            {
                "target_year_id": int(batch.school_year_id),
                "commune_id": int(batch.commune_id),
            },
        ).mappings().first()
    except Exception:
        return None
    return dict(row) if row is not None else None


def _completed_merger_plans_for_year(target_year_id: int) -> list[dict[str, Any]]:
    if not RECONCILIATION_PLAN_FILE.exists():
        return []
    try:
        payload = json.loads(RECONCILIATION_PLAN_FILE.read_text(encoding="utf-8"))
    except Exception:
        return []
    result = []
    for raw in payload.get("plans") or []:
        if not isinstance(raw, dict):
            continue
        if str(raw.get("status") or "").upper() != "COMPLETED":
            continue
        try:
            plan_year = int(raw.get("school_year_id"))
            target = int(raw.get("target_school_id"))
            sources = [int(x) for x in raw.get("source_school_ids") or []]
        except Exception:
            continue
        if plan_year != int(target_year_id):
            continue
        result.append({"target": target, "sources": sources})
    return result


def _effective_school_id_for_year(school_id: int, target_year_id: int) -> int:
    sid = int(school_id)
    for plan in _completed_merger_plans_for_year(target_year_id):
        if sid in plan["sources"]:
            return int(plan["target"])
    return sid


def _baseline_school_ids_for_current_school(
    current_school_id: int,
    target_year_id: int,
) -> set[int]:
    current = int(current_school_id)
    effective = _effective_school_id_for_year(current, target_year_id)
    result = {current, effective}
    for plan in _completed_merger_plans_for_year(target_year_id):
        if int(plan["target"]) == effective:
            result.add(int(plan["target"]))
            result.update(int(x) for x in plan["sources"])
    return result


def _baseline_allowed_school_ids(
    db: Session,
    request: Request,
    batch: SurveyBatch,
) -> set[int]:
    user = lay_thong_tin_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))
    if role_code in {SCHOOL_ROLE_CODE, TEACHER_ROLE_CODE}:
        school_id = user.get("school_id")
        if school_id is None:
            return set()
        return _baseline_school_ids_for_current_school(
            int(school_id), int(batch.school_year_id)
        )
    rows = db.execute(
        text("SELECT id FROM schools WHERE commune_id=:commune_id"),
        {"commune_id": int(batch.commune_id)},
    ).all()
    return {int(r[0]) for r in rows}


def _baseline_year_id(db: Session, batch: SurveyBatch) -> int | None:
    state = _baseline_state_for_batch(db, batch)
    if state is None:
        return None
    return int(state["source_school_year_id"])


def _schools_equivalent_across_baseline(
    current_school_id: int,
    baseline_school_id: int,
    target_year_id: int,
) -> bool:
    return _effective_school_id_for_year(
        int(current_school_id), int(target_year_id)
    ) == _effective_school_id_for_year(
        int(baseline_school_id), int(target_year_id)
    )
# === V12_BASELINE_SOURCE_HELPERS_END ===


# =========================================================
# NẠP DỮ LIỆU TRONG PHẠM VI
# =========================================================


def _batch_in_scope(
    db: Session,
    request: Request,
    batch_id: int,
) -> SurveyBatch | None:
    filters = tao_bo_loc_dot_theo_nguoi_dung(request)
    return db.scalar(
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
        )
        .where(
            SurveyBatch.id == batch_id,
            *filters,
        )
    )


def _load_survey_rows(
    db: Session,
    request: Request,
    batch: SurveyBatch,
) -> list[dict[str, Any]]:
    # === BAI31C2A_SCHOOL_FULL_COMMUNE_SURVEY_SCOPE_START ===
    # Đối chiếu cấp TRƯỜNG phải được tìm học sinh của trường trong toàn bộ
    # dữ liệu điều tra của xã, không chỉ các phiếu do GV của trường được giao.
    # Các vai trò khác giữ nguyên bộ lọc cũ.
    _comparison_user = lay_thong_tin_nguoi_dung(request)
    _comparison_role = normalize_role_code(_comparison_user.get("role_code"))
    if _comparison_role == SCHOOL_ROLE_CODE:
        form_filters = []
    else:
        form_filters = tao_bo_loc_phieu_theo_nguoi_dung(request)
    # === BAI31C2A_SCHOOL_FULL_COMMUNE_SURVEY_SCOPE_END ===

    statement = (
        select(
            SurveyForm,
            Household,
            SurveyPerson,
            SurveyPersonYearRecord,
        )
        .join(
            Household,
            Household.id == SurveyForm.household_id,
        )
        .join(
            SurveyPerson,
            SurveyPerson.household_id == Household.id,
        )
        .outerjoin(
            SurveyPersonYearRecord,
            and_(
                SurveyPersonYearRecord.survey_form_id == SurveyForm.id,
                SurveyPersonYearRecord.survey_person_id == SurveyPerson.id,
                SurveyPersonYearRecord.school_year_id
                == batch.school_year_id,
            ),
        )
        .where(
            SurveyForm.survey_batch_id == batch.id,
            SurveyPerson.is_active.is_(True),
            *form_filters,
        )
        .order_by(
            Household.hamlet_name.asc(),
            Household.head_name.asc(),
            SurveyPerson.full_name.asc(),
            SurveyPerson.id.asc(),
        )
    )

    rows = db.execute(statement).all()

    record_ids = [
        row[3].id
        for row in rows
        if row[3] is not None
    ]
    record_map: dict[int, SurveyPersonYearRecord] = {}

    for chunk in _chunked(record_ids):
        records = db.scalars(
            select(SurveyPersonYearRecord)
            .options(
                selectinload(SurveyPersonYearRecord.school),
                selectinload(SurveyPersonYearRecord.classroom),
            )
            .where(SurveyPersonYearRecord.id.in_(chunk))
        ).all()
        record_map.update({record.id: record for record in records})

    result: list[dict[str, Any]] = []
    for form, household, person, record in rows:
        if record is not None:
            record = record_map.get(record.id, record)
        result.append(
            {
                "form": form,
                "household": household,
                "person": person,
                "record": record,
            }
        )
    return result


def _load_scope_enrollments(
    db: Session,
    request: Request,
    batch: SurveyBatch,
) -> list[StudentEnrollment]:
    baseline_year_id = _baseline_year_id(db, batch)
    if baseline_year_id is None:
        return []
    allowed_school_ids = _baseline_allowed_school_ids(db, request, batch)
    if not allowed_school_ids:
        return []

    return db.scalars(
        select(StudentEnrollment)
        .join(Student, Student.id == StudentEnrollment.student_id)
        .join(School, School.id == StudentEnrollment.school_id)
        .options(
            selectinload(StudentEnrollment.student),
            selectinload(StudentEnrollment.school).selectinload(School.commune),
            selectinload(StudentEnrollment.classroom),
            selectinload(StudentEnrollment.school_year),
        )
        .where(
            StudentEnrollment.school_year_id == baseline_year_id,
            StudentEnrollment.school_id.in_(sorted(allowed_school_ids)),
            Student.is_active.is_(True),
        )
        .order_by(
            Student.full_name.asc(),
            Student.code.asc(),
            StudentEnrollment.is_current.desc(),
            StudentEnrollment.id.desc(),
        )
    ).all()



def _load_extra_candidate_students(
    db: Session,
    survey_rows: list[dict[str, Any]],
    existing_student_ids: set[int],
) -> list[Student]:
    linked_ids: set[int] = set()
    personal_ids: set[str] = set()
    student_codes: set[str] = set()

    for item in survey_rows:
        person: SurveyPerson = item["person"]
        if person.student_id is not None:
            linked_ids.add(int(person.student_id))
        personal_id = _normalize_code(person.personal_id)
        if personal_id:
            personal_ids.add(personal_id)
        student_code = _normalize_code(person.ministry_student_code)
        if student_code:
            student_codes.add(student_code)

    result: dict[int, Student] = {}

    ids_to_load = linked_ids - existing_student_ids
    for chunk in _chunked(sorted(ids_to_load)):
        students = db.scalars(
            select(Student).where(
                Student.id.in_(chunk),
                Student.is_active.is_(True),
            )
        ).all()
        result.update({student.id: student for student in students})

    for chunk in _chunked(sorted(personal_ids)):
        students = db.scalars(
            select(Student).where(
                Student.personal_id.in_(chunk),
                Student.is_active.is_(True),
            )
        ).all()
        result.update({student.id: student for student in students})

    for chunk in _chunked(sorted(student_codes)):
        students = db.scalars(
            select(Student).where(
                Student.code.in_(chunk),
                Student.is_active.is_(True),
            )
        ).all()
        result.update({student.id: student for student in students})

    return list(result.values())


def _load_target_enrollments_for_students(
    db: Session,
    school_year_id: int,
    student_ids: set[int],
    allowed_school_ids: set[int] | None = None,
) -> dict[int, list[StudentEnrollment]]:
    result: dict[int, list[StudentEnrollment]] = defaultdict(list)
    if not student_ids:
        return result
    filters: list[Any] = [
        StudentEnrollment.school_year_id == int(school_year_id),
    ]
    if allowed_school_ids is not None:
        if not allowed_school_ids:
            return result
        filters.append(
            StudentEnrollment.school_id.in_(sorted(allowed_school_ids))
        )
    for chunk in _chunked(sorted(student_ids)):
        enrollments = db.scalars(
            select(StudentEnrollment)
            .options(
                selectinload(StudentEnrollment.student),
                selectinload(StudentEnrollment.school).selectinload(School.commune),
                selectinload(StudentEnrollment.classroom),
                selectinload(StudentEnrollment.school_year),
            )
            .where(
                *filters,
                StudentEnrollment.student_id.in_(chunk),
            )
            .order_by(
                StudentEnrollment.is_current.desc(),
                StudentEnrollment.id.desc(),
            )
        ).all()
        for enrollment in enrollments:
            result[enrollment.student_id].append(enrollment)
    return result



def _choose_enrollment(
    enrollments: list[StudentEnrollment],
    preferred_school_id: int | None = None,
) -> StudentEnrollment | None:
    if not enrollments:
        return None

    if preferred_school_id is not None:
        for enrollment in enrollments:
            if enrollment.school_id == preferred_school_id:
                return enrollment

    return sorted(
        enrollments,
        key=lambda item: (
            bool(item.is_current),
            item.id,
        ),
        reverse=True,
    )[0]


# =========================================================
# ĐỐI CHIẾU
# =========================================================


def _build_student_indexes(
    students: Iterable[Student],
) -> dict[str, Any]:
    by_id: dict[int, Student] = {}
    by_personal: dict[str, list[Student]] = defaultdict(list)
    by_code: dict[str, list[Student]] = defaultdict(list)
    by_name_dob_gender: dict[tuple[Any, ...], list[Student]] = defaultdict(list)
    by_name_dob: dict[tuple[Any, ...], list[Student]] = defaultdict(list)

    for student in students:
        by_id[student.id] = student

        personal_id = _normalize_code(student.personal_id)
        if personal_id:
            by_personal[personal_id].append(student)

        code = _normalize_code(student.code)
        if code:
            by_code[code].append(student)

        name = _normalize_text(student.full_name)
        if name and student.date_of_birth is not None:
            by_name_dob[(name, student.date_of_birth)].append(student)
            gender = _normalize_gender(student.gender)
            if gender:
                by_name_dob_gender[
                    (name, student.date_of_birth, gender)
                ].append(student)

    return {
        "by_id": by_id,
        "by_personal": by_personal,
        "by_code": by_code,
        "by_name_dob_gender": by_name_dob_gender,
        "by_name_dob": by_name_dob,
    }


def _find_candidates(
    person: SurveyPerson,
    indexes: dict[str, Any],
) -> tuple[list[tuple[Student, int, str]], bool]:
    scores: dict[int, tuple[Student, int, str]] = {}
    exact_identifier_student_ids: set[int] = set()

    def add_candidates(
        candidates: Iterable[Student],
        score: int,
        basis: str,
        exact_identifier: bool = False,
    ) -> None:
        for student in candidates:
            current = scores.get(student.id)
            if current is None or score > current[1]:
                scores[student.id] = (student, score, basis)
            if exact_identifier:
                exact_identifier_student_ids.add(student.id)

    if person.student_id is not None:
        student = indexes["by_id"].get(int(person.student_id))
        if student is not None:
            add_candidates([student], 100, "LIEN_KET", True)

    personal_id = _normalize_code(person.personal_id)
    if personal_id:
        add_candidates(
            indexes["by_personal"].get(personal_id, []),
            95,
            "SO_DINH_DANH",
            True,
        )

    ministry_code = _normalize_code(person.ministry_student_code)
    if ministry_code:
        add_candidates(
            indexes["by_code"].get(ministry_code, []),
            90,
            "MA_HOC_SINH",
            True,
        )

    normalized_name = _normalize_text(person.full_name)
    if normalized_name and person.date_of_birth is not None:
        normalized_gender = _normalize_gender(person.gender)
        if normalized_gender:
            add_candidates(
                indexes["by_name_dob_gender"].get(
                    (
                        normalized_name,
                        person.date_of_birth,
                        normalized_gender,
                    ),
                    [],
                ),
                80,
                "HO_TEN_NGAY_SINH_GIOI_TINH",
            )
        add_candidates(
            indexes["by_name_dob"].get(
                (normalized_name, person.date_of_birth),
                [],
            ),
            75,
            "HO_TEN_NGAY_SINH",
        )

    ordered = sorted(
        scores.values(),
        key=lambda item: (-item[1], item[0].full_name, item[0].id),
    )
    has_identifier_conflict = len(exact_identifier_student_ids) > 1
    return ordered, has_identifier_conflict


# === BAI31C2A_CROSS_YEAR_PROGRESSION_START ===
def _class_grade_number_for_cross_year(value: str | None) -> int | None:
    """
    Chỉ nhận dạng khối lớp phổ thông rõ ràng: 1..12, ví dụ 5A, Lớp 6A1.
    Không hiểu '5 tuổi' của mầm non là lớp 5.
    """
    raw = str(value or "").strip()
    if not raw:
        return None

    lowered = raw.casefold()
    if "tuổi" in lowered or "tuoi" in lowered:
        return None

    match = re.match(
        r"^\s*(?:lớp\s*)?(1[0-2]|[1-9])(?=[^\d]|$)",
        raw,
        flags=re.IGNORECASE,
    )
    if not match:
        return None

    try:
        return int(match.group(1))
    except Exception:
        return None


def _compare_school_class(
    record: SurveyPersonYearRecord | None,
    enrollment: StudentEnrollment | None,
) -> list[str]:
    if record is None or enrollment is None:
        return []

    reasons: list[str] = []
    cross_year = int(record.school_year_id) != int(enrollment.school_year_id)

    if cross_year:
        # Nguồn là năm trước, vì vậy thay đổi trường là bình thường khi:
        # - lên lớp/chuyển cấp;
        # - chuyển trường;
        # - trường nguồn đã sáp nhập.
        # Không coi "khác trường" tự thân là lỗi liên năm.
        #
        # Chỉ kiểm tra tiến trình lớp nếu CẢ HAI tên lớp đều nhận dạng chắc chắn
        # được khối phổ thông 1..12. Nếu tên lớp không chuẩn (đặc biệt mầm non),
        # không tự suy diễn và không sinh lỗi giả.
        baseline_class_name = (
            enrollment.classroom.name
            if enrollment.classroom is not None
            else ""
        )
        current_class_name = ""
        if record.classroom is not None:
            current_class_name = record.classroom.name or ""
        elif getattr(record, "class_name_reported", None):
            current_class_name = record.class_name_reported or ""

        baseline_grade = _class_grade_number_for_cross_year(
            baseline_class_name
        )
        current_grade = _class_grade_number_for_cross_year(
            current_class_name
        )

        if (
            baseline_grade is not None
            and current_grade is not None
            and baseline_grade < 12
        ):
            expected_grade = baseline_grade + 1
            is_repeating = bool(
                getattr(record, "is_repeating_grade", False)
            )

            normal_progression = current_grade == expected_grade
            repeated_as_recorded = (
                current_grade == baseline_grade and is_repeating
            )

            if not normal_progression and not repeated_as_recorded:
                reasons.append(
                    "Lớp hiện tại chưa phù hợp tiến trình liên năm "
                    f"(nguồn năm trước: lớp {baseline_grade}; "
                    f"điều tra hiện tại: lớp {current_grade})"
                )

        return reasons

    # Cùng một năm học: giữ nguyên logic cũ, trường/lớp phải thống nhất.
    if record.school_id is not None:
        if int(record.school_id) != int(enrollment.school_id):
            reasons.append(
                "Trường điều tra khác trường trong dữ liệu học sinh"
            )
    else:
        survey_school_name = _normalize_text(
            record.school_name_reported
        )
        student_school_name = _normalize_text(
            enrollment.school.name if enrollment.school else ""
        )
        if (
            survey_school_name
            and student_school_name
            and survey_school_name != student_school_name
        ):
            reasons.append(
                "Tên trường ghi trên phiếu khác dữ liệu học sinh"
            )

    if record.class_id is not None:
        if int(record.class_id) != int(enrollment.class_id):
            reasons.append(
                "Lớp điều tra khác lớp trong dữ liệu học sinh"
            )
    else:
        survey_class_name = _normalize_text(
            record.class_name_reported
        )
        student_class_name = _normalize_text(
            enrollment.classroom.name if enrollment.classroom else ""
        )
        if (
            survey_class_name
            and student_class_name
            and survey_class_name != student_class_name
        ):
            reasons.append(
                "Tên lớp ghi trên phiếu khác dữ liệu học sinh"
            )

    return reasons
# === BAI31C2A_CROSS_YEAR_PROGRESSION_END ===


def _compare_learning_status(
    record: SurveyPersonYearRecord | None,
    enrollment: StudentEnrollment | None,
) -> list[str]:
    reasons: list[str] = []
    if record is None:
        return ["Đối tượng chưa có hồ sơ năm học hiện tại"]
    survey_code = str(record.learning_status or "CHUA_XAC_DINH")
    if survey_code == "CHUA_XAC_DINH":
        return ["Điều tra chưa xác định trạng thái học tập"]
    if enrollment is None:
        if survey_code in ACTIVE_SURVEY_STATUSES:
            return ["Điều tra ghi đang học nhưng không có hồ sơ trong nguồn học sinh năm trước"]
        return []

    student_code = str(enrollment.status or "")
    survey_active = survey_code in ACTIVE_SURVEY_STATUSES
    student_active = bool(enrollment.is_current) and student_code in ACTIVE_STUDENT_STATUSES
    cross_year = int(record.school_year_id) != int(enrollment.school_year_id)
    if cross_year:
        if survey_active != student_active:
            reasons.append(
                "Trạng thái học tập hiện tại thay đổi so với nguồn học sinh năm trước"
            )
        return reasons

    if survey_active != student_active:
        reasons.append("Trạng thái đang học/không còn học giữa hai nguồn không thống nhất")
    elif survey_code == "CHUYEN_DI" and student_code != "CHUYEN_DI":
        reasons.append("Điều tra ghi chuyển đi nhưng dữ liệu học sinh chưa ghi chuyển đi")
    elif survey_code in {"THOI_HOC", "BO_HOC"} and student_code not in {"THOI_HOC", "TAM_NGHI"}:
        reasons.append("Trạng thái thôi học/bỏ học giữa hai nguồn chưa thống nhất")
    return reasons



def _build_comparison_rows(
    db: Session,
    request: Request,
    batch: SurveyBatch,
) -> tuple[list[dict[str, Any]], list[School]]:
    baseline_state = _baseline_state_for_batch(db, batch)
    if baseline_state is None:
        return [], []
    baseline_year_id = int(baseline_state["source_school_year_id"])
    allowed_school_ids = _baseline_allowed_school_ids(db, request, batch)
    survey_rows = _load_survey_rows(db, request, batch)
    scope_enrollments = _load_scope_enrollments(db, request, batch)

    scope_enrollment_by_student: dict[int, StudentEnrollment] = {}
    all_students: dict[int, Student] = {}

    for enrollment in scope_enrollments:
        all_students[enrollment.student.id] = enrollment.student
        current = scope_enrollment_by_student.get(enrollment.student_id)
        if current is None:
            scope_enrollment_by_student[enrollment.student_id] = enrollment
        elif bool(enrollment.is_current) and not bool(current.is_current):
            scope_enrollment_by_student[enrollment.student_id] = enrollment

    # Candidate chỉ lấy từ nguồn baseline của đúng địa bàn; không kéo Student toàn tỉnh vào.
    enrollments_by_student = _load_target_enrollments_for_students(
        db,
        baseline_year_id,
        set(all_students),
        allowed_school_ids=allowed_school_ids,
    )

    indexes = _build_student_indexes(all_students.values())
    rows: list[dict[str, Any]] = []
    matched_student_to_row_indexes: dict[int, list[int]] = defaultdict(list)

    # === BAI31C2A_SCHOOL_RELEVANT_SURVEY_ROWS_START ===
    # Với tài khoản TRƯỜNG:
    # - tìm học sinh nguồn của chính trường trong toàn xã;
    # - chỉ đưa vào bảng các đối tượng khớp nguồn của trường hoặc đang được
    #   điều tra là học tại chính trường. Không kéo các đối tượng không liên quan.
    _comparison_user = lay_thong_tin_nguoi_dung(request)
    _comparison_role = normalize_role_code(_comparison_user.get("role_code"))
    _current_school_id: int | None = None
    _current_school_name = ""
    if _comparison_role == SCHOOL_ROLE_CODE:
        _raw_school_id = _comparison_user.get("school_id")
        if _raw_school_id not in (None, ""):
            try:
                _current_school_id = int(_raw_school_id)
            except Exception:
                _current_school_id = None
        if _current_school_id is not None:
            _current_school = db.get(School, _current_school_id)
            if _current_school is not None:
                _current_school_name = _normalize_text(_current_school.name)
    # === BAI31C2A_SCHOOL_RELEVANT_SURVEY_ROWS_END ===

    for survey_item in survey_rows:
        form: SurveyForm = survey_item["form"]
        household: Household = survey_item["household"]
        person: SurveyPerson = survey_item["person"]
        record: SurveyPersonYearRecord | None = survey_item["record"]

        candidates, identifier_conflict = _find_candidates(person, indexes)

        # === BAI31C2A_SCHOOL_RELEVANCE_FILTER_START ===
        if _comparison_role == SCHOOL_ROLE_CODE:
            _survey_points_to_current_school = False
            if record is not None and _current_school_id is not None:
                if record.school_id is not None:
                    _survey_points_to_current_school = (
                        _effective_school_id_for_year(
                            int(record.school_id),
                            int(batch.school_year_id),
                        )
                        == _effective_school_id_for_year(
                            int(_current_school_id),
                            int(batch.school_year_id),
                        )
                    )
                elif (
                    _current_school_name
                    and getattr(record, "school_name_reported", None)
                ):
                    _survey_points_to_current_school = (
                        _normalize_text(record.school_name_reported)
                        == _current_school_name
                    )

            if not candidates and not _survey_points_to_current_school:
                continue
        # === BAI31C2A_SCHOOL_RELEVANCE_FILTER_END ===
        selected_student: Student | None = None
        selected_enrollment: StudentEnrollment | None = None
        result_code = ""
        detail_parts: list[str] = []
        match_basis = "CHUA_KHOP"

        if identifier_conflict:
            result_code = "NGHI_TRUNG"
            match_basis = "XUNG_DOT"
            detail_parts.append(
                "Các trường nhận diện đang trỏ tới nhiều hồ sơ học sinh khác nhau"
            )
        elif not candidates:
            has_minimum_identity = bool(
                _normalize_code(person.personal_id)
                or _normalize_code(person.ministry_student_code)
                or (
                    _normalize_text(person.full_name)
                    and person.date_of_birth is not None
                )
            )
            result_code = (
                "DIEU_TRA_CHUA_CO_HOC_SINH"
                if has_minimum_identity
                else "THIEU_THONG_TIN"
            )
            if has_minimum_identity:
                detail_parts.append(
                    "Không tìm thấy hồ sơ học sinh khớp theo các thông tin nhận diện hiện có"
                )
            else:
                detail_parts.append(
                    "Thiếu số định danh, mã học sinh hoặc ngày sinh để đối chiếu tin cậy"
                )
        else:
            top_score = candidates[0][1]
            top_candidates = [
                item for item in candidates if item[1] == top_score
            ]

            if len(top_candidates) > 1:
                result_code = "NGHI_TRUNG"
                match_basis = "XUNG_DOT"
                detail_parts.append(
                    f"Có {len(top_candidates)} hồ sơ học sinh cùng mức độ khớp"
                )
            else:
                selected_student, _, match_basis = top_candidates[0]
                preferred_school_id = (
                    record.school_id if record is not None else None
                )
                selected_enrollment = _choose_enrollment(
                    enrollments_by_student.get(selected_student.id, []),
                    preferred_school_id=preferred_school_id,
                )

                status_reasons = _compare_learning_status(
                    record,
                    selected_enrollment,
                )
                school_class_reasons = _compare_school_class(
                    record,
                    selected_enrollment,
                )

                if status_reasons:
                    result_code = "SAI_KHAC_TRANG_THAI"
                    detail_parts.extend(status_reasons)
                    detail_parts.extend(school_class_reasons)
                elif school_class_reasons:
                    result_code = "SAI_KHAC_TRUONG_LOP"
                    detail_parts.extend(school_class_reasons)
                else:
                    result_code = "KHOP_HOAN_TOAN"
                    detail_parts.append("Thông tin nhận diện, trạng thái, trường và lớp thống nhất")

        row = {
            "row_type": "SURVEY",
            "result_code": result_code,
            "result_label": RESULT_LABELS[result_code],
            "result_class": RESULT_CLASSES[result_code],
            "match_basis": match_basis,
            "match_basis_label": MATCH_BASIS_LABELS[match_basis],
            "detail": "; ".join(detail_parts),
            "form": form,
            "household": household,
            "person": person,
            "record": record,
            "student": selected_student,
            "enrollment": selected_enrollment,
            "survey_status_label": _survey_status_label(record),
            "student_status_label": _student_status_label(selected_enrollment),
            "survey_school_name": _school_name(record),
            "survey_class_name": _class_name(record),
            "student_school_name": (
                selected_enrollment.school.name
                if selected_enrollment is not None
                and selected_enrollment.school is not None
                else "—"
            ),
            "student_class_name": (
                selected_enrollment.classroom.name
                if selected_enrollment is not None
                and selected_enrollment.classroom is not None
                else "—"
            ),
            "school_id": (
                selected_enrollment.school_id
                if selected_enrollment is not None
                else (record.school_id if record is not None else None)
            ),
            "search_text": " ".join(
                [
                    person.code,
                    person.full_name,
                    person.personal_id or "",
                    person.ministry_student_code or "",
                    household.code,
                    household.head_name,
                    form.form_number,
                    selected_student.code if selected_student else "",
                    selected_student.full_name if selected_student else "",
                    selected_student.personal_id or ""
                    if selected_student
                    else "",
                ]
            ).lower(),
        }
        rows.append(row)

        if selected_student is not None:
            matched_student_to_row_indexes[selected_student.id].append(
                len(rows) - 1
            )

    duplicated_student_ids = {
        student_id
        for student_id, indexes_list in matched_student_to_row_indexes.items()
        if len(indexes_list) > 1
    }
    for student_id in duplicated_student_ids:
        for row_index in matched_student_to_row_indexes[student_id]:
            rows[row_index]["result_code"] = "NGHI_TRUNG"
            rows[row_index]["result_label"] = RESULT_LABELS["NGHI_TRUNG"]
            rows[row_index]["result_class"] = RESULT_CLASSES["NGHI_TRUNG"]
            rows[row_index]["match_basis"] = "XUNG_DOT"
            rows[row_index]["match_basis_label"] = MATCH_BASIS_LABELS[
                "XUNG_DOT"
            ]
            rows[row_index]["detail"] = (
                "Một hồ sơ học sinh đang khớp với nhiều đối tượng điều tra; "
                "cần kiểm tra nguy cơ trùng người"
            )

    matched_student_ids = {
        row["student"].id
        for row in rows
        if row.get("student") is not None
    }

    for student_id, enrollment in scope_enrollment_by_student.items():
        # Chỉ enrollment còn hiệu lực cuối năm nguồn mới là danh sách kỳ vọng.
        if not bool(enrollment.is_current):
            continue
        if student_id in matched_student_ids:
            continue
        student = enrollment.student
        rows.append(
            {
                "row_type": "STUDENT",
                "result_code": "HOC_SINH_CHUA_CO_DIEU_TRA",
                "result_label": RESULT_LABELS[
                    "HOC_SINH_CHUA_CO_DIEU_TRA"
                ],
                "result_class": RESULT_CLASSES[
                    "HOC_SINH_CHUA_CO_DIEU_TRA"
                ],
                "match_basis": "CHUA_KHOP",
                "match_basis_label": MATCH_BASIS_LABELS["CHUA_KHOP"],
                "detail": (
                    "Học sinh có trong danh sách năm học nhưng chưa tìm thấy "
                    "đối tượng tương ứng trong đợt điều tra"
                ),
                "form": None,
                "household": None,
                "person": None,
                "record": None,
                "student": student,
                "enrollment": enrollment,
                "survey_status_label": "—",
                "student_status_label": _student_status_label(enrollment),
                "survey_school_name": "—",
                "survey_class_name": "—",
                "student_school_name": (
                    enrollment.school.name if enrollment.school else "—"
                ),
                "student_class_name": (
                    enrollment.classroom.name
                    if enrollment.classroom
                    else "—"
                ),
                "school_id": enrollment.school_id,
                "search_text": " ".join(
                    [
                        student.code,
                        student.full_name,
                        student.personal_id or "",
                        enrollment.school.name if enrollment.school else "",
                        enrollment.classroom.name
                        if enrollment.classroom
                        else "",
                    ]
                ).lower(),
            }
        )

    school_map: dict[int, School] = {}
    for row in rows:
        enrollment = row.get("enrollment")
        record = row.get("record")
        if enrollment is not None and enrollment.school is not None:
            school_map[enrollment.school.id] = enrollment.school
        if record is not None and record.school is not None:
            school_map[record.school.id] = record.school

    rows.sort(
        key=lambda row: (
            row["result_code"] == "KHOP_HOAN_TOAN",
            row["result_label"],
            _normalize_text(
                row["person"].full_name
                if row.get("person") is not None
                else row["student"].full_name
            ),
        )
    )

    schools = sorted(school_map.values(), key=lambda item: item.name)
    return rows, schools


def _apply_filters(
    rows: list[dict[str, Any]],
    result_group: str,
    school_id: int | None,
    q: str,
    resolution_state: str = "",
) -> list[dict[str, Any]]:
    q_normalized = q.strip().lower()
    result: list[dict[str, Any]] = []

    for row in rows:
        if result_group and row["result_code"] != result_group:
            continue
        if school_id is not None and row.get("school_id") != school_id:
            continue
        if q_normalized and q_normalized not in row["search_text"]:
            continue
        if resolution_state == "RECORDED" and not row.get("resolution_done"):
            continue
        if resolution_state == "PENDING" and (
            row["result_code"] == "KHOP_HOAN_TOAN"
            or row.get("resolution_done")
        ):
            continue
        result.append(row)

    return result


def _build_stats(rows: list[dict[str, Any]]) -> dict[str, int]:
    stats = {code: 0 for code in RESULT_LABELS}
    for row in rows:
        stats[row["result_code"]] += 1
    stats["TOTAL"] = len(rows)
    stats["CAN_XU_LY"] = sum(
        value
        for code, value in stats.items()
        if code not in {"TOTAL", "KHOP_HOAN_TOAN", "CAN_XU_LY"}
    )
    stats["DA_GHI_NHAN_XU_LY"] = sum(
        1
        for row in rows
        if row["result_code"] != "KHOP_HOAN_TOAN"
        and bool(row.get("resolution_done"))
    )
    stats["CON_CHO_XU_LY"] = max(
        0,
        stats["CAN_XU_LY"] - stats["DA_GHI_NHAN_XU_LY"],
    )
    return stats


def _page_url(
    batch_id: int,
    *,
    result_group: str,
    school_id: int | None,
    q: str,
    resolution_state: str,
    page_size: int,
    page: int,
) -> str:
    params: dict[str, Any] = {
        "result_group": result_group,
        "q": q,
        "resolution_state": resolution_state,
        "page_size": page_size,
        "page": page,
    }
    if school_id is not None:
        params["school_id"] = school_id
    return (
        f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?"
        + urlencode(params)
    )



# =========================================================
# BÀI 12D-17 - CHỐT KẾT QUẢ ĐỐI CHIẾU
# =========================================================


def _can_manage_signoff(request: Request) -> bool:
    user = lay_thong_tin_nguoi_dung(request)
    return is_admin_role(user.get("role_code"))


def _load_signoff_history(
    db: Session,
    batch_id: int,
) -> list[StudentSurveyComparisonSignoff]:
    return list(
        db.scalars(
            select(StudentSurveyComparisonSignoff)
            .where(
                StudentSurveyComparisonSignoff.survey_batch_id == batch_id
            )
            .order_by(
                StudentSurveyComparisonSignoff.created_at.desc(),
                StudentSurveyComparisonSignoff.id.desc(),
            )
        ).all()
    )


def _latest_signoff(
    db: Session,
    batch_id: int,
) -> StudentSurveyComparisonSignoff | None:
    return db.scalar(
        select(StudentSurveyComparisonSignoff)
        .where(StudentSurveyComparisonSignoff.survey_batch_id == batch_id)
        .order_by(
            StudentSurveyComparisonSignoff.created_at.desc(),
            StudentSurveyComparisonSignoff.id.desc(),
        )
        .limit(1)
    )


def _comparison_is_signed(db: Session, batch_id: int) -> bool:
    latest = _latest_signoff(db, batch_id)
    return latest is not None and latest.action_code == "CHOT"


def _signoff_status_message(status: str) -> dict[str, str] | None:
    item = SIGNOFF_STATUS_MESSAGES.get(str(status or "").strip())
    if item is None:
        return None
    return {"kind": item[0], "text": item[1]}


def _signoff_snapshot(
    stats: dict[str, int],
) -> dict[str, int]:
    return {
        "total_rows": int(stats.get("TOTAL", 0)),
        "matched_rows": int(stats.get("KHOP_HOAN_TOAN", 0)),
        "issue_rows": int(stats.get("CAN_XU_LY", 0)),
        "resolved_rows": int(stats.get("DA_GHI_NHAN_XU_LY", 0)),
        "pending_rows": int(stats.get("CON_CHO_XU_LY", 0)),
    }


def _add_signoff_log(
    db: Session,
    request: Request,
    batch: SurveyBatch,
    *,
    action_code: str,
    note: str,
    stats: dict[str, int],
) -> None:
    user = lay_thong_tin_nguoi_dung(request)
    snapshot = _signoff_snapshot(stats)
    db.add(
        StudentSurveyComparisonSignoff(
            survey_batch_id=batch.id,
            action_code=action_code,
            note=_clean_text(note)[:1500],
            actor_user_id=user.get("id"),
            actor_name_snapshot=_clean_text(user.get("full_name")) or "Không xác định",
            actor_role_snapshot=_clean_text(user.get("role_name")) or _clean_text(user.get("role_code")) or "Không xác định",
            **snapshot,
        )
    )


def _signoff_redirect_url(batch_id: int, status: str = "") -> str:
    params = urlencode({"status": status}) if status else ""
    base = f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh/chot-ket-qua"
    return base + ("?" + params if params else "")


# =========================================================
# GIAO DIỆN
# =========================================================


@router.get(
    "/{batch_id}/doi-chieu-hoc-sinh",
    response_class=HTMLResponse,
)
def comparison_page(
    request: Request,
    batch_id: int,
    result_group: str = "",
    school_id: str = "",
    q: str = "",
    resolution_state: str = "",
    status: str = "",
    page_size: int = 50,
    page: int = 1,
    db: Session = Depends(get_db),
):
    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return templates.TemplateResponse(
            request=request,
            name="surveys/student_survey_comparison.html",
            context={
                "nguoi_dung": lay_thong_tin_nguoi_dung(request),
                "batch": None,
                "error_message": (
                    "Không tìm thấy đợt điều tra hoặc tài khoản không có quyền xem."
                ),
            },
            status_code=404,
        )

    school_id = _parse_optional_int(school_id)

    result_group = result_group.strip().upper()
    if result_group not in RESULT_LABELS:
        result_group = ""

    q = q.strip()[:120]
    resolution_state = resolution_state.strip().upper()
    if resolution_state not in {"", "PENDING", "RECORDED"}:
        resolution_state = ""
    if page_size not in PAGE_SIZE_OPTIONS:
        page_size = 50

    baseline_state = _baseline_state_for_batch(db, batch)
    all_rows, schools = _build_comparison_rows(db, request, batch)
    _attach_latest_resolutions(db, batch.id, all_rows)
    stats = _build_stats(all_rows)
    signoff_history = _load_signoff_history(db, batch.id)
    latest_signoff = signoff_history[0] if signoff_history else None
    comparison_signed = bool(
        latest_signoff is not None
        and latest_signoff.action_code == "CHOT"
    )
    can_resolve = _can_write_resolution(request) and not comparison_signed
    filtered_rows = _apply_filters(
        all_rows,
        result_group=result_group,
        school_id=school_id,
        q=q,
        resolution_state=resolution_state,
    )

    total_records = len(filtered_rows)
    total_pages = max(1, ceil(total_records / page_size))
    page = max(1, min(page, total_pages))
    start_index = (page - 1) * page_size
    end_index = start_index + page_size
    page_rows = filtered_rows[start_index:end_index]

    page_links = [
        {
            "number": number,
            "url": _page_url(
                batch.id,
                result_group=result_group,
                school_id=school_id,
                q=q,
                resolution_state=resolution_state,
                page_size=page_size,
                page=number,
            ),
        }
        for number in range(max(1, page - 2), min(total_pages, page + 2) + 1)
    ]

    return templates.TemplateResponse(
        request=request,
        name="surveys/student_survey_comparison.html",
        context={
            "nguoi_dung": lay_thong_tin_nguoi_dung(request),
            "batch": batch,
            "scope_label": _scope_label(request),
            "baseline_ready": baseline_state is not None,
            "baseline_source_year_id": (int(baseline_state["source_school_year_id"]) if baseline_state else None),
            "baseline_source_year_code": (str(baseline_state.get("source_year_code") or "") if baseline_state else ""),
            "error_message": None,
            "result_labels": RESULT_LABELS,
            "result_group": result_group,
            "schools": schools,
            "school_id": school_id,
            "q": q,
            "resolution_state": resolution_state,
            "resolution_state_labels": RESOLUTION_STATE_LABELS,
            "status_message": _comparison_status_message(status),
            "page_size": page_size,
            "page_size_options": PAGE_SIZE_OPTIONS,
            "stats": stats,
            "rows": page_rows,
            "total_records": total_records,
            "page": page,
            "total_pages": total_pages,
            "page_links": page_links,
            "previous_url": (
                _page_url(
                    batch.id,
                    result_group=result_group,
                    school_id=school_id,
                    q=q,
                    resolution_state=resolution_state,
                    page_size=page_size,
                    page=page - 1,
                )
                if page > 1
                else None
            ),
            "next_url": (
                _page_url(
                    batch.id,
                    result_group=result_group,
                    school_id=school_id,
                    q=q,
                    resolution_state=resolution_state,
                    page_size=page_size,
                    page=page + 1,
                )
                if page < total_pages
                else None
            ),
            "can_open_students": is_admin_role(
                lay_thong_tin_nguoi_dung(request).get("role_code")
            ),
            "can_resolve": can_resolve,
            "comparison_signed": comparison_signed,
            "latest_signoff": latest_signoff,
            "signoff_action_labels": SIGNOFF_ACTION_LABELS,
            "signoff_url": (
                f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh/chot-ket-qua"
            ),
            "export_url": (
                f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh/xuat-excel?"
                + urlencode(
                    {
                        "result_group": result_group,
                        "school_id": school_id or "",
                        "q": q,
                        "resolution_state": resolution_state,
                    }
                )
            ),
        },
    )


# =========================================================
# XUẤT EXCEL
# =========================================================


def _style_worksheet(worksheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="C7D3DF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    widths = {
        1: 7,
        2: 24,
        3: 18,
        4: 21,
        5: 20,
        6: 18,
        7: 18,
        8: 24,
        9: 18,
        10: 22,
        11: 18,
        12: 24,
        13: 18,
        14: 25,
        15: 48,
        16: 27,
        17: 42,
        18: 24,
        19: 20,
    }
    for index, width in widths.items():
        worksheet.column_dimensions[get_column_letter(index)].width = width


def _append_detail_sheet(
    workbook: Workbook,
    title: str,
    rows: list[dict[str, Any]],
) -> None:
    worksheet = workbook.create_sheet(title=title[:31])
    headers = [
        "STT",
        "Nhóm kết quả",
        "Căn cứ khớp",
        "Mã đối tượng điều tra",
        "Họ tên điều tra",
        "Ngày sinh",
        "Số định danh",
        "Mã học sinh trên phiếu",
        "Số phiếu / Mã hộ",
        "Trạng thái điều tra",
        "Trường / lớp điều tra",
        "Mã học sinh",
        "Trạng thái học sinh",
        "Trường / lớp học sinh",
        "Nội dung phát hiện",
        "Trạng thái xử lý",
        "Nội dung ghi nhận",
        "Người ghi nhận",
        "Thời điểm ghi nhận",
    ]
    worksheet.append(headers)

    for index, row in enumerate(rows, start=1):
        person: SurveyPerson | None = row.get("person")
        student: Student | None = row.get("student")
        form: SurveyForm | None = row.get("form")
        household: Household | None = row.get("household")

        worksheet.append(
            [
                index,
                row["result_label"],
                row["match_basis_label"],
                person.code if person else "",
                person.full_name if person else "",
                (
                    person.date_of_birth.strftime("%d/%m/%Y")
                    if person and person.date_of_birth
                    else ""
                ),
                person.personal_id if person else "",
                person.ministry_student_code if person else "",
                (
                    f"{form.form_number} / {household.code}"
                    if form and household
                    else ""
                ),
                row["survey_status_label"],
                (
                    f"{row['survey_school_name']} / "
                    f"{row['survey_class_name']}"
                ),
                student.code if student else "",
                row["student_status_label"],
                (
                    f"{row['student_school_name']} / "
                    f"{row['student_class_name']}"
                ),
                row["detail"],
                row.get("resolution_action_label", "Chưa ghi nhận"),
                row.get("resolution_note", ""),
                row.get("resolution_actor_name", ""),
                row.get("resolution_created_at_text", ""),
            ]
        )

    _style_worksheet(worksheet)


@router.get("/{batch_id}/doi-chieu-hoc-sinh/xuat-excel")
def export_comparison_excel(
    request: Request,
    batch_id: int,
    result_group: str = "",
    school_id: str = "",
    q: str = "",
    resolution_state: str = "",
    db: Session = Depends(get_db),
):
    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return StreamingResponse(
            iter([b"Khong tim thay dot dieu tra."]),
            status_code=404,
            media_type="text/plain",
        )

    school_id = _parse_optional_int(school_id)

    result_group = result_group.strip().upper()
    if result_group not in RESULT_LABELS:
        result_group = ""

    resolution_state = resolution_state.strip().upper()
    if resolution_state not in {"", "PENDING", "RECORDED"}:
        resolution_state = ""

    all_rows, _ = _build_comparison_rows(db, request, batch)
    _attach_latest_resolutions(db, batch.id, all_rows)
    filtered_rows = _apply_filters(
        all_rows,
        result_group=result_group,
        school_id=school_id,
        q=q.strip()[:120],
        resolution_state=resolution_state,
    )
    stats = _build_stats(all_rows)

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Tong quan"
    overview.append(["Thông tin", "Giá trị"])
    overview_rows = [
        ("Mã đợt", batch.code),
        ("Tên đợt", batch.name),
        ("Năm học", batch.school_year.code),
        ("Xã/phường", batch.commune.name),
        ("Phạm vi", _scope_label(request)),
        ("Tổng số dòng đối chiếu", stats["TOTAL"]),
        ("Khớp hoàn toàn", stats["KHOP_HOAN_TOAN"]),
        (
            "Điều tra chưa có học sinh",
            stats["DIEU_TRA_CHUA_CO_HOC_SINH"],
        ),
        (
            "Học sinh chưa có trong điều tra",
            stats["HOC_SINH_CHUA_CO_DIEU_TRA"],
        ),
        ("Sai khác trường/lớp", stats["SAI_KHAC_TRUONG_LOP"]),
        ("Sai khác trạng thái", stats["SAI_KHAC_TRANG_THAI"]),
        ("Nghi trùng", stats["NGHI_TRUNG"]),
        ("Thiếu thông tin", stats["THIEU_THONG_TIN"]),
        ("Tổng cần xử lý", stats["CAN_XU_LY"]),
        ("Đã ghi nhận xử lý", stats["DA_GHI_NHAN_XU_LY"]),
        ("Còn chờ xử lý", stats["CON_CHO_XU_LY"]),
    ]
    for label, value in overview_rows:
        overview.append([label, value])
    _style_worksheet(overview)
    overview.column_dimensions["A"].width = 38
    overview.column_dimensions["B"].width = 55

    _append_detail_sheet(workbook, "Doi chieu chi tiet", filtered_rows)
    issue_rows = [
        row
        for row in all_rows
        if row["result_code"] != "KHOP_HOAN_TOAN"
    ]
    _append_detail_sheet(workbook, "Can xu ly", issue_rows)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    safe_code = re.sub(r"[^0-9A-Za-z_-]+", "_", batch.code)
    filename = f"doi_chieu_hoc_sinh_{safe_code}.xlsx"

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )



# =========================================================
# BÀI 12D-17 - BIÊN BẢN VÀ CHỐT KẾT QUẢ ĐỐI CHIẾU
# =========================================================


@router.get(
    "/{batch_id}/doi-chieu-hoc-sinh/chot-ket-qua",
    response_class=HTMLResponse,
)
def comparison_signoff_page(
    request: Request,
    batch_id: int,
    status: str = "",
    db: Session = Depends(get_db),
):
    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )

    rows, _ = _build_comparison_rows(db, request, batch)
    _attach_latest_resolutions(db, batch.id, rows)
    stats = _build_stats(rows)
    history = _load_signoff_history(db, batch.id)
    latest = history[0] if history else None
    is_signed = bool(latest is not None and latest.action_code == "CHOT")
    ready_to_sign = stats["TOTAL"] > 0 and stats["CON_CHO_XU_LY"] == 0

    return templates.TemplateResponse(
        request=request,
        name="surveys/student_survey_signoff.html",
        context={
            "nguoi_dung": lay_thong_tin_nguoi_dung(request),
            "batch": batch,
            "scope_label": _scope_label(request),
            "stats": stats,
            "history": history,
            "latest_signoff": latest,
            "comparison_signed": is_signed,
            "ready_to_sign": ready_to_sign,
            "can_manage_signoff": _can_manage_signoff(request),
            "status_message": _signoff_status_message(status),
            "signoff_action_labels": SIGNOFF_ACTION_LABELS,
        },
    )


@router.post(
    "/{batch_id}/doi-chieu-hoc-sinh/chot-ket-qua/chot"
)
def sign_comparison_result(
    request: Request,
    batch_id: int,
    note: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    if not _can_manage_signoff(request):
        return RedirectResponse(
            url=_signoff_redirect_url(batch_id, "forbidden"),
            status_code=303,
        )

    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )

    if _comparison_is_signed(db, batch.id):
        return RedirectResponse(
            url=_signoff_redirect_url(batch.id, "already_signed"),
            status_code=303,
        )

    clean_note = _clean_text(note)
    if len(clean_note) < 5:
        return RedirectResponse(
            url=_signoff_redirect_url(batch.id, "invalid_note"),
            status_code=303,
        )

    rows, _ = _build_comparison_rows(db, request, batch)
    _attach_latest_resolutions(db, batch.id, rows)
    stats = _build_stats(rows)
    if stats["TOTAL"] <= 0 or stats["CON_CHO_XU_LY"] > 0:
        return RedirectResponse(
            url=_signoff_redirect_url(batch.id, "pending"),
            status_code=303,
        )

    _add_signoff_log(
        db,
        request,
        batch,
        action_code="CHOT",
        note=clean_note,
        stats=stats,
    )
    try:
        db.commit()
    except Exception:
        db.rollback()
        return RedirectResponse(
            url=_signoff_redirect_url(batch.id, "database_error"),
            status_code=303,
        )

    return RedirectResponse(
        url=_signoff_redirect_url(batch.id, "signed"),
        status_code=303,
    )


@router.post(
    "/{batch_id}/doi-chieu-hoc-sinh/chot-ket-qua/mo-lai"
)
def reopen_comparison_result(
    request: Request,
    batch_id: int,
    note: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    if not _can_manage_signoff(request):
        return RedirectResponse(
            url=_signoff_redirect_url(batch_id, "forbidden"),
            status_code=303,
        )

    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )

    if not _comparison_is_signed(db, batch.id):
        return RedirectResponse(
            url=_signoff_redirect_url(batch.id, "not_signed"),
            status_code=303,
        )

    clean_note = _clean_text(note)
    if len(clean_note) < 5:
        return RedirectResponse(
            url=_signoff_redirect_url(batch.id, "invalid_note"),
            status_code=303,
        )

    rows, _ = _build_comparison_rows(db, request, batch)
    _attach_latest_resolutions(db, batch.id, rows)
    stats = _build_stats(rows)
    _add_signoff_log(
        db,
        request,
        batch,
        action_code="MO_LAI",
        note=clean_note,
        stats=stats,
    )
    try:
        db.commit()
    except Exception:
        db.rollback()
        return RedirectResponse(
            url=_signoff_redirect_url(batch.id, "database_error"),
            status_code=303,
        )

    return RedirectResponse(
        url=_signoff_redirect_url(batch.id, "reopened"),
        status_code=303,
    )


def _style_signoff_overview(worksheet: Any) -> None:
    title_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="C7D3DF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    worksheet.merge_cells("A1:D1")
    worksheet["A1"] = "BIÊN BẢN XÁC NHẬN KẾT QUẢ ĐỐI CHIẾU DỮ LIỆU"
    worksheet["A1"].fill = title_fill
    worksheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 28

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row_number in (3, 10, 18):
        for cell in worksheet[row_number]:
            cell.fill = section_fill
            cell.font = Font(bold=True)

    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 42
    worksheet.column_dimensions["C"].width = 30
    worksheet.column_dimensions["D"].width = 42
    worksheet.freeze_panes = "A2"


def _append_signoff_history_sheet(
    workbook: Workbook,
    history: list[StudentSurveyComparisonSignoff],
) -> None:
    worksheet = workbook.create_sheet("Lich su chot")
    worksheet.append(
        [
            "STT",
            "Thao tác",
            "Người thực hiện",
            "Vai trò",
            "Thời điểm",
            "Tổng dòng",
            "Khớp hoàn toàn",
            "Cần xử lý",
            "Đã ghi nhận",
            "Còn chờ",
            "Lý do / căn cứ",
        ]
    )
    for index, item in enumerate(history, start=1):
        worksheet.append(
            [
                index,
                SIGNOFF_ACTION_LABELS.get(item.action_code, item.action_code),
                item.actor_name_snapshot,
                item.actor_role_snapshot,
                item.created_at.strftime("%d/%m/%Y %H:%M"),
                item.total_rows,
                item.matched_rows,
                item.issue_rows,
                item.resolved_rows,
                item.pending_rows,
                item.note,
            ]
        )
    _style_worksheet(worksheet)
    worksheet.column_dimensions["B"].width = 28
    worksheet.column_dimensions["C"].width = 24
    worksheet.column_dimensions["D"].width = 22
    worksheet.column_dimensions["E"].width = 20
    worksheet.column_dimensions["K"].width = 60


@router.get(
    "/{batch_id}/doi-chieu-hoc-sinh/chot-ket-qua/xuat-bien-ban"
)
def export_comparison_signoff_minutes(
    request: Request,
    batch_id: int,
    db: Session = Depends(get_db),
):
    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return StreamingResponse(
            iter([b"Khong tim thay dot dieu tra."]),
            status_code=404,
            media_type="text/plain",
        )

    rows, _ = _build_comparison_rows(db, request, batch)
    _attach_latest_resolutions(db, batch.id, rows)
    current_stats = _build_stats(rows)
    history = _load_signoff_history(db, batch.id)
    latest = history[0] if history else None
    is_signed = bool(latest is not None and latest.action_code == "CHOT")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Bien ban doi chieu"
    worksheet.append([""])
    worksheet.append(["Trạng thái biên bản", "CHÍNH THỨC" if is_signed else "DỰ THẢO", "Mã đợt", batch.code])
    worksheet.append(["I. THÔNG TIN ĐỢT ĐIỀU TRA", "", "", ""])
    worksheet.append(["Tên đợt", batch.name, "Năm học", batch.school_year.code])
    worksheet.append(["Xã/phường", batch.commune.name, "Phạm vi", _scope_label(request)])
    worksheet.append(["Trạng thái đối chiếu", "Đã chốt" if is_signed else "Chưa chốt", "Ngày xuất", "Tự động theo thời điểm tải tệp"])
    worksheet.append(["", "", "", ""])
    worksheet.append(["", "", "", ""])
    worksheet.append(["", "", "", ""])
    worksheet.append(["II. SỐ LIỆU TỔNG HỢP HIỆN TẠI", "", "", ""])
    worksheet.append(["Tổng số dòng", current_stats["TOTAL"], "Khớp hoàn toàn", current_stats["KHOP_HOAN_TOAN"]])
    worksheet.append(["Tổng cần xử lý", current_stats["CAN_XU_LY"], "Đã ghi nhận xử lý", current_stats["DA_GHI_NHAN_XU_LY"]])
    worksheet.append(["Còn chờ xử lý", current_stats["CON_CHO_XU_LY"], "Điều tra chưa có học sinh", current_stats["DIEU_TRA_CHUA_CO_HOC_SINH"]])
    worksheet.append(["Học sinh chưa có điều tra", current_stats["HOC_SINH_CHUA_CO_DIEU_TRA"], "Sai trường/lớp", current_stats["SAI_KHAC_TRUONG_LOP"]])
    worksheet.append(["Sai trạng thái", current_stats["SAI_KHAC_TRANG_THAI"], "Nghi trùng/thiếu thông tin", current_stats["NGHI_TRUNG"] + current_stats["THIEU_THONG_TIN"]])
    worksheet.append(["", "", "", ""])
    worksheet.append(["", "", "", ""])
    worksheet.append(["III. LẦN XÁC NHẬN GẦN NHẤT", "", "", ""])
    if latest is None:
        worksheet.append(["Kết quả", "Chưa có lịch sử chốt", "", ""])
        worksheet.append(["Căn cứ", "", "", ""])
    else:
        worksheet.append(["Thao tác", SIGNOFF_ACTION_LABELS.get(latest.action_code, latest.action_code), "Thời điểm", latest.created_at.strftime("%d/%m/%Y %H:%M")])
        worksheet.append(["Người thực hiện", latest.actor_name_snapshot, "Vai trò", latest.actor_role_snapshot])
        worksheet.append(["Căn cứ / lý do", latest.note, "Còn chờ tại thời điểm thao tác", latest.pending_rows])
    worksheet.append(["", "", "", ""])
    worksheet.append(["Người lập biểu", "", "Người xác nhận", ""])
    worksheet.append(["(Ký, ghi rõ họ tên)", "", "(Ký, ghi rõ họ tên)", ""])
    _style_signoff_overview(worksheet)
    _append_signoff_history_sheet(workbook, history)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    safe_code = re.sub(r"[^0-9A-Za-z_-]+", "_", batch.code)
    filename = f"bien_ban_doi_chieu_{safe_code}.xlsx"
    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )


# =========================================================
# BÀI 12D-16 - XỬ LÝ KẾT QUẢ ĐỐI CHIẾU
# =========================================================


def _can_write_resolution(request: Request) -> bool:
    user = lay_thong_tin_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))
    return is_admin_role(role_code) or role_code in {SCHOOL_ROLE_CODE, TEACHER_ROLE_CODE}


def _comparison_status_message(status: str) -> dict[str, str] | None:
    messages = {
        "linked": (
            "success",
            "Đã liên kết đối tượng điều tra với hồ sơ học sinh. "
            "Kết quả đối chiếu đã được tính lại.",
        ),
        "unlinked": (
            "success",
            "Đã bỏ liên kết học sinh và ghi nhật ký xử lý.",
        ),
        "recorded": (
            "success",
            "Đã lưu nội dung rà soát vào nhật ký xử lý.",
        ),
        "not_found": (
            "warning",
            "Không tìm thấy dòng đối chiếu trong phạm vi được phép xem.",
        ),
        "forbidden": (
            "warning",
            "Tài khoản hiện tại chỉ được xem, không được ghi nhận xử lý.",
        ),
        "invalid": (
            "warning",
            "Thông tin gửi lên không hợp lệ. Dữ liệu chưa được thay đổi.",
        ),
        "duplicate_link": (
            "warning",
            "Hồ sơ học sinh đã được liên kết với một đối tượng khác.",
        ),
        "conflict_not_confirmed": (
            "warning",
            "Thông tin nhận diện có xung đột. Hãy đánh dấu xác nhận "
            "sau khi đã kiểm tra hồ sơ gốc.",
        ),
        "student_out_of_scope": (
            "warning",
            "Hồ sơ học sinh không thuộc phạm vi tài khoản được phép xử lý.",
        ),
        "database_error": (
            "warning",
            "Không thể lưu thay đổi vào cơ sở dữ liệu. Hệ thống đã hoàn tác.",
        ),
        "signoff_locked": (
            "warning",
            "Kết quả đối chiếu đã được chốt. Hãy mở lại trước khi thay đổi liên kết hoặc ghi nhận xử lý.",
        ),
    }
    item = messages.get(str(status or "").strip())
    if item is None:
        return None
    return {"kind": item[0], "text": item[1]}


def _load_latest_resolution_maps(
    db: Session,
    batch_id: int,
) -> tuple[dict[int, dict[str, Any]], dict[int, dict[str, Any]]]:
    person_map: dict[int, dict[str, Any]] = {}
    student_map: dict[int, dict[str, Any]] = {}

    records = db.execute(
        select(StudentSurveyResolutionLog, User)
        .outerjoin(
            User,
            User.id == StudentSurveyResolutionLog.actor_user_id,
        )
        .where(StudentSurveyResolutionLog.survey_batch_id == batch_id)
        .order_by(
            StudentSurveyResolutionLog.created_at.asc(),
            StudentSurveyResolutionLog.id.asc(),
        )
    ).all()

    for log, actor in records:
        item = {
            "log": log,
            "action_label": RESOLUTION_ACTION_LABELS.get(
                log.action_code,
                log.action_code,
            ),
            "done": log.action_code in RESOLUTION_DONE_ACTIONS,
            "actor_name": actor.full_name if actor is not None else "Hệ thống",
            "created_at_text": (
                log.created_at.strftime("%d/%m/%Y %H:%M")
                if log.created_at is not None
                else ""
            ),
        }
        if log.survey_person_id is not None:
            person_map[int(log.survey_person_id)] = item
        if log.student_id is not None:
            student_map[int(log.student_id)] = item

    return person_map, student_map


def _attach_latest_resolutions(
    db: Session,
    batch_id: int,
    rows: list[dict[str, Any]],
) -> None:
    person_map, student_map = _load_latest_resolution_maps(db, batch_id)

    for row in rows:
        resolution: dict[str, Any] | None = None
        person: SurveyPerson | None = row.get("person")
        student: Student | None = row.get("student")

        if person is not None:
            resolution = person_map.get(person.id)
        elif student is not None:
            resolution = student_map.get(student.id)

        row["resolution"] = resolution
        row["resolution_done"] = bool(
            resolution is not None and resolution.get("done")
        )
        row["resolution_action_label"] = (
            resolution["action_label"]
            if resolution is not None
            else "Chưa ghi nhận"
        )
        row["resolution_note"] = (
            _clean_text(resolution["log"].note)
            if resolution is not None
            else ""
        )
        row["resolution_actor_name"] = (
            resolution["actor_name"] if resolution is not None else ""
        )
        row["resolution_created_at_text"] = (
            resolution["created_at_text"] if resolution is not None else ""
        )


def _load_resolution_history(
    db: Session,
    batch_id: int,
    *,
    person_id: int | None = None,
    student_id: int | None = None,
) -> list[dict[str, Any]]:
    statement = (
        select(StudentSurveyResolutionLog, User)
        .outerjoin(
            User,
            User.id == StudentSurveyResolutionLog.actor_user_id,
        )
        .where(StudentSurveyResolutionLog.survey_batch_id == batch_id)
    )

    if person_id is not None:
        statement = statement.where(
            StudentSurveyResolutionLog.survey_person_id == person_id
        )
    elif student_id is not None:
        statement = statement.where(
            StudentSurveyResolutionLog.student_id == student_id
        )
    else:
        return []

    records = db.execute(
        statement.order_by(
            StudentSurveyResolutionLog.created_at.desc(),
            StudentSurveyResolutionLog.id.desc(),
        )
    ).all()

    return [
        {
            "log": log,
            "action_label": RESOLUTION_ACTION_LABELS.get(
                log.action_code,
                log.action_code,
            ),
            "actor_name": actor.full_name if actor is not None else "Hệ thống",
            "created_at_text": (
                log.created_at.strftime("%d/%m/%Y %H:%M")
                if log.created_at is not None
                else ""
            ),
        }
        for log, actor in records
    ]


def _person_student_match_info(
    person: SurveyPerson,
    student: Student,
) -> dict[str, Any]:
    score = 0
    basis = "Không có căn cứ tự động"

    if person.student_id == student.id:
        score = 100
        basis = "Liên kết hiện tại"
    elif (
        _normalize_code(person.personal_id)
        and _normalize_code(person.personal_id)
        == _normalize_code(student.personal_id)
    ):
        score = 95
        basis = "Trùng số định danh cá nhân"
    elif (
        _normalize_code(person.ministry_student_code)
        and _normalize_code(person.ministry_student_code)
        == _normalize_code(student.code)
    ):
        score = 90
        basis = "Trùng mã học sinh"
    elif (
        _normalize_text(person.full_name)
        == _normalize_text(student.full_name)
        and person.date_of_birth is not None
        and person.date_of_birth == student.date_of_birth
        and _normalize_gender(person.gender)
        and _normalize_gender(person.gender)
        == _normalize_gender(student.gender)
    ):
        score = 80
        basis = "Trùng họ tên, ngày sinh và giới tính"
    elif (
        _normalize_text(person.full_name)
        == _normalize_text(student.full_name)
        and person.date_of_birth is not None
        and person.date_of_birth == student.date_of_birth
    ):
        score = 75
        basis = "Trùng họ tên và ngày sinh"
    elif _normalize_text(person.full_name) == _normalize_text(student.full_name):
        score = 45
        basis = "Trùng họ tên"

    conflicts: list[str] = []
    if (
        _normalize_code(person.personal_id)
        and _normalize_code(student.personal_id)
        and _normalize_code(person.personal_id)
        != _normalize_code(student.personal_id)
    ):
        conflicts.append("Số định danh khác nhau")
    if (
        _normalize_code(person.ministry_student_code)
        and _normalize_code(person.ministry_student_code)
        != _normalize_code(student.code)
    ):
        conflicts.append("Mã học sinh trên phiếu khác hồ sơ được chọn")
    if (
        person.date_of_birth is not None
        and student.date_of_birth is not None
        and person.date_of_birth != student.date_of_birth
    ):
        conflicts.append("Ngày sinh khác nhau")
    if (
        _normalize_gender(person.gender)
        and _normalize_gender(student.gender)
        and _normalize_gender(person.gender)
        != _normalize_gender(student.gender)
    ):
        conflicts.append("Giới tính khác nhau")

    return {
        "score": score,
        "basis": basis,
        "conflicts": conflicts,
        "requires_confirmation": bool(conflicts),
    }


def _student_search_text(student: Student) -> str:
    return _normalize_text(
        " ".join(
            [
                student.code,
                student.full_name,
                student.personal_id or "",
                student.father_name or "",
                student.mother_name or "",
                student.contact_phone or "",
            ]
        )
    )


def _person_search_text(person: SurveyPerson) -> str:
    return _normalize_text(
        " ".join(
            [
                person.code,
                person.full_name,
                person.personal_id or "",
                person.ministry_student_code or "",
                person.father_name or "",
                person.mother_name or "",
                person.contact_phone or "",
            ]
        )
    )


def _load_person_student_candidates(
    db: Session,
    request: Request,
    batch: SurveyBatch,
    survey_item: dict[str, Any],
    q: str,
) -> list[dict[str, Any]]:
    person: SurveyPerson = survey_item["person"]
    candidates: dict[int, Student] = {}

    scope_enrollments = _load_scope_enrollments(db, request, batch)
    for enrollment in scope_enrollments:
        candidates[enrollment.student_id] = enrollment.student

    extra_students = _load_extra_candidate_students(
        db,
        [survey_item],
        set(candidates),
    )
    for student in extra_students:
        candidates[student.id] = student

    user = lay_thong_tin_nguoi_dung(request)
    if q and is_admin_role(user.get("role_code")):
        pattern = f"%{q.strip()}%"
        searched = db.scalars(
            select(Student)
            .where(
                Student.is_active.is_(True),
                or_(
                    Student.full_name.ilike(pattern),
                    Student.code.ilike(pattern),
                    Student.personal_id.ilike(pattern),
                    Student.father_name.ilike(pattern),
                    Student.mother_name.ilike(pattern),
                    Student.contact_phone.ilike(pattern),
                ),
            )
            .order_by(Student.full_name.asc(), Student.code.asc())
            .limit(120)
        ).all()
        for student in searched:
            candidates[student.id] = student

    if person.student_id is not None and person.student_id not in candidates:
        linked_student = db.get(Student, int(person.student_id))
        if linked_student is not None and linked_student.is_active:
            candidates[linked_student.id] = linked_student

    baseline_year_id = _baseline_year_id(db, batch)
    allowed_school_ids = _baseline_allowed_school_ids(db, request, batch)
    enrollment_map = _load_target_enrollments_for_students(
        db,
        baseline_year_id if baseline_year_id is not None else -1,
        set(candidates),
        allowed_school_ids=allowed_school_ids,
    )

    linked_map: dict[int, SurveyPerson] = {}
    if candidates:
        linked_people = db.scalars(
            select(SurveyPerson).where(
                SurveyPerson.student_id.in_(list(candidates)),
                SurveyPerson.is_active.is_(True),
            )
        ).all()
        for linked_person in linked_people:
            linked_map[int(linked_person.student_id)] = linked_person

    q_normalized = _normalize_text(q)
    result: list[dict[str, Any]] = []
    for student in candidates.values():
        match_info = _person_student_match_info(person, student)
        if q_normalized:
            if q_normalized not in _student_search_text(student):
                continue
        elif match_info["score"] <= 0 and person.student_id != student.id:
            continue

        enrollment = _choose_enrollment(
            enrollment_map.get(student.id, []),
            preferred_school_id=(
                survey_item["record"].school_id
                if survey_item.get("record") is not None
                else None
            ),
        )
        linked_person = linked_map.get(student.id)
        linked_to_other = (
            linked_person is not None and linked_person.id != person.id
        )

        result.append(
            {
                "student": student,
                "enrollment": enrollment,
                "score": match_info["score"],
                "basis": match_info["basis"],
                "conflicts": match_info["conflicts"],
                "requires_confirmation": match_info[
                    "requires_confirmation"
                ],
                "linked_to_other": linked_to_other,
                "linked_person": linked_person,
                "can_link": not linked_to_other,
                "school_name": (
                    enrollment.school.name
                    if enrollment is not None and enrollment.school is not None
                    else "Chưa có xếp lớp năm học này"
                ),
                "class_name": (
                    enrollment.classroom.name
                    if enrollment is not None
                    and enrollment.classroom is not None
                    else "—"
                ),
                "status_label": _student_status_label(enrollment),
            }
        )

    result.sort(
        key=lambda item: (
            item["linked_to_other"],
            -item["score"],
            _normalize_text(item["student"].full_name),
            item["student"].id,
        )
    )
    return result[:60]


def _load_student_person_candidates(
    db: Session,
    request: Request,
    batch: SurveyBatch,
    student: Student,
    q: str,
) -> list[dict[str, Any]]:
    survey_rows = _load_survey_rows(db, request, batch)
    q_normalized = _normalize_text(q)
    result: list[dict[str, Any]] = []

    for item in survey_rows:
        person: SurveyPerson = item["person"]
        match_info = _person_student_match_info(person, student)

        if q_normalized:
            if q_normalized not in _person_search_text(person):
                continue
        elif match_info["score"] <= 0:
            continue

        linked_to_other = (
            person.student_id is not None and person.student_id != student.id
        )
        result.append(
            {
                "survey_item": item,
                "person": person,
                "record": item.get("record"),
                "household": item["household"],
                "form": item["form"],
                "score": match_info["score"],
                "basis": match_info["basis"],
                "conflicts": match_info["conflicts"],
                "requires_confirmation": match_info[
                    "requires_confirmation"
                ],
                "linked_to_other": linked_to_other,
                "can_link": not linked_to_other,
            }
        )

    result.sort(
        key=lambda item: (
            item["linked_to_other"],
            -item["score"],
            _normalize_text(item["person"].full_name),
            item["person"].id,
        )
    )
    return result[:60]


def _find_survey_row(
    rows: list[dict[str, Any]],
    person_id: int,
) -> dict[str, Any] | None:
    return next(
        (
            row
            for row in rows
            if row.get("person") is not None
            and row["person"].id == person_id
        ),
        None,
    )


def _find_student_only_row(
    rows: list[dict[str, Any]],
    student_id: int,
) -> dict[str, Any] | None:
    return next(
        (
            row
            for row in rows
            if row.get("row_type") == "STUDENT"
            and row.get("student") is not None
            and row["student"].id == student_id
        ),
        None,
    )


def _student_is_writeable_in_scope(
    db: Session,
    request: Request,
    batch: SurveyBatch,
    student_id: int,
) -> bool:
    user = lay_thong_tin_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))
    baseline_year_id = _baseline_year_id(db, batch)
    if baseline_year_id is None:
        return False
    allowed_school_ids = _baseline_allowed_school_ids(db, request, batch)
    if not allowed_school_ids:
        return False
    if is_admin_role(role_code):
        return db.scalar(
            select(StudentEnrollment.id).where(
                StudentEnrollment.student_id == student_id,
                StudentEnrollment.school_year_id == baseline_year_id,
                StudentEnrollment.school_id.in_(sorted(allowed_school_ids)),
            )
        ) is not None
    if role_code not in {SCHOOL_ROLE_CODE, TEACHER_ROLE_CODE}:
        return False
    return db.scalar(
        select(StudentEnrollment.id).where(
            StudentEnrollment.student_id == student_id,
            StudentEnrollment.school_year_id == baseline_year_id,
            StudentEnrollment.school_id.in_(sorted(allowed_school_ids)),
        )
    ) is not None



def _identity_conflicts_for_link(
    person: SurveyPerson,
    student: Student,
) -> list[str]:
    return _person_student_match_info(person, student)["conflicts"]


def _add_resolution_log(
    db: Session,
    request: Request,
    batch: SurveyBatch,
    *,
    action_code: str,
    result_before: str | None,
    person_id: int | None,
    student_id: int | None,
    previous_student_id: int | None,
    note: str,
) -> None:
    user = lay_thong_tin_nguoi_dung(request)
    db.add(
        StudentSurveyResolutionLog(
            survey_batch_id=batch.id,
            survey_person_id=person_id,
            student_id=student_id,
            previous_student_id=previous_student_id,
            action_code=action_code,
            result_before=result_before,
            note=_clean_text(note)[:1000] or None,
            actor_user_id=user.get("id"),
        )
    )


def _resolution_detail_url(
    batch_id: int,
    *,
    person_id: int | None = None,
    student_id: int | None = None,
    status: str = "",
    q: str = "",
) -> str:
    if person_id is not None:
        base = (
            f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh/"
            f"xu-ly/doi-tuong/{person_id}"
        )
    else:
        base = (
            f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh/"
            f"xu-ly/hoc-sinh/{student_id}"
        )
    params = {"status": status, "q": q}
    return base + "?" + urlencode(params)


@router.get(
    "/{batch_id}/doi-chieu-hoc-sinh/xu-ly/doi-tuong/{person_id}",
    response_class=HTMLResponse,
)
def survey_person_resolution_page(
    request: Request,
    batch_id: int,
    person_id: int,
    q: str = "",
    status: str = "",
    db: Session = Depends(get_db),
):
    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )

    rows, _ = _build_comparison_rows(db, request, batch)
    _attach_latest_resolutions(db, batch.id, rows)
    comparison_row = _find_survey_row(rows, person_id)
    if comparison_row is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )

    survey_item = {
        "form": comparison_row["form"],
        "household": comparison_row["household"],
        "person": comparison_row["person"],
        "record": comparison_row["record"],
    }
    q = q.strip()[:120]
    candidates = _load_person_student_candidates(
        db,
        request,
        batch,
        survey_item,
        q,
    )
    history = _load_resolution_history(
        db,
        batch.id,
        person_id=person_id,
    )

    return templates.TemplateResponse(
        request=request,
        name="surveys/student_survey_resolution.html",
        context={
            "nguoi_dung": lay_thong_tin_nguoi_dung(request),
            "batch": batch,
            "mode": "SURVEY",
            "comparison_row": comparison_row,
            "person": comparison_row["person"],
            "student": comparison_row.get("student"),
            "enrollment": comparison_row.get("enrollment"),
            "household": comparison_row["household"],
            "form": comparison_row["form"],
            "record": comparison_row.get("record"),
            "candidates": candidates,
            "history": history,
            "q": q,
            "can_write": (
                _can_write_resolution(request)
                and not _comparison_is_signed(db, batch.id)
            ),
            "status_message": _comparison_status_message(status),
            "resolution_action_labels": RESOLUTION_ACTION_LABELS,
        },
    )


@router.get(
    "/{batch_id}/doi-chieu-hoc-sinh/xu-ly/hoc-sinh/{student_id}",
    response_class=HTMLResponse,
)
def student_resolution_page(
    request: Request,
    batch_id: int,
    student_id: int,
    q: str = "",
    status: str = "",
    db: Session = Depends(get_db),
):
    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )

    rows, _ = _build_comparison_rows(db, request, batch)
    _attach_latest_resolutions(db, batch.id, rows)
    comparison_row = _find_student_only_row(rows, student_id)
    if comparison_row is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )

    student: Student = comparison_row["student"]
    q = q.strip()[:120]
    candidates = _load_student_person_candidates(
        db,
        request,
        batch,
        student,
        q,
    )
    history = _load_resolution_history(
        db,
        batch.id,
        student_id=student.id,
    )

    return templates.TemplateResponse(
        request=request,
        name="surveys/student_survey_resolution.html",
        context={
            "nguoi_dung": lay_thong_tin_nguoi_dung(request),
            "batch": batch,
            "mode": "STUDENT",
            "comparison_row": comparison_row,
            "person": None,
            "student": student,
            "enrollment": comparison_row.get("enrollment"),
            "household": None,
            "form": None,
            "record": None,
            "candidates": candidates,
            "history": history,
            "q": q,
            "can_write": (
                _can_write_resolution(request)
                and not _comparison_is_signed(db, batch.id)
            ),
            "status_message": _comparison_status_message(status),
            "resolution_action_labels": RESOLUTION_ACTION_LABELS,
        },
    )


@router.post(
    "/{batch_id}/doi-chieu-hoc-sinh/xu-ly/doi-tuong/{person_id}/lien-ket"
)
def link_survey_person_to_student(
    request: Request,
    batch_id: int,
    person_id: int,
    student_id: Annotated[str, Form()] = "",
    note: Annotated[str, Form()] = "",
    confirm_conflict: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    if not _can_write_resolution(request):
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=forbidden",
            status_code=303,
        )

    batch = _batch_in_scope(db, request, batch_id)
    if batch is None or not student_id.isdigit():
        return RedirectResponse(
            url=_resolution_detail_url(
                batch_id,
                person_id=person_id,
                status="invalid",
            ),
            status_code=303,
        )
    if _comparison_is_signed(db, batch.id):
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                person_id=person_id,
                status="signoff_locked",
            ),
            status_code=303,
        )

    rows, _ = _build_comparison_rows(db, request, batch)
    comparison_row = _find_survey_row(rows, person_id)
    if comparison_row is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )

    person: SurveyPerson = comparison_row["person"]
    selected_student_id = int(student_id)
    student = db.get(Student, selected_student_id)
    if student is None or not student.is_active:
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                person_id=person.id,
                status="invalid",
            ),
            status_code=303,
        )

    if not _student_is_writeable_in_scope(
        db,
        request,
        batch,
        selected_student_id,
    ):
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                person_id=person.id,
                status="student_out_of_scope",
            ),
            status_code=303,
        )

    existing_link = db.scalar(
        select(SurveyPerson).where(
            SurveyPerson.student_id == selected_student_id,
            SurveyPerson.id != person.id,
            SurveyPerson.is_active.is_(True),
        )
    )
    if existing_link is not None:
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                person_id=person.id,
                status="duplicate_link",
            ),
            status_code=303,
        )

    conflicts = _identity_conflicts_for_link(person, student)
    if conflicts and confirm_conflict != "1":
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                person_id=person.id,
                status="conflict_not_confirmed",
            ),
            status_code=303,
        )

    previous_student_id = person.student_id
    person.student_id = selected_student_id
    _add_resolution_log(
        db,
        request,
        batch,
        action_code="LIEN_KET",
        result_before=comparison_row["result_code"],
        person_id=person.id,
        student_id=student.id,
        previous_student_id=previous_student_id,
        note=note or (
            "Liên kết thủ công sau khi đối chiếu hồ sơ."
            + (" Đã xác nhận xung đột nhận diện." if conflicts else "")
        ),
    )

    try:
        db.commit()
    except Exception:
        db.rollback()
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                person_id=person.id,
                status="database_error",
            ),
            status_code=303,
        )

    return RedirectResponse(
        url=(
            f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh?"
            + urlencode({"status": "linked", "q": person.full_name})
        ),
        status_code=303,
    )


@router.post(
    "/{batch_id}/doi-chieu-hoc-sinh/xu-ly/hoc-sinh/{student_id}/lien-ket"
)
def link_student_to_survey_person(
    request: Request,
    batch_id: int,
    student_id: int,
    person_id: Annotated[str, Form()] = "",
    note: Annotated[str, Form()] = "",
    confirm_conflict: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    if not _can_write_resolution(request):
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=forbidden",
            status_code=303,
        )

    batch = _batch_in_scope(db, request, batch_id)
    if batch is None or not person_id.isdigit():
        return RedirectResponse(
            url=_resolution_detail_url(
                batch_id,
                student_id=student_id,
                status="invalid",
            ),
            status_code=303,
        )
    if _comparison_is_signed(db, batch.id):
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                student_id=student_id,
                status="signoff_locked",
            ),
            status_code=303,
        )

    rows, _ = _build_comparison_rows(db, request, batch)
    student_row = _find_student_only_row(rows, student_id)
    person_row = _find_survey_row(rows, int(person_id))
    if student_row is None or person_row is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )

    if not _student_is_writeable_in_scope(
        db,
        request,
        batch,
        student_id,
    ):
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                student_id=student_id,
                status="student_out_of_scope",
            ),
            status_code=303,
        )

    student: Student = student_row["student"]
    person: SurveyPerson = person_row["person"]

    if person.student_id is not None and person.student_id != student.id:
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                student_id=student.id,
                status="duplicate_link",
            ),
            status_code=303,
        )

    existing_link = db.scalar(
        select(SurveyPerson).where(
            SurveyPerson.student_id == student.id,
            SurveyPerson.id != person.id,
            SurveyPerson.is_active.is_(True),
        )
    )
    if existing_link is not None:
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                student_id=student.id,
                status="duplicate_link",
            ),
            status_code=303,
        )

    conflicts = _identity_conflicts_for_link(person, student)
    if conflicts and confirm_conflict != "1":
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                student_id=student.id,
                status="conflict_not_confirmed",
            ),
            status_code=303,
        )

    previous_student_id = person.student_id
    person.student_id = student.id
    _add_resolution_log(
        db,
        request,
        batch,
        action_code="LIEN_KET",
        result_before=student_row["result_code"],
        person_id=person.id,
        student_id=student.id,
        previous_student_id=previous_student_id,
        note=note or (
            "Liên kết thủ công từ dòng học sinh chưa có trong điều tra."
            + (" Đã xác nhận xung đột nhận diện." if conflicts else "")
        ),
    )

    try:
        db.commit()
    except Exception:
        db.rollback()
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                student_id=student.id,
                status="database_error",
            ),
            status_code=303,
        )

    return RedirectResponse(
        url=(
            f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh?"
            + urlencode({"status": "linked", "q": person.full_name})
        ),
        status_code=303,
    )


@router.post(
    "/{batch_id}/doi-chieu-hoc-sinh/xu-ly/doi-tuong/{person_id}/ghi-nhan"
)
def record_survey_person_resolution(
    request: Request,
    batch_id: int,
    person_id: int,
    action_code: Annotated[str, Form()] = "",
    note: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    if not _can_write_resolution(request):
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=forbidden",
            status_code=303,
        )

    allowed_actions = {
        "DA_RA_SOAT",
        "KHONG_TIM_THAY",
        "XAC_NHAN_SAI_KHAC",
        "CHO_XAC_MINH",
    }
    action_code = action_code.strip().upper()
    if action_code not in allowed_actions:
        return RedirectResponse(
            url=_resolution_detail_url(
                batch_id,
                person_id=person_id,
                status="invalid",
            ),
            status_code=303,
        )

    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )
    if _comparison_is_signed(db, batch.id):
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                person_id=person_id,
                status="signoff_locked",
            ),
            status_code=303,
        )

    rows, _ = _build_comparison_rows(db, request, batch)
    comparison_row = _find_survey_row(rows, person_id)
    if comparison_row is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )

    person: SurveyPerson = comparison_row["person"]
    _add_resolution_log(
        db,
        request,
        batch,
        action_code=action_code,
        result_before=comparison_row["result_code"],
        person_id=person.id,
        student_id=person.student_id,
        previous_student_id=person.student_id,
        note=note,
    )

    try:
        db.commit()
    except Exception:
        db.rollback()
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                person_id=person.id,
                status="database_error",
            ),
            status_code=303,
        )

    return RedirectResponse(
        url=_resolution_detail_url(
            batch.id,
            person_id=person.id,
            status="recorded",
        ),
        status_code=303,
    )


@router.post(
    "/{batch_id}/doi-chieu-hoc-sinh/xu-ly/hoc-sinh/{student_id}/ghi-nhan"
)
def record_student_resolution(
    request: Request,
    batch_id: int,
    student_id: int,
    action_code: Annotated[str, Form()] = "",
    note: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    if not _can_write_resolution(request):
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=forbidden",
            status_code=303,
        )

    allowed_actions = {
        "DA_RA_SOAT",
        "XAC_NHAN_SAI_KHAC",
        "CHO_XAC_MINH",
    }
    action_code = action_code.strip().upper()
    if action_code not in allowed_actions:
        return RedirectResponse(
            url=_resolution_detail_url(
                batch_id,
                student_id=student_id,
                status="invalid",
            ),
            status_code=303,
        )

    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )
    if _comparison_is_signed(db, batch.id):
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                student_id=student_id,
                status="signoff_locked",
            ),
            status_code=303,
        )

    rows, _ = _build_comparison_rows(db, request, batch)
    comparison_row = _find_student_only_row(rows, student_id)
    if comparison_row is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )

    _add_resolution_log(
        db,
        request,
        batch,
        action_code=action_code,
        result_before=comparison_row["result_code"],
        person_id=None,
        student_id=student_id,
        previous_student_id=None,
        note=note,
    )

    try:
        db.commit()
    except Exception:
        db.rollback()
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                student_id=student_id,
                status="database_error",
            ),
            status_code=303,
        )

    return RedirectResponse(
        url=_resolution_detail_url(
            batch.id,
            student_id=student_id,
            status="recorded",
        ),
        status_code=303,
    )


@router.post(
    "/{batch_id}/doi-chieu-hoc-sinh/xu-ly/doi-tuong/{person_id}/bo-lien-ket"
)
def unlink_survey_person_student(
    request: Request,
    batch_id: int,
    person_id: int,
    note: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    if not _can_write_resolution(request):
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=forbidden",
            status_code=303,
        )

    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch_id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )
    if _comparison_is_signed(db, batch.id):
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                person_id=person_id,
                status="signoff_locked",
            ),
            status_code=303,
        )

    rows, _ = _build_comparison_rows(db, request, batch)
    comparison_row = _find_survey_row(rows, person_id)
    if comparison_row is None:
        return RedirectResponse(
            url=f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh?status=not_found",
            status_code=303,
        )

    person: SurveyPerson = comparison_row["person"]
    previous_student_id = person.student_id
    if previous_student_id is None:
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                person_id=person.id,
                status="invalid",
            ),
            status_code=303,
        )

    person.student_id = None
    _add_resolution_log(
        db,
        request,
        batch,
        action_code="BO_LIEN_KET",
        result_before=comparison_row["result_code"],
        person_id=person.id,
        student_id=previous_student_id,
        previous_student_id=previous_student_id,
        note=note or "Bỏ liên kết thủ công để tiếp tục xác minh.",
    )

    try:
        db.commit()
    except Exception:
        db.rollback()
        return RedirectResponse(
            url=_resolution_detail_url(
                batch.id,
                person_id=person.id,
                status="database_error",
            ),
            status_code=303,
        )

    return RedirectResponse(
        url=_resolution_detail_url(
            batch.id,
            person_id=person.id,
            status="unlinked",
        ),
        status_code=303,
    )
