from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.database import get_db
from app.permissions import (
    COMMUNE_ROLE_CODE,
    SCHOOL_ROLE_CODE,
    TEACHER_ROLE_CODE,
    normalize_role_code,
)
from app.survey_models import (
    SurveyBatch,
    SurveyForm,
    SurveyPersonYearRecord,
)
from app.routers.surveys import (
    FORM_STATUS_LABELS,
    LEARNING_STATUS_LABELS,
    b131133_chuan_hoa_khong_dau,
    b131133_xac_dinh_cap_hoc,
    lay_thong_tin_nguoi_dung,
    tao_bo_loc_bao_cao_theo_nguoi_dung,
    tao_bo_loc_dot_theo_nguoi_dung,
)


APP_DIR = Path(__file__).resolve().parent.parent

templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

router = APIRouter(
    prefix="/dieu-tra",
    tags=["Xu hướng dữ liệu điều tra nhiều năm"],
)


INDICATOR_LABELS = {
    # === BAI_13B_12_V2_4_4_6_3_TREND_INDICATOR_COMPAT ===
    "completed_preschool_by_age": "Hoàn thành Chương trình GDMN theo độ tuổi",
    "attends_two_sessions_per_day": "Học 2 buổi/ngày",
    "prepared_vietnamese": "Trẻ dân tộc được chuẩn bị tiếng Việt",
}

# === BAI_13B_12_V2_4_4_6_6_ACTIVE_TREND_INDICATORS ===
# Chi cac chi bao con theo doi moi xuat hien trong xu huong va Excel.
TRACKED_YEAR_FIELDS = (
    "completed_preschool_by_age",
    "attends_two_sessions_per_day",
    "prepared_vietnamese",
)

LEARNING_STATUS_WITH_UNKNOWN = {
    "CHUA_XAC_DINH": "Chưa xác định",
    **LEARNING_STATUS_LABELS,
}


# =========================================================
# HÀM DÙNG CHUNG
# =========================================================


def _school_year_key(code: str | None) -> tuple[int, ...]:
    """Chuyển mã năm học thành khóa số để sắp xếp ổn định."""

    if not code:
        return (0,)

    values: list[int] = []
    current = ""

    for char in str(code):
        if char.isdigit():
            current += char
        elif current:
            values.append(int(current))
            current = ""

    if current:
        values.append(int(current))

    return tuple(values) or (0,)


def _scope_label(request: Request) -> str:
    user = lay_thong_tin_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))

    if role_code in {SCHOOL_ROLE_CODE, TEACHER_ROLE_CODE}:
        return "PHẠM VI TOÀN TRƯỜNG"

    if role_code == COMMUNE_ROLE_CODE:
        return "PHẠM VI TOÀN XÃ/PHƯỜNG"

    return "PHẠM VI TOÀN TỈNH"


def _batch_in_scope(
    db: Session,
    request: Request,
    batch_id: int,
) -> SurveyBatch | None:
    filters = tao_bo_loc_dot_theo_nguoi_dung(request)

    return db.scalar(
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
        )
        .where(
            SurveyBatch.id == batch_id,
            *filters,
        )
    )


def _batch_choice_key(batch: SurveyBatch) -> tuple[int, Any, int]:
    """Ưu tiên đợt đã khóa, sau đó đến thời gian và mã nội bộ mới nhất."""

    reference_date = (
        batch.end_date
        or batch.start_date
        or batch.created_at.date()
    )
    return (1 if batch.is_locked else 0, reference_date, batch.id)


def _timeline_batches(
    db: Session,
    request: Request,
    target_batch: SurveyBatch,
) -> list[SurveyBatch]:
    """
    Lấy một đợt đại diện cho mỗi năm học cùng xã/phường.

    Đợt được chọn trên giao diện luôn là đại diện của năm đích. Với các năm
    trước, hệ thống ưu tiên đợt đã khóa và mới nhất để tránh trộn nhiều đợt
    trong cùng một năm học vào chuỗi xu hướng.
    """

    filters = tao_bo_loc_dot_theo_nguoi_dung(request)
    candidates = db.scalars(
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
        )
        .where(
            SurveyBatch.commune_id == target_batch.commune_id,
            *filters,
        )
    ).all()

    target_key = _school_year_key(target_batch.school_year.code)
    chosen: dict[str, SurveyBatch] = {}

    for batch in candidates:
        year_code = str(batch.school_year.code)
        year_key = _school_year_key(year_code)

        if year_key > target_key:
            continue

        if year_code == str(target_batch.school_year.code):
            if batch.id == target_batch.id:
                chosen[year_code] = target_batch
            continue

        current = chosen.get(year_code)
        if current is None or _batch_choice_key(batch) > _batch_choice_key(current):
            chosen[year_code] = batch

    chosen[str(target_batch.school_year.code)] = target_batch

    return sorted(
        chosen.values(),
        key=lambda item: (
            _school_year_key(item.school_year.code),
            item.id,
        ),
    )


def _load_forms(
    db: Session,
    request: Request,
    batch_id: int,
) -> list[SurveyForm]:
    form_filters = tao_bo_loc_bao_cao_theo_nguoi_dung(request)

    return db.scalars(
        select(SurveyForm)
        .options(
            selectinload(SurveyForm.household),
            selectinload(SurveyForm.person_year_records)
            .selectinload(SurveyPersonYearRecord.survey_person),
            selectinload(SurveyForm.person_year_records)
            .selectinload(SurveyPersonYearRecord.school),
            selectinload(SurveyForm.person_year_records)
            .selectinload(SurveyPersonYearRecord.classroom),
        )
        .where(
            SurveyForm.survey_batch_id == batch_id,
            *form_filters,
        )
        .order_by(SurveyForm.id.asc())
    ).all()


def _clean_text(value: Any) -> str:
    return " ".join(str(value or "").split())


def _learning_label(record: SurveyPersonYearRecord | None) -> str:
    if record is None:
        return "Chưa có hồ sơ"

    code = str(record.learning_status or "CHUA_XAC_DINH")
    return LEARNING_STATUS_WITH_UNKNOWN.get(code, code)


def _school_name(record: SurveyPersonYearRecord | None) -> str:
    if record is None:
        return "—"

    if record.school is not None:
        return record.school.name

    return _clean_text(record.school_name_reported) or "Chưa xác định"


def _class_name(record: SurveyPersonYearRecord | None) -> str:
    if record is None:
        return "—"

    if record.classroom is not None:
        return record.classroom.name

    return _clean_text(record.class_name_reported) or "Chưa xác định"


def _tracked_fields_for_record(
    record: SurveyPersonYearRecord | None,
    school_year: Any = None,
) -> tuple[str, ...]:
    """Chi tra ve cac chi bao thuc su phai theo doi cho doi tuong nay."""

    if record is None:
        return ()

    person = getattr(record, "survey_person", None)
    if person is None:
        return ()

    # Khoi chi bao nay chi ap dung cho doi tuong mam non.
    if b131133_xac_dinh_cap_hoc(person, record, school_year) != "MN":
        return ()

    fields = [
        "completed_preschool_by_age",
        "attends_two_sessions_per_day",
    ]

    ethnic_key = b131133_chuan_hoa_khong_dau(
        getattr(person, "ethnic_group", None)
    )
    if ethnic_key and ethnic_key not in {"KINH", "DAN TOC KINH"}:
        fields.append("prepared_vietnamese")

    return tuple(fields)


def _answered_indicator_total(
    record: SurveyPersonYearRecord | None,
    school_year: Any = None,
) -> int:
    if record is None:
        return 0

    return sum(
        getattr(record, field_name, None) is not None
        for field_name in _tracked_fields_for_record(record, school_year)
    )


def _record_needs_review(
    record: SurveyPersonYearRecord | None,
    school_year: Any = None,
) -> tuple[bool, list[str]]:
    reasons: list[str] = []

    if record is None:
        return True, ["Chưa có hồ sơ năm học"]

    learning_status = str(record.learning_status or "CHUA_XAC_DINH")
    if learning_status == "CHUA_XAC_DINH":
        reasons.append("Chưa xác định trạng thái học tập")

    if learning_status == "DANG_HOC" and _school_name(record) == "Chưa xác định":
        reasons.append("Đang học nhưng chưa xác định trường")

    required_fields = _tracked_fields_for_record(record, school_year)
    answered = _answered_indicator_total(record, school_year)
    if answered < len(required_fields):
        reasons.append(
            f"Mới nhập {answered}/{len(required_fields)} chỉ báo cần theo dõi"
        )

    return bool(reasons), reasons


def _percent(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return round(numerator * 100 / denominator, 1)


def _build_year_row(
    batch: SurveyBatch,
    forms: list[SurveyForm],
) -> dict[str, Any]:
    records: list[SurveyPersonYearRecord] = []

    for form in forms:
        records.extend(
            record
            for record in form.person_year_records
            if record.school_year_id == batch.school_year_id
        )

    unique_records: dict[int, SurveyPersonYearRecord] = {}
    for record in records:
        unique_records[record.survey_person_id] = record
    records = list(unique_records.values())

    learning_counts = {
        code: 0
        for code in LEARNING_STATUS_WITH_UNKNOWN
    }
    indicator_counts = {
        field_name: {"yes": 0, "no": 0, "unknown": 0}
        for field_name in TRACKED_YEAR_FIELDS
    }

    answered_total = 0
    possible_answers = 0
    review_total = 0

    for record in records:
        learning_code = str(record.learning_status or "CHUA_XAC_DINH")
        if learning_code not in learning_counts:
            learning_counts[learning_code] = 0
        learning_counts[learning_code] += 1

        active_fields = _tracked_fields_for_record(record, batch.school_year)
        answered_total += _answered_indicator_total(record, batch.school_year)
        possible_answers += len(active_fields)
        needs_review, _ = _record_needs_review(record, batch.school_year)
        review_total += int(needs_review)

        for field_name in active_fields:
            value = getattr(record, field_name, None)
            if value is True:
                indicator_counts[field_name]["yes"] += 1
            elif value is False:
                indicator_counts[field_name]["no"] += 1
            else:
                indicator_counts[field_name]["unknown"] += 1

    people_total = len(records)

    row = {
        "batch": batch,
        "year_code": str(batch.school_year.code),
        "forms_total": len(forms),
        "completed_forms": sum(
            form.status == "DA_HOAN_THANH"
            for form in forms
        ),
        "households_total": len({form.household_id for form in forms}),
        "people_total": people_total,
        "learning_known": people_total - learning_counts.get("CHUA_XAC_DINH", 0),
        "attending_total": learning_counts.get("DANG_HOC", 0),
        "preschool_completed": indicator_counts[
            "completed_preschool_by_age"
        ]["yes"],
        "answered_total": answered_total,
        "possible_answers": possible_answers,
        "indicator_completion_percent": _percent(answered_total, possible_answers),
        "review_total": review_total,
        "learning_counts": learning_counts,
        "indicator_counts": indicator_counts,
        "forms": forms,
        "records": records,
        "is_locked": bool(batch.is_locked),
        "status_label": "Đã khóa dữ liệu" if batch.is_locked else "Đang cập nhật",
    }

    return row


def _build_trend_data(
    batches: list[SurveyBatch],
    forms_by_batch: dict[int, list[SurveyForm]],
) -> dict[str, Any]:
    year_rows = [
        _build_year_row(batch, forms_by_batch.get(batch.id, []))
        for batch in batches
    ]

    previous: dict[str, Any] | None = None
    for row in year_rows:
        if previous is None:
            row["delta_households"] = None
            row["delta_people"] = None
            row["delta_attending"] = None
            row["delta_review"] = None
        else:
            row["delta_households"] = (
                row["households_total"] - previous["households_total"]
            )
            row["delta_people"] = row["people_total"] - previous["people_total"]
            row["delta_attending"] = (
                row["attending_total"] - previous["attending_total"]
            )
            row["delta_review"] = row["review_total"] - previous["review_total"]
        previous = row

    all_status_codes = list(LEARNING_STATUS_WITH_UNKNOWN)
    for row in year_rows:
        for code in row["learning_counts"]:
            if code not in all_status_codes:
                all_status_codes.append(code)

    learning_rows = [
        {
            "code": code,
            "label": LEARNING_STATUS_WITH_UNKNOWN.get(code, code),
            "values": [
                row["learning_counts"].get(code, 0)
                for row in year_rows
            ],
        }
        for code in all_status_codes
        if any(row["learning_counts"].get(code, 0) for row in year_rows)
        or code in {"CHUA_XAC_DINH", "DANG_HOC"}
    ]

    indicator_rows = []
    for field_name in TRACKED_YEAR_FIELDS:
        indicator_rows.append(
            {
                "field_name": field_name,
                "label": INDICATOR_LABELS.get(field_name, field_name),
                "years": [
                    {
                        "yes": row["indicator_counts"].get(field_name, {}).get("yes", 0),
                        "no": row["indicator_counts"].get(field_name, {}).get("no", 0),
                        "unknown": row["indicator_counts"].get(field_name, {}).get("unknown", 0),
                    }
                    for row in year_rows
                ],
            }
        )

    person_map: dict[int, dict[str, Any]] = {}

    for row in year_rows:
        year_code = row["year_code"]
        batch: SurveyBatch = row["batch"]

        for form in row["forms"]:
            for record in form.person_year_records:
                if record.school_year_id != batch.school_year_id:
                    continue

                person = record.survey_person
                if person is None:
                    continue

                item = person_map.setdefault(
                    person.id,
                    {
                        "person_id": person.id,
                        "person_code": person.code,
                        "full_name": person.full_name,
                        "date_of_birth": person.date_of_birth,
                        "gender": person.gender or "—",
                        "household_code": (
                            form.household.code if form.household else "—"
                        ),
                        "head_name": (
                            form.household.head_name if form.household else "—"
                        ),
                        "history": {},
                    },
                )

                needs_review, review_reasons = _record_needs_review(record, batch.school_year)
                item["history"][year_code] = {
                    "batch_id": batch.id,
                    "school_year_id": batch.school_year_id,
                    "household_id": form.household_id,
                    "learning": _learning_label(record),
                    "school": _school_name(record),
                    "classroom": _class_name(record),
                    "answered": _answered_indicator_total(record, batch.school_year),
                    "indicator_total": len(_tracked_fields_for_record(record, batch.school_year)),
                    "needs_review": needs_review,
                    "review_text": "; ".join(review_reasons) or "Đã nhập đủ dữ liệu chính",
                }

    person_rows: list[dict[str, Any]] = []
    year_codes = [row["year_code"] for row in year_rows]
    latest_year_code = year_codes[-1] if year_codes else ""

    for item in person_map.values():
        ordered_history = [
            (year_code, item["history"][year_code])
            for year_code in year_codes
            if year_code in item["history"]
        ]

        first_entry = ordered_history[0][1] if ordered_history else None
        latest_entry = (
            item["history"].get(latest_year_code)
            if latest_year_code
            else None
        )
        if latest_entry is None and ordered_history:
            latest_entry = ordered_history[-1][1]

        status_path = " → ".join(
            f"{year_code}: {entry['learning']}"
            for year_code, entry in ordered_history
        ) or "Chưa có lịch sử"

        item.update(
            {
                "years_tracked": len(ordered_history),
                "first_learning": (
                    first_entry["learning"] if first_entry else "—"
                ),
                "latest_learning": (
                    latest_entry["learning"] if latest_entry else "—"
                ),
                "latest_school": (
                    latest_entry["school"] if latest_entry else "—"
                ),
                "latest_classroom": (
                    latest_entry["classroom"] if latest_entry else "—"
                ),
                "latest_answered": (
                    latest_entry["answered"] if latest_entry else 0
                ),
                "latest_indicator_total": (
                    latest_entry["indicator_total"] if latest_entry else 0
                ),
                "latest_needs_review": bool(
                    latest_entry and latest_entry["needs_review"]
                ),
                "latest_review_text": (
                    latest_entry["review_text"]
                    if latest_entry
                    else "Chưa có hồ sơ năm học mới nhất"
                ),
                "latest_household_id": (
                    latest_entry["household_id"] if latest_entry else None
                ),
                "status_path": status_path,
            }
        )
        person_rows.append(item)

    person_rows.sort(
        key=lambda item: (
            str(item["head_name"]).casefold(),
            str(item["full_name"]).casefold(),
        )
    )

    latest_review_rows = [
        item for item in person_rows if item["latest_needs_review"]
    ]

    first_row = year_rows[0] if year_rows else None
    latest_row = year_rows[-1] if year_rows else None

    summary = {
        "year_total": len(year_rows),
        "first_year": first_row["year_code"] if first_row else "—",
        "latest_year": latest_row["year_code"] if latest_row else "—",
        "latest_households": latest_row["households_total"] if latest_row else 0,
        "latest_people": latest_row["people_total"] if latest_row else 0,
        "latest_attending": latest_row["attending_total"] if latest_row else 0,
        "latest_review": latest_row["review_total"] if latest_row else 0,
        "people_change": (
            latest_row["people_total"] - first_row["people_total"]
            if first_row and latest_row and len(year_rows) > 1
            else 0
        ),
        "household_change": (
            latest_row["households_total"] - first_row["households_total"]
            if first_row and latest_row and len(year_rows) > 1
            else 0
        ),
    }

    return {
        "summary": summary,
        "year_rows": year_rows,
        "year_codes": year_codes,
        "learning_rows": learning_rows,
        "indicator_rows": indicator_rows,
        "person_rows": person_rows,
        "latest_review_rows": latest_review_rows,
    }


# =========================================================
# EXCEL
# =========================================================


def _style_sheet(ws, widths: list[int]) -> None:
    thin = Side(style="thin", color="B8C7D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1976D2")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def _build_excel(
    request: Request,
    target_batch: SurveyBatch,
    trend_data: dict[str, Any],
) -> BytesIO:
    summary = trend_data["summary"]
    year_rows = trend_data["year_rows"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Tong quan lien nam"

    _write_header(ws, ["Thông tin", "Giá trị"])
    overview_rows = [
        ("Phạm vi tài khoản", _scope_label(request)),
        ("Xã/phường", target_batch.commune.name),
        ("Đợt đích đang xem", target_batch.code),
        ("Số năm học trong chuỗi", summary["year_total"]),
        ("Năm đầu chuỗi", summary["first_year"]),
        ("Năm cuối chuỗi", summary["latest_year"]),
        ("Hộ năm mới nhất", summary["latest_households"]),
        ("Đối tượng năm mới nhất", summary["latest_people"]),
        ("Đang học năm mới nhất", summary["latest_attending"]),
        ("Cần rà soát năm mới nhất", summary["latest_review"]),
        ("Biến động hộ đầu-cuối", summary["household_change"]),
        ("Biến động đối tượng đầu-cuối", summary["people_change"]),
    ]
    for row in overview_rows:
        ws.append(list(row))
    _style_sheet(ws, [34, 52])

    ws_years = wb.create_sheet("Theo nam")
    _write_header(
        ws_years,
        [
            "Năm học",
            "Mã đợt",
            "Tên đợt",
            "Trạng thái dữ liệu",
            "Số hộ",
            "Số phiếu",
            "Phiếu hoàn thành",
            "Đối tượng",
            "Đã xác định học tập",
            "Đang học",
            "Hoàn thành CTGDMN theo độ tuổi",
            "Tỷ lệ nhập chỉ báo đang theo dõi (%)",
            "Cần rà soát",
            "Biến động đối tượng",
        ],
    )
    for row in year_rows:
        ws_years.append(
            [
                row["year_code"],
                row["batch"].code,
                row["batch"].name,
                row["status_label"],
                row["households_total"],
                row["forms_total"],
                row["completed_forms"],
                row["people_total"],
                row["learning_known"],
                row["attending_total"],
                row["preschool_completed"],
                row["indicator_completion_percent"],
                row["review_total"],
                row["delta_people"] if row["delta_people"] is not None else "",
            ]
        )
    _style_sheet(
        ws_years,
        [14, 24, 36, 20, 10, 10, 16, 12, 19, 12, 30, 30, 14, 20],
    )

    ws_learning = wb.create_sheet("Trang thai hoc tap")
    _write_header(
        ws_learning,
        ["Năm học", "Mã trạng thái", "Trạng thái học tập", "Số đối tượng"],
    )
    for row in year_rows:
        for code, count in row["learning_counts"].items():
            if count <= 0 and code not in {"CHUA_XAC_DINH", "DANG_HOC"}:
                continue
            ws_learning.append(
                [
                    row["year_code"],
                    code,
                    LEARNING_STATUS_WITH_UNKNOWN.get(code, code),
                    count,
                ]
            )
    _style_sheet(ws_learning, [14, 24, 34, 16])

    ws_indicators = wb.create_sheet("Chi bao theo nam")
    _write_header(
        ws_indicators,
        ["Năm học", "Chỉ báo", "Có", "Không", "Chưa xác định", "Tổng"],
    )
    for row in year_rows:
        for field_name in TRACKED_YEAR_FIELDS:
            counts = row["indicator_counts"].get(
                field_name, {"yes": 0, "no": 0, "unknown": 0}
            )
            ws_indicators.append(
                [
                    row["year_code"],
                    INDICATOR_LABELS.get(field_name, field_name),
                    counts["yes"],
                    counts["no"],
                    counts["unknown"],
                    sum(counts.values()),
                ]
            )
    _style_sheet(ws_indicators, [14, 48, 10, 10, 18, 10])

    ws_people = wb.create_sheet("Lich su doi tuong")
    _write_header(
        ws_people,
        [
            "STT",
            "Mã đối tượng",
            "Họ và tên",
            "Ngày sinh",
            "Mã hộ",
            "Chủ hộ",
            "Năm học",
            "Trạng thái học tập",
            "Trường",
            "Lớp",
            "Chỉ báo đã nhập",
            "Kết quả rà soát",
        ],
    )
    index = 0
    for person in trend_data["person_rows"]:
        for year_code in trend_data["year_codes"]:
            entry = person["history"].get(year_code)
            if entry is None:
                continue
            index += 1
            ws_people.append(
                [
                    index,
                    person["person_code"],
                    person["full_name"],
                    person["date_of_birth"].strftime("%d/%m/%Y")
                    if person["date_of_birth"]
                    else "",
                    person["household_code"],
                    person["head_name"],
                    year_code,
                    entry["learning"],
                    entry["school"],
                    entry["classroom"],
                    f"{entry['answered']}/{entry['indicator_total']}",
                    entry["review_text"],
                ]
            )
    _style_sheet(
        ws_people,
        [7, 20, 25, 14, 18, 25, 14, 25, 32, 22, 18, 52],
    )

    ws_review = wb.create_sheet("Can ra soat moi nhat")
    _write_header(
        ws_review,
        [
            "STT",
            "Mã đối tượng",
            "Họ và tên",
            "Mã hộ",
            "Chủ hộ",
            "Trạng thái mới nhất",
            "Trường mới nhất",
            "Lớp mới nhất",
            "Chỉ báo đã nhập",
            "Nội dung cần rà soát",
        ],
    )
    for index, person in enumerate(trend_data["latest_review_rows"], start=1):
        ws_review.append(
            [
                index,
                person["person_code"],
                person["full_name"],
                person["household_code"],
                person["head_name"],
                person["latest_learning"],
                person["latest_school"],
                person["latest_classroom"],
                f"{person['latest_answered']}/{person['latest_indicator_total']}",
                person["latest_review_text"],
            ]
        )
    _style_sheet(ws_review, [7, 20, 25, 18, 25, 25, 32, 22, 18, 52])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# =========================================================
# ROUTES
# =========================================================


def _prepare_trend(
    db: Session,
    request: Request,
    target_batch: SurveyBatch,
) -> tuple[list[SurveyBatch], dict[str, Any]]:
    batches = _timeline_batches(db, request, target_batch)
    forms_by_batch = {
        batch.id: _load_forms(db, request, batch.id)
        for batch in batches
    }
    return batches, _build_trend_data(batches, forms_by_batch)


@router.get(
    "/{target_batch_id}/xu-huong-lien-nam",
    response_class=HTMLResponse,
)
def multi_year_trend(
    target_batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    target_batch = _batch_in_scope(db, request, target_batch_id)

    if target_batch is None:
        return RedirectResponse(
            url="/dieu-tra?status=forbidden",
            status_code=303,
        )

    batches, trend_data = _prepare_trend(
        db,
        request,
        target_batch,
    )

    return templates.TemplateResponse(
        request=request,
        name="surveys/multi_year_trend.html",
        context={
            "nguoi_dung": lay_thong_tin_nguoi_dung(request),
            "scope_label": _scope_label(request),
            "target_batch": target_batch,
            "timeline_batches": batches,
            "trend": trend_data,
            "form_status_labels": FORM_STATUS_LABELS,
        },
    )


@router.get(
    "/{target_batch_id}/xu-huong-lien-nam/xuat-excel",
)
def export_multi_year_trend(
    target_batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    target_batch = _batch_in_scope(db, request, target_batch_id)

    if target_batch is None:
        return RedirectResponse(
            url="/dieu-tra?status=forbidden",
            status_code=303,
        )

    _, trend_data = _prepare_trend(
        db,
        request,
        target_batch,
    )
    buffer = _build_excel(request, target_batch, trend_data)

    filename = (
        "xu_huong_lien_nam_"
        f"{target_batch.commune.code}_"
        f"den_{target_batch.school_year.code}.xlsx"
    )

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )

# === BAI_13B_11_17_V1_DASHBOARD_DAT_CHUAN_START ===
from fastapi import Request as _PCGDDashboardRequest
from fastapi.responses import JSONResponse as _PCGDDashboardJSONResponse

@router.get("/{batch_id}/dashboard-dat-chuan-data")
def _dashboard_dat_chuan_data_v1(
    batch_id: int,
    request: _PCGDDashboardRequest,
):
    """
    Dashboard 2.4.3:
    - Lộ trình PCGDMN trẻ 3-5 tuổi: lấy từ JSON đã sinh từ file lộ trình.
    - Trẻ 5 tuổi: so sánh 2 chỉ tiêu trẻ em theo NĐ 20/2014/NĐ-CP.
    - Không tự kết luận xã đạt chuẩn toàn diện khi chưa đủ điều kiện GV/CSVC.
    - Không ghi dữ liệu, chỉ đọc.
    """
    import json
    import sqlite3
    import unicodedata
    from pathlib import Path

    project_root = Path(__file__).resolve().parents[2]
    db_path = project_root / "data" / "phocap.db"
    roadmap_path = (
        project_root
        / "app"
        / "data"
        / "pcgd_recognition_roadmap_2026_2030.json"
    )

    def _norm(value):
        text = str(value or "").strip().lower()
        text = unicodedata.normalize("NFD", text)
        text = "".join(
            ch for ch in text
            if unicodedata.category(ch) != "Mn"
        )
        return " ".join(text.split())

    def _year(value):
        try:
            y = int(str(value).strip())
        except Exception:
            return None
        return y if 2026 <= y <= 2030 else None

    if not db_path.is_file():
        return _PCGDDashboardJSONResponse(
            {
                "ok": False,
                "message": "Không tìm thấy database.",
            },
            status_code=500,
        )

    con = sqlite3.connect(str(db_path))
    con.row_factory = sqlite3.Row

    try:
        batch = con.execute(
            """
            SELECT
                sb.id,
                sb.school_year_id,
                sb.commune_id,
                sy.code AS school_year_code
            FROM survey_batches sb
            JOIN school_years sy
              ON sy.id = sb.school_year_id
            WHERE sb.id = ?
            """,
            (batch_id,),
        ).fetchone()

        if batch is None:
            return _PCGDDashboardJSONResponse(
                {
                    "ok": False,
                    "message": "Không tìm thấy đợt điều tra.",
                },
                status_code=404,
            )

        school_year_id = int(batch["school_year_id"])
        school_year_code = str(
            batch["school_year_code"] or ""
        )

        try:
            start_year = int(
                school_year_code.split("-")[0]
            )
        except Exception:
            start_year = 2026

        user_id = request.session.get("user_id")
        viewer_commune_id = None
        viewer_school_id = None

        if user_id:
            viewer = con.execute(
                """
                SELECT commune_id, school_id
                FROM users
                WHERE id = ?
                """,
                (user_id,),
            ).fetchone()
            if viewer is not None:
                viewer_commune_id = viewer["commune_id"]
                viewer_school_id = viewer["school_id"]

        commune_rows = con.execute(
            """
            SELECT
                id,
                code,
                name,
                is_special_difficulty_area
            FROM communes
            WHERE is_active = 1
            ORDER BY name
            """
        ).fetchall()

        communes = []
        by_norm_name = {}
        by_code = {}

        for row in commune_rows:
            item = {
                "id": int(row["id"]),
                "code": str(row["code"] or ""),
                "name": str(row["name"] or ""),
                "is_special": bool(
                    row["is_special_difficulty_area"]
                ),
            }
            communes.append(item)
            by_norm_name[_norm(item["name"])] = item
            by_code[item["code"]] = item

        # ------------------------------------------------------------
        # Đọc lộ trình 3-5 tuổi theo cách linh hoạt.
        # Hỗ trợ cả cấu trúc danh sách bản ghi và cấu trúc nhóm theo năm.
        # ------------------------------------------------------------
        roadmap = {}
        roadmap_warning = None

        if roadmap_path.is_file():
            try:
                raw_roadmap = json.loads(
                    roadmap_path.read_text(
                        encoding="utf-8-sig"
                    )
                )

                def _match_commune(value):
                    s = str(value or "").strip()
                    if not s:
                        return None
                    if s in by_code:
                        return by_code[s]
                    return by_norm_name.get(_norm(s))

                def _walk(obj, inherited_year=None):
                    if isinstance(obj, dict):
                        local_year = inherited_year

                        for k, v in obj.items():
                            kn = _norm(k)
                            if (
                                "year" in kn
                                or kn in {
                                    "nam",
                                    "nam dat",
                                    "nam dat chuan",
                                    "nam dang ky",
                                    "nam lo trinh",
                                }
                            ):
                                y = _year(v)
                                if y:
                                    local_year = y

                        found = []
                        for k, v in obj.items():
                            kn = _norm(k)
                            if any(
                                token in kn
                                for token in (
                                    "commune",
                                    "xa",
                                    "phuong",
                                    "don vi",
                                    "unit",
                                    "name",
                                    "code",
                                    "ma",
                                )
                            ):
                                c = _match_commune(v)
                                if c is not None:
                                    found.append(c)

                        if local_year:
                            for c in found:
                                roadmap[c["id"]] = int(
                                    local_year
                                )

                        for k, v in obj.items():
                            key_year = _year(k)
                            _walk(
                                v,
                                key_year or local_year,
                            )

                    elif isinstance(obj, list):
                        for item in obj:
                            _walk(item, inherited_year)

                    elif isinstance(obj, str):
                        if inherited_year:
                            c = _match_commune(obj)
                            if c is not None:
                                roadmap[c["id"]] = int(
                                    inherited_year
                                )

                _walk(raw_roadmap)

                if len(roadmap) < len(communes):
                    roadmap_warning = (
                        "Đã đọc được lộ trình cho "
                        f"{len(roadmap)}/{len(communes)} xã/phường."
                    )
            except Exception as exc:
                roadmap_warning = (
                    "Không đọc được file lộ trình: "
                    + str(exc)
                )
        else:
            roadmap_warning = (
                "Chưa tìm thấy app/data/"
                "pcgd_recognition_roadmap_2026_2030.json"
            )

        # ------------------------------------------------------------
        # Tiến độ phiếu theo xã.
        # ------------------------------------------------------------
        progress_rows = con.execute(
            """
            SELECT
                sb.commune_id,
                COUNT(sf.id) AS form_total,
                SUM(
                    CASE
                        WHEN sf.status = 'DA_HOAN_THANH'
                        THEN 1 ELSE 0
                    END
                ) AS completed_total
            FROM survey_batches sb
            LEFT JOIN survey_forms sf
              ON sf.survey_batch_id = sb.id
            WHERE sb.school_year_id = ?
            GROUP BY sb.commune_id
            """,
            (school_year_id,),
        ).fetchall()

        progress = {
            int(r["commune_id"]): {
                "forms": int(r["form_total"] or 0),
                "completed": int(
                    r["completed_total"] or 0
                ),
            }
            for r in progress_rows
        }

        # ------------------------------------------------------------
        # Dữ liệu đối tượng theo năm học.
        # Không suy diễn "đạt chuẩn toàn diện".
        # Chỉ tính các chỉ báo trẻ em có nguồn trực tiếp trong CSDL.
        # ------------------------------------------------------------
        people_rows = con.execute(
            """
            SELECT DISTINCT
                sb.commune_id,
                p.id AS person_id,
                p.date_of_birth,
                p.residency_status,
                pyr.learning_status,
                pyr.completed_preschool_5,
                pyr.completed_preschool_by_age
            FROM survey_person_year_records pyr
            JOIN survey_people p
              ON p.id = pyr.survey_person_id
            JOIN survey_forms sf
              ON sf.id = pyr.survey_form_id
            JOIN survey_batches sb
              ON sb.id = sf.survey_batch_id
            WHERE
                pyr.school_year_id = ?
                AND sb.school_year_id = ?
                AND p.is_active = 1
            """,
            (
                school_year_id,
                school_year_id,
            ),
        ).fetchall()

        metrics = {}

        def _m(cid):
            if cid not in metrics:
                metrics[cid] = {
                    "age_3_5_total": 0,
                    "age_3_5_attending": 0,
                    "age_5_total": 0,
                    "age_5_attending": 0,
                    "age_5_completed": 0,
                }
            return metrics[cid]

        valid_residency = {
            "THUONG_TRU",
            "TAM_TRU",
            "DANG_CU_TRU",
            "CU_TRU",
        }

        birth_5 = start_year - 5
        birth_4 = start_year - 4
        birth_3 = start_year - 3

        for row in people_rows:
            dob = str(row["date_of_birth"] or "")
            if len(dob) < 4:
                continue

            try:
                birth_year = int(dob[:4])
            except Exception:
                continue

            residence = str(
                row["residency_status"] or ""
            ).strip().upper()

            if (
                residence
                and residence not in valid_residency
            ):
                continue

            cid = int(row["commune_id"])
            mm = _m(cid)

            attending = (
                str(
                    row["learning_status"] or ""
                ).strip().upper()
                == "DANG_HOC"
            )

            completed = bool(
                row["completed_preschool_by_age"]
            ) or bool(
                row["completed_preschool_5"]
            )

            if birth_year in {
                birth_3,
                birth_4,
                birth_5,
            }:
                mm["age_3_5_total"] += 1
                if attending:
                    mm["age_3_5_attending"] += 1

            if birth_year == birth_5:
                mm["age_5_total"] += 1
                if attending:
                    mm["age_5_attending"] += 1
                if completed:
                    mm["age_5_completed"] += 1

        def _pct(n, d):
            if not d:
                return None
            return round(n * 100.0 / d, 2)

        # ------------------------------------------------------------
        # Phạm vi xem:
        # - tài khoản Sở/Admin: commune_id trống -> toàn tỉnh.
        # - tài khoản Xã/Trường/GV: chỉ xã của tài khoản.
        # ------------------------------------------------------------
        scoped_communes = communes
        if viewer_commune_id:
            scoped_communes = [
                c
                for c in communes
                if c["id"] == int(viewer_commune_id)
            ]

        rows = []
        target_counts = {}
        summary = {
            "total_communes": 0,
            "roadmap_current_year": 0,
            "survey_complete": 0,
            "five_child_indicator_pass": 0,
            "special_difficulty": 0,
        }

        for c in scoped_communes:
            cid = c["id"]
            target_year = roadmap.get(cid)
            pp = progress.get(
                cid,
                {
                    "forms": 0,
                    "completed": 0,
                },
            )
            mm = metrics.get(
                cid,
                {
                    "age_3_5_total": 0,
                    "age_3_5_attending": 0,
                    "age_5_total": 0,
                    "age_5_attending": 0,
                    "age_5_completed": 0,
                },
            )

            forms = pp["forms"]
            completed_forms = pp["completed"]
            survey_complete = (
                forms > 0
                and completed_forms >= forms
            )

            rate_3_5 = _pct(
                mm["age_3_5_attending"],
                mm["age_3_5_total"],
            )
            rate_5_attend = _pct(
                mm["age_5_attending"],
                mm["age_5_total"],
            )
            rate_5_complete = _pct(
                mm["age_5_completed"],
                mm["age_5_total"],
            )

            # NĐ20 - chỉ 2 chỉ tiêu trẻ em đang có nguồn trực tiếp:
            # - đến lớp: 95%, ĐBKK 90%
            # - hoàn thành CTGDMN: 85%, ĐBKK 80%
            # Không dùng 2 chỉ tiêu này để thay thế kết luận đạt chuẩn toàn diện.
            attend_threshold = (
                90.0 if c["is_special"] else 95.0
            )
            complete_threshold = (
                80.0 if c["is_special"] else 85.0
            )

            five_pass = (
                survey_complete
                and rate_5_attend is not None
                and rate_5_complete is not None
                and rate_5_attend
                    >= attend_threshold
                and rate_5_complete
                    >= complete_threshold
            )

            if not survey_complete:
                five_status = "Chưa đủ dữ liệu"
                five_status_code = "DATA"
            elif (
                rate_5_attend is None
                or rate_5_complete is None
            ):
                five_status = "Chưa có đối tượng 5 tuổi"
                five_status_code = "EMPTY"
            elif five_pass:
                five_status = (
                    "Đạt 2 chỉ tiêu trẻ em NĐ20"
                )
                five_status_code = "PASS"
            else:
                five_status = (
                    "Chưa đạt 2 chỉ tiêu trẻ em NĐ20"
                )
                five_status_code = "FAIL"

            if target_year is None:
                roadmap_status = "Chưa có lộ trình"
                roadmap_code = "NONE"
            elif target_year < start_year:
                roadmap_status = (
                    "Đã quá năm đăng ký - cần rà soát"
                )
                roadmap_code = "LATE"
            elif target_year == start_year:
                roadmap_status = (
                    "Năm đăng ký đạt chuẩn"
                )
                roadmap_code = "DUE"
            else:
                roadmap_status = (
                    f"Còn {target_year - start_year} năm"
                )
                roadmap_code = "ON_TRACK"

            if target_year:
                target_counts[str(target_year)] = (
                    target_counts.get(
                        str(target_year),
                        0,
                    )
                    + 1
                )

            summary["total_communes"] += 1

            if target_year == start_year:
                summary[
                    "roadmap_current_year"
                ] += 1

            if survey_complete:
                summary["survey_complete"] += 1

            if five_pass:
                summary[
                    "five_child_indicator_pass"
                ] += 1

            if c["is_special"]:
                summary["special_difficulty"] += 1

            rows.append(
                {
                    "commune_id": cid,
                    "code": c["code"],
                    "name": c["name"],
                    "is_special": c["is_special"],
                    "target_year_3_5": target_year,
                    "roadmap_status": roadmap_status,
                    "roadmap_code": roadmap_code,
                    "forms": forms,
                    "completed_forms": completed_forms,
                    "survey_complete": survey_complete,
                    "age_3_5_total":
                        mm["age_3_5_total"],
                    "age_3_5_attending":
                        mm["age_3_5_attending"],
                    "rate_3_5_attending": rate_3_5,
                    "age_5_total":
                        mm["age_5_total"],
                    "age_5_attending":
                        mm["age_5_attending"],
                    "rate_5_attending":
                        rate_5_attend,
                    "age_5_completed":
                        mm["age_5_completed"],
                    "rate_5_completed":
                        rate_5_complete,
                    "attend_threshold":
                        attend_threshold,
                    "complete_threshold":
                        complete_threshold,
                    "five_status": five_status,
                    "five_status_code":
                        five_status_code,
                }
            )

        rows.sort(
            key=lambda x: (
                x["target_year_3_5"]
                if x["target_year_3_5"]
                is not None
                else 9999,
                x["name"],
            )
        )

        return _PCGDDashboardJSONResponse(
            {
                "ok": True,
                "school_year_id": school_year_id,
                "school_year_code": school_year_code,
                "analysis_year": start_year,
                "scope": (
                    "COMMUNE"
                    if viewer_commune_id
                    else "PROVINCE"
                ),
                "viewer_commune_id":
                    viewer_commune_id,
                "viewer_school_id":
                    viewer_school_id,
                "roadmap_loaded":
                    len(roadmap),
                "roadmap_warning":
                    roadmap_warning,
                "summary": summary,
                "target_counts": target_counts,
                "rows": rows,
                "notes": [
                    (
                        "Lộ trình 3-5 tuổi là kế hoạch đăng ký; "
                        "không tự động đồng nghĩa xã đã đạt chuẩn."
                    ),
                    (
                        "Phần 5 tuổi chỉ đánh giá 2 chỉ tiêu trẻ em "
                        "có nguồn trực tiếp trong CSDL theo NĐ20; "
                        "không thay thế kết luận công nhận toàn diện."
                    ),
                    (
                        "Cờ xã đặc biệt khó khăn lấy từ "
                        "communes.is_special_difficulty_area, "
                        "không lấy cờ màu trong file lộ trình."
                    ),
                ],
            }
        )
    finally:
        con.close()
# === BAI_13B_11_17_V1_DASHBOARD_DAT_CHUAN_END ===
