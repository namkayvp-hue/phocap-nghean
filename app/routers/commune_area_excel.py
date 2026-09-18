from datetime import datetime
from io import BytesIO
import json
from uuid import uuid4
from zipfile import ZipFile

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.database import get_db
from app.routers.survey_team_registration import (
    _B1322_AREA_TYPES, _b1322_ensure_area_schema, _forbidden, _role, _user, templates,
)

router = APIRouter(prefix="/dieu-tra/phan-cong-to-dieu-tra/dia-ban", tags=["Excel địa bàn"])
HEADERS = ["Loại địa bàn", "Tên địa bàn", "Số hộ dự kiến", "Ghi chú"]
MAX_SIZE = 5 * 1024 * 1024


def clean(value):
    return " ".join(str(value if value is not None else "").split())


def validate_rows(rows):
    if not isinstance(rows, list) or not 1 <= len(rows) <= 500:
        raise ValueError("File phải có từ 1 đến 500 dòng địa bàn.")
    types = {label.casefold(): code for code, label in _B1322_AREA_TYPES.items()}
    types.update({code.casefold(): code for code in _B1322_AREA_TYPES})
    result, seen = [], set()
    for number, row in enumerate(rows, 2):
        if not isinstance(row, dict):
            raise ValueError("Dữ liệu xem trước không hợp lệ.")
        kind = types.get(clean(row.get("area_type")).casefold())
        name = clean(row.get("name"))
        notes = clean(row.get("notes"))
        expected = clean(row.get("expected_households")) or "0"
        if not kind or not name or len(name) > 200 or len(notes) > 500:
            raise ValueError(f"Dòng {number}: kiểm tra loại, tên địa bàn (tối đa 200 ký tự) và ghi chú (tối đa 500 ký tự).")
        if not expected.isascii() or not expected.isdecimal() or len(expected) > 9:
            raise ValueError(f"Dòng {number}: số hộ dự kiến phải là số nguyên không âm, tối đa 999.999.999.")
        key = (kind, name.casefold())
        if key in seen:
            raise ValueError(f"Dòng {number}: địa bàn trùng loại và tên trong file. Hãy giữ một dòng duy nhất.")
        seen.add(key)
        result.append(dict(area_type=kind, name=name, expected_households=int(expected), notes=notes))
    return result


def read_excel(content):
    if len(content) > MAX_SIZE:
        raise ValueError("File không được vượt quá 5 MB.")
    book = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError("File Excel có dữ liệu giải nén quá lớn.")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
        sheet = book.active
        if sheet.max_row and sheet.max_row > 501:
            raise ValueError("File chỉ được chứa tối đa 500 dòng địa bàn.")
        cells = sheet.iter_rows(max_row=502, max_col=20)
        header = [clean(cell.value).casefold() for cell in next(cells, ())]
        required = [name.casefold() for name in HEADERS[:3]]
        if any(header.count(name) != 1 for name in required):
            raise ValueError("Thiếu hoặc trùng cột Loại địa bàn, Tên địa bàn, Số hộ dự kiến. Hãy dùng file mẫu.")
        indices = [header.index(name.casefold()) if name.casefold() in header else None for name in HEADERS]
        rows = []
        for number, line in enumerate(cells, 2):
            selected = [line[i] if i is not None else None for i in indices]
            if all(cell is None or not clean(cell.value) for cell in selected):
                continue
            if number > 501:
                raise ValueError("File chỉ được chứa tối đa 500 dòng địa bàn.")
            if any(cell is not None and cell.data_type in {"f", "e"} for cell in selected):
                raise ValueError(f"Dòng {number}: không dùng công thức hoặc ô lỗi.")
            values = [cell.value if cell is not None else None for cell in selected]
            rows.append(dict(zip(["area_type", "name", "expected_households", "notes"], values)))
        return validate_rows(rows)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Không đọc được file Excel. Hãy dùng file .xlsx theo mẫu.") from exc
    finally:
        if book is not None:
            book.close()


def area_scope(request, db, year_id):
    user = _user(request)
    if _role(request) != "XA" or not user.get("commune_id"):
        return None
    year = db.execute(text("SELECT id,code FROM school_years WHERE id=:id AND is_active=1"), {"id": year_id}).mappings().first()
    commune = db.execute(text("SELECT id,name FROM communes WHERE id=:id AND is_active=1"), {"id": user["commune_id"]}).mappings().first()
    return (user, year, commune) if year and commune else None


def plan_rows(db, rows, commune_id, year_id):
    existing = {(row.area_type, clean(row.name).casefold()) for row in db.execute(
        text("SELECT area_type,name FROM survey_commune_areas WHERE commune_id=:commune AND school_year_id=:year"),
        {"commune": commune_id, "year": year_id},
    )}
    return [{**row, "exists": (row["area_type"], row["name"].casefold()) in existing} for row in rows]


@router.get("/mau-excel")
def download_template(request: Request):
    if _role(request) != "XA" or not _user(request).get("commune_id"):
        return _forbidden()
    book = Workbook()
    sheet = book.active
    sheet.title = "Danh mục địa bàn"
    sheet.append(HEADERS)
    sheet.append(["Thôn", "Thôn 1", 100, ""])
    sheet.append(["Xóm", "Xóm 2", 120, ""])
    for col, width in [("A", 20), ("B", 35), ("C", 22), ("D", 45)]:
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A2"
    output = BytesIO()
    book.save(output)
    book.close()
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="mau_dia_ban.xlsx"'})


@router.post("/nhap-excel")
def preview(request: Request, school_year_id: int = Form(...), file: UploadFile = File(...), db: Session = Depends(get_db)):
    scope = area_scope(request, db, school_year_id)
    if scope is None:
        file.file.close()
        return _forbidden()
    user, year, commune = scope
    rows, error = [], ""
    try:
        if not str(file.filename or "").lower().endswith(".xlsx"):
            raise ValueError("Vui lòng chọn file .xlsx.")
        rows = read_excel(file.file.read(MAX_SIZE + 1))
        _b1322_ensure_area_schema(db)
        rows = plan_rows(db, rows, commune["id"], year["id"])
    except ValueError as exc:
        error = str(exc)
    finally:
        file.file.close()
    return templates.TemplateResponse(request=request, name="survey_teams/area_import.html",
        context=dict(nguoi_dung=user, year=year, commune=commune, rows=rows, payload=json.dumps(rows, ensure_ascii=False),
                     error=error, saved=False, area_types=_B1322_AREA_TYPES), status_code=400 if error else 200)


@router.post("/nhap-excel/luu")
def save(request: Request, school_year_id: int = Form(...), payload: str = Form(..., max_length=1000000), db: Session = Depends(get_db)):
    scope = area_scope(request, db, school_year_id)
    if scope is None:
        return _forbidden()
    user, year, commune = scope
    rows, error, created, skipped = [], "", 0, 0
    try:
        rows = validate_rows(json.loads(payload))
        _b1322_ensure_area_schema(db)
        rows = plan_rows(db, rows, commune["id"], year["id"])
        for row in rows:
            if row["exists"]:
                skipped += 1
                continue
            db.execute(text("""INSERT INTO survey_commune_areas
                (commune_id,school_year_id,code,name,area_type,expected_households,notes,is_active,created_by_user_id,created_at,updated_at)
                VALUES (:commune,:year,:code,:name,:area_type,:expected_households,:notes,1,:actor,:now,:now)"""),
                {**row, "commune": commune["id"], "year": year["id"], "code": uuid4().hex,
                 "actor": user.get("id"), "now": datetime.now()})
            created += 1
        db.commit()
    except (ValueError, IntegrityError) as exc:
        db.rollback()
        error = str(exc) if isinstance(exc, ValueError) else "Danh mục vừa thay đổi. Hãy tải file và xem trước lại; chưa lưu dòng nào."
    return templates.TemplateResponse(request=request, name="survey_teams/area_import.html",
        context=dict(nguoi_dung=user, year=year, commune=commune, rows=[], payload="", error=error,
                     saved=not error, created=created, skipped=skipped, area_types=_B1322_AREA_TYPES), status_code=400 if error else 200)
