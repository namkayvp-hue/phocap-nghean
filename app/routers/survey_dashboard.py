from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.database import get_db
from app.models import Commune, SchoolYear
from app.permissions import (
    ADMIN_ROLE_CODES,
    COMMUNE_ROLE_CODE,
    DEPARTMENT_ROLE_CODE,
    normalize_role_code,
)
from app.routers.surveys import (
    BATCH_STATUS_LABELS,
    BOOLEAN_YEAR_FIELDS,
    FORM_STATUS_LABELS,
    lay_thong_tin_nguoi_dung,
    tao_bo_loc_bao_cao_theo_nguoi_dung,
    tao_bo_loc_dot_theo_nguoi_dung,
)
from app.survey_models import (
    SurveyBatch,
    SurveyForm,
    SurveyPersonYearRecord,
)


APP_DIR = Path(__file__).resolve().parent.parent

templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

router = APIRouter(
    prefix="/dieu-tra",
    tags=["Bảng điều hành điều tra theo địa bàn"],
)


DASHBOARD_ROLE_CODES = frozenset({
    *ADMIN_ROLE_CODES,
    DEPARTMENT_ROLE_CODE,
    COMMUNE_ROLE_CODE,
})

VIEW_LABELS = {
    "": "Tất cả xã/phường",
    "attention": "Cần xử lý",
    "missing": "Chưa tạo đợt",
    "open": "Đang mở cập nhật",
    "locked": "Đã khóa dữ liệu",
    "ready": "Sẵn sàng/chất lượng tốt",
}

STATE_LABELS = {
    "missing": "Chưa tạo đợt",
    "empty": "Chưa có phiếu",
    "attention": "Cần xử lý",
    "ready": "Sẵn sàng",
    "locked": "Đã khóa dữ liệu",
}


# =========================================================
# HÀM DÙNG CHUNG
# =========================================================


def _scope_label(request: Request) -> str:
    user = lay_thong_tin_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))

    if role_code == COMMUNE_ROLE_CODE:
        return "PHẠM VI TOÀN XÃ/PHƯỜNG"

    return "PHẠM VI TOÀN TỈNH"


def _parse_positive_int(value: str | None) -> int | None:
    normalized = str(value or "").strip()
    if not normalized:
        return None

    try:
        number = int(normalized)
    except (TypeError, ValueError):
        return None

    return number if number > 0 else None


def _clean_text(value: Any) -> str:
    return " ".join(str(value or "").split())


def _batch_in_scope(
    db: Session,
    request: Request,
    batch_id: int,
) -> SurveyBatch | None:
    filters = tao_bo_loc_dot_theo_nguoi_dung(request)

    return db.scalar(
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
        )
        .where(
            SurveyBatch.id == batch_id,
            *filters,
        )
    )


def _school_year_key(code: str | None) -> tuple[int, ...]:
    values: list[int] = []
    current = ""

    for char in str(code or ""):
        if char.isdigit():
            current += char
        elif current:
            values.append(int(current))
            current = ""

    if current:
        values.append(int(current))

    return tuple(values) or (0,)


def _batch_choice_key(batch: SurveyBatch) -> tuple[int, Any, int]:
    reference_date = (
        batch.end_date
        or batch.start_date
        or batch.created_at.date()
    )
    return (1 if batch.is_locked else 0, reference_date, batch.id)


def _communes_in_scope(
    db: Session,
    request: Request,
) -> list[Commune]:
    user = lay_thong_tin_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))

    statement = select(Commune).where(Commune.is_active.is_(True))

    if role_code == COMMUNE_ROLE_CODE:
        commune_id = user.get("commune_id")
        if commune_id is None:
            return []
        statement = statement.where(Commune.id == int(commune_id))

    return db.scalars(
        statement.order_by(Commune.name.asc(), Commune.code.asc())
    ).all()


def _school_years(db: Session) -> list[SchoolYear]:
    years = db.scalars(select(SchoolYear)).all()
    return sorted(
        years,
        key=lambda item: (_school_year_key(item.code), item.id),
        reverse=True,
    )


def _selected_school_year(
    db: Session,
    target_batch: SurveyBatch,
    requested_id: str | None,
) -> SchoolYear:
    parsed_id = _parse_positive_int(requested_id)
    if parsed_id is None:
        return target_batch.school_year

    selected = db.get(SchoolYear, parsed_id)
    return selected or target_batch.school_year


def _representative_batches(
    db: Session,
    request: Request,
    school_year_id: int,
    commune_ids: list[int],
) -> dict[int, SurveyBatch]:
    if not commune_ids:
        return {}

    filters = tao_bo_loc_dot_theo_nguoi_dung(request)
    candidates = db.scalars(
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
        )
        .where(
            SurveyBatch.school_year_id == school_year_id,
            SurveyBatch.commune_id.in_(commune_ids),
            *filters,
        )
    ).all()

    result: dict[int, SurveyBatch] = {}
    for batch in candidates:
        current = result.get(batch.commune_id)
        if current is None or _batch_choice_key(batch) > _batch_choice_key(current):
            result[batch.commune_id] = batch

    return result


def _load_forms_by_batch(
    db: Session,
    request: Request,
    batch_ids: list[int],
) -> dict[int, list[SurveyForm]]:
    if not batch_ids:
        return {}

    form_filters = tao_bo_loc_bao_cao_theo_nguoi_dung(request)
    forms = db.scalars(
        select(SurveyForm)
        .options(
            selectinload(SurveyForm.household),
            selectinload(SurveyForm.investigators),
            selectinload(SurveyForm.person_year_records),
        )
        .where(
            SurveyForm.survey_batch_id.in_(batch_ids),
            *form_filters,
        )
        .order_by(SurveyForm.id.asc())
    ).all()

    grouped: dict[int, list[SurveyForm]] = defaultdict(list)
    for form in forms:
        grouped[form.survey_batch_id].append(form)

    return dict(grouped)


def _answered_indicator_total(record: SurveyPersonYearRecord) -> int:
    return sum(
        getattr(record, field_name) is not None
        for field_name in BOOLEAN_YEAR_FIELDS
    )


def _record_needs_review(record: SurveyPersonYearRecord) -> bool:
    learning_status = str(record.learning_status or "CHUA_XAC_DINH")

    if learning_status == "CHUA_XAC_DINH":
        return True

    if learning_status == "DANG_HOC":
        reported_school = _clean_text(record.school_name_reported)
        if record.school_id is None and not reported_school:
            return True

    return _answered_indicator_total(record) < len(BOOLEAN_YEAR_FIELDS)


def _percent(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return round(numerator * 100 / denominator, 1)


def _format_datetime(value: datetime | None) -> str:
    if value is None:
        return "—"
    return value.strftime("%d/%m/%Y %H:%M")


def _missing_row(commune: Commune) -> dict[str, Any]:
    return {
        "commune": commune,
        "batch": None,
        "has_batch": False,
        "state": "missing",
        "state_label": STATE_LABELS["missing"],
        "needs_attention": True,
        "ready": False,
        "forms_total": 0,
        "completed_forms": 0,
        "completion_percent": 0.0,
        "unassigned_forms": 0,
        "people_total": 0,
        "attending_total": 0,
        "learning_known": 0,
        "indicator_completion_percent": 0.0,
        "review_total": 0,
        "last_updated": None,
        "last_updated_label": "—",
        "attention_reasons": ["Chưa tạo đợt điều tra cho năm học này"],
        "attention_text": "Chưa tạo đợt điều tra cho năm học này",
    }


def _build_row(
    commune: Commune,
    batch: SurveyBatch,
    forms: list[SurveyForm],
    school_year_id: int,
) -> dict[str, Any]:
    forms_total = len(forms)
    completed_forms = sum(
        form.status == "DA_HOAN_THANH"
        for form in forms
    )
    unassigned_forms = sum(
        len(form.investigators) == 0
        for form in forms
    )

    records_by_person: dict[int, SurveyPersonYearRecord] = {}
    for form in forms:
        for record in form.person_year_records:
            if record.school_year_id == school_year_id:
                records_by_person[record.survey_person_id] = record

    records = list(records_by_person.values())
    people_total = len(records)
    learning_known = sum(
        str(record.learning_status or "CHUA_XAC_DINH") != "CHUA_XAC_DINH"
        for record in records
    )
    attending_total = sum(
        str(record.learning_status or "") == "DANG_HOC"
        for record in records
    )
    indicator_answered = sum(
        _answered_indicator_total(record)
        for record in records
    )
    indicator_denominator = people_total * len(BOOLEAN_YEAR_FIELDS)
    review_total = sum(
        _record_needs_review(record)
        for record in records
    )

    attention_reasons: list[str] = []
    if forms_total == 0:
        attention_reasons.append("Đợt chưa có phiếu điều tra")
    if completed_forms < forms_total:
        attention_reasons.append(
            f"Còn {forms_total - completed_forms} phiếu chưa hoàn thành"
        )
    if unassigned_forms:
        attention_reasons.append(
            f"Còn {unassigned_forms} phiếu chưa phân công"
        )
    if review_total:
        attention_reasons.append(
            f"Có {review_total} hồ sơ đối tượng cần rà soát"
        )

    needs_attention = bool(attention_reasons)
    ready = forms_total > 0 and not needs_attention

    if batch.is_locked:
        state = "locked"
    elif forms_total == 0:
        state = "empty"
    elif needs_attention:
        state = "attention"
    else:
        state = "ready"

    timestamps = [batch.updated_at]
    timestamps.extend(
        form.updated_at
        for form in forms
        if form.updated_at is not None
    )
    last_updated = max(timestamps) if timestamps else None

    return {
        "commune": commune,
        "batch": batch,
        "has_batch": True,
        "state": state,
        "state_label": STATE_LABELS[state],
        "needs_attention": needs_attention,
        "ready": ready,
        "forms_total": forms_total,
        "completed_forms": completed_forms,
        "completion_percent": _percent(completed_forms, forms_total),
        "unassigned_forms": unassigned_forms,
        "people_total": people_total,
        "attending_total": attending_total,
        "learning_known": learning_known,
        "indicator_completion_percent": _percent(
            indicator_answered,
            indicator_denominator,
        ),
        "review_total": review_total,
        "last_updated": last_updated,
        "last_updated_label": _format_datetime(last_updated),
        "attention_reasons": attention_reasons,
        "attention_text": "; ".join(attention_reasons) or "Không có vấn đề nổi bật",
    }


def _row_matches_view(row: dict[str, Any], view: str) -> bool:
    if view == "attention":
        return bool(row["needs_attention"])
    if view == "missing":
        return not row["has_batch"]
    if view == "open":
        return bool(row["batch"] and not row["batch"].is_locked)
    if view == "locked":
        return bool(row["batch"] and row["batch"].is_locked)
    if view == "ready":
        return bool(row["ready"])
    return True


def _build_dashboard(
    db: Session,
    request: Request,
    selected_year: SchoolYear,
    view: str,
    keyword: str,
) -> dict[str, Any]:
    communes = _communes_in_scope(db, request)
    commune_ids = [commune.id for commune in communes]
    batch_map = _representative_batches(
        db,
        request,
        selected_year.id,
        commune_ids,
    )
    forms_by_batch = _load_forms_by_batch(
        db,
        request,
        [batch.id for batch in batch_map.values()],
    )

    rows: list[dict[str, Any]] = []
    for commune in communes:
        batch = batch_map.get(commune.id)
        if batch is None:
            rows.append(_missing_row(commune))
            continue

        rows.append(
            _build_row(
                commune,
                batch,
                forms_by_batch.get(batch.id, []),
                selected_year.id,
            )
        )

    state_order = {
        "missing": 0,
        "empty": 1,
        "attention": 2,
        "ready": 3,
        "locked": 4,
    }
    rows.sort(
        key=lambda item: (
            state_order.get(item["state"], 9),
            item["commune"].name.casefold(),
            item["commune"].code,
        )
    )

    summary = {
        "communes_total": len(rows),
        "with_batch": sum(row["has_batch"] for row in rows),
        "missing": sum(not row["has_batch"] for row in rows),
        "locked": sum(
            bool(row["batch"] and row["batch"].is_locked)
            for row in rows
        ),
        "ready": sum(row["ready"] for row in rows),
        "attention": sum(row["needs_attention"] for row in rows),
        "forms_total": sum(row["forms_total"] for row in rows),
        "completed_forms": sum(row["completed_forms"] for row in rows),
        "unassigned_forms": sum(row["unassigned_forms"] for row in rows),
        "people_total": sum(row["people_total"] for row in rows),
        "attending_total": sum(row["attending_total"] for row in rows),
        "review_total": sum(row["review_total"] for row in rows),
    }
    summary["completion_percent"] = _percent(
        summary["completed_forms"],
        summary["forms_total"],
    )

    normalized_keyword = keyword.casefold()
    filtered_rows = [
        row
        for row in rows
        if _row_matches_view(row, view)
        and (
            not normalized_keyword
            or normalized_keyword in row["commune"].name.casefold()
            or normalized_keyword in row["commune"].code.casefold()
            or (
                row["batch"] is not None
                and normalized_keyword in row["batch"].code.casefold()
            )
        )
    ]

    return {
        "summary": summary,
        "rows": rows,
        "filtered_rows": filtered_rows,
        "attention_rows": [row for row in rows if row["needs_attention"]],
        "missing_rows": [row for row in rows if not row["has_batch"]],
    }


# =========================================================
# EXCEL
# =========================================================


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1976D2")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def _style_sheet(ws, widths: list[int]) -> None:
    thin = Side(style="thin", color="B8C7D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _append_commune_rows(ws, rows: list[dict[str, Any]]) -> None:
    for index, row in enumerate(rows, start=1):
        batch = row["batch"]
        ws.append(
            [
                index,
                row["commune"].code,
                row["commune"].name,
                batch.code if batch else "",
                batch.name if batch else "",
                (
                    "Đã khóa dữ liệu"
                    if batch and batch.is_locked
                    else BATCH_STATUS_LABELS.get(batch.status, batch.status)
                    if batch
                    else "Chưa tạo đợt"
                ),
                row["forms_total"],
                row["completed_forms"],
                row["completion_percent"],
                row["unassigned_forms"],
                row["people_total"],
                row["attending_total"],
                row["indicator_completion_percent"],
                row["review_total"],
                row["state_label"],
                row["attention_text"],
                row["last_updated_label"],
            ]
        )


def _build_excel(
    request: Request,
    selected_year: SchoolYear,
    dashboard: dict[str, Any],
) -> BytesIO:
    summary = dashboard["summary"]
    wb = Workbook()

    ws = wb.active
    ws.title = "Tong quan"
    _write_header(ws, ["Thông tin", "Giá trị"])
    overview_rows = [
        ("Phạm vi", _scope_label(request)),
        ("Năm học", selected_year.code),
        ("Thời điểm xuất", datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
        ("Tổng xã/phường trong phạm vi", summary["communes_total"]),
        ("Xã/phường đã có đợt", summary["with_batch"]),
        ("Xã/phường chưa có đợt", summary["missing"]),
        ("Đợt đã khóa", summary["locked"]),
        ("Xã/phường sẵn sàng", summary["ready"]),
        ("Xã/phường cần xử lý", summary["attention"]),
        ("Tổng số phiếu", summary["forms_total"]),
        ("Phiếu hoàn thành", summary["completed_forms"]),
        ("Tỷ lệ hoàn thành (%)", summary["completion_percent"]),
        ("Phiếu chưa phân công", summary["unassigned_forms"]),
        ("Tổng số đối tượng", summary["people_total"]),
        ("Đang học", summary["attending_total"]),
        ("Hồ sơ đối tượng cần rà soát", summary["review_total"]),
    ]
    for row in overview_rows:
        ws.append(list(row))
    _style_sheet(ws, [38, 54])

    headers = [
        "STT",
        "Mã xã/phường",
        "Tên xã/phường",
        "Mã đợt đại diện",
        "Tên đợt",
        "Trạng thái đợt",
        "Số phiếu",
        "Phiếu hoàn thành",
        "Tỷ lệ hoàn thành (%)",
        "Chưa phân công",
        "Đối tượng",
        "Đang học",
        "Tỷ lệ nhập chỉ báo (%)",
        "Cần rà soát",
        "Nhóm điều hành",
        "Nội dung cần xử lý",
        "Cập nhật gần nhất",
    ]

    ws_communes = wb.create_sheet("Theo xa phuong")
    _write_header(ws_communes, headers)
    _append_commune_rows(ws_communes, dashboard["rows"])
    _style_sheet(
        ws_communes,
        [7, 18, 28, 25, 36, 21, 11, 17, 20, 16, 12, 12, 23, 14, 19, 55, 20],
    )

    ws_attention = wb.create_sheet("Can xu ly")
    _write_header(ws_attention, headers)
    _append_commune_rows(ws_attention, dashboard["attention_rows"])
    _style_sheet(
        ws_attention,
        [7, 18, 28, 25, 36, 21, 11, 17, 20, 16, 12, 12, 23, 14, 19, 55, 20],
    )

    ws_missing = wb.create_sheet("Xa chua co dot")
    _write_header(ws_missing, ["STT", "Mã xã/phường", "Tên xã/phường", "Ghi chú"])
    for index, row in enumerate(dashboard["missing_rows"], start=1):
        ws_missing.append(
            [
                index,
                row["commune"].code,
                row["commune"].name,
                "Chưa tạo đợt điều tra cho năm học được chọn",
            ]
        )
    _style_sheet(ws_missing, [7, 20, 32, 55])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# =========================================================
# ROUTES
# =========================================================


def _prepare_dashboard(
    db: Session,
    request: Request,
    target_batch_id: int,
    school_year_id: str | None,
    view: str | None,
    q: str | None,
) -> tuple[SurveyBatch, SchoolYear, dict[str, Any], str, str] | None:
    user = lay_thong_tin_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))
    if role_code not in DASHBOARD_ROLE_CODES:
        return None

    target_batch = _batch_in_scope(db, request, target_batch_id)
    if target_batch is None:
        return None

    selected_year = _selected_school_year(
        db,
        target_batch,
        school_year_id,
    )
    selected_view = str(view or "").strip().lower()
    if selected_view not in VIEW_LABELS:
        selected_view = ""
    keyword = _clean_text(q)

    dashboard = _build_dashboard(
        db,
        request,
        selected_year,
        selected_view,
        keyword,
    )
    return target_batch, selected_year, dashboard, selected_view, keyword


@router.get(
    "/{target_batch_id}/bang-dieu-hanh-dia-ban",
    response_class=HTMLResponse,
)
def area_dashboard(
    target_batch_id: int,
    request: Request,
    school_year_id: str | None = None,
    view: str | None = None,
    q: str | None = None,
    db: Session = Depends(get_db),
):
    prepared = _prepare_dashboard(
        db,
        request,
        target_batch_id,
        school_year_id,
        view,
        q,
    )
    if prepared is None:
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    target_batch, selected_year, dashboard, selected_view, keyword = prepared

    return templates.TemplateResponse(
        request=request,
        name="surveys/area_dashboard.html",
        context={
            "nguoi_dung": lay_thong_tin_nguoi_dung(request),
            "scope_label": _scope_label(request),
            "target_batch": target_batch,
            "school_years": _school_years(db),
            "selected_year": selected_year,
            "selected_view": selected_view,
            "keyword": keyword,
            "view_labels": VIEW_LABELS,
            "dashboard": dashboard,
            "batch_status_labels": BATCH_STATUS_LABELS,
            "form_status_labels": FORM_STATUS_LABELS,
        },
    )


@router.get(
    "/{target_batch_id}/bang-dieu-hanh-dia-ban/xuat-excel",
)
def export_area_dashboard(
    target_batch_id: int,
    request: Request,
    school_year_id: str | None = None,
    db: Session = Depends(get_db),
):
    prepared = _prepare_dashboard(
        db,
        request,
        target_batch_id,
        school_year_id,
        "",
        "",
    )
    if prepared is None:
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    _, selected_year, dashboard, _, _ = prepared
    buffer = _build_excel(request, selected_year, dashboard)

    filename = (
        "bang_dieu_hanh_dia_ban_"
        f"{selected_year.code}.xlsx"
    )

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )
