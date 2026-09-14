from __future__ import annotations

import hashlib
import json
import math
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Annotated, Any
from uuid import uuid4

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Classroom, Commune, School, SchoolYear, Student, StudentEnrollment
from app.permissions import (
    ADMIN_ROLE_CODES,
    COMMUNE_ROLE_CODE,
    normalize_role_code,
)
from app.routers.surveys import lay_thong_tin_nguoi_dung
from app.services.school_name_history import school_names_for_year


APP_DIR = Path(__file__).resolve().parent.parent
PROJECT_DIR = APP_DIR.parent
UPLOAD_ROOT = PROJECT_DIR / "uploads" / "student_reconciliation_sources"
PREVIEW_ROOT = UPLOAD_ROOT / "previews"
MAX_UPLOAD_BYTES = 80 * 1024 * 1024

router = APIRouter(
    prefix="/dieu-tra/nguon-hoc-sinh-doi-chieu",
    tags=["Nguồn học sinh đối chiếu theo xã/phường"],
)
templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "school_name": ("Trường", "Tên trường"),
    "class_group": ("Nhóm/Lớp", "Nhóm lớp"),
    "class_name": ("Lớp", "Tên lớp"),
    "student_code": ("Mã định danh Bộ GD&ĐT", "Mã học sinh"),
    "full_name": ("Họ tên", "Họ và tên", "Tên học sinh"),
    "date_of_birth": ("Ngày sinh",),
    "gender": ("Giới tính",),
    "ethnic_group": ("Dân tộc",),
    "status": ("Trạng thái",),
    "province": ("Tỉnh/Thành phố", "Tỉnh"),
    "commune_name": ("Xã/Phường", "Xã phường"),
    "birth_place": ("Nơi sinh",),
    "current_address": ("Chỗ ở hiện nay", "Địa chỉ hiện nay"),
    "hamlet": ("Thôn/Xóm", "Thôn xóm"),
    "contact_phone": ("SĐT liên hệ", "Số điện thoại"),
    "disability_type": ("Loại khuyết tật",),
    "citizen_id": ("Số CCCD", "CCCD"),
    "personal_id": ("Số định danh cá nhân",),
    "father_name": ("Tên cha",),
    "father_job": ("Nghề nghiệp cha",),
    "father_birth_year": ("Năm sinh cha",),
    "mother_name": ("Tên mẹ",),
    "mother_job": ("Nghề nghiệp mẹ",),
    "mother_birth_year": ("Năm sinh mẹ",),
}
REQUIRED_FIELDS = {
    "school_name",
    "class_name",
    "student_code",
    "full_name",
    "date_of_birth",
}

STATUS_MAP = {
    "dang hoc": ("DANG_HOC", True),
    "chuyen den ky 1": ("CHUYEN_DEN_KY_1", True),
    "chuyen den hoc ky 1": ("CHUYEN_DEN_KY_1", True),
    "chuyen den trong he": ("CHUYEN_DEN_KY_1", True),
    "chuyen den": ("CHUYEN_DEN_KY_1", True),
    "chuyen den ky 2": ("CHUYEN_DEN_KY_2", True),
    "chuyen den hoc ky 2": ("CHUYEN_DEN_KY_2", True),
    "chuyen di": ("CHUYEN_DI", False),
    "thoi hoc": ("THOI_HOC", False),
    "bo hoc": ("THOI_HOC", False),
    "tam nghi": ("TAM_NGHI", False),
    "nghi hoc": ("TAM_NGHI", False),
}


def _clean(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _is_36_placeholder(value: Any) -> bool:
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        try:
            return float(value) == 36.0
        except Exception:
            return False
    return str(value).strip() in {"36", "36.0"}


def _optional_text(value: Any) -> str | None:
    if value is None or _is_36_placeholder(value):
        return None
    result = _clean(value)
    return result or None


def _normalize(value: Any) -> str:
    text_value = _clean(value).lower()
    text_value = unicodedata.normalize("NFD", text_value)
    text_value = "".join(
        ch for ch in text_value if unicodedata.category(ch) != "Mn"
    ).replace("đ", "d")
    text_value = re.sub(r"[^a-z0-9]+", " ", text_value)
    return " ".join(text_value.split())


def _school_key(value: Any) -> str:
    value_norm = _normalize(value)
    replacements = (
        (r"^truong trung hoc pho thong\s+", "thpt "),
        (r"^trung hoc pho thong\s+", "thpt "),
        (r"^truong thpt\s+", "thpt "),
        (r"^truong trung hoc co so\s+", "thcs "),
        (r"^trung hoc co so\s+", "thcs "),
        (r"^truong thcs\s+", "thcs "),
        (r"^truong tieu hoc\s+", "th "),
        (r"^tieu hoc\s+", "th "),
        (r"^truong th\s+", "th "),
        (r"^truong mam non\s+", "mn "),
        (r"^mam non\s+", "mn "),
        (r"^truong mn\s+", "mn "),
    )
    for pattern, replacement in replacements:
        value_norm = re.sub(pattern, replacement, value_norm)
    return " ".join(value_norm.split())


def _header_key(value: Any) -> str:
    return _clean(value).lower()


def _normalize_code(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer():
            return str(int(value))
    text_value = str(value).strip()
    if re.fullmatch(r"[+-]?\d+\.0+", text_value):
        text_value = text_value.split(".")[0]
    return re.sub(r"\s+", "", text_value).upper()


def _normalize_date(value: Any, workbook_epoch) -> date | None:
    if value in (None, "") or _is_36_placeholder(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch=workbook_epoch)
            if isinstance(converted, datetime):
                return converted.date()
            if isinstance(converted, date):
                return converted
        except Exception:
            return None
    text_value = _clean(value)
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text_value, fmt).date()
        except ValueError:
            pass
    return None


def _normalize_year(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        year = int(float(str(value).strip()))
    except Exception:
        return None
    return year if 1900 <= year <= date.today().year else None


def _normalize_gender(value: Any) -> str | None:
    text_value = _optional_text(value)
    if not text_value:
        return None
    key = _normalize(text_value)
    if key == "nam":
        return "Nam"
    if key in {"nu", "gai"}:
        return "Nữ"
    return text_value


# === V12_1_STUDENT_STATUS_NORMALIZATION_START ===
def _normalize_status(value: Any) -> tuple[str, bool] | None:
    text_value = _optional_text(value)
    if not text_value:
        return "DANG_HOC", True

    key = _normalize(text_value)

    exact = STATUS_MAP.get(key)
    if exact is not None:
        return exact

    if key == "dang hoc" or key.startswith("dang hoc "):
        return "DANG_HOC", True

    # Nguồn thực tế có 'Nghỉ học xin học lại kỳ 1':
    # đã xin học lại => coi là còn hiệu lực.
    if "nghi hoc xin hoc lai" in key:
        return "DANG_HOC", True

    if key.startswith("chuyen den"):
        if "ky 2" in key or "hoc ky 2" in key:
            return "CHUYEN_DEN_KY_2", True
        return "CHUYEN_DEN_KY_1", True

    if key.startswith("chuyen di"):
        return "CHUYEN_DI", False

    if key.startswith("thoi hoc") or key.startswith("bo hoc"):
        return "THOI_HOC", False

    if key.startswith("tam nghi") or key.startswith("nghi hoc"):
        return "TAM_NGHI", False

    return None
# === V12_1_STUDENT_STATUS_NORMALIZATION_END ===


def _current_user(request: Request) -> dict[str, Any]:
    return lay_thong_tin_nguoi_dung(request) or {}


def _role(request: Request) -> str:
    return normalize_role_code(_current_user(request).get("role_code"))


def _can_manage(request: Request) -> bool:
    return _role(request) in {*ADMIN_ROLE_CODES, COMMUNE_ROLE_CODE}


def _is_admin(request: Request) -> bool:
    return _role(request) in ADMIN_ROLE_CODES


def _scope_commune_id(request: Request, requested: int | None) -> int | None:
    user = _current_user(request)
    role = _role(request)
    if role == COMMUNE_ROLE_CODE:
        value = user.get("commune_id")
        return int(value) if value is not None else None
    if role in ADMIN_ROLE_CODES:
        return requested
    return None


def _previous_year(db: Session, target: SchoolYear) -> SchoolYear | None:
    match = re.search(r"(\d{4})\D+(\d{4})", str(target.code or target.name or ""))
    if not match:
        return None
    start = int(match.group(1))
    wanted_start = start - 1
    for item in db.scalars(select(SchoolYear).order_by(SchoolYear.id.desc())).all():
        m = re.search(r"(\d{4})\D+(\d{4})", str(item.code or item.name or ""))
        if m and int(m.group(1)) == wanted_start and int(m.group(2)) == start:
            return item
    return None


def _default_target_year(years: list[SchoolYear]) -> int | None:
    today = date.today()
    for item in years:
        if item.start_date and item.end_date and item.start_date <= today <= item.end_date:
            return item.id
    for item in years:
        if item.is_active:
            return item.id
    return years[0].id if years else None


def _state(db: Session, target_year_id: int, commune_id: int):
    return db.execute(
        text(
            "SELECT * FROM student_reconciliation_source_states "
            "WHERE target_school_year_id=:target AND commune_id=:commune LIMIT 1"
        ),
        {"target": int(target_year_id), "commune": int(commune_id)},
    ).mappings().first()


def _state_enrollment_ids(db: Session, state_id: int) -> set[int]:
    rows = db.execute(
        text(
            "SELECT enrollment_id FROM student_reconciliation_state_enrollments "
            "WHERE state_id=:state_id"
        ),
        {"state_id": int(state_id)},
    ).all()
    return {int(r[0]) for r in rows}


def _find_header_row(ws) -> int | None:
    required = {_header_key("Mã định danh Bộ GD&ĐT"), _header_key("Họ tên")}
    for row_no, row in enumerate(
        ws.iter_rows(min_row=1, max_row=min(int(ws.max_row or 1), 30), values_only=True),
        start=1,
    ):
        keys = {_header_key(v) for v in row if v is not None}
        if required.issubset(keys):
            return row_no
    return None


def _column_map(header_values: tuple[Any, ...]) -> dict[str, int]:
    normalized = {
        _header_key(value): idx
        for idx, value in enumerate(header_values)
        if value is not None
    }
    result: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            key = _header_key(alias)
            if key in normalized:
                result[field] = normalized[key]
                break
    return result


def _val(row: tuple[Any, ...], mapping: dict[str, int], field: str):
    idx = mapping.get(field)
    return row[idx] if idx is not None and idx < len(row) else None


def _school_lookup(
    db: Session,
    commune_id: int,
    school_year_id: int,
) -> tuple[dict[str, list[School]], dict[int, School], dict[int, str]]:
    schools = list(
        db.scalars(
            select(School)
            .where(School.commune_id == int(commune_id))
            .order_by(School.name.asc())
        ).all()
    )
    names_by_id = school_names_for_year(
        db,
        schools=schools,
        school_year_id=int(school_year_id),
    )
    by_key: dict[str, list[School]] = defaultdict(list)
    by_id: dict[int, School] = {}
    for school in schools:
        school_id = int(school.id)
        historical_name = names_by_id.get(school_id, school.name)
        by_key[_school_key(historical_name)].append(school)
        by_id[school_id] = school
    return by_key, by_id, names_by_id


def _issue(level: str, sheet: str, row: int, code: str, name: str, message: str) -> dict[str, Any]:
    return {
        "level": level,
        "sheet": sheet,
        "row": row,
        "student_code": code,
        "full_name": name,
        "message": message,
    }


def _parse_workbook(
    db: Session,
    content: bytes,
    commune_id: int,
    source_school_year_id: int,
) -> dict[str, Any]:
    result: dict[str, Any] = {
        "records": [],
        "errors": [],
        "warnings": [],
        "school_counts": Counter(),
        "class_keys": set(),
        "sheet_names": [],
    }
    by_school_key, school_by_id, school_name_by_id = _school_lookup(
        db,
        commune_id,
        source_school_year_id,
    )
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        result["sheet_names"] = list(wb.sheetnames)
        for ws in wb.worksheets:
            header_row = _find_header_row(ws)
            if header_row is None:
                continue
            headers = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
            mapping = _column_map(headers)
            missing = REQUIRED_FIELDS - set(mapping)
            if missing:
                result["errors"].append(
                    _issue("LỖI", ws.title, header_row, "", "", "Thiếu cột bắt buộc: " + ", ".join(sorted(missing)))
                )
                continue

            for row_no, row in enumerate(
                ws.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                if all(v is None for v in row):
                    continue
                school_name = _clean(_val(row, mapping, "school_name"))
                class_name = _clean(_val(row, mapping, "class_name"))
                code = _normalize_code(_val(row, mapping, "student_code"))
                full_name = _clean(_val(row, mapping, "full_name"))
                if not any((school_name, class_name, code, full_name)):
                    continue
                dob = _normalize_date(_val(row, mapping, "date_of_birth"), wb.epoch)
                gender = _normalize_gender(_val(row, mapping, "gender"))
                status_info = _normalize_status(_val(row, mapping, "status"))

                row_error = False
                for condition, message in (
                    (not school_name, "Thiếu tên trường."),
                    (not class_name, "Thiếu tên lớp."),
                    (not code, "Thiếu mã học sinh."),
                    (not full_name, "Thiếu họ tên học sinh."),
                    (dob is None, "Ngày sinh trống hoặc không hợp lệ."),
                    (status_info is None, "Trạng thái học sinh chưa được hỗ trợ."),
                ):
                    if condition:
                        result["errors"].append(_issue("LỖI", ws.title, row_no, code, full_name, message))
                        row_error = True
                if row_error:
                    continue

                matches = by_school_key.get(_school_key(school_name), [])
                if len(matches) == 0:
                    result["errors"].append(
                        _issue("LỖI", ws.title, row_no, code, full_name, f"Không khớp trường trong xã/phường: '{school_name}'.")
                    )
                    continue
                if len(matches) > 1:
                    result["errors"].append(
                        _issue("LỖI", ws.title, row_no, code, full_name, f"Tên trường '{school_name}' khớp nhiều trường; cần chuẩn hóa tên/mã trường.")
                    )
                    continue

                school = matches[0]
                personal_id = _optional_text(_val(row, mapping, "personal_id")) or _optional_text(_val(row, mapping, "citizen_id"))
                status_code, is_current = status_info
                province = _optional_text(_val(row, mapping, "province"))
                commune_name = _optional_text(_val(row, mapping, "commune_name"))
                hamlet = _optional_text(_val(row, mapping, "hamlet"))
                permanent_address = ", ".join(x for x in (hamlet, commune_name, province) if x) or None
                record = {
                    "sheet": ws.title,
                    "row": row_no,
                    "school_id": int(school.id),
                    "school_name_excel": school_name,
                    "school_name_db": school_name_by_id.get(int(school.id), school.name),
                    "class_name": class_name,
                    "class_group": _clean(_val(row, mapping, "class_group")) or None,
                    "student_code": code,
                    "full_name": full_name,
                    "date_of_birth": dob.isoformat(),
                    "gender": gender,
                    "ethnic_group": _optional_text(_val(row, mapping, "ethnic_group")),
                    "status": status_code,
                    "is_current": bool(is_current),
                    "birth_place": _optional_text(_val(row, mapping, "birth_place")),
                    "permanent_address": permanent_address,
                    "current_address": _optional_text(_val(row, mapping, "current_address")),
                    "contact_phone": (None if _is_36_placeholder(_val(row, mapping, "contact_phone")) else (_normalize_code(_val(row, mapping, "contact_phone")) or None)),
                    "disability_type": _optional_text(_val(row, mapping, "disability_type")),
                    "personal_id": personal_id,
                    "father_name": _optional_text(_val(row, mapping, "father_name")),
                    "father_job": _optional_text(_val(row, mapping, "father_job")),
                    "father_birth_year": _normalize_year(_val(row, mapping, "father_birth_year")),
                    "mother_name": _optional_text(_val(row, mapping, "mother_name")),
                    "mother_job": _optional_text(_val(row, mapping, "mother_job")),
                    "mother_birth_year": _normalize_year(_val(row, mapping, "mother_birth_year")),
                }
                result["records"].append(record)
                result["school_counts"][int(school.id)] += 1
                result["class_keys"].add((int(school.id), _normalize(class_name)))

        if not result["records"] and not result["errors"]:
            result["errors"].append(_issue("LỖI", "", 0, "", "", "Không tìm thấy sheet có cấu trúc danh sách học sinh."))

        by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
        exact_keys: dict[tuple[str, int, str], dict[str, Any]] = {}
        for record in result["records"]:
            by_code[record["student_code"]].append(record)
            key = (record["student_code"], record["school_id"], _normalize(record["class_name"]))
            if key in exact_keys:
                result["errors"].append(
                    _issue("LỖI", record["sheet"], record["row"], record["student_code"], record["full_name"], "Trùng đúng cùng học sinh + trường + lớp trong file.")
                )
            else:
                exact_keys[key] = record

        transfer_codes = 0
        for code, items in by_code.items():
            if len(items) <= 1:
                continue
            names = {_normalize(x["full_name"]) for x in items}
            dobs = {x["date_of_birth"] for x in items}
            genders = {_normalize(x["gender"]) for x in items if x.get("gender")}
            if len(names) > 1 or len(dobs) > 1 or len(genders) > 1:
                result["errors"].append(
                    _issue("LỖI", items[-1]["sheet"], items[-1]["row"], code, items[-1]["full_name"], "Cùng mã học sinh nhưng thông tin nhận diện không thống nhất.")
                )
                continue
            current_count = sum(1 for x in items if x["is_current"])
            if current_count > 1:
                result["errors"].append(
                    _issue("LỖI", items[-1]["sheet"], items[-1]["row"], code, items[-1]["full_name"], "Một học sinh có nhiều hơn 1 dòng còn hiệu lực trong cùng năm học.")
                )
                continue
            transfer_codes += 1
            result["warnings"].append(
                _issue("CẢNH BÁO", items[-1]["sheet"], items[-1]["row"], code, items[-1]["full_name"], f"Mã xuất hiện {len(items)} lần; được nhận diện là lịch sử chuyển trường/lớp.")
            )

        result["student_count"] = len(by_code)
        result["transfer_code_count"] = transfer_codes
        result["school_counts"] = [
            {
                "school_id": school_id,
                "school_name": school_name_by_id.get(
                    school_id,
                    school_by_id[school_id].name,
                ),
                "rows": count,
            }
            for school_id, count in sorted(
                result["school_counts"].items(),
                key=lambda item: school_name_by_id.get(
                    item[0],
                    school_by_id[item[0]].name,
                ),
            )
        ]
        result["class_count"] = len(result["class_keys"])
        result.pop("class_keys", None)
        return result
    finally:
        wb.close()


def _check_database_conflicts(
    db: Session,
    result: dict[str, Any],
    *,
    source_year_id: int,
    commune_id: int,
    state_row,
) -> None:
    records = result["records"]
    codes = sorted({x["student_code"] for x in records})
    existing: dict[str, Student] = {}
    for start in range(0, len(codes), 700):
        chunk = codes[start:start + 700]
        for student in db.scalars(select(Student).where(Student.code.in_(chunk))).all():
            existing[student.code] = student
    for record in records:
        student = existing.get(record["student_code"])
        if student is None:
            continue
        if _normalize(student.full_name) != _normalize(record["full_name"]):
            result["errors"].append(
                _issue("LỖI", record["sheet"], record["row"], record["student_code"], record["full_name"], "Mã học sinh đã tồn tại trong CSDL nhưng họ tên khác.")
            )
            continue
        if student.date_of_birth and student.date_of_birth.isoformat() != record["date_of_birth"]:
            result["errors"].append(
                _issue("LỖI", record["sheet"], record["row"], record["student_code"], record["full_name"], "Mã học sinh đã tồn tại trong CSDL nhưng ngày sinh khác.")
            )

    owned_ids: set[int] = set()
    if state_row is not None:
        owned_ids = _state_enrollment_ids(db, int(state_row["id"]))

    rows = db.execute(
        text(
            "SELECT e.id FROM student_enrollments e "
            "JOIN schools s ON s.id=e.school_id "
            "WHERE e.school_year_id=:year_id AND s.commune_id=:commune_id"
        ),
        {"year_id": int(source_year_id), "commune_id": int(commune_id)},
    ).all()
    all_existing_ids = {int(r[0]) for r in rows}
    external = all_existing_ids - owned_ids
    result["existing_owned_enrollments"] = len(owned_ids)
    result["existing_external_enrollments"] = len(external)
    if external:
        result["errors"].append(
            _issue(
                "LỖI",
                "",
                0,
                "",
                "",
                f"Địa bàn đã có {len(external)} enrollment năm nguồn không thuộc V12. Dừng để tránh ghi đè dữ liệu khác.",
            )
        )


def _safe_file_name(value: str) -> str:
    name = Path(value or "nguon_hoc_sinh.xlsx").name
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", name)
    return stem[:160] or "nguon_hoc_sinh.xlsx"


def _preview_path(token: str) -> Path:
    return PREVIEW_ROOT / f"{token}.json"


def _load_preview(token: str) -> dict[str, Any] | None:
    if not re.fullmatch(r"[0-9a-f]{32}", token or ""):
        return None
    path = _preview_path(token)
    if not path.exists():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None
    return payload if isinstance(payload, dict) else None


def _page_context(
    db: Session,
    request: Request,
    *,
    target_year_id: int | None = None,
    commune_id: int | None = None,
    status: str = "",
    preview: dict[str, Any] | None = None,
) -> dict[str, Any]:
    years = list(db.scalars(select(SchoolYear).order_by(SchoolYear.code.desc())).all())
    if target_year_id is None:
        target_year_id = _default_target_year(years)
    target_year = db.get(SchoolYear, target_year_id) if target_year_id else None
    source_year = _previous_year(db, target_year) if target_year else None

    scoped_commune = _scope_commune_id(request, commune_id)
    communes: list[Commune] = []
    if _is_admin(request):
        communes = list(db.scalars(select(Commune).where(Commune.is_active.is_(True)).order_by(Commune.name.asc())).all())
    elif scoped_commune is not None:
        item = db.get(Commune, scoped_commune)
        if item:
            communes = [item]

    state_row = None
    if target_year and scoped_commune:
        state_row = _state(db, target_year.id, scoped_commune)

    status_messages = {
        "imported": ("success", "Đã nạp dữ liệu học sinh đối chiếu thành công."),
        "locked": ("success", "Đã khóa nguồn đối chiếu của xã/phường."),
        "unlocked": ("success", "Đã mở khóa nguồn đối chiếu."),
        "forbidden": ("warning", "Tài khoản không có quyền thực hiện thao tác này."),
        "invalid": ("warning", "Thông tin lựa chọn không hợp lệ."),
        "database_error": ("warning", "Không thể hoàn tất giao dịch; hệ thống đã rollback."),
    }
    status_message = status_messages.get(status)

    return {
        "request": request,
        "nguoi_dung": _current_user(request),
        "years": years,
        "target_year_id": target_year_id,
        "target_year": target_year,
        "source_year": source_year,
        "communes": communes,
        "commune_id": scoped_commune,
        "selected_commune": db.get(Commune, scoped_commune) if scoped_commune else None,
        "state": state_row,
        "preview": preview,
        "status_message": {"kind": status_message[0], "text": status_message[1]} if status_message else None,
        "can_manage": _can_manage(request),
        "can_unlock": _is_admin(request),
    }


@router.get("", response_class=HTMLResponse)
def source_page(
    request: Request,
    target_school_year_id: int | None = None,
    commune_id: int | None = None,
    status: str = "",
    db: Session = Depends(get_db),
):
    if not _can_manage(request):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    context = _page_context(
        db,
        request,
        target_year_id=target_school_year_id,
        commune_id=commune_id,
        status=status,
    )
    return templates.TemplateResponse(
        request=request,
        name="surveys/student_reconciliation_source.html",
        context=context,
    )


@router.post("/xem-truoc", response_class=HTMLResponse)
def preview_source(
    request: Request,
    target_school_year_id: Annotated[int, Form()],
    commune_id: Annotated[int, Form()],
    source_file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    if not _can_manage(request):
        return RedirectResponse(url="/dieu-tra/nguon-hoc-sinh-doi-chieu?status=forbidden", status_code=303)

    scoped_commune = _scope_commune_id(request, commune_id)
    target_year = db.get(SchoolYear, target_school_year_id)
    source_year = _previous_year(db, target_year) if target_year else None
    if scoped_commune is None or target_year is None or source_year is None:
        return RedirectResponse(url="/dieu-tra/nguon-hoc-sinh-doi-chieu?status=invalid", status_code=303)

    existing_state = _state(db, target_year.id, scoped_commune)
    if existing_state is not None and str(existing_state["status"]).upper() == "LOCKED":
        preview = {
            "errors": [_issue("LỖI", "", 0, "", "", "Nguồn của xã/phường đã khóa. Sở phải mở khóa trước khi nạp lại.")],
            "warnings": [],
            "records": [],
            "school_counts": [],
            "student_count": 0,
            "class_count": 0,
            "transfer_code_count": 0,
        }
        context = _page_context(db, request, target_year_id=target_year.id, commune_id=scoped_commune, preview=preview)
        return templates.TemplateResponse(request=request, name="surveys/student_reconciliation_source.html", context=context)

    if Path(source_file.filename or "").suffix.lower() != ".xlsx":
        preview = {
            "errors": [_issue("LỖI", "", 0, "", "", "Hiện V12 nhận file Excel .xlsx.")],
            "warnings": [], "records": [], "school_counts": [], "student_count": 0, "class_count": 0, "transfer_code_count": 0,
        }
        context = _page_context(db, request, target_year_id=target_year.id, commune_id=scoped_commune, preview=preview)
        return templates.TemplateResponse(request=request, name="surveys/student_reconciliation_source.html", context=context)

    content = source_file.file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        preview = {
            "errors": [_issue("LỖI", "", 0, "", "", "File vượt quá 80 MB.")],
            "warnings": [], "records": [], "school_counts": [], "student_count": 0, "class_count": 0, "transfer_code_count": 0,
        }
        context = _page_context(db, request, target_year_id=target_year.id, commune_id=scoped_commune, preview=preview)
        return templates.TemplateResponse(request=request, name="surveys/student_reconciliation_source.html", context=context)

    try:
        parsed = _parse_workbook(
            db,
            content,
            scoped_commune,
            int(source_year.id),
        )
    except Exception as exc:
        parsed = {
            "errors": [_issue("LỖI", "", 0, "", "", f"Không đọc được Excel: {type(exc).__name__}: {exc}")],
            "warnings": [], "records": [], "school_counts": [], "student_count": 0, "class_count": 0, "transfer_code_count": 0,
        }

    _check_database_conflicts(
        db,
        parsed,
        source_year_id=source_year.id,
        commune_id=scoped_commune,
        state_row=existing_state,
    )

    token = uuid4().hex
    UPLOAD_ROOT.mkdir(parents=True, exist_ok=True)
    PREVIEW_ROOT.mkdir(parents=True, exist_ok=True)
    safe_name = _safe_file_name(source_file.filename or "nguon_hoc_sinh.xlsx")
    stored_path = UPLOAD_ROOT / f"{scoped_commune}_{source_year.code}_{token}_{safe_name}"
    stored_path.write_bytes(content)
    sha256 = hashlib.sha256(content).hexdigest()

    payload = {
        "token": token,
        "target_school_year_id": int(target_year.id),
        "target_year_code": target_year.code,
        "source_school_year_id": int(source_year.id),
        "source_year_code": source_year.code,
        "commune_id": int(scoped_commune),
        "file_name": safe_name,
        "stored_path": str(stored_path),
        "sha256": sha256,
        "created_at": datetime.now().isoformat(),
        **parsed,
    }
    _preview_path(token).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    preview_view = dict(payload)
    preview_view["records_sample"] = payload["records"][:20]
    preview_view["issues"] = (payload["errors"] + payload["warnings"])[:100]
    preview_view["total_rows"] = len(payload["records"])
    preview_view["can_import"] = len(payload["errors"]) == 0 and len(payload["records"]) > 0
    preview_view.pop("records", None)
    context = _page_context(db, request, target_year_id=target_year.id, commune_id=scoped_commune, preview=preview_view)
    return templates.TemplateResponse(request=request, name="surveys/student_reconciliation_source.html", context=context)


def _fill_blank(student: Student, field: str, value: Any) -> bool:
    if value in (None, ""):
        return False
    current = getattr(student, field)
    if current not in (None, ""):
        return False
    setattr(student, field, value)
    return True


@router.post("/nap")
def import_source(
    request: Request,
    preview_token: Annotated[str, Form()],
    db: Session = Depends(get_db),
):
    if not _can_manage(request):
        return RedirectResponse(url="/dieu-tra/nguon-hoc-sinh-doi-chieu?status=forbidden", status_code=303)
    payload = _load_preview(preview_token)
    if payload is None or payload.get("errors") or not payload.get("records"):
        return RedirectResponse(url="/dieu-tra/nguon-hoc-sinh-doi-chieu?status=invalid", status_code=303)

    commune_id = _scope_commune_id(request, int(payload["commune_id"]))
    if commune_id != int(payload["commune_id"]):
        return RedirectResponse(url="/dieu-tra/nguon-hoc-sinh-doi-chieu?status=forbidden", status_code=303)

    stored = Path(str(payload["stored_path"]))
    try:
        stored.resolve().relative_to(UPLOAD_ROOT.resolve())
    except Exception:
        return RedirectResponse(url="/dieu-tra/nguon-hoc-sinh-doi-chieu?status=invalid", status_code=303)
    if not stored.exists() or hashlib.sha256(stored.read_bytes()).hexdigest() != payload["sha256"]:
        return RedirectResponse(url="/dieu-tra/nguon-hoc-sinh-doi-chieu?status=invalid", status_code=303)

    target_year_id = int(payload["target_school_year_id"])
    source_year_id = int(payload["source_school_year_id"])
    state_row = _state(db, target_year_id, commune_id)
    if state_row is not None and str(state_row["status"]).upper() == "LOCKED":
        return RedirectResponse(
            url=f"/dieu-tra/nguon-hoc-sinh-doi-chieu?target_school_year_id={target_year_id}&commune_id={commune_id}&status=invalid",
            status_code=303,
        )

    actor = _current_user(request)
    now = datetime.now()
    records = payload["records"]
    inserted_students = 0
    updated_students = 0
    inserted_classes = 0
    inserted_enrollments = 0
    skipped_rows = 0

    try:
        if state_row is None:
            db.execute(
                text(
                    "INSERT INTO student_reconciliation_source_states "
                    "(target_school_year_id,source_school_year_id,commune_id,status,created_at,updated_at) "
                    "VALUES (:target,:source,:commune,'DRAFT',:now,:now)"
                ),
                {"target": target_year_id, "source": source_year_id, "commune": commune_id, "now": now},
            )
            db.flush()
            state_row = _state(db, target_year_id, commune_id)
        state_id = int(state_row["id"])

        # === V12_3_CUMULATIVE_IMPORT_START ===
        # Chỉ thay dữ liệu của các trường có trong file đang nạp.
        school_ids = sorted({int(x["school_id"]) for x in records})
        owned_ids = _state_enrollment_ids(db, state_id)

        replace_ids: set[int] = set()
        if owned_ids and school_ids:
            replace_rows = db.execute(
                text(
                    "SELECT se.enrollment_id "
                    "FROM student_reconciliation_state_enrollments se "
                    "JOIN student_enrollments e ON e.id=se.enrollment_id "
                    "WHERE se.state_id=:state_id "
                    "AND e.school_year_id=:source_year_id "
                    "AND e.school_id IN ("
                    + ",".join(str(int(x)) for x in school_ids)
                    + ")"
                ),
                {
                    "state_id": state_id,
                    "source_year_id": source_year_id,
                },
            ).all()
            replace_ids = {int(r[0]) for r in replace_rows}

        if replace_ids:
            for start in range(0, len(replace_ids), 700):
                chunk = sorted(replace_ids)[start:start + 700]
                id_sql = ",".join(str(int(x)) for x in chunk)

                db.execute(
                    text(
                        "DELETE FROM student_reconciliation_state_enrollments "
                        "WHERE state_id=:state_id "
                        "AND enrollment_id IN (" + id_sql + ")"
                    ),
                    {"state_id": state_id},
                )
                db.execute(
                    text(
                        "DELETE FROM student_enrollments "
                        "WHERE id IN (" + id_sql + ")"
                    )
                )

            db.flush()
        # === V12_3_CUMULATIVE_IMPORT_END ===
        classrooms = list(
            db.scalars(
                select(Classroom).where(
                    Classroom.school_year_id == source_year_id,
                    Classroom.school_id.in_(school_ids),
                )
            ).all()
        )
        class_map: dict[tuple[int, str], Classroom] = {
            (int(item.school_id), _normalize(item.name)): item for item in classrooms
        }
        for record in records:
            key = (int(record["school_id"]), _normalize(record["class_name"]))
            if key not in class_map:
                item = Classroom(
                    school_id=key[0],
                    school_year_id=source_year_id,
                    code=record.get("class_group") or None,
                    name=record["class_name"],
                    is_active=True,
                )
                db.add(item)
                db.flush()
                class_map[key] = item
                inserted_classes += 1

        codes = sorted({x["student_code"] for x in records})
        student_map: dict[str, Student] = {}
        for start in range(0, len(codes), 700):
            chunk = codes[start:start + 700]
            for item in db.scalars(select(Student).where(Student.code.in_(chunk))).all():
                student_map[item.code] = item

        by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for record in records:
            by_code[record["student_code"]].append(record)

        touched_student_ids: set[int] = set()
        for code, items in by_code.items():
            first = sorted(items, key=lambda x: (not bool(x["is_current"]), x["row"]))[0]
            student = student_map.get(code)
            if student is None:
                student = Student(
                    code=code,
                    full_name=first["full_name"],
                    date_of_birth=date.fromisoformat(first["date_of_birth"]),
                    gender=first.get("gender"),
                    ethnic_group=first.get("ethnic_group"),
                    birth_place=first.get("birth_place"),
                    permanent_address=first.get("permanent_address"),
                    current_address=first.get("current_address"),
                    father_name=first.get("father_name"),
                    father_birth_year=first.get("father_birth_year"),
                    father_job=first.get("father_job"),
                    mother_name=first.get("mother_name"),
                    mother_birth_year=first.get("mother_birth_year"),
                    mother_job=first.get("mother_job"),
                    contact_phone=first.get("contact_phone"),
                    personal_id=first.get("personal_id"),
                    disability_type=first.get("disability_type"),
                    notes=f"Nguồn đối chiếu {payload['source_year_code']} - xã/phường {commune_id}",
                    is_active=True,
                )
                db.add(student)
                db.flush()
                student_map[code] = student
                inserted_students += 1
            else:
                changed = False
                if student.date_of_birth is None and first.get("date_of_birth"):
                    student.date_of_birth = date.fromisoformat(first["date_of_birth"])
                    changed = True
                for field in (
                    "gender", "ethnic_group", "birth_place", "permanent_address", "current_address",
                    "father_name", "father_birth_year", "father_job", "mother_name", "mother_birth_year",
                    "mother_job", "contact_phone", "personal_id", "disability_type",
                ):
                    changed = _fill_blank(student, field, first.get(field)) or changed
                if changed:
                    updated_students += 1
            touched_student_ids.add(int(student.id))

        for record in records:
            student = student_map[record["student_code"]]
            classroom = class_map[(int(record["school_id"]), _normalize(record["class_name"]))]
            external = db.scalar(
                select(StudentEnrollment).where(
                    StudentEnrollment.student_id == student.id,
                    StudentEnrollment.school_year_id == source_year_id,
                    StudentEnrollment.school_id == int(record["school_id"]),
                    StudentEnrollment.class_id == classroom.id,
                )
            )
            if external is not None:
                raise RuntimeError(
                    f"Enrollment đã tồn tại ngoài lô V12: student={student.code}, school={record['school_id']}, class={classroom.name}"
                )
            enrollment = StudentEnrollment(
                student_id=student.id,
                school_year_id=source_year_id,
                school_id=int(record["school_id"]),
                class_id=classroom.id,
                status=record["status"],
                is_current=bool(record["is_current"]),
            )
            db.add(enrollment)
            db.flush()
            db.execute(
                text(
                    "INSERT INTO student_reconciliation_state_enrollments (state_id,enrollment_id) "
                    "VALUES (:state_id,:enrollment_id)"
                ),
                {"state_id": state_id, "enrollment_id": int(enrollment.id)},
            )
            inserted_enrollments += 1

        if touched_student_ids:
            for start in range(0, len(touched_student_ids), 700):
                chunk = sorted(touched_student_ids)[start:start + 700]
                bad = db.execute(
                    text(
                        "SELECT student_id,COUNT(*) AS n FROM student_enrollments "
                        "WHERE school_year_id=:year_id AND is_current=1 AND student_id IN ("
                        + ",".join(str(int(x)) for x in chunk)
                        + ") GROUP BY student_id HAVING COUNT(*)>1 LIMIT 1"
                    ),
                    {"year_id": source_year_id},
                ).first()
                if bad is not None:
                    raise RuntimeError(f"Học sinh id={bad[0]} có nhiều hơn 1 enrollment hiện hành trong năm nguồn.")

        # V12.3: card trạng thái phản ánh toàn bộ nguồn cộng dồn của xã.
        aggregate = db.execute(
            text(
                "SELECT "
                "COUNT(*) AS enrollment_count, "
                "COUNT(DISTINCT e.student_id) AS student_count, "
                "COUNT(DISTINCT e.class_id) AS class_count "
                "FROM student_reconciliation_state_enrollments se "
                "JOIN student_enrollments e ON e.id=se.enrollment_id "
                "WHERE se.state_id=:state_id"
            ),
            {"state_id": state_id},
        ).mappings().one()

        aggregate_enrollment_count = int(
            aggregate["enrollment_count"] or 0
        )
        aggregate_student_count = int(
            aggregate["student_count"] or 0
        )
        aggregate_class_count = int(
            aggregate["class_count"] or 0
        )

        db.execute(
            text(
                "UPDATE student_reconciliation_source_states SET "
                "source_school_year_id=:source,status='LOADED',file_name=:file_name,total_rows=:total_rows,"
                "valid_rows=:valid_rows,student_count=:student_count,enrollment_count=:enrollment_count,"
                "class_count=:class_count,import_count=COALESCE(import_count,0)+1,last_imported_at=:now,"
                "imported_by_user_id=:actor,locked_at=NULL,locked_by_user_id=NULL,updated_at=:now "
                "WHERE id=:state_id"
            ),
            {
                "source": source_year_id,
                "file_name": payload["file_name"],
                "total_rows": aggregate_enrollment_count,
                "valid_rows": aggregate_enrollment_count,
                "student_count": aggregate_student_count,
                "enrollment_count": aggregate_enrollment_count,
                "class_count": aggregate_class_count,
                "now": now,
                "actor": actor.get("id"),
                "state_id": state_id,
            },
        )
        db.execute(
            text(
                "INSERT INTO student_reconciliation_import_logs "
                "(state_id,file_name,file_sha256,total_rows,valid_rows,inserted_students,updated_students,"
                "inserted_classes,inserted_enrollments,skipped_rows,warning_rows,error_rows,actor_user_id,"
                "actor_name_snapshot,created_at) VALUES "
                "(:state_id,:file_name,:sha,:total,:valid,:ins_students,:upd_students,:ins_classes,:ins_enrollments,"
                ":skipped,:warnings,0,:actor_id,:actor_name,:created_at)"
            ),
            {
                "state_id": state_id,
                "file_name": payload["file_name"],
                "sha": payload["sha256"],
                "total": len(records),
                "valid": len(records),
                "ins_students": inserted_students,
                "upd_students": updated_students,
                "ins_classes": inserted_classes,
                "ins_enrollments": inserted_enrollments,
                "skipped": skipped_rows,
                "warnings": len(payload.get("warnings") or []),
                "actor_id": actor.get("id"),
                "actor_name": _clean(actor.get("full_name")) or _clean(actor.get("username")) or "Không xác định",
                "created_at": now,
            },
        )
        db.commit()
    except Exception:
        db.rollback()
        return RedirectResponse(
            url=f"/dieu-tra/nguon-hoc-sinh-doi-chieu?target_school_year_id={target_year_id}&commune_id={commune_id}&status=database_error",
            status_code=303,
        )

    try:
        _preview_path(preview_token).unlink(missing_ok=True)
    except Exception:
        pass
    return RedirectResponse(
        url=f"/dieu-tra/nguon-hoc-sinh-doi-chieu?target_school_year_id={target_year_id}&commune_id={commune_id}&status=imported",
        status_code=303,
    )


@router.post("/khoa")
def lock_source(
    request: Request,
    target_school_year_id: Annotated[int, Form()],
    commune_id: Annotated[int, Form()],
    db: Session = Depends(get_db),
):
    if not _can_manage(request):
        return RedirectResponse(url="/dieu-tra/nguon-hoc-sinh-doi-chieu?status=forbidden", status_code=303)
    scoped_commune = _scope_commune_id(request, commune_id)
    if scoped_commune is None:
        return RedirectResponse(url="/dieu-tra/nguon-hoc-sinh-doi-chieu?status=forbidden", status_code=303)
    state_row = _state(db, target_school_year_id, scoped_commune)
    if state_row is None or str(state_row["status"]).upper() != "LOADED":
        return RedirectResponse(url="/dieu-tra/nguon-hoc-sinh-doi-chieu?status=invalid", status_code=303)
    user = _current_user(request)
    now = datetime.now()
    db.execute(
        text(
            "UPDATE student_reconciliation_source_states SET status='LOCKED',locked_at=:now,"
            "locked_by_user_id=:uid,updated_at=:now WHERE id=:id"
        ),
        {"now": now, "uid": user.get("id"), "id": int(state_row["id"])},
    )
    db.commit()
    return RedirectResponse(
        url=f"/dieu-tra/nguon-hoc-sinh-doi-chieu?target_school_year_id={target_school_year_id}&commune_id={scoped_commune}&status=locked",
        status_code=303,
    )


@router.post("/mo-khoa")
def unlock_source(
    request: Request,
    target_school_year_id: Annotated[int, Form()],
    commune_id: Annotated[int, Form()],
    db: Session = Depends(get_db),
):
    if not _is_admin(request):
        return RedirectResponse(url="/dieu-tra/nguon-hoc-sinh-doi-chieu?status=forbidden", status_code=303)
    state_row = _state(db, target_school_year_id, commune_id)
    if state_row is None:
        return RedirectResponse(url="/dieu-tra/nguon-hoc-sinh-doi-chieu?status=invalid", status_code=303)
    now = datetime.now()
    db.execute(
        text(
            "UPDATE student_reconciliation_source_states SET status='LOADED',locked_at=NULL,"
            "locked_by_user_id=NULL,updated_at=:now WHERE id=:id"
        ),
        {"now": now, "id": int(state_row["id"])},
    )
    db.commit()
    return RedirectResponse(
        url=f"/dieu-tra/nguon-hoc-sinh-doi-chieu?target_school_year_id={target_school_year_id}&commune_id={commune_id}&status=unlocked",
        status_code=303,
    )
