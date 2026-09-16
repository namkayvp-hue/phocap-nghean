from __future__ import annotations

from io import BytesIO
from typing import Any
from urllib.parse import urlencode

from fastapi import APIRouter, Form, Query, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from app.permissions import is_admin_role, normalize_role_code
from app.services.school_merger_official_registry import list_official_registry_plans, official_registry_summary
from app.services.school_merger_execution_service import annotate_official_plan_execution_status
from app.services.school_merger_source_audit_service import audit_official_plans
from app.services.school_merger_service import list_communes, list_school_years, list_schools
from app.services.school_merger_future_service import (
    FUTURE_CONFIRM_PHRASE,
    SchoolMergerFutureError,
    build_future_preview,
    execute_future_merge,
    get_future_execution_result,
    list_future_executions,
)


router = APIRouter(
    prefix="/cong-cu-du-lieu/sap-nhap-truong",
    tags=["Sáp nhập trường"],
)
templates = Jinja2Templates(directory="app/templates")

LEVEL_CHOICES = (("MN", "Mầm non"), ("TH", "Tiểu học"), ("THCS", "THCS"), ("THPT", "THPT"))
DASHBOARD_LEVEL_CHOICES = (("MN", "Mầm non"), ("TH", "Tiểu học"), ("THCS", "THCS"))


def _auth_user(request: Request) -> dict[str, Any]:
    return dict(request.scope.get("auth_user") or {})


def _admin_only(request: Request) -> bool:
    return is_admin_role(normalize_role_code(_auth_user(request).get("role_code")))


def _safe_int(value: Any) -> int | None:
    try:
        text = str(value or "").strip()
        return int(text) if text else None
    except (TypeError, ValueError):
        return None


def _default_year_id(years: list[dict[str, Any]]) -> int | None:
    wanted = next((x for x in years if str(x.get("code") or "").replace("–", "-").replace("—", "-") == "2026-2027"), None)
    if wanted is not None:
        return int(wanted["id"])
    active = next((x for x in years if x.get("is_active")), None)
    if active is not None:
        return int(active["id"])
    return int(years[0]["id"]) if years else None


def _dashboard_payload(*, school_year_id: int | None, commune_id: int | None, level_code: str) -> dict[str, Any]:
    level = str(level_code or "TH").upper().strip()
    if level not in {"MN", "TH", "THCS"}:
        level = "TH"
    plans = list_official_registry_plans(
        school_year_id=school_year_id,
        commune_id=commune_id,
        level_code=level,
        display_mode="all",
    )
    plans = annotate_official_plan_execution_status(plans)
    audit = audit_official_plans(
        school_year_id=school_year_id,
        commune_id=commune_id,
        level_code=level,
    )
    audit_map = {str((row.get("plan") or {}).get("id") or ""): row for row in (audit.get("rows") or [])}

    rows: list[dict[str, Any]] = []
    counts = {"KEEP": 0, "DONE": 0, "READY": 0, "BLOCK": 0}
    for plan in plans:
        item = dict(plan)
        plan_id = str(item.get("id") or "")
        audit_row = audit_map.get(plan_id) or {}
        audit_item = dict(audit_row.get("audit") or {})
        execution_status = str(item.get("execution_status") or "")
        if str(item.get("action_type") or "").upper() == "KEEP":
            state = "KEEP"
            label = "GIỮ NGUYÊN"
            reasons: list[str] = []
        elif execution_status.startswith("COMPLETED"):
            state = "DONE"
            label = "ĐÃ THỰC HIỆN"
            reasons = []
        elif (
            str(item.get("action_type") or "").upper() == "MAPPED"
            and str(item.get("match_status") or "") == "MATCHED"
            and bool(audit_item.get("pass"))
        ):
            state = "READY"
            label = "PASS – CÓ THỂ XEM TRƯỚC"
            reasons = []
        else:
            state = "BLOCK"
            label = "BLOCK – CẦN XỬ LÝ"
            reasons = [str(x) for x in (audit_item.get("blockers") or [])]
            if str(item.get("match_status") or "") != "MATCHED":
                reasons.insert(0, str(item.get("match_label") or "Chưa khớp trường trong database."))
            if not reasons:
                reasons = ["Phương án chưa đủ điều kiện kỹ thuật để thực hiện."]
        counts[state] += 1
        item.update({
            "dashboard_state": state,
            "dashboard_label": label,
            "dashboard_reasons": reasons,
            "source_audit": audit_item,
        })
        rows.append(item)

    return {
        "rows": rows,
        "counts": counts,
        "total": len(rows),
        "audit": audit,
        "registry_summary": official_registry_summary(),
        "level_code": level,
    }


@router.get("/bang-dieu-khien", response_class=HTMLResponse)
def merger_dashboard(
    request: Request,
    school_year_id: str | None = Query(default=None),
    commune_id: str | None = Query(default=None),
    level_code: str = Query(default="TH"),
    state: str = Query(default=""),
):
    if not _admin_only(request):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    years = list_school_years()
    year_id = _safe_int(school_year_id) or _default_year_id(years)
    commune_id_int = _safe_int(commune_id)
    payload = _dashboard_payload(school_year_id=year_id, commune_id=commune_id_int, level_code=level_code)
    wanted_state = str(state or "").upper().strip()
    visible_rows = payload["rows"] if wanted_state not in {"KEEP", "DONE", "READY", "BLOCK"} else [
        x for x in payload["rows"] if x.get("dashboard_state") == wanted_state
    ]
    selected_commune = next((x for x in list_communes() if int(x["id"]) == int(commune_id_int or -1)), None)
    return templates.TemplateResponse(
        request=request,
        name="data_tools/school_merger_dashboard_v45.html",
        context={
            "nguoi_dung": _auth_user(request),
            "school_years": years,
            "communes": list_communes(),
            "selected_school_year_id": year_id,
            "selected_commune_id": commune_id_int,
            "selected_commune_name": str((selected_commune or {}).get("name") or ""),
            "selected_level_code": payload["level_code"],
            "selected_state": wanted_state,
            "level_choices": DASHBOARD_LEVEL_CHOICES,
            "rows": visible_rows,
            "counts": payload["counts"],
            "total": payload["total"],
            "registry_summary": payload["registry_summary"],
            "future_executions": list_future_executions(20),
        },
    )


@router.get("/bang-dieu-khien/xuat-block")
def export_block_xlsx(
    request: Request,
    school_year_id: str | None = Query(default=None),
    commune_id: str | None = Query(default=None),
    level_code: str = Query(default="TH"),
):
    if not _admin_only(request):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    years = list_school_years()
    year_id = _safe_int(school_year_id) or _default_year_id(years)
    commune_id_int = _safe_int(commune_id)
    payload = _dashboard_payload(school_year_id=year_id, commune_id=commune_id_int, level_code=level_code)
    blocks = [x for x in payload["rows"] if x.get("dashboard_state") == "BLOCK"]

    wb = Workbook()
    ws = wb.active
    ws.title = "BLOCK"
    headers = ["STT", "Mã phương án", "Địa bàn", "Cấp", "Nhóm trường", "Phương án", "Trạng thái DB", "Nguyên nhân BLOCK"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for idx, p in enumerate(blocks, start=1):
        group = " + ".join(str(x.get("excel_name") or x.get("name") or "") for x in (p.get("member_schools") or []))
        reasons = " | ".join(str(x) for x in (p.get("dashboard_reasons") or []))
        ws.append([
            idx,
            p.get("id"),
            p.get("commune_excel"),
            " + ".join(p.get("level_codes") or []),
            group,
            p.get("plan_text"),
            p.get("match_label"),
            reasons,
        ])
    widths = [8, 24, 22, 12, 50, 40, 28, 90]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    level = str(payload["level_code"] or "TH")
    filename = f"Danh_sach_BLOCK_sap_nhap_{level}_{year_id or 'nam'}.xlsx"
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/phat-sinh", response_class=HTMLResponse)
def future_merge_page(
    request: Request,
    school_year_id: str | None = Query(default=None),
    commune_id: str | None = Query(default=None),
    level_code: str = Query(default="TH"),
    status: str = Query(default=""),
):
    if not _admin_only(request):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    years = list_school_years()
    year_id = _safe_int(school_year_id) or _default_year_id(years)
    commune_id_int = _safe_int(commune_id)
    communes = list_communes()
    selected_commune = next((x for x in communes if int(x["id"]) == int(commune_id_int or -1)), None)
    level = str(level_code or "TH").upper().strip()
    if level not in {x[0] for x in LEVEL_CHOICES}:
        level = "TH"
    schools = (
        list_schools(
            year_id=year_id,
            commune_id=commune_id_int,
            level_code=level,
            include_inactive=False,
        )
        if year_id is not None and commune_id_int is not None
        else []
    )
    message = ""
    if status == "stale":
        message = "Dữ liệu đã thay đổi; hãy XEM TRƯỚC lại."
    return templates.TemplateResponse(
        request=request,
        name="data_tools/school_merger_future_v45.html",
        context={
            "nguoi_dung": _auth_user(request),
            "school_years": years,
            "selected_school_year_id": year_id,
            "selected_school_year_label": next((str(y.get("code") or y.get("name") or "") for y in years if int(y.get("id") or -1) == int(year_id or -1)), ""),
            "communes": communes,
            "selected_commune_id": commune_id_int,
            "selected_commune_name": str((selected_commune or {}).get("name") or ""),
            "selected_level_code": level,
            "selected_level_label": next((label for code, label in LEVEL_CHOICES if code == level), level),
            "level_choices": LEVEL_CHOICES,
            "schools": schools,
            "status_message": message,
            "future_executions": list_future_executions(30),
        },
    )


@router.post("/phat-sinh/xem-truoc", response_class=HTMLResponse)
def future_merge_preview(
    request: Request,
    school_year_id: str = Form(default=""),
    commune_id: str = Form(default=""),
    level_code: str = Form(default=""),
    target_school_id: str = Form(default=""),
    source_school_ids: list[str] = Form(default=[]),
):
    if not _admin_only(request):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    year_id = _safe_int(school_year_id)
    target_id = _safe_int(target_school_id)
    source_ids = [int(x) for x in source_school_ids if _safe_int(x) is not None]
    if year_id is None or target_id is None or not source_ids:
        q = urlencode({"school_year_id": school_year_id, "commune_id": commune_id, "level_code": level_code})
        return RedirectResponse(url=f"/cong-cu-du-lieu/sap-nhap-truong/phat-sinh?{q}", status_code=303)
    try:
        preview = build_future_preview(
            school_year_id=year_id,
            level_code=level_code,
            target_school_id=target_id,
            source_ids=source_ids,
        )
        error_message = ""
        status_code = 200
    except Exception as exc:
        preview = None
        error_message = str(exc)
        status_code = 409
    return templates.TemplateResponse(
        request=request,
        name="data_tools/school_merger_future_preview_v45.html",
        context={
            "nguoi_dung": _auth_user(request),
            "preview": preview,
            "error_message": error_message,
            "confirm_phrase": FUTURE_CONFIRM_PHRASE,
            "selected_commune_id": _safe_int(commune_id),
        },
        status_code=status_code,
    )


@router.post("/phat-sinh/thuc-hien")
def future_merge_execute(
    request: Request,
    school_year_id: str = Form(default=""),
    level_code: str = Form(default=""),
    target_school_id: str = Form(default=""),
    source_school_ids: list[str] = Form(default=[]),
    expected_impact_state_fingerprint: str = Form(default=""),
    expected_source_audit_fingerprint: str = Form(default=""),
    confirmation_scope: str = Form(default=""),
    confirmation_text: str = Form(default=""),
):
    if not _admin_only(request):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    year_id = _safe_int(school_year_id)
    target_id = _safe_int(target_school_id)
    source_ids = [int(x) for x in source_school_ids if _safe_int(x) is not None]
    if year_id is None or target_id is None or not source_ids:
        return RedirectResponse(url="/cong-cu-du-lieu/sap-nhap-truong/phat-sinh", status_code=303)
    try:
        result = execute_future_merge(
            school_year_id=year_id,
            level_code=level_code,
            target_school_id=target_id,
            source_ids=source_ids,
            actor=_auth_user(request),
            expected_impact_state_fingerprint=expected_impact_state_fingerprint,
            expected_source_audit_fingerprint=expected_source_audit_fingerprint,
            confirmation_scope=confirmation_scope,
            confirmation_text=confirmation_text,
        )
    except SchoolMergerFutureError as exc:
        try:
            preview = build_future_preview(
                school_year_id=year_id,
                level_code=level_code,
                target_school_id=target_id,
                source_ids=source_ids,
            )
        except Exception:
            preview = None
        return templates.TemplateResponse(
            request=request,
            name="data_tools/school_merger_future_preview_v45.html",
            context={
                "nguoi_dung": _auth_user(request),
                "preview": preview,
                "error_message": str(exc),
                "confirm_phrase": FUTURE_CONFIRM_PHRASE,
            },
            status_code=409,
        )
    return RedirectResponse(
        url=f"/cong-cu-du-lieu/sap-nhap-truong/phat-sinh/ket-qua/{result['execution_id']}",
        status_code=303,
    )


@router.get("/phat-sinh/ket-qua/{execution_id}", response_class=HTMLResponse)
def future_merge_result(request: Request, execution_id: int):
    if not _admin_only(request):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    result = get_future_execution_result(int(execution_id))
    return templates.TemplateResponse(
        request=request,
        name="data_tools/school_merger_future_result_v45.html",
        context={
            "nguoi_dung": _auth_user(request),
            "result": result,
            "error_message": "" if result else "Không tìm thấy nhật ký sáp nhập phát sinh.",
        },
        status_code=200 if result else 404,
    )
