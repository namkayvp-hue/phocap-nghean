from __future__ import annotations

import re
import shutil
import unicodedata
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from app.database import DATABASE_PATH, get_db
from app.models import Commune, SchoolYear
from app.permissions import (
    ADMIN_ROLE_CODES,
    COMMUNE_ROLE_CODE,
    DEPARTMENT_ROLE_CODE,
    is_admin_role,
    normalize_role_code,
)
from app.survey_models import (
    HistoricalDataLog,
    HistoricalDataset,
    HistoricalHousehold,
    HistoricalPerson,
)


APP_DIR = Path(__file__).resolve().parent.parent
PROJECT_DIR = APP_DIR.parent
BACKUP_DIR = PROJECT_DIR / "data" / "backups"

templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

router = APIRouter(
    prefix="/dieu-tra/du-lieu-lich-su",
    tags=["Dữ liệu lịch sử"],
)


DATASET_STATUS_LABELS = {
    "DANG_CAP_NHAT": "Đang cập nhật",
    "DA_GUI_KIEM_TRA": "Đã gửi kiểm tra",
    "DA_KHOA_LAM_NEN": "Đã khóa làm dữ liệu nền",
}

DATASET_STATUS_CLASSES = {
    "DANG_CAP_NHAT": "status-editing",
    "DA_GUI_KIEM_TRA": "status-review",
    "DA_KHOA_LAM_NEN": "status-locked",
}

LOG_ACTION_LABELS = {
    "TAO_BO_DU_LIEU": "Tạo bộ dữ liệu",
    "TAO_BO_DU_LIEU_TOAN_TINH": "Tạo bộ dữ liệu toàn tỉnh",
    "NHAP_EXCEL": "Nhập Excel",
    "NHAP_EXCEL_TOAN_TINH": "Nhập Excel toàn tỉnh",
    "XUAT_EXCEL": "Xuất Excel",
    "THEM_HO": "Thêm hộ",
    "SUA_HO": "Sửa hộ",
    "THEM_DOI_TUONG": "Thêm đối tượng",
    "SUA_DOI_TUONG": "Sửa đối tượng",
    "GUI_KIEM_TRA": "Gửi kiểm tra",
    "KHOA_LAM_NEN": "Khóa làm dữ liệu nền",
    "MO_KHOA": "Mở lại cập nhật",
}

RESIDENCY_STATUS_LABELS = {
    "THUONG_TRU": "Thường trú",
    "TAM_TRU": "Tạm trú",
    "TAM_VANG": "Tạm vắng",
    "CHUYEN_DEN": "Chuyển đến",
    "CHUYEN_DI": "Chuyển đi",
    "CHUA_XAC_DINH": "Chưa xác định",
}

LEARNING_STATUS_LABELS = {
    "DANG_HOC": "Đang học",
    "CHUA_DI_HOC": "Chưa đi học",
    "CHUYEN_DEN": "Chuyển đến",
    "CHUYEN_DI": "Chuyển đi",
    "TAM_NGHI": "Tạm nghỉ",
    "THOI_HOC": "Thôi học",
    "BO_HOC": "Bỏ học",
    "DA_TOT_NGHIEP": "Đã tốt nghiệp",
    "KHONG_THUOC_DIEN": "Không thuộc diện theo dõi",
    "CHUA_XAC_DINH": "Chưa xác định",
}

BOOLEAN_FIELDS = (
    "completed_preschool_5",
    "attends_required_days",
    "attends_regularly",
    "prepared_vietnamese",
    "weight_monitored",
    "underweight",
    "height_monitored",
    "stunted",
)

HISTORICAL_HEADERS = [
    "Mã xã/phường",
    "Tên xã/phường",
    "Năm học",
    "Mã hộ",
    "Chủ hộ",
    "Thôn/xóm",
    "Địa chỉ hộ",
    "Điện thoại",
    "Ghi chú hộ",
    "Mã đối tượng",
    "Họ và tên",
    "Ngày sinh",
    "Giới tính",
    "Dân tộc",
    "Nơi sinh",
    "Họ tên cha",
    "Họ tên mẹ",
    "Điện thoại liên hệ đối tượng",
    "Quan hệ với chủ hộ",
    "Số định danh cá nhân",
    "Mã học sinh Bộ GD&ĐT",
    "Địa chỉ thường trú",
    "Địa chỉ hiện tại",
    "Trạng thái cư trú",
    "Trạng thái học tập",
    "Trường gần nhất",
    "Lớp gần nhất",
    "Hoàn thành CT MN 5 tuổi",
    "Đi học đủ ngày",
    "Đi học chuyên cần",
    "Được chuẩn bị tiếng Việt",
    "Theo dõi cân nặng",
    "Suy dinh dưỡng nhẹ cân",
    "Theo dõi chiều cao",
    "Suy dinh dưỡng thấp còi",
    "Hoàn cảnh đặc biệt",
    "Ghi chú đối tượng",
]

HEADER_KEYS = {
    "Mã xã/phường": "commune_code",
    "Tên xã/phường": "commune_name",
    "Năm học": "school_year_code",
    "Mã hộ": "household_code",
    "Chủ hộ": "head_name",
    "Thôn/xóm": "hamlet_name",
    "Địa chỉ hộ": "household_address",
    "Điện thoại": "phone",
    "Ghi chú hộ": "household_notes",
    "Mã đối tượng": "person_code",
    "Họ và tên": "full_name",
    "Ngày sinh": "date_of_birth",
    "Giới tính": "gender",
    "Dân tộc": "ethnic_group",
    "Nơi sinh": "birth_place",
    "Họ tên cha": "father_name",
    "Họ tên mẹ": "mother_name",
    "Điện thoại liên hệ đối tượng": "contact_phone",
    "Quan hệ với chủ hộ": "relationship_to_head",
    "Số định danh cá nhân": "personal_id",
    "Mã học sinh Bộ GD&ĐT": "ministry_student_code",
    "Địa chỉ thường trú": "permanent_address",
    "Địa chỉ hiện tại": "current_address",
    "Trạng thái cư trú": "residency_status",
    "Trạng thái học tập": "learning_status",
    "Trường gần nhất": "school_name_reported",
    "Lớp gần nhất": "class_name_reported",
    "Hoàn thành CT MN 5 tuổi": "completed_preschool_5",
    "Đi học đủ ngày": "attends_required_days",
    "Đi học chuyên cần": "attends_regularly",
    "Được chuẩn bị tiếng Việt": "prepared_vietnamese",
    "Theo dõi cân nặng": "weight_monitored",
    "Suy dinh dưỡng nhẹ cân": "underweight",
    "Theo dõi chiều cao": "height_monitored",
    "Suy dinh dưỡng thấp còi": "stunted",
    "Hoàn cảnh đặc biệt": "special_circumstances",
    "Ghi chú đối tượng": "person_notes",
}

STATUS_MESSAGES = {
    "dataset_created": "Đã tạo bộ dữ liệu lịch sử thành công.",
    "datasets_created": "Đã tạo các bộ dữ liệu lịch sử còn thiếu cho toàn tỉnh.",
    "household_created": "Đã thêm hộ lịch sử thành công.",
    "household_updated": "Đã cập nhật hộ lịch sử thành công.",
    "person_created": "Đã thêm đối tượng lịch sử thành công.",
    "person_updated": "Đã cập nhật đối tượng lịch sử thành công.",
    "imported": "Đã kiểm tra và cập nhật dữ liệu lịch sử từ Excel thành công.",
    "province_imported": "Đã nhập dữ liệu lịch sử toàn tỉnh từ Excel thành công.",
    "submitted": "Đã gửi bộ dữ liệu lịch sử để Sở kiểm tra.",
    "locked": "Đã khóa bộ dữ liệu để làm dữ liệu nền.",
    "unlocked": "Đã mở lại bộ dữ liệu lịch sử để cập nhật.",
}

ERROR_MESSAGES = {
    "forbidden": "Tài khoản không có quyền thực hiện thao tác này.",
    "dataset_locked": "Bộ dữ liệu đã khóa hoặc đang chờ kiểm tra nên không thể cập nhật.",
    "missing_file": "Chưa chọn tệp Excel cần nhập.",
    "invalid_file": "Tệp Excel không đúng mẫu hoặc không thể đọc.",
    "wrong_scope": "Tệp có dữ liệu ngoài phạm vi xã/phường hoặc năm học đã chọn.",
    "missing_required": "Tệp có dòng thiếu Chủ hộ, Địa chỉ hộ hoặc Họ và tên đối tượng.",
    "duplicate": "Mã hộ hoặc mã đối tượng bị trùng không hợp lệ.",
    "reason_required": "Cần nhập lý do trước khi khóa hoặc mở khóa.",
}


def current_user(request: Request) -> dict[str, Any]:
    return dict(request.scope.get("auth_user") or {})


def role_code(request: Request) -> str:
    return normalize_role_code(current_user(request).get("role_code"))


def can_view_center(request: Request) -> bool:
    return role_code(request) in {
        *ADMIN_ROLE_CODES,
        DEPARTMENT_ROLE_CODE,
        COMMUNE_ROLE_CODE,
    }


def can_update_history(request: Request) -> bool:
    return role_code(request) in {*ADMIN_ROLE_CODES, COMMUNE_ROLE_CODE}


def can_manage_province(request: Request) -> bool:
    return is_admin_role(role_code(request))


def commune_scope_id(request: Request) -> int | None:
    if role_code(request) != COMMUNE_ROLE_CODE:
        return None
    value = current_user(request).get("commune_id")
    return int(value) if value is not None else -1


def dataset_in_scope(request: Request, dataset: HistoricalDataset) -> bool:
    scope_id = commune_scope_id(request)
    return scope_id is None or dataset.commune_id == scope_id


def dataset_editable(request: Request, dataset: HistoricalDataset) -> bool:
    return (
        can_update_history(request)
        and dataset_in_scope(request, dataset)
        and dataset.status == "DANG_CAP_NHAT"
    )


HISTORICAL_YEAR_MIN = "2020-2021"
HISTORICAL_YEAR_MAX = "2025-2026"


def historical_year_filter() -> Any:
    return SchoolYear.code.between(HISTORICAL_YEAR_MIN, HISTORICAL_YEAR_MAX)


def is_historical_year_code(code: str) -> bool:
    normalized = normalize_text(code)
    return HISTORICAL_YEAR_MIN <= normalized <= HISTORICAL_YEAR_MAX


def normalize_text(value: Any) -> str:
    return str(value or "").strip()


def normalize_header(value: Any) -> str:
    text_value = normalize_text(value).lower()
    text_value = unicodedata.normalize("NFD", text_value)
    text_value = "".join(ch for ch in text_value if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", "", text_value)


def parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text_value = normalize_text(value)
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text_value, pattern).date()
        except ValueError:
            continue
    return None


def parse_bool(value: Any) -> bool | None:
    if value is None or normalize_text(value) == "":
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(int(value))
    normalized = normalize_text(value).lower()
    if normalized in {"1", "x", "co", "có", "yes", "true", "dung", "đúng"}:
        return True
    if normalized in {"0", "khong", "không", "no", "false", "sai"}:
        return False
    return None


def bool_to_excel(value: bool | None) -> str:
    if value is True:
        return "Có"
    if value is False:
        return "Không"
    return ""


def normalize_status(value: Any, labels: dict[str, str], default: str) -> str:
    raw = normalize_text(value)
    if not raw:
        return default
    upper = raw.upper().replace(" ", "_")
    if upper in labels:
        return upper
    normalized_raw = normalize_header(raw)
    for code, label in labels.items():
        if normalize_header(label) == normalized_raw:
            return code
    return default


def school_year_digits(code: str) -> str:
    return re.sub(r"\D", "", code)


def backup_database(prefix: str) -> Path | None:
    if not DATABASE_PATH.exists():
        return None
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    target = BACKUP_DIR / f"{prefix}_{timestamp}.db"
    shutil.copy2(DATABASE_PATH, target)
    return target


def dataset_counts(db: Session, dataset_id: int) -> tuple[int, int]:
    household_total = db.scalar(
        select(func.count(HistoricalHousehold.id)).where(
            HistoricalHousehold.dataset_id == dataset_id,
            HistoricalHousehold.is_active.is_(True),
        )
    )
    person_total = db.scalar(
        select(func.count(HistoricalPerson.id)).where(
            HistoricalPerson.dataset_id == dataset_id,
            HistoricalPerson.is_active.is_(True),
        )
    )
    return int(household_total or 0), int(person_total or 0)


def add_log(
    db: Session,
    request: Request,
    dataset: HistoricalDataset,
    action: str,
    *,
    file_name: str | None = None,
    inserted_households: int = 0,
    updated_households: int = 0,
    inserted_people: int = 0,
    updated_people: int = 0,
    notes: str | None = None,
) -> None:
    user = current_user(request)
    household_total, person_total = dataset_counts(db, dataset.id)
    db.add(
        HistoricalDataLog(
            dataset_id=dataset.id,
            action=action,
            actor_user_id=user.get("id"),
            actor_name_snapshot=normalize_text(user.get("full_name")) or None,
            actor_role_snapshot=normalize_text(user.get("role_name")) or role_code(request),
            file_name=file_name,
            household_total=household_total,
            person_total=person_total,
            inserted_household_total=inserted_households,
            updated_household_total=updated_households,
            inserted_person_total=inserted_people,
            updated_person_total=updated_people,
            notes=notes,
        )
    )


def next_code(
    db: Session,
    dataset: HistoricalDataset,
    model: type[HistoricalHousehold] | type[HistoricalPerson],
) -> str:
    prefix = "LSH" if model is HistoricalHousehold else "LSDT"
    code_column = (
        HistoricalHousehold.household_code
        if model is HistoricalHousehold
        else HistoricalPerson.person_code
    )
    count_value = db.scalar(
        select(func.count()).select_from(model).where(model.dataset_id == dataset.id)
    )
    sequence = int(count_value or 0) + 1
    year_part = school_year_digits(dataset.school_year.code)
    commune_part = re.sub(r"[^A-Za-z0-9]", "", dataset.commune.code)
    while True:
        candidate = f"{prefix}-{commune_part}-{year_part}-{sequence:06d}"
        exists = db.scalar(
            select(func.count()).select_from(model).where(
                model.dataset_id == dataset.id,
                code_column == candidate,
            )
        )
        if not exists:
            return candidate
        sequence += 1


def get_dataset(db: Session, dataset_id: int) -> HistoricalDataset | None:
    return db.scalar(
        select(HistoricalDataset)
        .options(
            selectinload(HistoricalDataset.school_year),
            selectinload(HistoricalDataset.commune),
            selectinload(HistoricalDataset.logs),
        )
        .where(HistoricalDataset.id == dataset_id)
    )


def redirect_center(**params: Any) -> RedirectResponse:
    clean = {key: value for key, value in params.items() if value not in (None, "")}
    suffix = f"?{urlencode(clean)}" if clean else ""
    return RedirectResponse(
        url=f"/dieu-tra/du-lieu-lich-su{suffix}",
        status_code=303,
    )


def redirect_dataset(dataset_id: int, **params: Any) -> RedirectResponse:
    clean = {key: value for key, value in params.items() if value not in (None, "")}
    suffix = f"?{urlencode(clean)}" if clean else ""
    return RedirectResponse(
        url=f"/dieu-tra/du-lieu-lich-su/{dataset_id}{suffix}",
        status_code=303,
    )


def apply_excel_style(workbook: Workbook) -> None:
    thin = Side(style="thin", color="B7C7D6")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for index, header in enumerate(HISTORICAL_HEADERS, start=1):
            width = 18
            if header in {"Chủ hộ", "Họ và tên", "Địa chỉ hộ", "Địa chỉ thường trú", "Địa chỉ hiện tại", "Trường gần nhất"}:
                width = 28
            if header in {"Ghi chú hộ", "Ghi chú đối tượng", "Hoàn cảnh đặc biệt"}:
                width = 32
            worksheet.column_dimensions[get_column_letter(index)].width = width
        worksheet.row_dimensions[1].height = 42


def workbook_for_datasets(datasets: Iterable[HistoricalDataset], include_blank: bool = True) -> Workbook:
    workbook = Workbook()
    instruction = workbook.active
    instruction.title = "HUONG_DAN"
    instruction.append(["HƯỚNG DẪN CẬP NHẬT DỮ LIỆU LỊCH SỬ"])
    instruction.append(["1", "Giữ nguyên tên sheet DU_LIEU_LICH_SU và dòng tiêu đề."])
    instruction.append(["2", "Mỗi dòng là một đối tượng; thông tin hộ được lặp lại cho các thành viên cùng hộ."])
    instruction.append(["3", "Mã hộ và Mã đối tượng nên được giữ ổn định qua các lần cập nhật."])
    instruction.append(["4", "Các cột Có/Không có thể nhập: Có, Không hoặc để trống."])
    instruction.append(["5", "Dữ liệu khóa làm nền không thể nhập lại cho đến khi Sở mở khóa."])
    instruction.column_dimensions["A"].width = 12
    instruction.column_dimensions["B"].width = 110
    instruction["A1"].font = Font(bold=True, size=14)

    worksheet = workbook.create_sheet("DU_LIEU_LICH_SU")
    worksheet.append(HISTORICAL_HEADERS)

    wrote_row = False
    for dataset in datasets:
        households = sorted(dataset.households, key=lambda item: (item.hamlet_name or "", item.head_name, item.household_code))
        for household in households:
            people = sorted(household.people, key=lambda item: (item.date_of_birth or date.max, item.full_name))
            if not people:
                worksheet.append([
                    dataset.commune.code,
                    dataset.commune.name,
                    dataset.school_year.code,
                    household.household_code,
                    household.head_name,
                    household.hamlet_name or "",
                    household.address,
                    household.phone or "",
                    household.notes or "",
                ] + [""] * (len(HISTORICAL_HEADERS) - 9))
                wrote_row = True
                continue
            for person in people:
                worksheet.append([
                    dataset.commune.code,
                    dataset.commune.name,
                    dataset.school_year.code,
                    household.household_code,
                    household.head_name,
                    household.hamlet_name or "",
                    household.address,
                    household.phone or "",
                    household.notes or "",
                    person.person_code,
                    person.full_name,
                    person.date_of_birth.strftime("%d/%m/%Y") if person.date_of_birth else "",
                    person.gender or "",
                    person.ethnic_group or "",
                    person.birth_place or "",
                    person.father_name or "",
                    person.mother_name or "",
                    person.contact_phone or "",
                    person.relationship_to_head or "",
                    person.personal_id or "",
                    person.ministry_student_code or "",
                    person.permanent_address or "",
                    person.current_address or "",
                    RESIDENCY_STATUS_LABELS.get(person.residency_status, person.residency_status),
                    LEARNING_STATUS_LABELS.get(person.learning_status, person.learning_status),
                    person.school_name_reported or "",
                    person.class_name_reported or "",
                    bool_to_excel(person.completed_preschool_5),
                    bool_to_excel(person.attends_required_days),
                    bool_to_excel(person.attends_regularly),
                    bool_to_excel(person.prepared_vietnamese),
                    bool_to_excel(person.weight_monitored),
                    bool_to_excel(person.underweight),
                    bool_to_excel(person.height_monitored),
                    bool_to_excel(person.stunted),
                    person.special_circumstances or "",
                    person.notes or "",
                ])
                wrote_row = True

    if include_blank and not wrote_row:
        dataset_list = list(datasets)
        if dataset_list:
            first = dataset_list[0]
            worksheet.append([first.commune.code, first.commune.name, first.school_year.code] + [""] * (len(HISTORICAL_HEADERS) - 3))

    apply_excel_style(workbook)
    return workbook


def load_rows(upload: UploadFile) -> list[dict[str, Any]]:
    upload.file.seek(0)
    workbook = load_workbook(upload.file, data_only=True)
    if "DU_LIEU_LICH_SU" not in workbook.sheetnames:
        raise ValueError("missing_sheet")
    worksheet = workbook["DU_LIEU_LICH_SU"]
    normalized_expected = {normalize_header(header): key for header, key in HEADER_KEYS.items()}
    column_map: dict[int, str] = {}
    for column_index, cell in enumerate(worksheet[1], start=1):
        key = normalized_expected.get(normalize_header(cell.value))
        if key:
            column_map[column_index] = key
    required_keys = {"commune_code", "school_year_code", "household_code", "head_name", "household_address", "person_code", "full_name"}
    if not required_keys.issubset(set(column_map.values())):
        raise ValueError("missing_headers")
    rows: list[dict[str, Any]] = []
    identity_keys = {
        "commune_code",
        "commune_name",
        "school_year_code",
    }
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        data = {key: row[index - 1] for index, key in column_map.items()}
        if not any(normalize_text(value) for value in data.values()):
            continue

        # Các mẫu Excel có thể chứa sẵn dòng xã/phường và năm học để
        # hướng dẫn nhập. Nếu chưa có thông tin hộ hoặc đối tượng thì bỏ qua.
        content_values = [
            value
            for key, value in data.items()
            if key not in identity_keys
        ]
        if not any(normalize_text(value) for value in content_values):
            continue

        rows.append(data)
    return rows


def import_rows_into_dataset(
    db: Session,
    dataset: HistoricalDataset,
    rows: Iterable[dict[str, Any]],
) -> dict[str, int]:
    result = {
        "inserted_households": 0,
        "updated_households": 0,
        "inserted_people": 0,
        "updated_people": 0,
    }
    household_cache: dict[str, HistoricalHousehold] = {}
    person_cache: dict[str, HistoricalPerson] = {}
    inserted_household_codes: set[str] = set()
    updated_household_codes: set[str] = set()
    inserted_person_codes: set[str] = set()
    updated_person_codes: set[str] = set()

    existing_households = db.scalars(
        select(HistoricalHousehold).where(HistoricalHousehold.dataset_id == dataset.id)
    ).all()
    for item in existing_households:
        household_cache[item.household_code] = item

    existing_people = db.scalars(
        select(HistoricalPerson).where(HistoricalPerson.dataset_id == dataset.id)
    ).all()
    for item in existing_people:
        person_cache[item.person_code] = item

    for data in rows:
        commune_code = normalize_text(data.get("commune_code"))
        year_code = normalize_text(data.get("school_year_code"))
        if commune_code != dataset.commune.code or year_code != dataset.school_year.code:
            raise PermissionError("wrong_scope")

        head_name = normalize_text(data.get("head_name"))
        address = normalize_text(data.get("household_address"))
        full_name = normalize_text(data.get("full_name"))
        if not head_name or not address or not full_name:
            raise ValueError("missing_required")

        household_code = normalize_text(data.get("household_code"))
        if not household_code:
            household_code = next_code(db, dataset, HistoricalHousehold)
        household = household_cache.get(household_code)
        if household is None:
            household = HistoricalHousehold(
                dataset_id=dataset.id,
                household_code=household_code,
                head_name=head_name,
                hamlet_name=normalize_text(data.get("hamlet_name")) or None,
                address=address,
                phone=normalize_text(data.get("phone")) or None,
                notes=normalize_text(data.get("household_notes")) or None,
                is_active=True,
            )
            db.add(household)
            db.flush()
            household_cache[household_code] = household
            inserted_household_codes.add(household_code)
        else:
            household.head_name = head_name
            household.hamlet_name = normalize_text(data.get("hamlet_name")) or None
            household.address = address
            household.phone = normalize_text(data.get("phone")) or None
            household.notes = normalize_text(data.get("household_notes")) or None
            household.is_active = True
            if household_code not in inserted_household_codes:
                updated_household_codes.add(household_code)

        person_code = normalize_text(data.get("person_code"))
        if not person_code:
            person_code = next_code(db, dataset, HistoricalPerson)
        person = person_cache.get(person_code)
        values = {
            "historical_household_id": household.id,
            "full_name": full_name,
            "date_of_birth": parse_date(data.get("date_of_birth")),
            "gender": normalize_text(data.get("gender")) or None,
            "ethnic_group": normalize_text(data.get("ethnic_group")) or None,
            "birth_place": normalize_text(data.get("birth_place")) or None,
            "father_name": normalize_text(data.get("father_name")) or None,
            "mother_name": normalize_text(data.get("mother_name")) or None,
            "contact_phone": normalize_text(data.get("contact_phone")) or None,
            "relationship_to_head": normalize_text(data.get("relationship_to_head")) or None,
            "personal_id": normalize_text(data.get("personal_id")) or None,
            "ministry_student_code": normalize_text(data.get("ministry_student_code")) or None,
            "permanent_address": normalize_text(data.get("permanent_address")) or None,
            "current_address": normalize_text(data.get("current_address")) or None,
            "residency_status": normalize_status(data.get("residency_status"), RESIDENCY_STATUS_LABELS, "CHUA_XAC_DINH"),
            "learning_status": normalize_status(data.get("learning_status"), LEARNING_STATUS_LABELS, "CHUA_XAC_DINH"),
            "school_name_reported": normalize_text(data.get("school_name_reported")) or None,
            "class_name_reported": normalize_text(data.get("class_name_reported")) or None,
            "special_circumstances": normalize_text(data.get("special_circumstances")) or None,
            "notes": normalize_text(data.get("person_notes")) or None,
            "is_active": True,
        }
        for field_name in BOOLEAN_FIELDS:
            values[field_name] = parse_bool(data.get(field_name))

        if person is None:
            person = HistoricalPerson(
                dataset_id=dataset.id,
                person_code=person_code,
                **values,
            )
            db.add(person)
            db.flush()
            person_cache[person_code] = person
            inserted_person_codes.add(person_code)
        else:
            for field_name, value in values.items():
                setattr(person, field_name, value)
            if person_code not in inserted_person_codes:
                updated_person_codes.add(person_code)

    result["inserted_households"] = len(inserted_household_codes)
    result["updated_households"] = len(updated_household_codes)
    result["inserted_people"] = len(inserted_person_codes)
    result["updated_people"] = len(updated_person_codes)
    dataset.updated_at = datetime.now()
    return result


@router.get("", response_class=HTMLResponse)
def historical_center(
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    status_filter: str = "",
    q: str = "",
    status: str | None = None,
    error: str | None = None,
    created_total: int | None = None,
    imported_datasets: int | None = None,
    imported_households: int | None = None,
    imported_people: int | None = None,
    db: Session = Depends(get_db),
):
    if not can_view_center(request):
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    scope_id = commune_scope_id(request)
    try:
        school_year_value = int(school_year_id) if normalize_text(school_year_id) else None
    except (TypeError, ValueError):
        school_year_value = None
    try:
        commune_value = int(commune_id) if normalize_text(commune_id) else None
    except (TypeError, ValueError):
        commune_value = None

    year_statement = select(SchoolYear).where(historical_year_filter()).order_by(SchoolYear.code.desc())
    school_years = list(db.scalars(year_statement).all())

    commune_statement = select(Commune).where(Commune.is_active.is_(True)).order_by(Commune.name)
    if scope_id is not None:
        commune_statement = commune_statement.where(Commune.id == scope_id)
    communes = list(db.scalars(commune_statement).all())

    statement = (
        select(HistoricalDataset)
        .options(
            selectinload(HistoricalDataset.school_year),
            selectinload(HistoricalDataset.commune),
        )
        .join(HistoricalDataset.school_year)
        .join(HistoricalDataset.commune)
        .where(historical_year_filter())
        .order_by(SchoolYear.code.desc(), Commune.name)
    )
    if scope_id is not None:
        statement = statement.where(HistoricalDataset.commune_id == scope_id)
    if school_year_value is not None:
        statement = statement.where(HistoricalDataset.school_year_id == school_year_value)
    if commune_value is not None and scope_id is None:
        statement = statement.where(HistoricalDataset.commune_id == commune_value)
    if status_filter in DATASET_STATUS_LABELS:
        statement = statement.where(HistoricalDataset.status == status_filter)
    query_text = normalize_text(q)
    if query_text:
        pattern = f"%{query_text}%"
        statement = statement.where(
            or_(
                Commune.name.ilike(pattern),
                Commune.code.ilike(pattern),
                SchoolYear.code.ilike(pattern),
            )
        )

    datasets = list(db.scalars(statement).all())
    rows: list[dict[str, Any]] = []
    for dataset in datasets:
        household_total, person_total = dataset_counts(db, dataset.id)
        rows.append({
            "dataset": dataset,
            "household_total": household_total,
            "person_total": person_total,
        })

    all_scope_statement = select(HistoricalDataset)
    if scope_id is not None:
        all_scope_statement = all_scope_statement.where(HistoricalDataset.commune_id == scope_id)
    all_scope = list(db.scalars(all_scope_statement).all())
    summary = {
        "dataset_total": len(all_scope),
        "editing_total": sum(1 for item in all_scope if item.status == "DANG_CAP_NHAT"),
        "review_total": sum(1 for item in all_scope if item.status == "DA_GUI_KIEM_TRA"),
        "locked_total": sum(1 for item in all_scope if item.status == "DA_KHOA_LAM_NEN"),
        "household_total": sum(dataset_counts(db, item.id)[0] for item in all_scope),
        "person_total": sum(dataset_counts(db, item.id)[1] for item in all_scope),
    }

    selected_year_id = school_year_value or (school_years[0].id if school_years else None)
    selected_commune_id = commune_value or (scope_id if scope_id not in (None, -1) else None)

    return templates.TemplateResponse(
        request=request,
        name="surveys/historical_center.html",
        context={
            "nguoi_dung": current_user(request),
            "school_years": school_years,
            "communes": communes,
            "rows": rows,
            "summary": summary,
            "status_labels": DATASET_STATUS_LABELS,
            "status_classes": DATASET_STATUS_CLASSES,
            "selected_year_id": selected_year_id,
            "selected_commune_id": selected_commune_id,
            "status_filter": status_filter,
            "q": query_text,
            "can_update": can_update_history(request),
            "can_manage_province": can_manage_province(request),
            "is_commune": role_code(request) == COMMUNE_ROLE_CODE,
            "message": STATUS_MESSAGES.get(status or ""),
            "error_message": ERROR_MESSAGES.get(error or ""),
            "created_total": created_total,
            "imported_datasets": imported_datasets,
            "imported_households": imported_households,
            "imported_people": imported_people,
        },
    )


@router.post("/tao")
def create_dataset(
    request: Request,
    school_year_id: int = Form(...),
    commune_id: int = Form(...),
    notes: str = Form(""),
    db: Session = Depends(get_db),
):
    if not can_update_history(request):
        return redirect_center(error="forbidden")

    scope_id = commune_scope_id(request)
    if scope_id is not None and commune_id != scope_id:
        return redirect_center(error="forbidden")

    school_year = db.get(SchoolYear, school_year_id)
    commune = db.get(Commune, commune_id)
    if (
        school_year is None
        or commune is None
        or not commune.is_active
        or not is_historical_year_code(school_year.code)
    ):
        return redirect_center(error="wrong_scope")

    existing = db.scalar(
        select(HistoricalDataset).where(
            HistoricalDataset.school_year_id == school_year_id,
            HistoricalDataset.commune_id == commune_id,
        )
    )
    if existing is not None:
        return redirect_dataset(existing.id)

    dataset = HistoricalDataset(
        school_year_id=school_year_id,
        commune_id=commune_id,
        status="DANG_CAP_NHAT",
        reconciliation_status="CHUA_DOI_CHIEU",
        notes=normalize_text(notes) or None,
        created_by_user_id=current_user(request).get("id"),
    )
    db.add(dataset)
    db.flush()
    add_log(db, request, dataset, "TAO_BO_DU_LIEU", notes=dataset.notes)
    db.commit()
    return redirect_dataset(dataset.id, status="dataset_created")


@router.post("/tao-toan-tinh")
def create_province_datasets(
    request: Request,
    school_year_id: int = Form(...),
    db: Session = Depends(get_db),
):
    if not can_manage_province(request):
        return redirect_center(error="forbidden")

    school_year = db.get(SchoolYear, school_year_id)
    if school_year is None or not is_historical_year_code(school_year.code):
        return redirect_center(error="wrong_scope")

    communes = list(
        db.scalars(
            select(Commune).where(Commune.is_active.is_(True)).order_by(Commune.name)
        ).all()
    )
    existing_ids = set(
        db.scalars(
            select(HistoricalDataset.commune_id).where(
                HistoricalDataset.school_year_id == school_year_id
            )
        ).all()
    )
    created_total = 0
    for commune in communes:
        if commune.id in existing_ids:
            continue
        dataset = HistoricalDataset(
            school_year_id=school_year_id,
            commune_id=commune.id,
            status="DANG_CAP_NHAT",
            reconciliation_status="CHUA_DOI_CHIEU",
            notes=f"Khởi tạo bộ dữ liệu lịch sử năm học {school_year.code}.",
            created_by_user_id=current_user(request).get("id"),
        )
        db.add(dataset)
        db.flush()
        add_log(db, request, dataset, "TAO_BO_DU_LIEU_TOAN_TINH")
        created_total += 1
    db.commit()
    return redirect_center(
        school_year_id=school_year_id,
        status="datasets_created",
        created_total=created_total,
    )


@router.get("/mau-excel-toan-tinh")
def province_template(
    request: Request,
    school_year_id: int,
    db: Session = Depends(get_db),
):
    if not can_manage_province(request):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    school_year = db.get(SchoolYear, school_year_id)
    if school_year is None or not is_historical_year_code(school_year.code):
        return redirect_center(error="wrong_scope")

    workbook = Workbook()
    instruction = workbook.active
    instruction.title = "HUONG_DAN"
    instruction.append(["MẪU NHẬP DỮ LIỆU LỊCH SỬ TOÀN TỈNH"])
    instruction.append(["Năm học", school_year.code])
    instruction.append(["Lưu ý", "Mỗi dòng là một đối tượng. Mã xã/phường phải đúng danh mục hệ thống."])
    instruction.column_dimensions["A"].width = 18
    instruction.column_dimensions["B"].width = 100
    instruction["A1"].font = Font(bold=True, size=14)

    worksheet = workbook.create_sheet("DU_LIEU_LICH_SU")
    worksheet.append(HISTORICAL_HEADERS)
    communes = list(
        db.scalars(select(Commune).where(Commune.is_active.is_(True)).order_by(Commune.name)).all()
    )
    for commune in communes:
        worksheet.append([
            commune.code,
            commune.name,
            school_year.code,
        ] + [""] * (len(HISTORICAL_HEADERS) - 3))
    apply_excel_style(workbook)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"mau_du_lieu_lich_su_toan_tinh_{school_year.code.replace('-', '')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/nhap-excel-toan-tinh")
def import_province_excel(
    request: Request,
    school_year_id: int = Form(...),
    excel_file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    if not can_manage_province(request):
        return redirect_center(error="forbidden")
    if not excel_file.filename:
        return redirect_center(error="missing_file")

    school_year = db.get(SchoolYear, school_year_id)
    if school_year is None or not is_historical_year_code(school_year.code):
        return redirect_center(error="wrong_scope")

    try:
        rows = load_rows(excel_file)
    except Exception:
        return redirect_center(error="invalid_file")

    communes = {
        item.code: item
        for item in db.scalars(select(Commune).where(Commune.is_active.is_(True))).all()
    }
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        commune_code = normalize_text(row.get("commune_code"))
        year_code = normalize_text(row.get("school_year_code"))
        if commune_code not in communes or year_code != school_year.code:
            return redirect_center(error="wrong_scope")
        grouped.setdefault(commune_code, []).append(row)

    backup_database("before_historical_province_import")
    imported_datasets = 0
    imported_households = 0
    imported_people = 0
    try:
        for commune_code, commune_rows in grouped.items():
            commune = communes[commune_code]
            dataset = db.scalar(
                select(HistoricalDataset)
                .options(
                    selectinload(HistoricalDataset.school_year),
                    selectinload(HistoricalDataset.commune),
                )
                .where(
                    HistoricalDataset.school_year_id == school_year.id,
                    HistoricalDataset.commune_id == commune.id,
                )
            )
            if dataset is None:
                dataset = HistoricalDataset(
                    school_year_id=school_year.id,
                    commune_id=commune.id,
                    status="DANG_CAP_NHAT",
                    reconciliation_status="CHUA_DOI_CHIEU",
                    source_file_name=excel_file.filename,
                    created_by_user_id=current_user(request).get("id"),
                )
                db.add(dataset)
                db.flush()
                dataset.school_year = school_year
                dataset.commune = commune
            if dataset.status != "DANG_CAP_NHAT":
                raise PermissionError("dataset_locked")
            result = import_rows_into_dataset(db, dataset, commune_rows)
            dataset.source_file_name = excel_file.filename
            add_log(
                db,
                request,
                dataset,
                "NHAP_EXCEL_TOAN_TINH",
                file_name=excel_file.filename,
                inserted_households=result["inserted_households"],
                updated_households=result["updated_households"],
                inserted_people=result["inserted_people"],
                updated_people=result["updated_people"],
            )
            imported_datasets += 1
            imported_households += result["inserted_households"]
            imported_people += result["inserted_people"]
        db.commit()
    except PermissionError:
        db.rollback()
        return redirect_center(error="dataset_locked")
    except (ValueError, IntegrityError):
        db.rollback()
        return redirect_center(error="invalid_file")

    return redirect_center(
        school_year_id=school_year_id,
        status="province_imported",
        imported_datasets=imported_datasets,
        imported_households=imported_households,
        imported_people=imported_people,
    )


@router.get("/mau-excel-xa")
def commune_template(
    request: Request,
    school_year_id: int,
    commune_id: int,
    db: Session = Depends(get_db),
):
    """Tải mẫu Excel riêng của một xã/phường.

    Nếu xã/phường đã có bộ dữ liệu, tệp tải về chứa dữ liệu hiện có để rà soát.
    Nếu chưa có bộ dữ liệu, hệ thống tạo mẫu trống đã điền sẵn mã xã và năm học.
    """
    if not can_update_history(request):
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    scope_id = commune_scope_id(request)
    if scope_id is not None and commune_id != scope_id:
        return redirect_center(error="forbidden")

    school_year = db.get(SchoolYear, school_year_id)
    commune = db.get(Commune, commune_id)
    if (
        school_year is None
        or commune is None
        or not commune.is_active
        or not is_historical_year_code(school_year.code)
    ):
        return redirect_center(error="wrong_scope")

    dataset = db.scalar(
        select(HistoricalDataset)
        .options(
            selectinload(HistoricalDataset.school_year),
            selectinload(HistoricalDataset.commune),
            selectinload(HistoricalDataset.households).selectinload(HistoricalHousehold.people),
        )
        .where(
            HistoricalDataset.school_year_id == school_year_id,
            HistoricalDataset.commune_id == commune_id,
        )
    )

    if dataset is not None:
        workbook = workbook_for_datasets([dataset])
    else:
        workbook = Workbook()
        instruction = workbook.active
        instruction.title = "HUONG_DAN"
        instruction.append(["MẪU CẬP NHẬT DỮ LIỆU LỊCH SỬ THEO XÃ/PHƯỜNG"])
        instruction.append(["Xã/phường", f"{commune.code} – {commune.name}"])
        instruction.append(["Năm học", school_year.code])
        instruction.append(["Lưu ý", "Mỗi dòng là một đối tượng. Không sửa mã xã/phường và năm học đã điền sẵn."])
        instruction.append(["Mã hộ / Mã đối tượng", "Có thể để trống khi thêm mới; hệ thống sẽ tự sinh mã."])
        instruction.column_dimensions["A"].width = 24
        instruction.column_dimensions["B"].width = 105
        instruction["A1"].font = Font(bold=True, size=14)

        worksheet = workbook.create_sheet("DU_LIEU_LICH_SU")
        worksheet.append(HISTORICAL_HEADERS)
        worksheet.append(
            [commune.code, commune.name, school_year.code]
            + [""] * (len(HISTORICAL_HEADERS) - 3)
        )
        apply_excel_style(workbook)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = (
        f"du_lieu_lich_su_{commune.code}_"
        f"{school_year.code.replace('-', '')}.xlsx"
    )
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/nhap-excel-xa")
def import_commune_excel(
    request: Request,
    school_year_id: int = Form(...),
    commune_id: int = Form(...),
    notes: str = Form(""),
    excel_file: UploadFile | None = File(None),
    db: Session = Depends(get_db),
):
    """Nhập trực tiếp Excel của một xã/phường ngay tại Trung tâm dữ liệu lịch sử."""
    if not can_update_history(request):
        return redirect_center(error="forbidden")

    scope_id = commune_scope_id(request)
    if scope_id is not None and commune_id != scope_id:
        return redirect_center(error="forbidden")

    school_year = db.get(SchoolYear, school_year_id)
    commune = db.get(Commune, commune_id)
    if (
        school_year is None
        or commune is None
        or not commune.is_active
        or not is_historical_year_code(school_year.code)
    ):
        return redirect_center(error="wrong_scope")
    if excel_file is None or not excel_file.filename:
        return redirect_center(
            school_year_id=school_year_id,
            commune_id=commune_id,
            error="missing_file",
        )

    try:
        rows = load_rows(excel_file)
    except Exception:
        return redirect_center(
            school_year_id=school_year_id,
            commune_id=commune_id,
            error="invalid_file",
        )

    dataset = db.scalar(
        select(HistoricalDataset)
        .options(
            selectinload(HistoricalDataset.school_year),
            selectinload(HistoricalDataset.commune),
        )
        .where(
            HistoricalDataset.school_year_id == school_year_id,
            HistoricalDataset.commune_id == commune_id,
        )
    )
    if dataset is not None and not dataset_editable(request, dataset):
        return redirect_dataset(dataset.id, error="dataset_locked")

    backup_database("before_historical_commune_import")
    try:
        if dataset is None:
            dataset = HistoricalDataset(
                school_year_id=school_year_id,
                commune_id=commune_id,
                status="DANG_CAP_NHAT",
                reconciliation_status="CHUA_DOI_CHIEU",
                notes=normalize_text(notes) or None,
                created_by_user_id=current_user(request).get("id"),
            )
            db.add(dataset)
            db.flush()
            # Gắn quan hệ để các hàm sinh mã và kiểm tra phạm vi dùng ngay trong giao dịch.
            dataset.school_year = school_year
            dataset.commune = commune
            add_log(db, request, dataset, "TAO_BO_DU_LIEU", notes=dataset.notes)

        result = import_rows_into_dataset(db, dataset, rows)
        dataset.source_file_name = excel_file.filename
        if normalize_text(notes):
            dataset.notes = normalize_text(notes)
        add_log(
            db,
            request,
            dataset,
            "NHAP_EXCEL",
            file_name=excel_file.filename,
            inserted_households=result["inserted_households"],
            updated_households=result["updated_households"],
            inserted_people=result["inserted_people"],
            updated_people=result["updated_people"],
            notes="Nhập trực tiếp theo xã/phường từ Trung tâm dữ liệu lịch sử.",
        )
        db.commit()
    except PermissionError:
        db.rollback()
        return redirect_center(
            school_year_id=school_year_id,
            commune_id=commune_id,
            error="wrong_scope",
        )
    except (ValueError, IntegrityError):
        db.rollback()
        return redirect_center(
            school_year_id=school_year_id,
            commune_id=commune_id,
            error="invalid_file",
        )

    return redirect_dataset(dataset.id, status="imported", **result)


@router.get("/{dataset_id}", response_class=HTMLResponse)
def dataset_detail(
    request: Request,
    dataset_id: int,
    q: str = "",
    status: str | None = None,
    error: str | None = None,
    inserted_households: int | None = None,
    updated_households: int | None = None,
    inserted_people: int | None = None,
    updated_people: int | None = None,
    db: Session = Depends(get_db),
):
    dataset = db.scalar(
        select(HistoricalDataset)
        .options(
            selectinload(HistoricalDataset.school_year),
            selectinload(HistoricalDataset.commune),
            selectinload(HistoricalDataset.logs),
        )
        .where(HistoricalDataset.id == dataset_id)
    )
    if dataset is None or not dataset_in_scope(request, dataset):
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    statement = (
        select(HistoricalHousehold)
        .options(selectinload(HistoricalHousehold.people))
        .where(HistoricalHousehold.dataset_id == dataset.id)
        .order_by(HistoricalHousehold.hamlet_name, HistoricalHousehold.head_name)
    )
    query_text = normalize_text(q)
    if query_text:
        pattern = f"%{query_text}%"
        statement = statement.where(
            or_(
                HistoricalHousehold.household_code.ilike(pattern),
                HistoricalHousehold.head_name.ilike(pattern),
                HistoricalHousehold.hamlet_name.ilike(pattern),
                HistoricalHousehold.address.ilike(pattern),
                HistoricalHousehold.people.any(
                    or_(
                        HistoricalPerson.person_code.ilike(pattern),
                        HistoricalPerson.full_name.ilike(pattern),
                        HistoricalPerson.personal_id.ilike(pattern),
                    )
                ),
            )
        )
    households = list(db.scalars(statement).unique().all())
    household_total, person_total = dataset_counts(db, dataset.id)
    logs = sorted(dataset.logs, key=lambda item: item.created_at, reverse=True)[:20]

    return templates.TemplateResponse(
        request=request,
        name="surveys/historical_dataset.html",
        context={
            "nguoi_dung": current_user(request),
            "dataset": dataset,
            "households": households,
            "household_total": household_total,
            "person_total": person_total,
            "logs": logs,
            "q": query_text,
            "can_edit": dataset_editable(request, dataset),
            "can_manage_province": can_manage_province(request),
            "is_commune": role_code(request) == COMMUNE_ROLE_CODE,
            "status_labels": DATASET_STATUS_LABELS,
            "status_classes": DATASET_STATUS_CLASSES,
            "residency_labels": RESIDENCY_STATUS_LABELS,
            "learning_labels": LEARNING_STATUS_LABELS,
            "log_action_labels": LOG_ACTION_LABELS,
            "message": STATUS_MESSAGES.get(status or ""),
            "error_message": ERROR_MESSAGES.get(error or ""),
            "inserted_households": inserted_households,
            "updated_households": updated_households,
            "inserted_people": inserted_people,
            "updated_people": updated_people,
        },
    )


@router.get("/{dataset_id}/xuat-excel")
def export_dataset_excel(
    request: Request,
    dataset_id: int,
    db: Session = Depends(get_db),
):
    dataset = db.scalar(
        select(HistoricalDataset)
        .options(
            selectinload(HistoricalDataset.school_year),
            selectinload(HistoricalDataset.commune),
            selectinload(HistoricalDataset.households).selectinload(HistoricalHousehold.people),
        )
        .where(HistoricalDataset.id == dataset_id)
    )
    if dataset is None or not dataset_in_scope(request, dataset):
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    workbook = workbook_for_datasets([dataset])
    add_log(db, request, dataset, "XUAT_EXCEL", file_name=None)
    db.commit()
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = (
        f"du_lieu_lich_su_{dataset.commune.code}_"
        f"{dataset.school_year.code.replace('-', '')}.xlsx"
    )
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/{dataset_id}/nhap-excel")
def import_dataset_excel(
    request: Request,
    dataset_id: int,
    excel_file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    dataset = db.scalar(
        select(HistoricalDataset)
        .options(
            selectinload(HistoricalDataset.school_year),
            selectinload(HistoricalDataset.commune),
        )
        .where(HistoricalDataset.id == dataset_id)
    )
    if dataset is None or not dataset_in_scope(request, dataset):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if not dataset_editable(request, dataset):
        return redirect_dataset(dataset_id, error="dataset_locked")
    if not excel_file.filename:
        return redirect_dataset(dataset_id, error="missing_file")

    try:
        rows = load_rows(excel_file)
    except Exception:
        return redirect_dataset(dataset_id, error="invalid_file")

    backup_database("before_historical_import")
    try:
        result = import_rows_into_dataset(db, dataset, rows)
        dataset.source_file_name = excel_file.filename
        add_log(
            db,
            request,
            dataset,
            "NHAP_EXCEL",
            file_name=excel_file.filename,
            inserted_households=result["inserted_households"],
            updated_households=result["updated_households"],
            inserted_people=result["inserted_people"],
            updated_people=result["updated_people"],
        )
        db.commit()
    except PermissionError:
        db.rollback()
        return redirect_dataset(dataset_id, error="wrong_scope")
    except (ValueError, IntegrityError):
        db.rollback()
        return redirect_dataset(dataset_id, error="invalid_file")

    return redirect_dataset(dataset_id, status="imported", **result)


def household_form_context(
    request: Request,
    dataset: HistoricalDataset,
    household: HistoricalHousehold | None,
    error_message: str | None = None,
) -> dict[str, Any]:
    return {
        "nguoi_dung": current_user(request),
        "dataset": dataset,
        "household": household,
        "error_message": error_message,
    }


@router.get("/{dataset_id}/ho/them", response_class=HTMLResponse)
def add_household_form(
    request: Request,
    dataset_id: int,
    db: Session = Depends(get_db),
):
    dataset = get_dataset(db, dataset_id)
    if dataset is None or not dataset_in_scope(request, dataset):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if not dataset_editable(request, dataset):
        return redirect_dataset(dataset_id, error="dataset_locked")
    return templates.TemplateResponse(
        request=request,
        name="surveys/historical_household_form.html",
        context=household_form_context(request, dataset, None),
    )


@router.post("/{dataset_id}/ho/them")
def add_household(
    request: Request,
    dataset_id: int,
    household_code: str = Form(""),
    head_name: str = Form(...),
    hamlet_name: str = Form(""),
    address: str = Form(...),
    phone: str = Form(""),
    notes: str = Form(""),
    db: Session = Depends(get_db),
):
    dataset = get_dataset(db, dataset_id)
    if dataset is None or not dataset_in_scope(request, dataset):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if not dataset_editable(request, dataset):
        return redirect_dataset(dataset_id, error="dataset_locked")
    if not normalize_text(head_name) or not normalize_text(address):
        return templates.TemplateResponse(
            request=request,
            name="surveys/historical_household_form.html",
            context=household_form_context(
                request,
                dataset,
                None,
                "Cần nhập Chủ hộ và Địa chỉ hộ.",
            ),
            status_code=400,
        )

    code_value = normalize_text(household_code) or next_code(db, dataset, HistoricalHousehold)
    household = HistoricalHousehold(
        dataset_id=dataset.id,
        household_code=code_value,
        head_name=normalize_text(head_name),
        hamlet_name=normalize_text(hamlet_name) or None,
        address=normalize_text(address),
        phone=normalize_text(phone) or None,
        notes=normalize_text(notes) or None,
        is_active=True,
    )
    db.add(household)
    try:
        db.flush()
        add_log(db, request, dataset, "THEM_HO", inserted_households=1, notes=f"Mã hộ: {code_value}")
        db.commit()
    except IntegrityError:
        db.rollback()
        return redirect_dataset(dataset_id, error="duplicate")
    return redirect_dataset(dataset_id, status="household_created")


@router.get("/{dataset_id}/ho/{household_id}/sua", response_class=HTMLResponse)
def edit_household_form(
    request: Request,
    dataset_id: int,
    household_id: int,
    db: Session = Depends(get_db),
):
    dataset = get_dataset(db, dataset_id)
    household = db.get(HistoricalHousehold, household_id)
    if (
        dataset is None
        or household is None
        or household.dataset_id != dataset.id
        or not dataset_in_scope(request, dataset)
    ):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if not dataset_editable(request, dataset):
        return redirect_dataset(dataset_id, error="dataset_locked")
    return templates.TemplateResponse(
        request=request,
        name="surveys/historical_household_form.html",
        context=household_form_context(request, dataset, household),
    )


@router.post("/{dataset_id}/ho/{household_id}/sua")
def edit_household(
    request: Request,
    dataset_id: int,
    household_id: int,
    household_code: str = Form(...),
    head_name: str = Form(...),
    hamlet_name: str = Form(""),
    address: str = Form(...),
    phone: str = Form(""),
    notes: str = Form(""),
    db: Session = Depends(get_db),
):
    dataset = get_dataset(db, dataset_id)
    household = db.get(HistoricalHousehold, household_id)
    if (
        dataset is None
        or household is None
        or household.dataset_id != dataset.id
        or not dataset_in_scope(request, dataset)
    ):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if not dataset_editable(request, dataset):
        return redirect_dataset(dataset_id, error="dataset_locked")
    code_value = normalize_text(household_code)
    head_value = normalize_text(head_name)
    address_value = normalize_text(address)
    if not code_value or not head_value or not address_value:
        return templates.TemplateResponse(
            request=request,
            name="surveys/historical_household_form.html",
            context=household_form_context(
                request,
                dataset,
                household,
                "Cần nhập Mã hộ, Chủ hộ và Địa chỉ hộ.",
            ),
            status_code=400,
        )
    household.household_code = code_value
    household.head_name = head_value
    household.hamlet_name = normalize_text(hamlet_name) or None
    household.address = address_value
    household.phone = normalize_text(phone) or None
    household.notes = normalize_text(notes) or None
    try:
        add_log(db, request, dataset, "SUA_HO", updated_households=1, notes=f"Mã hộ: {household.household_code}")
        db.commit()
    except IntegrityError:
        db.rollback()
        return redirect_dataset(dataset_id, error="duplicate")
    return redirect_dataset(dataset_id, status="household_updated")


def person_form_context(
    request: Request,
    dataset: HistoricalDataset,
    household: HistoricalHousehold,
    person: HistoricalPerson | None,
    error_message: str | None = None,
) -> dict[str, Any]:
    return {
        "nguoi_dung": current_user(request),
        "dataset": dataset,
        "household": household,
        "person": person,
        "error_message": error_message,
        "residency_labels": RESIDENCY_STATUS_LABELS,
        "learning_labels": LEARNING_STATUS_LABELS,
    }


def populate_person_from_form(person: HistoricalPerson, form: dict[str, Any]) -> None:
    person.person_code = normalize_text(form.get("person_code"))
    person.full_name = normalize_text(form.get("full_name"))
    person.date_of_birth = parse_date(form.get("date_of_birth"))
    person.gender = normalize_text(form.get("gender")) or None
    person.ethnic_group = normalize_text(form.get("ethnic_group")) or None
    person.birth_place = normalize_text(form.get("birth_place")) or None
    person.father_name = normalize_text(form.get("father_name")) or None
    person.mother_name = normalize_text(form.get("mother_name")) or None
    person.contact_phone = normalize_text(form.get("contact_phone")) or None
    person.relationship_to_head = normalize_text(form.get("relationship_to_head")) or None
    person.personal_id = normalize_text(form.get("personal_id")) or None
    person.ministry_student_code = normalize_text(form.get("ministry_student_code")) or None
    person.permanent_address = normalize_text(form.get("permanent_address")) or None
    person.current_address = normalize_text(form.get("current_address")) or None
    person.residency_status = normalize_status(form.get("residency_status"), RESIDENCY_STATUS_LABELS, "CHUA_XAC_DINH")
    person.learning_status = normalize_status(form.get("learning_status"), LEARNING_STATUS_LABELS, "CHUA_XAC_DINH")
    person.school_name_reported = normalize_text(form.get("school_name_reported")) or None
    person.class_name_reported = normalize_text(form.get("class_name_reported")) or None
    person.completed_preschool_5 = parse_bool(form.get("completed_preschool_5"))
    person.attends_required_days = parse_bool(form.get("attends_required_days"))
    person.attends_regularly = parse_bool(form.get("attends_regularly"))
    person.prepared_vietnamese = parse_bool(form.get("prepared_vietnamese"))
    person.weight_monitored = parse_bool(form.get("weight_monitored"))
    person.underweight = parse_bool(form.get("underweight"))
    person.height_monitored = parse_bool(form.get("height_monitored"))
    person.stunted = parse_bool(form.get("stunted"))
    person.special_circumstances = normalize_text(form.get("special_circumstances")) or None
    person.notes = normalize_text(form.get("notes")) or None
    person.is_active = True


@router.get("/{dataset_id}/ho/{household_id}/doi-tuong/them", response_class=HTMLResponse)
def add_person_form(
    request: Request,
    dataset_id: int,
    household_id: int,
    db: Session = Depends(get_db),
):
    dataset = get_dataset(db, dataset_id)
    household = db.get(HistoricalHousehold, household_id)
    if (
        dataset is None
        or household is None
        or household.dataset_id != dataset.id
        or not dataset_in_scope(request, dataset)
    ):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if not dataset_editable(request, dataset):
        return redirect_dataset(dataset_id, error="dataset_locked")
    return templates.TemplateResponse(
        request=request,
        name="surveys/historical_person_form.html",
        context=person_form_context(request, dataset, household, None),
    )


@router.post("/{dataset_id}/ho/{household_id}/doi-tuong/them")
async def add_person(
    request: Request,
    dataset_id: int,
    household_id: int,
    db: Session = Depends(get_db),
):
    dataset = get_dataset(db, dataset_id)
    household = db.get(HistoricalHousehold, household_id)
    if (
        dataset is None
        or household is None
        or household.dataset_id != dataset.id
        or not dataset_in_scope(request, dataset)
    ):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if not dataset_editable(request, dataset):
        return redirect_dataset(dataset_id, error="dataset_locked")
    form = dict(await request.form())
    full_name = normalize_text(form.get("full_name"))
    if not full_name:
        return templates.TemplateResponse(
            request=request,
            name="surveys/historical_person_form.html",
            context=person_form_context(request, dataset, household, None, "Cần nhập Họ và tên đối tượng."),
            status_code=400,
        )
    person = HistoricalPerson(
        dataset_id=dataset.id,
        historical_household_id=household.id,
        person_code=normalize_text(form.get("person_code")) or next_code(db, dataset, HistoricalPerson),
        full_name=full_name,
    )
    populate_person_from_form(person, {**form, "person_code": person.person_code})
    db.add(person)
    try:
        db.flush()
        add_log(db, request, dataset, "THEM_DOI_TUONG", inserted_people=1, notes=f"Mã đối tượng: {person.person_code}")
        db.commit()
    except IntegrityError:
        db.rollback()
        return redirect_dataset(dataset_id, error="duplicate")
    return redirect_dataset(dataset_id, status="person_created")


@router.get("/{dataset_id}/doi-tuong/{person_id}/sua", response_class=HTMLResponse)
def edit_person_form(
    request: Request,
    dataset_id: int,
    person_id: int,
    db: Session = Depends(get_db),
):
    dataset = get_dataset(db, dataset_id)
    person = db.get(HistoricalPerson, person_id)
    household = db.get(HistoricalHousehold, person.historical_household_id) if person else None
    if (
        dataset is None
        or person is None
        or household is None
        or person.dataset_id != dataset.id
        or not dataset_in_scope(request, dataset)
    ):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if not dataset_editable(request, dataset):
        return redirect_dataset(dataset_id, error="dataset_locked")
    return templates.TemplateResponse(
        request=request,
        name="surveys/historical_person_form.html",
        context=person_form_context(request, dataset, household, person),
    )


@router.post("/{dataset_id}/doi-tuong/{person_id}/sua")
async def edit_person(
    request: Request,
    dataset_id: int,
    person_id: int,
    db: Session = Depends(get_db),
):
    dataset = get_dataset(db, dataset_id)
    person = db.get(HistoricalPerson, person_id)
    if dataset is None or person is None or person.dataset_id != dataset.id or not dataset_in_scope(request, dataset):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if not dataset_editable(request, dataset):
        return redirect_dataset(dataset_id, error="dataset_locked")
    form = dict(await request.form())
    if not normalize_text(form.get("person_code")) or not normalize_text(form.get("full_name")):
        household = db.get(HistoricalHousehold, person.historical_household_id)
        return templates.TemplateResponse(
            request=request,
            name="surveys/historical_person_form.html",
            context=person_form_context(
                request,
                dataset,
                household,
                person,
                "Cần nhập Mã đối tượng và Họ và tên.",
            ),
            status_code=400,
        )
    populate_person_from_form(person, form)
    try:
        add_log(db, request, dataset, "SUA_DOI_TUONG", updated_people=1, notes=f"Mã đối tượng: {person.person_code}")
        db.commit()
    except IntegrityError:
        db.rollback()
        return redirect_dataset(dataset_id, error="duplicate")
    return redirect_dataset(dataset_id, status="person_updated")


@router.post("/{dataset_id}/gui-kiem-tra")
def submit_dataset(
    request: Request,
    dataset_id: int,
    db: Session = Depends(get_db),
):
    dataset = get_dataset(db, dataset_id)
    if dataset is None or not dataset_in_scope(request, dataset):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if not can_update_history(request) or dataset.status != "DANG_CAP_NHAT":
        return redirect_dataset(dataset_id, error="dataset_locked")
    dataset.status = "DA_GUI_KIEM_TRA"
    dataset.submitted_at = datetime.now()
    dataset.submitted_by_user_id = current_user(request).get("id")
    add_log(db, request, dataset, "GUI_KIEM_TRA")
    db.commit()
    return redirect_dataset(dataset_id, status="submitted")


@router.post("/{dataset_id}/khoa")
def lock_dataset(
    request: Request,
    dataset_id: int,
    reason: str = Form(...),
    db: Session = Depends(get_db),
):
    dataset = get_dataset(db, dataset_id)
    if dataset is None or not dataset_in_scope(request, dataset):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if not can_manage_province(request):
        return redirect_dataset(dataset_id, error="forbidden")
    reason_value = normalize_text(reason)
    if not reason_value:
        return redirect_dataset(dataset_id, error="reason_required")
    dataset.status = "DA_KHOA_LAM_NEN"
    dataset.locked_at = datetime.now()
    dataset.locked_by_user_id = current_user(request).get("id")
    dataset.lock_reason = reason_value
    add_log(db, request, dataset, "KHOA_LAM_NEN", notes=reason_value)
    db.commit()
    return redirect_dataset(dataset_id, status="locked")


@router.post("/{dataset_id}/mo-khoa")
def unlock_dataset(
    request: Request,
    dataset_id: int,
    reason: str = Form(...),
    db: Session = Depends(get_db),
):
    dataset = get_dataset(db, dataset_id)
    if dataset is None or not dataset_in_scope(request, dataset):
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if not can_manage_province(request):
        return redirect_dataset(dataset_id, error="forbidden")
    reason_value = normalize_text(reason)
    if not reason_value:
        return redirect_dataset(dataset_id, error="reason_required")
    dataset.status = "DANG_CAP_NHAT"
    dataset.locked_at = None
    dataset.locked_by_user_id = None
    dataset.lock_reason = None
    dataset.submitted_at = None
    dataset.submitted_by_user_id = None
    dataset.reconciliation_status = "CHUA_DOI_CHIEU"
    add_log(db, request, dataset, "MO_KHOA", notes=reason_value)
    db.commit()
    return redirect_dataset(dataset_id, status="unlocked")
