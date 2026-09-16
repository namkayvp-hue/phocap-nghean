from __future__ import annotations

from datetime import datetime
from io import BytesIO
from math import ceil
from pathlib import Path
from typing import Any
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.comparison_models import StudentSurveyComparisonSignoff
from app.database import get_db
from app.models import Commune, SchoolYear
from app.permissions import is_province_reader
from app.routers.surveys import BATCH_STATUS_LABELS, lay_thong_tin_nguoi_dung
from app.survey_models import SurveyBatch


APP_DIR = Path(__file__).resolve().parent.parent
templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

router = APIRouter(
    prefix="/dieu-tra",
    tags=["Tổng hợp chốt đối chiếu học sinh"],
)

PAGE_SIZE_OPTIONS = (50, 100, 200)

STATE_LABELS = {
    "": "Tất cả trạng thái",
    "SIGNED": "Đã chốt",
    "REOPENED": "Đã mở lại",
    "UNSIGNED": "Chưa chốt",
    "MISSING_BATCH": "Chưa có đợt",
}

STATE_BADGES = {
    "SIGNED": "state-signed",
    "REOPENED": "state-reopened",
    "UNSIGNED": "state-unsigned",
    "MISSING_BATCH": "state-missing",
}

ACTION_LABELS = {
    "CHOT": "Đã chốt kết quả đối chiếu",
    "MO_LAI": "Đã mở lại để xử lý",
}


def _clean_text(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _parse_optional_int(value: Any) -> int | None:
    text = str(value or "").strip()
    if not text or not text.isdigit():
        return None
    return int(text)


def _format_datetime(value: datetime | None) -> str:
    return value.strftime("%d/%m/%Y %H:%M") if value is not None else "—"


def _load_school_years(db: Session) -> list[SchoolYear]:
    return list(
        db.scalars(
            select(SchoolYear).order_by(
                SchoolYear.code.desc(),
                SchoolYear.id.desc(),
            )
        ).all()
    )


def _default_year_id(
    school_years: list[SchoolYear],
    requested_year_id: int | None,
) -> int | None:
    valid_ids = {item.id for item in school_years}
    if requested_year_id in valid_ids:
        return requested_year_id
    return school_years[0].id if school_years else None


def _load_latest_signoffs(
    db: Session,
    batch_ids: list[int],
) -> tuple[
    dict[int, StudentSurveyComparisonSignoff],
    list[StudentSurveyComparisonSignoff],
]:
    if not batch_ids:
        return {}, []

    history = list(
        db.scalars(
            select(StudentSurveyComparisonSignoff)
            .where(
                StudentSurveyComparisonSignoff.survey_batch_id.in_(batch_ids)
            )
            .order_by(
                StudentSurveyComparisonSignoff.created_at.desc(),
                StudentSurveyComparisonSignoff.id.desc(),
            )
        ).all()
    )

    latest: dict[int, StudentSurveyComparisonSignoff] = {}
    for item in history:
        latest.setdefault(item.survey_batch_id, item)
    return latest, history


def _build_rows(
    db: Session,
    school_year_id: int | None,
) -> tuple[list[dict[str, Any]], list[StudentSurveyComparisonSignoff]]:
    communes = list(
        db.scalars(
            select(Commune)
            .where(Commune.is_active.is_(True))
            .order_by(Commune.name.asc(), Commune.id.asc())
        ).all()
    )

    if school_year_id is None:
        return [], []

    batches = list(
        db.scalars(
            select(SurveyBatch)
            .options(
                selectinload(SurveyBatch.commune),
                selectinload(SurveyBatch.school_year),
            )
            .where(SurveyBatch.school_year_id == school_year_id)
            .order_by(
                SurveyBatch.commune_id.asc(),
                SurveyBatch.id.desc(),
            )
        ).all()
    )

    batches_by_commune: dict[int, list[SurveyBatch]] = {}
    for batch in batches:
        batches_by_commune.setdefault(batch.commune_id, []).append(batch)

    batch_ids = [batch.id for batch in batches]
    latest_signoffs, history = _load_latest_signoffs(db, batch_ids)

    rows: list[dict[str, Any]] = []
    for commune in communes:
        commune_batches = batches_by_commune.get(commune.id, [])
        batch = commune_batches[0] if commune_batches else None
        latest = latest_signoffs.get(batch.id) if batch is not None else None

        if batch is None:
            state = "MISSING_BATCH"
        elif latest is None:
            state = "UNSIGNED"
        elif latest.action_code == "CHOT":
            state = "SIGNED"
        else:
            state = "REOPENED"

        rows.append(
            {
                "commune": commune,
                "batch": batch,
                "extra_batch_count": max(0, len(commune_batches) - 1),
                "state": state,
                "state_label": STATE_LABELS[state],
                "state_badge": STATE_BADGES[state],
                "latest_signoff": latest,
                "latest_action_label": (
                    ACTION_LABELS.get(latest.action_code, latest.action_code)
                    if latest is not None
                    else "Chưa có lịch sử chốt"
                ),
                "latest_time_label": _format_datetime(
                    latest.created_at if latest is not None else None
                ),
                "batch_status_label": (
                    BATCH_STATUS_LABELS.get(batch.status, batch.status)
                    if batch is not None
                    else "—"
                ),
                "total_rows": latest.total_rows if latest is not None else None,
                "matched_rows": latest.matched_rows if latest is not None else None,
                "issue_rows": latest.issue_rows if latest is not None else None,
                "resolved_rows": latest.resolved_rows if latest is not None else None,
                "pending_rows": latest.pending_rows if latest is not None else None,
            }
        )

    priority = {
        "REOPENED": 0,
        "UNSIGNED": 1,
        "MISSING_BATCH": 2,
        "SIGNED": 3,
    }
    rows.sort(
        key=lambda item: (
            priority.get(item["state"], 9),
            item["commune"].name.casefold(),
            item["commune"].id,
        )
    )
    return rows, history


def _apply_filters(
    rows: list[dict[str, Any]],
    *,
    state: str,
    q: str,
) -> list[dict[str, Any]]:
    query = _clean_text(q).casefold()
    result: list[dict[str, Any]] = []
    for row in rows:
        if state and row["state"] != state:
            continue
        if query:
            batch = row["batch"]
            search_text = " ".join(
                (
                    row["commune"].code,
                    row["commune"].name,
                    batch.code if batch is not None else "",
                    batch.name if batch is not None else "",
                    row["latest_signoff"].actor_name_snapshot
                    if row["latest_signoff"] is not None
                    else "",
                    row["latest_signoff"].note
                    if row["latest_signoff"] is not None
                    else "",
                )
            ).casefold()
            if query not in search_text:
                continue
        result.append(row)
    return result


def _build_summary(rows: list[dict[str, Any]]) -> dict[str, int]:
    summary = {
        "commune_total": len(rows),
        "batch_total": 0,
        "signed_total": 0,
        "reopened_total": 0,
        "unsigned_total": 0,
        "missing_batch_total": 0,
        "snapshot_rows": 0,
        "snapshot_pending": 0,
    }
    for row in rows:
        if row["batch"] is not None:
            summary["batch_total"] += 1
        if row["state"] == "SIGNED":
            summary["signed_total"] += 1
        elif row["state"] == "REOPENED":
            summary["reopened_total"] += 1
        elif row["state"] == "UNSIGNED":
            summary["unsigned_total"] += 1
        else:
            summary["missing_batch_total"] += 1

        if row["total_rows"] is not None:
            summary["snapshot_rows"] += int(row["total_rows"] or 0)
            summary["snapshot_pending"] += int(row["pending_rows"] or 0)
    return summary


def _page_url(
    *,
    school_year_id: int | None,
    state: str,
    q: str,
    page_size: int,
    page: int,
) -> str:
    params = {
        "school_year_id": school_year_id or "",
        "state": state,
        "q": q,
        "page_size": page_size,
        "page": page,
    }
    return "/dieu-tra/tong-hop-doi-chieu-hoc-sinh?" + urlencode(params)


def _forbidden_response(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="surveys/student_survey_province_dashboard.html",
        context={
            "nguoi_dung": lay_thong_tin_nguoi_dung(request),
            "error_message": (
                "Chức năng tổng hợp toàn tỉnh chỉ dành cho tài khoản "
                "quản trị hoặc phòng ban cấp Sở."
            ),
            "school_years": [],
            "selected_year": None,
            "rows": [],
            "summary": _build_summary([]),
        },
        status_code=403,
    )


@router.get(
    "/tong-hop-doi-chieu-hoc-sinh",
    response_class=HTMLResponse,
)
def province_comparison_dashboard(
    request: Request,
    school_year_id: str = "",
    state: str = "",
    q: str = "",
    page_size: int = 50,
    page: int = 1,
    db: Session = Depends(get_db),
):
    user = lay_thong_tin_nguoi_dung(request)
    if not is_province_reader(user.get("role_code")):
        return _forbidden_response(request)

    school_years = _load_school_years(db)
    selected_year_id = _default_year_id(
        school_years,
        _parse_optional_int(school_year_id),
    )
    selected_year = next(
        (item for item in school_years if item.id == selected_year_id),
        None,
    )

    normalized_state = _clean_text(state).upper()
    if normalized_state not in STATE_LABELS:
        normalized_state = ""
    q = _clean_text(q)[:120]
    if page_size not in PAGE_SIZE_OPTIONS:
        page_size = 50

    all_rows, _ = _build_rows(db, selected_year_id)
    summary = _build_summary(all_rows)
    filtered_rows = _apply_filters(
        all_rows,
        state=normalized_state,
        q=q,
    )

    total_records = len(filtered_rows)
    total_pages = max(1, ceil(total_records / page_size))
    page = max(1, min(page, total_pages))
    start_index = (page - 1) * page_size
    page_rows = filtered_rows[start_index:start_index + page_size]

    page_links = [
        {
            "number": number,
            "url": _page_url(
                school_year_id=selected_year_id,
                state=normalized_state,
                q=q,
                page_size=page_size,
                page=number,
            ),
        }
        for number in range(max(1, page - 2), min(total_pages, page + 2) + 1)
    ]

    export_params = urlencode(
        {
            "school_year_id": selected_year_id or "",
            "state": normalized_state,
            "q": q,
        }
    )

    return templates.TemplateResponse(
        request=request,
        name="surveys/student_survey_province_dashboard.html",
        context={
            "nguoi_dung": user,
            "error_message": None,
            "school_years": school_years,
            "selected_year": selected_year,
            "selected_year_id": selected_year_id,
            "state": normalized_state,
            "state_labels": STATE_LABELS,
            "q": q,
            "page_size": page_size,
            "page_size_options": PAGE_SIZE_OPTIONS,
            "summary": summary,
            "rows": page_rows,
            "total_records": total_records,
            "page": page,
            "total_pages": total_pages,
            "page_links": page_links,
            "previous_url": (
                _page_url(
                    school_year_id=selected_year_id,
                    state=normalized_state,
                    q=q,
                    page_size=page_size,
                    page=page - 1,
                )
                if page > 1
                else None
            ),
            "next_url": (
                _page_url(
                    school_year_id=selected_year_id,
                    state=normalized_state,
                    q=q,
                    page_size=page_size,
                    page=page + 1,
                )
                if page < total_pages
                else None
            ),
            "export_url": (
                "/dieu-tra/tong-hop-doi-chieu-hoc-sinh/xuat-excel?"
                + export_params
            ),
        },
    )


def _style_sheet(worksheet: Any) -> None:
    dark_fill = PatternFill("solid", fgColor="1F4E78")
    light_fill = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)
    thin = Side(style="thin", color="C9D4DF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in worksheet[1]:
        cell.fill = dark_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for cell in worksheet[2]:
        if cell.value is not None:
            cell.fill = light_fill
            cell.font = bold_font

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def _set_widths(worksheet: Any, widths: dict[int, float]) -> None:
    for index, width in widths.items():
        worksheet.column_dimensions[get_column_letter(index)].width = width


@router.get(
    "/tong-hop-doi-chieu-hoc-sinh/xuat-excel",
)
def export_province_comparison_dashboard(
    request: Request,
    school_year_id: str = "",
    state: str = "",
    q: str = "",
    db: Session = Depends(get_db),
):
    user = lay_thong_tin_nguoi_dung(request)
    if not is_province_reader(user.get("role_code")):
        return _forbidden_response(request)

    school_years = _load_school_years(db)
    selected_year_id = _default_year_id(
        school_years,
        _parse_optional_int(school_year_id),
    )
    selected_year = next(
        (item for item in school_years if item.id == selected_year_id),
        None,
    )
    normalized_state = _clean_text(state).upper()
    if normalized_state not in STATE_LABELS:
        normalized_state = ""
    q = _clean_text(q)[:120]

    all_rows, history = _build_rows(db, selected_year_id)
    rows = _apply_filters(all_rows, state=normalized_state, q=q)

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Tong hop dia ban"
    overview.append([
        "STT",
        "Mã xã",
        "Xã/phường",
        "Mã đợt",
        "Tên đợt",
        "Trạng thái đợt",
        "Khóa dữ liệu",
        "Trạng thái chốt",
        "Thao tác gần nhất",
        "Người thực hiện",
        "Vai trò",
        "Thời điểm",
        "Căn cứ / lý do",
        "Tổng dòng",
        "Khớp hoàn toàn",
        "Cần xử lý",
        "Đã ghi nhận",
        "Còn chờ",
        "Số đợt khác cùng xã/năm",
    ])

    for index, row in enumerate(rows, start=1):
        batch = row["batch"]
        signoff = row["latest_signoff"]
        overview.append([
            index,
            row["commune"].code,
            row["commune"].name,
            batch.code if batch is not None else "",
            batch.name if batch is not None else "",
            row["batch_status_label"],
            "Đã khóa" if batch is not None and batch.is_locked else "Chưa khóa",
            row["state_label"],
            row["latest_action_label"],
            signoff.actor_name_snapshot if signoff is not None else "",
            signoff.actor_role_snapshot if signoff is not None else "",
            row["latest_time_label"] if signoff is not None else "",
            signoff.note if signoff is not None else "",
            row["total_rows"] if row["total_rows"] is not None else "",
            row["matched_rows"] if row["matched_rows"] is not None else "",
            row["issue_rows"] if row["issue_rows"] is not None else "",
            row["resolved_rows"] if row["resolved_rows"] is not None else "",
            row["pending_rows"] if row["pending_rows"] is not None else "",
            row["extra_batch_count"],
        ])

    _style_sheet(overview)
    _set_widths(
        overview,
        {
            1: 7,
            2: 13,
            3: 28,
            4: 25,
            5: 42,
            6: 16,
            7: 13,
            8: 16,
            9: 28,
            10: 22,
            11: 22,
            12: 19,
            13: 45,
            14: 12,
            15: 14,
            16: 12,
            17: 14,
            18: 11,
            19: 18,
        },
    )

    history_sheet = workbook.create_sheet("Lich su chot")
    history_sheet.append([
        "STT",
        "Mã xã",
        "Xã/phường",
        "Mã đợt",
        "Thao tác",
        "Người thực hiện",
        "Vai trò",
        "Thời điểm",
        "Căn cứ / lý do",
        "Tổng dòng",
        "Khớp hoàn toàn",
        "Cần xử lý",
        "Đã ghi nhận",
        "Còn chờ",
    ])

    batch_by_id = {
        row["batch"].id: row
        for row in all_rows
        if row["batch"] is not None
    }
    history_index = 0
    for item in history:
        row = batch_by_id.get(item.survey_batch_id)
        if row is None:
            continue
        if row not in rows:
            continue
        history_index += 1
        history_sheet.append([
            history_index,
            row["commune"].code,
            row["commune"].name,
            row["batch"].code,
            ACTION_LABELS.get(item.action_code, item.action_code),
            item.actor_name_snapshot,
            item.actor_role_snapshot,
            _format_datetime(item.created_at),
            item.note,
            item.total_rows,
            item.matched_rows,
            item.issue_rows,
            item.resolved_rows,
            item.pending_rows,
        ])

    _style_sheet(history_sheet)
    _set_widths(
        history_sheet,
        {
            1: 7,
            2: 13,
            3: 28,
            4: 25,
            5: 30,
            6: 22,
            7: 22,
            8: 19,
            9: 50,
            10: 12,
            11: 14,
            12: 12,
            13: 14,
            14: 11,
        },
    )

    metadata = workbook.create_sheet("Thong tin bao cao")
    metadata.append(["Nội dung", "Giá trị"])
    metadata.append([
        "Năm học",
        selected_year.code if selected_year is not None else "Không xác định",
    ])
    metadata.append(["Trạng thái lọc", STATE_LABELS.get(normalized_state, "Tất cả")])
    metadata.append(["Từ khóa", q or "Không lọc"])
    metadata.append(["Người xuất", _clean_text(user.get("full_name"))])
    metadata.append(["Vai trò", _clean_text(user.get("role_name"))])
    metadata.append(["Thời điểm xuất", datetime.now().strftime("%d/%m/%Y %H:%M")])
    _style_sheet(metadata)
    _set_widths(metadata, {1: 24, 2: 55})

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    year_code = selected_year.code if selected_year is not None else "khong-xac-dinh"
    filename = f"tong_hop_chot_doi_chieu_{year_code}.xlsx"
    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"'
    }
    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers=headers,
    )
