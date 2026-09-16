from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Annotated, Any
from urllib.parse import quote
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import APIRouter, Depends, Form, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from app.database import get_db
from app.models import Commune, SchoolYear
from app.permissions import (
    COMMUNE_ROLE_CODE,
    DEPARTMENT_ROLE_CODE,
    is_admin_role,
    is_province_reader,
    normalize_role_code,
)
from app.survey_models import (
    Household,
    SurveyBatch,
    SurveyFileExchangeLog,
    SurveyForm,
    SurveyPerson,
)
from app.routers.surveys import (
    lay_dot_dieu_tra,
    xuat_excel_di_dieu_tra,
)


APP_DIR = Path(__file__).resolve().parent.parent

templates = Jinja2Templates(
    directory=str(APP_DIR / "templates")
)

router = APIRouter(
    prefix="/dieu-tra",
    tags=["Trung tâm phiếu điều tra"],
)


EXCHANGE_STATUS_LABELS = {
    "": "Tất cả trạng thái",
    "CHUA_CO_DOT": "Chưa có đợt",
    "CHUA_PHAT_HANH": "Chưa phát hành",
    "CHO_TIEP_NHAN": "Đã phát hành, chờ tiếp nhận",
    "DA_TIEP_NHAN": "Đã tiếp nhận file cập nhật",
}

ACTION_LABELS = {
    "XUAT_DON": "Xuất file điều tra",
    "XUAT_HANG_LOAT": "Xuất hàng loạt",
    "NHAP_EXCEL": "Tiếp nhận file cập nhật",
    "IN_PHIEU": "In phiếu thực địa",
}


def lay_nguoi_dung(request: Request) -> dict[str, Any]:
    return dict(request.scope.get("auth_user") or {})


def chuyen_so_nguyen(value: str | None) -> int | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        result = int(text)
    except (TypeError, ValueError):
        return None
    return result if result > 0 else None


def co_quyen_mo_trung_tam(request: Request) -> bool:
    user = lay_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))
    return (
        is_province_reader(role_code)
        or role_code == COMMUNE_ROLE_CODE
    )


def lay_pham_vi_xa(
    request: Request,
    db: Session,
) -> list[Commune]:
    user = lay_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))

    statement = (
        select(Commune)
        .where(Commune.is_active.is_(True))
        .order_by(Commune.name.asc())
    )

    if role_code == COMMUNE_ROLE_CODE:
        commune_id = user.get("commune_id")
        if commune_id is None:
            return []
        statement = statement.where(Commune.id == int(commune_id))

    return list(db.scalars(statement).all())


def lay_nam_hoc(
    db: Session,
    school_year_id: int | None,
) -> tuple[list[SchoolYear], SchoolYear | None]:
    years = list(
        db.scalars(
            select(SchoolYear)
            .where(SchoolYear.is_active.is_(True))
            .order_by(SchoolYear.code.desc())
        ).all()
    )

    selected = None
    if school_year_id is not None:
        selected = next(
            (item for item in years if item.id == school_year_id),
            None,
        )
    if selected is None and years:
        selected = years[0]

    return years, selected


def lay_dot_dai_dien_theo_xa(
    db: Session,
    *,
    school_year_id: int,
    commune_ids: list[int],
) -> dict[int, SurveyBatch]:
    if not commune_ids:
        return {}

    batches = list(
        db.scalars(
            select(SurveyBatch)
            .options(
                selectinload(SurveyBatch.commune),
                selectinload(SurveyBatch.school_year),
            )
            .where(
                SurveyBatch.school_year_id == school_year_id,
                SurveyBatch.commune_id.in_(commune_ids),
            )
            .order_by(
                SurveyBatch.commune_id.asc(),
                SurveyBatch.is_locked.desc(),
                SurveyBatch.id.desc(),
            )
        ).all()
    )

    result: dict[int, SurveyBatch] = {}
    for batch in batches:
        result.setdefault(batch.commune_id, batch)
    return result


def lay_so_lieu_dot(
    db: Session,
    batch_ids: list[int],
) -> tuple[dict[int, int], dict[int, int], dict[int, int]]:
    if not batch_ids:
        return {}, {}, {}

    form_counts = {
        int(batch_id): int(total or 0)
        for batch_id, total in db.execute(
            select(
                SurveyForm.survey_batch_id,
                func.count(SurveyForm.id),
            )
            .where(SurveyForm.survey_batch_id.in_(batch_ids))
            .group_by(SurveyForm.survey_batch_id)
        ).all()
    }

    completed_counts = {
        int(batch_id): int(total or 0)
        for batch_id, total in db.execute(
            select(
                SurveyForm.survey_batch_id,
                func.count(SurveyForm.id),
            )
            .where(
                SurveyForm.survey_batch_id.in_(batch_ids),
                SurveyForm.status == "DA_HOAN_THANH",
            )
            .group_by(SurveyForm.survey_batch_id)
        ).all()
    }

    person_counts = {
        int(batch_id): int(total or 0)
        for batch_id, total in db.execute(
            select(
                SurveyForm.survey_batch_id,
                func.count(func.distinct(SurveyPerson.id)),
            )
            .join(Household, Household.id == SurveyForm.household_id)
            .join(
                SurveyPerson,
                SurveyPerson.household_id == Household.id,
            )
            .where(
                SurveyForm.survey_batch_id.in_(batch_ids),
                SurveyPerson.is_active.is_(True),
            )
            .group_by(SurveyForm.survey_batch_id)
        ).all()
    }

    return form_counts, completed_counts, person_counts


def lay_nhat_ky_moi_nhat(
    db: Session,
    batch_ids: list[int],
) -> tuple[
    dict[int, SurveyFileExchangeLog],
    dict[int, SurveyFileExchangeLog],
    list[SurveyFileExchangeLog],
]:
    if not batch_ids:
        return {}, {}, []

    logs = list(
        db.scalars(
            select(SurveyFileExchangeLog)
            .where(SurveyFileExchangeLog.survey_batch_id.in_(batch_ids))
            .order_by(SurveyFileExchangeLog.created_at.desc())
        ).all()
    )

    latest_export: dict[int, SurveyFileExchangeLog] = {}
    latest_import: dict[int, SurveyFileExchangeLog] = {}

    for log in logs:
        if log.action in {"XUAT_DON", "XUAT_HANG_LOAT", "IN_PHIEU"}:
            latest_export.setdefault(log.survey_batch_id, log)
        elif log.action == "NHAP_EXCEL":
            latest_import.setdefault(log.survey_batch_id, log)

    return latest_export, latest_import, logs


def xac_dinh_trang_thai_trao_doi(
    *,
    batch: SurveyBatch | None,
    export_log: SurveyFileExchangeLog | None,
    import_log: SurveyFileExchangeLog | None,
) -> str:
    if batch is None:
        return "CHUA_CO_DOT"
    if export_log is None:
        return "CHUA_PHAT_HANH"
    if import_log is None:
        return "CHO_TIEP_NHAN"
    if import_log.created_at >= export_log.created_at:
        return "DA_TIEP_NHAN"
    return "CHO_TIEP_NHAN"


def tao_du_lieu_trung_tam(
    *,
    request: Request,
    db: Session,
    selected_year: SchoolYear,
    q: str,
    exchange_status: str,
) -> dict[str, Any]:
    communes = lay_pham_vi_xa(request, db)
    commune_ids = [item.id for item in communes]
    batches_by_commune = lay_dot_dai_dien_theo_xa(
        db,
        school_year_id=selected_year.id,
        commune_ids=commune_ids,
    )
    batch_ids = [item.id for item in batches_by_commune.values()]
    form_counts, completed_counts, person_counts = lay_so_lieu_dot(
        db,
        batch_ids,
    )
    latest_export, latest_import, logs = lay_nhat_ky_moi_nhat(
        db,
        batch_ids,
    )

    q_key = q.strip().casefold()
    rows: list[dict[str, Any]] = []

    for commune in communes:
        batch = batches_by_commune.get(commune.id)
        export_log = latest_export.get(batch.id) if batch else None
        import_log = latest_import.get(batch.id) if batch else None
        state = xac_dinh_trang_thai_trao_doi(
            batch=batch,
            export_log=export_log,
            import_log=import_log,
        )

        haystack = " ".join(
            [
                commune.code,
                commune.name,
                batch.code if batch else "",
                batch.name if batch else "",
            ]
        ).casefold()

        if q_key and q_key not in haystack:
            continue
        if exchange_status and state != exchange_status:
            continue

        form_total = form_counts.get(batch.id, 0) if batch else 0
        completed_total = completed_counts.get(batch.id, 0) if batch else 0
        person_total = person_counts.get(batch.id, 0) if batch else 0

        rows.append({
            "commune": commune,
            "batch": batch,
            "state": state,
            "state_label": EXCHANGE_STATUS_LABELS.get(state, state),
            "form_total": form_total,
            "completed_total": completed_total,
            "person_total": person_total,
            "latest_export": export_log,
            "latest_import": import_log,
        })

    all_states = [
        xac_dinh_trang_thai_trao_doi(
            batch=batches_by_commune.get(commune.id),
            export_log=(
                latest_export.get(batches_by_commune[commune.id].id)
                if commune.id in batches_by_commune
                else None
            ),
            import_log=(
                latest_import.get(batches_by_commune[commune.id].id)
                if commune.id in batches_by_commune
                else None
            ),
        )
        for commune in communes
    ]

    summary = {
        "commune_total": len(communes),
        "batch_total": len(batches_by_commune),
        "not_issued_total": all_states.count("CHUA_PHAT_HANH")
        + all_states.count("CHUA_CO_DOT"),
        "waiting_total": all_states.count("CHO_TIEP_NHAN"),
        "received_total": all_states.count("DA_TIEP_NHAN"),
        "form_total": sum(form_counts.values()),
    }

    return {
        "rows": rows,
        "summary": summary,
        "all_logs": logs,
    }


async def doc_noi_dung_streaming(response: StreamingResponse) -> bytes:
    chunks: list[bytes] = []
    async for chunk in response.body_iterator:
        if isinstance(chunk, str):
            chunks.append(chunk.encode("utf-8"))
        else:
            chunks.append(bytes(chunk))
    return b"".join(chunks)


def lay_ten_file_tu_header(response: StreamingResponse) -> str:
    disposition = str(response.headers.get("content-disposition") or "")
    marker = 'filename="'
    if marker in disposition:
        return disposition.split(marker, 1)[1].split('"', 1)[0]
    return f"phieu_dieu_tra_{datetime.now():%Y%m%d_%H%M%S}.xlsx"


def ghi_nhat_ky_xuat(
    *,
    request: Request,
    db: Session,
    batch: SurveyBatch,
    action: str,
    file_name: str,
    form_total: int,
    person_total: int,
    notes: str | None = None,
) -> None:
    user = lay_nguoi_dung(request)
    log = SurveyFileExchangeLog(
        school_year_id=batch.school_year_id,
        survey_batch_id=batch.id,
        commune_id=batch.commune_id,
        action=action,
        status="THANH_CONG",
        file_name=file_name,
        actor_user_id=user.get("id"),
        actor_name_snapshot=user.get("full_name"),
        actor_role_snapshot=user.get("role_name") or user.get("role_code"),
        form_total=form_total,
        person_total=person_total,
        notes=notes,
    )
    db.add(log)


@router.get(
    "/trung-tam-phieu",
    response_class=HTMLResponse,
)
def trung_tam_phieu_dieu_tra(
    request: Request,
    school_year_id: str | None = None,
    q: str = "",
    exchange_status: str = "",
    status: str | None = None,
    db: Session = Depends(get_db),
):
    if not co_quyen_mo_trung_tam(request):
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    years, selected_year = lay_nam_hoc(
        db,
        chuyen_so_nguyen(school_year_id),
    )

    if selected_year is None:
        return templates.TemplateResponse(
            request=request,
            name="surveys/survey_exchange_center.html",
            context={
                "nguoi_dung": lay_nguoi_dung(request),
                "school_years": years,
                "selected_year": None,
                "rows": [],
                "summary": {
                    "commune_total": 0,
                    "batch_total": 0,
                    "not_issued_total": 0,
                    "waiting_total": 0,
                    "received_total": 0,
                    "form_total": 0,
                },
                "q": q,
                "exchange_status": exchange_status,
                "exchange_status_labels": EXCHANGE_STATUS_LABELS,
                "co_quyen_nhap": False,
                "co_quyen_xuat_hang_loat": False,
                "message": "Chưa có năm học đang hoạt động.",
            },
        )

    data = tao_du_lieu_trung_tam(
        request=request,
        db=db,
        selected_year=selected_year,
        q=q,
        exchange_status=exchange_status,
    )

    user = lay_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))
    messages = {
        "bulk_exported": "Đã tạo gói ZIP phát hành hàng loạt và ghi nhật ký.",
        "no_batch_selected": "Hãy chọn ít nhất một xã/phường đã có đợt.",
    }

    return templates.TemplateResponse(
        request=request,
        name="surveys/survey_exchange_center.html",
        context={
            "nguoi_dung": user,
            "school_years": years,
            "selected_year": selected_year,
            "rows": data["rows"],
            "summary": data["summary"],
            "q": q,
            "exchange_status": exchange_status,
            "exchange_status_labels": EXCHANGE_STATUS_LABELS,
            "co_quyen_nhap": is_admin_role(role_code),
            "co_quyen_xuat_hang_loat": is_admin_role(role_code),
            "message": messages.get(str(status or "")),
        },
    )


@router.get("/{batch_id}/trung-tam-phieu/xuat")
async def xuat_file_tu_trung_tam(
    batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    batch = lay_dot_dieu_tra(db, batch_id)
    if batch is None:
        return RedirectResponse(
            url="/dieu-tra/trung-tam-phieu",
            status_code=303,
        )

    original = xuat_excel_di_dieu_tra(
        batch_id=batch.id,
        request=request,
        db=db,
    )
    if not isinstance(original, StreamingResponse):
        return original

    content = await doc_noi_dung_streaming(original)
    file_name = lay_ten_file_tu_header(original)
    form_total, _, person_total = lay_so_lieu_dot(db, [batch.id])

    ghi_nhat_ky_xuat(
        request=request,
        db=db,
        batch=batch,
        action="XUAT_DON",
        file_name=file_name,
        form_total=form_total.get(batch.id, 0),
        person_total=person_total.get(batch.id, 0),
        notes="Phát hành file từ Trung tâm phiếu điều tra.",
    )
    db.commit()

    return Response(
        content=content,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f"attachment; filename*=UTF-8''{quote(file_name)}"
            )
        },
    )


@router.post("/trung-tam-phieu/xuat-hang-loat")
async def xuat_hang_loat_tu_trung_tam(
    request: Request,
    batch_ids: Annotated[list[int], Form()] = [],
    school_year_id: Annotated[int, Form()] = 0,
    db: Session = Depends(get_db),
):
    user = lay_nguoi_dung(request)
    if not is_admin_role(user.get("role_code")):
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    unique_ids = sorted({int(item) for item in batch_ids if int(item) > 0})
    if not unique_ids:
        return RedirectResponse(
            url=(
                "/dieu-tra/trung-tam-phieu"
                f"?school_year_id={school_year_id}"
                "&status=no_batch_selected"
            ),
            status_code=303,
        )

    batches = list(
        db.scalars(
            select(SurveyBatch)
            .options(
                selectinload(SurveyBatch.commune),
                selectinload(SurveyBatch.school_year),
            )
            .where(
                SurveyBatch.id.in_(unique_ids),
                SurveyBatch.school_year_id == school_year_id,
            )
            .order_by(SurveyBatch.commune_id.asc())
        ).all()
    )

    if not batches:
        return RedirectResponse(
            url=(
                "/dieu-tra/trung-tam-phieu"
                f"?school_year_id={school_year_id}"
                "&status=no_batch_selected"
            ),
            status_code=303,
        )

    form_counts, _, person_counts = lay_so_lieu_dot(
        db,
        [item.id for item in batches],
    )

    zip_output = BytesIO()
    manifest_rows: list[list[Any]] = []

    with ZipFile(zip_output, mode="w", compression=ZIP_DEFLATED) as archive:
        for batch in batches:
            original = xuat_excel_di_dieu_tra(
                batch_id=batch.id,
                request=request,
                db=db,
            )
            if not isinstance(original, StreamingResponse):
                continue

            content = await doc_noi_dung_streaming(original)
            file_name = lay_ten_file_tu_header(original)
            archive.writestr(file_name, content)

            form_total = form_counts.get(batch.id, 0)
            person_total = person_counts.get(batch.id, 0)
            ghi_nhat_ky_xuat(
                request=request,
                db=db,
                batch=batch,
                action="XUAT_HANG_LOAT",
                file_name=file_name,
                form_total=form_total,
                person_total=person_total,
                notes="Phát hành trong gói ZIP hàng loạt.",
            )
            manifest_rows.append([
                batch.commune.code,
                batch.commune.name,
                batch.code,
                form_total,
                person_total,
                file_name,
            ])

        manifest = Workbook()
        sheet = manifest.active
        sheet.title = "DANH_SACH_PHAT_HANH"
        headers = [
            "Mã xã/phường",
            "Tên xã/phường",
            "Mã đợt",
            "Số phiếu",
            "Số đối tượng",
            "Tên file",
        ]
        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        for row_index, values in enumerate(manifest_rows, start=2):
            for col, value in enumerate(values, start=1):
                sheet.cell(row=row_index, column=col, value=value)
        widths = [18, 30, 28, 12, 14, 42]
        for col, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        manifest_bytes = BytesIO()
        manifest.save(manifest_bytes)
        archive.writestr(
            "DANH_SACH_PHAT_HANH.xlsx",
            manifest_bytes.getvalue(),
        )

    db.commit()
    zip_output.seek(0)

    selected_year = db.get(SchoolYear, school_year_id)
    year_code = (
        selected_year.code.replace("-", "")
        if selected_year is not None
        else str(school_year_id)
    )
    file_name = f"phat_hanh_phieu_toan_tinh_{year_code}.zip"

    return Response(
        content=zip_output.getvalue(),
        media_type="application/zip",
        headers={
            "Content-Disposition": (
                f"attachment; filename*=UTF-8''{quote(file_name)}"
            )
        },
    )


@router.get("/trung-tam-phieu/xuat-bao-cao")
def xuat_bao_cao_trung_tam(
    request: Request,
    school_year_id: str | None = None,
    db: Session = Depends(get_db),
):
    if not co_quyen_mo_trung_tam(request):
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    _, selected_year = lay_nam_hoc(
        db,
        chuyen_so_nguyen(school_year_id),
    )
    if selected_year is None:
        return RedirectResponse(
            url="/dieu-tra/trung-tam-phieu",
            status_code=303,
        )

    data = tao_du_lieu_trung_tam(
        request=request,
        db=db,
        selected_year=selected_year,
        q="",
        exchange_status="",
    )

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Tong quan"
    overview.append(["TRUNG TÂM PHÁT HÀNH VÀ TIẾP NHẬN PHIẾU"])
    overview.append(["Năm học", selected_year.code])
    overview.append(["Xã/phường trong phạm vi", data["summary"]["commune_total"]])
    overview.append(["Đã có đợt", data["summary"]["batch_total"]])
    overview.append(["Chưa phát hành", data["summary"]["not_issued_total"]])
    overview.append(["Đã phát hành, chờ tiếp nhận", data["summary"]["waiting_total"]])
    overview.append(["Đã tiếp nhận", data["summary"]["received_total"]])
    overview.append(["Tổng số phiếu", data["summary"]["form_total"]])
    overview.column_dimensions["A"].width = 38
    overview.column_dimensions["B"].width = 22

    detail = workbook.create_sheet("Theo xa phuong")
    detail_headers = [
        "STT",
        "Mã xã/phường",
        "Tên xã/phường",
        "Mã đợt",
        "Số phiếu",
        "Hoàn thành",
        "Đối tượng",
        "Trạng thái trao đổi",
        "Lần xuất gần nhất",
        "Người xuất",
        "Lần nhập gần nhất",
        "Người nhập",
    ]
    detail.append(detail_headers)
    for index, row in enumerate(data["rows"], start=1):
        export_log = row["latest_export"]
        import_log = row["latest_import"]
        detail.append([
            index,
            row["commune"].code,
            row["commune"].name,
            row["batch"].code if row["batch"] else "",
            row["form_total"],
            row["completed_total"],
            row["person_total"],
            row["state_label"],
            export_log.created_at if export_log else "",
            export_log.actor_name_snapshot if export_log else "",
            import_log.created_at if import_log else "",
            import_log.actor_name_snapshot if import_log else "",
        ])

    logs_sheet = workbook.create_sheet("Nhat ky")
    logs_sheet.append([
        "Thời gian",
        "Thao tác",
        "Mã đợt",
        "Tên file",
        "Người thực hiện",
        "Vai trò",
        "Số phiếu",
        "Số đối tượng",
        "Ghi chú",
    ])
    for log in data["all_logs"]:
        logs_sheet.append([
            log.created_at,
            ACTION_LABELS.get(log.action, log.action),
            log.survey_batch_id,
            log.file_name or "",
            log.actor_name_snapshot or "",
            log.actor_role_snapshot or "",
            log.form_total,
            log.person_total,
            log.notes or "",
        ])

    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 22

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    file_name = (
        "trung_tam_phieu_"
        f"{selected_year.code.replace('-', '')}.xlsx"
    )

    return Response(
        content=output.getvalue(),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f"attachment; filename*=UTF-8''{quote(file_name)}"
            )
        },
    )
