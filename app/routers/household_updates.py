from __future__ import annotations

import hashlib
import json
import re
import threading
import unicodedata
import zipfile
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

from fastapi import APIRouter, Depends, File, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import select, text
from sqlalchemy.orm import Session, selectinload

from app.database import get_db
from app.household_import_models import HouseholdImportError, HouseholdImportJob
from app.models import Classroom, School, SchoolYear
from app.permissions import (
    ADMIN_ROLE_CODES,
    COMMUNE_ROLE_CODE,
    DEPARTMENT_ROLE_CODE,
    SCHOOL_ROLE_CODE,
    TEACHER_ROLE_CODE,
    normalize_role_code,
)
from app.survey_models import (
    Household,
    SurveyBatch,
    SurveyForm,
    SurveyPerson,
    SurveyPersonYearRecord,
)
from app.routers.surveys import (
    BATCH_STATUS_LABELS,
    FORM_STATUS_LABELS,
    LEARNING_STATUS_LABELS,
    RESIDENCY_STATUS_LABELS,
    lay_thong_tin_nguoi_dung,
    tao_bo_loc_bao_cao_theo_nguoi_dung,
    tao_bo_loc_dot_theo_nguoi_dung,
)


APP_DIR = Path(__file__).resolve().parent.parent
PROJECT_DIR = APP_DIR.parent
REFERENCE_DIR = APP_DIR / "reference" / "bai_13b_9"
IMPORT_ROOT = PROJECT_DIR / "exports" / "household_imports"

templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

router = APIRouter(
    prefix="/dieu-tra",
    tags=["Cập nhật dữ liệu hộ dân"],
)

UPLOAD_ROLE_CODES = frozenset(
    {
        *ADMIN_ROLE_CODES,
        COMMUNE_ROLE_CODE,
        SCHOOL_ROLE_CODE,
    }
)

VIEW_ROLE_CODES = frozenset(
    {
        *ADMIN_ROLE_CODES,
        DEPARTMENT_ROLE_CODE,
        COMMUNE_ROLE_CODE,
        SCHOOL_ROLE_CODE,
    }
)

IMPORT_STATUS_LABELS = {
    "DA_NHAN": "Đã nhận file",
    "DANG_KIEM_TRA": "Đang kiểm tra",
    "DA_KIEM_TRA_CAU_TRUC": "Đã kiểm tra cấu trúc",
    "DANG_CAP_NHAT": "Đang cập nhật dữ liệu",
    "LOI_DU_LIEU": "Có lỗi dữ liệu – chưa cập nhật",
    "LOI_CAP_NHAT": "Cập nhật thất bại – đã rollback",
    "CHO_XU_LY_XLS": "Đã nhận – chờ xử lý XLS",
    "TRUNG_FILE": "File đã gửi trước đó",
    "LOI_DINH_DANG": "Có lỗi định dạng",
    "LOI_FILE": "Có lỗi file",
    "LOI_DOT_KHOA": "Đợt không nhận cập nhật",
    "DA_CAP_NHAT": "Đã cập nhật thành công",
    "CO_XUNG_DOT": "Có xung đột dữ liệu",
}

IMPORT_STATUS_CLASSES = {
    "DA_NHAN": "info",
    "DANG_KIEM_TRA": "info",
    "DA_KIEM_TRA_CAU_TRUC": "success",
    "DANG_CAP_NHAT": "info",
    "LOI_DU_LIEU": "danger",
    "LOI_CAP_NHAT": "danger",
    "CHO_XU_LY_XLS": "warning",
    "TRUNG_FILE": "warning",
    "LOI_DINH_DANG": "danger",
    "LOI_FILE": "danger",
    "LOI_DOT_KHOA": "danger",
    "DA_CAP_NHAT": "success",
    "CO_XUNG_DOT": "warning",
}


# === BAI_13B_9_V1B_1_DELETABLE_STATUSES_START ===
DELETEABLE_IMPORT_STATUSES = frozenset(
    {
        "DA_KIEM_TRA_CAU_TRUC",
        "LOI_DU_LIEU",
        "LOI_CAP_NHAT",
        "CHO_XU_LY_XLS",
        "TRUNG_FILE",
        "LOI_DINH_DANG",
        "LOI_FILE",
        "LOI_DOT_KHOA",
        "CO_XUNG_DOT",
    }
)
# === BAI_13B_9_V1B_1_DELETABLE_STATUSES_END ===

def _user(request: Request) -> dict:
    return dict(request.scope.get("auth_user") or {})


def _role(request: Request) -> str:
    return normalize_role_code(_user(request).get("role_code"))


def _scope_label(request: Request) -> str:
    role_code = _role(request)
    if role_code == COMMUNE_ROLE_CODE:
        return "PHẠM VI XÃ/PHƯỜNG"
    if role_code == SCHOOL_ROLE_CODE:
        return "PHẠM VI TRƯỜNG ĐƯỢC PHÂN CÔNG"
    if role_code == DEPARTMENT_ROLE_CODE:
        return "PHẠM VI TOÀN TỈNH – CHỈ XEM"
    return "PHẠM VI TOÀN TỈNH"


def _batch_in_scope(
    db: Session,
    request: Request,
    batch_id: int,
) -> SurveyBatch | None:
    filters = tao_bo_loc_dot_theo_nguoi_dung(request)
    return db.scalar(
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
        )
        .where(
            SurveyBatch.id == batch_id,
            *filters,
        )
    )


def _safe_file_name(name: str) -> str:
    base = Path(name or "upload.xlsx").name
    base = re.sub(r"[^0-9A-Za-zÀ-ỹ._() -]+", "_", base)
    base = re.sub(r"\s+", "_", base).strip("._")
    return base[:220] or "upload.xlsx"


def _new_job_code() -> str:
    now = datetime.now()
    return (
        "IMP-"
        + now.strftime("%Y%m%d-%H%M%S")
        + "-"
        + uuid4().hex[:6].upper()
    )


def _add_issue(
    db: Session,
    job: HouseholdImportJob,
    *,
    severity: str,
    code: str,
    message: str,
    sheet_name: str | None = None,
    row_number: int | None = None,
    column_name: str | None = None,
) -> None:
    db.add(
        HouseholdImportError(
            job_id=job.id,
            severity=severity,
            sheet_name=sheet_name,
            row_number=row_number,
            column_name=column_name,
            error_code=code,
            message=message,
        )
    )


def _xlsx_structure(content: bytes) -> tuple[list[str], bool]:
    if not content.startswith(b"PK\x03\x04"):
        raise ValueError("Tệp .xlsx không có cấu trúc ZIP hợp lệ.")

    with zipfile.ZipFile(BytesIO(content)) as archive:
        names = set(archive.namelist())
        if "xl/workbook.xml" not in names:
            raise ValueError("Không tìm thấy xl/workbook.xml.")

        has_macro = any(
            item.lower().endswith("vbaproject.bin")
            for item in names
        )

        xml = archive.read("xl/workbook.xml")
        root = ElementTree.fromstring(xml)
        sheet_names: list[str] = []
        for item in root.iter():
            if item.tag.endswith("}sheet") or item.tag == "sheet":
                name = str(item.attrib.get("name") or "").strip()
                if name:
                    sheet_names.append(name)

        return sheet_names, has_macro


def _xls_macro_hint(content: bytes) -> bool:
    marker_ascii = b"_VBA_PROJECT_CUR"
    marker_utf16 = "_VBA_PROJECT_CUR".encode("utf-16le")
    return marker_ascii in content or marker_utf16 in content


def _store_received_file(
    job_code: str,
    file_name: str,
    content: bytes,
) -> Path:
    folder = IMPORT_ROOT / job_code
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / _safe_file_name(file_name)
    path.write_bytes(content)
    return path


def _json_result(
    *,
    ok: bool,
    job: HouseholdImportJob | None,
    message: str,
    status_code: int = 200,
) -> JSONResponse:
    return JSONResponse(
        status_code=status_code,
        content={
            "ok": ok,
            "job_code": job.job_code if job is not None else None,
            "status": job.status if job is not None else None,
            "status_label": (
                IMPORT_STATUS_LABELS.get(job.status, job.status)
                if job is not None
                else None
            ),
            "message": message,
            "redirect_url": (
                f"/dieu-tra/cap-nhat-ho-dan/lan-gui/{job.job_code}"
                if job is not None
                else None
            ),
        },
    )



_IMPORT_WRITE_LOCK = threading.Lock()

HEADER_ALIASES = {
    "tt": {"TT", "STT"},
    "ho_dem": {"HO DEM", "HO VA TEN DEM", "HO TEN DEM"},
    "ten": {"TEN"},
    "ngay": {"NGAY"},
    "thang": {"THANG"},
    "nam_sinh": {"NAM SINH", "NAM"},
    "nu": {"NU", "GIOI TINH NU"},
    "dan_toc": {"DAN TOC"},
    "ton_giao": {"TON GIAO"},
    "dien_uu_tien": {"DIEN UU TIEN", "UU TIEN"},
    "chu_ho_ho_dem": {"CHU HO - HO DEM", "HO DEM CHU HO", "CHU HO HO DEM"},
    "chu_ho_ten": {"CHU HO - TEN", "TEN CHU HO", "CHU HO TEN"},
    "dia_chi": {"DIA CHI: SO NHA, TEN DUONG, TO (NEU CO)", "DIA CHI", "DIA CHI HO"},
    "so_phieu": {"SO PHIEU", "PHIEU"},
    "dien_cu_tru": {"DIEN CU TRU", "CU TRU"},
    "tinh_trang_cu_tru": {"TINH TRANG CU TRU", "TRANG THAI CU TRU"},
    "khoi_hoc": {"KHOI HOC", "CAP HOC", "NHOM CAP HOC"},
    "lop_hoc": {"LOP HOC", "LOP"},
    "ten_truong": {"TEN TRUONG", "TRUONG"},
    "dia_ban_truong": {"DIA BAN TRUONG", "DIA BAN"},
    "ma_truong": {"MA TRUONG"},
    "so_dinh_danh": {"SO DINH DANH", "SO DINH DANH CA NHAN"},
    "cccd": {"CAN CUOC CONG DAN", "CCCD", "CAN CUOC"},
}

REQUIRED_FIELDS = {
    "ho_dem", "ten", "ngay", "thang", "nam_sinh",
    "chu_ho_ho_dem", "chu_ho_ten", "dia_chi", "so_phieu",
}


def _norm(value) -> str:
    raw = str(value or "").strip()
    raw = unicodedata.normalize("NFD", raw)
    raw = "".join(ch for ch in raw if unicodedata.category(ch) != "Mn")
    raw = raw.upper().replace("Đ", "D")
    raw = re.sub(r"[^0-9A-Z]+", " ", raw)
    return re.sub(r"\s+", " ", raw).strip()


def _plain(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        fmt = str(cell.number_format or "")
        if re.fullmatch(r"0+", fmt):
            width = len(fmt)
            try:
                return str(int(value)).zfill(width)
            except Exception:
                pass
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
    return str(value).strip()


def _batch_start_year(batch) -> int | None:
    match = re.search(r"(20\d{2})", str(batch.school_year.code or ""))
    return int(match.group(1)) if match else None


def _find_header_row(ws) -> tuple[int, dict[str, int]]:
    alias_to_key = {}
    for key, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            alias_to_key[_norm(alias)] = key

    best_row = 0
    best_map: dict[str, int] = {}
    limit = min(max(ws.max_row, 1), 20)
    for row_no in range(1, limit + 1):
        current: dict[str, int] = {}
        for col_no in range(1, min(ws.max_column, 80) + 1):
            key = alias_to_key.get(_norm(ws.cell(row_no, col_no).value))
            if key and key not in current:
                current[key] = col_no
        if len(current) > len(best_map):
            best_row, best_map = row_no, current
        if REQUIRED_FIELDS.issubset(current):
            return row_no, current

    missing = sorted(REQUIRED_FIELDS - set(best_map))
    raise ValueError(
        "Không tìm thấy hàng tiêu đề đúng mẫu. Thiếu: " + ", ".join(missing)
    )


def _metadata_value(ws, labels: set[str], *, max_rows: int = 6, max_cols: int = 12):
    wanted = {_norm(item) for item in labels}
    for row in range(1, min(ws.max_row, max_rows) + 1):
        for col in range(1, min(ws.max_column, max_cols) + 1):
            if _norm(ws.cell(row, col).value) in wanted:
                for offset in (1, 2):
                    if col + offset <= ws.max_column:
                        value = _plain(ws.cell(row, col + offset).value)
                        if value:
                            return value
    return ""


def _looks_like_numbering_row(ws, row_no: int, header_map: dict[str, int]) -> bool:
    values = []
    for col in header_map.values():
        text_value = _plain(ws.cell(row_no, col).value)
        if text_value:
            values.append(text_value)
    if not values:
        return False
    numeric = sum(1 for item in values if re.fullmatch(r"\d+", item))
    return numeric >= max(3, int(len(values) * 0.6))


def _map_residency(dien_cu_tru: str, tinh_trang: str) -> str:
    status_key = _norm(tinh_trang)
    base_key = _norm(dien_cu_tru)
    if "CHUYEN DEN" in status_key:
        return "CHUYEN_DEN"
    if "CHUYEN DI" in status_key:
        return "CHUYEN_DI"
    if "TAM VANG" in status_key or "TAM VANG" in base_key:
        return "TAM_VANG"
    if "TAM TRU" in base_key:
        return "TAM_TRU"
    if "THUONG TRU" in base_key:
        return "THUONG_TRU"
    return "CHUA_XAC_DINH"


def _map_learning(khoi_hoc: str, school_name: str, class_name: str, residency_status: str) -> str:
    if residency_status == "CHUYEN_DEN":
        return "CHUYEN_DEN"
    if residency_status == "CHUYEN_DI":
        return "CHUYEN_DI"
    key = _norm(khoi_hoc)
    if key in {"XMC", "XOA MU CHU"} and not (school_name or class_name):
        return "KHONG_THUOC_DIEN"
    if school_name or class_name or key in {"MN", "TH", "THCS", "THPT"}:
        return "DANG_HOC"
    if "CHUA DI HOC" in key:
        return "CHUA_DI_HOC"
    if "BO HOC" in key:
        return "BO_HOC"
    return "CHUA_XAC_DINH"


def _merge_notes(*parts: str) -> str | None:
    items = [str(item).strip() for item in parts if str(item or "").strip()]
    return "; ".join(items) if items else None


def _parse_xlsx_rows(content: bytes, batch) -> tuple[list[dict], list[dict], dict]:
    wb = load_workbook(BytesIO(content), data_only=True, read_only=False)
    ws = wb["MauNhapLieu"] if "MauNhapLieu" in wb.sheetnames else wb[wb.sheetnames[0]]
    header_row, header_map = _find_header_row(ws)

    metadata = {
        "sheet": ws.title,
        "year": _metadata_value(ws, {"Năm điều tra", "Năm học", "Năm"}),
        "commune": _metadata_value(ws, {"Xã/phường", "Xã", "Phường"}),
    }
    issues: list[dict] = []

    expected_year = _batch_start_year(batch)
    if metadata["year"]:
        match = re.search(r"20\d{2}", metadata["year"])
        if match and expected_year is not None and int(match.group(0)) != expected_year:
            issues.append({
                "severity": "ERROR", "code": "YEAR_MISMATCH", "row": None,
                "column": "Năm điều tra",
                "message": f"Năm trong file ({metadata['year']}) không khớp đợt {batch.school_year.code}.",
            })

    if metadata["commune"]:
        file_commune = _norm(metadata["commune"])
        batch_commune = _norm(batch.commune.name)
        if file_commune != batch_commune and file_commune not in batch_commune and batch_commune not in file_commune:
            issues.append({
                "severity": "ERROR", "code": "COMMUNE_MISMATCH", "row": None,
                "column": "Xã/phường",
                "message": f"Xã/phường trong file ({metadata['commune']}) không khớp đợt nhận ({batch.commune.name}).",
            })

    data_start = header_row + 1
    if data_start <= ws.max_row and _looks_like_numbering_row(ws, data_start, header_map):
        data_start += 1

    def get_text(row_no: int, key: str) -> str:
        col = header_map.get(key)
        return _cell_text(ws.cell(row_no, col)) if col else ""

    rows: list[dict] = []
    seen_pid: dict[str, tuple[int, str, str]] = {}
    seen_cccd: dict[str, tuple[int, str, str]] = {}
    form_identity: dict[str, tuple[str, str, int]] = {}
    blank_run = 0

    for row_no in range(data_start, ws.max_row + 1):
        raw_values = [get_text(row_no, key) for key in header_map]
        if not any(raw_values):
            blank_run += 1
            if blank_run >= 25:
                break
            continue
        blank_run = 0

        full_name = " ".join(filter(None, [get_text(row_no, "ho_dem"), get_text(row_no, "ten")])).strip()
        head_name = " ".join(filter(None, [get_text(row_no, "chu_ho_ho_dem"), get_text(row_no, "chu_ho_ten")])).strip()
        form_number = get_text(row_no, "so_phieu")
        address = get_text(row_no, "dia_chi")
        personal_id = get_text(row_no, "so_dinh_danh")
        citizen_id = get_text(row_no, "cccd")

        row_errors_before = len(issues)
        if not full_name:
            issues.append({"severity":"ERROR","code":"MISSING_NAME","row":row_no,"column":"Họ tên","message":"Thiếu họ tên thành viên."})
        if not head_name:
            issues.append({"severity":"ERROR","code":"MISSING_HEAD","row":row_no,"column":"Chủ hộ","message":"Thiếu tên chủ hộ."})
        if not address:
            issues.append({"severity":"ERROR","code":"MISSING_ADDRESS","row":row_no,"column":"Địa chỉ","message":"Thiếu địa chỉ hộ."})
        if not form_number:
            issues.append({"severity":"ERROR","code":"MISSING_FORM","row":row_no,"column":"Số phiếu","message":"Thiếu Số phiếu; không xác định được hộ cần cập nhật."})

        try:
            day = int(float(get_text(row_no, "ngay")))
            month = int(float(get_text(row_no, "thang")))
            year = int(float(get_text(row_no, "nam_sinh")))
            birth = date(year, month, day)
        except Exception:
            birth = None
            issues.append({"severity":"ERROR","code":"INVALID_BIRTH_DATE","row":row_no,"column":"Ngày/Tháng/Năm sinh","message":"Ngày sinh không hợp lệ."})

        identity = (_norm(head_name), _norm(address), row_no)
        if form_number:
            previous = form_identity.get(form_number)
            if previous and (previous[0], previous[1]) != (identity[0], identity[1]):
                issues.append({"severity":"ERROR","code":"FORM_INCONSISTENT","row":row_no,"column":"Số phiếu","message":f"Số phiếu {form_number} có thông tin chủ hộ/địa chỉ không thống nhất với dòng {previous[2]}."})
            else:
                form_identity.setdefault(form_number, identity)

        birth_key = birth.isoformat() if birth else ""
        for value, seen, label, code in [
            (personal_id, seen_pid, "Số định danh", "DUPLICATE_PERSONAL_ID_IN_FILE"),
            (citizen_id, seen_cccd, "CCCD", "DUPLICATE_CCCD_IN_FILE"),
        ]:
            if not value:
                continue
            previous = seen.get(value)
            if previous and (previous[1], previous[2]) != (_norm(full_name), birth_key):
                issues.append({"severity":"ERROR","code":code,"row":row_no,"column":label,"message":f"{label} {value} bị dùng cho người khác ở dòng {previous[0]}."})
            else:
                seen.setdefault(value, (row_no, _norm(full_name), birth_key))

        residency = _map_residency(get_text(row_no, "dien_cu_tru"), get_text(row_no, "tinh_trang_cu_tru"))
        school_name = get_text(row_no, "ten_truong")
        class_name = get_text(row_no, "lop_hoc")
        level = get_text(row_no, "khoi_hoc")
        learning = _map_learning(level, school_name, class_name, residency)

        rows.append({
            "row_number": row_no,
            "full_name": full_name,
            "date_of_birth": birth,
            "gender": "Nữ" if _norm(get_text(row_no, "nu")) in {"X", "1", "CO", "NU"} else "Nam",
            "ethnic_group": get_text(row_no, "dan_toc") or None,
            "religion": get_text(row_no, "ton_giao"),
            "special": get_text(row_no, "dien_uu_tien"),
            "head_name": head_name,
            "address": address,
            "hamlet_name": address.split(",", 1)[0].strip() if address else None,
            "form_number": form_number,
            "residency_status": residency,
            "residency_text": get_text(row_no, "tinh_trang_cu_tru"),
            "level": level,
            "class_name": class_name,
            "school_name": school_name,
            "school_area": get_text(row_no, "dia_ban_truong"),
            "school_code": get_text(row_no, "ma_truong"),
            "personal_id": personal_id,
            "citizen_id": citizen_id,
            "learning_status": learning,
            "row_has_error": len(issues) > row_errors_before,
        })

    if not rows:
        issues.append({"severity":"ERROR","code":"NO_DATA_ROWS","row":None,"column":None,"message":"Không tìm thấy dòng dữ liệu thành viên trong file."})

    return rows, issues, metadata


def _add_issue_dict(db: Session, job: HouseholdImportJob, item: dict, sheet_name: str | None = None) -> None:
    _add_issue(
        db, job,
        severity=item.get("severity", "ERROR"),
        code=item.get("code", "DATA_ERROR"),
        message=item.get("message", "Lỗi dữ liệu."),
        sheet_name=sheet_name,
        row_number=item.get("row"),
        column_name=item.get("column"),
    )


def _find_existing_person(db: Session, row: dict, household: Household):
    person = None
    if row["personal_id"]:
        person = db.scalar(
            select(SurveyPerson).where(
                SurveyPerson.personal_id == row["personal_id"],
                SurveyPerson.is_active.is_(True),
            )
        )
    if person is None and row["citizen_id"] and hasattr(SurveyPerson, "citizen_id"):
        person = db.scalar(
            select(SurveyPerson).where(
                SurveyPerson.citizen_id == row["citizen_id"],
                SurveyPerson.is_active.is_(True),
            )
        )
    if person is None:
        candidates = list(db.scalars(
            select(SurveyPerson).where(
                SurveyPerson.household_id == household.id,
                SurveyPerson.is_active.is_(True),
                SurveyPerson.date_of_birth == row["date_of_birth"],
            )
        ).all())
        wanted = _norm(row["full_name"])
        person = next((item for item in candidates if _norm(item.full_name) == wanted), None)
    return person


def _precheck_database_conflicts(db: Session, rows: list[dict], batch) -> list[dict]:
    issues: list[dict] = []
    for row in rows:
        if row.get("row_has_error"):
            continue
        for field, model_field, label in [
            ("personal_id", SurveyPerson.personal_id, "Số định danh"),
        ]:
            value = row.get(field)
            if not value:
                continue
            person = db.scalar(select(SurveyPerson).where(model_field == value, SurveyPerson.is_active.is_(True)))
            if person is not None:
                if _norm(person.full_name) != _norm(row["full_name"]) or person.date_of_birth != row["date_of_birth"]:
                    issues.append({"severity":"ERROR","code":"IDENTIFIER_CONFLICT","row":row["row_number"],"column":label,"message":f"{label} {value} đã thuộc về người khác trong CSDL: {person.full_name}."})
        if row.get("citizen_id") and hasattr(SurveyPerson, "citizen_id"):
            person = db.scalar(select(SurveyPerson).where(SurveyPerson.citizen_id == row["citizen_id"], SurveyPerson.is_active.is_(True)))
            if person is not None and (_norm(person.full_name) != _norm(row["full_name"]) or person.date_of_birth != row["date_of_birth"]):
                issues.append({"severity":"ERROR","code":"CCCD_CONFLICT","row":row["row_number"],"column":"CCCD","message":f"CCCD {row['citizen_id']} đã thuộc về người khác trong CSDL: {person.full_name}."})
    return issues


def _apply_rows(db: Session, job: HouseholdImportJob, batch, rows: list[dict]) -> dict:
    households_created = households_updated = people_created = people_updated = 0
    touched_forms: dict[str, tuple[SurveyForm, Household, bool]] = {}

    for row in rows:
        form_number = row["form_number"]
        if form_number not in touched_forms:
            survey_form = db.scalar(
                select(SurveyForm)
                .options(selectinload(SurveyForm.household))
                .where(
                    SurveyForm.survey_batch_id == batch.id,
                    SurveyForm.form_number == form_number,
                )
            )
            created_house = False
            if survey_form is None:
                household = Household(
                    code="TMP-HO-" + uuid4().hex[:16].upper(),
                    commune_id=batch.commune_id,
                    head_name=row["head_name"],
                    hamlet_name=row["hamlet_name"],
                    address=row["address"],
                    phone=None,
                    is_active=True,
                )
                db.add(household)
                db.flush()
                commune_code = re.sub(r"[^0-9A-Za-z]+", "", str(batch.commune.code or batch.commune_id))
                household.code = f"HO-{commune_code}-{household.id:06d}"
                survey_form = SurveyForm(
                    survey_batch_id=batch.id,
                    household_id=household.id,
                    form_number=form_number,
                    head_name_snapshot=row["head_name"],
                    address_snapshot=row["address"],
                    hamlet_name_snapshot=row["hamlet_name"],
                    status="CHUA_DIEU_TRA",
                )
                db.add(survey_form)
                db.flush()
                created_house = True
                households_created += 1
            else:
                household = survey_form.household
                household.head_name = row["head_name"]
                household.address = row["address"]
                household.hamlet_name = row["hamlet_name"]
                household.is_active = True
                survey_form.head_name_snapshot = row["head_name"]
                survey_form.address_snapshot = row["address"]
                survey_form.hamlet_name_snapshot = row["hamlet_name"]
                households_updated += 1
            touched_forms[form_number] = (survey_form, household, created_house)

        survey_form, household, _ = touched_forms[form_number]
        person = _find_existing_person(db, row, household)
        if person is not None and int(person.household_id) != int(household.id):
            raise ValueError(
                f"Dòng {row['row_number']}: định danh của {row['full_name']} đang thuộc hộ khác ({person.household.code})."
            )

        if person is None:
            person = SurveyPerson(
                code="TMP-DT-" + uuid4().hex[:16].upper(),
                household_id=household.id,
                full_name=row["full_name"],
                date_of_birth=row["date_of_birth"],
                gender=row["gender"],
                ethnic_group=row["ethnic_group"],
                relationship_to_head=("Chủ hộ" if _norm(row["full_name"]) == _norm(row["head_name"]) else "Thành viên"),
                personal_id=row["personal_id"] or None,
                permanent_address=row["address"],
                current_address=row["address"],
                residency_status=row["residency_status"],
                special_circumstances=row["special"] or None,
                notes=_merge_notes(
                    f"Tôn giáo: {row['religion']}" if row["religion"] else "",
                    "Nhập từ Excel " + job.job_code,
                ),
                is_active=True,
            )
            if hasattr(person, "citizen_id"):
                person.citizen_id = row["citizen_id"] or None
            db.add(person)
            db.flush()
            person.code = f"DT-{person.id:08d}"
            people_created += 1
        else:
            person.full_name = row["full_name"]
            person.date_of_birth = row["date_of_birth"]
            person.gender = row["gender"]
            person.ethnic_group = row["ethnic_group"]
            person.personal_id = row["personal_id"] or person.personal_id
            person.permanent_address = row["address"]
            person.current_address = row["address"]
            person.residency_status = row["residency_status"]
            person.special_circumstances = row["special"] or person.special_circumstances
            if hasattr(person, "citizen_id") and row["citizen_id"]:
                person.citizen_id = row["citizen_id"]
            people_updated += 1

        school = None
        if row["school_code"]:
            school = db.scalar(select(School).where(School.code == row["school_code"]))
        if school is None and row["school_name"]:
            candidates = list(db.scalars(select(School).where(School.is_active.is_(True))).all())
            wanted_school = _norm(row["school_name"])
            school = next((item for item in candidates if _norm(item.name) == wanted_school), None)

        record = db.scalar(
            select(SurveyPersonYearRecord).where(
                SurveyPersonYearRecord.survey_form_id == survey_form.id,
                SurveyPersonYearRecord.survey_person_id == person.id,
                SurveyPersonYearRecord.school_year_id == batch.school_year_id,
            )
        )
        if record is None:
            record = SurveyPersonYearRecord(
                survey_form_id=survey_form.id,
                survey_person_id=person.id,
                school_year_id=batch.school_year_id,
            )
            db.add(record)
        record.school_id = school.id if school is not None else None
        record.class_id = None
        record.school_name_reported = row["school_name"] or (school.name if school is not None else None)
        record.class_name_reported = row["class_name"] or None
        record.learning_status = row["learning_status"]
        record.special_circumstances = row["special"] or None
        record.notes = _merge_notes(
            f"Cấp/nhóm: {row['level']}" if row["level"] else "",
            f"Địa bàn trường: {row['school_area']}" if row["school_area"] else "",
            f"Tình trạng cư trú theo file: {row['residency_text']}" if row["residency_text"] else "",
            "Cập nhật từ Excel " + job.job_code,
        )

    return {
        "households_created": households_created,
        "households_updated": households_updated,
        "people_created": people_created,
        "people_updated": people_updated,
        "households_total": len(touched_forms),
    }


def _process_xlsx_job(db: Session, job: HouseholdImportJob, batch, content: bytes) -> tuple[bool, str]:
    db.query(HouseholdImportError).filter(HouseholdImportError.job_id == job.id).delete(synchronize_session=False)
    try:
        rows, issues, metadata = _parse_xlsx_rows(content, batch)
    except Exception as error:
        job.status = "LOI_DU_LIEU"
        job.stage = "KIEM_TRA_NOI_DUNG"
        job.rows_total = 0
        job.rows_valid = 0
        job.rows_error = 1
        job.checked_at = datetime.now()
        job.message = f"Không đọc được nội dung file theo mẫu: {error}"
        _add_issue(db, job, severity="ERROR", code="READ_DATA_FAILED", message=job.message)
        db.commit()
        return False, job.message

    job.rows_total = len(rows)
    db_conflicts = _precheck_database_conflicts(db, rows, batch)
    issues.extend(db_conflicts)
    error_issues = [item for item in issues if item.get("severity", "ERROR") == "ERROR"]
    job.rows_error = len({item.get("row") for item in error_issues if item.get("row") is not None})
    job.rows_valid = max(0, len(rows) - job.rows_error)
    job.conflict_count = sum(1 for item in error_issues if "CONFLICT" in item.get("code", ""))
    job.checked_at = datetime.now()

    for item in issues:
        _add_issue_dict(db, job, item, metadata.get("sheet"))

    if error_issues:
        job.status = "CO_XUNG_DOT" if job.conflict_count else "LOI_DU_LIEU"
        job.stage = "KIEM_TRA_NOI_DUNG"
        job.message = (
            f"File có {len(error_issues)} lỗi. Dữ liệu hộ dân CHƯA được cập nhật. "
            "Hãy xem chi tiết từng dòng/cột, sửa file rồi gửi lại."
        )
        db.commit()
        return False, job.message

    job.status = "DANG_CAP_NHAT"
    job.stage = "CHO_GHI_CSDL"
    job.message = "File hợp lệ, đang chờ lượt ghi dữ liệu an toàn."
    db.commit()

    with _IMPORT_WRITE_LOCK:
        try:
            # Kiểm tra lại sau khi chờ khóa để bắt xung đột từ file vừa ghi trước đó.
            late_conflicts = _precheck_database_conflicts(db, rows, batch)
            if late_conflicts:
                for item in late_conflicts:
                    _add_issue_dict(db, job, item, metadata.get("sheet"))
                job.status = "CO_XUNG_DOT"
                job.stage = "KIEM_TRA_LAI_TRUOC_GHI"
                job.conflict_count = len(late_conflicts)
                job.message = "Phát hiện xung đột với dữ liệu vừa được cập nhật bởi lần gửi khác. File này chưa được ghi."
                db.commit()
                return False, job.message

            summary = _apply_rows(db, job, batch, rows)
            job.households_created = summary["households_created"]
            job.households_updated = summary["households_updated"]
            job.people_created = summary["people_created"]
            job.people_updated = summary["people_updated"]
            job.rows_valid = len(rows)
            job.rows_error = 0
            job.status = "DA_CAP_NHAT"
            job.stage = "HOAN_THANH"
            job.processed_at = datetime.now()
            job.message = (
                f"Đã cập nhật thành công {summary['households_total']} hộ / {len(rows)} thành viên. "
                f"Hộ mới: {summary['households_created']}; hộ cập nhật: {summary['households_updated']}; "
                f"thành viên mới: {summary['people_created']}; thành viên cập nhật: {summary['people_updated']}."
            )
            db.commit()
            return True, job.message
        except Exception as error:
            db.rollback()
            job = db.get(HouseholdImportJob, job.id)
            job.status = "LOI_CAP_NHAT"
            job.stage = "ROLLBACK"
            job.processed_at = datetime.now()
            job.message = (
                "Có lỗi trong khi ghi CSDL; toàn bộ thay đổi của file đã rollback. "
                f"Lỗi: {type(error).__name__}: {error}"
            )
            _add_issue(db, job, severity="ERROR", code="UPDATE_ROLLBACK", message=job.message)
            db.commit()
            return False, job.message


@router.get("/cap-nhat-ho-dan", response_class=HTMLResponse)
def household_update_center(
    request: Request,
    school_year_id: int | None = None,
    batch_id: int | None = None,
    db: Session = Depends(get_db),
):
    role_code = _role(request)
    if role_code not in VIEW_ROLE_CODES:
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    school_years = list(
        db.scalars(
            select(SchoolYear)
            .where(SchoolYear.is_active.is_(True))
            .order_by(SchoolYear.code.desc())
        ).all()
    )

    filters = tao_bo_loc_dot_theo_nguoi_dung(request)
    statement = (
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
        )
        .where(*filters)
        .order_by(
            SurveyBatch.school_year_id.desc(),
            SurveyBatch.created_at.desc(),
            SurveyBatch.id.desc(),
        )
    )
    if school_year_id is not None:
        statement = statement.where(
            SurveyBatch.school_year_id == school_year_id
        )
    batches = list(db.scalars(statement).all())

    batch_ids = [item.id for item in batches]
    selected_batch_id = (
        batch_id
        if batch_id in batch_ids
        else (batch_ids[0] if batch_ids else None)
    )

    jobs: list[HouseholdImportJob] = []
    if batch_ids:
        jobs = list(
            db.scalars(
                select(HouseholdImportJob)
                .where(
                    HouseholdImportJob.survey_batch_id.in_(batch_ids)
                )
                .order_by(
                    HouseholdImportJob.created_at.desc(),
                    HouseholdImportJob.id.desc(),
                )
                .limit(100)
            ).all()
        )

    can_upload = role_code in UPLOAD_ROLE_CODES

    return templates.TemplateResponse(
        request=request,
        name="surveys/household_update_center_v1.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "school_years": school_years,
            "selected_school_year_id": school_year_id,
            "selected_batch_id": selected_batch_id,
            "batches": batches,
            "jobs": jobs,
            "can_upload": can_upload,
            "scope_label": _scope_label(request),
            "batch_status_labels": BATCH_STATUS_LABELS,
            "import_status_labels": IMPORT_STATUS_LABELS,
            "import_status_classes": IMPORT_STATUS_CLASSES,
        },
    )


@router.post("/{batch_id}/cap-nhat-ho-dan/gui-file")
async def receive_household_update_file(
    batch_id: int,
    request: Request,
    excel_file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    role_code = _role(request)
    if role_code not in UPLOAD_ROLE_CODES:
        return _json_result(ok=False, job=None, message="Tài khoản không có quyền gửi file cập nhật hộ dân.", status_code=403)

    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return _json_result(ok=False, job=None, message="Đợt điều tra không thuộc phạm vi của tài khoản.", status_code=403)

    try:
        db.execute(text("PRAGMA busy_timeout=30000"))
    except Exception:
        pass

    original_name = excel_file.filename or "khong_ten"
    ext = Path(original_name).suffix.lower()
    content = await excel_file.read()
    user = _user(request)

    job = HouseholdImportJob(
        job_code=_new_job_code(),
        survey_batch_id=batch.id,
        school_year_id=batch.school_year_id,
        commune_id=batch.commune_id,
        uploaded_by_user_id=user.get("id"),
        actor_name_snapshot=user.get("full_name"),
        actor_role_snapshot=user.get("role_name") or user.get("role_code"),
        actor_unit_snapshot=user.get("unit_name"),
        original_file_name=original_name,
        file_ext=ext,
        file_size=len(content),
        status="DA_NHAN",
        stage="TIEP_NHAN",
        message="Máy chủ đã nhận file và cấp mã lần gửi.",
        received_at=datetime.now(),
    )
    db.add(job)
    db.commit()
    db.refresh(job)

    if not content:
        job.status = "LOI_FILE"; job.message = "File tải lên đang trống."; job.checked_at = datetime.now()
        _add_issue(db, job, severity="ERROR", code="EMPTY_FILE", message=job.message); db.commit()
        return _json_result(ok=False, job=job, message=job.message)

    if len(content) > 25 * 1024 * 1024:
        job.status = "LOI_FILE"; job.message = "File lớn quá 25 MB."; job.checked_at = datetime.now()
        _add_issue(db, job, severity="ERROR", code="FILE_TOO_LARGE", message=job.message); db.commit()
        return _json_result(ok=False, job=job, message=job.message)

    digest = hashlib.sha256(content).hexdigest()
    job.sha256 = digest
    duplicate = db.scalar(
        select(HouseholdImportJob).where(
            HouseholdImportJob.survey_batch_id == batch.id,
            HouseholdImportJob.sha256 == digest,
            HouseholdImportJob.id != job.id,
            HouseholdImportJob.status != "LOI_DU_LIEU",
            HouseholdImportJob.status != "CO_XUNG_DOT",
            HouseholdImportJob.status != "LOI_CAP_NHAT",
        ).order_by(HouseholdImportJob.created_at.desc())
    )
    if duplicate is not None:
        job.status = "TRUNG_FILE"; job.stage = "KIEM_TRA"; job.duplicate_of_job_code = duplicate.job_code
        job.message = f"Nội dung file đã được gửi trước ở lần {duplicate.job_code}. Hệ thống không cập nhật lặp."
        job.checked_at = datetime.now(); _add_issue(db, job, severity="WARNING", code="DUPLICATE_FILE", message=job.message); db.commit()
        return _json_result(ok=False, job=job, message=job.message)

    if ext not in {".xls", ".xlsx"}:
        job.status = "LOI_DINH_DANG"; job.message = "Chỉ chấp nhận .xls hoặc .xlsx."; job.checked_at = datetime.now()
        _add_issue(db, job, severity="ERROR", code="INVALID_EXTENSION", message=job.message); db.commit()
        return _json_result(ok=False, job=job, message=job.message)

    try:
        stored = _store_received_file(job.job_code, original_name, content)
        job.stored_path = str(stored.relative_to(PROJECT_DIR))
    except Exception as error:
        job.status = "LOI_FILE"; job.message = f"Máy chủ nhận file nhưng lưu bản an toàn thất bại: {error}"
        _add_issue(db, job, severity="ERROR", code="STORE_FAILED", message=job.message); db.commit()
        return _json_result(ok=False, job=job, message=job.message)

    if bool(getattr(batch, "is_locked", False)) or batch.status == "DA_KET_THUC":
        job.status = "LOI_DOT_KHOA"; job.message = "Đợt điều tra đang khóa/đã kết thúc; file đã lưu nhưng không cập nhật CSDL."; job.checked_at = datetime.now()
        _add_issue(db, job, severity="ERROR", code="BATCH_LOCKED", message=job.message); db.commit()
        return _json_result(ok=False, job=job, message=job.message)

    if ext == ".xls":
        job.status = "CHO_XU_LY_XLS"; job.stage = "CHO_BO_DOC_XLS"; job.has_macro = 1 if _xls_macro_hint(content) else 0
        job.checked_at = datetime.now(); job.message = "File .xls đã được nhận và lưu an toàn. V1B cập nhật trực tiếp .xlsx; bộ đọc .xls chính thức sẽ được bổ sung ở V1C, không chạy macro."
        db.commit(); return _json_result(ok=False, job=job, message=job.message)

    try:
        sheet_names, has_macro = _xlsx_structure(content)
        job.sheet_names = json.dumps(sheet_names, ensure_ascii=False)
        job.has_macro = 1 if has_macro else 0
    except Exception as error:
        job.status = "LOI_DINH_DANG"; job.checked_at = datetime.now(); job.message = f"File .xlsx không hợp lệ: {error}"
        _add_issue(db, job, severity="ERROR", code="STRUCTURE_ERROR", message=job.message); db.commit()
        return _json_result(ok=False, job=job, message=job.message)

    ok, message = _process_xlsx_job(db, job, batch, content)
    return _json_result(ok=ok, job=job, message=message)

def _job_in_scope(
    db: Session,
    request: Request,
    job_code: str,
) -> HouseholdImportJob | None:
    filters = tao_bo_loc_dot_theo_nguoi_dung(request)
    batch_ids = list(
        db.scalars(
            select(SurveyBatch.id).where(*filters)
        ).all()
    )
    if not batch_ids:
        return None
    return db.scalar(
        select(HouseholdImportJob).where(
            HouseholdImportJob.job_code == job_code,
            HouseholdImportJob.survey_batch_id.in_(batch_ids),
        )
    )


@router.get(
    "/cap-nhat-ho-dan/lan-gui/{job_code}",
    response_class=HTMLResponse,
)
def household_import_detail(
    job_code: str,
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) not in VIEW_ROLE_CODES:
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    job = _job_in_scope(db, request, job_code)
    if job is None:
        return RedirectResponse(
            url="/dieu-tra/cap-nhat-ho-dan",
            status_code=303,
        )

    batch = db.scalar(
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
        )
        .where(SurveyBatch.id == job.survey_batch_id)
    )
    errors = list(
        db.scalars(
            select(HouseholdImportError)
            .where(HouseholdImportError.job_id == job.id)
            .order_by(
                HouseholdImportError.id.asc()
            )
        ).all()
    )

    sheet_names = []
    if job.sheet_names:
        try:
            sheet_names = json.loads(job.sheet_names)
        except Exception:
            sheet_names = [job.sheet_names]

    return templates.TemplateResponse(
        request=request,
        name="surveys/household_import_detail_v1.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "job": job,
            "batch": batch,
            "errors": errors,
            "sheet_names": sheet_names,
            "status_label": IMPORT_STATUS_LABELS.get(
                job.status,
                job.status,
            ),
            "status_class": IMPORT_STATUS_CLASSES.get(
                job.status,
                "info",
            ),
            "can_delete": (
                _role(request) in UPLOAD_ROLE_CODES
                and job.status in DELETEABLE_IMPORT_STATUSES
            ),
        },
    )



# === BAI_13B_9_V1B_1_DELETE_ROUTE_START ===
@router.post("/cap-nhat-ho-dan/lan-gui/{job_code}/xoa")
def delete_household_import_job(
    job_code: str,
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) not in UPLOAD_ROLE_CODES:
        return RedirectResponse(
            url="/?status=forbidden",
            status_code=303,
        )

    job = _job_in_scope(
        db,
        request,
        job_code,
    )
    if job is None:
        return RedirectResponse(
            url="/dieu-tra/cap-nhat-ho-dan",
            status_code=303,
        )

    if job.status not in DELETEABLE_IMPORT_STATUSES:
        return RedirectResponse(
            url=(
                f"/dieu-tra/cap-nhat-ho-dan/lan-gui/{job.job_code}"
                "?delete_error=not_allowed"
            ),
            status_code=303,
        )

    batch_id = int(job.survey_batch_id)
    stored_path_value = (job.stored_path or "").strip()

    other_file_reference = None
    if stored_path_value:
        other_file_reference = db.scalar(
            select(HouseholdImportJob.id)
            .where(
                HouseholdImportJob.id != job.id,
                HouseholdImportJob.stored_path == stored_path_value,
            )
            .limit(1)
        )

    issue_rows = list(
        db.scalars(
            select(HouseholdImportError).where(
                HouseholdImportError.job_id == job.id
            )
        ).all()
    )
    for issue in issue_rows:
        db.delete(issue)

    db.delete(job)
    db.commit()

    if stored_path_value and other_file_reference is None:
        try:
            file_path = (
                PROJECT_DIR / Path(stored_path_value)
            ).resolve()
            import_root = IMPORT_ROOT.resolve()

            if (
                file_path == import_root
                or import_root in file_path.parents
            ):
                if file_path.is_file():
                    file_path.unlink()
        except Exception:
            pass

    return RedirectResponse(
        url=(
            "/dieu-tra/cap-nhat-ho-dan"
            f"?batch_id={batch_id}&deleted=1"
        ),
        status_code=303,
    )
# === BAI_13B_9_V1B_1_DELETE_ROUTE_END ===


@router.post("/cap-nhat-ho-dan/lan-gui/{job_code}/xu-ly")
def reprocess_household_import(
    job_code: str,
    request: Request,
    db: Session = Depends(get_db),
):
    if _role(request) not in UPLOAD_ROLE_CODES:
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    job = _job_in_scope(db, request, job_code)
    if job is None:
        return RedirectResponse(url="/dieu-tra/cap-nhat-ho-dan", status_code=303)
    batch = _batch_in_scope(db, request, job.survey_batch_id)
    if batch is None:
        return RedirectResponse(url="/?status=forbidden", status_code=303)
    if job.file_ext != ".xlsx" or not job.stored_path:
        return RedirectResponse(url=f"/dieu-tra/cap-nhat-ho-dan/lan-gui/{job_code}", status_code=303)
    path = PROJECT_DIR / job.stored_path
    if not path.exists():
        job.status = "LOI_FILE"; job.message = "Không còn bản file đã lưu trên máy chủ."; db.commit()
        return RedirectResponse(url=f"/dieu-tra/cap-nhat-ho-dan/lan-gui/{job_code}", status_code=303)
    _process_xlsx_job(db, job, batch, path.read_bytes())
    return RedirectResponse(url=f"/dieu-tra/cap-nhat-ho-dan/lan-gui/{job_code}", status_code=303)


@router.get("/{batch_id}/cap-nhat-ho-dan/tai-mau-xlsx")
def download_neutral_update_template(
    batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    batch = _batch_in_scope(db, request, batch_id)
    if batch is None or _role(request) not in VIEW_ROLE_CODES:
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    wb = Workbook()
    ws = wb.active
    ws.title = "MauNhapLieu"
    year_value = _batch_start_year(batch) or ""
    ws["A1"] = "Năm điều tra"; ws["B1"] = year_value
    ws["A2"] = "Xã/phường"; ws["B2"] = batch.commune.name
    ws["D1"] = "Mã xã/phường"; ws["E1"] = batch.commune.code
    ws["D2"] = "Ghi chú"; ws["E2"] = "Mẫu trắng không tô màu; năm lấy từ đợt đang chọn."

    headers = [
        "TT", "Họ đệm", "Tên", "Ngày", "Tháng", "Năm sinh", "Nữ",
        "Dân tộc", "Tôn giáo", "Diện ưu tiên", "Chủ hộ - Họ đệm", "Chủ hộ - Tên",
        "Địa chỉ: Số nhà, tên đường, tổ (nếu có)", "Số phiếu", "Diện cư trú",
        "Tình trạng cư trú", "Khối học", "Lớp học", "Tên trường", "Địa bàn trường",
        "Mã Trường", "Số định danh", "Căn cước công dân",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col).value = value
        ws.cell(5, col).value = col
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=105, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for cell in ws[4] + ws[5]:
        cell.font = Font(bold=True)
    widths = [6,18,12,7,8,10,6,10,12,16,18,12,34,13,14,18,12,12,28,20,16,18,18]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+idx) if idx <= 26 else "A"].width = width
    ws.freeze_panes = "A6"

    guide = wb.create_sheet("README")
    guide.append(["Nội dung", "Hướng dẫn"])
    guide.append(["Năm điều tra", "Tự lấy theo đợt đã chọn; không cố định năm 2025."])
    guide.append(["Xã/phường", batch.commune.name])
    guide.append(["Số phiếu", "Các dòng cùng một hộ dùng cùng Số phiếu."])
    guide.append(["Đa cấp", "Một hộ có thể có MN, TH, THCS, THPT, XMC."])
    guide.append(["Định danh", "Có thể để trống Số định danh/CCCD; hệ thống vẫn có mã nội bộ."])
    guide.append(["An toàn", "Không đổi tên các cột. File lỗi sẽ không cập nhật nửa chừng."])
    guide.column_dimensions["A"].width = 22; guide.column_dimensions["B"].width = 75
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    guide["A1"].font = Font(bold=True); guide["B1"].font = Font(bold=True)

    output = BytesIO(); wb.save(output); output.seek(0)
    safe_commune = re.sub(r"[^0-9A-Za-z]+", "_", _norm(batch.commune.name)).strip("_").lower()
    filename = f"mau_cap_nhat_ho_dan_{safe_commune}_{year_value}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/cap-nhat-ho-dan/mau-cap-nhat-goc")
def download_original_update_sample(
    request: Request,
):
    if _role(request) not in VIEW_ROLE_CODES:
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    path = REFERENCE_DIR / "pcgd_nhapphieudieutra_2025.xls"
    if not path.exists():
        return RedirectResponse(
            url="/dieu-tra/cap-nhat-ho-dan",
            status_code=303,
        )
    return FileResponse(
        path=str(path),
        filename="pcgd_nhapphieudieutra_2025.xls",
        media_type="application/vnd.ms-excel",
    )


@router.get("/cap-nhat-ho-dan/mau-phieu-dieu-tra")
def download_official_survey_form_sample(
    request: Request,
):
    if _role(request) not in VIEW_ROLE_CODES:
        return RedirectResponse(url="/?status=forbidden", status_code=303)

    path = (
        REFERENCE_DIR
        / "mau-phieu-dieu-tra-pho-cap-giao-duc-xoa-mu-chu.xlsx"
    )
    if not path.exists():
        return RedirectResponse(
            url="/dieu-tra/cap-nhat-ho-dan",
            status_code=303,
        )
    return FileResponse(
        path=str(path),
        filename="mau-phieu-dieu-tra-pho-cap-giao-duc-xoa-mu-chu.xlsx",
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


def _age_on(value, reference_date: date) -> int | None:
    if value is None:
        return None
    return (
        reference_date.year
        - value.year
        - (
            (reference_date.month, reference_date.day)
            < (value.month, value.day)
        )
    )


def _normalize_no_accent(value: str) -> str:
    import unicodedata

    text_value = unicodedata.normalize(
        "NFD",
        str(value or "").lower(),
    )
    text_value = "".join(
        char
        for char in text_value
        if unicodedata.category(char) != "Mn"
    )
    return re.sub(r"\s+", " ", text_value).strip()


def _level_from_record(
    record: SurveyPersonYearRecord | None,
    school: School | None,
    classroom: Classroom | None,
) -> str:
    if record is None:
        return "CHUA_XAC_DINH"

    school_name = (
        record.school_name_reported
        or (school.name if school is not None else "")
        or ""
    )
    class_name = (
        record.class_name_reported
        or (classroom.name if classroom is not None else "")
        or ""
    )
    key = _normalize_no_accent(f"{school_name} {class_name}")

    if any(x in key for x in ("mam non", "mau giao", "nha tre")):
        return "MN"
    if any(x in key for x in ("trung hoc pho thong", "thpt")):
        return "THPT"
    if any(x in key for x in ("trung hoc co so", "thcs")):
        return "THCS"
    if "tieu hoc" in key:
        return "TH"

    match = re.search(r"(?:lop\s*)?(\d{1,2})(?:\D|$)", key)
    if match:
        grade = int(match.group(1))
        if 1 <= grade <= 5:
            return "TH"
        if 6 <= grade <= 9:
            return "THCS"
        if 10 <= grade <= 12:
            return "THPT"

    if record.learning_status in {
        "CHUA_DI_HOC",
        "BO_HOC",
        "THOI_HOC",
    }:
        return "NGOAI_NHA_TRUONG"

    return "CHUA_XAC_DINH"


@router.get(
    "/{batch_id}/cap-nhat-ho-dan/thong-ke",
    response_class=HTMLResponse,
)
def household_related_statistics(
    batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    batch = _batch_in_scope(db, request, batch_id)
    if batch is None:
        return RedirectResponse(
            url="/dieu-tra/cap-nhat-ho-dan",
            status_code=303,
        )

    scope_filters = tao_bo_loc_bao_cao_theo_nguoi_dung(request)
    forms = list(
        db.scalars(
            select(SurveyForm)
            .join(
                Household,
                Household.id == SurveyForm.household_id,
            )
            .options(
                selectinload(SurveyForm.household)
                .selectinload(Household.people),
            )
            .where(
                SurveyForm.survey_batch_id == batch.id,
                *scope_filters,
            )
        ).all()
    )

    people = [
        person
        for form in forms
        for person in form.household.people
        if person.is_active
    ]
    form_ids = [item.id for item in forms]
    person_ids = [item.id for item in people]

    records = []
    if form_ids and person_ids:
        records = list(
            db.scalars(
                select(SurveyPersonYearRecord).where(
                    SurveyPersonYearRecord.survey_form_id.in_(form_ids),
                    SurveyPersonYearRecord.survey_person_id.in_(person_ids),
                    SurveyPersonYearRecord.school_year_id
                    == batch.school_year_id,
                )
            ).all()
        )

    record_by_person = {
        item.survey_person_id: item
        for item in records
    }

    school_ids = {
        item.school_id
        for item in records
        if item.school_id is not None
    }
    class_ids = {
        item.class_id
        for item in records
        if item.class_id is not None
    }
    schools = {
        item.id: item
        for item in db.scalars(
            select(School).where(School.id.in_(school_ids))
        ).all()
    } if school_ids else {}
    classrooms = {
        item.id: item
        for item in db.scalars(
            select(Classroom).where(Classroom.id.in_(class_ids))
        ).all()
    } if class_ids else {}

    reference_date = batch.start_date or date.today()
    total_pcgd = 0
    total_xmc = 0
    total_age_18 = 0
    missing_personal_id = 0
    missing_citizen_id = 0
    missing_year_record = 0

    levels = {
        "MN": 0,
        "TH": 0,
        "THCS": 0,
        "THPT": 0,
        "NGOAI_NHA_TRUONG": 0,
        "CHUA_XAC_DINH": 0,
    }
    learning = {
        code: 0
        for code in LEARNING_STATUS_LABELS
    }
    residency = {
        code: 0
        for code in RESIDENCY_STATUS_LABELS
    }
    form_status = {
        code: 0
        for code in FORM_STATUS_LABELS
    }

    for form in forms:
        form_status[form.status] = form_status.get(form.status, 0) + 1

    for person in people:
        age = _age_on(person.date_of_birth, reference_date)
        if age is not None and 0 <= age <= 18:
            total_pcgd += 1
        if age is not None and 18 <= age <= 60:
            total_xmc += 1
        if age == 18:
            total_age_18 += 1

        if not (person.personal_id or "").strip():
            missing_personal_id += 1
        if not (getattr(person, "citizen_id", None) or "").strip():
            missing_citizen_id += 1

        residency[person.residency_status] = (
            residency.get(person.residency_status, 0) + 1
        )

        record = record_by_person.get(person.id)
        if record is None:
            missing_year_record += 1
            levels["CHUA_XAC_DINH"] += 1
            continue

        learning[record.learning_status] = (
            learning.get(record.learning_status, 0) + 1
        )
        level = _level_from_record(
            record,
            schools.get(record.school_id),
            classrooms.get(record.class_id),
        )
        levels[level] = levels.get(level, 0) + 1

    return templates.TemplateResponse(
        request=request,
        name="surveys/household_statistics_v1.html",
        context={
            "nguoi_dung": request.scope.get("auth_user"),
            "batch": batch,
            "reference_date": reference_date,
            "scope_label": _scope_label(request),
            "total_households": len(forms),
            "total_people": len(people),
            "total_pcgd": total_pcgd,
            "total_xmc": total_xmc,
            "total_age_18": total_age_18,
            "missing_personal_id": missing_personal_id,
            "missing_citizen_id": missing_citizen_id,
            "missing_year_record": missing_year_record,
            "level_counts": levels,
            "level_labels": {
                "MN": "Mầm non",
                "TH": "Tiểu học",
                "THCS": "THCS",
                "THPT": "THPT",
                "NGOAI_NHA_TRUONG": "Ngoài nhà trường",
                "CHUA_XAC_DINH": "Chưa xác định",
            },
            "learning_counts": learning,
            "learning_labels": LEARNING_STATUS_LABELS,
            "residency_counts": residency,
            "residency_labels": RESIDENCY_STATUS_LABELS,
            "form_status_counts": form_status,
            "form_status_labels": FORM_STATUS_LABELS,
        },
    )
