from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.database import get_db
from app.models import Classroom, School
from app.permissions import (
    COMMUNE_ROLE_CODE,
    SCHOOL_ROLE_CODE,
    TEACHER_ROLE_CODE,
    normalize_role_code,
)
from app.survey_models import (
    Household,
    SurveyBatch,
    SurveyForm,
    SurveyPerson,
    SurveyPersonYearRecord,
)
from app.routers.surveys import (
    BOOLEAN_YEAR_FIELDS,
    LEARNING_STATUS_LABELS,
    lay_thong_tin_nguoi_dung,
    tao_bo_loc_dot_theo_nguoi_dung,
    tao_bo_loc_phieu_theo_nguoi_dung,
)


APP_DIR = Path(__file__).resolve().parent.parent

templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

router = APIRouter(
    prefix="/dieu-tra",
    tags=["Đối chiếu dữ liệu điều tra"],
)


HOUSEHOLD_CHANGE_LABELS = {
    "head_name": "Tên chủ hộ",
    "address": "Địa chỉ",
    "hamlet": "Thôn/xóm/khối/bản",
}

PERSON_CHANGE_LABELS = {
    "learning_status": "Trạng thái học tập",
    "school": "Trường học",
    "classroom": "Lớp học",
    "indicators": "Chỉ báo năm học",
}


# =========================================================
# HÀM DÙNG CHUNG
# =========================================================


def _school_year_key(code: str | None) -> tuple[int, ...]:
    """Chuyển mã 2026-2027 thành khóa số để sắp xếp an toàn."""

    if not code:
        return (0,)

    values: list[int] = []
    current = ""

    for char in str(code):
        if char.isdigit():
            current += char
        elif current:
            values.append(int(current))
            current = ""

    if current:
        values.append(int(current))

    return tuple(values) or (0,)


def _batch_in_scope(
    db: Session,
    request: Request,
    batch_id: int,
) -> SurveyBatch | None:
    """Lấy đợt điều tra và áp dụng đúng phạm vi tài khoản."""

    filters = tao_bo_loc_dot_theo_nguoi_dung(request)

    return db.scalar(
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
        )
        .where(
            SurveyBatch.id == batch_id,
            *filters,
        )
    )


def _candidate_source_batches(
    db: Session,
    request: Request,
    target_batch: SurveyBatch,
) -> list[SurveyBatch]:
    """Tìm các đợt cũ cùng xã/phường mà tài khoản được quyền xem."""

    filters = tao_bo_loc_dot_theo_nguoi_dung(request)
    candidates = db.scalars(
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
        )
        .where(
            SurveyBatch.commune_id == target_batch.commune_id,
            SurveyBatch.id != target_batch.id,
            *filters,
        )
    ).all()

    target_key = _school_year_key(target_batch.school_year.code)
    previous = [
        item
        for item in candidates
        if _school_year_key(item.school_year.code) < target_key
    ]

    previous.sort(
        key=lambda item: (
            _school_year_key(item.school_year.code),
            item.start_date or item.created_at.date(),
            item.id,
        ),
        reverse=True,
    )

    return previous


def _load_forms(
    db: Session,
    request: Request,
    batch_id: int,
) -> list[SurveyForm]:
    """Nạp phiếu, hộ và hồ sơ năm học trong phạm vi tài khoản."""

    form_filters = tao_bo_loc_phieu_theo_nguoi_dung(request)

    return db.scalars(
        select(SurveyForm)
        .options(
            selectinload(SurveyForm.household),
            selectinload(SurveyForm.person_year_records)
            .selectinload(SurveyPersonYearRecord.survey_person),
            selectinload(SurveyForm.person_year_records)
            .selectinload(SurveyPersonYearRecord.school),
            selectinload(SurveyForm.person_year_records)
            .selectinload(SurveyPersonYearRecord.classroom),
        )
        .where(
            SurveyForm.survey_batch_id == batch_id,
            *form_filters,
        )
        .order_by(SurveyForm.id.asc())
    ).all()


def _scope_label(request: Request) -> str:
    user = lay_thong_tin_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))

    if role_code in {SCHOOL_ROLE_CODE, TEACHER_ROLE_CODE}:
        return "PHẠM VI TOÀN TRƯỜNG"

    if role_code == COMMUNE_ROLE_CODE:
        return "PHẠM VI TOÀN XÃ/PHƯỜNG"

    return "PHẠM VI TOÀN TỈNH"


def _clean_text(value: Any) -> str:
    return " ".join(str(value or "").split())


def _school_name(record: SurveyPersonYearRecord | None) -> str:
    if record is None:
        return "—"

    if record.school is not None:
        return record.school.name

    return _clean_text(record.school_name_reported) or "Chưa xác định"


def _class_name(record: SurveyPersonYearRecord | None) -> str:
    if record is None:
        return "—"

    if record.classroom is not None:
        return record.classroom.name

    return _clean_text(record.class_name_reported) or "Chưa xác định"


def _learning_label(record: SurveyPersonYearRecord | None) -> str:
    if record is None:
        return "—"

    code = str(record.learning_status or "CHUA_XAC_DINH")
    return LEARNING_STATUS_LABELS.get(code, code)


def _answered_indicator_total(
    record: SurveyPersonYearRecord | None,
) -> int:
    if record is None:
        return 0

    return sum(
        getattr(record, field_name) is not None
        for field_name in BOOLEAN_YEAR_FIELDS
    )


def _indicator_change_total(
    source: SurveyPersonYearRecord | None,
    target: SurveyPersonYearRecord | None,
) -> int:
    if source is None or target is None:
        return 0

    return sum(
        getattr(source, field_name) != getattr(target, field_name)
        for field_name in BOOLEAN_YEAR_FIELDS
    )


def _target_record_needs_review(
    record: SurveyPersonYearRecord | None,
) -> tuple[bool, list[str]]:
    reasons: list[str] = []

    if record is None:
        return True, ["Chưa có hồ sơ năm học mới"]

    learning_status = str(record.learning_status or "CHUA_XAC_DINH")

    if learning_status == "CHUA_XAC_DINH":
        reasons.append("Chưa xác định trạng thái học tập")

    if learning_status == "DANG_HOC" and _school_name(record) == "Chưa xác định":
        reasons.append("Đang học nhưng chưa xác định trường")

    answered = _answered_indicator_total(record)
    if answered < len(BOOLEAN_YEAR_FIELDS):
        reasons.append(
            f"Mới nhập {answered}/{len(BOOLEAN_YEAR_FIELDS)} chỉ báo"
        )

    return bool(reasons), reasons


def _household_name(form: SurveyForm | None) -> str:
    if form is None:
        return "—"

    return _clean_text(form.head_name_snapshot) or (
        form.household.head_name if form.household else "—"
    )


def _hamlet_name(form: SurveyForm | None) -> str:
    if form is None:
        return "—"

    return _clean_text(form.hamlet_name_snapshot) or (
        _clean_text(form.household.hamlet_name)
        if form.household
        else "—"
    ) or "—"


def _build_comparison(
    source_batch: SurveyBatch,
    target_batch: SurveyBatch,
    source_forms: list[SurveyForm],
    target_forms: list[SurveyForm],
) -> dict[str, Any]:
    source_form_map = {item.household_id: item for item in source_forms}
    target_form_map = {item.household_id: item for item in target_forms}

    source_household_ids = set(source_form_map)
    target_household_ids = set(target_form_map)

    household_rows: list[dict[str, Any]] = []

    for household_id in sorted(source_household_ids | target_household_ids):
        source_form = source_form_map.get(household_id)
        target_form = target_form_map.get(household_id)

        if source_form and target_form:
            category_code = "TIEP_TUC"
            category_label = "Tiếp tục theo dõi"
        elif target_form:
            category_code = "HO_MOI"
            category_label = "Hộ mới trong năm học đích"
        else:
            category_code = "KHONG_CON"
            category_label = "Không còn trong đợt mới"

        changes: list[str] = []

        if source_form and target_form:
            if _clean_text(source_form.head_name_snapshot).casefold() != _clean_text(
                target_form.head_name_snapshot
            ).casefold():
                changes.append(HOUSEHOLD_CHANGE_LABELS["head_name"])

            if _clean_text(source_form.address_snapshot).casefold() != _clean_text(
                target_form.address_snapshot
            ).casefold():
                changes.append(HOUSEHOLD_CHANGE_LABELS["address"])

            if _clean_text(source_form.hamlet_name_snapshot).casefold() != _clean_text(
                target_form.hamlet_name_snapshot
            ).casefold():
                changes.append(HOUSEHOLD_CHANGE_LABELS["hamlet"])

        reference_form = target_form or source_form
        household_rows.append(
            {
                "household_id": household_id,
                "household_code": (
                    reference_form.household.code
                    if reference_form and reference_form.household
                    else "—"
                ),
                "head_name": _household_name(reference_form),
                "hamlet_name": _hamlet_name(reference_form),
                "source_form_number": (
                    source_form.form_number if source_form else "—"
                ),
                "target_form_number": (
                    target_form.form_number if target_form else "—"
                ),
                "category_code": category_code,
                "category_label": category_label,
                "changes": changes,
                "change_text": ", ".join(changes) or "Không thay đổi ảnh chụp hộ",
            }
        )

    household_rows.sort(
        key=lambda row: (
            str(row["hamlet_name"]).casefold(),
            str(row["head_name"]).casefold(),
        )
    )

    source_record_map: dict[int, SurveyPersonYearRecord] = {}
    target_record_map: dict[int, SurveyPersonYearRecord] = {}
    source_record_form_map: dict[int, SurveyForm] = {}
    target_record_form_map: dict[int, SurveyForm] = {}

    for form in source_forms:
        for record in form.person_year_records:
            source_record_map[record.survey_person_id] = record
            source_record_form_map[record.survey_person_id] = form

    for form in target_forms:
        for record in form.person_year_records:
            target_record_map[record.survey_person_id] = record
            target_record_form_map[record.survey_person_id] = form

    source_person_ids = set(source_record_map)
    target_person_ids = set(target_record_map)

    person_rows: list[dict[str, Any]] = []

    for person_id in sorted(source_person_ids | target_person_ids):
        source_record = source_record_map.get(person_id)
        target_record = target_record_map.get(person_id)
        person = (
            target_record.survey_person
            if target_record is not None
            else source_record.survey_person
        )

        source_form = source_record_form_map.get(person_id)
        target_form = target_record_form_map.get(person_id)
        reference_form = target_form or source_form

        if source_record and target_record:
            category_code = "TIEP_TUC"
            category_label = "Tiếp tục theo dõi"
        elif target_record:
            category_code = "DOI_TUONG_MOI"
            category_label = "Đối tượng mới"
        else:
            category_code = "KHONG_CON"
            category_label = "Không còn trong năm học mới"

        changes: list[str] = []
        indicator_change_total = 0

        if source_record and target_record:
            if source_record.learning_status != target_record.learning_status:
                changes.append(PERSON_CHANGE_LABELS["learning_status"])

            if _school_name(source_record).casefold() != _school_name(
                target_record
            ).casefold():
                changes.append(PERSON_CHANGE_LABELS["school"])

            if _class_name(source_record).casefold() != _class_name(
                target_record
            ).casefold():
                changes.append(PERSON_CHANGE_LABELS["classroom"])

            indicator_change_total = _indicator_change_total(
                source_record,
                target_record,
            )
            if indicator_change_total:
                changes.append(
                    f"{PERSON_CHANGE_LABELS['indicators']} ({indicator_change_total})"
                )

        needs_review, review_reasons = _target_record_needs_review(target_record)
        target_answered = _answered_indicator_total(target_record)

        person_rows.append(
            {
                "person_id": person_id,
                "person_code": person.code,
                "full_name": person.full_name,
                "date_of_birth": person.date_of_birth,
                "gender": person.gender or "—",
                "is_active": bool(person.is_active),
                "household_code": (
                    reference_form.household.code
                    if reference_form and reference_form.household
                    else "—"
                ),
                "head_name": _household_name(reference_form),
                "hamlet_name": _hamlet_name(reference_form),
                "category_code": category_code,
                "category_label": category_label,
                "source_learning": _learning_label(source_record),
                "target_learning": _learning_label(target_record),
                "source_school": _school_name(source_record),
                "target_school": _school_name(target_record),
                "source_class": _class_name(source_record),
                "target_class": _class_name(target_record),
                "changes": changes,
                "change_text": ", ".join(changes) or "Không phát hiện thay đổi",
                "indicator_change_total": indicator_change_total,
                "target_answered": target_answered,
                "target_indicator_total": len(BOOLEAN_YEAR_FIELDS),
                "needs_review": needs_review,
                "review_reasons": review_reasons,
                "review_text": "; ".join(review_reasons) or "Đã nhập đủ dữ liệu chính",
                "target_household_id": target_form.household_id if target_form else None,
            }
        )

    person_rows.sort(
        key=lambda row: (
            str(row["hamlet_name"]).casefold(),
            str(row["head_name"]).casefold(),
            str(row["full_name"]).casefold(),
        )
    )

    continuing_people = [
        row for row in person_rows if row["category_code"] == "TIEP_TUC"
    ]
    changed_people = [row for row in continuing_people if row["changes"]]
    review_rows = [row for row in person_rows if row["needs_review"]]

    summary = {
        "source_households": len(source_household_ids),
        "target_households": len(target_household_ids),
        "continuing_households": len(source_household_ids & target_household_ids),
        "new_households": len(target_household_ids - source_household_ids),
        "missing_households": len(source_household_ids - target_household_ids),
        "changed_households": sum(bool(row["changes"]) for row in household_rows),
        "source_people": len(source_person_ids),
        "target_people": len(target_person_ids),
        "continuing_people": len(source_person_ids & target_person_ids),
        "new_people": len(target_person_ids - source_person_ids),
        "missing_people": len(source_person_ids - target_person_ids),
        "changed_people": len(changed_people),
        "review_people": len(review_rows),
    }

    return {
        "source_batch": source_batch,
        "target_batch": target_batch,
        "summary": summary,
        "household_rows": household_rows,
        "person_rows": person_rows,
        "changed_people": changed_people,
        "review_rows": review_rows,
    }


# =========================================================
# EXCEL
# =========================================================


def _style_sheet(ws, widths: list[int]) -> None:
    thin = Side(style="thin", color="B8C7D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1976D2")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _build_excel(
    request: Request,
    comparison: dict[str, Any],
) -> BytesIO:
    source_batch: SurveyBatch = comparison["source_batch"]
    target_batch: SurveyBatch = comparison["target_batch"]
    summary = comparison["summary"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Tong quan"

    _write_header(ws, ["Thông tin", "Giá trị"])
    overview_rows = [
        ("Phạm vi", _scope_label(request)),
        ("Xã/phường", target_batch.commune.name),
        ("Đợt nguồn", source_batch.code),
        ("Năm học nguồn", source_batch.school_year.code),
        ("Đợt đích", target_batch.code),
        ("Năm học đích", target_batch.school_year.code),
        ("Trạng thái đợt đích", "Đã khóa" if target_batch.is_locked else "Đang cập nhật"),
        ("Hộ nguồn", summary["source_households"]),
        ("Hộ đích", summary["target_households"]),
        ("Hộ tiếp tục", summary["continuing_households"]),
        ("Hộ mới", summary["new_households"]),
        ("Hộ không còn", summary["missing_households"]),
        ("Đối tượng nguồn", summary["source_people"]),
        ("Đối tượng đích", summary["target_people"]),
        ("Đối tượng mới", summary["new_people"]),
        ("Đối tượng không còn", summary["missing_people"]),
        ("Đối tượng có thay đổi", summary["changed_people"]),
        ("Hồ sơ năm mới cần rà soát", summary["review_people"]),
    ]
    for row in overview_rows:
        ws.append(list(row))
    _style_sheet(ws, [32, 48])

    ws_households = wb.create_sheet("Bien dong ho")
    _write_header(
        ws_households,
        [
            "STT",
            "Mã hộ",
            "Chủ hộ",
            "Thôn/xóm/khối/bản",
            "Phiếu nguồn",
            "Phiếu đích",
            "Nhóm biến động",
            "Nội dung thay đổi",
        ],
    )
    for index, row in enumerate(comparison["household_rows"], start=1):
        ws_households.append(
            [
                index,
                row["household_code"],
                row["head_name"],
                row["hamlet_name"],
                row["source_form_number"],
                row["target_form_number"],
                row["category_label"],
                row["change_text"],
            ]
        )
    _style_sheet(ws_households, [7, 20, 24, 24, 24, 24, 28, 42])

    ws_people = wb.create_sheet("Bien dong doi tuong")
    _write_header(
        ws_people,
        [
            "STT",
            "Mã đối tượng",
            "Họ và tên",
            "Ngày sinh",
            "Mã hộ",
            "Chủ hộ",
            "Nhóm biến động",
            "Trạng thái nguồn",
            "Trạng thái đích",
            "Trường nguồn",
            "Trường đích",
            "Lớp nguồn",
            "Lớp đích",
            "Nội dung thay đổi",
            "Chỉ báo đã nhập",
            "Kết quả rà soát",
        ],
    )
    for index, row in enumerate(comparison["person_rows"], start=1):
        ws_people.append(
            [
                index,
                row["person_code"],
                row["full_name"],
                row["date_of_birth"].strftime("%d/%m/%Y") if row["date_of_birth"] else "",
                row["household_code"],
                row["head_name"],
                row["category_label"],
                row["source_learning"],
                row["target_learning"],
                row["source_school"],
                row["target_school"],
                row["source_class"],
                row["target_class"],
                row["change_text"],
                f"{row['target_answered']}/{row['target_indicator_total']}",
                row["review_text"],
            ]
        )
    _style_sheet(
        ws_people,
        [7, 20, 24, 14, 18, 24, 26, 22, 22, 28, 28, 20, 20, 42, 16, 46],
    )

    ws_changes = wb.create_sheet("Thay doi hoc tap")
    _write_header(
        ws_changes,
        [
            "STT",
            "Mã đối tượng",
            "Họ và tên",
            "Trạng thái nguồn",
            "Trạng thái đích",
            "Trường nguồn",
            "Trường đích",
            "Lớp nguồn",
            "Lớp đích",
            "Nội dung thay đổi",
        ],
    )
    for index, row in enumerate(comparison["changed_people"], start=1):
        ws_changes.append(
            [
                index,
                row["person_code"],
                row["full_name"],
                row["source_learning"],
                row["target_learning"],
                row["source_school"],
                row["target_school"],
                row["source_class"],
                row["target_class"],
                row["change_text"],
            ]
        )
    _style_sheet(ws_changes, [7, 20, 25, 22, 22, 28, 28, 20, 20, 46])

    ws_review = wb.create_sheet("Can ra soat nam moi")
    _write_header(
        ws_review,
        [
            "STT",
            "Mã đối tượng",
            "Họ và tên",
            "Mã hộ",
            "Chủ hộ",
            "Trạng thái đích",
            "Trường đích",
            "Lớp đích",
            "Chỉ báo đã nhập",
            "Nội dung cần rà soát",
        ],
    )
    for index, row in enumerate(comparison["review_rows"], start=1):
        ws_review.append(
            [
                index,
                row["person_code"],
                row["full_name"],
                row["household_code"],
                row["head_name"],
                row["target_learning"],
                row["target_school"],
                row["target_class"],
                f"{row['target_answered']}/{row['target_indicator_total']}",
                row["review_text"],
            ]
        )
    _style_sheet(ws_review, [7, 20, 25, 18, 25, 22, 28, 20, 16, 50])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# =========================================================
# ROUTES
# =========================================================


@router.get(
    "/{target_batch_id}/doi-chieu-nam-hoc",
    response_class=HTMLResponse,
)
def compare_school_years(
    target_batch_id: int,
    request: Request,
    source_batch_id: int = 0,
    db: Session = Depends(get_db),
):
    target_batch = _batch_in_scope(db, request, target_batch_id)

    if target_batch is None:
        return RedirectResponse(url="/dieu-tra?status=forbidden", status_code=303)

    source_candidates = _candidate_source_batches(db, request, target_batch)

    selected_source: SurveyBatch | None = None

    if source_batch_id > 0:
        selected_source = next(
            (item for item in source_candidates if item.id == source_batch_id),
            None,
        )

    if selected_source is None and source_candidates:
        selected_source = source_candidates[0]

    comparison: dict[str, Any] | None = None

    if selected_source is not None:
        source_forms = _load_forms(db, request, selected_source.id)
        target_forms = _load_forms(db, request, target_batch.id)
        comparison = _build_comparison(
            selected_source,
            target_batch,
            source_forms,
            target_forms,
        )

    return templates.TemplateResponse(
        request=request,
        name="surveys/year_comparison.html",
        context={
            "nguoi_dung": lay_thong_tin_nguoi_dung(request),
            "scope_label": _scope_label(request),
            "target_batch": target_batch,
            "source_candidates": source_candidates,
            "selected_source": selected_source,
            "comparison": comparison,
            "learning_status_labels": LEARNING_STATUS_LABELS,
        },
    )


@router.get(
    "/{target_batch_id}/doi-chieu-nam-hoc/xuat-excel",
)
def export_school_year_comparison(
    target_batch_id: int,
    request: Request,
    source_batch_id: int,
    db: Session = Depends(get_db),
):
    target_batch = _batch_in_scope(db, request, target_batch_id)
    source_batch = _batch_in_scope(db, request, source_batch_id)

    if target_batch is None or source_batch is None:
        return RedirectResponse(url="/dieu-tra?status=forbidden", status_code=303)

    if target_batch.commune_id != source_batch.commune_id:
        return RedirectResponse(
            url=f"/dieu-tra/{target_batch_id}/doi-chieu-nam-hoc",
            status_code=303,
        )

    if _school_year_key(source_batch.school_year.code) >= _school_year_key(
        target_batch.school_year.code
    ):
        return RedirectResponse(
            url=f"/dieu-tra/{target_batch_id}/doi-chieu-nam-hoc",
            status_code=303,
        )

    source_forms = _load_forms(db, request, source_batch.id)
    target_forms = _load_forms(db, request, target_batch.id)
    comparison = _build_comparison(
        source_batch,
        target_batch,
        source_forms,
        target_forms,
    )
    buffer = _build_excel(request, comparison)

    filename = (
        "doi_chieu_nam_hoc_"
        f"{source_batch.school_year.code}_"
        f"{target_batch.school_year.code}_"
        f"{target_batch.commune.code}.xlsx"
    )

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )
