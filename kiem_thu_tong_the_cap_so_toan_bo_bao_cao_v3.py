# -*- coding: utf-8 -*-
"""
KIỂM THỬ TỔNG THỂ HỆ THỐNG BÁO CÁO - CẤP SỞ - V3

KHÔNG GẮN SỐ BÀI MỚI.

PHẠM VI
-------
Tài khoản kiểm thử mang vai trò SO, phạm vi Toàn tỉnh.
Kiểm thử 18 đầu ra Excel hiện có:

MẦM NON (5)
  1. MN-M1
  2. MN-02
  3. MN-01-GV
  4. MN-01-CSVC
  5. MN-TÀI CHÍNH

TIỂU HỌC (4)
  6. TH-M1
  7. TH-02
  8. TH-01-GV
  9. TH-01-CSVC

THCS (5)
 10. THCS-M1
 11. THCS-M2
 12. THCS-TK
 13. THCS-M5
 14. THCS-CSVC

XÓA MÙ CHỮ (4)
 15. XMC-3
 16. CMC-2
 17. CMC-1
 18. XMC-4

NGUYÊN TẮC AN TOÀN
------------------
- Không sửa source.
- Không UPDATE / INSERT / DELETE database.
- SQLAlchemy connection bật PRAGMA query_only = ON.
- Kiểm tra integrity_check + foreign_key_check.
- SHA256 phocap.db trước/sau phải giữ nguyên.
- Nếu một biểu lỗi, vẫn tiếp tục kiểm các biểu còn lại.

MỨC KẾT QUẢ
-----------
PASS : đường xuất chạy được và các kiểm tra bắt buộc đạt.
WARN : file xuất được nhưng có điểm trình bày/phạm vi cần xem bằng mắt.
FAIL : lỗi runtime, file XLSX hỏng, lỗi Excel, kết luận sai kiểu,
       hoặc định dạng phần trăm đã khóa bị mất.
"""

from __future__ import annotations

import asyncio
import ast
import hashlib
import inspect
import os
import re
import sqlite3
import sys
import traceback
import unicodedata
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any


PROJECT = Path(
    os.environ.get("PHOCAP_PROJECT", r"C:\PhoCap")
).resolve()

APP = PROJECT / "app"
DB_PATH = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"

MN_BUILDER_PATH = APP / "pcgd_mn_report_builders_v1.py"
XMC_BUILDER_PATH = APP / "pcgd_xmc_report_builders_v1.py"
REPORT_CENTER_PATH = APP / "routers" / "report_center.py"
RULES_PATH = APP / "services" / "pcgd_business_rules.py"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

OUT_DIR = EXPORTS / f"kiem_thu_tong_the_cap_so_{STAMP}"
REPORT_PATH = OUT_DIR / f"bao_cao_kiem_thu_tong_the_cap_so_{STAMP}.txt"
SUMMARY_CSV = OUT_DIR / f"tong_hop_kiem_thu_cap_so_{STAMP}.csv"

EXPECTED_COMMUNES = 130
EXPECTED_SPECIAL = 51
EXPECTED_PERCENT_FORMAT = r"0\%"

EXCEL_ERROR_VALUES = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
}


REPORT_SPECS = (
    # ------------------------------------------------------------------
    # MẦM NON - export_mn_report
    # ------------------------------------------------------------------
    {
        "group": "MẦM NON",
        "short": "MN-M1",
        "code": "PCGD_MN_M1_2025",
        "exporter": "mn",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "MẦM NON",
        "short": "MN-02",
        "code": "PCGD_MN_02_2025",
        "exporter": "mn",
        "conclusion": {
            "R7": {"Đạt", "Không đạt", "Chưa đủ dữ liệu"},
        },
        "percent_formats": ("H7", "K7", "O7"),
    },
    {
        "group": "MẦM NON",
        "short": "MN-01-GV",
        "code": "PCGD_MN_01_GV_2025",
        "exporter": "mn",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "MẦM NON",
        "short": "MN-01-CSVC",
        "code": "PCGD_MN_01_CSVC_2025",
        "exporter": "mn",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "MẦM NON",
        "short": "MN-TÀI-CHÍNH",
        "code": "PCGD_MN_TAICHINH_2025",
        "exporter": "mn",
        "conclusion": {},
        "percent_formats": (),
    },

    # ------------------------------------------------------------------
    # TIỂU HỌC
    # TH-M1 dùng exporter riêng trong report_center.py.
    # ------------------------------------------------------------------
    {
        "group": "TIỂU HỌC",
        "short": "TH-M1",
        "code": "PCGD_TH_M1_2025",
        "exporter": "th_m1",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "TIỂU HỌC",
        "short": "TH-02",
        "code": "PCGD_TH_02_2025",
        "exporter": "xmc",
        "conclusion": {
            "T8": {1, 2, 3, "Không đạt", "Chưa đủ dữ liệu"},
        },
        "percent_formats": ("I8", "K8", "M8", "Q8"),
    },
    {
        "group": "TIỂU HỌC",
        "short": "TH-01-GV",
        "code": "PCGD_TH_01_GV_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "TIỂU HỌC",
        "short": "TH-01-CSVC",
        "code": "PCGD_TH_01_CSVC_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },

    # ------------------------------------------------------------------
    # THCS
    # ------------------------------------------------------------------
    {
        "group": "THCS",
        "short": "THCS-M1",
        "code": "PCGD_THCS_M1_2025",
        "exporter": "xmc",
        "conclusion": {
            "G36": {1, 2, 3, "Không đạt", "Chưa đủ dữ liệu"},
            "G37": {1, 2, "Không đạt", "Chưa đủ dữ liệu"},
        },
        "percent_formats": ("H39", "H40", "H41"),
    },
    {
        "group": "THCS",
        "short": "THCS-M2",
        "code": "PCGD_THCS_M2_2025",
        "exporter": "xmc",
        "conclusion": {
            "W14": {1, 2, 3, "Không đạt", "Chưa đủ dữ liệu"},
        },
        "percent_formats": (
            "E14", "J14", "M14", "Q14", "V14",
            "E15", "J15", "M15", "Q15", "V15",
            "K20", "K21", "K22", "K23", "K24",
        ),
    },
    {
        "group": "THCS",
        "short": "THCS-TK",
        "code": "PCGD_THCS_TK_2025",
        "exporter": "xmc",
        "conclusion": {
            "F8": {1, 2, 3, "Không đạt", "Chưa đủ dữ liệu"},
            "G8": {1, 2, "Không đạt", "Chưa đủ dữ liệu"},
            "R8": {1, 2, 3, "Không đạt", "Chưa đủ dữ liệu"},
        },
        "percent_formats": ("I8", "K8", "O8"),
    },
    {
        "group": "THCS",
        "short": "THCS-M5",
        "code": "PCGD_THCS_M5_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "THCS",
        "short": "THCS-CSVC",
        "code": "PCGD_THCS_CSVC_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },

    # ------------------------------------------------------------------
    # XÓA MÙ CHỮ - giữ nguyên định dạng hiện hành.
    # ------------------------------------------------------------------
    {
        "group": "XÓA MÙ CHỮ",
        "short": "XMC-3",
        "code": "PCGD_XMC_3_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "XÓA MÙ CHỮ",
        "short": "CMC-2",
        "code": "PCGD_CMC_2_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "XÓA MÙ CHỮ",
        "short": "CMC-1",
        "code": "PCGD_CMC_1_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "XÓA MÙ CHỮ",
        "short": "XMC-4",
        "code": "PCGD_XMC_4_2025",
        "exporter": "xmc",
        "conclusion": {
            "R8": {1, 2, "Không đạt", "Chưa đủ dữ liệu"},
        },
        "percent_formats": (),
    },
)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def norm_text(value: Any) -> str:
    text = str(value or "").strip().upper()
    if not text:
        return ""

    text = unicodedata.normalize("NFD", text)
    text = "".join(
        ch
        for ch in text
        if unicodedata.category(ch) != "Mn"
    )

    return " ".join(text.split())


def safe_filename(value: str) -> str:
    text = norm_text(value)
    text = re.sub(r"[^A-Z0-9_-]+", "_", text)
    return text.strip("_") or "BAO_CAO"


def normalize_cell_value(value: Any) -> Any:
    if value is None:
        return None

    if isinstance(value, bool):
        return value

    if isinstance(value, int):
        return value

    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        return value

    return " ".join(str(value).strip().split())


def db_state() -> dict[str, Any]:
    if not DB_PATH.exists():
        raise RuntimeError(
            f"Không tìm thấy database: {DB_PATH}"
        )

    uri = DB_PATH.resolve().as_uri() + "?mode=ro"
    con = sqlite3.connect(uri, uri=True)

    try:
        integrity = str(
            con.execute(
                "PRAGMA integrity_check"
            ).fetchone()[0]
        )

        fk_rows = con.execute(
            "PRAGMA foreign_key_check"
        ).fetchall()

        total = int(
            con.execute(
                "SELECT COUNT(*) FROM communes"
            ).fetchone()[0]
        )

        special = int(
            con.execute(
                """
                SELECT COUNT(*)
                FROM communes
                WHERE COALESCE(
                    is_special_difficulty_area,
                    0
                ) = 1
                """
            ).fetchone()[0]
        )

        return {
            "integrity": integrity,
            "fk_count": len(fk_rows),
            "total": total,
            "special": special,
        }

    finally:
        con.close()


def verify_db_state(state: dict[str, Any]) -> None:
    if str(state["integrity"]).lower() != "ok":
        raise RuntimeError(
            f"integrity_check={state['integrity']!r}"
        )

    if int(state["fk_count"]) != 0:
        raise RuntimeError(
            f"foreign_key_check có {state['fk_count']} lỗi."
        )

    if int(state["total"]) != EXPECTED_COMMUNES:
        raise RuntimeError(
            f"communes={state['total']}; "
            f"yêu cầu={EXPECTED_COMMUNES}."
        )

    if int(state["special"]) != EXPECTED_SPECIAL:
        raise RuntimeError(
            f"xã ĐBKK={state['special']}; "
            f"yêu cầu={EXPECTED_SPECIAL}."
        )


def verify_source_files() -> list[str]:
    warnings = []

    required = (
        MN_BUILDER_PATH,
        XMC_BUILDER_PATH,
        REPORT_CENTER_PATH,
        RULES_PATH,
    )

    for path in required:
        if not path.exists():
            raise RuntimeError(
                f"Không tìm thấy source: {path}"
            )

        source = path.read_text(
            encoding="utf-8-sig",
            errors="strict",
        )
        ast.parse(source)

    mn = MN_BUILDER_PATH.read_text(
        encoding="utf-8-sig",
        errors="strict",
    )
    xmc = XMC_BUILDER_PATH.read_text(
        encoding="utf-8-sig",
        errors="strict",
    )
    rc = REPORT_CENTER_PATH.read_text(
        encoding="utf-8-sig",
        errors="strict",
    )
    rules = RULES_PATH.read_text(
        encoding="utf-8-sig",
        errors="strict",
    )

    # Các mốc đã khóa ở phần kết luận.
    must_have = (
        (
            mn,
            "BAI_13B_11_15_2_4_9_1_MN_HELPERS_START",
            "MN thiếu marker kết luận cấp xã 4.9.1",
        ),
        (
            mn,
            "BAI_13B_11_15_2_4_9_2_MN_HELPERS_START",
            "MN thiếu marker kết luận cấp tỉnh 4.9.2",
        ),
        (
            mn,
            "BAI_13B_11_15_2_4_9_2A_1_MN_HELPERS_START",
            "MN thiếu marker định dạng 100%",
        ),
        (
            xmc,
            "BAI_13B_11_15_2_4_9_1_XMC_HELPERS_START",
            "TH/XMC/THCS thiếu marker 4.9.1",
        ),
        (
            xmc,
            "BAI_13B_11_15_2_4_9_2_XMC_HELPERS_START",
            "TH/XMC/THCS thiếu marker 4.9.2",
        ),
        (
            xmc,
            "BAI_13B_11_15_2_4_9_2A_1_XMC_HELPERS_START",
            "TH/XMC/THCS thiếu marker định dạng 100%",
        ),
        (
            rules,
            "evaluate_primary_province",
            "Rules thiếu evaluator cấp tỉnh",
        ),
        (
            rc,
            "def _export_primary_th_m1(",
            "report_center thiếu exporter TH-M1",
        ),
    )

    for source, token, message in must_have:
        if token not in source:
            raise RuntimeError(message)

    # Inventory sanity.
    for code in (
        "PCGD_MN_M1_2025",
        "PCGD_MN_02_2025",
        "PCGD_MN_01_GV_2025",
        "PCGD_MN_01_CSVC_2025",
        "PCGD_MN_TAICHINH_2025",
    ):
        if code not in mn:
            warnings.append(
                f"Không thấy mã {code} trong MN builder."
            )

    for code in (
        "PCGD_TH_02_2025",
        "PCGD_TH_01_GV_2025",
        "PCGD_TH_01_CSVC_2025",
        "PCGD_THCS_M1_2025",
        "PCGD_THCS_M2_2025",
        "PCGD_THCS_TK_2025",
        "PCGD_THCS_M5_2025",
        "PCGD_THCS_CSVC_2025",
        "PCGD_XMC_3_2025",
        "PCGD_CMC_2_2025",
        "PCGD_CMC_1_2025",
        "PCGD_XMC_4_2025",
    ):
        if code not in xmc:
            warnings.append(
                f"Không thấy mã {code} trong TH/XMC/THCS builder."
            )

    return warnings


def make_so_request():
    from starlette.requests import Request

    scope = {
        "type": "http",
        "asgi": {
            "version": "3.0",
            "spec_version": "2.3",
        },
        "http_version": "1.1",
        "server": ("127.0.0.1", 80),
        "client": ("127.0.0.1", 0),
        "scheme": "http",
        "method": "GET",
        "root_path": "",
        "path": "/bao-cao",
        "raw_path": b"/bao-cao",
        "query_string": b"",
        "headers": [],
        "auth_user": {
            "id": 0,
            "username": "kiem_thu_cap_so",
            "full_name": "Kiểm thử tổng thể cấp Sở",
            "role_code": "SO",
            "role_name": "Sở",
            "commune_id": None,
            "school_id": None,
            "unit_name": "Toàn tỉnh",
        },
    }

    return Request(scope)


async def response_to_bytes(response) -> bytes:
    body = getattr(response, "body", None)

    if isinstance(body, (bytes, bytearray)) and body:
        return bytes(body)

    iterator = getattr(response, "body_iterator", None)

    if iterator is None:
        raise RuntimeError(
            f"Response {type(response).__name__} "
            "không có body/body_iterator."
        )

    chunks = []

    async for chunk in iterator:
        if isinstance(chunk, str):
            chunk = chunk.encode("utf-8")
        chunks.append(bytes(chunk))

    return b"".join(chunks)


def resolve_result(value):
    if inspect.isawaitable(value):
        return asyncio.run(value)
    return value


def call_with_known_args(
    func,
    *,
    db,
    request,
    school_year,
    report_type: str | None = None,
):
    candidates = {
        "db": db,
        "request": request,
        "school_year": school_year,
        "report_type": report_type,
        "selected_commune_id": None,
        "selected_school_id": None,
        "scope_label": "Toàn tỉnh",
    }

    signature = inspect.signature(func)
    kwargs = {}

    for name, parameter in signature.parameters.items():
        if name in candidates:
            if (
                name == "report_type"
                and report_type is None
            ):
                if (
                    parameter.default
                    is not inspect.Parameter.empty
                ):
                    continue

            kwargs[name] = candidates[name]
            continue

        if (
            parameter.default
            is not inspect.Parameter.empty
        ):
            continue

        if parameter.kind in (
            inspect.Parameter.VAR_POSITIONAL,
            inspect.Parameter.VAR_KEYWORD,
        ):
            continue

        raise RuntimeError(
            f"{func.__module__}.{func.__name__}: "
            f"tham số bắt buộc chưa biết: {name}"
        )

    return resolve_result(func(**kwargs))


def choose_school_year(db):
    from sqlalchemy import select, text
    from app.models import SchoolYear

    requested = str(
        os.environ.get("PHOCAP_TEST_YEAR", "")
        or ""
    ).strip()

    if requested:
        year = db.scalar(
            select(SchoolYear).where(
                SchoolYear.code == requested
            )
        )

        if year is None:
            raise RuntimeError(
                f"PHOCAP_TEST_YEAR={requested!r} "
                "nhưng DB không có năm học này."
            )

        return year

    rows = db.execute(
        text(
            """
            SELECT
                sy.id,
                sy.code,
                COUNT(spr.id) AS cnt
            FROM school_years sy
            LEFT JOIN survey_person_year_records spr
              ON spr.school_year_id = sy.id
            GROUP BY sy.id, sy.code
            ORDER BY cnt DESC, sy.code DESC
            """
        )
    ).all()

    if not rows:
        raise RuntimeError(
            "Database chưa có school_years."
        )

    best_id = int(rows[0][0])

    year = db.scalar(
        select(SchoolYear).where(
            SchoolYear.id == best_id
        )
    )

    if year is None:
        raise RuntimeError(
            f"Không lấy được SchoolYear ID={best_id}."
        )

    return year


def year_record_count(db, school_year_id: int) -> int:
    from sqlalchemy import text

    value = db.execute(
        text(
            """
            SELECT COUNT(*)
            FROM survey_person_year_records
            WHERE school_year_id=:id
            """
        ),
        {"id": int(school_year_id)},
    ).scalar_one()

    return int(value or 0)


def expected_reference_year(school_year) -> str:
    code = str(
        getattr(school_year, "code", "")
        or ""
    )

    found = re.findall(r"\d{4}", code)

    if found:
        return found[0]

    for attr in ("start_date", "end_date"):
        value = getattr(school_year, attr, None)
        found = re.findall(r"\d{4}", str(value or ""))
        if found:
            return found[0]

    return ""


def workbook_scan(
    data: bytes,
    spec: dict,
    reference_year: str,
):
    from openpyxl import load_workbook

    if not data.startswith(b"PK"):
        raise RuntimeError(
            "Response không có chữ ký ZIP/XLSX (PK)."
        )

    workbook = load_workbook(
        BytesIO(data),
        data_only=False,
    )

    try:
        if not workbook.sheetnames:
            raise RuntimeError(
                "Workbook không có sheet."
            )

        all_text = []
        errors = []
        formula_ref_errors = []
        nonempty = 0
        scanned_cells = 0
        truncated = False

        for ws in workbook.worksheets:
            max_row = min(
                max(int(ws.max_row or 1), 1),
                250,
            )
            max_col = min(
                max(int(ws.max_column or 1), 1),
                120,
            )

            if (
                int(ws.max_row or 1) > max_row
                or int(ws.max_column or 1) > max_col
            ):
                truncated = True

            for row in ws.iter_rows(
                min_row=1,
                max_row=max_row,
                min_col=1,
                max_col=max_col,
            ):
                for cell in row:
                    scanned_cells += 1
                    value = cell.value

                    if value not in (None, ""):
                        nonempty += 1

                    if isinstance(value, str):
                        stripped = value.strip()

                        if stripped:
                            all_text.append(stripped)

                        upper = stripped.upper()

                        if upper in EXCEL_ERROR_VALUES:
                            errors.append(
                                f"{ws.title}!{cell.coordinate}={stripped}"
                            )

                        if (
                            stripped.startswith("=")
                            and "#REF!" in upper
                        ):
                            formula_ref_errors.append(
                                f"{ws.title}!{cell.coordinate}={stripped}"
                            )

                    if (
                        getattr(cell, "data_type", None) == "e"
                        and str(value or "").upper()
                        in EXCEL_ERROR_VALUES
                    ):
                        errors.append(
                            f"{ws.title}!{cell.coordinate}={value}"
                        )

        normalized_texts = [
            norm_text(text)
            for text in all_text
        ]

        has_toan_tinh = any(
            "TOAN TINH" in text
            for text in normalized_texts
        )

        has_tinh_nghe_an = any(
            (
                "TINH: NGHE AN" in text
                or "TINH NGHE AN" in text
            )
            for text in normalized_texts
        )

        has_year = (
            not reference_year
            or any(
                reference_year in str(text)
                for text in all_text
            )
        )

        has_department_signature = any(
            (
                "SO GIAO DUC" in text
                or "GIAM DOC" in text
            )
            for text in normalized_texts
        )

        conclusion_values = {}

        for ref, allowed in spec["conclusion"].items():
            ws = workbook[
                workbook.sheetnames[0]
            ]

            value = normalize_cell_value(
                ws[ref].value
            )

            conclusion_values[ref] = value

            if value not in allowed:
                raise RuntimeError(
                    f"Ô kết luận {ref}={value!r}; "
                    f"không thuộc tập hợp hợp lệ {sorted(map(str, allowed))}."
                )

        percent_formats = {}

        for ref in spec["percent_formats"]:
            ws = workbook[
                workbook.sheetnames[0]
            ]

            fmt = str(
                ws[ref].number_format
                or ""
            )

            percent_formats[ref] = fmt

            if fmt != EXPECTED_PERCENT_FORMAT:
                raise RuntimeError(
                    f"Ô tỷ lệ {ref} format={fmt!r}; "
                    f"yêu cầu={EXPECTED_PERCENT_FORMAT!r}."
                )

        if errors:
            raise RuntimeError(
                "Phát hiện lỗi Excel: "
                + "; ".join(errors[:10])
            )

        if formula_ref_errors:
            raise RuntimeError(
                "Phát hiện công thức #REF!: "
                + "; ".join(formula_ref_errors[:10])
            )

        if nonempty < 10:
            raise RuntimeError(
                f"Workbook chỉ có {nonempty} ô có dữ liệu."
            )

        warnings = []

        if not has_toan_tinh:
            warnings.append(
                "Không tìm thấy chuỗi 'Toàn tỉnh' trong vùng quét."
            )

        if not has_tinh_nghe_an:
            warnings.append(
                "Không tìm thấy chuỗi 'Tỉnh: Nghệ An' trong vùng quét."
            )

        if not has_year:
            warnings.append(
                f"Không tìm thấy năm {reference_year} trong vùng quét."
            )

        if not has_department_signature:
            warnings.append(
                "Không phát hiện chữ ký/cụm 'Sở Giáo dục' hoặc 'Giám đốc'."
            )

        if truncated:
            warnings.append(
                "Workbook lớn; kiểm tra lỗi Excel chỉ quét tối đa "
                "250 hàng x 120 cột mỗi sheet."
            )

        return {
            "sheetnames": tuple(workbook.sheetnames),
            "nonempty": nonempty,
            "scanned_cells": scanned_cells,
            "has_toan_tinh": has_toan_tinh,
            "has_tinh_nghe_an": has_tinh_nghe_an,
            "has_year": has_year,
            "has_department_signature": has_department_signature,
            "conclusions": conclusion_values,
            "percent_formats": percent_formats,
            "warnings": warnings,
        }

    finally:
        workbook.close()


def export_one(
    spec: dict,
    *,
    db,
    request,
    school_year,
    mn_builder,
    xmc_builder,
    report_center,
):
    if spec["exporter"] == "mn":
        response = call_with_known_args(
            mn_builder.export_mn_report,
            db=db,
            request=request,
            school_year=school_year,
            report_type=spec["code"],
        )

    elif spec["exporter"] == "xmc":
        response = call_with_known_args(
            xmc_builder.export_additional_report,
            db=db,
            request=request,
            school_year=school_year,
            report_type=spec["code"],
        )

    elif spec["exporter"] == "th_m1":
        exporter = getattr(
            report_center,
            "_export_primary_th_m1",
            None,
        )

        if not callable(exporter):
            raise RuntimeError(
                "Không tìm thấy "
                "report_center._export_primary_th_m1."
            )

        response = call_with_known_args(
            exporter,
            db=db,
            request=request,
            school_year=school_year,
            report_type=None,
        )

    else:
        raise RuntimeError(
            f"Loại exporter không biết: {spec['exporter']}"
        )

    return asyncio.run(
        response_to_bytes(response)
    )


def csv_escape(value: Any) -> str:
    text = str(value or "")
    if any(ch in text for ch in (",", '"', "\n", "\r")):
        return '"' + text.replace('"', '""') + '"'
    return text


def main() -> int:
    print("=" * 112)
    print(
        "KIỂM THỬ TỔNG THỂ HỆ THỐNG BÁO CÁO - CẤP SỞ - V3"
    )
    print("=" * 112)
    print()
    print(
        "Phạm vi: tài khoản SO / Toàn tỉnh / 18 đầu ra Excel."
    )
    print(
        "Chỉ đọc database; không sửa source."
    )
    print()

    required_paths = (
        DB_PATH,
        MN_BUILDER_PATH,
        XMC_BUILDER_PATH,
        REPORT_CENTER_PATH,
        RULES_PATH,
    )

    for path in required_paths:
        if not path.exists():
            print(
                f"[LỖI] Không tìm thấy: {path}"
            )
            return 1

    db_hash_before = sha256_file(DB_PATH)
    state_before = db_state()

    try:
        verify_db_state(state_before)
    except Exception as exc:
        print(
            "[LỖI] Database không đạt preflight: "
            + str(exc)
        )
        return 1

    print(
        "[1/8] Database: "
        f"{state_before['total']} xã/phường; "
        f"ĐBKK={state_before['special']}; "
        "integrity=ok; FK=0."
    )

    try:
        source_warnings = verify_source_files()
    except Exception as exc:
        print(
            "[LỖI] Source không đạt preflight: "
            + str(exc)
        )
        return 1

    print(
        "[2/8] Source: AST OK; "
        "4.9.1 + 4.9.2 + định dạng 100% còn đủ marker."
    )

    if source_warnings:
        for warning in source_warnings:
            print(
                "      [WARN] " + warning
            )

    old_cwd = Path.cwd()
    results = []

    try:
        os.chdir(PROJECT)
        sys.path.insert(0, str(PROJECT))

        from sqlalchemy import text
        from app.database import SessionLocal

        import app.pcgd_mn_report_builders_v1 as mn_builder
        import app.pcgd_xmc_report_builders_v1 as xmc_builder
        import app.routers.report_center as report_center

        db = SessionLocal()

        try:
            # Không cho phép bất kỳ thao tác ghi nào trong phiên kiểm thử.
            db.execute(
                text(
                    "PRAGMA query_only = ON"
                )
            )

            school_year = choose_school_year(db)

            # V3_FIX_DETACHED_SCHOOL_YEAR:
            # Chụp primitive TRƯỚC khi Session đóng.
            school_year_id = int(school_year.id)
            school_year_code = str(
                getattr(school_year, "code", "")
                or ""
            )

            record_count = year_record_count(
                db,
                school_year_id,
            )
            reference_year = expected_reference_year(
                school_year
            )

            print(
                "[3/8] Năm học kiểm thử: "
                f"{school_year_code} "
                f"(ID={school_year_id}); "
                f"bản ghi điều tra={record_count}."
            )

            request = make_so_request()

            OUT_DIR.mkdir(
                parents=True,
                exist_ok=False,
            )

            print(
                "[4/8] Xuất và kiểm tra 18 báo cáo..."
            )

            for index, spec in enumerate(
                REPORT_SPECS,
                start=1,
            ):
                print(
                    f"      [{index:02d}/18] "
                    f"{spec['group']} / {spec['short']} ..."
                )

                row = {
                    "group": spec["group"],
                    "short": spec["short"],
                    "code": spec["code"],
                    "status": "FAIL",
                    "file": "",
                    "message": "",
                    "warnings": [],
                    "details": {},
                }

                try:
                    data = export_one(
                        spec,
                        db=db,
                        request=request,
                        school_year=school_year,
                        mn_builder=mn_builder,
                        xmc_builder=xmc_builder,
                        report_center=report_center,
                    )

                    if len(data) < 1000:
                        raise RuntimeError(
                            f"File quá nhỏ: {len(data)} bytes."
                        )

                    if not data.startswith(b"PK"):
                        raise RuntimeError(
                            "Dữ liệu trả về không phải XLSX."
                        )

                    file_path = (
                        OUT_DIR
                        / (
                            f"{index:02d}_"
                            f"{safe_filename(spec['short'])}_"
                            f"{school_year_code}_"
                            "SO_TOAN_TINH.xlsx"
                        )
                    )

                    file_path.write_bytes(data)

                    details = workbook_scan(
                        data,
                        spec,
                        reference_year,
                    )

                    row["file"] = str(file_path)
                    row["details"] = details
                    row["warnings"] = list(
                        details["warnings"]
                    )

                    if row["warnings"]:
                        row["status"] = "WARN"
                        row["message"] = (
                            f"Xuất OK ({len(data)} bytes); "
                            f"{len(row['warnings'])} cảnh báo."
                        )
                    else:
                        row["status"] = "PASS"
                        row["message"] = (
                            f"Xuất OK ({len(data)} bytes); "
                            "kiểm tra bắt buộc đạt."
                        )

                    print(
                        f"             [{row['status']}] "
                        + row["message"]
                    )

                    if details["conclusions"]:
                        print(
                            "             Kết luận: "
                            + ", ".join(
                                f"{k}={v!r}"
                                for k, v
                                in details["conclusions"].items()
                            )
                        )

                    if details["percent_formats"]:
                        print(
                            "             Format %: "
                            + ", ".join(
                                f"{k}={v}"
                                for k, v
                                in details["percent_formats"].items()
                            )
                        )

                    for warning in row["warnings"]:
                        print(
                            "             [WARN] "
                            + warning
                        )

                except Exception as exc:
                    row["status"] = "FAIL"
                    row["message"] = str(exc)

                    print(
                        "             [FAIL] "
                        + str(exc)
                    )

                results.append(row)

        finally:
            try:
                db.rollback()
            except Exception:
                pass
            db.close()

        print(
            "[5/8] Đã chạy đủ 18 nhánh báo cáo."
        )

        state_after = db_state()
        verify_db_state(state_after)

        db_hash_after = sha256_file(DB_PATH)

        if state_after != state_before:
            raise RuntimeError(
                "Trạng thái database trước/sau khác nhau."
            )

        if db_hash_after != db_hash_before:
            raise RuntimeError(
                "SHA256 phocap.db thay đổi. "
                "Hãy dừng Uvicorn rồi chạy lại kiểm thử."
            )

        print(
            "[6/8] Database: KHÔNG THAY ĐỔI; "
            "SHA256 giữ nguyên."
        )

        pass_count = sum(
            1 for row in results
            if row["status"] == "PASS"
        )
        warn_count = sum(
            1 for row in results
            if row["status"] == "WARN"
        )
        fail_count = sum(
            1 for row in results
            if row["status"] == "FAIL"
        )

        by_group = {}

        for row in results:
            group = by_group.setdefault(
                row["group"],
                {
                    "PASS": 0,
                    "WARN": 0,
                    "FAIL": 0,
                },
            )
            group[row["status"]] += 1

        report_lines = [
            "=" * 112,
            "KIỂM THỬ TỔNG THỂ HỆ THỐNG BÁO CÁO - CẤP SỞ",
            "=" * 112,
            "",
            "PHẠM VI:",
            " - Vai trò: SO",
            " - Đơn vị: Toàn tỉnh",
            " - commune_id: None",
            " - school_id: None",
            f" - Năm học: {school_year_code}",
            f" - Bản ghi điều tra: {record_count}",
            "",
            "DATABASE:",
            f" - integrity_check: {state_before['integrity']}",
            f" - foreign_key_check: {state_before['fk_count']}",
            f" - xã/phường: {state_before['total']}",
            f" - xã ĐBKK: {state_before['special']}",
            f" - SHA256 trước: {db_hash_before}",
            f" - SHA256 sau  : {db_hash_after}",
            " - Database: KHÔNG THAY ĐỔI.",
            "",
            "KẾT QUẢ TỔNG:",
            f" - PASS: {pass_count}",
            f" - WARN: {warn_count}",
            f" - FAIL: {fail_count}",
            "",
            "THEO NHÓM:",
        ]

        for group_name in (
            "MẦM NON",
            "TIỂU HỌC",
            "THCS",
            "XÓA MÙ CHỮ",
        ):
            counts = by_group.get(
                group_name,
                {
                    "PASS": 0,
                    "WARN": 0,
                    "FAIL": 0,
                },
            )

            report_lines.append(
                f" - {group_name}: "
                f"PASS={counts['PASS']}; "
                f"WARN={counts['WARN']}; "
                f"FAIL={counts['FAIL']}"
            )

        report_lines.extend(
            [
                "",
                "=" * 112,
                "CHI TIẾT 18 BÁO CÁO",
                "=" * 112,
            ]
        )

        for index, row in enumerate(
            results,
            start=1,
        ):
            report_lines.extend(
                [
                    "",
                    (
                        f"[{index:02d}] "
                        f"{row['group']} / {row['short']} / "
                        f"{row['code']}"
                    ),
                    f"Trạng thái: {row['status']}",
                    f"Thông tin : {row['message']}",
                    f"File      : {row['file'] or '(không tạo được)'}",
                ]
            )

            details = row.get(
                "details"
            ) or {}

            if details:
                report_lines.extend(
                    [
                        (
                            "Sheets    : "
                            + ", ".join(
                                details["sheetnames"]
                            )
                        ),
                        (
                            "Ô có dữ liệu: "
                            f"{details['nonempty']}"
                        ),
                        (
                            "Có 'Toàn tỉnh': "
                            f"{details['has_toan_tinh']}"
                        ),
                        (
                            "Có 'Tỉnh: Nghệ An': "
                            f"{details['has_tinh_nghe_an']}"
                        ),
                        (
                            f"Có năm {reference_year}: "
                            f"{details['has_year']}"
                        ),
                        (
                            "Có chữ ký Sở/Giám đốc: "
                            f"{details['has_department_signature']}"
                        ),
                    ]
                )

                if details["conclusions"]:
                    report_lines.append(
                        "Kết luận : "
                        + ", ".join(
                            f"{k}={v!r}"
                            for k, v
                            in details["conclusions"].items()
                        )
                    )

                if details["percent_formats"]:
                    report_lines.append(
                        "Format %  : "
                        + ", ".join(
                            f"{k}={v!r}"
                            for k, v
                            in details["percent_formats"].items()
                        )
                    )

            for warning in row["warnings"]:
                report_lines.append(
                    "WARN      : " + warning
                )

        report_lines.extend(
            [
                "",
                "=" * 112,
                "CÁCH ĐỌC KẾT QUẢ",
                "=" * 112,
                "PASS = đường xuất + file Excel + kiểm tra bắt buộc đạt.",
                "WARN = file xuất được, nhưng cần xem trực quan một điểm trình bày/phạm vi.",
                "FAIL = phải sửa trước khi chuyển sang kiểm thử cấp Xã.",
                "",
            ]
        )

        if fail_count == 0:
            report_lines.append(
                "KẾT LUẬN KỸ THUẬT: "
                "KHÔNG CÓ FAIL Ở CẤP SỞ."
            )

            if warn_count:
                report_lines.append(
                    "Cần xem các WARN trước khi chốt cấp Sở."
                )
            else:
                report_lines.append(
                    "18/18 báo cáo cấp Sở PASS."
                )
        else:
            report_lines.append(
                "KẾT LUẬN KỸ THUẬT: "
                f"CÒN {fail_count} FAIL; "
                "CHƯA CHUYỂN SANG CẤP XÃ."
            )

        REPORT_PATH.write_text(
            "\n".join(report_lines),
            encoding="utf-8-sig",
        )

        csv_rows = [
            "STT,NHOM,BAO_CAO,REPORT_CODE,TRANG_THAI,THONG_TIN,FILE"
        ]

        for index, row in enumerate(
            results,
            start=1,
        ):
            csv_rows.append(
                ",".join(
                    (
                        str(index),
                        csv_escape(row["group"]),
                        csv_escape(row["short"]),
                        csv_escape(row["code"]),
                        csv_escape(row["status"]),
                        csv_escape(row["message"]),
                        csv_escape(row["file"]),
                    )
                )
            )

        SUMMARY_CSV.write_text(
            "\n".join(csv_rows),
            encoding="utf-8-sig",
        )

        print(
            "[7/8] Tổng hợp: "
            f"PASS={pass_count}; "
            f"WARN={warn_count}; "
            f"FAIL={fail_count}."
        )

        print(
            f"[8/8] Báo cáo: {REPORT_PATH}"
        )
        print(
            f"      CSV    : {SUMMARY_CSV}"
        )
        print()

        print("=" * 112)

        if fail_count == 0:
            print(
                "KIỂM THỬ TỔNG THỂ CẤP SỞ: "
                "KHÔNG CÓ FAIL."
            )

            if warn_count:
                print(
                    "Có WARN cần xem trước khi chuyển sang cấp Xã."
                )
            else:
                print(
                    "18/18 BÁO CÁO PASS."
                )
        else:
            print(
                "KIỂM THỬ TỔNG THỂ CẤP SỞ: "
                f"CÒN {fail_count} FAIL."
            )
            print(
                "CHƯA CHUYỂN SANG CẤP XÃ."
            )

        print(
            "Database: KHÔNG THAY ĐỔI."
        )
        print(
            f"Thư mục Excel: {OUT_DIR}"
        )
        print("=" * 112)

        return 1 if fail_count else 0

    except Exception as exc:
        print()
        print("=" * 112)
        print(
            "[LỖI HỆ THỐNG KIỂM THỬ CẤP SỞ]"
        )
        print(str(exc))
        print("=" * 112)

        try:
            OUT_DIR.mkdir(
                parents=True,
                exist_ok=True,
            )

            error_path = (
                OUT_DIR
                / f"loi_he_thong_kiem_thu_cap_so_{STAMP}.txt"
            )

            error_path.write_text(
                (
                    f"Lỗi: {exc}\n\n"
                    + traceback.format_exc()
                ),
                encoding="utf-8-sig",
            )

            print(
                f"Chi tiết lỗi: {error_path}"
            )

        except Exception:
            pass

        print(
            "Script không có lệnh "
            "UPDATE/INSERT/DELETE database."
        )

        return 2

    finally:
        os.chdir(old_cwd)


if __name__ == "__main__":
    raise SystemExit(main())
