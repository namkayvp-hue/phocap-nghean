from __future__ import annotations

import json
import re
from copy import copy
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break


TRACKING_SHEET_NAME = "Sổ theo dõi PCGDMN"
META_SHEET_NAME = "_PCGDMN_META"
FIELD_SHEET_NAME = "Phiếu điều tra"
FIELD_MAP_SHEET_NAME = "_PCGDMN_FIELD_MAP"
FIELD_VERSION = "PCGDMN_FIELD_FORM_V240"
DATA_START_ROW = 11
DATA_END_ROW = 110
META_TABLE_ROW = 20
MULTI_MARKER_COLUMN = 21
MULTI_JSON_COLUMN = 22
MULTI_BLOCK_START = "__PCGDMN_BLOCK_START__"
MULTI_PERSON_ROW = "__PCGDMN_PERSON_ROW__"
MULTI_BLOCK_END = "__PCGDMN_BLOCK_END__"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value).strip()


def _meta_values(sheet: Any) -> dict[str, str]:
    values: dict[str, str] = {}
    for row in range(1, min(sheet.max_row, META_TABLE_ROW - 1) + 1):
        key = _text(sheet.cell(row, 1).value)
        if key:
            values[key] = _text(sheet.cell(row, 2).value)
    return values


def _single_row_meta(sheet: Any) -> dict[int, dict[str, Any]]:
    headers = [
        _text(sheet.cell(META_TABLE_ROW, col).value)
        for col in range(1, sheet.max_column + 1)
    ]
    result: dict[int, dict[str, Any]] = {}
    if "visible_row" not in headers:
        return result
    for row in range(META_TABLE_ROW + 1, sheet.max_row + 1):
        values = {
            header: sheet.cell(row, col).value
            for col, header in enumerate(headers, start=1)
            if header
        }
        visible = values.get("visible_row")
        try:
            visible_row = int(visible)
        except (TypeError, ValueError):
            continue
        result[visible_row] = values
    return result


def _find_tracking(workbook: Any) -> Any:
    for name in workbook.sheetnames:
        if str(name).strip() == TRACKING_SHEET_NAME:
            return workbook[name]
    raise ValueError("Không tìm thấy trang kỹ thuật Sổ theo dõi PCGDMN.")


def _parse_header_value(raw: Any, label: str) -> str:
    text = _text(raw)
    marker = f"{label}:"
    pos = text.lower().find(marker.lower())
    if pos < 0:
        return text
    return text[pos + len(marker):].strip()


def _extract_hamlet(header: Any) -> str:
    raw = _text(header)
    match = re.search(
        r"(?:Xóm/khối|Tổ,\s*thôn,\s*xóm|Thôn/xóm|Xóm|Khối|Tổ|Bản)\s*:\s*([^;]+)",
        raw,
        flags=re.IGNORECASE,
    )
    if match:
        return match.group(1).strip()
    return ""


def _extract_address(header: Any) -> str:
    return _parse_header_value(header, "Địa chỉ")


def _extract_head(header: Any) -> str:
    return _parse_header_value(header, "Họ và tên chủ hộ")


def _split_name(full_name: str) -> tuple[str, str]:
    parts = [part for part in str(full_name or "").split() if part]
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return " ".join(parts[:-1]), parts[-1]


def _line_value(raw: Any, labels: tuple[str, ...]) -> str:
    text = _text(raw)
    for label in labels:
        marker = label + ":"
        pos = text.lower().find(marker.lower())
        if pos >= 0:
            return text[pos + len(marker):].strip()
    return text


def _gender_ethnic(raw: Any) -> tuple[str, str]:
    text = _text(raw)
    gender = ""
    ethnic = ""
    match_gender = re.search(
        r"(?:Giới\s*tính|GT)\s*:\s*(Nam|Nữ|Nu)",
        text,
        flags=re.IGNORECASE,
    )
    if match_gender:
        gender = "Nữ" if match_gender.group(1).lower() in {"nữ", "nu"} else "Nam"
    else:
        if re.search(r"\bNữ\b", text, flags=re.IGNORECASE):
            gender = "Nữ"
        elif re.search(r"\bNam\b", text, flags=re.IGNORECASE):
            gender = "Nam"

    match_ethnic = re.search(
        r"(?:DT|Dân\s*tộc)\s*:\s*([^|;]+)",
        text,
        flags=re.IGNORECASE,
    )
    if match_ethnic:
        ethnic = match_ethnic.group(1).strip()
    return gender, ethnic


def _year_label(tracking: Any, block_start: int, column: int) -> str:
    top = _text(tracking.cell(block_start + 6, column).value)
    bottom = _text(tracking.cell(block_start + 7, column).value)
    if top and bottom:
        return f"{top}-{bottom}"
    return top or bottom


def _payload_from_marker(tracking: Any, row: int) -> dict[str, Any]:
    if _text(tracking.cell(row, MULTI_MARKER_COLUMN).value) != MULTI_PERSON_ROW:
        return {}
    raw = _text(tracking.cell(row, MULTI_JSON_COLUMN).value)
    if not raw:
        return {}
    try:
        value = json.loads(raw)
        return value if isinstance(value, dict) else {}
    except json.JSONDecodeError:
        return {}


def _select_rows(
    tracking: Any,
    start_row: int,
    end_row: int,
    *,
    extra_blank_rows: int = 5,
) -> list[int]:
    existing = []
    for row in range(start_row, end_row + 1):
        full_name = " ".join(
            part for part in (
                _text(tracking.cell(row, 3).value),
                _text(tracking.cell(row, 4).value),
            )
            if part
        ).strip()
        if full_name:
            existing.append(row)

    if existing:
        last = max(existing)
        stop = min(end_row, last + extra_blank_rows)
        return list(range(start_row, stop + 1))

    return list(range(start_row, min(end_row, start_row + extra_blank_rows - 1) + 1))


def _form_specs(workbook: Any, tracking: Any, meta: dict[str, str]) -> list[dict[str, Any]]:
    version = meta.get("VERSION", "")
    result: list[dict[str, Any]] = []

    if version == "PCGDMN_NHIEU_HO_V1":
        active_start: int | None = None
        active_payload: dict[str, Any] | None = None

        for row in range(1, tracking.max_row + 1):
            marker = _text(tracking.cell(row, MULTI_MARKER_COLUMN).value)
            if marker == MULTI_BLOCK_START:
                active_start = row
                raw = _text(tracking.cell(row, MULTI_JSON_COLUMN).value)
                try:
                    parsed = json.loads(raw) if raw else {}
                except json.JSONDecodeError:
                    parsed = {}
                active_payload = parsed if isinstance(parsed, dict) else {}
            elif marker == MULTI_BLOCK_END and active_start is not None:
                block_end = row - 1
                payload = active_payload or {}
                data_start = active_start + DATA_START_ROW - 1
                selected_rows = _select_rows(
                    tracking,
                    data_start,
                    block_end,
                )
                result.append({
                    "form_id": payload.get("form_id"),
                    "household_id": payload.get("household_id"),
                    "form_number": payload.get("form_number") or "",
                    "household_code": payload.get("household_code") or "",
                    "current_year_column": int(payload.get("current_year_column") or 18),
                    "school_year_code": meta.get("SCHOOL_YEAR_CODE", ""),
                    "commune_name": meta.get("COMMUNE_NAME", ""),
                    "block_start": active_start,
                    "tracking_rows": selected_rows,
                    "head_name": _extract_head(tracking.cell(active_start + 2, 1).value),
                    "hamlet": _extract_hamlet(tracking.cell(active_start + 1, 1).value),
                    "address": _extract_address(tracking.cell(active_start + 2, 11).value),
                    "row_payloads": {
                        item: _payload_from_marker(tracking, item)
                        for item in selected_rows
                    },
                })
                active_start = None
                active_payload = None
        return result

    if version != "PCGDMN_HO_V1":
        return result

    current_year_column = int(meta.get("CURRENT_YEAR_COLUMN") or 18)
    row_meta = _single_row_meta(workbook[META_SHEET_NAME])
    selected_rows = _select_rows(
        tracking,
        DATA_START_ROW,
        DATA_END_ROW,
    )
    result.append({
        "form_id": meta.get("FORM_ID"),
        "household_id": meta.get("HOUSEHOLD_ID"),
        "form_number": meta.get("FORM_NUMBER", ""),
        "household_code": meta.get("HOUSEHOLD_CODE", ""),
        "current_year_column": current_year_column,
        "school_year_code": meta.get("SCHOOL_YEAR_CODE", ""),
        "commune_name": meta.get("COMMUNE_NAME", ""),
        "block_start": 1,
        "tracking_rows": selected_rows,
        "head_name": meta.get("HOUSEHOLD_HEAD_NAME", "") or _extract_head(tracking["A3"].value),
        "hamlet": _extract_hamlet(tracking["A2"].value),
        "address": _extract_address(tracking["K3"].value),
        "row_payloads": {
            item: row_meta.get(item, {})
            for item in selected_rows
        },
    })
    return result


def _apply_all_borders(ws: Any, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell.border = border


def _fill_header(
    ws: Any,
    row: int,
    *,
    spec: dict[str, Any],
    page_index: int,
    page_total: int,
) -> int:
    start = row

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    ws.cell(row, 1).value = f"Phường/Xã (1): {spec.get('commune_name') or ''}"
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=3)
    ws.cell(row + 1, 1).value = f"Tổ, thôn, xóm (2): {spec.get('hamlet') or ''}"
    ws.merge_cells(start_row=row + 2, start_column=1, end_row=row + 2, end_column=3)
    ws.cell(row + 2, 1).value = f"Địa chỉ (3): {spec.get('address') or ''}"

    ws.merge_cells(start_row=row, start_column=4, end_row=row + 2, end_column=8)
    title = ws.cell(row, 4)
    title.value = (
        "PHIẾU ĐIỀU TRA PHỔ CẬP GIÁO DỤC - CHỐNG MÙ CHỮ\n"
        f"Năm học {spec.get('school_year_code') or ''}"
        + (f"\nTrang {page_index}/{page_total}" if page_total > 1 else "")
    )
    title.font = Font(name="Times New Roman", size=12, bold=True)
    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=11)
    ws.cell(row, 9).value = f"Họ và tên chủ hộ (4): {spec.get('head_name') or ''}"
    ws.merge_cells(start_row=row + 1, start_column=9, end_row=row + 1, end_column=11)
    ws.cell(row + 1, 9).value = f"Số phiếu (5): {spec.get('form_number') or ''}"
    ws.merge_cells(start_row=row + 2, start_column=9, end_row=row + 2, end_column=11)
    ws.cell(row + 2, 9).value = "Diện cư trú (6): □ Thường trú   □ Tạm trú   Điện thoại (7):"

    for r in range(row, row + 3):
        ws.row_dimensions[r].height = 17
        for c in range(1, 12):
            cell = ws.cell(r, c)
            if cell.coordinate not in ws.merged_cells:
                pass
            cell.font = Font(name="Times New Roman", size=7.5, bold=(c in {1, 9}))
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    row += 3
    headers = [
        "STT",
        "HỌ VÀ TÊN ĐỐI TƯỢNG\n(Lớn tuổi ghi trước)",
        "TÊN LỚP ĐANG HỌC\n(Theo năm học)",
        "TÊN TRƯỜNG - HUYỆN ĐANG HỌC\n(Tương ứng từng năm học)",
        "TỐT NGHIỆP / HOÀN THÀNH\nMN → THPT · Bổ túc · TN nghề",
        "HỌC XONG\nLớp / Năm",
        "BỎ HỌC\nLớp / Năm",
        "MÙ CHỮ\nHọc XMC / CN XMC / Tái mù",
        "KHUYẾT TẬT",
        "CHUYỂN ĐẾN, CHUYỂN ĐI, CHẾT",
        "GHI CHÚ",
    ]
    for c, value in enumerate(headers, start=1):
        cell = ws.cell(row, c, value)
        cell.font = Font(name="Times New Roman", size=6.5, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 33
    _apply_all_borders(ws, start, row, 1, 11)
    return row + 1


def _history_rows(
    tracking: Any,
    spec: dict[str, Any],
    tracking_row: int,
) -> list[tuple[str, str]]:
    block_start = int(spec.get("block_start") or 1)
    current_col = int(spec.get("current_year_column") or 18)

    values: list[tuple[str, str]] = []
    current_label = spec.get("school_year_code") or _year_label(
        tracking,
        block_start,
        current_col,
    )
    current_value = _text(tracking.cell(tracking_row, current_col).value)
    values.append((str(current_label), current_value))

    for col in range(current_col - 1, max(11, current_col - 6), -1):
        label = _year_label(tracking, block_start, col)
        value = _text(tracking.cell(tracking_row, col).value)
        if label or value:
            values.append((label, value))
        if len(values) >= 7:
            break

    while len(values) < 7:
        values.append(("20__-20__", ""))

    return values[:7]


def _fill_person_slot(
    ws: Any,
    start_row: int,
    *,
    tracking: Any,
    spec: dict[str, Any],
    tracking_row: int,
    slot_number: int,
    map_sheet: Any,
    map_row: int,
) -> int:
    end_row = start_row + 6

    payload = dict(spec.get("row_payloads", {}).get(tracking_row, {}) or {})
    full_name = " ".join(
        part for part in (
            _text(tracking.cell(tracking_row, 3).value),
            _text(tracking.cell(tracking_row, 4).value),
        )
        if part
    ).strip()
    dob = _text(tracking.cell(tracking_row, 5).value)
    gender = _text(tracking.cell(tracking_row, 6).value)
    ethnic = _text(tracking.cell(tracking_row, 9).value)
    relationship = _text(tracking.cell(tracking_row, 10).value)
    parent_name = _text(tracking.cell(tracking_row, 11).value)
    school_name = _text(tracking.cell(tracking_row, 19).value)
    notes = _text(tracking.cell(tracking_row, 20).value)
    person_code = _text(payload.get("person_code"))

    ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
    ws.cell(start_row, 1).value = slot_number
    ws.cell(start_row, 1).alignment = Alignment(horizontal="center", vertical="center")

    person_lines = [
        full_name,
        f"QH chủ hộ: {relationship}" if relationship else "QH chủ hộ:",
        f"Sinh: {dob}" if dob else "Sinh:",
        (
            f"Giới tính: {gender or ''} | DT: {ethnic or ''}"
            if (gender or ethnic)
            else "Giới tính: | DT:"
        ),
        f"Cha, mẹ/Ng. đỡ đầu: {parent_name}" if parent_name else "Cha, mẹ/Ng. đỡ đầu:",
        "TG:",
        f"Mã nội bộ: {person_code}" if person_code else "Mã nội bộ:",
    ]

    history = _history_rows(tracking, spec, tracking_row)

    for offset in range(7):
        row = start_row + offset
        b = ws.cell(row, 2)
        b.value = person_lines[offset]
        b.font = Font(
            name="Times New Roman",
            size=7.0,
            bold=(offset == 0),
        )
        b.alignment = Alignment(vertical="center", wrap_text=True)

        year_label, class_value = history[offset]
        c = ws.cell(row, 3)
        c.value = f"{year_label}   {class_value}".rstrip()
        c.font = Font(name="Times New Roman", size=6.4)
        c.alignment = Alignment(vertical="center", wrap_text=True)

        d = ws.cell(row, 4)
        if offset == 0:
            d.value = f"{spec.get('school_year_code') or year_label}   {school_name}".rstrip()
        else:
            d.value = f"{year_label}"
        d.font = Font(name="Times New Roman", size=6.4)
        d.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 8.7

    for col in range(5, 12):
        ws.merge_cells(
            start_row=start_row,
            start_column=col,
            end_row=end_row,
            end_column=col,
        )
        cell = ws.cell(start_row, col)
        cell.font = Font(name="Times New Roman", size=6.5)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.cell(start_row, 11).value = notes

    _apply_all_borders(ws, start_row, end_row, 1, 11)

    headers = [
        "VERSION",
        "VISIBLE_SHEET",
        "VISIBLE_START_ROW",
        "TRACKING_ROW",
        "FORM_ID",
        "HOUSEHOLD_ID",
        "PERSON_ID",
        "PERSON_CODE",
        "CURRENT_YEAR_COLUMN",
        "SCHOOL_YEAR_CODE",
        "HEADER_ROW",
    ]
    if map_row == 2:
        for col, name in enumerate(headers, start=1):
            map_sheet.cell(1, col).value = name

    values = [
        FIELD_VERSION,
        FIELD_SHEET_NAME,
        start_row,
        tracking_row,
        spec.get("form_id"),
        spec.get("household_id"),
        payload.get("person_id"),
        payload.get("person_code"),
        spec.get("current_year_column"),
        spec.get("school_year_code"),
        spec.get("visible_header_row"),
    ]
    for col, value in enumerate(values, start=1):
        map_sheet.cell(map_row, col).value = value

    return end_row + 1


def _fill_signature(ws: Any, row: int, head_name: str) -> int:
    ranges = [(1, 2), (3, 5), (6, 8), (9, 11)]
    labels = [
        f"HỌ, TÊN CHỦ HỘ\n{head_name or ''}\n(Ký, ghi rõ họ tên)",
        "CÁN BỘ, NHÂN VIÊN ĐIỀU TRA 1\n(Ký, ghi rõ họ tên)",
        "CÁN BỘ, NHÂN VIÊN ĐIỀU TRA 2\n(Ký, ghi rõ họ tên)",
        "XÁC NHẬN CỦA UBND PHƯỜNG/XÃ\n(Ký tên, đóng dấu)",
    ]
    for (c1, c2), value in zip(ranges, labels):
        ws.merge_cells(start_row=row, start_column=c1, end_row=row + 2, end_column=c2)
        cell = ws.cell(row, c1)
        cell.value = value
        cell.font = Font(name="Times New Roman", size=7, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    for r in range(row, row + 3):
        ws.row_dimensions[r].height = 14
    return row + 3


def add_field_form_view_to_bytes(content: bytes) -> bytes:
    workbook = load_workbook(BytesIO(content))
    tracking = _find_tracking(workbook)
    if META_SHEET_NAME not in workbook.sheetnames:
        raise ValueError("File thiếu trang kỹ thuật ẩn.")

    meta = _meta_values(workbook[META_SHEET_NAME])
    specs = _form_specs(workbook, tracking, meta)
    if not specs:
        raise ValueError("Không xác định được phiếu để dựng mẫu Excel điều tra.")

    for name in (FIELD_SHEET_NAME, FIELD_MAP_SHEET_NAME):
        if name in workbook.sheetnames:
            workbook.remove(workbook[name])

    ws = workbook.create_sheet(FIELD_SHEET_NAME, 0)
    map_sheet = workbook.create_sheet(FIELD_MAP_SHEET_NAME)
    map_sheet.sheet_state = "veryHidden"

    widths = {
        "A": 4.5,
        "B": 29,
        "C": 19,
        "D": 28,
        "E": 18,
        "F": 12,
        "G": 12,
        "H": 14,
        "I": 13,
        "J": 18,
        "K": 19,
    }
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width

    current_row = 1
    map_row = 2

    for spec_index, spec in enumerate(specs, start=1):
        tracking_rows = list(spec.get("tracking_rows") or [])
        pages = [
            tracking_rows[i:i + 7]
            for i in range(0, len(tracking_rows), 7)
        ] or [[]]

        for page_index, page_rows in enumerate(pages, start=1):
            header_row = current_row
            spec["visible_header_row"] = header_row
            current_row = _fill_header(
                ws,
                current_row,
                spec=spec,
                page_index=page_index,
                page_total=len(pages),
            )

            for slot_index in range(7):
                if slot_index < len(page_rows):
                    tracking_row = page_rows[slot_index]
                    current_row = _fill_person_slot(
                        ws,
                        current_row,
                        tracking=tracking,
                        spec=spec,
                        tracking_row=tracking_row,
                        slot_number=(page_index - 1) * 7 + slot_index + 1,
                        map_sheet=map_sheet,
                        map_row=map_row,
                    )
                    map_row += 1
                else:
                    # Vẫn tạo một dòng trắng đúng mẫu để in, nhưng không có mapping nhập.
                    blank_start = current_row
                    blank_end = blank_start + 6
                    ws.merge_cells(
                        start_row=blank_start,
                        start_column=1,
                        end_row=blank_end,
                        end_column=1,
                    )
                    ws.cell(blank_start, 1).value = (
                        (page_index - 1) * 7 + slot_index + 1
                    )
                    for offset in range(7):
                        r = blank_start + offset
                        ws.cell(r, 2).value = (
                            "" if offset == 0 else
                            (
                                "QH chủ hộ:" if offset == 1 else
                                "Sinh:" if offset == 2 else
                                "Giới tính: | DT:" if offset == 3 else
                                "Cha, mẹ/Ng. đỡ đầu:" if offset == 4 else
                                "TG:" if offset == 5 else
                                "Mã nội bộ:"
                            )
                        )
                        ws.cell(r, 3).value = "20__-20__"
                        ws.cell(r, 4).value = "20__-20__"
                        ws.row_dimensions[r].height = 8.7
                    for col in range(5, 12):
                        ws.merge_cells(
                            start_row=blank_start,
                            start_column=col,
                            end_row=blank_end,
                            end_column=col,
                        )
                    _apply_all_borders(ws, blank_start, blank_end, 1, 11)
                    current_row = blank_end + 1

            if page_index == len(pages):
                current_row = _fill_signature(
                    ws,
                    current_row,
                    str(spec.get("head_name") or ""),
                )

            footer_row = current_row
            ws.merge_cells(
                start_row=footer_row,
                start_column=1,
                end_row=footer_row,
                end_column=11,
            )
            ws.cell(footer_row, 1).value = (
                f"{meta.get('BATCH_CODE') or ''} · "
                f"{spec.get('form_number') or ''} · "
                f"{spec.get('household_code') or ''}"
            )
            ws.cell(footer_row, 1).font = Font(
                name="Times New Roman",
                size=6.5,
                italic=True,
            )
            ws.cell(footer_row, 1).alignment = Alignment(horizontal="left")
            current_row += 1

            ws.row_breaks.append(Break(id=current_row - 1))
            current_row += 1

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(
        left=0.12,
        right=0.12,
        top=0.18,
        bottom=0.18,
        header=0.0,
        footer=0.0,
    )
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:K{max(1, current_row - 1)}"

    tracking.sheet_state = "veryHidden"
    workbook[META_SHEET_NAME].sheet_state = "veryHidden"
    ws.sheet_state = "visible"
    workbook.active = 0

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _map_headers(sheet: Any) -> dict[str, int]:
    result = {}
    for col in range(1, sheet.max_column + 1):
        name = _text(sheet.cell(1, col).value)
        if name:
            result[name] = col
    return result


def _parse_current_value(raw: Any, year_code: str) -> str:
    text = _text(raw)
    if not text:
        return ""
    if year_code and text.startswith(year_code):
        return text[len(year_code):].strip()
    match = re.match(r"^\d{4}-\d{4}\s+(.*)$", text)
    if match:
        return match.group(1).strip()
    return text


def _parse_person_fields(ws: Any, start_row: int, year_code: str) -> dict[str, str]:
    full_name = _text(ws.cell(start_row, 2).value)
    relationship = _line_value(
        ws.cell(start_row + 1, 2).value,
        ("QH chủ hộ", "Quan hệ với chủ hộ"),
    )
    birth = _line_value(
        ws.cell(start_row + 2, 2).value,
        ("Sinh", "Ngày sinh"),
    )
    gender, ethnic = _gender_ethnic(ws.cell(start_row + 3, 2).value)
    parent = _line_value(
        ws.cell(start_row + 4, 2).value,
        ("Cha, mẹ/Ng. đỡ đầu", "Cha/mẹ/Ng. đỡ đầu", "Cha/mẹ/người đỡ đầu"),
    )
    program_class = _parse_current_value(
        ws.cell(start_row, 3).value,
        year_code,
    )
    school = _parse_current_value(
        ws.cell(start_row, 4).value,
        year_code,
    )
    notes = _text(ws.cell(start_row, 11).value)

    return {
        "full_name": full_name,
        "relationship": relationship,
        "birth": birth,
        "gender": gender,
        "ethnic": ethnic,
        "parent": parent,
        "program_class": program_class,
        "school": school,
        "notes": notes,
    }


def sync_field_form_to_tracking(workbook: Any) -> bool:
    if FIELD_MAP_SHEET_NAME not in workbook.sheetnames:
        return False
    if FIELD_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(
            "File có thông tin ánh xạ Phiếu điều tra nhưng trang Phiếu điều tra đã bị xóa/đổi tên."
        )

    tracking = _find_tracking(workbook)
    map_sheet = workbook[FIELD_MAP_SHEET_NAME]
    visible = workbook[FIELD_SHEET_NAME]
    headers = _map_headers(map_sheet)

    required = {
        "VERSION",
        "VISIBLE_START_ROW",
        "TRACKING_ROW",
        "CURRENT_YEAR_COLUMN",
        "SCHOOL_YEAR_CODE",
        "HEADER_ROW",
    }
    if not required.issubset(headers):
        raise ValueError("Trang ánh xạ kỹ thuật của Phiếu điều tra không đúng phiên bản.")

    for row in range(2, map_sheet.max_row + 1):
        if _text(map_sheet.cell(row, headers["VERSION"]).value) != FIELD_VERSION:
            continue

        try:
            visible_start = int(
                map_sheet.cell(row, headers["VISIBLE_START_ROW"]).value
            )
            tracking_row = int(
                map_sheet.cell(row, headers["TRACKING_ROW"]).value
            )
            current_year_col = int(
                map_sheet.cell(row, headers["CURRENT_YEAR_COLUMN"]).value
            )
            header_row = int(
                map_sheet.cell(row, headers["HEADER_ROW"]).value
            )
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"Ánh xạ kỹ thuật tại dòng {row} bị hỏng."
            ) from exc

        year_code = _text(
            map_sheet.cell(row, headers["SCHOOL_YEAR_CODE"]).value
        )
        fields = _parse_person_fields(
            visible,
            visible_start,
            year_code,
        )

        # Đồng bộ thôn/xóm từ đầu phiếu vào từng dòng như bộ importer cũ.
        hamlet_text = _extract_hamlet(visible.cell(header_row + 1, 1).value)

        left, right = _split_name(fields["full_name"])
        tracking.cell(tracking_row, 3).value = left
        tracking.cell(tracking_row, 4).value = right
        tracking.cell(tracking_row, 5).value = fields["birth"]
        tracking.cell(tracking_row, 6).value = fields["gender"]
        if hamlet_text:
            tracking.cell(tracking_row, 7).value = hamlet_text
        tracking.cell(tracking_row, 9).value = fields["ethnic"]
        tracking.cell(tracking_row, 10).value = fields["relationship"]
        tracking.cell(tracking_row, 11).value = fields["parent"]
        tracking.cell(tracking_row, current_year_col).value = fields[
            "program_class"
        ]
        tracking.cell(tracking_row, 19).value = fields["school"]
        tracking.cell(tracking_row, 20).value = fields["notes"]

    return True
