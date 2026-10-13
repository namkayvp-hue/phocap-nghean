import json
import unittest
from collections import Counter
from datetime import date
from io import BytesIO
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from app.services.appendix_pagination import column_maps
from app.xmc_appendices import TEMPLATE_DIR
from app.models import School
from app.survey_models import SurveyPerson, SurveyPersonYearRecord
from app.report_input_models import SchoolNetworkYearData, SchoolStructuredReportInput
from app.xmc_appendices import COMPULSORY_FORMS
from tools.test_xmc_scope import ScopeTests


class NewAppendixTests(unittest.TestCase):
    def setUp(self):
        self.fixture=ScopeTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.doCleanups)
        self.db=self.fixture.db
        self.client=self.fixture.client
        self.fixture.user['role_code']='ADMIN'
        for i, age, grade in [(1,6,1),(2,11,5),(3,15,9)]:
            p=self.db.get(SurveyPerson,i)
            p.date_of_birth=date(2026-age,1,1)
            r=self.db.get(SurveyPersonYearRecord,i)
            r.learning_status='DANG_HOC'
            r.class_name_reported=f'{grade}A'
            r.school_name_reported='THCS Test' if grade>=6 else 'Tieu hoc Test'
            r.completed_primary_program=i>=2
            r.completed_lower_secondary_program=i==3
            r.disability_status='KHONG_KHUYET_TAT'
            r.study_location_scope='TAI_CHO'
            r.current_education_program='GDTX' if i==3 else 'CHUA_XAC_DINH'
            r.post_lower_secondary_path='CHUA_XAC_DINH'
        self.db.commit()

    def response(self, code, query=''):
        response=self.client.get(f'/bao-cao/xoa-mu-chu/phu-luc/{code}/tong-hop?school_year_id=1'+query)
        self.assertEqual(response.status_code,200,response.text[:200] if response.status_code!=200 else '')
        return response

    def workbook(self,code,query=''):
        return load_workbook(BytesIO(self.response(code,query).content))

    def age_cell(self, wb, page, address):
        source = load_workbook(TEMPLATE_DIR / 'Phu_luc_2_Giao_duc_bat_buoc.xlsx')
        names = ['Trang_04', 'Trang_05'] if page in (4,5) else ['Trang_09', 'Trang_10']
        sheets = [source[name] for name in names]
        _, maps = column_maps(sheets)
        index = names.index(f'Trang_{page:02d}')
        row, col = coordinate_to_tuple(address)
        return wb[names[0]].cell(row + (sheets[0].max_row if index else 0), maps[index][col-1])

    def resources(self):
        for i in range(1,8):
            s=self.db.get(School,i)
            if not s:
                s=School(id=i,code=f'NEW{i}',name=f'School {i}',commune_id=1)
                self.db.add(s)
                self.db.flush()
            for level in ['TH','THCS']:
                self.db.add(SchoolNetworkYearData(school_id=i,school_year_id=1,level_code=level,
                    data_json=json.dumps(dict(class_total=i,teachers=2*i,rooms_permanent=i,rooms_semi_permanent=0,rooms_temporary=0))))
                self.db.add(SchoolStructuredReportInput(school_id=i,school_year_id=1,form_code=level+'_01_CSVC',
                    data_json=json.dumps(dict(room_rent_borrow=0,school_site_count=1))))
        # A different year must not override current figures.
        self.db.add(SchoolNetworkYearData(school_id=1,school_year_id=2,level_code='TH',data_json='{"teachers":999,"class_total":999}'))
        self.db.commit()

    def test_blank_downloads_keep_all_content_in_one_sheet_per_form(self):
        for code in ['01-TH_TN','05-THCS_TTN','giao-duc-bat-buoc']:
            response=self.client.get(f'/bao-cao/xoa-mu-chu/phu-luc/{code}/excel')
            self.assertEqual(response.status_code,200)
            wb=load_workbook(BytesIO(response.content))
            source=load_workbook(TEMPLATE_DIR / ('Phu_luc_2_Giao_duc_bat_buoc.xlsx' if code=='giao-duc-bat-buoc' else f'Mau_{code}.xlsx'))
            from app.services.report_export_date import stamp_workbook
            stamp_workbook(source)
            values=lambda book: Counter(str(c.value) for ws in book for row in ws for c in row if c.value is not None)
            self.assertEqual(values(wb),values(source))
            self.assertEqual(len(wb.sheetnames),9 if code=='giao-duc-bat-buoc' else 1)
            for ws in wb:
                if ws.title=='Trang_03': continue
                self.assertTrue(ws.print_title_rows)
                self.assertEqual(ws.page_setup.fitToWidth,1)
                self.assertEqual(ws.page_setup.fitToHeight,0)
                self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)

    def test_all_eight_forms_bundle_and_population_values(self):
        for code,_,pages in COMPULSORY_FORMS:
            wb=self.workbook(code)
            self.assertEqual(len(wb.sheetnames),2)
            self.assertTrue(wb.worksheets[0].sheet_properties.pageSetUpPr.fitToPage)
            self.assertEqual(wb.worksheets[0].page_setup.fitToHeight,0)
        wb=self.workbook('bo-pl2')
        self.assertEqual(len(wb.sheetnames),10)
        self.assertNotIn('Trang_05',wb.sheetnames)
        self.assertNotIn('Trang_10',wb.sheetnames)
        for page,cell in [(4,'F9'),(4,'L9'),(5,'F11'),(5,'G11'),(9,'N9'),(9,'O26'),(10,'L20'),(10,'M20')]:
            self.assertEqual(self.age_cell(wb,page,cell).value,1)
        self.assertEqual(self.age_cell(wb,5,'G11').number_format,'0%')

    def test_school_rows_expand_and_weighted_totals(self):
        self.resources()
        wb=self.workbook('bo-pl2')
        # 7 rows exceed every original template capacity.
        self.assertEqual(wb['Trang_06']['I19'].value,56)
        self.assertEqual(wb['Trang_06']['P19'].value,2)
        self.assertEqual(wb['Trang_07']['J18'].value,28)
        self.assertEqual(wb['Trang_07']['AB18'].value,1)
        self.assertEqual(wb['Trang_11']['H18'].value,56)
        self.assertEqual(wb['Trang_11']['O18'].value,2)
        self.assertEqual(wb['Trang_12']['D19'].value,28)
        self.assertEqual(wb['Trang_12']['H19'].value,1)
        for page in [6,7,11,12]:
            self.assertEqual(wb[f'Trang_{page:02d}'].page_setup.fitToHeight,0)

    def test_missing_data_does_not_become_zero(self):
        self.db.get(SurveyPersonYearRecord,1).completed_primary_program=None
        self.db.commit()
        wb=self.workbook('01-TH_TN')
        self.assertIsNone(self.age_cell(wb,4,'F31').value)
        self.assertEqual(self.age_cell(wb,4,'F12').value,0)  # explicit no disability
        self.db.get(SurveyPerson,1).date_of_birth=None
        self.db.commit()
        self.assertIsNone(self.age_cell(self.workbook('01-TH_TN'),4,'F9').value)

    def test_completion_list_is_confirmed_and_safe(self):
        self.db.get(SurveyPerson,2).full_name='=1+1'
        self.db.commit()
        ws=self.workbook('04-TH_HT')['Trang_08']
        self.assertEqual(ws['B9'].value,'=1+1')
        self.assertEqual(ws['B9'].data_type,'s')
        self.assertIsNone(ws['H9'].value)
        self.assertIsNone(ws['I9'].value)
        self.assertNotIn('Person 1',[c.value for row in ws for c in row])
        self.db.get(SurveyPersonYearRecord,2).primary_completion_year=2026
        self.db.commit()
        self.assertEqual(self.workbook('04-TH_HT')['Trang_08']['H9'].value,2026)

    def test_word_is_valid_and_has_numbers_and_template_sections(self):
        for code in ['pl4','pl5']:
            response=self.response(code)
            self.assertIn('wordprocessingml',response.headers['content-type'])
            with ZipFile(BytesIO(response.content)) as z:
                self.assertIsNone(z.testzip())
                root=ET.fromstring(z.read('word/document.xml'))
                text=' '.join(root.itertext())
                self.assertIn('1/1',text)
                self.assertIn('100%',text)
                self.assertIn('THÔNG TIN CẦN HOÀN THIỆN',text)
                self.assertIn('Trong đó số người mù chữ:',text)
                self.assertIn('KIẾN NGHỊ',text)

    def test_all_exports_reject_forged_scope_and_year(self):
        self.fixture.user['role_code']='GIAO_VIEN'
        for code in [c for c,_,_ in COMPULSORY_FORMS]+['bo-pl2','pl4','pl5']:
            base=f'/bao-cao/xoa-mu-chu/phu-luc/{code}/tong-hop'
            self.assertEqual(self.client.get(base+'?school_year_id=1&school_id=2').status_code,403)
            self.assertEqual(self.client.get(base+'?school_year_id=999').status_code,422)
        self.resources()
        ws=self.workbook('02-TH_GV')['Trang_06']
        self.assertEqual(ws['I12'].value,2)
        self.assertIsNone(ws['B13'].value)
        self.fixture.user['role_code']='PHONG_BAN'
        self.assertEqual(self.workbook('02-TH_GV')['Trang_06']['I19'].value,56)


if __name__=='__main__':
    unittest.main()
