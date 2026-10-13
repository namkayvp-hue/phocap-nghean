import unittest
from datetime import date, timedelta
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from openpyxl import load_workbook
from app.routers import household_updates as module


class BirthValidationTests(unittest.TestCase):
    def setUp(self):
        self.batch = SimpleNamespace(school_year=SimpleNamespace(code='2026-2027'),
                                     commune=SimpleNamespace(name='Phường Thành Vinh'))
        self.content = Path('du_lieu_test_50_ho_phuong_thanh_vinh.xlsx').read_bytes()

    def test_original_file_reports_every_bad_birth_year(self):
        rows, issues, _ = module._parse_xlsx_rows(self.content,self.batch)
        errors = [i for i in issues if i['code']=='INVALID_BIRTH_DATE']
        self.assertEqual(len(rows),250)
        self.assertEqual({i['row'] for i in errors},set(range(6,256)))
        self.assertTrue(all(r['row_has_error'] for r in rows))

    def test_bad_dates_and_valid_leap_day(self):
        for parts,valid in [((29,2,2024),True),((29,2,2023),False),((1,1,11),False),
                            ((1.5,1,2010),False),((1,1,2020.5),False),
                            ((1,1,1899),False),((1,1,2027),False)]:
            wb=load_workbook(BytesIO(self.content)); ws=wb['MauNhapLieu']
            for c,value in zip((4,5,6),parts): ws.cell(6,c,value)
            output=BytesIO();wb.save(output)
            _,issues,_=module._parse_xlsx_rows(output.getvalue(),self.batch)
            self.assertEqual(not any(i['row']==6 and i['code']=='INVALID_BIRTH_DATE' for i in issues),valid,parts)

    def test_error_stops_before_applying_households(self):
        with patch.object(module,'_precheck_database_conflicts',return_value=[]), \
             patch.object(module,'_add_issue_dict'), patch.object(module,'_apply_rows') as apply:
            job=SimpleNamespace(id=1)
            ok,_=module._process_xlsx_job(MagicMock(),job,self.batch,self.content)
            self.assertFalse(ok)
            self.assertEqual(job.rows_error,250)
            apply.assert_not_called()

    def test_corrected_sample_passes(self):
        content=Path('du_lieu_test_50_ho_phuong_thanh_vinh_da_sua.xlsx').read_bytes()
        rows,issues,_=module._parse_xlsx_rows(content,self.batch)
        self.assertEqual(len(rows),250)
        self.assertEqual(issues,[])
        self.assertEqual(len({r['form_number'] for r in rows}),50)


if __name__ == '__main__': unittest.main()
