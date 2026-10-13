from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

OUT = Path("C:/phocap/du_lieu_test_50_ho_phuong_thanh_vinh_da_sua.xlsx")
commune = "Phường Thành Vinh"
year = 2026
headers = [
    "TT", "Họ đệm", "Tên", "Ngày", "Tháng", "Năm sinh", "Nữ",
    "Dân tộc", "Tôn giáo", "Diện ưu tiên", "Chủ hộ - Họ đệm", "Chủ hộ - Tên",
    "Địa chỉ: Số nhà, tên đường, tổ (nếu có)", "Số phiếu", "Diện cư trú",
    "Tình trạng cư trú", "Khối học", "Lớp học", "Tên trường", "Địa bàn trường",
    "Mã Trường", "Số định danh", "Căn cước công dân",
]
hamlets = ["Khối 1", "Khối 2", "Khối 3", "Khối 4", "Khối 5"]
streets = ["Đường Lê Lợi", "Đường Nguyễn Du", "Đường Trần Phú", "Đường Quang Trung", "Đường Vinh Tân"]

wb = Workbook()
ws = wb.active
ws.title = "MauNhapLieu"
ws["A1"] = "Năm điều tra"; ws["B1"] = year
ws["A2"] = "Xã/phường"; ws["B2"] = commune
ws["D1"] = "Mã xã/phường"; ws["E1"] = "TV"
ws["D2"] = "Ghi chú"; ws["E2"] = "Dữ liệu kiểm thử 50 hộ dân; có thể nhập vào đợt điều tra năm 2026 của phường Thành Vinh."
for col, value in enumerate(headers, 1):
    ws.cell(4, col).value = value
    ws.cell(5, col).value = col

row = 6
for household_no in range(1, 51):
    form = f"TEST-TV-2026-{household_no:03d}"
    head_first = f"Nguyễn Thị {['An','Bình','Châu','Dung','Hạnh'][household_no % 5]}"
    head_last, head_name = head_first.rsplit(" ", 1)
    address = f"{household_no} {streets[(household_no - 1) % len(streets)]}, {hamlets[(household_no - 1) % len(hamlets)]}, {commune}"
    members = [
        (head_last, head_name, 35 + household_no % 18, "X", "", "", "", "", ""),
        ("Nguyễn Văn", f"Minh {household_no:02d}", 15 + household_no % 3, "", "THPT", f"Lớp {10 + household_no % 3}", "Trường THPT Thành Vinh", "THPT-TV", ""),
        ("Nguyễn Thị", f"Lan {household_no:02d}", 11 + household_no % 4, "X", "THCS", f"Lớp {6 + household_no % 4}", "Trường THCS Thành Vinh", "THCS-TV", ""),
        ("Nguyễn Văn", f"Nam {household_no:02d}", 6 + household_no % 5, "", "TH", f"Lớp {1 + household_no % 5}", "Trường Tiểu học Thành Vinh", "TH-TV", ""),
        ("Nguyễn Thị", f"Mai {household_no:02d}", 5, "X", "MN", "Mẫu giáo 5 tuổi", "Trường Mầm non Thành Vinh", "MN-TV", ""),
    ]
    for member_index, (surname, given, age, female, level, klass, school, school_code, _unused) in enumerate(members):
        pid = f"TV{household_no:03d}{member_index + 1:02d}"
        cccd = f"405026{household_no:04d}{member_index + 1:03d}"
        birth = date(year - age, (household_no % 12) + 1, (household_no % 27) + 1)
        ws.cell(row, 1).value = row - 5
        ws.cell(row, 2).value = surname
        ws.cell(row, 3).value = given
        ws.cell(row, 4).value = birth.day
        ws.cell(row, 5).value = birth.month
        ws.cell(row, 6).value = birth.year
        ws.cell(row, 7).value = female
        ws.cell(row, 8).value = "Kinh"
        ws.cell(row, 11).value = head_last
        ws.cell(row, 12).value = head_name
        ws.cell(row, 13).value = address
        ws.cell(row, 14).value = form
        ws.cell(row, 15).value = "Thường trú"
        ws.cell(row, 16).value = "Đang cư trú"
        ws.cell(row, 17).value = level
        ws.cell(row, 18).value = klass
        ws.cell(row, 19).value = school
        ws.cell(row, 20).value = commune
        ws.cell(row, 21).value = school_code
        ws.cell(row, 22).value = pid
        ws.cell(row, 23).value = cccd
        row += 1

thin = Side(style="thin", color="B7C9D6")
for cells in ws.iter_rows(min_row=1, max_row=row - 1, min_col=1, max_col=len(headers)):
    for cell in cells:
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
for cell in ws[4] + ws[5]:
    cell.font = Font(bold=True)
ws.freeze_panes = "A6"
widths = [6, 18, 16, 8, 8, 10, 6, 10, 12, 15, 18, 14, 46, 20, 15, 18, 12, 12, 30, 22, 15, 18, 20]
for index, width in enumerate(widths, 1):
    ws.column_dimensions[chr(64 + index) if index <= 26 else "A"].width = width

guide = wb.create_sheet("README")
guide.append(["Mục", "Nội dung"])
guide.append(["Phạm vi", "50 hộ dân, mỗi hộ 5 thành viên: chủ hộ, THPT, THCS, Tiểu học và Mầm non; phường Thành Vinh."])
guide.append(["Số phiếu", "TEST-TV-2026-001 đến TEST-TV-2026-050."])
guide.append(["Năm", "2026; chọn đợt điều tra năm học bắt đầu 2026."])
guide.append(["Lưu ý", "Đây là dữ liệu giả lập để kiểm thử, không phải dữ liệu dân cư thật."])
guide.column_dimensions["A"].width = 18
guide.column_dimensions["B"].width = 90
for cells in guide.iter_rows():
    for cell in cells:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
guide["A1"].font = guide["B1"].font = Font(bold=True)
wb.save(OUT)
print(OUT)
