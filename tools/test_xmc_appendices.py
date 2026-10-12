import unittest
from collections import Counter
from pathlib import Path

from fastapi import FastAPI, Request
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.routers.report_center import router
from app.xmc_appendices import TEMPLATE_DIR, XMC_FORMS, COMPULSORY_FORMS, template_files


class AppendixTests(unittest.TestCase):
    def test_long_forms_are_continuous_and_paginate(self):
        for code, page, ages in (("01", 15, range(15, 61)), ("07", 23, range(0, 61))):
            wb = load_workbook(TEMPLATE_DIR / f"Mau_{code}_XMC.xlsx")
            self.assertEqual(wb.sheetnames, [f"Trang_{page}"])
            sheet = wb.active
            values = [str(row[1].value or "").strip() for row in sheet]
            for age in ages:
                value = f"{age} tuổi" if code == "01" else str(age)
                self.assertEqual(values.count(value), 1, (code, age))
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.page_setup.fitToHeight, 0)
            self.assertTrue(sheet.print_title_rows)
            self.assertEqual(sum(c.value == "NGƯỜI LẬP BIỂU" for row in sheet for c in row), 1)

    def test_workbook_pages_and_editable_cells(self):
        for code, name in template_files().items():
            wb = load_workbook(TEMPLATE_DIR / name)
            expected = {"tat-ca": 24, "bo-xmc": 10, "phieu-dieu-tra": 2, "giao-duc-bat-buoc": 11}.get(code)
            if expected is None:
                expected = len(next((p for c, _, p in COMPULSORY_FORMS if c == code), (1,)))
            self.assertEqual(len(wb.sheetnames), expected)
            for sheet in wb:
                self.assertFalse(sheet.protection.sheet)
                self.assertFalse(sheet._images)
                self.assertGreater(sum(c.value is not None for row in sheet for c in row), 10)

    def test_download_routes(self):
        app = FastAPI()
        role = "PHONG_BAN"

        @app.middleware("http")
        async def auth(request: Request, call_next):
            request.scope["auth_user"] = {"role_code": role}
            return await call_next(request)

        app.include_router(router)
        with TestClient(app) as client:
            for role in ("PHONG_BAN", "XA", "ADMIN", "TRUONG", "GIAO_VIEN"):
                for code in template_files():
                    response = client.get(f"/bao-cao/xoa-mu-chu/phu-luc/{code}/excel")
                    self.assertEqual(response.status_code, 200)
                    self.assertTrue(response.content.startswith(b"PK"))
                    self.assertIn("attachment", response.headers["content-disposition"])
            self.assertEqual(client.get("/bao-cao/xoa-mu-chu/phu-luc/unknown/excel").status_code, 404)
            role = "UNKNOWN"
            self.assertEqual(client.get("/bao-cao/xoa-mu-chu/phu-luc/01/excel").status_code, 403)

    def test_source_text_preserved(self):
        try:
            import pymupdf
        except ImportError:
            self.skipTest("Build-only PDF dependency not installed")
        pdf = Path(__file__).resolve().parents[1] / "Các phụ lục.pdf"
        if not pdf.exists():
            self.skipTest("Source PDF not available")
        wb = load_workbook(TEMPLATE_DIR / "Cac_phu_luc_2026.xlsx")
        with pymupdf.open(pdf) as doc:
            for i, page in enumerate(doc, 1):
                with self.subTest(page=i):
                    first = next((pages[0] for _, _, pages in XMC_FORMS if i in pages), i)
                    sheet = wb[f"Trang_{first:02d}"]
                    text = " ".join(str(c.value) for row in sheet for c in row if c.value is not None)
                    missing = Counter(page.get_text().split()) - Counter(text.split())
                    self.assertFalse(missing, f"Missing PDF text: {missing}")


if __name__ == "__main__":
    unittest.main()
