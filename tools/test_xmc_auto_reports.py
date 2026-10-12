import unittest
from datetime import date
from types import SimpleNamespace as NS
from openpyxl import load_workbook

from app.services.xmc_appendix_reports import build_report, collect, counts, unit_values, conclusion
from app.xmc_appendices import XMC_FORMS


def person(i, age=20, g1=True, g2=False):
    p = NS(id=i, code=f"P{i}", full_name=f"Người {i}", date_of_birth=date(2026-age, 12, 31),
           gender="NU" if i % 2 else "NAM", ethnic_group="Thái" if i % 2 else "Kinh",
           household=NS(commune_id=1, address="Địa chỉ", head_name="Chủ hộ"),
           residency_status="THUONG_TRU", personal_id=f"001234{i:06d}")
    r = NS(id=i, completed_grade_3=g1, completed_grade_5=g2, disability_status="KHONG_KHUYET_TAT",
           learning_status="DANG_HOC", school=None, classroom=None, school_name_reported="THPT",
           class_name_reported="12A", completed_primary_program=g2, completed_lower_secondary_program=False)
    return p, r


class AutoReportTests(unittest.TestCase):
    def setUp(self):
        self.people = [person(i) for i in range(1, 8)]
        self.communes = [NS(id=1, name="Xã Một", is_special_difficulty_area=False)]
        self.evidence = {i: dict(enrolled1=True, enrolled2=False, recognized1=True,
                                 recognized2=False, relapsed=False, illiterate=False) for i in range(1,8)}

    def build(self, code, **kwargs):
        return load_workbook(build_report(code, self.people, self.evidence, 2026, "2026-2027",
                                         "Xã Một", self.communes, **kwargs))

    def test_all_forms_and_large_lists(self):
        for code, _, pages in XMC_FORMS:
            with self.subTest(code=code):
                wb = self.build(code, literacy_basis="g1")
                self.assertIn("Kiem_tra_du_lieu", wb.sheetnames)
                self.assertEqual(len(wb.sheetnames), 2)
        wb = self.build("02")
        ws = wb["Trang_18"]
        self.assertEqual(ws["B16"].value, "Người 1")
        self.assertEqual(ws["B22"].value, "Người 7")
        self.assertEqual(ws["J16"].value, "001234000001")
        self.assertTrue(any(c.value == "NGƯỜI LẬP BIỂU" for row in ws for c in row))
        ws = self.build("01", literacy_basis="g1")["Trang_15"]
        self.assertEqual(ws["D23"].value, 7)  # age 20
        self.assertEqual(ws["H23"].value, 0)  # fail level 1
        self.assertEqual(ws["L23"].value, 7)  # fail level 2
        self.assertEqual(ws["P23"].value, 7)
        self.assertEqual(ws["T23"].value, 1)
        self.assertEqual(ws["T23"].number_format, "0%")
        self.assertEqual(self.build("01")["Trang_15"]["P23"].value, 7)
        self.assertIsNone(self.build("01", literacy_basis=None)["Trang_15"]["P23"].value)

    def test_missing_is_not_zero_and_boundary_ages(self):
        people = [person(1, 15, None, None), person(2,35), person(3,60), person(4,61)]
        items = collect(people, {}, 2026)
        values = unit_values(items, "09")
        self.assertEqual(values[0], 2)
        self.assertEqual(values[5], 3)
        self.assertIsNone(values[1])
        self.assertEqual(conclusion(items, self.communes[0], True), "Chưa đủ dữ liệu")

    def test_enrollment_not_inferred_from_population(self):
        items = collect(self.people, {}, 2026)
        self.assertTrue(all(v is None for v in unit_values(items,"05")))
        items = collect(self.people, self.evidence, 2026)
        values = unit_values(items,"05")
        self.assertEqual(values[0],7)
        self.assertEqual(values[4],7)
        self.assertEqual(values[5],100)

    def test_unit_rows_extend_and_totals_recalculate(self):
        self.communes = [NS(id=i,name=f"Xã {i}",is_special_difficulty_area=False) for i in range(1,7)]
        for i, (p, _) in enumerate(self.people):
            p.household.commune_id = i % 6 + 1
        wb = self.build("09")
        ws = wb["Trang_28"]
        self.assertEqual(ws["B23"].value,"Xã 6")
        self.assertEqual(ws["B24"].value,"Cộng")
        self.assertEqual(ws["C24"].value,7)
        self.assertEqual(ws["F24"].value,1)
        self.assertEqual(ws["F24"].number_format,"0%")

    def test_percentage_cells_for_all_rate_forms(self):
        self.people = [person(i, g1=(i != 1)) for i in range(1,17)]
        for code, page in (("03",19),("04",20),("05",21),("06",22)):
            ws = self.build(code)[f"Trang_{page}"]
            for col in ("H","J","L","N","T","V","X","Z"):
                self.assertIn(ws[f"{col}18"].number_format,{"0%", "0.00%"})
                self.assertIn(ws[f"{col}21"].number_format,{"0%", "0.00%"})
        ws = self.build("03")["Trang_19"]
        self.assertEqual(ws["H18"].value,0.0625)
        self.assertEqual(ws["H18"].number_format,"0.00%")
        self.assertEqual(ws["H21"].value,0.0625)
        self.assertEqual(ws["C18"].value,16)
        self.assertNotIn("%",ws["C18"].number_format)
        self.assertIsNone(ws["H19"].value)

    def test_whole_percentage_has_no_decimal_places(self):
        self.people = [person(i, g1=(i != 1)) for i in range(1,6)]
        ws = self.build("03")["Trang_19"]
        self.assertEqual(ws["H18"].value,0.2)
        self.assertEqual(ws["H18"].number_format,"0%")
        self.assertEqual(ws["H21"].number_format,"0%")

    def test_relapser_and_excel_formula_text(self):
        self.people[0][0].full_name = "=1+1"
        self.evidence[1].update(relapsed=True, recognized_year="2024")
        ws = self.build("08")["Trang_27"]
        self.assertEqual(ws["B16"].value,"=1+1")
        self.assertEqual(ws["B16"].data_type,"s")
        self.assertEqual(collect(self.people,self.evidence,2026)[0]["g1"],False)


if __name__ == "__main__":
    unittest.main()
