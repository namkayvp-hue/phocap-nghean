from __future__ import annotations

import csv
import re
import sqlite3
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from app.database import DATABASE_PATH


PROJECT_DIR = Path(__file__).resolve().parent
INPUT_FILE = PROJECT_DIR / "uploads" / "Danh_sach_giao_vien.xlsx"
EXPORT_ROOT = PROJECT_DIR / "exports"

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "commune_name": (
        "UBND Xã",
        "UBND xã",
        "Xã/Phường",
        "Xã/phường",
        "Tên xã",
        "Tên xã/phường",
    ),
    "school_name": (
        "Tên trường",
        "Trường",
        "Đơn vị công tác",
        "Tên đơn vị",
    ),
    "staff_code": (
        "Mã định danh Bộ GD&ĐT",
        "Mã định danh Bộ GDĐT",
        "Mã định danh",
        "Mã CBGVNV",
        "Mã cán bộ",
    ),
    "full_name": (
        "Họ tên",
        "Họ và tên",
        "Họ tên cán bộ",
        "Họ tên giáo viên",
    ),
    "date_of_birth": (
        "Ngày sinh",
        "Ngày tháng năm sinh",
    ),
    "employment_status": (
        "Trạng thái",
        "Trạng thái làm việc",
        "Tình trạng công tác",
    ),
    "job_position": (
        "Vị trí việc làm",
        "Vị trí",
        "Chức danh",
        "Chức vụ",
    ),
}

REQUIRED_FIELDS = {
    "commune_name",
    "school_name",
    "staff_code",
    "full_name",
    "employment_status",
    "job_position",
}

ACTIVE_STATUS_KEYS = {
    "dang lam viec",
    "dang cong tac",
    "cong tac",
    "dang lam",
}


@dataclass(frozen=True)
class CommuneItem:
    id: int
    code: str
    name: str


@dataclass(frozen=True)
class SchoolItem:
    id: int
    name: str
    commune_id: int | None
    commune_name: str


def configure_console() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except (AttributeError, OSError):
        pass


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def comparison_key(value: Any) -> str:
    text = clean_text(value).replace("Đ", "D").replace("đ", "d")
    text = unicodedata.normalize("NFD", text)
    text = "".join(
        char for char in text if unicodedata.category(char) != "Mn"
    )
    text = text.lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def commune_key(value: Any) -> str:
    key = comparison_key(value)
    for prefix in ("uy ban nhan dan ", "ubnd "):
        if key.startswith(prefix):
            key = key[len(prefix):].strip()
    return key


def school_key(value: Any) -> str:
    key = comparison_key(value)
    replacements = (
        (r"\btruong mam non\b", " mn "),
        (r"\bco so mam non\b", " cs mn "),
        (r"\btruong tieu hoc\b", " th "),
        (r"\btruong trung hoc co so\b", " thcs "),
        (r"\btruong trung hoc pho thong\b", " thpt "),
    )
    for pattern, replacement in replacements:
        key = re.sub(pattern, replacement, key)
    return " ".join(key.split())


def normalize_staff_code(value: Any) -> str:
    text = clean_text(value).upper()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return re.sub(r"\s+", "", text)


def format_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return clean_text(value)


def determine_role(job_position: str) -> tuple[str, str] | None:
    key = comparison_key(job_position)
    if "giao vien" in key:
        return "GIAO_VIEN", "Giáo viên"

    manager_terms = (
        "can bo quan ly",
        "hieu truong",
        "pho hieu truong",
        "giam doc",
        "pho giam doc",
    )
    if any(term in key for term in manager_terms):
        return "TRUONG", "Trường"
    return None


def is_active_status(value: str) -> bool:
    key = comparison_key(value)
    return key in ACTIVE_STATUS_KEYS or any(
        status_key in key for status_key in ACTIVE_STATUS_KEYS
    )


def make_username(role_code: str, staff_code: str) -> str:
    prefix = "gv" if role_code == "GIAO_VIEN" else "cbql"
    safe_code = comparison_key(staff_code).replace(" ", ".")
    safe_code = re.sub(r"[^a-z0-9.]+", "", safe_code).strip(".")
    return f"{prefix}.{safe_code}"[:100]


def discover_headers() -> tuple[Any, Any, int, dict[str, int]]:
    """Mở workbook và tìm dòng tiêu đề thực tế.

    Một số file Excel xuất từ hệ thống có vùng ``dimension`` khai báo sai
    (ví dụ chỉ A1:BS4 dù dữ liệu có hơn 16.000 dòng). Ở chế độ read-only,
    openpyxl tin vào vùng này và chỉ đọc bốn dòng đầu. ``reset_dimensions``
    buộc openpyxl đọc toàn bộ các dòng thực tế trong XML.
    """
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"Không tìm thấy tệp: {INPUT_FILE}")

    workbook = load_workbook(INPUT_FILE, read_only=True, data_only=True)

    for worksheet in workbook.worksheets:
        if hasattr(worksheet, "reset_dimensions"):
            worksheet.reset_dimensions()

        rows = worksheet.iter_rows(
            min_row=1,
            max_row=40,
            values_only=True,
        )
        for row_number, row in enumerate(rows, start=1):
            values = [clean_text(value) for value in row]
            keys = [comparison_key(value) for value in values]

            mapping: dict[str, int] = {}
            for field, aliases in HEADER_ALIASES.items():
                alias_keys = {comparison_key(alias) for alias in aliases}
                for index, key in enumerate(keys):
                    if key in alias_keys:
                        mapping[field] = index
                        break

            if REQUIRED_FIELDS.issubset(mapping):
                return workbook, worksheet, row_number, mapping

    workbook.close()
    raise ValueError(
        "Không tìm thấy dòng tiêu đề có đủ các cột bắt buộc. "
        "Hãy kiểm tra lại tên cột trong file Excel."
    )


def load_database_catalogs() -> tuple[list[CommuneItem], list[SchoolItem], set[str]]:
    connection = sqlite3.connect(str(DATABASE_PATH))
    connection.row_factory = sqlite3.Row

    communes = [
        CommuneItem(
            id=int(row["id"]),
            code=clean_text(row["code"]),
            name=clean_text(row["name"]),
        )
        for row in connection.execute(
            "SELECT id, code, name FROM communes ORDER BY name"
        )
    ]

    schools = [
        SchoolItem(
            id=int(row["id"]),
            name=clean_text(row["name"]),
            commune_id=(int(row["commune_id"]) if row["commune_id"] is not None else None),
            commune_name=clean_text(row["commune_name"]),
        )
        for row in connection.execute(
            """
            SELECT s.id, s.name, s.commune_id, c.name AS commune_name
            FROM schools AS s
            LEFT JOIN communes AS c ON c.id = s.commune_id
            ORDER BY s.name
            """
        )
    ]

    usernames = {
        comparison_key(row["username"]).replace(" ", ".")
        for row in connection.execute("SELECT username FROM users")
        if clean_text(row["username"])
    }

    connection.close()
    return communes, schools, usernames


def build_indexes(
    communes: Iterable[CommuneItem],
    schools: Iterable[SchoolItem],
) -> tuple[
    dict[str, list[CommuneItem]],
    dict[tuple[int, str], list[SchoolItem]],
    dict[str, list[SchoolItem]],
    dict[str, list[SchoolItem]],
]:
    commune_index: dict[str, list[CommuneItem]] = defaultdict(list)
    school_by_commune: dict[tuple[int, str], list[SchoolItem]] = defaultdict(list)
    school_global: dict[str, list[SchoolItem]] = defaultdict(list)
    school_exact_global: dict[str, list[SchoolItem]] = defaultdict(list)

    for commune in communes:
        commune_index[commune_key(commune.name)].append(commune)

    for school in schools:
        fuzzy_key = school_key(school.name)
        exact_key = comparison_key(school.name)
        school_global[fuzzy_key].append(school)
        school_exact_global[exact_key].append(school)
        if school.commune_id is not None:
            school_by_commune[(school.commune_id, fuzzy_key)].append(school)

    return (
        commune_index,
        school_by_commune,
        school_global,
        school_exact_global,
    )


def write_csv(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    configure_console()

    print("=" * 78)
    print("BÀI 11B-6A.4 - KIỂM TRA DANH SÁCH GIÁO VIÊN")
    print("=" * 78)
    print(f"Tệp đầu vào: {INPUT_FILE}")
    print(f"Cơ sở dữ liệu: {DATABASE_PATH}")
    print()
    print("Chế độ hiện tại: CHỈ KIỂM TRA, KHÔNG GHI VÀO CƠ SỞ DỮ LIỆU.")
    print()

    workbook, worksheet, header_row, columns = discover_headers()
    print(
        f"Đã tìm thấy tiêu đề tại sheet '{worksheet.title}', "
        f"dòng {header_row}."
    )

    communes, schools, existing_usernames = load_database_catalogs()
    (
        commune_index,
        school_by_commune,
        school_global,
        school_exact_global,
    ) = build_indexes(communes, schools)

    def value_at(values: tuple[Any, ...], field: str) -> Any:
        index = columns[field]
        return values[index] if index < len(values) else None

    raw_rows: list[dict[str, Any]] = []
    data_rows = worksheet.iter_rows(
        min_row=header_row + 1,
        values_only=True,
    )
    for excel_row, values in enumerate(
        data_rows,
        start=header_row + 1,
    ):
        full_name = clean_text(value_at(values, "full_name"))
        school_name = clean_text(value_at(values, "school_name"))
        commune_name = clean_text(value_at(values, "commune_name"))
        staff_code = normalize_staff_code(value_at(values, "staff_code"))

        if not any((full_name, school_name, commune_name, staff_code)):
            continue

        raw_rows.append(
            {
                "excel_row": excel_row,
                "commune_name": commune_name,
                "school_name": school_name,
                "staff_code": staff_code,
                "full_name": full_name,
                "date_of_birth": format_date(
                    value_at(values, "date_of_birth")
                    if "date_of_birth" in columns
                    else ""
                ),
                "employment_status": clean_text(
                    value_at(values, "employment_status")
                ),
                "job_position": clean_text(
                    value_at(values, "job_position")
                ),
            }
        )

    workbook.close()

    code_counts = Counter(
        row["staff_code"] for row in raw_rows if row["staff_code"]
    )

    ready_rows: list[dict[str, Any]] = []
    error_rows: list[dict[str, Any]] = []
    duplicate_rows: list[dict[str, Any]] = []
    ignored_rows: list[dict[str, Any]] = []

    for row in raw_rows:
        reasons: list[str] = []

        if not is_active_status(row["employment_status"]):
            ignored_rows.append({**row, "reason": "Không ở trạng thái đang làm việc"})
            continue

        role = determine_role(row["job_position"])
        if role is None:
            ignored_rows.append(
                {
                    **row,
                    "reason": "Không thuộc nhóm Giáo viên hoặc Cán bộ quản lý",
                }
            )
            continue

        role_code, role_name = role

        if not row["staff_code"]:
            reasons.append("Thiếu mã định danh Bộ GD&ĐT")
        elif code_counts[row["staff_code"]] > 1:
            reasons.append("Mã định danh Bộ GD&ĐT bị trùng trong Excel")
            duplicate_rows.append(
                {**row, "duplicate_count": code_counts[row["staff_code"]]}
            )

        commune_candidates = commune_index.get(commune_key(row["commune_name"]), [])
        commune: CommuneItem | None = None
        if len(commune_candidates) == 1:
            commune = commune_candidates[0]
        elif len(commune_candidates) > 1:
            reasons.append("Tên xã/phường khớp nhiều danh mục")
        elif row["commune_name"]:
            reasons.append("Không tìm thấy xã/phường trong cơ sở dữ liệu")

        school_candidates: list[SchoolItem] = []
        school: SchoolItem | None = None
        current_school_key = school_key(row["school_name"])
        current_school_exact_key = comparison_key(row["school_name"])

        # Ưu tiên 1: khớp theo xã/phường và tên trường đã chuẩn hóa.
        if commune is not None:
            school_candidates = school_by_commune.get(
                (commune.id, current_school_key), []
            )

        # Ưu tiên 2: khớp chính xác tên trường trên toàn hệ thống.
        # Quy tắc này xử lý đúng trường trực thuộc Sở không có xã/phường,
        # ví dụ "Trường Mầm non Hoa Sen".
        if not school_candidates:
            exact_candidates = school_exact_global.get(
                current_school_exact_key, []
            )

            if commune is not None:
                exact_same_commune = [
                    candidate
                    for candidate in exact_candidates
                    if candidate.commune_id == commune.id
                ]
                if len(exact_same_commune) == 1:
                    school_candidates = exact_same_commune
            else:
                if len(exact_candidates) == 1:
                    school_candidates = exact_candidates
                elif len(exact_candidates) > 1:
                    direct_schools = [
                        candidate
                        for candidate in exact_candidates
                        if candidate.commune_id is None
                    ]
                    if len(direct_schools) == 1:
                        school_candidates = direct_schools

        # Ưu tiên 3: khớp tên rút gọn như MN/Trường Mầm non.
        if not school_candidates:
            global_candidates = school_global.get(current_school_key, [])

            if commune is None:
                direct_schools = [
                    candidate
                    for candidate in global_candidates
                    if candidate.commune_id is None
                ]
                if len(direct_schools) == 1:
                    school_candidates = direct_schools
                elif len(global_candidates) == 1:
                    candidate = global_candidates[0]
                    school_candidates = [candidate]
                    if candidate.commune_id is not None:
                        commune = next(
                            (
                                item
                                for item in communes
                                if item.id == candidate.commune_id
                            ),
                            None,
                        )
            else:
                same_commune = [
                    candidate
                    for candidate in global_candidates
                    if candidate.commune_id == commune.id
                ]
                if len(same_commune) == 1:
                    school_candidates = same_commune

        if len(school_candidates) == 1:
            school = school_candidates[0]
        elif len(school_candidates) > 1:
            reasons.append("Tên trường khớp nhiều danh mục")
        else:
            reasons.append("Không tìm thấy trường tương ứng với xã/phường")

        username = (
            make_username(role_code, row["staff_code"])
            if row["staff_code"]
            else ""
        )
        if username and comparison_key(username).replace(" ", ".") in existing_usernames:
            reasons.append("Tên đăng nhập dự kiến đã tồn tại trong users")

        result = {
            **row,
            "role_code": role_code,
            "role_name": role_name,
            "username_proposed": username,
            "commune_id": commune.id if commune else "",
            "commune_code": commune.code if commune else "",
            "commune_name_db": commune.name if commune else "",
            "school_id": school.id if school else "",
            "school_name_db": school.name if school else "",
            "school_commune_name_db": school.commune_name if school else "",
        }

        if reasons:
            error_rows.append(
                {**result, "reason": "; ".join(dict.fromkeys(reasons))}
            )
        else:
            ready_rows.append(result)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    export_dir = EXPORT_ROOT / f"kiem_tra_giao_vien_{timestamp}"
    export_dir.mkdir(parents=True, exist_ok=True)

    base_headers = [
        "excel_row",
        "commune_name",
        "school_name",
        "staff_code",
        "full_name",
        "date_of_birth",
        "employment_status",
        "job_position",
    ]
    ready_headers = base_headers + [
        "role_code",
        "role_name",
        "username_proposed",
        "commune_id",
        "commune_code",
        "commune_name_db",
        "school_id",
        "school_name_db",
        "school_commune_name_db",
    ]

    write_csv(
        export_dir / "01_du_kien_tao_tai_khoan.csv",
        ready_headers,
        ready_rows,
    )
    write_csv(
        export_dir / "02_loi_can_kiem_tra.csv",
        ready_headers + ["reason"],
        error_rows,
    )
    write_csv(
        export_dir / "03_ma_dinh_danh_bi_trung.csv",
        base_headers + ["duplicate_count"],
        duplicate_rows,
    )
    write_csv(
        export_dir / "04_dong_khong_nhap.csv",
        base_headers + ["reason"],
        ignored_rows,
    )

    summary_rows = [
        {"chi_tieu": "Tổng dòng có dữ liệu", "so_luong": len(raw_rows)},
        {
            "chi_tieu": "Đủ điều kiện dự kiến tạo tài khoản",
            "so_luong": len(ready_rows),
        },
        {"chi_tieu": "Có lỗi cần kiểm tra", "so_luong": len(error_rows)},
        {
            "chi_tieu": "Dòng có mã định danh bị trùng",
            "so_luong": len(duplicate_rows),
        },
        {"chi_tieu": "Không thuộc phạm vi nhập", "so_luong": len(ignored_rows)},
    ]
    write_csv(
        export_dir / "00_tong_hop.csv",
        ["chi_tieu", "so_luong"],
        summary_rows,
    )

    print("KẾT QUẢ KIỂM TRA")
    print("-" * 78)
    for item in summary_rows:
        print(f"{item['chi_tieu']}: {item['so_luong']}")
    print()
    print(f"Đã tạo báo cáo tại: {export_dir}")
    print()
    print("CÁC TỆP BÁO CÁO:")
    print("  00_tong_hop.csv")
    print("  01_du_kien_tao_tai_khoan.csv")
    print("  02_loi_can_kiem_tra.csv")
    print("  03_ma_dinh_danh_bi_trung.csv")
    print("  04_dong_khong_nhap.csv")
    print()
    print("HOÀN TẤT: chưa có dữ liệu nào được ghi vào bảng users.")


if __name__ == "__main__":
    main()
