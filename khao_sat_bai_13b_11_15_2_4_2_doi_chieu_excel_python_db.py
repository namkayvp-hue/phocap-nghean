from __future__ import annotations

import ast
import os
import re
import sqlite3
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


PROJECT = Path(
    os.environ.get("PHOCAP_PROJECT", r"C:\PhoCap")
).resolve()

APP = PROJECT / "app"
TEMPLATE_ROOT = APP / "report_templates"
DB = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT = (
    EXPORTS
    / f"doi_chieu_cong_thuc_nguon_db_bai_13b_11_15_2_4_2_{STAMP}.txt"
)

SOURCE_FILES = (
    APP / "pcgd_mn_report_builders_v1.py",
    APP / "pcgd_xmc_report_builders_v1.py",
    APP / "routers" / "report_center.py",
    APP / "routers" / "pcgdmn_template_report.py",
    APP / "routers" / "mn_official_reports.py",
    APP / "routers" / "preschool_3_5_report.py",
    APP / "routers" / "report_inputs.py",
)

SKIP_DB_COLUMNS = {
    "id",
    "name",
    "code",
    "created_at",
    "updated_at",
    "created_by",
    "updated_by",
    "note",
    "notes",
    "status",
    "is_active",
}

RATIO_PHRASES = (
    "tỉ lệ gv/l",
    "tỷ lệ gv/l",
    "gv/lớp",
    "gv/lop",
    "giáo viên/lớp",
    "giao vien/lop",
    "tỉ lệ ph/lớp",
    "tỷ lệ ph/lớp",
    "tỉ lệ p/l",
    "tỷ lệ p/l",
    "ph/lớp",
    "ph/lop",
    "phòng/lớp",
    "phong/lop",
    "phòng học/lớp",
    "phong hoc/lop",
)

PERCENT_PHRASES = (
    "tỷ lệ %",
    "tỉ lệ %",
    "tỷ lệ huy động",
    "tỉ lệ huy động",
    "tỷ lệ hoàn thành",
    "tỉ lệ hoàn thành",
    "tỷ lệ gv đạt",
    "tỉ lệ gv đạt",
    "tỷ lệ trẻ",
    "tỉ lệ trẻ",
    "tỷ lệ biết chữ",
    "tỉ lệ biết chữ",
    "tiếp cận gd",
    "đạt chuẩn",
)

AVERAGE_PHRASES = (
    "bình quân",
    "b. quân",
    "trung bình",
    "bình quân 5",
)

REPORT_HINTS = {
    "MN_M1": (
        "_fill_mn_m1",
        "mn_m1",
        "pcgd_2025_mn_m1",
    ),
    "MN_02": (
        "_fill_mn02",
        "mn02",
        "pcgd_2025_mn_02",
    ),
    "MN_01_GV": (
        "mn_01_gv",
        "staff",
        "gv_report",
        "teacher",
    ),
    "MN_01_CSVC": (
        "mn_01_csvc",
        "csvc",
        "facility",
    ),
    "MN_TAICHINH": (
        "mn_taichinh",
        "tai_chinh",
        "taichinh",
        "financial",
    ),
    "TH_M1": (
        "_th_m1",
        "th_m1",
        "pcgd_2025_th_m1",
    ),
    "TH_02": (
        "_build_th02",
        "th02",
        "th_02",
    ),
    "TH_01_CSVC": (
        "th_01_csvc",
        "csvc",
        "facility",
    ),
    "THCS_M1": (
        "_build_thcs_m1",
        "thcs_m1",
    ),
    "THCS_M2": (
        "_build_thcs_m2",
        "thcs_m2",
    ),
    "THCS_M5": (
        "thcs_m5",
        "staff",
        "teacher",
    ),
    "THCS_TK": (
        "_build_thcs_tk",
        "thcs_tk",
    ),
    "THCS_CSVC": (
        "thcs_csvc",
        "csvc",
        "facility",
    ),
    "XMC_3": (
        "xmc_3",
        "xmc3",
        "literacy",
    ),
    "XMC_4": (
        "xmc_4",
        "xmc4",
        "literacy",
    ),
    "CMC_1": (
        "cmc_1",
        "cmc1",
        "literacy",
        "chong_mu",
    ),
}


def norm(value: object) -> str:
    return " ".join(
        str(value or "")
        .strip()
        .lower()
        .split()
    )


def report_key(path: Path) -> str:
    name = path.stem.upper()

    if "THCS_M1" in name:
        return "THCS_M1"

    if "THCS_M2" in name:
        return "THCS_M2"

    if "THCS_M5" in name:
        return "THCS_M5"

    if "THCS_TK" in name:
        return "THCS_TK"

    if "THCS_CSVC" in name:
        return "THCS_CSVC"

    if "TH_M1" in name:
        return "TH_M1"

    if "TH_02" in name:
        return "TH_02"

    if "TH_01_CSVC" in name:
        return "TH_01_CSVC"

    if "MN_M1" in name:
        return "MN_M1"

    if "MN_02" in name:
        return "MN_02"

    if "MN_01_GV" in name:
        return "MN_01_GV"

    if "MN_01_CSVC" in name:
        return "MN_01_CSVC"

    if "MN_TAICHINH" in name:
        return "MN_TAICHINH"

    if "XMC_3" in name:
        return "XMC_3"

    if "XMC_4" in name:
        return "XMC_4"

    if "CMC_1" in name:
        return "CMC_1"

    if "PCGDMN" in name:
        return "PCGDMN_MASTER"

    return name


def group_from_key(key: str) -> str:
    if key.startswith("MN") or key == "PCGDMN_MASTER":
        return "MẦM NON"

    if key.startswith("THCS"):
        return "THCS"

    if key.startswith("TH_"):
        return "TIỂU HỌC"

    if key.startswith("XMC") or key.startswith("CMC"):
        return "XÓA MÙ CHỮ"

    return "KHÁC"


def canonical_files() -> list[Path]:
    files: list[Path] = []

    master = TEMPLATE_ROOT / "Bieu_mau_PCGDMN_2025.xlsx"

    if master.exists():
        files.append(master)

    for folder_name in (
        "pcgd_mn_2025",
        "pcgd_xmc_2025",
    ):
        folder = TEMPLATE_ROOT / folder_name

        if folder.exists():
            files.extend(
                sorted(
                    path
                    for path in folder.glob("*.xlsx")
                    if path.is_file()
                )
            )

    seen = set()
    result = []

    for path in files:
        key = str(path.resolve()).lower()

        if key in seen:
            continue

        seen.add(key)
        result.append(path)

    return result


def merged_value(
    ws,
    row: int,
    col: int,
):
    cell = ws.cell(
        row=row,
        column=col,
    )

    if not isinstance(
        cell,
        MergedCell,
    ):
        return cell.value

    coord = cell.coordinate

    for rng in ws.merged_cells.ranges:
        if coord in rng:
            return ws.cell(
                row=rng.min_row,
                column=rng.min_col,
            ).value

    return None


def semantic_context(
    ws,
    row: int,
    col: int,
) -> str:
    texts = []

    # Header dọc cùng cột.
    for rr in range(
        1,
        row,
    ):
        value = merged_value(
            ws,
            rr,
            col,
        )

        text = str(
            value or ""
        ).strip()

        if text and not text.startswith("="):
            texts.append(text)

    # Nhãn trái cùng dòng.
    for cc in range(
        1,
        col,
    ):
        value = merged_value(
            ws,
            row,
            cc,
        )

        text = str(
            value or ""
        ).strip()

        if text and not text.startswith("="):
            texts.append(text)

    # Chỉ giữ các mục cuối và mục chứa keyword nghiệp vụ.
    selected = []
    keywords = (
        "tỷ",
        "tỉ",
        "lớp",
        "phòng",
        "gv",
        "giáo viên",
        "huy động",
        "hoàn thành",
        "khuyết tật",
        "tiếp cận",
        "chuẩn",
        "biết chữ",
        "mù chữ",
        "bình quân",
        "trung bình",
        "tổng số",
    )

    for text in texts:
        n = norm(text)

        if any(
            kw in n
            for kw in keywords
        ):
            selected.append(text)

    for text in texts[-8:]:
        if text not in selected:
            selected.append(text)

    return " | ".join(
        selected[-12:]
    )


def parse_formula(
    formula: str,
) -> tuple[str, str, bool]:
    text = str(
        formula or ""
    ).strip()

    if text.startswith("="):
        text = text[1:].strip()

    percent_signal = False

    if text.endswith("%"):
        percent_signal = True
        text = text[:-1].strip()

    m = re.fullmatch(
        r"(?is)(.+?)\*\s*100(?:\.0+)?",
        text,
    )

    if m:
        percent_signal = True
        text = m.group(1).strip()

    outer = re.fullmatch(
        r"(?is)SUM\((.*)\)",
        text,
    )

    if outer and "/" in outer.group(1):
        text = outer.group(1).strip()

    depth = 0

    for idx, ch in enumerate(text):
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(
                0,
                depth - 1,
            )
        elif ch == "/" and depth == 0:
            numerator = text[:idx].strip()
            denominator = text[idx + 1:].strip()

            numerator = re.sub(
                r"(?is)^SUM\((.*)\)$",
                r"\1",
                numerator,
            )

            return (
                numerator,
                denominator,
                percent_signal,
            )

    return (
        "",
        "",
        percent_signal,
    )


def classify_formula(
    context: str,
    number_format: str,
    percent_signal: bool,
    denominator: str,
) -> str:
    text = norm(context)
    fmt = norm(number_format)

    # Ratio phải ưu tiên trước percent vì mẫu có thể dùng chữ "tỉ lệ Ph/Lớp".
    if any(
        phrase in text
        for phrase in RATIO_PHRASES
    ):
        return "RATIO"

    if any(
        phrase in text
        for phrase in AVERAGE_PHRASES
    ):
        return "AVERAGE"

    # Chia cho hằng số 5 ở báo cáo tài chính thường là bình quân.
    if denominator.strip() == "5":
        return "AVERAGE_OR_DERIVED"

    if (
        percent_signal
        or "%" in fmt
        or any(
            phrase in text
            for phrase in PERCENT_PHRASES
        )
    ):
        return "PERCENT"

    return "DERIVED_DIVISION"


def load_db_schema() -> dict[str, set[str]]:
    result: dict[str, set[str]] = {}

    if not DB.exists():
        return result

    conn = sqlite3.connect(
        str(DB)
    )

    try:
        tables = [
            row[0]
            for row in conn.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type='table'
                ORDER BY name
                """
            ).fetchall()
        ]

        for table in tables:
            try:
                columns = {
                    str(row[1])
                    for row in conn.execute(
                        f'PRAGMA table_info("{table}")'
                    ).fetchall()
                }
            except sqlite3.DatabaseError:
                continue

            if columns:
                result[
                    table
                ] = columns
    finally:
        conn.close()

    return result


def source_functions() -> list[dict]:
    records = []

    for path in SOURCE_FILES:
        if not path.exists():
            continue

        try:
            source = path.read_text(
                encoding="utf-8-sig",
            )
        except UnicodeDecodeError:
            source = path.read_text(
                encoding="utf-8",
                errors="replace",
            )

        try:
            tree = ast.parse(
                source
            )
        except SyntaxError:
            continue

        lines = source.splitlines()

        for node in ast.walk(tree):
            if not isinstance(
                node,
                (
                    ast.FunctionDef,
                    ast.AsyncFunctionDef,
                ),
            ):
                continue

            start = int(
                node.lineno
            )
            end = int(
                node.end_lineno
                or node.lineno
            )

            body_lines = lines[
                start - 1:
                end
            ]

            body = "\n".join(
                body_lines
            )

            records.append(
                {
                    "file": str(
                        path.relative_to(
                            PROJECT
                        )
                    ),
                    "name": node.name,
                    "start": start,
                    "end": end,
                    "body": body,
                    "lines": body_lines,
                }
            )

    return records


def cell_assignments(
    function: dict,
    cell_ref: str,
) -> list[str]:
    rows = []

    q1 = f'["{cell_ref}"]'
    q2 = f"['{cell_ref}']"

    for offset, line in enumerate(
        function[
            "lines"
        ],
        start=function[
            "start"
        ],
    ):
        if (
            q1 not in line
            and q2 not in line
        ):
            continue

        rows.append(
            f"{function['file']}:{offset}: {line.strip()}"
        )

    return rows


def relevance_score(
    function: dict,
    report: str,
) -> int:
    name = norm(
        function[
            "name"
        ]
    )

    body = norm(
        function[
            "body"
        ]
    )

    score = 0

    hints = REPORT_HINTS.get(
        report,
        (),
    )

    for hint in hints:
        h = norm(
            hint
        )

        if h in name:
            score += 120

        if h in body:
            score += 35

    # Report code literals give high confidence.
    literal = norm(
        report
    )

    if literal in body:
        score += 80

    # Generic group-aware scoring.
    group = group_from_key(
        report
    )

    file_name = norm(
        function[
            "file"
        ]
    )

    if group == "MẦM NON":
        if (
            "pcgd_mn_report_builders_v1.py" in file_name
            or "pcgdmn_template_report.py" in file_name
            or "mn_official_reports.py" in file_name
        ):
            score += 20

    elif group in {
        "TIỂU HỌC",
        "THCS",
        "XÓA MÙ CHỮ",
    }:
        if (
            "pcgd_xmc_report_builders_v1.py" in file_name
            or "report_center.py" in file_name
        ):
            score += 20

    return score


def db_dependencies(
    function: dict,
    schema: dict[str, set[str]],
) -> list[str]:
    body = function[
        "body"
    ]

    found = []

    for table, columns in schema.items():
        for column in columns:
            if (
                column in SKIP_DB_COLUMNS
                or len(
                    column
                ) < 4
            ):
                continue

            if re.search(
                rf"\b{re.escape(column)}\b",
                body,
            ):
                found.append(
                    f"{table}.{column}"
                )

    # Ưu tiên bảng điều tra/năm học, staff, csvc, finance.
    priority_terms = (
        "survey",
        "year",
        "person",
        "staff",
        "teacher",
        "class",
        "school",
        "facility",
        "csvc",
        "finance",
        "report",
    )

    found.sort(
        key=lambda item: (
            -sum(
                1
                for term in priority_terms
                if term in item.lower()
            ),
            item,
        )
    )

    return found[:45]


def classify_assignment(
    lines: list[str],
    formula_kind: str,
) -> str:
    joined = " ".join(
        lines
    ).lower()

    if not lines:
        return "NO_PYTHON_ASSIGNMENT"

    if re.search(
        r"=\s*none(?:\s*;|$)",
        joined,
    ):
        return "EXPLICIT_NONE_MISSING"

    if formula_kind == "PERCENT":
        if any(
            token in joined
            for token in (
                "_percent(",
                "_percentage(",
                "_safe_percent(",
                "_th_m1_percent(",
            )
        ):
            return "PYTHON_PERCENT_HELPER"

    if formula_kind == "RATIO":
        if any(
            token in joined
            for token in (
                "_ratio(",
                "_safe_ratio(",
                "teacher_class_ratio",
            )
        ):
            return "PYTHON_RATIO_HELPER"

    if formula_kind.startswith(
        "AVERAGE"
    ):
        return "PYTHON_ASSIGNMENT_REVIEW_AVERAGE"

    return "PYTHON_ASSIGNMENT_NEEDS_FORMULA_COMPARE"


def main() -> None:
    print("=" * 148)
    print(
        "BÀI 13B-11.15.2.4.2 - "
        "ĐỐI CHIẾU CÔNG THỨC EXCEL ↔ PYTHON ↔ NGUỒN DATABASE"
    )
    print("=" * 148)
    print()
    print("MỤC TIÊU:")
    print(
        " - Ghép theo đúng BIỂU/HÀM, không chỉ cùng tọa độ ô."
    )
    print(
        " - Tách PERCENT khỏi RATIO GV/lớp, phòng/lớp."
    )
    print(
        " - Tách phép chia trung bình/derived không phải tỷ lệ."
    )
    print(
        " - Phát hiện Python đang để None ở ô cần tính."
    )
    print(
        " - Liệt kê các cột DB được hàm báo cáo tham chiếu."
    )
    print(
        " - Chưa sửa source/database/Excel."
    )
    print()

    files = canonical_files()

    if not files:
        raise SystemExit(
            "Không tìm thấy file mẫu chính thức."
        )

    functions = source_functions()
    schema = load_db_schema()

    rows = []

    for path in files:
        report = report_key(
            path
        )

        group = group_from_key(
            report
        )

        wb = load_workbook(
            path,
            data_only=False,
            read_only=False,
        )

        try:
            for ws in wb.worksheets:
                for excel_row in ws.iter_rows():
                    for cell in excel_row:
                        formula = cell.value

                        if (
                            not isinstance(
                                formula,
                                str,
                            )
                            or not formula.startswith("=")
                            or "/" not in formula
                        ):
                            continue

                        numerator, denominator, percent_signal = parse_formula(
                            formula
                        )

                        context = semantic_context(
                            ws,
                            cell.row,
                            cell.column,
                        )

                        kind = classify_formula(
                            context,
                            cell.number_format,
                            percent_signal,
                            denominator,
                        )

                        candidates = []

                        for fn in functions:
                            assignments = cell_assignments(
                                fn,
                                cell.coordinate,
                            )

                            if not assignments:
                                continue

                            score = relevance_score(
                                fn,
                                report,
                            )

                            candidates.append(
                                {
                                    "score": score,
                                    "function": fn,
                                    "assignments": assignments,
                                }
                            )

                        candidates.sort(
                            key=lambda item: item[
                                "score"
                            ],
                            reverse=True,
                        )

                        strong = [
                            item
                            for item in candidates
                            if item[
                                "score"
                            ] >= 80
                        ]

                        selected = (
                            strong[:3]
                            if strong
                            else candidates[:2]
                        )

                        assignment_lines = []

                        for item in selected:
                            assignment_lines.extend(
                                item[
                                    "assignments"
                                ]
                            )

                        assignment_status = classify_assignment(
                            assignment_lines,
                            kind,
                        )

                        deps = []

                        if selected:
                            deps = db_dependencies(
                                selected[0][
                                    "function"
                                ],
                                schema,
                            )

                        confidence = (
                            "CAO"
                            if strong
                            else (
                                "THẤP"
                                if candidates
                                else "KHÔNG_CÓ"
                            )
                        )

                        rows.append(
                            {
                                "group": group,
                                "report": report,
                                "file": str(
                                    path.relative_to(
                                        PROJECT
                                    )
                                ),
                                "sheet": ws.title,
                                "cell": cell.coordinate,
                                "formula": formula,
                                "numerator": numerator,
                                "denominator": denominator,
                                "kind": kind,
                                "context": context,
                                "confidence": confidence,
                                "assignment_status": assignment_status,
                                "selected": selected,
                                "db_dependencies": deps,
                            }
                        )
        finally:
            wb.close()

    by_group = defaultdict(
        list
    )

    for row in rows:
        by_group[
            row[
                "group"
            ]
        ].append(
            row
        )

    report_lines = []

    report_lines.append(
        "=" * 148
        + "\n"
    )
    report_lines.append(
        "BÀI 13B-11.15.2.4.2 - ĐỐI CHIẾU CÔNG THỨC EXCEL / PYTHON / DATABASE\n"
    )
    report_lines.append(
        "=" * 148
        + "\n"
    )
    report_lines.append(
        f"Project: {PROJECT}\n"
    )
    report_lines.append(
        f"Thời điểm: {datetime.now():%d/%m/%Y %H:%M:%S}\n"
    )
    report_lines.append(
        f"File mẫu: {len(files)}\n"
    )
    report_lines.append(
        f"Hàm Python khảo sát: {len(functions)}\n"
    )
    report_lines.append(
        f"Bảng DB khảo sát schema: {len(schema)}\n\n"
    )

    report_lines.append(
        "QUY TẮC\n"
    )
    report_lines.append(
        "-" * 148
        + "\n"
    )
    report_lines.append(
        "PERCENT = tử số / mẫu số * 100; mẫu số 0/chưa có -> trống.\n"
    )
    report_lines.append(
        "RATIO = tử số / mẫu số; không nhân 100 (ví dụ GV/lớp, phòng/lớp).\n"
    )
    report_lines.append(
        "AVERAGE/DERIVED = phép chia phục vụ tính trung bình/chỉ tiêu khác, không gọi là tỷ lệ %.\n"
    )
    report_lines.append(
        "Chỉ assignment có độ tin cậy CAO mới được xem là ứng viên cùng biểu.\n"
    )
    report_lines.append(
        "Bài này chưa tự kết luận assignment đúng tử số/mẫu số; chỉ chuẩn bị dữ liệu để chốt Bài 15.2.4.3.\n"
    )

    critical = []

    for group in (
        "MẦM NON",
        "TIỂU HỌC",
        "THCS",
        "XÓA MÙ CHỮ",
        "KHÁC",
    ):
        items = by_group.get(
            group,
            [],
        )

        if not items:
            continue

        report_lines.append(
            "\n\n"
            + "=" * 148
            + "\n"
        )
        report_lines.append(
            f"NHÓM: {group}\n"
        )
        report_lines.append(
            "=" * 148
            + "\n"
        )

        type_counter = Counter(
            item[
                "kind"
            ]
            for item in items
        )

        status_counter = Counter(
            item[
                "assignment_status"
            ]
            for item in items
        )

        confidence_counter = Counter(
            item[
                "confidence"
            ]
            for item in items
        )

        report_lines.append(
            f"Tổng phép chia: {len(items)}\n"
        )

        report_lines.append(
            "Phân loại:\n"
        )

        for key, count in sorted(
            type_counter.items()
        ):
            report_lines.append(
                f" - {key}: {count}\n"
            )

        report_lines.append(
            "Python assignment:\n"
        )

        for key, count in sorted(
            status_counter.items()
        ):
            report_lines.append(
                f" - {key}: {count}\n"
            )

        report_lines.append(
            "Độ tin cậy ghép biểu/hàm:\n"
        )

        for key, count in sorted(
            confidence_counter.items()
        ):
            report_lines.append(
                f" - {key}: {count}\n"
            )

        for index, item in enumerate(
            items,
            start=1,
        ):
            report_lines.append(
                f"\n[{index}] {item['report']} | "
                f"{item['sheet']}!{item['cell']}\n"
            )
            report_lines.append(
                f"    File       : {item['file']}\n"
            )
            report_lines.append(
                f"    Loại       : {item['kind']}\n"
            )
            report_lines.append(
                f"    Excel      : {item['formula']}\n"
            )
            report_lines.append(
                f"    Tử số      : {item['numerator'] or '(chưa tách)'}\n"
            )
            report_lines.append(
                f"    Mẫu số     : {item['denominator'] or '(chưa tách)'}\n"
            )
            report_lines.append(
                f"    Ngữ cảnh   : {item['context'] or '(không có)'}\n"
            )
            report_lines.append(
                f"    Ghép Python: {item['confidence']}\n"
            )
            report_lines.append(
                f"    Trạng thái : {item['assignment_status']}\n"
            )

            if item[
                "selected"
            ]:
                for cand in item[
                    "selected"
                ]:
                    fn = cand[
                        "function"
                    ]

                    report_lines.append(
                        f"    Hàm ứng viên [{cand['score']}]: "
                        f"{fn['file']}::{fn['name']} "
                        f"(dòng {fn['start']}-{fn['end']})\n"
                    )

                    for assignment in cand[
                        "assignments"
                    ]:
                        report_lines.append(
                            f"      - {assignment}\n"
                        )
            else:
                report_lines.append(
                    "    Hàm ứng viên: không phát hiện assignment đúng ô\n"
                )

            if item[
                "db_dependencies"
            ]:
                report_lines.append(
                    "    Cột DB xuất hiện trong hàm ứng viên:\n"
                )

                for dep in item[
                    "db_dependencies"
                ]:
                    report_lines.append(
                        f"      - {dep}\n"
                    )
            else:
                report_lines.append(
                    "    Cột DB: chưa xác định từ source ứng viên\n"
                )

            if item[
                "assignment_status"
            ] in {
                "EXPLICIT_NONE_MISSING",
                "NO_PYTHON_ASSIGNMENT",
            }:
                critical.append(
                    item
                )

    report_lines.append(
        "\n\n"
        + "=" * 148
        + "\n"
    )
    report_lines.append(
        "DANH SÁCH ƯU TIÊN CHO BÀI 15.2.4.3\n"
    )
    report_lines.append(
        "=" * 148
        + "\n"
    )

    if not critical:
        report_lines.append(
            "Không phát hiện ô None/không có assignment trong phạm vi quét.\n"
        )
    else:
        for item in critical:
            report_lines.append(
                f" - {item['group']} | {item['report']} | "
                f"{item['sheet']}!{item['cell']} | "
                f"{item['kind']} | {item['assignment_status']} | "
                f"{item['formula']}\n"
            )

    report_lines.append(
        "\nBƯỚC SAU:\n"
    )
    report_lines.append(
        "1. Chốt những ô có ghép Python CAO: so sánh helper/biến Python với tử số-mẫu số của Excel.\n"
    )
    report_lines.append(
        "2. Với ô NO_PYTHON_ASSIGNMENT: xác định nguồn DB rồi mới cài công thức Python.\n"
    )
    report_lines.append(
        "3. Với EXPLICIT_NONE_MISSING: ưu tiên sửa trước vì đây là khoảng trống đã xác nhận.\n"
    )
    report_lines.append(
        "4. Sau khi tỷ lệ nguồn đủ mới làm Đạt/Không đạt và Bài 15.3 kiểm tra chất lượng.\n"
    )

    OUT.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    OUT.write_text(
        "".join(
            report_lines
        ),
        encoding="utf-8",
    )

    print("TÓM TẮT:")

    for group in (
        "MẦM NON",
        "TIỂU HỌC",
        "THCS",
        "XÓA MÙ CHỮ",
    ):
        items = by_group.get(
            group,
            [],
        )

        if not items:
            continue

        missing = sum(
            1
            for item in items
            if item[
                "assignment_status"
            ] in {
                "EXPLICIT_NONE_MISSING",
                "NO_PYTHON_ASSIGNMENT",
            }
        )

        print(
            f" - {group}: {len(items)} công thức chia; "
            f"{missing} ô ưu tiên chưa có/đang None"
        )

    print()
    print(
        "Source/database/Excel: KHÔNG THAY ĐỔI."
    )
    print(
        "Báo cáo:",
        OUT,
    )
    print()
    print("=" * 148)
    print(
        "BÀI 13B-11.15.2.4.2 KHẢO SÁT THÀNH CÔNG"
    )
    print("=" * 148)


if __name__ == "__main__":
    main()
