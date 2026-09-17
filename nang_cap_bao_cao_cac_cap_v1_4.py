from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import sys
import traceback
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

PROJECT = Path(r"C:\PhoCap")
VERSION = "PCGD-XMC-V1.4-CAP-XA-TRUONG-GIAO-VIEN"
MAIN = PROJECT / "app" / "main.py"
DB = PROJECT / "data" / "phocap.db"
ROUTER = PROJECT / "app" / "routers" / "report_center.py"
HTML = PROJECT / "app" / "templates" / "reports" / "report_center.html"
BUILDER = PROJECT / "app" / "pcgd_xmc_report_builders_v1.py"
PAYLOAD = Path(__file__).resolve().parent / "pcgd_xmc_report_builders_v1_4_payload.py"
TEMPLATE_ROOT = PROJECT / "app" / "report_templates" / "pcgd_xmc_2025"
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
BACKUP = PROJECT / "exports" / f"backup_pcgd_xmc_v1_4_cac_cap_{STAMP}"
POINTER = PROJECT / "exports" / "pcgd_xmc_v1_4_cac_cap_backup_moi_nhat.txt"
MANIFEST = BACKUP / "manifest.json"

REPORT_TEMPLATE_FILES = [
    "PCGD_2025_TH_02.xlsx",
    "PCGD_2025_TH_01_GV.xlsx",
    "PCGD_2025_TH_01_CSVC.xlsx",
    "PCGD_2025_THCS_M1.xlsx",
    "PCGD_2025_THCS_M2.xlsx",
    "PCGD_2025_THCS_TK.xlsx",
    "PCGD_2025_THCS_M5.xlsx",
    "PCGD_2025_THCS_CSVC.xlsx",
    "PCGD_2025_XMC_3.xlsx",
    "PCGD_2025_CMC_2.xlsx",
    "PCGD_2025_CMC_1.xlsx",
    "PCGD_2025_XMC_4.xlsx",
]


def sha(path: Path) -> str | None:
    if not path.exists() or not path.is_file():
        return None
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def route_signature(root: Path) -> list[tuple[str, str, str]]:
    pattern = re.compile(r'@router\.(get|post|put|delete|patch)\(\s*["\']([^"\']+)["\']')
    result = []
    for path in sorted((root / "app" / "routers").glob("*.py")):
        text = path.read_text(encoding="utf-8-sig")
        for method, route in pattern.findall(text):
            result.append((path.name, method, route))
    return result


def backup_file(path: Path, manifest: dict) -> None:
    rel = path.relative_to(PROJECT).as_posix()
    existed = path.exists()
    manifest[rel] = {"existed": existed}
    if existed:
        dest = BACKUP / rel
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, dest)


def restore(manifest: dict) -> None:
    print("\nDANG KHOI PHUC TRANG THAI TRUOC V1.4...")
    for rel, info in manifest.items():
        target = PROJECT / rel
        saved = BACKUP / rel
        if info.get("existed"):
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(saved, target)
        elif target.exists():
            target.unlink()
    print("DA KHOI PHUC AN TOAN.")


def patch_th_m1_lower_scopes(text: str) -> str:
    if "# === BAI_13C_1_TH_M1_V2_TOAN_TINH ===" not in text:
        raise RuntimeError(
            "Chua tim thay TH-M1 V2. Can giu nguyen ban dang xuat cap tinh "
            "truoc khi nang cap cac cap."
        )

    pattern = re.compile(
        r"    is_province_scope = selected_commune_id is None and selected_school_id is None\n"
        r".*?"
        r"    birth_years = \{",
        re.S,
    )
    match = pattern.search(text)
    if not match:
        raise RuntimeError("Khong tim thay khoi pham vi TH-M1 V2 de mo rong an toan.")

    replacement = '''    is_province_scope = selected_commune_id is None and selected_school_id is None
    is_school_scope = selected_school_id is not None
    scope_title = (
        "Toàn tỉnh"
        if is_province_scope
        else _th_m1_scope_title(scope_label, selected_commune_id)
    )

    if is_school_scope and scope_title.lower().startswith("trường trường "):
        scope_title = scope_title[len("Trường "):]

    if is_province_scope and ws.title != "Toàn tỉnh":
        ws.title = "Toàn tỉnh"

    ws["A1"] = "Tỉnh: Nghệ An"
    ws["A2"] = scope_title
    ws["E2"] = f"Thời điểm: ngày 30 tháng 9 năm {reference_year}"

    if is_province_scope:
        ws["J44"] = f"Nghệ An, ngày      tháng      năm {reference_year}"
        ws["J45"] = "XÁC NHẬN CỦA SỞ GIÁO DỤC VÀ ĐÀO TẠO"
        ws["J46"] = "GIÁM ĐỐC"
    elif is_school_scope:
        ws["J44"] = f"{scope_title}, ngày      tháng      năm {reference_year}"
        ws["J45"] = "XÁC NHẬN CỦA NHÀ TRƯỜNG"
        ws["J46"] = "HIỆU TRƯỞNG"
    else:
        ws["J44"] = f"{scope_title}, ngày      tháng      năm {reference_year}"
        ws["J45"] = "XÁC NHẬN CỦA UBND XÃ/PHƯỜNG"
        ws["J46"] = "PHÓ CHỦ TỊCH"

    birth_years = {'''
    return text[:match.start()] + replacement + text[match.end():]


def read_only_scope_checks() -> None:
    sys.path.insert(0, str(PROJECT))
    from sqlalchemy import select
    from app.database import SessionLocal
    from app.models import Role, User
    from app.routers import report_center as rc
    import app.pcgd_xmc_report_builders_v1 as builder

    with SessionLocal() as db:
        users = {}
        for code in ("XA", "TRUONG", "GIAO_VIEN"):
            stmt = (
                select(User)
                .join(Role, Role.id == User.role_id)
                .where(Role.code == code, User.is_active.is_(True))
                .order_by(User.id.asc())
            )
            user = db.scalars(stmt).first()
            if user is None:
                raise RuntimeError(f"Khong tim thay tai khoan mau cap {code} de kiem tra.")
            users[code] = user

        for code, user in users.items():
            auth_user = {
                "id": user.id,
                "username": user.username,
                "full_name": user.full_name,
                "role_code": code,
                "role_name": code,
                "commune_id": user.commune_id,
                "school_id": user.school_id,
                "unit_name": user.full_name,
            }
            request = SimpleNamespace(scope={"auth_user": auth_user})
            _communes, _schools, commune_id, school_id, _scope_label = rc._load_scope_options(
                db, auth_user, 99999999, 99999999
            )
            if code == "XA":
                if commune_id != user.commune_id:
                    raise RuntimeError("Pham vi cap XA khong khoa dung xa cua tai khoan.")
                if school_id is not None:
                    raise RuntimeError("Cap XA da nhan school_id khong hop le ngoai danh muc.")
            else:
                if school_id != user.school_id:
                    raise RuntimeError(f"Pham vi cap {code} khong khoa dung truong cua tai khoan.")

            eff_commune, eff_school, _eff_label, _kind = builder._effective_scope_for_user(
                db=db,
                request=request,
                selected_commune_id=99999999,
                selected_school_id=99999999,
                scope_label="Pham vi thu nghiem",
            )
            if code == "XA":
                if eff_commune != user.commune_id or eff_school is not None:
                    raise RuntimeError("Bo may xuat Excel cap XA chua khoa pham vi an toan.")
            else:
                if eff_school != user.school_id:
                    raise RuntimeError(f"Bo may xuat Excel cap {code} chua khoa dung truong.")

        from openpyxl import load_workbook
        for report_type, filename in builder.REPORT_TEMPLATE_FILES.items():
            path = builder.TEMPLATE_ROOT / filename
            if not path.exists():
                raise RuntimeError(f"Thieu mau: {filename}")
            for kind, commune_id, school_id, label in (
                ("xa", 1, None, "Phường kiểm tra"),
                ("truong", 1, 1, "Trường kiểm tra"),
            ):
                wb = load_workbook(path)
                ws = wb[wb.sheetnames[0]]
                meta = builder._scope_meta(
                    scope_label=label,
                    selected_commune_id=commune_id,
                    selected_school_id=school_id,
                    year=2026,
                )
                builder._write_scope_signature(ws, report_type, meta)
                if kind == "xa" and meta["signer"] != "PHÓ CHỦ TỊCH":
                    raise RuntimeError(f"Chu ky cap xa sai o {report_type}")
                if kind == "truong" and meta["signer"] != "HIỆU TRƯỞNG":
                    raise RuntimeError(f"Chu ky cap truong sai o {report_type}")


def main() -> int:
    print("=" * 78)
    print("BO BAO CAO PCGD & XMC V1.4 - CAP XA / TRUONG / GIAO VIEN")
    print("=" * 78)
    print("GIU NGUYEN: CHUC NANG CU - ROUTE - CSDL - LUONG NGHIEP VU")

    required = [PROJECT, MAIN, DB, ROUTER, HTML, BUILDER, PAYLOAD, TEMPLATE_ROOT]
    for path in required:
        if not path.exists():
            raise RuntimeError(f"Khong tim thay: {path}")
    for filename in REPORT_TEMPLATE_FILES:
        if not (TEMPLATE_ROOT / filename).exists():
            raise RuntimeError(f"Thieu mau Excel dang hoat dong: {filename}")

    before_main = sha(MAIN)
    before_db = sha(DB)
    before_html = sha(HTML)
    before_routes = route_signature(PROJECT)

    BACKUP.mkdir(parents=True, exist_ok=False)
    manifest: dict = {}
    for path in (ROUTER, BUILDER):
        backup_file(path, manifest)
    MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    POINTER.write_text(str(BACKUP), encoding="utf-8")
    print("Ban sao an toan:", BACKUP)

    try:
        print("\nBUOC 1 - MO RONG TH-M1 CHO XA / TRUONG / GIAO VIEN")
        old_router = ROUTER.read_text(encoding="utf-8-sig")
        ROUTER.write_text(patch_th_m1_lower_scopes(old_router), encoding="utf-8")
        print("TH-M1: cap tinh GIU NGUYEN; cap xa/truong/giao vien DA CO TIEU DE + CHU KY DUNG CAP")

        print("\nBUOC 2 - KHOA PHAM VI 12 BIEU THEO TAI KHOAN")
        shutil.copy2(PAYLOAD, BUILDER)
        print("Cap XA: chi xa cua tai khoan")
        print("Cap TRUONG: chi truong cua tai khoan")
        print("Cap GIAO VIEN: du lieu dieu tra theo phan cong; thong tin truong theo truong cua tai khoan")

        print("\nBUOC 3 - KIEM TRA CU PHAP")
        python = PROJECT / ".venv" / "Scripts" / "python.exe"
        if not python.exists():
            python = Path(sys.executable)
        subprocess.run(
            [str(python), "-m", "py_compile", str(ROUTER), str(BUILDER)],
            cwd=PROJECT,
            check=True,
        )
        print("Python: DAT")

        print("\nBUOC 4 - KIEM TRA PHAM VI 3 CAP BANG CSDL HIEN CO (CHI DOC)")
        read_only_scope_checks()
        print("Cap XA: DAT")
        print("Cap TRUONG: DAT")
        print("Cap GIAO VIEN: DAT")
        print("12 mau - chu ky cap xa/truong: DAT")

        print("\nBUOC 5 - KIEM TRA NGUYEN TAC BAT BUOC")
        if sha(MAIN) != before_main:
            raise RuntimeError("app/main.py bi thay doi.")
        if sha(DB) != before_db:
            raise RuntimeError("data/phocap.db bi thay doi.")
        if sha(HTML) != before_html:
            raise RuntimeError("report_center.html bi thay doi.")
        if route_signature(PROJECT) != before_routes:
            raise RuntimeError("Danh sach route bi thay doi.")
        router_text = ROUTER.read_text(encoding="utf-8-sig")
        builder_text = BUILDER.read_text(encoding="utf-8-sig")
        for token in (
            'ws["J45"] = "XÁC NHẬN CỦA NHÀ TRƯỜNG"',
            'ws["J46"] = "HIỆU TRƯỞNG"',
            'ws["J45"] = "XÁC NHẬN CỦA UBND XÃ/PHƯỜNG"',
        ):
            if token not in router_text:
                raise RuntimeError(f"TH-M1 thieu noi dung: {token}")
        for token in (
            'LOWER_SCOPE_REPORT_VERSION = "PCGD-XMC-REPORT-LOWER-SCOPES-V1.4"',
            'def _effective_scope_for_user(',
            'role_code == COMMUNE_ROLE_CODE',
            'role_code in {SCHOOL_ROLE_CODE, TEACHER_ROLE_CODE}',
        ):
            if token not in builder_text:
                raise RuntimeError(f"Bo may 12 bieu thieu: {token}")

        print("\n" + "=" * 78)
        print("NANG CAP BAO CAO CAC CAP V1.4 THANH CONG")
        print("=" * 78)
        print("CAP XA: XUAT BAO CAO DUNG PHAM VI XA")
        print("CAP TRUONG: XUAT BAO CAO DUNG PHAM VI TRUONG")
        print("CAP GIAO VIEN: XUAT THEO PHAM VI TAI KHOAN/PHAN CONG HIEN CO")
        print("TH-M1 + 12 BIEU CON LAI: DA SAN SANG CAC CAP")
        print("app/main.py: KHONG DOI")
        print("report_center.html: KHONG DOI")
        print("Route: KHONG DOI")
        print("data/phocap.db: KHONG DOI")
        print("Luong nghiep vu: KHONG DOI")
        print("Chuc nang cu: GIU NGUYEN")
        print("Ban sao an toan:", BACKUP)
        return 0

    except Exception as exc:
        print("\nNANG CAP KHONG THANH CONG:", exc)
        traceback.print_exc()
        restore(manifest)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
