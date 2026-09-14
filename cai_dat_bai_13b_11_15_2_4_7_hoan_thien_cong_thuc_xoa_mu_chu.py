from __future__ import annotations

import ast
import importlib.util
import os
import py_compile
import shutil
import sqlite3
from datetime import datetime
from pathlib import Path


PROJECT = Path(
    os.environ.get(
        "PHOCAP_PROJECT",
        r"C:\PhoCap",
    )
).resolve()

APP = PROJECT / "app"
EXPORTS = PROJECT / "exports"
DB = PROJECT / "data" / "phocap.db"

BUILDERS = (
    APP
    / "pcgd_xmc_report_builders_v1.py"
)

SERVICE = (
    APP
    / "services"
    / "xmc_report_v247.py"
)

STAMP = datetime.now().strftime(
    "%Y%m%d_%H%M%S"
)

BACKUP = (
    EXPORTS
    / f"backup_bai_13b_11_15_2_4_7_{STAMP}"
)

REPORT = (
    EXPORTS
    / f"bao_cao_bai_13b_11_15_2_4_7_{STAMP}.txt"
)

MARKER = (
    "# === BAI_13B_11_15_2_4_7_"
    "XMC_FORMULA_CONTRACT ==="
)

CALL_MARKER = (
    "# === BAI_13B_11_15_2_4_7_"
    "APPLY_XMC_BEFORE_SAVE ==="
)

SERVICE_CONTENT = 'from __future__ import annotations\n\n"""\nBÀI 13B-11.15.2.4.7\nBộ tính trực tiếp 79 tỷ lệ Xóa mù chữ theo mẫu Excel chính thức.\n\nNguyên tắc:\n- Chỉ tính từ các ô nguồn đã có trong workbook.\n- Không tự suy diễn/điền số liệu nguồn.\n- Mẫu số <= 0 hoặc thiếu nguồn -> để trống.\n- Giá trị tỷ lệ giữ thang 0..100.\n- Hiển thị theo yêu cầu: 100%, 50%, 0% (không phần thập phân).\n- Không phụ thuộc Excel recalculation.\n"""\n\nfrom typing import Any\n\n\nXMC_REPORT_TYPES = {\n    "PCGD_CMC_1_2025",\n    "PCGD_CMC_2_2025",\n    "PCGD_XMC_3_2025",\n    "PCGD_XMC_4_2025",\n}\n\n# 24 công thức CMC-1.\nCMC1_PERCENT_FORMULAS: tuple[\n    tuple[str, str, str],\n    ...,\n] = (\n    # 15-25 tuổi, mẫu số G8.\n    ("L8", "K8", "G8"),\n    ("N8", "M8", "G8"),\n    ("P8", "O8", "G8"),\n    ("R8", "Q8", "G8"),\n    ("T8", "S8", "G8"),\n    ("V8", "U8", "G8"),\n    ("X8", "W8", "G8"),\n    ("Z8", "Y8", "G8"),\n\n    # 15-35 tuổi, mẫu số AA8.\n    ("AF8", "AE8", "AA8"),\n    ("AH8", "AG8", "AA8"),\n    ("AJ8", "AI8", "AA8"),\n    ("AL8", "AK8", "AA8"),\n    ("AN8", "AM8", "AA8"),\n    ("AP8", "AO8", "AA8"),\n    ("AR8", "AQ8", "AA8"),\n    ("AT8", "AS8", "AA8"),\n\n    # 15-60 tuổi, mẫu số AU8.\n    ("AZ8", "AY8", "AU8"),\n    ("BB8", "BA8", "AU8"),\n    ("BD8", "BC8", "AU8"),\n    ("BF8", "BE8", "AU8"),\n    ("BH8", "BG8", "AU8"),\n    ("BJ8", "BI8", "AU8"),\n    ("BL8", "BK8", "AU8"),\n    ("BN8", "BM8", "AU8"),\n)\n\n# XMC-3: S9:S57 = O/C, gồm cả các dòng cộng 20, 31, 57.\nXMC3_PERCENT_FORMULAS: tuple[\n    tuple[str, str, str],\n    ...,\n] = tuple(\n    (\n        f"S{row}",\n        f"O{row}",\n        f"C{row}",\n    )\n    for row in range(9, 58)\n)\n\n# 6 công thức XMC-4.\nXMC4_PERCENT_FORMULAS: tuple[\n    tuple[str, str, str],\n    ...,\n] = (\n    ("E8", "D8", "C8"),\n    ("G8", "F8", "C8"),\n    ("J8", "I8", "H8"),\n    ("L8", "K8", "H8"),\n    ("O8", "N8", "M8"),\n    ("Q8", "P8", "M8"),\n)\n\n\ndef _number(\n    value: Any,\n) -> float | None:\n    if value is None:\n        return None\n\n    if isinstance(\n        value,\n        bool,\n    ):\n        return float(\n            int(value)\n        )\n\n    if isinstance(\n        value,\n        (int, float),\n    ):\n        return float(value)\n\n    text = str(\n        value\n    ).strip()\n\n    if not text:\n        return None\n\n    # Không dùng cached Excel formula; Python phải lấy\n    # từ các ô nguồn đã được builder ghi bằng số.\n    if text.startswith("="):\n        return None\n\n    compact = text.replace(\n        " ",\n        "",\n    )\n\n    if (\n        "," in compact\n        and "." in compact\n    ):\n        if (\n            compact.rfind(",")\n            > compact.rfind(".")\n        ):\n            compact = (\n                compact\n                .replace(\n                    ".",\n                    "",\n                )\n                .replace(\n                    ",",\n                    ".",\n                )\n            )\n        else:\n            compact = compact.replace(\n                ",",\n                "",\n            )\n\n    elif "," in compact:\n        compact = compact.replace(\n            ",",\n            ".",\n        )\n\n    try:\n        return float(\n            compact\n        )\n    except (\n        TypeError,\n        ValueError,\n    ):\n        return None\n\n\ndef _percent(\n    numerator: Any,\n    denominator: Any,\n) -> float | None:\n    n = _number(\n        numerator\n    )\n\n    d = _number(\n        denominator\n    )\n\n    if (\n        n is None\n        or d is None\n        or d <= 0\n    ):\n        return None\n\n    return round(\n        n * 100.0 / d,\n        6,\n    )\n\n\ndef _write_percent(\n    ws: Any,\n    target_ref: str,\n    numerator_ref: str,\n    denominator_ref: str,\n) -> None:\n    value = _percent(\n        ws[numerator_ref].value,\n        ws[denominator_ref].value,\n    )\n\n    cell = ws[target_ref]\n    cell.value = value\n\n    # Giá trị nội bộ 100 => hiển thị 100%.\n    # Dấu % là literal, không để Excel nhân thêm 100.\n    cell.number_format = r\'0\\%\'\n\n\ndef _apply_formulas(\n    ws: Any,\n    formulas: tuple[\n        tuple[str, str, str],\n        ...,\n    ],\n) -> None:\n    for (\n        target_ref,\n        numerator_ref,\n        denominator_ref,\n    ) in formulas:\n        _write_percent(\n            ws,\n            target_ref,\n            numerator_ref,\n            denominator_ref,\n        )\n\n\ndef apply_xmc_formula_contract(\n    report_type: str,\n    ws: Any,\n) -> None:\n    """\n    Tính toàn bộ tỷ lệ XMC đã được xác minh trong mẫu.\n\n    CMC-2 hiện không có phép chia/tỷ lệ trong ma trận 79\n    công thức, nên giữ nguyên builder hiện tại.\n    """\n    code = str(\n        report_type or ""\n    ).strip().upper()\n\n    if code == "PCGD_CMC_1_2025":\n        _apply_formulas(\n            ws,\n            CMC1_PERCENT_FORMULAS,\n        )\n        return\n\n    if code == "PCGD_XMC_3_2025":\n        _apply_formulas(\n            ws,\n            XMC3_PERCENT_FORMULAS,\n        )\n        return\n\n    if code == "PCGD_XMC_4_2025":\n        _apply_formulas(\n            ws,\n            XMC4_PERCENT_FORMULAS,\n        )\n        return\n\n    # PCGD_CMC_2_2025: không có tỷ lệ trong ma trận,\n    # không sửa dữ liệu.\n    return\n\n\ndef formula_count() -> dict[str, int]:\n    return {\n        "CMC_1": len(\n            CMC1_PERCENT_FORMULAS\n        ),\n        "XMC_3": len(\n            XMC3_PERCENT_FORMULAS\n        ),\n        "XMC_4": len(\n            XMC4_PERCENT_FORMULAS\n        ),\n        "TOTAL": (\n            len(\n                CMC1_PERCENT_FORMULAS\n            )\n            + len(\n                XMC3_PERCENT_FORMULAS\n            )\n            + len(\n                XMC4_PERCENT_FORMULAS\n            )\n        ),\n    }\n'


def read_text(
    path: Path,
) -> str:
    if not path.exists():
        raise RuntimeError(
            f"Không tìm thấy: {path}"
        )

    return path.read_text(
        encoding="utf-8-sig",
        errors="strict",
    )


def write_text(
    path: Path,
    value: str,
) -> None:
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        value,
        encoding="utf-8",
    )


def function_node(
    source: str,
    name: str,
):
    tree = ast.parse(
        source
    )

    matches = [
        node
        for node in tree.body
        if isinstance(
            node,
            (
                ast.FunctionDef,
                ast.AsyncFunctionDef,
            ),
        )
        and node.name == name
    ]

    if len(matches) != 1:
        raise RuntimeError(
            f"Cần đúng 1 hàm {name}; "
            f"tìm thấy {len(matches)}."
        )

    return matches[0]


def function_span(
    source: str,
    name: str,
) -> tuple[int, int]:
    node = function_node(
        source,
        name,
    )

    lines = source.splitlines(
        keepends=True
    )

    offsets = [0]

    for line in lines:
        offsets.append(
            offsets[-1]
            + len(line)
        )

    return (
        offsets[
            int(node.lineno) - 1
        ],
        offsets[
            int(node.end_lineno)
        ],
    )


def get_function(
    source: str,
    name: str,
) -> str:
    start, end = function_span(
        source,
        name,
    )

    return source[
        start:end
    ]


def replace_function(
    source: str,
    name: str,
    replacement: str,
) -> str:
    start, end = function_span(
        source,
        name,
    )

    result = (
        source[:start]
        + replacement.rstrip()
        + "\n\n"
        + source[end:].lstrip("\n")
    )

    ast.parse(
        result
    )

    return result


def db_state() -> dict:
    if not DB.exists():
        return {
            "exists": False,
        }

    conn = sqlite3.connect(
        str(DB)
    )

    try:
        integrity = str(
            conn.execute(
                "PRAGMA integrity_check"
            ).fetchone()[0]
        )

        fk_count = len(
            conn.execute(
                "PRAGMA foreign_key_check"
            ).fetchall()
        )

        columns = [
            str(row[1])
            for row in conn.execute(
                "PRAGMA table_info("
                "survey_person_year_records"
                ")"
            ).fetchall()
        ]

        count = int(
            conn.execute(
                "SELECT COUNT(*) "
                "FROM survey_person_year_records"
            ).fetchone()[0]
        )

        # Thống kê chỉ đọc để biết mức độ dữ liệu XMC hiện có.
        stats = {}

        for key, sql in (
            (
                "is_literacy_target_true",
                """
                SELECT COUNT(*)
                FROM survey_person_year_records
                WHERE is_literacy_target = 1
                """,
            ),
            (
                "grade3_true",
                """
                SELECT COUNT(*)
                FROM survey_person_year_records
                WHERE completed_grade_3 = 1
                """,
            ),
            (
                "grade3_false",
                """
                SELECT COUNT(*)
                FROM survey_person_year_records
                WHERE completed_grade_3 = 0
                """,
            ),
            (
                "grade5_true",
                """
                SELECT COUNT(*)
                FROM survey_person_year_records
                WHERE completed_grade_5 = 1
                """,
            ),
            (
                "grade5_false",
                """
                SELECT COUNT(*)
                FROM survey_person_year_records
                WHERE completed_grade_5 = 0
                """,
            ),
        ):
            stats[key] = int(
                conn.execute(
                    sql
                ).fetchone()[0]
            )

        return {
            "exists": True,
            "integrity": integrity,
            "fk_count": fk_count,
            "columns": columns,
            "count": count,
            "stats": stats,
        }

    finally:
        conn.close()


def backup_db() -> None:
    if not DB.exists():
        return

    src = sqlite3.connect(
        str(DB)
    )

    dst = sqlite3.connect(
        str(
            BACKUP
            / "phocap.db"
        )
    )

    try:
        src.backup(
            dst
        )
    finally:
        dst.close()
        src.close()


def restore_sources() -> None:
    builder_backup = (
        BACKUP
        / BUILDERS.name
    )

    if builder_backup.exists():
        shutil.copy2(
            builder_backup,
            BUILDERS,
        )

    service_backup = (
        BACKUP
        / SERVICE.name
    )

    if service_backup.exists():
        shutil.copy2(
            service_backup,
            SERVICE,
        )
    elif SERVICE.exists():
        SERVICE.unlink()


def clear_cache() -> None:
    for cache in APP.rglob(
        "__pycache__"
    ):
        if cache.is_dir():
            shutil.rmtree(
                cache,
                ignore_errors=True,
            )


def patch_export(
    source: str,
) -> tuple[str, list[str]]:
    notes: list[str] = []

    required_functions = (
        "_build_xmc3",
        "_build_cmc2",
        "_build_cmc1",
        "_build_xmc4",
        "_write_scope_signature",
        "export_additional_report",
    )

    for name in required_functions:
        function_node(
            source,
            name,
        )

    scope_fn = get_function(
        source,
        "_write_scope_signature",
    )

    for report_code in (
        "PCGD_XMC_3_2025",
        "PCGD_CMC_2_2025",
        "PCGD_CMC_1_2025",
        "PCGD_XMC_4_2025",
    ):
        if report_code not in scope_fn:
            raise RuntimeError(
                "Thiếu mapping phạm vi/tên xã cho "
                + report_code
            )

    fn = get_function(
        source,
        "export_additional_report",
    )

    # Xác nhận 4 builder XMC đều được dispatch trước khi cài.
    dispatch_tokens = (
        '_build_xmc3(ws,people,year)',
        '_build_cmc2(ws,people,year)',
        '_build_cmc1(ws,people,year,meta)',
        '_build_xmc4(ws,people,year,meta)',
    )

    for token in dispatch_tokens:
        if token not in fn:
            raise RuntimeError(
                "Không tìm thấy dispatch XMC hiện tại: "
                + token
            )

    if CALL_MARKER in fn:
        notes.append(
            "Lời gọi bộ tính XMC 4.7 đã tồn tại."
        )

        return (
            source,
            notes,
        )

    # Source hiện có thể ở dạng một dòng:
    # output=BytesIO(); workbook.save(output); output.seek(0)
    # hoặc đã được formatter xuống nhiều dòng.
    save_token = "workbook.save(output)"

    save_pos = fn.find(
        save_token
    )

    if save_pos < 0:
        raise RuntimeError(
            "Không tìm thấy workbook.save(output) "
            "trong export_additional_report."
        )

    line_start = (
        fn.rfind(
            "\n",
            0,
            save_pos,
        )
        + 1
    )

    indent = fn[
        line_start:save_pos
    ]

    # Nếu save nằm sau dấu ; trên cùng dòng,
    # indent lấy cả text trước save. Khi đó chèn vào trước
    # toàn bộ dòng output=BytesIO(); ... để an toàn.
    if ";" in indent:
        actual_indent = indent[
            : len(indent)
            - len(
                indent.lstrip()
            )
        ]
    else:
        actual_indent = indent

    call = (
        actual_indent
        + CALL_MARKER
        + "\n"
        + actual_indent
        + "if report_type in {\n"
        + actual_indent
        + '    "PCGD_CMC_1_2025",\n'
        + actual_indent
        + '    "PCGD_CMC_2_2025",\n'
        + actual_indent
        + '    "PCGD_XMC_3_2025",\n'
        + actual_indent
        + '    "PCGD_XMC_4_2025",\n'
        + actual_indent
        + "}:\n"
        + actual_indent
        + "    from app.services.xmc_report_v247 import (\n"
        + actual_indent
        + "        apply_xmc_formula_contract,\n"
        + actual_indent
        + "    )\n"
        + actual_indent
        + "    apply_xmc_formula_contract(\n"
        + actual_indent
        + "        report_type,\n"
        + actual_indent
        + "        ws,\n"
        + actual_indent
        + "    )\n"
    )

    # Chèn trước dòng chứa workbook.save.
    fn = (
        fn[:line_start]
        + call
        + fn[line_start:]
    )

    pos_call = fn.find(
        "apply_xmc_formula_contract("
    )

    pos_save = fn.find(
        "workbook.save(output)"
    )

    if not (
        0 <= pos_call < pos_save
    ):
        raise RuntimeError(
            "Sai thứ tự: bộ tính XMC phải chạy "
            "trước workbook.save(output)."
        )

    source = replace_function(
        source,
        "export_additional_report",
        fn,
    )

    if MARKER not in source:
        start, _end = function_span(
            source,
            "_build_xmc3",
        )

        source = (
            source[:start]
            + MARKER
            + "\n"
            + source[start:]
        )

    notes.append(
        "Đã nối bộ tính 79 tỷ lệ XMC "
        "trước workbook.save(output)."
    )

    notes.append(
        "Tên xã/phường tiếp tục lấy từ "
        "_write_scope_signature(meta['title'])."
    )

    ast.parse(
        source
    )

    return (
        source,
        notes,
    )


def verify_service_text(
    text: str,
) -> None:
    ast.parse(
        text
    )

    required = (
        "CMC1_PERCENT_FORMULAS",
        "XMC3_PERCENT_FORMULAS",
        "XMC4_PERCENT_FORMULAS",
        "def apply_xmc_formula_contract(",
        'cell.number_format = r\'0\\%\'',
        '("L8", "K8", "G8")',
        '("BN8", "BM8", "AU8")',
        'f"S{row}"',
        '("E8", "D8", "C8")',
        '("Q8", "P8", "M8")',
    )

    for token in required:
        if token not in text:
            raise RuntimeError(
                "Verifier service thiếu: "
                + token
            )


def verify_builder_text(
    text: str,
) -> None:
    ast.parse(
        text
    )

    fn = get_function(
        text,
        "export_additional_report",
    )

    required = (
        CALL_MARKER,
        "apply_xmc_formula_contract(",
        "workbook.save(output)",
        "PCGD_CMC_1_2025",
        "PCGD_CMC_2_2025",
        "PCGD_XMC_3_2025",
        "PCGD_XMC_4_2025",
    )

    for token in required:
        if token not in fn:
            raise RuntimeError(
                "Verifier export thiếu: "
                + token
            )

    if (
        fn.find(
            "apply_xmc_formula_contract("
        )
        >= fn.find(
            "workbook.save(output)"
        )
    ):
        raise RuntimeError(
            "Bộ tính XMC đang nằm sau save."
        )


def runtime_self_test() -> None:
    """
    Test độc lập service vừa ghi:
    - đủ 79 công thức;
    - 100/100 -> 100;
    - 0/100 -> 0;
    - denominator 0 -> blank;
    - format 0\\%.
    """
    from openpyxl import Workbook

    spec = importlib.util.spec_from_file_location(
        "xmc_report_v247_test",
        SERVICE,
    )

    if (
        spec is None
        or spec.loader is None
    ):
        raise RuntimeError(
            "Không nạp được service XMC để self-test."
        )

    module = importlib.util.module_from_spec(
        spec
    )

    spec.loader.exec_module(
        module
    )

    counts = module.formula_count()

    if counts != {
        "CMC_1": 24,
        "XMC_3": 49,
        "XMC_4": 6,
        "TOTAL": 79,
    }:
        raise RuntimeError(
            "Số công thức XMC không đúng: "
            + repr(counts)
        )

    wb = Workbook()
    ws = wb.active

    ws["K8"] = 100
    ws["G8"] = 100

    module.apply_xmc_formula_contract(
        "PCGD_CMC_1_2025",
        ws,
    )

    if ws["L8"].value != 100:
        raise RuntimeError(
            "Self-test 100% thất bại."
        )

    if ws["L8"].number_format != r"0\%":
        raise RuntimeError(
            "Self-test format 100% thất bại."
        )

    ws["D8"] = 0
    ws["C8"] = 100

    module.apply_xmc_formula_contract(
        "PCGD_XMC_4_2025",
        ws,
    )

    if ws["E8"].value != 0:
        raise RuntimeError(
            "Self-test 0% thất bại."
        )

    ws["O9"] = 1
    ws["C9"] = 0

    module.apply_xmc_formula_contract(
        "PCGD_XMC_3_2025",
        ws,
    )

    if ws["S9"].value is not None:
        raise RuntimeError(
            "Self-test mẫu số 0 phải để trống."
        )


def main() -> None:
    print("=" * 146)
    print(
        "BÀI 13B-11.15.2.4.7 - "
        "HOÀN THIỆN CÔNG THỨC XÓA MÙ CHỮ"
    )
    print("=" * 146)
    print()

    print("SẼ LÀM:")
    print(
        " - CMC-1: nối đủ 24 tỷ lệ chính thức."
    )
    print(
        " - XMC-3: nối đủ 49 tỷ lệ S9:S57 = O/C."
    )
    print(
        " - XMC-4: nối đủ 6 tỷ lệ chính thức."
    )
    print(
        " - Tổng cộng: 79/79 công thức tỷ lệ XMC."
    )
    print(
        " - CMC-2: giữ nguyên vì ma trận chính thức "
        "không có phép chia/tỷ lệ thuộc nhóm 79 này."
    )
    print(
        " - Hiển thị tỷ lệ dạng 100%, 50%, 0%."
    )
    print(
        " - Tên xã/phường tiếp tục lấy theo phạm vi "
        "người dùng đã chọn."
    )
    print()

    print("KHÔNG LÀM:")
    print(
        " - Không tự suy diễn số liệu nguồn còn trống."
    )
    print(
        " - Không tự gán completed_grade_3/5 "
        "vào cột chưa xác minh ngữ nghĩa."
    )
    print(
        " - Không tự kết luận Đạt/Không đạt theo NĐ 20 "
        "ở bài này."
    )
    print(
        " - Không ALTER/UPDATE database."
    )
    print(
        " - Không đổi menu/route/template."
    )
    print()

    for path in (
        BUILDERS,
    ):
        if not path.exists():
            raise SystemExit(
                f"Không tìm thấy: {path}"
            )

    before = db_state()

    if before.get(
        "exists"
    ):
        if (
            before.get(
                "integrity"
            )
            != "ok"
        ):
            raise SystemExit(
                "Database integrity_check != ok."
            )

        if (
            before.get(
                "fk_count"
            )
            != 0
        ):
            raise SystemExit(
                "Database có lỗi foreign key."
            )

        required_columns = {
            "is_literacy_target",
            "literacy_status",
            "completed_grade_3",
            "completed_grade_5",
        }

        missing = (
            required_columns
            - set(
                before[
                    "columns"
                ]
            )
        )

        if missing:
            raise SystemExit(
                "DB thiếu field XMC bắt buộc: "
                + ", ".join(
                    sorted(
                        missing
                    )
                )
            )

    old_builder = read_text(
        BUILDERS
    )

    ast.parse(
        old_builder
    )

    ast.parse(
        SERVICE_CONTENT
    )

    new_builder, notes = patch_export(
        old_builder
    )

    verify_builder_text(
        new_builder
    )

    verify_service_text(
        SERVICE_CONTENT
    )

    # Preflight hoàn tất mới backup/ghi.
    BACKUP.mkdir(
        parents=True,
        exist_ok=False,
    )

    shutil.copy2(
        BUILDERS,
        BACKUP / BUILDERS.name,
    )

    if SERVICE.exists():
        shutil.copy2(
            SERVICE,
            BACKUP / SERVICE.name,
        )

    backup_db()

    try:
        write_text(
            SERVICE,
            SERVICE_CONTENT,
        )

        if (
            new_builder
            != old_builder
        ):
            write_text(
                BUILDERS,
                new_builder,
            )

        for path in (
            BUILDERS,
            SERVICE,
        ):
            py_compile.compile(
                str(path),
                doraise=True,
            )

        verify_builder_text(
            read_text(
                BUILDERS
            )
        )

        verify_service_text(
            read_text(
                SERVICE
            )
        )

        runtime_self_test()

        after = db_state()

        if before != after:
            raise RuntimeError(
                "Database/schema/count/nguồn XMC "
                "thay đổi ngoài dự kiến."
            )

        clear_cache()

    except Exception:
        restore_sources()
        clear_cache()
        raise

    lines = []

    lines.append(
        "=" * 146
        + "\n"
    )

    lines.append(
        "BÀI 13B-11.15.2.4.7 - "
        "BÁO CÁO CÀI ĐẶT\n"
    )

    lines.append(
        "=" * 146
        + "\n\n"
    )

    lines.append(
        "CÔNG THỨC ĐÃ NỐI:\n"
    )

    lines.append(
        " - CMC-1: 24 công thức.\n"
    )

    lines.append(
        " - XMC-3: 49 công thức S9:S57 = O/C.\n"
    )

    lines.append(
        " - XMC-4: 6 công thức.\n"
    )

    lines.append(
        " - TỔNG: 79/79 công thức tỷ lệ XMC.\n\n"
    )

    lines.append(
        "HIỂN THỊ:\n"
    )

    lines.append(
        " - Tỷ lệ hiển thị dạng 100%, 50%, 0%.\n"
    )

    lines.append(
        " - Giá trị nội bộ vẫn là số thang 0..100.\n"
    )

    lines.append(
        " - Mẫu số 0/chưa có -> để trống.\n\n"
    )

    lines.append(
        "PHẠM VI:\n"
    )

    lines.append(
        " - Tên xã/phường tiếp tục do "
        "_write_scope_signature lấy từ meta['title'].\n"
    )

    lines.append(
        " - Toàn tỉnh giữ nhãn Toàn tỉnh.\n\n"
    )

    lines.append(
        "NGUỒN DỮ LIỆU XMC:\n"
    )

    lines.append(
        " - is_literacy_target: ĐÃ CÓ.\n"
    )

    lines.append(
        " - literacy_status: ĐÃ CÓ.\n"
    )

    lines.append(
        " - completed_grade_3: ĐÃ CÓ.\n"
    )

    lines.append(
        " - completed_grade_5: ĐÃ CÓ.\n"
    )

    if before.get(
        "exists"
    ):
        stats = before[
            "stats"
        ]

        lines.append(
            " - Hiện có is_literacy_target=True: "
            + str(
                stats[
                    "is_literacy_target_true"
                ]
            )
            + "\n"
        )

        lines.append(
            " - completed_grade_3=True/False: "
            + str(
                stats[
                    "grade3_true"
                ]
            )
            + "/"
            + str(
                stats[
                    "grade3_false"
                ]
            )
            + "\n"
        )

        lines.append(
            " - completed_grade_5=True/False: "
            + str(
                stats[
                    "grade5_true"
                ]
            )
            + "/"
            + str(
                stats[
                    "grade5_false"
                ]
            )
            + "\n"
        )

    lines.append(
        "\nLƯU Ý NGHIỆP VỤ:\n"
    )

    lines.append(
        " - Bài 4.7 chỉ tính từ tử/mẫu đã có trên biểu.\n"
    )

    lines.append(
        " - Không suy diễn ô nguồn chưa được xác minh.\n"
    )

    lines.append(
        " - Kết luận mức độ/Đạt-Chưa đạt theo NĐ 20 "
        "để bài sau sau khi kiểm thử số liệu nguồn.\n\n"
    )

    lines.append(
        "AN TOÀN:\n"
    )

    lines.append(
        " - Database không thay đổi.\n"
    )

    lines.append(
        " - Schema/số bản ghi không thay đổi.\n"
    )

    lines.append(
        " - Không đổi menu/route/template.\n"
    )

    lines.append(
        " - runtime self-test: 79 công thức + "
        "100% + 0% + mẫu số 0.\n"
    )

    lines.append(
        f" - Backup: {BACKUP}\n"
    )

    for note in notes:
        lines.append(
            " - "
            + note
            + "\n"
        )

    REPORT.write_text(
        "".join(
            lines
        ),
        encoding="utf-8",
    )

    print("KIỂM TRA:")
    print(" - AST: OK")
    print(" - py_compile: OK")
    print(" - 79/79 công thức: OK")
    print(" - Self-test 100% / 0% / mẫu số 0: OK")
    print(" - Database invariant: OK")
    print(" - Tên xã/phường: giữ theo scope hiện có")
    print()
    print("Backup:", BACKUP)
    print("Báo cáo:", REPORT)
    print("Service:", SERVICE)
    print()
    print("=" * 146)
    print(
        "BÀI 13B-11.15.2.4.7 THÀNH CÔNG"
    )
    print("=" * 146)


if __name__ == "__main__":
    main()
