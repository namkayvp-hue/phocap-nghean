# -*- coding: utf-8 -*-
from __future__ import annotations

import re
import sqlite3
from datetime import datetime
from pathlib import Path
from collections import Counter

ROOT = Path(r"C:\PhoCap")
DB = ROOT / "data" / "phocap.db"
EXPORTS = ROOT / "exports"
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
REPORT = EXPORTS / f"bao_cao_khao_sat_v2_4_3_2_khoa_mapping_route_7_bieu_mn_{STAMP}.txt"

FILES = [
    ROOT / "app" / "routers" / "survey_summary_report.py",
    ROOT / "app" / "routers" / "pcgdmn_template_report.py",
    ROOT / "app" / "routers" / "mn_official_reports.py",
    ROOT / "app" / "services" / "mn_report_v246.py",
    ROOT / "app" / "main.py",
]
TEMPLATE = ROOT / "app" / "report_templates" / "Bieu_mau_PCGDMN_2025.xlsx"

YEAR_CODE = "2026-2027"
COMMUNE_CODE = "17827"
SCHOOL_CODE = "40429325"

def add(lines, s=""):
    lines.append(str(s))

def extract_func(text: str, name: str):
    raw = text.splitlines()
    pat = re.compile(rf"^(\s*)(?:async\s+)?def\s+{re.escape(name)}\s*\(")
    start = None
    indent = None
    for i, line in enumerate(raw):
        m = pat.match(line)
        if m:
            start = i
            indent = len(m.group(1))
            break
    if start is None:
        return []
    end = min(len(raw), start + 350)
    for j in range(start + 1, min(len(raw), start + 350)):
        line = raw[j]
        if not line.strip():
            continue
        lead = len(line) - len(line.lstrip())
        if lead <= indent and (
            re.match(r"^(?:async\s+)?def\s+", line.lstrip())
            or re.match(r"^class\s+", line.lstrip())
            or line.lstrip().startswith("@router.")
        ):
            end = j
            break
    return [(k + 1, raw[k]) for k in range(start, end)]

def route_blocks(text: str):
    raw = text.splitlines()
    hits = []
    keywords = (
        "tre-khuyet-tat", "so-theo-doi", "xuat-bieu", "pcgdmn",
        "MN-01", "MN-02", "tai-chinh", "csvc", "khuyet", "theo-doi"
    )
    for i, line in enumerate(raw):
        lo = line.lower()
        if "@router." in line and any(k.lower() in lo for k in keywords):
            start = max(0, i)
            end = min(len(raw), i + 160)
            for j in range(i + 1, min(len(raw), i + 160)):
                if j > i + 3 and raw[j].lstrip().startswith("@router."):
                    end = j
                    break
            hits.append((i + 1, [(k + 1, raw[k]) for k in range(start, end)]))
    return hits

def rowdict(row):
    return dict(row) if row else None

def main():
    lines = []
    add(lines, "=" * 124)
    add(lines, "BÀI 13B-12 V2.4.3.2 - KHẢO SÁT CHỈ ĐỌC KHÓA MAPPING + ROUTE 7 BIỂU MẦM NON")
    add(lines, "=" * 124)
    add(lines, f"Thời gian: {datetime.now():%d/%m/%Y %H:%M:%S}")
    add(lines)
    add(lines, "KHÔNG sửa source/database/template. KHÔNG đụng Excel điều tra 2.2.3/2.2.4.")
    add(lines, "Mục tiêu: đủ căn cứ viết bộ cài V2.4.3 thật, không suy diễn số liệu.")
    add(lines)

    if not DB.is_file():
        raise FileNotFoundError(DB)
    conn = sqlite3.connect(DB.resolve().as_uri() + "?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    try:
        add(lines, "=" * 124)
        add(lines, "0. DATABASE")
        add(lines, "=" * 124)
        add(lines, f"integrity_check: {conn.execute('PRAGMA integrity_check').fetchone()[0]}")
        add(lines, f"foreign_key_check: {len(conn.execute('PRAGMA foreign_key_check').fetchall())} lỗi")
        years = [dict(r) for r in conn.execute("SELECT * FROM school_years ORDER BY id")]
        add(lines, "school_years:")
        for r in years:
            add(lines, "  " + repr(r))
        sy = rowdict(conn.execute("SELECT * FROM school_years WHERE code=?", (YEAR_CODE,)).fetchone())
        commune = rowdict(conn.execute("SELECT * FROM communes WHERE code=?", (COMMUNE_CODE,)).fetchone())
        school = rowdict(conn.execute("SELECT * FROM schools WHERE code=?", (SCHOOL_CODE,)).fetchone())
        add(lines, f"TARGET YEAR: {sy}")
        add(lines, f"TARGET COMMUNE: {commune}")
        add(lines, f"TARGET SCHOOL: {school}")

        if sy and commune and school:
            syid, cid, sid = sy["id"], commune["id"], school["id"]
            add(lines)
            add(lines, "=" * 124)
            add(lines, "I. DỮ LIỆU TRẺ MN CẦN ĐỔ VÀO BIỂU")
            add(lines, "=" * 124)
            rows = [dict(r) for r in conn.execute(
                """SELECT yr.*, p.full_name, p.date_of_birth, p.gender, p.ethnic_group,
                          p.residency_status, p.relationship_to_head,
                          h.commune_id AS household_commune_id, h.hamlet_name,
                          f.form_number, s.name AS school_name_linked
                   FROM survey_person_year_records yr
                   JOIN survey_people p ON p.id=yr.survey_person_id
                   JOIN households h ON h.id=p.household_id
                   LEFT JOIN survey_forms f ON f.id=yr.survey_form_id
                   LEFT JOIN schools s ON s.id=yr.school_id
                   WHERE yr.school_year_id=? AND h.commune_id=?
                   ORDER BY yr.id""",
                (syid, cid)
            )]
            start_year = int(YEAR_CODE.split("-")[0])
            by_age = Counter()
            for r in rows:
                dob = r["date_of_birth"]
                age = None
                if dob:
                    try:
                        age = start_year - int(str(dob)[:4])
                    except Exception:
                        pass
                by_age[age] += 1
            add(lines, f"Phân bố tuổi theo năm chuẩn {start_year}: {dict(sorted(by_age.items(), key=lambda x: (999 if x[0] is None else x[0])))}")
            mn = []
            for r in rows:
                dob = r["date_of_birth"]
                age = None
                if dob:
                    try:
                        age = start_year - int(str(dob)[:4])
                    except Exception:
                        pass
                if isinstance(age, int) and 0 <= age <= 6:
                    r["age_2026"] = age
                    mn.append(r)
            add(lines, f"0-6 tuổi trong địa bàn: {len(mn)}")
            for r in mn:
                keep = {
                    "id": r["id"], "person": r["full_name"], "dob": r["date_of_birth"],
                    "age": r["age_2026"], "school_id": r["school_id"],
                    "school_linked": r["school_name_linked"],
                    "school_reported": r["school_name_reported"],
                    "class_reported": r["class_name_reported"],
                    "learning_status": r["learning_status"],
                    "completed_preschool_5": r["completed_preschool_5"],
                    "completed_preschool_by_age": r["completed_preschool_by_age"],
                    "attends_two_sessions_per_day": r["attends_two_sessions_per_day"],
                    "prepared_vietnamese": r["prepared_vietnamese"],
                    "disability_status": r["disability_status"],
                    "disability_type": r["disability_type"],
                    "disability_can_learn": r["disability_can_learn"],
                    "disability_access_education": r["disability_access_education"],
                    "study_location_scope": r["study_location_scope"],
                    "school_commune_name_reported": r["school_commune_name_reported"],
                    "school_province_name_reported": r["school_province_name_reported"],
                    "form_number": r["form_number"],
                }
                add(lines, "  " + repr(keep))

        add(lines)
        add(lines, "=" * 124)
        add(lines, "II. TEMPLATE CHÍNH THỨC - NHÃN Ô / MERGE / PRINT AREA")
        add(lines, "=" * 124)
        if not TEMPLATE.is_file():
            add(lines, f"KHÔNG TÌM THẤY TEMPLATE: {TEMPLATE}")
        else:
            from openpyxl import load_workbook
            wb = load_workbook(TEMPLATE, data_only=False, read_only=False)
            add(lines, f"Template: {TEMPLATE}")
            add(lines, f"Sheets: {wb.sheetnames}")
            for ws in wb.worksheets:
                add(lines)
                add(lines, "-" * 124)
                add(lines, f"SHEET: {ws.title} | max_row={ws.max_row} max_col={ws.max_column}")
                add(lines, f"print_area={ws.print_area}")
                add(lines, f"orientation={ws.page_setup.orientation} fitToWidth={ws.page_setup.fitToWidth} fitToHeight={ws.page_setup.fitToHeight}")
                add(lines, "MERGES:")
                for m in list(ws.merged_cells.ranges)[:200]:
                    add(lines, f"  {m}")
                add(lines, "NONEMPTY CELLS (giới hạn khu vực có nội dung):")
                maxr = min(ws.max_row, 140 if ws.title != "Sổ theo dõi PCGDMN" else 20)
                maxc = min(ws.max_column, 40)
                for row in ws.iter_rows(min_row=1, max_row=maxr, min_col=1, max_col=maxc):
                    vals = []
                    for cell in row:
                        if cell.value not in (None, ""):
                            val = str(cell.value).replace("\r", " ").replace("\n", " | ")
                            if len(val) > 240:
                                val = val[:237] + "..."
                            vals.append(f"{cell.coordinate}={val!r}")
                    if vals:
                        add(lines, "  " + " ; ".join(vals))
    finally:
        conn.close()

    add(lines)
    add(lines, "=" * 124)
    add(lines, "III. SOURCE - HÀM SCOPE, TUỔI, ROUTE XUẤT BIỂU")
    add(lines, "=" * 124)
    wanted = {
        "survey_summary_report.py": [
            "_reference_date", "_school_batch_condition", "_form_permission_filters",
            "_load_batches", "_load_person_rows", "_finalize_people", "build_report_data",
        ],
        "pcgdmn_template_report.py": [
            "_reference_years", "_load_staff_report_data", "_write_staff_data_row",
            "_staff_ratios_by_age", "_write_mn01_te", "_write_mn02",
            "_write_mn01_gv", "_write_disability_sheet", "_write_tracking_book",
            "_blank_future_modules", "build_template_workbook",
        ],
        "mn_official_reports.py": [],
        "mn_report_v246.py": [],
        "main.py": [],
    }
    for p in FILES:
        add(lines)
        add(lines, "#" * 124)
        add(lines, f"FILE: {p}")
        if not p.is_file():
            add(lines, "KHÔNG TỒN TẠI")
            continue
        text = p.read_text(encoding="utf-8", errors="replace")
        names = wanted.get(p.name, [])
        for name in names:
            block = extract_func(text, name)
            add(lines, "-" * 110)
            add(lines, f"FUNCTION {name}: {'CÓ' if block else 'KHÔNG'}")
            for ln, content in block:
                add(lines, f"{ln:05d}: {content}")
        add(lines, "-" * 110)
        add(lines, "ROUTE BLOCKS LIÊN QUAN:")
        blocks = route_blocks(text)
        if not blocks:
            add(lines, "  Không tìm thấy route theo từ khóa.")
        for _, block in blocks:
            add(lines, "  " + "-" * 100)
            for ln, content in block:
                add(lines, f"{ln:05d}: {content}")

    add(lines)
    add(lines, "=" * 124)
    add(lines, "IV. TỪ KHÓA/ĐIỂM PATCH CẦN KHÓA")
    add(lines, "=" * 124)
    needles = [
        "year_record.school_id == school_id",
        "completed_preschool_5",
        "completed_preschool_by_age",
        "attends_two_sessions_per_day",
        "prepared_vietnamese",
        "disability_can_learn",
        "disability_access_education",
        "_blank_future_modules",
        "tre-khuyet-tat",
        "so-theo-doi",
    ]
    for p in FILES:
        if not p.is_file():
            continue
        text = p.read_text(encoding="utf-8", errors="replace")
        raw = text.splitlines()
        add(lines, f"[{p.name}]")
        for needle in needles:
            hits = [i+1 for i, line in enumerate(raw) if needle in line]
            if hits:
                add(lines, f"  {needle!r}: {hits}")

    EXPORTS.mkdir(parents=True, exist_ok=True)
    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")
    print("=" * 124)
    print("BÀI 13B-12 V2.4.3.2 - KHẢO SÁT CHỈ ĐỌC")
    print("=" * 124)
    print("ĐÃ TẠO:")
    print(REPORT)
    print()
    print("Không có source/database/template nào bị sửa.")

if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print()
        print("KHẢO SÁT GẶP LỖI - KHÔNG CÓ SOURCE/DATABASE/TEMPLATE NÀO BỊ SỬA.")
        print(type(exc).__name__ + ":", exc)
        raise
