# -*- coding: utf-8 -*-
from __future__ import annotations

import ast
import importlib.util
import py_compile
import shutil
import sqlite3
from datetime import datetime
from pathlib import Path

PROJECT = Path(r"C:\PhoCap")
APP = PROJECT / "app"
DB = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"
SERVICE = APP / "services" / "mn_report_v246.py"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
BACKUP = EXPORTS / f"backup_bai_13b_12_v2_4_2_1_{STAMP}"
DB_BACKUP = BACKUP / "phocap.db"
REPORT = EXPORTS / f"bao_cao_cai_bai_13b_12_v2_4_2_1_{STAMP}.txt"

MARKER = "BAI_13B_12_V2_4_2_1_SAFE_MERGEDCELL_SET"

NEW_SET = '# === BAI_13B_12_V2_4_2_1_SAFE_MERGEDCELL_SET ===\ndef _set(ws: Any, ref: str, value: Any) -> None:\n    """\n    Ghi giá trị an toàn vào ô thường.\n\n    Nếu ref là MergedCell không phải ô neo trên-trái của vùng merge:\n    - KHÔNG ghi;\n    - KHÔNG chuyển giá trị về ô neo;\n    vì làm vậy có thể ghi đè tiêu đề/nhãn của biểu.\n    """\n    cell = ws[ref]\n\n    if isinstance(cell, MergedCell):\n        return\n\n    cell.value = value\n\n\n'


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig", errors="replace")


def write_text(path: Path, value: str) -> None:
    path.write_text(value, encoding="utf-8")


def db_state() -> dict:
    con = sqlite3.connect(str(DB))
    try:
        result = {
            "integrity": str(
                con.execute("PRAGMA integrity_check").fetchone()[0]
            ),
            "fk_count": len(
                con.execute("PRAGMA foreign_key_check").fetchall()
            ),
        }
        for table in (
            "survey_forms",
            "survey_people",
            "survey_person_year_records",
        ):
            result[table] = int(
                con.execute(
                    f"SELECT COUNT(*) FROM {table}"
                ).fetchone()[0]
            )
        return result
    finally:
        con.close()


def backup_file(path: Path) -> None:
    target = BACKUP / path.relative_to(PROJECT)
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, target)


def restore_file(path: Path) -> None:
    source = BACKUP / path.relative_to(PROJECT)
    if source.exists():
        shutil.copy2(source, path)


def backup_db() -> None:
    src = sqlite3.connect(str(DB))
    dst = sqlite3.connect(str(DB_BACKUP))
    try:
        src.backup(dst)
        dst.commit()
    finally:
        dst.close()
        src.close()


def restore_db() -> None:
    if not DB_BACKUP.exists():
        return
    src = sqlite3.connect(str(DB_BACKUP))
    dst = sqlite3.connect(str(DB))
    try:
        src.backup(dst)
        dst.commit()
    finally:
        dst.close()
        src.close()


def function_node(tree: ast.Module, name: str):
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            if node.name == name:
                return node
    raise RuntimeError(f"Không xác định được hàm {name} bằng AST.")


def replace_function_ast(
    source: str,
    name: str,
    replacement: str,
) -> str:
    tree = ast.parse(source)
    node = function_node(tree, name)

    lines = source.splitlines(keepends=True)
    start = int(node.lineno) - 1
    end = int(getattr(node, "end_lineno", node.lineno) or node.lineno)

    text = replacement
    if not text.endswith("\n"):
        text += "\n"

    result = "".join(lines[:start] + [text] + lines[end:])
    ast.parse(result)
    return result


def ensure_mergedcell_import(source: str) -> str:
    if "from openpyxl.cell.cell import MergedCell" in source:
        return source

    lines = source.splitlines(keepends=True)

    insert_at = 0
    for index, line in enumerate(lines):
        if line.startswith("from __future__ import "):
            insert_at = index + 1
            break

    lines.insert(
        insert_at,
        "from openpyxl.cell.cell import MergedCell\n",
    )
    result = "".join(lines)
    ast.parse(result)
    return result


def patch_service(source: str) -> str:
    if MARKER in source:
        return source

    required = (
        "def _set(",
        "def _apply_old_gv(",
        "def apply_formula_contract(",
        "FORMULA_CONTRACT",
        "REFERENCE_SHA256",
    )
    for token in required:
        if token not in source:
            raise RuntimeError(
                "mn_report_v246.py không đúng nền mong đợi. Thiếu: "
                + token
            )

    tree = ast.parse(source)
    old_set = function_node(tree, "_set")
    old_set_text = ast.get_source_segment(source, old_set) or ""

    if "ws[ref]" not in old_set_text and "cell.value" not in old_set_text:
        raise RuntimeError(
            "Hàm _set hiện tại khác cấu trúc đã khảo sát; không tự động vá."
        )

    source = ensure_mergedcell_import(source)
    source = replace_function_ast(
        source,
        "_set",
        NEW_SET,
    )
    ast.parse(source)
    return source


def load_service():
    spec = importlib.util.spec_from_file_location(
        "mn_report_v246_v2421_test",
        SERVICE,
    )
    if spec is None or spec.loader is None:
        raise RuntimeError("Không nạp được mn_report_v246.py.")

    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def smoke_test() -> None:
    from openpyxl import Workbook

    module = load_service()

    wb = Workbook()
    ws = wb.active
    ws.title = "MN-01 GV"

    ws.merge_cells("A8:S8")
    ws["A8"] = "TRƯỜNG MẦM NON"

    ws["J9"] = 2
    ws["K9"] = 4
    ws["O9"] = 3
    ws["P9"] = 1
    ws["R9"] = 2

    module.apply_formula_contract(wb)

    if ws["A8"].value != "TRƯỜNG MẦM NON":
        raise RuntimeError(
            "Smoke test: tiêu đề vùng merge bị ghi đè."
        )

    if ws["N9"].value != 2:
        raise RuntimeError(
            f"Smoke test: N9 mong đợi 2, nhận {ws['N9'].value!r}."
        )

    if ws["Q9"].value != 100:
        raise RuntimeError(
            f"Smoke test: Q9 mong đợi 100, nhận {ws['Q9'].value!r}."
        )

    if ws["S9"].value != 50:
        raise RuntimeError(
            f"Smoke test: S9 mong đợi 50, nhận {ws['S9'].value!r}."
        )

    module._set(ws, "N8", 999)

    if ws["A8"].value != "TRƯỜNG MẦM NON":
        raise RuntimeError(
            "Smoke test: _set đã chuyển giá trị về ô neo của vùng merge."
        )


def verify_source(source: str) -> None:
    required = (
        MARKER,
        "from openpyxl.cell.cell import MergedCell",
        "isinstance(cell, MergedCell)",
        "cell.value = value",
        "def apply_formula_contract(",
    )
    for token in required:
        if token not in source:
            raise RuntimeError("Verifier thiếu: " + token)

    ast.parse(source)
    py_compile.compile(str(SERVICE), doraise=True)


def main() -> int:
    print("=" * 126)
    print(
        "BÀI 13B-12 V2.4.2.1 - "
        "SỬA LỖI MERGEDCELL KHI XUẤT MN-TRẺ KT / SỔ THEO DÕI"
    )
    print("=" * 126)

    for path in (DB, SERVICE):
        if not path.exists():
            print("DỪNG AN TOÀN: thiếu", path)
            return 2

    before = db_state()

    print("integrity_check:", before["integrity"])
    print("foreign_key_check:", before["fk_count"], "lỗi")

    if before["integrity"].lower() != "ok" or before["fk_count"] != 0:
        print("DỪNG: database chưa đạt kiểm tra an toàn.")
        return 3

    source = read_text(SERVICE)

    if MARKER in source:
        print("V2.4.2.1 đã cài. Không cài lặp.")
        return 0

    try:
        patched = patch_service(source)
        ast.parse(patched)
    except Exception as exc:
        print()
        print("DỪNG AN TOÀN TRƯỚC KHI GHI SOURCE:")
        print(type(exc).__name__ + ":", exc)
        return 4

    EXPORTS.mkdir(parents=True, exist_ok=True)
    BACKUP.mkdir(parents=True, exist_ok=False)

    backup_file(SERVICE)
    backup_db()

    print("Backup:", BACKUP)

    try:
        write_text(SERVICE, patched)
        verify_source(read_text(SERVICE))
        smoke_test()

        after = db_state()

        if after["integrity"].lower() != "ok":
            raise RuntimeError(
                "integrity_check không đạt sau cài."
            )
        if after["fk_count"] != 0:
            raise RuntimeError(
                "foreign_key_check có lỗi sau cài."
            )

        for table in (
            "survey_forms",
            "survey_people",
            "survey_person_year_records",
        ):
            if after[table] != before[table]:
                raise RuntimeError(
                    f"Số bản ghi {table} thay đổi ngoài dự kiến."
                )

        for cache in APP.rglob("__pycache__"):
            if cache.is_dir():
                shutil.rmtree(cache, ignore_errors=True)

    except Exception as exc:
        print()
        print("CÓ LỖI - ĐANG ROLLBACK SOURCE + DATABASE...")
        print(type(exc).__name__ + ":", exc)
        restore_file(SERVICE)
        restore_db()
        print("ĐÃ KHÔI PHỤC.")
        return 9

    REPORT.write_text(
        "\n".join(
            [
                "=" * 126,
                "BÁO CÁO CÀI BÀI 13B-12 V2.4.2.1",
                "=" * 126,
                f"Thời gian: {datetime.now():%d/%m/%Y %H:%M:%S}",
                "",
                "NGUYÊN NHÂN:",
                "- apply_formula_contract chạy cho toàn workbook trước khi tách biểu.",
                "- _apply_old_gv quét qua cả các dòng tiêu đề đã merge.",
                "- _set cũ ghi trực tiếp ws[ref] nên gặp MergedCell read-only.",
                "",
                "ĐÃ SỬA:",
                "- _set bỏ qua MergedCell không phải ô neo.",
                "- Không chuyển giá trị về ô neo để tránh ghi đè tiêu đề.",
                "- Ô dữ liệu bình thường vẫn tính/ghi như cũ.",
                "- Smoke test vùng merge + MN-01-GV: ĐẠT.",
                "",
                "KHÔNG THAY:",
                "- Excel điều tra 2.2.3/2.2.4;",
                "- menu 7 biểu V2.4.2;",
                "- dữ liệu điều tra;",
                "- PCGD/XMC;",
                "- phân công/giao phiếu/nhập nhanh.",
                "",
                f"integrity_check: {after['integrity']}",
                f"foreign_key_check: {after['fk_count']} lỗi",
                f"Backup: {BACKUP}",
            ]
        ) + "\n",
        encoding="utf-8-sig",
    )

    print()
    print("=" * 126)
    print("CÀI ĐẶT THÀNH CÔNG BÀI 13B-12 V2.4.2.1")
    print("=" * 126)
    print(" - Sửa MergedCell: ĐẠT.")
    print(" - Smoke test công thức + vùng merge: ĐẠT.")
    print(" - Không sửa dữ liệu hiện có.")
    print("Báo cáo:", REPORT)
    print("Backup:", BACKUP)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
