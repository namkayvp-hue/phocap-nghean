from __future__ import annotations

from collections import defaultdict
from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Classroom, School, SchoolYear
from app.routers.report_center import (
    _choose_year_id,
    _load_school_years,
    _load_scope_options,
    _parse_optional_int,
)
from app.routers.survey_summary_report import build_report_data
from app.routers.surveys import lay_thong_tin_nguoi_dung
from app.survey_models import SurveyPerson, SurveyPersonYearRecord
from app.staff_models import StaffMember, StaffYearRecord, SchoolStaffYearSummary
from app.services.finance_school_service import fill_finance_worksheet
from sqlalchemy import text


APP_DIR = Path(__file__).resolve().parent.parent
TEMPLATE_DIR = APP_DIR / "report_templates"
TEMPLATE_PATH = TEMPLATE_DIR / "Bieu_mau_PCGDMN_2025.xlsx"
templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

router = APIRouter(
    prefix="/bao-cao/pcgdmn-mau-2025",
    tags=["Bộ biểu mẫu PCGDMN 2025"],
)

PROVINCE_NAME = "Nghệ An"
DATA_STATE_LABELS = {
    "": "Tất cả trạng thái dữ liệu",
    "HAS_DATA": "Đã có dữ liệu",
    "NO_DATA": "Chưa có dữ liệu",
    "LOCKED": "Đã khóa số liệu",
    "OPEN": "Dữ liệu đang mở",
}


# === BAI_13B_12_V2_4_2_MN_SINGLE_SHEET_EXPORT_START ===
# === BAI_13B_12_V2_4_3_4_SINGLE_PIPELINE_7_MN ===
SINGLE_SHEET_EXPORTS = {
    "mn-01-te": {
        "sheet_name": "MN-01 TE",
        "label": "MN-01-TE – Thống kê trẻ em Mầm non theo độ tuổi",
        "filename": "MN-01-TE",
    },
    "mn-02": {
        "sheet_name": "MN-02",
        "label": "MN-02 – Kết quả PCGD Mầm non",
        "filename": "MN-02",
    },
    "mn-01-gv": {
        "sheet_name": "MN-01 GV",
        "label": "MN-01-GV – Thống kê đội ngũ CBQL, giáo viên, nhân viên",
        "filename": "MN-01-GV",
    },
    "mn-01-csvc": {
        "sheet_name": "MN-01 CSVC",
        "label": "MN-01-CSVC – Cơ sở vật chất",
        "filename": "MN-01-CSVC",
    },
    "mn-tc": {
        "sheet_name": "MN - Tài chính",
        "label": "MN-TC – Báo cáo tài chính",
        "filename": "MN-TC",
    },
    "tre-khuyet-tat": {
        "sheet_name": "MN- Trẻ KT",
        "label": "MN-Trẻ KT – Thống kê trẻ khuyết tật",
        "filename": "MN_Tre_KT",
    },
    "so-theo-doi": {
        "sheet_name": "Sổ theo dõi PCGDMN",
        "label": "Sổ theo dõi PCGDMN",
        "filename": "So_theo_doi_PCGDMN",
    },
}
# === BAI_13B_12_V2_4_2_MN_SINGLE_SHEET_EXPORT_END ===

STUDYING_CODES = {"DANG_HOC", "CHUYEN_DEN"}
TRANSFER_OUT_CODES = {"CHUYEN_DI"}
DEAD_CODES = {"DA_CHET", "CHET"}


def _safe_percent(numerator: int | float, denominator: int | float) -> float | None:
    if not denominator:
        return None
    return round(float(numerator) * 100.0 / float(denominator), 2)


def _safe_ratio(numerator: int | float, denominator: int | float) -> float | None:
    if not denominator:
        return None
    return float(numerator) / float(denominator)


def _is_female(row: dict[str, Any]) -> bool:
    return str(row.get("gender_code") or "").upper() == "NU"


def _is_minority(row: dict[str, Any]) -> bool:
    value = " ".join(str(row.get("ethnic_group") or "").strip().upper().split())
    return bool(value and value not in {"KINH", "DÂN TỘC KINH"})


def _is_disabled(row: dict[str, Any]) -> bool:
    return str(row.get("disability_status_code") or "").upper() == "CO_KHUYET_TAT"


def _is_studying(row: dict[str, Any]) -> bool:
    return str(row.get("learning_status_code") or "").upper() in STUDYING_CODES


def _split_name(value: str | None) -> tuple[str, str]:
    parts = [part for part in str(value or "").strip().split() if part]
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return " ".join(parts[:-1]), parts[-1]


def _clear_range_values(ws: Any, cell_range: str) -> None:
    for row in ws[cell_range]:
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None


def _reference_years(school_year: SchoolYear) -> tuple[int, int]:
    try:
        start_text, end_text = str(school_year.code).split("-", 1)
        return int(start_text), int(end_text)
    except (TypeError, ValueError):
        current = datetime.now().year
        return current, current + 1


def _load_context(
    *,
    db: Session,
    request: Request,
    school_year_id: int | None,
    commune_id: int | None,
    school_id: int | None,
    data_state: str,
) -> dict[str, Any]:
    user = lay_thong_tin_nguoi_dung(request)
    school_years = _load_school_years(db)
    selected_year_id = _choose_year_id(school_years, school_year_id)
    (
        communes,
        schools,
        selected_commune_id,
        selected_school_id,
        scope_label,
    ) = _load_scope_options(
        db,
        user,
        commune_id,
        school_id,
    )

    report_data = build_report_data(
        db=db,
        request=request,
        school_year_id=selected_year_id,
        commune_id=selected_commune_id,
        school_id=selected_school_id,
        data_state=data_state,
        q="",
    )

    selected_school = next(
        (item for item in schools if item.id == selected_school_id),
        None,
    )
    selected_commune = next(
        (item for item in communes if item.id == selected_commune_id),
        None,
    )
    if selected_school is not None and selected_commune is None:
        selected_commune = db.get(School, selected_school.id).commune

    person_rows = list(report_data.get("person_rows") or [])
    age_3_5_rows = [
        row for row in person_rows if row.get("age") in {3, 4, 5}
    ]
    stats = {
        "children_3_5": len(age_3_5_rows),
        "age_3": sum(row.get("age") == 3 for row in age_3_5_rows),
        "age_4": sum(row.get("age") == 4 for row in age_3_5_rows),
        "age_5": sum(row.get("age") == 5 for row in age_3_5_rows),
        "studying": sum(_is_studying(row) for row in age_3_5_rows),
        "disabled": sum(_is_disabled(row) for row in age_3_5_rows),
        "missing_personal_id": sum(not row.get("personal_id") for row in age_3_5_rows),
        "completed_age_5": sum(
            row.get("age") == 5 and row.get("completed_preschool_5") is True
            for row in age_3_5_rows
        ),
    }

    return {
        "user": user,
        "school_years": school_years,
        "selected_year_id": selected_year_id,
        "communes": communes,
        "schools": schools,
        "selected_commune_id": selected_commune_id,
        "selected_school_id": selected_school_id,
        "selected_commune": selected_commune,
        "selected_school": selected_school,
        "scope_label": scope_label,
        "data_state": data_state,
        "report_data": report_data,
        "stats": stats,
    }


def _load_extra_models(
    db: Session,
    person_rows: list[dict[str, Any]],
) -> tuple[dict[int, SurveyPerson], dict[int, SurveyPersonYearRecord]]:
    person_ids = sorted({int(row["person_id"]) for row in person_rows if row.get("person_id")})
    record_ids = sorted({int(row["year_record_id"]) for row in person_rows if row.get("year_record_id")})

    people: dict[int, SurveyPerson] = {}
    records: dict[int, SurveyPersonYearRecord] = {}
    if person_ids:
        people = {
            item.id: item
            for item in db.scalars(
                select(SurveyPerson).where(SurveyPerson.id.in_(person_ids))
            ).all()
        }
    if record_ids:
        records = {
            item.id: item
            for item in db.scalars(
                select(SurveyPersonYearRecord).where(
                    SurveyPersonYearRecord.id.in_(record_ids)
                )
            ).all()
        }
    return people, records


def _write_header(ws: Any, commune_label: str, year_code: str) -> None:
    sheet_title = ws.title.strip()
    if sheet_title in {"MN-01 TE", "MN-01 GV", "MN-01 CSVC"}:
        ws["A1"] = f"Xã: {commune_label}"
        ws["A2"] = f"Tỉnh: {PROVINCE_NAME}"
        if sheet_title == "MN-01 TE":
            ws["C2"] = f"Thời điểm: năm học {year_code}"
        else:
            ws["C2"] = f"Thời điểm: năm học {year_code}"
    elif sheet_title == "MN-02":
        ws["A1"] = f"Xã: {commune_label}"
        ws["A2"] = f"Tỉnh: {PROVINCE_NAME}"
        ws["D2"] = f"Thời điểm: năm học {year_code}"
    elif sheet_title == "MN - Tài chính":
        ws["A1"] = f"Tỉnh {PROVINCE_NAME}"
        ws["A2"] = f"Xã/phường: {commune_label}"
        ws["A4"] = f"Năm học: {year_code}"
    elif sheet_title == "MN- Trẻ KT":
        ws["A1"] = f"Tỉnh/TP: {PROVINCE_NAME}"
        ws["A2"] = f"Xã/phường: {commune_label}"
        ws["A5"] = f"Năm học: {year_code}"
    elif sheet_title == "Sổ theo dõi PCGDMN":
        ws["A1"] = f"Tỉnh/TP: {PROVINCE_NAME}"
        ws["A2"] = f"Xã/phường: {commune_label}"
        ws["A3"] = f"Năm học: {year_code}"



def _normalize_text(value: Any) -> str:
    import unicodedata
    text = " ".join(str(value or "").strip().lower().split())
    text = unicodedata.normalize("NFD", text)
    return "".join(ch for ch in text if unicodedata.category(ch) != "Mn")


def _infer_institution_group(name: str) -> str:
    normalized = _normalize_text(name)
    tokens = ("nhom tre", "lop mam non", "co so mn", "csmn", "doc lap", "tgd ", "nhom lop")
    return "CO_SO_GDMN_DOC_LAP" if any(token in normalized for token in tokens) else "TRUONG_MAM_NON"


def _infer_class_counts(classes: list[Classroom]) -> tuple[int, int, int]:
    total = len([item for item in classes if item.is_active])
    age_34 = 0
    age_5 = 0
    for item in classes:
        if not item.is_active:
            continue
        text = _normalize_text(f"{item.code or ''} {item.name}")
        if "5 tuoi" in text or "lop lon" in text or "lon " in text or text.startswith("5"):
            age_5 += 1
        elif any(token in text for token in ("3 tuoi", "4 tuoi", "be ", "nho ", "mau giao")):
            age_34 += 1
    return total, age_34, age_5


def _load_staff_report_data(db: Session, context: dict[str, Any]) -> dict[str, Any]:
    year_id = int(context["selected_year_id"])
    school_statement = select(School).where(School.is_active.is_(True)).order_by(School.name)
    if context.get("selected_school_id") is not None:
        school_statement = school_statement.where(School.id == int(context["selected_school_id"]))
    elif context.get("selected_commune_id") is not None:
        school_statement = school_statement.where(School.commune_id == int(context["selected_commune_id"]))
    schools = list(db.scalars(school_statement).all())
    school_ids = [item.id for item in schools]
    records = []
    summaries = []
    classes = []
    if school_ids:
        records = list(db.scalars(
            select(StaffYearRecord)
            .join(StaffYearRecord.staff_member)
            .where(
                StaffYearRecord.school_year_id == year_id,
                StaffYearRecord.school_id.in_(school_ids),
                StaffYearRecord.is_active.is_(True),
                StaffYearRecord.status_code == "DANG_LAM_VIEC",
                StaffMember.is_active.is_(True),
            )
        ).all())
        summaries = list(db.scalars(
            select(SchoolStaffYearSummary).where(
                SchoolStaffYearSummary.school_year_id == year_id,
                SchoolStaffYearSummary.school_id.in_(school_ids),
            )
        ).all())
        classes = list(db.scalars(
            select(Classroom).where(
                Classroom.school_year_id == year_id,
                Classroom.school_id.in_(school_ids),
                Classroom.is_active.is_(True),
            )
        ).all())
    records_by_school: dict[int, list[StaffYearRecord]] = defaultdict(list)
    for record in records:
        records_by_school[record.school_id].append(record)
    summary_by_school = {item.school_id: item for item in summaries}
    classes_by_school: dict[int, list[Classroom]] = defaultdict(list)
    for classroom in classes:
        classes_by_school[classroom.school_id].append(classroom)

    rows: list[dict[str, Any]] = []
    for school in schools:
        summary = summary_by_school.get(school.id)
        if summary is not None:
            total_classes = int(summary.total_groups_classes or 0)
            classes_34 = int(summary.preschool_classes_3_4 or 0)
            classes_5 = int(summary.preschool_classes_5 or 0)
            group = summary.institution_group
        else:
            total_classes, classes_34, classes_5 = _infer_class_counts(classes_by_school.get(school.id, []))
            group = _infer_institution_group(school.name)
        school_records = records_by_school.get(school.id, [])
        rows.append({
            "school": school,
            "records": school_records,
            "institution_group": group,
            "total_classes": total_classes,
            "classes_34": classes_34,
            "classes_5": classes_5,
        })
    active_records = [record for row in rows for record in row["records"]]
    return {
        "rows": rows,
        "records": active_records,
        "summary": {
            "staff_total": len(active_records),
            "managers": sum(r.position_group == "CBQL" for r in active_records),
            "teachers": sum(r.position_group == "GIAO_VIEN" for r in active_records),
            "employees": sum(r.position_group == "NHAN_VIEN" for r in active_records),
            "configured_schools": len(summaries),
        },
    }


def _copy_row_style(ws: Any, source_row: int, target_row: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for column in range(1, 20):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)
        if source.has_style:
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.number_format = source.number_format
            target.protection = copy(source.protection)


def _ratio(numerator: int, denominator: int) -> float | None:
    if not denominator:
        return None
    return round(float(numerator) / float(denominator), 2)


def _percent_value(numerator: int, denominator: int) -> float | None:
    if not denominator:
        return None
    return round(float(numerator) * 100.0 / float(denominator), 2)


def _staff_counts(records: list[StaffYearRecord], age_group: str | None = None) -> dict[str, int]:
    if age_group is None:
        selected = records
    else:
        selected = [
            item for item in records
            if item.position_group == "GIAO_VIEN"
            and item.teaching_level == "MAU_GIAO"
            and item.teaching_age_group == age_group
        ]
    return {
        "total": len(selected),
        "contract": sum(item.employment_type in {"HOP_DONG_LAM_VIEC", "HOP_DONG_LAO_DONG"} for item in selected),
        "manager": sum(item.position_group == "CBQL" for item in selected),
        "teacher": sum(item.position_group == "GIAO_VIEN" for item in selected),
        "employee": sum(item.position_group == "NHAN_VIEN" for item in selected),
        "policy": sum(bool(item.receives_policy) for item in selected),
        "standard": sum(item.qualification_standard == "DAT_CHUAN" for item in selected),
        "above": sum(item.qualification_standard == "TREN_CHUAN" for item in selected),
        "professional": sum(item.professional_standard in {"DAT", "KHA", "TOT"} for item in selected),
    }


def _write_staff_data_row(ws: Any, row_number: int, school_row: dict[str, Any], index: int) -> None:
    records = school_row["records"]
    all_counts = _staff_counts(records)
    ws.cell(row_number, 1).value = index
    ws.cell(row_number, 2).value = school_row["school"].name
    ws.cell(row_number, 3).value = all_counts["total"]
    ws.cell(row_number, 4).value = all_counts["contract"]
    ws.cell(row_number, 5).value = all_counts["manager"]
    ws.cell(row_number, 6).value = all_counts["teacher"]
    ws.cell(row_number, 7).value = _ratio(all_counts["teacher"], school_row["total_classes"])
    ws.cell(row_number, 8).value = all_counts["employee"]

    for offset, (label, age_code, class_key) in enumerate((
        ("3,4 tuổi", "TUOI_3_4", "classes_34"),
        ("5 tuổi", "TUOI_5", "classes_5"),
    )):
        target_row = row_number + offset
        if offset:
            _copy_row_style(ws, 11, target_row)
        counts = _staff_counts(records, age_code)
        class_count = int(school_row[class_key] or 0)
        ws.cell(target_row, 9).value = label
        ws.cell(target_row, 10).value = class_count
        ws.cell(target_row, 11).value = counts["total"]
        ws.cell(target_row, 12).value = counts["contract"]
        ws.cell(target_row, 13).value = counts["policy"]
        ws.cell(target_row, 14).value = _ratio(counts["total"], class_count)
        ws.cell(target_row, 15).value = counts["standard"]
        ws.cell(target_row, 16).value = counts["above"]
        ws.cell(target_row, 17).value = _percent_value(counts["standard"] + counts["above"], counts["total"])
        ws.cell(target_row, 18).value = counts["professional"]
        ws.cell(target_row, 19).value = _percent_value(counts["professional"], counts["total"])


def _aggregate_staff_rows(rows: list[dict[str, Any]]) -> dict[str, Any]:
    records = [record for row in rows for record in row["records"]]
    return {
        "school": type("SummarySchool", (), {"name": ""})(),
        "records": records,
        "total_classes": sum(int(row["total_classes"] or 0) for row in rows),
        "classes_34": sum(int(row["classes_34"] or 0) for row in rows),
        "classes_5": sum(int(row["classes_5"] or 0) for row in rows),
    }


def _write_staff_summary_row(ws: Any, row_number: int, label: str, rows: list[dict[str, Any]]) -> None:
    aggregate = _aggregate_staff_rows(rows)
    records = aggregate["records"]
    all_counts = _staff_counts(records)
    if isinstance(ws.cell(row_number, 2), MergedCell):
        ws.cell(row_number, 1).value = label
    else:
        ws.cell(row_number, 2).value = label
    ws.cell(row_number, 3).value = all_counts["total"]
    ws.cell(row_number, 4).value = all_counts["contract"]
    ws.cell(row_number, 5).value = all_counts["manager"]
    ws.cell(row_number, 6).value = all_counts["teacher"]
    ws.cell(row_number, 7).value = _ratio(all_counts["teacher"], aggregate["total_classes"])
    ws.cell(row_number, 8).value = all_counts["employee"]
    ws.cell(row_number, 9).value = None
    ws.cell(row_number, 10).value = aggregate["classes_34"] + aggregate["classes_5"]
    preschool = [item for item in records if item.position_group == "GIAO_VIEN" and item.teaching_level == "MAU_GIAO" and item.teaching_age_group in {"TUOI_3_4", "TUOI_5"}]
    ws.cell(row_number, 11).value = len(preschool)
    ws.cell(row_number, 12).value = sum(item.employment_type in {"HOP_DONG_LAM_VIEC", "HOP_DONG_LAO_DONG"} for item in preschool)
    ws.cell(row_number, 13).value = sum(bool(item.receives_policy) for item in preschool)
    ws.cell(row_number, 14).value = _ratio(len(preschool), aggregate["classes_34"] + aggregate["classes_5"])
    standard = sum(item.qualification_standard == "DAT_CHUAN" for item in preschool)
    above = sum(item.qualification_standard == "TREN_CHUAN" for item in preschool)
    professional = sum(item.professional_standard in {"DAT", "KHA", "TOT"} for item in preschool)
    ws.cell(row_number, 15).value = standard
    ws.cell(row_number, 16).value = above
    ws.cell(row_number, 17).value = _percent_value(standard + above, len(preschool))
    ws.cell(row_number, 18).value = professional
    ws.cell(row_number, 19).value = _percent_value(professional, len(preschool))


def _write_mn01_gv(ws: Any, staff_data: dict[str, Any]) -> None:
    if "A52:B52" in {str(item) for item in ws.merged_cells.ranges}:
        ws.unmerge_cells("A52:B52")
    _clear_range_values(ws, "A8:S68")
    all_rows = staff_data["rows"]
    public_rows = [row for row in all_rows if row["institution_group"] == "TRUONG_MAM_NON"]
    independent_rows = [row for row in all_rows if row["institution_group"] == "CO_SO_GDMN_DOC_LAP"]
    _write_staff_summary_row(ws, 8, "TOÀN XÃ/PHƯỜNG", all_rows)
    current_row = 9
    index = 1
    for label, group_rows, style_source in (
        ("TRƯỜNG MẦM NON", public_rows, 9),
        ("CƠ SỞ GDMN ĐỘC LẬP", independent_rows, 52),
    ):
        if not group_rows:
            continue
        if current_row > 68:
            break
        _copy_row_style(ws, style_source, current_row)
        _write_staff_summary_row(ws, current_row, label, group_rows)
        current_row += 1
        for school_row in group_rows:
            if current_row + 1 > 68:
                ws["B68"] = "CÒN ĐƠN VỊ CHƯA HIỂN THỊ – HÃY CHỌN XÃ/PHƯỜNG HẸP HƠN"
                ws["B68"].font = Font(color="C00000", bold=True)
                return
            _copy_row_style(ws, 10, current_row)
            _write_staff_data_row(ws, current_row, school_row, index)
            index += 1
            current_row += 2
    if not any(row["records"] for row in all_rows):
        ws["B10"] = "CHƯA NHẬP DỮ LIỆU ĐỘI NGŨ CHO NĂM HỌC VÀ PHẠM VI ĐÃ CHỌN"
        ws["B10"].font = Font(color="C00000", bold=True)
    for row_number in range(8, 69):
        ws.cell(row_number, 7).number_format = "0.00"
        ws.cell(row_number, 14).number_format = "0.00"
        ws.cell(row_number, 17).number_format = "0.00"
        ws.cell(row_number, 19).number_format = "0.00"


def _staff_ratios_by_age(staff_data: dict[str, Any]) -> dict[str, float | None]:
    rows = staff_data["rows"]
    records = [record for row in rows for record in row["records"]]
    result = {}
    for key, code, class_key in (("age_34", "TUOI_3_4", "classes_34"), ("age_5", "TUOI_5", "classes_5")):
        teachers = sum(
            item.position_group == "GIAO_VIEN"
            and item.teaching_level == "MAU_GIAO"
            and item.teaching_age_group == code
            for item in records
        )
        classes = sum(int(row[class_key] or 0) for row in rows)
        result[key] = _ratio(teachers, classes)
    return result



# === BAI_13B_12_V2_4_3_TEMPLATE_HELPERS_START ===
def _v243_true(value: Any) -> bool:
    if value is True:
        return True
    if isinstance(value, (int, float)):
        return value == 1
    return str(value or "").strip().upper() in {
        "1",
        "TRUE",
        "YES",
        "CO",
        "CÓ",
    }


def _v243_false(value: Any) -> bool:
    if value is False:
        return True
    if isinstance(value, (int, float)):
        return value == 0
    return str(value or "").strip().upper() in {
        "0",
        "FALSE",
        "NO",
        "KHONG",
        "KHÔNG",
    }


def _v243_completed(row: dict[str, Any]) -> bool:
    canonical = row.get("completed_preschool_by_age")
    if canonical is not None:
        return _v243_true(canonical)
    return row.get("completed_preschool_5") is True


def _v243_can_learn(row: dict[str, Any]) -> bool:
    return _is_disabled(row) and _v243_true(
        row.get("disability_can_learn")
    )


def _v243_access(row: dict[str, Any]) -> bool:
    return _is_disabled(row) and _v243_true(
        row.get("disability_access_education")
    )


def _v243_mobilizable(row: dict[str, Any]) -> bool:
    if not _is_disabled(row):
        return True
    # Chỉ loại khỏi "phải huy động" khi đã xác nhận rõ là không có khả năng học.
    return not _v243_false(row.get("disability_can_learn"))


def _v243_two_sessions(row: dict[str, Any]) -> bool:
    return _v243_true(row.get("attends_two_sessions_per_day"))


def _v243_local_study(row: dict[str, Any]) -> bool:
    # BATCH22_MN01_LOCATION_SCOPE
    # Nơi học khác được phân loại độc lập với biến động cư trú.
    code = str(row.get("study_location_scope") or "").strip().upper()

    if code == "DI_HOC_TRONG_TINH":
        return True
    if code == "DI_HOC_NGOAI_TINH":
        return False

    if code in {"TAI_CHO", "TRONG_XA", "CUNG_XA"}:
        return True
    if code in {
        "KHAC_XA",
        "TRONG_TINH_KHAC_XA",
        "NGOAI_DIA_BAN",
        "TRAI_TUYEN",
        "DI_HOC_NOI_KHAC",
    }:
        return False

    school_commune_id = row.get("school_commune_id")
    return (
        not school_commune_id
        or school_commune_id == row.get("commune_id")
    )


def _v243_cross_study(row: dict[str, Any]) -> bool:
    return _is_studying(row) and not _v243_local_study(row)


def _v243_normalize_year_labels(
    workbook: Any,
    school_year: SchoolYear,
) -> None:
    start_year, _end_year = _reference_years(school_year)

    ws = _get_sheet(workbook, "MN-01 TE")
    ws["C2"] = f"Thời điểm: Năm {start_year}"

    ws = _get_sheet(workbook, "MN-01 GV")
    ws["C2"] = f"Thời điểm: Năm {start_year}"

    ws = _get_sheet(workbook, "MN-01 CSVC")
    ws["C2"] = f"Thời điểm: Năm {start_year}"

    ws = _get_sheet(workbook, "MN - Tài chính")
    ws["A4"] = f"Năm: {start_year}"
    for offset, column in enumerate(range(5, 10)):
        ws.cell(8, column).value = start_year + offset

    ws = _get_sheet(workbook, "MN- Trẻ KT")
    ws["A5"] = f"Năm: {school_year.code}"

    ws = _get_sheet(workbook, "Sổ theo dõi PCGDMN")
    ws["A3"] = f"Năm học: {school_year.code}"
# === BAI_13B_12_V2_4_3_TEMPLATE_HELPERS_END ===


def _write_mn01_te(
    ws: Any,
    person_rows: list[dict[str, Any]],
    record_map: dict[int, SurveyPersonYearRecord],
    school_year: SchoolYear,
) -> None:
    start_year, _end_year = _reference_years(school_year)
    age_cols = {age: 5 + age for age in range(7)}

    # Tuổi PCGD theo năm dương lịch bắt đầu của năm học.
    for age, column in age_cols.items():
        ws.cell(5, column).value = start_year - age
        ws.cell(6, column).value = f"{age} tuổi"

    _clear_range_values(ws, "E7:L29")

    age_rows: dict[int, list[dict[str, Any]]] = {
        age: [
            row
            for row in person_rows
            if row.get("age") == age
        ]
        for age in range(7)
    }

    row_values: dict[int, dict[int, Any]] = defaultdict(dict)

    for age, rows in age_rows.items():
        disabled = [row for row in rows if _is_disabled(row)]
        can_learn = [row for row in disabled if _v243_can_learn(row)]
        access = [row for row in disabled if _v243_access(row)]
        mobilizable = [row for row in rows if _v243_mobilizable(row)]
        studying = [row for row in rows if _is_studying(row)]

        local_studying = [
            row
            for row in studying
            if _v243_local_study(row)
        ]
        cross_studying = [
            row
            for row in studying
            if _v243_cross_study(row)
        ]

        transferred_in = [
            row
            for row in rows
            if str(
                row.get("learning_status_code") or ""
            ).upper() == "CHUYEN_DEN"
        ]
        transferred_out = [
            row
            for row in rows
            if str(
                row.get("learning_status_code") or ""
            ).upper() in TRANSFER_OUT_CODES
        ]
        dead = [
            row
            for row in rows
            if str(
                row.get("learning_status_code") or ""
            ).upper() in DEAD_CODES
        ]

        two_sessions = [
            row
            for row in studying
            if _v243_two_sessions(row)
        ]
        completed = [
            row
            for row in rows
            if _v243_completed(row)
        ]

        row_values[7][age] = len(rows)
        row_values[8][age] = sum(_is_female(row) for row in rows)
        row_values[9][age] = sum(_is_minority(row) for row in rows)
        row_values[10][age] = len(disabled)
        row_values[11][age] = len(can_learn)
        row_values[12][age] = len(access)
        row_values[13][age] = len(mobilizable)
        row_values[14][age] = len(studying)
        row_values[15][age] = len(local_studying)
        row_values[16][age] = len(cross_studying)
        row_values[17][age] = _safe_ratio(
            len(studying),
            len(mobilizable),
        )
        row_values[18][age] = sum(
            _is_female(row)
            for row in studying
        )
        row_values[19][age] = sum(
            _is_minority(row)
            for row in studying
        )
        row_values[20][age] = sum(
            _is_minority(row)
            and _v243_true(row.get("prepared_vietnamese"))
            for row in studying
        )
        row_values[21][age] = len(transferred_in)
        row_values[22][age] = len(two_sessions)
        row_values[23][age] = _safe_percent(
            len(two_sessions),
            len(studying),
        )
        row_values[24][age] = len(dead)
        row_values[25][age] = len(transferred_out)
        row_values[26][age] = len(transferred_in)

        if age == 5:
            row_values[27][age] = len(completed)
            row_values[28][age] = _safe_percent(
                len(completed),
                len(studying),
            )
            row_values[29][age] = sum(
                row in transferred_in
                for row in completed
            )
        else:
            row_values[27][age] = None
            row_values[28][age] = None
            row_values[29][age] = None

    for row_number, values in row_values.items():
        for age, value in values.items():
            ws.cell(
                row_number,
                age_cols[age],
            ).value = value

    # Tổng cộng của biểu gốc là nhóm 0-5 tuổi.
    for row_number in range(7, 30):
        values = [
            row_values.get(row_number, {}).get(age)
            for age in range(6)
        ]
        numeric_values = [
            value
            for value in values
            if isinstance(value, (int, float))
        ]

        if row_number == 17:
            ws.cell(row_number, 12).value = _safe_ratio(
                sum(
                    row_values[14].get(age, 0) or 0
                    for age in range(6)
                ),
                sum(
                    row_values[13].get(age, 0) or 0
                    for age in range(6)
                ),
            )
        elif row_number == 23:
            ws.cell(row_number, 12).value = _safe_percent(
                sum(
                    row_values[22].get(age, 0) or 0
                    for age in range(6)
                ),
                sum(
                    row_values[14].get(age, 0) or 0
                    for age in range(6)
                ),
            )
        elif row_number == 28:
            ws.cell(row_number, 12).value = None
        elif numeric_values:
            ws.cell(row_number, 12).value = sum(numeric_values)

    for column in list(age_cols.values()) + [12]:
        ws.cell(17, column).number_format = "0.00%"
        ws.cell(23, column).number_format = "0.00"

    age_5 = age_rows[5]
    age_34 = age_rows[3] + age_rows[4]

    age_5_mobilizable = sum(_v243_mobilizable(row) for row in age_5)
    age_34_mobilizable = sum(_v243_mobilizable(row) for row in age_34)

    age_5_studying = sum(_is_studying(row) for row in age_5)
    age_34_studying = sum(_is_studying(row) for row in age_34)

    age_5_completed = sum(_v243_completed(row) for row in age_5)
    age_34_completed = sum(_v243_completed(row) for row in age_34)

    age_5_can_learn = sum(_v243_can_learn(row) for row in age_5)
    age_5_access = sum(_v243_access(row) for row in age_5)

    age_5_two_sessions = sum(
        _is_studying(row) and _v243_two_sessions(row)
        for row in age_5
    )
    age_34_two_sessions = sum(
        _is_studying(row) and _v243_two_sessions(row)
        for row in age_34
    )

    ws["D33"] = age_5_studying
    ws["E33"] = _safe_percent(
        age_5_studying,
        age_5_mobilizable,
    )

    ws["D34"] = age_5_completed
    ws["E34"] = _safe_percent(
        age_5_completed,
        age_5_studying,
    )

    ws["D35"] = age_5_access
    ws["E35"] = _safe_percent(
        age_5_access,
        age_5_can_learn,
    )

    ws["D36"] = age_5_two_sessions
    ws["E36"] = _safe_percent(
        age_5_two_sessions,
        age_5_studying,
    )

    ws["D38"] = age_34_studying
    ws["E38"] = _safe_percent(
        age_34_studying,
        age_34_mobilizable,
    )

    ws["D39"] = age_34_completed
    ws["E39"] = _safe_percent(
        age_34_completed,
        age_34_studying,
    )

    ws["D40"] = age_34_two_sessions
    ws["E40"] = _safe_percent(
        age_34_two_sessions,
        age_34_studying,
    )




def _write_mn02(
    ws: Any,
    person_rows: list[dict[str, Any]],
    school_rows: list[dict[str, Any]],
) -> None:
    _clear_range_values(ws, "D8:V9")

    groups = {
        8: [
            row
            for row in person_rows
            if row.get("age") in {3, 4}
        ],
        9: [
            row
            for row in person_rows
            if row.get("age") == 5
        ],
    }

    for row_number, rows in groups.items():
        mobilizable = [
            row
            for row in rows
            if _v243_mobilizable(row)
        ]
        studying = [
            row
            for row in rows
            if _is_studying(row)
        ]
        disabled = [
            row
            for row in rows
            if _is_disabled(row)
        ]
        can_learn = [
            row
            for row in disabled
            if _v243_can_learn(row)
        ]
        access = [
            row
            for row in disabled
            if _v243_access(row)
        ]

        class_keys = {
            (
                row.get("school_display"),
                row.get("class_display"),
            )
            for row in studying
            if (
                row.get("school_display")
                or row.get("class_display")
            )
        }

        completed = sum(
            _v243_completed(row)
            for row in rows
        )

        ws.cell(row_number, 3).value = None
        ws.cell(row_number, 4).value = len(school_rows)
        ws.cell(row_number, 5).value = None
        ws.cell(row_number, 6).value = None
        ws.cell(row_number, 7).value = len(class_keys)
        ws.cell(row_number, 8).value = len(class_keys)
        ws.cell(row_number, 9).value = 0
        ws.cell(row_number, 10).value = len(mobilizable)
        ws.cell(row_number, 11).value = len(studying)
        ws.cell(row_number, 12).value = _safe_percent(
            len(studying),
            len(mobilizable),
        )
        ws.cell(row_number, 13).value = completed
        ws.cell(row_number, 14).value = _safe_percent(
            completed,
            len(studying),
        )
        ws.cell(row_number, 15).value = len(disabled)
        ws.cell(row_number, 16).value = len(can_learn)
        ws.cell(row_number, 17).value = len(access)
        ws.cell(row_number, 18).value = _safe_percent(
            len(access),
            len(can_learn),
        )
        ws.cell(row_number, 19).value = None
        ws.cell(row_number, 20).value = None
        ws.cell(row_number, 21).value = None
        ws.cell(row_number, 22).value = "Thiếu dữ liệu CSVC"




def _write_disability_sheet(
    ws: Any,
    person_rows: list[dict[str, Any]],
    school_year: SchoolYear,
) -> None:
    start_year, _ = _reference_years(school_year)
    _clear_range_values(ws, "C9:M15")

    type_columns = {
        "VAN_DONG": 4,
        "NGHE_NOI": 5,
        "NHIN": 6,
        "THAN_KINH_TAM_THAN": 7,
        "TRI_TUE": 8,
        "TU_KY": 9,
        "HOC_TAP": 10,
        "DA_TAT": 11,
        "KHAC": 11,
    }

    totals_by_col: dict[int, int] = defaultdict(int)

    for age in range(6):
        row_number = 9 + age
        rows = [
            row
            for row in person_rows
            if row.get("age") == age
            and _is_disabled(row)
        ]

        ws.cell(row_number, 1).value = start_year - age
        ws.cell(row_number, 2).value = age
        ws.cell(row_number, 3).value = len(rows)
        totals_by_col[3] += len(rows)

        for row in rows:
            code = str(
                row.get("disability_type") or "KHAC"
            ).upper()
            column = type_columns.get(code, 11)
            current = ws.cell(row_number, column).value or 0
            ws.cell(row_number, column).value = current + 1
            totals_by_col[column] += 1

        access = sum(
            _v243_access(row)
            for row in rows
        )
        ws.cell(row_number, 12).value = access
        ws.cell(row_number, 13).value = _safe_percent(
            access,
            len(rows),
        )
        totals_by_col[12] += access

    ws["B15"] = "Tổng:"
    for column in range(3, 13):
        ws.cell(15, column).value = totals_by_col.get(column, 0)

    ws.cell(15, 13).value = _safe_percent(
        totals_by_col.get(12, 0),
        totals_by_col.get(3, 0),
    )




def _write_tracking_book(
    ws: Any,
    person_rows: list[dict[str, Any]],
    person_map: dict[int, SurveyPerson],
    record_map: dict[int, SurveyPersonYearRecord],
    school_year: SchoolYear,
) -> None:
    _clear_range_values(ws, "A11:T8540")
    start_year, _end_year = _reference_years(school_year)
    first_start = start_year - 6
    for offset, column in enumerate(range(12, 19)):
        ws.cell(7, column).value = first_start + offset
        ws.cell(8, column).value = first_start + offset + 1
    target_column = 18

    sorted_rows = sorted(
        [row for row in person_rows if isinstance(row.get("age"), int) and 0 <= row.get("age") <= 6],
        key=lambda item: (
            str(item.get("commune_name") or ""),
            str(item.get("hamlet_name") or ""),
            str(item.get("full_name") or ""),
        ),
    )
    for index, row in enumerate(sorted_rows, start=1):
        excel_row = 10 + index
        if excel_row > 8540:
            break
        person = person_map.get(int(row["person_id"]))
        record = record_map.get(int(row["year_record_id"])) if row.get("year_record_id") else None
        name_left, name_right = _split_name(row.get("full_name"))
        parent_name = ""
        if person is not None:
            parent_name = person.mother_name or person.father_name or ""
        parent_left, parent_right = _split_name(parent_name)

        ws.cell(excel_row, 1).value = index
        ws.cell(excel_row, 2).value = row.get("form_number")
        ws.cell(excel_row, 3).value = name_left
        ws.cell(excel_row, 4).value = name_right
        ws.cell(excel_row, 5).value = row.get("date_of_birth")
        ws.cell(excel_row, 5).number_format = "dd/mm/yyyy"
        ws.cell(excel_row, 6).value = "x" if _is_female(row) else None
        ws.cell(excel_row, 7).value = row.get("hamlet_name")
        ws.cell(excel_row, 8).value = row.get("commune_name")
        ws.cell(excel_row, 9).value = row.get("ethnic_group")
        ws.cell(excel_row, 10).value = parent_left
        ws.cell(excel_row, 11).value = parent_right

        age = row.get("age")
        program = "CTNT" if isinstance(age, int) and age <= 2 else "CTMG"
        class_text = row.get("class_display") or ""
        ws.cell(excel_row, target_column).value = (
            f"{program}\n{class_text}" if class_text else program
        )
        ws.cell(excel_row, 19).value = row.get("school_display")

        notes: list[str] = []
        if _is_disabled(row):
            disability = row.get("disability_type_label") or "Trẻ khuyết tật"
            notes.append(str(disability))
        if str(row.get("residency_status") or "").upper() != "THUONG_TRU":
            notes.append(str(row.get("residency_status_label") or "Tạm trú"))
        learning_code = str(row.get("learning_status_code") or "").upper()
        if learning_code in {"CHUYEN_DEN", "CHUYEN_DI", "BO_HOC", "THOI_HOC"}:
            notes.append(str(row.get("learning_status_label") or learning_code))
        if not row.get("personal_id"):
            notes.append("Thiếu số định danh")
        if record is not None and record.notes:
            notes.append(str(record.notes))
        ws.cell(excel_row, 20).value = "; ".join(notes)

        for column in range(1, 21):
            ws.cell(excel_row, column).alignment = Alignment(
                horizontal="center" if column not in {3, 4, 7, 8, 9, 10, 11, 19, 20} else "left",
                vertical="center",
                wrap_text=True,
            )




def _get_sheet(workbook: Any, expected_name: str) -> Any:
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip() == expected_name.strip():
            return workbook[sheet_name]
    raise KeyError(f"Không tìm thấy trang mẫu: {expected_name}")

def _blank_future_modules(
    workbook: Any,
    commune_label: str,
    year_code: str,
) -> None:
    """
    Không tạo số liệu giả khi phân hệ hiện chưa có dữ liệu nguồn.

    Giữ đúng mẫu nhưng ghi rõ trạng thái "chưa nhập dữ liệu" thay vì
    ghi "sẽ triển khai", vì các phân hệ đã có cấu trúc nhưng scope hiện tại rỗng.
    """
    ws = _get_sheet(workbook, "MN-01 CSVC")
    _write_header(ws, commune_label, year_code)
    _clear_range_values(ws, "C9:X70")
    _clear_range_values(ws, "A11:B70")
    ws["B11"] = (
        "CHƯA NHẬP DỮ LIỆU CƠ SỞ VẬT CHẤT "
        f"CHO NĂM HỌC {year_code}"
    )
    ws["B11"].font = Font(
        color="C00000",
        bold=True,
    )





# === BAI_13B_12_V2_4_3_11_MN02_CSVC_REAL_DATA_START ===
def _b24311_int(value: Any) -> int:
    try:
        return max(0, int(value or 0))
    except (TypeError, ValueError):
        return 0


def _b24311_ratio(numerator: int, denominator: int) -> float | None:
    if int(denominator or 0) <= 0:
        return None
    return round(float(numerator or 0) / float(denominator), 2)


def _b24311_load_mn02_csvc(
    db: Session,
    context: dict[str, Any],
) -> dict[str, Any]:
    report_data = context.get("report_data") or {}
    school_year = report_data.get("school_year")

    year_id = context.get("selected_year_id")
    if year_id is None and school_year is not None:
        year_id = getattr(school_year, "id", None)

    school_id = context.get("selected_school_id")
    commune_id = context.get("selected_commune_id")

    result = {
        "has_csvc": False,
        "record_count": 0,
        "school_count_with_csvc": 0,
        "independent_facility_count": None,
        "satellite_site_count": None,
        "classes_34": None,
        "classes_5": None,
        "classes_total": None,
        "preschool_rooms": None,
        "room_ratio": None,
        "equipped_total": None,
        "equipped_34": None,
        "equipped_5": None,
        "equipment_age_split_known": False,
    }

    if year_id is None:
        return result

    filters = ["c.school_year_id = :year_id"]
    params: dict[str, Any] = {"year_id": int(year_id)}

    if school_id is not None:
        filters.append("c.school_id = :school_id")
        params["school_id"] = int(school_id)
    elif commune_id is not None:
        filters.append("s.commune_id = :commune_id")
        params["commune_id"] = int(commune_id)

    sql = """
        SELECT
            c.school_id,
            c.satellite_site_count,
            c.preschool_class_3_4_count,
            c.preschool_class_5_count,
            c.permanent_preschool_room_count,
            c.semi_permanent_preschool_room_count,
            c.temporary_preschool_room_count,
            c.equipped_preschool_class_count
        FROM school_mn01_csvc_inputs AS c
        JOIN schools AS s
          ON s.id = c.school_id
        WHERE {where_sql}
    """.format(where_sql=" AND ".join(filters))

    rows = list(
        db.execute(
            text(sql),
            params,
        ).mappings().all()
    )

    if not rows:
        return result

    result["has_csvc"] = True
    result["record_count"] = len(rows)
    result["school_count_with_csvc"] = len(
        {int(row["school_id"]) for row in rows}
    )

    satellite = sum(
        _b24311_int(row["satellite_site_count"])
        for row in rows
    )
    classes_34 = sum(
        _b24311_int(row["preschool_class_3_4_count"])
        for row in rows
    )
    classes_5 = sum(
        _b24311_int(row["preschool_class_5_count"])
        for row in rows
    )
    rooms = sum(
        _b24311_int(row["permanent_preschool_room_count"])
        + _b24311_int(row["semi_permanent_preschool_room_count"])
        + _b24311_int(row["temporary_preschool_room_count"])
        for row in rows
    )
    equipped = sum(
        _b24311_int(row["equipped_preschool_class_count"])
        for row in rows
    )
    total_classes = classes_34 + classes_5

    result.update(
        {
            "satellite_site_count": satellite,
            "classes_34": classes_34,
            "classes_5": classes_5,
            "classes_total": total_classes,
            "preschool_rooms": rooms,
            "room_ratio": _b24311_ratio(rooms, total_classes),
            "equipped_total": equipped,
        }
    )

    if total_classes > 0 and equipped >= total_classes:
        result["equipped_34"] = classes_34
        result["equipped_5"] = classes_5
        result["equipment_age_split_known"] = True
    elif equipped == 0:
        result["equipped_34"] = 0
        result["equipped_5"] = 0
        result["equipment_age_split_known"] = True

    table_exists = db.execute(
        text(
            """
            SELECT 1
            FROM sqlite_master
            WHERE type='table'
              AND name='school_site_year_records'
            LIMIT 1
            """
        )
    ).scalar()

    if table_exists is not None:
        site_filters = ["school_year_id = :year_id"]
        site_params: dict[str, Any] = {"year_id": int(year_id)}

        if school_id is not None:
            site_filters.append("school_id = :site_school_id")
            site_params["site_school_id"] = int(school_id)
        elif commune_id is not None:
            site_filters.append(
                "school_id IN (SELECT id FROM schools "
                "WHERE commune_id = :site_commune_id)"
            )
            site_params["site_commune_id"] = int(commune_id)

        independent = db.execute(
            text(
                """
                SELECT COUNT(*)
                FROM school_site_year_records
                WHERE {where_sql}
                  AND COALESCE(is_active, 1) = 1
                  AND COALESCE(is_independent, 0) = 1
                """.format(
                    where_sql=" AND ".join(site_filters)
                )
            ),
            site_params,
        ).scalar()

        result["independent_facility_count"] = int(independent or 0)

    return result


def _b24311_apply_mn02_csvc(
    ws: Any,
    csvc: dict[str, Any],
    context: dict[str, Any],
    staff_ratios: dict[str, float | None],
) -> None:
    if not csvc.get("has_csvc"):
        ws["V8"] = "Thiếu dữ liệu CSVC"
        ws["V9"] = "Thiếu dữ liệu CSVC"
        return

    independent = csvc.get("independent_facility_count")
    satellite = csvc.get("satellite_site_count")

    if independent is not None:
        ws["E8"] = independent
        ws["E9"] = independent

    if satellite is not None:
        ws["F8"] = satellite
        ws["F9"] = satellite

    row_specs = (
        (8, "classes_34", "age_34", "equipped_34"),
        (9, "classes_5", "age_5", "equipped_5"),
    )

    for row_number, class_key, staff_key, equipment_key in row_specs:
        class_count = csvc.get(class_key)
        if class_count is not None:
            class_count = _b24311_int(class_count)
            ws.cell(row_number, 7).value = class_count

            single = ws.cell(row_number, 8).value
            mixed = ws.cell(row_number, 9).value
            try:
                split_total = int(single or 0) + int(mixed or 0)
            except (TypeError, ValueError):
                split_total = -1

            if split_total != class_count:
                ws.cell(row_number, 8).value = None
                ws.cell(row_number, 9).value = None

        staff_ratio = staff_ratios.get(staff_key)
        ws.cell(row_number, 19).value = staff_ratio
        if staff_ratio is not None:
            ws.cell(row_number, 19).number_format = "0.00"

        room_ratio = csvc.get("room_ratio")
        ws.cell(row_number, 20).value = room_ratio
        if room_ratio is not None:
            ws.cell(row_number, 20).number_format = "0.00"

        equipped_value = csvc.get(equipment_key)
        ws.cell(row_number, 21).value = equipped_value

        missing: list[str] = []
        if staff_ratio is None:
            missing.append("phân nhóm GV")
        if room_ratio is None:
            missing.append("phòng/lớp")
        if not csvc.get("equipment_age_split_known"):
            missing.append("TBDH theo độ tuổi")

        if missing:
            ws.cell(row_number, 22).value = (
                "Chưa đủ dữ liệu: " + ", ".join(missing)
            )
        elif context.get("selected_school_id") is not None:
            ws.cell(row_number, 22).value = "Chưa kết luận cấp xã"
        else:
            ws.cell(row_number, 22).value = "Chờ kết luận cấp xã"
# === BAI_13B_12_V2_4_3_11_MN02_CSVC_REAL_DATA_END ===


# === BAI_13B_12_V2_4_3_14_MN02_COMMUNE_CONCLUSION_START ===
def _b24314_mn02_completion_rates_by_age(
    person_rows: list[dict[str, Any]],
) -> dict[int, float | None]:
    """Tính tỷ lệ hoàn thành CTGDMN theo từng tuổi 3, 4, 5.

    Chỉ kết luận khi toàn bộ trẻ đang học của từng tuổi đã có giá trị
    hoàn thành/chưa hoàn thành rõ ràng. Không biến dữ liệu chưa nhập
    thành "Không đạt".
    """
    result: dict[int, float | None] = {}

    for age in (3, 4, 5):
        attending = [
            row
            for row in person_rows
            if row.get("age") == age and _is_studying(row)
        ]
        if not attending:
            result[age] = None
            continue

        flags: list[bool] = []
        has_unknown = False
        for row in attending:
            raw = row.get("completed_preschool_by_age")
            if raw is None and age == 5:
                raw = row.get("completed_preschool_5")

            if _v243_true(raw):
                flags.append(True)
            elif _v243_false(raw):
                flags.append(False)
            else:
                has_unknown = True
                break

        if has_unknown:
            result[age] = None
        else:
            result[age] = _safe_percent(
                sum(1 for flag in flags if flag),
                len(attending),
            )

    return result


def _b24314_apply_mn02_commune_conclusion(
    ws: Any,
    person_rows: list[dict[str, Any]],
    context: dict[str, Any],
) -> None:
    """Tự động ghi Đạt/Không đạt cho cột 22 (V) của MN-02 cấp xã.

    V2.4.3.16:
    - Kết luận theo đúng số liệu ĐÃ TỔNG HỢP trên MN-02 cho toàn bộ trẻ 3-5 tuổi.
    - Tỷ lệ huy động = (K8 + K9) / (J8 + J9).
    - Tỷ lệ hoàn thành = (M8 + M9) / (K8 + K9).
    - Không yêu cầu phải có riêng dữ liệu từng tuổi 3, 4, 5 mới được kết luận,
      vì mẫu MN-02 đang gộp 3-4 tuổi ở dòng 8 và 5 tuổi ở dòng 9.
    - Nếu GV/CSVC/TBDH còn thiếu, giữ nguyên cảnh báo của bước trước.
    """
    if context.get("selected_school_id") is not None:
        return

    selected_commune = context.get("selected_commune")
    if selected_commune is None or context.get("selected_commune_id") is None:
        return

    # Chỉ kết luận khi bước GV/CSVC/TBDH đã xác nhận đủ dữ liệu đầu vào.
    for row_number in (8, 9):
        current = str(ws.cell(row_number, 22).value or "").strip()
        if current.startswith("Thiếu dữ liệu") or current.startswith("Chưa đủ dữ liệu:"):
            return

    def _cell_number(row: int, column: int) -> float | None:
        value = ws.cell(row, column).value
        if value is None or isinstance(value, bool):
            return None
        try:
            return float(value)
        except (TypeError, ValueError):
            return None

    mobilizable_34 = _cell_number(8, 10)
    mobilizable_5 = _cell_number(9, 10)
    attending_34 = _cell_number(8, 11)
    attending_5 = _cell_number(9, 11)
    completed_34 = _cell_number(8, 13)
    completed_5 = _cell_number(9, 13)

    values = (
        mobilizable_34, mobilizable_5,
        attending_34, attending_5,
        completed_34, completed_5,
    )
    if any(value is None for value in values):
        conclusion = "Chưa đủ dữ liệu"
    else:
        mobilizable_total = float(mobilizable_34) + float(mobilizable_5)
        attending_total = float(attending_34) + float(attending_5)
        completed_total = float(completed_34) + float(completed_5)

        attendance_rate = _safe_percent(attending_total, mobilizable_total)
        completion_rate = _safe_percent(completed_total, attending_total)

        from app.services.pcgd_business_rules import evaluate_preschool_3_5_commune

        result = evaluate_preschool_3_5_commune(
            {
                "attendance_rate_3_5": attendance_rate,
                # Dùng một tỷ lệ tổng 3-5 tuổi, đúng với tiêu chuẩn công nhận cấp xã.
                "completion_rate_3_5_by_age": completion_rate,
            },
            is_special_difficult=bool(
                getattr(selected_commune, "is_special_difficulty_area", False)
            ),
        )

        if result.passed:
            conclusion = "Đạt"
        elif tuple(result.missing_fields or ()):
            conclusion = "Chưa đủ dữ liệu"
        else:
            conclusion = "Không đạt"

    # Mẫu MN-02 có hai dòng nhóm tuổi nhưng cùng là kết luận của một xã.
    ws["V8"] = conclusion
    ws["V9"] = conclusion

# === BAI_13B_12_V2_4_3_16_MN02_CONCLUSION_FROM_VISIBLE_TOTALS ===
# === BAI_13B_12_V2_4_3_14_MN02_COMMUNE_CONCLUSION_END ===


# === BAI_13B_12_V2_4_3_12_MN02_STAFF_RATIO_START ===
# === BAI_13B_12_V2_4_3_13_MN02_STAFF_RATIO_CSVCDENOM_START ===
def _b24312_complete_mn02_staff_ratios(
    staff_data: dict[str, Any],
    csvc: dict[str, Any],
    current_ratios: dict[str, float | None],
) -> dict[str, float | None]:
    """
    V2.4.3.13:
    - Tử số GV lấy từ staff_year_records theo teaching_age_group.
    - Mẫu số lớp lấy TRỰC TIẾP từ MN-01-CSVC năm hiện tại.
    - Không còn phụ thuộc school_staff_year_summaries/classes_34/classes_5
      khi tính cột 19 của MN-02.

    Nếu đã có GV phân nhóm tuổi:
      TUOI_3_4 / classes_34
      TUOI_5   / classes_5

    Nếu hoàn toàn chưa phân nhóm tuổi:
      tổng GV Mẫu giáo / tổng lớp Mẫu giáo
      làm fallback cho các dòng có lớp.

    Nếu chỉ phân nhóm một phần:
      không tự phân bổ phần còn thiếu.
    """
    result = dict(current_ratios or {})
    result.setdefault("age_34", None)
    result.setdefault("age_5", None)

    rows = list(staff_data.get("rows") or [])
    records = [
        record
        for row in rows
        for record in list(row.get("records") or [])
    ]

    preschool_teachers = [
        item
        for item in records
        if getattr(item, "position_group", None) == "GIAO_VIEN"
        and getattr(item, "teaching_level", None) == "MAU_GIAO"
    ]

    specs = (
        ("age_34", "TUOI_3_4", "classes_34"),
        ("age_5", "TUOI_5", "classes_5"),
    )

    known_counts: dict[str, int] = {}
    for key, age_code, _class_key in specs:
        known_counts[key] = sum(
            1
            for item in preschool_teachers
            if getattr(item, "teaching_age_group", None) == age_code
        )

    has_any_age_assignment = any(
        count > 0
        for count in known_counts.values()
    )

    if has_any_age_assignment:
        for key, _age_code, class_key in specs:
            teachers = int(known_counts.get(key, 0) or 0)
            classes = _b24311_int(csvc.get(class_key))

            if teachers > 0 and classes > 0:
                result[key] = round(
                    float(teachers) / float(classes),
                    2,
                )
            elif teachers <= 0:
                # Có phân nhóm một phần: không suy đoán nhóm còn thiếu.
                result[key] = None

        return result

    total_teachers = len(preschool_teachers)
    total_classes = _b24311_int(csvc.get("classes_total"))

    if total_teachers <= 0 or total_classes <= 0:
        return result

    pooled_ratio = round(
        float(total_teachers) / float(total_classes),
        2,
    )

    for key, _age_code, class_key in specs:
        if _b24311_int(csvc.get(class_key)) > 0:
            result[key] = pooled_ratio

    return result
# === BAI_13B_12_V2_4_3_13_MN02_STAFF_RATIO_CSVCDENOM_END ===

# === BAI_13B_12_V2_4_3_12_MN02_STAFF_RATIO_END ===


def build_template_workbook(
    *,
    db: Session,
    context: dict[str, Any],
) -> Any:
    if not TEMPLATE_PATH.is_file():
        raise FileNotFoundError(f"Không tìm thấy tệp mẫu: {TEMPLATE_PATH}")
    workbook = load_workbook(TEMPLATE_PATH)
    required_sheets = {
        "MN-01 TE",
        "MN-01 GV",
        "MN-01 CSVC",
        "MN-02",
        "MN - Tài chính",
        "MN- Trẻ KT",
        "Sổ theo dõi PCGDMN",
    }
    normalized_sheet_names = {name.strip() for name in workbook.sheetnames}
    missing = sorted(required_sheets - normalized_sheet_names)
    if missing:
        raise RuntimeError("Tệp mẫu thiếu trang: " + ", ".join(missing))

    report_data = context["report_data"]
    person_rows = list(report_data.get("person_rows") or [])
    person_map, record_map = _load_extra_models(db, person_rows)
    staff_data = _load_staff_report_data(db, context)
    school_year = report_data["school_year"]
    commune_label = (
        context["selected_commune"].name
        if context.get("selected_commune") is not None
        else context["scope_label"]
    )
    if context.get("selected_school") is not None:
        commune_label = context["selected_school"].commune.name

    for ws in workbook.worksheets:
        _write_header(ws, commune_label, school_year.code)

    _write_mn01_te(
        _get_sheet(workbook, "MN-01 TE"),
        person_rows,
        record_map,
        school_year,
    )
    _write_mn02(
        _get_sheet(workbook, "MN-02"),
        person_rows,
        list(report_data.get("school_rows") or []),
    )
    _write_mn01_gv(
        _get_sheet(workbook, "MN-01 GV"),
        staff_data,
    )
    staff_ratios = _staff_ratios_by_age(staff_data)
    mn02_ws = _get_sheet(workbook, "MN-02")
    mn02_csvc = _b24311_load_mn02_csvc(db, context)
    staff_ratios = _b24312_complete_mn02_staff_ratios(
        staff_data,
        mn02_csvc,
        staff_ratios,
    )
    _b24311_apply_mn02_csvc(
        mn02_ws,
        mn02_csvc,
        context,
        staff_ratios,
    )
    _b24314_apply_mn02_commune_conclusion(
        mn02_ws,
        person_rows,
        context,
    )
    _write_disability_sheet(
        _get_sheet(workbook, "MN- Trẻ KT"),
        person_rows,
        school_year,
    )
    _write_tracking_book(
        _get_sheet(workbook, "Sổ theo dõi PCGDMN"),
        person_rows,
        person_map,
        record_map,
        school_year,
    )
    _blank_future_modules(workbook, commune_label, school_year.code)
    _v243_normalize_year_labels(workbook, school_year)
    start_year, _end_year = _reference_years(school_year)
    finance_ws = _get_sheet(workbook, "MN - Tài chính")
    finance_summary = fill_finance_worksheet(
        db,
        finance_ws,
        start_year=start_year,
        commune_id=context.get("selected_commune_id") if context.get("selected_school_id") is None else None,
        school_id=context.get("selected_school_id"),
    )
    if not finance_summary.get("has_data"):
        finance_ws["D10"].font = Font(color="C00000", bold=True)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    # === BAI_13B_11_15_2_4_6_MASTER_POSTPROCESS ===
    from app.services.mn_report_v246 import apply_formula_contract
    apply_formula_contract(workbook)
    return workbook

@router.get("", response_class=HTMLResponse)
def report_page(
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
    data_state: str = "",
    sheet_code: str = "",
    db: Session = Depends(get_db),
) -> HTMLResponse:
    normalized_state = str(data_state or "").strip().upper()
    if normalized_state not in DATA_STATE_LABELS:
        normalized_state = ""

    normalized_sheet_code = str(sheet_code or "").strip().lower()
    selected_sheet = SINGLE_SHEET_EXPORTS.get(normalized_sheet_code)
    if selected_sheet is None:
        normalized_sheet_code = ""

    context = _load_context(
        db=db,
        request=request,
        school_year_id=_parse_optional_int(school_year_id),
        commune_id=_parse_optional_int(commune_id),
        school_id=_parse_optional_int(school_id),
        data_state=normalized_state,
    )
    staff_data = _load_staff_report_data(db, context)
    export_params: dict[str, Any] = {
        "school_year_id": context["selected_year_id"],
    }
    if context["selected_commune_id"] is not None:
        export_params["commune_id"] = context["selected_commune_id"]
    if context["selected_school_id"] is not None:
        export_params["school_id"] = context["selected_school_id"]
    if normalized_state:
        export_params["data_state"] = normalized_state
    export_query = urlencode(export_params)

    if normalized_sheet_code:
        export_url = (
            "/bao-cao/pcgdmn-mau-2025/xuat-bieu/"
            + normalized_sheet_code
            + "?"
            + export_query
        )
    else:
        export_url = (
            "/bao-cao/pcgdmn-mau-2025/xuat-excel?"
            + export_query
        )

    return templates.TemplateResponse(
        request=request,
        name="reports/pcgdmn_template_report.html",
        context={
            "nguoi_dung": context["user"],
            "school_years": context["school_years"],
            "selected_year_id": context["selected_year_id"],
            "communes": context["communes"],
            "schools": context["schools"],
            "selected_commune_id": context["selected_commune_id"],
            "selected_school_id": context["selected_school_id"],
            "scope_label": context["scope_label"],
            "data_state": normalized_state,
            "data_state_labels": DATA_STATE_LABELS,
            "stats": context["stats"],
            "staff_summary": staff_data["summary"],
            "summary": context["report_data"]["summary"],
            "is_official": context["report_data"]["is_official"],
            "export_url": export_url,
            "sheet_code": normalized_sheet_code,
            "selected_sheet_label": (
                selected_sheet["label"] if selected_sheet is not None else ""
            ),
        },
    )


@router.get("/xuat-excel")
def export_excel(
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
    data_state: str = "",
    db: Session = Depends(get_db),
) -> StreamingResponse:
    normalized_state = str(data_state or "").strip().upper()
    if normalized_state not in DATA_STATE_LABELS:
        normalized_state = ""
    context = _load_context(
        db=db,
        request=request,
        school_year_id=_parse_optional_int(school_year_id),
        commune_id=_parse_optional_int(commune_id),
        school_id=_parse_optional_int(school_id),
        data_state=normalized_state,
    )
    workbook = build_template_workbook(db=db, context=context)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    year_code = context["report_data"]["school_year"].code.replace("-", "_")
    scope_code = "toan_tinh"
    if context.get("selected_school") is not None:
        scope_code = f"truong_{context['selected_school'].code}"
    elif context.get("selected_commune") is not None:
        scope_code = f"xa_{context['selected_commune'].code}"
    filename = f"bieu_mau_PCGDMN_2025_{year_code}_{scope_code}.xlsx"
    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )


# === BAI_13B_12_V2_4_2_MN_SINGLE_SHEET_EXPORT_ROUTE_START ===
@router.get("/xuat-bieu/{sheet_code}")
def export_single_sheet_excel(
    sheet_code: str,
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
    data_state: str = "",
    db: Session = Depends(get_db),
) -> StreamingResponse:
    normalized_sheet_code = str(sheet_code or "").strip().lower()
    sheet_spec = SINGLE_SHEET_EXPORTS.get(normalized_sheet_code)
    if sheet_spec is None:
        raise ValueError("Mã biểu Mầm non không hợp lệ.")

    normalized_state = str(data_state or "").strip().upper()
    if normalized_state not in DATA_STATE_LABELS:
        normalized_state = ""

    context = _load_context(
        db=db,
        request=request,
        school_year_id=_parse_optional_int(school_year_id),
        commune_id=_parse_optional_int(commune_id),
        school_id=_parse_optional_int(school_id),
        data_state=normalized_state,
    )

    workbook = build_template_workbook(
        db=db,
        context=context,
    )

    target_sheet = next(
        (
            ws
            for ws in workbook.worksheets
            if ws.title.strip() == sheet_spec["sheet_name"]
        ),
        None,
    )
    if target_sheet is None:
        raise RuntimeError(
            "Tệp mẫu không có trang " + sheet_spec["sheet_name"]
        )

    for ws in list(workbook.worksheets):
        if ws is not target_sheet:
            workbook.remove(ws)

    target_sheet.sheet_state = "visible"
    workbook.active = 0

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    year_code = context["report_data"]["school_year"].code.replace("-", "_")
    scope_code = "toan_tinh"
    if context.get("selected_school") is not None:
        scope_code = f"truong_{context['selected_school'].code}"
    elif context.get("selected_commune") is not None:
        scope_code = f"xa_{context['selected_commune'].code}"

    filename = (
        f"{sheet_spec['filename']}_{year_code}_{scope_code}.xlsx"
    )

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )
# === BAI_13B_12_V2_4_2_MN_SINGLE_SHEET_EXPORT_ROUTE_END ===
