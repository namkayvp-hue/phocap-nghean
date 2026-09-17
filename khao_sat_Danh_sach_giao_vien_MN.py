import sys
import json
import hashlib
import unicodedata
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from openpyxl import load_workbook

ROOT = Path(r"C:\PhoCap")

DB = ROOT / "data" / "phocap.db"

TH_JSON = (
    ROOT
    / "data"
    / "school_merger_source_rosters"
    / "TH_2025_2026.json"
)

EXCEL = (
    ROOT
    / "uploads"
    / "Danh_sach_giao_vien.xlsx"
)

def sha256(path):
    h = hashlib.sha256()

    with path.open("rb") as f:
        while True:
            b = f.read(1024 * 1024)

            if not b:
                break

            h.update(b)

    return h.hexdigest()


def norm(value):
    s = str(value or "").strip().lower()

    s = "".join(
        c
        for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )

    s = s.replace("đ", "d")

    return " ".join(s.split())


for p in (DB, TH_JSON, EXCEL):
    if not p.exists():
        raise SystemExit(
            "Không tìm thấy file bắt buộc: " + str(p)
        )

db_before = sha256(DB)

print("=" * 118)
print("KHẢO SÁT FILE NGUỒN ĐỘI NGŨ MN 2025-2026")
print("CHỈ ĐỌC - KHÔNG TẠO REGISTRY - KHÔNG GHI DATABASE")
print("=" * 118)

# ==========================================================
# 1. Đọc cấu trúc registry TH hiện đang hoạt động
# ==========================================================

th = json.loads(
    TH_JSON.read_text(encoding="utf-8")
)

print()
print("=" * 118)
print("1. CẤU TRÚC TH_2025_2026.json ĐANG HOẠT ĐỘNG")
print("=" * 118)

for key in (
    "dataset_id",
    "level_code",
    "school_year_code",
    "source_file",
    "source_sha256",
    "raw_row_count",
    "row_count",
    "excluded_row_count",
    "school_count",
):
    print(f"{key:22} =", th.get(key))

rows = th.get("rows") or []

print()
print("Số rows thực tế      =", len(rows))

if rows:
    print(
        "Các khóa của 1 row   =",
        sorted(rows[0].keys()),
    )

    print()
    print("MẪU ROW TH ĐẦU TIÊN:")

    print(
        json.dumps(
            rows[0],
            ensure_ascii=False,
            indent=2,
            default=str,
        )
    )

# ==========================================================
# 2. Đọc workbook ứng viên
# ==========================================================

print()
print("=" * 118)
print("2. FILE EXCEL ỨNG VIÊN")
print("=" * 118)

print("File   :", EXCEL)
print("SHA256 :", sha256(EXCEL))
print(
    "Dung lượng:",
    EXCEL.stat().st_size,
    "bytes",
)

wb = load_workbook(
    EXCEL,
    read_only=True,
    data_only=True,
)

print("Sheets :", wb.sheetnames)

# ==========================================================
# 3. Khảo sát từng sheet
# ==========================================================

for ws in wb.worksheets:

    print()
    print("=" * 118)
    print("SHEET:", ws.title)
    print("=" * 118)

    print(
        "max_row =",
        ws.max_row,
        "| max_column =",
        ws.max_column,
    )

    # ------------------------------------------------------
    # In tối đa 20 dòng đầu có dữ liệu để xác định header
    # ------------------------------------------------------

    print()
    print("20 DÒNG ĐẦU CÓ DỮ LIỆU:")

    printed = 0

    for row_no, row in enumerate(
        ws.iter_rows(values_only=True),
        1,
    ):
        vals = [
            "" if v is None else str(v).strip()
            for v in row
        ]

        nonempty = [
            v
            for v in vals
            if v != ""
        ]

        if not nonempty:
            continue

        short = vals[:20]

        print(
            f"Dòng {row_no:>4}:",
            json.dumps(
                short,
                ensure_ascii=False,
            ),
        )

        printed += 1

        if printed >= 20:
            break

    # ------------------------------------------------------
    # Quét toàn sheet để xác định cấp học/năm học
    # ------------------------------------------------------

    mn_rows = 0
    th_rows = 0
    thcs_rows = 0

    year_2526_rows = 0
    year_2627_rows = 0

    sample_mn = []
    sample_th = []

    total_nonempty_rows = 0

    for row_no, row in enumerate(
        ws.iter_rows(values_only=True),
        1,
    ):
        vals = [
            norm(v)
            for v in row
            if v is not None
            and str(v).strip() != ""
        ]

        if not vals:
            continue

        total_nonempty_rows += 1

        joined = " | ".join(vals)

        is_mn = (
            "mam non" in joined
            or "truong mn" in joined
            or "cap mn" in joined
        )

        is_thcs = (
            "thcs" in joined
            or "trung hoc co so" in joined
        )

        is_th = (
            (
                "tieu hoc" in joined
                or "truong th " in joined
                or "| th |" in f"| {joined} |"
            )
            and not is_thcs
        )

        if is_mn:
            mn_rows += 1

            if len(sample_mn) < 5:
                sample_mn.append(
                    (row_no, joined[:500])
                )

        if is_th:
            th_rows += 1

            if len(sample_th) < 5:
                sample_th.append(
                    (row_no, joined[:500])
                )

        if is_thcs:
            thcs_rows += 1

        if (
            "2025-2026" in joined
            or "2025–2026" in joined
        ):
            year_2526_rows += 1

        if (
            "2026-2027" in joined
            or "2026–2027" in joined
        ):
            year_2627_rows += 1

    print()
    print("THỐNG KÊ NHẬN DIỆN NỘI DUNG:")
    print(
        "Tổng dòng có dữ liệu :",
        total_nonempty_rows,
    )
    print(
        "Dòng có dấu hiệu MN  :",
        mn_rows,
    )
    print(
        "Dòng có dấu hiệu TH  :",
        th_rows,
    )
    print(
        "Dòng có dấu hiệu THCS:",
        thcs_rows,
    )
    print(
        "Dòng có 2025-2026    :",
        year_2526_rows,
    )
    print(
        "Dòng có 2026-2027    :",
        year_2627_rows,
    )

    if sample_mn:
        print()
        print("MẪU 5 DÒNG CÓ DẤU HIỆU MẦM NON:")

        for row_no, text in sample_mn:
            print(
                f"  Dòng {row_no}:",
                text,
            )

    if sample_th:
        print()
        print("MẪU 5 DÒNG CÓ DẤU HIỆU TIỂU HỌC:")

        for row_no, text in sample_th:
            print(
                f"  Dòng {row_no}:",
                text,
            )

wb.close()

# ==========================================================
# 4. Kiểm tra DB
# ==========================================================

db_after = sha256(DB)

print()
print("=" * 118)
print("KIỂM TRA AN TOÀN")
print("=" * 118)

print("DB hash trước:", db_before)
print("DB hash sau  :", db_after)

if db_before != db_after:
    raise RuntimeError(
        "DỪNG: phocap.db thay đổi ngoài dự kiến."
    )

print("Database     : KHÔNG THAY ĐỔI")
print("=" * 118)
