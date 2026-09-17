from __future__ import annotations

import os
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT = Path(
    os.environ.get("PHOCAP_PROJECT", r"C:\PhoCap")
).resolve()

APP = PROJECT / "app"
TEMPLATE_ROOT = APP / "report_templates"
EXPORTS = PROJECT / "exports"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT = (
    EXPORTS
    / f"ma_tran_cong_thuc_ty_le_bai_13b_11_15_2_4_{STAMP}.txt"
)

CANONICAL_DIRS = (
    TEMPLATE_ROOT / "pcgd_mn_2025",
    TEMPLATE_ROOT / "pcgd_xmc_2025",
)

CANONICAL_FILES = (
    TEMPLATE_ROOT / "Bieu_mau_PCGDMN_2025.xlsx",
)

SOURCE_FILES = (
    APP / "pcgd_mn_report_builders_v1.py",
    APP / "pcgd_xmc_report_builders_v1.py",
    APP / "routers" / "mn_official_reports.py",
    APP / "routers" / "report_center.py",
    APP / "routers" / "pcgdmn_template_report.py",
    APP / "routers" / "preschool_3_5_report.py",
    APP / "routers" / "report_inputs.py",
)

PERCENT_WORDS = (
    "tỷ lệ",
    "tỉ lệ",
    "ty le",
    "ti le",
    "percent",
    "%",
)

RATIO_WORDS = (
    "gv/lớp",
    "gv/lop",
    "giáo viên/lớp",
    "giao vien/lop",
    "phòng/lớp",
    "phong/lop",
    "phòng học/lớp",
    "phong hoc/lop",
    "tỷ số",
    "tỉ số",
    "ratio",
)


def norm(value: object) -> str:
    return str(value or "").strip().lower()


def group_from_path(path: Path) -> str:
    name = path.name.upper()

    if "THCS" in name:
        return "THCS"

    if "XMC" in name:
        return "XÓA MÙ CHỮ"

    if "PCGD_2025_TH_" in name or "_TH_" in name:
        return "TIỂU HỌC"

    if "MN" in name or "PCGDMN" in name:
        return "MẦM NON"

    return "CHƯA PHÂN NHÓM"


def canonical_xlsx_files() -> list[Path]:
    found: list[Path] = []

    for path in CANONICAL_FILES:
        if path.exists():
            found.append(path)

    for folder in CANONICAL_DIRS:
        if not folder.exists():
            continue

        found.extend(
            sorted(
                path
                for path in folder.glob("*.xlsx")
                if path.is_file()
            )
        )

    seen = set()
    result = []

    for path in found:
        resolved = str(path.resolve()).lower()

        if resolved in seen:
            continue

        seen.add(resolved)
        result.append(path)

    return result


def nearest_label(
    ws,
    row: int,
    col: int,
) -> str:
    for delta in range(1, 9):
        check_col = col - delta

        if check_col < 1:
            break

        value = ws.cell(
            row=row,
            column=check_col,
        ).value

        if (
            value is not None
            and not str(value).startswith("=")
            and str(value).strip()
        ):
            return str(value).strip()

    for delta in range(1, 6):
        check_row = row - delta

        if check_row < 1:
            break

        value = ws.cell(
            row=check_row,
            column=col,
        ).value

        if (
            value is not None
            and not str(value).startswith("=")
            and str(value).strip()
        ):
            return str(value).strip()

    for rr in range(
        max(1, row - 2),
        row + 1,
    ):
        for cc in range(
            max(1, col - 5),
            col,
        ):
            value = ws.cell(
                row=rr,
                column=cc,
            ).value

            if (
                value is not None
                and not str(value).startswith("=")
                and str(value).strip()
            ):
                return str(value).strip()

    return ""


def classify_formula(
    formula: str,
    number_format: str,
    label: str,
) -> str:
    f = norm(formula)
    fmt = norm(number_format)
    lab = norm(label)

    if "%" in fmt:
        return "PERCENT"

    if any(
        word in lab
        for word in PERCENT_WORDS
    ):
        return "PERCENT"

    compact = f.replace(" ", "")

    if (
        f.endswith("%")
        or "*100" in compact
        or "*100.0" in compact
    ):
        return "PERCENT"

    if any(
        word in lab
        for word in RATIO_WORDS
    ):
        return "RATIO"

    if "/" in f:
        return "DIVISION_CẦN_XÁC_MINH"

    return "OTHER"


def raw_fraction(
    formula: str,
) -> tuple[str, str]:
    text = str(formula or "").strip()

    if text.startswith("="):
        text = text[1:]

    text = text.rstrip("%").strip()

    depth = 0
    slash_index = -1

    for idx, char in enumerate(text):
        if char == "(":
            depth += 1
        elif char == ")":
            depth = max(0, depth - 1)
        elif char == "/" and depth == 0:
            if slash_index >= 0:
                return (
                    "CÔNG_THỨC_PHỨC_HỢP",
                    "CÔNG_THỨC_PHỨC_HỢP",
                )

            slash_index = idx

    if slash_index < 0:
        return (
            "",
            "",
        )

    numerator = text[
        :slash_index
    ].strip()

    denominator = text[
        slash_index + 1:
    ].strip()

    return (
        numerator,
        denominator,
    )


def load_sources() -> dict[str, list[str]]:
    result: dict[str, list[str]] = {}

    for path in SOURCE_FILES:
        if not path.exists():
            continue

        try:
            text = path.read_text(
                encoding="utf-8-sig",
                errors="strict",
            )
        except UnicodeDecodeError:
            text = path.read_text(
                encoding="utf-8",
                errors="replace",
            )

        result[
            str(path.relative_to(PROJECT))
        ] = text.splitlines()

    return result


def source_matches_for_cell(
    sources: dict[str, list[str]],
    cell_ref: str,
) -> list[str]:
    matches = []

    q1 = f'"{cell_ref}"'
    q2 = f"'{cell_ref}'"

    for rel, lines in sources.items():
        for number, line in enumerate(
            lines,
            start=1,
        ):
            if q1 in line or q2 in line:
                matches.append(
                    f"{rel}:{number}: {line.strip()}"
                )

    return matches[:8]


def source_percent_helpers(
    sources: dict[str, list[str]],
) -> list[str]:
    rows = []

    helper_names = (
        "_percent",
        "_percentage",
        "_safe_percent",
        "_pct",
        "_ratio",
        "_safe_ratio",
        "_th_m1_percent",
    )

    for rel, lines in sources.items():
        for number, line in enumerate(
            lines,
            start=1,
        ):
            stripped = line.strip()

            if not (
                stripped.startswith("def ")
                or stripped.startswith("async def ")
            ):
                continue

            if any(
                f"{name}(" in stripped
                for name in helper_names
            ):
                rows.append(
                    f"{rel}:{number}: {stripped}"
                )

    return rows


def collect_matrix():
    sources = load_sources()
    rows = []

    for path in canonical_xlsx_files():
        group = group_from_path(
            path
        )

        try:
            wb = load_workbook(
                path,
                data_only=False,
                read_only=False,
            )
        except Exception as exc:
            rows.append(
                {
                    "group": group,
                    "file": str(
                        path.relative_to(
                            PROJECT
                        )
                    ),
                    "sheet": "",
                    "cell": "",
                    "label": "",
                    "formula": f"LỖI_MỞ_FILE: {exc}",
                    "kind": "ERROR",
                    "numerator": "",
                    "denominator": "",
                    "number_format": "",
                    "python_matches": [],
                    "status": "LỖI",
                }
            )
            continue

        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        value = cell.value

                        if (
                            not isinstance(
                                value,
                                str,
                            )
                            or not value.startswith("=")
                        ):
                            continue

                        if "/" not in value:
                            continue

                        label = nearest_label(
                            ws,
                            cell.row,
                            cell.column,
                        )

                        kind = classify_formula(
                            value,
                            cell.number_format,
                            label,
                        )

                        numerator, denominator = raw_fraction(
                            value
                        )

                        python_matches = source_matches_for_cell(
                            sources,
                            cell.coordinate,
                        )

                        if python_matches:
                            status = "PYTHON_CÓ_GHI_Ô_ĐÍCH"
                        elif kind in (
                            "PERCENT",
                            "RATIO",
                        ):
                            status = "MẪU_CÓ_CT_CHƯA_THẤY_PYTHON_GHI_Ô"
                        else:
                            status = "CẦN_XÁC_MINH_CÓ_PHẢI_TỶ_LỆ"

                        rows.append(
                            {
                                "group": group,
                                "file": str(
                                    path.relative_to(
                                        PROJECT
                                    )
                                ),
                                "sheet": ws.title,
                                "cell": cell.coordinate,
                                "label": label,
                                "formula": value,
                                "kind": kind,
                                "numerator": numerator,
                                "denominator": denominator,
                                "number_format": cell.number_format,
                                "python_matches": python_matches,
                                "status": status,
                            }
                        )
        finally:
            wb.close()

    return (
        rows,
        sources,
    )


def main() -> None:
    print("=" * 138)
    print(
        "BÀI 13B-11.15.2.4 - "
        "CHUẨN HÓA MA TRẬN CÔNG THỨC TỶ LỆ "
        "MẦM NON / TIỂU HỌC / THCS / XÓA MÙ CHỮ"
    )
    print("=" * 138)
    print()
    print("MỤC TIÊU:")
    print(
        " - Chỉ dùng file mẫu chính thức trong app/report_templates."
    )
    print(
        " - Loại nhiễu từ exports/payload/backup."
    )
    print(
        " - Liệt kê mọi công thức có phép chia trong 4 nhóm."
    )
    print(
        " - Phân loại PERCENT / RATIO / CẦN XÁC MINH."
    )
    print(
        " - Tách tử số/mẫu số khi công thức đủ đơn giản."
    )
    print(
        " - Đối chiếu Python hiện tại có ghi trực tiếp ô kết quả hay chưa."
    )
    print()
    print("AN TOÀN:")
    print(
        " - CHỈ ĐỌC source và Excel."
    )
    print(
        " - KHÔNG sửa source."
    )
    print(
        " - KHÔNG sửa database."
    )
    print(
        " - KHÔNG sửa mẫu Excel."
    )
    print()

    if not TEMPLATE_ROOT.exists():
        raise SystemExit(
            f"Không tìm thấy: {TEMPLATE_ROOT}"
        )

    files = canonical_xlsx_files()

    if not files:
        raise SystemExit(
            "Không tìm thấy file mẫu Excel chính thức."
        )

    print(
        "Số file mẫu chính thức:",
        len(files),
    )

    rows, sources = collect_matrix()

    by_group = defaultdict(
        list
    )

    for row in rows:
        by_group[
            row["group"]
        ].append(
            row
        )

    lines = []

    lines.append(
        "=" * 138
        + "\n"
    )
    lines.append(
        "BÀI 13B-11.15.2.4 - MA TRẬN CÔNG THỨC TỶ LỆ 4 NHÓM\n"
    )
    lines.append(
        "=" * 138
        + "\n"
    )
    lines.append(
        f"Project: {PROJECT}\n"
    )
    lines.append(
        f"Thời điểm: {datetime.now():%d/%m/%Y %H:%M:%S}\n"
    )
    lines.append(
        f"File mẫu chính thức đã đọc: {len(files)}\n"
    )
    lines.append(
        f"Source Python đã đọc: {len(sources)}\n\n"
    )

    lines.append(
        "NGUYÊN TẮC CHUẨN HÓA\n"
    )
    lines.append(
        "-" * 138
        + "\n"
    )
    lines.append(
        "1. Tỷ lệ % và tỷ số là dữ liệu tự tính, không nhập tay.\n"
    )
    lines.append(
        "2. Công thức Excel chính thức là chuẩn đối chiếu đầu tiên.\n"
    )
    lines.append(
        "3. Mẫu số = 0 hoặc chưa đủ dữ liệu: phần mềm phải để trống.\n"
    )
    lines.append(
        "4. Python sau này phải tính trực tiếp; không phụ thuộc Excel recalculation.\n"
    )
    lines.append(
        "5. Đạt/Không đạt là lớp logic riêng, không nhập tay.\n"
    )
    lines.append(
        "6. Công thức phức hợp chưa đủ ngữ nghĩa sẽ đánh dấu CẦN XÁC MINH.\n\n"
    )

    lines.append(
        "HÀM PYTHON TỶ LỆ/TỶ SỐ ĐANG CÓ\n"
    )
    lines.append(
        "-" * 138
        + "\n"
    )

    helpers = source_percent_helpers(
        sources
    )

    if helpers:
        for item in helpers:
            lines.append(
                f" - {item}\n"
            )
    else:
        lines.append(
            " - Chưa phát hiện helper tỷ lệ/tỷ số.\n"
        )

    for group in (
        "MẦM NON",
        "TIỂU HỌC",
        "THCS",
        "XÓA MÙ CHỮ",
        "CHƯA PHÂN NHÓM",
    ):
        group_rows = by_group.get(
            group,
            [],
        )

        if not group_rows:
            continue

        lines.append(
            "\n\n"
            + "=" * 138
            + "\n"
        )
        lines.append(
            f"NHÓM: {group}\n"
        )
        lines.append(
            "=" * 138
            + "\n"
        )

        kinds = Counter(
            row["kind"]
            for row in group_rows
        )

        lines.append(
            f"Tổng công thức có phép chia: {len(group_rows)}\n"
        )

        for key in (
            "PERCENT",
            "RATIO",
            "DIVISION_CẦN_XÁC_MINH",
            "ERROR",
        ):
            if kinds.get(
                key
            ):
                lines.append(
                    f" - {key}: {kinds[key]}\n"
                )

        lines.append(
            "\nMA TRẬN CHI TIẾT\n"
        )
        lines.append(
            "-" * 138
            + "\n"
        )

        for index, row in enumerate(
            group_rows,
            start=1,
        ):
            lines.append(
                f"\n[{index}] {row['file']} | "
                f"{row['sheet']}!{row['cell']}\n"
            )
            lines.append(
                f"    Nhãn gần nhất : {row['label'] or '(không xác định)'}\n"
            )
            lines.append(
                f"    Loại           : {row['kind']}\n"
            )
            lines.append(
                f"    Công thức Excel: {row['formula']}\n"
            )
            lines.append(
                f"    Tử số          : {row['numerator'] or '(chưa tách)'}\n"
            )
            lines.append(
                f"    Mẫu số         : {row['denominator'] or '(chưa tách)'}\n"
            )
            lines.append(
                f"    Number format  : {row['number_format']}\n"
            )
            lines.append(
                f"    Trạng thái     : {row['status']}\n"
            )

            if row[
                "python_matches"
            ]:
                lines.append(
                    "    Dấu vết Python :\n"
                )

                for match in row[
                    "python_matches"
                ]:
                    lines.append(
                        f"      - {match}\n"
                    )
            else:
                lines.append(
                    "    Dấu vết Python : chưa thấy ghi trực tiếp đúng ô này\n"
                )

    lines.append(
        "\n\n"
        + "=" * 138
        + "\n"
    )
    lines.append(
        "TỔNG KẾT ĐỂ LÀM BÀI 15.2.4.1\n"
    )
    lines.append(
        "=" * 138
        + "\n"
    )
    lines.append(
        "Chỉ những dòng PERCENT/RATIO đã xác định rõ tử số và mẫu số mới được đưa vào Python tính trực tiếp.\n"
    )
    lines.append(
        "Các dòng DIVISION_CẦN_XÁC_MINH phải đối chiếu nhãn/mẫu chính thức trước khi triển khai.\n"
    )
    lines.append(
        "Sau khi chốt ma trận, mới nối nguồn dữ liệu DB cho từng tử số/mẫu số và kiểm tra đủ dữ liệu đầu vào.\n"
    )

    OUT.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    OUT.write_text(
        "".join(
            lines
        ),
        encoding="utf-8",
    )

    print()
    print("TÓM TẮT:")

    for group in (
        "MẦM NON",
        "TIỂU HỌC",
        "THCS",
        "XÓA MÙ CHỮ",
    ):
        total = len(
            by_group.get(
                group,
                [],
            )
        )

        print(
            f" - {group}: {total} công thức có phép chia"
        )

    print()
    print(
        "Source/database/Excel: KHÔNG THAY ĐỔI."
    )
    print(
        "Ma trận:",
        OUT,
    )
    print()
    print("=" * 138)
    print(
        "BÀI 13B-11.15.2.4 TẠO MA TRẬN THÀNH CÔNG"
    )
    print("=" * 138)


if __name__ == "__main__":
    main()
