from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
root=Path(r'C:\\PhoCap\\app\\report_templates\\pcgd_mn_2025')
files=['PCGD_2025_MN_M1.xlsx','PCGD_2025_MN_02.xlsx','PCGD_2025_MN_01_GV.xlsx','PCGD_2025_MN_01_CSVC.xlsx','PCGD_2025_MN_TAICHINH.xlsx']
for name in files:
    p=root/name
    wb=load_workbook(p)
    out=BytesIO(); wb.save(out)
    assert len(out.getvalue())>1000
    print('DAT:',name,'->',wb.sheetnames)
print('5 MAU XLSX: DAT')
