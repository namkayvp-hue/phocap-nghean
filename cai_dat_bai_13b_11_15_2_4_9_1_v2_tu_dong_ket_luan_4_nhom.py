# -*- coding: utf-8 -*-
# =============================================================================
# BÀI 13B-11.15.2.4.9.1
# TỰ ĐỘNG KẾT LUẬN ĐẠT CHUẨN / MỨC ĐỘ - 4 NHÓM
# Mầm non / Tiểu học / Xóa mù chữ / THCS
#
# PHẠM VI:
# - Chỉ tự kết luận ở cấp XÃ/PHƯỜNG.
# - Không dùng evaluator cấp xã để kết luận Toàn tỉnh.
# - Không tự kết luận ở phạm vi Trường.
# - Cờ pháp lý chọn ngưỡng:
#       communes.is_special_difficulty_area
#
# AN TOÀN:
# - Backup source trước khi sửa.
# - Không UPDATE database.
# - Nếu lỗi sau khi sửa source: tự rollback source.
#
# SOURCE SỬA:
#   app/pcgd_mn_report_builders_v1.py
#   app/pcgd_xmc_report_builders_v1.py
#
# Ô KẾT LUẬN:
# - MN-02       : R7
# - TH-02       : T8
# - THCS-M1     : G36, G37
# - THCS-M2     : W14
# - THCS-TK     : F8, G8, R8
# - XMC-4       : R8
# =============================================================================

from __future__ import annotations

import ast
import os
import shutil
import sqlite3
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Iterable


PROJECT = Path(os.environ.get("PHOCAP_PROJECT", r"C:\PhoCap")).resolve()
APP = PROJECT / "app"
DB = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"

MN_BUILDER = APP / "pcgd_mn_report_builders_v1.py"
XMC_BUILDER = APP / "pcgd_xmc_report_builders_v1.py"
RULES = APP / "services" / "pcgd_business_rules.py"

MN_TEMPLATE = APP / "report_templates" / "pcgd_mn_2025" / "PCGD_2025_MN_02.xlsx"
TH_TEMPLATE = APP / "report_templates" / "pcgd_xmc_2025" / "PCGD_2025_TH_02.xlsx"
THCS_M2_TEMPLATE = APP / "report_templates" / "pcgd_xmc_2025" / "PCGD_2025_THCS_M2.xlsx"
THCS_TK_TEMPLATE = APP / "report_templates" / "pcgd_xmc_2025" / "PCGD_2025_THCS_TK.xlsx"
XMC4_TEMPLATE = APP / "report_templates" / "pcgd_xmc_2025" / "PCGD_2025_XMC_4.xlsx"

EXPECTED_COMMUNES = 130
EXPECTED_SPECIAL = 51

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
BACKUP = EXPORTS / f"backup_source_bai_13b_11_15_2_4_9_1_{STAMP}"
REPORT = EXPORTS / f"bao_cao_cai_dat_bai_13b_11_15_2_4_9_1_{STAMP}.txt"

MN_HELPER_START = "# === BAI_13B_11_15_2_4_9_1_MN_HELPERS_START ==="
MN_HELPER_END = "# === BAI_13B_11_15_2_4_9_1_MN_HELPERS_END ==="
MN_CALL_START = "# === BAI_13B_11_15_2_4_9_1_MN_CALL_START ==="
MN_CALL_END = "# === BAI_13B_11_15_2_4_9_1_MN_CALL_END ==="

XMC_HELPER_START = "# === BAI_13B_11_15_2_4_9_1_XMC_HELPERS_START ==="
XMC_HELPER_END = "# === BAI_13B_11_15_2_4_9_1_XMC_HELPERS_END ==="
XMC_CALL_START = "# === BAI_13B_11_15_2_4_9_1_XMC_CALL_START ==="
XMC_CALL_END = "# === BAI_13B_11_15_2_4_9_1_XMC_CALL_END ==="


MN_HELPERS = r'''
# === BAI_13B_11_15_2_4_9_1_MN_HELPERS_START ===
def _b491_mn_special_difficult(db, commune_id: int | None) -> bool:
    if commune_id is None:
        return False

    from sqlalchemy import text as _b491_sql_text

    value = db.execute(
        _b491_sql_text(
            "SELECT is_special_difficulty_area "
            "FROM communes WHERE id=:commune_id"
        ),
        {"commune_id": int(commune_id)},
    ).scalar_one_or_none()

    try:
        return bool(int(value or 0))
    except (TypeError, ValueError):
        return bool(value)


def _b491_mn_tri_bool(value):
    if value is True or value == 1:
        return True
    if value is False or value == 0:
        return False

    text = str(value or "").strip().upper()
    if text in {"CO", "CÓ", "TRUE", "YES", "1"}:
        return True
    if text in {"KHONG", "KHÔNG", "FALSE", "NO", "0"}:
        return False
    return None


def _b491_mn_completion_flag(record, age: int):
    if record is None:
        return None

    values = []
    if int(age) == 5:
        values.append(getattr(record, "completed_preschool_5", None))

    values.append(getattr(record, "completed_preschool_by_age", None))

    for raw in values:
        parsed = _b491_mn_tri_bool(raw)
        if parsed is not None:
            return parsed

    return None


def _b491_mn_age(person, year: int):
    """
    Helper tuổi riêng của Bài 4.9.1 cho Mầm non.

    pcgd_mn_report_builders_v1.py không định nghĩa _age.
    Tính tuổi theo năm tham chiếu của builder:
        tuổi = năm tham chiếu - năm sinh.

    Không ghi database.
    """
    birth = getattr(person, "date_of_birth", None)
    birth_year = getattr(birth, "year", None)

    if birth_year is None:
        return None

    try:
        age = int(year) - int(birth_year)
    except (TypeError, ValueError):
        return None

    if age < 0 or age > 120:
        return None

    return age


def _b491_mn_completion_rates_by_age(people, year: int):
    metrics = _metrics(people, year)
    result = {}

    for age in (3, 4, 5):
        attending_expected = int(
            metrics.get(age, {}).get("attending", 0) or 0
        )

        if attending_expected <= 0:
            result[age] = None
            continue

        flags = []

        for person, record in people:
            if _b491_mn_age(person, year) != age:
                continue

            status = str(
                getattr(record, "learning_status", "") or ""
            ).strip().upper().replace(" ", "_").replace("-", "_")

            if status in {
                "BO_HOC",
                "THOI_HOC",
                "CHUA_DI_HOC",
                "KHONG_THUOC_DIEN",
            }:
                continue

            if record is None:
                continue

            flags.append(
                _b491_mn_completion_flag(record, age)
            )

        if len(flags) != attending_expected:
            result[age] = None
            continue

        if any(flag is None for flag in flags):
            result[age] = None
            continue

        completed = sum(1 for flag in flags if flag is True)
        result[age] = _percent(completed, attending_expected)

    return result


def _b491_mn_result_text(result) -> str:
    if bool(getattr(result, "passed", False)):
        return "Đạt"

    if tuple(getattr(result, "missing_fields", ()) or ()):
        return "Chưa đủ dữ liệu"

    return "Không đạt"


def _b491_apply_mn_conclusion(
    *,
    ws,
    report_type: str,
    db,
    selected_commune_id: int | None,
    selected_school_id: int | None,
    meta,
    people,
    year: int,
) -> None:
    if report_type != "PCGD_MN_02_2025":
        return

    ws["R7"] = None

    if (
        selected_commune_id is None
        or selected_school_id is not None
        or str(meta.get("kind") or "") != "commune"
    ):
        return

    from app.services.pcgd_business_rules import (
        evaluate_preschool_3_5_commune,
    )

    metrics_by_age = _metrics(people, year)

    ppc_3_5 = sum(
        int(metrics_by_age.get(age, {}).get("ppc", 0) or 0)
        for age in (3, 4, 5)
    )
    attending_3_5 = sum(
        int(metrics_by_age.get(age, {}).get("attending", 0) or 0)
        for age in (3, 4, 5)
    )

    metrics = {
        "attendance_rate_3_5": _percent(attending_3_5, ppc_3_5),
        "completion_rate_3_5_by_age": _b491_mn_completion_rates_by_age(
            people,
            year,
        ),
    }

    result = evaluate_preschool_3_5_commune(
        metrics,
        is_special_difficult=_b491_mn_special_difficult(
            db,
            selected_commune_id,
        ),
    )

    ws["R7"] = _b491_mn_result_text(result)
# === BAI_13B_11_15_2_4_9_1_MN_HELPERS_END ===
'''.strip()


MN_CALL = r'''
# === BAI_13B_11_15_2_4_9_1_MN_CALL_START ===
_b491_apply_mn_conclusion(
    ws=ws,
    report_type=report_type,
    db=db,
    selected_commune_id=selected_commune_id,
    selected_school_id=selected_school_id,
    meta=meta,
    people=people,
    year=year,
)
# === BAI_13B_11_15_2_4_9_1_MN_CALL_END ===
'''.strip()


XMC_HELPERS = r'''
# === BAI_13B_11_15_2_4_9_1_XMC_HELPERS_START ===
def _b491_special_difficult(db, commune_id: int | None) -> bool:
    if commune_id is None:
        return False

    from sqlalchemy import text as _b491_sql_text

    value = db.execute(
        _b491_sql_text(
            "SELECT is_special_difficulty_area "
            "FROM communes WHERE id=:commune_id"
        ),
        {"commune_id": int(commune_id)},
    ).scalar_one_or_none()

    try:
        return bool(int(value or 0))
    except (TypeError, ValueError):
        return bool(value)


def _b491_result_level_or_status(result):
    if bool(getattr(result, "passed", False)):
        level = getattr(result, "level", None)
        if level is not None:
            return int(level)

    if tuple(getattr(result, "missing_fields", ()) or ()):
        return "Chưa đủ dữ liệu"

    return "Không đạt"


def _b491_primary_metrics(people, year: int):
    age6 = [
        (person, record)
        for person, record in people
        if _age(person, year) == 6 and _is_ppc(person, record)
    ]

    age11 = [
        (person, record)
        for person, record in people
        if _age(person, year) == 11 and _is_ppc(person, record)
    ]

    age14 = [
        (person, record)
        for person, record in people
        if _age(person, year) == 14 and _is_ppc(person, record)
    ]

    grade1_age6 = sum(
        1 for _person, record in age6 if _grade(record) == 1
    )
    completed_age11 = sum(
        1
        for person, record in age11
        if _completed_primary(person, record)
    )
    completed_age14 = sum(
        1
        for person, record in age14
        if _completed_primary(person, record)
    )

    remaining_age11 = [
        (person, record)
        for person, record in age11
        if not _completed_primary(person, record)
    ]

    def _still_primary(record) -> bool:
        grade = _grade(record)
        status = str(
            getattr(record, "learning_status", "") or ""
        ).strip().upper().replace(" ", "_").replace("-", "_")

        return (
            grade in {1, 2, 3, 4, 5}
            and status not in {
                "BO_HOC",
                "THOI_HOC",
                "CHUA_DI_HOC",
                "KHONG_THUOC_DIEN",
            }
        )

    return {
        "age6_grade1_rate": _percent(grade1_age6, len(age6)),
        "age14_completed_rate": _percent(completed_age14, len(age14)),
        "age11_completed_rate": _percent(completed_age11, len(age11)),
        "age11_remaining_all_study_primary": all(
            _still_primary(record)
            for _person, record in remaining_age11
        ),
    }


def _b491_literacy_metrics(people, year: int):
    metrics = _xmc1471b_metrics(people, year)

    def _rate(ages, true_key: str, false_key: str):
        ages = list(ages)
        population = _xmc1471b_sum(metrics, ages, "total")

        if int(population or 0) <= 0:
            return None

        positive = _xmc1471b_sum(metrics, ages, true_key)
        negative = _xmc1471b_sum(metrics, ages, false_key)

        if int(positive or 0) + int(negative or 0) != int(population):
            return None

        return _percent(positive, population)

    return {
        "literacy_level1_rate_15_25": _rate(
            range(15, 26),
            "bc1_total",
            "mc1_total",
        ),
        "literacy_level1_rate_15_35": _rate(
            range(15, 36),
            "bc1_total",
            "mc1_total",
        ),
        "literacy_level2_rate_15_35": _rate(
            range(15, 36),
            "bc2_total",
            "mc2_total",
        ),
        "literacy_level2_rate_15_60": _rate(
            range(15, 61),
            "bc2_total",
            "mc2_total",
        ),
    }


def _b491_thcs_metrics(people, year: int):
    age15_18 = [
        (person, record)
        for person, record in people
        if (_age(person, year) or -1) in range(15, 19)
        and _is_ppc(person, record)
    ]

    graduated = sum(
        1
        for person, record in age15_18
        if _completed_secondary(person, record)
    )
    upper = sum(
        1
        for person, record in age15_18
        if _upper_secondary(person, record)
    )

    return {
        "age15_18_graduated_thcs_rate": _percent(
            graduated,
            len(age15_18),
        ),
        "age15_18_upper_program_rate": _percent(
            upper,
            len(age15_18),
        ),
    }


def _b491_evaluate_all(*, db, commune_id: int, people, year: int):
    from app.services.pcgd_business_rules import (
        evaluate_literacy_commune,
        evaluate_primary_commune,
        evaluate_thcs_commune,
    )

    special = _b491_special_difficult(db, commune_id)

    primary = evaluate_primary_commune(
        _b491_primary_metrics(people, year),
        is_special_difficult=special,
    )

    literacy = evaluate_literacy_commune(
        _b491_literacy_metrics(people, year),
        is_special_difficult=special,
    )

    primary_level = (
        int(primary.level)
        if bool(primary.passed) and primary.level is not None
        else None
    )
    literacy_level = (
        int(literacy.level)
        if bool(literacy.passed) and literacy.level is not None
        else None
    )

    primary_missing = tuple(
        getattr(primary, "missing_fields", ()) or ()
    )
    literacy_missing = tuple(
        getattr(literacy, "missing_fields", ()) or ()
    )

    prereq_missing = bool(primary_missing or literacy_missing)
    prereq_failed = bool(
        (not bool(primary.passed) and not primary_missing)
        or (not bool(literacy.passed) and not literacy_missing)
    )

    if prereq_missing:
        thcs = None
        thcs_value = "Chưa đủ dữ liệu"
    elif prereq_failed:
        thcs = None
        thcs_value = "Không đạt"
    else:
        thcs = evaluate_thcs_commune(
            _b491_thcs_metrics(people, year),
            primary_level=primary_level,
            literacy_level=literacy_level,
            is_special_difficult=special,
        )
        thcs_value = _b491_result_level_or_status(thcs)

    return {
        "special": special,
        "primary": primary,
        "primary_value": _b491_result_level_or_status(primary),
        "literacy": literacy,
        "literacy_value": _b491_result_level_or_status(literacy),
        "thcs": thcs,
        "thcs_value": thcs_value,
    }


def _b491_clear_conclusion_cells(ws, report_type: str) -> None:
    mappings = {
        "PCGD_TH_02_2025": ("T8",),
        "PCGD_THCS_M1_2025": ("G36", "G37"),
        "PCGD_THCS_M2_2025": ("W14",),
        "PCGD_THCS_TK_2025": ("F8", "G8", "R8"),
        "PCGD_XMC_4_2025": ("R8",),
    }

    for ref in mappings.get(report_type, ()):
        ws[ref] = None


def _b491_apply_xmc_conclusions(
    *,
    ws,
    report_type: str,
    db,
    selected_commune_id: int | None,
    selected_school_id: int | None,
    meta,
    people,
    year: int,
) -> None:
    supported = {
        "PCGD_TH_02_2025",
        "PCGD_THCS_M1_2025",
        "PCGD_THCS_M2_2025",
        "PCGD_THCS_TK_2025",
        "PCGD_XMC_4_2025",
    }

    if report_type not in supported:
        return

    _b491_clear_conclusion_cells(ws, report_type)

    if (
        selected_commune_id is None
        or selected_school_id is not None
        or str(meta.get("kind") or "") != "commune"
    ):
        return

    result = _b491_evaluate_all(
        db=db,
        commune_id=int(selected_commune_id),
        people=people,
        year=year,
    )

    if report_type == "PCGD_TH_02_2025":
        ws["T8"] = result["primary_value"]

    elif report_type == "PCGD_THCS_M1_2025":
        ws["G36"] = result["primary_value"]
        ws["G37"] = result["literacy_value"]

    elif report_type == "PCGD_THCS_M2_2025":
        ws["W14"] = result["thcs_value"]

    elif report_type == "PCGD_THCS_TK_2025":
        ws["F8"] = result["primary_value"]
        ws["G8"] = result["literacy_value"]
        ws["R8"] = result["thcs_value"]

    elif report_type == "PCGD_XMC_4_2025":
        ws["R8"] = result["literacy_value"]
# === BAI_13B_11_15_2_4_9_1_XMC_HELPERS_END ===
'''.strip()


XMC_CALL = r'''
# === BAI_13B_11_15_2_4_9_1_XMC_CALL_START ===
_b491_apply_xmc_conclusions(
    ws=ws,
    report_type=report_type,
    db=db,
    selected_commune_id=selected_commune_id,
    selected_school_id=selected_school_id,
    meta=meta,
    people=people,
    year=year,
)
# === BAI_13B_11_15_2_4_9_1_XMC_CALL_END ===
'''.strip()


def read_text(path: Path) -> str:
    if not path.exists():
        raise RuntimeError(f"Không tìm thấy file bắt buộc: {path}")

    return path.read_text(
        encoding="utf-8-sig",
        errors="strict",
    )


def write_text(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8")


def function_node(source: str, name: str):
    tree = ast.parse(source)

    matches = [
        node
        for node in tree.body
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
        and node.name == name
    ]

    if len(matches) != 1:
        raise RuntimeError(
            f"Phải tìm thấy đúng 1 hàm {name}; hiện có {len(matches)}."
        )

    return matches[0]


def function_span(source: str, name: str) -> tuple[int, int]:
    node = function_node(source, name)

    lines = source.splitlines(keepends=True)
    offsets = [0]

    for line in lines:
        offsets.append(offsets[-1] + len(line))

    return (
        offsets[int(node.lineno) - 1],
        offsets[int(node.end_lineno)],
    )


def get_function(source: str, name: str) -> str:
    start, end = function_span(source, name)
    return source[start:end]


def replace_function(
    source: str,
    name: str,
    replacement: str,
) -> str:
    start, end = function_span(source, name)

    result = (
        source[:start]
        + replacement.rstrip()
        + "\n\n"
        + source[end:].lstrip("\n")
    )

    ast.parse(result)
    return result


def remove_marker_block(
    source: str,
    start_marker: str,
    end_marker: str,
) -> str:
    while start_marker in source:
        start = source.find(start_marker)
        end = source.find(end_marker, start)

        if end < 0:
            raise RuntimeError(
                "Có marker đầu nhưng thiếu marker cuối: "
                + start_marker
            )

        end += len(end_marker)
        line_start = source.rfind("\n", 0, start) + 1
        line_end = source.find("\n", end)

        if line_end < 0:
            line_end = len(source)
        else:
            line_end += 1

        source = source[:line_start] + source[line_end:]

    return source


def indent_block(block: str, indent: str) -> str:
    return "\n".join(
        indent + line if line.strip() else ""
        for line in block.splitlines()
    )


def inject_before_save(
    source: str,
    function_name: str,
    call_block: str,
) -> str:
    fn = get_function(source, function_name)
    lines = fn.splitlines(keepends=True)

    matches = [
        index
        for index, line in enumerate(lines)
        if "workbook.save(" in line
    ]

    if len(matches) != 1:
        raise RuntimeError(
            f"Hàm {function_name}: cần đúng 1 dòng workbook.save(...), "
            f"hiện có {len(matches)}."
        )

    index = matches[0]
    save_line = lines[index]
    indent = save_line[: len(save_line) - len(save_line.lstrip())]

    lines.insert(
        index,
        indent_block(call_block, indent) + "\n",
    )

    return replace_function(
        source,
        function_name,
        "".join(lines),
    )


def patch_mn_source(source: str) -> str:
    source = remove_marker_block(
        source,
        MN_HELPER_START,
        MN_HELPER_END,
    )
    source = remove_marker_block(
        source,
        MN_CALL_START,
        MN_CALL_END,
    )

    for name in (
        "_metrics",
        "_percent",
        "export_mn_report",
    ):
        function_node(source, name)

    export_fn = get_function(source, "export_mn_report")

    for token in (
        "report_type",
        "selected_commune_id",
        "selected_school_id",
        "meta",
        "people",
        "year",
        "workbook.save(",
    ):
        if token not in export_fn:
            raise RuntimeError(
                "export_mn_report thiếu nền cần thiết: " + token
            )

    source = inject_before_save(
        source,
        "export_mn_report",
        MN_CALL,
    )

    source = source.rstrip() + "\n\n\n" + MN_HELPERS + "\n"
    ast.parse(source)
    return source


def patch_xmc_source(source: str) -> str:
    source = remove_marker_block(
        source,
        XMC_HELPER_START,
        XMC_HELPER_END,
    )
    source = remove_marker_block(
        source,
        XMC_CALL_START,
        XMC_CALL_END,
    )

    for name in (
        "_age",
        "_is_ppc",
        "_grade",
        "_completed_primary",
        "_completed_secondary",
        "_upper_secondary",
        "_percent",
        "_xmc1471b_metrics",
        "_xmc1471b_sum",
        "_build_th02",
        "_build_thcs_m1",
        "_build_thcs_m2",
        "_build_thcs_tk",
        "_build_xmc4",
        "export_additional_report",
    ):
        function_node(source, name)

    export_fn = get_function(
        source,
        "export_additional_report",
    )

    for token in (
        "report_type",
        "selected_commune_id",
        "selected_school_id",
        "meta",
        "people",
        "year",
        "workbook.save(",
    ):
        if token not in export_fn:
            raise RuntimeError(
                "export_additional_report thiếu nền cần thiết: "
                + token
            )

    for report_code in (
        "PCGD_TH_02_2025",
        "PCGD_THCS_M1_2025",
        "PCGD_THCS_M2_2025",
        "PCGD_THCS_TK_2025",
        "PCGD_XMC_4_2025",
    ):
        if report_code not in export_fn:
            raise RuntimeError(
                "export_additional_report chưa có nhánh: "
                + report_code
            )

    source = inject_before_save(
        source,
        "export_additional_report",
        XMC_CALL,
    )

    source = source.rstrip() + "\n\n\n" + XMC_HELPERS + "\n"
    ast.parse(source)
    return source


def backup_sources() -> None:
    BACKUP.mkdir(parents=True, exist_ok=False)

    for path in (MN_BUILDER, XMC_BUILDER):
        dst = BACKUP / path.relative_to(PROJECT)
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, dst)


def restore_sources() -> None:
    for target in (MN_BUILDER, XMC_BUILDER):
        src = BACKUP / target.relative_to(PROJECT)
        if src.exists():
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, target)


def db_state() -> dict:
    if not DB.exists():
        raise RuntimeError(f"Không tìm thấy database: {DB}")

    conn = sqlite3.connect(str(DB))

    try:
        integrity = str(
            conn.execute(
                "PRAGMA integrity_check"
            ).fetchone()[0]
        )

        fk = len(
            conn.execute(
                "PRAGMA foreign_key_check"
            ).fetchall()
        )

        columns = [
            str(row[1])
            for row in conn.execute(
                "PRAGMA table_info(communes)"
            ).fetchall()
        ]

        if "is_special_difficulty_area" not in columns:
            raise RuntimeError(
                "Bảng communes chưa có cột "
                "is_special_difficulty_area."
            )

        total = int(
            conn.execute(
                "SELECT COUNT(*) FROM communes"
            ).fetchone()[0]
        )

        special = int(
            conn.execute(
                "SELECT COUNT(*) "
                "FROM communes "
                "WHERE COALESCE(is_special_difficulty_area, 0)=1"
            ).fetchone()[0]
        )

        return {
            "integrity": integrity,
            "fk": fk,
            "total": total,
            "special": special,
            "columns": columns,
        }

    finally:
        conn.close()


def verify_rules() -> None:
    source = read_text(RULES)
    tree = ast.parse(source)

    functions = {
        node.name
        for node in tree.body
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
    }

    for name in (
        "evaluate_preschool_3_5_commune",
        "evaluate_primary_commune",
        "evaluate_thcs_commune",
        "evaluate_literacy_commune",
    ):
        if name not in functions:
            raise RuntimeError(
                "Bộ quy tắc nghiệp vụ thiếu evaluator: " + name
            )

    if "is_special_difficult" not in source:
        raise RuntimeError(
            "Evaluator hiện tại không có tham số "
            "is_special_difficult."
        )


def verify_template_cell(
    path: Path,
    cell_ref: str,
    expected_words: Iterable[str],
) -> None:
    if not path.exists():
        raise RuntimeError(f"Không tìm thấy template: {path}")

    from openpyxl import load_workbook

    workbook = load_workbook(
        path,
        read_only=False,
        data_only=False,
    )

    try:
        ws = workbook[workbook.sheetnames[0]]
        value = str(ws[cell_ref].value or "").lower()

        if not all(word.lower() in value for word in expected_words):
            raise RuntimeError(
                f"Template {path.name}: ô {cell_ref}="
                f"'{ws[cell_ref].value}' không đúng nền dự kiến."
            )
    finally:
        workbook.close()


def verify_templates() -> None:
    verify_template_cell(MN_TEMPLATE, "R3", ("đạt",))
    verify_template_cell(TH_TEMPLATE, "T4", ("đạt",))
    verify_template_cell(THCS_M2_TEMPLATE, "W9", ("đạt",))
    verify_template_cell(THCS_TK_TEMPLATE, "F4", ("đạt",))
    verify_template_cell(THCS_TK_TEMPLATE, "G4", ("đạt",))
    verify_template_cell(THCS_TK_TEMPLATE, "R4", ("đạt",))
    verify_template_cell(XMC4_TEMPLATE, "R5", ("đạt",))


def verify_patched_sources() -> None:
    mn = read_text(MN_BUILDER)
    xmc = read_text(XMC_BUILDER)

    ast.parse(mn)
    ast.parse(xmc)

    for token in (
        MN_HELPER_START,
        MN_HELPER_END,
        MN_CALL_START,
        MN_CALL_END,
        "def _b491_apply_mn_conclusion(",
        'ws["R7"]',
        "evaluate_preschool_3_5_commune",
        "is_special_difficulty_area",
    ):
        if token not in mn:
            raise RuntimeError("Verifier MN thiếu: " + token)

    for token in (
        XMC_HELPER_START,
        XMC_HELPER_END,
        XMC_CALL_START,
        XMC_CALL_END,
        "def _b491_apply_xmc_conclusions(",
        "def _b491_primary_metrics(",
        "def _b491_literacy_metrics(",
        "def _b491_thcs_metrics(",
        "evaluate_primary_commune",
        "evaluate_literacy_commune",
        "evaluate_thcs_commune",
        'ws["T8"]',
        'ws["G36"]',
        'ws["G37"]',
        'ws["W14"]',
        'ws["F8"]',
        'ws["G8"]',
        'ws["R8"]',
        "is_special_difficulty_area",
    ):
        if token not in xmc:
            raise RuntimeError(
                "Verifier TH/XMC/THCS thiếu: " + token
            )

    mn_export = get_function(mn, "export_mn_report")
    xmc_export = get_function(xmc, "export_additional_report")

    if mn_export.find(MN_CALL_START) >= mn_export.find("workbook.save("):
        raise RuntimeError(
            "Call kết luận MN phải chạy trước workbook.save."
        )

    if xmc_export.find(XMC_CALL_START) >= xmc_export.find("workbook.save("):
        raise RuntimeError(
            "Call kết luận TH/XMC/THCS phải chạy trước workbook.save."
        )

    for path in (MN_BUILDER, XMC_BUILDER, RULES):
        subprocess.run(
            [sys.executable, "-m", "py_compile", str(path)],
            cwd=str(PROJECT),
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )


def clear_cache() -> None:
    for cache in APP.rglob("__pycache__"):
        if cache.is_dir():
            shutil.rmtree(cache, ignore_errors=True)


def build_report(before: dict, after: dict) -> str:
    lines = [
        "=" * 96,
        "BÀI 13B-11.15.2.4.9.1 - BÁO CÁO CÀI ĐẶT",
        "=" * 96,
        "",
        "PHẠM VI:",
        " - Tự động kết luận XÃ/PHƯỜNG.",
        " - Không kết luận Toàn tỉnh bằng evaluator cấp xã.",
        " - Không kết luận phạm vi Trường.",
        "",
        "CỜ PHÁP LÝ:",
        " - communes.is_special_difficulty_area",
        f" - Tổng xã/phường: {after['total']}",
        f" - Xã đặc biệt khó khăn: {after['special']}",
        "",
        "SOURCE ĐÃ NỐI:",
        f" - {MN_BUILDER.relative_to(PROJECT)}",
        f" - {XMC_BUILDER.relative_to(PROJECT)}",
        "",
        "Ô KẾT LUẬN KHI XUẤT EXCEL:",
        " - MN-02: R7",
        " - TH-02: T8",
        " - THCS-M1: G36, G37",
        " - THCS-M2: W14",
        " - THCS-TK: F8, G8, R8",
        " - XMC-4: R8",
        "",
        "AN TOÀN:",
        f" - Backup source: {BACKUP}",
        " - Database: KHÔNG UPDATE.",
        f" - integrity_check trước/sau: {before['integrity']} / {after['integrity']}",
        f" - foreign_key_check trước/sau: {before['fk']} / {after['fk']}",
        f" - Cờ đặc biệt khó khăn trước/sau: {before['special']} / {after['special']}",
        "",
        "GHI CHÚ:",
        " - Thiếu chỉ tiêu bắt buộc => Chưa đủ dữ liệu.",
        " - Không suy đoán Đạt.",
        " - Bộ 7 sheet Bieu_mau_PCGDMN_2025.xlsx không bị sửa ở bài này.",
        "",
        "KẾT LUẬN: CÀI ĐẶT THÀNH CÔNG.",
    ]

    return "\n".join(lines)


def main() -> int:
    print("=" * 96)
    print(
        "BÀI 13B-11.15.2.4.9.1 - "
        "TỰ ĐỘNG KẾT LUẬN ĐẠT CHUẨN / MỨC ĐỘ - 4 NHÓM"
    )
    print("=" * 96)
    print()

    backup_created = False

    try:
        for path in (
            APP,
            DB,
            MN_BUILDER,
            XMC_BUILDER,
            RULES,
        ):
            if not path.exists():
                raise RuntimeError(f"Không tìm thấy: {path}")

        print("[1/8] Kiểm tra database...")
        before = db_state()

        if before["integrity"].lower() != "ok":
            raise RuntimeError("integrity_check != ok")

        if before["fk"] != 0:
            raise RuntimeError("foreign_key_check có lỗi.")

        if before["total"] != EXPECTED_COMMUNES:
            raise RuntimeError(
                f"Tổng communes={before['total']}, "
                f"không phải {EXPECTED_COMMUNES}."
            )

        if before["special"] != EXPECTED_SPECIAL:
            raise RuntimeError(
                "Cờ is_special_difficulty_area chưa đúng nền Bài 4.9.0.1: "
                f"TRUE={before['special']}, yêu cầu {EXPECTED_SPECIAL}."
            )

        print(
            "[OK] Database: "
            f"{before['total']} xã/phường; "
            f"đặc biệt khó khăn={before['special']}."
        )

        print("[2/8] Kiểm tra 4 evaluator...")
        verify_rules()
        print("[OK] Đủ 4 evaluator.")

        print("[3/8] Kiểm tra template...")
        verify_templates()
        print("[OK] Template đúng nền.")

        print("[4/8] Backup source...")
        backup_sources()
        backup_created = True
        print(f"[OK] {BACKUP}")

        print("[5/8] Nối kết luận Mầm non...")
        write_text(
            MN_BUILDER,
            patch_mn_source(read_text(MN_BUILDER)),
        )
        print("[OK] MN-02 -> R7.")

        print("[6/8] Nối kết luận TH / XMC / THCS...")
        write_text(
            XMC_BUILDER,
            patch_xmc_source(read_text(XMC_BUILDER)),
        )
        print(
            "[OK] TH-02 / THCS-M1 / THCS-M2 / "
            "THCS-TK / XMC-4."
        )

        print("[7/8] Kiểm tra source sau cài...")
        verify_patched_sources()
        print("[OK] AST + py_compile + marker.")

        print("[8/8] Kiểm tra database sau cài...")
        after = db_state()

        if after["integrity"].lower() != "ok":
            raise RuntimeError(
                "integrity_check sau cài != ok"
            )

        if after["fk"] != 0:
            raise RuntimeError(
                "foreign_key_check sau cài có lỗi."
            )

        if (
            after["total"] != before["total"]
            or after["special"] != before["special"]
        ):
            raise RuntimeError(
                "Database thay đổi ngoài dự kiến."
            )

        clear_cache()

        EXPORTS.mkdir(parents=True, exist_ok=True)
        REPORT.write_text(
            build_report(before, after),
            encoding="utf-8-sig",
        )

        print()
        print("=" * 96)
        print("BÀI 4.9.1 HOÀN THÀNH.")
        print("=" * 96)
        print("Database               : KHÔNG THAY ĐỔI")
        print(
            "Cờ đặc biệt khó khăn   : "
            f"{after['special']}/{after['total']}"
        )
        print("MN-02                  : R7 tự kết luận")
        print("TH-02                  : T8 tự xác định mức")
        print("XMC-4                  : R8 tự xác định mức")
        print("THCS                    : M1/M2/TK tự nối mức")
        print(f"Backup source           : {BACKUP}")
        print(f"Báo cáo cài đặt         : {REPORT}")
        print()
        print(
            "Thiếu chỉ tiêu bắt buộc -> "
            "Chưa đủ dữ liệu; không suy đoán Đạt."
        )
        print("=" * 96)

        return 0

    except Exception as exc:
        print()
        print("=" * 96)
        print("[LỖI] CÀI ĐẶT BÀI 4.9.1 KHÔNG HOÀN THÀNH")
        print(str(exc))
        print("=" * 96)

        if backup_created:
            try:
                restore_sources()
                clear_cache()
                print(
                    "[ROLLBACK] Đã khôi phục source từ backup."
                )
            except Exception as restore_exc:
                print(
                    "[CẢNH BÁO] Khôi phục source lỗi: "
                    + str(restore_exc)
                )

        print("Database không có lệnh UPDATE trong bộ cài này.")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
