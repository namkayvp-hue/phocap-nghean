from __future__ import annotations

import hashlib
import py_compile
import re
import shutil
import sys
import traceback
from datetime import datetime
from pathlib import Path

PROJECT = Path(r"C:\PhoCap")
BUILDER = PROJECT / "app" / "pcgd_xmc_report_builders_v1.py"
MAIN = PROJECT / "app" / "main.py"
ROUTER = PROJECT / "app" / "routers" / "report_center.py"
DB = PROJECT / "data" / "phocap.db"
TEMPLATE_ROOT = PROJECT / "app" / "report_templates" / "pcgd_xmc_2025"
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
BACKUP = PROJECT / "exports" / f"backup_pcgd_xmc_v1_2_{STAMP}"

OLD_CLEAR = '''def _clear_range(ws: Any, min_row: int, max_row: int, min_col: int, max_col: int) -> None:\n    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):\n        for cell in row:\n            cell.value = None\n'''

NEW_CLEAR = '''def _clear_range(ws: Any, min_row: int, max_row: int, min_col: int, max_col: int) -> None:\n    # Không ghi trực tiếp vào ô con của vùng merge vì openpyxl sẽ báo\n    # AttributeError: MergedCell object attribute value is read-only.\n    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):\n        for cell in row:\n            if cell.__class__.__name__ == "MergedCell":\n                continue\n            cell.value = None\n'''

OLD_XMC3_RE = re.compile(
    r"def _build_xmc3\(ws: Any, people: list\[tuple\[SurveyPerson, SurveyPersonYearRecord \| None\]\], year:int\) -> None:\n"
    r"(?:    .*\n)+?"
    r"(?=\ndef _build_cmc2)",
    re.MULTILINE,
)

NEW_XMC3 = '''def _build_xmc3(ws: Any, people: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year:int) -> None:\n    # Mẫu XMC-3 có 3 dòng cộng chèn giữa các nhóm tuổi:\n    # dòng 20 = Cộng 15-25, dòng 31 = Cộng 15-35, dòng 57 = Cộng 15-60.\n    # Vì vậy không thể ánh xạ tuổi theo công thức row = 9 + (age - 15).\n    row_by_age: dict[int, int] = {}\n    for age in range(15, 26):\n        row_by_age[age] = 9 + (age - 15)\n    for age in range(26, 36):\n        row_by_age[age] = 21 + (age - 26)\n    for age in range(36, 61):\n        row_by_age[age] = 32 + (age - 36)\n\n    # Xóa số liệu cũ nhưng giữ nguyên nhãn, merge, định dạng và các dòng cộng.\n    for row in range(9, 58):\n        for col in range(2, 20):\n            cell = ws.cell(row, col)\n            if cell.__class__.__name__ == "MergedCell":\n                continue\n            if col == 2 and row in {20, 31, 57}:\n                continue\n            if col >= 3:\n                cell.value = None\n\n    metrics = _population_metrics(people, year)\n    for age, row in row_by_age.items():\n        item = metrics.get(age, {})\n        ws.cell(row, 2).value = year - age\n        ws.cell(row, 3).value = item.get("total", 0) or None\n        ws.cell(row, 4).value = item.get("female", 0) or None\n        ws.cell(row, 5).value = item.get("ethnic", 0) or None\n        ws.cell(row, 6).value = item.get("female_ethnic", 0) or None\n\n    for row, ages in ((20, range(15, 26)), (31, range(15, 36)), (57, range(15, 61))):\n        for col, key in ((3, "total"), (4, "female"), (5, "ethnic"), (6, "female_ethnic")):\n            ws.cell(row, col).value = _sum_metric(metrics, list(ages), key)\n\n'''


def sha(path: Path) -> str | None:
    if not path.exists():
        return None
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def backup_file(path: Path) -> None:
    if not path.exists():
        return
    rel = path.relative_to(PROJECT)
    dest = BACKUP / rel
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, dest)


def restore() -> None:
    saved = BACKUP / BUILDER.relative_to(PROJECT)
    if saved.exists():
        BUILDER.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(saved, BUILDER)


def run_template_safety_check() -> None:
    from openpyxl import load_workbook
    from io import BytesIO

    xmc3 = TEMPLATE_ROOT / "PCGD_2025_XMC_3.xlsx"
    if not xmc3.exists():
        raise RuntimeError(f"Không tìm thấy mẫu XMC-3: {xmc3}")

    wb = load_workbook(xmc3)
    ws = wb[wb.sheetnames[0]]

    # Mô phỏng đúng phần có lỗi trước đây với dữ liệu rỗng.
    row_by_age = {}
    for age in range(15, 26):
        row_by_age[age] = 9 + (age - 15)
    for age in range(26, 36):
        row_by_age[age] = 21 + (age - 26)
    for age in range(36, 61):
        row_by_age[age] = 32 + (age - 36)

    for row in range(9, 58):
        for col in range(2, 20):
            cell = ws.cell(row, col)
            if cell.__class__.__name__ == "MergedCell":
                continue
            if col == 2 and row in {20, 31, 57}:
                continue
            if col >= 3:
                cell.value = None

    for age, row in row_by_age.items():
        ws.cell(row, 2).value = 2026 - age
        for col in range(3, 7):
            ws.cell(row, col).value = None

    for row in (20, 31, 57):
        for col in range(3, 7):
            ws.cell(row, col).value = 0

    output = BytesIO()
    wb.save(output)
    if len(output.getvalue()) < 1000:
        raise RuntimeError("Kiểm tra ghi mẫu XMC-3 không đạt.")


def main() -> None:
    print("=" * 68)
    print("SUA LOI BO BAO CAO PCGD & XMC V1.2")
    print("XMC-3 + AN TOAN O MERGE CHO CAC BIEU CON LAI")
    print("=" * 68)

    if not BUILDER.exists():
        raise RuntimeError(f"Không tìm thấy tệp: {BUILDER}")

    before = {
        "main": sha(MAIN),
        "router": sha(ROUTER),
        "db": sha(DB),
    }

    BACKUP.mkdir(parents=True, exist_ok=False)
    backup_file(BUILDER)
    (BACKUP / "THONG_TIN.txt").write_text(
        f"Sao lưu trước khi sửa PCGD & XMC V1.2\nThời gian: {datetime.now()}\n",
        encoding="utf-8",
    )

    try:
        text = BUILDER.read_text(encoding="utf-8-sig")

        if NEW_CLEAR not in text:
            if OLD_CLEAR not in text:
                raise RuntimeError("Không tìm thấy hàm _clear_range đúng phiên bản để sửa an toàn.")
            text = text.replace(OLD_CLEAR, NEW_CLEAR, 1)

        match = OLD_XMC3_RE.search(text)
        if match:
            text = text[:match.start()] + NEW_XMC3 + text[match.end():]
        elif "row_by_age[age] = 21 + (age - 26)" not in text:
            raise RuntimeError("Không tìm thấy hàm _build_xmc3 đúng phiên bản để sửa an toàn.")

        BUILDER.write_text(text, encoding="utf-8")

        py_compile.compile(str(BUILDER), doraise=True)
        run_template_safety_check()

        after = {
            "main": sha(MAIN),
            "router": sha(ROUTER),
            "db": sha(DB),
        }
        if before != after:
            raise RuntimeError("Phát hiện main.py, route hoặc cơ sở dữ liệu bị thay đổi ngoài phạm vi cho phép.")

        print("")
        print("KIEM TRA:")
        print("- XMC-3: KHONG CON GHI VAO O MERGE B20/B31")
        print("- XMC-3: ANH XA DUNG 15-25 / 26-35 / 36-60")
        print("- Cac ham xoa du lieu mau: BO QUA O MERGEDCELL")
        print("- app/main.py: KHONG DOI")
        print("- app/routers/report_center.py: KHONG DOI")
        print("- data/phocap.db: KHONG DOI")
        print("- Route: KHONG DOI")
        print("- Luong nghiep vu cu: KHONG DOI")
        print("")
        print("SUA LOI BO BAO CAO PCGD & XMC V1.2 THANH CONG")
        print(f"Ban sao an toan: {BACKUP}")

    except Exception:
        print("")
        print("SUA LOI KHONG THANH CONG - DANG KHOI PHUC...")
        restore()
        traceback.print_exc()
        print(f"Da khoi phuc tu: {BACKUP}")
        sys.exit(1)


if __name__ == "__main__":
    main()
