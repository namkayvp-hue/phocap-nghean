# -*- coding: utf-8 -*-
# =============================================================================
# BÀI 13B-11.15.2.4.9.2 V2
# TỰ ĐỘNG KẾT LUẬN / CÔNG NHẬN CẤP TỈNH - 4 NHÓM
#
# - Cấp xã/phường giữ nguyên Bài 4.9.1.
# - Toàn tỉnh: đánh giá từng xã/phường rồi tổng hợp tỷ lệ.
# - Không tự kết luận phạm vi Trường.
# - Backup + rollback source; không ghi database.
# =============================================================================

from __future__ import annotations

import ast
import hashlib
import os
import shutil
import sqlite3
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Iterable

PROJECT = Path(os.environ.get("PHOCAP_PROJECT", r"C:\PhoCap")).resolve()
APP = PROJECT / "app"
DB = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"

RULES = APP / "services" / "pcgd_business_rules.py"
MN_BUILDER = APP / "pcgd_mn_report_builders_v1.py"
XMC_BUILDER = APP / "pcgd_xmc_report_builders_v1.py"

MN_TEMPLATE = APP / "report_templates" / "pcgd_mn_2025" / "PCGD_2025_MN_02.xlsx"
TH_TEMPLATE = APP / "report_templates" / "pcgd_xmc_2025" / "PCGD_2025_TH_02.xlsx"
THCS_M2_TEMPLATE = APP / "report_templates" / "pcgd_xmc_2025" / "PCGD_2025_THCS_M2.xlsx"
THCS_TK_TEMPLATE = APP / "report_templates" / "pcgd_xmc_2025" / "PCGD_2025_THCS_TK.xlsx"
XMC4_TEMPLATE = APP / "report_templates" / "pcgd_xmc_2025" / "PCGD_2025_XMC_4.xlsx"

EXPECTED_COMMUNES = 130
EXPECTED_SPECIAL = 51

EXPECTED_SOURCE_SHA256 = {
    RULES: "6acd5cf321bdc6575b1e24eed5832cee97eda9b5a4301781ca4311d51de98c35",
    MN_BUILDER: "72d63041adf1435c44709663c7a105f543fbe4f486398c1f03ecd3bd6eb7b7f5",
    XMC_BUILDER: "696b2a5457284bed0077b2892e3def5fe8efca4b7ca8adb64d5f540ea03c0810",
}

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
BACKUP = EXPORTS / f"backup_source_bai_13b_11_15_2_4_9_2_{STAMP}"
REPORT = EXPORTS / f"bao_cao_cai_dat_bai_13b_11_15_2_4_9_2_{STAMP}.txt"

RULES_START = "# === BAI_13B_11_15_2_4_9_2_PROVINCE_RULES_START ==="
RULES_END = "# === BAI_13B_11_15_2_4_9_2_PROVINCE_RULES_END ==="
MN_HELPER_START = "# === BAI_13B_11_15_2_4_9_2_MN_HELPERS_START ==="
MN_HELPER_END = "# === BAI_13B_11_15_2_4_9_2_MN_HELPERS_END ==="
MN_CALL_START = "# === BAI_13B_11_15_2_4_9_2_MN_CALL_START ==="
MN_CALL_END = "# === BAI_13B_11_15_2_4_9_2_MN_CALL_END ==="
XMC_HELPER_START = "# === BAI_13B_11_15_2_4_9_2_XMC_HELPERS_START ==="
XMC_HELPER_END = "# === BAI_13B_11_15_2_4_9_2_XMC_HELPERS_END ==="
XMC_CALL_START = "# === BAI_13B_11_15_2_4_9_2_XMC_CALL_START ==="
XMC_CALL_END = "# === BAI_13B_11_15_2_4_9_2_XMC_CALL_END ==="

RULES_BLOCK = '# === BAI_13B_11_15_2_4_9_2_PROVINCE_RULES_START ===\n# Bài 13B-11.15.2.4.9.2 - quy tắc công nhận cấp tỉnh.\n# Cấp tỉnh đánh giá từng xã/phường rồi tổng hợp tỷ lệ, không dùng\n# dữ liệu toàn tỉnh làm đầu vào trực tiếp cho evaluator cấp xã.\n\nPROVINCE_RECOGNITION_ND142_2025 = {\n    "primary": {1: 90.0, 2: 90.0, 3: 90.0},\n    "thcs": {1: 90.0, 2: 95.0, 3: 100.0},\n    "literacy": {1: 81.0, 2: 90.0},\n}\n\n\ndef _b492_bool_status(value):\n    if value is True or value == 1:\n        return True\n    if value is False or value == 0:\n        return False\n    return None\n\n\ndef _b492_group_check(*, key: str, statuses, required: float):\n    values = [_b492_bool_status(v) for v in list(statuses or ())]\n    total = len(values)\n\n    if total <= 0:\n        return (\n            RuleCheck(key, None, float(required), None, "Không có đơn vị cấp xã."),\n            True,\n        )\n\n    recognized = sum(1 for v in values if v is True)\n    unknown = sum(1 for v in values if v is None)\n\n    min_rate = recognized * 100.0 / total\n    max_rate = (recognized + unknown) * 100.0 / total\n\n    if min_rate >= float(required):\n        passed = True\n        uncertain = False\n    elif max_rate < float(required):\n        passed = False\n        uncertain = False\n    else:\n        passed = None\n        uncertain = True\n\n    note = (\n        f"Đạt chắc chắn={recognized}/{total}; chưa đủ dữ liệu={unknown}; "\n        f"khoảng tỷ lệ={min_rate:.2f}%..{max_rate:.2f}%."\n    )\n\n    return (\n        RuleCheck(key, round(min_rate, 6), float(required), passed, note),\n        uncertain,\n    )\n\n\ndef evaluate_preschool_3_5_province(\n    *,\n    special_difficult_commune_statuses,\n    other_commune_statuses,\n) -> RecognitionResult:\n    special_required = float(\n        PRESCHOOL_3_5_ND277["province_recognition"][\n            "special_difficult_communes_recognized_min_rate"\n        ]\n    )\n    other_required = float(\n        PRESCHOOL_3_5_ND277["province_recognition"][\n            "other_communes_recognized_min_rate"\n        ]\n    )\n\n    special_check, special_uncertain = _b492_group_check(\n        key="special_difficult_communes_recognized_rate",\n        statuses=special_difficult_commune_statuses,\n        required=special_required,\n    )\n    other_check, other_uncertain = _b492_group_check(\n        key="other_communes_recognized_rate",\n        statuses=other_commune_statuses,\n        required=other_required,\n    )\n\n    checks = (special_check, other_check)\n    definitive_fail = any(c.passed is False for c in checks)\n\n    if definitive_fail:\n        return RecognitionResult(\n            program="PRESCHOOL_3_5_ND277_PROVINCE",\n            level=None,\n            passed=False,\n            checks=checks,\n            missing_fields=(),\n        )\n\n    missing = []\n    if special_uncertain:\n        missing.append("special_difficult_communes_status")\n    if other_uncertain:\n        missing.append("other_communes_status")\n\n    passed = all(c.passed is True for c in checks) and not missing\n\n    return RecognitionResult(\n        program="PRESCHOOL_3_5_ND277_PROVINCE",\n        level=1 if passed else None,\n        passed=passed,\n        checks=checks,\n        missing_fields=tuple(missing),\n    )\n\n\ndef _b492_normalize_level_item(item):\n    if isinstance(item, dict):\n        raw_level = item.get("level")\n        incomplete = bool(item.get("incomplete") or item.get("missing"))\n    elif isinstance(item, (tuple, list)) and len(item) >= 2:\n        raw_level = item[0]\n        incomplete = bool(item[1])\n    else:\n        raw_level = item\n        incomplete = item is None\n\n    if raw_level is None:\n        return None, bool(incomplete)\n\n    try:\n        level = int(raw_level)\n    except (TypeError, ValueError):\n        return None, True\n\n    if level < 0:\n        return None, True\n\n    return level, bool(incomplete)\n\n\ndef _b492_level_check(*, items, target_level: int, required: float, key: str):\n    normalized = [_b492_normalize_level_item(item) for item in list(items or ())]\n    total = len(normalized)\n\n    if total <= 0:\n        return (\n            RuleCheck(key, None, float(required), None, "Không có đơn vị cấp xã."),\n            True,\n        )\n\n    recognized = 0\n    unknown = 0\n\n    for confirmed_level, incomplete in normalized:\n        if confirmed_level is None:\n            unknown += 1\n        elif confirmed_level >= int(target_level):\n            recognized += 1\n        elif incomplete:\n            unknown += 1\n\n    min_rate = recognized * 100.0 / total\n    max_rate = (recognized + unknown) * 100.0 / total\n\n    if min_rate >= float(required):\n        passed = True\n        uncertain = False\n    elif max_rate < float(required):\n        passed = False\n        uncertain = False\n    else:\n        passed = None\n        uncertain = True\n\n    note = (\n        f"Đạt chắc chắn M{target_level}+={recognized}/{total}; "\n        f"chưa xác định={unknown}; khoảng tỷ lệ={min_rate:.2f}%..{max_rate:.2f}%."\n    )\n\n    return (\n        RuleCheck(key, round(min_rate, 6), float(required), passed, note),\n        uncertain,\n    )\n\n\ndef _b492_evaluate_level_province(*, program: str, commune_levels, thresholds):\n    items = list(commune_levels or ())\n    checks = []\n    confirmed_level = None\n\n    for target_level in sorted(int(x) for x in thresholds):\n        check, uncertain = _b492_level_check(\n            items=items,\n            target_level=target_level,\n            required=float(thresholds[target_level]),\n            key=f"communes_recognized_level_{target_level}_rate",\n        )\n        checks.append(check)\n\n        if check.passed is True:\n            confirmed_level = target_level\n            continue\n\n        if uncertain:\n            missing = (f"commune_levels_L{target_level}",)\n            if confirmed_level is not None:\n                return RecognitionResult(\n                    program=program,\n                    level=confirmed_level,\n                    passed=True,\n                    checks=tuple(checks),\n                    missing_fields=missing,\n                )\n            return RecognitionResult(\n                program=program,\n                level=None,\n                passed=False,\n                checks=tuple(checks),\n                missing_fields=missing,\n            )\n\n        if confirmed_level is not None:\n            return RecognitionResult(\n                program=program,\n                level=confirmed_level,\n                passed=True,\n                checks=tuple(checks),\n                missing_fields=(),\n            )\n\n        return RecognitionResult(\n            program=program,\n            level=None,\n            passed=False,\n            checks=tuple(checks),\n            missing_fields=(),\n        )\n\n    if confirmed_level is not None:\n        return RecognitionResult(\n            program=program,\n            level=confirmed_level,\n            passed=True,\n            checks=tuple(checks),\n            missing_fields=(),\n        )\n\n    return RecognitionResult(\n        program=program,\n        level=None,\n        passed=False,\n        checks=tuple(checks),\n        missing_fields=(),\n    )\n\n\ndef evaluate_primary_province(commune_levels) -> RecognitionResult:\n    return _b492_evaluate_level_province(\n        program="PRIMARY_ND20_ND142_PROVINCE",\n        commune_levels=commune_levels,\n        thresholds=PROVINCE_RECOGNITION_ND142_2025["primary"],\n    )\n\n\ndef evaluate_literacy_province(commune_levels) -> RecognitionResult:\n    return _b492_evaluate_level_province(\n        program="LITERACY_ND20_ND142_PROVINCE",\n        commune_levels=commune_levels,\n        thresholds=PROVINCE_RECOGNITION_ND142_2025["literacy"],\n    )\n\n\ndef evaluate_thcs_province(commune_levels) -> RecognitionResult:\n    return _b492_evaluate_level_province(\n        program="THCS_ND20_ND142_PROVINCE",\n        commune_levels=commune_levels,\n        thresholds=PROVINCE_RECOGNITION_ND142_2025["thcs"],\n    )\n# === BAI_13B_11_15_2_4_9_2_PROVINCE_RULES_END ==='
MN_HELPERS = '# === BAI_13B_11_15_2_4_9_2_MN_HELPERS_START ===\ndef _b492_mn_province_result_text(result) -> str:\n    if bool(getattr(result, "passed", False)):\n        return "Đạt"\n    if tuple(getattr(result, "missing_fields", ()) or ()):\n        return "Chưa đủ dữ liệu"\n    return "Không đạt"\n\n\ndef _b492_mn_commune_status(\n    *,\n    db,\n    request,\n    school_year_id: int,\n    commune_id: int,\n    is_special_difficult: bool,\n    year: int,\n):\n    from app.services.pcgd_business_rules import evaluate_preschool_3_5_commune\n\n    people = _load_people_and_records(\n        db=db,\n        request=request,\n        school_year_id=int(school_year_id),\n        selected_commune_id=int(commune_id),\n        selected_school_id=None,\n    )\n\n    metrics_by_age = _metrics(people, year)\n\n    ppc_3_5 = sum(\n        int(metrics_by_age.get(age, {}).get("ppc", 0) or 0)\n        for age in (3, 4, 5)\n    )\n    attending_3_5 = sum(\n        int(metrics_by_age.get(age, {}).get("attending", 0) or 0)\n        for age in (3, 4, 5)\n    )\n\n    metrics = {\n        "attendance_rate_3_5": _percent(attending_3_5, ppc_3_5),\n        "completion_rate_3_5_by_age": _b491_mn_completion_rates_by_age(\n            people, year\n        ),\n    }\n\n    result = evaluate_preschool_3_5_commune(\n        metrics,\n        is_special_difficult=bool(is_special_difficult),\n    )\n\n    if tuple(getattr(result, "missing_fields", ()) or ()):\n        return None\n\n    return bool(getattr(result, "passed", False))\n\n\ndef _b492_apply_mn_province_conclusion(\n    *,\n    ws,\n    report_type: str,\n    db,\n    request,\n    school_year,\n    selected_commune_id: int | None,\n    selected_school_id: int | None,\n    meta,\n    year: int,\n) -> None:\n    if report_type != "PCGD_MN_02_2025":\n        return\n\n    if (\n        selected_commune_id is not None\n        or selected_school_id is not None\n        or str(meta.get("kind") or "") != "province"\n    ):\n        return\n\n    from sqlalchemy import text as _b492_sql_text\n    from app.services.pcgd_business_rules import evaluate_preschool_3_5_province\n\n    rows = db.execute(\n        _b492_sql_text(\n            """\n            SELECT id, COALESCE(is_special_difficulty_area, 0) AS is_special\n            FROM communes\n            WHERE COALESCE(is_active, 1) = 1\n            ORDER BY id\n            """\n        )\n    ).all()\n\n    if len(rows) != 130:\n        ws["R7"] = "Chưa đủ dữ liệu"\n        return\n\n    special_statuses = []\n    other_statuses = []\n\n    for row in rows:\n        commune_id = int(row[0])\n        try:\n            is_special = bool(int(row[1] or 0))\n        except (TypeError, ValueError):\n            is_special = bool(row[1])\n\n        status = _b492_mn_commune_status(\n            db=db,\n            request=request,\n            school_year_id=int(school_year.id),\n            commune_id=commune_id,\n            is_special_difficult=is_special,\n            year=year,\n        )\n\n        if is_special:\n            special_statuses.append(status)\n        else:\n            other_statuses.append(status)\n\n    if len(special_statuses) != 51 or len(other_statuses) != 79:\n        ws["R7"] = "Chưa đủ dữ liệu"\n        return\n\n    province = evaluate_preschool_3_5_province(\n        special_difficult_commune_statuses=special_statuses,\n        other_commune_statuses=other_statuses,\n    )\n\n    ws["R7"] = _b492_mn_province_result_text(province)\n# === BAI_13B_11_15_2_4_9_2_MN_HELPERS_END ==='
MN_CALL = '_b492_apply_mn_province_conclusion(\n    ws=ws,\n    report_type=report_type,\n    db=db,\n    request=request,\n    school_year=school_year,\n    selected_commune_id=selected_commune_id,\n    selected_school_id=selected_school_id,\n    meta=meta,\n    year=year,\n)'
XMC_HELPERS = '# === BAI_13B_11_15_2_4_9_2_XMC_HELPERS_START ===\ndef _b492_province_level_item(result):\n    if result is None:\n        return {"level": None, "incomplete": True}\n\n    missing = bool(tuple(getattr(result, "missing_fields", ()) or ()))\n    passed = bool(getattr(result, "passed", False))\n    level = getattr(result, "level", None)\n\n    if passed and level is not None:\n        try:\n            return {"level": int(level), "incomplete": missing}\n        except (TypeError, ValueError):\n            return {"level": None, "incomplete": True}\n\n    if missing:\n        return {"level": None, "incomplete": True}\n\n    return {"level": 0, "incomplete": False}\n\n\ndef _b492_thcs_level_item(result, value):\n    if result is not None:\n        return _b492_province_level_item(result)\n\n    text = str(value or "").strip()\n    if text == "Chưa đủ dữ liệu":\n        return {"level": None, "incomplete": True}\n    if text == "Không đạt":\n        return {"level": 0, "incomplete": False}\n\n    try:\n        return {"level": int(value), "incomplete": False}\n    except (TypeError, ValueError):\n        return {"level": None, "incomplete": True}\n\n\ndef _b492_province_result_value(result):\n    if bool(getattr(result, "passed", False)):\n        level = getattr(result, "level", None)\n        if level is not None:\n            return int(level)\n\n    if tuple(getattr(result, "missing_fields", ()) or ()):\n        return "Chưa đủ dữ liệu"\n\n    return "Không đạt"\n\n\ndef _b492_evaluate_province_all(\n    *,\n    db,\n    request,\n    school_year_id: int,\n    year: int,\n):\n    from sqlalchemy import text as _b492_sql_text\n    from app.services.pcgd_business_rules import (\n        evaluate_primary_province,\n        evaluate_literacy_province,\n        evaluate_thcs_province,\n    )\n\n    rows = db.execute(\n        _b492_sql_text(\n            """\n            SELECT id\n            FROM communes\n            WHERE COALESCE(is_active, 1) = 1\n            ORDER BY id\n            """\n        )\n    ).all()\n\n    if len(rows) != 130:\n        return None\n\n    primary_items = []\n    literacy_items = []\n    thcs_items = []\n\n    for row in rows:\n        commune_id = int(row[0])\n\n        people = _load_people_and_records(\n            db=db,\n            request=request,\n            school_year_id=int(school_year_id),\n            selected_commune_id=commune_id,\n            selected_school_id=None,\n        )\n\n        commune = _b491_evaluate_all(\n            db=db,\n            commune_id=commune_id,\n            people=people,\n            year=year,\n        )\n\n        primary_items.append(\n            _b492_province_level_item(commune["primary"])\n        )\n        literacy_items.append(\n            _b492_province_level_item(commune["literacy"])\n        )\n        thcs_items.append(\n            _b492_thcs_level_item(\n                commune["thcs"],\n                commune["thcs_value"],\n            )\n        )\n\n    return {\n        "primary": evaluate_primary_province(primary_items),\n        "literacy": evaluate_literacy_province(literacy_items),\n        "thcs": evaluate_thcs_province(thcs_items),\n    }\n\n\ndef _b492_apply_xmc_province_conclusions(\n    *,\n    ws,\n    report_type: str,\n    db,\n    request,\n    school_year,\n    selected_commune_id: int | None,\n    selected_school_id: int | None,\n    meta,\n    year: int,\n) -> None:\n    supported = {\n        "PCGD_TH_02_2025",\n        "PCGD_THCS_M1_2025",\n        "PCGD_THCS_M2_2025",\n        "PCGD_THCS_TK_2025",\n        "PCGD_XMC_4_2025",\n    }\n\n    if report_type not in supported:\n        return\n\n    if (\n        selected_commune_id is not None\n        or selected_school_id is not None\n        or str(meta.get("kind") or "") != "province"\n    ):\n        return\n\n    result = _b492_evaluate_province_all(\n        db=db,\n        request=request,\n        school_year_id=int(school_year.id),\n        year=year,\n    )\n\n    if result is None:\n        primary_value = "Chưa đủ dữ liệu"\n        literacy_value = "Chưa đủ dữ liệu"\n        thcs_value = "Chưa đủ dữ liệu"\n    else:\n        primary_value = _b492_province_result_value(result["primary"])\n        literacy_value = _b492_province_result_value(result["literacy"])\n        thcs_value = _b492_province_result_value(result["thcs"])\n\n    if report_type == "PCGD_TH_02_2025":\n        ws["T8"] = primary_value\n    elif report_type == "PCGD_THCS_M1_2025":\n        ws["G36"] = primary_value\n        ws["G37"] = literacy_value\n    elif report_type == "PCGD_THCS_M2_2025":\n        ws["W14"] = thcs_value\n    elif report_type == "PCGD_THCS_TK_2025":\n        ws["F8"] = primary_value\n        ws["G8"] = literacy_value\n        ws["R8"] = thcs_value\n    elif report_type == "PCGD_XMC_4_2025":\n        ws["R8"] = literacy_value\n# === BAI_13B_11_15_2_4_9_2_XMC_HELPERS_END ==='
XMC_CALL = '_b492_apply_xmc_province_conclusions(\n    ws=ws,\n    report_type=report_type,\n    db=db,\n    request=request,\n    school_year=school_year,\n    selected_commune_id=selected_commune_id,\n    selected_school_id=selected_school_id,\n    meta=meta,\n    year=year,\n)'


def read_text(path: Path) -> str:
    if not path.exists():
        raise RuntimeError(f"Không tìm thấy file bắt buộc: {path}")
    return path.read_text(encoding="utf-8-sig", errors="strict")


def write_text(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8")


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def function_node(source: str, name: str):
    tree = ast.parse(source)
    matches = [
        n for n in tree.body
        if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))
        and n.name == name
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Phải tìm thấy đúng 1 hàm {name}; hiện có {len(matches)}."
        )
    return matches[0]


def get_function(source: str, name: str) -> str:
    node = function_node(source, name)
    lines = source.splitlines(keepends=True)
    return "".join(lines[node.lineno - 1:node.end_lineno])


def replace_function(source: str, name: str, new_function: str) -> str:
    node = function_node(source, name)
    lines = source.splitlines(keepends=True)
    before = "".join(lines[:node.lineno - 1])
    after = "".join(lines[node.end_lineno:])
    result = before + new_function.rstrip() + "\n\n" + after.lstrip("\n")
    ast.parse(result)
    return result


def indent_block(block: str, indent: str) -> str:
    return "\n".join(
        indent + row if row.strip() else ""
        for row in block.splitlines()
    )


def insert_call_before_save(
    source: str,
    *,
    function_name: str,
    start_marker: str,
    end_marker: str,
    call_block: str,
) -> str:
    fn = get_function(source, function_name)

    if start_marker in fn and end_marker in fn:
        return source
    if (start_marker in fn) != (end_marker in fn):
        raise RuntimeError(f"Marker call dở dang trong {function_name}.")

    fn_lines = fn.splitlines(keepends=True)
    indexes = [
        i for i, line in enumerate(fn_lines)
        if "workbook.save(output)" in line
    ]
    if len(indexes) != 1:
        raise RuntimeError(
            f"{function_name}: cần đúng 1 workbook.save(output); "
            f"hiện có {len(indexes)}."
        )

    idx = indexes[0]
    save_line = fn_lines[idx]
    indent = save_line[:len(save_line) - len(save_line.lstrip())]

    payload = (
        indent + start_marker + "\n"
        + indent_block(call_block, indent) + "\n"
        + indent + end_marker + "\n"
    )

    fn_lines.insert(idx, payload)
    new_fn = "".join(fn_lines)
    return replace_function(source, function_name, new_fn)


def append_block(
    source: str,
    *,
    start_marker: str,
    end_marker: str,
    block: str,
) -> str:
    has_start = start_marker in source
    has_end = end_marker in source

    if has_start and has_end:
        return source
    if has_start != has_end:
        raise RuntimeError("Marker block dở dang: " + start_marker)

    result = source.rstrip() + "\n\n" + block.rstrip() + "\n"
    ast.parse(result)
    return result


def all_markers_present() -> bool:
    rules = read_text(RULES)
    mn = read_text(MN_BUILDER)
    xmc = read_text(XMC_BUILDER)

    return (
        all(x in rules for x in (RULES_START, RULES_END))
        and all(
            x in mn
            for x in (MN_HELPER_START, MN_HELPER_END, MN_CALL_START, MN_CALL_END)
        )
        and all(
            x in xmc
            for x in (
                XMC_HELPER_START,
                XMC_HELPER_END,
                XMC_CALL_START,
                XMC_CALL_END,
            )
        )
    )


def any_marker_present() -> bool:
    text = "\n".join(
        (read_text(RULES), read_text(MN_BUILDER), read_text(XMC_BUILDER))
    )
    return any(
        x in text
        for x in (
            RULES_START, RULES_END,
            MN_HELPER_START, MN_HELPER_END, MN_CALL_START, MN_CALL_END,
            XMC_HELPER_START, XMC_HELPER_END, XMC_CALL_START, XMC_CALL_END,
        )
    )


def verify_hashes() -> None:
    wrong = []
    for path, expected in EXPECTED_SOURCE_SHA256.items():
        actual = sha256_file(path)
        if actual != expected:
            wrong.append((path, expected, actual))

    if wrong:
        rows = ["Source hiện tại khác nền đã khảo sát 4.9.2.0/4.9.2.1."]
        for path, expected, actual in wrong:
            rows.extend(
                [
                    f" - {path}",
                    f"   expected={expected}",
                    f"   actual  ={actual}",
                ]
            )
        raise RuntimeError("\n".join(rows))


def db_state() -> dict:
    uri = DB.resolve().as_uri() + "?mode=ro"
    con = sqlite3.connect(uri, uri=True)
    try:
        integrity = str(con.execute("PRAGMA integrity_check").fetchone()[0])
        fk = len(con.execute("PRAGMA foreign_key_check").fetchall())
        columns = {
            str(r[1])
            for r in con.execute("PRAGMA table_info(communes)").fetchall()
        }
        total = int(con.execute("SELECT COUNT(*) FROM communes").fetchone()[0])
        active = int(
            con.execute(
                "SELECT COUNT(*) FROM communes WHERE COALESCE(is_active,1)=1"
            ).fetchone()[0]
        )
        special = int(
            con.execute(
                """
                SELECT COUNT(*) FROM communes
                WHERE COALESCE(is_active,1)=1
                  AND COALESCE(is_special_difficulty_area,0)=1
                """
            ).fetchone()[0]
        )
        return {
            "integrity": integrity,
            "fk": fk,
            "columns": columns,
            "total": total,
            "active": active,
            "special": special,
        }
    finally:
        con.close()


def verify_db(state: dict) -> None:
    if str(state["integrity"]).lower() != "ok":
        raise RuntimeError("integrity_check != ok")
    if int(state["fk"]) != 0:
        raise RuntimeError("foreign_key_check có lỗi.")
    if "is_special_difficulty_area" not in state["columns"]:
        raise RuntimeError("communes thiếu is_special_difficulty_area.")
    if int(state["total"]) != EXPECTED_COMMUNES:
        raise RuntimeError(
            f"Tổng communes={state['total']}, không phải {EXPECTED_COMMUNES}."
        )
    if int(state["active"]) != EXPECTED_COMMUNES:
        raise RuntimeError(
            f"Xã/phường active={state['active']}, không phải {EXPECTED_COMMUNES}."
        )
    if int(state["special"]) != EXPECTED_SPECIAL:
        raise RuntimeError(
            f"ĐBKK active={state['special']}, không phải {EXPECTED_SPECIAL}."
        )


def verify_template(path: Path, cell: str, words: Iterable[str]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=False, data_only=False)
    try:
        ws = wb[wb.sheetnames[0]]
        value = str(ws[cell].value or "").lower()
        if not all(w.lower() in value for w in words):
            raise RuntimeError(
                f"Template {path.name}: ô {cell}={ws[cell].value!r} "
                "không đúng nền."
            )
    finally:
        wb.close()


def verify_templates() -> None:
    verify_template(MN_TEMPLATE, "R3", ("đạt",))
    verify_template(TH_TEMPLATE, "T4", ("đạt",))
    verify_template(THCS_M2_TEMPLATE, "W9", ("đạt",))
    verify_template(THCS_TK_TEMPLATE, "F4", ("đạt",))
    verify_template(THCS_TK_TEMPLATE, "G4", ("đạt",))
    verify_template(THCS_TK_TEMPLATE, "R4", ("đạt",))
    verify_template(XMC4_TEMPLATE, "R5", ("đạt",))


def verify_b491_base() -> None:
    mn = read_text(MN_BUILDER)
    xmc = read_text(XMC_BUILDER)

    for token in (
        "BAI_13B_11_15_2_4_9_1_MN_HELPERS_START",
        "def _b491_apply_mn_conclusion(",
        "def _b491_mn_completion_rates_by_age(",
        "_load_people_and_records",
    ):
        if token not in mn:
            raise RuntimeError("Nền 4.9.1 MN thiếu: " + token)

    for token in (
        "BAI_13B_11_15_2_4_9_1_XMC_HELPERS_START",
        "def _b491_evaluate_all(",
        "def _b491_primary_metrics(",
        "def _b491_literacy_metrics(",
        "def _b491_thcs_metrics(",
        "def _load_people_and_records(",
    ):
        if token not in xmc:
            raise RuntimeError("Nền 4.9.1 TH/XMC/THCS thiếu: " + token)


def backup_sources() -> None:
    for path in (RULES, MN_BUILDER, XMC_BUILDER):
        dst = BACKUP / path.relative_to(PROJECT)
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, dst)


def restore_sources() -> None:
    for path in (RULES, MN_BUILDER, XMC_BUILDER):
        src = BACKUP / path.relative_to(PROJECT)
        if src.exists():
            shutil.copy2(src, path)


def patch_all() -> None:
    rules = append_block(
        read_text(RULES),
        start_marker=RULES_START,
        end_marker=RULES_END,
        block=RULES_BLOCK,
    )

    mn = append_block(
        read_text(MN_BUILDER),
        start_marker=MN_HELPER_START,
        end_marker=MN_HELPER_END,
        block=MN_HELPERS,
    )
    mn = insert_call_before_save(
        mn,
        function_name="export_mn_report",
        start_marker=MN_CALL_START,
        end_marker=MN_CALL_END,
        call_block=MN_CALL,
    )

    xmc = append_block(
        read_text(XMC_BUILDER),
        start_marker=XMC_HELPER_START,
        end_marker=XMC_HELPER_END,
        block=XMC_HELPERS,
    )
    xmc = insert_call_before_save(
        xmc,
        function_name="export_additional_report",
        start_marker=XMC_CALL_START,
        end_marker=XMC_CALL_END,
        call_block=XMC_CALL,
    )

    ast.parse(rules)
    ast.parse(mn)
    ast.parse(xmc)

    write_text(RULES, rules)
    write_text(MN_BUILDER, mn)
    write_text(XMC_BUILDER, xmc)


def verify_installed() -> None:
    rules = read_text(RULES)
    mn = read_text(MN_BUILDER)
    xmc = read_text(XMC_BUILDER)

    ast.parse(rules)
    ast.parse(mn)
    ast.parse(xmc)

    for name in (
        "evaluate_preschool_3_5_province",
        "evaluate_primary_province",
        "evaluate_literacy_province",
        "evaluate_thcs_province",
    ):
        function_node(rules, name)

    for token in (
        RULES_START,
        RULES_END,
        "PROVINCE_RECOGNITION_ND142_2025",
    ):
        if token not in rules:
            raise RuntimeError("Verifier rules thiếu: " + token)

    for token in (
        MN_HELPER_START, MN_HELPER_END, MN_CALL_START, MN_CALL_END,
        "def _b492_apply_mn_province_conclusion(",
        'ws["R7"]',
    ):
        if token not in mn:
            raise RuntimeError("Verifier MN thiếu: " + token)

    for token in (
        XMC_HELPER_START, XMC_HELPER_END, XMC_CALL_START, XMC_CALL_END,
        "def _b492_apply_xmc_province_conclusions(",
        'ws["T8"]', 'ws["G36"]', 'ws["G37"]', 'ws["W14"]',
        'ws["F8"]', 'ws["G8"]', 'ws["R8"]',
    ):
        if token not in xmc:
            raise RuntimeError("Verifier TH/XMC/THCS thiếu: " + token)

    mn_export = get_function(mn, "export_mn_report")
    xmc_export = get_function(xmc, "export_additional_report")

    mn_call = mn_export.find(MN_CALL_START)
    mn_save = mn_export.find("workbook.save(output)")
    xmc_call = xmc_export.find(XMC_CALL_START)
    xmc_save = xmc_export.find("workbook.save(output)")

    if mn_call < 0 or mn_save < 0 or mn_call >= mn_save:
        raise RuntimeError("Call 4.9.2 MN chưa đứng đúng trước workbook.save.")
    if xmc_call < 0 or xmc_save < 0 or xmc_call >= xmc_save:
        raise RuntimeError(
            "Call 4.9.2 TH/XMC/THCS chưa đứng đúng trước workbook.save."
        )

    for path in (RULES, MN_BUILDER, XMC_BUILDER):
        subprocess.run(
            [sys.executable, "-m", "py_compile", str(path)],
            cwd=str(PROJECT),
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )


def smoke_test() -> None:
    code = r"""
from app.services.pcgd_business_rules import (
    evaluate_preschool_3_5_province,
    evaluate_primary_province,
    evaluate_literacy_province,
    evaluate_thcs_province,
)

mn = evaluate_preschool_3_5_province(
    special_difficult_commune_statuses=[True]*44 + [False]*7,
    other_commune_statuses=[True]*72 + [False]*7,
)
assert mn.passed is True

primary = evaluate_primary_province(
    [{"level": 3, "incomplete": False}]*117
    + [{"level": 0, "incomplete": False}]*13
)
assert primary.passed is True and primary.level == 3

literacy = evaluate_literacy_province(
    [{"level": 2, "incomplete": False}]*117
    + [{"level": 0, "incomplete": False}]*13
)
assert literacy.passed is True and literacy.level == 2

thcs = evaluate_thcs_province(
    [{"level": 2, "incomplete": False}]*124
    + [{"level": 0, "incomplete": False}]*6
)
assert thcs.passed is True and thcs.level == 2

uncertain = evaluate_primary_province(
    [{"level": 1, "incomplete": False}]*116
    + [{"level": None, "incomplete": True}]
    + [{"level": 0, "incomplete": False}]*13
)
assert uncertain.passed is False
assert uncertain.level is None
assert uncertain.missing_fields

print("B492_SMOKE_OK")
"""
    result = subprocess.run(
        [sys.executable, "-c", code],
        cwd=str(PROJECT),
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    if "B492_SMOKE_OK" not in result.stdout:
        raise RuntimeError("Smoke test 4.9.2 không trả B492_SMOKE_OK.")


def clear_cache() -> None:
    for cache in APP.rglob("__pycache__"):
        if cache.is_dir():
            shutil.rmtree(cache, ignore_errors=True)


def main() -> int:
    print("=" * 104)
    print(
        "BÀI 13B-11.15.2.4.9.2 V2 - "
        "TỰ ĐỘNG KẾT LUẬN / CÔNG NHẬN CẤP TỈNH - 4 NHÓM"
    )
    print("=" * 104)
    print()

    backup_created = False
    already_installed = False

    try:
        for path in (
            DB, RULES, MN_BUILDER, XMC_BUILDER,
            MN_TEMPLATE, TH_TEMPLATE, THCS_M2_TEMPLATE,
            THCS_TK_TEMPLATE, XMC4_TEMPLATE,
        ):
            if not path.exists():
                raise RuntimeError(f"Không tìm thấy: {path}")

        db_hash_before = sha256_file(DB)

        print("[1/9] Kiểm tra database...")
        before = db_state()
        verify_db(before)
        print(
            f"[OK] {before['active']} xã/phường; "
            f"ĐBKK={before['special']}; integrity=ok."
        )

        print("[2/9] Kiểm tra nền Bài 4.9.1 + loader từng xã...")
        verify_b491_base()
        print("[OK] Đúng nền.")

        print("[3/9] Kiểm tra template...")
        verify_templates()
        print("[OK] Template đúng nền.")

        if all_markers_present():
            already_installed = True
            print("[4/9] Marker 4.9.2 đã đầy đủ.")
            print("[5/9] Không sửa source lặp.")
            print("[6/9] Kiểm tra source hiện có...")
            verify_installed()
            smoke_test()
            print("[OK] Source hiện có đạt verifier.")
        else:
            if any_marker_present():
                raise RuntimeError(
                    "Phát hiện marker Bài 4.9.2 dở dang. "
                    "Dừng để tránh sửa chồng."
                )

            print("[4/9] Khóa SHA256 source sau khảo sát...")
            verify_hashes()
            print("[OK] 3/3 source đúng hash.")

            print("[5/9] Backup source...")
            BACKUP.mkdir(parents=True, exist_ok=False)
            backup_created = True
            backup_sources()
            print(f"[OK] {BACKUP}")

            print("[6/9] Cài rules + nối exporter...")
            patch_all()
            verify_installed()
            smoke_test()
            print("[OK] Đã cài 4 nhóm.")

        print("[7/9] AST + py_compile + smoke test...")
        verify_installed()
        smoke_test()
        print("[OK] Đạt.")

        print("[8/9] Kiểm tra database sau cài...")
        after = db_state()
        verify_db(after)
        db_hash_after = sha256_file(DB)

        if db_hash_after != db_hash_before:
            raise RuntimeError(
                "SHA256 phocap.db thay đổi. Bài 4.9.2 không được sửa DB."
            )
        print("[OK] Database KHÔNG THAY ĐỔI; SHA256 giữ nguyên.")

        print("[9/9] Ghi báo cáo...")
        EXPORTS.mkdir(parents=True, exist_ok=True)

        lines = [
            "=" * 104,
            "BÀI 13B-11.15.2.4.9.2 V2 - BÁO CÁO CÀI ĐẶT",
            "=" * 104,
            "",
            "PHẠM VI:",
            " - Xã/phường: giữ nguyên Bài 4.9.1.",
            " - Toàn tỉnh: đánh giá từng xã rồi tổng hợp tỷ lệ.",
            " - Trường: không tự kết luận.",
            "",
            "NGƯỠNG CẤP TỈNH:",
            " - PCGDMN 3-5: ĐBKK >=85%; vùng còn lại >=90%.",
            " - Tiểu học M1/M2/M3: >=90% số xã.",
            " - XMC M1 >=81%; M2 >=90%.",
            " - THCS M1 >=90%; M2 >=95%; M3 =100%.",
            "",
            "Ô TOÀN TỈNH:",
            " - MN-02 R7",
            " - TH-02 T8",
            " - THCS-M1 G36/G37",
            " - THCS-M2 W14",
            " - THCS-TK F8/G8/R8",
            " - XMC-4 R8",
            "",
            f"Database SHA256 trước: {db_hash_before}",
            f"Database SHA256 sau  : {db_hash_after}",
            "Database: KHÔNG THAY ĐỔI.",
            "",
            (
                "Source đã có từ trước; verifier đạt."
                if already_installed
                else f"Backup source: {BACKUP}"
            ),
            "",
            "AST + py_compile + smoke test: ĐẠT.",
            "KẾT LUẬN: BÀI 4.9.2 V2 CÀI ĐẶT THÀNH CÔNG.",
        ]

        REPORT.write_text("\n".join(lines), encoding="utf-8-sig")
        clear_cache()

        print(f"[OK] {REPORT}")
        print()
        print("=" * 104)
        print("BÀI 4.9.2 V2 HOÀN THÀNH.")
        print("=" * 104)
        print("Cấp xã/phường        : GIỮ Bài 4.9.1")
        print("Toàn tỉnh / MN-02    : R7 tự kết luận")
        print("Toàn tỉnh / TH-02    : T8 tự xác định mức")
        print("Toàn tỉnh / XMC-4    : R8 tự xác định mức")
        print("Toàn tỉnh / THCS     : M1/M2/TK tự nối mức")
        print("Database             : KHÔNG THAY ĐỔI")
        print(f"Báo cáo              : {REPORT}")
        print("=" * 104)
        return 0

    except Exception as exc:
        print()
        print("=" * 104)
        print("[LỖI] BÀI 4.9.2 V2 KHÔNG HOÀN THÀNH")
        print(str(exc))
        print("=" * 104)

        if backup_created:
            try:
                restore_sources()
                clear_cache()
                print("[ROLLBACK] Đã khôi phục 3 source từ backup.")
            except Exception as rollback_exc:
                print("[CẢNH BÁO] Rollback lỗi: " + str(rollback_exc))

        print("Bài này không có lệnh UPDATE/INSERT/DELETE database.")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
