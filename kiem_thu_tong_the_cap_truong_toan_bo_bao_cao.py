# -*- coding: utf-8 -*-
"""
KIỂM THỬ TỔNG THỂ HỆ THỐNG BÁO CÁO - CẤP TRƯỜNG

KHÔNG GẮN SỐ BÀI MỚI.

PHẠM VI
-------
Tự chọn:
- 1 trường MẦM NON
- 1 trường TIỂU HỌC
- 1 trường THCS

Mỗi trường chỉ kiểm đúng nhóm báo cáo thuộc cấp học của mình:

MẦM NON: 5 biểu
  1. MN-M1
  2. MN-02
  3. MN-01-GV
  4. MN-01-CSVC
  5. MN-TÀI CHÍNH

TIỂU HỌC: 4 biểu
  6. TH-M1
  7. TH-02
  8. TH-01-GV
  9. TH-01-CSVC

THCS: 5 biểu
 10. THCS-M1
 11. THCS-M2
 12. THCS-TK
 13. THCS-M5
 14. THCS-CSVC

XÓA MÙ CHỮ không gắn vào tài khoản trường trong ma trận kiểm thử này.

KIỂM TRA
--------
- Exporter thực tế chạy được.
- File XLSX hợp lệ, không có #REF!, #DIV/0!, #VALUE!...
- Có tên đúng trường trong file.
- Không xuất nhầm chữ "Toàn tỉnh".
- Kiểm dấu hiệu chữ ký Nhà trường/Hiệu trưởng.
- Các ô kết luận tự động phải TRỐNG ở scope trường:
    MN-02 R7
    TH-02 T8
    THCS-M1 G36/G37
    THCS-M2 W14
    THCS-TK F8/G8/R8
  (4.9.1 chỉ xã; 4.9.2 chỉ tỉnh.)
- Định dạng % MN/TH/THCS vẫn là 0\\%.
- Database không thay đổi.

AN TOÀN
-------
- Không sửa source.
- Không UPDATE / INSERT / DELETE database.
- PRAGMA query_only = ON.
- integrity_check + foreign_key_check.
- SHA256 phocap.db trước/sau phải giữ nguyên.
"""

from __future__ import annotations

import asyncio
import ast
import hashlib
import inspect
import os
import re
import sqlite3
import sys
import traceback
import unicodedata
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any


PROJECT = Path(
    os.environ.get(
        "PHOCAP_PROJECT",
        r"C:\PhoCap",
    )
).resolve()

APP = PROJECT / "app"
DB_PATH = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"

MN_BUILDER_PATH = APP / "pcgd_mn_report_builders_v1.py"
XMC_BUILDER_PATH = APP / "pcgd_xmc_report_builders_v1.py"
REPORT_CENTER_PATH = APP / "routers" / "report_center.py"
RULES_PATH = APP / "services" / "pcgd_business_rules.py"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

OUT_DIR = EXPORTS / f"kiem_thu_tong_the_cap_truong_{STAMP}"
REPORT_PATH = (
    OUT_DIR
    / f"bao_cao_kiem_thu_tong_the_cap_truong_{STAMP}.txt"
)
SUMMARY_CSV = (
    OUT_DIR
    / f"tong_hop_kiem_thu_cap_truong_{STAMP}.csv"
)

EXPECTED_PERCENT_FORMAT = r"0\%"

EXCEL_ERROR_VALUES = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
}


REPORTS_BY_LEVEL = {
    "MN": (
        {
            "group": "MẦM NON",
            "short": "MN-M1",
            "code": "PCGD_MN_M1_2025",
            "exporter": "mn",
            "blank_conclusion": (),
            "percent_formats": (),
        },
        {
            "group": "MẦM NON",
            "short": "MN-02",
            "code": "PCGD_MN_02_2025",
            "exporter": "mn",
            "blank_conclusion": ("R7",),
            "percent_formats": (
                "H7",
                "K7",
                "O7",
            ),
        },
        {
            "group": "MẦM NON",
            "short": "MN-01-GV",
            "code": "PCGD_MN_01_GV_2025",
            "exporter": "mn",
            "blank_conclusion": (),
            "percent_formats": (),
        },
        {
            "group": "MẦM NON",
            "short": "MN-01-CSVC",
            "code": "PCGD_MN_01_CSVC_2025",
            "exporter": "mn",
            "blank_conclusion": (),
            "percent_formats": (),
        },
        {
            "group": "MẦM NON",
            "short": "MN-TÀI-CHÍNH",
            "code": "PCGD_MN_TAICHINH_2025",
            "exporter": "mn",
            "blank_conclusion": (),
            "percent_formats": (),
        },
    ),

    "TH": (
        {
            "group": "TIỂU HỌC",
            "short": "TH-M1",
            "code": "PCGD_TH_M1_2025",
            "exporter": "th_m1",
            "blank_conclusion": (),
            "percent_formats": (),
        },
        {
            "group": "TIỂU HỌC",
            "short": "TH-02",
            "code": "PCGD_TH_02_2025",
            "exporter": "xmc",
            "blank_conclusion": ("T8",),
            "percent_formats": (
                "I8",
                "K8",
                "M8",
                "Q8",
            ),
        },
        {
            "group": "TIỂU HỌC",
            "short": "TH-01-GV",
            "code": "PCGD_TH_01_GV_2025",
            "exporter": "xmc",
            "blank_conclusion": (),
            "percent_formats": (),
        },
        {
            "group": "TIỂU HỌC",
            "short": "TH-01-CSVC",
            "code": "PCGD_TH_01_CSVC_2025",
            "exporter": "xmc",
            "blank_conclusion": (),
            "percent_formats": (),
        },
    ),

    "THCS": (
        {
            "group": "THCS",
            "short": "THCS-M1",
            "code": "PCGD_THCS_M1_2025",
            "exporter": "xmc",
            "blank_conclusion": (
                "G36",
                "G37",
            ),
            "percent_formats": (
                "H39",
                "H40",
                "H41",
            ),
        },
        {
            "group": "THCS",
            "short": "THCS-M2",
            "code": "PCGD_THCS_M2_2025",
            "exporter": "xmc",
            "blank_conclusion": ("W14",),
            "percent_formats": (
                "E14",
                "J14",
                "M14",
                "Q14",
                "V14",
                "E15",
                "J15",
                "M15",
                "Q15",
                "V15",
                "K20",
                "K21",
                "K22",
                "K23",
                "K24",
            ),
        },
        {
            "group": "THCS",
            "short": "THCS-TK",
            "code": "PCGD_THCS_TK_2025",
            "exporter": "xmc",
            "blank_conclusion": (
                "F8",
                "G8",
                "R8",
            ),
            "percent_formats": (
                "I8",
                "K8",
                "O8",
            ),
        },
        {
            "group": "THCS",
            "short": "THCS-M5",
            "code": "PCGD_THCS_M5_2025",
            "exporter": "xmc",
            "blank_conclusion": (),
            "percent_formats": (),
        },
        {
            "group": "THCS",
            "short": "THCS-CSVC",
            "code": "PCGD_THCS_CSVC_2025",
            "exporter": "xmc",
            "blank_conclusion": (),
            "percent_formats": (),
        },
    ),
}


def sha256_file(
    path: Path,
) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(
                1024 * 1024
            ),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


def norm_text(
    value: Any,
) -> str:
    text = str(
        value
        or ""
    ).strip().upper()

    if not text:
        return ""

    text = unicodedata.normalize(
        "NFD",
        text,
    )

    text = "".join(
        ch
        for ch in text
        if unicodedata.category(ch)
        != "Mn"
    )

    return " ".join(
        text.split()
    )


def safe_filename(
    value: str,
) -> str:
    text = norm_text(
        value
    )

    text = re.sub(
        r"[^A-Z0-9_-]+",
        "_",
        text,
    )

    return (
        text.strip("_")
        or "BAO_CAO"
    )


def normalize_cell_value(
    value: Any,
) -> Any:
    if value is None:
        return None

    if isinstance(
        value,
        bool,
    ):
        return value

    if isinstance(
        value,
        int,
    ):
        return value

    if isinstance(
        value,
        float,
    ):
        if value.is_integer():
            return int(
                value
            )

        return value

    return " ".join(
        str(value)
        .strip()
        .split()
    )


def db_state() -> dict[str, Any]:
    uri = (
        DB_PATH.resolve().as_uri()
        + "?mode=ro"
    )

    con = sqlite3.connect(
        uri,
        uri=True,
    )

    try:
        integrity = str(
            con.execute(
                "PRAGMA integrity_check"
            ).fetchone()[0]
        )

        fk_rows = con.execute(
            "PRAGMA foreign_key_check"
        ).fetchall()

        communes = int(
            con.execute(
                "SELECT COUNT(*) "
                "FROM communes "
                "WHERE is_active=1"
            ).fetchone()[0]
        )

        schools = int(
            con.execute(
                "SELECT COUNT(*) "
                "FROM schools "
                "WHERE is_active=1"
            ).fetchone()[0]
        )

        return {
            "integrity": integrity,
            "fk_count": len(
                fk_rows
            ),
            "communes": communes,
            "schools": schools,
        }

    finally:
        con.close()


def verify_db_state(
    state: dict[str, Any],
) -> None:
    if (
        str(
            state["integrity"]
        ).lower()
        != "ok"
    ):
        raise RuntimeError(
            "integrity_check != ok"
        )

    if int(
        state["fk_count"]
    ) != 0:
        raise RuntimeError(
            "foreign_key_check có lỗi."
        )

    if int(
        state["communes"]
    ) <= 0:
        raise RuntimeError(
            "Không có xã hoạt động."
        )

    if int(
        state["schools"]
    ) <= 0:
        raise RuntimeError(
            "Không có trường hoạt động."
        )


def verify_source_files() -> None:
    for path in (
        MN_BUILDER_PATH,
        XMC_BUILDER_PATH,
        REPORT_CENTER_PATH,
        RULES_PATH,
    ):
        if not path.exists():
            raise RuntimeError(
                "Không tìm thấy source: "
                + str(
                    path
                )
            )

        ast.parse(
            path.read_text(
                encoding="utf-8-sig",
                errors="strict",
            )
        )

    mn = MN_BUILDER_PATH.read_text(
        encoding="utf-8-sig",
    )

    xmc = XMC_BUILDER_PATH.read_text(
        encoding="utf-8-sig",
    )

    rc = REPORT_CENTER_PATH.read_text(
        encoding="utf-8-sig",
    )

    required = (
        (
            mn,
            "BAI_13B_11_15_2_4_9_1_MN_HELPERS_START",
        ),
        (
            mn,
            "BAI_13B_11_15_2_4_9_2_MN_HELPERS_START",
        ),
        (
            mn,
            "BAI_13B_11_15_2_4_9_2A_1_MN_HELPERS_START",
        ),
        (
            xmc,
            "BAI_13B_11_15_2_4_9_1_XMC_HELPERS_START",
        ),
        (
            xmc,
            "BAI_13B_11_15_2_4_9_2_XMC_HELPERS_START",
        ),
        (
            xmc,
            "BAI_13B_11_15_2_4_9_2A_1_XMC_HELPERS_START",
        ),
        (
            rc,
            "def _export_primary_th_m1(",
        ),
    )

    missing = [
        token
        for source, token
        in required
        if token not in source
    ]

    if missing:
        raise RuntimeError(
            "Source thiếu nền kiểm thử trường: "
            + "; ".join(
                missing
            )
        )


def make_school_request(
    school: dict,
):
    from starlette.requests import Request

    scope = {
        "type": "http",
        "asgi": {
            "version": "3.0",
            "spec_version": "2.3",
        },
        "http_version": "1.1",
        "server": (
            "127.0.0.1",
            80,
        ),
        "client": (
            "127.0.0.1",
            0,
        ),
        "scheme": "http",
        "method": "GET",
        "root_path": "",
        "path": "/bao-cao",
        "raw_path": b"/bao-cao",
        "query_string": b"",
        "headers": [],
        "auth_user": {
            "id": 0,
            "username": (
                "kiem_thu_cap_truong"
            ),
            "full_name": (
                "Kiểm thử tổng thể cấp Trường"
            ),
            "role_code": "TRUONG",
            "role_name": "Trường",
            "commune_id": int(
                school["commune_id"]
            ),
            "school_id": int(
                school["id"]
            ),
            "unit_name": str(
                school["name"]
            ),
        },
    }

    return Request(
        scope
    )


async def response_to_bytes(
    response,
) -> bytes:
    body = getattr(
        response,
        "body",
        None,
    )

    if isinstance(
        body,
        (
            bytes,
            bytearray,
        ),
    ) and body:
        return bytes(
            body
        )

    iterator = getattr(
        response,
        "body_iterator",
        None,
    )

    if iterator is None:
        raise RuntimeError(
            "Response không có "
            "body/body_iterator."
        )

    chunks = []

    async for chunk in iterator:
        if isinstance(
            chunk,
            str,
        ):
            chunk = chunk.encode(
                "utf-8"
            )

        chunks.append(
            bytes(
                chunk
            )
        )

    return b"".join(
        chunks
    )


def resolve_result(
    value,
):
    if inspect.isawaitable(
        value
    ):
        return asyncio.run(
            value
        )

    return value


def call_with_known_args(
    func,
    *,
    db,
    request,
    school_year,
    school: dict,
    report_type: str | None = None,
):
    candidates = {
        "db": db,
        "request": request,
        "school_year": school_year,
        "report_type": report_type,
        "selected_commune_id": int(
            school["commune_id"]
        ),
        "selected_school_id": int(
            school["id"]
        ),
        "scope_label": str(
            school["name"]
        ),
    }

    signature = inspect.signature(
        func
    )

    kwargs = {}

    for (
        name,
        parameter,
    ) in signature.parameters.items():
        if name in candidates:
            if (
                name == "report_type"
                and report_type is None
                and parameter.default
                is not inspect.Parameter.empty
            ):
                continue

            kwargs[
                name
            ] = candidates[
                name
            ]

            continue

        if (
            parameter.default
            is not inspect.Parameter.empty
        ):
            continue

        if parameter.kind in (
            inspect.Parameter.VAR_POSITIONAL,
            inspect.Parameter.VAR_KEYWORD,
        ):
            continue

        raise RuntimeError(
            f"{func.__module__}."
            f"{func.__name__}: "
            "tham số bắt buộc chưa biết: "
            + name
        )

    return resolve_result(
        func(
            **kwargs
        )
    )


def choose_school_year(
    db,
):
    from sqlalchemy import (
        select,
        text,
    )

    from app.models import (
        SchoolYear,
    )

    requested = str(
        os.environ.get(
            "PHOCAP_TEST_YEAR",
            "",
        )
        or ""
    ).strip()

    if requested:
        year = db.scalar(
            select(
                SchoolYear
            ).where(
                SchoolYear.code
                == requested
            )
        )

        if year is None:
            raise RuntimeError(
                "Không có năm học "
                + requested
            )

        return year

    rows = db.execute(
        text(
            """
            SELECT
                sy.id,
                sy.code,
                COUNT(spr.id) AS cnt
            FROM school_years sy
            LEFT JOIN
                survey_person_year_records spr
              ON spr.school_year_id=sy.id
            GROUP BY
                sy.id,
                sy.code
            ORDER BY
                cnt DESC,
                sy.code DESC
            """
        )
    ).all()

    if not rows:
        raise RuntimeError(
            "Database chưa có năm học."
        )

    best_id = int(
        rows[0][0]
    )

    year = db.scalar(
        select(
            SchoolYear
        ).where(
            SchoolYear.id
            == best_id
        )
    )

    if year is None:
        raise RuntimeError(
            "Không lấy được năm học."
        )

    return year


def expected_reference_year(
    school_year,
) -> str:
    code = str(
        getattr(
            school_year,
            "code",
            "",
        )
        or ""
    )

    years = re.findall(
        r"\d{4}",
        code,
    )

    return (
        years[0]
        if years
        else ""
    )


def table_columns(
    db,
    table: str,
) -> set[str]:
    from sqlalchemy import text

    rows = db.execute(
        text(
            f'PRAGMA table_info("{table}")'
        )
    ).all()

    return {
        str(
            row[1]
        )
        for row in rows
    }


def school_person_count(
    db,
    school_id: int,
    school_year_id: int,
) -> int:
    from sqlalchemy import text

    columns = table_columns(
        db,
        "survey_person_year_records",
    )

    if "school_id" not in columns:
        return 0

    value = db.execute(
        text(
            """
            SELECT COUNT(*)
            FROM survey_person_year_records
            WHERE school_year_id=:year_id
              AND school_id=:school_id
            """
        ),
        {
            "year_id": int(
                school_year_id
            ),
            "school_id": int(
                school_id
            ),
        },
    ).scalar_one()

    return int(
        value
        or 0
    )


def school_class_count(
    grades: dict,
    school_id: int,
) -> int:
    try:
        value = grades.get(
            int(
                school_id
            ),
            [],
        )
    except Exception:
        return 0

    try:
        return len(
            value
        )
    except Exception:
        return 0


def classify_school_name(
    name: str,
) -> str | None:
    n = norm_text(
        name
    )

    if (
        "THCS" in n
        or "TRUNG HOC CO SO"
        in n
    ):
        return "THCS"

    if (
        "MAM NON" in n
        or "MAU GIAO" in n
        or re.search(
            r"(^|\s)MN(\s|$)",
            n,
        )
    ):
        return "MN"

    if (
        "TIEU HOC"
        in n
        or re.search(
            r"(^|\s)TH(\s|$)",
            n,
        )
    ):
        return "TH"

    return None


def normalize_level(
    value: Any,
) -> str | None:
    n = norm_text(
        value
    )

    if not n:
        return None

    if (
        n == "THCS"
        or "TRUNG HOC CO SO"
        in n
        or "CAP 2"
        in n
    ):
        return "THCS"

    if (
        n == "MN"
        or "MAM NON"
        in n
        or "MAU GIAO"
        in n
    ):
        return "MN"

    if (
        n == "TH"
        or "TIEU HOC"
        in n
        or "CAP 1"
        in n
    ):
        return "TH"

    return None


def choose_test_schools(
    db,
    *,
    school_year_id: int,
    xmc_builder,
) -> tuple[dict, dict, dict]:
    from sqlalchemy import select
    from app.models import School

    # Dùng chính helper mà exporter TH/THCS đang dùng để đọc
    # danh sách trường/lớp. Nếu helper không nhận diện được MN,
    # fallback theo tên trường.
    try:
        all_schools, grades = (
            xmc_builder
            ._schools_and_classes(
                db,
                int(
                    school_year_id
                ),
                None,
                None,
            )
        )
    except Exception:
        all_schools = list(
            db.scalars(
                select(
                    School
                ).where(
                    School.is_active.is_(
                        True
                    )
                )
            ).all()
        )
        grades = {}

    pools = {
        "MN": [],
        "TH": [],
        "THCS": [],
    }

    for school in all_schools:
        if not bool(
            getattr(
                school,
                "is_active",
                True,
            )
        ):
            continue

        school_id = int(
            school.id
        )

        school_name = str(
            getattr(
                school,
                "name",
                "",
            )
            or ""
        )

        commune_id = getattr(
            school,
            "commune_id",
            None,
        )

        if commune_id is None:
            continue

        level = None

        # Ưu tiên helper của source thực tế.
        try:
            helper = getattr(
                xmc_builder,
                "_school_level",
                None,
            )

            if callable(
                helper
            ):
                level = normalize_level(
                    helper(
                        school,
                        grades,
                    )
                )
        except Exception:
            level = None

        # Kiểm các thuộc tính thường dùng nếu helper chưa đủ.
        if level is None:
            for attr in (
                "level_code",
                "school_level",
                "education_level",
                "level",
                "school_type",
                "type_code",
            ):
                level = normalize_level(
                    getattr(
                        school,
                        attr,
                        None,
                    )
                )

                if level is not None:
                    break

        # Cuối cùng fallback theo tên.
        if level is None:
            level = classify_school_name(
                school_name
            )

        if level not in pools:
            continue

        classes = school_class_count(
            grades,
            school_id,
        )

        people = school_person_count(
            db,
            school_id,
            int(
                school_year_id
            ),
        )

        item = {
            "id": school_id,
            "name": school_name,
            "commune_id": int(
                commune_id
            ),
            "level": level,
            "class_count": int(
                classes
            ),
            "person_count": int(
                people
            ),
        }

        # Dữ liệu điều tra ưu tiên trước, rồi đến số lớp.
        item["score"] = (
            int(
                people
            )
            * 10000
            + int(
                classes
            )
        )

        pools[
            level
        ].append(
            item
        )

    env_map = {
        "MN": "PHOCAP_TEST_MN_SCHOOL_ID",
        "TH": "PHOCAP_TEST_TH_SCHOOL_ID",
        "THCS": "PHOCAP_TEST_THCS_SCHOOL_ID",
    }

    selected = {}

    for level in (
        "MN",
        "TH",
        "THCS",
    ):
        pool = pools[
            level
        ]

        if not pool:
            raise RuntimeError(
                "Không tìm thấy trường "
                f"cấp {level} để kiểm thử."
            )

        requested = str(
            os.environ.get(
                env_map[
                    level
                ],
                "",
            )
            or ""
        ).strip()

        if requested:
            try:
                wanted = int(
                    requested
                )
            except ValueError:
                raise RuntimeError(
                    f"{env_map[level]} "
                    "không phải số."
                )

            matches = [
                item
                for item in pool
                if item["id"]
                == wanted
            ]

            if len(
                matches
            ) != 1:
                raise RuntimeError(
                    f"Không tìm thấy school_id="
                    f"{wanted} đúng cấp {level}."
                )

            selected[
                level
            ] = dict(
                matches[0]
            )

        else:
            ranked = sorted(
                pool,
                key=lambda item: (
                    int(
                        item["score"]
                    ),
                    norm_text(
                        item["name"]
                    ),
                ),
                reverse=True,
            )

            selected[
                level
            ] = dict(
                ranked[0]
            )

    return (
        selected["MN"],
        selected["TH"],
        selected["THCS"],
    )


def workbook_scan(
    data: bytes,
    spec: dict,
    school: dict,
    reference_year: str,
):
    from openpyxl import load_workbook

    if not data.startswith(
        b"PK"
    ):
        raise RuntimeError(
            "Response không phải XLSX."
        )

    workbook = load_workbook(
        BytesIO(
            data
        ),
        data_only=False,
    )

    try:
        if not workbook.sheetnames:
            raise RuntimeError(
                "Workbook không có sheet."
            )

        all_text = []
        errors = []
        nonempty = 0
        truncated = False

        for ws in workbook.worksheets:
            max_row = min(
                max(
                    int(
                        ws.max_row
                        or 1
                    ),
                    1,
                ),
                250,
            )

            max_col = min(
                max(
                    int(
                        ws.max_column
                        or 1
                    ),
                    1,
                ),
                120,
            )

            if (
                int(
                    ws.max_row
                    or 1
                )
                > max_row
                or int(
                    ws.max_column
                    or 1
                )
                > max_col
            ):
                truncated = True

            for row in ws.iter_rows(
                min_row=1,
                max_row=max_row,
                min_col=1,
                max_col=max_col,
            ):
                for cell in row:
                    value = cell.value

                    if value not in (
                        None,
                        "",
                    ):
                        nonempty += 1

                    if isinstance(
                        value,
                        str,
                    ):
                        stripped = value.strip()

                        if stripped:
                            all_text.append(
                                stripped
                            )

                        upper = stripped.upper()

                        if (
                            upper
                            in EXCEL_ERROR_VALUES
                            or (
                                stripped.startswith(
                                    "="
                                )
                                and "#REF!"
                                in upper
                            )
                        ):
                            errors.append(
                                f"{ws.title}!"
                                f"{cell.coordinate}="
                                f"{stripped}"
                            )

                    if (
                        getattr(
                            cell,
                            "data_type",
                            None,
                        )
                        == "e"
                    ):
                        errors.append(
                            f"{ws.title}!"
                            f"{cell.coordinate}="
                            f"{value}"
                        )

        if errors:
            raise RuntimeError(
                "Lỗi Excel: "
                + "; ".join(
                    errors[:10]
                )
            )

        if nonempty < 10:
            raise RuntimeError(
                f"Chỉ có {nonempty} ô dữ liệu."
            )

        normalized_texts = [
            norm_text(
                text
            )
            for text in all_text
        ]

        school_norm = norm_text(
            school["name"]
        )

        has_school = any(
            (
                school_norm
                and school_norm
                in text
            )
            for text
            in normalized_texts
        )

        has_toan_tinh = any(
            "TOAN TINH"
            in text
            for text
            in normalized_texts
        )

        has_school_signature = any(
            (
                "HIEU TRUONG"
                in text
                or "NHA TRUONG"
                in text
            )
            for text
            in normalized_texts
        )

        has_year = (
            not reference_year
            or any(
                reference_year
                in str(
                    text
                )
                for text
                in all_text
            )
        )

        first_ws = workbook[
            workbook.sheetnames[0]
        ]

        blank_results = {}

        for ref in spec[
            "blank_conclusion"
        ]:
            value = normalize_cell_value(
                first_ws[
                    ref
                ].value
            )

            blank_results[
                ref
            ] = value

            if value not in (
                None,
                "",
            ):
                raise RuntimeError(
                    f"Scope trường nhưng "
                    f"ô kết luận {ref}="
                    f"{value!r}; phải trống."
                )

        percent_formats = {}

        for ref in spec[
            "percent_formats"
        ]:
            fmt = str(
                first_ws[
                    ref
                ].number_format
                or ""
            )

            percent_formats[
                ref
            ] = fmt

            if (
                fmt
                != EXPECTED_PERCENT_FORMAT
            ):
                raise RuntimeError(
                    f"{ref} format={fmt!r}; "
                    f"yêu cầu="
                    f"{EXPECTED_PERCENT_FORMAT!r}."
                )

        warnings = []

        if not has_school:
            warnings.append(
                "Không tìm thấy tên trường "
                f"{school['name']!r} "
                "trong vùng quét."
            )

        if has_toan_tinh:
            warnings.append(
                "Phát hiện chữ 'Toàn tỉnh' "
                "trong file cấp Trường."
            )

        if not has_school_signature:
            warnings.append(
                "Không phát hiện chữ "
                "'Hiệu trưởng'/'Nhà trường' "
                "trong vùng quét."
            )

        if not has_year:
            warnings.append(
                f"Không tìm thấy năm "
                f"{reference_year}."
            )

        if truncated:
            warnings.append(
                "Workbook lớn; quét tối đa "
                "250 hàng x 120 cột/sheet."
            )

        return {
            "sheetnames": tuple(
                workbook.sheetnames
            ),
            "nonempty": nonempty,
            "has_school": has_school,
            "has_toan_tinh": has_toan_tinh,
            "has_school_signature": (
                has_school_signature
            ),
            "has_year": has_year,
            "blank_conclusions": (
                blank_results
            ),
            "percent_formats": (
                percent_formats
            ),
            "warnings": warnings,
        }

    finally:
        workbook.close()


def export_one(
    spec: dict,
    *,
    db,
    request,
    school_year,
    school: dict,
    mn_builder,
    xmc_builder,
    report_center,
) -> bytes:
    if (
        spec["exporter"]
        == "mn"
    ):
        response = (
            call_with_known_args(
                mn_builder
                .export_mn_report,
                db=db,
                request=request,
                school_year=(
                    school_year
                ),
                school=school,
                report_type=(
                    spec["code"]
                ),
            )
        )

    elif (
        spec["exporter"]
        == "xmc"
    ):
        response = (
            call_with_known_args(
                xmc_builder
                .export_additional_report,
                db=db,
                request=request,
                school_year=(
                    school_year
                ),
                school=school,
                report_type=(
                    spec["code"]
                ),
            )
        )

    elif (
        spec["exporter"]
        == "th_m1"
    ):
        exporter = getattr(
            report_center,
            "_export_primary_th_m1",
            None,
        )

        if not callable(
            exporter
        ):
            raise RuntimeError(
                "Không tìm thấy "
                "_export_primary_th_m1."
            )

        response = (
            call_with_known_args(
                exporter,
                db=db,
                request=request,
                school_year=(
                    school_year
                ),
                school=school,
                report_type=None,
            )
        )

    else:
        raise RuntimeError(
            "Exporter không biết."
        )

    return asyncio.run(
        response_to_bytes(
            response
        )
    )


def csv_escape(
    value: Any,
) -> str:
    text = str(
        value
        or ""
    )

    if any(
        ch in text
        for ch in (
            ",",
            '"',
            "\n",
            "\r",
        )
    ):
        return (
            '"'
            + text.replace(
                '"',
                '""',
            )
            + '"'
        )

    return text


def main() -> int:
    print(
        "=" * 116
    )
    print(
        "KIỂM THỬ TỔNG THỂ "
        "HỆ THỐNG BÁO CÁO - CẤP TRƯỜNG"
    )
    print(
        "=" * 116
    )
    print()
    print(
        "Phạm vi: 1 trường MN + "
        "1 trường TH + 1 trường THCS."
    )
    print(
        "Tổng cộng 14 báo cáo đúng cấp học."
    )
    print(
        "Chỉ đọc database; "
        "không sửa source."
    )
    print()

    for path in (
        DB_PATH,
        MN_BUILDER_PATH,
        XMC_BUILDER_PATH,
        REPORT_CENTER_PATH,
        RULES_PATH,
    ):
        if not path.exists():
            print(
                "[LỖI] Không tìm thấy: "
                + str(
                    path
                )
            )
            return 1

    db_hash_before = (
        sha256_file(
            DB_PATH
        )
    )

    state_before = (
        db_state()
    )

    try:
        verify_db_state(
            state_before
        )

        verify_source_files()

    except Exception as exc:
        print(
            "[LỖI PREFLIGHT] "
            + str(
                exc
            )
        )
        return 1

    print(
        "[1/9] Database: "
        f"{state_before['communes']} xã/phường; "
        f"{state_before['schools']} trường; "
        "integrity=ok; FK=0."
    )

    print(
        "[2/9] Source: AST OK; "
        "scope xã/tỉnh + format % "
        "còn đủ marker."
    )

    old_cwd = Path.cwd()
    results = []

    try:
        os.chdir(
            PROJECT
        )

        sys.path.insert(
            0,
            str(
                PROJECT
            ),
        )

        from sqlalchemy import text
        from app.database import SessionLocal

        import app.pcgd_mn_report_builders_v1 as mn_builder
        import app.pcgd_xmc_report_builders_v1 as xmc_builder
        import app.routers.report_center as report_center

        db = SessionLocal()

        try:
            db.execute(
                text(
                    "PRAGMA query_only = ON"
                )
            )

            school_year = (
                choose_school_year(
                    db
                )
            )

            school_year_id = int(
                school_year.id
            )

            school_year_code = str(
                getattr(
                    school_year,
                    "code",
                    "",
                )
                or ""
            )

            reference_year = (
                expected_reference_year(
                    school_year
                )
            )

            print(
                "[3/9] Năm học: "
                f"{school_year_code} "
                f"(ID={school_year_id})."
            )

            (
                mn_school,
                th_school,
                thcs_school,
            ) = choose_test_schools(
                db,
                school_year_id=(
                    school_year_id
                ),
                xmc_builder=(
                    xmc_builder
                ),
            )

            schools = (
                mn_school,
                th_school,
                thcs_school,
            )

            print(
                "[4/9] Trường tự chọn:"
            )

            for school in schools:
                print(
                    "      - "
                    f"{school['level']}: "
                    f"ID={school['id']} | "
                    f"{school['name']} | "
                    f"commune_id="
                    f"{school['commune_id']} | "
                    f"lớp="
                    f"{school['class_count']} | "
                    f"record_school_id="
                    f"{school['person_count']}"
                )

            OUT_DIR.mkdir(
                parents=True,
                exist_ok=False,
            )

            print(
                "[5/9] Xuất và kiểm "
                "14 file Excel..."
            )

            global_index = 0

            for school in schools:
                level = school[
                    "level"
                ]

                specs = REPORTS_BY_LEVEL[
                    level
                ]

                request = (
                    make_school_request(
                        school
                    )
                )

                school_dir = (
                    OUT_DIR
                    / (
                        level
                        + "_"
                        + safe_filename(
                            school["name"]
                        )
                    )
                )

                school_dir.mkdir(
                    parents=True,
                    exist_ok=True,
                )

                print()
                print(
                    "      >>> "
                    f"{level} / "
                    f"{school['name']} "
                    f"(ID={school['id']})"
                )

                for local_index, spec in enumerate(
                    specs,
                    start=1,
                ):
                    global_index += 1

                    row = {
                        "level": level,
                        "school_id": (
                            school["id"]
                        ),
                        "school_name": (
                            school["name"]
                        ),
                        "commune_id": (
                            school[
                                "commune_id"
                            ]
                        ),
                        "group": (
                            spec["group"]
                        ),
                        "short": (
                            spec["short"]
                        ),
                        "code": (
                            spec["code"]
                        ),
                        "status": "FAIL",
                        "message": "",
                        "file": "",
                        "warnings": [],
                        "details": {},
                    }

                    print(
                        "          "
                        f"[{local_index:02d}/"
                        f"{len(specs):02d}] "
                        f"{spec['short']} ..."
                    )

                    try:
                        data = export_one(
                            spec,
                            db=db,
                            request=request,
                            school_year=(
                                school_year
                            ),
                            school=school,
                            mn_builder=(
                                mn_builder
                            ),
                            xmc_builder=(
                                xmc_builder
                            ),
                            report_center=(
                                report_center
                            ),
                        )

                        if len(
                            data
                        ) < 1000:
                            raise RuntimeError(
                                "File quá nhỏ: "
                                f"{len(data)} bytes."
                            )

                        file_path = (
                            school_dir
                            / (
                                f"{global_index:02d}_"
                                f"{safe_filename(spec['short'])}_"
                                f"{school_year_code}_"
                                f"TRUONG_{school['id']}.xlsx"
                            )
                        )

                        file_path.write_bytes(
                            data
                        )

                        details = workbook_scan(
                            data,
                            spec,
                            school,
                            reference_year,
                        )

                        row[
                            "file"
                        ] = str(
                            file_path
                        )

                        row[
                            "details"
                        ] = details

                        row[
                            "warnings"
                        ] = list(
                            details[
                                "warnings"
                            ]
                        )

                        if row[
                            "warnings"
                        ]:
                            row[
                                "status"
                            ] = "WARN"

                            row[
                                "message"
                            ] = (
                                "Xuất OK; "
                                f"{len(row['warnings'])} "
                                "cảnh báo."
                            )

                        else:
                            row[
                                "status"
                            ] = "PASS"

                            row[
                                "message"
                            ] = (
                                "Xuất OK; "
                                "kiểm tra bắt buộc đạt."
                            )

                        print(
                            "                 "
                            f"[{row['status']}] "
                            f"{row['message']}"
                        )

                        if details[
                            "blank_conclusions"
                        ]:
                            print(
                                "                 "
                                "Kết luận trường: "
                                + ", ".join(
                                    f"{k}={v!r}"
                                    for k, v
                                    in details[
                                        "blank_conclusions"
                                    ].items()
                                )
                                + " (đúng: phải trống)"
                            )

                        if details[
                            "percent_formats"
                        ]:
                            print(
                                "                 "
                                "Format %: "
                                + ", ".join(
                                    f"{k}={v}"
                                    for k, v
                                    in details[
                                        "percent_formats"
                                    ].items()
                                )
                            )

                        for warning in row[
                            "warnings"
                        ]:
                            print(
                                "                 "
                                "[WARN] "
                                + warning
                            )

                    except Exception as exc:
                        row[
                            "status"
                        ] = "FAIL"

                        row[
                            "message"
                        ] = str(
                            exc
                        )

                        print(
                            "                 "
                            "[FAIL] "
                            + str(
                                exc
                            )
                        )

                    results.append(
                        row
                    )

        finally:
            try:
                db.rollback()
            except Exception:
                pass

            db.close()

        print()
        print(
            "[6/9] Đã chạy đủ "
            f"{len(results)}/14 đầu ra."
        )

        state_after = (
            db_state()
        )

        verify_db_state(
            state_after
        )

        db_hash_after = (
            sha256_file(
                DB_PATH
            )
        )

        if (
            state_after
            != state_before
        ):
            raise RuntimeError(
                "Trạng thái DB trước/sau "
                "khác nhau."
            )

        if (
            db_hash_after
            != db_hash_before
        ):
            raise RuntimeError(
                "SHA256 DB thay đổi. "
                "Hãy dừng Uvicorn "
                "rồi chạy lại."
            )

        print(
            "[7/9] Database: "
            "KHÔNG THAY ĐỔI; "
            "SHA256 giữ nguyên."
        )

        pass_count = sum(
            row["status"]
            == "PASS"
            for row
            in results
        )

        warn_count = sum(
            row["status"]
            == "WARN"
            for row
            in results
        )

        fail_count = sum(
            row["status"]
            == "FAIL"
            for row
            in results
        )

        report_lines = [
            "=" * 116,
            "KIỂM THỬ TỔNG THỂ "
            "HỆ THỐNG BÁO CÁO - CẤP TRƯỜNG",
            "=" * 116,
            "",
            f"Năm học: {school_year_code}",
            "",
            "TRƯỜNG KIỂM THỬ:",
        ]

        for school in schools:
            report_lines.append(
                " - "
                f"{school['level']}: "
                f"ID={school['id']} | "
                f"{school['name']} | "
                f"commune_id="
                f"{school['commune_id']} | "
                f"classes="
                f"{school['class_count']} | "
                f"person_school_id="
                f"{school['person_count']}"
            )

        report_lines.extend(
            [
                "",
                "DATABASE:",
                (
                    " - SHA256 trước: "
                    f"{db_hash_before}"
                ),
                (
                    " - SHA256 sau  : "
                    f"{db_hash_after}"
                ),
                " - Database: KHÔNG THAY ĐỔI.",
                "",
                "KẾT QUẢ TỔNG:",
                f" - PASS: {pass_count}",
                f" - WARN: {warn_count}",
                f" - FAIL: {fail_count}",
                "",
            ]
        )

        for school in schools:
            subset = [
                row
                for row
                in results
                if row[
                    "school_id"
                ]
                == school[
                    "id"
                ]
            ]

            p = sum(
                row["status"]
                == "PASS"
                for row
                in subset
            )

            w = sum(
                row["status"]
                == "WARN"
                for row
                in subset
            )

            f = sum(
                row["status"]
                == "FAIL"
                for row
                in subset
            )

            report_lines.append(
                f"{school['level']} / "
                f"{school['name']}: "
                f"PASS={p}; "
                f"WARN={w}; "
                f"FAIL={f}"
            )

        report_lines.extend(
            [
                "",
                "=" * 116,
                "CHI TIẾT 14 ĐẦU RA",
                "=" * 116,
            ]
        )

        for index, row in enumerate(
            results,
            start=1,
        ):
            report_lines.extend(
                [
                    "",
                    (
                        f"[{index:02d}] "
                        f"{row['level']} / "
                        f"{row['school_name']} / "
                        f"{row['short']}"
                    ),
                    (
                        "Trạng thái: "
                        f"{row['status']}"
                    ),
                    (
                        "Thông tin : "
                        f"{row['message']}"
                    ),
                    (
                        "File      : "
                        f"{row['file'] or '(không tạo được)'}"
                    ),
                ]
            )

            details = row.get(
                "details"
            ) or {}

            if details:
                report_lines.extend(
                    [
                        (
                            "Có tên trường: "
                            f"{details['has_school']}"
                        ),
                        (
                            "Có Toàn tỉnh: "
                            f"{details['has_toan_tinh']}"
                        ),
                        (
                            "Có chữ ký trường: "
                            f"{details['has_school_signature']}"
                        ),
                        (
                            "Có năm: "
                            f"{details['has_year']}"
                        ),
                    ]
                )

                if details[
                    "blank_conclusions"
                ]:
                    report_lines.append(
                        "Ô kết luận scope trường: "
                        + ", ".join(
                            f"{k}={v!r}"
                            for k, v
                            in details[
                                "blank_conclusions"
                            ].items()
                        )
                    )

                if details[
                    "percent_formats"
                ]:
                    report_lines.append(
                        "Format %: "
                        + ", ".join(
                            f"{k}={v!r}"
                            for k, v
                            in details[
                                "percent_formats"
                            ].items()
                        )
                    )

            for warning in row[
                "warnings"
            ]:
                report_lines.append(
                    "WARN: "
                    + warning
                )

        if fail_count:
            final_text = (
                f"CÒN {fail_count} FAIL. "
                "CHƯA CHỐT KIỂM THỬ "
                "TỔNG THỂ BÁO CÁO."
            )
        elif warn_count:
            final_text = (
                "KHÔNG CÓ FAIL, "
                f"NHƯNG CÒN {warn_count} WARN "
                "CẦN RÀ SOÁT."
            )
        else:
            final_text = (
                "14/14 ĐẦU RA PASS. "
                "CẤP TRƯỜNG ĐẠT."
            )

        report_lines.extend(
            [
                "",
                "=" * 116,
                "KẾT LUẬN",
                "=" * 116,
                final_text,
            ]
        )

        REPORT_PATH.write_text(
            "\n".join(
                report_lines
            ),
            encoding="utf-8-sig",
        )

        csv_rows = [
            (
                "STT,CAP_HOC,SCHOOL_ID,"
                "TEN_TRUONG,COMMUNE_ID,"
                "BAO_CAO,REPORT_CODE,"
                "TRANG_THAI,THONG_TIN,FILE"
            )
        ]

        for index, row in enumerate(
            results,
            start=1,
        ):
            csv_rows.append(
                ",".join(
                    (
                        str(
                            index
                        ),
                        csv_escape(
                            row[
                                "level"
                            ]
                        ),
                        str(
                            row[
                                "school_id"
                            ]
                        ),
                        csv_escape(
                            row[
                                "school_name"
                            ]
                        ),
                        str(
                            row[
                                "commune_id"
                            ]
                        ),
                        csv_escape(
                            row[
                                "short"
                            ]
                        ),
                        csv_escape(
                            row[
                                "code"
                            ]
                        ),
                        csv_escape(
                            row[
                                "status"
                            ]
                        ),
                        csv_escape(
                            row[
                                "message"
                            ]
                        ),
                        csv_escape(
                            row[
                                "file"
                            ]
                        ),
                    )
                )
            )

        SUMMARY_CSV.write_text(
            "\n".join(
                csv_rows
            ),
            encoding="utf-8-sig",
        )

        print(
            "[8/9] Tổng hợp: "
            f"PASS={pass_count}; "
            f"WARN={warn_count}; "
            f"FAIL={fail_count}."
        )

        print(
            f"[9/9] Báo cáo: "
            f"{REPORT_PATH}"
        )

        print(
            f"      CSV    : "
            f"{SUMMARY_CSV}"
        )

        print()
        print(
            "=" * 116
        )

        print(
            final_text
        )

        print(
            "Database: KHÔNG THAY ĐỔI."
        )

        print(
            f"Thư mục Excel: "
            f"{OUT_DIR}"
        )

        print(
            "=" * 116
        )

        return (
            1
            if fail_count
            else 0
        )

    except Exception as exc:
        print()
        print(
            "=" * 116
        )

        print(
            "[LỖI HỆ THỐNG "
            "KIỂM THỬ CẤP TRƯỜNG]"
        )

        print(
            str(
                exc
            )
        )

        print(
            "=" * 116
        )

        try:
            OUT_DIR.mkdir(
                parents=True,
                exist_ok=True,
            )

            error_path = (
                OUT_DIR
                / (
                    "loi_he_thong_"
                    f"kiem_thu_cap_truong_"
                    f"{STAMP}.txt"
                )
            )

            error_path.write_text(
                (
                    f"Lỗi: {exc}\n\n"
                    + traceback.format_exc()
                ),
                encoding="utf-8-sig",
            )

            print(
                "Chi tiết lỗi: "
                f"{error_path}"
            )

        except Exception:
            pass

        print(
            "Script không có lệnh "
            "UPDATE/INSERT/DELETE."
        )

        return 2

    finally:
        os.chdir(
            old_cwd
        )


if __name__ == "__main__":
    raise SystemExit(
        main()
    )
