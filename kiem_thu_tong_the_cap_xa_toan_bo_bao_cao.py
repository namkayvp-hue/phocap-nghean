# -*- coding: utf-8 -*-
"""
KIỂM THỬ TỔNG THỂ HỆ THỐNG BÁO CÁO - CẤP XÃ

KHÔNG GẮN SỐ BÀI MỚI.

MỤC TIÊU
--------
Kiểm thử cùng 18 đầu ra Excel đã PASS ở cấp Sở, nhưng chuyển sang
phạm vi tài khoản XÃ/PHƯỜNG.

Tự chọn 2 địa bàn trong năm học có nhiều dữ liệu nhất:
1) Một xã/phường BÌNH THƯỜNG có nhiều bản ghi điều tra nhất.
2) Một xã/phường ĐẶC BIỆT KHÓ KHĂN có nhiều bản ghi điều tra nhất.

Hiện dữ liệu ĐBKK có thể bằng 0; vẫn phải kiểm:
- đường xuất;
- phạm vi xã;
- chữ ký cấp xã;
- kết luận "Chưa đủ dữ liệu" khi thiếu dữ liệu;
- định dạng %;
- file XLSX;
- database không đổi.

Tổng đầu ra: 18 biểu x 2 xã = 36 file Excel.

AN TOÀN
-------
- Không sửa source.
- Không UPDATE / INSERT / DELETE database.
- SQLAlchemy: PRAGMA query_only = ON.
- integrity_check + foreign_key_check.
- SHA256 database trước/sau phải giữ nguyên.
"""

from __future__ import annotations

import asyncio
import ast
import hashlib
import inspect
import os
import re
import sqlite3
import sys
import traceback
import unicodedata
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any


PROJECT = Path(
    os.environ.get(
        "PHOCAP_PROJECT",
        r"C:\PhoCap",
    )
).resolve()

APP = PROJECT / "app"
DB_PATH = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"

MN_BUILDER_PATH = APP / "pcgd_mn_report_builders_v1.py"
XMC_BUILDER_PATH = APP / "pcgd_xmc_report_builders_v1.py"
REPORT_CENTER_PATH = APP / "routers" / "report_center.py"
RULES_PATH = APP / "services" / "pcgd_business_rules.py"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

OUT_DIR = EXPORTS / f"kiem_thu_tong_the_cap_xa_{STAMP}"
REPORT_PATH = OUT_DIR / f"bao_cao_kiem_thu_tong_the_cap_xa_{STAMP}.txt"
SUMMARY_CSV = OUT_DIR / f"tong_hop_kiem_thu_cap_xa_{STAMP}.csv"

EXPECTED_COMMUNES = 130
EXPECTED_SPECIAL = 51
EXPECTED_PERCENT_FORMAT = r"0\%"

EXCEL_ERROR_VALUES = {
    "#REF!",
    "#DIV/0!",
    "#VALUE!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
}


REPORT_SPECS = (
    # MẦM NON
    {
        "group": "MẦM NON",
        "short": "MN-M1",
        "code": "PCGD_MN_M1_2025",
        "exporter": "mn",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "MẦM NON",
        "short": "MN-02",
        "code": "PCGD_MN_02_2025",
        "exporter": "mn",
        "conclusion": {
            "R7": {
                "Đạt",
                "Không đạt",
                "Chưa đủ dữ liệu",
            },
        },
        "percent_formats": (
            "H7",
            "K7",
            "O7",
        ),
    },
    {
        "group": "MẦM NON",
        "short": "MN-01-GV",
        "code": "PCGD_MN_01_GV_2025",
        "exporter": "mn",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "MẦM NON",
        "short": "MN-01-CSVC",
        "code": "PCGD_MN_01_CSVC_2025",
        "exporter": "mn",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "MẦM NON",
        "short": "MN-TÀI-CHÍNH",
        "code": "PCGD_MN_TAICHINH_2025",
        "exporter": "mn",
        "conclusion": {},
        "percent_formats": (),
    },

    # TIỂU HỌC
    {
        "group": "TIỂU HỌC",
        "short": "TH-M1",
        "code": "PCGD_TH_M1_2025",
        "exporter": "th_m1",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "TIỂU HỌC",
        "short": "TH-02",
        "code": "PCGD_TH_02_2025",
        "exporter": "xmc",
        "conclusion": {
            "T8": {
                1,
                2,
                3,
                "Không đạt",
                "Chưa đủ dữ liệu",
            },
        },
        "percent_formats": (
            "I8",
            "K8",
            "M8",
            "Q8",
        ),
    },
    {
        "group": "TIỂU HỌC",
        "short": "TH-01-GV",
        "code": "PCGD_TH_01_GV_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "TIỂU HỌC",
        "short": "TH-01-CSVC",
        "code": "PCGD_TH_01_CSVC_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },

    # THCS
    {
        "group": "THCS",
        "short": "THCS-M1",
        "code": "PCGD_THCS_M1_2025",
        "exporter": "xmc",
        "conclusion": {
            "G36": {
                1,
                2,
                3,
                "Không đạt",
                "Chưa đủ dữ liệu",
            },
            "G37": {
                1,
                2,
                "Không đạt",
                "Chưa đủ dữ liệu",
            },
        },
        "percent_formats": (
            "H39",
            "H40",
            "H41",
        ),
    },
    {
        "group": "THCS",
        "short": "THCS-M2",
        "code": "PCGD_THCS_M2_2025",
        "exporter": "xmc",
        "conclusion": {
            "W14": {
                1,
                2,
                3,
                "Không đạt",
                "Chưa đủ dữ liệu",
            },
        },
        "percent_formats": (
            "E14",
            "J14",
            "M14",
            "Q14",
            "V14",
            "E15",
            "J15",
            "M15",
            "Q15",
            "V15",
            "K20",
            "K21",
            "K22",
            "K23",
            "K24",
        ),
    },
    {
        "group": "THCS",
        "short": "THCS-TK",
        "code": "PCGD_THCS_TK_2025",
        "exporter": "xmc",
        "conclusion": {
            "F8": {
                1,
                2,
                3,
                "Không đạt",
                "Chưa đủ dữ liệu",
            },
            "G8": {
                1,
                2,
                "Không đạt",
                "Chưa đủ dữ liệu",
            },
            "R8": {
                1,
                2,
                3,
                "Không đạt",
                "Chưa đủ dữ liệu",
            },
        },
        "percent_formats": (
            "I8",
            "K8",
            "O8",
        ),
    },
    {
        "group": "THCS",
        "short": "THCS-M5",
        "code": "PCGD_THCS_M5_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "THCS",
        "short": "THCS-CSVC",
        "code": "PCGD_THCS_CSVC_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },

    # XÓA MÙ CHỮ
    {
        "group": "XÓA MÙ CHỮ",
        "short": "XMC-3",
        "code": "PCGD_XMC_3_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "XÓA MÙ CHỮ",
        "short": "CMC-2",
        "code": "PCGD_CMC_2_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "XÓA MÙ CHỮ",
        "short": "CMC-1",
        "code": "PCGD_CMC_1_2025",
        "exporter": "xmc",
        "conclusion": {},
        "percent_formats": (),
    },
    {
        "group": "XÓA MÙ CHỮ",
        "short": "XMC-4",
        "code": "PCGD_XMC_4_2025",
        "exporter": "xmc",
        "conclusion": {
            "R8": {
                1,
                2,
                "Không đạt",
                "Chưa đủ dữ liệu",
            },
        },
        "percent_formats": (),
    },
)


def sha256_file(
    path: Path,
) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(
                1024 * 1024
            ),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


def norm_text(
    value: Any,
) -> str:
    text = str(
        value
        or ""
    ).strip().upper()

    if not text:
        return ""

    text = unicodedata.normalize(
        "NFD",
        text,
    )

    text = "".join(
        ch
        for ch in text
        if unicodedata.category(ch)
        != "Mn"
    )

    return " ".join(
        text.split()
    )


def safe_filename(
    value: str,
) -> str:
    text = norm_text(
        value
    )

    text = re.sub(
        r"[^A-Z0-9_-]+",
        "_",
        text,
    )

    return (
        text.strip("_")
        or "BAO_CAO"
    )


def normalize_cell_value(
    value: Any,
) -> Any:
    if value is None:
        return None

    if isinstance(
        value,
        bool,
    ):
        return value

    if isinstance(
        value,
        int,
    ):
        return value

    if isinstance(
        value,
        float,
    ):
        if value.is_integer():
            return int(value)

        return value

    return " ".join(
        str(value)
        .strip()
        .split()
    )


def db_state() -> dict[str, Any]:
    uri = (
        DB_PATH.resolve().as_uri()
        + "?mode=ro"
    )

    con = sqlite3.connect(
        uri,
        uri=True,
    )

    try:
        integrity = str(
            con.execute(
                "PRAGMA integrity_check"
            ).fetchone()[0]
        )

        fk_rows = con.execute(
            "PRAGMA foreign_key_check"
        ).fetchall()

        total = int(
            con.execute(
                "SELECT COUNT(*) "
                "FROM communes "
                "WHERE is_active=1"
            ).fetchone()[0]
        )

        special = int(
            con.execute(
                """
                SELECT COUNT(*)
                FROM communes
                WHERE is_active=1
                  AND COALESCE(
                    is_special_difficulty_area,
                    0
                  )=1
                """
            ).fetchone()[0]
        )

        return {
            "integrity": integrity,
            "fk_count": len(
                fk_rows
            ),
            "total": total,
            "special": special,
        }

    finally:
        con.close()


def verify_db_state(
    state: dict[str, Any],
) -> None:
    if (
        str(
            state["integrity"]
        ).lower()
        != "ok"
    ):
        raise RuntimeError(
            "integrity_check != ok"
        )

    if int(
        state["fk_count"]
    ) != 0:
        raise RuntimeError(
            "foreign_key_check có lỗi."
        )

    if int(
        state["total"]
    ) != EXPECTED_COMMUNES:
        raise RuntimeError(
            f"Xã hoạt động={state['total']}; "
            f"yêu cầu={EXPECTED_COMMUNES}."
        )

    if int(
        state["special"]
    ) != EXPECTED_SPECIAL:
        raise RuntimeError(
            f"ĐBKK={state['special']}; "
            f"yêu cầu={EXPECTED_SPECIAL}."
        )


def verify_source_files() -> None:
    for path in (
        MN_BUILDER_PATH,
        XMC_BUILDER_PATH,
        REPORT_CENTER_PATH,
        RULES_PATH,
    ):
        if not path.exists():
            raise RuntimeError(
                f"Không tìm thấy source: {path}"
            )

        ast.parse(
            path.read_text(
                encoding="utf-8-sig",
                errors="strict",
            )
        )

    mn = MN_BUILDER_PATH.read_text(
        encoding="utf-8-sig",
    )

    xmc = XMC_BUILDER_PATH.read_text(
        encoding="utf-8-sig",
    )

    rules = RULES_PATH.read_text(
        encoding="utf-8-sig",
    )

    rc = REPORT_CENTER_PATH.read_text(
        encoding="utf-8-sig",
    )

    required = (
        (
            mn,
            "BAI_13B_11_15_2_4_9_1_MN_HELPERS_START",
        ),
        (
            mn,
            "BAI_13B_11_15_2_4_9_2A_1_MN_HELPERS_START",
        ),
        (
            xmc,
            "BAI_13B_11_15_2_4_9_1_XMC_HELPERS_START",
        ),
        (
            xmc,
            "BAI_13B_11_15_2_4_9_2A_1_XMC_HELPERS_START",
        ),
        (
            rules,
            "evaluate_preschool_3_5_commune",
        ),
        (
            rules,
            "evaluate_primary_commune",
        ),
        (
            rules,
            "evaluate_literacy_commune",
        ),
        (
            rules,
            "evaluate_thcs_commune",
        ),
        (
            rc,
            "def _export_primary_th_m1(",
        ),
    )

    missing = [
        token
        for source, token
        in required
        if token not in source
    ]

    if missing:
        raise RuntimeError(
            "Source thiếu nền cấp Xã: "
            + "; ".join(
                missing
            )
        )


def make_xa_request(
    commune: dict,
):
    from starlette.requests import Request

    scope = {
        "type": "http",
        "asgi": {
            "version": "3.0",
            "spec_version": "2.3",
        },
        "http_version": "1.1",
        "server": (
            "127.0.0.1",
            80,
        ),
        "client": (
            "127.0.0.1",
            0,
        ),
        "scheme": "http",
        "method": "GET",
        "root_path": "",
        "path": "/bao-cao",
        "raw_path": b"/bao-cao",
        "query_string": b"",
        "headers": [],
        "auth_user": {
            "id": 0,
            "username": (
                "kiem_thu_cap_xa"
            ),
            "full_name": (
                "Kiểm thử tổng thể cấp Xã"
            ),
            "role_code": "XA",
            "role_name": "Xã/phường",
            "commune_id": int(
                commune["id"]
            ),
            "school_id": None,
            "unit_name": str(
                commune["name"]
            ),
        },
    }

    return Request(
        scope
    )


async def response_to_bytes(
    response,
) -> bytes:
    body = getattr(
        response,
        "body",
        None,
    )

    if isinstance(
        body,
        (
            bytes,
            bytearray,
        ),
    ) and body:
        return bytes(
            body
        )

    iterator = getattr(
        response,
        "body_iterator",
        None,
    )

    if iterator is None:
        raise RuntimeError(
            "Response không có "
            "body/body_iterator."
        )

    chunks = []

    async for chunk in iterator:
        if isinstance(
            chunk,
            str,
        ):
            chunk = chunk.encode(
                "utf-8"
            )

        chunks.append(
            bytes(
                chunk
            )
        )

    return b"".join(
        chunks
    )


def resolve_result(
    value,
):
    if inspect.isawaitable(
        value
    ):
        return asyncio.run(
            value
        )

    return value


def call_with_known_args(
    func,
    *,
    db,
    request,
    school_year,
    commune: dict,
    report_type: str | None = None,
):
    candidates = {
        "db": db,
        "request": request,
        "school_year": school_year,
        "report_type": report_type,
        "selected_commune_id": int(
            commune["id"]
        ),
        "selected_school_id": None,
        "scope_label": str(
            commune["name"]
        ),
    }

    signature = inspect.signature(
        func
    )

    kwargs = {}

    for (
        name,
        parameter,
    ) in signature.parameters.items():
        if name in candidates:
            if (
                name == "report_type"
                and report_type is None
                and parameter.default
                is not inspect.Parameter.empty
            ):
                continue

            kwargs[name] = (
                candidates[name]
            )
            continue

        if (
            parameter.default
            is not inspect.Parameter.empty
        ):
            continue

        if parameter.kind in (
            inspect.Parameter.VAR_POSITIONAL,
            inspect.Parameter.VAR_KEYWORD,
        ):
            continue

        raise RuntimeError(
            f"{func.__module__}."
            f"{func.__name__}: "
            "tham số bắt buộc chưa biết: "
            + name
        )

    return resolve_result(
        func(
            **kwargs
        )
    )


def choose_school_year(
    db,
):
    from sqlalchemy import (
        select,
        text,
    )

    from app.models import (
        SchoolYear,
    )

    requested = str(
        os.environ.get(
            "PHOCAP_TEST_YEAR",
            "",
        )
        or ""
    ).strip()

    if requested:
        year = db.scalar(
            select(
                SchoolYear
            ).where(
                SchoolYear.code
                == requested
            )
        )

        if year is None:
            raise RuntimeError(
                "Không có năm học "
                + requested
            )

        return year

    rows = db.execute(
        text(
            """
            SELECT
                sy.id,
                sy.code,
                COUNT(spr.id) AS cnt
            FROM school_years sy
            LEFT JOIN
                survey_person_year_records spr
              ON spr.school_year_id=sy.id
            GROUP BY
                sy.id,
                sy.code
            ORDER BY
                cnt DESC,
                sy.code DESC
            """
        )
    ).all()

    if not rows:
        raise RuntimeError(
            "Database chưa có năm học."
        )

    best_id = int(
        rows[0][0]
    )

    year = db.scalar(
        select(
            SchoolYear
        ).where(
            SchoolYear.id
            == best_id
        )
    )

    if year is None:
        raise RuntimeError(
            "Không lấy được năm học."
        )

    return year


def year_record_count(
    db,
    school_year_id: int,
) -> int:
    from sqlalchemy import text

    return int(
        db.execute(
            text(
                """
                SELECT COUNT(*)
                FROM survey_person_year_records
                WHERE school_year_id=:id
                """
            ),
            {
                "id": int(
                    school_year_id
                )
            },
        ).scalar_one()
        or 0
    )


def commune_data_count(
    db,
    commune_id: int,
    school_year_id: int,
) -> int:
    from sqlalchemy import text

    value = db.execute(
        text(
            """
            SELECT COUNT(spr.id)
            FROM
                survey_person_year_records spr
            JOIN
                survey_forms sf
              ON sf.id=spr.survey_form_id
            JOIN
                survey_batches sb
              ON sb.id=sf.survey_batch_id
            WHERE
                sb.commune_id=:commune_id
              AND
                spr.school_year_id=:school_year_id
            """
        ),
        {
            "commune_id": int(
                commune_id
            ),
            "school_year_id": int(
                school_year_id
            ),
        },
    ).scalar_one()

    return int(
        value
        or 0
    )


def choose_test_communes(
    db,
    school_year_id: int,
) -> tuple[dict, dict]:
    from sqlalchemy import select
    from app.models import Commune

    rows = list(
        db.scalars(
            select(
                Commune
            )
            .where(
                Commune.is_active.is_(
                    True
                )
            )
            .order_by(
                Commune.name.asc(),
                Commune.id.asc(),
            )
        ).all()
    )

    if len(rows) != EXPECTED_COMMUNES:
        raise RuntimeError(
            f"Có {len(rows)} xã hoạt động; "
            f"yêu cầu {EXPECTED_COMMUNES}."
        )

    items = []

    for commune in rows:
        item = {
            "id": int(
                commune.id
            ),
            "name": str(
                commune.name
            ),
            "special": bool(
                getattr(
                    commune,
                    "is_special_difficulty_area",
                    False,
                )
            ),
        }

        item["data_count"] = (
            commune_data_count(
                db,
                item["id"],
                int(
                    school_year_id
                ),
            )
        )

        items.append(
            item
        )

    specials = [
        item
        for item in items
        if item["special"]
    ]

    normals = [
        item
        for item in items
        if not item["special"]
    ]

    if len(
        specials
    ) != EXPECTED_SPECIAL:
        raise RuntimeError(
            f"ĐBKK={len(specials)}, "
            f"yêu cầu={EXPECTED_SPECIAL}."
        )

    if not normals:
        raise RuntimeError(
            "Không có xã thường."
        )

    # Có thể ép ID khi cần kiểm lại một xã cụ thể.
    requested_normal = str(
        os.environ.get(
            "PHOCAP_TEST_COMMUNE_ID",
            "",
        )
        or ""
    ).strip()

    requested_special = str(
        os.environ.get(
            "PHOCAP_TEST_SPECIAL_COMMUNE_ID",
            "",
        )
        or ""
    ).strip()

    def choose(
        pool: list[dict],
        requested: str,
        label: str,
    ) -> dict:
        if requested:
            try:
                wanted = int(
                    requested
                )
            except ValueError:
                raise RuntimeError(
                    f"{label}: ID không hợp lệ "
                    f"{requested!r}."
                )

            found = [
                item
                for item in pool
                if item["id"]
                == wanted
            ]

            if len(
                found
            ) != 1:
                raise RuntimeError(
                    f"{label}: không tìm thấy "
                    f"commune_id={wanted} "
                    "đúng nhóm."
                )

            return dict(
                found[0]
            )

        ranked = sorted(
            pool,
            key=lambda x: (
                int(
                    x["data_count"]
                ),
                norm_text(
                    x["name"]
                ),
            ),
            reverse=True,
        )

        return dict(
            ranked[0]
        )

    normal = choose(
        normals,
        requested_normal,
        "Xã thường",
    )

    special = choose(
        specials,
        requested_special,
        "Xã ĐBKK",
    )

    return (
        normal,
        special,
    )


def expected_reference_year(
    school_year,
) -> str:
    code = str(
        getattr(
            school_year,
            "code",
            "",
        )
        or ""
    )

    found = re.findall(
        r"\d{4}",
        code,
    )

    if found:
        return found[0]

    return ""


def workbook_scan(
    data: bytes,
    spec: dict,
    commune: dict,
    reference_year: str,
):
    from openpyxl import (
        load_workbook,
    )

    if not data.startswith(
        b"PK"
    ):
        raise RuntimeError(
            "Không phải XLSX."
        )

    workbook = load_workbook(
        BytesIO(
            data
        ),
        data_only=False,
    )

    try:
        if not workbook.sheetnames:
            raise RuntimeError(
                "Workbook không có sheet."
            )

        all_text = []
        errors = []
        nonempty = 0
        truncated = False

        for ws in workbook.worksheets:
            max_row = min(
                max(
                    int(
                        ws.max_row
                        or 1
                    ),
                    1,
                ),
                250,
            )

            max_col = min(
                max(
                    int(
                        ws.max_column
                        or 1
                    ),
                    1,
                ),
                120,
            )

            if (
                int(
                    ws.max_row
                    or 1
                )
                > max_row
                or int(
                    ws.max_column
                    or 1
                )
                > max_col
            ):
                truncated = True

            for row in ws.iter_rows(
                min_row=1,
                max_row=max_row,
                min_col=1,
                max_col=max_col,
            ):
                for cell in row:
                    value = cell.value

                    if value not in (
                        None,
                        "",
                    ):
                        nonempty += 1

                    if isinstance(
                        value,
                        str,
                    ):
                        stripped = (
                            value.strip()
                        )

                        if stripped:
                            all_text.append(
                                stripped
                            )

                        upper = (
                            stripped.upper()
                        )

                        if (
                            upper
                            in EXCEL_ERROR_VALUES
                            or (
                                stripped.startswith("=")
                                and "#REF!"
                                in upper
                            )
                        ):
                            errors.append(
                                f"{ws.title}!"
                                f"{cell.coordinate}="
                                f"{stripped}"
                            )

                    if (
                        getattr(
                            cell,
                            "data_type",
                            None,
                        )
                        == "e"
                    ):
                        errors.append(
                            f"{ws.title}!"
                            f"{cell.coordinate}="
                            f"{value}"
                        )

        if errors:
            raise RuntimeError(
                "Lỗi Excel: "
                + "; ".join(
                    errors[:10]
                )
            )

        if nonempty < 10:
            raise RuntimeError(
                f"Chỉ có {nonempty} ô dữ liệu."
            )

        normalized_texts = [
            norm_text(
                text
            )
            for text in all_text
        ]

        commune_norm = norm_text(
            commune["name"]
        )

        has_commune = any(
            (
                commune_norm
                and commune_norm
                in text
            )
            for text
            in normalized_texts
        )

        has_toan_tinh = any(
            "TOAN TINH"
            in text
            for text
            in normalized_texts
        )

        has_year = (
            not reference_year
            or any(
                reference_year
                in str(
                    text
                )
                for text
                in all_text
            )
        )

        has_commune_signature = any(
            (
                "UBND XA" in text
                or "UBND PHUONG"
                in text
                or "XA/PHUONG"
                in text
                or "PHO CHU TICH"
                in text
            )
            for text
            in normalized_texts
        )

        first_ws = workbook[
            workbook.sheetnames[0]
        ]

        conclusion_values = {}

        for (
            ref,
            allowed,
        ) in spec[
            "conclusion"
        ].items():
            value = (
                normalize_cell_value(
                    first_ws[
                        ref
                    ].value
                )
            )

            conclusion_values[
                ref
            ] = value

            if value not in allowed:
                raise RuntimeError(
                    f"Kết luận {ref}="
                    f"{value!r} không hợp lệ."
                )

        percent_formats = {}

        for ref in spec[
            "percent_formats"
        ]:
            fmt = str(
                first_ws[
                    ref
                ].number_format
                or ""
            )

            percent_formats[
                ref
            ] = fmt

            if (
                fmt
                != EXPECTED_PERCENT_FORMAT
            ):
                raise RuntimeError(
                    f"{ref} format={fmt!r}; "
                    f"yêu cầu="
                    f"{EXPECTED_PERCENT_FORMAT!r}."
                )

        warnings = []

        if not has_commune:
            warnings.append(
                "Không tìm thấy tên địa bàn "
                f"{commune['name']!r} "
                "trong vùng quét."
            )

        if has_toan_tinh:
            warnings.append(
                "Phát hiện chuỗi 'Toàn tỉnh' "
                "trong file cấp Xã."
            )

        if not has_year:
            warnings.append(
                f"Không tìm thấy năm "
                f"{reference_year}."
            )

        if not has_commune_signature:
            warnings.append(
                "Không phát hiện cụm chữ ký "
                "UBND xã/phường/Phó Chủ tịch."
            )

        if truncated:
            warnings.append(
                "Workbook lớn; quét tối đa "
                "250 hàng x 120 cột/sheet."
            )

        return {
            "sheetnames": tuple(
                workbook.sheetnames
            ),
            "nonempty": nonempty,
            "has_commune": has_commune,
            "has_toan_tinh": has_toan_tinh,
            "has_year": has_year,
            "has_commune_signature": (
                has_commune_signature
            ),
            "conclusions": (
                conclusion_values
            ),
            "percent_formats": (
                percent_formats
            ),
            "warnings": warnings,
        }

    finally:
        workbook.close()


def export_one(
    spec: dict,
    *,
    db,
    request,
    school_year,
    commune: dict,
    mn_builder,
    xmc_builder,
    report_center,
) -> bytes:
    if (
        spec["exporter"]
        == "mn"
    ):
        response = (
            call_with_known_args(
                mn_builder
                .export_mn_report,
                db=db,
                request=request,
                school_year=(
                    school_year
                ),
                commune=commune,
                report_type=(
                    spec["code"]
                ),
            )
        )

    elif (
        spec["exporter"]
        == "xmc"
    ):
        response = (
            call_with_known_args(
                xmc_builder
                .export_additional_report,
                db=db,
                request=request,
                school_year=(
                    school_year
                ),
                commune=commune,
                report_type=(
                    spec["code"]
                ),
            )
        )

    elif (
        spec["exporter"]
        == "th_m1"
    ):
        exporter = getattr(
            report_center,
            "_export_primary_th_m1",
            None,
        )

        if not callable(
            exporter
        ):
            raise RuntimeError(
                "Không tìm thấy "
                "_export_primary_th_m1."
            )

        response = (
            call_with_known_args(
                exporter,
                db=db,
                request=request,
                school_year=(
                    school_year
                ),
                commune=commune,
                report_type=None,
            )
        )

    else:
        raise RuntimeError(
            "Exporter không biết."
        )

    return asyncio.run(
        response_to_bytes(
            response
        )
    )


def csv_escape(
    value: Any,
) -> str:
    text = str(
        value
        or ""
    )

    if any(
        ch in text
        for ch in (
            ",",
            '"',
            "\n",
            "\r",
        )
    ):
        return (
            '"'
            + text.replace(
                '"',
                '""',
            )
            + '"'
        )

    return text


def main() -> int:
    print(
        "=" * 116
    )
    print(
        "KIỂM THỬ TỔNG THỂ "
        "HỆ THỐNG BÁO CÁO - CẤP XÃ"
    )
    print(
        "=" * 116
    )
    print()
    print(
        "Phạm vi: 2 xã/phường "
        "(1 thường + 1 ĐBKK), "
        "18 biểu mỗi xã = 36 file."
    )
    print(
        "Chỉ đọc database; "
        "không sửa source."
    )
    print()

    for path in (
        DB_PATH,
        MN_BUILDER_PATH,
        XMC_BUILDER_PATH,
        REPORT_CENTER_PATH,
        RULES_PATH,
    ):
        if not path.exists():
            print(
                "[LỖI] Không tìm thấy: "
                + str(
                    path
                )
            )
            return 1

    db_hash_before = (
        sha256_file(
            DB_PATH
        )
    )

    state_before = (
        db_state()
    )

    try:
        verify_db_state(
            state_before
        )
        verify_source_files()

    except Exception as exc:
        print(
            "[LỖI PREFLIGHT] "
            + str(
                exc
            )
        )
        return 1

    print(
        "[1/9] Database: "
        f"{state_before['total']} xã/phường; "
        f"ĐBKK={state_before['special']}; "
        "integrity=ok; FK=0."
    )

    print(
        "[2/9] Source: AST OK; "
        "evaluator cấp xã + format % "
        "còn đủ marker."
    )

    old_cwd = Path.cwd()
    results = []

    try:
        os.chdir(
            PROJECT
        )

        sys.path.insert(
            0,
            str(
                PROJECT
            ),
        )

        from sqlalchemy import text
        from app.database import SessionLocal

        import app.pcgd_mn_report_builders_v1 as mn_builder
        import app.pcgd_xmc_report_builders_v1 as xmc_builder
        import app.routers.report_center as report_center

        db = SessionLocal()

        try:
            db.execute(
                text(
                    "PRAGMA query_only = ON"
                )
            )

            school_year = (
                choose_school_year(
                    db
                )
            )

            school_year_id = int(
                school_year.id
            )

            school_year_code = str(
                getattr(
                    school_year,
                    "code",
                    "",
                )
                or ""
            )

            reference_year = (
                expected_reference_year(
                    school_year
                )
            )

            record_count = (
                year_record_count(
                    db,
                    school_year_id,
                )
            )

            print(
                "[3/9] Năm học: "
                f"{school_year_code} "
                f"(ID={school_year_id}); "
                f"toàn tỉnh có "
                f"{record_count} bản ghi."
            )

            (
                normal,
                special,
            ) = choose_test_communes(
                db,
                school_year_id,
            )

            communes = (
                {
                    **normal,
                    "kind": (
                        "BÌNH THƯỜNG"
                    ),
                },
                {
                    **special,
                    "kind": (
                        "ĐBKK"
                    ),
                },
            )

            print(
                "[4/9] Địa bàn tự chọn:"
            )

            for commune in communes:
                print(
                    "      - "
                    f"{commune['kind']}: "
                    f"ID={commune['id']} | "
                    f"{commune['name']} | "
                    f"bản ghi="
                    f"{commune['data_count']}"
                )

            OUT_DIR.mkdir(
                parents=True,
                exist_ok=False,
            )

            print(
                "[5/9] Xuất và kiểm "
                "36 file Excel..."
            )

            running_index = 0

            for commune in communes:
                request = (
                    make_xa_request(
                        commune
                    )
                )

                commune_dir = (
                    OUT_DIR
                    / (
                        safe_filename(
                            commune["kind"]
                        )
                        + "_"
                        + safe_filename(
                            commune["name"]
                        )
                    )
                )

                commune_dir.mkdir(
                    parents=True,
                    exist_ok=True,
                )

                print()
                print(
                    "      >>> "
                    f"{commune['kind']} / "
                    f"{commune['name']} "
                    f"(ID={commune['id']}; "
                    f"records="
                    f"{commune['data_count']})"
                )

                for local_index, spec in enumerate(
                    REPORT_SPECS,
                    start=1,
                ):
                    running_index += 1

                    row = {
                        "commune_kind": (
                            commune["kind"]
                        ),
                        "commune_id": (
                            commune["id"]
                        ),
                        "commune_name": (
                            commune["name"]
                        ),
                        "data_count": (
                            commune[
                                "data_count"
                            ]
                        ),
                        "group": spec[
                            "group"
                        ],
                        "short": spec[
                            "short"
                        ],
                        "code": spec[
                            "code"
                        ],
                        "status": "FAIL",
                        "message": "",
                        "file": "",
                        "warnings": [],
                        "details": {},
                    }

                    print(
                        f"          "
                        f"[{local_index:02d}/18] "
                        f"{spec['group']} / "
                        f"{spec['short']} ..."
                    )

                    try:
                        data = export_one(
                            spec,
                            db=db,
                            request=request,
                            school_year=(
                                school_year
                            ),
                            commune=commune,
                            mn_builder=(
                                mn_builder
                            ),
                            xmc_builder=(
                                xmc_builder
                            ),
                            report_center=(
                                report_center
                            ),
                        )

                        if len(
                            data
                        ) < 1000:
                            raise RuntimeError(
                                "File quá nhỏ: "
                                f"{len(data)} bytes."
                            )

                        file_path = (
                            commune_dir
                            / (
                                f"{local_index:02d}_"
                                f"{safe_filename(spec['short'])}_"
                                f"{school_year_code}_"
                                f"XA_{commune['id']}.xlsx"
                            )
                        )

                        file_path.write_bytes(
                            data
                        )

                        details = workbook_scan(
                            data,
                            spec,
                            commune,
                            reference_year,
                        )

                        row["file"] = str(
                            file_path
                        )

                        row["details"] = (
                            details
                        )

                        row["warnings"] = list(
                            details[
                                "warnings"
                            ]
                        )

                        if row[
                            "warnings"
                        ]:
                            row[
                                "status"
                            ] = "WARN"

                            row[
                                "message"
                            ] = (
                                "Xuất OK; "
                                f"{len(row['warnings'])} "
                                "cảnh báo."
                            )

                        else:
                            row[
                                "status"
                            ] = "PASS"

                            row[
                                "message"
                            ] = (
                                "Xuất OK; "
                                "kiểm tra bắt buộc đạt."
                            )

                        print(
                            "                 "
                            f"[{row['status']}] "
                            f"{row['message']}"
                        )

                        if details[
                            "conclusions"
                        ]:
                            print(
                                "                 "
                                "Kết luận: "
                                + ", ".join(
                                    f"{k}={v!r}"
                                    for k, v
                                    in details[
                                        "conclusions"
                                    ].items()
                                )
                            )

                        for warning in row[
                            "warnings"
                        ]:
                            print(
                                "                 "
                                "[WARN] "
                                + warning
                            )

                    except Exception as exc:
                        row[
                            "status"
                        ] = "FAIL"

                        row[
                            "message"
                        ] = str(
                            exc
                        )

                        print(
                            "                 "
                            "[FAIL] "
                            + str(
                                exc
                            )
                        )

                    results.append(
                        row
                    )

        finally:
            try:
                db.rollback()
            except Exception:
                pass

            db.close()

        print()
        print(
            "[6/9] Đã chạy đủ "
            f"{len(results)}/36 đầu ra."
        )

        state_after = (
            db_state()
        )

        verify_db_state(
            state_after
        )

        db_hash_after = (
            sha256_file(
                DB_PATH
            )
        )

        if (
            state_after
            != state_before
        ):
            raise RuntimeError(
                "Trạng thái DB trước/sau "
                "khác nhau."
            )

        if (
            db_hash_after
            != db_hash_before
        ):
            raise RuntimeError(
                "SHA256 DB thay đổi. "
                "Hãy dừng Uvicorn "
                "rồi chạy lại."
            )

        print(
            "[7/9] Database: "
            "KHÔNG THAY ĐỔI; "
            "SHA256 giữ nguyên."
        )

        pass_count = sum(
            row["status"]
            == "PASS"
            for row
            in results
        )

        warn_count = sum(
            row["status"]
            == "WARN"
            for row
            in results
        )

        fail_count = sum(
            row["status"]
            == "FAIL"
            for row
            in results
        )

        report_lines = [
            "=" * 116,
            "KIỂM THỬ TỔNG THỂ "
            "HỆ THỐNG BÁO CÁO - CẤP XÃ",
            "=" * 116,
            "",
            f"Năm học: {school_year_code}",
            (
                "Toàn tỉnh bản ghi điều tra: "
                f"{record_count}"
            ),
            "",
            "ĐỊA BÀN KIỂM THỬ:",
        ]

        for commune in communes:
            report_lines.append(
                " - "
                f"{commune['kind']}: "
                f"ID={commune['id']} | "
                f"{commune['name']} | "
                f"records="
                f"{commune['data_count']}"
            )

        report_lines.extend(
            [
                "",
                "DATABASE:",
                (
                    " - SHA256 trước: "
                    f"{db_hash_before}"
                ),
                (
                    " - SHA256 sau  : "
                    f"{db_hash_after}"
                ),
                " - Database: KHÔNG THAY ĐỔI.",
                "",
                "KẾT QUẢ TỔNG:",
                f" - PASS: {pass_count}",
                f" - WARN: {warn_count}",
                f" - FAIL: {fail_count}",
                "",
            ]
        )

        for commune in communes:
            subset = [
                row
                for row
                in results
                if row[
                    "commune_id"
                ]
                == commune[
                    "id"
                ]
            ]

            p = sum(
                row["status"]
                == "PASS"
                for row in subset
            )

            w = sum(
                row["status"]
                == "WARN"
                for row in subset
            )

            f = sum(
                row["status"]
                == "FAIL"
                for row in subset
            )

            report_lines.append(
                f"{commune['kind']} / "
                f"{commune['name']}: "
                f"PASS={p}; WARN={w}; FAIL={f}"
            )

        report_lines.extend(
            [
                "",
                "=" * 116,
                "CHI TIẾT 36 ĐẦU RA",
                "=" * 116,
            ]
        )

        for index, row in enumerate(
            results,
            start=1,
        ):
            report_lines.extend(
                [
                    "",
                    (
                        f"[{index:02d}] "
                        f"{row['commune_kind']} / "
                        f"{row['commune_name']} / "
                        f"{row['group']} / "
                        f"{row['short']}"
                    ),
                    (
                        "Trạng thái: "
                        f"{row['status']}"
                    ),
                    (
                        "Thông tin : "
                        f"{row['message']}"
                    ),
                    (
                        "File      : "
                        f"{row['file'] or '(không tạo được)'}"
                    ),
                ]
            )

            details = row.get(
                "details"
            ) or {}

            if details:
                report_lines.extend(
                    [
                        (
                            "Có tên xã: "
                            f"{details['has_commune']}"
                        ),
                        (
                            "Có Toàn tỉnh: "
                            f"{details['has_toan_tinh']}"
                        ),
                        (
                            "Có năm: "
                            f"{details['has_year']}"
                        ),
                        (
                            "Có chữ ký xã: "
                            f"{details['has_commune_signature']}"
                        ),
                    ]
                )

                if details[
                    "conclusions"
                ]:
                    report_lines.append(
                        "Kết luận: "
                        + ", ".join(
                            f"{k}={v!r}"
                            for k, v
                            in details[
                                "conclusions"
                            ].items()
                        )
                    )

                if details[
                    "percent_formats"
                ]:
                    report_lines.append(
                        "Format %: "
                        + ", ".join(
                            f"{k}={v!r}"
                            for k, v
                            in details[
                                "percent_formats"
                            ].items()
                        )
                    )

            for warning in row[
                "warnings"
            ]:
                report_lines.append(
                    "WARN: "
                    + warning
                )

        if fail_count:
            final_text = (
                f"CÒN {fail_count} FAIL. "
                "CHƯA CHUYỂN SANG CẤP TRƯỜNG."
            )
        elif warn_count:
            final_text = (
                "KHÔNG CÓ FAIL, "
                f"NHƯNG CÒN {warn_count} WARN "
                "CẦN RÀ SOÁT."
            )
        else:
            final_text = (
                "36/36 ĐẦU RA PASS. "
                "CẤP XÃ ĐẠT."
            )

        report_lines.extend(
            [
                "",
                "=" * 116,
                "KẾT LUẬN",
                "=" * 116,
                final_text,
            ]
        )

        REPORT_PATH.write_text(
            "\n".join(
                report_lines
            ),
            encoding="utf-8-sig",
        )

        csv_rows = [
            (
                "STT,LOAI_XA,COMMUNE_ID,"
                "TEN_XA,NHOM,BAO_CAO,"
                "REPORT_CODE,TRANG_THAI,"
                "THONG_TIN,FILE"
            )
        ]

        for index, row in enumerate(
            results,
            start=1,
        ):
            csv_rows.append(
                ",".join(
                    (
                        str(
                            index
                        ),
                        csv_escape(
                            row[
                                "commune_kind"
                            ]
                        ),
                        str(
                            row[
                                "commune_id"
                            ]
                        ),
                        csv_escape(
                            row[
                                "commune_name"
                            ]
                        ),
                        csv_escape(
                            row[
                                "group"
                            ]
                        ),
                        csv_escape(
                            row[
                                "short"
                            ]
                        ),
                        csv_escape(
                            row[
                                "code"
                            ]
                        ),
                        csv_escape(
                            row[
                                "status"
                            ]
                        ),
                        csv_escape(
                            row[
                                "message"
                            ]
                        ),
                        csv_escape(
                            row[
                                "file"
                            ]
                        ),
                    )
                )
            )

        SUMMARY_CSV.write_text(
            "\n".join(
                csv_rows
            ),
            encoding="utf-8-sig",
        )

        print(
            "[8/9] Tổng hợp: "
            f"PASS={pass_count}; "
            f"WARN={warn_count}; "
            f"FAIL={fail_count}."
        )

        print(
            f"[9/9] Báo cáo: "
            f"{REPORT_PATH}"
        )

        print(
            f"      CSV    : "
            f"{SUMMARY_CSV}"
        )

        print()
        print(
            "=" * 116
        )

        print(
            final_text
        )

        print(
            "Database: KHÔNG THAY ĐỔI."
        )

        print(
            f"Thư mục Excel: "
            f"{OUT_DIR}"
        )

        print(
            "=" * 116
        )

        return (
            1
            if fail_count
            else 0
        )

    except Exception as exc:
        print()
        print(
            "=" * 116
        )
        print(
            "[LỖI HỆ THỐNG "
            "KIỂM THỬ CẤP XÃ]"
        )
        print(
            str(
                exc
            )
        )
        print(
            "=" * 116
        )

        try:
            OUT_DIR.mkdir(
                parents=True,
                exist_ok=True,
            )

            error_path = (
                OUT_DIR
                / (
                    "loi_he_thong_"
                    f"kiem_thu_cap_xa_"
                    f"{STAMP}.txt"
                )
            )

            error_path.write_text(
                (
                    f"Lỗi: {exc}\n\n"
                    + traceback.format_exc()
                ),
                encoding="utf-8-sig",
            )

            print(
                "Chi tiết lỗi: "
                f"{error_path}"
            )

        except Exception:
            pass

        print(
            "Script không có lệnh "
            "UPDATE/INSERT/DELETE."
        )

        return 2

    finally:
        os.chdir(
            old_cwd
        )


if __name__ == "__main__":
    raise SystemExit(
        main()
    )
