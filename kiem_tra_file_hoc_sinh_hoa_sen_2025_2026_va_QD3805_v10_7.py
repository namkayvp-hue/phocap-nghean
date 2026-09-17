# -*- coding: utf-8 -*-
r"""
DRY-RUN V10.7 - KIỂM TRA FILE hoc_sinh_hoa_sen_2025_2026.xlsx
VÀ ĐỐI CHIẾU VỚI CÁC TRƯỜNG QĐ 3805
================================================================

MỤC TIÊU
--------
Đọc CHỈ ĐỌC file:
    C:\PhoCap\uploads\hoc_sinh_hoa_sen_2025_2026.xlsx

Kiểm tra:
1. Các sheet, dòng tiêu đề, số dòng học sinh.
2. Danh sách tên trường xuất hiện trong file.
3. Số học sinh theo trường / sheet / lớp.
4. Đối chiếu tên trường trong file với toàn bộ trường nguồn + đích QĐ 3805
   lấy trực tiếp từ database và file plan.
5. Kết luận file này có phải nguồn 2025-2026 cho nhóm QĐ 3805 hay chỉ là
   dữ liệu riêng của trường Hoa Sen.
6. Không import, không sửa file, không sửa database.

CHẠY
----
    cd C:\PhoCap
    .\.venv\Scripts\python.exe .\kiem_tra_file_hoc_sinh_hoa_sen_2025_2026_va_QD3805_v10_7.py
"""

from __future__ import annotations

import csv
import json
import re
import sqlite3
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path


ROOT = Path(r"C:\PhoCap")
XLSX = ROOT / "uploads" / "hoc_sinh_hoa_sen_2025_2026.xlsx"
DB_PATH = ROOT / "data" / "phocap.db"
PLAN_FILE = ROOT / "data" / "school_merger_approved_plans.json"


class AuditAbort(RuntimeError):
    pass


def qident(name: str) -> str:
    return '"' + str(name).replace('"', '""') + '"'


def normalize(value) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(
        ch for ch in text
        if unicodedata.category(ch) != "Mn"
    )
    text = text.replace("đ", "d")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def write_csv(path: Path, headers: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(
            f,
            fieldnames=headers,
            extrasaction="ignore",
        )
        w.writeheader()
        w.writerows(rows)


def connect_ro(path: Path) -> sqlite3.Connection:
    uri = path.resolve().as_uri() + "?mode=ro"
    con = sqlite3.connect(uri, uri=True)
    con.row_factory = sqlite3.Row
    return con


def table_exists(con: sqlite3.Connection, table: str) -> bool:
    return con.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
        (table,),
    ).fetchone() is not None


def colset(con: sqlite3.Connection, table: str) -> set[str]:
    return {
        str(r[1])
        for r in con.execute(
            f"PRAGMA table_info({qident(table)})"
        ).fetchall()
    }


def load_qd3805_school_ids() -> list[int]:
    if not PLAN_FILE.exists():
        raise AuditAbort(f"Không tìm thấy {PLAN_FILE}")

    payload = json.loads(
        PLAN_FILE.read_text(encoding="utf-8")
    )

    ids = set()
    plans = []

    for raw in payload.get("plans") or []:
        if not isinstance(raw, dict):
            continue
        pid = str(raw.get("id") or "")
        doc = str(raw.get("document_code") or "")
        if not (pid.startswith("QD3805-") or "3805" in doc):
            continue

        plans.append(raw)

        try:
            ids.add(int(raw.get("target_school_id")))
        except (TypeError, ValueError):
            pass

        for value in raw.get("source_school_ids") or []:
            try:
                ids.add(int(value))
            except (TypeError, ValueError):
                pass

    if len(plans) != 6:
        raise AuditAbort(
            f"QĐ 3805 phải có 6 nhóm, hiện tìm thấy {len(plans)}."
        )

    return sorted(ids)


def get_school_names(
    con: sqlite3.Connection,
    ids: list[int],
) -> dict[int, str]:
    if not table_exists(con, "schools"):
        raise AuditAbort("Không có bảng schools.")

    cs = colset(con, "schools")
    name_col = next(
        (
            c for c in (
                "name", "school_name", "ten_truong"
            )
            if c in cs
        ),
        None,
    )
    if not name_col:
        raise AuditAbort(
            "Không xác định được cột tên trường."
        )

    marks = ",".join("?" for _ in ids)
    rows = con.execute(
        f"SELECT id,{qident(name_col)} AS name "
        f"FROM schools WHERE id IN ({marks}) ORDER BY id",
        ids,
    ).fetchall()

    return {
        int(r["id"]): str(r["name"] or "")
        for r in rows
    }


def simplified_school_name(value: str) -> str:
    """
    Bỏ các tiền tố hành chính phổ biến để so tên trường mềm hơn,
    nhưng không đoán/sửa tên chính thức.
    """
    n = normalize(value)
    prefixes = [
        "truong mam non ",
        "mam non ",
        "truong tieu hoc va thcs ",
        "truong tieu hoc ",
        "tieu hoc va thcs ",
        "tieu hoc ",
        "truong thcs ",
        "thcs ",
    ]
    for p in prefixes:
        if n.startswith(p):
            return n[len(p):].strip()
    return n


def match_qd_school(
    excel_school_name: str,
    qd_names: dict[int, str],
) -> tuple[int | None, str, str]:
    raw_norm = normalize(excel_school_name)
    raw_simple = simplified_school_name(excel_school_name)

    exact = []
    soft = []

    for sid, official in qd_names.items():
        off_norm = normalize(official)
        off_simple = simplified_school_name(official)

        if raw_norm and raw_norm == off_norm:
            exact.append((sid, official))
        elif (
            raw_simple
            and off_simple
            and raw_simple == off_simple
        ):
            soft.append((sid, official))

    if len(exact) == 1:
        return exact[0][0], exact[0][1], "EXACT"
    if len(soft) == 1:
        return soft[0][0], soft[0][1], "SOFT_EXACT_CORE"

    return None, "", "NO_MATCH"


def find_header_row(ws) -> tuple[int | None, dict[str, int]]:
    """
    Tìm dòng có các cột Trường, Họ tên, Ngày sinh, Lớp.
    """
    best_row = None
    best_map = {}
    best_score = -1

    for ridx in range(1, min(int(ws.max_row or 1), 30) + 1):
        values = [
            ws.cell(ridx, c).value
            for c in range(1, min(int(ws.max_column or 1), 80) + 1)
        ]

        header_map = {}
        score = 0

        for cidx, value in enumerate(values, start=1):
            n = normalize(value)

            if n in {"truong", "ten truong"}:
                header_map["school"] = cidx
                score += 5
            elif n in {"ho ten", "ho va ten", "ten hoc sinh"}:
                header_map["student_name"] = cidx
                score += 5
            elif n in {"ngay sinh"}:
                header_map["birth_date"] = cidx
                score += 2
            elif n in {"lop"}:
                if "class" not in header_map:
                    header_map["class"] = cidx
                    score += 2
            elif n in {"nhom lop", "nhom/lop"}:
                header_map["group_class"] = cidx
            elif "ma dinh danh" in n or n == "ma hoc sinh":
                header_map["student_code"] = cidx
                score += 2
            elif n == "gioi tinh":
                header_map["gender"] = cidx

        if score > best_score:
            best_score = score
            best_row = ridx
            best_map = header_map

    if (
        best_row is None
        or "school" not in best_map
        or "student_name" not in best_map
    ):
        return None, {}

    return best_row, best_map


def main() -> None:
    print("=" * 124)
    print("DRY-RUN V10.7 - KIỂM TRA FILE HỌC SINH 2025-2026 VÀ QĐ3805")
    print("=" * 124)
    print("Chế độ: CHỈ ĐỌC")

    if not XLSX.exists():
        raise AuditAbort(f"Không tìm thấy file: {XLSX}")
    if not DB_PATH.exists():
        raise AuditAbort(f"Không tìm thấy DB: {DB_PATH}")

    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise AuditAbort(
            f"Không import được openpyxl: {exc}"
        )

    con = connect_ro(DB_PATH)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = ROOT / f"dry_run_kiem_tra_file_hoc_sinh_2025_2026_v10_7_{ts}"
    out_dir.mkdir(parents=True, exist_ok=False)

    try:
        integrity = con.execute(
            "PRAGMA integrity_check"
        ).fetchone()[0]

        qd_ids = load_qd3805_school_ids()
        qd_names = get_school_names(con, qd_ids)

        wb = load_workbook(
            XLSX,
            read_only=True,
            data_only=True,
        )

        sheet_rows = []
        student_rows = []
        school_counter = Counter()
        class_counter = Counter()
        matched_qd_counter = Counter()
        duplicate_code_counter = Counter()

        for ws in wb.worksheets:
            header_row, hmap = find_header_row(ws)

            sheet_total_students = 0
            sheet_school_names = set()

            if header_row is not None:
                for ridx in range(
                    header_row + 1,
                    int(ws.max_row or 0) + 1
                ):
                    student_name = ws.cell(
                        ridx, hmap["student_name"]
                    ).value
                    school_name = ws.cell(
                        ridx, hmap["school"]
                    ).value

                    # Dòng không có học sinh -> bỏ qua.
                    if not str(student_name or "").strip():
                        continue

                    student_name = str(student_name).strip()
                    school_name = str(school_name or "").strip()

                    class_name = ""
                    if "class" in hmap:
                        class_name = str(
                            ws.cell(ridx, hmap["class"]).value or ""
                        ).strip()

                    group_class = ""
                    if "group_class" in hmap:
                        group_class = str(
                            ws.cell(ridx, hmap["group_class"]).value or ""
                        ).strip()

                    student_code = ""
                    if "student_code" in hmap:
                        student_code = str(
                            ws.cell(ridx, hmap["student_code"]).value or ""
                        ).strip()

                    birth_date = ""
                    if "birth_date" in hmap:
                        birth_date = str(
                            ws.cell(ridx, hmap["birth_date"]).value or ""
                        ).strip()

                    gender = ""
                    if "gender" in hmap:
                        gender = str(
                            ws.cell(ridx, hmap["gender"]).value or ""
                        ).strip()

                    match_id, match_name, match_type = match_qd_school(
                        school_name,
                        qd_names,
                    )

                    school_counter[school_name] += 1
                    sheet_school_names.add(school_name)
                    class_counter[(school_name, class_name)] += 1

                    if match_id is not None:
                        matched_qd_counter[match_id] += 1

                    if student_code:
                        duplicate_code_counter[student_code] += 1

                    student_rows.append({
                        "sheet": ws.title,
                        "excel_row": ridx,
                        "student_code": student_code,
                        "student_name": student_name,
                        "birth_date": birth_date,
                        "gender": gender,
                        "school_name_excel": school_name,
                        "group_class": group_class,
                        "class_name": class_name,
                        "qd3805_school_id_match": (
                            match_id if match_id is not None else ""
                        ),
                        "qd3805_school_name_match": match_name,
                        "match_type": match_type,
                    })

                    sheet_total_students += 1

            sheet_rows.append({
                "sheet": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "header_row": (
                    header_row if header_row is not None else ""
                ),
                "school_column": hmap.get("school", ""),
                "student_name_column": hmap.get("student_name", ""),
                "class_column": hmap.get("class", ""),
                "student_code_column": hmap.get("student_code", ""),
                "student_rows_detected": sheet_total_students,
                "unique_school_names": len(sheet_school_names),
                "school_names": " | ".join(
                    sorted(x for x in sheet_school_names if x)
                ),
            })

        wb.close()

        school_rows = []
        for name, count in school_counter.most_common():
            sid, official, match_type = match_qd_school(
                name, qd_names
            )
            school_rows.append({
                "school_name_excel": name,
                "student_rows": count,
                "normalized": normalize(name),
                "qd3805_school_id_match": (
                    sid if sid is not None else ""
                ),
                "qd3805_school_name_match": official,
                "match_type": match_type,
            })

        class_rows = []
        for (school_name, class_name), count in sorted(
            class_counter.items(),
            key=lambda x: (
                normalize(x[0][0]),
                normalize(x[0][1]),
            ),
        ):
            class_rows.append({
                "school_name_excel": school_name,
                "class_name": class_name,
                "student_rows": count,
            })

        qd_rows = []
        for sid, official in sorted(qd_names.items()):
            qd_rows.append({
                "school_id": sid,
                "official_school_name": official,
                "matched_student_rows_in_excel": (
                    matched_qd_counter.get(sid, 0)
                ),
                "found_in_excel": (
                    "YES"
                    if matched_qd_counter.get(sid, 0) > 0
                    else "NO"
                ),
            })

        duplicate_codes = [
            {
                "student_code": code,
                "rows": count,
            }
            for code, count in duplicate_code_counter.items()
            if count > 1
        ]

        total_students = len(student_rows)
        qd_matched_rows = sum(matched_qd_counter.values())
        unique_schools = len(
            [x for x in school_counter if str(x).strip()]
        )

        if total_students == 0:
            conclusion = "NO_STUDENT_ROWS"
        elif qd_matched_rows == 0:
            conclusion = "NOT_QD3805_SOURCE"
        elif qd_matched_rows == total_students:
            conclusion = "ALL_ROWS_MATCH_QD3805"
        else:
            conclusion = "PARTIAL_QD3805_SOURCE"

        gate_rows = [
            {
                "check": "integrity_check",
                "result": "PASS" if integrity == "ok" else "FAIL",
                "detail": str(integrity),
            },
            {
                "check": "file_exists",
                "result": "PASS",
                "detail": str(XLSX),
            },
            {
                "check": "sheet_count",
                "result": "INFO",
                "detail": str(len(sheet_rows)),
            },
            {
                "check": "student_rows_detected",
                "result": "INFO",
                "detail": str(total_students),
            },
            {
                "check": "unique_school_names",
                "result": "INFO",
                "detail": str(unique_schools),
            },
            {
                "check": "qd3805_matched_student_rows",
                "result": (
                    "FOUND" if qd_matched_rows > 0 else "NO_DATA"
                ),
                "detail": str(qd_matched_rows),
            },
            {
                "check": "duplicate_student_codes",
                "result": (
                    "REVIEW" if duplicate_codes else "PASS"
                ),
                "detail": str(len(duplicate_codes)),
            },
            {
                "check": "SOURCE_SCOPE_CONCLUSION",
                "result": conclusion,
                "detail": (
                    "File chỉ được coi là nguồn QĐ3805 nếu tên trường "
                    "khớp trực tiếp với danh mục trường QĐ3805."
                ),
            },
        ]

        write_csv(
            out_dir / "490_SHEET_SUMMARY.csv",
            [
                "sheet", "max_row", "max_column",
                "header_row", "school_column",
                "student_name_column", "class_column",
                "student_code_column",
                "student_rows_detected",
                "unique_school_names", "school_names",
            ],
            sheet_rows,
        )

        write_csv(
            out_dir / "491_SCHOOL_NAMES_IN_EXCEL.csv",
            [
                "school_name_excel",
                "student_rows",
                "normalized",
                "qd3805_school_id_match",
                "qd3805_school_name_match",
                "match_type",
            ],
            school_rows,
        )

        write_csv(
            out_dir / "492_QD3805_SCHOOL_MATCH.csv",
            [
                "school_id",
                "official_school_name",
                "matched_student_rows_in_excel",
                "found_in_excel",
            ],
            qd_rows,
        )

        write_csv(
            out_dir / "493_CLASS_COUNTS.csv",
            [
                "school_name_excel",
                "class_name",
                "student_rows",
            ],
            class_rows,
        )

        write_csv(
            out_dir / "494_STUDENT_SAMPLE_AND_MATCH.csv",
            [
                "sheet", "excel_row",
                "student_code", "student_name",
                "birth_date", "gender",
                "school_name_excel",
                "group_class", "class_name",
                "qd3805_school_id_match",
                "qd3805_school_name_match",
                "match_type",
            ],
            student_rows[:2000],
        )

        write_csv(
            out_dir / "495_DUPLICATE_STUDENT_CODES.csv",
            ["student_code", "rows"],
            duplicate_codes,
        )

        write_csv(
            out_dir / "496_GATE_V10_7.csv",
            ["check", "result", "detail"],
            gate_rows,
        )

        summary = out_dir / "00_TONG_QUAN_V10_7.txt"
        with summary.open("w", encoding="utf-8") as f:
            f.write(
                "DRY-RUN V10.7 - KIỂM TRA FILE HỌC SINH 2025-2026\n"
            )
            f.write("=" * 124 + "\n")
            f.write("CHỈ ĐỌC - KHÔNG IMPORT\n\n")

            f.write(f"File: {XLSX}\n")
            f.write(f"Số sheet: {len(sheet_rows)}\n")
            f.write(f"Dòng học sinh phát hiện: {total_students}\n")
            f.write(f"Số tên trường khác nhau: {unique_schools}\n")
            f.write(
                f"Dòng khớp trường QĐ3805: {qd_matched_rows}\n"
            )
            f.write(
                f"Duplicate student codes: {len(duplicate_codes)}\n\n"
            )

            f.write("TÊN TRƯỜNG TRONG FILE\n")
            for r in school_rows:
                f.write(
                    f"- {r['school_name_excel']}: "
                    f"{r['student_rows']} học sinh | "
                    f"match={r['match_type']} "
                    f"{r['qd3805_school_name_match']}\n"
                )

            f.write("\nĐỐI CHIẾU QĐ3805\n")
            for r in qd_rows:
                f.write(
                    f"- {r['school_id']} {r['official_school_name']}: "
                    f"{r['matched_student_rows_in_excel']} dòng\n"
                )

            f.write("\nKẾT LUẬN\n")
            f.write(
                f"SOURCE_SCOPE_CONCLUSION = {conclusion}\n"
            )

            if conclusion == "NOT_QD3805_SOURCE":
                f.write(
                    "- File này có dữ liệu học sinh 2025-2026 nhưng "
                    "không thuộc các trường QĐ3805/Nghi Lộc.\n"
                )
                f.write(
                    "- Không được dùng file này để dựng baseline QĐ3805.\n"
                )
            elif conclusion == "PARTIAL_QD3805_SOURCE":
                f.write(
                    "- File có một phần dữ liệu khớp QĐ3805; "
                    "cần kiểm tra chi tiết trước khi import.\n"
                )
            elif conclusion == "ALL_ROWS_MATCH_QD3805":
                f.write(
                    "- Toàn bộ dòng học sinh trong file thuộc QĐ3805.\n"
                )
                f.write(
                    "- Có thể chuyển sang dry-run mapping/import baseline.\n"
                )

        zip_path = ROOT / (
            f"dry_run_kiem_tra_file_hoc_sinh_2025_2026_v10_7_{ts}.zip"
        )
        with zipfile.ZipFile(
            zip_path,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as zf:
            for p in sorted(out_dir.iterdir()):
                zf.write(p, arcname=p.name)

        print()
        print("=" * 124)
        print("HOÀN THÀNH DRY-RUN V10.7")
        print("=" * 124)
        print(f"Dòng học sinh phát hiện: {total_students}")
        print(f"Tên trường khác nhau: {unique_schools}")
        print(f"Dòng khớp QĐ3805: {qd_matched_rows}")
        print(f"Duplicate student codes: {len(duplicate_codes)}")
        print(f"SOURCE_SCOPE_CONCLUSION: {conclusion}")
        print(f"ZIP: {zip_path}")
        print("Không có file/database nào bị thay đổi.")

    finally:
        con.close()


if __name__ == "__main__":
    try:
        main()
    except AuditAbort as exc:
        print()
        print("=" * 124)
        print("ĐÃ DỪNG DRY-RUN V10.7")
        print("=" * 124)
        print(str(exc))
        print("Không có file/database nào bị thay đổi.")
        sys.exit(2)
    except Exception as exc:
        print()
        print("=" * 124)
        print("LỖI KHÔNG DỰ KIẾN TRONG V10.7")
        print("=" * 124)
        print(repr(exc))
        print("Không có file/database nào bị thay đổi.")
        sys.exit(4)
