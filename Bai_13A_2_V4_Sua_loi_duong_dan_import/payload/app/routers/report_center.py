from __future__ import annotations

from datetime import datetime
from io import BytesIO
from math import ceil
from pathlib import Path
from typing import Any
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import distinct, func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.comparison_models import StudentSurveyComparisonSignoff
from app.database import get_db
from app.models import (
    Commune,
    School,
    SchoolYear,
    StudentEnrollment,
    User,
)
from app.permissions import (
    COMMUNE_ROLE_CODE,
    SCHOOL_ROLE_CODE,
    TEACHER_ROLE_CODE,
    is_province_reader,
    normalize_role_code,
)
from app.routers.surveys import (
    BATCH_STATUS_LABELS,
    lay_thong_tin_nguoi_dung,
    tao_bo_loc_bao_cao_theo_nguoi_dung,
    tao_bo_loc_dot_theo_nguoi_dung,
)
from app.survey_models import (
    Household,
    SurveyBatch,
    SurveyForm,
    SurveyFormInvestigator,
    SurveyPerson,
    SurveyPersonYearRecord,
)


APP_DIR = Path(__file__).resolve().parent.parent
templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

router = APIRouter(
    prefix="/bao-cao",
    tags=["Trung tâm báo cáo"],
)

PAGE_SIZE_OPTIONS = (20, 50, 100, 200)

REPORT_TYPES: dict[str, dict[str, Any]] = {
    "TONG_HOP_DIEU_TRA": {
        "title": "Báo cáo tổng hợp điều tra",
        "short_title": "Tổng hợp điều tra",
        "description": (
            "Tổng hợp hộ, phiếu, đối tượng, độ tuổi, giới tính, tình trạng "
            "học tập, khuyết tật và dữ liệu theo trường."
        ),
        "icon": "📊",
        "stage": "Sẵn sàng",
        "stage_class": "ready",
        "lesson": "Bài 13A-2",
    },
    "PCMN_5_TUOI": {
        "title": "Báo cáo phổ cập mầm non 5 tuổi",
        "short_title": "Phổ cập 5 tuổi",
        "description": (
            "Tổng hợp trẻ 5 tuổi, huy động ra lớp, hoàn thành chương trình, "
            "khuyết tật và học hòa nhập."
        ),
        "icon": "🧒",
        "stage": "Sắp triển khai",
        "stage_class": "next",
        "lesson": "Bài 13A-3",
    },
    "BIEN_DONG_THEO_DOI": {
        "title": "Báo cáo biến động và đối tượng cần theo dõi",
        "short_title": "Biến động – theo dõi",
        "description": (
            "Theo dõi chuyển đến, chuyển đi, chưa đi học, bỏ học, thiếu số "
            "định danh và thiếu dữ liệu năm học."
        ),
        "icon": "🔎",
        "stage": "Sắp triển khai",
        "stage_class": "next",
        "lesson": "Bài 13A-4",
    },
    "DOI_CHIEU_HOC_SINH": {
        "title": "Báo cáo đối chiếu điều tra với học sinh",
        "short_title": "Đối chiếu học sinh",
        "description": (
            "Phát hiện hồ sơ chưa khớp, sai trường/lớp, sai trạng thái và "
            "theo dõi kết quả xử lý."
        ),
        "icon": "🧩",
        "stage": "Sẵn sàng",
        "stage_class": "ready",
        "lesson": "Bài 12D-15 đến 12D-18",
    },
    "TIEN_DO_CHOT": {
        "title": "Báo cáo tiến độ chốt và đôn đốc",
        "short_title": "Tiến độ – đôn đốc",
        "description": (
            "Giám sát trạng thái chốt, lịch sử mở lại, thời hạn và danh sách "
            "địa bàn cần nhắc việc."
        ),
        "icon": "⏱️",
        "stage": "Sẵn sàng cấp tỉnh",
        "stage_class": "ready",
        "lesson": "Bài 12D-18 và 12D-19",
    },
}

STATUS_LABELS = {
    "": "Tất cả trạng thái",
    "HAS_DATA": "Đã có dữ liệu",
    "NO_DATA": "Chưa có dữ liệu",
    "LOCKED": "Đã khóa số liệu",
    "OPEN": "Dữ liệu đang mở",
    "SIGNED": "Đã chốt đối chiếu",
    "UNSIGNED": "Chưa chốt đối chiếu",
}

SIGNOFF_ACTION_LABELS = {
    "CHOT": "Đã chốt",
    "MO_LAI": "Đã mở lại",
}


def _clean_text(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _parse_optional_int(value: Any) -> int | None:
    text = str(value or "").strip()
    return int(text) if text.isdigit() else None


def _load_school_years(db: Session) -> list[SchoolYear]:
    return list(
        db.scalars(
            select(SchoolYear).order_by(
                SchoolYear.code.desc(),
                SchoolYear.id.desc(),
            )
        ).all()
    )


def _choose_year_id(
    school_years: list[SchoolYear],
    requested_id: int | None,
) -> int | None:
    valid_ids = {item.id for item in school_years}
    if requested_id in valid_ids:
        return requested_id
    active = next((item for item in school_years if item.is_active), None)
    if active is not None:
        return active.id
    return school_years[0].id if school_years else None


def _load_scope_options(
    db: Session,
    user: dict[str, Any],
    requested_commune_id: int | None,
    requested_school_id: int | None,
) -> tuple[list[Commune], list[School], int | None, int | None, str]:
    role_code = normalize_role_code(user.get("role_code"))
    account_commune_id = _parse_optional_int(user.get("commune_id"))
    account_school_id = _parse_optional_int(user.get("school_id"))

    if is_province_reader(role_code):
        communes = list(
            db.scalars(
                select(Commune)
                .where(Commune.is_active.is_(True))
                .order_by(Commune.name.asc(), Commune.id.asc())
            ).all()
        )
        valid_commune_ids = {item.id for item in communes}
        selected_commune_id = (
            requested_commune_id
            if requested_commune_id in valid_commune_ids
            else None
        )

        school_statement = (
            select(School)
            .where(School.is_active.is_(True))
            .order_by(School.name.asc(), School.id.asc())
        )
        if selected_commune_id is not None:
            school_statement = school_statement.where(
                School.commune_id == selected_commune_id
            )
        schools = list(db.scalars(school_statement).all())
        valid_school_ids = {item.id for item in schools}
        selected_school_id = (
            requested_school_id
            if requested_school_id in valid_school_ids
            else None
        )

        if selected_school_id is not None:
            school = next(
                item for item in schools if item.id == selected_school_id
            )
            scope_label = f"Trường {school.name}"
        elif selected_commune_id is not None:
            commune = next(
                item for item in communes if item.id == selected_commune_id
            )
            scope_label = commune.name
        else:
            scope_label = "Toàn tỉnh"
        return (
            communes,
            schools,
            selected_commune_id,
            selected_school_id,
            scope_label,
        )

    if role_code == COMMUNE_ROLE_CODE:
        communes = list(
            db.scalars(
                select(Commune)
                .where(
                    Commune.id == account_commune_id,
                    Commune.is_active.is_(True),
                )
            ).all()
        )
        schools = list(
            db.scalars(
                select(School)
                .where(
                    School.commune_id == account_commune_id,
                    School.is_active.is_(True),
                )
                .order_by(School.name.asc(), School.id.asc())
            ).all()
        )
        valid_school_ids = {item.id for item in schools}
        selected_school_id = (
            requested_school_id
            if requested_school_id in valid_school_ids
            else None
        )
        if selected_school_id is not None:
            school = next(
                item for item in schools if item.id == selected_school_id
            )
            scope_label = f"Trường {school.name}"
        else:
            scope_label = communes[0].name if communes else "Cấp xã/phường"
        return (
            communes,
            schools,
            account_commune_id,
            selected_school_id,
            scope_label,
        )

    school = db.get(School, account_school_id) if account_school_id else None
    commune_id = school.commune_id if school is not None else account_commune_id
    commune = db.get(Commune, commune_id) if commune_id else None
    communes = [commune] if commune is not None else []
    schools = [school] if school is not None else []
    scope_label = school.name if school is not None else user.get("unit_name", "Phạm vi tài khoản")
    return (
        communes,
        schools,
        commune_id,
        account_school_id,
        f"Trường {scope_label}" if school is not None else str(scope_label),
    )


def _load_latest_signoffs(
    db: Session,
    batch_ids: list[int],
) -> dict[int, StudentSurveyComparisonSignoff]:
    if not batch_ids:
        return {}
    history = list(
        db.scalars(
            select(StudentSurveyComparisonSignoff)
            .where(StudentSurveyComparisonSignoff.survey_batch_id.in_(batch_ids))
            .order_by(
                StudentSurveyComparisonSignoff.created_at.desc(),
                StudentSurveyComparisonSignoff.id.desc(),
            )
        ).all()
    )
    latest: dict[int, StudentSurveyComparisonSignoff] = {}
    for item in history:
        latest.setdefault(item.survey_batch_id, item)
    return latest


def _form_scope_filters(
    request: Request,
    selected_school_id: int | None,
) -> list[Any]:
    filters = list(tao_bo_loc_bao_cao_theo_nguoi_dung(request))
    user = lay_thong_tin_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))
    account_school_id = _parse_optional_int(user.get("school_id"))

    # Với tài khoản cấp tỉnh/xã, lựa chọn trường thu hẹp theo hồ sơ năm học
    # hoặc theo phiếu đã được giao cho nhân sự của trường đó.
    if selected_school_id is not None and role_code not in {
        SCHOOL_ROLE_CODE,
        TEACHER_ROLE_CODE,
    }:
        filters.append(
            or_(
                SurveyForm.person_year_records.any(
                    SurveyPersonYearRecord.school_id == selected_school_id
                ),
                SurveyForm.investigators.any(
                    SurveyFormInvestigator.user.has(
                        User.school_id == selected_school_id
                    )
                ),
            )
        )
    elif role_code in {SCHOOL_ROLE_CODE, TEACHER_ROLE_CODE} and account_school_id:
        # tao_bo_loc_bao_cao_theo_nguoi_dung đã áp dụng phạm vi trường.
        pass
    return filters


def _aggregate_batch_counts(
    db: Session,
    request: Request,
    batch_ids: list[int],
    school_year_id: int | None,
    selected_school_id: int | None,
) -> dict[int, dict[str, int]]:
    result = {
        batch_id: {
            "forms": 0,
            "households": 0,
            "people": 0,
            "year_records": 0,
            "reviewed_records": 0,
            "linked_people": 0,
        }
        for batch_id in batch_ids
    }
    if not batch_ids:
        return result

    scope_filters = _form_scope_filters(request, selected_school_id)

    form_rows = db.execute(
        select(
            SurveyForm.survey_batch_id,
            func.count(distinct(SurveyForm.id)).label("forms"),
            func.count(distinct(SurveyForm.household_id)).label("households"),
        )
        .where(
            SurveyForm.survey_batch_id.in_(batch_ids),
            *scope_filters,
        )
        .group_by(SurveyForm.survey_batch_id)
    ).all()
    for batch_id, forms, households in form_rows:
        result[int(batch_id)]["forms"] = int(forms or 0)
        result[int(batch_id)]["households"] = int(households or 0)

    people_rows = db.execute(
        select(
            SurveyForm.survey_batch_id,
            func.count(distinct(SurveyPerson.id)).label("people"),
            func.count(
                distinct(SurveyPerson.id)
            ).filter(
                SurveyPerson.student_id.is_not(None)
            ).label("linked_people"),
        )
        .select_from(SurveyForm)
        .join(Household, Household.id == SurveyForm.household_id)
        .join(SurveyPerson, SurveyPerson.household_id == Household.id)
        .where(
            SurveyForm.survey_batch_id.in_(batch_ids),
            SurveyPerson.is_active.is_(True),
            *scope_filters,
        )
        .group_by(SurveyForm.survey_batch_id)
    ).all()
    for batch_id, people, linked_people in people_rows:
        result[int(batch_id)]["people"] = int(people or 0)
        result[int(batch_id)]["linked_people"] = int(linked_people or 0)

    if school_year_id is not None:
        year_filters: list[Any] = [
            SurveyPersonYearRecord.school_year_id == school_year_id
        ]
        if selected_school_id is not None:
            year_filters.append(
                SurveyPersonYearRecord.school_id == selected_school_id
            )
        year_rows = db.execute(
            select(
                SurveyForm.survey_batch_id,
                func.count(distinct(SurveyPersonYearRecord.id)).label(
                    "year_records"
                ),
                func.count(
                    distinct(SurveyPersonYearRecord.id)
                ).filter(
                    SurveyPersonYearRecord.is_reviewed.is_(True)
                ).label("reviewed_records"),
            )
            .select_from(SurveyForm)
            .join(
                SurveyPersonYearRecord,
                SurveyPersonYearRecord.survey_form_id == SurveyForm.id,
            )
            .where(
                SurveyForm.survey_batch_id.in_(batch_ids),
                *scope_filters,
                *year_filters,
            )
            .group_by(SurveyForm.survey_batch_id)
        ).all()
        for batch_id, year_records, reviewed_records in year_rows:
            result[int(batch_id)]["year_records"] = int(year_records or 0)
            result[int(batch_id)]["reviewed_records"] = int(
                reviewed_records or 0
            )
    return result


def _student_total(
    db: Session,
    school_year_id: int | None,
    selected_commune_id: int | None,
    selected_school_id: int | None,
    user: dict[str, Any],
) -> int:
    if school_year_id is None:
        return 0
    role_code = normalize_role_code(user.get("role_code"))
    account_school_id = _parse_optional_int(user.get("school_id"))
    account_commune_id = _parse_optional_int(user.get("commune_id"))

    statement = (
        select(func.count(distinct(StudentEnrollment.student_id)))
        .select_from(StudentEnrollment)
        .join(School, School.id == StudentEnrollment.school_id)
        .where(StudentEnrollment.school_year_id == school_year_id)
    )
    if selected_school_id is not None:
        statement = statement.where(
            StudentEnrollment.school_id == selected_school_id
        )
    elif role_code in {SCHOOL_ROLE_CODE, TEACHER_ROLE_CODE} and account_school_id:
        statement = statement.where(
            StudentEnrollment.school_id == account_school_id
        )
    elif selected_commune_id is not None:
        statement = statement.where(School.commune_id == selected_commune_id)
    elif role_code == COMMUNE_ROLE_CODE and account_commune_id:
        statement = statement.where(School.commune_id == account_commune_id)
    return int(db.scalar(statement) or 0)


def _report_action(
    report_type: str,
    batch: SurveyBatch,
    latest_signoff: StudentSurveyComparisonSignoff | None,
    province_reader: bool,
) -> dict[str, Any]:
    if report_type == "TONG_HOP_DIEU_TRA":
        report_params = urlencode(
            {
                "school_year_id": batch.school_year_id,
                "commune_id": batch.commune_id,
            }
        )
        return {
            "view_url": (
                "/bao-cao/tong-hop-dieu-tra?"
                + report_params
            ),
            "view_label": "Lập báo cáo tổng hợp",
            "export_url": (
                "/bao-cao/tong-hop-dieu-tra/xuat-excel?"
                + report_params
            ),
            "export_label": (
                "Xuất Excel chính thức"
                if batch.is_locked
                else "Xuất Excel dự thảo"
            ),
            "availability": (
                "Đủ điều kiện xuất chính thức"
                if batch.is_locked
                else "Có thể xem và xuất dự thảo; cần khóa đợt để chính thức"
            ),
            "availability_class": "ready" if batch.is_locked else "draft",
            "secondary_url": f"/dieu-tra/{batch.id}/bao-cao-tong-hop",
            "secondary_label": "Xem báo cáo theo đợt",
        }
    if report_type == "DOI_CHIEU_HOC_SINH":
        signed = latest_signoff is not None and latest_signoff.action_code == "CHOT"
        return {
            "view_url": f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh",
            "view_label": "Mở đối chiếu",
            "export_url": f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh/xuat-excel",
            "export_label": "Xuất Excel đối chiếu",
            "availability": (
                "Đã chốt – có biên bản"
                if signed
                else "Đang xử lý hoặc chưa chốt"
            ),
            "availability_class": "ready" if signed else "draft",
            "secondary_url": (
                f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh/chot-ket-qua"
            ),
            "secondary_label": "Biên bản / chốt",
        }
    if report_type == "TIEN_DO_CHOT":
        if province_reader:
            return {
                "view_url": "/dieu-tra/tong-hop-doi-chieu-hoc-sinh",
                "view_label": "Mở tổng hợp chốt",
                "export_url": "/dieu-tra/don-doc-doi-chieu-hoc-sinh/xuat-excel",
                "export_label": "Xuất danh sách đôn đốc",
                "secondary_url": "/dieu-tra/don-doc-doi-chieu-hoc-sinh",
                "secondary_label": "Mở giám sát đôn đốc",
                "availability": "Sẵn sàng ở phạm vi toàn tỉnh",
                "availability_class": "ready",
            }
        return {
            "view_url": f"/dieu-tra/{batch.id}/doi-chieu-hoc-sinh/chot-ket-qua",
            "view_label": "Xem trạng thái chốt",
            "export_url": None,
            "availability": "Xem trạng thái trong phạm vi tài khoản",
            "availability_class": "draft",
        }
    meta = REPORT_TYPES[report_type]
    return {
        "view_url": None,
        "export_url": None,
        "availability": f"Sẽ triển khai tại {meta['lesson']}",
        "availability_class": "next",
    }


def _build_report_rows(
    db: Session,
    request: Request,
    school_year_id: int | None,
    selected_commune_id: int | None,
    selected_school_id: int | None,
    report_type: str,
    status_filter: str,
    q: str,
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    user = lay_thong_tin_nguoi_dung(request)
    role_code = normalize_role_code(user.get("role_code"))
    province_reader = is_province_reader(role_code)

    if school_year_id is None:
        empty_summary = {
            "batches": 0,
            "with_data": 0,
            "locked": 0,
            "signed": 0,
            "forms": 0,
            "households": 0,
            "people": 0,
            "year_records": 0,
            "reviewed_records": 0,
            "linked_people": 0,
            "students": 0,
        }
        return [], empty_summary

    statement = (
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.commune),
            selectinload(SurveyBatch.school_year),
        )
        .where(
            SurveyBatch.school_year_id == school_year_id,
            *tao_bo_loc_dot_theo_nguoi_dung(request),
        )
        .order_by(
            SurveyBatch.commune_id.asc(),
            SurveyBatch.created_at.desc(),
            SurveyBatch.id.desc(),
        )
    )
    if selected_commune_id is not None:
        statement = statement.where(
            SurveyBatch.commune_id == selected_commune_id
        )
    if selected_school_id is not None:
        statement = statement.where(
            SurveyBatch.survey_forms.any(
                or_(
                    SurveyForm.person_year_records.any(
                        SurveyPersonYearRecord.school_id == selected_school_id
                    ),
                    SurveyForm.investigators.any(
                        SurveyFormInvestigator.user.has(
                            User.school_id == selected_school_id
                        )
                    ),
                )
            )
        )
    if q:
        q_like = f"%{q}%"
        statement = statement.join(
            Commune,
            Commune.id == SurveyBatch.commune_id,
        ).where(
            or_(
                SurveyBatch.code.ilike(q_like),
                SurveyBatch.name.ilike(q_like),
                Commune.name.ilike(q_like),
                Commune.code.ilike(q_like),
            )
        )

    batches = list(db.scalars(statement).unique().all())
    batch_ids = [item.id for item in batches]
    latest_signoffs = _load_latest_signoffs(db, batch_ids)
    counts = _aggregate_batch_counts(
        db,
        request,
        batch_ids,
        school_year_id,
        selected_school_id,
    )

    rows: list[dict[str, Any]] = []
    for batch in batches:
        count_item = counts.get(batch.id, {})
        latest = latest_signoffs.get(batch.id)
        has_data = int(count_item.get("forms", 0)) > 0
        signed = latest is not None and latest.action_code == "CHOT"

        include = True
        if status_filter == "HAS_DATA":
            include = has_data
        elif status_filter == "NO_DATA":
            include = not has_data
        elif status_filter == "LOCKED":
            include = bool(batch.is_locked)
        elif status_filter == "OPEN":
            include = not bool(batch.is_locked)
        elif status_filter == "SIGNED":
            include = signed
        elif status_filter == "UNSIGNED":
            include = not signed
        if not include:
            continue

        action = _report_action(
            report_type,
            batch,
            latest,
            province_reader,
        )
        rows.append(
            {
                "batch": batch,
                "batch_status_label": BATCH_STATUS_LABELS.get(
                    batch.status,
                    batch.status,
                ),
                "data_status_label": "Đã khóa" if batch.is_locked else "Đang mở",
                "has_data": has_data,
                "signoff_label": (
                    SIGNOFF_ACTION_LABELS.get(
                        latest.action_code,
                        latest.action_code,
                    )
                    if latest is not None
                    else "Chưa có lịch sử chốt"
                ),
                "signoff_class": (
                    "signed"
                    if signed
                    else "reopened"
                    if latest is not None
                    else "unsigned"
                ),
                "signoff_actor": (
                    latest.actor_name_snapshot if latest is not None else ""
                ),
                "signoff_time": (
                    latest.created_at.strftime("%d/%m/%Y %H:%M")
                    if latest is not None
                    else ""
                ),
                "signoff_note": latest.note if latest is not None else "",
                **count_item,
                **action,
            }
        )

    summary = {
        "batches": len(batches),
        "with_data": sum(
            1 for batch in batches if counts.get(batch.id, {}).get("forms", 0)
        ),
        "locked": sum(1 for batch in batches if batch.is_locked),
        "signed": sum(
            1
            for batch in batches
            if latest_signoffs.get(batch.id) is not None
            and latest_signoffs[batch.id].action_code == "CHOT"
        ),
        "forms": sum(item.get("forms", 0) for item in counts.values()),
        "households": sum(
            item.get("households", 0) for item in counts.values()
        ),
        "people": sum(item.get("people", 0) for item in counts.values()),
        "year_records": sum(
            item.get("year_records", 0) for item in counts.values()
        ),
        "reviewed_records": sum(
            item.get("reviewed_records", 0) for item in counts.values()
        ),
        "linked_people": sum(
            item.get("linked_people", 0) for item in counts.values()
        ),
        "students": _student_total(
            db,
            school_year_id,
            selected_commune_id,
            selected_school_id,
            user,
        ),
    }
    return rows, summary


def _build_query_url(
    *,
    school_year_id: int | None,
    report_type: str,
    commune_id: int | None,
    school_id: int | None,
    status_filter: str,
    q: str,
    page_size: int,
    page: int,
) -> str:
    params = urlencode(
        {
            "school_year_id": school_year_id or "",
            "report_type": report_type,
            "commune_id": commune_id or "",
            "school_id": school_id or "",
            "status_filter": status_filter,
            "q": q,
            "page_size": page_size,
            "page": page,
        }
    )
    return "/bao-cao?" + params


def _style_sheet(worksheet: Any) -> None:
    dark_fill = PatternFill("solid", fgColor="1F4E78")
    light_fill = PatternFill("solid", fgColor="D9EAF7")
    white_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)
    thin = Side(style="thin", color="C9D4DF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in worksheet[1]:
        cell.fill = dark_fill
        cell.font = white_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[2]:
        if cell.value is not None:
            cell.fill = light_fill
            cell.font = bold_font


def _set_widths(worksheet: Any, widths: dict[int, float]) -> None:
    for index, width in widths.items():
        worksheet.column_dimensions[get_column_letter(index)].width = width


@router.get("", response_class=HTMLResponse)
def report_center_page(
    request: Request,
    school_year_id: str = "",
    report_type: str = "TONG_HOP_DIEU_TRA",
    commune_id: str = "",
    school_id: str = "",
    status_filter: str = "",
    q: str = "",
    page_size: int = 20,
    page: int = 1,
    db: Session = Depends(get_db),
):
    user = lay_thong_tin_nguoi_dung(request)
    school_years = _load_school_years(db)
    selected_year_id = _choose_year_id(
        school_years,
        _parse_optional_int(school_year_id),
    )
    selected_year = next(
        (item for item in school_years if item.id == selected_year_id),
        None,
    )

    normalized_report_type = _clean_text(report_type).upper()
    if normalized_report_type not in REPORT_TYPES:
        normalized_report_type = "TONG_HOP_DIEU_TRA"
    normalized_status = _clean_text(status_filter).upper()
    if normalized_status not in STATUS_LABELS:
        normalized_status = ""
    q = _clean_text(q)[:120]
    if page_size not in PAGE_SIZE_OPTIONS:
        page_size = 20

    (
        communes,
        schools,
        selected_commune_id,
        selected_school_id,
        scope_label,
    ) = _load_scope_options(
        db,
        user,
        _parse_optional_int(commune_id),
        _parse_optional_int(school_id),
    )

    all_rows, summary = _build_report_rows(
        db,
        request,
        selected_year_id,
        selected_commune_id,
        selected_school_id,
        normalized_report_type,
        normalized_status,
        q,
    )

    total_records = len(all_rows)
    total_pages = max(1, ceil(total_records / page_size))
    page = max(1, min(page, total_pages))
    start = (page - 1) * page_size
    page_rows = all_rows[start:start + page_size]

    page_links = [
        {
            "number": number,
            "url": _build_query_url(
                school_year_id=selected_year_id,
                report_type=normalized_report_type,
                commune_id=selected_commune_id,
                school_id=selected_school_id,
                status_filter=normalized_status,
                q=q,
                page_size=page_size,
                page=number,
            ),
        }
        for number in range(max(1, page - 2), min(total_pages, page + 2) + 1)
    ]

    export_params = urlencode(
        {
            "school_year_id": selected_year_id or "",
            "report_type": normalized_report_type,
            "commune_id": selected_commune_id or "",
            "school_id": selected_school_id or "",
            "status_filter": normalized_status,
            "q": q,
        }
    )

    return templates.TemplateResponse(
        request=request,
        name="reports/report_center.html",
        context={
            "nguoi_dung": user,
            "school_years": school_years,
            "selected_year": selected_year,
            "selected_year_id": selected_year_id,
            "report_types": REPORT_TYPES,
            "selected_report_type": normalized_report_type,
            "selected_report": REPORT_TYPES[normalized_report_type],
            "communes": communes,
            "schools": schools,
            "selected_commune_id": selected_commune_id,
            "selected_school_id": selected_school_id,
            "scope_label": scope_label,
            "status_labels": STATUS_LABELS,
            "status_filter": normalized_status,
            "q": q,
            "summary": summary,
            "rows": page_rows,
            "total_records": total_records,
            "page": page,
            "total_pages": total_pages,
            "page_links": page_links,
            "page_size": page_size,
            "page_size_options": PAGE_SIZE_OPTIONS,
            "previous_url": (
                _build_query_url(
                    school_year_id=selected_year_id,
                    report_type=normalized_report_type,
                    commune_id=selected_commune_id,
                    school_id=selected_school_id,
                    status_filter=normalized_status,
                    q=q,
                    page_size=page_size,
                    page=page - 1,
                )
                if page > 1
                else None
            ),
            "next_url": (
                _build_query_url(
                    school_year_id=selected_year_id,
                    report_type=normalized_report_type,
                    commune_id=selected_commune_id,
                    school_id=selected_school_id,
                    status_filter=normalized_status,
                    q=q,
                    page_size=page_size,
                    page=page + 1,
                )
                if page < total_pages
                else None
            ),
            "export_url": "/bao-cao/xuat-danh-muc-excel?" + export_params,
            "province_reader": is_province_reader(user.get("role_code")),
            "generated_at": datetime.now().strftime("%d/%m/%Y %H:%M"),
        },
    )


@router.get("/xuat-danh-muc-excel")
def export_report_catalog_excel(
    request: Request,
    school_year_id: str = "",
    report_type: str = "TONG_HOP_DIEU_TRA",
    commune_id: str = "",
    school_id: str = "",
    status_filter: str = "",
    q: str = "",
    db: Session = Depends(get_db),
):
    user = lay_thong_tin_nguoi_dung(request)
    school_years = _load_school_years(db)
    selected_year_id = _choose_year_id(
        school_years,
        _parse_optional_int(school_year_id),
    )
    selected_year = next(
        (item for item in school_years if item.id == selected_year_id),
        None,
    )
    normalized_report_type = _clean_text(report_type).upper()
    if normalized_report_type not in REPORT_TYPES:
        normalized_report_type = "TONG_HOP_DIEU_TRA"
    normalized_status = _clean_text(status_filter).upper()
    if normalized_status not in STATUS_LABELS:
        normalized_status = ""
    q = _clean_text(q)[:120]

    (
        _communes,
        _schools,
        selected_commune_id,
        selected_school_id,
        scope_label,
    ) = _load_scope_options(
        db,
        user,
        _parse_optional_int(commune_id),
        _parse_optional_int(school_id),
    )
    rows, summary = _build_report_rows(
        db,
        request,
        selected_year_id,
        selected_commune_id,
        selected_school_id,
        normalized_report_type,
        normalized_status,
        q,
    )

    workbook = Workbook()
    catalog = workbook.active
    catalog.title = "Danh muc bao cao"
    catalog.append(
        [
            "STT",
            "Mã đợt",
            "Tên đợt",
            "Năm học",
            "Xã/phường",
            "Trạng thái đợt",
            "Trạng thái dữ liệu",
            "Phiếu",
            "Hộ",
            "Đối tượng",
            "Hồ sơ năm học",
            "Đã rà soát",
            "Đã liên kết học sinh",
            "Trạng thái chốt đối chiếu",
            "Người thao tác gần nhất",
            "Thời điểm gần nhất",
            "Căn cứ/lý do",
            "Mức sẵn sàng báo cáo",
        ]
    )
    for index, row in enumerate(rows, start=1):
        batch = row["batch"]
        catalog.append(
            [
                index,
                batch.code,
                batch.name,
                batch.school_year.code,
                batch.commune.name,
                row["batch_status_label"],
                row["data_status_label"],
                row.get("forms", 0),
                row.get("households", 0),
                row.get("people", 0),
                row.get("year_records", 0),
                row.get("reviewed_records", 0),
                row.get("linked_people", 0),
                row["signoff_label"],
                row["signoff_actor"],
                row["signoff_time"],
                row["signoff_note"],
                row["availability"],
            ]
        )
    _style_sheet(catalog)
    _set_widths(
        catalog,
        {
            1: 8,
            2: 24,
            3: 42,
            4: 14,
            5: 25,
            6: 18,
            7: 18,
            8: 10,
            9: 10,
            10: 12,
            11: 16,
            12: 14,
            13: 18,
            14: 24,
            15: 24,
            16: 20,
            17: 42,
            18: 42,
        },
    )

    overview = workbook.create_sheet("Tong hop pham vi")
    overview.append(["Chỉ tiêu", "Giá trị"])
    overview_rows = [
        ("Loại báo cáo", REPORT_TYPES[normalized_report_type]["title"]),
        ("Năm học", selected_year.code if selected_year else "Chưa xác định"),
        ("Phạm vi", scope_label),
        ("Số đợt trong phạm vi", summary["batches"]),
        ("Đợt đã có dữ liệu", summary["with_data"]),
        ("Đợt đã khóa số liệu", summary["locked"]),
        ("Đợt đã chốt đối chiếu", summary["signed"]),
        ("Tổng phiếu", summary["forms"]),
        ("Tổng hộ", summary["households"]),
        ("Tổng đối tượng", summary["people"]),
        ("Hồ sơ năm học", summary["year_records"]),
        ("Hồ sơ đã rà soát", summary["reviewed_records"]),
        ("Đối tượng đã liên kết học sinh", summary["linked_people"]),
        ("Học sinh theo năm học", summary["students"]),
    ]
    for item in overview_rows:
        overview.append(list(item))
    _style_sheet(overview)
    _set_widths(overview, {1: 38, 2: 50})

    info = workbook.create_sheet("Thong tin bo loc")
    info.append(["Thông tin", "Giá trị"])
    info_rows = [
        ("Thời điểm lập", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Người lập", user.get("full_name", "")),
        ("Vai trò", user.get("role_name", "")),
        ("Đơn vị", user.get("unit_name", "")),
        ("Năm học", selected_year.code if selected_year else ""),
        ("Loại báo cáo", REPORT_TYPES[normalized_report_type]["title"]),
        ("Phạm vi", scope_label),
        ("Trạng thái lọc", STATUS_LABELS[normalized_status]),
        ("Từ khóa", q or "Không"),
        (
            "Ghi chú",
            "Tệp này là danh mục và ảnh chụp mức sẵn sàng báo cáo; "
            "không tự sửa dữ liệu gốc.",
        ),
    ]
    for item in info_rows:
        info.append(list(item))
    _style_sheet(info)
    _set_widths(info, {1: 30, 2: 70})

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    year_code = selected_year.code if selected_year else "khong_nam_hoc"
    filename = f"trung_tam_bao_cao_{year_code}_{normalized_report_type.lower()}.xlsx"
    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )
