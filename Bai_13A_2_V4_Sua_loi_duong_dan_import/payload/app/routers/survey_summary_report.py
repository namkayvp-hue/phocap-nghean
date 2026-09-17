from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from math import ceil
from pathlib import Path
from typing import Any
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, distinct, func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.database import get_db
from app.models import Classroom, Commune, School, SchoolYear, User
from app.permissions import (
    COMMUNE_ROLE_CODE,
    SCHOOL_ROLE_CODE,
    TEACHER_ROLE_CODE,
    is_province_reader,
    normalize_role_code,
)
from app.routers.report_center import (
    _choose_year_id,
    _clean_text,
    _load_school_years,
    _load_scope_options,
    _parse_optional_int,
)
from app.routers.surveys import (
    BATCH_STATUS_LABELS,
    DISABILITY_LEVEL_LABELS,
    DISABILITY_STATUS_LABELS,
    DISABILITY_TYPE_LABELS,
    LEARNING_STATUS_LABELS,
    RESIDENCY_STATUS_LABELS,
    FORM_STATUS_LABELS,
    lay_thong_tin_nguoi_dung,
    tao_bo_loc_bao_cao_theo_nguoi_dung,
    tao_bo_loc_dot_theo_nguoi_dung,
    tinh_tuoi_tai_ngay,
)
from app.survey_models import (
    Household,
    SurveyBatch,
    SurveyForm,
    SurveyFormInvestigator,
    SurveyPerson,
    SurveyPersonYearRecord,
)


APP_DIR = Path(__file__).resolve().parent.parent
templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

router = APIRouter(
    prefix="/bao-cao/tong-hop-dieu-tra",
    tags=["Báo cáo tổng hợp điều tra"],
)

PAGE_SIZE_OPTIONS = (20, 50, 100, 200)
DATA_STATE_LABELS = {
    "": "Tất cả trạng thái dữ liệu",
    "HAS_DATA": "Đã có dữ liệu",
    "NO_DATA": "Chưa có dữ liệu",
    "LOCKED": "Đã khóa số liệu",
    "OPEN": "Dữ liệu đang mở",
}

GENDER_LABELS = {
    "NAM": "Nam",
    "NU": "Nữ",
    "KHAC": "Khác",
    "CHUA_XAC_DINH": "Chưa xác định",
}

AGE_GROUP_ORDER = ("0", "1", "2", "3", "4", "5", "6+", "Chưa rõ")


def _normalize_gender(value: Any) -> tuple[str, str]:
    text = _clean_text(value).upper()
    if text in {"NAM", "MALE", "M"}:
        return "NAM", "Nam"
    if text in {"NỮ", "NU", "FEMALE", "F"}:
        return "NU", "Nữ"
    if text:
        return "KHAC", _clean_text(value)
    return "CHUA_XAC_DINH", "Chưa xác định"


def _reference_date(school_year: SchoolYear | None) -> date:
    if school_year is not None and school_year.start_date is not None:
        return school_year.start_date
    if school_year is not None:
        try:
            ending_year = int(str(school_year.code).split("-")[-1])
            return date(ending_year, 7, 15)
        except (TypeError, ValueError):
            pass
    return date.today()


def _age_group(age: int | None) -> str:
    if age is None:
        return "Chưa rõ"
    if age <= 5:
        return str(max(age, 0))
    return "6+"


def _safe_date(value: date | None) -> str:
    return value.strftime("%d/%m/%Y") if value is not None else ""


def _safe_datetime(value: datetime | None) -> str:
    return value.strftime("%d/%m/%Y %H:%M") if value is not None else ""


def _form_permission_filters(request: Request) -> list[Any]:
    return list(tao_bo_loc_bao_cao_theo_nguoi_dung(request))


def _school_batch_condition(
    school_year_id: int,
    school_id: int,
) -> Any:
    return SurveyBatch.survey_forms.any(
        SurveyForm.person_year_records.any(
            and_(
                SurveyPersonYearRecord.school_year_id == school_year_id,
                SurveyPersonYearRecord.school_id == school_id,
            )
        )
    )


def _load_batches(
    *,
    db: Session,
    request: Request,
    school_year_id: int | None,
    commune_id: int | None,
    school_id: int | None,
    data_state: str,
    q: str,
) -> tuple[list[SurveyBatch], dict[int, dict[str, int]]]:
    if school_year_id is None:
        return [], {}

    statement = (
        select(SurveyBatch)
        .options(
            selectinload(SurveyBatch.school_year),
            selectinload(SurveyBatch.commune),
        )
        .where(
            SurveyBatch.school_year_id == school_year_id,
            *tao_bo_loc_dot_theo_nguoi_dung(request),
        )
        .order_by(
            SurveyBatch.commune_id.asc(),
            SurveyBatch.id.asc(),
        )
    )
    if commune_id is not None:
        statement = statement.where(SurveyBatch.commune_id == commune_id)
    if school_id is not None:
        statement = statement.where(
            _school_batch_condition(school_year_id, school_id)
        )
    if data_state == "LOCKED":
        statement = statement.where(SurveyBatch.is_locked.is_(True))
    elif data_state == "OPEN":
        statement = statement.where(SurveyBatch.is_locked.is_(False))
    if q:
        like_value = f"%{q}%"
        statement = statement.join(Commune, Commune.id == SurveyBatch.commune_id)
        statement = statement.where(
            or_(
                SurveyBatch.code.ilike(like_value),
                SurveyBatch.name.ilike(like_value),
                Commune.code.ilike(like_value),
                Commune.name.ilike(like_value),
            )
        )

    batches = list(db.scalars(statement).unique().all())
    batch_ids = [item.id for item in batches]
    counts = {
        batch_id: {"forms": 0, "households": 0}
        for batch_id in batch_ids
    }
    if not batch_ids:
        return batches, counts

    form_statement = (
        select(
            SurveyForm.survey_batch_id,
            func.count(distinct(SurveyForm.id)).label("forms"),
            func.count(distinct(SurveyForm.household_id)).label("households"),
        )
        .where(
            SurveyForm.survey_batch_id.in_(batch_ids),
            *_form_permission_filters(request),
        )
        .group_by(SurveyForm.survey_batch_id)
    )
    if school_id is not None:
        form_statement = form_statement.where(
            SurveyForm.person_year_records.any(
                and_(
                    SurveyPersonYearRecord.school_year_id == school_year_id,
                    SurveyPersonYearRecord.school_id == school_id,
                )
            )
        )

    for batch_id, form_total, household_total in db.execute(form_statement):
        counts[int(batch_id)] = {
            "forms": int(form_total or 0),
            "households": int(household_total or 0),
        }

    if data_state == "HAS_DATA":
        batches = [item for item in batches if counts[item.id]["forms"] > 0]
    elif data_state == "NO_DATA":
        batches = [item for item in batches if counts[item.id]["forms"] == 0]
    return batches, counts


def _load_person_rows(
    *,
    db: Session,
    request: Request,
    batch_ids: list[int],
    school_year_id: int | None,
    school_id: int | None,
) -> list[dict[str, Any]]:
    if not batch_ids or school_year_id is None:
        return []

    year_record = SurveyPersonYearRecord
    statement = (
        select(
            SurveyBatch.id.label("batch_id"),
            SurveyBatch.code.label("batch_code"),
            SurveyBatch.name.label("batch_name"),
            SurveyBatch.is_locked.label("batch_is_locked"),
            SurveyBatch.status.label("batch_status"),
            Commune.id.label("commune_id"),
            Commune.code.label("commune_code"),
            Commune.name.label("commune_name"),
            SurveyForm.id.label("form_id"),
            SurveyForm.form_number.label("form_number"),
            SurveyForm.status.label("form_status"),
            SurveyForm.survey_date.label("survey_date"),
            Household.id.label("household_id"),
            Household.code.label("household_code"),
            Household.head_name.label("head_name"),
            Household.hamlet_name.label("hamlet_name"),
            Household.address.label("address"),
            SurveyPerson.id.label("person_id"),
            SurveyPerson.code.label("person_code"),
            SurveyPerson.full_name.label("full_name"),
            SurveyPerson.date_of_birth.label("date_of_birth"),
            SurveyPerson.gender.label("gender"),
            SurveyPerson.ethnic_group.label("ethnic_group"),
            SurveyPerson.personal_id.label("personal_id"),
            SurveyPerson.ministry_student_code.label("ministry_student_code"),
            SurveyPerson.residency_status.label("residency_status"),
            SurveyPerson.relationship_to_head.label("relationship_to_head"),
            year_record.id.label("year_record_id"),
            year_record.learning_status.label("learning_status"),
            year_record.is_reviewed.label("is_reviewed"),
            year_record.completed_preschool_5.label("completed_preschool_5"),
            year_record.disability_status.label("disability_status"),
            year_record.disability_type.label("disability_type"),
            year_record.disability_level.label("disability_level"),
            year_record.disability_certificate.label("disability_certificate"),
            year_record.inclusive_education.label("inclusive_education"),
            year_record.disability_support.label("disability_support"),
            year_record.school_id.label("school_id"),
            year_record.school_name_reported.label("school_name_reported"),
            year_record.class_name_reported.label("class_name_reported"),
            School.code.label("school_code"),
            School.name.label("school_name"),
            School.commune_id.label("school_commune_id"),
            Classroom.name.label("class_name"),
        )
        .select_from(SurveyForm)
        .join(SurveyBatch, SurveyBatch.id == SurveyForm.survey_batch_id)
        .join(Commune, Commune.id == SurveyBatch.commune_id)
        .join(Household, Household.id == SurveyForm.household_id)
        .join(SurveyPerson, SurveyPerson.household_id == Household.id)
        .outerjoin(
            year_record,
            and_(
                year_record.survey_form_id == SurveyForm.id,
                year_record.survey_person_id == SurveyPerson.id,
                year_record.school_year_id == school_year_id,
            ),
        )
        .outerjoin(School, School.id == year_record.school_id)
        .outerjoin(Classroom, Classroom.id == year_record.class_id)
        .where(
            SurveyForm.survey_batch_id.in_(batch_ids),
            SurveyPerson.is_active.is_(True),
            *_form_permission_filters(request),
        )
        .order_by(
            Commune.name.asc(),
            Household.hamlet_name.asc(),
            Household.head_name.asc(),
            SurveyPerson.full_name.asc(),
            SurveyPerson.id.asc(),
        )
    )
    if school_id is not None:
        statement = statement.where(year_record.school_id == school_id)

    raw_rows = [dict(item) for item in db.execute(statement).mappings().all()]
    seen: set[tuple[int, int]] = set()
    rows: list[dict[str, Any]] = []
    for row in raw_rows:
        key = (int(row["batch_id"]), int(row["person_id"]))
        if key in seen:
            continue
        seen.add(key)
        rows.append(row)
    return rows


def _new_commune_item(batch: SurveyBatch, count_item: dict[str, int]) -> dict[str, Any]:
    return {
        "commune_id": batch.commune_id,
        "commune_code": batch.commune.code,
        "commune_name": batch.commune.name,
        "batch_total": 1,
        "locked_batch_total": 1 if batch.is_locked else 0,
        "forms": int(count_item.get("forms", 0)),
        "households": int(count_item.get("households", 0)),
        "people": 0,
        "male": 0,
        "female": 0,
        "age_5": 0,
        "studying": 0,
        "not_started": 0,
        "dropout": 0,
        "completed_preschool_5": 0,
        "disability": 0,
        "inclusive": 0,
        "missing_personal_id": 0,
        "missing_year_record": 0,
    }


def _finalize_people(
    *,
    batches: list[SurveyBatch],
    counts: dict[int, dict[str, int]],
    raw_rows: list[dict[str, Any]],
    reference_date: date,
) -> dict[str, Any]:
    batch_by_id = {item.id: item for item in batches}
    commune_grouped: dict[int, dict[str, Any]] = {}
    for batch in batches:
        existing = commune_grouped.get(batch.commune_id)
        if existing is None:
            commune_grouped[batch.commune_id] = _new_commune_item(
                batch,
                counts.get(batch.id, {}),
            )
        else:
            existing["batch_total"] += 1
            existing["locked_batch_total"] += 1 if batch.is_locked else 0
            existing["forms"] += int(counts.get(batch.id, {}).get("forms", 0))
            existing["households"] += int(
                counts.get(batch.id, {}).get("households", 0)
            )

    gender_counts = {code: 0 for code in GENDER_LABELS}
    residency_counts = {code: 0 for code in RESIDENCY_STATUS_LABELS}
    learning_counts = {
        **{code: 0 for code in LEARNING_STATUS_LABELS},
        "CHUA_CO_DU_LIEU": 0,
    }
    age_counts = {key: 0 for key in AGE_GROUP_ORDER}
    disability_status_counts = {
        code: 0 for code in DISABILITY_STATUS_LABELS
    }
    school_grouped: dict[str, dict[str, Any]] = {}
    details: list[dict[str, Any]] = []

    summary = {
        "batches": len(batches),
        "with_data": sum(1 for item in batches if counts.get(item.id, {}).get("forms", 0) > 0),
        "locked": sum(1 for item in batches if item.is_locked),
        "forms": sum(counts.get(item.id, {}).get("forms", 0) for item in batches),
        "households": sum(counts.get(item.id, {}).get("households", 0) for item in batches),
        "people": 0,
        "year_records": 0,
        "reviewed_records": 0,
        "missing_year_records": 0,
        "missing_personal_id": 0,
        "male": 0,
        "female": 0,
        "age_5": 0,
        "studying": 0,
        "not_started": 0,
        "dropout": 0,
        "completed_preschool_5": 0,
        "disability": 0,
        "inclusive": 0,
        "permanent": 0,
        "temporary": 0,
    }

    for row in raw_rows:
        batch = batch_by_id.get(int(row["batch_id"]))
        if batch is None:
            continue
        commune_item = commune_grouped[batch.commune_id]

        age = tinh_tuoi_tai_ngay(row.get("date_of_birth"), reference_date)
        age_group = _age_group(age)
        gender_code, gender_label = _normalize_gender(row.get("gender"))
        residency_code = str(row.get("residency_status") or "CHUA_XAC_DINH")
        if residency_code not in residency_counts:
            residency_code = "CHUA_XAC_DINH"
        learning_code = (
            str(row.get("learning_status"))
            if row.get("year_record_id") is not None and row.get("learning_status")
            else "CHUA_CO_DU_LIEU"
        )
        if learning_code not in learning_counts:
            learning_counts[learning_code] = 0
        disability_code = str(
            row.get("disability_status") or "CHUA_XAC_DINH"
        )
        if disability_code not in disability_status_counts:
            disability_code = "CHUA_XAC_DINH"

        school_display = _clean_text(
            row.get("school_name_reported") or row.get("school_name")
        )
        class_display = _clean_text(
            row.get("class_name_reported") or row.get("class_name")
        )
        school_key = ""
        if school_display:
            school_key = (
                f"ID:{row.get('school_id')}"
                if row.get("school_id") is not None
                else f"REPORTED:{school_display.upper()}"
            )
            school_item = school_grouped.setdefault(
                school_key,
                {
                    "school_id": row.get("school_id"),
                    "school_code": row.get("school_code") or "",
                    "school_name": school_display,
                    "commune_name": row.get("commune_name") or "",
                    "forms_set": set(),
                    "households_set": set(),
                    "people": 0,
                    "male": 0,
                    "female": 0,
                    "age_5": 0,
                    "studying": 0,
                    "completed_preschool_5": 0,
                    "disability": 0,
                    "inclusive": 0,
                    "missing_personal_id": 0,
                },
            )
            school_item["forms_set"].add(row["form_id"])
            school_item["households_set"].add(row["household_id"])
            school_item["people"] += 1
            school_item["male"] += 1 if gender_code == "NAM" else 0
            school_item["female"] += 1 if gender_code == "NU" else 0
            school_item["age_5"] += 1 if age == 5 else 0
            school_item["studying"] += 1 if learning_code == "DANG_HOC" else 0
            school_item["completed_preschool_5"] += (
                1 if row.get("completed_preschool_5") is True else 0
            )
            school_item["disability"] += 1 if disability_code == "CO_KHUYET_TAT" else 0
            school_item["inclusive"] += 1 if row.get("inclusive_education") is True else 0
            school_item["missing_personal_id"] += 1 if not _clean_text(row.get("personal_id")) else 0

        summary["people"] += 1
        summary["year_records"] += 1 if row.get("year_record_id") is not None else 0
        summary["reviewed_records"] += 1 if row.get("is_reviewed") is True else 0
        summary["missing_year_records"] += 1 if row.get("year_record_id") is None else 0
        summary["missing_personal_id"] += 1 if not _clean_text(row.get("personal_id")) else 0
        summary["male"] += 1 if gender_code == "NAM" else 0
        summary["female"] += 1 if gender_code == "NU" else 0
        summary["age_5"] += 1 if age == 5 else 0
        summary["studying"] += 1 if learning_code == "DANG_HOC" else 0
        summary["not_started"] += 1 if learning_code == "CHUA_DI_HOC" else 0
        summary["dropout"] += 1 if learning_code in {"BO_HOC", "THOI_HOC"} else 0
        summary["completed_preschool_5"] += 1 if row.get("completed_preschool_5") is True else 0
        summary["disability"] += 1 if disability_code == "CO_KHUYET_TAT" else 0
        summary["inclusive"] += 1 if row.get("inclusive_education") is True else 0
        summary["permanent"] += 1 if residency_code == "THUONG_TRU" else 0
        summary["temporary"] += 1 if residency_code == "TAM_TRU" else 0

        gender_counts[gender_code] += 1
        residency_counts[residency_code] += 1
        learning_counts[learning_code] += 1
        age_counts[age_group] += 1
        disability_status_counts[disability_code] += 1

        commune_item["people"] += 1
        commune_item["male"] += 1 if gender_code == "NAM" else 0
        commune_item["female"] += 1 if gender_code == "NU" else 0
        commune_item["age_5"] += 1 if age == 5 else 0
        commune_item["studying"] += 1 if learning_code == "DANG_HOC" else 0
        commune_item["not_started"] += 1 if learning_code == "CHUA_DI_HOC" else 0
        commune_item["dropout"] += 1 if learning_code in {"BO_HOC", "THOI_HOC"} else 0
        commune_item["completed_preschool_5"] += 1 if row.get("completed_preschool_5") is True else 0
        commune_item["disability"] += 1 if disability_code == "CO_KHUYET_TAT" else 0
        commune_item["inclusive"] += 1 if row.get("inclusive_education") is True else 0
        commune_item["missing_personal_id"] += 1 if not _clean_text(row.get("personal_id")) else 0
        commune_item["missing_year_record"] += 1 if row.get("year_record_id") is None else 0

        details.append(
            {
                **row,
                "age": age,
                "age_group": age_group,
                "date_of_birth_display": _safe_date(row.get("date_of_birth")),
                "survey_date_display": _safe_date(row.get("survey_date")),
                "gender_code": gender_code,
                "gender_label": gender_label,
                "residency_status_label": RESIDENCY_STATUS_LABELS.get(
                    residency_code,
                    residency_code,
                ),
                "learning_status_code": learning_code,
                "learning_status_label": (
                    "Chưa có dữ liệu năm học"
                    if learning_code == "CHUA_CO_DU_LIEU"
                    else LEARNING_STATUS_LABELS.get(learning_code, learning_code)
                ),
                "disability_status_code": disability_code,
                "disability_status_label": DISABILITY_STATUS_LABELS.get(
                    disability_code,
                    disability_code,
                ),
                "disability_type_label": DISABILITY_TYPE_LABELS.get(
                    row.get("disability_type"),
                    "",
                ),
                "disability_level_label": DISABILITY_LEVEL_LABELS.get(
                    row.get("disability_level"),
                    "",
                ),
                "form_status_label": FORM_STATUS_LABELS.get(
                    row.get("form_status"),
                    row.get("form_status") or "",
                ),
                "batch_status_label": BATCH_STATUS_LABELS.get(
                    row.get("batch_status"),
                    row.get("batch_status") or "",
                ),
                "school_display": school_display,
                "class_display": class_display,
            }
        )

    commune_rows = sorted(
        commune_grouped.values(),
        key=lambda item: (item["commune_name"], item["commune_code"]),
    )
    school_rows: list[dict[str, Any]] = []
    for item in school_grouped.values():
        item_copy = {
            key: value
            for key, value in item.items()
            if key not in {"forms_set", "households_set"}
        }
        item_copy["forms"] = len(item["forms_set"])
        item_copy["households"] = len(item["households_set"])
        school_rows.append(item_copy)
    school_rows.sort(
        key=lambda item: (
            item["commune_name"],
            item["school_name"],
            item["school_code"],
        )
    )
    details.sort(
        key=lambda item: (
            item.get("commune_name") or "",
            item.get("hamlet_name") or "",
            item.get("head_name") or "",
            item.get("full_name") or "",
            item.get("person_id") or 0,
        )
    )

    gender_rows = [
        {"code": code, "label": GENDER_LABELS[code], "total": gender_counts[code]}
        for code in GENDER_LABELS
    ]
    residency_rows = [
        {"code": code, "label": label, "total": residency_counts[code]}
        for code, label in RESIDENCY_STATUS_LABELS.items()
    ]
    learning_rows = [
        {
            "code": code,
            "label": (
                "Chưa có dữ liệu năm học"
                if code == "CHUA_CO_DU_LIEU"
                else LEARNING_STATUS_LABELS.get(code, code)
            ),
            "total": total,
        }
        for code, total in learning_counts.items()
        if total > 0
    ]
    learning_rows.sort(key=lambda item: (-item["total"], item["label"]))
    age_rows = [
        {"age_group": key, "total": age_counts[key]}
        for key in AGE_GROUP_ORDER
    ]
    disability_rows = [
        {"code": code, "label": label, "total": disability_status_counts[code]}
        for code, label in DISABILITY_STATUS_LABELS.items()
    ]

    return {
        "summary": summary,
        "commune_rows": commune_rows,
        "school_rows": school_rows,
        "person_rows": details,
        "gender_rows": gender_rows,
        "residency_rows": residency_rows,
        "learning_rows": learning_rows,
        "age_rows": age_rows,
        "disability_rows": disability_rows,
    }


def build_report_data(
    *,
    db: Session,
    request: Request,
    school_year_id: int | None,
    commune_id: int | None,
    school_id: int | None,
    data_state: str,
    q: str,
) -> dict[str, Any]:
    school_year = db.get(SchoolYear, school_year_id) if school_year_id else None
    batches, counts = _load_batches(
        db=db,
        request=request,
        school_year_id=school_year_id,
        commune_id=commune_id,
        school_id=school_id,
        data_state=data_state,
        q=q,
    )
    batch_ids = [item.id for item in batches]
    person_rows = _load_person_rows(
        db=db,
        request=request,
        batch_ids=batch_ids,
        school_year_id=school_year_id,
        school_id=school_id,
    )
    result = _finalize_people(
        batches=batches,
        counts=counts,
        raw_rows=person_rows,
        reference_date=_reference_date(school_year),
    )
    result.update(
        {
            "batches": batches,
            "counts": counts,
            "school_year": school_year,
            "reference_date": _reference_date(school_year),
            "is_official": bool(batches) and all(item.is_locked for item in batches),
        }
    )
    return result


def _build_url(
    *,
    school_year_id: int | None,
    commune_id: int | None,
    school_id: int | None,
    data_state: str,
    q: str,
    page_size: int,
    page: int,
) -> str:
    return "/bao-cao/tong-hop-dieu-tra?" + urlencode(
        {
            "school_year_id": school_year_id or "",
            "commune_id": commune_id or "",
            "school_id": school_id or "",
            "data_state": data_state,
            "q": q,
            "page_size": page_size,
            "page": page,
        }
    )


def _style_title(worksheet: Any, title: str, subtitle: str, last_column: int) -> None:
    dark_fill = PatternFill("solid", fgColor="1F4E78")
    light_fill = PatternFill("solid", fgColor="D9EAF7")
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    worksheet.cell(1, 1, title)
    worksheet.cell(1, 1).fill = dark_fill
    worksheet.cell(1, 1).font = Font(color="FFFFFF", bold=True, size=14)
    worksheet.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 28
    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    worksheet.cell(2, 1, subtitle)
    worksheet.cell(2, 1).fill = light_fill
    worksheet.cell(2, 1).font = Font(italic=True, color="365F91")
    worksheet.cell(2, 1).alignment = Alignment(horizontal="left", vertical="center")


def _style_table(
    worksheet: Any,
    *,
    header_row: int,
    widths: dict[int, float],
) -> None:
    dark_fill = PatternFill("solid", fgColor="1F4E78")
    white_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="C9D4DF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in worksheet[header_row]:
        cell.fill = dark_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in worksheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.auto_filter.ref = f"A{header_row}:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    for index, width in widths.items():
        worksheet.column_dimensions[get_column_letter(index)].width = width


def _append_overview(
    workbook: Workbook,
    *,
    data: dict[str, Any],
    scope_label: str,
    user: dict[str, Any],
    data_state_label: str,
    q: str,
) -> None:
    sheet = workbook.active
    sheet.title = "Tong quan"
    _style_title(
        sheet,
        "BÁO CÁO TỔNG HỢP KẾT QUẢ ĐIỀU TRA PHỔ CẬP GIÁO DỤC",
        "Bài 13A-2 · Tổng hợp theo phạm vi, xã/phường, trường và danh sách đối tượng",
        4,
    )
    sheet.append([])
    sheet.append(["Thông tin", "Giá trị", "Thông tin", "Giá trị"])
    year = data.get("school_year")
    summary = data["summary"]
    report_status = "CHÍNH THỨC" if data["is_official"] else "DỰ THẢO"
    metadata = [
        ("Trạng thái báo cáo", report_status, "Năm học", year.code if year else ""),
        ("Phạm vi", scope_label, "Ngày tham chiếu tính tuổi", _safe_date(data["reference_date"])),
        ("Người lập", user.get("full_name", ""), "Vai trò", user.get("role_name", user.get("role_code", ""))),
        ("Thời điểm xuất", datetime.now().strftime("%d/%m/%Y %H:%M"), "Bộ lọc dữ liệu", data_state_label),
        ("Từ khóa", q or "Không", "Điều kiện chính thức", "Tất cả đợt trong phạm vi đã khóa"),
    ]
    for row in metadata:
        sheet.append(list(row))
    sheet.append([])
    sheet.append(["Chỉ tiêu", "Giá trị", "Chỉ tiêu", "Giá trị"])
    summary_rows = [
        ("Số đợt", summary["batches"], "Đợt đã có dữ liệu", summary["with_data"]),
        ("Đợt đã khóa", summary["locked"], "Tổng số phiếu", summary["forms"]),
        ("Tổng số hộ", summary["households"], "Tổng số đối tượng", summary["people"]),
        ("Có hồ sơ năm học", summary["year_records"], "Thiếu hồ sơ năm học", summary["missing_year_records"]),
        ("Thiếu số định danh", summary["missing_personal_id"], "Đang học", summary["studying"]),
        ("Chưa đi học", summary["not_started"], "Bỏ/thôi học", summary["dropout"]),
        ("Trẻ 5 tuổi", summary["age_5"], "Hoàn thành MN 5 tuổi", summary["completed_preschool_5"]),
        ("Trẻ có khuyết tật", summary["disability"], "Học hòa nhập", summary["inclusive"]),
        ("Nam", summary["male"], "Nữ", summary["female"]),
        ("Thường trú", summary["permanent"], "Tạm trú", summary["temporary"]),
    ]
    for row in summary_rows:
        sheet.append(list(row))
    _style_table(sheet, header_row=4, widths={1: 28, 2: 42, 3: 30, 4: 42})
    # Tô lại hàng tiêu đề nhóm chỉ tiêu thứ hai.
    for cell in sheet[11]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True, color="1F4E78")


def _append_commune_sheet(workbook: Workbook, data: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Theo xa phuong")
    headers = [
        "STT", "Mã xã", "Xã/phường", "Số đợt", "Đợt đã khóa", "Phiếu", "Hộ",
        "Đối tượng", "Nam", "Nữ", "Trẻ 5 tuổi", "Đang học", "Chưa đi học",
        "Bỏ/thôi học", "Hoàn thành MN 5 tuổi", "Khuyết tật", "Học hòa nhập",
        "Thiếu số định danh", "Thiếu hồ sơ năm học",
    ]
    sheet.append(headers)
    for index, row in enumerate(data["commune_rows"], start=1):
        sheet.append([
            index, row["commune_code"], row["commune_name"], row["batch_total"],
            row["locked_batch_total"], row["forms"], row["households"], row["people"],
            row["male"], row["female"], row["age_5"], row["studying"], row["not_started"],
            row["dropout"], row["completed_preschool_5"], row["disability"],
            row["inclusive"], row["missing_personal_id"], row["missing_year_record"],
        ])
    _style_table(
        sheet,
        header_row=1,
        widths={1: 7, 2: 12, 3: 28, 4: 10, 5: 13, 6: 10, 7: 10, 8: 12,
                9: 9, 10: 9, 11: 12, 12: 12, 13: 14, 14: 14, 15: 19,
                16: 12, 17: 15, 18: 18, 19: 20},
    )


def _append_school_sheet(workbook: Workbook, data: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Theo truong")
    headers = [
        "STT", "Mã trường", "Trường", "Xã/phường", "Phiếu", "Hộ", "Đối tượng",
        "Nam", "Nữ", "Trẻ 5 tuổi", "Đang học", "Hoàn thành MN 5 tuổi",
        "Khuyết tật", "Học hòa nhập", "Thiếu số định danh",
    ]
    sheet.append(headers)
    for index, row in enumerate(data["school_rows"], start=1):
        sheet.append([
            index, row["school_code"], row["school_name"], row["commune_name"],
            row["forms"], row["households"], row["people"], row["male"], row["female"],
            row["age_5"], row["studying"], row["completed_preschool_5"],
            row["disability"], row["inclusive"], row["missing_personal_id"],
        ])
    _style_table(
        sheet,
        header_row=1,
        widths={1: 7, 2: 16, 3: 38, 4: 28, 5: 10, 6: 10, 7: 12, 8: 9,
                9: 9, 10: 12, 11: 12, 12: 20, 13: 12, 14: 15, 15: 18},
    )


def _append_detail_sheet(workbook: Workbook, data: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Chi tiet doi tuong")
    headers = [
        "STT", "Mã đợt", "Xã/phường", "Số phiếu", "Mã hộ", "Chủ hộ", "Thôn/xóm",
        "Mã đối tượng", "Họ và tên", "Ngày sinh", "Tuổi", "Giới tính", "Dân tộc",
        "Số định danh", "Mã học sinh trên phiếu", "Cư trú", "Tình trạng học tập",
        "Trường", "Lớp", "Hoàn thành MN 5 tuổi", "Tình trạng khuyết tật",
        "Dạng khuyết tật", "Mức độ khuyết tật", "Học hòa nhập", "Đã rà soát",
        "Thiếu hồ sơ năm học", "Trạng thái phiếu", "Trạng thái dữ liệu đợt",
    ]
    sheet.append(headers)
    for index, row in enumerate(data["person_rows"], start=1):
        sheet.append([
            index, row["batch_code"], row["commune_name"], row["form_number"],
            row["household_code"], row["head_name"], row["hamlet_name"], row["person_code"],
            row["full_name"], row["date_of_birth_display"], row["age"], row["gender_label"],
            row["ethnic_group"], row["personal_id"], row["ministry_student_code"],
            row["residency_status_label"], row["learning_status_label"], row["school_display"],
            row["class_display"], "Có" if row.get("completed_preschool_5") is True else "Không" if row.get("completed_preschool_5") is False else "Chưa xác định",
            row["disability_status_label"], row["disability_type_label"], row["disability_level_label"],
            "Có" if row.get("inclusive_education") is True else "Không" if row.get("inclusive_education") is False else "Chưa xác định",
            "Có" if row.get("is_reviewed") is True else "Chưa",
            "Có" if row.get("year_record_id") is None else "Không",
            row["form_status_label"], "Đã khóa" if row.get("batch_is_locked") else "Đang mở",
        ])
    _style_table(
        sheet,
        header_row=1,
        widths={1: 7, 2: 24, 3: 25, 4: 24, 5: 20, 6: 24, 7: 20, 8: 18,
                9: 25, 10: 14, 11: 9, 12: 12, 13: 15, 14: 18, 15: 20, 16: 16,
                17: 20, 18: 34, 19: 18, 20: 20, 21: 22, 22: 22, 23: 20, 24: 15,
                25: 14, 26: 20, 27: 18, 28: 20},
    )


def _append_info_sheet(
    workbook: Workbook,
    *,
    data: dict[str, Any],
    scope_label: str,
    data_state_label: str,
    q: str,
) -> None:
    sheet = workbook.create_sheet("Thong tin bao cao")
    sheet.append(["Nội dung", "Giá trị"])
    year = data.get("school_year")
    items = [
        ("Bài", "13A-2 V1 – Báo cáo tổng hợp điều tra"),
        ("Năm học", year.code if year else ""),
        ("Phạm vi", scope_label),
        ("Trạng thái dữ liệu", data_state_label),
        ("Từ khóa", q or "Không"),
        ("Ngày tham chiếu tính tuổi", _safe_date(data["reference_date"])),
        ("Trạng thái tệp", "CHÍNH THỨC" if data["is_official"] else "DỰ THẢO"),
        ("Nguyên tắc", "Chỉ đọc và tổng hợp dữ liệu hiện có; không tự sửa dữ liệu gốc."),
        ("Ghi chú chính thức", "Chỉ được coi là chính thức khi tất cả đợt trong phạm vi đã khóa số liệu."),
        ("Thời điểm tạo tệp", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]
    for item in items:
        sheet.append(list(item))
    _style_table(sheet, header_row=1, widths={1: 30, 2: 95})


@router.get("", response_class=HTMLResponse)
def summary_report_page(
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
    data_state: str = "",
    q: str = "",
    page_size: int = 20,
    page: int = 1,
    db: Session = Depends(get_db),
):
    user = lay_thong_tin_nguoi_dung(request)
    school_years = _load_school_years(db)
    selected_year_id = _choose_year_id(
        school_years,
        _parse_optional_int(school_year_id),
    )
    selected_year = next(
        (item for item in school_years if item.id == selected_year_id),
        None,
    )
    normalized_state = _clean_text(data_state).upper()
    if normalized_state not in DATA_STATE_LABELS:
        normalized_state = ""
    q = _clean_text(q)[:120]
    if page_size not in PAGE_SIZE_OPTIONS:
        page_size = 20

    (
        communes,
        schools,
        selected_commune_id,
        selected_school_id,
        scope_label,
    ) = _load_scope_options(
        db,
        user,
        _parse_optional_int(commune_id),
        _parse_optional_int(school_id),
    )

    data = build_report_data(
        db=db,
        request=request,
        school_year_id=selected_year_id,
        commune_id=selected_commune_id,
        school_id=selected_school_id,
        data_state=normalized_state,
        q=q,
    )

    all_people = data["person_rows"]
    total_records = len(all_people)
    total_pages = max(1, ceil(total_records / page_size))
    page = max(1, min(page, total_pages))
    start = (page - 1) * page_size
    page_rows = all_people[start:start + page_size]
    page_links = [
        {
            "number": number,
            "url": _build_url(
                school_year_id=selected_year_id,
                commune_id=selected_commune_id,
                school_id=selected_school_id,
                data_state=normalized_state,
                q=q,
                page_size=page_size,
                page=number,
            ),
        }
        for number in range(max(1, page - 2), min(total_pages, page + 2) + 1)
    ]
    export_params = urlencode(
        {
            "school_year_id": selected_year_id or "",
            "commune_id": selected_commune_id or "",
            "school_id": selected_school_id or "",
            "data_state": normalized_state,
            "q": q,
        }
    )

    return templates.TemplateResponse(
        request=request,
        name="reports/survey_summary_report.html",
        context={
            "nguoi_dung": user,
            "school_years": school_years,
            "selected_year": selected_year,
            "selected_year_id": selected_year_id,
            "communes": communes,
            "schools": schools,
            "selected_commune_id": selected_commune_id,
            "selected_school_id": selected_school_id,
            "scope_label": scope_label,
            "data_state_labels": DATA_STATE_LABELS,
            "data_state": normalized_state,
            "q": q,
            "page_size": page_size,
            "page_size_options": PAGE_SIZE_OPTIONS,
            "page": page,
            "total_pages": total_pages,
            "page_links": page_links,
            "previous_url": (
                _build_url(
                    school_year_id=selected_year_id,
                    commune_id=selected_commune_id,
                    school_id=selected_school_id,
                    data_state=normalized_state,
                    q=q,
                    page_size=page_size,
                    page=page - 1,
                )
                if page > 1
                else None
            ),
            "next_url": (
                _build_url(
                    school_year_id=selected_year_id,
                    commune_id=selected_commune_id,
                    school_id=selected_school_id,
                    data_state=normalized_state,
                    q=q,
                    page_size=page_size,
                    page=page + 1,
                )
                if page < total_pages
                else None
            ),
            "export_url": "/bao-cao/tong-hop-dieu-tra/xuat-excel?" + export_params,
            "province_reader": is_province_reader(user.get("role_code")),
            "person_rows": page_rows,
            "total_person_rows": total_records,
            "commune_rows": data["commune_rows"],
            "school_rows": data["school_rows"],
            "summary": data["summary"],
            "gender_rows": data["gender_rows"],
            "residency_rows": data["residency_rows"],
            "learning_rows": data["learning_rows"],
            "age_rows": data["age_rows"],
            "disability_rows": data["disability_rows"],
            "reference_date": data["reference_date"],
            "is_official": data["is_official"],
            "generated_at": datetime.now().strftime("%d/%m/%Y %H:%M"),
        },
    )


@router.get("/xuat-excel")
def export_summary_report_excel(
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
    data_state: str = "",
    q: str = "",
    db: Session = Depends(get_db),
):
    user = lay_thong_tin_nguoi_dung(request)
    school_years = _load_school_years(db)
    selected_year_id = _choose_year_id(
        school_years,
        _parse_optional_int(school_year_id),
    )
    normalized_state = _clean_text(data_state).upper()
    if normalized_state not in DATA_STATE_LABELS:
        normalized_state = ""
    q = _clean_text(q)[:120]

    (
        _communes,
        _schools,
        selected_commune_id,
        selected_school_id,
        scope_label,
    ) = _load_scope_options(
        db,
        user,
        _parse_optional_int(commune_id),
        _parse_optional_int(school_id),
    )
    data = build_report_data(
        db=db,
        request=request,
        school_year_id=selected_year_id,
        commune_id=selected_commune_id,
        school_id=selected_school_id,
        data_state=normalized_state,
        q=q,
    )

    workbook = Workbook()
    _append_overview(
        workbook,
        data=data,
        scope_label=scope_label,
        user=user,
        data_state_label=DATA_STATE_LABELS[normalized_state],
        q=q,
    )
    _append_commune_sheet(workbook, data)
    _append_school_sheet(workbook, data)
    _append_detail_sheet(workbook, data)
    _append_info_sheet(
        workbook,
        data=data,
        scope_label=scope_label,
        data_state_label=DATA_STATE_LABELS[normalized_state],
        q=q,
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    year_code = data["school_year"].code if data.get("school_year") else "khong_ro"
    safe_year = str(year_code).replace("/", "-")
    suffix = "chinh_thuc" if data["is_official"] else "du_thao"
    filename = f"bao_cao_tong_hop_dieu_tra_{safe_year}_{suffix}.xlsx"
    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )
