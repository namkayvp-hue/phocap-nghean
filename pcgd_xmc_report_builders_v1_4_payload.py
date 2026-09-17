from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO
from pathlib import Path
import re
import unicodedata
from typing import Any

from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models import Classroom, Commune, School, SchoolYear
from app.staff_models import StaffMember, StaffYearRecord
from app.permissions import (
    COMMUNE_ROLE_CODE,
    SCHOOL_ROLE_CODE,
    TEACHER_ROLE_CODE,
    is_province_reader,
    normalize_role_code,
)
from app.survey_models import SurveyPerson, SurveyPersonYearRecord


APP_DIR = Path(__file__).resolve().parent
TEMPLATE_ROOT = APP_DIR / "report_templates" / "pcgd_xmc_2025"

REPORT_TEMPLATE_FILES = {
    "PCGD_TH_02_2025": "PCGD_2025_TH_02.xlsx",
    "PCGD_TH_01_GV_2025": "PCGD_2025_TH_01_GV.xlsx",
    "PCGD_TH_01_CSVC_2025": "PCGD_2025_TH_01_CSVC.xlsx",
    "PCGD_THCS_M1_2025": "PCGD_2025_THCS_M1.xlsx",
    "PCGD_THCS_M2_2025": "PCGD_2025_THCS_M2.xlsx",
    "PCGD_THCS_TK_2025": "PCGD_2025_THCS_TK.xlsx",
    "PCGD_THCS_M5_2025": "PCGD_2025_THCS_M5.xlsx",
    "PCGD_THCS_CSVC_2025": "PCGD_2025_THCS_CSVC.xlsx",
    "PCGD_XMC_3_2025": "PCGD_2025_XMC_3.xlsx",
    "PCGD_CMC_2_2025": "PCGD_2025_CMC_2.xlsx",
    "PCGD_CMC_1_2025": "PCGD_2025_CMC_1.xlsx",
    "PCGD_XMC_4_2025": "PCGD_2025_XMC_4.xlsx",
}

SUPPORTED_REPORT_TYPES = frozenset(REPORT_TEMPLATE_FILES)
LOWER_SCOPE_REPORT_VERSION = "PCGD-XMC-REPORT-LOWER-SCOPES-V1.4"


def _norm(value: Any) -> str:
    text = str(value or "").strip().upper().replace("_", " ")
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return " ".join(text.split())


def _reference_year(school_year: SchoolYear | None) -> int:
    if school_year is not None:
        match = re.search(r"(20\d{2})", str(school_year.code or ""))
        if match:
            return int(match.group(1))
    return datetime.now().year


def _safe_filename_text(value: Any) -> str:
    text = _norm(value)
    text = re.sub(r"[^A-Z0-9]+", "_", text).strip("_")
    return text or "TOAN_TINH"


def _clean_scope_label(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _school_scope_label(name: Any) -> str:
    text = _clean_scope_label(name)
    if not text:
        return "Trường"
    return text if _norm(text).startswith("TRUONG ") else f"Trường {text}"


def _effective_scope_for_user(
    *,
    db: Session,
    request: Any,
    selected_commune_id: int | None,
    selected_school_id: int | None,
    scope_label: str,
) -> tuple[int | None, int | None, str, str]:
    """
    Cố định phạm vi xuất theo tài khoản mà không thay đổi route/luồng cũ.

    - Sở/Phòng ban: giữ phạm vi đã được report_center kiểm tra.
    - Xã/phường: luôn giới hạn trong xã của tài khoản; vẫn cho phép chọn
      một trường thuộc đúng xã nếu giao diện hiện tại đã cho chọn.
    - Trường/Giáo viên: luôn khóa về đúng trường của tài khoản.

    Hàm chỉ đọc auth_user và danh mục; không ghi cơ sở dữ liệu.
    """
    user = getattr(request, "scope", {}).get("auth_user") or {}
    role_code = normalize_role_code(user.get("role_code"))
    account_commune_id = user.get("commune_id")
    account_school_id = user.get("school_id")

    try:
        account_commune_id = int(account_commune_id) if account_commune_id is not None else None
    except (TypeError, ValueError):
        account_commune_id = None
    try:
        account_school_id = int(account_school_id) if account_school_id is not None else None
    except (TypeError, ValueError):
        account_school_id = None

    if is_province_reader(role_code):
        return (
            selected_commune_id,
            selected_school_id,
            _clean_scope_label(scope_label) or "Toàn tỉnh",
            "province",
        )

    if role_code == COMMUNE_ROLE_CODE:
        commune = db.get(Commune, account_commune_id) if account_commune_id else None
        effective_commune_id = int(commune.id) if commune is not None else account_commune_id
        effective_school_id: int | None = None
        if selected_school_id is not None:
            school = db.get(School, int(selected_school_id))
            if (
                school is not None
                and school.is_active
                and effective_commune_id is not None
                and int(school.commune_id) == int(effective_commune_id)
            ):
                effective_school_id = int(school.id)
        if effective_school_id is not None:
            school = db.get(School, effective_school_id)
            label = _school_scope_label(school.name) if school is not None else _clean_scope_label(scope_label)
        else:
            label = commune.name if commune is not None else (_clean_scope_label(scope_label) or "Cấp xã/phường")
        return effective_commune_id, effective_school_id, label, "commune"

    if role_code in {SCHOOL_ROLE_CODE, TEACHER_ROLE_CODE}:
        school = db.get(School, account_school_id) if account_school_id else None
        if school is None:
            return account_commune_id, account_school_id, (_clean_scope_label(scope_label) or str(user.get("unit_name") or "Phạm vi tài khoản")), "school"
        return (
            int(school.commune_id) if school.commune_id is not None else account_commune_id,
            int(school.id),
            _school_scope_label(school.name),
            "school",
        )

    # Vai trò khác: không nới phạm vi ngoài giá trị đã được router kiểm tra.
    return selected_commune_id, selected_school_id, (_clean_scope_label(scope_label) or "Phạm vi tài khoản"), "account"


def _scope_meta(
    *,
    scope_label: str,
    selected_commune_id: int | None,
    selected_school_id: int | None,
    year: int,
) -> dict[str, str]:
    raw = str(scope_label or "").strip() or "Toàn tỉnh"
    province = selected_commune_id is None and selected_school_id is None
    school = selected_school_id is not None
    if province:
        return {
            "title": "Toàn tỉnh",
            "place": "Nghệ An",
            "signature": "XÁC NHẬN CỦA SỞ GIÁO DỤC VÀ ĐÀO TẠO",
            "signer": "GIÁM ĐỐC",
            "year": str(year),
            "kind": "province",
        }
    if school:
        return {
            "title": raw,
            "place": raw,
            "signature": "XÁC NHẬN CỦA NHÀ TRƯỜNG",
            "signer": "HIỆU TRƯỞNG",
            "year": str(year),
            "kind": "school",
        }
    return {
        "title": raw,
        "place": raw,
        "signature": "XÁC NHẬN CỦA UBND XÃ/PHƯỜNG",
        "signer": "PHÓ CHỦ TỊCH",
        "year": str(year),
        "kind": "commune",
    }


def _age(person: SurveyPerson, year: int) -> int | None:
    if person.date_of_birth is None:
        return None
    return year - int(person.date_of_birth.year)


def _female(person: SurveyPerson) -> bool:
    return _norm(person.gender) in {"NU", "F", "FEMALE", "GIRL"}


def _ethnic(person: SurveyPerson) -> bool:
    value = _norm(person.ethnic_group)
    return bool(value) and value not in {"KINH", "K"}


def _female_ethnic(person: SurveyPerson) -> bool:
    return _female(person) and _ethnic(person)


def _grade(record: SurveyPersonYearRecord | None) -> int | None:
    if record is None:
        return None
    name = ""
    if record.classroom is not None:
        name = str(record.classroom.name or "")
    if not name:
        name = str(record.class_name_reported or "")
    text = _norm(name)
    match = re.search(r"(?<!\d)(1[0-2]|[1-9])(?=[A-Z]|\b|\s|/|\-|$)", text)
    if not match:
        return None
    value = int(match.group(1))
    return value if 1 <= value <= 12 else None


def _status(record: SurveyPersonYearRecord | None) -> str:
    return _norm(record.learning_status if record is not None else "")


def _is_disabled(person: SurveyPerson, record: SurveyPersonYearRecord | None) -> bool:
    if record is not None and _norm(record.disability_status) in {"CO KHUYET TAT", "CO_KHUYET_TAT"}:
        return True
    student = getattr(person, "student", None)
    return bool(student is not None and str(getattr(student, "disability_type", "") or "").strip())


def _disabled_access(person: SurveyPerson, record: SurveyPersonYearRecord | None) -> bool:
    if not _is_disabled(person, record):
        return False
    return _status(record) in {"DANG HOC", "DANG_HOC", "CHUYEN DEN", "CHUYEN_DEN", "CHUYEN DI", "CHUYEN_DI", "DA TOT NGHIEP", "DA_TOT_NGHIEP"}


def _completed_primary(person: SurveyPerson, record: SurveyPersonYearRecord | None) -> bool:
    if record is None:
        return False
    status = _status(record)
    if status in {"DA TOT NGHIEP", "DA_TOT_NGHIEP", "HOAN THANH", "HTCTTH"}:
        return True
    grade = _grade(record)
    if grade is not None and grade >= 6:
        return True
    note = _norm(" ".join([str(record.notes or ""), str(person.notes or ""), str(person.special_circumstances or "")]))
    return "HOAN THANH CHUONG TRINH TIEU HOC" in note or "HTCTTH" in note


def _completed_secondary(person: SurveyPerson, record: SurveyPersonYearRecord | None) -> bool:
    if record is None:
        return False
    status = _status(record)
    if status in {"TN THCS", "TN_THCS", "TOT NGHIEP THCS", "DA TOT NGHIEP THCS", "DA_TOT_NGHIEP_THCS"}:
        return True
    grade = _grade(record)
    if grade is not None and grade >= 10:
        return True
    note = _norm(" ".join([str(record.notes or ""), str(person.notes or ""), str(person.special_circumstances or "")]))
    return "TOT NGHIEP THCS" in note or "TNTHCS" in note


def _upper_secondary(person: SurveyPerson, record: SurveyPersonYearRecord | None) -> bool:
    if record is None:
        return False
    grade = _grade(record)
    status = _status(record)
    if status in {"BO HOC", "BO_HOC", "THOI HOC", "THOI_HOC", "CHUA DI HOC", "CHUA_DI_HOC"}:
        return False
    if grade is not None and grade >= 10:
        return True
    text = _norm(" ".join([str(record.school_name_reported or ""), str(record.notes or "")]))
    return any(token in text for token in ("THPT", "GDTX", "GDNN", "TRUNG HOC PHO THONG"))


def _is_ppc(person: SurveyPerson, record: SurveyPersonYearRecord | None) -> bool:
    if _norm(person.residency_status) not in {"THUONG TRU", "THUONG_TRU"}:
        return False
    return _status(record) not in {"KHONG THUOC DIEN", "KHONG_THUOC_DIEN"}


def _location_bucket(person: SurveyPerson, record: SurveyPersonYearRecord | None, reporting_commune_id: int | None) -> str:
    home_commune_id = None
    household = getattr(person, "household", None)
    if household is not None and getattr(household, "commune_id", None) is not None:
        home_commune_id = int(household.commune_id)
    school_commune_id = None
    if record is not None and record.school is not None and record.school.commune_id is not None:
        school_commune_id = int(record.school.commune_id)
    if reporting_commune_id is None:
        if record is not None and record.school_id is None and str(record.school_name_reported or "").strip():
            return "other"
        return "local"
    if home_commune_id == reporting_commune_id:
        if school_commune_id is None:
            return "other" if record is not None and str(record.school_name_reported or "").strip() else "local"
        return "local" if school_commune_id == reporting_commune_id else "other"
    if school_commune_id == reporting_commune_id:
        return "inbound"
    return "other"


def _percent(numerator: int | float, denominator: int | float) -> float | None:
    if not denominator:
        return None
    return round(float(numerator) * 100.0 / float(denominator), 2)


def _load_people_and_records(
    *, db: Session, request: Any, school_year_id: int, selected_commune_id: int | None, selected_school_id: int | None
) -> list[tuple[SurveyPerson, SurveyPersonYearRecord | None]]:
    from app.routers import report_center as rc
    return rc._th_m1_load_people_and_records(
        db=db,
        request=request,
        school_year_id=school_year_id,
        selected_commune_id=selected_commune_id,
        selected_school_id=selected_school_id,
    )


def _population_metrics(people_and_records: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year: int) -> dict[int, dict[str, int]]:
    metrics: dict[int, dict[str, int]] = defaultdict(lambda: {"total": 0, "female": 0, "ethnic": 0, "female_ethnic": 0, "disabled": 0, "disabled_access": 0, "ppc": 0})
    for person, record in people_and_records:
        age = _age(person, year)
        if age is None:
            continue
        item = metrics[age]
        item["total"] += 1
        item["female"] += int(_female(person))
        item["ethnic"] += int(_ethnic(person))
        item["female_ethnic"] += int(_female_ethnic(person))
        item["disabled"] += int(_is_disabled(person, record))
        item["disabled_access"] += int(_disabled_access(person, record))
        item["ppc"] += int(_is_ppc(person, record))
    return metrics


def _sum_metric(metrics: dict[int, dict[str, int]], ages: range | list[int] | tuple[int, ...], key: str) -> int:
    return sum(int(metrics.get(age, {}).get(key, 0)) for age in ages)


def _write_scope_signature(ws: Any, report_type: str, meta: dict[str, str]) -> None:
    year = meta["year"]
    title = meta["title"]
    place = meta["place"]
    sig = meta["signature"]
    signer = meta["signer"]
    province = meta["kind"] == "province"
    mappings = {
        "PCGD_TH_02_2025": ("A2", "K10", "K11", "K12", ("A18", "K18")),
        "PCGD_TH_01_GV_2025": ("A2", "O25", "O26", "O27", ("B33", "O33")),
        "PCGD_TH_01_CSVC_2025": ("A2", "T20", "T21", "T22", ("C27", "T27")),
        "PCGD_THCS_M1_2025": ("A2", "M43", "M44", "M45", ("D51", "M51")),
        "PCGD_THCS_M2_2025": ("A3", "K27", "K28", "K29", ("A35", "K35")),
        "PCGD_THCS_TK_2025": ("A2", "K10", "K11", "K12", ("A18", "K18")),
        "PCGD_THCS_M5_2025": ("A3", "AB25", "AB26", "AB27", ("F33", "AB33")),
        "PCGD_THCS_CSVC_2025": ("A3", "K20", "K21", "K22", ("A28", "K28")),
        "PCGD_XMC_3_2025": ("A2", "H59", "H60", "H61", ("A67", "H67")),
        "PCGD_CMC_2_2025": ("A2", "K13", "K14", "K15", ("A21", "K21")),
        "PCGD_CMC_1_2025": ("A2", "X11", "X12", "X13", ("C19", "X19")),
        "PCGD_XMC_4_2025": ("A2", "K10", "K11", "K12", ("A18", "K18")),
    }
    loc_cell, date_cell, sig_cell, signer_cell, clear_cells = mappings[report_type]
    if report_type == "PCGD_CMC_2_2025":
        ws[loc_cell] = f"{title}, Tỉnh: Nghệ An" if not province else "Toàn tỉnh, Tỉnh: Nghệ An"
    elif report_type == "PCGD_CMC_1_2025":
        ws[loc_cell] = f"{title}, tỉnh Nghệ An" if not province else "Toàn tỉnh, tỉnh Nghệ An"
    elif report_type in {"PCGD_THCS_M1_2025", "PCGD_THCS_M2_2025", "PCGD_THCS_M5_2025", "PCGD_THCS_CSVC_2025", "PCGD_XMC_3_2025", "PCGD_XMC_4_2025"}:
        ws[loc_cell] = title if province else title
    else:
        ws[loc_cell] = title
    ws[date_cell] = f"{place}, ngày      tháng      năm {year}"
    ws[sig_cell] = sig
    ws[signer_cell] = signer
    for ref in clear_cells:
        ws[ref] = None


def _schools_and_classes(db: Session, school_year_id: int, selected_commune_id: int | None, selected_school_id: int | None) -> tuple[list[School], dict[int, list[int]]]:
    stmt = select(School).where(School.is_active.is_(True)).order_by(School.name.asc())
    if selected_school_id is not None:
        stmt = stmt.where(School.id == selected_school_id)
    elif selected_commune_id is not None:
        stmt = stmt.where(School.commune_id == selected_commune_id)
    schools = list(db.scalars(stmt).all())
    if not schools:
        return [], {}
    class_stmt = select(Classroom).where(Classroom.school_year_id == school_year_id, Classroom.school_id.in_([s.id for s in schools]), Classroom.is_active.is_(True))
    classes = list(db.scalars(class_stmt).all())
    grades: dict[int, list[int]] = defaultdict(list)
    for item in classes:
        text = _norm(item.name)
        m = re.search(r"(?<!\d)(1[0-2]|[1-9])(?=[A-Z]|\b|\s|/|\-|$)", text)
        if m:
            grades[int(item.school_id)].append(int(m.group(1)))
    return schools, grades


def _school_level(school: School, grades: dict[int, list[int]]) -> str | None:
    name = _norm(school.name)
    g = grades.get(int(school.id), [])
    if any(6 <= x <= 9 for x in g) or "THCS" in name or "TRUNG HOC CO SO" in name:
        return "THCS"
    if any(1 <= x <= 5 for x in g) or "TIEU HOC" in name or re.search(r"(^| )TH( |$)", name):
        return "TH"
    return None


def _staff_rows(db: Session, school_year_id: int, school_ids: list[int]) -> dict[int, list[StaffYearRecord]]:
    if not school_ids:
        return {}
    stmt = (
        select(StaffYearRecord)
        .where(StaffYearRecord.school_year_id == school_year_id, StaffYearRecord.school_id.in_(school_ids), StaffYearRecord.is_active.is_(True))
        .options(selectinload(StaffYearRecord.staff_member))
    )
    result: dict[int, list[StaffYearRecord]] = defaultdict(list)
    for row in db.scalars(stmt).all():
        result[int(row.school_id)].append(row)
    return result


def _class_count(grades: list[int], level: str) -> int:
    if level == "TH":
        return sum(1 for x in grades if 1 <= x <= 5)
    return sum(1 for x in grades if 6 <= x <= 9)


def _staff_summary(records: list[StaffYearRecord], class_count: int) -> dict[str, Any]:
    teachers = [r for r in records if _norm(r.position_group) == "GIAO VIEN"]
    managers = [r for r in records if _norm(r.position_group) == "CBQL"]
    employees = [r for r in records if _norm(r.position_group) == "NHAN VIEN"]
    def member(r: StaffYearRecord) -> StaffMember | None:
        return getattr(r, "staff_member", None)
    principal = sum(1 for r in managers if "HIEU TRUONG" in _norm(r.position_title) and "PHO" not in _norm(r.position_title))
    vice = sum(1 for r in managers if "PHO HIEU TRUONG" in _norm(r.position_title) or "P HIEU TRUONG" in _norm(r.position_title))
    female = sum(1 for r in teachers if member(r) is not None and _norm(member(r).gender) == "NU")
    ethnic = sum(
        1
        for r in teachers
        if member(r) is not None
        and bool(_norm(getattr(member(r), "ethnic_group", "")))
        and _norm(getattr(member(r), "ethnic_group", "")) != "KINH"
    )
    bienche = sum(1 for r in teachers if _norm(r.employment_type) == "BIEN CHE")
    hopdong = len(teachers) - bienche
    qual = defaultdict(int)
    prof = defaultdict(int)
    subjects = defaultdict(int)
    emp_roles = defaultdict(int)
    for r in teachers:
        q = _norm(r.qualification_level)
        if q in {"THAC SI", "TIEN SI"}:
            qual["postgrad"] += 1
        elif q == "DAI HOC":
            qual["univ"] += 1
        elif q == "CAO DANG":
            qual["college"] += 1
        elif q == "TRUNG CAP":
            qual["secondary"] += 1
        ps = _norm(r.professional_standard)
        if ps == "TOT": prof["excellent"] += 1
        elif ps == "KHA": prof["good"] += 1
        elif ps == "DAT": prof["average"] += 1
        elif ps == "CHUA DAT": prof["poor"] += 1
        text = _norm(" ".join([str(r.position_title or ""), str(r.notes or "")]))
        mapping = {
            "TOAN": "math", "NGU VAN": "literature", "VAN": "literature", "VAT LY": "physics", "LY": "physics", "HOA": "chemistry", "SINH": "biology", "LICH SU": "history", "SU": "history", "DIA": "geography", "NHAC": "music", "AM NHAC": "music", "MY THUAT": "art", "MT": "art", "THE DUC": "pe", "TD": "pe", "GDCD": "civic", "CONG NGHE": "technology", "TIN HOC": "it", "TIN": "it", "TIENG ANH": "english", "ANH": "english", "TIENG NGA": "russian", "TIENG PHAP": "french", "NGOAI NGU": "foreign",
        }
        matched = False
        for token, key in mapping.items():
            if re.search(rf"(^| )({re.escape(token)})( |$)", text):
                subjects[key] += 1
                matched = True
                break
        if not matched:
            subjects["general"] += 1
    for r in employees:
        text = _norm(" ".join([str(r.position_title or ""), str(r.notes or "")]))
        if "THU VIEN" in text: emp_roles["library"] += 1
        elif "THIET BI" in text or "THI NGHIEM" in text: emp_roles["equipment"] += 1
        elif "Y TE" in text: emp_roles["health"] += 1
        elif "VAN PHONG" in text or "VAN THU" in text: emp_roles["office"] += 1
        else: emp_roles["other"] += 1
    return {
        "principal": principal, "vice": vice, "teachers": len(teachers), "bienche": bienche, "hopdong": hopdong,
        "female": female, "ethnic": ethnic, "ratio": round(len(teachers) / class_count, 2) if class_count else None,
        "qual": qual, "prof": prof, "subjects": subjects, "employees": emp_roles,
    }


def _clear_range(ws: Any, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if cell.__class__.__name__ == "MergedCell":
                continue
            cell.value = None


def _build_th02(ws: Any, people: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year: int, meta: dict[str, str], school_count: int) -> None:
    _clear_range(ws, 8, 8, 1, 20)
    age6 = [pr for pr in people if _age(pr[0], year) == 6]
    age11 = [pr for pr in people if _age(pr[0], year) == 11]
    age11_14 = [pr for pr in people if (_age(pr[0], year) or -1) in range(11, 15)]
    students = [pr for pr in people if (_grade(pr[1]) or 0) in range(1, 6) and _status(pr[1]) not in {"BO HOC", "THOI HOC", "CHUA DI HOC"}]
    disabled_students = sum(1 for p, r in students if _is_disabled(p, r))
    ppc6 = sum(1 for p, r in age6 if _is_ppc(p, r))
    grade1_age6 = sum(1 for p, r in age6 if _is_ppc(p, r) and _grade(r) == 1)
    ppc11 = sum(1 for p, r in age11 if _is_ppc(p, r))
    comp11 = sum(1 for p, r in age11 if _is_ppc(p, r) and _completed_primary(p, r))
    ppc1114 = sum(1 for p, r in age11_14 if _is_ppc(p, r))
    comp1114 = sum(1 for p, r in age11_14 if _is_ppc(p, r) and _completed_primary(p, r))
    disabled = sum(1 for p, r in people if 6 <= (_age(p, year) or -1) <= 14 and _is_disabled(p, r))
    disabled_access = sum(1 for p, r in people if 6 <= (_age(p, year) or -1) <= 14 and _disabled_access(p, r))
    row = 8
    values = {1: 1, 2: meta["title"], 3: 1 if meta["kind"] != "province" else None, 4: school_count, 6: len(students), 7: disabled_students, 8: grade1_age6, 9: _percent(grade1_age6, ppc6), 10: comp11, 11: _percent(comp11, ppc11), 12: comp1114, 13: _percent(comp1114, ppc1114), 14: disabled, 16: disabled_access, 17: _percent(disabled_access, disabled)}
    for col, value in values.items(): ws.cell(row, col).value = value


def _build_thcs_m1(ws: Any, people: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year: int, meta: dict[str, str], commune_id: int | None) -> None:
    age_cols = {11: 6, 12: 8, 13: 10, 14: 12, 15: 14, 16: 17, 17: 20, 18: 23}
    for row in list(range(6, 13)) + list(range(14, 35)):
        for col in range(6, 27):
            if row == 13:
                continue
            cell = ws.cell(row, col)
            if cell.__class__.__name__ == "MergedCell":
                continue
            cell.value = None
    by_age = defaultdict(list)
    for p, r in people:
        age = _age(p, year)
        if age in age_cols: by_age[age].append((p, r))
    for age, col in age_cols.items():
        group = by_age[age]
        ws.cell(4, col).value = year - age
        ws.cell(6, col).value = len(group) or None
        ws.cell(7, col).value = sum(1 for p, r in group if _female(p)) or None
        ws.cell(8, col).value = sum(1 for p, r in group if _ethnic(p)) or None
        ws.cell(9, col).value = sum(1 for p, r in group if _is_disabled(p, r)) or None
        ws.cell(11, col).value = sum(1 for p, r in group if _disabled_access(p, r)) or None
        ws.cell(12, col).value = sum(1 for p, r in group if _is_ppc(p, r)) or None
    for row_start, grade in ((14,6),(17,7),(20,8),(23,9)):
        for p, r in people:
            age = _age(p, year)
            if age not in age_cols or _grade(r) != grade or not _is_ppc(p, r): continue
            bucket = _location_bucket(p, r, commune_id)
            rr = row_start + {"local":0,"other":1,"inbound":2}[bucket]
            col = age_cols[age]
            ws.cell(rr, col).value = int(ws.cell(rr, col).value or 0) + 1
    for p, r in people:
        age = _age(p, year)
        if age not in age_cols: continue
        col = age_cols[age]
        bucket = _location_bucket(p, r, commune_id)
        off = {"local":0,"other":1,"inbound":2}[bucket]
        if _completed_secondary(p, r):
            rr = 26 + off; ws.cell(rr,col).value = int(ws.cell(rr,col).value or 0)+1
        if _upper_secondary(p, r):
            rr = 29 + off; ws.cell(rr,col).value = int(ws.cell(rr,col).value or 0)+1
        if _is_ppc(p, r) and _status(r) in {"BO HOC","BO_HOC","THOI HOC","THOI_HOC"}:
            rr = 32 + off; ws.cell(rr,col).value = int(ws.cell(rr,col).value or 0)+1
    for row in range(6,35):
        if row == 13: continue
        ws.cell(row,26).value = sum(int(ws.cell(row,c).value or 0) for c in (14,17,20,23)) or None
    age15_18 = [(p,r) for p,r in people if (_age(p,year) or -1) in range(15,19) and _is_ppc(p,r)]
    sec = sum(1 for p,r in age15_18 if _completed_secondary(p,r))
    upper = sum(1 for p,r in age15_18 if _upper_secondary(p,r))
    disabled = sum(1 for p,r in age15_18 if _is_disabled(p,r))
    access = sum(1 for p,r in age15_18 if _disabled_access(p,r))
    ws["G39"] = sec; ws["H39"] = _percent(sec, len(age15_18))
    ws["G40"] = upper; ws["H40"] = _percent(upper, len(age15_18))
    ws["G41"] = access; ws["H41"] = _percent(access, disabled)
    ws["G42"] = sum(1 for p,r in people if (_grade(r) or 0) in range(6,10) and _status(r) not in {"BO HOC","THOI HOC","CHUA DI HOC"})
    ws["G36"] = None; ws["G37"] = None


def _build_thcs_m2(ws: Any, people: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year: int, meta: dict[str,str]) -> None:
    _clear_range(ws, 14, 15, 1, 23)
    age6=[pr for pr in people if _age(pr[0],year)==6 and _is_ppc(*pr)]
    age11_14=[pr for pr in people if (_age(pr[0],year) or -1) in range(11,15) and _is_ppc(*pr)]
    age15_18=[pr for pr in people if (_age(pr[0],year) or -1) in range(15,19) and _is_ppc(*pr)]
    grade1=sum(1 for p,r in age6 if _grade(r)==1)
    htth=sum(1 for p,r in age11_14 if _completed_primary(p,r))
    tnthcs=sum(1 for p,r in age15_18 if _completed_secondary(p,r))
    row=14
    vals={1:1,2:meta["title"],3:len(age6),4:grade1,5:_percent(grade1,len(age6)),11:len(age11_14),12:htth,13:_percent(htth,len(age11_14)),18:len(age15_18),19:tnthcs,21:tnthcs,22:_percent(tnthcs,len(age15_18))}
    for c,v in vals.items(): ws.cell(row,c).value=v
    ws.cell(15,1).value="Tổng"
    for c in range(3,23): ws.cell(15,c).value=ws.cell(14,c).value
    for ref in ("P20","P21","P22","P23","P24","P25","W14"): ws[ref]=None


def _build_thcs_tk(ws: Any, people: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year: int, meta: dict[str,str], school_count:int) -> None:
    _clear_range(ws,8,8,1,18)
    age1118=[pr for pr in people if (_age(pr[0],year) or -1) in range(11,19)]
    age1518=[pr for pr in age1118 if (_age(pr[0],year) or -1) in range(15,19) and _is_ppc(*pr)]
    disabled=[pr for pr in age1118 if _is_disabled(*pr)]
    tn=sum(1 for p,r in age1518 if _completed_secondary(p,r))
    upper=sum(1 for p,r in age1518 if _upper_secondary(p,r))
    access=sum(1 for p,r in disabled if _disabled_access(p,r))
    vals={1:1,2:meta["title"],3:school_count,4:len(age1118),5:sum(1 for p,r in age1118 if _is_disabled(p,r)),8:tn,9:_percent(tn,len(age1518)),10:upper,11:_percent(upper,len(age1518)),12:len(disabled),14:access,15:_percent(access,len(disabled))}
    for c,v in vals.items(): ws.cell(8,c).value=v


def _build_xmc3(ws: Any, people: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year:int) -> None:
    # Mẫu XMC-3 có 3 dòng cộng chèn giữa các nhóm tuổi:
    # dòng 20 = Cộng 15-25, dòng 31 = Cộng 15-35, dòng 57 = Cộng 15-60.
    # Vì vậy không thể ánh xạ tuổi theo công thức row = 9 + (age - 15).
    row_by_age: dict[int, int] = {}
    for age in range(15, 26):
        row_by_age[age] = 9 + (age - 15)
    for age in range(26, 36):
        row_by_age[age] = 21 + (age - 26)
    for age in range(36, 61):
        row_by_age[age] = 32 + (age - 36)

    # Xóa số liệu cũ nhưng giữ nguyên nhãn, merge, định dạng và các dòng cộng.
    for row in range(9, 58):
        for col in range(2, 20):
            cell = ws.cell(row, col)
            if cell.__class__.__name__ == "MergedCell":
                continue
            if col == 2 and row in {20, 31, 57}:
                continue
            if col >= 3:
                cell.value = None

    metrics = _population_metrics(people, year)
    for age, row in row_by_age.items():
        item = metrics.get(age, {})
        ws.cell(row, 2).value = year - age
        ws.cell(row, 3).value = item.get("total", 0) or None
        ws.cell(row, 4).value = item.get("female", 0) or None
        ws.cell(row, 5).value = item.get("ethnic", 0) or None
        ws.cell(row, 6).value = item.get("female_ethnic", 0) or None

    for row, ages in ((20, range(15, 26)), (31, range(15, 36)), (57, range(15, 61))):
        for col, key in ((3, "total"), (4, "female"), (5, "ethnic"), (6, "female_ethnic")):
            ws.cell(row, col).value = _sum_metric(metrics, list(ages), key)


def _build_cmc2(ws: Any, people: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year:int) -> None:
    _clear_range(ws,8,11,2,19)
    metrics=_population_metrics(people,year)
    groups=[(8,range(15,26)),(9,range(26,36)),(10,range(36,61)),(11,range(15,61))]
    for row,ages in groups:
        ws.cell(row,2).value=_sum_metric(metrics,list(ages),"total")
        ws.cell(row,3).value=_sum_metric(metrics,list(ages),"female")
        ws.cell(row,4).value=_sum_metric(metrics,list(ages),"ethnic")


def _build_cmc1(ws: Any, people: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year:int, meta:dict[str,str]) -> None:
    _clear_range(ws,8,8,3,66)
    metrics=_population_metrics(people,year)
    all_people=[p for p,r in people]
    ws["A8"]=1; ws["B8"]=meta["title"]
    ws["C8"]=len(all_people); ws["D8"]=sum(_female(p) for p in all_people); ws["E8"]=sum(_ethnic(p) for p in all_people); ws["F8"]=sum(_female_ethnic(p) for p in all_people)
    for start_col,ages in ((7,range(15,26)),(27,range(15,36)),(47,range(15,61))):
        ws.cell(8,start_col).value=_sum_metric(metrics,list(ages),"total")
        ws.cell(8,start_col+1).value=_sum_metric(metrics,list(ages),"female")
        ws.cell(8,start_col+2).value=_sum_metric(metrics,list(ages),"ethnic")
        ws.cell(8,start_col+3).value=_sum_metric(metrics,list(ages),"female_ethnic")


def _build_xmc4(ws: Any, people: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year:int, meta:dict[str,str]) -> None:
    _clear_range(ws,8,8,1,18)
    metrics=_population_metrics(people,year)
    ws["A8"]=1; ws["B8"]=meta["title"]
    for col,ages in ((3,range(15,26)),(8,range(15,36)),(13,range(15,61))):
        ws.cell(8,col).value=_sum_metric(metrics,list(ages),"total")
    # Các cột biết chữ mức độ 1/2 và mức đạt chuẩn để trống vì CSDL hiện chưa có trường chuyên biệt.


def _write_staff_detail(ws: Any, report_type:str, db:Session, school_year_id:int, selected_commune_id:int|None, selected_school_id:int|None, level:str, meta:dict[str,str]) -> None:
    schools, grades=_schools_and_classes(db,school_year_id,selected_commune_id,selected_school_id)
    schools=[s for s in schools if _school_level(s,grades)==level]
    records=_staff_rows(db,school_year_id,[int(s.id) for s in schools])
    if report_type=="PCGD_TH_01_GV_2025": start,total_row,max_rows=8,18,10
    else: start,total_row,max_rows=9,19,10
    _clear_range(ws,start,total_row,1,43 if level=="THCS" else 33)
    display=schools if len(schools)<=max_rows else []
    if not display:
        # Phạm vi lớn: một dòng tổng hợp để không phá bố cục mẫu gốc.
        class_count=sum(_class_count(grades.get(int(s.id),[]),level) for s in schools)
        merged=[]
        for s in schools: merged.extend(records.get(int(s.id),[]))
        summary=_staff_summary(merged,class_count)
        display_data=[(meta["title"],summary,class_count)]
    else:
        display_data=[]
        for s in display:
            cc=_class_count(grades.get(int(s.id),[]),level)
            display_data.append((s.name,_staff_summary(records.get(int(s.id),[]),cc),cc))
    for idx,(name,s,cc) in enumerate(display_data,start=1):
        row=start+idx-1
        ws.cell(row,1).value=idx; ws.cell(row,2).value=name
        if level=="TH":
            mapping={7:s["principal"],8:s["vice"],9:s["teachers"],10:s["bienche"],11:s["hopdong"],12:s["female"],13:s["ethnic"],14:s["ratio"],15:s["qual"]["postgrad"],16:s["qual"]["univ"],17:s["qual"]["college"],18:s["qual"]["secondary"],20:s["subjects"]["general"],21:s["subjects"]["music"],22:s["subjects"]["art"],23:s["subjects"]["pe"],24:s["subjects"]["it"],25:s["subjects"]["english"]+s["subjects"]["foreign"],27:s["prof"]["excellent"],28:s["prof"]["good"],29:s["prof"]["average"],30:s["prof"]["poor"],32:s["employees"]["office"],33:s["employees"]["library"]+s["employees"]["equipment"]}
        else:
            mapping={6:s["principal"],7:s["vice"],8:s["teachers"],9:s["bienche"],10:s["hopdong"],11:s["female"],12:s["ethnic"],13:s["ratio"],14:s["qual"]["postgrad"],15:s["qual"]["univ"],16:s["qual"]["college"],17:s["qual"]["secondary"],18:s["subjects"]["math"],19:s["subjects"]["literature"],20:s["subjects"]["physics"],21:s["subjects"]["chemistry"],22:s["subjects"]["biology"],23:s["subjects"]["history"],24:s["subjects"]["geography"],25:s["subjects"]["music"],26:s["subjects"]["art"],27:s["subjects"]["pe"],28:s["subjects"]["civic"],29:s["subjects"]["technology"],30:s["subjects"]["it"],31:s["subjects"]["english"],32:s["subjects"]["russian"],33:s["subjects"]["french"],34:s["subjects"]["other"],36:s["prof"]["excellent"],37:s["prof"]["good"],38:s["prof"]["average"],39:s["prof"]["poor"],40:s["employees"]["library"],41:s["employees"]["equipment"],42:s["employees"]["office"],43:s["employees"]["health"]}
        for c,v in mapping.items(): ws.cell(row,c).value=v or None
    # Dòng cộng
    label_col=1; ws.cell(total_row,label_col).value="Cộng tỉnh:" if meta["kind"]=="province" else "Cộng:"
    for c in range(6 if level=="THCS" else 7,44 if level=="THCS" else 34):
        vals=[ws.cell(r,c).value for r in range(start,start+len(display_data))]
        nums=[v for v in vals if isinstance(v,(int,float))]
        if c==13 and level=="THCS" or c==14 and level=="TH":
            continue
        if nums: ws.cell(total_row,c).value=sum(nums)


def _write_csvc_detail(ws:Any, report_type:str, db:Session, school_year_id:int, selected_commune_id:int|None, selected_school_id:int|None, level:str, meta:dict[str,str]) -> None:
    schools,grades=_schools_and_classes(db,school_year_id,selected_commune_id,selected_school_id)
    schools=[s for s in schools if _school_level(s,grades)==level]
    if level=="TH": start,total_row,max_rows=8,18,10
    else: start,total_row,max_rows=8,18,10
    _clear_range(ws,start,total_row,1,35 if level=="TH" else 24)
    if len(schools)>max_rows or meta["kind"]=="province":
        rows=[(meta["title"],sum(_class_count(grades.get(int(s.id),[]),level) for s in schools))]
    else:
        rows=[(s.name,_class_count(grades.get(int(s.id),[]),level)) for s in schools]
    for idx,(name,cc) in enumerate(rows,start=1):
        row=start+idx-1; ws.cell(row,1).value=idx; ws.cell(row,2).value=name; ws.cell(row,4).value=cc or None
    ws.cell(total_row,2).value="Cộng tỉnh:" if meta["kind"]=="province" else "Cộng:"
    ws.cell(total_row,4).value=sum(cc for _,cc in rows) or None
    # Các chỉ tiêu phòng học/phòng chức năng/vệ sinh/sân chơi/bãi tập để trống: CSDL hiện chưa có trường dữ liệu tương ứng.


def export_additional_report(
    *,
    report_type: str,
    db: Session,
    request: Any,
    school_year: SchoolYear | None,
    selected_commune_id: int | None,
    selected_school_id: int | None,
    scope_label: str,
) -> StreamingResponse:
    if report_type not in SUPPORTED_REPORT_TYPES:
        raise ValueError(f"Loại báo cáo chưa được hỗ trợ: {report_type}")
    if school_year is None:
        raise ValueError("Chưa xác định được năm học để lập báo cáo.")
    template_path=TEMPLATE_ROOT/REPORT_TEMPLATE_FILES[report_type]
    if not template_path.exists():
        raise FileNotFoundError(f"Không tìm thấy mẫu Excel: {template_path}")
    year=_reference_year(school_year)
    (
        selected_commune_id,
        selected_school_id,
        scope_label,
        _scope_source,
    ) = _effective_scope_for_user(
        db=db,
        request=request,
        selected_commune_id=selected_commune_id,
        selected_school_id=selected_school_id,
        scope_label=scope_label,
    )
    meta=_scope_meta(scope_label=scope_label,selected_commune_id=selected_commune_id,selected_school_id=selected_school_id,year=year)
    people=_load_people_and_records(db=db,request=request,school_year_id=int(school_year.id),selected_commune_id=selected_commune_id,selected_school_id=selected_school_id)
    workbook=load_workbook(template_path)
    ws=workbook[workbook.sheetnames[0]]
    _write_scope_signature(ws,report_type,meta)
    # Thời điểm báo cáo
    date_cells={"PCGD_TH_02_2025":"C2","PCGD_TH_01_GV_2025":"C2","PCGD_TH_01_CSVC_2025":"C2","PCGD_THCS_M1_2025":"F2","PCGD_THCS_M2_2025":"A7","PCGD_THCS_TK_2025":"C2","PCGD_THCS_M5_2025":"G3","PCGD_THCS_CSVC_2025":"C2","PCGD_XMC_4_2025":"A4"}
    if report_type in date_cells: ws[date_cells[report_type]]=f"Thời điểm: ngày 30 tháng 9 năm {year}" if report_type not in {"PCGD_THCS_M2_2025","PCGD_THCS_M5_2025","PCGD_THCS_CSVC_2025"} else (f"Thời điểm điều tra: ngày 30 tháng 9 năm {year}" if report_type=="PCGD_THCS_M2_2025" else f"Tính đến thời điểm: ngày 30 tháng 9 năm {year}")
    schools,grades=_schools_and_classes(db,int(school_year.id),selected_commune_id,selected_school_id)
    primary_count=sum(1 for s in schools if _school_level(s,grades)=="TH")
    thcs_count=sum(1 for s in schools if _school_level(s,grades)=="THCS")
    if report_type=="PCGD_TH_02_2025": _build_th02(ws,people,year,meta,primary_count)
    elif report_type=="PCGD_TH_01_GV_2025": _write_staff_detail(ws,report_type,db,int(school_year.id),selected_commune_id,selected_school_id,"TH",meta)
    elif report_type=="PCGD_TH_01_CSVC_2025": _write_csvc_detail(ws,report_type,db,int(school_year.id),selected_commune_id,selected_school_id,"TH",meta)
    elif report_type=="PCGD_THCS_M1_2025": _build_thcs_m1(ws,people,year,meta,selected_commune_id)
    elif report_type=="PCGD_THCS_M2_2025": _build_thcs_m2(ws,people,year,meta)
    elif report_type=="PCGD_THCS_TK_2025": _build_thcs_tk(ws,people,year,meta,thcs_count)
    elif report_type=="PCGD_THCS_M5_2025": _write_staff_detail(ws,report_type,db,int(school_year.id),selected_commune_id,selected_school_id,"THCS",meta)
    elif report_type=="PCGD_THCS_CSVC_2025": _write_csvc_detail(ws,report_type,db,int(school_year.id),selected_commune_id,selected_school_id,"THCS",meta)
    elif report_type=="PCGD_XMC_3_2025": _build_xmc3(ws,people,year)
    elif report_type=="PCGD_CMC_2_2025": _build_cmc2(ws,people,year)
    elif report_type=="PCGD_CMC_1_2025": _build_cmc1(ws,people,year,meta)
    elif report_type=="PCGD_XMC_4_2025": _build_xmc4(ws,people,year,meta)
    if meta["kind"]=="province":
        try: ws.title="Toàn tỉnh"
        except Exception: pass
    output=BytesIO(); workbook.save(output); output.seek(0)
    year_code=str(school_year.code or year).replace("/","-")
    code=report_type.replace("PCGD_","").replace("_2025","")
    filename=f"PCGD_{code}_{year_code}_{_safe_filename_text(meta['title'])}.xlsx"
    return StreamingResponse(output,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="{filename}"'})
