# -*- coding: utf-8 -*-
"""
BÀI 13B-12 V2.4.3 - KHẢO SÁT CHỈ ĐỌC NGUỒN DỮ LIỆU 7 BIỂU MẦM NON

Mục tiêu:
- KHÔNG sửa source.
- KHÔNG sửa database.
- KHÔNG đụng Excel điều tra 2.2.3 / 2.2.4.
- Đối chiếu dữ liệu gốc và logic hiện tại của 7 biểu:
  MN-01-TE, MN-02, MN-01-GV, MN-01-CSVC, MN-TC, MN-Trẻ KT, Sổ theo dõi PCGDMN.
"""
from __future__ import annotations

import hashlib
import os
import re
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any

PROJECT = Path(r"C:\PhoCap")
APP = PROJECT / "app"
DB = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"

# Phạm vi kiểm thử hiện tại theo ảnh đã dùng.
DEFAULT_SCHOOL_YEAR_ID = 2
DEFAULT_COMMUNE_ID = 114
DEFAULT_SCHOOL_ID = 750

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
REPORT = EXPORTS / f"bao_cao_khao_sat_v2_4_3_nguon_du_lieu_7_bieu_mn_{STAMP}.txt"

SOURCE_FILES = [
    APP / "routers" / "pcgdmn_template_report.py",
    APP / "services" / "mn_report_v246.py",
    APP / "routers" / "mn_official_reports.py",
    APP / "routers" / "survey_summary_report.py",
    APP / "survey_models.py",
    APP / "staff_models.py",
    APP / "report_input_models.py",
    APP / "models.py",
]

SOURCE_FUNCTIONS = {
    "pcgdmn_template_report.py": [
        "_write_mn01_te",
        "_write_mn02",
        "_write_mn01_gv",
        "_write_disability_sheet",
        "_write_tracking_book",
        "_blank_future_modules",
        "build_template_workbook",
    ],
    "mn_report_v246.py": [
        "_apply_old_te",
        "_apply_old_gv",
        "apply_formula_contract",
    ],
    "survey_summary_report.py": [
        "build_report_data",
    ],
}

CANDIDATE_TABLE_PATTERNS = (
    "survey",
    "person",
    "house",
    "staff",
    "class",
    "school",
    "commune",
    "finance",
    "financial",
    "csvc",
    "facility",
    "facilities",
    "report_input",
    "disability",
    "disabled",
)


def qident(name: str) -> str:
    return '"' + str(name).replace('"', '""') + '"'


def short(value: Any, limit: int = 180) -> str:
    text = str(value)
    if len(text) <= limit:
        return text
    return text[: limit - 3] + "..."


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def table_names(con: sqlite3.Connection) -> list[str]:
    return [
        row[0]
        for row in con.execute(
            "SELECT name FROM sqlite_master "
            "WHERE type='table' AND name NOT LIKE 'sqlite_%' "
            "ORDER BY name"
        ).fetchall()
    ]


def columns(con: sqlite3.Connection, table: str) -> list[dict[str, Any]]:
    rows = con.execute(f"PRAGMA table_info({qident(table)})").fetchall()
    return [
        {
            "cid": row[0],
            "name": row[1],
            "type": row[2],
            "notnull": row[3],
            "default": row[4],
            "pk": row[5],
        }
        for row in rows
    ]


def column_names(con: sqlite3.Connection, table: str) -> list[str]:
    return [item["name"] for item in columns(con, table)]


def foreign_keys(con: sqlite3.Connection, table: str) -> list[dict[str, Any]]:
    rows = con.execute(f"PRAGMA foreign_key_list({qident(table)})").fetchall()
    result = []
    for row in rows:
        result.append(
            {
                "id": row[0],
                "seq": row[1],
                "table": row[2],
                "from": row[3],
                "to": row[4],
                "on_update": row[5],
                "on_delete": row[6],
            }
        )
    return result


def has_table(tables: list[str], name: str) -> bool:
    return name in tables


def first_existing(cols: list[str], candidates: list[str]) -> str | None:
    lower = {c.lower(): c for c in cols}
    for candidate in candidates:
        if candidate.lower() in lower:
            return lower[candidate.lower()]
    return None


def safe_scalar(
    con: sqlite3.Connection,
    sql: str,
    params: tuple[Any, ...] = (),
) -> Any:
    try:
        row = con.execute(sql, params).fetchone()
        return None if row is None else row[0]
    except Exception as exc:
        return f"<LỖI SQL: {type(exc).__name__}: {exc}>"


def safe_rows(
    con: sqlite3.Connection,
    sql: str,
    params: tuple[Any, ...] = (),
    limit: int = 100,
) -> list[sqlite3.Row]:
    try:
        return con.execute(sql, params).fetchmany(limit)
    except Exception:
        return []


def add(lines: list[str], text: str = "") -> None:
    lines.append(text)


def heading(lines: list[str], title: str) -> None:
    add(lines)
    add(lines, "=" * 120)
    add(lines, title)
    add(lines, "=" * 120)


def subheading(lines: list[str], title: str) -> None:
    add(lines)
    add(lines, "-" * 100)
    add(lines, title)
    add(lines, "-" * 100)


def source_excerpt(path: Path, function_name: str, max_lines: int = 90) -> str:
    if not path.exists():
        return "<FILE KHÔNG TỒN TẠI>"
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    lines = text.splitlines()

    pattern = re.compile(
        rf"^(?:async\s+def|def)\s+{re.escape(function_name)}\s*\(",
        re.M,
    )
    match = pattern.search(text)
    if not match:
        return "<KHÔNG TÌM THẤY HÀM>"

    start_line = text[: match.start()].count("\n")
    base_indent = len(lines[start_line]) - len(lines[start_line].lstrip())
    end_line = min(len(lines), start_line + max_lines)

    # Tìm hàm/class top-level tiếp theo.
    for idx in range(start_line + 1, min(len(lines), start_line + max_lines * 3)):
        line = lines[idx]
        stripped = line.lstrip()
        indent = len(line) - len(stripped)
        if (
            stripped.startswith("def ")
            or stripped.startswith("async def ")
            or stripped.startswith("class ")
        ) and indent <= base_indent:
            end_line = idx
            break

    excerpt = []
    for idx in range(start_line, min(end_line, start_line + max_lines)):
        excerpt.append(f"{idx+1:05d}: {lines[idx]}")
    if end_line - start_line > max_lines:
        excerpt.append("... <đã cắt>")
    return "\n".join(excerpt)


def report_scope(
    con: sqlite3.Connection,
    tables: list[str],
    lines: list[str],
) -> tuple[int, int, int]:
    school_year_id = DEFAULT_SCHOOL_YEAR_ID
    commune_id = DEFAULT_COMMUNE_ID
    school_id = DEFAULT_SCHOOL_ID

    subheading(lines, "PHẠM VI KIỂM THỬ")

    if has_table(tables, "school_years"):
        cols = column_names(con, "school_years")
        id_col = first_existing(cols, ["id"])
        code_col = first_existing(cols, ["code", "name"])
        if id_col and code_col:
            row = con.execute(
                f"SELECT {qident(id_col)}, {qident(code_col)} "
                f"FROM school_years WHERE {qident(id_col)}=?",
                (school_year_id,),
            ).fetchone()
            add(lines, f"Năm học ID={school_year_id}: {dict(row) if row else 'KHÔNG TÌM THẤY'}")

    if has_table(tables, "communes"):
        cols = column_names(con, "communes")
        fields = [c for c in ("id", "code", "name") if c in cols]
        if fields:
            row = con.execute(
                f"SELECT {', '.join(qident(c) for c in fields)} "
                f"FROM communes WHERE id=?",
                (commune_id,),
            ).fetchone()
            add(lines, f"Xã/phường ID={commune_id}: {dict(row) if row else 'KHÔNG TÌM THẤY'}")

    if has_table(tables, "schools"):
        cols = column_names(con, "schools")
        fields = [
            c
            for c in (
                "id",
                "code",
                "name",
                "commune_id",
                "education_level",
                "level",
                "school_level",
            )
            if c in cols
        ]
        if fields:
            row = con.execute(
                f"SELECT {', '.join(qident(c) for c in fields)} "
                f"FROM schools WHERE id=?",
                (school_id,),
            ).fetchone()
            add(lines, f"Trường ID={school_id}: {dict(row) if row else 'KHÔNG TÌM THẤY'}")

    return school_year_id, commune_id, school_id


def report_schema(
    con: sqlite3.Connection,
    tables: list[str],
    lines: list[str],
) -> None:
    heading(lines, "I. SƠ ĐỒ CSDL LIÊN QUAN ĐẾN 7 BIỂU")

    candidates = [
        t
        for t in tables
        if any(token in t.lower() for token in CANDIDATE_TABLE_PATTERNS)
    ]

    for table in candidates:
        cols = columns(con, table)
        fks = foreign_keys(con, table)
        count = safe_scalar(con, f"SELECT COUNT(*) FROM {qident(table)}")
        add(lines, f"\n[{table}] rows={count}")
        add(
            lines,
            "  columns: "
            + ", ".join(
                f"{c['name']}:{c['type'] or '?'}"
                for c in cols
            ),
        )
        if fks:
            add(
                lines,
                "  FK: "
                + "; ".join(
                    f"{fk['from']} -> {fk['table']}.{fk['to']}"
                    for fk in fks
                ),
            )


def report_people_year_records(
    con: sqlite3.Connection,
    tables: list[str],
    lines: list[str],
    school_year_id: int,
    commune_id: int,
    school_id: int,
) -> None:
    heading(lines, "II. DỮ LIỆU TRẺ/ĐỐI TƯỢNG VÀ YEAR-RECORD")

    for table in ("survey_people", "survey_person_year_records", "survey_forms", "households"):
        if has_table(tables, table):
            add(lines, f"{table}: {safe_scalar(con, f'SELECT COUNT(*) FROM {qident(table)}')} bản ghi")
            add(lines, "  " + ", ".join(column_names(con, table)))

    if not (
        has_table(tables, "survey_people")
        and has_table(tables, "survey_person_year_records")
    ):
        add(lines, "Không đủ hai bảng survey_people + survey_person_year_records để đối chiếu sâu.")
        return

    pcols = column_names(con, "survey_people")
    rcols = column_names(con, "survey_person_year_records")

    person_id_col = first_existing(pcols, ["id"])
    record_person_col = first_existing(rcols, ["person_id", "survey_person_id"])
    year_col = first_existing(rcols, ["school_year_id"])
    school_col = first_existing(rcols, ["school_id", "current_school_id"])
    dob_col = first_existing(pcols, ["date_of_birth", "birth_date", "dob"])
    gender_col = first_existing(pcols, ["gender", "sex"])
    full_name_col = first_existing(pcols, ["full_name", "name"])
    disability_col = first_existing(
        rcols,
        [
            "disability_status",
            "is_disabled",
            "has_disability",
            "disability_type",
        ],
    )
    learning_col = first_existing(
        rcols,
        ["learning_status", "study_status", "education_status"],
    )
    class_col = first_existing(
        rcols,
        ["class_name", "current_class_name", "grade_name", "classroom_name"],
    )
    completed_mn_col = first_existing(
        rcols,
        [
            "completed_preschool_by_age",
            "completed_preschool_5",
            "completed_preschool_program",
            "completed_preschool_program_by_age",
        ],
    )
    two_sessions_col = first_existing(
        rcols,
        ["attends_two_sessions_per_day", "two_sessions_per_day"],
    )
    prepared_vn_col = first_existing(
        rcols,
        [
            "ethnic_minority_prepared_vietnamese",
            "prepared_vietnamese",
            "prepared_vietnamese_language",
        ],
    )

    add(lines, f"Khóa nối person: {person_id_col} -> year_record.{record_person_col}")
    add(lines, f"Năm học: {year_col}; trường: {school_col}; ngày sinh: {dob_col}")
    add(lines, f"Tình trạng học: {learning_col}; lớp: {class_col}")
    add(lines, f"Hoàn thành CTGDMN: {completed_mn_col}")
    add(lines, f"Học 2 buổi/ngày: {two_sessions_col}")
    add(lines, f"Chuẩn bị tiếng Việt: {prepared_vn_col}")
    add(lines, f"Khuyết tật: {disability_col}")

    if not (person_id_col and record_person_col and year_col):
        add(lines, "Thiếu cột khóa/năm học để chạy truy vấn chi tiết.")
        return

    conditions = [f"r.{qident(year_col)}=?"]
    params: list[Any] = [school_year_id]
    if school_col:
        conditions.append(f"r.{qident(school_col)}=?")
        params.append(school_id)

    select_fields = [
        f"p.{qident(person_id_col)} AS person_id",
    ]
    if full_name_col:
        select_fields.append(f"p.{qident(full_name_col)} AS full_name")
    if dob_col:
        select_fields.append(f"p.{qident(dob_col)} AS dob")
    if gender_col:
        select_fields.append(f"p.{qident(gender_col)} AS gender")
    if learning_col:
        select_fields.append(f"r.{qident(learning_col)} AS learning_status")
    if class_col:
        select_fields.append(f"r.{qident(class_col)} AS class_name")
    if school_col:
        select_fields.append(f"r.{qident(school_col)} AS school_id")
    if completed_mn_col:
        select_fields.append(f"r.{qident(completed_mn_col)} AS completed_mn")
    if two_sessions_col:
        select_fields.append(f"r.{qident(two_sessions_col)} AS two_sessions")
    if prepared_vn_col:
        select_fields.append(f"r.{qident(prepared_vn_col)} AS prepared_vn")
    if disability_col:
        select_fields.append(f"r.{qident(disability_col)} AS disability")

    sql = (
        f"SELECT {', '.join(select_fields)} "
        f"FROM survey_person_year_records r "
        f"JOIN survey_people p ON p.{qident(person_id_col)}=r.{qident(record_person_col)} "
        f"WHERE {' AND '.join(conditions)} "
        f"ORDER BY p.{qident(person_id_col)}"
    )
    rows = safe_rows(con, sql, tuple(params), limit=5000)

    add(lines, f"Year-record đúng trường/năm hiện tại: {len(rows)}")
    for row in rows[:80]:
        add(lines, "  " + short(dict(row), 360))
    if len(rows) > 80:
        add(lines, f"  ... còn {len(rows)-80} dòng")

    # Phân bố năm sinh / tuổi quy ước 2026 - birth_year.
    if dob_col:
        stats: dict[int, int] = {}
        for row in rows:
            dob = str(row["dob"] or "")
            m = re.search(r"(19|20)\d{2}", dob)
            if not m:
                continue
            birth_year = int(m.group(0))
            age = 2026 - birth_year
            stats[age] = stats.get(age, 0) + 1
        add(lines, "Phân bố tuổi quy ước tại 2026:")
        for age in sorted(stats):
            add(lines, f"  tuổi {age}: {stats[age]}")

    # Đặc biệt tuổi 5.
    if dob_col:
        five_rows = []
        for row in rows:
            dob = str(row["dob"] or "")
            m = re.search(r"(19|20)\d{2}", dob)
            if m and int(m.group(0)) == 2021:
                five_rows.append(row)
        add(lines, f"Đối tượng sinh 2021 (tuổi 5 trong năm 2026): {len(five_rows)}")
        for row in five_rows:
            add(lines, "  " + short(dict(row), 360))


def report_school_rows_and_scope(
    con: sqlite3.Connection,
    tables: list[str],
    lines: list[str],
    school_year_id: int,
    commune_id: int,
    school_id: int,
) -> None:
    heading(lines, "III. KIỂM TRA NGUỒN SCOPE TRƯỜNG/XÃ CHO MN-01-TE VÀ MN-02")

    source = APP / "routers" / "pcgdmn_template_report.py"
    if source.exists():
        text = source.read_text(encoding="utf-8-sig", errors="replace")
        for token in (
            'report_data.get("person_rows")',
            'report_data.get("school_rows")',
            "_write_mn01_te(",
            "_write_mn02(",
        ):
            add(lines, f"{token}: {'CÓ' if token in text else 'KHÔNG'}")

    if has_table(tables, "schools"):
        cols = column_names(con, "schools")
        if "commune_id" in cols:
            rows = safe_rows(
                con,
                "SELECT id, code, name, commune_id "
                "FROM schools WHERE commune_id=? ORDER BY id",
                (commune_id,),
                limit=500,
            )
            add(lines, f"Số trường thuộc xã ID={commune_id}: {len(rows)}")
            for row in rows:
                add(lines, "  " + short(dict(row), 300))

    add(lines)
    add(lines, "MỤC TIÊU ĐỐI CHIẾU:")
    add(lines, "- MN-01-TE và MN-02 cùng phạm vi trường phải xuất phát từ cùng tập đối tượng.")
    add(lines, "- Nếu MN-02 có trẻ 5 tuổi nhưng MN-01-TE = 0, cần xác định school_rows và person_rows đang lọc khác nhau ở đâu.")


def report_staff(
    con: sqlite3.Connection,
    tables: list[str],
    lines: list[str],
    school_year_id: int,
    school_id: int,
) -> None:
    heading(lines, "IV. NGUỒN DỮ LIỆU MN-01-GV")

    staff_tables = [t for t in tables if "staff" in t.lower()]
    if not staff_tables:
        add(lines, "Không tìm thấy bảng có tên chứa 'staff'.")
        return

    for table in staff_tables:
        cols = column_names(con, table)
        add(lines, f"\n[{table}] rows={safe_scalar(con, f'SELECT COUNT(*) FROM {qident(table)}')}")
        add(lines, "  " + ", ".join(cols))

        school_col = first_existing(cols, ["school_id"])
        year_col = first_existing(cols, ["school_year_id"])
        where = []
        params: list[Any] = []
        if school_col:
            where.append(f"{qident(school_col)}=?")
            params.append(school_id)
        if year_col:
            where.append(f"{qident(year_col)}=?")
            params.append(school_year_id)
        if where:
            count = safe_scalar(
                con,
                f"SELECT COUNT(*) FROM {qident(table)} WHERE {' AND '.join(where)}",
                tuple(params),
            )
            add(lines, f"  scope school/year: {count}")

            rows = safe_rows(
                con,
                f"SELECT * FROM {qident(table)} WHERE {' AND '.join(where)} LIMIT 40",
                tuple(params),
                limit=40,
            )
            for row in rows[:15]:
                add(lines, "    " + short(dict(row), 380))

    if has_table(tables, "classes"):
        cols = column_names(con, "classes")
        add(lines, f"\n[classes] columns: {', '.join(cols)}")
        school_col = first_existing(cols, ["school_id"])
        year_col = first_existing(cols, ["school_year_id"])
        where = []
        params = []
        if school_col:
            where.append(f"{qident(school_col)}=?")
            params.append(school_id)
        if year_col:
            where.append(f"{qident(year_col)}=?")
            params.append(school_year_id)
        if where:
            rows = safe_rows(
                con,
                f"SELECT * FROM classes WHERE {' AND '.join(where)} ORDER BY id",
                tuple(params),
                limit=200,
            )
            add(lines, f"Lớp của trường/năm: {len(rows)}")
            for row in rows:
                add(lines, "  " + short(dict(row), 360))


def report_csvc_finance(
    con: sqlite3.Connection,
    tables: list[str],
    lines: list[str],
    school_year_id: int,
    school_id: int,
) -> None:
    heading(lines, "V. NGUỒN DỮ LIỆU MN-01-CSVC VÀ MN-TC")

    candidate_tokens = (
        "csvc",
        "facility",
        "facilities",
        "finance",
        "financial",
        "report_input",
    )
    candidates = [
        t
        for t in tables
        if any(token in t.lower() for token in candidate_tokens)
    ]

    if not candidates:
        add(lines, "Không tìm thấy bảng ứng viên CSVC/Tài chính.")
    for table in candidates:
        cols = column_names(con, table)
        add(lines, f"\n[{table}] rows={safe_scalar(con, f'SELECT COUNT(*) FROM {qident(table)}')}")
        add(lines, "  " + ", ".join(cols))
        school_col = first_existing(cols, ["school_id"])
        year_col = first_existing(cols, ["school_year_id"])
        where = []
        params = []
        if school_col:
            where.append(f"{qident(school_col)}=?")
            params.append(school_id)
        if year_col:
            where.append(f"{qident(year_col)}=?")
            params.append(school_year_id)
        if where:
            rows = safe_rows(
                con,
                f"SELECT * FROM {qident(table)} WHERE {' AND '.join(where)} LIMIT 50",
                tuple(params),
                limit=50,
            )
            add(lines, f"  scope school/year: {len(rows)}")
            for row in rows[:20]:
                add(lines, "    " + short(dict(row), 420))

    source = APP / "routers" / "pcgdmn_template_report.py"
    if source.exists():
        text = source.read_text(encoding="utf-8-sig", errors="replace")
        add(lines)
        add(lines, f"_blank_future_modules: {'CÓ' if 'def _blank_future_modules' in text else 'KHÔNG'}")
        add(lines, f"'Thiếu dữ liệu CSVC': {'CÓ' if 'Thiếu dữ liệu CSVC' in text else 'KHÔNG'}")
        add(lines, f"'CHƯA CÓ DỮ LIỆU PHÂN HỆ CƠ SỞ VẬT CHẤT': {'CÓ' if 'CHƯA CÓ DỮ LIỆU PHÂN HỆ CƠ SỞ VẬT CHẤT' in text else 'KHÔNG'}")
        add(lines, f"'CHƯA CÓ DỮ LIỆU PHÂN HỆ TÀI CHÍNH': {'CÓ' if 'CHƯA CÓ DỮ LIỆU PHÂN HỆ TÀI CHÍNH' in text else 'KHÔNG'}")


def report_disability(
    con: sqlite3.Connection,
    tables: list[str],
    lines: list[str],
    school_year_id: int,
    school_id: int,
) -> None:
    heading(lines, "VI. NGUỒN DỮ LIỆU MN-TRẺ KT")

    if not (
        has_table(tables, "survey_people")
        and has_table(tables, "survey_person_year_records")
    ):
        add(lines, "Thiếu bảng survey_people/year_records.")
        return

    rcols = column_names(con, "survey_person_year_records")
    pcols = column_names(con, "survey_people")
    disability_cols = [
        c
        for c in rcols
        if any(
            token in c.lower()
            for token in ("disab", "khuyet", "special", "inclusive")
        )
    ]
    add(lines, "Các cột year-record liên quan KT/hòa nhập: " + ", ".join(disability_cols or ["<không tìm thấy>"]))

    year_col = first_existing(rcols, ["school_year_id"])
    school_col = first_existing(rcols, ["school_id", "current_school_id"])
    person_col = first_existing(rcols, ["person_id", "survey_person_id"])
    pid = first_existing(pcols, ["id"])
    name_col = first_existing(pcols, ["full_name", "name"])

    if year_col and person_col and pid and disability_cols:
        select_cols = [f"p.{qident(pid)} AS person_id"]
        if name_col:
            select_cols.append(f"p.{qident(name_col)} AS full_name")
        for c in disability_cols:
            select_cols.append(f"r.{qident(c)} AS {qident(c)}")

        where = [f"r.{qident(year_col)}=?"]
        params: list[Any] = [school_year_id]
        if school_col:
            where.append(f"r.{qident(school_col)}=?")
            params.append(school_id)

        sql = (
            f"SELECT {', '.join(select_cols)} "
            f"FROM survey_person_year_records r "
            f"JOIN survey_people p ON p.{qident(pid)}=r.{qident(person_col)} "
            f"WHERE {' AND '.join(where)}"
        )
        rows = safe_rows(con, sql, tuple(params), limit=1000)
        add(lines, f"Year-record trong scope để soi KT: {len(rows)}")
        for row in rows[:100]:
            values = dict(row)
            if any(
                str(values.get(c) or "").strip().lower()
                not in ("", "0", "false", "không", "khong", "none", "chưa xác định")
                for c in disability_cols
            ):
                add(lines, "  CÓ TÍN HIỆU KT: " + short(values, 420))


def report_tracking(
    con: sqlite3.Connection,
    tables: list[str],
    lines: list[str],
    school_year_id: int,
    school_id: int,
) -> None:
    heading(lines, "VII. NGUỒN DỮ LIỆU SỔ THEO DÕI PCGDMN")

    source = APP / "routers" / "pcgdmn_template_report.py"
    if source.exists():
        text = source.read_text(encoding="utf-8-sig", errors="replace")
        add(lines, f"_write_tracking_book có trong source: {'CÓ' if 'def _write_tracking_book' in text else 'KHÔNG'}")
        add(lines, f"Lời gọi _write_tracking_book: {text.count('_write_tracking_book(')}")

    if has_table(tables, "survey_person_year_records"):
        cols = column_names(con, "survey_person_year_records")
        year_col = first_existing(cols, ["school_year_id"])
        school_col = first_existing(cols, ["school_id", "current_school_id"])
        where = []
        params = []
        if year_col:
            where.append(f"{qident(year_col)}=?")
            params.append(school_year_id)
        if school_col:
            where.append(f"{qident(school_col)}=?")
            params.append(school_id)
        if where:
            count = safe_scalar(
                con,
                f"SELECT COUNT(*) FROM survey_person_year_records WHERE {' AND '.join(where)}",
                tuple(params),
            )
            add(lines, f"Số year-record trong scope trường/năm có thể cấp dữ liệu cho Sổ: {count}")

    if has_table(tables, "survey_people"):
        add(lines, f"Tổng survey_people toàn CSDL: {safe_scalar(con, 'SELECT COUNT(*) FROM survey_people')}")

    add(lines)
    add(lines, "CẦN CHỐT SAU KHẢO SÁT:")
    add(lines, "- Nếu Sổ ra 0 dòng nhưng year-record trong scope > 0, lỗi nằm ở lọc person_rows / writer.")
    add(lines, "- Nếu school scope chỉ lấy trẻ đang học tại trường thì cần đối chiếu nghiệp vụ Sổ theo dõi: theo địa bàn điều tra hay theo cơ sở đang học.")


def report_source(
    lines: list[str],
) -> None:
    heading(lines, "VIII. SOURCE HIỆN TẠI - CÁC HÀM QUYẾT ĐỊNH 7 BIỂU")

    for path in SOURCE_FILES:
        add(lines, f"{path}: {'CÓ' if path.exists() else 'THIẾU'}")
        if path.exists():
            add(lines, f"  size={path.stat().st_size:,} bytes")
            add(lines, f"  sha256={sha256_file(path)}")

    for file_name, funcs in SOURCE_FUNCTIONS.items():
        path = None
        for candidate in SOURCE_FILES:
            if candidate.name == file_name:
                path = candidate
                break
        if path is None or not path.exists():
            continue
        for func in funcs:
            subheading(lines, f"{file_name} :: {func}")
            add(lines, source_excerpt(path, func, max_lines=110))


def report_template(lines: list[str]) -> None:
    heading(lines, "IX. FILE MẪU 7 BIỂU")

    template = APP / "report_templates" / "Bieu_mau_PCGDMN_2025.xlsx"
    add(lines, f"Đường dẫn: {template}")
    add(lines, f"Tồn tại: {'CÓ' if template.exists() else 'KHÔNG'}")
    if not template.exists():
        return

    add(lines, f"SHA256: {sha256_file(template)}")
    try:
        from openpyxl import load_workbook

        wb = load_workbook(template, read_only=False, data_only=False)
        add(lines, f"Số sheet: {len(wb.sheetnames)}")
        for idx, name in enumerate(wb.sheetnames, start=1):
            ws = wb[name]
            add(lines, f"  {idx}. {name} | max_row={ws.max_row} max_col={ws.max_column}")
            # Ghi các ô có 2025/2026 trong 40 dòng đầu để phát hiện năm cố định.
            hits = []
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
                for cell in row:
                    text = str(cell.value or "")
                    if "2025" in text or "2026" in text:
                        hits.append(f"{cell.coordinate}={short(text, 120)}")
            if hits:
                add(lines, "     ô chứa 2025/2026: " + " | ".join(hits[:20]))
    except Exception as exc:
        add(lines, f"LỖI đọc template: {type(exc).__name__}: {exc}")


def main() -> int:
    print("=" * 122)
    print("BÀI 13B-12 V2.4.3 - KHẢO SÁT CHỈ ĐỌC NGUỒN DỮ LIỆU 7 BIỂU MẦM NON")
    print("=" * 122)

    if not DB.exists():
        print("DỪNG: Không tìm thấy database:", DB)
        return 2

    EXPORTS.mkdir(parents=True, exist_ok=True)

    lines: list[str] = []
    add(lines, "=" * 120)
    add(lines, "KHẢO SÁT BÀI 13B-12 V2.4.3 - NGUỒN DỮ LIỆU 7 BIỂU MẦM NON")
    add(lines, "=" * 120)
    add(lines, f"Thời gian: {datetime.now():%d/%m/%Y %H:%M:%S}")
    add(lines)
    add(lines, "NGUYÊN TẮC:")
    add(lines, "- CHỈ ĐỌC source + database.")
    add(lines, "- KHÔNG UPDATE / INSERT / DELETE.")
    add(lines, "- KHÔNG sửa file source.")
    add(lines, "- KHÔNG đụng Excel điều tra 2.2.3 / 2.2.4.")
    add(lines, "- Mục tiêu là khóa đúng nguồn dữ liệu trước khi viết bộ cài V2.4.3.")

    # Mở database ở chế độ read-only.
    uri = DB.resolve().as_uri() + "?mode=ro"
    con = sqlite3.connect(uri, uri=True)
    con.row_factory = sqlite3.Row

    try:
        heading(lines, "0. KIỂM TRA AN TOÀN DATABASE")
        integrity = con.execute("PRAGMA integrity_check").fetchone()[0]
        fk_rows = con.execute("PRAGMA foreign_key_check").fetchall()
        add(lines, f"integrity_check: {integrity}")
        add(lines, f"foreign_key_check: {len(fk_rows)} lỗi")

        tables = table_names(con)
        add(lines, f"Tổng số bảng: {len(tables)}")

        school_year_id, commune_id, school_id = report_scope(con, tables, lines)

        report_schema(con, tables, lines)
        report_people_year_records(
            con,
            tables,
            lines,
            school_year_id,
            commune_id,
            school_id,
        )
        report_school_rows_and_scope(
            con,
            tables,
            lines,
            school_year_id,
            commune_id,
            school_id,
        )
        report_staff(
            con,
            tables,
            lines,
            school_year_id,
            school_id,
        )
        report_csvc_finance(
            con,
            tables,
            lines,
            school_year_id,
            school_id,
        )
        report_disability(
            con,
            tables,
            lines,
            school_year_id,
            school_id,
        )
        report_tracking(
            con,
            tables,
            lines,
            school_year_id,
            school_id,
        )
        report_source(lines)
        report_template(lines)

        heading(lines, "X. KẾT LUẬN TỰ ĐỘNG / ĐIỂM CẦN KHÓA CHO V2.4.3")
        add(lines, "1. So sánh trực tiếp tập person_rows dùng MN-01-TE với school_rows dùng MN-02.")
        add(lines, "2. Xác định vì sao cùng trường/năm MN-02 có trẻ 5 tuổi nhưng MN-01-TE đang = 0.")
        add(lines, "3. Xác định mapping StaffMember/StaffYearRecord/SchoolStaffYearSummary sang lớp mẫu giáo, GV dạy lớp, chuẩn/trên chuẩn.")
        add(lines, "4. Xác định dữ liệu CSVC/Tài chính hiện có thật hay vẫn đang bị _blank_future_modules xóa/trắng.")
        add(lines, "5. Đối chiếu số đối tượng khuyết tật trong year-record với MN-Trẻ KT.")
        add(lines, "6. Đối chiếu số year-record/person trong scope với số dòng Sổ theo dõi PCGDMN.")
        add(lines, "7. Sửa năm hiển thị từ giá trị cố định của mẫu sang đúng năm học/năm báo cáo hiện hành.")
        add(lines)
        add(lines, "KHÔNG ĐƯỢC làm trong V2.4.3:")
        add(lines, "- Không đổi mẫu Excel điều tra.")
        add(lines, "- Không thay nghiệp vụ PCGD/XMC.")
        add(lines, "- Không tạo dữ liệu giả để làm đẹp báo cáo.")
        add(lines, "- Không tự suy diễn chỉ tiêu chưa có nguồn dữ liệu.")
    finally:
        con.close()

    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")

    print("ĐÃ TẠO BÁO CÁO CHỈ ĐỌC:")
    print(REPORT)
    print()
    print("Không có source/database nào bị sửa.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
