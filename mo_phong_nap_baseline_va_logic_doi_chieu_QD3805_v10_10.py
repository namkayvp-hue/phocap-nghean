# -*- coding: utf-8 -*-
r"""
DRY-RUN V10.10 - MÔ PHỎNG NẠP BASELINE HS 2025-2026 + LOGIC ĐỐI CHIẾU QĐ3805
================================================================================

MỤC TIÊU
--------
1. Đọc 3 nguồn:
   C:\PhoCap\uploads\MN_Nghi_Loc_2025_2026.xlsx
   C:\PhoCap\uploads\TH_Nghi_Loc_2025_2026.xlsx
   C:\PhoCap\uploads\THCS_Nghi_Loc_2025_2026.xlsx

2. Chỉ lấy 13 trường nguồn/đích thuộc QĐ3805.

3. Sao chép DB hiện tại vào SQLite :memory: bằng Backup API.

4. TRÊN BẢN SAO TRONG RAM:
   - tạo lớp 2025-2026 theo đúng trường cũ;
   - tạo Student theo mã định danh Bộ GD&ĐT;
   - tạo StudentEnrollment 2025-2026;
   - giữ nhiều enrollment cùng năm nếu học sinh chuyển trường;
   - map trạng thái gốc sang trạng thái hệ thống;
   - kiểm tra FK + integrity + UNIQUE.

5. KHÔNG ghi DB thật.

6. Kiểm tra route:
   app\routers\student_survey_comparison.py
   để xác định các chỗ cần sửa khi dùng nguồn năm trước:
   - load scope enrollment;
   - load target enrollment;
   - quyền giáo viên khi liên kết;
   - trường nguồn -> trường đích QĐ3805;
   - không được so class_id giữa hai năm khác nhau;
   - không được so trạng thái cùng-năm một cách máy móc giữa 2025-26 và 2026-27.

KẾT LUẬN
--------
SAFE_SIMULATED_IMPORT = YES/NO
COMPARISON_PATCH_REQUIRED = YES/NO
READY_FOR_V11_OFFICIAL_INSTALL = YES/NO

CHỈ ĐỌC - DB thật và mã nguồn thật KHÔNG bị thay đổi.
"""

from __future__ import annotations

import csv
import re
import sqlite3
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any


ROOT = Path(r"C:\PhoCap")
DB_PATH = ROOT / "data" / "phocap.db"
UPLOAD_DIR = ROOT / "uploads"
ROUTE_PATH = ROOT / "app" / "routers" / "student_survey_comparison.py"

SOURCE_FILES = {
    "MN": UPLOAD_DIR / "MN_Nghi_Loc_2025_2026.xlsx",
    "TH": UPLOAD_DIR / "TH_Nghi_Loc_2025_2026.xlsx",
    "THCS": UPLOAD_DIR / "THCS_Nghi_Loc_2025_2026.xlsx",
}

MERGE_MAP = {
    750: 748,
    751: 748,
    749: 747,
    1181: 1178,
    1180: 1179,
    1608: 1605,
    1607: 1606,
}
SOURCE_IDS = sorted(MERGE_MAP)
TARGET_IDS = sorted(set(MERGE_MAP.values()))
QD_IDS = sorted(set(SOURCE_IDS) | set(TARGET_IDS))

EXCEL_SCHOOL_TO_ID = {
    "Trường Mầm non Nghi Diên": 748,
    "Trường Mầm non Nghi Hoa": 750,
    "Trường Mầm non Nghi Vạn": 751,
    "Trường Mầm non TT Quán Hành": 747,
    "Trường MN Nghi Trung": 749,

    "Trường TH Nghi Diên": 1178,
    "Trường Tiểu học Nghi Vạn": 1181,
    "Trường TH Nghi Trung": 1179,
    "Trường Tiểu học Nghi Hoa": 1180,

    "THCS Nghi Diên": 1605,
    "THCS Nghi Vạn": 1608,
    "THCS Nghi Trung": 1606,
    "THCS Nghi Hoa": 1607,
}

EXPECTED_TOTAL_ROWS = 7781
EXPECTED_UNIQUE_CODES = 7777
EXPECTED_DUPLICATE_CODES = 4
EXPECTED_CLASS_COUNT = 217


class AuditAbort(RuntimeError):
    pass


def qident(name: str) -> str:
    return '"' + str(name).replace('"', '""') + '"'


def markers(n: int) -> str:
    return ",".join("?" for _ in range(n))


def normalize_text(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("đ", "d")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def clean(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def code_text(value: Any) -> str:
    if value is None:
        return ""
    # openpyxl có thể trả số.
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", "", str(value).strip())


def to_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def to_year(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        y = int(float(str(value).strip()))
    except Exception:
        return None
    return y if 1900 <= y <= 2100 else None


def write_csv(path: Path, headers: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def connect_ro(path: Path) -> sqlite3.Connection:
    uri = path.resolve().as_uri() + "?mode=ro"
    con = sqlite3.connect(uri, uri=True)
    con.row_factory = sqlite3.Row
    return con


def table_exists(con: sqlite3.Connection, table: str) -> bool:
    return con.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
        (table,),
    ).fetchone() is not None


def colset(con: sqlite3.Connection, table: str) -> set[str]:
    if not table_exists(con, table):
        return set()
    return {
        str(r[1])
        for r in con.execute(f"PRAGMA table_info({qident(table)})").fetchall()
    }


def scalar(con: sqlite3.Connection, sql: str, params=()) -> int:
    row = con.execute(sql, params).fetchone()
    return int(row[0] or 0) if row else 0


def school_year_id(con: sqlite3.Connection, code: str) -> int:
    row = con.execute(
        "SELECT id FROM school_years WHERE code=?",
        (code,),
    ).fetchone()
    if row is None:
        # fallback name contains both years
        a, b = code.split("-")
        rows = con.execute(
            "SELECT id,code,name FROM school_years"
        ).fetchall()
        matches = [
            int(r["id"])
            for r in rows
            if a in str(r["code"] or "") + str(r["name"] or "")
            and b in str(r["code"] or "") + str(r["name"] or "")
        ]
        if len(matches) != 1:
            raise AuditAbort(f"Không xác định được năm học {code}: {matches}")
        return matches[0]
    return int(row["id"])


def map_status(raw: str) -> tuple[str, int]:
    s = normalize_text(raw)

    if s == "dang hoc":
        return "DANG_HOC", 1

    if "chuyen den" in s:
        if "ky 2" in s:
            return "CHUYEN_DEN_KY_2", 1
        # kỳ 1 hoặc trong hè -> đầu năm/kỳ 1
        return "CHUYEN_DEN_KY_1", 1

    if "chuyen di" in s:
        return "CHUYEN_DI", 0

    if "thoi hoc" in s:
        return "THOI_HOC", 0

    if "nghi hoc" in s or "tam nghi" in s:
        return "TAM_NGHI", 0

    raise AuditAbort(f"Chưa có mapping trạng thái nguồn: {raw!r}")


def read_source_rows() -> list[dict]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise AuditAbort(f"Không import được openpyxl: {exc}")

    missing = [str(p) for p in SOURCE_FILES.values() if not p.exists()]
    if missing:
        raise AuditAbort(
            "Thiếu file nguồn:\n- " + "\n- ".join(missing)
        )

    result = []

    for level, path in SOURCE_FILES.items():
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if "Sheet1" not in wb.sheetnames:
                raise AuditAbort(f"{path.name} không có Sheet1.")
            ws = wb["Sheet1"]

            iterator = ws.iter_rows(values_only=True)

            # Dòng 1-5 tiêu đề lớn.
            for _ in range(5):
                try:
                    next(iterator)
                except StopIteration:
                    raise AuditAbort(f"{path.name} không đủ dòng.")

            headers = next(iterator)
            h = {
                clean(value): idx
                for idx, value in enumerate(headers)
                if clean(value)
            }

            required = [
                "Trường",
                "Lớp",
                "Mã định danh Bộ GD&ĐT",
                "Họ tên",
                "Ngày sinh",
                "Giới tính",
                "Dân tộc",
                "Trạng thái",
            ]
            missing_headers = [x for x in required if x not in h]
            if missing_headers:
                raise AuditAbort(
                    f"{path.name} thiếu cột: {missing_headers}"
                )

            def val(row, header):
                idx = h.get(header)
                if idx is None or idx >= len(row):
                    return None
                return row[idx]

            for excel_row, row in enumerate(iterator, start=7):
                school_name = clean(val(row, "Trường"))
                school_id = EXCEL_SCHOOL_TO_ID.get(school_name)
                if school_id not in QD_IDS:
                    continue

                student_code = code_text(
                    val(row, "Mã định danh Bộ GD&ĐT")
                )
                full_name = clean(val(row, "Họ tên"))
                class_name = clean(val(row, "Lớp"))
                status_raw = clean(val(row, "Trạng thái"))

                if not student_code or not full_name or not class_name:
                    raise AuditAbort(
                        f"{path.name} dòng {excel_row}: thiếu mã/tên/lớp."
                    )

                status, is_current = map_status(status_raw)

                personal_id = clean(val(row, "Số định danh cá nhân"))
                if not personal_id:
                    personal_id = clean(val(row, "Số CCCD"))

                record = {
                    "level": level,
                    "file": path.name,
                    "excel_row": excel_row,
                    "school_id": school_id,
                    "school_name_excel": school_name,
                    "class_name": class_name,
                    "student_code": student_code,
                    "full_name": full_name,
                    "date_of_birth": to_date(val(row, "Ngày sinh")),
                    "gender": clean(val(row, "Giới tính")) or None,
                    "ethnic_group": clean(val(row, "Dân tộc")) or None,
                    "birth_place": clean(val(row, "Nơi sinh")) or None,
                    "current_address": clean(val(row, "Chỗ ở hiện nay")) or None,
                    "contact_phone": code_text(val(row, "SĐT liên hệ")) or None,
                    "personal_id": personal_id or None,
                    "disability_type": clean(val(row, "Loại khuyết tật")) or None,
                    "father_name": clean(val(row, "Tên cha")) or None,
                    "father_job": clean(val(row, "Nghề nghiệp cha")) or None,
                    "father_birth_year": to_year(val(row, "Năm sinh cha")),
                    "mother_name": clean(val(row, "Tên mẹ")) or None,
                    "mother_job": clean(val(row, "Nghề nghiệp mẹ")) or None,
                    "mother_birth_year": to_year(val(row, "Năm sinh mẹ")),
                    "status_raw": status_raw,
                    "status": status,
                    "is_current": is_current,
                }
                result.append(record)
        finally:
            wb.close()

    return result


def critical_identity_conflicts(records: list[dict]) -> list[dict]:
    by_code = defaultdict(list)
    for r in records:
        by_code[r["student_code"]].append(r)

    conflicts = []
    for code, items in by_code.items():
        if len(items) <= 1:
            continue

        names = {normalize_text(x["full_name"]) for x in items}
        dobs = {x["date_of_birth"] for x in items if x["date_of_birth"]}
        genders = {normalize_text(x["gender"]) for x in items if x["gender"]}

        if len(names) > 1 or len(dobs) > 1 or len(genders) > 1:
            conflicts.append({
                "student_code": code,
                "rows": len(items),
                "names": " | ".join(sorted({x["full_name"] for x in items})),
                "dobs": " | ".join(sorted(dobs)),
                "genders": " | ".join(sorted(
                    {str(x["gender"] or "") for x in items}
                )),
            })
    return conflicts


def profile_for_code(items: list[dict]) -> dict:
    # Ưu tiên enrollment current rồi dòng nhiều dữ liệu nhất.
    def score(r):
        fields = [
            "date_of_birth", "gender", "ethnic_group", "birth_place",
            "current_address", "contact_phone", "personal_id",
            "disability_type", "father_name", "father_job",
            "father_birth_year", "mother_name", "mother_job",
            "mother_birth_year",
        ]
        completeness = sum(bool(r.get(x)) for x in fields)
        return (int(r["is_current"]), completeness, -int(r["excel_row"]))

    ordered = sorted(items, key=score, reverse=True)
    base = dict(ordered[0])

    # Điền chỗ trống từ các dòng còn lại, không tự ghi đè.
    merge_fields = [
        "date_of_birth", "gender", "ethnic_group", "birth_place",
        "current_address", "contact_phone", "personal_id",
        "disability_type", "father_name", "father_job",
        "father_birth_year", "mother_name", "mother_job",
        "mother_birth_year",
    ]
    for other in ordered[1:]:
        for field in merge_fields:
            if not base.get(field) and other.get(field):
                base[field] = other[field]

    traces = []
    for item in items:
        traces.append(
            f"{item['school_name_excel']} / {item['class_name']} / "
            f"{item['status_raw']}"
        )
    base["notes"] = (
        "Nguồn HS 2025-2026 QĐ3805 Nghi Lộc. "
        "Lịch sử nguồn: " + " || ".join(traces)
    )[:4000]
    return base


def inspect_route() -> tuple[list[dict], bool]:
    if not ROUTE_PATH.exists():
        return [{
            "check": "route_exists",
            "result": "FAIL",
            "detail": str(ROUTE_PATH),
        }], False

    text = ROUTE_PATH.read_text(encoding="utf-8")
    checks = []

    def add(name, found, detail):
        checks.append({
            "check": name,
            "result": "FOUND" if found else "NOT_FOUND",
            "detail": detail,
        })

    add(
        "scope_uses_batch_current_year",
        "StudentEnrollment.school_year_id == batch.school_year_id" in text,
        "Phải đổi sang năm nguồn trước.",
    )
    add(
        "build_rows_passes_batch_year_to_target_enrollments",
        bool(re.search(
            r"_load_target_enrollments_for_students\(\s*db,\s*batch\.school_year_id,",
            text,
            flags=re.S,
        )),
        "Phải truyền năm baseline.",
    )
    add(
        "school_scope_exact_school_id",
        "StudentEnrollment.school_id == int(school_id)" in text,
        "Trường đích QĐ3805 phải nhìn target + source schools.",
    )
    add(
        "compare_class_id_directly",
        "if int(record.class_id) != int(enrollment.class_id)" in text,
        "Không hợp lệ khi record=2026-27, enrollment=2025-26.",
    )
    add(
        "compare_school_id_directly",
        "if int(record.school_id) != int(enrollment.school_id)" in text,
        "Phải coi source->target QĐ3805 là cùng phạm vi trường.",
    )
    add(
        "same_year_status_comparison",
        "survey_active != student_active" in text,
        "Không nên áp dụng máy móc cho hai năm học khác nhau.",
    )
    add(
        "teacher_write_scope_uses_batch_year",
        bool(re.search(
            r"def _student_is_writeable_in_scope[\s\S]{0,1600}"
            r"StudentEnrollment\.school_year_id == batch\.school_year_id",
            text,
        )),
        "Phải dùng baseline year và group scope.",
    )

    patch_required = any(
        x["result"] == "FOUND"
        for x in checks
        if x["check"] != "route_exists"
    )
    return checks, patch_required


def main():
    print("=" * 128)
    print("DRY-RUN V10.10 - MÔ PHỎNG NẠP BASELINE + LOGIC ĐỐI CHIẾU QĐ3805")
    print("=" * 128)
    print("Chế độ: CHỈ ĐỌC - DB thật KHÔNG bị thay đổi")

    if not DB_PATH.exists():
        raise AuditAbort(f"Không tìm thấy DB: {DB_PATH}")

    records = read_source_rows()

    by_code = defaultdict(list)
    for r in records:
        by_code[r["student_code"]].append(r)

    duplicate_codes = {
        code: items
        for code, items in by_code.items()
        if len(items) > 1
    }
    identity_conflicts = critical_identity_conflicts(records)

    class_keys = sorted({
        (r["school_id"], r["class_name"])
        for r in records
    })

    if len(records) != EXPECTED_TOTAL_ROWS:
        raise AuditAbort(
            f"Số dòng nguồn thay đổi: {len(records)} != {EXPECTED_TOTAL_ROWS}"
        )
    if len(by_code) != EXPECTED_UNIQUE_CODES:
        raise AuditAbort(
            f"Số mã duy nhất thay đổi: {len(by_code)} != {EXPECTED_UNIQUE_CODES}"
        )
    if len(duplicate_codes) != EXPECTED_DUPLICATE_CODES:
        raise AuditAbort(
            f"Số mã lặp thay đổi: {len(duplicate_codes)} != {EXPECTED_DUPLICATE_CODES}"
        )
    if len(class_keys) != EXPECTED_CLASS_COUNT:
        raise AuditAbort(
            f"Số lớp nguồn thay đổi: {len(class_keys)} != {EXPECTED_CLASS_COUNT}"
        )
    if identity_conflicts:
        raise AuditAbort(
            f"Có {len(identity_conflicts)} mã học sinh xung đột danh tính."
        )

    real = connect_ro(DB_PATH)
    mem = sqlite3.connect(":memory:")
    mem.row_factory = sqlite3.Row
    mem.execute("PRAGMA foreign_keys=ON")

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = ROOT / f"dry_run_simulate_baseline_QD3805_v10_10_{ts}"
    out_dir.mkdir(parents=True, exist_ok=False)

    try:
        pre_integrity = real.execute("PRAGMA integrity_check").fetchone()[0]
        pre_fk = real.execute("PRAGMA foreign_key_check").fetchall()

        if pre_integrity != "ok" or pre_fk:
            raise AuditAbort(
                f"DB thật không đạt preflight: integrity={pre_integrity}, fk={len(pre_fk)}"
            )

        baseline_year_id = school_year_id(real, "2025-2026")
        current_year_id = school_year_id(real, "2026-2027")

        pre_baseline_enrollments = scalar(
            real,
            f"SELECT COUNT(*) FROM student_enrollments "
            f"WHERE school_year_id=? AND school_id IN ({markers(len(QD_IDS))})",
            [baseline_year_id] + QD_IDS,
        )
        pre_baseline_classes = scalar(
            real,
            f"SELECT COUNT(*) FROM classes "
            f"WHERE school_year_id=? AND school_id IN ({markers(len(QD_IDS))})",
            [baseline_year_id] + QD_IDS,
        )
        existing_students = scalar(real, "SELECT COUNT(*) FROM students")

        if pre_baseline_enrollments != 0:
            raise AuditAbort(
                f"DB thật đã có {pre_baseline_enrollments} enrollment baseline QĐ3805."
            )
        if pre_baseline_classes != 0:
            raise AuditAbort(
                f"DB thật đã có {pre_baseline_classes} lớp baseline QĐ3805."
            )
        if existing_students != 0:
            raise AuditAbort(
                f"DB thật đã có {existing_students} students; cần chiến lược merge thay vì fresh import."
            )

        # Copy DB thật -> RAM. Không mở ghi DB thật.
        real.backup(mem)

        # Sau backup phải bật FK lại.
        mem.execute("PRAGMA foreign_keys=ON")

        now = datetime.now().isoformat(sep=" ", timespec="seconds")

        # School names để report.
        school_rows = mem.execute(
            f"SELECT id,name,is_active FROM schools "
            f"WHERE id IN ({markers(len(QD_IDS))})",
            QD_IDS,
        ).fetchall()
        school_names = {
            int(r["id"]): str(r["name"])
            for r in school_rows
        }

        # Transaction trên RAM.
        mem.execute("BEGIN")

        class_id_by_key = {}
        for school_id, class_name in class_keys:
            cur = mem.execute(
                """
                INSERT INTO classes
                    (school_id,school_year_id,code,name,is_active,created_at)
                VALUES (?,?,?,?,?,?)
                """,
                (
                    school_id,
                    baseline_year_id,
                    None,
                    class_name,
                    1,
                    now,
                ),
            )
            class_id_by_key[(school_id, class_name)] = int(cur.lastrowid)

        student_id_by_code = {}

        for student_code in sorted(by_code):
            items = by_code[student_code]
            p = profile_for_code(items)

            cur = mem.execute(
                """
                INSERT INTO students (
                    code,full_name,date_of_birth,gender,ethnic_group,
                    birth_place,permanent_address,current_address,
                    father_name,father_birth_year,father_job,
                    mother_name,mother_birth_year,mother_job,
                    contact_phone,personal_id,disability_type,notes,
                    is_active,created_at,updated_at
                )
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    student_code,
                    p["full_name"],
                    p.get("date_of_birth"),
                    p.get("gender"),
                    p.get("ethnic_group"),
                    p.get("birth_place"),
                    None,
                    p.get("current_address"),
                    p.get("father_name"),
                    p.get("father_birth_year"),
                    p.get("father_job"),
                    p.get("mother_name"),
                    p.get("mother_birth_year"),
                    p.get("mother_job"),
                    p.get("contact_phone"),
                    p.get("personal_id"),
                    p.get("disability_type"),
                    p.get("notes"),
                    1,
                    now,
                    now,
                ),
            )
            student_id_by_code[student_code] = int(cur.lastrowid)

        for r in records:
            student_id = student_id_by_code[r["student_code"]]
            class_id = class_id_by_key[(r["school_id"], r["class_name"])]

            mem.execute(
                """
                INSERT INTO student_enrollments (
                    student_id,school_year_id,school_id,class_id,
                    status,enrollment_date,transfer_date,is_current,
                    created_at,updated_at
                )
                VALUES (?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    student_id,
                    baseline_year_id,
                    r["school_id"],
                    class_id,
                    r["status"],
                    None,
                    None,
                    int(r["is_current"]),
                    now,
                    now,
                ),
            )

        mem.commit()

        post_integrity = mem.execute("PRAGMA integrity_check").fetchone()[0]
        post_fk = mem.execute("PRAGMA foreign_key_check").fetchall()

        sim_students = scalar(mem, "SELECT COUNT(*) FROM students")
        sim_classes = scalar(
            mem,
            f"SELECT COUNT(*) FROM classes "
            f"WHERE school_year_id=? AND school_id IN ({markers(len(QD_IDS))})",
            [baseline_year_id] + QD_IDS,
        )
        sim_enrollments = scalar(
            mem,
            f"SELECT COUNT(*) FROM student_enrollments "
            f"WHERE school_year_id=? AND school_id IN ({markers(len(QD_IDS))})",
            [baseline_year_id] + QD_IDS,
        )
        sim_current_enrollments = scalar(
            mem,
            f"SELECT COUNT(*) FROM student_enrollments "
            f"WHERE school_year_id=? "
            f"AND school_id IN ({markers(len(QD_IDS))}) "
            f"AND is_current=1",
            [baseline_year_id] + QD_IDS,
        )

        # Kiểm tra mỗi enrollment trỏ đúng class cùng school/year.
        bad_enrollment_class = scalar(
            mem,
            """
            SELECT COUNT(*)
            FROM student_enrollments e
            JOIN classes c ON c.id=e.class_id
            WHERE e.school_year_id=?
              AND e.school_id IN ({})
              AND (
                  c.school_id<>e.school_id
                  OR c.school_year_id<>e.school_year_id
              )
            """.format(markers(len(QD_IDS))),
            [baseline_year_id] + QD_IDS,
        )

        # Một student có bao nhiêu current enrollment trong baseline.
        multi_current_rows = mem.execute(
            """
            SELECT s.code,COUNT(*) AS n
            FROM student_enrollments e
            JOIN students s ON s.id=e.student_id
            WHERE e.school_year_id=?
              AND e.is_current=1
              AND e.school_id IN ({})
            GROUP BY e.student_id
            HAVING COUNT(*)>1
            ORDER BY s.code
            """.format(markers(len(QD_IDS))),
            [baseline_year_id] + QD_IDS,
        ).fetchall()

        duplicate_report = []
        for code, items in sorted(duplicate_codes.items()):
            current_n = sum(int(x["is_current"]) for x in items)
            duplicate_report.append({
                "student_code": code,
                "full_name": items[0]["full_name"],
                "enrollment_rows": len(items),
                "current_rows": current_n,
                "history": " || ".join(
                    f"{school_names.get(x['school_id'], x['school_id'])}"
                    f" / {x['class_name']} / {x['status_raw']}"
                    f" -> {x['status']} current={x['is_current']}"
                    for x in items
                ),
                "check": "PASS" if current_n <= 1 else "FAIL",
            })

        status_report = []
        source_status_counts = Counter(
            (r["status_raw"], r["status"], r["is_current"])
            for r in records
        )
        for key, n in sorted(source_status_counts.items()):
            raw, mapped, current = key
            status_report.append({
                "source_status": raw,
                "mapped_status": mapped,
                "is_current": current,
                "rows": n,
            })

        # Scope counts: chỉ current baseline là danh sách kỳ vọng mang sang năm sau.
        scope_report = []
        for target_id in TARGET_IDS:
            group_ids = sorted(
                [target_id]
                + [sid for sid, tid in MERGE_MAP.items() if tid == target_id]
            )
            current_count = scalar(
                mem,
                f"""
                SELECT COUNT(*)
                FROM student_enrollments
                WHERE school_year_id=?
                  AND school_id IN ({markers(len(group_ids))})
                  AND is_current=1
                """,
                [baseline_year_id] + group_ids,
            )
            all_count = scalar(
                mem,
                f"""
                SELECT COUNT(*)
                FROM student_enrollments
                WHERE school_year_id=?
                  AND school_id IN ({markers(len(group_ids))})
                """,
                [baseline_year_id] + group_ids,
            )
            unique_current = scalar(
                mem,
                f"""
                SELECT COUNT(DISTINCT student_id)
                FROM student_enrollments
                WHERE school_year_id=?
                  AND school_id IN ({markers(len(group_ids))})
                  AND is_current=1
                """,
                [baseline_year_id] + group_ids,
            )
            scope_report.append({
                "target_id": target_id,
                "target_name": school_names[target_id],
                "baseline_school_ids": ",".join(map(str, group_ids)),
                "all_history_enrollments": all_count,
                "current_expected_enrollments": current_count,
                "current_expected_unique_students": unique_current,
            })

        route_checks, patch_required = inspect_route()

        safe_sim = (
            post_integrity == "ok"
            and not post_fk
            and sim_students == EXPECTED_UNIQUE_CODES
            and sim_classes == EXPECTED_CLASS_COUNT
            and sim_enrollments == EXPECTED_TOTAL_ROWS
            and bad_enrollment_class == 0
            and not multi_current_rows
        )

        # Ta kỳ vọng current baseline nhỏ hơn total vì chuyển đi/thôi học.
        expected_scope_ok = (
            0 < sim_current_enrollments < sim_enrollments
        )

        ready = safe_sim and patch_required and expected_scope_ok

        gate_rows = [
            {
                "check": "real_db_pre_integrity",
                "result": "PASS" if pre_integrity == "ok" else "FAIL",
                "detail": str(pre_integrity),
            },
            {
                "check": "real_db_pre_fk",
                "result": "PASS" if not pre_fk else "FAIL",
                "detail": str(len(pre_fk)),
            },
            {
                "check": "baseline_year",
                "result": "PASS",
                "detail": f"2025-2026 id={baseline_year_id}",
            },
            {
                "check": "current_year",
                "result": "PASS",
                "detail": f"2026-2027 id={current_year_id}",
            },
            {
                "check": "source_rows",
                "result": "PASS" if len(records) == EXPECTED_TOTAL_ROWS else "FAIL",
                "detail": str(len(records)),
            },
            {
                "check": "source_unique_students",
                "result": "PASS" if len(by_code) == EXPECTED_UNIQUE_CODES else "FAIL",
                "detail": str(len(by_code)),
            },
            {
                "check": "source_classes",
                "result": "PASS" if len(class_keys) == EXPECTED_CLASS_COUNT else "FAIL",
                "detail": str(len(class_keys)),
            },
            {
                "check": "simulated_students",
                "result": "PASS" if sim_students == EXPECTED_UNIQUE_CODES else "FAIL",
                "detail": str(sim_students),
            },
            {
                "check": "simulated_classes_qd3805",
                "result": "PASS" if sim_classes == EXPECTED_CLASS_COUNT else "FAIL",
                "detail": str(sim_classes),
            },
            {
                "check": "simulated_enrollments_qd3805",
                "result": "PASS" if sim_enrollments == EXPECTED_TOTAL_ROWS else "FAIL",
                "detail": str(sim_enrollments),
            },
            {
                "check": "simulated_current_baseline_scope",
                "result": "PASS" if expected_scope_ok else "FAIL",
                "detail": str(sim_current_enrollments),
            },
            {
                "check": "bad_enrollment_class_refs",
                "result": "PASS" if bad_enrollment_class == 0 else "FAIL",
                "detail": str(bad_enrollment_class),
            },
            {
                "check": "students_with_multiple_current_enrollments",
                "result": "PASS" if not multi_current_rows else "FAIL",
                "detail": str(len(multi_current_rows)),
            },
            {
                "check": "memory_db_integrity",
                "result": "PASS" if post_integrity == "ok" else "FAIL",
                "detail": str(post_integrity),
            },
            {
                "check": "memory_db_fk",
                "result": "PASS" if not post_fk else "FAIL",
                "detail": str(len(post_fk)),
            },
            {
                "check": "SAFE_SIMULATED_IMPORT",
                "result": "YES" if safe_sim else "NO",
                "detail": (
                    "Toàn bộ import chạy thành công trên bản sao DB trong RAM."
                    if safe_sim
                    else "Mô phỏng import chưa đạt."
                ),
            },
            {
                "check": "COMPARISON_PATCH_REQUIRED",
                "result": "YES" if patch_required else "NO",
                "detail": (
                    "Route hiện là logic cùng năm; phải sửa trước khi bật baseline năm trước."
                    if patch_required
                    else "Không phát hiện marker cũ."
                ),
            },
            {
                "check": "READY_FOR_V11_OFFICIAL_INSTALL",
                "result": "YES" if ready else "NO",
                "detail": (
                    "Có thể xây V11: backup + transaction import + patch route."
                    if ready
                    else "Chưa đủ điều kiện."
                ),
            },
        ]

        write_csv(
            out_dir / "520_SIMULATION_GATE.csv",
            ["check", "result", "detail"],
            gate_rows,
        )
        write_csv(
            out_dir / "521_DUPLICATE_TRANSFER_CODES.csv",
            [
                "student_code", "full_name", "enrollment_rows",
                "current_rows", "history", "check",
            ],
            duplicate_report,
        )
        write_csv(
            out_dir / "522_STATUS_MAPPING.csv",
            ["source_status", "mapped_status", "is_current", "rows"],
            status_report,
        )
        write_csv(
            out_dir / "523_TARGET_BASELINE_SCOPE.csv",
            [
                "target_id", "target_name", "baseline_school_ids",
                "all_history_enrollments",
                "current_expected_enrollments",
                "current_expected_unique_students",
            ],
            scope_report,
        )
        write_csv(
            out_dir / "524_ROUTE_PATCH_CHECK.csv",
            ["check", "result", "detail"],
            route_checks,
        )

        patch_plan = out_dir / "525_PATCH_PLAN_V11.txt"
        patch_plan.write_text(
            """V11 - KẾ HOẠCH SỬA ROUTE ĐỐI CHIẾU
====================================

1. Nguồn đối chiếu:
   - SurveyPersonYearRecord: giữ năm đợt điều tra (2026-2027).
   - StudentEnrollment: lấy năm học liền trước (2025-2026).

2. Phạm vi trường:
   - Trường không sáp nhập: baseline của chính trường.
   - Trường đích QĐ3805: baseline của trường đích + tất cả trường nguồn
     đã nhập vào trường đó.
   - Không đổi school_id lịch sử của enrollment 2025-2026.

3. Danh sách kỳ vọng:
   - _load_scope_enrollments chỉ lấy is_current=True để không biến hồ sơ
     chuyển đi/thôi học năm trước thành dòng "HS chưa có điều tra".
   - Những Student không-current vẫn giữ trong DB và vẫn có thể được tìm
     theo mã/định danh nếu xuất hiện trong điều tra năm hiện tại.

4. So trường:
   - Khi hai năm khác nhau, source school và target school QĐ3805 được
     xem là cùng nhóm.
   - Trường ngoài nhóm vẫn được báo là thay đổi trường để rà soát.

5. So lớp:
   - Không so class_id giữa 2025-2026 và 2026-2027 vì class_id là bản ghi
     theo năm và tất nhiên khác nhau.
   - V11 không tự suy diễn lên lớp; logic tiến lớp sẽ làm ở bước riêng nếu cần.

6. So trạng thái:
   - Không so trực tiếp trạng thái 2025-2026 với trạng thái 2026-2027 như
     hai nguồn cùng năm.
   - Trạng thái baseline vẫn hiển thị/lưu lịch sử.

7. Quyền liên kết:
   - Giáo viên/trường đích phải nhận diện được Student thuộc baseline của
     các trường nguồn QĐ3805.

8. An toàn:
   - Backup DB + file route trước khi thay đổi.
   - Import trong transaction.
   - Patch route có kiểm marker/cú pháp.
   - FK + integrity sau import.
   - Nếu lỗi: rollback DB và khôi phục route.
""",
            encoding="utf-8",
        )

        summary = out_dir / "00_TONG_QUAN_V10_10.txt"
        with summary.open("w", encoding="utf-8") as f:
            f.write(
                "DRY-RUN V10.10 - MÔ PHỎNG NẠP BASELINE + LOGIC ĐỐI CHIẾU QĐ3805\n"
            )
            f.write("=" * 128 + "\n")
            f.write("DB THẬT KHÔNG BỊ THAY ĐỔI\n\n")
            f.write(f"Baseline year: 2025-2026 id={baseline_year_id}\n")
            f.write(f"Current/test year: 2026-2027 id={current_year_id}\n\n")
            f.write(f"Source rows: {len(records)}\n")
            f.write(f"Unique student codes: {len(by_code)}\n")
            f.write(f"Duplicate transfer codes: {len(duplicate_codes)}\n")
            f.write(f"Classes source: {len(class_keys)}\n\n")

            f.write("MÔ PHỎNG TRONG RAM\n")
            f.write(f"- students: {sim_students}\n")
            f.write(f"- classes baseline QĐ3805: {sim_classes}\n")
            f.write(f"- enrollments baseline QĐ3805: {sim_enrollments}\n")
            f.write(
                f"- current baseline expected enrollments: {sim_current_enrollments}\n"
            )
            f.write(f"- bad class refs: {bad_enrollment_class}\n")
            f.write(
                f"- students with >1 current enrollment: {len(multi_current_rows)}\n"
            )
            f.write(f"- integrity: {post_integrity}\n")
            f.write(f"- FK errors: {len(post_fk)}\n\n")

            f.write("ROUTE\n")
            for r in route_checks:
                f.write(
                    f"- {r['check']}: {r['result']} | {r['detail']}\n"
                )

            f.write("\nKẾT LUẬN\n")
            f.write(
                "SAFE_SIMULATED_IMPORT = "
                + ("YES" if safe_sim else "NO") + "\n"
            )
            f.write(
                "COMPARISON_PATCH_REQUIRED = "
                + ("YES" if patch_required else "NO") + "\n"
            )
            f.write(
                "READY_FOR_V11_OFFICIAL_INSTALL = "
                + ("YES" if ready else "NO") + "\n"
            )

        zip_path = ROOT / f"dry_run_simulate_baseline_QD3805_v10_10_{ts}.zip"
        with zipfile.ZipFile(
            zip_path,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as zf:
            for p in sorted(out_dir.iterdir()):
                zf.write(p, arcname=p.name)

        print()
        print("=" * 128)
        print("HOÀN THÀNH DRY-RUN V10.10")
        print("=" * 128)
        print(f"Source rows: {len(records)}")
        print(f"Unique students: {len(by_code)}")
        print(f"Classes simulated: {sim_classes}")
        print(f"Enrollments simulated: {sim_enrollments}")
        print(f"Current baseline expected: {sim_current_enrollments}")
        print(f"Memory integrity: {post_integrity}")
        print(f"Memory FK errors: {len(post_fk)}")
        print(
            "SAFE_SIMULATED_IMPORT: "
            + ("YES" if safe_sim else "NO")
        )
        print(
            "COMPARISON_PATCH_REQUIRED: "
            + ("YES" if patch_required else "NO")
        )
        print(
            "READY_FOR_V11_OFFICIAL_INSTALL: "
            + ("YES" if ready else "NO")
        )
        print(f"ZIP: {zip_path}")
        print("DB thật / file nguồn / mã nguồn: KHÔNG bị thay đổi.")

    finally:
        real.close()
        mem.close()


if __name__ == "__main__":
    try:
        main()
    except AuditAbort as exc:
        print()
        print("=" * 128)
        print("ĐÃ DỪNG DRY-RUN V10.10")
        print("=" * 128)
        print(str(exc))
        print("DB thật / file nguồn / mã nguồn: KHÔNG bị thay đổi.")
        sys.exit(2)
    except Exception as exc:
        print()
        print("=" * 128)
        print("LỖI KHÔNG DỰ KIẾN TRONG V10.10")
        print("=" * 128)
        print(repr(exc))
        print("DB thật / file nguồn / mã nguồn: KHÔNG bị thay đổi.")
        sys.exit(4)
