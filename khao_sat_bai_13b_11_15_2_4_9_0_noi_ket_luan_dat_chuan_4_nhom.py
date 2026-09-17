from __future__ import annotations

import ast
import json
import os
import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT = Path(
    os.environ.get("PHOCAP_PROJECT", r"C:\PhoCap")
).resolve()

APP = PROJECT / "app"
EXPORTS = PROJECT / "exports"
DB = PROJECT / "data" / "phocap.db"

RULES = APP / "services" / "pcgd_business_rules.py"
ROADMAP = APP / "data" / "pcgd_recognition_roadmap_2026_2030.json"
BUILDERS = APP / "pcgd_xmc_report_builders_v1.py"
REPORT_CENTER = APP / "routers" / "report_center.py"
MN_REPORTS = APP / "routers" / "mn_official_reports.py"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT = EXPORTS / (
    "bao_cao_khao_sat_bai_13b_11_15_2_4_9_0_"
    f"noi_ket_luan_dat_chuan_4_nhom_{STAMP}.txt"
)

EVALUATORS = (
    "evaluate_preschool_3_5_commune",
    "evaluate_primary_commune",
    "evaluate_thcs_commune",
    "evaluate_literacy_commune",
)

LABEL_HINTS = (
    "đạt chuẩn",
    "đạt hay",
    "không đạt",
    "mức độ",
    "kết quả đánh giá",
    "tiêu chuẩn",
    "công nhận",
)

SOURCE_HINTS = (
    "pcgd_business_rules",
    "evaluate_preschool_3_5_commune",
    "evaluate_primary_commune",
    "evaluate_thcs_commune",
    "evaluate_literacy_commune",
    "RecognitionResult",
    "is_special_difficult",
    "is_difficult_area_from_roadmap",
    "Đạt chuẩn",
    "Mức độ",
)


def read_text(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(
        encoding="utf-8-sig",
        errors="replace",
    )


def db_check() -> dict:
    conn = sqlite3.connect(str(DB))
    try:
        integrity = str(
            conn.execute(
                "PRAGMA integrity_check"
            ).fetchone()[0]
        )
        fk_count = len(
            conn.execute(
                "PRAGMA foreign_key_check"
            ).fetchall()
        )

        tables = {
            row[0]
            for row in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            ).fetchall()
        }

        commune_columns = []
        if "communes" in tables:
            commune_columns = [
                row[1]
                for row in conn.execute(
                    "PRAGMA table_info(communes)"
                ).fetchall()
            ]

        return {
            "integrity": integrity,
            "fk_count": fk_count,
            "commune_columns": commune_columns,
        }
    finally:
        conn.close()


def evaluator_status() -> dict:
    source = read_text(RULES)

    if not source:
        return {
            "exists": False,
            "evaluators": {},
        }

    tree = ast.parse(source)

    functions = {
        node.name: node
        for node in tree.body
        if isinstance(node, ast.FunctionDef)
    }

    return {
        "exists": True,
        "evaluators": {
            name: (
                f"L{functions[name].lineno}-L{functions[name].end_lineno}"
                if name in functions
                else "THIẾU"
            )
            for name in EVALUATORS
        },
        "has_special_difficult": "is_special_difficult" in source,
        "has_roadmap_difficult": (
            "is_difficult_area_from_roadmap" in source
        ),
    }


def roadmap_info() -> dict:
    if not ROADMAP.exists():
        return {"exists": False}

    data = json.loads(
        ROADMAP.read_text(
            encoding="utf-8-sig"
        )
    )

    if isinstance(data, dict):
        records = (
            data.get("records")
            or data.get("items")
            or data.get("communes")
            or []
        )
    else:
        records = data

    if not isinstance(records, list):
        records = []

    difficult = sum(
        1
        for item in records
        if isinstance(item, dict)
        and item.get("is_difficult_area_from_roadmap") is True
    )

    years = {}
    for item in records:
        if not isinstance(item, dict):
            continue
        year = item.get("target_year")
        years[year] = years.get(year, 0) + 1

    return {
        "exists": True,
        "count": len(records),
        "difficult_count": difficult,
        "year_counts": years,
    }


def source_hits(path: Path) -> list[tuple[int, str]]:
    source = read_text(path)
    hits = []

    for lineno, line in enumerate(
        source.splitlines(),
        start=1,
    ):
        if any(
            hint.lower() in line.lower()
            for hint in SOURCE_HINTS
        ):
            hits.append((lineno, line.rstrip()))

    return hits


def extract_functions(path: Path) -> list[str]:
    source = read_text(path)
    if not source:
        return []

    tree = ast.parse(source)
    found = []

    for node in ast.walk(tree):
        if isinstance(
            node,
            (ast.FunctionDef, ast.AsyncFunctionDef),
        ):
            segment = ast.get_source_segment(
                source,
                node,
            ) or ""

            if any(
                hint.lower() in segment.lower()
                for hint in SOURCE_HINTS
            ):
                found.append(
                    f"{node.name} (L{node.lineno}-L{node.end_lineno})"
                )

    return found


def find_xlsx_files() -> list[Path]:
    result = []

    root = APP / "report_templates"
    if root.exists():
        for path in root.rglob("*.xlsx"):
            lowered = str(path).lower()

            if "backup" in lowered or "truoc_sua" in lowered:
                continue

            result.append(path)

    return sorted(set(result))


def workbook_labels(
    path: Path,
) -> list[tuple[str, str, str]]:
    rows = []

    try:
        wb = load_workbook(
            path,
            read_only=False,
            data_only=False,
        )
    except Exception as exc:
        return [("<LỖI>", "", str(exc))]

    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value

                    if not isinstance(value, str):
                        continue

                    normalized = value.strip()

                    if any(
                        hint in normalized.lower()
                        for hint in LABEL_HINTS
                    ):
                        rows.append(
                            (
                                ws.title,
                                cell.coordinate,
                                normalized.replace("\n", " | "),
                            )
                        )
    finally:
        wb.close()

    return rows


def main() -> None:
    print("=" * 148)
    print(
        "BÀI 13B-11.15.2.4.9.0 - "
        "KHẢO SÁT NỐI KẾT LUẬN ĐẠT CHUẨN/MỨC ĐỘ 4 NHÓM"
    )
    print("=" * 148)
    print()
    print("CHỈ ĐỌC:")
    print(" - Không sửa source.")
    print(" - Không sửa database.")
    print(" - Không sửa file Excel mẫu.")
    print()

    if not APP.exists():
        raise RuntimeError(f"Không tìm thấy: {APP}")

    if not DB.exists():
        raise RuntimeError(f"Không tìm thấy: {DB}")

    db = db_check()

    if db["integrity"] != "ok":
        raise RuntimeError(
            "Database integrity_check != ok."
        )

    if db["fk_count"] != 0:
        raise RuntimeError(
            "Database có lỗi foreign key."
        )

    rules = evaluator_status()
    roadmap = roadmap_info()

    source_files = [
        path
        for path in (
            RULES,
            BUILDERS,
            REPORT_CENTER,
            MN_REPORTS,
        )
        if path.exists()
    ]

    xlsx_files = find_xlsx_files()

    lines = [
        "=" * 148,
        "BÀI 13B-11.15.2.4.9.0 - KẾT QUẢ KHẢO SÁT",
        "=" * 148,
        "",
        "MỤC TIÊU:",
        " 1. Tự động kết luận Mầm non.",
        " 2. Tự động kết luận Tiểu học.",
        " 3. Tự động kết luận THCS.",
        " 4. Tự động kết luận Xóa mù chữ.",
        "",
        "NGUYÊN TẮC:",
        " - Kết luận là dữ liệu tự tính, không nhập tay.",
        " - Thiếu chỉ tiêu bắt buộc: không tự kết luận đạt.",
        " - Không đồng nhất 'vùng khó khăn' trong lộ trình "
        "với 'đặc biệt khó khăn' pháp lý.",
        "",
        "=" * 148,
        "1. BỘ QUY TẮC NGHIỆP VỤ",
        "=" * 148,
        f"Rules: {RULES}",
        f"Tồn tại: {rules.get('exists')}",
    ]

    for name, status in rules.get(
        "evaluators",
        {},
    ).items():
        lines.append(
            f" - {name}: {status}"
        )

    lines.extend(
        [
            f"Có tham số is_special_difficult: "
            f"{rules.get('has_special_difficult')}",
            f"Có cờ roadmap khó khăn trong service: "
            f"{rules.get('has_roadmap_difficult')}",
            "",
            "=" * 148,
            "2. LỘ TRÌNH CÔNG NHẬN",
            "=" * 148,
            repr(roadmap),
            "",
            "=" * 148,
            "3. DATABASE - PHÂN LOẠI ĐỊA BÀN",
            "=" * 148,
            "Các cột communes:",
            ", ".join(db["commune_columns"]),
            "",
            "=" * 148,
            "4. SOURCE PYTHON LIÊN QUAN",
            "=" * 148,
        ]
    )

    for path in source_files:
        lines.append("")
        lines.append(
            f"[{path.relative_to(PROJECT)}]"
        )

        try:
            functions = extract_functions(path)
        except Exception as exc:
            functions = []
            lines.append("AST lỗi: " + str(exc))

        if functions:
            lines.append("Hàm liên quan:")
            for item in functions:
                lines.append(" - " + item)

        hits = source_hits(path)

        if hits:
            lines.append("Dòng liên quan:")
            for lineno, text in hits[:220]:
                lines.append(
                    f" L{lineno}: {text}"
                )

    lines.extend(
        [
            "",
            "=" * 148,
            "5. Ô/NHÃN KẾT LUẬN TRONG FILE EXCEL MẪU",
            "=" * 148,
        ]
    )

    for path in xlsx_files:
        labels = workbook_labels(path)

        if not labels:
            continue

        lines.append("")
        lines.append(
            f"[{path.relative_to(PROJECT)}]"
        )

        for sheet, coord, label in labels:
            lines.append(
                f" - {sheet}!{coord}: {label}"
            )

    lines.extend(
        [
            "",
            "=" * 148,
            "6. ĐIỀU KIỆN ĐỂ LÀM BÀI 4.9.1",
            "=" * 148,
            " A. Có đủ 4 evaluator.",
            " B. Xác định đúng ô Đạt chuẩn/Mức độ trên từng biểu.",
            " C. Xác định đủ metric đầu vào từ builder hiện tại.",
            " D. Xác định nguồn pháp lý của 'đặc biệt khó khăn'.",
            "",
            "Nếu thiếu metric hoặc thiếu phân loại pháp lý: "
            "để trống/Chưa đủ dữ liệu, không suy đoán.",
            "",
            f"integrity_check: {db['integrity']}",
            f"foreign_key_check: {db['fk_count']}",
            "Source/database/Excel: KHÔNG THAY ĐỔI.",
        ]
    )

    EXPORTS.mkdir(
        parents=True,
        exist_ok=True,
    )

    OUT.write_text(
        "\n".join(lines),
        encoding="utf-8",
    )

    print("KIỂM TRA:")
    print(
        " - Bộ quy tắc:",
        "CÓ" if rules.get("exists") else "THIẾU",
    )
    print(
        " - Lộ trình:",
        "CÓ" if roadmap.get("exists") else "THIẾU",
    )
    print(
        " - File Excel mẫu đã quét:",
        len(xlsx_files),
    )
    print(" - integrity_check: OK")
    print(" - foreign_key_check: 0")
    print(" - Source/database/Excel: KHÔNG THAY ĐỔI")
    print()
    print("Báo cáo:", OUT)
    print()
    print("=" * 148)
    print(
        "KHẢO SÁT BÀI 13B-11.15.2.4.9.0 THÀNH CÔNG"
    )
    print("=" * 148)


if __name__ == "__main__":
    main()
