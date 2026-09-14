import sys
import hashlib
import unicodedata
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from openpyxl import load_workbook

ROOTS = [
    Path(r"C:\PhoCap"),
    Path.home() / "Desktop",
    Path.home() / "Downloads",
    Path.home() / "Documents",
]

DB = Path(r"C:\PhoCap\data\phocap.db")

EXCLUDE_PARTS = {
    ".venv",
    "__pycache__",
    ".git",
    "backups",
    "backup",
}

def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        while True:
            b = f.read(1024 * 1024)
            if not b:
                break
            h.update(b)
    return h.hexdigest()

def norm(v):
    s = str(v or "").strip().lower()
    s = "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )
    s = s.replace("đ", "d")
    return " ".join(s.split())

def excluded(path):
    parts = {x.lower() for x in path.parts}
    return any(x in parts for x in EXCLUDE_PARTS)

db_before = sha256(DB)

print("=" * 120)
print("TÌM FILE NGUỒN ĐỘI NGŨ MẦM NON 2025-2026")
print("CHỈ ĐỌC - KHÔNG SỬA FILE - KHÔNG GHI DATABASE")
print("=" * 120)

files = set()

for root in ROOTS:
    if not root.exists():
        continue

    for pattern in ("*.xlsx", "*.xlsm", "*.xls"):
        try:
            for p in root.rglob(pattern):
                if excluded(p):
                    continue
                files.add(p.resolve())
        except Exception:
            pass

print("Tổng file Excel tìm thấy:", len(files))

xlsx_candidates = []
xls_candidates = []

for p in sorted(files):
    suffix = p.suffix.lower()

    if suffix == ".xls":
        name = norm(p.name)

        if any(
            k in name
            for k in (
                "giao vien",
                "doi ngu",
                "mam non",
                "mn",
                "can bo",
                "nhan vien",
            )
        ):
            xls_candidates.append(p)

        continue

    try:
        wb = load_workbook(
            p,
            read_only=True,
            data_only=True,
        )
    except Exception:
        continue

    score = 0
    evidence = []
    sheet_info = []

    for ws in wb.worksheets:
        max_row = int(ws.max_row or 0)
        max_col = int(ws.max_column or 0)

        if max_row >= 20:
            score += 2
            evidence.append(
                f"{ws.title}: có {max_row} dòng"
            )

        texts = []

        # Quét tối đa 80 dòng đầu, đủ để nhận header và dữ liệu mẫu.
        for row_no, row in enumerate(
            ws.iter_rows(values_only=True),
            1,
        ):
            if row_no > 80:
                break

            vals = [
                norm(v)
                for v in row
                if v is not None
                and str(v).strip() != ""
            ]

            if vals:
                texts.extend(vals)

        joined = " | ".join(texts)

        if "2025-2026" in joined:
            score += 3
            evidence.append(
                f"{ws.title}: có năm 2025-2026"
            )

        if (
            "mam non" in joined
            or "cap mn" in joined
            or "truong mn" in joined
        ):
            score += 4
            evidence.append(
                f"{ws.title}: có dấu hiệu Mầm non"
            )

        if "giao vien" in joined:
            score += 2
            evidence.append(
                f"{ws.title}: có Giáo viên"
            )

        if (
            "ma can bo" in joined
            or "ma giao vien" in joined
            or "ma nhan su" in joined
            or "staff code" in joined
        ):
            score += 3
            evidence.append(
                f"{ws.title}: có dấu hiệu mã nhân sự"
            )

        if (
            "truong" in joined
            and (
                "ho ten" in joined
                or "ho va ten" in joined
            )
        ):
            score += 2
            evidence.append(
                f"{ws.title}: có Trường + Họ tên"
            )

        sheet_info.append(
            (ws.title, max_row, max_col)
        )

    wb.close()

    # Chỉ giữ file có dấu hiệu đáng xem.
    if score >= 3:
        xlsx_candidates.append({
            "score": score,
            "path": p,
            "size": p.stat().st_size,
            "sheets": sheet_info,
            "evidence": evidence,
        })

xlsx_candidates.sort(
    key=lambda x: (-x["score"], str(x["path"]))
)

print()
print("=" * 120)
print("ỨNG VIÊN XLSX/XLSM")
print("=" * 120)

if not xlsx_candidates:
    print("<KHÔNG TÌM THẤY FILE XLSX/XLSM PHÙ HỢP>")
else:
    for i, item in enumerate(
        xlsx_candidates,
        1,
    ):
        print()
        print(
            f"{i}. SCORE={item['score']}"
        )
        print("   File :", item["path"])
        print("   Size :", item["size"], "bytes")

        print("   Sheets:")
        for title, rows, cols in item["sheets"]:
            print(
                f"      - {title}: "
                f"{rows} dòng x {cols} cột"
            )

        print("   Dấu hiệu:")
        for x in item["evidence"]:
            print("      -", x)

print()
print("=" * 120)
print("ỨNG VIÊN .XLS CŨ - CHỈ LIỆT KÊ TÊN")
print("=" * 120)

if not xls_candidates:
    print("<KHÔNG CÓ>")
else:
    for p in xls_candidates:
        print(p)

db_after = sha256(DB)

print()
print("=" * 120)
print("KIỂM TRA AN TOÀN")
print("=" * 120)
print("DB hash trước:", db_before)
print("DB hash sau  :", db_after)

if db_before != db_after:
    raise RuntimeError(
        "DỪNG: phocap.db thay đổi ngoài dự kiến."
    )

print("Database     : KHÔNG THAY ĐỔI")
print("=" * 120)
