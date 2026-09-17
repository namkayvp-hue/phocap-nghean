from __future__ import annotations

import ast
import os
import re
import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


PROJECT = Path(
    os.environ.get("PHOCAP_PROJECT", r"C:\PhoCap")
).resolve()

APP = PROJECT / "app"
EXPORTS = PROJECT / "exports"
DB = PROJECT / "data" / "phocap.db"

RULES = APP / "services" / "pcgd_business_rules.py"
MODELS = APP / "models.py"
MN_BUILDERS = APP / "pcgd_mn_report_builders_v1.py"
XMC_BUILDERS = APP / "pcgd_xmc_report_builders_v1.py"
REPORT_CENTER = APP / "routers" / "report_center.py"
MN_OFFICIAL = APP / "routers" / "mn_official_reports.py"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT = EXPORTS / (
    "bao_cao_bai_13b_11_15_2_4_9_0_1_"
    f"khoa_nguon_metric_ket_luan_4_nhom_{STAMP}.txt"
)

FUNCTIONS = {
    MN_BUILDERS: (
        "_fill_mn_m1",
        "_fill_mn02",
        "export_mn_report",
    ),
    XMC_BUILDERS: (
        "_build_th02",
        "_build_thcs_m1",
        "_build_thcs_m2",
        "_build_thcs_tk",
        "_build_xmc4",
        "export_additional_report",
    ),
}

TEMPLATES = (
    (
        "MN_NEW",
        APP / "report_templates" / "Bieu_mau_PCGDMN_2025.xlsx",
        ("MN-02",),
        ("đạt chuẩn",),
    ),
    (
        "MN_02",
        APP / "report_templates" / "pcgd_mn_2025" / "PCGD_2025_MN_02.xlsx",
        None,
        ("đạt chuẩn",),
    ),
    (
        "TH_02",
        APP / "report_templates" / "pcgd_xmc_2025" / "PCGD_2025_TH_02.xlsx",
        None,
        ("đạt chuẩn", "mức độ"),
    ),
    (
        "THCS_M2",
        APP / "report_templates" / "pcgd_xmc_2025" / "PCGD_2025_THCS_M2.xlsx",
        None,
        ("đạt hay", "kết quả đánh giá"),
    ),
    (
        "THCS_TK",
        APP / "report_templates" / "pcgd_xmc_2025" / "PCGD_2025_THCS_TK.xlsx",
        None,
        ("đạt chuẩn", "mức độ"),
    ),
    (
        "XMC_4",
        APP / "report_templates" / "pcgd_xmc_2025" / "PCGD_2025_XMC_4.xlsx",
        None,
        ("đạt", "mức độ"),
    ),
)

CELL_HINTS = {
    "MN_NEW": (
        "V4",
    ),
    "MN_02": (
        "R3",
        "H7",
        "K7",
        "O7",
    ),
    "TH_02": (
        "T4",
        "I8",
        "K8",
        "M8",
        "Q8",
        "T8",
    ),
    "THCS_M2": (
        "W9",
        "E14",
        "M14",
        "V14",
        "W14",
        "E15",
        "M15",
        "V15",
        "W15",
    ),
    "THCS_TK": (
        "F4",
        "G4",
        "R4",
        "F8",
        "G8",
        "R8",
    ),
    "XMC_4": (
        "R5",
        "E8",
        "G8",
        "J8",
        "L8",
        "O8",
        "Q8",
        "R8",
    ),
}


def read_text(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(
        encoding="utf-8-sig",
        errors="replace",
    )


def function_source(
    path: Path,
    name: str,
) -> tuple[int, int, str] | None:
    source = read_text(path)
    if not source:
        return None

    tree = ast.parse(source)

    matches = [
        node
        for node in tree.body
        if isinstance(
            node,
            (ast.FunctionDef, ast.AsyncFunctionDef),
        )
        and node.name == name
    ]

    if len(matches) != 1:
        return None

    node = matches[0]
    segment = ast.get_source_segment(
        source,
        node,
    ) or ""

    return (
        int(node.lineno),
        int(node.end_lineno),
        segment,
    )


def compact_lines(
    text: str,
    *,
    max_lines: int = 90,
) -> list[str]:
    result = []

    keywords = (
        "ws[",
        "ws.cell",
        "selected_commune_id",
        "commune_id",
        "meta",
        "kind",
        "is_special",
        "is_difficult",
        "percent",
        "completed",
        "upper",
        "grade1",
        "age11",
        "age14",
        "age15",
        "literacy",
        "workbook.save",
        "return StreamingResponse",
    )

    for line in text.splitlines():
        if any(
            keyword.lower() in line.lower()
            for keyword in keywords
        ):
            result.append(
                line.rstrip()
            )

    return result[:max_lines]


def db_info() -> dict:
    conn = sqlite3.connect(
        str(DB)
    )

    try:
        integrity = str(
            conn.execute(
                "PRAGMA integrity_check"
            ).fetchone()[0]
        )
        fk = len(
            conn.execute(
                "PRAGMA foreign_key_check"
            ).fetchall()
        )

        cols = [
            row[1]
            for row in conn.execute(
                "PRAGMA table_info(communes)"
            ).fetchall()
        ]

        count = int(
            conn.execute(
                "SELECT COUNT(*) FROM communes"
            ).fetchone()[0]
        )

        special_count = None

        if "is_special_difficulty_area" in cols:
            special_count = int(
                conn.execute(
                    "SELECT COUNT(*) FROM communes "
                    "WHERE is_special_difficulty_area = 1"
                ).fetchone()[0]
            )

        return {
            "integrity": integrity,
            "fk": fk,
            "commune_columns": cols,
            "commune_count": count,
            "special_count": special_count,
        }

    finally:
        conn.close()


def merged_range_for(
    ws,
    coord: str,
) -> str | None:
    for rng in ws.merged_cells.ranges:
        if coord in rng:
            return str(rng)
    return None


def value_repr(value) -> str:
    if value is None:
        return "<trống>"
    text = str(value).replace(
        "\n",
        " | ",
    )
    if len(text) > 180:
        text = text[:177] + "..."
    return text


def inspect_template(
    code: str,
    path: Path,
    sheet_filter,
    label_hints,
) -> list[str]:
    lines = []

    if not path.exists():
        return [
            f"[{code}] THIẾU FILE: {path}"
        ]

    try:
        wb = load_workbook(
            path,
            data_only=False,
            read_only=False,
        )
    except Exception as exc:
        return [
            f"[{code}] LỖI MỞ FILE: {exc}"
        ]

    try:
        sheets = []

        for ws in wb.worksheets:
            if (
                sheet_filter
                and ws.title not in sheet_filter
            ):
                continue
            sheets.append(ws)

        if not sheets:
            sheets = list(
                wb.worksheets
            )

        lines.append(
            f"[{code}] {path.relative_to(PROJECT)}"
        )

        for ws in sheets:
            lines.append(
                f"  Sheet: {ws.title}"
            )

            hits = []

            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(
                        cell,
                        MergedCell,
                    ):
                        continue

                    value = cell.value

                    if not isinstance(
                        value,
                        str,
                    ):
                        continue

                    lower = value.lower()

                    if any(
                        hint in lower
                        for hint in label_hints
                    ):
                        hits.append(
                            cell.coordinate
                        )

            for coord in hits[:30]:
                cell = ws[coord]
                lines.append(
                    "    Nhãn "
                    + coord
                    + " = "
                    + value_repr(
                        cell.value
                    )
                )

                merged = merged_range_for(
                    ws,
                    coord,
                )

                if merged:
                    lines.append(
                        "      merged: "
                        + merged
                    )

            hints = CELL_HINTS.get(
                code,
                (),
            )

            if hints:
                lines.append(
                    "    Các ô khóa:"
                )

                for coord in hints:
                    cell = ws[coord]

                    lines.append(
                        f"      {coord}: "
                        f"{value_repr(cell.value)} "
                        f"| format={cell.number_format!r}"
                    )

                    merged = merged_range_for(
                        ws,
                        coord,
                    )

                    if merged:
                        lines.append(
                            "        merged: "
                            + merged
                        )

        return lines

    finally:
        wb.close()


def source_assignment_hits(
    coord: str,
) -> list[str]:
    patterns = (
        f'["{coord}"]',
        f"['{coord}']",
    )

    result = []

    for path in (
        MN_BUILDERS,
        XMC_BUILDERS,
        REPORT_CENTER,
        MN_OFFICIAL,
    ):
        source = read_text(
            path
        )

        if not source:
            continue

        for lineno, line in enumerate(
            source.splitlines(),
            start=1,
        ):
            if any(
                token in line
                for token in patterns
            ):
                result.append(
                    f"{path.relative_to(PROJECT)}:"
                    f"{lineno}: {line.strip()}"
                )

    return result


def main() -> None:
    print("=" * 150)
    print(
        "BÀI 13B-11.15.2.4.9.0.1 - "
        "KHÓA NGUỒN METRIC KẾT LUẬN 4 NHÓM"
    )
    print("=" * 150)
    print()
    print("CHỈ ĐỌC - KHÔNG SỬA SOURCE/DATABASE/EXCEL.")
    print()

    required = (
        APP,
        DB,
        RULES,
        MODELS,
        XMC_BUILDERS,
    )

    for path in required:
        if not path.exists():
            raise RuntimeError(
                f"Không tìm thấy: {path}"
            )

    db = db_info()

    if db["integrity"] != "ok":
        raise RuntimeError(
            "integrity_check != ok"
        )

    if db["fk"] != 0:
        raise RuntimeError(
            "foreign_key_check có lỗi"
        )

    rules_source = read_text(
        RULES
    )

    evaluators = (
        "evaluate_preschool_3_5_commune",
        "evaluate_primary_commune",
        "evaluate_thcs_commune",
        "evaluate_literacy_commune",
    )

    lines = [
        "=" * 150,
        "BÀI 13B-11.15.2.4.9.0.1 - KẾT QUẢ",
        "=" * 150,
        "",
        "I. KẾT QUẢ BÀI 4.9.0",
        " - 4 evaluator cấp XÃ/PHƯỜNG phải có đủ.",
        " - Không dùng evaluator cấp xã để kết luận Toàn tỉnh.",
        " - Cờ pháp lý ưu tiên: communes.is_special_difficulty_area.",
        "",
        "II. DATABASE",
        f" - integrity_check: {db['integrity']}",
        f" - foreign_key_check: {db['fk']}",
        f" - Tổng communes: {db['commune_count']}",
        " - Cột communes: "
        + ", ".join(
            db["commune_columns"]
        ),
        " - Số xã/phường is_special_difficulty_area=1: "
        + str(
            db["special_count"]
        ),
        "",
        "III. EVALUATOR",
    ]

    for name in evaluators:
        lines.append(
            " - "
            + name
            + ": "
            + (
                "CÓ"
                if f"def {name}(" in rules_source
                else "THIẾU"
            )
        )

    lines.extend(
        [
            "",
            "IV. SOURCE BUILDER / EXPORT",
        ]
    )

    for path, names in FUNCTIONS.items():
        lines.append("")
        lines.append(
            f"[{path.relative_to(PROJECT)}]"
        )

        for name in names:
            info = function_source(
                path,
                name,
            )

            if info is None:
                lines.append(
                    f" - {name}: KHÔNG TÌM THẤY DUY NHẤT"
                )
                continue

            start, end, body = info

            lines.append(
                f" - {name}: L{start}-L{end}"
            )

            for item in compact_lines(
                body
            ):
                lines.append(
                    "      "
                    + item
                )

    lines.extend(
        [
            "",
            "V. TEMPLATE + Ô KẾT LUẬN/METRIC",
        ]
    )

    for (
        code,
        path,
        sheet_filter,
        label_hints,
    ) in TEMPLATES:
        lines.append("")

        lines.extend(
            inspect_template(
                code,
                path,
                sheet_filter,
                label_hints,
            )
        )

        for coord in CELL_HINTS.get(
            code,
            (),
        ):
            hits = source_assignment_hits(
                coord
            )

            if hits:
                lines.append(
                    f"    Dấu vết source ghi {coord}:"
                )

                for hit in hits[:20]:
                    lines.append(
                        "      - "
                        + hit
                    )

    lines.extend(
        [
            "",
            "VI. ĐIỂM KHÓA TRƯỚC KHI CÀI 4.9.1",
            "",
            "1. MẦM NON:",
            " - Xác định report nào là biểu kết luận chính hiện hành:",
            "   Bieu_mau_PCGDMN_2025.xlsx/MN-02 hay PCGD_2025_MN_02.xlsx.",
            " - Chỉ dùng evaluate_preschool_3_5_commune cho chuẩn 3-5 tuổi NĐ277.",
            "",
            "2. TIỂU HỌC:",
            " - Metric evaluator cần: age6_grade1_rate, age14_completed_rate,",
            "   age11_completed_rate, age11_remaining_all_study_primary.",
            " - Đối chiếu với TH_02/M1 trước khi ghi Mức độ.",
            "",
            "3. THCS:",
            " - Metric evaluator cần primary_level + literacy_level +",
            "   age15_18_graduated_thcs_rate + age15_18_upper_program_rate.",
            " - Không kết luận THCS nếu thiếu level TH hoặc XMC.",
            "",
            "4. XMC:",
            " - Metric evaluator dùng mức biết chữ theo nhóm tuổi.",
            " - XMC-4 là biểu kết luận phù hợp để ghi Mức độ khi đủ tỷ lệ.",
            "",
            "5. PHẠM VI:",
            " - Bài 4.9.1 trước hết chỉ tự kết luận XÃ/PHƯỜNG.",
            " - Toàn tỉnh cần bộ quy tắc công nhận cấp tỉnh riêng, làm bước sau.",
            "",
            "Source/database/Excel: KHÔNG THAY ĐỔI.",
        ]
    )

    EXPORTS.mkdir(
        parents=True,
        exist_ok=True,
    )

    OUT.write_text(
        "\n".join(
            lines
        ),
        encoding="utf-8",
    )

    print("KIỂM TRA:")
    print(" - integrity_check: OK")
    print(" - foreign_key_check: 0")
    print(
        " - communes.is_special_difficulty_area:",
        (
            "CÓ"
            if "is_special_difficulty_area"
            in db["commune_columns"]
            else "THIẾU"
        ),
    )
    print(" - Source/database/Excel: KHÔNG THAY ĐỔI")
    print()
    print("Báo cáo:", OUT)
    print()
    print("=" * 150)
    print(
        "KHẢO SÁT BÀI 13B-11.15.2.4.9.0.1 THÀNH CÔNG"
    )
    print("=" * 150)


if __name__ == "__main__":
    main()
