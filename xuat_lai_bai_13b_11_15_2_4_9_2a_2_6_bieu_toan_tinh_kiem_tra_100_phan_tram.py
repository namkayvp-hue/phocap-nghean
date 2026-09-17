# -*- coding: utf-8 -*-
"""
BÀI 13B-11.15.2.4.9.2A.2
XUẤT LẠI 6 BIỂU TOÀN TỈNH - KIỂM TRA HIỂN THỊ 100%

Sau Bài 4.9.2A.1:
- MN / TH / THCS: tỷ lệ phải dùng number_format = 0\\%
  Ví dụ giá trị 100 phải hiển thị: 100%
- XMC: giữ nguyên cách hiển thị hiện tại.

Script:
- Gọi exporter thực tế.
- Xuất lại 6 file Excel Toàn tỉnh.
- Kiểm tra các ô kết luận.
- Kiểm tra các ô tỷ lệ MN / TH / THCS.
- Không sửa source.
- Không ghi database.
- Kiểm tra SHA256 database trước/sau.
"""

from __future__ import annotations

import asyncio
import hashlib
import inspect
import os
import sqlite3
import sys
import traceback
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any


PROJECT = Path(
    os.environ.get(
        "PHOCAP_PROJECT",
        r"C:\PhoCap",
    )
).resolve()

APP = PROJECT / "app"
DB_PATH = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"

STAMP = datetime.now().strftime(
    "%Y%m%d_%H%M%S"
)

OUT_DIR = (
    EXPORTS
    / (
        "xuat_lai_bai_13b_11_15_2_4_9_2a_2_"
        f"kiem_tra_100_phan_tram_{STAMP}"
    )
)

REPORT = (
    OUT_DIR
    / (
        "bao_cao_xuat_lai_6_bieu_"
        f"kiem_tra_100_phan_tram_{STAMP}.txt"
    )
)

EXPECTED_COMMUNES = 130
EXPECTED_SPECIAL = 51
EXPECTED_PERCENT_FORMAT = r"0\%"


REPORTS = (
    {
        "code": "PCGD_MN_02_2025",
        "short": "MN-02",
        "exporter": "mn",
        "conclusion_cells": ("R7",),
        "percent_cells": (
            "H7",
            "K7",
            "O7",
        ),
        "check_percent_format": True,
    },
    {
        "code": "PCGD_TH_02_2025",
        "short": "TH-02",
        "exporter": "xmc",
        "conclusion_cells": ("T8",),
        "percent_cells": (
            "I8",
            "K8",
            "M8",
            "Q8",
        ),
        "check_percent_format": True,
    },
    {
        "code": "PCGD_XMC_4_2025",
        "short": "XMC-4",
        "exporter": "xmc",
        "conclusion_cells": ("R8",),
        "percent_cells": (
            "E8",
            "G8",
            "J8",
            "L8",
            "O8",
            "Q8",
        ),
        "check_percent_format": False,
    },
    {
        "code": "PCGD_THCS_M1_2025",
        "short": "THCS-M1",
        "exporter": "xmc",
        "conclusion_cells": (
            "G36",
            "G37",
        ),
        "percent_cells": (
            "H39",
            "H40",
            "H41",
        ),
        "check_percent_format": True,
    },
    {
        "code": "PCGD_THCS_M2_2025",
        "short": "THCS-M2",
        "exporter": "xmc",
        "conclusion_cells": ("W14",),
        "percent_cells": (
            "E14",
            "J14",
            "M14",
            "Q14",
            "V14",
            "E15",
            "J15",
            "M15",
            "Q15",
            "V15",
            "K20",
            "K21",
            "K22",
            "K23",
            "K24",
        ),
        "check_percent_format": True,
    },
    {
        "code": "PCGD_THCS_TK_2025",
        "short": "THCS-TK",
        "exporter": "xmc",
        "conclusion_cells": (
            "F8",
            "G8",
            "R8",
        ),
        "percent_cells": (
            "I8",
            "K8",
            "O8",
        ),
        "check_percent_format": True,
    },
)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(
                1024 * 1024
            ),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


def db_state() -> dict[str, Any]:
    uri = (
        DB_PATH.resolve().as_uri()
        + "?mode=ro"
    )

    con = sqlite3.connect(
        uri,
        uri=True,
    )

    try:
        integrity = str(
            con.execute(
                "PRAGMA integrity_check"
            ).fetchone()[0]
        )

        fk = len(
            con.execute(
                "PRAGMA foreign_key_check"
            ).fetchall()
        )

        total = int(
            con.execute(
                "SELECT COUNT(*) FROM communes"
            ).fetchone()[0]
        )

        special = int(
            con.execute(
                """
                SELECT COUNT(*)
                FROM communes
                WHERE COALESCE(
                    is_special_difficulty_area,
                    0
                ) = 1
                """
            ).fetchone()[0]
        )

        return {
            "integrity": integrity,
            "fk": fk,
            "total": total,
            "special": special,
        }

    finally:
        con.close()


def verify_db(
    state: dict[str, Any],
) -> None:
    if (
        str(
            state["integrity"]
        ).lower()
        != "ok"
    ):
        raise RuntimeError(
            "integrity_check != ok"
        )

    if int(
        state["fk"]
    ) != 0:
        raise RuntimeError(
            "foreign_key_check có lỗi."
        )

    if int(
        state["total"]
    ) != EXPECTED_COMMUNES:
        raise RuntimeError(
            f"communes={state['total']}, "
            f"không phải {EXPECTED_COMMUNES}."
        )

    if int(
        state["special"]
    ) != EXPECTED_SPECIAL:
        raise RuntimeError(
            f"ĐBKK={state['special']}, "
            f"không phải {EXPECTED_SPECIAL}."
        )


def verify_source_markers() -> None:
    mn = (
        APP
        / "pcgd_mn_report_builders_v1.py"
    ).read_text(
        encoding="utf-8-sig",
        errors="strict",
    )

    xmc = (
        APP
        / "pcgd_xmc_report_builders_v1.py"
    ).read_text(
        encoding="utf-8-sig",
        errors="strict",
    )

    required_mn = (
        "BAI_13B_11_15_2_4_9_2_"
        "MN_HELPERS_START",
        "BAI_13B_11_15_2_4_9_2A_1_"
        "MN_HELPERS_START",
        "_b492a1_format_mn_percent_display",
    )

    required_xmc = (
        "BAI_13B_11_15_2_4_9_2_"
        "XMC_HELPERS_START",
        "BAI_13B_11_15_2_4_9_2A_1_"
        "XMC_HELPERS_START",
        "_b492a1_format_th_thcs_percent_display",
    )

    missing = (
        [
            "MN:" + token
            for token in required_mn
            if token not in mn
        ]
        + [
            "XMC:" + token
            for token in required_xmc
            if token not in xmc
        ]
    )

    if missing:
        raise RuntimeError(
            "Thiếu marker Bài 4.9.2/"
            "4.9.2A.1: "
            + "; ".join(missing)
        )


def make_request():
    from starlette.requests import (
        Request,
    )

    scope = {
        "type": "http",
        "asgi": {
            "version": "3.0",
            "spec_version": "2.3",
        },
        "http_version": "1.1",
        "server": (
            "127.0.0.1",
            80,
        ),
        "client": (
            "127.0.0.1",
            0,
        ),
        "scheme": "http",
        "method": "GET",
        "root_path": "",
        "path": (
            "/kiem-tra-"
            "bai-4-9-2a-2"
        ),
        "raw_path": (
            b"/kiem-tra-"
            b"bai-4-9-2a-2"
        ),
        "query_string": b"",
        "headers": [],
        "auth_user": {
            "id": 0,
            "username": (
                "kiem_thu_4_9_2a_2"
            ),
            "full_name": (
                "Kiểm thử 4.9.2A.2"
            ),
            "role_code": "ADMIN",
            "commune_id": None,
            "school_id": None,
        },
    }

    return Request(scope)


async def response_to_bytes(
    response,
) -> bytes:
    body = getattr(
        response,
        "body",
        None,
    )

    if isinstance(
        body,
        (
            bytes,
            bytearray,
        ),
    ) and body:
        return bytes(body)

    iterator = getattr(
        response,
        "body_iterator",
        None,
    )

    if iterator is None:
        raise RuntimeError(
            "Response không có "
            "body/body_iterator."
        )

    chunks = []

    async for chunk in iterator:
        if isinstance(
            chunk,
            str,
        ):
            chunk = chunk.encode(
                "utf-8"
            )

        chunks.append(
            bytes(chunk)
        )

    return b"".join(
        chunks
    )


def choose_year(db):
    from sqlalchemy import text
    from sqlalchemy import select
    from app.models import SchoolYear

    requested = str(
        os.environ.get(
            "PHOCAP_TEST_YEAR",
            "",
        )
        or ""
    ).strip()

    if requested:
        school_year = db.scalar(
            select(
                SchoolYear
            ).where(
                SchoolYear.code
                == requested
            )
        )

        if school_year is None:
            raise RuntimeError(
                "Không có năm học "
                + requested
            )

        return school_year

    rows = db.execute(
        text(
            """
            SELECT
                sy.id,
                sy.code,
                COUNT(spr.id) AS cnt
            FROM school_years sy
            LEFT JOIN
                survey_person_year_records spr
              ON spr.school_year_id = sy.id
            GROUP BY sy.id, sy.code
            ORDER BY cnt DESC, sy.code DESC
            """
        )
    ).all()

    if not rows:
        raise RuntimeError(
            "Database chưa có năm học."
        )

    best_id = int(
        rows[0][0]
    )

    school_year = db.scalar(
        select(
            SchoolYear
        ).where(
            SchoolYear.id
            == best_id
        )
    )

    if school_year is None:
        raise RuntimeError(
            f"Không lấy được "
            f"SchoolYear ID={best_id}."
        )

    return school_year


def year_record_count(
    db,
    school_year_id: int,
) -> int:
    from sqlalchemy import text

    value = db.execute(
        text(
            """
            SELECT COUNT(*)
            FROM survey_person_year_records
            WHERE school_year_id=:id
            """
        ),
        {
            "id": int(
                school_year_id
            )
        },
    ).scalar_one()

    return int(
        value
        or 0
    )


def call_exporter(
    func,
    *,
    report_type: str,
    db,
    request,
    school_year,
):
    candidates = {
        "report_type": report_type,
        "db": db,
        "request": request,
        "school_year": school_year,
        "selected_commune_id": None,
        "selected_school_id": None,
        "scope_label": "Toàn tỉnh",
    }

    signature = inspect.signature(
        func
    )

    kwargs = {}

    for (
        name,
        parameter,
    ) in signature.parameters.items():
        if name in candidates:
            kwargs[name] = (
                candidates[name]
            )
            continue

        if (
            parameter.default
            is not inspect.Parameter.empty
        ):
            continue

        if parameter.kind in (
            inspect.Parameter.VAR_POSITIONAL,
            inspect.Parameter.VAR_KEYWORD,
        ):
            continue

        raise RuntimeError(
            f"Exporter {func.__name__}: "
            "chưa biết tham số bắt buộc "
            + name
        )

    return func(
        **kwargs
    )


def normalized_value(value):
    if value is None:
        return None

    if isinstance(
        value,
        bool,
    ):
        return value

    if isinstance(
        value,
        int,
    ):
        return value

    if isinstance(
        value,
        float,
    ):
        if value.is_integer():
            return int(value)
        return value

    return " ".join(
        str(value)
        .strip()
        .split()
    )


def display_percent(
    value,
    number_format: str,
):
    """
    Chỉ dùng để in ra terminal
    cách người dùng sẽ nhìn thấy
    đối với format 0\\%.
    """
    value = normalized_value(
        value
    )

    if (
        number_format
        == EXPECTED_PERCENT_FORMAT
        and isinstance(
            value,
            (
                int,
                float,
            ),
        )
    ):
        return (
            f"{float(value):.0f}%"
        )

    if value is None:
        return "(trống)"

    return str(value)


def conclusion_allowed(
    short: str,
    cell: str,
    value,
) -> bool:
    value = normalized_value(
        value
    )

    status = {
        "Không đạt",
        "Chưa đủ dữ liệu",
    }

    if (
        short == "MN-02"
        and cell == "R7"
    ):
        return value in {
            "Đạt",
            *status,
        }

    if short == "TH-02":
        return value in {
            1,
            2,
            3,
            *status,
        }

    if short == "XMC-4":
        return value in {
            1,
            2,
            *status,
        }

    if short == "THCS-M1":
        if cell == "G36":
            return value in {
                1,
                2,
                3,
                *status,
            }

        if cell == "G37":
            return value in {
                1,
                2,
                *status,
            }

    if short == "THCS-M2":
        return value in {
            1,
            2,
            3,
            *status,
        }

    if short == "THCS-TK":
        if cell == "F8":
            return value in {
                1,
                2,
                3,
                *status,
            }

        if cell == "G8":
            return value in {
                1,
                2,
                *status,
            }

        if cell == "R8":
            return value in {
                1,
                2,
                3,
                *status,
            }

    return False


def inspect_xlsx(
    data: bytes,
    spec: dict,
):
    from openpyxl import (
        load_workbook,
    )

    if not data.startswith(
        b"PK"
    ):
        raise RuntimeError(
            f"{spec['short']}: "
            "response không phải XLSX."
        )

    workbook = load_workbook(
        BytesIO(data),
        data_only=False,
    )

    try:
        ws = workbook[
            workbook.sheetnames[0]
        ]

        conclusions = {}

        for ref in (
            spec[
                "conclusion_cells"
            ]
        ):
            value = normalized_value(
                ws[ref].value
            )

            if not conclusion_allowed(
                spec["short"],
                ref,
                value,
            ):
                raise RuntimeError(
                    f"{spec['short']}: "
                    f"{ref}={value!r} "
                    "không hợp lệ."
                )

            conclusions[
                ref
            ] = value

        percentages = {}

        for ref in (
            spec[
                "percent_cells"
            ]
        ):
            cell = ws[ref]

            item = {
                "value": (
                    normalized_value(
                        cell.value
                    )
                ),
                "format": str(
                    cell.number_format
                    or ""
                ),
            }

            item["display"] = (
                display_percent(
                    item["value"],
                    item["format"],
                )
            )

            if (
                spec[
                    "check_percent_format"
                ]
                and item["format"]
                != EXPECTED_PERCENT_FORMAT
            ):
                raise RuntimeError(
                    f"{spec['short']}: "
                    f"{ref} format="
                    f"{item['format']!r}; "
                    f"yêu cầu="
                    f"{EXPECTED_PERCENT_FORMAT!r}."
                )

            percentages[
                ref
            ] = item

        return (
            ws.title,
            conclusions,
            percentages,
        )

    finally:
        workbook.close()


def conclusion_text(
    values: dict,
) -> str:
    return ", ".join(
        f"{ref}={value!r}"
        for (
            ref,
            value,
        ) in values.items()
    )


def percent_text(
    values: dict,
) -> str:
    return ", ".join(
        (
            f"{ref}="
            f"{item['display']}"
            f"[fmt={item['format']}]"
        )
        for (
            ref,
            item,
        ) in values.items()
    )


def main() -> int:
    print("=" * 110)
    print(
        "BÀI 13B-11.15.2.4.9.2A.2 - "
        "XUẤT LẠI 6 BIỂU TOÀN TỈNH / "
        "KIỂM TRA 100%"
    )
    print("=" * 110)
    print()

    if not DB_PATH.exists():
        print(
            "[LỖI] Không tìm thấy: "
            f"{DB_PATH}"
        )
        return 1

    db_hash_before = (
        sha256_file(
            DB_PATH
        )
    )

    before = db_state()
    verify_db(
        before
    )

    old_cwd = Path.cwd()

    try:
        os.chdir(
            PROJECT
        )

        sys.path.insert(
            0,
            str(PROJECT),
        )

        print(
            "[1/7] Database: "
            f"{before['total']} xã/phường; "
            f"ĐBKK={before['special']}; "
            "integrity=ok."
        )

        verify_source_markers()

        print(
            "[2/7] Source: "
            "Bài 4.9.2 + 4.9.2A.1 "
            "ĐỦ marker."
        )

        from sqlalchemy import text
        from app.database import SessionLocal
        import app.pcgd_mn_report_builders_v1 as mn_builder
        import app.pcgd_xmc_report_builders_v1 as xmc_builder

        db = SessionLocal()

        try:
            db.execute(
                text(
                    "PRAGMA query_only = ON"
                )
            )

            school_year = (
                choose_year(
                    db
                )
            )

            record_count = (
                year_record_count(
                    db,
                    int(
                        school_year.id
                    ),
                )
            )

            print(
                "[3/7] Năm học: "
                f"{school_year.code} "
                f"(ID={school_year.id}); "
                f"bản ghi={record_count}."
            )

            request = make_request()

            OUT_DIR.mkdir(
                parents=True,
                exist_ok=False,
            )

            report_lines = [
                "=" * 110,
                "BÀI 13B-11.15.2.4.9.2A.2 - "
                "XUẤT LẠI 6 BIỂU TOÀN TỈNH",
                "=" * 110,
                "",
                f"Năm học: {school_year.code}",
                (
                    "Số bản ghi: "
                    f"{record_count}"
                ),
                (
                    "Format yêu cầu "
                    "MN/TH/THCS: "
                    f"{EXPECTED_PERCENT_FORMAT}"
                ),
                "XMC: giữ nguyên.",
                "",
            ]

            print(
                "[4/7] Xuất lại 6 biểu..."
            )

            for spec in REPORTS:
                exporter = (
                    mn_builder
                    .export_mn_report
                    if (
                        spec[
                            "exporter"
                        ]
                        == "mn"
                    )
                    else
                    xmc_builder
                    .export_additional_report
                )

                response = call_exporter(
                    exporter,
                    report_type=(
                        spec["code"]
                    ),
                    db=db,
                    request=request,
                    school_year=(
                        school_year
                    ),
                )

                data = asyncio.run(
                    response_to_bytes(
                        response
                    )
                )

                if len(data) < 1000:
                    raise RuntimeError(
                        f"{spec['short']}: "
                        "file quá nhỏ: "
                        f"{len(data)} bytes."
                    )

                file_path = (
                    OUT_DIR
                    / (
                        f"{spec['short']}_"
                        f"{school_year.code}_"
                        "TOAN_TINH_"
                        "KIEM_TRA_100_PERCENT.xlsx"
                    )
                )

                file_path.write_bytes(
                    data
                )

                (
                    sheet_name,
                    conclusions,
                    percentages,
                ) = inspect_xlsx(
                    data,
                    spec,
                )

                print(
                    "      [OK] "
                    f"{spec['short']}"
                )

                print(
                    "           Kết luận: "
                    + conclusion_text(
                        conclusions
                    )
                )

                if (
                    spec[
                        "check_percent_format"
                    ]
                ):
                    print(
                        "           Tỷ lệ   : "
                        + percent_text(
                            percentages
                        )
                    )
                else:
                    print(
                        "           XMC     : "
                        "GIỮ NGUYÊN | "
                        + percent_text(
                            percentages
                        )
                    )

                report_lines.extend(
                    [
                        "",
                        (
                            f"[{spec['short']}] "
                            f"sheet={sheet_name}"
                        ),
                        (
                            "Kết luận: "
                            + conclusion_text(
                                conclusions
                            )
                        ),
                        (
                            (
                                "Tỷ lệ: "
                                if spec[
                                    "check_percent_format"
                                ]
                                else (
                                    "XMC giữ nguyên: "
                                )
                            )
                            + percent_text(
                                percentages
                            )
                        ),
                        (
                            "File: "
                            + str(
                                file_path
                            )
                        ),
                    ]
                )

        finally:
            try:
                db.rollback()
            except Exception:
                pass

            db.close()

        print(
            "[5/7] Đã xuất đủ "
            "6 file Excel mới."
        )

        after = db_state()
        verify_db(
            after
        )

        db_hash_after = (
            sha256_file(
                DB_PATH
            )
        )

        if (
            before
            != after
        ):
            raise RuntimeError(
                "Trạng thái database "
                "trước/sau khác nhau."
            )

        if (
            db_hash_before
            != db_hash_after
        ):
            raise RuntimeError(
                "SHA256 phocap.db thay đổi. "
                "Hãy dừng Uvicorn và chạy lại."
            )

        print(
            "[6/7] Database: "
            "KHÔNG THAY ĐỔI; "
            "SHA256 giữ nguyên."
        )

        report_lines.extend(
            [
                "",
                "AN TOÀN:",
                (
                    "DB SHA256 trước: "
                    f"{db_hash_before}"
                ),
                (
                    "DB SHA256 sau  : "
                    f"{db_hash_after}"
                ),
                "Database: KHÔNG THAY ĐỔI.",
                "",
                "KẾT LUẬN:",
                (
                    " - MN/TH/THCS: "
                    "các ô tỷ lệ đã có "
                    r"number_format=0\%."
                ),
                (
                    " - XMC: không kiểm ép "
                    "format; giữ nguyên."
                ),
                (
                    " - 6 file đã xuất để "
                    "mở Excel kiểm tra trực quan."
                ),
            ]
        )

        REPORT.write_text(
            "\n".join(
                report_lines
            ),
            encoding="utf-8-sig",
        )

        print(
            f"[7/7] Báo cáo: "
            f"{REPORT}"
        )

        print()
        print("=" * 110)
        print(
            "BÀI 4.9.2A.2 HOÀN THÀNH."
        )
        print(
            "MN / TH / THCS: "
            "đã kiểm tra format 100%."
        )
        print(
            "XMC: GIỮ NGUYÊN."
        )
        print(
            f"6 file Excel: {OUT_DIR}"
        )
        print(
            "Database: KHÔNG THAY ĐỔI."
        )
        print("=" * 110)

        return 0

    except Exception as exc:
        print()
        print("=" * 110)
        print(
            "[LỖI] BÀI 4.9.2A.2 "
            "CHƯA ĐẠT"
        )
        print(
            str(exc)
        )
        print("=" * 110)

        try:
            OUT_DIR.mkdir(
                parents=True,
                exist_ok=True,
            )

            error_path = (
                OUT_DIR
                / (
                    "loi_kiem_tra_"
                    f"4_9_2a_2_{STAMP}.txt"
                )
            )

            error_path.write_text(
                (
                    f"Lỗi: {exc}\n\n"
                    + traceback.format_exc()
                ),
                encoding="utf-8-sig",
            )

            print(
                "Chi tiết lỗi: "
                f"{error_path}"
            )

        except Exception:
            pass

        print(
            "Script không có lệnh "
            "UPDATE/INSERT/DELETE "
            "và bật PRAGMA query_only."
        )

        return 1

    finally:
        os.chdir(
            old_cwd
        )


if __name__ == "__main__":
    raise SystemExit(
        main()
    )
