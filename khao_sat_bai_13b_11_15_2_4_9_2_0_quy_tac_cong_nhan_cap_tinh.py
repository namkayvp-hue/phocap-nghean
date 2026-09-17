# -*- coding: utf-8 -*-
"""
BÀI 13B-11.15.2.4.9.2.0
KHẢO SÁT / KHÓA NGUỒN QUY TẮC CÔNG NHẬN CẤP TỈNH

CHỈ ĐỌC:
- Không sửa source.
- Không sửa database.
- Không sửa Excel template.
- Không đổi menu/route.

MỤC TIÊU:
1. Xác nhận nền Bài 4.9.1 đã có.
2. Xác định trong pcgd_business_rules.py đã có hay chưa evaluator cấp tỉnh.
3. Trích đúng quy tắc province_recognition hiện có.
4. Đọc vùng ghi chú/quy tắc cấp tỉnh trong PCGD_2025_THCS_M2.xlsx/Sheet1.
5. Khóa các ô kết luận của 4 nhóm khi scope = Toàn tỉnh.
6. Khóa đường exporter hiện tại và guard Bài 4.9.1 chỉ chạy cấp xã.
7. Kiểm tra DB 130 xã/phường, 51 xã ĐBKK, integrity/FK.
8. Xuất báo cáo TXT để làm nền cho bộ cài Bài 4.9.2.
"""

from __future__ import annotations

import ast
import hashlib
import os
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any


PROJECT = Path(
    os.environ.get("PHOCAP_PROJECT", r"C:\PhoCap")
).resolve()

APP = PROJECT / "app"
DB = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"

RULES = APP / "services" / "pcgd_business_rules.py"
MN_BUILDER = APP / "pcgd_mn_report_builders_v1.py"
XMC_BUILDER = APP / "pcgd_xmc_report_builders_v1.py"

THCS_M2 = (
    APP
    / "report_templates"
    / "pcgd_xmc_2025"
    / "PCGD_2025_THCS_M2.xlsx"
)

MN02 = (
    APP
    / "report_templates"
    / "pcgd_mn_2025"
    / "PCGD_2025_MN_02.xlsx"
)

TH02 = (
    APP
    / "report_templates"
    / "pcgd_xmc_2025"
    / "PCGD_2025_TH_02.xlsx"
)

XMC4 = (
    APP
    / "report_templates"
    / "pcgd_xmc_2025"
    / "PCGD_2025_XMC_4.xlsx"
)

THCS_TK = (
    APP
    / "report_templates"
    / "pcgd_xmc_2025"
    / "PCGD_2025_THCS_TK.xlsx"
)

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT = (
    EXPORTS
    / (
        "bao_cao_khao_sat_bai_13b_11_15_2_4_9_2_0_"
        f"quy_tac_cong_nhan_cap_tinh_{STAMP}.txt"
    )
)

EXPECTED_COMMUNES = 130
EXPECTED_SPECIAL = 51


def read_text(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(
        encoding="utf-8-sig",
        errors="replace",
    )


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            h.update(chunk)
    return h.hexdigest()


def db_state() -> dict[str, Any]:
    if not DB.exists():
        raise RuntimeError(
            f"Không tìm thấy database: {DB}"
        )

    uri = DB.resolve().as_uri() + "?mode=ro"
    conn = sqlite3.connect(uri, uri=True)

    try:
        integrity = str(
            conn.execute(
                "PRAGMA integrity_check"
            ).fetchone()[0]
        )
        fk = len(
            conn.execute(
                "PRAGMA foreign_key_check"
            ).fetchall()
        )
        total = int(
            conn.execute(
                "SELECT COUNT(*) FROM communes"
            ).fetchone()[0]
        )
        special = int(
            conn.execute(
                """
                SELECT COUNT(*)
                FROM communes
                WHERE COALESCE(
                    is_special_difficulty_area,
                    0
                ) = 1
                """
            ).fetchone()[0]
        )

        years = []
        try:
            rows = conn.execute(
                """
                SELECT
                    sy.id,
                    sy.code,
                    COUNT(spr.id) AS record_count,
                    SUM(
                        CASE
                        WHEN COALESCE(
                            c.is_special_difficulty_area,
                            0
                        ) = 1
                        THEN 1 ELSE 0 END
                    ) AS special_records,
                    SUM(
                        CASE
                        WHEN COALESCE(
                            c.is_special_difficulty_area,
                            0
                        ) = 0
                        THEN 1 ELSE 0 END
                    ) AS normal_records
                FROM school_years sy
                LEFT JOIN survey_person_year_records spr
                  ON spr.school_year_id = sy.id
                LEFT JOIN survey_forms sf
                  ON sf.id = spr.survey_form_id
                LEFT JOIN survey_batches sb
                  ON sb.id = sf.survey_batch_id
                LEFT JOIN communes c
                  ON c.id = sb.commune_id
                GROUP BY sy.id, sy.code
                ORDER BY sy.code DESC
                """
            ).fetchall()

            years = [
                {
                    "id": int(row[0]),
                    "code": str(row[1]),
                    "records": int(row[2] or 0),
                    "special_records": int(
                        row[3] or 0
                    ),
                    "normal_records": int(
                        row[4] or 0
                    ),
                }
                for row in rows
            ]
        except Exception as exc:
            years = [
                {
                    "error": str(exc),
                }
            ]

        return {
            "integrity": integrity,
            "fk": fk,
            "total": total,
            "special": special,
            "years": years,
        }

    finally:
        conn.close()


def module_functions(
    source: str,
) -> list[str]:
    tree = ast.parse(source)
    return [
        node.name
        for node in tree.body
        if isinstance(
            node,
            (
                ast.FunctionDef,
                ast.AsyncFunctionDef,
            ),
        )
    ]


def top_level_assignment_source(
    source: str,
    name: str,
) -> str:
    tree = ast.parse(source)
    lines = source.splitlines()

    for node in tree.body:
        targets = []

        if isinstance(node, ast.Assign):
            targets = list(node.targets)
        elif isinstance(node, ast.AnnAssign):
            targets = [node.target]
        else:
            continue

        if any(
            isinstance(target, ast.Name)
            and target.id == name
            for target in targets
        ):
            start = int(node.lineno) - 1
            end = int(node.end_lineno)
            return "\n".join(
                lines[start:end]
            )

    return ""


def function_source(
    source: str,
    name: str,
) -> str:
    tree = ast.parse(source)
    lines = source.splitlines()

    matches = [
        node
        for node in tree.body
        if isinstance(
            node,
            (
                ast.FunctionDef,
                ast.AsyncFunctionDef,
            ),
        )
        and node.name == name
    ]

    if len(matches) != 1:
        return ""

    node = matches[0]
    return "\n".join(
        lines[
            int(node.lineno) - 1:
            int(node.end_lineno)
        ]
    )


def compact(
    text: str,
    max_lines: int = 120,
) -> list[str]:
    rows = []
    for line in text.splitlines():
        if line.strip():
            rows.append(line.rstrip())
        if len(rows) >= max_lines:
            rows.append("... [rút gọn]")
            break
    return rows


def inspect_thcs_m2_sheet1() -> list[str]:
    from openpyxl import load_workbook

    if not THCS_M2.exists():
        return [
            f"KHÔNG TÌM THẤY: {THCS_M2}"
        ]

    wb = load_workbook(
        THCS_M2,
        read_only=False,
        data_only=False,
    )

    try:
        if "Sheet1" not in wb.sheetnames:
            return [
                "Không có sheet Sheet1.",
                "Sheets: "
                + ", ".join(wb.sheetnames),
            ]

        ws = wb["Sheet1"]
        lines = []

        lines.append(
            "Vùng Sheet1 hàng 33-52, cột A-H:"
        )

        for row in range(33, 53):
            cells = []
            for col in range(1, 9):
                cell = ws.cell(
                    row=row,
                    column=col,
                )
                value = cell.value
                if value not in (None, ""):
                    cells.append(
                        f"{cell.coordinate}={value!r}"
                    )

            if cells:
                lines.append(
                    "  " + " | ".join(cells)
                )

        lines.append("")
        lines.append(
            "Merged ranges giao với hàng 33-52:"
        )

        hit = False
        for merged in ws.merged_cells.ranges:
            if (
                merged.max_row >= 33
                and merged.min_row <= 52
            ):
                hit = True
                lines.append(
                    "  " + str(merged)
                )

        if not hit:
            lines.append("  (không có)")

        return lines

    finally:
        wb.close()


def inspect_template(
    label: str,
    path: Path,
    cells: tuple[str, ...],
) -> list[str]:
    from openpyxl import load_workbook

    lines = [
        f"[{label}] {path}"
    ]

    if not path.exists():
        lines.append("  KHÔNG TÌM THẤY")
        return lines

    wb = load_workbook(
        path,
        read_only=False,
        data_only=False,
    )

    try:
        ws = wb[wb.sheetnames[0]]

        lines.append(
            f"  Sheet đầu: {ws.title}"
        )

        for coord in cells:
            lines.append(
                f"  {coord}: "
                f"{ws[coord].value!r}"
            )

        return lines

    finally:
        wb.close()


def source_hits(
    source: str,
    needles: tuple[str, ...],
) -> list[str]:
    result = []

    for lineno, line in enumerate(
        source.splitlines(),
        start=1,
    ):
        if any(
            needle in line
            for needle in needles
        ):
            result.append(
                f"L{lineno}: {line.strip()}"
            )

    return result


def main() -> int:
    print("=" * 108)
    print(
        "BÀI 13B-11.15.2.4.9.2.0 - "
        "KHẢO SÁT QUY TẮC CÔNG NHẬN CẤP TỈNH"
    )
    print("=" * 108)
    print()
    print(
        "CHỈ ĐỌC - KHÔNG SỬA SOURCE / DATABASE / EXCEL."
    )
    print()

    for path in (
        RULES,
        MN_BUILDER,
        XMC_BUILDER,
        DB,
    ):
        if not path.exists():
            raise RuntimeError(
                f"Không tìm thấy: {path}"
            )

    before_rules_hash = sha256(RULES)
    before_mn_hash = sha256(MN_BUILDER)
    before_xmc_hash = sha256(XMC_BUILDER)

    db = db_state()

    if db["integrity"].lower() != "ok":
        raise RuntimeError(
            "integrity_check != ok"
        )

    if db["fk"] != 0:
        raise RuntimeError(
            "foreign_key_check có lỗi."
        )

    if db["total"] != EXPECTED_COMMUNES:
        raise RuntimeError(
            f"Tổng communes={db['total']}, "
            f"không phải {EXPECTED_COMMUNES}."
        )

    if db["special"] != EXPECTED_SPECIAL:
        raise RuntimeError(
            f"ĐBKK={db['special']}, "
            f"không phải {EXPECTED_SPECIAL}."
        )

    rules = read_text(RULES)
    mn = read_text(MN_BUILDER)
    xmc = read_text(XMC_BUILDER)

    ast.parse(rules)
    ast.parse(mn)
    ast.parse(xmc)

    functions = module_functions(rules)

    province_evaluators = [
        name
        for name in functions
        if "province" in name.lower()
        or "tinh" in name.lower()
    ]

    commune_evaluators = [
        name
        for name in (
            "evaluate_preschool_3_5_commune",
            "evaluate_primary_commune",
            "evaluate_literacy_commune",
            "evaluate_thcs_commune",
        )
        if name in functions
    ]

    lines = [
        "=" * 108,
        "BÀI 13B-11.15.2.4.9.2.0 - "
        "BÁO CÁO KHẢO SÁT QUY TẮC CÔNG NHẬN CẤP TỈNH",
        "=" * 108,
        "",
        "I. DATABASE",
        f"integrity_check: {db['integrity']}",
        f"foreign_key_check: {db['fk']}",
        f"Tổng xã/phường: {db['total']}",
        f"Xã ĐBKK: {db['special']}",
        "",
        "Dữ liệu theo năm học:",
    ]

    for item in db["years"]:
        if "error" in item:
            lines.append(
                "  Không đếm được: "
                + item["error"]
            )
        else:
            lines.append(
                "  - "
                f"{item['code']} (ID={item['id']}): "
                f"records={item['records']}; "
                f"ĐBKK={item['special_records']}; "
                f"thường={item['normal_records']}"
            )

    lines.extend(
        [
            "",
            "II. EVALUATOR HIỆN CÓ",
            "Evaluator cấp xã/phường:",
        ]
    )

    for name in commune_evaluators:
        lines.append(
            "  - " + name
        )

    lines.append(
        "Evaluator/hàm tên gợi ý cấp tỉnh:"
    )

    if province_evaluators:
        for name in province_evaluators:
            lines.append(
                "  - " + name
            )
    else:
        lines.append(
            "  - KHÔNG THẤY evaluator cấp tỉnh."
        )

    lines.extend(
        [
            "",
            "III. QUY TẮC TOP-LEVEL TRONG SERVICE",
        ]
    )

    for name in (
        "PRIMARY_THRESHOLDS",
        "THCS_THRESHOLDS",
        "LITERACY_THRESHOLDS",
        "PRESCHOOL_3_5_ND277",
    ):
        block = top_level_assignment_source(
            rules,
            name,
        )

        lines.append("")
        lines.append(
            f"[{name}]"
        )

        if block:
            lines.extend(
                "  " + row
                for row in compact(
                    block,
                    160,
                )
            )
        else:
            lines.append(
                "  KHÔNG TÌM THẤY."
            )

    lines.extend(
        [
            "",
            "IV. DẤU VẾT TỪ KHÓA CẤP TỈNH TRONG SERVICE",
        ]
    )

    hits = source_hits(
        rules,
        (
            "province",
            "tỉnh",
            "Tinh",
            "recognized_min_rate",
            "national_roadmap",
        ),
    )

    if hits:
        lines.extend(
            "  " + item
            for item in hits
        )
    else:
        lines.append(
            "  Không có dấu vết."
        )

    lines.extend(
        [
            "",
            "V. MẪU THCS_M2 - VÙNG QUY TẮC CẤP TỈNH",
        ]
    )
    lines.extend(
        inspect_thcs_m2_sheet1()
    )

    lines.extend(
        [
            "",
            "VI. Ô KẾT LUẬN 4 NHÓM",
        ]
    )

    for block in (
        inspect_template(
            "MN-02",
            MN02,
            ("R3", "R7"),
        ),
        inspect_template(
            "TH-02",
            TH02,
            ("T4", "T8"),
        ),
        inspect_template(
            "XMC-4",
            XMC4,
            ("R5", "R8"),
        ),
        inspect_template(
            "THCS-TK",
            THCS_TK,
            (
                "F4",
                "G4",
                "R4",
                "F8",
                "G8",
                "R8",
            ),
        ),
    ):
        lines.extend(block)
        lines.append("")

    lines.extend(
        [
            "VII. NỀN BÀI 4.9.1 TRONG BUILDER",
            "",
            "[Mầm non]",
        ]
    )

    mn_hits = source_hits(
        mn,
        (
            "BAI_13B_11_15_2_4_9_1",
            "_b491_apply_mn_conclusion",
            'meta.get("kind")',
            "workbook.save(",
        ),
    )

    lines.extend(
        "  " + item
        for item in mn_hits
    )

    lines.extend(
        [
            "",
            "[TH / XMC / THCS]",
        ]
    )

    xmc_hits = source_hits(
        xmc,
        (
            "BAI_13B_11_15_2_4_9_1",
            "_b491_apply_xmc_conclusions",
            'meta.get("kind")',
            "workbook.save(",
        ),
    )

    lines.extend(
        "  " + item
        for item in xmc_hits
    )

    lines.extend(
        [
            "",
            "VIII. CÁC HÀM EXPORTER HIỆN TẠI",
        ]
    )

    for path_label, source, fn_name in (
        (
            "MN",
            mn,
            "export_mn_report",
        ),
        (
            "TH/XMC/THCS",
            xmc,
            "export_additional_report",
        ),
    ):
        fn = function_source(
            source,
            fn_name,
        )
        lines.append("")
        lines.append(
            f"[{path_label}] {fn_name}"
        )
        if fn:
            for row in compact(
                fn,
                180,
            ):
                if any(
                    token in row
                    for token in (
                        "selected_commune_id",
                        "selected_school_id",
                        "meta",
                        "province",
                        "workbook.save",
                        "_b491_",
                    )
                ):
                    lines.append(
                        "  " + row
                    )
        else:
            lines.append(
                "  KHÔNG TÌM THẤY."
            )

    after_rules_hash = sha256(RULES)
    after_mn_hash = sha256(MN_BUILDER)
    after_xmc_hash = sha256(XMC_BUILDER)

    if (
        before_rules_hash
        != after_rules_hash
        or before_mn_hash
        != after_mn_hash
        or before_xmc_hash
        != after_xmc_hash
    ):
        raise RuntimeError(
            "Source thay đổi trong lúc khảo sát."
        )

    lines.extend(
        [
            "",
            "IX. AN TOÀN",
            f"SHA256 rules : {after_rules_hash}",
            f"SHA256 MN    : {after_mn_hash}",
            f"SHA256 XMC   : {after_xmc_hash}",
            "Source: KHÔNG THAY ĐỔI.",
            "Database: chỉ mở mode=ro.",
            "Excel: chỉ đọc.",
            "",
            "X. KẾT LUẬN TỰ ĐỘNG",
            (
                " - Có evaluator cấp tỉnh: "
                + (
                    "CÓ"
                    if province_evaluators
                    else "CHƯA CÓ"
                )
            ),
            (
                " - Có đủ 4 evaluator cấp xã: "
                + (
                    "CÓ"
                    if len(
                        commune_evaluators
                    ) == 4
                    else "KHÔNG"
                )
            ),
            " - Bài 4.9.2 KHÔNG được tái dùng trực tiếp evaluator cấp xã cho Toàn tỉnh.",
            " - Báo cáo này dùng để khóa chính xác bộ quy tắc trước khi tạo bộ cài 4.9.2.",
            "",
            "CHỈ ĐỌC - KHÔNG SỬA SOURCE/DATABASE/EXCEL.",
        ]
    )

    EXPORTS.mkdir(
        parents=True,
        exist_ok=True,
    )

    OUT.write_text(
        "\n".join(lines),
        encoding="utf-8-sig",
    )

    print(
        f"[1/6] Database: "
        f"{db['total']} xã/phường; "
        f"ĐBKK={db['special']}; "
        "integrity=ok."
    )
    print(
        f"[2/6] Evaluator cấp xã: "
        f"{len(commune_evaluators)}/4."
    )
    print(
        "[3/6] Evaluator/hàm cấp tỉnh: "
        + (
            ", ".join(province_evaluators)
            if province_evaluators
            else "CHƯA CÓ"
        )
    )
    print(
        "[4/6] Đã đọc quy tắc Sheet1 "
        "của PCGD_2025_THCS_M2.xlsx."
    )
    print(
        "[5/6] Source/DB/Excel: KHÔNG THAY ĐỔI."
    )
    print(
        f"[6/6] Báo cáo: {OUT}"
    )
    print()
    print("=" * 108)
    print(
        "BÀI 4.9.2.0 KHẢO SÁT HOÀN THÀNH."
    )
    print(
        "Gửi file TXT báo cáo cho ChatGPT "
        "để tạo BỘ CÀI DUY NHẤT Bài 4.9.2."
    )
    print("=" * 108)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
