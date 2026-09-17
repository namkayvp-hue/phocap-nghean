from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from math import ceil
from pathlib import Path
from typing import Any
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Commune, SchoolYear
from app.permissions import is_province_reader
from app.routers.report_center import (
    _choose_year_id,
    _clean_text,
    _load_school_years,
    _load_scope_options,
    _parse_optional_int,
)
from app.routers.survey_summary_report import (
    DATA_STATE_LABELS,
    PAGE_SIZE_OPTIONS,
    _load_batches,
    _load_person_rows,
    _reference_date,
)
from app.routers.surveys import (
    DISABILITY_LEVEL_LABELS,
    DISABILITY_STATUS_LABELS,
    DISABILITY_TYPE_LABELS,
    LEARNING_STATUS_LABELS,
    RESIDENCY_STATUS_LABELS,
    lay_thong_tin_nguoi_dung,
    tinh_tuoi_tai_ngay,
)
from app.survey_models import SurveyPersonYearRecord


APP_DIR = Path(__file__).resolve().parent.parent
templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

router = APIRouter(
    prefix="/bao-cao/pho-cap-3-5-tuoi",
    tags=["Báo cáo phổ cập GDMN trẻ em từ 3 đến 5 tuổi"],
)

AGE_ORDER = (3, 4, 5)
AGE_LABELS = {3: "3 tuổi", 4: "4 tuổi", 5: "5 tuổi"}
ATTENDING_CODES = {"DANG_HOC", "CHUYEN_DEN", "TAM_NGHI"}
EXCLUDED_LEARNING_CODES = {"KHONG_THUOC_DIEN", "CHUYEN_DI"}

ISSUE_GROUP_LABELS = {
    "": "Tất cả trẻ 3-5 tuổi",
    "CAN_KIEM_TRA": "Tất cả trường hợp cần kiểm tra",
    "CHUA_RA_LOP": "Chưa ra lớp",
    "HOC_NGOAI_DIA_BAN": "Học ngoài địa bàn",
    "CHUA_XAC_DINH_TRUONG": "Đang học nhưng chưa xác định trường",
    "CHUA_HOAN_THANH_THEO_TUOI": "Chưa xác nhận hoàn thành theo độ tuổi",
    "KHUYET_TAT_CHUA_HOA_NHAP": "Khuyết tật chưa xác nhận học hòa nhập",
    "THIEU_SO_DINH_DANH": "Thiếu số định danh",
    "THIEU_DU_LIEU_NAM_HOC": "Thiếu dữ liệu năm học",
}

GENDER_LABELS = {
    "NAM": "Nam",
    "NU": "Nữ",
    "KHAC": "Khác",
    "CHUA_XAC_DINH": "Chưa xác định",
}


def _safe_date(value: date | None) -> str:
    return value.strftime("%d/%m/%Y") if value is not None else ""


def _normalize_gender(value: Any) -> tuple[str, str]:
    text = _clean_text(value).upper()
    if text in {"NAM", "MALE", "M"}:
        return "NAM", "Nam"
    if text in {"NỮ", "NU", "FEMALE", "F"}:
        return "NU", "Nữ"
    if text:
        return "KHAC", _clean_text(value)
    return "CHUA_XAC_DINH", "Chưa xác định"


def _percentage(numerator: int, denominator: int) -> float:
    if denominator <= 0:
        return 0.0
    return round(numerator * 100.0 / denominator, 2)


def _thresholds(is_special: bool) -> tuple[float, float]:
    # Điều 12 Nghị định 277/2025/NĐ-CP.
    return (85.0, 80.0) if is_special else (90.0, 85.0)


def _new_age_row(age: int) -> dict[str, Any]:
    return {
        "age": age,
        "label": AGE_LABELS.get(age, str(age)),
        "total": 0,
        "must_mobilize": 0,
        "attending": 0,
        "not_in_school": 0,
        "completed": 0,
        "not_completed": 0,
        "completion_unknown": 0,
        "disability": 0,
        "inclusive": 0,
    }


def _new_commune_row(batch: Any, count_item: dict[str, int]) -> dict[str, Any]:
    is_special = bool(
        getattr(batch.commune, "is_special_difficulty_area", False)
    )
    mobilization_threshold, completion_threshold = _thresholds(is_special)
    return {
        "commune_id": batch.commune_id,
        "commune_code": batch.commune.code,
        "commune_name": batch.commune.name,
        "is_special_difficulty_area": is_special,
        "area_type_label": (
            "KTXH đặc biệt khó khăn" if is_special else "Địa bàn thông thường"
        ),
        "mobilization_threshold": mobilization_threshold,
        "completion_threshold": completion_threshold,
        "batch_total": 1,
        "locked_batch_total": 1 if batch.is_locked else 0,
        "forms": int(count_item.get("forms", 0)),
        "households": int(count_item.get("households", 0)),
        "total_3_5": 0,
        "must_mobilize": 0,
        "attending": 0,
        "not_in_school": 0,
        "in_area": 0,
        "outside_area": 0,
        "unknown_school": 0,
        "completed": 0,
        "not_completed": 0,
        "completion_unknown": 0,
        "disability": 0,
        "inclusive": 0,
        "missing_personal_id": 0,
        "missing_year_record": 0,
        "needs_attention": 0,
        "age_rows": {age: _new_age_row(age) for age in AGE_ORDER},
    }


def _issue_flags(item: dict[str, Any]) -> list[str]:
    flags: list[str] = []
    if item["missing_year_record"]:
        flags.append("THIEU_DU_LIEU_NAM_HOC")
    if item["missing_personal_id"]:
        flags.append("THIEU_SO_DINH_DANH")
    if item["not_in_school"]:
        flags.append("CHUA_RA_LOP")
    if item["outside_area"]:
        flags.append("HOC_NGOAI_DIA_BAN")
    if item["unknown_school"]:
        flags.append("CHUA_XAC_DINH_TRUONG")
    if item["attending"] and item["completion_unknown"]:
        flags.append("CHUA_HOAN_THANH_THEO_TUOI")
    if item["disability"] and not item["inclusive"]:
        flags.append("KHUYET_TAT_CHUA_HOA_NHAP")
    return flags


def _matches_issue_group(item: dict[str, Any], issue_group: str) -> bool:
    if not issue_group:
        return True
    if issue_group == "CAN_KIEM_TRA":
        return bool(item["issue_flags"])
    return issue_group in item["issue_flags"]


def _attention_reason(item: dict[str, Any]) -> str:
    return "; ".join(
        ISSUE_GROUP_LABELS.get(code, code)
        for code in item["issue_flags"]
    )


def _load_completion_by_age(
    db: Session,
    record_ids: list[int],
) -> dict[int, bool | None]:
    if not record_ids:
        return {}
    rows = db.execute(
        select(
            SurveyPersonYearRecord.id,
            SurveyPersonYearRecord.completed_preschool_by_age,
        ).where(SurveyPersonYearRecord.id.in_(record_ids))
    ).all()
    return {int(record_id): value for record_id, value in rows}


def _finalize_age_row(item: dict[str, Any]) -> None:
    item["mobilization_rate"] = _percentage(
        item["attending"], item["must_mobilize"]
    )
    item["completion_rate"] = _percentage(
        item["completed"], item["attending"]
    )
    item["inclusive_rate"] = _percentage(
        item["inclusive"], item["disability"]
    )


def _finalize_commune(item: dict[str, Any]) -> None:
    item["mobilization_rate"] = _percentage(
        item["attending"], item["must_mobilize"]
    )
    item["completion_rate"] = _percentage(
        item["completed"], item["attending"]
    )
    for age_item in item["age_rows"].values():
        _finalize_age_row(age_item)

    if item["must_mobilize"] <= 0:
        item["child_indicator_code"] = "NO_DATA"
        item["child_indicator_label"] = "Chưa có dữ liệu trẻ 3-5"
        item["child_indicator_class"] = "neutral"
    elif item["completion_unknown"] > 0 or item["missing_year_record"] > 0:
        item["child_indicator_code"] = "INCOMPLETE"
        item["child_indicator_label"] = "Chưa đủ dữ liệu kết luận"
        item["child_indicator_class"] = "warn"
    elif (
        item["mobilization_rate"] >= item["mobilization_threshold"]
        and item["completion_rate"] >= item["completion_threshold"]
    ):
        item["child_indicator_code"] = "PASS"
        item["child_indicator_label"] = "Đạt chỉ tiêu trẻ em"
        item["child_indicator_class"] = "good"
    else:
        item["child_indicator_code"] = "FAIL"
        item["child_indicator_label"] = "Chưa đạt chỉ tiêu trẻ em"
        item["child_indicator_class"] = "danger"


def build_preschool_3_5_report_data(
    *,
    db: Session,
    request: Request,
    school_year_id: int | None,
    commune_id: int | None,
    school_id: int | None,
    data_state: str,
    q: str,
) -> dict[str, Any]:
    school_year = db.get(SchoolYear, school_year_id) if school_year_id else None
    reference_date = _reference_date(school_year)
    batches, counts = _load_batches(
        db=db,
        request=request,
        school_year_id=school_year_id,
        commune_id=commune_id,
        school_id=school_id,
        data_state=data_state,
        q=q,
    )
    raw_rows = _load_person_rows(
        db=db,
        request=request,
        batch_ids=[item.id for item in batches],
        school_year_id=school_year_id,
        school_id=school_id,
    )
    completion_map = _load_completion_by_age(
        db,
        [
            int(row["year_record_id"])
            for row in raw_rows
            if row.get("year_record_id") is not None
        ],
    )

    commune_map = {
        item.id: item
        for item in db.scalars(select(Commune)).all()
    }
    batch_by_id = {item.id: item for item in batches}
    commune_grouped: dict[int, dict[str, Any]] = {}
    for batch in batches:
        existing = commune_grouped.get(batch.commune_id)
        if existing is None:
            commune_grouped[batch.commune_id] = _new_commune_row(
                batch,
                counts.get(batch.id, {}),
            )
        else:
            existing["batch_total"] += 1
            existing["locked_batch_total"] += 1 if batch.is_locked else 0
            existing["forms"] += int(counts.get(batch.id, {}).get("forms", 0))
            existing["households"] += int(
                counts.get(batch.id, {}).get("households", 0)
            )

    summary = {
        "batches": len(batches),
        "with_data": sum(
            1 for item in batches if counts.get(item.id, {}).get("forms", 0) > 0
        ),
        "locked": sum(1 for item in batches if item.is_locked),
        "forms": sum(counts.get(item.id, {}).get("forms", 0) for item in batches),
        "households": sum(
            counts.get(item.id, {}).get("households", 0) for item in batches
        ),
        "total_3_5": 0,
        "must_mobilize": 0,
        "attending": 0,
        "not_in_school": 0,
        "in_area": 0,
        "outside_area": 0,
        "unknown_school": 0,
        "completed": 0,
        "not_completed": 0,
        "completion_unknown": 0,
        "disability": 0,
        "inclusive": 0,
        "missing_personal_id": 0,
        "missing_year_record": 0,
        "needs_attention": 0,
        "age_rows": {age: _new_age_row(age) for age in AGE_ORDER},
    }

    school_grouped: dict[str, dict[str, Any]] = {}
    details: list[dict[str, Any]] = []
    gender_counts = {code: 0 for code in GENDER_LABELS}
    residency_counts: defaultdict[str, int] = defaultdict(int)

    for row in raw_rows:
        age = tinh_tuoi_tai_ngay(row.get("date_of_birth"), reference_date)
        if age not in AGE_ORDER:
            continue
        batch = batch_by_id.get(int(row["batch_id"]))
        if batch is None:
            continue

        gender_code, gender_label = _normalize_gender(row.get("gender"))
        learning_code = (
            str(row.get("learning_status") or "")
            if row.get("year_record_id") is not None
            else ""
        )
        residency_code = str(row.get("residency_status") or "CHUA_XAC_DINH")
        must_mobilize = (
            residency_code != "CHUYEN_DI"
            and learning_code not in EXCLUDED_LEARNING_CODES
        )
        attending = must_mobilize and learning_code in ATTENDING_CODES
        missing_year_record = row.get("year_record_id") is None
        missing_personal_id = not bool(_clean_text(row.get("personal_id")))
        school_commune_id = _parse_optional_int(row.get("school_commune_id"))
        in_area = attending and school_commune_id == batch.commune_id
        outside_area = (
            attending
            and school_commune_id is not None
            and school_commune_id != batch.commune_id
        )
        unknown_school = attending and school_commune_id is None

        completed_value = completion_map.get(
            int(row["year_record_id"])
            if row.get("year_record_id") is not None
            else -1
        )
        # Dữ liệu cũ chỉ được dùng thay thế cho trẻ đúng 5 tuổi.
        if completed_value is None and age == 5:
            completed_value = row.get("completed_preschool_5")
        completed = attending and completed_value is True
        not_completed = attending and completed_value is False
        completion_unknown = attending and completed_value is None

        disability = row.get("disability_status") == "CO_KHUYET_TAT"
        inclusive = disability and row.get("inclusive_education") is True
        not_in_school = must_mobilize and not attending
        school_display = _clean_text(
            row.get("school_name_reported") or row.get("school_name")
        )
        class_display = _clean_text(
            row.get("class_name_reported") or row.get("class_name")
        )
        school_commune = commune_map.get(school_commune_id)
        school_commune_name = (
            school_commune.name if school_commune is not None else ""
        )

        item = {
            **row,
            "age": age,
            "age_label": AGE_LABELS[age],
            "date_of_birth_display": _safe_date(row.get("date_of_birth")),
            "gender_code": gender_code,
            "gender_label": gender_label,
            "residency_status_label": RESIDENCY_STATUS_LABELS.get(
                residency_code,
                residency_code or "Chưa xác định",
            ),
            "learning_status_code": learning_code,
            "learning_status_label": (
                LEARNING_STATUS_LABELS.get(learning_code, learning_code)
                if learning_code
                else "Chưa có dữ liệu năm học"
            ),
            "school_display": school_display,
            "class_display": class_display,
            "school_commune_name": school_commune_name,
            "must_mobilize": must_mobilize,
            "attending": attending,
            "not_in_school": not_in_school,
            "in_area": in_area,
            "outside_area": outside_area,
            "unknown_school": unknown_school,
            "completed_value": completed_value,
            "completed": completed,
            "not_completed": not_completed,
            "completion_unknown": completion_unknown,
            "disability": disability,
            "inclusive": inclusive,
            "missing_personal_id": missing_personal_id,
            "missing_year_record": missing_year_record,
            "disability_status_label": DISABILITY_STATUS_LABELS.get(
                row.get("disability_status"),
                row.get("disability_status") or "Chưa xác định",
            ),
            "disability_type_label": DISABILITY_TYPE_LABELS.get(
                row.get("disability_type"),
                "",
            ),
            "disability_level_label": DISABILITY_LEVEL_LABELS.get(
                row.get("disability_level"),
                "",
            ),
        }
        item["issue_flags"] = _issue_flags(item)
        item["attention_reason"] = _attention_reason(item)
        details.append(item)

        summary["total_3_5"] += 1
        summary["must_mobilize"] += 1 if must_mobilize else 0
        for key in (
            "attending", "not_in_school", "in_area", "outside_area",
            "unknown_school", "completed", "not_completed",
            "completion_unknown", "disability", "inclusive",
            "missing_personal_id", "missing_year_record",
        ):
            summary[key] += 1 if item[key] else 0
        summary["needs_attention"] += 1 if item["issue_flags"] else 0
        age_summary = summary["age_rows"][age]
        age_summary["total"] += 1
        age_summary["must_mobilize"] += 1 if must_mobilize else 0
        for key in (
            "attending", "not_in_school", "completed", "not_completed",
            "completion_unknown", "disability", "inclusive",
        ):
            age_summary[key] += 1 if item[key] else 0
        gender_counts[gender_code] += 1
        residency_counts[residency_code] += 1

        commune_item = commune_grouped[batch.commune_id]
        commune_item["total_3_5"] += 1
        commune_item["must_mobilize"] += 1 if must_mobilize else 0
        for key in (
            "attending", "not_in_school", "in_area", "outside_area",
            "unknown_school", "completed", "not_completed",
            "completion_unknown", "disability", "inclusive",
            "missing_personal_id", "missing_year_record",
        ):
            commune_item[key] += 1 if item[key] else 0
        commune_item["needs_attention"] += 1 if item["issue_flags"] else 0
        commune_age = commune_item["age_rows"][age]
        commune_age["total"] += 1
        commune_age["must_mobilize"] += 1 if must_mobilize else 0
        for key in (
            "attending", "not_in_school", "completed", "not_completed",
            "completion_unknown", "disability", "inclusive",
        ):
            commune_age[key] += 1 if item[key] else 0

        if attending and school_display:
            school_key = (
                f"ID:{row.get('school_id')}"
                if row.get("school_id") is not None
                else f"REPORTED:{school_display.upper()}"
            )
            school_item = school_grouped.setdefault(
                school_key,
                {
                    "school_id": row.get("school_id"),
                    "school_code": row.get("school_code") or "",
                    "school_name": school_display,
                    "school_commune_name": school_commune_name,
                    "total_3_5": 0,
                    "age_3": 0,
                    "age_4": 0,
                    "age_5": 0,
                    "completed": 0,
                    "completion_unknown": 0,
                    "disability": 0,
                    "inclusive": 0,
                    "home_communes": set(),
                },
            )
            school_item["total_3_5"] += 1
            school_item[f"age_{age}"] += 1
            school_item["completed"] += 1 if completed else 0
            school_item["completion_unknown"] += 1 if completion_unknown else 0
            school_item["disability"] += 1 if disability else 0
            school_item["inclusive"] += 1 if inclusive else 0
            school_item["home_communes"].add(row.get("commune_name") or "")

    summary["mobilization_rate"] = _percentage(
        summary["attending"], summary["must_mobilize"]
    )
    summary["completion_rate"] = _percentage(
        summary["completed"], summary["attending"]
    )
    summary["inclusive_rate"] = _percentage(
        summary["inclusive"], summary["disability"]
    )
    for age_item in summary["age_rows"].values():
        _finalize_age_row(age_item)

    summary["group_3_4"] = _new_age_row(34)
    summary["group_3_4"]["label"] = "3-4 tuổi"
    for age in (3, 4):
        for key in (
            "total", "must_mobilize", "attending", "not_in_school",
            "completed", "not_completed", "completion_unknown",
            "disability", "inclusive",
        ):
            summary["group_3_4"][key] += summary["age_rows"][age][key]
    _finalize_age_row(summary["group_3_4"])

    commune_rows = sorted(
        commune_grouped.values(),
        key=lambda item: (item["commune_name"], item["commune_code"]),
    )
    for item in commune_rows:
        _finalize_commune(item)

    school_rows: list[dict[str, Any]] = []
    for item in school_grouped.values():
        school_row = dict(item)
        school_row["home_communes"] = ", ".join(
            sorted(x for x in item["home_communes"] if x)
        )
        school_row["completion_rate"] = _percentage(
            school_row["completed"], school_row["total_3_5"]
        )
        school_rows.append(school_row)
    school_rows.sort(key=lambda item: item["school_name"])

    details.sort(
        key=lambda item: (
            item.get("commune_name") or "",
            item.get("hamlet_name") or "",
            item.get("full_name") or "",
            item.get("person_id") or 0,
        )
    )
    attention_rows = [item for item in details if item["issue_flags"]]

    gender_rows = [
        {"code": code, "label": label, "total": gender_counts.get(code, 0)}
        for code, label in GENDER_LABELS.items()
        if gender_counts.get(code, 0) > 0
    ]
    residency_rows = [
        {
            "code": code,
            "label": RESIDENCY_STATUS_LABELS.get(code, code),
            "total": total,
        }
        for code, total in sorted(residency_counts.items())
        if total > 0
    ]

    regular_rows = [
        item for item in commune_rows
        if not item["is_special_difficulty_area"]
        and item["child_indicator_code"] in {"PASS", "FAIL"}
    ]
    special_rows = [
        item for item in commune_rows
        if item["is_special_difficulty_area"]
        and item["child_indicator_code"] in {"PASS", "FAIL"}
    ]
    province_reference = {
        "regular_evaluable": len(regular_rows),
        "regular_pass": sum(1 for item in regular_rows if item["child_indicator_code"] == "PASS"),
        "regular_rate": _percentage(
            sum(1 for item in regular_rows if item["child_indicator_code"] == "PASS"),
            len(regular_rows),
        ),
        "special_evaluable": len(special_rows),
        "special_pass": sum(1 for item in special_rows if item["child_indicator_code"] == "PASS"),
        "special_rate": _percentage(
            sum(1 for item in special_rows if item["child_indicator_code"] == "PASS"),
            len(special_rows),
        ),
    }

    return {
        "school_year": school_year,
        "reference_date": reference_date,
        "summary": summary,
        "commune_rows": commune_rows,
        "school_rows": school_rows,
        "person_rows": details,
        "attention_rows": attention_rows,
        "gender_rows": gender_rows,
        "residency_rows": residency_rows,
        "province_reference": province_reference,
        "is_official": bool(batches) and all(item.is_locked for item in batches),
    }


def _build_url(
    *,
    school_year_id: int | None,
    commune_id: int | None,
    school_id: int | None,
    data_state: str,
    issue_group: str,
    q: str,
    page_size: int,
    page: int,
) -> str:
    return "/bao-cao/pho-cap-3-5-tuoi?" + urlencode(
        {
            "school_year_id": school_year_id or "",
            "commune_id": commune_id or "",
            "school_id": school_id or "",
            "data_state": data_state,
            "issue_group": issue_group,
            "q": q,
            "page_size": page_size,
            "page": page,
        }
    )


def _style_title(worksheet: Any, title: str, subtitle: str, last_column: int) -> None:
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    worksheet.cell(1, 1, title)
    worksheet.cell(1, 1).font = Font(bold=True, color="FFFFFF", size=15)
    worksheet.cell(1, 1).fill = PatternFill("solid", fgColor="1F4E78")
    worksheet.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 30
    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    worksheet.cell(2, 1, subtitle)
    worksheet.cell(2, 1).font = Font(italic=True, color="48657D")
    worksheet.cell(2, 1).alignment = Alignment(horizontal="left", vertical="center")


def _style_table(
    worksheet: Any,
    *,
    header_row: int,
    widths: dict[int, float],
    percent_columns: set[int] | None = None,
) -> None:
    thin = Side(style="thin", color="D4DFEA")
    for cell in worksheet[header_row]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in worksheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for column, width in widths.items():
        worksheet.column_dimensions[get_column_letter(column)].width = width
    if percent_columns:
        for column in percent_columns:
            for row in range(header_row + 1, worksheet.max_row + 1):
                worksheet.cell(row, column).number_format = '0.00"%"'
    worksheet.freeze_panes = worksheet.cell(header_row + 1, 1)
    worksheet.auto_filter.ref = worksheet.dimensions


def _append_overview(
    workbook: Workbook,
    *,
    data: dict[str, Any],
    scope_label: str,
    user: dict[str, Any],
) -> None:
    ws = workbook.active
    ws.title = "Tong quan ND277"
    _style_title(
        ws,
        "BÁO CÁO CHỈ TIÊU TRẺ EM TỪ 3 ĐẾN 5 TUỔI - NGHỊ ĐỊNH 277/2025/NĐ-CP",
        f"Phạm vi: {scope_label} | Năm học: {data['school_year'].code if data.get('school_year') else ''} | Tham chiếu tuổi: {_safe_date(data['reference_date'])}",
        5,
    )
    ws.append([])
    ws.append(["Chỉ tiêu", "3 tuổi", "4 tuổi", "5 tuổi", "Tổng 3-5 tuổi"])
    age = data["summary"]["age_rows"]
    rows = [
        ("Tổng số trẻ", age[3]["total"], age[4]["total"], age[5]["total"], data["summary"]["total_3_5"]),
        ("Số trẻ phải huy động", age[3]["must_mobilize"], age[4]["must_mobilize"], age[5]["must_mobilize"], data["summary"]["must_mobilize"]),
        ("Số trẻ đến lớp", age[3]["attending"], age[4]["attending"], age[5]["attending"], data["summary"]["attending"]),
        ("Tỷ lệ huy động (%)", age[3]["mobilization_rate"], age[4]["mobilization_rate"], age[5]["mobilization_rate"], data["summary"]["mobilization_rate"]),
        ("Hoàn thành CT GDMN theo độ tuổi", age[3]["completed"], age[4]["completed"], age[5]["completed"], data["summary"]["completed"]),
        ("Tỷ lệ hoàn thành (%)", age[3]["completion_rate"], age[4]["completion_rate"], age[5]["completion_rate"], data["summary"]["completion_rate"]),
        ("Khuyết tật", age[3]["disability"], age[4]["disability"], age[5]["disability"], data["summary"]["disability"]),
        ("Học hòa nhập", age[3]["inclusive"], age[4]["inclusive"], age[5]["inclusive"], data["summary"]["inclusive"]),
        ("Cần kiểm tra", "", "", "", data["summary"]["needs_attention"]),
    ]
    for row in rows:
        ws.append(list(row))
    _style_table(ws, header_row=4, widths={1: 42, 2: 16, 3: 16, 4: 16, 5: 20})
    for row in (8, 10):
        for col in range(2, 6):
            ws.cell(row, col).number_format = '0.00"%"'
    ws.append([])
    ws.append(["Ghi chú pháp lý"])
    ws.append(["Ngưỡng cấp xã: địa bàn thông thường huy động >=90%, hoàn thành >=85%; địa bàn KTXH đặc biệt khó khăn huy động >=85%, hoàn thành >=80% (Điều 12 Nghị định 277/2025/NĐ-CP)."])
    ws.append(["Kết quả 'Đạt chỉ tiêu trẻ em' không thay thế kết luận xã đạt chuẩn; việc công nhận còn phải kiểm tra chương trình giáo dục, đội ngũ và điều kiện bảo đảm theo Điều 12-13."])
    ws.append([f"Người lập: {user.get('full_name', '')} - {user.get('role_name', '')} - {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    ws.merge_cells(start_row=15, start_column=1, end_row=15, end_column=5)
    ws.merge_cells(start_row=16, start_column=1, end_row=16, end_column=5)
    ws.merge_cells(start_row=17, start_column=1, end_row=17, end_column=5)
    ws.merge_cells(start_row=18, start_column=1, end_row=18, end_column=5)
    for row in range(15, 19):
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")


def _append_mn01(workbook: Workbook, data: dict[str, Any], scope_label: str) -> None:
    ws = workbook.create_sheet("MN-01 TE")
    _style_title(
        ws,
        "THỐNG KÊ TRẺ EM MẦM NON THEO ĐỘ TUỔI",
        f"Phạm vi: {scope_label} | Năm học: {data['school_year'].code if data.get('school_year') else ''}",
        5,
    )
    ws.append([])
    ws.append(["Chỉ tiêu", "3 tuổi", "4 tuổi", "5 tuổi", "Tổng 3-5 tuổi"])
    age = data["summary"]["age_rows"]
    mn_rows = [
        ["Tổng số trẻ trong độ tuổi", age[3]["total"], age[4]["total"], age[5]["total"], data["summary"]["total_3_5"]],
        ["Số trẻ phải huy động", age[3]["must_mobilize"], age[4]["must_mobilize"], age[5]["must_mobilize"], data["summary"]["must_mobilize"]],
        ["Số trẻ đến trường, nhóm, lớp", age[3]["attending"], age[4]["attending"], age[5]["attending"], data["summary"]["attending"]],
        ["Tỷ lệ huy động (%)", age[3]["mobilization_rate"], age[4]["mobilization_rate"], age[5]["mobilization_rate"], data["summary"]["mobilization_rate"]],
        ["Số trẻ chưa ra lớp", age[3]["not_in_school"], age[4]["not_in_school"], age[5]["not_in_school"], data["summary"]["not_in_school"]],
        ["Số trẻ hoàn thành CT GDMN theo độ tuổi", age[3]["completed"], age[4]["completed"], age[5]["completed"], data["summary"]["completed"]],
        ["Tỷ lệ hoàn thành CT GDMN (%)", age[3]["completion_rate"], age[4]["completion_rate"], age[5]["completion_rate"], data["summary"]["completion_rate"]],
        ["Trẻ khuyết tật", age[3]["disability"], age[4]["disability"], age[5]["disability"], data["summary"]["disability"]],
        ["Trẻ khuyết tật học hòa nhập", age[3]["inclusive"], age[4]["inclusive"], age[5]["inclusive"], data["summary"]["inclusive"]],
    ]
    for row in mn_rows:
        ws.append(row)
    _style_table(ws, header_row=4, widths={1: 48, 2: 16, 3: 16, 4: 16, 5: 20})
    for row in (8, 11):
        for col in range(2, 6):
            ws.cell(row, col).number_format = '0.00"%"'
    ws.append([])
    ws.append(["Nhóm chỉ tiêu", "Số trẻ phải huy động", "Số trẻ đến lớp", "Tỷ lệ huy động (%)", "Tỷ lệ hoàn thành theo độ tuổi (%)"])
    g34 = data["summary"]["group_3_4"]
    a5 = data["summary"]["age_rows"][5]
    ws.append(["Trẻ 3-4 tuổi", g34["must_mobilize"], g34["attending"], g34["mobilization_rate"], g34["completion_rate"]])
    ws.append(["Trẻ 5 tuổi", a5["must_mobilize"], a5["attending"], a5["mobilization_rate"], a5["completion_rate"]])
    _style_table(ws, header_row=15, widths={1: 48, 2: 22, 3: 20, 4: 22, 5: 30})
    for row in (16, 17):
        for col in (4, 5):
            ws.cell(row, col).number_format = '0.00"%"'


def _append_mn02(workbook: Workbook, data: dict[str, Any]) -> None:
    ws = workbook.create_sheet("MN-02")
    headers = [
        "STT", "Mã xã", "Xã/phường", "Loại địa bàn", "Đợt", "Đã khóa",
        "Trẻ 3-5", "Phải huy động", "Đến lớp", "Tỷ lệ huy động (%)",
        "Hoàn thành theo tuổi", "Tỷ lệ hoàn thành (%)", "Khuyết tật",
        "Hòa nhập", "Ngưỡng huy động", "Ngưỡng hoàn thành",
        "Kết quả chỉ tiêu trẻ em", "Cần kiểm tra", "Ghi chú",
    ]
    ws.append(headers)
    for index, item in enumerate(data["commune_rows"], start=1):
        ws.append([
            index, item["commune_code"], item["commune_name"], item["area_type_label"],
            item["batch_total"], item["locked_batch_total"], item["total_3_5"],
            item["must_mobilize"], item["attending"], item["mobilization_rate"],
            item["completed"], item["completion_rate"], item["disability"],
            item["inclusive"], item["mobilization_threshold"], item["completion_threshold"],
            item["child_indicator_label"], item["needs_attention"],
            "Chỉ kết luận chỉ tiêu trẻ em; chưa phải kết luận đạt chuẩn toàn diện.",
        ])
    _style_table(
        ws,
        header_row=1,
        widths={1: 7, 2: 12, 3: 28, 4: 28, 5: 9, 6: 10, 7: 12, 8: 16, 9: 12, 10: 18, 11: 22, 12: 20, 13: 12, 14: 12, 15: 18, 16: 20, 17: 30, 18: 14, 19: 52},
    )
    for row in range(2, ws.max_row + 1):
        for col in (10, 12, 15, 16):
            ws.cell(row, col).number_format = '0.00"%"'


def _append_schools(workbook: Workbook, data: dict[str, Any]) -> None:
    ws = workbook.create_sheet("Theo truong")
    ws.append([
        "STT", "Mã trường", "Trường", "Địa bàn trường", "Địa bàn cư trú",
        "Trẻ 3 tuổi", "Trẻ 4 tuổi", "Trẻ 5 tuổi", "Tổng trẻ 3-5 đang học",
        "Hoàn thành theo tuổi", "Chưa xác định hoàn thành", "Tỷ lệ hoàn thành (%)",
        "Khuyết tật", "Hòa nhập",
    ])
    for index, item in enumerate(data["school_rows"], start=1):
        ws.append([
            index, item["school_code"], item["school_name"], item["school_commune_name"],
            item["home_communes"], item["age_3"], item["age_4"], item["age_5"],
            item["total_3_5"], item["completed"], item["completion_unknown"],
            item["completion_rate"], item["disability"], item["inclusive"],
        ])
    _style_table(ws, header_row=1, widths={1: 7, 2: 16, 3: 34, 4: 26, 5: 38, 6: 13, 7: 13, 8: 13, 9: 22, 10: 22, 11: 24, 12: 20, 13: 12, 14: 12})
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 12).number_format = '0.00"%"'


def _detail_headers() -> list[str]:
    return [
        "STT", "Xã/phường", "Mã đợt", "Số phiếu", "Mã hộ", "Chủ hộ",
        "Mã đối tượng", "Họ và tên", "Ngày sinh", "Tuổi", "Giới tính",
        "Số định danh", "Cư trú", "Tình trạng học tập", "Trường", "Địa bàn trường",
        "Lớp", "Trong diện huy động", "Đến lớp", "Hoàn thành theo độ tuổi",
        "Khuyết tật", "Học hòa nhập", "Nội dung cần kiểm tra",
    ]


def _detail_values(index: int, item: dict[str, Any]) -> list[Any]:
    completion_label = (
        "Có" if item["completed_value"] is True
        else "Không" if item["completed_value"] is False
        else "Chưa xác định"
    )
    return [
        index, item.get("commune_name", ""), item.get("batch_code", ""),
        item.get("form_number", ""), item.get("household_code", ""),
        item.get("head_name", ""), item.get("person_code", ""), item.get("full_name", ""),
        item["date_of_birth_display"], item["age"], item["gender_label"],
        item.get("personal_id", "") or "", item["residency_status_label"],
        item["learning_status_label"], item["school_display"], item["school_commune_name"],
        item["class_display"], "Có" if item["must_mobilize"] else "Không",
        "Có" if item["attending"] else "Không", completion_label,
        item["disability_status_label"], "Có" if item["inclusive"] else "Không/Chưa xác định",
        item["attention_reason"],
    ]


def _append_details(workbook: Workbook, data: dict[str, Any]) -> None:
    ws = workbook.create_sheet("Danh sach tre 3-5")
    ws.append(_detail_headers())
    for index, item in enumerate(data["person_rows"], start=1):
        ws.append(_detail_values(index, item))
    _style_table(ws, header_row=1, widths={1: 7, 2: 26, 3: 24, 4: 24, 5: 20, 6: 24, 7: 18, 8: 28, 9: 14, 10: 9, 11: 12, 12: 18, 13: 18, 14: 22, 15: 34, 16: 28, 17: 18, 18: 18, 19: 12, 20: 26, 21: 22, 22: 18, 23: 58})


def _append_attention(workbook: Workbook, data: dict[str, Any]) -> None:
    ws = workbook.create_sheet("Can kiem tra")
    ws.append(_detail_headers())
    for index, item in enumerate(data["attention_rows"], start=1):
        ws.append(_detail_values(index, item))
    _style_table(ws, header_row=1, widths={1: 7, 2: 26, 3: 24, 4: 24, 5: 20, 6: 24, 7: 18, 8: 28, 9: 14, 10: 9, 11: 12, 12: 18, 13: 18, 14: 22, 15: 34, 16: 28, 17: 18, 18: 18, 19: 12, 20: 26, 21: 22, 22: 18, 23: 58})


def _append_info(
    workbook: Workbook,
    *,
    data: dict[str, Any],
    scope_label: str,
    user: dict[str, Any],
    data_state_label: str,
    q: str,
) -> None:
    ws = workbook.create_sheet("Thong tin bao cao")
    ws.append(["Thông tin", "Giá trị"])
    rows = [
        ("Tên báo cáo", "Báo cáo chỉ tiêu trẻ em từ 3 đến 5 tuổi theo Nghị định 277/2025/NĐ-CP"),
        ("Năm học", data["school_year"].code if data.get("school_year") else ""),
        ("Phạm vi", scope_label),
        ("Ngày tham chiếu tính tuổi", _safe_date(data["reference_date"])),
        ("Trạng thái", "CHÍNH THỨC" if data["is_official"] else "DỰ THẢO"),
        ("Trạng thái dữ liệu lọc", data_state_label),
        ("Từ khóa", q or "Không"),
        ("Người lập", user.get("full_name", "")),
        ("Vai trò", user.get("role_name", "")),
        ("Đơn vị", user.get("unit_name", "")),
        ("Thời điểm lập", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Căn cứ", "Nghị định 277/2025/NĐ-CP, đặc biệt Điều 12 về tiêu chuẩn công nhận đạt chuẩn"),
        ("Ngưỡng địa bàn thông thường", "Tỷ lệ huy động >= 90%; tỷ lệ hoàn thành CT GDMN theo độ tuổi >= 85%"),
        ("Ngưỡng địa bàn KTXH đặc biệt khó khăn", "Tỷ lệ huy động >= 85%; tỷ lệ hoàn thành CT GDMN theo độ tuổi >= 80%"),
        ("Cách tính tỷ lệ huy động", "Số trẻ 3-5 tuổi đến lớp / số trẻ 3-5 tuổi phải huy động"),
        ("Cách tính tỷ lệ hoàn thành", "Số trẻ được xác nhận hoàn thành CT GDMN theo độ tuổi / số trẻ 3-5 tuổi đến lớp"),
        ("Giới hạn kết luận", "Cột đạt/chưa đạt chỉ phản ánh chỉ tiêu trẻ em. Công nhận đạt chuẩn còn phải kiểm tra chương trình giáo dục, đội ngũ và điều kiện bảo đảm."),
        ("Biểu mẫu tham chiếu", "MN-01 TE và MN-02 trong Biểu mẫu PCGDMN 2025 do người dùng cung cấp"),
        ("Nguyên tắc an toàn", "Báo cáo chỉ đọc dữ liệu; không tự sửa hồ sơ điều tra hoặc hồ sơ học sinh."),
    ]
    for item in rows:
        ws.append(list(item))
    _style_table(ws, header_row=1, widths={1: 42, 2: 105})


@router.get("", response_class=HTMLResponse)
def preschool_3_5_report_page(
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
    data_state: str = "",
    issue_group: str = "",
    q: str = "",
    page_size: int = 20,
    page: int = 1,
    db: Session = Depends(get_db),
):
    user = lay_thong_tin_nguoi_dung(request)
    school_years = _load_school_years(db)
    selected_year_id = _choose_year_id(
        school_years,
        _parse_optional_int(school_year_id),
    )
    selected_year = next(
        (item for item in school_years if item.id == selected_year_id),
        None,
    )
    normalized_state = _clean_text(data_state).upper()
    if normalized_state not in DATA_STATE_LABELS:
        normalized_state = ""
    normalized_issue = _clean_text(issue_group).upper()
    if normalized_issue not in ISSUE_GROUP_LABELS:
        normalized_issue = ""
    q = _clean_text(q)[:120]
    if page_size not in PAGE_SIZE_OPTIONS:
        page_size = 20

    (
        communes,
        schools,
        selected_commune_id,
        selected_school_id,
        scope_label,
    ) = _load_scope_options(
        db,
        user,
        _parse_optional_int(commune_id),
        _parse_optional_int(school_id),
    )

    data = build_preschool_3_5_report_data(
        db=db,
        request=request,
        school_year_id=selected_year_id,
        commune_id=selected_commune_id,
        school_id=selected_school_id,
        data_state=normalized_state,
        q=q,
    )
    filtered_people = [
        item for item in data["person_rows"]
        if _matches_issue_group(item, normalized_issue)
    ]
    total_records = len(filtered_people)
    total_pages = max(1, ceil(total_records / page_size))
    page = max(1, min(page, total_pages))
    start = (page - 1) * page_size
    page_rows = filtered_people[start:start + page_size]

    page_links = [
        {
            "number": number,
            "url": _build_url(
                school_year_id=selected_year_id,
                commune_id=selected_commune_id,
                school_id=selected_school_id,
                data_state=normalized_state,
                issue_group=normalized_issue,
                q=q,
                page_size=page_size,
                page=number,
            ),
        }
        for number in range(max(1, page - 2), min(total_pages, page + 2) + 1)
    ]
    export_params = urlencode(
        {
            "school_year_id": selected_year_id or "",
            "commune_id": selected_commune_id or "",
            "school_id": selected_school_id or "",
            "data_state": normalized_state,
            "q": q,
        }
    )

    return templates.TemplateResponse(
        request=request,
        name="reports/preschool_3_5_report.html",
        context={
            "nguoi_dung": user,
            "school_years": school_years,
            "selected_year": selected_year,
            "selected_year_id": selected_year_id,
            "communes": communes,
            "schools": schools,
            "selected_commune_id": selected_commune_id,
            "selected_school_id": selected_school_id,
            "scope_label": scope_label,
            "data_state_labels": DATA_STATE_LABELS,
            "data_state": normalized_state,
            "issue_group_labels": ISSUE_GROUP_LABELS,
            "issue_group": normalized_issue,
            "q": q,
            "page_size": page_size,
            "page_size_options": PAGE_SIZE_OPTIONS,
            "page": page,
            "total_pages": total_pages,
            "page_links": page_links,
            "previous_url": (
                _build_url(
                    school_year_id=selected_year_id,
                    commune_id=selected_commune_id,
                    school_id=selected_school_id,
                    data_state=normalized_state,
                    issue_group=normalized_issue,
                    q=q,
                    page_size=page_size,
                    page=page - 1,
                ) if page > 1 else None
            ),
            "next_url": (
                _build_url(
                    school_year_id=selected_year_id,
                    commune_id=selected_commune_id,
                    school_id=selected_school_id,
                    data_state=normalized_state,
                    issue_group=normalized_issue,
                    q=q,
                    page_size=page_size,
                    page=page + 1,
                ) if page < total_pages else None
            ),
            "export_url": "/bao-cao/pho-cap-3-5-tuoi/xuat-excel?" + export_params,
            "province_reader": is_province_reader(user.get("role_code")),
            "person_rows": page_rows,
            "total_person_rows": total_records,
            "commune_rows": data["commune_rows"],
            "school_rows": data["school_rows"],
            "summary": data["summary"],
            "gender_rows": data["gender_rows"],
            "residency_rows": data["residency_rows"],
            "province_reference": data["province_reference"],
            "reference_date": data["reference_date"],
            "is_official": data["is_official"],
            "generated_at": datetime.now().strftime("%d/%m/%Y %H:%M"),
        },
    )


@router.get("/xuat-excel")
def export_preschool_3_5_report_excel(
    request: Request,
    school_year_id: str = "",
    commune_id: str = "",
    school_id: str = "",
    data_state: str = "",
    q: str = "",
    db: Session = Depends(get_db),
):
    user = lay_thong_tin_nguoi_dung(request)
    school_years = _load_school_years(db)
    selected_year_id = _choose_year_id(
        school_years,
        _parse_optional_int(school_year_id),
    )
    normalized_state = _clean_text(data_state).upper()
    if normalized_state not in DATA_STATE_LABELS:
        normalized_state = ""
    q = _clean_text(q)[:120]
    (
        _communes,
        _schools,
        selected_commune_id,
        selected_school_id,
        scope_label,
    ) = _load_scope_options(
        db,
        user,
        _parse_optional_int(commune_id),
        _parse_optional_int(school_id),
    )
    data = build_preschool_3_5_report_data(
        db=db,
        request=request,
        school_year_id=selected_year_id,
        commune_id=selected_commune_id,
        school_id=selected_school_id,
        data_state=normalized_state,
        q=q,
    )

    workbook = Workbook()
    _append_overview(workbook, data=data, scope_label=scope_label, user=user)
    _append_mn01(workbook, data, scope_label)
    _append_mn02(workbook, data)
    _append_schools(workbook, data)
    _append_details(workbook, data)
    _append_attention(workbook, data)
    _append_info(
        workbook,
        data=data,
        scope_label=scope_label,
        user=user,
        data_state_label=DATA_STATE_LABELS[normalized_state],
        q=q,
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    year_code = data["school_year"].code if data.get("school_year") else "khong_ro"
    suffix = "chinh_thuc" if data["is_official"] else "du_thao"
    filename = f"bao_cao_tre_3_5_nd277_{str(year_code).replace('/', '-')}_{suffix}.xlsx"
    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
