from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from pathlib import Path
import re
from typing import Any

from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models import Classroom, School, SchoolYear
from app.staff_models import StaffYearRecord
from app.survey_models import SurveyPerson, SurveyPersonYearRecord
from app.pcgd_xmc_report_builders_v1 import (
    _age,
    _disabled_access,
    _effective_scope_for_user,
    _ethnic,
    _female,
    _is_disabled,
    _is_ppc,
    _load_people_and_records,
    _norm,
    _percent,
    _safe_filename_text,
    _status,
)

APP_DIR = Path(__file__).resolve().parent
TEMPLATE_ROOT = APP_DIR / "report_templates" / "pcgd_mn_2025"

REPORT_TEMPLATE_FILES = {
    "PCGD_MN_M1_2025": "PCGD_2025_MN_M1.xlsx",
    "PCGD_MN_02_2025": "PCGD_2025_MN_02.xlsx",
    "PCGD_MN_01_GV_2025": "PCGD_2025_MN_01_GV.xlsx",
    "PCGD_MN_01_CSVC_2025": "PCGD_2025_MN_01_CSVC.xlsx",
    "PCGD_MN_TAICHINH_2025": "PCGD_2025_MN_TAICHINH.xlsx",
}
SUPPORTED_REPORT_TYPES = frozenset(REPORT_TEMPLATE_FILES)

ATTENDING = {"DANG HOC", "DANG_HOC", "CHUYEN DEN", "CHUYEN_DEN", "TAM NGHI", "TAM_NGHI"}
MOVE_IN = {"CHUYEN DEN", "CHUYEN_DEN"}
MOVE_OUT = {"CHUYEN DI", "CHUYEN_DI"}
DECEASED = {"CHET", "DA CHET", "DA_CHET"}


def _reference_year(school_year: SchoolYear | None) -> int:
    if school_year is not None:
        match = re.search(r"(20\d{2})", str(getattr(school_year, "code", "") or ""))
        if match:
            return int(match.group(1))
    from datetime import datetime
    return datetime.now().year


def _clean_label(value: Any) -> str:
    return " ".join(str(value or "").strip().split())


def _scope_meta(*, scope_label: str, commune_id: int | None, school_id: int | None, year: int) -> dict[str, str]:
    label = _clean_label(scope_label) or "Toàn tỉnh"
    if commune_id is None and school_id is None:
        return {
            "title": "Toàn tỉnh",
            "place": "Nghệ An",
            "signature": "XÁC NHẬN CỦA SỞ GIÁO DỤC VÀ ĐÀO TẠO",
            "signer": "GIÁM ĐỐC",
            "kind": "province",
            "year": str(year),
        }
    if school_id is not None:
        return {
            "title": label,
            "place": label,
            "signature": "XÁC NHẬN CỦA NHÀ TRƯỜNG",
            "signer": "HIỆU TRƯỞNG",
            "kind": "school",
            "year": str(year),
        }
    return {
        "title": label,
        "place": label,
        "signature": "XÁC NHẬN CỦA UBND XÃ/PHƯỜNG",
        "signer": "PHÓ CHỦ TỊCH",
        "kind": "commune",
        "year": str(year),
    }


def _safe_set(ws: Any, ref: str, value: Any) -> None:
    cell = ws[ref]
    if cell.__class__.__name__ != "MergedCell":
        cell.value = value


def _clear_range(ws: Any, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if cell.__class__.__name__ != "MergedCell":
                cell.value = None


def _replace_bottom(ws: Any, tokens: tuple[str, ...], value: str) -> bool:
    start = max(1, ws.max_row - 20)
    for row in range(ws.max_row, start - 1, -1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            if cell.__class__.__name__ == "MergedCell":
                continue
            text = str(cell.value or "")
            if text and any(token in text for token in tokens):
                cell.value = value
                return True
    return False


def _write_scope(ws: Any, report_type: str, meta: dict[str, str]) -> None:
    year = meta["year"]
    title = meta["title"]
    kind = meta["kind"]
    _safe_set(ws, "A1", "Tỉnh: Nghệ An")
    if kind == "province":
        _safe_set(ws, "A2", "Toàn tỉnh")
    elif kind == "school":
        _safe_set(ws, "A2", title if _norm(title).startswith("TRUONG ") else f"Trường: {title}")
    else:
        _safe_set(ws, "A2", title)

    # Update any date caption near the top without touching formatting/merges.
    for r in range(1, min(ws.max_row, 7) + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            if cell.__class__.__name__ == "MergedCell":
                continue
            text = str(cell.value or "")
            if "tháng 9 năm" in text or "Thời điểm" in text or "Tính đến" in text:
                if "Tính đến" in text:
                    cell.value = f"Tính đến thời điểm: ngày 30 tháng 9 năm {year}"
                else:
                    cell.value = f"Thời điểm: ngày 30 tháng 9 năm {year}"

    _replace_bottom(ws, ("ngày      tháng", "ngày     tháng", "ngày       tháng"), f"{meta['place']}, ngày      tháng      năm {year}")
    _replace_bottom(ws, ("XÁC NHẬN CỦA UBND", "XÁC NHẬN CỦA SỞ", "XÁC NHẬN CỦA NHÀ TRƯỜNG"), meta["signature"])
    _replace_bottom(ws, ("PHÓ CHỦ TỊCH", "GIÁM ĐỐC", "HIỆU TRƯỞNG"), meta["signer"])
    if kind == "province":
        try:
            ws.title = "Toàn tỉnh"
        except Exception:
            pass


def _is_preschool_school(school: School) -> bool:
    text = _norm(getattr(school, "name", ""))
    return (
        "MAM NON" in text
        or "MAU GIAO" in text
        or re.search(r"(^| )MN( |$)", text) is not None
    )


def _schools_and_classes(db: Session, school_year_id: int, commune_id: int | None, school_id: int | None) -> tuple[list[School], dict[int, list[Classroom]]]:
    stmt = select(School).where(School.is_active.is_(True)).order_by(School.name.asc())
    if school_id is not None:
        stmt = stmt.where(School.id == school_id)
    elif commune_id is not None:
        stmt = stmt.where(School.commune_id == commune_id)
    schools = [s for s in db.scalars(stmt).all() if _is_preschool_school(s)]
    grouped: dict[int, list[Classroom]] = defaultdict(list)
    if schools:
        cstmt = select(Classroom).where(
            Classroom.school_year_id == school_year_id,
            Classroom.school_id.in_([int(s.id) for s in schools]),
            Classroom.is_active.is_(True),
        )
        for item in db.scalars(cstmt).all():
            grouped[int(item.school_id)].append(item)
    return schools, grouped


def _is_five_class(name: Any) -> bool:
    text = _norm(name)
    return bool(
        "5 TUOI" in text
        or "MAU GIAO 5" in text
        or "LOP LON" in text
        or "MG 5" in text
        or re.search(r"(^| )5T( |$)", text)
    )


def _record_attending(record: SurveyPersonYearRecord | None) -> bool:
    return _status(record) in ATTENDING


def _completed_preschool(record: SurveyPersonYearRecord | None) -> bool:
    if record is None:
        return False
    return bool(
        getattr(record, "completed_preschool_5", None) is True
        or getattr(record, "completed_preschool_by_age", None) is True
    )


def _prepared_vietnamese(record: SurveyPersonYearRecord | None) -> bool:
    return bool(record is not None and getattr(record, "prepared_vietnamese", None) is True)


def _metrics(people: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year: int) -> dict[int, dict[str, int]]:
    out: dict[int, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for person, record in people:
        age = _age(person, year)
        if age is None or age < 0 or age > 6:
            continue
        m = out[age]
        m["total"] += 1
        m["female"] += int(_female(person))
        m["ethnic"] += int(_ethnic(person))
        m["disabled"] += int(_is_disabled(person, record))
        m["disabled_access"] += int(_disabled_access(person, record))
        m["ppc"] += int(_is_ppc(person, record))
        attending = _record_attending(record)
        m["attending"] += int(attending)
        if attending:
            m["attending_female"] += int(_female(person))
            m["attending_ethnic"] += int(_ethnic(person))
            m["prepared_vietnamese"] += int(_ethnic(person) and _prepared_vietnamese(record))
        status = _status(record)
        m["move_in"] += int(status in MOVE_IN)
        m["move_out"] += int(status in MOVE_OUT)
        m["deceased"] += int(status in DECEASED)
        m["completed"] += int(_completed_preschool(record))
    return out


def _fill_mn_m1(ws: Any, people: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year: int) -> None:
    metrics = _metrics(people, year)
    # Template columns E:K correspond to ages 0..6; L is total 0..5.
    age_cols = {age: 5 + age for age in range(0, 7)}
    _clear_range(ws, 7, 29, 5, 12)
    _clear_range(ws, 32, 35, 4, 5)

    row_keys = {
        7: "total", 8: "female", 9: "ethnic", 10: "disabled",
        12: "disabled_access", 13: "ppc", 14: "attending",
        18: "attending_female", 19: "attending_ethnic",
        20: "prepared_vietnamese", 24: "deceased", 25: "move_out",
        26: "move_in", 27: "completed",
    }
    for row, key in row_keys.items():
        for age, col in age_cols.items():
            ws.cell(row=row, column=col).value = int(metrics[age].get(key, 0))
        ws.cell(row=row, column=12).value = sum(int(metrics[a].get(key, 0)) for a in range(0, 6))

    # Rows 11 (disabled capable of learning), 15/16/21 location buckets,
    # 22/23 two sessions per day are left blank unless a direct source is defined.
    for age, col in age_cols.items():
        ws.cell(row=17, column=col).value = _percent(
            int(metrics[age].get("attending", 0)),
            int(metrics[age].get("ppc", 0)),
        )
        ws.cell(row=28, column=col).value = _percent(
            int(metrics[age].get("completed", 0)),
            int(metrics[age].get("attending", 0)),
        )
    ws["L17"] = _percent(
        sum(int(metrics[a].get("attending", 0)) for a in range(0, 6)),
        sum(int(metrics[a].get("ppc", 0)) for a in range(0, 6)),
    )
    ws["L28"] = _percent(
        sum(int(metrics[a].get("completed", 0)) for a in range(0, 6)),
        sum(int(metrics[a].get("attending", 0)) for a in range(0, 6)),
    )

    five = metrics[5]
    ws["D32"] = int(five.get("attending", 0))
    ws["E32"] = _percent(int(five.get("attending", 0)), int(five.get("ppc", 0)))
    ws["D33"] = int(five.get("completed", 0))
    ws["E33"] = _percent(int(five.get("completed", 0)), int(five.get("attending", 0)))
    ws["D34"] = int(five.get("disabled_access", 0))
    ws["E34"] = _percent(int(five.get("disabled_access", 0)), int(five.get("disabled", 0)))


def _fill_mn02(ws: Any, people: list[tuple[SurveyPerson, SurveyPersonYearRecord | None]], year: int, school_count: int, five_class_count: int, label: str) -> None:
    metrics = _metrics(people, year)
    five = metrics[5]
    six = metrics[6]
    _clear_range(ws, 7, 7, 1, 18)
    ws["A7"] = 1
    ws["B7"] = label
    ws["C7"] = school_count
    # D: school sites has no dedicated source -> blank
    ws["E7"] = five_class_count
    ws["F7"] = int(five.get("ppc", 0))
    ws["G7"] = int(five.get("attending", 0))
    ws["H7"] = _percent(int(five.get("attending", 0)), int(five.get("ppc", 0)))
    ws["I7"] = int(six.get("ppc", 0))
    ws["J7"] = int(six.get("completed", 0))
    ws["K7"] = _percent(int(six.get("completed", 0)), int(six.get("ppc", 0)))
    ws["L7"] = int(five.get("disabled", 0))
    # M capable-of-learning -> blank
    ws["N7"] = int(five.get("disabled_access", 0))
    ws["O7"] = _percent(int(five.get("disabled_access", 0)), int(five.get("disabled", 0)))
    # P/Q/R require dedicated staff/facility/standard sources -> blank in V1.6B


def _load_staff(db: Session, school_year_id: int, school_ids: list[int]) -> dict[int, list[StaffYearRecord]]:
    result: dict[int, list[StaffYearRecord]] = defaultdict(list)
    if not school_ids:
        return result
    stmt = (
        select(StaffYearRecord)
        .where(
            StaffYearRecord.school_year_id == school_year_id,
            StaffYearRecord.school_id.in_(school_ids),
            StaffYearRecord.is_active.is_(True),
        )
        .options(selectinload(StaffYearRecord.staff_member))
    )
    for row in db.scalars(stmt).all():
        result[int(row.school_id)].append(row)
    return result


def _staff_summary(records: list[StaffYearRecord], class_count: int, five_classes: int) -> list[Any]:
    managers = [r for r in records if _norm(getattr(r, "position_group", "")) == "CBQL"]
    teachers = [r for r in records if _norm(getattr(r, "position_group", "")) == "GIAO VIEN"]
    employees = [r for r in records if _norm(getattr(r, "position_group", "")) == "NHAN VIEN"]

    def contract(r: StaffYearRecord) -> bool:
        return "HOP DONG" in _norm(getattr(r, "employment_type", ""))

    def five_teacher(r: StaffYearRecord) -> bool:
        text = _norm(" ".join([
            str(getattr(r, "teaching_age_group", "") or ""),
            str(getattr(r, "position_title", "") or ""),
            str(getattr(r, "notes", "") or ""),
        ]))
        return "5 TUOI" in text or "MG 5" in text or "LOP LON" in text

    five_teachers = [r for r in teachers if five_teacher(r)]
    labor_all = [r for r in records if contract(r)]
    five_labor = [r for r in five_teachers if contract(r)]
    vice = sum(1 for r in managers if "PHO HIEU TRUONG" in _norm(getattr(r, "position_title", "")))
    standard = sum(
        1 for r in five_teachers
        if _norm(getattr(r, "qualification_standard", "")) in {"DAT CHUAN", "TREN CHUAN"}
        or _norm(getattr(r, "qualification_level", "")) in {"CAO DANG", "DAI HOC", "THAC SI", "TIEN SI"}
    )
    above = sum(
        1 for r in five_teachers
        if _norm(getattr(r, "qualification_standard", "")) == "TREN CHUAN"
        or _norm(getattr(r, "qualification_level", "")) in {"DAI HOC", "THAC SI", "TIEN SI"}
    )
    prof = sum(
        1 for r in five_teachers
        if _norm(getattr(r, "professional_standard", "")) in {"DAT", "KHA", "TOT"}
    )
    return [
        len(records),
        len(records) - len(labor_all),
        len(labor_all),
        sum(1 for r in labor_all if bool(getattr(r, "receives_policy", False))),
        len(managers),
        vice,
        len(teachers),
        None,  # ethnicity of staff is not a direct current field
        round(len(teachers) / class_count, 2) if class_count else None,
        len(employees),
        len(five_teachers),
        len(five_teachers) - len(five_labor),
        len(five_labor),
        sum(1 for r in five_labor if bool(getattr(r, "receives_policy", False))),
        round(len(five_teachers) / five_classes, 2) if five_classes else None,
        standard,
        above,
        prof,
    ]


def _fill_staff(ws: Any, db: Session, school_year_id: int, schools: list[School], classes: dict[int, list[Classroom]], province: bool) -> None:
    _clear_range(ws, 9, 19, 1, 20)
    staff = _load_staff(db, school_year_id, [int(s.id) for s in schools])
    rows: list[tuple[str, list[Any]]] = []
    if province:
        all_records: list[StaffYearRecord] = []
        total_classes = 0
        five_classes = 0
        for school in schools:
            cls = classes.get(int(school.id), [])
            total_classes += len(cls)
            five_classes += sum(1 for c in cls if _is_five_class(getattr(c, "name", "")))
            all_records.extend(staff.get(int(school.id), []))
        rows = [("Toàn tỉnh", _staff_summary(all_records, total_classes, five_classes))]
    else:
        for school in schools[:10]:
            cls = classes.get(int(school.id), [])
            rows.append((school.name, _staff_summary(
                staff.get(int(school.id), []),
                len(cls),
                sum(1 for c in cls if _is_five_class(getattr(c, "name", ""))),
            )))
    for row_index, (name, values) in enumerate(rows, start=9):
        ws.cell(row_index, 1).value = row_index - 8
        ws.cell(row_index, 2).value = name
        for offset, value in enumerate(values, start=3):
            ws.cell(row_index, offset).value = value
    ws["B19"] = "Cộng phạm vi:"
    for col in range(3, 21):
        if col in {11, 17}:
            continue
        vals = [ws.cell(r, col).value for r in range(9, 19)]
        vals = [v for v in vals if isinstance(v, (int, float))]
        if vals:
            ws.cell(19, col).value = sum(vals)


def _fill_csvc(ws: Any, schools: list[School], classes: dict[int, list[Classroom]], province: bool) -> None:
    _clear_range(ws, 9, 19, 1, 22)
    rows: list[tuple[str, int, int]] = []
    if province:
        five = 0
        under = 0
        for school in schools:
            cls = classes.get(int(school.id), [])
            f = sum(1 for c in cls if _is_five_class(getattr(c, "name", "")))
            five += f
            under += max(0, len(cls) - f)
        rows = [("Toàn tỉnh", five, under)]
    else:
        for school in schools[:10]:
            cls = classes.get(int(school.id), [])
            f = sum(1 for c in cls if _is_five_class(getattr(c, "name", "")))
            rows.append((school.name, f, max(0, len(cls) - f)))
    for row_index, (name, five, under) in enumerate(rows, start=9):
        ws.cell(row_index, 1).value = row_index - 8
        ws.cell(row_index, 2).value = name
        ws.cell(row_index, 5).value = five
        ws.cell(row_index, 8).value = under
    ws["B19"] = "Cộng phạm vi:"
    ws["E19"] = sum(int(ws.cell(r, 5).value or 0) for r in range(9, 19))
    ws["H19"] = sum(int(ws.cell(r, 8).value or 0) for r in range(9, 19))
    # Other facility indicators deliberately remain blank: no dedicated data source yet.


def _fill_finance(ws: Any, year: int) -> None:
    # No dedicated finance model exists in current DB. Remove sample data rather than infer zeros.
    _clear_range(ws, 8, 24, 4, 9)
    # Keep/update visible year headings if this is the first sheet of the finance template.
    for col, y in enumerate(range(year - 4, year + 1), start=5):
        if col <= ws.max_column:
            ws.cell(6, col).value = y


def export_mn_report(*, report_type: str, db: Session, request: Any, school_year: SchoolYear | None, selected_commune_id: int | None, selected_school_id: int | None, scope_label: str) -> StreamingResponse:
    if report_type not in SUPPORTED_REPORT_TYPES:
        raise ValueError(f"Loại báo cáo Mầm non chưa hỗ trợ: {report_type}")
    if school_year is None:
        raise ValueError("Chưa xác định năm học.")

    selected_commune_id, selected_school_id, scope_label, _ = _effective_scope_for_user(
        db=db,
        request=request,
        selected_commune_id=selected_commune_id,
        selected_school_id=selected_school_id,
        scope_label=scope_label,
    )
    year = _reference_year(school_year)
    meta = _scope_meta(
        scope_label=scope_label,
        commune_id=selected_commune_id,
        school_id=selected_school_id,
        year=year,
    )

    template = TEMPLATE_ROOT / REPORT_TEMPLATE_FILES[report_type]
    if not template.exists():
        raise FileNotFoundError(f"Không tìm thấy mẫu Mầm non: {template}")

    workbook = load_workbook(template)
    ws = workbook[workbook.sheetnames[0]]
    _write_scope(ws, report_type, meta)

    people = _load_people_and_records(
        db=db,
        request=request,
        school_year_id=int(school_year.id),
        selected_commune_id=selected_commune_id,
        selected_school_id=selected_school_id,
    )
    schools, classes = _schools_and_classes(
        db,
        int(school_year.id),
        selected_commune_id,
        selected_school_id,
    )

    if report_type == "PCGD_MN_M1_2025":
        _fill_mn_m1(ws, people, year)
    elif report_type == "PCGD_MN_02_2025":
        five_classes = sum(
            1 for school in schools
            for c in classes.get(int(school.id), [])
            if _is_five_class(getattr(c, "name", ""))
        )
        _fill_mn02(ws, people, year, len(schools), five_classes, meta["title"])
    elif report_type == "PCGD_MN_01_GV_2025":
        _fill_staff(ws, db, int(school_year.id), schools, classes, meta["kind"] == "province")
    elif report_type == "PCGD_MN_01_CSVC_2025":
        _fill_csvc(ws, schools, classes, meta["kind"] == "province")
    elif report_type == "PCGD_MN_TAICHINH_2025":
        _fill_finance(ws, year)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    year_code = str(getattr(school_year, "code", year) or year).replace("/", "-")
    code = report_type.replace("PCGD_", "").replace("_2025", "")
    filename = f"PCGD_{code}_{year_code}_{_safe_filename_text(meta['title'])}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
