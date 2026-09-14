# -*- coding: utf-8 -*-
"""
BÀI 13B-11.15.2.4.9.2A
KIỂM THỬ THỰC TẾ XUẤT BÁO CÁO "TOÀN TỈNH"

MỤC TIÊU
--------
- Gọi đúng exporter thực tế sau khi cài Bài 4.9.2 V2.
- Phạm vi: Toàn tỉnh.
- Xuất 6 file Excel:
    1) MN-02
    2) TH-02
    3) XMC-4
    4) THCS-M1
    5) THCS-M2
    6) THCS-TK
- Kiểm tra các ô kết luận:
    MN-02   : R7
    TH-02   : T8
    XMC-4   : R8
    THCS-M1 : G36, G37
    THCS-M2 : W14
    THCS-TK : F8, G8, R8
- Không sửa source.
- Không có lệnh UPDATE/INSERT/DELETE database.
- Kiểm tra SHA256 database trước/sau.
"""

from __future__ import annotations

import asyncio
import hashlib
import inspect
import os
import sqlite3
import sys
import traceback
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any


PROJECT = Path(
    os.environ.get("PHOCAP_PROJECT", r"C:\PhoCap")
).resolve()

APP = PROJECT / "app"
DB_PATH = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

OUT_DIR = (
    EXPORTS
    / f"kiem_thu_bai_13b_11_15_2_4_9_2a_{STAMP}"
)

REPORT = (
    OUT_DIR
    / f"bao_cao_kiem_thu_bai_4_9_2a_{STAMP}.txt"
)

EXPECTED_COMMUNES = 130
EXPECTED_SPECIAL = 51


REPORTS = (
    {
        "code": "PCGD_MN_02_2025",
        "short": "MN-02",
        "exporter": "mn",
        "cells": ("R7",),
    },
    {
        "code": "PCGD_TH_02_2025",
        "short": "TH-02",
        "exporter": "xmc",
        "cells": ("T8",),
    },
    {
        "code": "PCGD_XMC_4_2025",
        "short": "XMC-4",
        "exporter": "xmc",
        "cells": ("R8",),
    },
    {
        "code": "PCGD_THCS_M1_2025",
        "short": "THCS-M1",
        "exporter": "xmc",
        "cells": ("G36", "G37"),
    },
    {
        "code": "PCGD_THCS_M2_2025",
        "short": "THCS-M2",
        "exporter": "xmc",
        "cells": ("W14",),
    },
    {
        "code": "PCGD_THCS_TK_2025",
        "short": "THCS-TK",
        "exporter": "xmc",
        "cells": ("F8", "G8", "R8"),
    },
)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            h.update(chunk)

    return h.hexdigest()


def db_state() -> dict[str, Any]:
    uri = DB_PATH.resolve().as_uri() + "?mode=ro"
    con = sqlite3.connect(uri, uri=True)

    try:
        integrity = str(
            con.execute(
                "PRAGMA integrity_check"
            ).fetchone()[0]
        )

        fk = len(
            con.execute(
                "PRAGMA foreign_key_check"
            ).fetchall()
        )

        total = int(
            con.execute(
                "SELECT COUNT(*) FROM communes"
            ).fetchone()[0]
        )

        special = int(
            con.execute(
                """
                SELECT COUNT(*)
                FROM communes
                WHERE COALESCE(
                    is_special_difficulty_area,
                    0
                ) = 1
                """
            ).fetchone()[0]
        )

        return {
            "integrity": integrity,
            "fk": fk,
            "total": total,
            "special": special,
        }

    finally:
        con.close()


def verify_db(state: dict[str, Any]) -> None:
    if str(state["integrity"]).lower() != "ok":
        raise RuntimeError(
            "integrity_check != ok"
        )

    if int(state["fk"]) != 0:
        raise RuntimeError(
            f"foreign_key_check có {state['fk']} lỗi."
        )

    if int(state["total"]) != EXPECTED_COMMUNES:
        raise RuntimeError(
            f"communes={state['total']}, "
            f"không phải {EXPECTED_COMMUNES}."
        )

    if int(state["special"]) != EXPECTED_SPECIAL:
        raise RuntimeError(
            f"ĐBKK={state['special']}, "
            f"không phải {EXPECTED_SPECIAL}."
        )


def make_request():
    from starlette.requests import Request

    scope = {
        "type": "http",
        "asgi": {
            "version": "3.0",
            "spec_version": "2.3",
        },
        "http_version": "1.1",
        "server": ("127.0.0.1", 80),
        "client": ("127.0.0.1", 0),
        "scheme": "http",
        "method": "GET",
        "root_path": "",
        "path": "/kiem-thu-bai-4-9-2a",
        "raw_path": b"/kiem-thu-bai-4-9-2a",
        "query_string": b"",
        "headers": [],
        "auth_user": {
            "id": 0,
            "username": "kiem_thu_4_9_2a",
            "full_name": "Kiểm thử Bài 4.9.2A",
            "role_code": "ADMIN",
            "commune_id": None,
            "school_id": None,
        },
    }

    return Request(scope)


async def response_to_bytes(response) -> bytes:
    body = getattr(
        response,
        "body",
        None,
    )

    if isinstance(
        body,
        (bytes, bytearray),
    ) and body:
        return bytes(body)

    iterator = getattr(
        response,
        "body_iterator",
        None,
    )

    if iterator is None:
        raise RuntimeError(
            f"Response {type(response).__name__} "
            "không có body/body_iterator."
        )

    chunks = []

    async for chunk in iterator:
        if isinstance(chunk, str):
            chunk = chunk.encode("utf-8")
        chunks.append(bytes(chunk))

    return b"".join(chunks)


def choose_year(db):
    from sqlalchemy import select
    from app.models import SchoolYear

    requested = str(
        os.environ.get(
            "PHOCAP_TEST_YEAR",
            "",
        )
        or ""
    ).strip()

    if requested:
        year = db.scalar(
            select(SchoolYear).where(
                SchoolYear.code
                == requested
            )
        )

        if year is None:
            raise RuntimeError(
                "PHOCAP_TEST_YEAR="
                f"{requested} nhưng DB không có năm này."
            )

        return year

    # Ưu tiên năm có nhiều dữ liệu điều tra nhất.
    from sqlalchemy import text

    rows = db.execute(
        text(
            """
            SELECT
                sy.id,
                sy.code,
                COUNT(spr.id) AS cnt
            FROM school_years sy
            LEFT JOIN survey_person_year_records spr
              ON spr.school_year_id = sy.id
            GROUP BY sy.id, sy.code
            ORDER BY cnt DESC, sy.code DESC
            """
        )
    ).all()

    if not rows:
        raise RuntimeError(
            "Database chưa có năm học."
        )

    best_id = int(rows[0][0])

    year = db.scalar(
        select(SchoolYear).where(
            SchoolYear.id == best_id
        )
    )

    if year is None:
        raise RuntimeError(
            f"Không lấy được SchoolYear ID={best_id}."
        )

    return year


def year_data_count(
    db,
    school_year_id: int,
) -> int:
    from sqlalchemy import text

    value = db.execute(
        text(
            """
            SELECT COUNT(*)
            FROM survey_person_year_records
            WHERE school_year_id=:school_year_id
            """
        ),
        {
            "school_year_id": int(
                school_year_id
            )
        },
    ).scalar_one()

    return int(value or 0)


def verify_b492_markers() -> None:
    rules = (
        APP
        / "services"
        / "pcgd_business_rules.py"
    ).read_text(
        encoding="utf-8-sig",
        errors="strict",
    )

    mn = (
        APP
        / "pcgd_mn_report_builders_v1.py"
    ).read_text(
        encoding="utf-8-sig",
        errors="strict",
    )

    xmc = (
        APP
        / "pcgd_xmc_report_builders_v1.py"
    ).read_text(
        encoding="utf-8-sig",
        errors="strict",
    )

    required_rules = (
        "BAI_13B_11_15_2_4_9_2_"
        "PROVINCE_RULES_START",
        "evaluate_preschool_3_5_province",
        "evaluate_primary_province",
        "evaluate_literacy_province",
        "evaluate_thcs_province",
    )

    required_mn = (
        "BAI_13B_11_15_2_4_9_2_"
        "MN_HELPERS_START",
        "BAI_13B_11_15_2_4_9_2_"
        "MN_CALL_START",
        "_b492_apply_mn_province_conclusion",
    )

    required_xmc = (
        "BAI_13B_11_15_2_4_9_2_"
        "XMC_HELPERS_START",
        "BAI_13B_11_15_2_4_9_2_"
        "XMC_CALL_START",
        "_b492_apply_xmc_province_conclusions",
    )

    missing = (
        [
            "RULES:" + x
            for x in required_rules
            if x not in rules
        ]
        + [
            "MN:" + x
            for x in required_mn
            if x not in mn
        ]
        + [
            "XMC:" + x
            for x in required_xmc
            if x not in xmc
        ]
    )

    if missing:
        raise RuntimeError(
            "Thiếu dấu vết Bài 4.9.2: "
            + "; ".join(missing)
        )


def call_exporter(
    func,
    *,
    report_type: str,
    db,
    request,
    school_year,
):
    values = {
        "report_type": report_type,
        "db": db,
        "request": request,
        "school_year": school_year,
        "selected_commune_id": None,
        "selected_school_id": None,
        "scope_label": "Toàn tỉnh",
    }

    sig = inspect.signature(func)
    kwargs = {}

    for name, parameter in sig.parameters.items():
        if name in values:
            kwargs[name] = values[name]
            continue

        if (
            parameter.default
            is not inspect.Parameter.empty
        ):
            continue

        if parameter.kind in (
            inspect.Parameter.VAR_POSITIONAL,
            inspect.Parameter.VAR_KEYWORD,
        ):
            continue

        raise RuntimeError(
            f"Exporter {func.__name__} "
            f"có tham số bắt buộc chưa biết: {name}."
        )

    return func(**kwargs)


def normalize(value):
    if value is None:
        return None

    if isinstance(value, bool):
        return value

    if isinstance(value, int):
        return value

    if (
        isinstance(value, float)
        and value.is_integer()
    ):
        return int(value)

    return " ".join(
        str(value).strip().split()
    )


def allowed(
    report_short: str,
    cell: str,
    value,
) -> bool:
    value = normalize(value)

    status = {
        "Không đạt",
        "Chưa đủ dữ liệu",
    }

    if (
        report_short == "MN-02"
        and cell == "R7"
    ):
        return value in {
            "Đạt",
            *status,
        }

    if report_short in {
        "TH-02",
        "THCS-M1",
        "THCS-TK",
    }:
        if cell in {
            "T8",
            "G36",
            "F8",
        }:
            return value in {
                1,
                2,
                3,
                *status,
            }

        if cell in {
            "G37",
            "G8",
        }:
            return value in {
                1,
                2,
                *status,
            }

        if cell == "R8":
            return value in {
                1,
                2,
                3,
                *status,
            }

    if (
        report_short == "XMC-4"
        and cell == "R8"
    ):
        return value in {
            1,
            2,
            *status,
        }

    if (
        report_short == "THCS-M2"
        and cell == "W14"
    ):
        return value in {
            1,
            2,
            3,
            *status,
        }

    return False


def inspect_xlsx(
    data: bytes,
    short: str,
    cells: tuple[str, ...],
):
    from openpyxl import load_workbook

    if not data.startswith(b"PK"):
        raise RuntimeError(
            f"{short}: response không phải XLSX."
        )

    wb = load_workbook(
        BytesIO(data),
        data_only=False,
    )

    try:
        ws = wb[
            wb.sheetnames[0]
        ]

        values = {
            cell: normalize(
                ws[cell].value
            )
            for cell in cells
        }

        bad = [
            f"{cell}={values[cell]!r}"
            for cell in cells
            if not allowed(
                short,
                cell,
                values[cell],
            )
        ]

        if bad:
            raise RuntimeError(
                f"{short}: giá trị kết luận "
                "không hợp lệ: "
                + ", ".join(bad)
            )

        return ws.title, values

    finally:
        wb.close()


def values_text(values: dict) -> str:
    return ", ".join(
        f"{key}={value!r}"
        for key, value in values.items()
    )


def run() -> int:
    print("=" * 104)
    print(
        "BÀI 13B-11.15.2.4.9.2A - "
        'KIỂM THỬ XUẤT "TOÀN TỈNH"'
    )
    print("=" * 104)
    print()

    if not DB_PATH.exists():
        print(
            f"[LỖI] Không tìm thấy DB: "
            f"{DB_PATH}"
        )
        return 1

    db_hash_before = (
        sha256_file(DB_PATH)
    )

    before = db_state()
    verify_db(before)

    old_cwd = Path.cwd()

    try:
        os.chdir(PROJECT)
        sys.path.insert(
            0,
            str(PROJECT),
        )

        print(
            "[1/7] Database: "
            f"{before['total']} xã/phường; "
            f"ĐBKK={before['special']}; "
            "integrity=ok."
        )

        verify_b492_markers()

        print(
            "[2/7] Source Bài 4.9.2 V2: "
            "ĐỦ marker."
        )

        from sqlalchemy import text
        from app.database import SessionLocal
        import app.pcgd_mn_report_builders_v1 as mn_builder
        import app.pcgd_xmc_report_builders_v1 as xmc_builder

        db = SessionLocal()

        try:
            # Chặn ghi DB trên connection kiểm thử.
            db.execute(
                text(
                    "PRAGMA query_only = ON"
                )
            )

            school_year = (
                choose_year(db)
            )

            count = year_data_count(
                db,
                int(school_year.id),
            )

            print(
                "[3/7] Năm học kiểm thử: "
                f"{school_year.code} "
                f"(ID={school_year.id}); "
                f"bản ghi={count}."
            )

            request = make_request()

            OUT_DIR.mkdir(
                parents=True,
                exist_ok=False,
            )

            report_rows = [
                "=" * 104,
                "BÀI 13B-11.15.2.4.9.2A - "
                'KIỂM THỬ XUẤT "TOÀN TỈNH"',
                "=" * 104,
                "",
                f"Năm học: {school_year.code}",
                f"Số bản ghi năm học: {count}",
                "",
            ]

            print(
                "[4/7] Xuất 6 báo cáo..."
            )

            for spec in REPORTS:
                exporter = (
                    mn_builder.export_mn_report
                    if spec["exporter"] == "mn"
                    else xmc_builder.export_additional_report
                )

                response = call_exporter(
                    exporter,
                    report_type=spec["code"],
                    db=db,
                    request=request,
                    school_year=school_year,
                )

                data = asyncio.run(
                    response_to_bytes(
                        response
                    )
                )

                if len(data) < 1000:
                    raise RuntimeError(
                        f"{spec['short']}: "
                        f"file chỉ có {len(data)} bytes."
                    )

                path = (
                    OUT_DIR
                    / (
                        f"{spec['short']}_"
                        f"{school_year.code}_"
                        "TOAN_TINH.xlsx"
                    )
                )

                path.write_bytes(data)

                sheet, values = (
                    inspect_xlsx(
                        data,
                        spec["short"],
                        spec["cells"],
                    )
                )

                print(
                    "      [OK] "
                    f"{spec['short']} / "
                    f"{values_text(values)}"
                )

                report_rows.extend(
                    [
                        (
                            f"{spec['short']} | "
                            f"sheet={sheet} | "
                            f"{values_text(values)}"
                        ),
                        f"  File: {path}",
                    ]
                )

        finally:
            try:
                db.rollback()
            except Exception:
                pass
            db.close()

        print(
            "[5/7] Đã xuất đủ 6 file Excel."
        )

        after = db_state()
        verify_db(after)

        db_hash_after = (
            sha256_file(DB_PATH)
        )

        if before != after:
            raise RuntimeError(
                "Trạng thái DB trước/sau khác nhau."
            )

        if (
            db_hash_before
            != db_hash_after
        ):
            raise RuntimeError(
                "SHA256 phocap.db thay đổi. "
                "Hãy dừng Uvicorn rồi kiểm thử lại."
            )

        print(
            "[6/7] Database trước/sau: "
            "KHÔNG THAY ĐỔI; SHA256 giữ nguyên."
        )

        report_rows.extend(
            [
                "",
                "AN TOÀN DATABASE:",
                f"SHA256 trước: {db_hash_before}",
                f"SHA256 sau  : {db_hash_after}",
                "Database: KHÔNG THAY ĐỔI.",
                "",
                "GHI CHÚ:",
                (
                    "Nếu nhiều xã chưa có dữ liệu, "
                    "'Chưa đủ dữ liệu' là kết quả hợp lệ."
                ),
                (
                    "Bài 4.9.2A chỉ xác nhận "
                    "đường xuất Toàn tỉnh hoạt động đúng."
                ),
                "",
                "KẾT LUẬN: BÀI 4.9.2A ĐẠT KIỂM THỬ KỸ THUẬT.",
            ]
        )

        REPORT.write_text(
            "\n".join(report_rows),
            encoding="utf-8-sig",
        )

        print(
            f"[7/7] Báo cáo: {REPORT}"
        )

        print()
        print("=" * 104)
        print(
            "BÀI 4.9.2A HOÀN THÀNH."
        )
        print(
            f"6 file Excel: {OUT_DIR}"
        )
        print(
            "Database: KHÔNG THAY ĐỔI."
        )
        print("=" * 104)

        return 0

    except Exception as exc:
        print()
        print("=" * 104)
        print(
            "[LỖI] BÀI 4.9.2A CHƯA ĐẠT"
        )
        print(str(exc))
        print("=" * 104)

        try:
            OUT_DIR.mkdir(
                parents=True,
                exist_ok=True,
            )

            error_path = (
                OUT_DIR
                / f"loi_kiem_thu_4_9_2a_{STAMP}.txt"
            )

            error_path.write_text(
                (
                    f"Lỗi: {exc}\n\n"
                    + traceback.format_exc()
                ),
                encoding="utf-8-sig",
            )

            print(
                f"Chi tiết lỗi: {error_path}"
            )

        except Exception:
            pass

        print(
            "Script bật PRAGMA query_only "
            "và không có lệnh UPDATE/INSERT/DELETE."
        )

        return 1

    finally:
        os.chdir(old_cwd)


if __name__ == "__main__":
    raise SystemExit(run())
