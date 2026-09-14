from __future__ import annotations

import csv
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


PROJECT_DIR = Path(__file__).resolve().parent
EXPORT_ROOT = PROJECT_DIR / "exports"
SUMMARY_FILENAME = "11_ma_bat_buoc_kiem_tra.csv"
DETAIL_FILENAME = "12_chi_tiet_day_du_ma_trung.csv"
OUTPUT_FILENAME = "13_phieu_xac_minh_22_ma.xlsx"

DECISIONS = [
    "GIỮ MÃ - CHỌN MỘT DÒNG",
    "SỬA MÃ ĐỊNH DANH",
    "LOẠI KHỎI DANH SÁCH",
    "CHƯA ĐỦ THÔNG TIN",
]

SUMMARY_COLUMNS = [
    ("STT", None),
    ("Mã định danh Bộ GD&ĐT", "staff_code"),
    ("Phân loại", "phan_loai"),
    ("Khuyến nghị", "khuyen_nghi"),
    ("Giải thích", "giai_thich"),
    ("Số dòng đủ điều kiện", "so_dong_du_dieu_kien"),
    ("Số dòng không nhập", "so_dong_khong_nhap"),
    ("Tổng số dòng đối chiếu", "tong_so_dong_doi_chieu"),
    ("Các dòng Excel", "excel_rows"),
    ("Các họ tên", "full_names"),
    ("Các ngày sinh", "dates_of_birth"),
    ("Các xã/phường", "communes"),
    ("Các trường", "schools"),
    ("Các trạng thái", "employment_statuses"),
    ("Các vị trí việc làm", "job_positions"),
    ("Lý do dòng không nhập", "ignored_reasons"),
    ("QUYẾT ĐỊNH XÁC MINH", "__decision__"),
    ("DÒNG EXCEL ĐƯỢC CHỌN", "__selected_excel_row__"),
    ("MÃ ĐỊNH DANH ĐÚNG", "__correct_staff_code__"),
    ("HỌ TÊN ĐÃ XÁC MINH", "__verified_full_name__"),
    ("NGÀY SINH ĐÃ XÁC MINH", "__verified_date_of_birth__"),
    ("GHI CHÚ XÁC MINH", "__verification_note__"),
    ("NGƯỜI XÁC MINH", "__verified_by__"),
    ("NGÀY XÁC MINH", "__verified_at__"),
]


THIN_GRAY = Side(style="thin", color="D9E2F3")
BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="2F75B5")
EDITABLE_FILL = PatternFill("solid", fgColor="FFF2CC")
INFO_FILL = PatternFill("solid", fgColor="DDEBF7")
SUCCESS_FILL = PatternFill("solid", fgColor="E2F0D9")
WARNING_FILL = PatternFill("solid", fgColor="FCE4D6")


def configure_console() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except (AttributeError, OSError):
        pass


def clean(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    encodings = ("utf-8-sig", "utf-8", "cp1258")
    last_error: Exception | None = None
    for encoding in encodings:
        try:
            with path.open("r", encoding=encoding, newline="") as file:
                return [
                    {clean(key): clean(value) for key, value in row.items()}
                    for row in csv.DictReader(file)
                ]
        except UnicodeDecodeError as error:
            last_error = error
    if last_error:
        raise last_error
    return []


def find_latest_report_dir() -> Path:
    if not EXPORT_ROOT.exists():
        raise FileNotFoundError(f"Không tìm thấy thư mục báo cáo: {EXPORT_ROOT}")

    candidates = [
        path
        for path in EXPORT_ROOT.glob("kiem_tra_giao_vien_*")
        if path.is_dir()
        and (path / SUMMARY_FILENAME).exists()
        and (path / DETAIL_FILENAME).exists()
    ]
    if not candidates:
        raise FileNotFoundError(
            "Không tìm thấy thư mục báo cáo có đủ hai tệp: "
            f"{SUMMARY_FILENAME} và {DETAIL_FILENAME}."
        )
    return max(candidates, key=lambda item: item.stat().st_mtime)


def apply_common_cell_style(cell: Any) -> None:
    cell.border = BORDER
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_instruction_sheet(workbook: Workbook, report_dir: Path) -> None:
    sheet = workbook.active
    sheet.title = "HUONG_DAN"
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A1:H1")
    sheet["A1"] = "PHIẾU XÁC MINH 22 MÃ ĐỊNH DANH BỊ XUNG ĐỘT"
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=16)
    sheet["A1"].fill = TITLE_FILL
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 32

    instructions = [
        "Mục đích: xác định đúng người và đúng mã định danh trước khi tạo tài khoản hàng loạt.",
        "Chỉ nhập dữ liệu tại các cột nền vàng trong sheet XAC_MINH_22_MA.",
        "GIỮ MÃ - CHỌN MỘT DÒNG: chọn đúng số dòng Excel tại cột DÒNG EXCEL ĐƯỢC CHỌN.",
        "SỬA MÃ ĐỊNH DANH: nhập mã đúng tại cột MÃ ĐỊNH DANH ĐÚNG và ghi rõ căn cứ.",
        "LOẠI KHỎI DANH SÁCH: dùng khi không tạo tài khoản cho mã này.",
        "CHƯA ĐỦ THÔNG TIN: dùng khi cần liên hệ đơn vị để xác minh thêm.",
        "Sheet CHI_TIET_22_MA chứa toàn bộ dòng nguồn để đối chiếu; không sửa sheet này.",
        "Bước này không ghi, sửa hoặc xóa dữ liệu trong bảng users.",
    ]

    sheet["A3"] = "Thư mục báo cáo nguồn"
    sheet["B3"] = str(report_dir)
    sheet["A3"].font = Font(bold=True)
    sheet["B3"].font = Font(color="1F4E78")

    start_row = 5
    for index, text in enumerate(instructions, start=1):
        row = start_row + index - 1
        sheet[f"A{row}"] = index
        sheet[f"B{row}"] = text
        sheet[f"A{row}"].alignment = Alignment(horizontal="center", vertical="top")
        sheet[f"B{row}"].alignment = Alignment(vertical="top", wrap_text=True)
        sheet[f"A{row}"].border = BORDER
        sheet[f"B{row}"].border = BORDER
        if index in (2, 3, 4, 5, 6):
            sheet[f"B{row}"].fill = EDITABLE_FILL

    sheet.column_dimensions["A"].width = 10
    sheet.column_dimensions["B"].width = 105
    sheet.freeze_panes = "A5"


def add_summary_sheet(
    workbook: Workbook,
    summary_rows: list[dict[str, str]],
) -> None:
    sheet = workbook.create_sheet("XAC_MINH_22_MA")
    sheet.sheet_view.showGridLines = False

    last_column = len(SUMMARY_COLUMNS)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title_cell = sheet.cell(1, 1, "DANH SÁCH 22 MÃ BẮT BUỘC XÁC MINH")
    title_cell.font = Font(bold=True, color="FFFFFF", size=15)
    title_cell.fill = TITLE_FILL
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

    for column_index, (title, _) in enumerate(SUMMARY_COLUMNS, start=1):
        cell = sheet.cell(3, column_index, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for row_index, source in enumerate(summary_rows, start=4):
        for column_index, (_, key) in enumerate(SUMMARY_COLUMNS, start=1):
            if key is None:
                value: Any = row_index - 3
            elif key.startswith("__"):
                value = ""
            else:
                value = source.get(key, "")
            cell = sheet.cell(row_index, column_index, value)
            apply_common_cell_style(cell)
            if column_index >= 17:
                cell.fill = EDITABLE_FILL
            elif column_index in (2, 3, 4, 5):
                cell.fill = INFO_FILL

    first_data_row = 4
    last_data_row = max(4, 3 + len(summary_rows))

    decision_validation = DataValidation(
        type="list",
        formula1="=DANH_MUC!$A$2:$A$5",
        allow_blank=True,
    )
    decision_validation.error = "Hãy chọn một quyết định trong danh sách."
    decision_validation.errorTitle = "Giá trị không hợp lệ"
    decision_validation.prompt = "Chọn phương án xử lý mã định danh."
    decision_validation.promptTitle = "Quyết định xác minh"
    sheet.add_data_validation(decision_validation)
    decision_validation.add(f"Q{first_data_row}:Q{last_data_row}")

    date_validation = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2000,1,1)",
        formula2="DATE(2100,12,31)",
        allow_blank=True,
    )
    sheet.add_data_validation(date_validation)
    date_validation.add(f"X{first_data_row}:X{last_data_row}")

    sheet.conditional_formatting.add(
        f"Q{first_data_row}:Q{last_data_row}",
        FormulaRule(
            formula=[f'Q{first_data_row}=""'],
            stopIfTrue=False,
            fill=WARNING_FILL,
        ),
    )
    sheet.conditional_formatting.add(
        f"Q{first_data_row}:Q{last_data_row}",
        FormulaRule(
            formula=[f'Q{first_data_row}="GIỮ MÃ - CHỌN MỘT DÒNG"'],
            stopIfTrue=False,
            fill=SUCCESS_FILL,
        ),
    )

    sheet.auto_filter.ref = f"A3:X{last_data_row}"
    sheet.freeze_panes = "A4"

    widths = {
        "A": 7,
        "B": 22,
        "C": 20,
        "D": 30,
        "E": 45,
        "F": 14,
        "G": 14,
        "H": 16,
        "I": 20,
        "J": 35,
        "K": 25,
        "L": 28,
        "M": 36,
        "N": 25,
        "O": 30,
        "P": 35,
        "Q": 30,
        "R": 22,
        "S": 22,
        "T": 28,
        "U": 22,
        "V": 40,
        "W": 24,
        "X": 18,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["X"].number_format = "dd/mm/yyyy"


def add_detail_sheet(
    workbook: Workbook,
    detail_rows: list[dict[str, str]],
    conflict_codes: set[str],
) -> None:
    sheet = workbook.create_sheet("CHI_TIET_22_MA")
    sheet.sheet_view.showGridLines = False

    filtered_rows = [
        row for row in detail_rows if clean(row.get("staff_code")) in conflict_codes
    ]
    filtered_rows.sort(
        key=lambda item: (
            clean(item.get("staff_code")),
            int(item.get("excel_row") or 0),
        )
    )

    headers: list[str] = []
    for row in filtered_rows:
        for key in row.keys():
            if key not in headers:
                headers.append(key)

    if not headers:
        headers = ["staff_code", "ghi_chu"]
        filtered_rows = [
            {
                "staff_code": "",
                "ghi_chu": "Không có dữ liệu chi tiết để hiển thị.",
            }
        ]

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(1, 1, "CHI TIẾT CÁC DÒNG NGUỒN CỦA 22 MÃ XUNG ĐỘT")
    title_cell.font = Font(bold=True, color="FFFFFF", size=14)
    title_cell.fill = TITLE_FILL
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(3, column_index, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for row_index, row in enumerate(filtered_rows, start=4):
        for column_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row_index, column_index, row.get(header, ""))
            apply_common_cell_style(cell)
            if header == "staff_code":
                cell.fill = INFO_FILL

    last_row = max(4, 3 + len(filtered_rows))
    last_column_letter = sheet.cell(3, len(headers)).column_letter
    sheet.auto_filter.ref = f"A3:{last_column_letter}{last_row}"
    sheet.freeze_panes = "A4"

    for index, header in enumerate(headers, start=1):
        column_letter = sheet.cell(3, index).column_letter
        key = header.lower()
        if "reason" in key or "giai" in key or "note" in key:
            width = 40
        elif "name" in key or "school" in key or "commune" in key:
            width = 28
        elif "status" in key or "position" in key:
            width = 24
        else:
            width = 18
        sheet.column_dimensions[column_letter].width = width


def add_catalog_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("DANH_MUC")
    sheet["A1"] = "QUYẾT ĐỊNH XÁC MINH"
    sheet["A1"].font = Font(bold=True)
    for index, value in enumerate(DECISIONS, start=2):
        sheet.cell(index, 1, value)
    sheet.sheet_state = "hidden"


def main() -> None:
    configure_console()

    print("=" * 78)
    print("BÀI 11B-6A.7 - TẠO PHIẾU XÁC MINH 22 MÃ XUNG ĐỘT")
    print("=" * 78)
    print("Chế độ: chỉ đọc các tệp báo cáo, không ghi vào cơ sở dữ liệu.")
    print()

    report_dir = find_latest_report_dir()
    summary_path = report_dir / SUMMARY_FILENAME
    detail_path = report_dir / DETAIL_FILENAME

    summary_rows = read_csv_rows(summary_path)
    detail_rows = read_csv_rows(detail_path)

    if not summary_rows:
        raise ValueError(f"Tệp không có dữ liệu: {summary_path}")

    conflict_codes = {
        clean(row.get("staff_code"))
        for row in summary_rows
        if clean(row.get("staff_code"))
    }

    workbook = Workbook()
    add_instruction_sheet(workbook, report_dir)
    add_summary_sheet(workbook, summary_rows)
    add_detail_sheet(workbook, detail_rows, conflict_codes)
    add_catalog_sheet(workbook)

    output_path = report_dir / OUTPUT_FILENAME
    workbook.save(output_path)

    print(f"Thư mục báo cáo: {report_dir}")
    print(f"Số mã cần xác minh: {len(conflict_codes)}")
    print(f"Số dòng chi tiết nguồn: {sum(1 for row in detail_rows if clean(row.get('staff_code')) in conflict_codes)}")
    print()
    print(f"ĐÃ TẠO PHIẾU: {output_path}")
    print()
    print("HOÀN TẤT: chưa tạo, sửa hoặc xóa bất kỳ tài khoản nào.")


if __name__ == "__main__":
    main()
