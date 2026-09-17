# -*- coding: utf-8 -*-
from __future__ import annotations

import ast
import importlib.util
import py_compile
import shutil
import sqlite3
from datetime import datetime
from io import BytesIO
from pathlib import Path

from jinja2 import Environment
from openpyxl import Workbook, load_workbook

PROJECT = Path(r"C:\PhoCap")
APP = PROJECT / "app"
DB = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"

EXCHANGE = APP / "household_excel_exchange.py"
SURVEYS = APP / "routers" / "surveys.py"
IMPORT_TEMPLATE = APP / "templates" / "surveys" / "excel_import.html"
HELPER = APP / "household_field_excel_v240.py"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
BACKUP = EXPORTS / f"backup_bai_13b_12_v2_4_1_{STAMP}"
DB_BACKUP = BACKUP / "phocap.db"
REPORT = EXPORTS / f"bao_cao_cai_bai_13b_12_v2_4_1_{STAMP}.txt"

MARKER_EXPORT = "BAI_13B_12_V2_4_1_FIELD_EXCEL_EXPORT"
MARKER_IMPORT = "BAI_13B_12_V2_4_1_FIELD_EXCEL_IMPORT"
MARKER_HELPER = "PCGDMN_FIELD_FORM_V240"

HELPER_SOURCE = 'from __future__ import annotations\n\nimport json\nimport re\nfrom copy import copy\nfrom datetime import date, datetime\nfrom io import BytesIO\nfrom typing import Any\n\nfrom openpyxl import load_workbook\nfrom openpyxl.cell.cell import MergedCell\nfrom openpyxl.styles import Alignment, Border, Font, Side\nfrom openpyxl.worksheet.page import PageMargins\nfrom openpyxl.worksheet.pagebreak import Break\n\n\nTRACKING_SHEET_NAME = "Sổ theo dõi PCGDMN"\nMETA_SHEET_NAME = "_PCGDMN_META"\nFIELD_SHEET_NAME = "Phiếu điều tra"\nFIELD_MAP_SHEET_NAME = "_PCGDMN_FIELD_MAP"\nFIELD_VERSION = "PCGDMN_FIELD_FORM_V240"\nDATA_START_ROW = 11\nDATA_END_ROW = 110\nMETA_TABLE_ROW = 20\nMULTI_MARKER_COLUMN = 21\nMULTI_JSON_COLUMN = 22\nMULTI_BLOCK_START = "__PCGDMN_BLOCK_START__"\nMULTI_PERSON_ROW = "__PCGDMN_PERSON_ROW__"\nMULTI_BLOCK_END = "__PCGDMN_BLOCK_END__"\n\n\ndef _text(value: Any) -> str:\n    if value is None:\n        return ""\n    if isinstance(value, datetime):\n        return value.strftime("%d/%m/%Y")\n    if isinstance(value, date):\n        return value.strftime("%d/%m/%Y")\n    return str(value).strip()\n\n\ndef _meta_values(sheet: Any) -> dict[str, str]:\n    values: dict[str, str] = {}\n    for row in range(1, min(sheet.max_row, META_TABLE_ROW - 1) + 1):\n        key = _text(sheet.cell(row, 1).value)\n        if key:\n            values[key] = _text(sheet.cell(row, 2).value)\n    return values\n\n\ndef _single_row_meta(sheet: Any) -> dict[int, dict[str, Any]]:\n    headers = [\n        _text(sheet.cell(META_TABLE_ROW, col).value)\n        for col in range(1, sheet.max_column + 1)\n    ]\n    result: dict[int, dict[str, Any]] = {}\n    if "visible_row" not in headers:\n        return result\n    for row in range(META_TABLE_ROW + 1, sheet.max_row + 1):\n        values = {\n            header: sheet.cell(row, col).value\n            for col, header in enumerate(headers, start=1)\n            if header\n        }\n        visible = values.get("visible_row")\n        try:\n            visible_row = int(visible)\n        except (TypeError, ValueError):\n            continue\n        result[visible_row] = values\n    return result\n\n\ndef _find_tracking(workbook: Any) -> Any:\n    for name in workbook.sheetnames:\n        if str(name).strip() == TRACKING_SHEET_NAME:\n            return workbook[name]\n    raise ValueError("Không tìm thấy trang kỹ thuật Sổ theo dõi PCGDMN.")\n\n\ndef _parse_header_value(raw: Any, label: str) -> str:\n    text = _text(raw)\n    marker = f"{label}:"\n    pos = text.lower().find(marker.lower())\n    if pos < 0:\n        return text\n    return text[pos + len(marker):].strip()\n\n\ndef _extract_hamlet(header: Any) -> str:\n    raw = _text(header)\n    match = re.search(\n        r"(?:Xóm/khối|Tổ,\\s*thôn,\\s*xóm|Thôn/xóm|Xóm|Khối|Tổ|Bản)\\s*:\\s*([^;]+)",\n        raw,\n        flags=re.IGNORECASE,\n    )\n    if match:\n        return match.group(1).strip()\n    return ""\n\n\ndef _extract_address(header: Any) -> str:\n    return _parse_header_value(header, "Địa chỉ")\n\n\ndef _extract_head(header: Any) -> str:\n    return _parse_header_value(header, "Họ và tên chủ hộ")\n\n\ndef _split_name(full_name: str) -> tuple[str, str]:\n    parts = [part for part in str(full_name or "").split() if part]\n    if not parts:\n        return "", ""\n    if len(parts) == 1:\n        return parts[0], ""\n    return " ".join(parts[:-1]), parts[-1]\n\n\ndef _line_value(raw: Any, labels: tuple[str, ...]) -> str:\n    text = _text(raw)\n    for label in labels:\n        marker = label + ":"\n        pos = text.lower().find(marker.lower())\n        if pos >= 0:\n            return text[pos + len(marker):].strip()\n    return text\n\n\ndef _gender_ethnic(raw: Any) -> tuple[str, str]:\n    text = _text(raw)\n    gender = ""\n    ethnic = ""\n    match_gender = re.search(\n        r"(?:Giới\\s*tính|GT)\\s*:\\s*(Nam|Nữ|Nu)",\n        text,\n        flags=re.IGNORECASE,\n    )\n    if match_gender:\n        gender = "Nữ" if match_gender.group(1).lower() in {"nữ", "nu"} else "Nam"\n    else:\n        if re.search(r"\\bNữ\\b", text, flags=re.IGNORECASE):\n            gender = "Nữ"\n        elif re.search(r"\\bNam\\b", text, flags=re.IGNORECASE):\n            gender = "Nam"\n\n    match_ethnic = re.search(\n        r"(?:DT|Dân\\s*tộc)\\s*:\\s*([^|;]+)",\n        text,\n        flags=re.IGNORECASE,\n    )\n    if match_ethnic:\n        ethnic = match_ethnic.group(1).strip()\n    return gender, ethnic\n\n\ndef _year_label(tracking: Any, block_start: int, column: int) -> str:\n    top = _text(tracking.cell(block_start + 6, column).value)\n    bottom = _text(tracking.cell(block_start + 7, column).value)\n    if top and bottom:\n        return f"{top}-{bottom}"\n    return top or bottom\n\n\ndef _payload_from_marker(tracking: Any, row: int) -> dict[str, Any]:\n    if _text(tracking.cell(row, MULTI_MARKER_COLUMN).value) != MULTI_PERSON_ROW:\n        return {}\n    raw = _text(tracking.cell(row, MULTI_JSON_COLUMN).value)\n    if not raw:\n        return {}\n    try:\n        value = json.loads(raw)\n        return value if isinstance(value, dict) else {}\n    except json.JSONDecodeError:\n        return {}\n\n\ndef _select_rows(\n    tracking: Any,\n    start_row: int,\n    end_row: int,\n    *,\n    extra_blank_rows: int = 5,\n) -> list[int]:\n    existing = []\n    for row in range(start_row, end_row + 1):\n        full_name = " ".join(\n            part for part in (\n                _text(tracking.cell(row, 3).value),\n                _text(tracking.cell(row, 4).value),\n            )\n            if part\n        ).strip()\n        if full_name:\n            existing.append(row)\n\n    if existing:\n        last = max(existing)\n        stop = min(end_row, last + extra_blank_rows)\n        return list(range(start_row, stop + 1))\n\n    return list(range(start_row, min(end_row, start_row + extra_blank_rows - 1) + 1))\n\n\ndef _form_specs(workbook: Any, tracking: Any, meta: dict[str, str]) -> list[dict[str, Any]]:\n    version = meta.get("VERSION", "")\n    result: list[dict[str, Any]] = []\n\n    if version == "PCGDMN_NHIEU_HO_V1":\n        active_start: int | None = None\n        active_payload: dict[str, Any] | None = None\n\n        for row in range(1, tracking.max_row + 1):\n            marker = _text(tracking.cell(row, MULTI_MARKER_COLUMN).value)\n            if marker == MULTI_BLOCK_START:\n                active_start = row\n                raw = _text(tracking.cell(row, MULTI_JSON_COLUMN).value)\n                try:\n                    parsed = json.loads(raw) if raw else {}\n                except json.JSONDecodeError:\n                    parsed = {}\n                active_payload = parsed if isinstance(parsed, dict) else {}\n            elif marker == MULTI_BLOCK_END and active_start is not None:\n                block_end = row - 1\n                payload = active_payload or {}\n                data_start = active_start + DATA_START_ROW - 1\n                selected_rows = _select_rows(\n                    tracking,\n                    data_start,\n                    block_end,\n                )\n                result.append({\n                    "form_id": payload.get("form_id"),\n                    "household_id": payload.get("household_id"),\n                    "form_number": payload.get("form_number") or "",\n                    "household_code": payload.get("household_code") or "",\n                    "current_year_column": int(payload.get("current_year_column") or 18),\n                    "school_year_code": meta.get("SCHOOL_YEAR_CODE", ""),\n                    "commune_name": meta.get("COMMUNE_NAME", ""),\n                    "block_start": active_start,\n                    "tracking_rows": selected_rows,\n                    "head_name": _extract_head(tracking.cell(active_start + 2, 1).value),\n                    "hamlet": _extract_hamlet(tracking.cell(active_start + 1, 1).value),\n                    "address": _extract_address(tracking.cell(active_start + 2, 11).value),\n                    "row_payloads": {\n                        item: _payload_from_marker(tracking, item)\n                        for item in selected_rows\n                    },\n                })\n                active_start = None\n                active_payload = None\n        return result\n\n    if version != "PCGDMN_HO_V1":\n        return result\n\n    current_year_column = int(meta.get("CURRENT_YEAR_COLUMN") or 18)\n    row_meta = _single_row_meta(workbook[META_SHEET_NAME])\n    selected_rows = _select_rows(\n        tracking,\n        DATA_START_ROW,\n        DATA_END_ROW,\n    )\n    result.append({\n        "form_id": meta.get("FORM_ID"),\n        "household_id": meta.get("HOUSEHOLD_ID"),\n        "form_number": meta.get("FORM_NUMBER", ""),\n        "household_code": meta.get("HOUSEHOLD_CODE", ""),\n        "current_year_column": current_year_column,\n        "school_year_code": meta.get("SCHOOL_YEAR_CODE", ""),\n        "commune_name": meta.get("COMMUNE_NAME", ""),\n        "block_start": 1,\n        "tracking_rows": selected_rows,\n        "head_name": meta.get("HOUSEHOLD_HEAD_NAME", "") or _extract_head(tracking["A3"].value),\n        "hamlet": _extract_hamlet(tracking["A2"].value),\n        "address": _extract_address(tracking["K3"].value),\n        "row_payloads": {\n            item: row_meta.get(item, {})\n            for item in selected_rows\n        },\n    })\n    return result\n\n\ndef _apply_all_borders(ws: Any, min_row: int, max_row: int, min_col: int, max_col: int) -> None:\n    thin = Side(style="thin", color="000000")\n    border = Border(left=thin, right=thin, top=thin, bottom=thin)\n    for row in ws.iter_rows(\n        min_row=min_row,\n        max_row=max_row,\n        min_col=min_col,\n        max_col=max_col,\n    ):\n        for cell in row:\n            cell.border = border\n\n\ndef _fill_header(\n    ws: Any,\n    row: int,\n    *,\n    spec: dict[str, Any],\n    page_index: int,\n    page_total: int,\n) -> int:\n    start = row\n\n    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)\n    ws.cell(row, 1).value = f"Phường/Xã (1): {spec.get(\'commune_name\') or \'\'}"\n    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=3)\n    ws.cell(row + 1, 1).value = f"Tổ, thôn, xóm (2): {spec.get(\'hamlet\') or \'\'}"\n    ws.merge_cells(start_row=row + 2, start_column=1, end_row=row + 2, end_column=3)\n    ws.cell(row + 2, 1).value = f"Địa chỉ (3): {spec.get(\'address\') or \'\'}"\n\n    ws.merge_cells(start_row=row, start_column=4, end_row=row + 2, end_column=8)\n    title = ws.cell(row, 4)\n    title.value = (\n        "PHIẾU ĐIỀU TRA PHỔ CẬP GIÁO DỤC - CHỐNG MÙ CHỮ\\n"\n        f"Năm học {spec.get(\'school_year_code\') or \'\'}"\n        + (f"\\nTrang {page_index}/{page_total}" if page_total > 1 else "")\n    )\n    title.font = Font(name="Times New Roman", size=12, bold=True)\n    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)\n\n    ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=11)\n    ws.cell(row, 9).value = f"Họ và tên chủ hộ (4): {spec.get(\'head_name\') or \'\'}"\n    ws.merge_cells(start_row=row + 1, start_column=9, end_row=row + 1, end_column=11)\n    ws.cell(row + 1, 9).value = f"Số phiếu (5): {spec.get(\'form_number\') or \'\'}"\n    ws.merge_cells(start_row=row + 2, start_column=9, end_row=row + 2, end_column=11)\n    ws.cell(row + 2, 9).value = "Diện cư trú (6): □ Thường trú   □ Tạm trú   Điện thoại (7):"\n\n    for r in range(row, row + 3):\n        ws.row_dimensions[r].height = 17\n        for c in range(1, 12):\n            cell = ws.cell(r, c)\n            if cell.coordinate not in ws.merged_cells:\n                pass\n            cell.font = Font(name="Times New Roman", size=7.5, bold=(c in {1, 9}))\n            cell.alignment = Alignment(vertical="center", wrap_text=True)\n\n    row += 3\n    headers = [\n        "STT",\n        "HỌ VÀ TÊN ĐỐI TƯỢNG\\n(Lớn tuổi ghi trước)",\n        "TÊN LỚP ĐANG HỌC\\n(Theo năm học)",\n        "TÊN TRƯỜNG - HUYỆN ĐANG HỌC\\n(Tương ứng từng năm học)",\n        "TỐT NGHIỆP / HOÀN THÀNH\\nMN → THPT · Bổ túc · TN nghề",\n        "HỌC XONG\\nLớp / Năm",\n        "BỎ HỌC\\nLớp / Năm",\n        "MÙ CHỮ\\nHọc XMC / CN XMC / Tái mù",\n        "KHUYẾT TẬT",\n        "CHUYỂN ĐẾN, CHUYỂN ĐI, CHẾT",\n        "GHI CHÚ",\n    ]\n    for c, value in enumerate(headers, start=1):\n        cell = ws.cell(row, c, value)\n        cell.font = Font(name="Times New Roman", size=6.5, bold=True)\n        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)\n    ws.row_dimensions[row].height = 33\n    _apply_all_borders(ws, start, row, 1, 11)\n    return row + 1\n\n\ndef _history_rows(\n    tracking: Any,\n    spec: dict[str, Any],\n    tracking_row: int,\n) -> list[tuple[str, str]]:\n    block_start = int(spec.get("block_start") or 1)\n    current_col = int(spec.get("current_year_column") or 18)\n\n    values: list[tuple[str, str]] = []\n    current_label = spec.get("school_year_code") or _year_label(\n        tracking,\n        block_start,\n        current_col,\n    )\n    current_value = _text(tracking.cell(tracking_row, current_col).value)\n    values.append((str(current_label), current_value))\n\n    for col in range(current_col - 1, max(11, current_col - 6), -1):\n        label = _year_label(tracking, block_start, col)\n        value = _text(tracking.cell(tracking_row, col).value)\n        if label or value:\n            values.append((label, value))\n        if len(values) >= 7:\n            break\n\n    while len(values) < 7:\n        values.append(("20__-20__", ""))\n\n    return values[:7]\n\n\ndef _fill_person_slot(\n    ws: Any,\n    start_row: int,\n    *,\n    tracking: Any,\n    spec: dict[str, Any],\n    tracking_row: int,\n    slot_number: int,\n    map_sheet: Any,\n    map_row: int,\n) -> int:\n    end_row = start_row + 6\n\n    payload = dict(spec.get("row_payloads", {}).get(tracking_row, {}) or {})\n    full_name = " ".join(\n        part for part in (\n            _text(tracking.cell(tracking_row, 3).value),\n            _text(tracking.cell(tracking_row, 4).value),\n        )\n        if part\n    ).strip()\n    dob = _text(tracking.cell(tracking_row, 5).value)\n    gender = _text(tracking.cell(tracking_row, 6).value)\n    ethnic = _text(tracking.cell(tracking_row, 9).value)\n    relationship = _text(tracking.cell(tracking_row, 10).value)\n    parent_name = _text(tracking.cell(tracking_row, 11).value)\n    school_name = _text(tracking.cell(tracking_row, 19).value)\n    notes = _text(tracking.cell(tracking_row, 20).value)\n    person_code = _text(payload.get("person_code"))\n\n    ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)\n    ws.cell(start_row, 1).value = slot_number\n    ws.cell(start_row, 1).alignment = Alignment(horizontal="center", vertical="center")\n\n    person_lines = [\n        full_name,\n        f"QH chủ hộ: {relationship}" if relationship else "QH chủ hộ:",\n        f"Sinh: {dob}" if dob else "Sinh:",\n        (\n            f"Giới tính: {gender or \'\'} | DT: {ethnic or \'\'}"\n            if (gender or ethnic)\n            else "Giới tính: | DT:"\n        ),\n        f"Cha, mẹ/Ng. đỡ đầu: {parent_name}" if parent_name else "Cha, mẹ/Ng. đỡ đầu:",\n        "TG:",\n        f"Mã nội bộ: {person_code}" if person_code else "Mã nội bộ:",\n    ]\n\n    history = _history_rows(tracking, spec, tracking_row)\n\n    for offset in range(7):\n        row = start_row + offset\n        b = ws.cell(row, 2)\n        b.value = person_lines[offset]\n        b.font = Font(\n            name="Times New Roman",\n            size=7.0,\n            bold=(offset == 0),\n        )\n        b.alignment = Alignment(vertical="center", wrap_text=True)\n\n        year_label, class_value = history[offset]\n        c = ws.cell(row, 3)\n        c.value = f"{year_label}   {class_value}".rstrip()\n        c.font = Font(name="Times New Roman", size=6.4)\n        c.alignment = Alignment(vertical="center", wrap_text=True)\n\n        d = ws.cell(row, 4)\n        if offset == 0:\n            d.value = f"{spec.get(\'school_year_code\') or year_label}   {school_name}".rstrip()\n        else:\n            d.value = f"{year_label}"\n        d.font = Font(name="Times New Roman", size=6.4)\n        d.alignment = Alignment(vertical="center", wrap_text=True)\n        ws.row_dimensions[row].height = 8.7\n\n    for col in range(5, 12):\n        ws.merge_cells(\n            start_row=start_row,\n            start_column=col,\n            end_row=end_row,\n            end_column=col,\n        )\n        cell = ws.cell(start_row, col)\n        cell.font = Font(name="Times New Roman", size=6.5)\n        cell.alignment = Alignment(vertical="top", wrap_text=True)\n\n    ws.cell(start_row, 11).value = notes\n\n    _apply_all_borders(ws, start_row, end_row, 1, 11)\n\n    headers = [\n        "VERSION",\n        "VISIBLE_SHEET",\n        "VISIBLE_START_ROW",\n        "TRACKING_ROW",\n        "FORM_ID",\n        "HOUSEHOLD_ID",\n        "PERSON_ID",\n        "PERSON_CODE",\n        "CURRENT_YEAR_COLUMN",\n        "SCHOOL_YEAR_CODE",\n        "HEADER_ROW",\n    ]\n    if map_row == 2:\n        for col, name in enumerate(headers, start=1):\n            map_sheet.cell(1, col).value = name\n\n    values = [\n        FIELD_VERSION,\n        FIELD_SHEET_NAME,\n        start_row,\n        tracking_row,\n        spec.get("form_id"),\n        spec.get("household_id"),\n        payload.get("person_id"),\n        payload.get("person_code"),\n        spec.get("current_year_column"),\n        spec.get("school_year_code"),\n        spec.get("visible_header_row"),\n    ]\n    for col, value in enumerate(values, start=1):\n        map_sheet.cell(map_row, col).value = value\n\n    return end_row + 1\n\n\ndef _fill_signature(ws: Any, row: int, head_name: str) -> int:\n    ranges = [(1, 2), (3, 5), (6, 8), (9, 11)]\n    labels = [\n        f"HỌ, TÊN CHỦ HỘ\\n{head_name or \'\'}\\n(Ký, ghi rõ họ tên)",\n        "CÁN BỘ, NHÂN VIÊN ĐIỀU TRA 1\\n(Ký, ghi rõ họ tên)",\n        "CÁN BỘ, NHÂN VIÊN ĐIỀU TRA 2\\n(Ký, ghi rõ họ tên)",\n        "XÁC NHẬN CỦA UBND PHƯỜNG/XÃ\\n(Ký tên, đóng dấu)",\n    ]\n    for (c1, c2), value in zip(ranges, labels):\n        ws.merge_cells(start_row=row, start_column=c1, end_row=row + 2, end_column=c2)\n        cell = ws.cell(row, c1)\n        cell.value = value\n        cell.font = Font(name="Times New Roman", size=7, bold=True)\n        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)\n    for r in range(row, row + 3):\n        ws.row_dimensions[r].height = 14\n    return row + 3\n\n\ndef add_field_form_view_to_bytes(content: bytes) -> bytes:\n    workbook = load_workbook(BytesIO(content))\n    tracking = _find_tracking(workbook)\n    if META_SHEET_NAME not in workbook.sheetnames:\n        raise ValueError("File thiếu trang kỹ thuật ẩn.")\n\n    meta = _meta_values(workbook[META_SHEET_NAME])\n    specs = _form_specs(workbook, tracking, meta)\n    if not specs:\n        raise ValueError("Không xác định được phiếu để dựng mẫu Excel điều tra.")\n\n    for name in (FIELD_SHEET_NAME, FIELD_MAP_SHEET_NAME):\n        if name in workbook.sheetnames:\n            workbook.remove(workbook[name])\n\n    ws = workbook.create_sheet(FIELD_SHEET_NAME, 0)\n    map_sheet = workbook.create_sheet(FIELD_MAP_SHEET_NAME)\n    map_sheet.sheet_state = "veryHidden"\n\n    widths = {\n        "A": 4.5,\n        "B": 29,\n        "C": 19,\n        "D": 28,\n        "E": 18,\n        "F": 12,\n        "G": 12,\n        "H": 14,\n        "I": 13,\n        "J": 18,\n        "K": 19,\n    }\n    for letter, width in widths.items():\n        ws.column_dimensions[letter].width = width\n\n    current_row = 1\n    map_row = 2\n\n    for spec_index, spec in enumerate(specs, start=1):\n        tracking_rows = list(spec.get("tracking_rows") or [])\n        pages = [\n            tracking_rows[i:i + 7]\n            for i in range(0, len(tracking_rows), 7)\n        ] or [[]]\n\n        for page_index, page_rows in enumerate(pages, start=1):\n            header_row = current_row\n            spec["visible_header_row"] = header_row\n            current_row = _fill_header(\n                ws,\n                current_row,\n                spec=spec,\n                page_index=page_index,\n                page_total=len(pages),\n            )\n\n            for slot_index in range(7):\n                if slot_index < len(page_rows):\n                    tracking_row = page_rows[slot_index]\n                    current_row = _fill_person_slot(\n                        ws,\n                        current_row,\n                        tracking=tracking,\n                        spec=spec,\n                        tracking_row=tracking_row,\n                        slot_number=(page_index - 1) * 7 + slot_index + 1,\n                        map_sheet=map_sheet,\n                        map_row=map_row,\n                    )\n                    map_row += 1\n                else:\n                    # Vẫn tạo một dòng trắng đúng mẫu để in, nhưng không có mapping nhập.\n                    blank_start = current_row\n                    blank_end = blank_start + 6\n                    ws.merge_cells(\n                        start_row=blank_start,\n                        start_column=1,\n                        end_row=blank_end,\n                        end_column=1,\n                    )\n                    ws.cell(blank_start, 1).value = (\n                        (page_index - 1) * 7 + slot_index + 1\n                    )\n                    for offset in range(7):\n                        r = blank_start + offset\n                        ws.cell(r, 2).value = (\n                            "" if offset == 0 else\n                            (\n                                "QH chủ hộ:" if offset == 1 else\n                                "Sinh:" if offset == 2 else\n                                "Giới tính: | DT:" if offset == 3 else\n                                "Cha, mẹ/Ng. đỡ đầu:" if offset == 4 else\n                                "TG:" if offset == 5 else\n                                "Mã nội bộ:"\n                            )\n                        )\n                        ws.cell(r, 3).value = "20__-20__"\n                        ws.cell(r, 4).value = "20__-20__"\n                        ws.row_dimensions[r].height = 8.7\n                    for col in range(5, 12):\n                        ws.merge_cells(\n                            start_row=blank_start,\n                            start_column=col,\n                            end_row=blank_end,\n                            end_column=col,\n                        )\n                    _apply_all_borders(ws, blank_start, blank_end, 1, 11)\n                    current_row = blank_end + 1\n\n            if page_index == len(pages):\n                current_row = _fill_signature(\n                    ws,\n                    current_row,\n                    str(spec.get("head_name") or ""),\n                )\n\n            footer_row = current_row\n            ws.merge_cells(\n                start_row=footer_row,\n                start_column=1,\n                end_row=footer_row,\n                end_column=11,\n            )\n            ws.cell(footer_row, 1).value = (\n                f"{meta.get(\'BATCH_CODE\') or \'\'} · "\n                f"{spec.get(\'form_number\') or \'\'} · "\n                f"{spec.get(\'household_code\') or \'\'}"\n            )\n            ws.cell(footer_row, 1).font = Font(\n                name="Times New Roman",\n                size=6.5,\n                italic=True,\n            )\n            ws.cell(footer_row, 1).alignment = Alignment(horizontal="left")\n            current_row += 1\n\n            ws.row_breaks.append(Break(id=current_row - 1))\n            current_row += 1\n\n    ws.sheet_view.showGridLines = False\n    ws.freeze_panes = None\n    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE\n    ws.page_setup.paperSize = ws.PAPERSIZE_A4\n    ws.page_setup.fitToWidth = 1\n    ws.page_setup.fitToHeight = 0\n    ws.sheet_properties.pageSetUpPr.fitToPage = True\n    ws.page_margins = PageMargins(\n        left=0.12,\n        right=0.12,\n        top=0.18,\n        bottom=0.18,\n        header=0.0,\n        footer=0.0,\n    )\n    ws.print_options.horizontalCentered = True\n    ws.print_area = f"A1:K{max(1, current_row - 1)}"\n\n    tracking.sheet_state = "veryHidden"\n    workbook[META_SHEET_NAME].sheet_state = "veryHidden"\n    ws.sheet_state = "visible"\n    workbook.active = 0\n\n    output = BytesIO()\n    workbook.save(output)\n    return output.getvalue()\n\n\ndef _map_headers(sheet: Any) -> dict[str, int]:\n    result = {}\n    for col in range(1, sheet.max_column + 1):\n        name = _text(sheet.cell(1, col).value)\n        if name:\n            result[name] = col\n    return result\n\n\ndef _parse_current_value(raw: Any, year_code: str) -> str:\n    text = _text(raw)\n    if not text:\n        return ""\n    if year_code and text.startswith(year_code):\n        return text[len(year_code):].strip()\n    match = re.match(r"^\\d{4}-\\d{4}\\s+(.*)$", text)\n    if match:\n        return match.group(1).strip()\n    return text\n\n\ndef _parse_person_fields(ws: Any, start_row: int, year_code: str) -> dict[str, str]:\n    full_name = _text(ws.cell(start_row, 2).value)\n    relationship = _line_value(\n        ws.cell(start_row + 1, 2).value,\n        ("QH chủ hộ", "Quan hệ với chủ hộ"),\n    )\n    birth = _line_value(\n        ws.cell(start_row + 2, 2).value,\n        ("Sinh", "Ngày sinh"),\n    )\n    gender, ethnic = _gender_ethnic(ws.cell(start_row + 3, 2).value)\n    parent = _line_value(\n        ws.cell(start_row + 4, 2).value,\n        ("Cha, mẹ/Ng. đỡ đầu", "Cha/mẹ/Ng. đỡ đầu", "Cha/mẹ/người đỡ đầu"),\n    )\n    program_class = _parse_current_value(\n        ws.cell(start_row, 3).value,\n        year_code,\n    )\n    school = _parse_current_value(\n        ws.cell(start_row, 4).value,\n        year_code,\n    )\n    notes = _text(ws.cell(start_row, 11).value)\n\n    return {\n        "full_name": full_name,\n        "relationship": relationship,\n        "birth": birth,\n        "gender": gender,\n        "ethnic": ethnic,\n        "parent": parent,\n        "program_class": program_class,\n        "school": school,\n        "notes": notes,\n    }\n\n\ndef sync_field_form_to_tracking(workbook: Any) -> bool:\n    if FIELD_MAP_SHEET_NAME not in workbook.sheetnames:\n        return False\n    if FIELD_SHEET_NAME not in workbook.sheetnames:\n        raise ValueError(\n            "File có thông tin ánh xạ Phiếu điều tra nhưng trang Phiếu điều tra đã bị xóa/đổi tên."\n        )\n\n    tracking = _find_tracking(workbook)\n    map_sheet = workbook[FIELD_MAP_SHEET_NAME]\n    visible = workbook[FIELD_SHEET_NAME]\n    headers = _map_headers(map_sheet)\n\n    required = {\n        "VERSION",\n        "VISIBLE_START_ROW",\n        "TRACKING_ROW",\n        "CURRENT_YEAR_COLUMN",\n        "SCHOOL_YEAR_CODE",\n        "HEADER_ROW",\n    }\n    if not required.issubset(headers):\n        raise ValueError("Trang ánh xạ kỹ thuật của Phiếu điều tra không đúng phiên bản.")\n\n    for row in range(2, map_sheet.max_row + 1):\n        if _text(map_sheet.cell(row, headers["VERSION"]).value) != FIELD_VERSION:\n            continue\n\n        try:\n            visible_start = int(\n                map_sheet.cell(row, headers["VISIBLE_START_ROW"]).value\n            )\n            tracking_row = int(\n                map_sheet.cell(row, headers["TRACKING_ROW"]).value\n            )\n            current_year_col = int(\n                map_sheet.cell(row, headers["CURRENT_YEAR_COLUMN"]).value\n            )\n            header_row = int(\n                map_sheet.cell(row, headers["HEADER_ROW"]).value\n            )\n        except (TypeError, ValueError) as exc:\n            raise ValueError(\n                f"Ánh xạ kỹ thuật tại dòng {row} bị hỏng."\n            ) from exc\n\n        year_code = _text(\n            map_sheet.cell(row, headers["SCHOOL_YEAR_CODE"]).value\n        )\n        fields = _parse_person_fields(\n            visible,\n            visible_start,\n            year_code,\n        )\n\n        # Đồng bộ thôn/xóm từ đầu phiếu vào từng dòng như bộ importer cũ.\n        hamlet_text = _extract_hamlet(visible.cell(header_row + 1, 1).value)\n\n        left, right = _split_name(fields["full_name"])\n        tracking.cell(tracking_row, 3).value = left\n        tracking.cell(tracking_row, 4).value = right\n        tracking.cell(tracking_row, 5).value = fields["birth"]\n        tracking.cell(tracking_row, 6).value = fields["gender"]\n        if hamlet_text:\n            tracking.cell(tracking_row, 7).value = hamlet_text\n        tracking.cell(tracking_row, 9).value = fields["ethnic"]\n        tracking.cell(tracking_row, 10).value = fields["relationship"]\n        tracking.cell(tracking_row, 11).value = fields["parent"]\n        tracking.cell(tracking_row, current_year_col).value = fields[\n            "program_class"\n        ]\n        tracking.cell(tracking_row, 19).value = fields["school"]\n        tracking.cell(tracking_row, 20).value = fields["notes"]\n\n    return True\n'
EXPORT_NEEDLE = '        action = "XUAT_HANG_LOAT_XLSX"\n\n    _log_exchange(\n'
EXPORT_REPLACEMENT = '        action = "XUAT_HANG_LOAT_XLSX"\n\n    # === BAI_13B_12_V2_4_1_FIELD_EXCEL_EXPORT ===\n    try:\n        from app.household_field_excel_v240 import (\n            add_field_form_view_to_bytes,\n        )\n        content = add_field_form_view_to_bytes(content)\n    except Exception as exc:\n        raise HouseholdExcelError(\n            "Không dựng được trang Phiếu điều tra A4 trong file Excel."\n        ) from exc\n\n    _log_exchange(\n'
IMPORT_NEEDLE = '    try:\n        workbook = load_workbook(BytesIO(content), data_only=False)\n    except Exception as exc:\n        raise HouseholdExcelError("Không đọc được file Excel hoặc file đã bị hỏng.") from exc\n    tracking = _find_tracking_sheet(workbook)\n'
IMPORT_REPLACEMENT = '    try:\n        workbook = load_workbook(BytesIO(content), data_only=False)\n    except Exception as exc:\n        raise HouseholdExcelError("Không đọc được file Excel hoặc file đã bị hỏng.") from exc\n\n    # === BAI_13B_12_V2_4_1_FIELD_EXCEL_IMPORT ===\n    try:\n        from app.household_field_excel_v240 import (\n            sync_field_form_to_tracking,\n        )\n        sync_field_form_to_tracking(workbook)\n    except Exception as exc:\n        raise HouseholdExcelError(\n            "Trang Phiếu điều tra trong file Excel bị sai cấu trúc hoặc đã bị đổi tên."\n        ) from exc\n\n    tracking = _find_tracking_sheet(workbook)\n'
DOC_OLD = '    """Xuất mỗi hộ thành một phiếu Excel đúng mẫu Sổ theo dõi PCGDMN."""\n'
DOC_NEW = '    """Xuất Excel điều tra theo mẫu Phiếu điều tra PCGD-XMC A4 ngang; file xuất dùng lại để nhập."""\n'
INSTRUCTION_OLD = '            <li>Giữ nguyên tên trang <strong>Sổ theo dõi PCGDMN</strong> và trang kỹ thuật ẩn.</li>\n'
INSTRUCTION_NEW = '            <li>Nhập trực tiếp trên trang <strong>Phiếu điều tra</strong>; giữ nguyên tên trang này và các trang kỹ thuật ẩn.</li>\n'
LOG_OLD = '            "Xuất/nhập phiếu điều tra một hoặc nhiều hộ bằng đúng trang Sổ theo dõi PCGDMN."\n'
LOG_NEW = '            "Xuất/nhập cùng một file Phiếu điều tra A4; dữ liệu kỹ thuật giữ ở trang ẩn."\n'


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def write_text(path: Path, value: str) -> None:
    path.write_text(value, encoding="utf-8")


def backup_file(path: Path) -> None:
    if not path.exists():
        return
    target = BACKUP / path.relative_to(PROJECT)
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, target)


def restore_file(path: Path) -> None:
    source = BACKUP / path.relative_to(PROJECT)
    if source.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, path)


def backup_db() -> None:
    src = sqlite3.connect(str(DB))
    dst = sqlite3.connect(str(DB_BACKUP))
    try:
        src.backup(dst)
        dst.commit()
    finally:
        dst.close()
        src.close()


def restore_db() -> None:
    if not DB_BACKUP.exists():
        return
    src = sqlite3.connect(str(DB_BACKUP))
    dst = sqlite3.connect(str(DB))
    try:
        src.backup(dst)
        dst.commit()
    finally:
        dst.close()
        src.close()


def db_state() -> dict:
    con = sqlite3.connect(str(DB))
    try:
        result = {
            "integrity": str(
                con.execute("PRAGMA integrity_check").fetchone()[0]
            ),
            "fk_count": len(
                con.execute("PRAGMA foreign_key_check").fetchall()
            ),
        }
        for table in (
            "survey_forms",
            "survey_people",
            "survey_person_year_records",
        ):
            result[table] = int(
                con.execute(
                    f"SELECT COUNT(*) FROM {table}"
                ).fetchone()[0]
            )
        return result
    finally:
        con.close()


def patch_once(source: str, old: str, new: str, label: str) -> str:
    count = source.count(old)
    if count != 1:
        raise RuntimeError(
            f"{label}: số marker khớp = {count}, mong đợi 1."
        )
    return source.replace(old, new, 1)


def load_helper_module():
    spec = importlib.util.spec_from_file_location(
        "household_field_excel_v240_smoke",
        HELPER,
    )
    if spec is None or spec.loader is None:
        raise RuntimeError("Không nạp được helper V2.4.1.")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def smoke_test_helper() -> None:
    module = load_helper_module()

    wb = Workbook()
    tracking = wb.active
    tracking.title = "Sổ theo dõi PCGDMN"
    tracking["A1"] = "PHIẾU ĐIỀU TRA PHỔ CẬP GIÁO DỤC"
    tracking["A2"] = (
        "Xóm/khối: Test-Xóm1; Xã/phường: Xã Test; "
        "Tỉnh: Nghệ An; Năm học: 2026-2027"
    )
    tracking["A3"] = "Họ và tên chủ hộ: TEST Chủ hộ"
    tracking["K3"] = "Địa chỉ: Test-Xóm1, số 01"

    for column, start_year in zip(range(12, 19), range(2020, 2027)):
        tracking.cell(7, column).value = start_year
        tracking.cell(8, column).value = start_year + 1

    tracking.cell(11, 3).value = "TEST Trần Văn"
    tracking.cell(11, 4).value = "A"
    tracking.cell(11, 5).value = "01/01/2010"
    tracking.cell(11, 6).value = "Nam"
    tracking.cell(11, 7).value = "Test-Xóm1"
    tracking.cell(11, 8).value = "Xã Test"
    tracking.cell(11, 9).value = "Kinh"
    tracking.cell(11, 10).value = "Con"
    tracking.cell(11, 11).value = "TEST Cha/Mẹ"
    tracking.cell(11, 18).value = "10"
    tracking.cell(11, 19).value = "Trường THPT Test"
    tracking.cell(11, 20).value = "Ghi chú test"

    meta = wb.create_sheet("_PCGDMN_META")
    values = [
        ("VERSION", "PCGDMN_HO_V1"),
        ("BATCH_ID", 1),
        ("BATCH_CODE", "DT-TEST"),
        ("FORM_ID", 1),
        ("FORM_NUMBER", "TEST-001"),
        ("HOUSEHOLD_ID", 1),
        ("HOUSEHOLD_CODE", "HO-TEST"),
        ("HOUSEHOLD_HEAD_NAME", "TEST Chủ hộ"),
        ("COMMUNE_NAME", "Xã Test"),
        ("SCHOOL_YEAR_ID", 2),
        ("SCHOOL_YEAR_CODE", "2026-2027"),
        ("CURRENT_YEAR_COLUMN", 18),
    ]
    for row, (key, value) in enumerate(values, start=1):
        meta.cell(row, 1).value = key
        meta.cell(row, 2).value = value

    for col, header in enumerate(
        ("visible_row", "person_id", "person_code"),
        start=1,
    ):
        meta.cell(20, col).value = header
    meta.cell(21, 1).value = 11
    meta.cell(21, 2).value = 99
    meta.cell(21, 3).value = "DT-TEST-99"
    meta.sheet_state = "veryHidden"

    raw = BytesIO()
    wb.save(raw)

    exported = module.add_field_form_view_to_bytes(raw.getvalue())
    wb2 = load_workbook(BytesIO(exported))

    if "Phiếu điều tra" not in wb2.sheetnames:
        raise RuntimeError("Smoke test: thiếu trang Phiếu điều tra.")
    if wb2["Sổ theo dõi PCGDMN"].sheet_state != "veryHidden":
        raise RuntimeError("Smoke test: trang kỹ thuật chưa được ẩn.")

    visible = wb2["Phiếu điều tra"]
    visible["B5"] = "TEST Trần Văn A Sửa"
    visible["C5"] = "2026-2027   11"
    visible["D5"] = "2026-2027   Trường THPT Mới"
    visible["K5"] = "Ghi chú mới"

    if not module.sync_field_form_to_tracking(wb2):
        raise RuntimeError("Smoke test: không đồng bộ ngược.")

    technical = wb2["Sổ theo dõi PCGDMN"]
    if technical.cell(11, 18).value != "11":
        raise RuntimeError("Smoke test: lớp mới chưa đồng bộ.")
    if technical.cell(11, 19).value != "Trường THPT Mới":
        raise RuntimeError("Smoke test: trường mới chưa đồng bộ.")
    if technical.cell(11, 20).value != "Ghi chú mới":
        raise RuntimeError("Smoke test: ghi chú mới chưa đồng bộ.")


def verify_source() -> None:
    exchange = read_text(EXCHANGE)
    surveys = read_text(SURVEYS)
    html = read_text(IMPORT_TEMPLATE)
    helper = read_text(HELPER)

    for token in (
        MARKER_EXPORT,
        MARKER_IMPORT,
        "add_field_form_view_to_bytes",
        "sync_field_form_to_tracking",
    ):
        if token not in exchange:
            raise RuntimeError("Verifier exchange thiếu: " + token)

    if "mẫu Phiếu điều tra PCGD-XMC A4 ngang" not in surveys:
        raise RuntimeError("Verifier surveys thiếu docstring V2.4.1.")

    if "Nhập trực tiếp trên trang" not in html:
        raise RuntimeError("Verifier template nhập Excel chưa đổi hướng dẫn.")

    if MARKER_HELPER not in helper:
        raise RuntimeError("Verifier helper thiếu marker.")

    ast.parse(exchange)
    ast.parse(surveys)
    ast.parse(helper)
    py_compile.compile(str(EXCHANGE), doraise=True)
    py_compile.compile(str(SURVEYS), doraise=True)
    py_compile.compile(str(HELPER), doraise=True)
    Environment().parse(html)
    smoke_test_helper()


def main() -> int:
    print("=" * 122)
    print(
        "BÀI 13B-12 V2.4.1 - "
        "EXCEL ĐIỀU TRA GIỐNG PHIẾU + CHÍNH FILE ĐÓ NHẬP LẠI"
    )
    print("=" * 122)

    for path in (DB, EXCHANGE, SURVEYS, IMPORT_TEMPLATE):
        if not path.exists():
            print("DỪNG AN TOÀN: thiếu", path)
            return 2

    before = db_state()
    print("integrity_check:", before["integrity"])
    print("foreign_key_check:", before["fk_count"], "lỗi")

    if before["integrity"].lower() != "ok" or before["fk_count"] != 0:
        print("DỪNG: database chưa đạt kiểm tra.")
        return 3

    exchange = read_text(EXCHANGE)
    surveys = read_text(SURVEYS)
    html = read_text(IMPORT_TEMPLATE)

    if (
        MARKER_EXPORT in exchange
        and MARKER_IMPORT in exchange
        and HELPER.exists()
        and MARKER_HELPER in read_text(HELPER)
    ):
        print("V2.4.1 đã cài. Không cài lặp.")
        return 0

    try:
        new_exchange = patch_once(
            exchange, EXPORT_NEEDLE, EXPORT_REPLACEMENT, "Patch export"
        )
        new_exchange = patch_once(
            new_exchange, IMPORT_NEEDLE, IMPORT_REPLACEMENT, "Patch import"
        )
        new_exchange = patch_once(
            new_exchange, LOG_OLD, LOG_NEW, "Patch nhật ký"
        )
        new_surveys = patch_once(
            surveys, DOC_OLD, DOC_NEW, "Patch docstring route"
        )
        new_html = patch_once(
            html,
            INSTRUCTION_OLD,
            INSTRUCTION_NEW,
            "Patch hướng dẫn nhập",
        )

        ast.parse(new_exchange)
        ast.parse(new_surveys)
        ast.parse(HELPER_SOURCE)
        Environment().parse(new_html)

    except Exception as exc:
        print()
        print("DỪNG AN TOÀN TRƯỚC KHI GHI FILE:")
        print(type(exc).__name__ + ":", exc)
        return 4

    EXPORTS.mkdir(parents=True, exist_ok=True)
    BACKUP.mkdir(parents=True, exist_ok=False)
    helper_existed = HELPER.exists()

    for path in (EXCHANGE, SURVEYS, IMPORT_TEMPLATE, HELPER):
        backup_file(path)
    backup_db()
    print("Backup:", BACKUP)

    try:
        write_text(EXCHANGE, new_exchange)
        write_text(SURVEYS, new_surveys)
        write_text(IMPORT_TEMPLATE, new_html)
        write_text(HELPER, HELPER_SOURCE)

        verify_source()

        after = db_state()
        if after["integrity"].lower() != "ok":
            raise RuntimeError("integrity_check không đạt sau cài.")
        if after["fk_count"] != 0:
            raise RuntimeError("foreign_key_check có lỗi sau cài.")

        for key in (
            "survey_forms",
            "survey_people",
            "survey_person_year_records",
        ):
            if after[key] != before[key]:
                raise RuntimeError(
                    f"Số bản ghi {key} thay đổi ngoài dự kiến."
                )

        for cache in APP.rglob("__pycache__"):
            if cache.is_dir():
                shutil.rmtree(cache, ignore_errors=True)

    except Exception as exc:
        print()
        print("CÓ LỖI - ĐANG ROLLBACK SOURCE + DATABASE...")
        print(type(exc).__name__ + ":", exc)
        restore_file(EXCHANGE)
        restore_file(SURVEYS)
        restore_file(IMPORT_TEMPLATE)
        if helper_existed:
            restore_file(HELPER)
        else:
            HELPER.unlink(missing_ok=True)
        restore_db()
        print("ĐÃ KHÔI PHỤC.")
        return 9

    REPORT.write_text(
        "\n".join(
            [
                "=" * 122,
                "BÁO CÁO CÀI BÀI 13B-12 V2.4.1",
                "=" * 122,
                f"Thời gian: {datetime.now():%d/%m/%Y %H:%M:%S}",
                "",
                "ĐÃ SỬA:",
                "- 2.2.3 mở ra trang Excel đầu tiên tên 'Phiếu điều tra'.",
                "- Hình thức A4 ngang bám Phiếu điều tra PCGD-XMC 2.2.2.",
                "- Sổ theo dõi PCGDMN cũ được giữ nguyên nhưng chuyển veryHidden.",
                "- _PCGDMN_META và _PCGDMN_FIELD_MAP là trang kỹ thuật ẩn.",
                "- Người dùng sửa trực tiếp trên trang Phiếu điều tra.",
                "- Khi upload 2.2.4, hệ thống đồng bộ các ô đã sửa về cấu trúc import cũ trước khi kiểm tra.",
                "- File cũ trước V2.4.1 vẫn nhập được theo cơ chế cũ.",
                "",
                "PHẠM VI NHẬP EXCEL GIỮ NGUYÊN SO VỚI BỘ IMPORT HIỆN CÓ:",
                "- Họ tên, ngày sinh, giới tính, dân tộc, quan hệ chủ hộ, cha/mẹ/người đỡ đầu.",
                "- Thôn/xóm, lớp/chương trình năm hiện tại, trường hiện tại, ghi chú/biến động.",
                "- Có 5 dòng dự phòng để thêm thành viên phát sinh theo cơ chế cũ.",
                "",
                "AN TOÀN:",
                "- Không thay đổi dữ liệu lúc cài.",
                "- Smoke test xuất -> sửa trang Phiếu điều tra -> đồng bộ ngược đã đạt.",
                f"integrity_check: {after['integrity']}",
                f"foreign_key_check: {after['fk_count']} lỗi",
                f"Backup: {BACKUP}",
            ]
        )
        + "\n",
        encoding="utf-8-sig",
    )

    print()
    print("=" * 122)
    print("CÀI ĐẶT THÀNH CÔNG BÀI 13B-12 V2.4.1")
    print("=" * 122)
    print(" - Excel điều tra đã chuyển sang mẫu Phiếu điều tra A4 ngang.")
    print(" - Chính file xuất đó dùng lại tại 2.2.4 để nhập.")
    print(" - Import cũ vẫn được giữ làm lớp kỹ thuật ẩn.")
    print("Báo cáo:", REPORT)
    print("Backup:", BACKUP)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
