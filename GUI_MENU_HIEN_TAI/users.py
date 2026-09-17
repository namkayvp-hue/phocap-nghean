from __future__ import annotations

import json
import math
import os
import re
import secrets
import string
import time
import unicodedata
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Annotated, Any
from urllib.parse import urlencode

from fastapi import APIRouter, BackgroundTasks, Depends, Form, HTTPException, Query, Request
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import and_, func, or_, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.database import SessionLocal, get_db
from app.models import Commune, Role, School, User
from app.permissions import (
    ADMIN_ROLE_CODES,
    ADMIN_ROLE_CODE,
    COMMUNE_ROLE_CODE,
    DEPARTMENT_ROLE_CODE,
    SCHOOL_ROLE_CODE,
    TEACHER_ROLE_CODE,
    normalize_role_code,
)
from app.security import ma_hoa_mat_khau


APP_DIR = Path(__file__).resolve().parent.parent
templates = Jinja2Templates(directory=str(APP_DIR / "templates"))
router = APIRouter(prefix="/tai-khoan", tags=["Quản lý tài khoản"])

STATUS_MESSAGES = {
    "created": "Đã tạo tài khoản mới thành công.",
    "updated": "Đã cập nhật tài khoản thành công.",
    "locked": "Đã khóa tài khoản thành công.",
    "unlocked": "Đã mở lại tài khoản thành công.",
    "duplicate": "Tên đăng nhập đã tồn tại.",
    "invalid": "Tên đăng nhập và họ tên không được để trống.",
    "invalid_password": "Mật khẩu phải có ít nhất 8 ký tự.",
    "invalid_role": "Vai trò được chọn không hợp lệ với quyền của bạn.",
    "invalid_scope": "Đơn vị được chọn không hợp lệ hoặc đã bị khóa.",
    "forbidden": "Bạn không có quyền quản lý tài khoản này.",
    "cannot_self": "Không thể tự khóa tài khoản đang đăng nhập.",
    "password_reset": "Đã đặt lại mật khẩu thành công.",
    "password_mismatch": "Hai lần nhập mật khẩu mới không giống nhau.",
    "bulk_none": "Không còn đơn vị phù hợp để cấp tài khoản đồng loạt.",
    "bulk_invalid_confirm": "Chưa xác nhận đúng thao tác cấp tài khoản đồng loạt.",
    "bulk_forbidden": "Bạn không có quyền cấp tài khoản đồng loạt.",
    "invalid_reset_password": (
        "Mật khẩu mới phải có từ 8 đến 128 ký tự và "
        "không có khoảng trắng ở đầu hoặc cuối."
    ),
}

MIN_PASSWORD_LENGTH = 8
MAX_PASSWORD_LENGTH = 128
DEFAULT_PAGE_SIZE = 20
ALLOWED_PAGE_SIZES = {20, 50, 100}
ADMIN_SECTIONS = {"tong_quan", "phong_ban", "xa", "truong", "tra_cuu"}
UNIT_ACCOUNT_EXCLUDED_PREFIXES = ("cbql.",)
BULK_CONFIRM_PHRASE = "CAP DONG LOAT"
BULK_SCOPE_FILTERED = "filtered"
BULK_SCOPE_ALL = "all"
BULK_UNIT_TYPES = {"xa", "truong"}
BULK_PREVIEW_LIMIT = 100
BULK_STATUS_PENDING = "PENDING"
BULK_STATUS_RUNNING = "RUNNING"
BULK_STATUS_COMPLETED = "COMPLETED"
BULK_STATUS_COMPLETED_WITH_ERRORS = "COMPLETED_WITH_ERRORS"
BULK_STATUS_FAILED = "FAILED"
BULK_ACTIVE_STATUSES = {BULK_STATUS_PENDING, BULK_STATUS_RUNNING}
BULK_FINISHED_STATUSES = {
    BULK_STATUS_COMPLETED,
    BULK_STATUS_COMPLETED_WITH_ERRORS,
    BULK_STATUS_FAILED,
}
BULK_EXPORT_DIR = APP_DIR.parent / "data" / "account_provision_exports"
BULK_EXPORT_DIR.mkdir(parents=True, exist_ok=True)

SECTION_ROLE_CODES: dict[str, set[str]] = {
    "phong_ban": {DEPARTMENT_ROLE_CODE},
    "xa": {COMMUNE_ROLE_CODE},
    "truong": {SCHOOL_ROLE_CODE},
}

SECTION_TITLES = {
    "tong_quan": "Truy cập nhanh",
    "phong_ban": "Tài khoản Phòng ban",
    "xa": "Tài khoản Xã/phường",
    "truong": "Tài khoản Trường mầm non",
    "tra_cuu": "Tra cứu toàn bộ tài khoản",
    "giao_vien": "Giáo viên của trường",
}


def chuan_hoa_ten_dang_nhap(value: str) -> str:
    return value.strip().lower()


def chuan_hoa_ho_ten(value: str) -> str:
    return " ".join(value.strip().split())


def chuyen_sang_id(value: str | int | None) -> int | None:
    if value is None:
        return None
    text_value = str(value).strip()
    if not text_value:
        return None
    try:
        return int(text_value)
    except ValueError:
        return None


def lay_nguoi_dung(request: Request) -> dict[str, Any]:
    return dict(request.scope.get("auth_user") or {})


def vai_tro_duoc_tao(role_code: str) -> set[str]:
    if role_code in ADMIN_ROLE_CODES:
        # Giáo viên được tạo và quản lý tại cấp trường.
        return {DEPARTMENT_ROLE_CODE, COMMUNE_ROLE_CODE, SCHOOL_ROLE_CODE}
    if role_code == SCHOOL_ROLE_CODE:
        return {TEACHER_ROLE_CODE}
    return set()


def muc_duoc_phep(actor_role: str, requested_section: str | None) -> str:
    if actor_role == SCHOOL_ROLE_CODE:
        return "giao_vien"
    if actor_role not in ADMIN_ROLE_CODES:
        return "tong_quan"
    section = (requested_section or "tong_quan").strip().lower()
    return section if section in ADMIN_SECTIONS else "tong_quan"


def url_quay_lai_hop_le(value: str | None) -> str:
    value = (value or "").strip()
    if not value.startswith("/tai-khoan") or value.startswith("//"):
        return "/tai-khoan"
    return value


def chuyen_huong_trang_thai(return_to: str | None, status: str) -> RedirectResponse:
    base_url = url_quay_lai_hop_le(return_to)
    separator = "&" if "?" in base_url else "?"
    return RedirectResponse(f"{base_url}{separator}status={status}", 303)


def tai_khoan_thuoc_pham_vi(*, actor: dict[str, Any], target: User) -> bool:
    actor_role = normalize_role_code(actor.get("role_code"))
    if actor_role in ADMIN_ROLE_CODES:
        return (
            target.role is not None
            and normalize_role_code(target.role.code)
            in {DEPARTMENT_ROLE_CODE, COMMUNE_ROLE_CODE, SCHOOL_ROLE_CODE}
        )
    if actor_role == SCHOOL_ROLE_CODE:
        return (
            target.role is not None
            and normalize_role_code(target.role.code) == TEACHER_ROLE_CODE
            and actor.get("school_id") is not None
            and target.school_id == int(actor["school_id"])
        )
    return False


def co_quyen_dat_lai_mat_khau(*, actor: dict[str, Any], target: User) -> bool:
    """
    ADMIN/SO đặt lại mật khẩu cho Phòng ban, Xã, Trường và Giáo viên.
    Tài khoản Trường đặt lại mật khẩu cho giáo viên thuộc chính trường.
    """
    if target.role is None:
        return False

    actor_role = normalize_role_code(actor.get("role_code"))
    target_role = normalize_role_code(target.role.code)

    if actor_role in ADMIN_ROLE_CODES:
        return target_role in {
            DEPARTMENT_ROLE_CODE,
            COMMUNE_ROLE_CODE,
            SCHOOL_ROLE_CODE,
            TEACHER_ROLE_CODE,
        }

    if actor_role == SCHOOL_ROLE_CODE:
        return (
            target_role == TEACHER_ROLE_CODE
            and actor.get("school_id") is not None
            and target.school_id == int(actor["school_id"])
        )

    return False


def xac_dinh_don_vi_quan_ly(
    db: Session,
    vai_tro: Role,
    commune_value: str,
    school_value: str,
    tai_khoan_hien_tai: User | None = None,
    school_id_bat_buoc: int | None = None,
) -> tuple[int | None, int | None, str | None]:
    role_code = normalize_role_code(vai_tro.code)
    commune_id = chuyen_sang_id(commune_value)
    school_id = chuyen_sang_id(school_value)

    if role_code in {*ADMIN_ROLE_CODES, DEPARTMENT_ROLE_CODE}:
        return None, None, None

    if role_code == COMMUNE_ROLE_CODE:
        if commune_id is None:
            return None, None, "invalid_scope"
        commune = db.get(Commune, commune_id)
        if commune is None:
            return None, None, "invalid_scope"
        keeping = (
            tai_khoan_hien_tai is not None
            and tai_khoan_hien_tai.commune_id == commune.id
        )
        if not commune.is_active and not keeping:
            return None, None, "invalid_scope"
        return commune.id, None, None

    if role_code in {SCHOOL_ROLE_CODE, TEACHER_ROLE_CODE}:
        if school_id_bat_buoc is not None:
            school_id = int(school_id_bat_buoc)
        if school_id is None:
            return None, None, "invalid_scope"

        school = db.get(School, school_id)
        if school is None:
            return None, None, "invalid_scope"

        keeping = (
            tai_khoan_hien_tai is not None
            and tai_khoan_hien_tai.school_id == school.id
        )
        if not school.is_active and not keeping:
            return None, None, "invalid_scope"

        commune = db.get(Commune, school.commune_id)
        if commune is None or (not commune.is_active and not keeping):
            return None, None, "invalid_scope"

        if (
            school_id_bat_buoc is None
            and commune_id is not None
            and commune_id != school.commune_id
        ):
            return None, None, "invalid_scope"

        return school.commune_id, school.id, None

    return None, None, "invalid_role"


def lay_danh_sach_vai_tro_tao(
    db: Session,
    actor_role: str,
    section: str,
) -> list[Role]:
    allowed = vai_tro_duoc_tao(actor_role)
    if not allowed:
        return []

    if actor_role in ADMIN_ROLE_CODES and section in SECTION_ROLE_CODES:
        allowed = allowed.intersection(SECTION_ROLE_CODES[section])

    return list(
        db.scalars(
            select(Role)
            .where(Role.code.in_(sorted(allowed)))
            .order_by(Role.id.asc())
        ).all()
    )


def dem_theo_vai_tro(db: Session, role_code: str) -> int:
    return int(
        db.scalar(
            select(func.count(User.id))
            .select_from(User)
            .join(Role, Role.id == User.role_id)
            .where(Role.code == role_code)
        )
        or 0
    )


def tong_hop_tai_khoan(db: Session) -> dict[str, int]:
    school_unit_conditions = [Role.code == SCHOOL_ROLE_CODE]
    for prefix in UNIT_ACCOUNT_EXCLUDED_PREFIXES:
        school_unit_conditions.append(
            ~func.lower(User.username).like(f"{prefix.lower()}%")
        )

    return {
        "phong_ban": dem_theo_vai_tro(db, DEPARTMENT_ROLE_CODE),
        "xa": dem_theo_vai_tro(db, COMMUNE_ROLE_CODE),
        "truong": int(
            db.scalar(
                select(func.count(User.id))
                .select_from(User)
                .join(Role, Role.id == User.role_id)
                .where(*school_unit_conditions)
            )
            or 0
        ),
        "giao_vien": dem_theo_vai_tro(db, TEACHER_ROLE_CODE),
        "tat_ca": int(db.scalar(select(func.count(User.id))) or 0),
        "tong_xa": int(
            db.scalar(
                select(func.count(Commune.id)).where(Commune.is_active.is_(True))
            )
            or 0
        ),
        "tong_truong": int(
            db.scalar(
                select(func.count(School.id)).where(School.is_active.is_(True))
            )
            or 0
        ),
    }


def _khop_tu_khoa(tu_khoa: str, *gia_tri: Any) -> bool:
    if not tu_khoa:
        return True
    needle = tu_khoa.casefold()
    return any(needle in str(value or "").casefold() for value in gia_tri)


def _la_tai_khoan_don_vi_truong(account: User) -> bool:
    username = (account.username or "").strip().lower()
    return not any(username.startswith(prefix) for prefix in UNIT_ACCOUNT_EXCLUDED_PREFIXES)


def lay_danh_sach_don_vi_phan_trang(
    *,
    db: Session,
    section: str,
    keyword: str,
    commune_id: int | None,
    school_id: int | None,
    account_status: str,
    page: int,
    page_size: int,
) -> tuple[list[dict[str, Any]], int]:
    """
    Hiển thị toàn bộ xã hoặc trường, kể cả đơn vị chưa được cấp tài khoản.

    Dữ liệu chỉ vài trăm đơn vị nên lọc trong bộ nhớ giúp tránh truy vấn
    OUTER JOIN phức tạp và không làm trùng dòng khi dữ liệu cũ có nhiều
    tài khoản cán bộ quản lý mang vai trò TRUONG.
    """
    rows: list[dict[str, Any]] = []

    if section == "xa":
        units = list(
            db.scalars(
                select(Commune)
                .where(Commune.is_active.is_(True))
                .order_by(Commune.name.asc(), Commune.code.asc())
            ).all()
        )
        accounts = list(
            db.scalars(
                select(User)
                .join(Role, Role.id == User.role_id)
                .options(selectinload(User.role), selectinload(User.commune))
                .where(Role.code == COMMUNE_ROLE_CODE)
                .order_by(User.is_active.desc(), User.id.asc())
            ).all()
        )
        account_map: dict[int, User] = {}
        for account in accounts:
            if account.commune_id is not None:
                account_map.setdefault(int(account.commune_id), account)

        for unit in units:
            if commune_id is not None and unit.id != commune_id:
                continue
            account = account_map.get(unit.id)
            if not _khop_tu_khoa(
                keyword,
                unit.code,
                unit.name,
                account.username if account else "",
                account.full_name if account else "",
            ):
                continue
            if account_status == "active" and not (account and account.is_active):
                continue
            if account_status == "locked" and not (account and not account.is_active):
                continue
            if account_status == "unassigned" and account is not None:
                continue
            rows.append(
                {
                    "loai": "xa",
                    "don_vi": unit,
                    "xa": unit,
                    "truong": None,
                    "tai_khoan": account,
                }
            )

    elif section == "truong":
        unit_statement = (
            select(School)
            .join(School.commune)
            .options(selectinload(School.commune))
            .where(School.is_active.is_(True), Commune.is_active.is_(True))
            .order_by(Commune.name.asc(), School.name.asc(), School.code.asc())
        )
        if commune_id is not None:
            unit_statement = unit_statement.where(School.commune_id == commune_id)
        if school_id is not None:
            unit_statement = unit_statement.where(School.id == school_id)
        units = list(db.scalars(unit_statement).all())

        account_statement = (
            select(User)
            .join(Role, Role.id == User.role_id)
            .options(
                selectinload(User.role),
                selectinload(User.school).selectinload(School.commune),
            )
            .where(Role.code == SCHOOL_ROLE_CODE, User.school_id.is_not(None))
            .order_by(User.is_active.desc(), User.id.asc())
        )
        accounts = [
            account
            for account in db.scalars(account_statement).all()
            if _la_tai_khoan_don_vi_truong(account)
        ]
        account_map: dict[int, User] = {}
        for account in accounts:
            if account.school_id is not None:
                account_map.setdefault(int(account.school_id), account)

        for unit in units:
            account = account_map.get(unit.id)
            if not _khop_tu_khoa(
                keyword,
                unit.code,
                unit.name,
                unit.commune.code if unit.commune else "",
                unit.commune.name if unit.commune else "",
                account.username if account else "",
                account.full_name if account else "",
            ):
                continue
            if account_status == "active" and not (account and account.is_active):
                continue
            if account_status == "locked" and not (account and not account.is_active):
                continue
            if account_status == "unassigned" and account is not None:
                continue
            rows.append(
                {
                    "loai": "truong",
                    "don_vi": unit,
                    "xa": unit.commune,
                    "truong": unit,
                    "tai_khoan": account,
                }
            )

    total = len(rows)
    start = max(page - 1, 0) * page_size
    return rows[start : start + page_size], total


def tao_dieu_kien_loc(
    *,
    actor: dict[str, Any],
    section: str,
    keyword: str,
    commune_id: int | None,
    school_id: int | None,
    account_status: str,
) -> list[Any]:
    actor_role = normalize_role_code(actor.get("role_code"))
    conditions: list[Any] = []

    if actor_role == SCHOOL_ROLE_CODE:
        actor_school_id = chuyen_sang_id(actor.get("school_id"))
        if actor_school_id is None:
            # Điều kiện không thể đúng, bảo đảm không lộ dữ liệu.
            conditions.append(User.id == -1)
        else:
            conditions.extend(
                [
                    Role.code == TEACHER_ROLE_CODE,
                    User.school_id == actor_school_id,
                ]
            )
    elif actor_role in ADMIN_ROLE_CODES:
        if section in SECTION_ROLE_CODES:
            conditions.append(Role.code.in_(SECTION_ROLE_CODES[section]))
        elif section == "tra_cuu":
            # Tra cứu toàn hệ thống, nhưng các nút thao tác vẫn được giới hạn
            # theo đúng quyền ở từng tuyến POST phía máy chủ.
            pass
        else:
            conditions.append(User.id == -1)
    else:
        conditions.append(User.id == -1)

    if keyword:
        pattern = f"%{keyword.strip()}%"
        conditions.append(
            or_(
                User.username.ilike(pattern),
                User.full_name.ilike(pattern),
                Commune.code.ilike(pattern),
                Commune.name.ilike(pattern),
                School.code.ilike(pattern),
                School.name.ilike(pattern),
            )
        )

    if commune_id is not None:
        conditions.append(
            or_(
                User.commune_id == commune_id,
                School.commune_id == commune_id,
            )
        )

    if school_id is not None:
        conditions.append(User.school_id == school_id)

    if account_status == "active":
        conditions.append(User.is_active.is_(True))
    elif account_status == "locked":
        conditions.append(User.is_active.is_(False))

    return conditions


def lay_danh_sach_tai_khoan_phan_trang(
    *,
    db: Session,
    actor: dict[str, Any],
    section: str,
    keyword: str,
    commune_id: int | None,
    school_id: int | None,
    account_status: str,
    page: int,
    page_size: int,
) -> tuple[list[User], int]:
    conditions = tao_dieu_kien_loc(
        actor=actor,
        section=section,
        keyword=keyword,
        commune_id=commune_id,
        school_id=school_id,
        account_status=account_status,
    )

    base_from = (
        select(User)
        .join(Role, Role.id == User.role_id)
        .outerjoin(Commune, Commune.id == User.commune_id)
        .outerjoin(School, School.id == User.school_id)
    )

    count_statement = (
        select(func.count(User.id))
        .select_from(User)
        .join(Role, Role.id == User.role_id)
        .outerjoin(Commune, Commune.id == User.commune_id)
        .outerjoin(School, School.id == User.school_id)
        .where(*conditions)
    )
    total = int(db.scalar(count_statement) or 0)

    statement = (
        base_from
        .options(
            selectinload(User.role),
            selectinload(User.commune),
            selectinload(User.school).selectinload(School.commune),
        )
        .where(*conditions)
        .order_by(User.full_name.asc(), User.username.asc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    accounts = list(db.scalars(statement).unique().all())
    return accounts, total


def tao_url_bo_loc(
    *,
    section: str,
    keyword: str,
    commune_id: int | None,
    school_id: int | None,
    account_status: str,
    page_size: int,
) -> str:
    params: dict[str, str | int] = {
        "muc": section,
        "trang_thai_tk": account_status,
        "so_dong": page_size,
    }
    if keyword:
        params["q"] = keyword
    if commune_id is not None:
        params["xa_id"] = commune_id
    if school_id is not None:
        params["truong_id"] = school_id
    return f"/tai-khoan?{urlencode(params)}"



def _bao_dam_bang_nhat_ky_cap_dong_loat(db: Session) -> None:
    """Tạo và nâng cấp bảng nhật ký cấp tài khoản đồng loạt."""
    db.execute(
        text(
            """
            CREATE TABLE IF NOT EXISTS account_provision_batches (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                unit_type VARCHAR(20) NOT NULL,
                scope_type VARCHAR(20) NOT NULL,
                filters_json TEXT,
                total_candidates INTEGER NOT NULL DEFAULT 0,
                created_count INTEGER NOT NULL DEFAULT 0,
                skipped_count INTEGER NOT NULL DEFAULT 0,
                processed_count INTEGER NOT NULL DEFAULT 0,
                failed_count INTEGER NOT NULL DEFAULT 0,
                job_status VARCHAR(40) NOT NULL DEFAULT 'COMPLETED',
                current_unit_name VARCHAR(300),
                result_file VARCHAR(500),
                error_message TEXT,
                created_by_user_id INTEGER,
                created_by_name VARCHAR(200),
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                started_at DATETIME,
                completed_at DATETIME,
                updated_at DATETIME
            )
            """
        )
    )
    db.execute(
        text(
            """
            CREATE TABLE IF NOT EXISTS account_provision_batch_items (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                batch_id INTEGER NOT NULL,
                unit_type VARCHAR(20) NOT NULL,
                unit_id INTEGER NOT NULL,
                commune_id INTEGER,
                school_id INTEGER,
                unit_code VARCHAR(50),
                unit_name VARCHAR(300),
                username VARCHAR(100),
                result_status VARCHAR(30) NOT NULL,
                result_note VARCHAR(500),
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY(batch_id) REFERENCES account_provision_batches(id)
            )
            """
        )
    )

    columns = {
        str(row[1])
        for row in db.execute(text("PRAGMA table_info(account_provision_batches)"))
    }
    additions = {
        "processed_count": "INTEGER NOT NULL DEFAULT 0",
        "failed_count": "INTEGER NOT NULL DEFAULT 0",
        "job_status": "VARCHAR(40) NOT NULL DEFAULT 'COMPLETED'",
        "current_unit_name": "VARCHAR(300)",
        "result_file": "VARCHAR(500)",
        "error_message": "TEXT",
        "started_at": "DATETIME",
        "completed_at": "DATETIME",
        "updated_at": "DATETIME",
    }
    for column_name, definition in additions.items():
        if column_name not in columns:
            db.execute(
                text(
                    f"ALTER TABLE account_provision_batches "
                    f"ADD COLUMN {column_name} {definition}"
                )
            )

    db.execute(
        text(
            """
            UPDATE account_provision_batches
            SET processed_count = CASE
                    WHEN processed_count = 0
                    THEN created_count + skipped_count
                    ELSE processed_count
                END,
                job_status = COALESCE(NULLIF(job_status, ''), 'COMPLETED'),
                completed_at = CASE
                    WHEN completed_at IS NULL
                         AND COALESCE(NULLIF(job_status, ''), 'COMPLETED') = 'COMPLETED'
                    THEN created_at
                    ELSE completed_at
                END,
                updated_at = COALESCE(updated_at, created_at)
            """
        )
    )
    db.execute(
        text(
            """
            CREATE INDEX IF NOT EXISTS ix_account_provision_batches_status
            ON account_provision_batches(job_status, id)
            """
        )
    )
    db.execute(
        text(
            """
            CREATE INDEX IF NOT EXISTS ix_account_provision_items_batch_status
            ON account_provision_batch_items(batch_id, result_status, id)
            """
        )
    )
    db.commit()


def _doc_dot_cap_dong_loat(db: Session, batch_id: int) -> dict[str, Any] | None:
    _bao_dam_bang_nhat_ky_cap_dong_loat(db)
    row = db.execute(
        text(
            """
            SELECT id, unit_type, scope_type, filters_json, total_candidates,
                   created_count, skipped_count, processed_count, failed_count,
                   job_status, current_unit_name, result_file, error_message,
                   created_by_user_id, created_by_name, created_at,
                   started_at, completed_at, updated_at
            FROM account_provision_batches
            WHERE id = :batch_id
            """
        ),
        {"batch_id": batch_id},
    ).first()
    return dict(row._mapping) if row is not None else None


def _parse_datetime_value(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if value is None:
        return None
    raw = str(value).strip()
    if not raw:
        return None
    try:
        return datetime.fromisoformat(raw)
    except ValueError:
        return None


def _duong_dan_file_ban_giao(filename: str | None) -> Path | None:
    safe_name = Path(str(filename or "")).name
    if not safe_name or safe_name != str(filename or ""):
        return None
    candidate = (BULK_EXPORT_DIR / safe_name).resolve()
    try:
        candidate.relative_to(BULK_EXPORT_DIR.resolve())
    except ValueError:
        return None
    return candidate


def _tai_khoan_don_vi_da_co(
    *,
    db: Session,
    section: str,
    unit_id: int,
) -> bool:
    statement = (
        select(User.id)
        .join(Role, Role.id == User.role_id)
        .where(Role.code == (COMMUNE_ROLE_CODE if section == "xa" else SCHOOL_ROLE_CODE))
    )
    if section == "xa":
        statement = statement.where(User.commune_id == unit_id)
    else:
        statement = statement.where(User.school_id == unit_id)
        for prefix in UNIT_ACCOUNT_EXCLUDED_PREFIXES:
            statement = statement.where(
                ~func.lower(User.username).like(f"{prefix.lower()}%")
            )
    return db.scalar(statement.limit(1)) is not None


def _ghi_file_excel_ban_giao(
    *,
    records: list[dict[str, Any]],
    section: str,
    scope_label: str,
    actor_name: str,
    created_at: datetime,
    batch_id: int,
) -> str:
    content = _tao_excel_ban_giao_tai_khoan(
        records=records,
        section=section,
        scope_label=scope_label,
        actor_name=actor_name,
        created_at=created_at,
        batch_id=batch_id,
    )
    timestamp = created_at.strftime("%Y%m%d_%H%M%S")
    filename = f"ban_giao_tai_khoan_{section}_dot_{batch_id}_{timestamp}.xlsx"
    final_path = BULK_EXPORT_DIR / filename
    temporary_path = BULK_EXPORT_DIR / f".{filename}.tmp"
    temporary_path.write_bytes(content)
    os.replace(temporary_path, final_path)
    return filename


def _cap_nhat_dot_cap(
    db: Session,
    *,
    batch_id: int,
    created_count: int,
    skipped_count: int,
    failed_count: int,
    current_unit_name: str | None,
) -> None:
    processed_count = created_count + skipped_count + failed_count
    db.execute(
        text(
            """
            UPDATE account_provision_batches
            SET created_count = :created_count,
                skipped_count = :skipped_count,
                failed_count = :failed_count,
                processed_count = :processed_count,
                current_unit_name = :current_unit_name,
                updated_at = :updated_at
            WHERE id = :batch_id
            """
        ),
        {
            "created_count": created_count,
            "skipped_count": skipped_count,
            "failed_count": failed_count,
            "processed_count": processed_count,
            "current_unit_name": current_unit_name,
            "updated_at": datetime.now(),
            "batch_id": batch_id,
        },
    )


def _xu_ly_cap_tai_khoan_dong_loat_nen(batch_id: int) -> None:
    """Chạy trong background thread; cập nhật tiến độ sau từng đơn vị."""
    db = SessionLocal()
    handover_records: list[dict[str, Any]] = []
    batch: dict[str, Any] | None = None
    created_count = 0
    skipped_count = 0
    failed_count = 0
    current_item_id: int | None = None
    try:
        _bao_dam_bang_nhat_ky_cap_dong_loat(db)
        batch = _doc_dot_cap_dong_loat(db, batch_id)
        if batch is None:
            return
        if batch["job_status"] in BULK_FINISHED_STATUSES:
            return

        claimed = db.execute(
            text(
                """
                UPDATE account_provision_batches
                SET job_status = :running,
                    started_at = COALESCE(started_at, :now),
                    updated_at = :now,
                    error_message = NULL
                WHERE id = :batch_id
                  AND job_status = :pending
                """
            ),
            {
                "running": BULK_STATUS_RUNNING,
                "pending": BULK_STATUS_PENDING,
                "now": datetime.now(),
                "batch_id": batch_id,
            },
        )
        db.commit()
        if claimed.rowcount != 1:
            return

        section = str(batch["unit_type"])
        role_code = COMMUNE_ROLE_CODE if section == "xa" else SCHOOL_ROLE_CODE
        role = db.scalar(select(Role).where(Role.code == role_code))
        if role is None:
            raise RuntimeError("Không tìm thấy vai trò cần cấp tài khoản.")

        filters_data: dict[str, Any] = {}
        try:
            filters_data = json.loads(batch.get("filters_json") or "{}")
        except (TypeError, json.JSONDecodeError):
            filters_data = {}
        scope_label = str(
            filters_data.get("scope_label")
            or ("Toàn bộ đơn vị chưa có tài khoản" if batch["scope_type"] == BULK_SCOPE_ALL else "Theo bộ lọc hiện tại")
        )
        actor_name = str(batch.get("created_by_name") or "Quản trị hệ thống")
        created_at = _parse_datetime_value(batch.get("created_at")) or datetime.now()

        existing_usernames = {
            str(username).casefold()
            for username in db.scalars(select(User.username)).all()
        }
        item_rows = list(
            db.execute(
                text(
                    """
                    SELECT id, unit_id, commune_id, school_id, unit_code, unit_name
                    FROM account_provision_batch_items
                    WHERE batch_id = :batch_id
                      AND result_status = 'QUEUED'
                    ORDER BY id ASC
                    """
                ),
                {"batch_id": batch_id},
            )
        )

        created_count = int(batch.get("created_count") or 0)
        skipped_count = int(batch.get("skipped_count") or 0)
        failed_count = int(batch.get("failed_count") or 0)

        for raw_item in item_rows:
            item = dict(raw_item._mapping)
            current_item_id = int(item["id"])
            unit_id = int(item["unit_id"])
            unit_name = str(item.get("unit_name") or f"Đơn vị {unit_id}")

            db.execute(
                text(
                    """
                    UPDATE account_provision_batches
                    SET current_unit_name = :unit_name, updated_at = :now
                    WHERE id = :batch_id
                    """
                ),
                {"unit_name": unit_name, "now": datetime.now(), "batch_id": batch_id},
            )
            db.commit()

            if _tai_khoan_don_vi_da_co(db=db, section=section, unit_id=unit_id):
                skipped_count += 1
                db.execute(
                    text(
                        """
                        UPDATE account_provision_batch_items
                        SET result_status = 'SKIPPED',
                            result_note = 'Đơn vị đã có tài khoản trước khi đến lượt xử lý'
                        WHERE id = :item_id
                        """
                    ),
                    {"item_id": current_item_id},
                )
                _cap_nhat_dot_cap(
                    db,
                    batch_id=batch_id,
                    created_count=created_count,
                    skipped_count=skipped_count,
                    failed_count=failed_count,
                    current_unit_name=unit_name,
                )
                db.commit()
                continue

            commune: Commune | None
            school: School | None
            if section == "xa":
                commune = db.get(Commune, unit_id)
                school = None
                valid_unit = commune is not None and bool(commune.is_active)
                unit_code = commune.code if commune is not None else item.get("unit_code")
                unit_name = commune.name if commune is not None else unit_name
            else:
                school = db.scalar(
                    select(School)
                    .options(selectinload(School.commune))
                    .where(School.id == unit_id)
                )
                commune = school.commune if school is not None else None
                valid_unit = (
                    school is not None
                    and bool(school.is_active)
                    and commune is not None
                    and bool(commune.is_active)
                )
                unit_code = school.code if school is not None else item.get("unit_code")
                unit_name = school.name if school is not None else unit_name

            if not valid_unit or commune is None:
                failed_count += 1
                db.execute(
                    text(
                        """
                        UPDATE account_provision_batch_items
                        SET result_status = 'FAILED',
                            result_note = 'Đơn vị không còn hoạt động hoặc thiếu thông tin địa bàn'
                        WHERE id = :item_id
                        """
                    ),
                    {"item_id": current_item_id},
                )
                _cap_nhat_dot_cap(
                    db,
                    batch_id=batch_id,
                    created_count=created_count,
                    skipped_count=skipped_count,
                    failed_count=failed_count,
                    current_unit_name=unit_name,
                )
                db.commit()
                continue

            username = _tao_ten_dang_nhap_don_vi(
                unit_type=section,
                unit_code=unit_code,
                unit_id=unit_id,
                existing_usernames=existing_usernames,
            )
            temporary_password = _tao_mat_khau_tam()

            try:
                password_hash = ma_hoa_mat_khau(temporary_password)
                account = User(
                    username=username,
                    password_hash=password_hash,
                    full_name=unit_name,
                    role_id=role.id,
                    commune_id=commune.id,
                    school_id=school.id if school is not None else None,
                    is_active=True,
                )
                db.add(account)
                db.flush()

                created_count += 1
                db.execute(
                    text(
                        """
                        UPDATE account_provision_batch_items
                        SET username = :username,
                            result_status = 'CREATED',
                            result_note = 'Tạo mới tài khoản đơn vị'
                        WHERE id = :item_id
                        """
                    ),
                    {"username": username, "item_id": current_item_id},
                )
                _cap_nhat_dot_cap(
                    db,
                    batch_id=batch_id,
                    created_count=created_count,
                    skipped_count=skipped_count,
                    failed_count=failed_count,
                    current_unit_name=unit_name,
                )
                db.commit()
                handover_records.append(
                    {
                        "unit_id": unit_id,
                        "unit_name": unit_name,
                        "commune_id": commune.id,
                        "commune_code": commune.code,
                        "commune_name": commune.name,
                        "school_id": school.id if school is not None else None,
                        "school_code": school.code if school is not None else "",
                        "username": username,
                        "password": temporary_password,
                    }
                )
            except IntegrityError:
                db.rollback()
                if _tai_khoan_don_vi_da_co(db=db, section=section, unit_id=unit_id):
                    skipped_count += 1
                    result_status = "SKIPPED"
                    result_note = "Đơn vị đã được cấp tài khoản bởi thao tác khác"
                else:
                    failed_count += 1
                    result_status = "FAILED"
                    result_note = "Trùng tên đăng nhập hoặc xung đột dữ liệu"
                db.execute(
                    text(
                        """
                        UPDATE account_provision_batch_items
                        SET result_status = :result_status,
                            result_note = :result_note
                        WHERE id = :item_id
                        """
                    ),
                    {
                        "result_status": result_status,
                        "result_note": result_note,
                        "item_id": current_item_id,
                    },
                )
                _cap_nhat_dot_cap(
                    db,
                    batch_id=batch_id,
                    created_count=created_count,
                    skipped_count=skipped_count,
                    failed_count=failed_count,
                    current_unit_name=unit_name,
                )
                db.commit()
            except Exception as item_error:
                db.rollback()
                failed_count += 1
                db.execute(
                    text(
                        """
                        UPDATE account_provision_batch_items
                        SET result_status = 'FAILED',
                            result_note = :result_note
                        WHERE id = :item_id
                        """
                    ),
                    {
                        "result_note": f"Lỗi: {str(item_error)[:430]}",
                        "item_id": current_item_id,
                    },
                )
                _cap_nhat_dot_cap(
                    db,
                    batch_id=batch_id,
                    created_count=created_count,
                    skipped_count=skipped_count,
                    failed_count=failed_count,
                    current_unit_name=unit_name,
                )
                db.commit()

        result_filename = _ghi_file_excel_ban_giao(
            records=handover_records,
            section=section,
            scope_label=scope_label,
            actor_name=actor_name,
            created_at=created_at,
            batch_id=batch_id,
        )
        final_status = (
            BULK_STATUS_COMPLETED_WITH_ERRORS
            if failed_count > 0
            else BULK_STATUS_COMPLETED
        )
        db.execute(
            text(
                """
                UPDATE account_provision_batches
                SET job_status = :job_status,
                    current_unit_name = NULL,
                    result_file = :result_file,
                    completed_at = :now,
                    updated_at = :now,
                    error_message = NULL
                WHERE id = :batch_id
                """
            ),
            {
                "job_status": final_status,
                "result_file": result_filename,
                "now": datetime.now(),
                "batch_id": batch_id,
            },
        )
        db.commit()
    except Exception as error:
        db.rollback()
        result_filename: str | None = None
        if batch is not None and handover_records:
            try:
                filters_data = json.loads(batch.get("filters_json") or "{}")
                scope_label = str(filters_data.get("scope_label") or "Theo phạm vi đã chọn")
                created_at = _parse_datetime_value(batch.get("created_at")) or datetime.now()
                result_filename = _ghi_file_excel_ban_giao(
                    records=handover_records,
                    section=str(batch["unit_type"]),
                    scope_label=scope_label,
                    actor_name=str(batch.get("created_by_name") or "Quản trị hệ thống"),
                    created_at=created_at,
                    batch_id=batch_id,
                )
            except Exception:
                result_filename = None
        try:
            _bao_dam_bang_nhat_ky_cap_dong_loat(db)
            if current_item_id is not None:
                db.execute(
                    text(
                        """
                        UPDATE account_provision_batch_items
                        SET result_status = CASE
                                WHEN result_status = 'QUEUED' THEN 'FAILED'
                                ELSE result_status
                            END,
                            result_note = CASE
                                WHEN result_status = 'QUEUED' THEN :note
                                ELSE result_note
                            END
                        WHERE id = :item_id
                        """
                    ),
                    {"note": f"Tiến trình dừng: {str(error)[:420]}", "item_id": current_item_id},
                )
            db.execute(
                text(
                    """
                    UPDATE account_provision_batches
                    SET job_status = :failed,
                        current_unit_name = NULL,
                        result_file = COALESCE(:result_file, result_file),
                        error_message = :error_message,
                        completed_at = :now,
                        updated_at = :now
                    WHERE id = :batch_id
                    """
                ),
                {
                    "failed": BULK_STATUS_FAILED,
                    "result_file": result_filename,
                    "error_message": str(error)[:2000],
                    "now": datetime.now(),
                    "batch_id": batch_id,
                },
            )
            db.commit()
        except Exception:
            db.rollback()
    finally:
        db.close()

def _lay_don_vi_theo_pham_vi_cap(
    *,
    db: Session,
    section: str,
    scope_type: str,
    keyword: str,
    commune_id: int | None,
    school_id: int | None,
    account_status: str = "unassigned",
) -> list[dict[str, Any]]:
    if section not in BULK_UNIT_TYPES:
        return []

    if scope_type == BULK_SCOPE_ALL:
        keyword = ""
        commune_id = None
        school_id = None

    rows, _ = lay_danh_sach_don_vi_phan_trang(
        db=db,
        section=section,
        keyword=keyword,
        commune_id=commune_id,
        school_id=school_id,
        account_status=account_status,
        page=1,
        page_size=100_000,
    )
    return rows


def _chuan_hoa_ma_ten_dang_nhap(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "", ascii_value.lower())


def _tao_ten_dang_nhap_don_vi(
    *,
    unit_type: str,
    unit_code: Any,
    unit_id: int,
    existing_usernames: set[str],
) -> str:
    clean_code = _chuan_hoa_ma_ten_dang_nhap(unit_code) or str(unit_id)
    prefix = "xa" if unit_type == "xa" else "truong"
    base = f"{prefix}_{clean_code}"[:100]
    candidate = base
    index = 2
    while candidate.casefold() in existing_usernames:
        suffix = f"_{index}"
        candidate = f"{base[:100-len(suffix)]}{suffix}"
        index += 1
    existing_usernames.add(candidate.casefold())
    return candidate


def _tao_mat_khau_tam() -> str:
    # Tránh các ký tự dễ nhầm khi bàn giao bằng văn bản hoặc Excel.
    upper = "ABCDEFGHJKLMNPQRSTUVWXYZ"
    lower = "abcdefghijkmnopqrstuvwxyz"
    digits = "23456789"
    symbols = "@#$%"
    required = [
        secrets.choice(upper),
        secrets.choice(lower),
        secrets.choice(digits),
        secrets.choice(symbols),
    ]
    alphabet = upper + lower + digits + symbols
    required.extend(secrets.choice(alphabet) for _ in range(8))
    secrets.SystemRandom().shuffle(required)
    return "".join(required)


def _style_excel_header(ws, row: int, start_col: int, end_col: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7C9D6")
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _tao_excel_ban_giao_tai_khoan(
    *,
    records: list[dict[str, Any]],
    section: str,
    scope_label: str,
    actor_name: str,
    created_at: datetime,
    batch_id: int,
) -> bytes:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Tai khoan da cap"

    unit_label = "XÃ/PHƯỜNG" if section == "xa" else "TRƯỜNG MẦM NON"
    ws.merge_cells("A1:J1")
    ws["A1"] = f"DANH SÁCH BÀN GIAO TÀI KHOẢN {unit_label}"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws["A2"] = "Mã đợt cấp"
    ws["B2"] = batch_id
    ws["D2"] = "Phạm vi"
    ws["E2"] = scope_label
    ws["G2"] = "Người thực hiện"
    ws["H2"] = actor_name
    ws["A3"] = "Thời điểm"
    ws["B3"] = created_at.strftime("%d/%m/%Y %H:%M:%S")
    ws["D3"] = "Số tài khoản"
    ws["E3"] = len(records)
    ws.merge_cells("H3:J3")
    ws["H3"] = "Mật khẩu chỉ xuất tại lần cấp này. Hãy bàn giao riêng và yêu cầu đổi ngay sau đăng nhập."
    ws["H3"].alignment = Alignment(wrap_text=True)
    ws["H3"].font = Font(italic=True, color="9C5700")

    headers = [
        "STT",
        "Mã xã/phường",
        "Tên xã/phường",
        "Mã trường",
        "Tên trường/đơn vị",
        "Tên đăng nhập",
        "Mật khẩu ban đầu",
        "Trạng thái",
        "Thời điểm cấp",
        "Người cấp",
    ]
    header_row = 5
    for col, value in enumerate(headers, 1):
        ws.cell(row=header_row, column=col, value=value)
    _style_excel_header(ws, header_row, 1, len(headers))
    ws.row_dimensions[header_row].height = 34

    thin = Side(style="thin", color="D9E2F3")
    for index, record in enumerate(records, 1):
        values = [
            index,
            record.get("commune_code", ""),
            record.get("commune_name", ""),
            record.get("school_code", ""),
            record.get("unit_name", ""),
            record.get("username", ""),
            record.get("password", ""),
            "Đang hoạt động",
            created_at.strftime("%d/%m/%Y %H:%M:%S"),
            actor_name,
        ]
        row_number = header_row + index
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_number, column=col, value=value)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if col in {2, 4, 6, 7}:
                cell.number_format = "@"
        if index % 2 == 0:
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_number, column=col).fill = PatternFill("solid", fgColor="F4F8FC")

    widths = [7, 15, 24, 18, 38, 24, 22, 18, 22, 24]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:J{max(5, 5 + len(records))}"
    ws.sheet_view.showGridLines = False

    info = workbook.create_sheet("Thong tin dot cap")
    info.append(["Nội dung", "Thông tin"])
    _style_excel_header(info, 1, 1, 2)
    info_rows = [
        ("Mã đợt cấp", batch_id),
        ("Loại đơn vị", unit_label),
        ("Phạm vi", scope_label),
        ("Số tài khoản đã cấp", len(records)),
        ("Người thực hiện", actor_name),
        ("Thời điểm", created_at.strftime("%d/%m/%Y %H:%M:%S")),
        ("Lưu ý", "Mật khẩu không được lưu dạng đọc được trong cơ sở dữ liệu."),
        ("Khuyến nghị", "Yêu cầu đơn vị đổi mật khẩu ngay sau lần đăng nhập đầu tiên."),
    ]
    for row in info_rows:
        info.append(list(row))
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 80
    for row in info.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    info.sheet_view.showGridLines = False

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _tao_excel_don_vi_chua_cap(
    *,
    rows: list[dict[str, Any]],
    section: str,
    scope_label: str,
) -> bytes:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Don vi chua cap"
    title = "DANH SÁCH XÃ/PHƯỜNG CHƯA CẤP TÀI KHOẢN" if section == "xa" else "DANH SÁCH TRƯỜNG CHƯA CẤP TÀI KHOẢN"
    ws.merge_cells("A1:F1")
    ws["A1"] = title
    ws["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="17365D")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = "Phạm vi"
    ws["B2"] = scope_label
    ws["D2"] = "Tổng số"
    ws["E2"] = len(rows)
    headers = ["STT", "Mã xã/phường", "Tên xã/phường", "Mã trường", "Tên trường/đơn vị", "Trạng thái"]
    for col, value in enumerate(headers, 1):
        ws.cell(row=4, column=col, value=value)
    _style_excel_header(ws, 4, 1, len(headers))
    for index, row in enumerate(rows, 1):
        unit = row["don_vi"]
        commune = row.get("xa")
        school = row.get("truong")
        values = [
            index,
            commune.code if commune else "",
            commune.name if commune else "",
            school.code if school else "",
            unit.name,
            "Chưa cấp tài khoản",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=4 + index, column=col, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if col in {2, 4}:
                cell.number_format = "@"
    widths = [7, 16, 28, 18, 46, 22]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:F{max(4, 4 + len(rows))}"
    ws.sheet_view.showGridLines = False
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _lay_nhat_ky_cap_dong_loat(db: Session, limit: int = 8) -> list[dict[str, Any]]:
    try:
        _bao_dam_bang_nhat_ky_cap_dong_loat(db)
        result = db.execute(
            text(
                """
                SELECT id, unit_type, scope_type, total_candidates,
                       created_count, skipped_count, processed_count,
                       failed_count, job_status, result_file,
                       created_by_name, created_at, completed_at
                FROM account_provision_batches
                ORDER BY id DESC
                LIMIT :limit
                """
            ),
            {"limit": limit},
        )
        return [dict(row._mapping) for row in result]
    except Exception:
        db.rollback()
        return []

def _tao_url_bulk(
    *,
    action: str,
    section: str,
    scope_type: str,
    keyword: str,
    commune_id: int | None,
    school_id: int | None,
) -> str:
    params: dict[str, str | int] = {"muc": section, "pham_vi": scope_type}
    if scope_type == BULK_SCOPE_FILTERED:
        if keyword:
            params["q"] = keyword
        if commune_id is not None:
            params["xa_id"] = commune_id
        if school_id is not None:
            params["truong_id"] = school_id
    return f"{action}?{urlencode(params)}"


@router.get("", response_class=HTMLResponse)
def danh_sach_tai_khoan(
    request: Request,
    db: Session = Depends(get_db),
    status: str | None = None,
    muc: str | None = Query(default=None),
    q: str = Query(default="", max_length=200),
    xa_id: str = Query(default=""),
    truong_id: str = Query(default=""),
    trang_thai_tk: str = Query(default="all"),
    trang: int = Query(default=1, ge=1),
    so_dong: int = Query(default=DEFAULT_PAGE_SIZE),
):
    actor = lay_nguoi_dung(request)
    actor_role = normalize_role_code(actor.get("role_code"))
    section = muc_duoc_phep(actor_role, muc)

    page_size = so_dong if so_dong in ALLOWED_PAGE_SIZES else DEFAULT_PAGE_SIZE
    valid_statuses = {"all", "active", "locked", "unassigned"}
    account_status = trang_thai_tk if trang_thai_tk in valid_statuses else "all"
    commune_id = chuyen_sang_id(xa_id)
    school_id = chuyen_sang_id(truong_id)
    keyword = " ".join(q.strip().split())

    communes = (
        list(
            db.scalars(
                select(Commune)
                .where(Commune.is_active.is_(True))
                .order_by(Commune.name.asc())
            ).all()
        )
        if actor_role in ADMIN_ROLE_CODES
        else []
    )

    if actor_role in ADMIN_ROLE_CODES:
        school_statement = (
            select(School)
            .join(School.commune)
            .where(School.is_active.is_(True), Commune.is_active.is_(True))
            .options(selectinload(School.commune))
            .order_by(School.name.asc())
        )
        if commune_id is not None:
            school_statement = school_statement.where(School.commune_id == commune_id)
        schools = list(db.scalars(school_statement).all())
    elif actor_role == SCHOOL_ROLE_CODE and actor.get("school_id") is not None:
        school = db.get(School, int(actor["school_id"]))
        schools = [school] if school is not None else []
    else:
        schools = []

    accounts: list[User] = []
    unit_rows: list[dict[str, Any]] = []
    total_accounts = 0
    total_pages = 1
    current_page = trang
    is_unit_list = actor_role in ADMIN_ROLE_CODES and section in {"xa", "truong"}

    should_load_list = section != "tong_quan" or actor_role == SCHOOL_ROLE_CODE
    if should_load_list:
        if is_unit_list:
            unit_rows, total_accounts = lay_danh_sach_don_vi_phan_trang(
                db=db,
                section=section,
                keyword=keyword,
                commune_id=commune_id,
                school_id=school_id,
                account_status=account_status,
                page=current_page,
                page_size=page_size,
            )
        else:
            accounts, total_accounts = lay_danh_sach_tai_khoan_phan_trang(
                db=db,
                actor=actor,
                section=section,
                keyword=keyword,
                commune_id=commune_id,
                school_id=school_id,
                account_status=(
                    "all" if account_status == "unassigned" else account_status
                ),
                page=current_page,
                page_size=page_size,
            )

        total_pages = max(1, math.ceil(total_accounts / page_size))
        if current_page > total_pages:
            current_page = total_pages
            if is_unit_list:
                unit_rows, total_accounts = lay_danh_sach_don_vi_phan_trang(
                    db=db,
                    section=section,
                    keyword=keyword,
                    commune_id=commune_id,
                    school_id=school_id,
                    account_status=account_status,
                    page=current_page,
                    page_size=page_size,
                )
            else:
                accounts, total_accounts = lay_danh_sach_tai_khoan_phan_trang(
                    db=db,
                    actor=actor,
                    section=section,
                    keyword=keyword,
                    commune_id=commune_id,
                    school_id=school_id,
                    account_status=(
                        "all" if account_status == "unassigned" else account_status
                    ),
                    page=current_page,
                    page_size=page_size,
                )

    roles = lay_danh_sach_vai_tro_tao(db, actor_role, section)
    summary = (
        tong_hop_tai_khoan(db)
        if actor_role in ADMIN_ROLE_CODES
        else {
            "phong_ban": 0,
            "xa": 0,
            "truong": 0,
            "giao_vien": total_accounts,
            "tat_ca": total_accounts,
            "tong_xa": 0,
            "tong_truong": 0,
        }
    )

    filter_url = tao_url_bo_loc(
        section=section,
        keyword=keyword,
        commune_id=commune_id,
        school_id=school_id,
        account_status=account_status,
        page_size=page_size,
    )

    start_index = (current_page - 1) * page_size + 1 if total_accounts else 0
    end_index = min(current_page * page_size, total_accounts)

    bulk_filtered_count = 0
    bulk_all_count = 0
    bulk_preview_filtered_url = ""
    bulk_preview_all_url = ""
    bulk_export_unassigned_url = ""
    bulk_history: list[dict[str, Any]] = []
    if actor_role in ADMIN_ROLE_CODES and is_unit_list:
        bulk_filtered_count = len(
            _lay_don_vi_theo_pham_vi_cap(
                db=db,
                section=section,
                scope_type=BULK_SCOPE_FILTERED,
                keyword=keyword,
                commune_id=commune_id,
                school_id=school_id,
            )
        )
        bulk_all_count = len(
            _lay_don_vi_theo_pham_vi_cap(
                db=db,
                section=section,
                scope_type=BULK_SCOPE_ALL,
                keyword="",
                commune_id=None,
                school_id=None,
            )
        )
        bulk_preview_filtered_url = _tao_url_bulk(
            action="/tai-khoan/cap-dong-loat/xem-truoc",
            section=section,
            scope_type=BULK_SCOPE_FILTERED,
            keyword=keyword,
            commune_id=commune_id,
            school_id=school_id,
        )
        bulk_preview_all_url = _tao_url_bulk(
            action="/tai-khoan/cap-dong-loat/xem-truoc",
            section=section,
            scope_type=BULK_SCOPE_ALL,
            keyword="",
            commune_id=None,
            school_id=None,
        )
        bulk_export_unassigned_url = _tao_url_bulk(
            action="/tai-khoan/xuat-don-vi-chua-cap-excel",
            section=section,
            scope_type=BULK_SCOPE_FILTERED,
            keyword=keyword,
            commune_id=commune_id,
            school_id=school_id,
        )
        bulk_history = _lay_nhat_ky_cap_dong_loat(db)

    return templates.TemplateResponse(
        request=request,
        name="users/list.html",
        context={
            "nguoi_dung": actor,
            "danh_sach_vai_tro": roles,
            "danh_sach_xa": communes,
            "danh_sach_truong": schools,
            "danh_sach_tai_khoan": accounts,
            "danh_sach_don_vi": unit_rows,
            "la_danh_sach_don_vi": is_unit_list,
            "thong_bao": STATUS_MESSAGES.get(status),
            "trang_thai": status,
            "la_quan_tri": actor_role in ADMIN_ROLE_CODES,
            "la_cap_truong": actor_role == SCHOOL_ROLE_CODE,
            "muc_hien_tai": section,
            "tieu_de_muc": SECTION_TITLES[section],
            "tong_hop": summary,
            "tu_khoa": keyword,
            "xa_da_chon": commune_id,
            "truong_da_chon": school_id,
            "trang_thai_da_chon": account_status,
            "trang_hien_tai": current_page,
            "tong_so_trang": total_pages,
            "so_dong": page_size,
            "tong_so_tai_khoan": total_accounts,
            "chi_so_bat_dau": start_index,
            "chi_so_ket_thuc": end_index,
            "url_bo_loc": filter_url,
            "return_to": filter_url,
            "so_don_vi_chua_cap_trong_bo_loc": bulk_filtered_count,
            "so_don_vi_chua_cap_toan_bo": bulk_all_count,
            "url_xem_truoc_cap_theo_bo_loc": bulk_preview_filtered_url,
            "url_xem_truoc_cap_tat_ca": bulk_preview_all_url,
            "url_xuat_don_vi_chua_cap": bulk_export_unassigned_url,
            "nhat_ky_cap_dong_loat": bulk_history,
        },
    )


@router.get("/cap-dong-loat/xem-truoc", response_class=HTMLResponse)
def xem_truoc_cap_tai_khoan_dong_loat(
    request: Request,
    db: Session = Depends(get_db),
    muc: str = Query(default="truong"),
    pham_vi: str = Query(default=BULK_SCOPE_FILTERED),
    q: str = Query(default="", max_length=200),
    xa_id: str = Query(default=""),
    truong_id: str = Query(default=""),
    status: str | None = None,
):
    actor = lay_nguoi_dung(request)
    actor_role = normalize_role_code(actor.get("role_code"))
    if actor_role not in ADMIN_ROLE_CODES:
        return RedirectResponse("/tai-khoan?status=bulk_forbidden", 303)

    section = muc if muc in BULK_UNIT_TYPES else "truong"
    scope_type = pham_vi if pham_vi in {BULK_SCOPE_FILTERED, BULK_SCOPE_ALL} else BULK_SCOPE_FILTERED
    keyword = " ".join(q.strip().split())
    commune_id = chuyen_sang_id(xa_id)
    school_id = chuyen_sang_id(truong_id)

    candidates = _lay_don_vi_theo_pham_vi_cap(
        db=db,
        section=section,
        scope_type=scope_type,
        keyword=keyword,
        commune_id=commune_id,
        school_id=school_id,
        account_status="unassigned",
    )
    all_rows = _lay_don_vi_theo_pham_vi_cap(
        db=db,
        section=section,
        scope_type=scope_type,
        keyword=keyword,
        commune_id=commune_id,
        school_id=school_id,
        account_status="all",
    )

    if scope_type == BULK_SCOPE_ALL:
        scope_label = "Toàn bộ đơn vị chưa có tài khoản"
    else:
        scope_parts = ["Bộ lọc hiện tại"]
        if commune_id is not None:
            commune = db.get(Commune, commune_id)
            if commune is not None:
                scope_parts.append(commune.name)
        if school_id is not None:
            school = db.get(School, school_id)
            if school is not None:
                scope_parts.append(school.name)
        if keyword:
            scope_parts.append(f'Từ khóa: "{keyword}"')
        scope_label = " · ".join(scope_parts)

    return_url = tao_url_bo_loc(
        section=section,
        keyword=keyword if scope_type == BULK_SCOPE_FILTERED else "",
        commune_id=commune_id if scope_type == BULK_SCOPE_FILTERED else None,
        school_id=school_id if scope_type == BULK_SCOPE_FILTERED else None,
        account_status="all",
        page_size=DEFAULT_PAGE_SIZE,
    )
    export_url = _tao_url_bulk(
        action="/tai-khoan/xuat-don-vi-chua-cap-excel",
        section=section,
        scope_type=scope_type,
        keyword=keyword,
        commune_id=commune_id,
        school_id=school_id,
    )

    return templates.TemplateResponse(
        request=request,
        name="users/bulk_preview.html",
        context={
            "nguoi_dung": actor,
            "muc_hien_tai": section,
            "ten_loai_don_vi": "Xã/phường" if section == "xa" else "Trường mầm non",
            "pham_vi": scope_type,
            "mo_ta_pham_vi": scope_label,
            "tu_khoa": keyword,
            "xa_da_chon": commune_id,
            "truong_da_chon": school_id,
            "tong_don_vi": len(all_rows),
            "da_co_tai_khoan": max(0, len(all_rows) - len(candidates)),
            "chua_co_tai_khoan": len(candidates),
            "du_dieu_kien": len(candidates),
            "danh_sach_xem_truoc": candidates[:BULK_PREVIEW_LIMIT],
            "so_dong_an": max(0, len(candidates) - BULK_PREVIEW_LIMIT),
            "cum_xac_nhan": BULK_CONFIRM_PHRASE,
            "url_quay_lai": return_url,
            "url_xuat_chua_cap": export_url,
            "thong_bao": STATUS_MESSAGES.get(status),
            "trang_thai": status,
        },
    )


@router.get("/xuat-don-vi-chua-cap-excel")
def xuat_don_vi_chua_cap_excel(
    request: Request,
    db: Session = Depends(get_db),
    muc: str = Query(default="truong"),
    pham_vi: str = Query(default=BULK_SCOPE_FILTERED),
    q: str = Query(default="", max_length=200),
    xa_id: str = Query(default=""),
    truong_id: str = Query(default=""),
):
    actor = lay_nguoi_dung(request)
    if normalize_role_code(actor.get("role_code")) not in ADMIN_ROLE_CODES:
        return RedirectResponse("/tai-khoan?status=bulk_forbidden", 303)

    section = muc if muc in BULK_UNIT_TYPES else "truong"
    scope_type = pham_vi if pham_vi in {BULK_SCOPE_FILTERED, BULK_SCOPE_ALL} else BULK_SCOPE_FILTERED
    keyword = " ".join(q.strip().split())
    commune_id = chuyen_sang_id(xa_id)
    school_id = chuyen_sang_id(truong_id)
    rows = _lay_don_vi_theo_pham_vi_cap(
        db=db,
        section=section,
        scope_type=scope_type,
        keyword=keyword,
        commune_id=commune_id,
        school_id=school_id,
    )
    scope_label = "Toàn bộ" if scope_type == BULK_SCOPE_ALL else "Theo bộ lọc hiện tại"
    content = _tao_excel_don_vi_chua_cap(rows=rows, section=section, scope_label=scope_label)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"don_vi_chua_cap_tai_khoan_{section}_{timestamp}.xlsx"
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/cap-dong-loat/{batch_id}", response_class=HTMLResponse)
def tien_do_cap_tai_khoan_dong_loat(
    batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    actor = lay_nguoi_dung(request)
    if normalize_role_code(actor.get("role_code")) not in ADMIN_ROLE_CODES:
        return RedirectResponse("/tai-khoan?status=bulk_forbidden", 303)
    batch = _doc_dot_cap_dong_loat(db, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy đợt cấp tài khoản.")

    try:
        filters_data = json.loads(batch.get("filters_json") or "{}")
    except (TypeError, json.JSONDecodeError):
        filters_data = {}
    section = str(batch["unit_type"])
    return_url = str(filters_data.get("return_url") or f"/tai-khoan?muc={section}")
    return templates.TemplateResponse(
        request=request,
        name="users/bulk_progress.html",
        context={
            "nguoi_dung": actor,
            "dot_cap": batch,
            "ten_loai_don_vi": "Xã/phường" if section == "xa" else "Trường mầm non",
            "mo_ta_pham_vi": str(filters_data.get("scope_label") or "Phạm vi đã xác nhận"),
            "url_quay_lai": url_quay_lai_hop_le(return_url),
            "url_trang_thai": f"/tai-khoan/cap-dong-loat/{batch_id}/trang-thai",
            "url_tai_excel": f"/tai-khoan/cap-dong-loat/{batch_id}/tai-excel",
        },
    )


@router.get("/cap-dong-loat/{batch_id}/trang-thai")
def trang_thai_cap_tai_khoan_dong_loat(
    batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    actor = lay_nguoi_dung(request)
    if normalize_role_code(actor.get("role_code")) not in ADMIN_ROLE_CODES:
        return JSONResponse({"detail": "Không có quyền."}, status_code=403)
    batch = _doc_dot_cap_dong_loat(db, batch_id)
    if batch is None:
        return JSONResponse({"detail": "Không tìm thấy đợt cấp."}, status_code=404)

    total = int(batch.get("total_candidates") or 0)
    processed = int(batch.get("processed_count") or 0)
    created = int(batch.get("created_count") or 0)
    skipped = int(batch.get("skipped_count") or 0)
    failed = int(batch.get("failed_count") or 0)
    percent = round((processed / total * 100) if total else 100.0, 1)

    started_at = _parse_datetime_value(batch.get("started_at"))
    completed_at = _parse_datetime_value(batch.get("completed_at"))
    end_time = completed_at or datetime.now()
    elapsed_seconds = max(0, int((end_time - started_at).total_seconds())) if started_at else 0
    rate = (processed / elapsed_seconds) if elapsed_seconds > 0 else 0.0
    remaining = max(0, total - processed)
    eta_seconds = int(remaining / rate) if rate > 0 and remaining > 0 else None
    job_status = str(batch.get("job_status") or BULK_STATUS_PENDING)
    is_finished = job_status in BULK_FINISHED_STATUSES
    file_path = _duong_dan_file_ban_giao(batch.get("result_file"))
    file_ready = bool(file_path and file_path.is_file())

    return JSONResponse(
        {
            "batch_id": batch_id,
            "status": job_status,
            "is_finished": is_finished,
            "total": total,
            "processed": processed,
            "created": created,
            "skipped": skipped,
            "failed": failed,
            "percent": percent,
            "current_unit": batch.get("current_unit_name") or "",
            "elapsed_seconds": elapsed_seconds,
            "eta_seconds": eta_seconds,
            "file_ready": file_ready,
            "download_url": f"/tai-khoan/cap-dong-loat/{batch_id}/tai-excel" if file_ready else "",
            "error_message": batch.get("error_message") or "",
        }
    )


@router.get("/cap-dong-loat/{batch_id}/tai-excel")
def tai_excel_ban_giao_cap_dong_loat(
    batch_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    actor = lay_nguoi_dung(request)
    if normalize_role_code(actor.get("role_code")) not in ADMIN_ROLE_CODES:
        return RedirectResponse("/tai-khoan?status=bulk_forbidden", 303)
    batch = _doc_dot_cap_dong_loat(db, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy đợt cấp tài khoản.")
    file_path = _duong_dan_file_ban_giao(batch.get("result_file"))
    if file_path is None or not file_path.is_file():
        raise HTTPException(
            status_code=409,
            detail="File Excel chưa sẵn sàng hoặc không còn trên máy chủ.",
        )
    return FileResponse(
        path=str(file_path),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=file_path.name,
    )


@router.post("/cap-dong-loat")
def cap_tai_khoan_dong_loat(
    request: Request,
    background_tasks: BackgroundTasks,
    muc: Annotated[str, Form()],
    pham_vi: Annotated[str, Form()],
    xac_nhan: Annotated[str, Form()] = "",
    cum_xac_nhan: Annotated[str, Form()] = "",
    q: Annotated[str, Form()] = "",
    xa_id: Annotated[str, Form()] = "",
    truong_id: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    actor = lay_nguoi_dung(request)
    actor_role = normalize_role_code(actor.get("role_code"))
    if actor_role not in ADMIN_ROLE_CODES:
        return RedirectResponse("/tai-khoan?status=bulk_forbidden", 303)

    section = muc if muc in BULK_UNIT_TYPES else "truong"
    scope_type = pham_vi if pham_vi in {BULK_SCOPE_FILTERED, BULK_SCOPE_ALL} else BULK_SCOPE_FILTERED
    keyword = " ".join(q.strip().split())
    commune_id = chuyen_sang_id(xa_id)
    school_id = chuyen_sang_id(truong_id)

    preview_url = _tao_url_bulk(
        action="/tai-khoan/cap-dong-loat/xem-truoc",
        section=section,
        scope_type=scope_type,
        keyword=keyword,
        commune_id=commune_id,
        school_id=school_id,
    )
    if xac_nhan != "yes" or cum_xac_nhan.strip().upper() != BULK_CONFIRM_PHRASE:
        return RedirectResponse(f"{preview_url}&status=bulk_invalid_confirm", 303)

    _bao_dam_bang_nhat_ky_cap_dong_loat(db)
    active_batch = db.execute(
        text(
            """
            SELECT id
            FROM account_provision_batches
            WHERE unit_type = :unit_type
              AND job_status IN ('PENDING', 'RUNNING')
            ORDER BY id DESC
            LIMIT 1
            """
        ),
        {"unit_type": section},
    ).first()
    if active_batch is not None:
        return RedirectResponse(
            f"/tai-khoan/cap-dong-loat/{int(active_batch[0])}",
            303,
        )

    candidates = _lay_don_vi_theo_pham_vi_cap(
        db=db,
        section=section,
        scope_type=scope_type,
        keyword=keyword,
        commune_id=commune_id,
        school_id=school_id,
    )
    if not candidates:
        return RedirectResponse(f"{preview_url}&status=bulk_none", 303)

    actor_name = str(actor.get("full_name") or actor.get("username") or "Quản trị hệ thống")
    actor_id = chuyen_sang_id(actor.get("id"))
    created_at = datetime.now()
    if scope_type == BULK_SCOPE_ALL:
        scope_label = "Toàn bộ đơn vị chưa có tài khoản"
    else:
        scope_parts = ["Bộ lọc hiện tại"]
        if commune_id is not None:
            commune = db.get(Commune, commune_id)
            if commune is not None:
                scope_parts.append(commune.name)
        if school_id is not None:
            school = db.get(School, school_id)
            if school is not None:
                scope_parts.append(school.name)
        if keyword:
            scope_parts.append(f'Từ khóa: "{keyword}"')
        scope_label = " · ".join(scope_parts)

    return_url = tao_url_bo_loc(
        section=section,
        keyword=keyword if scope_type == BULK_SCOPE_FILTERED else "",
        commune_id=commune_id if scope_type == BULK_SCOPE_FILTERED else None,
        school_id=school_id if scope_type == BULK_SCOPE_FILTERED else None,
        account_status="all",
        page_size=DEFAULT_PAGE_SIZE,
    )
    filters_data = {
        "q": keyword,
        "commune_id": commune_id,
        "school_id": school_id,
        "scope_label": scope_label,
        "return_url": return_url,
    }

    batch_result = db.execute(
        text(
            """
            INSERT INTO account_provision_batches (
                unit_type, scope_type, filters_json, total_candidates,
                created_count, skipped_count, processed_count, failed_count,
                job_status, created_by_user_id, created_by_name,
                created_at, updated_at
            ) VALUES (
                :unit_type, :scope_type, :filters_json, :total_candidates,
                0, 0, 0, 0, :job_status, :created_by_user_id,
                :created_by_name, :created_at, :updated_at
            )
            """
        ),
        {
            "unit_type": section,
            "scope_type": scope_type,
            "filters_json": json.dumps(filters_data, ensure_ascii=False),
            "total_candidates": len(candidates),
            "job_status": BULK_STATUS_PENDING,
            "created_by_user_id": actor_id,
            "created_by_name": actor_name,
            "created_at": created_at,
            "updated_at": created_at,
        },
    )
    batch_id = int(batch_result.lastrowid)
    for row in candidates:
        unit = row["don_vi"]
        commune = row.get("xa")
        school = row.get("truong")
        db.execute(
            text(
                """
                INSERT INTO account_provision_batch_items (
                    batch_id, unit_type, unit_id, commune_id, school_id,
                    unit_code, unit_name, username, result_status,
                    result_note, created_at
                ) VALUES (
                    :batch_id, :unit_type, :unit_id, :commune_id, :school_id,
                    :unit_code, :unit_name, NULL, 'QUEUED',
                    'Đang chờ xử lý', :created_at
                )
                """
            ),
            {
                "batch_id": batch_id,
                "unit_type": section,
                "unit_id": unit.id,
                "commune_id": commune.id if commune is not None else None,
                "school_id": school.id if school is not None else None,
                "unit_code": unit.code,
                "unit_name": unit.name,
                "created_at": created_at,
            },
        )
    db.commit()

    background_tasks.add_task(_xu_ly_cap_tai_khoan_dong_loat_nen, batch_id)
    return RedirectResponse(f"/tai-khoan/cap-dong-loat/{batch_id}", 303)


@router.post("/them")
def them_tai_khoan(
    request: Request,
    ten_dang_nhap: Annotated[str, Form()],
    ho_ten: Annotated[str, Form()],
    mat_khau: Annotated[str, Form()],
    role_id: Annotated[int, Form()],
    commune_id: Annotated[str, Form()] = "",
    school_id: Annotated[str, Form()] = "",
    return_to: Annotated[str, Form()] = "/tai-khoan",
    db: Session = Depends(get_db),
):
    actor = lay_nguoi_dung(request)
    actor_role = normalize_role_code(actor.get("role_code"))
    username = chuan_hoa_ten_dang_nhap(ten_dang_nhap)
    full_name = chuan_hoa_ho_ten(ho_ten)

    if not username or not full_name or " " in username:
        return chuyen_huong_trang_thai(return_to, "invalid")
    if len(username) > 100 or len(full_name) > 200:
        return chuyen_huong_trang_thai(return_to, "invalid")
    if len(mat_khau) < MIN_PASSWORD_LENGTH or len(mat_khau) > MAX_PASSWORD_LENGTH:
        return chuyen_huong_trang_thai(return_to, "invalid_password")

    role = db.get(Role, role_id)
    if role is None or normalize_role_code(role.code) not in vai_tro_duoc_tao(actor_role):
        return chuyen_huong_trang_thai(return_to, "invalid_role")

    if db.scalar(select(User).where(User.username == username)) is not None:
        return chuyen_huong_trang_thai(return_to, "duplicate")

    target_role_code = normalize_role_code(role.code)
    if (
        actor_role in ADMIN_ROLE_CODES
        and target_role_code == SCHOOL_ROLE_CODE
        and chuyen_sang_id(commune_id) is None
    ):
        return chuyen_huong_trang_thai(return_to, "invalid_scope")

    fixed_school_id = (
        int(actor["school_id"])
        if actor_role == SCHOOL_ROLE_CODE and actor.get("school_id") is not None
        else None
    )
    school_value = str(fixed_school_id) if fixed_school_id is not None else school_id

    assigned_commune, assigned_school, error = xac_dinh_don_vi_quan_ly(
        db=db,
        vai_tro=role,
        commune_value=commune_id,
        school_value=school_value,
        school_id_bat_buoc=fixed_school_id,
    )
    if error is not None:
        return chuyen_huong_trang_thai(return_to, error)

    account = User(
        username=username,
        password_hash=ma_hoa_mat_khau(mat_khau),
        full_name=full_name,
        role_id=role.id,
        commune_id=assigned_commune,
        school_id=assigned_school,
        is_active=True,
    )
    db.add(account)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return chuyen_huong_trang_thai(return_to, "duplicate")
    return chuyen_huong_trang_thai(return_to, "created")


@router.get("/{user_id}/sua", response_class=HTMLResponse)
def form_sua_tai_khoan(
    user_id: int,
    request: Request,
    db: Session = Depends(get_db),
    status: str | None = None,
):
    actor = lay_nguoi_dung(request)
    actor_role = normalize_role_code(actor.get("role_code"))
    account = db.scalar(
        select(User)
        .options(
            selectinload(User.role),
            selectinload(User.school),
            selectinload(User.commune),
        )
        .where(User.id == user_id)
    )
    if account is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản.")
    if not tai_khoan_thuoc_pham_vi(actor=actor, target=account):
        return RedirectResponse("/tai-khoan?status=forbidden", 303)

    if actor_role == SCHOOL_ROLE_CODE:
        roles = list(db.scalars(select(Role).where(Role.code == TEACHER_ROLE_CODE)).all())
        schools = [account.school] if account.school is not None else []
        communes = []
    else:
        roles = list(
            db.scalars(
                select(Role)
                .where(
                    Role.code.in_(
                        [
                            DEPARTMENT_ROLE_CODE,
                            COMMUNE_ROLE_CODE,
                            SCHOOL_ROLE_CODE,
                        ]
                    )
                )
                .order_by(Role.id.asc())
            ).all()
        )
        communes = list(
            db.scalars(
                select(Commune)
                .where(or_(Commune.is_active.is_(True), Commune.id == account.commune_id))
                .order_by(Commune.name.asc())
            ).all()
        )
        schools = list(
            db.scalars(
                select(School)
                .join(School.commune)
                .where(
                    or_(
                        and_(School.is_active.is_(True), Commune.is_active.is_(True)),
                        School.id == account.school_id,
                    )
                )
                .options(selectinload(School.commune))
                .order_by(School.name.asc())
            ).all()
        )

    return templates.TemplateResponse(
        request=request,
        name="users/edit.html",
        context={
            "nguoi_dung": actor,
            "tai_khoan": account,
            "danh_sach_vai_tro": roles,
            "danh_sach_xa": communes,
            "danh_sach_truong": schools,
            "thong_bao": STATUS_MESSAGES.get(status),
            "la_quan_tri": actor_role in ADMIN_ROLE_CODES,
            "la_cap_truong": actor_role == SCHOOL_ROLE_CODE,
        },
    )


@router.post("/{user_id}/sua")
def cap_nhat_tai_khoan(
    user_id: int,
    request: Request,
    ten_dang_nhap: Annotated[str, Form()],
    ho_ten: Annotated[str, Form()],
    role_id: Annotated[int, Form()],
    mat_khau_moi: Annotated[str, Form()] = "",
    commune_id: Annotated[str, Form()] = "",
    school_id: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    actor = lay_nguoi_dung(request)
    actor_role = normalize_role_code(actor.get("role_code"))
    account = db.scalar(
        select(User).options(selectinload(User.role)).where(User.id == user_id)
    )
    if account is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản.")
    if not tai_khoan_thuoc_pham_vi(actor=actor, target=account):
        return RedirectResponse("/tai-khoan?status=forbidden", 303)

    username = chuan_hoa_ten_dang_nhap(ten_dang_nhap)
    full_name = chuan_hoa_ho_ten(ho_ten)
    if not username or not full_name or " " in username:
        return RedirectResponse(f"/tai-khoan/{user_id}/sua?status=invalid", 303)
    if mat_khau_moi and (
        len(mat_khau_moi) < MIN_PASSWORD_LENGTH
        or len(mat_khau_moi) > MAX_PASSWORD_LENGTH
    ):
        return RedirectResponse(f"/tai-khoan/{user_id}/sua?status=invalid_password", 303)

    role = db.get(Role, role_id)
    if role is None:
        return RedirectResponse(f"/tai-khoan/{user_id}/sua?status=invalid_role", 303)

    target_role = normalize_role_code(role.code)
    if actor_role == SCHOOL_ROLE_CODE and target_role != TEACHER_ROLE_CODE:
        return RedirectResponse(f"/tai-khoan/{user_id}/sua?status=invalid_role", 303)
    if actor_role in ADMIN_ROLE_CODES and target_role == "SO":
        return RedirectResponse(f"/tai-khoan/{user_id}/sua?status=invalid_role", 303)

    duplicate = db.scalar(
        select(User).where(User.username == username, User.id != user_id)
    )
    if duplicate is not None:
        return RedirectResponse(f"/tai-khoan/{user_id}/sua?status=duplicate", 303)

    fixed_school_id = (
        int(actor["school_id"])
        if actor_role == SCHOOL_ROLE_CODE and actor.get("school_id") is not None
        else None
    )
    school_value = str(fixed_school_id) if fixed_school_id is not None else school_id
    assigned_commune, assigned_school, error = xac_dinh_don_vi_quan_ly(
        db=db,
        vai_tro=role,
        commune_value=commune_id,
        school_value=school_value,
        tai_khoan_hien_tai=account,
        school_id_bat_buoc=fixed_school_id,
    )
    if error is not None:
        return RedirectResponse(f"/tai-khoan/{user_id}/sua?status={error}", 303)

    account.username = username
    account.full_name = full_name
    account.role_id = role.id
    account.commune_id = assigned_commune
    account.school_id = assigned_school
    if mat_khau_moi:
        account.password_hash = ma_hoa_mat_khau(mat_khau_moi)

    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return RedirectResponse(f"/tai-khoan/{user_id}/sua?status=duplicate", 303)
    return RedirectResponse("/tai-khoan?status=updated", 303)


@router.post("/{user_id}/dat-lai-mat-khau")
def dat_lai_mat_khau_tai_khoan(
    user_id: int,
    request: Request,
    mat_khau_moi: Annotated[str, Form()],
    xac_nhan_mat_khau: Annotated[str, Form()],
    return_to: Annotated[str, Form()] = "/tai-khoan",
    db: Session = Depends(get_db),
):
    actor = lay_nguoi_dung(request)
    account = db.scalar(
        select(User)
        .options(selectinload(User.role))
        .where(User.id == user_id)
    )

    if account is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản.")

    if not co_quyen_dat_lai_mat_khau(actor=actor, target=account):
        return chuyen_huong_trang_thai(return_to, "forbidden")

    if (
        len(mat_khau_moi) < MIN_PASSWORD_LENGTH
        or len(mat_khau_moi) > MAX_PASSWORD_LENGTH
        or mat_khau_moi != mat_khau_moi.strip()
        or "\x00" in mat_khau_moi
    ):
        return chuyen_huong_trang_thai(return_to, "invalid_reset_password")

    if mat_khau_moi != xac_nhan_mat_khau:
        return chuyen_huong_trang_thai(return_to, "password_mismatch")

    account.password_hash = ma_hoa_mat_khau(mat_khau_moi)

    try:
        db.commit()
    except Exception:
        db.rollback()
        raise

    return chuyen_huong_trang_thai(return_to, "password_reset")


@router.post("/{user_id}/khoa-mo")
def khoa_hoac_mo_tai_khoan(
    user_id: int,
    request: Request,
    return_to: Annotated[str, Form()] = "/tai-khoan",
    db: Session = Depends(get_db),
):
    actor = lay_nguoi_dung(request)
    account = db.scalar(
        select(User).options(selectinload(User.role)).where(User.id == user_id)
    )
    if account is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy tài khoản.")
    if not tai_khoan_thuoc_pham_vi(actor=actor, target=account):
        return chuyen_huong_trang_thai(return_to, "forbidden")
    if int(actor.get("id") or -1) == account.id:
        return chuyen_huong_trang_thai(return_to, "cannot_self")

    account.is_active = not account.is_active
    db.commit()
    status = "unlocked" if account.is_active else "locked"
    return chuyen_huong_trang_thai(return_to, status)
