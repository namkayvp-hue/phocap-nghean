# -*- coding: utf-8 -*-
from __future__ import annotations

import ast
import importlib.util
import py_compile
import shutil
import sqlite3
import zipfile
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

PROJECT = Path(r"C:\PhoCap")
APP = PROJECT / "app"
DB = PROJECT / "data" / "phocap.db"
EXPORTS = PROJECT / "exports"
HELPER = APP / "household_field_excel_v240.py"

STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")
BACKUP = EXPORTS / f"backup_bai_13b_12_v2_4_1_2a_{STAMP}"
DB_BACKUP = BACKUP / "phocap.db"
REPORT = EXPORTS / f"bao_cao_cai_bai_13b_12_v2_4_1_2a_{STAMP}.txt"

MARKER = "BAI_13B_12_V2_4_1_2A_CLEAN_WORKBOOK_START"

NEW_SELECT_ROWS = 'def _select_rows(\n    tracking: Any,\n    start_row: int,\n    end_row: int,\n    *,\n    extra_blank_rows: int = 5,\n) -> list[int]:\n    """Giữ đúng số trang cần dùng; không sinh trang mới chỉ vì dòng dự phòng."""\n    existing: list[int] = []\n\n    for row in range(start_row, end_row + 1):\n        full_name = " ".join(\n            part\n            for part in (\n                _text(tracking.cell(row, 3).value),\n                _text(tracking.cell(row, 4).value),\n            )\n            if part\n        ).strip()\n        if full_name:\n            existing.append(row)\n\n    if existing:\n        used_rows = max(existing) - start_row + 1\n    else:\n        used_rows = 0\n\n    # Mỗi trang phiếu có 7 vị trí thành viên.\n    # Chỉ bù chỗ trống tới cuối trang đang dùng.\n    slot_count = max(7, ((max(1, used_rows) + 6) // 7) * 7)\n    stop = min(end_row, start_row + slot_count - 1)\n    return list(range(start_row, stop + 1))\n\n\n'
NEW_ADD_FUNCTION = 'def add_field_form_view_to_bytes(content: bytes) -> bytes:\n    """Dựng workbook sạch: Phiếu điều tra + các trang kỹ thuật chỉ chứa giá trị."""\n    # === BAI_13B_12_V2_4_1_2A_CLEAN_WORKBOOK_START ===\n    source_workbook = load_workbook(\n        BytesIO(content),\n        data_only=False,\n        keep_links=False,\n    )\n    value_workbook = load_workbook(\n        BytesIO(content),\n        data_only=True,\n        keep_links=False,\n    )\n\n    source_tracking = _find_tracking(source_workbook)\n    value_tracking = _find_tracking(value_workbook)\n\n    if META_SHEET_NAME not in source_workbook.sheetnames:\n        raise ValueError("File thiếu trang kỹ thuật ẩn.")\n    if META_SHEET_NAME not in value_workbook.sheetnames:\n        raise ValueError("File thiếu trang kỹ thuật ẩn.")\n\n    source_meta = source_workbook[META_SHEET_NAME]\n    value_meta = value_workbook[META_SHEET_NAME]\n\n    meta = _meta_values(source_meta)\n    specs = _form_specs(\n        source_workbook,\n        source_tracking,\n        meta,\n    )\n    if not specs:\n        raise ValueError(\n            "Không xác định được phiếu để dựng mẫu Excel điều tra."\n        )\n\n    # Tạo workbook hoàn toàn mới để không mang external links /\n    # defined names / relationship lỗi từ workbook mẫu.\n    workbook = Workbook()\n    default_sheet = workbook.active\n    workbook.remove(default_sheet)\n\n    tracking = workbook.create_sheet(source_tracking.title)\n    meta_sheet = workbook.create_sheet(META_SHEET_NAME)\n\n    def _copy_value_sheet(\n        source_sheet: Any,\n        value_sheet: Any,\n        target_sheet: Any,\n    ) -> None:\n        max_row = max(source_sheet.max_row, value_sheet.max_row)\n        max_col = max(source_sheet.max_column, value_sheet.max_column)\n\n        for row in range(1, max_row + 1):\n            for col in range(1, max_col + 1):\n                raw = source_sheet.cell(row, col).value\n                if isinstance(raw, str) and raw.startswith("="):\n                    raw = value_sheet.cell(row, col).value\n                target_sheet.cell(row, col).value = raw\n\n    _copy_value_sheet(\n        source_tracking,\n        value_tracking,\n        tracking,\n    )\n    _copy_value_sheet(\n        source_meta,\n        value_meta,\n        meta_sheet,\n    )\n    # === BAI_13B_12_V2_4_1_2A_CLEAN_WORKBOOK_END ===\n\n    ws = workbook.create_sheet(FIELD_SHEET_NAME, 0)\n    map_sheet = workbook.create_sheet(FIELD_MAP_SHEET_NAME)\n    map_sheet.sheet_state = "veryHidden"\n\n    widths = {\n        "A": 4.5,\n        "B": 29,\n        "C": 19,\n        "D": 28,\n        "E": 18,\n        "F": 12,\n        "G": 12,\n        "H": 14,\n        "I": 13,\n        "J": 18,\n        "K": 19,\n    }\n    for letter, width in widths.items():\n        ws.column_dimensions[letter].width = width\n\n    current_row = 1\n    map_row = 2\n\n    for spec_index, spec in enumerate(specs, start=1):\n        tracking_rows = list(spec.get("tracking_rows") or [])\n        pages = [\n            tracking_rows[i:i + 7]\n            for i in range(0, len(tracking_rows), 7)\n        ] or [[]]\n\n        for page_index, page_rows in enumerate(pages, start=1):\n            header_row = current_row\n            spec["visible_header_row"] = header_row\n\n            current_row = _fill_header(\n                ws,\n                current_row,\n                spec=spec,\n                page_index=page_index,\n                page_total=len(pages),\n            )\n\n            # _select_rows đã bù đủ 7 vị trí/trang nên mọi slot đều có\n            # tracking_row kỹ thuật tương ứng, kể cả slot trống để nhập phát sinh.\n            for slot_index in range(7):\n                if slot_index < len(page_rows):\n                    tracking_row = page_rows[slot_index]\n                    current_row = _fill_person_slot(\n                        ws,\n                        current_row,\n                        tracking=tracking,\n                        spec=spec,\n                        tracking_row=tracking_row,\n                        slot_number=(\n                            (page_index - 1) * 7\n                            + slot_index\n                            + 1\n                        ),\n                        map_sheet=map_sheet,\n                        map_row=map_row,\n                    )\n                    map_row += 1\n                else:\n                    current_row = _fill_blank_slot(\n                        ws,\n                        current_row,\n                        slot_number=(\n                            (page_index - 1) * 7\n                            + slot_index\n                            + 1\n                        ),\n                    )\n\n            if page_index == len(pages):\n                current_row = _fill_signature(\n                    ws,\n                    current_row,\n                    str(spec.get("head_name") or ""),\n                )\n\n            footer_row = current_row\n            ws.merge_cells(\n                start_row=footer_row,\n                start_column=1,\n                end_row=footer_row,\n                end_column=11,\n            )\n            ws.cell(footer_row, 1).value = (\n                f"{meta.get(\'BATCH_CODE\') or \'\'} · "\n                f"{spec.get(\'form_number\') or \'\'} · "\n                f"{spec.get(\'household_code\') or \'\'}"\n            )\n            ws.cell(footer_row, 1).font = Font(\n                name="Times New Roman",\n                size=6.5,\n                italic=True,\n            )\n            current_row += 1\n\n            is_last_page = (\n                spec_index == len(specs)\n                and page_index == len(pages)\n            )\n            if not is_last_page:\n                ws.row_breaks.append(Break(id=current_row - 1))\n\n            current_row += 1\n\n    ws.sheet_view.showGridLines = False\n    ws.freeze_panes = None\n    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE\n    ws.page_setup.paperSize = ws.PAPERSIZE_A4\n    ws.page_setup.fitToWidth = 1\n    ws.page_setup.fitToHeight = 0\n    ws.sheet_properties.pageSetUpPr.fitToPage = True\n    ws.page_margins = PageMargins(\n        left=0.12,\n        right=0.12,\n        top=0.18,\n        bottom=0.18,\n        header=0.0,\n        footer=0.0,\n    )\n    ws.print_options.horizontalCentered = True\n    ws.print_area = f"A1:K{max(1, current_row - 1)}"\n\n    tracking.sheet_state = "veryHidden"\n    meta_sheet.sheet_state = "veryHidden"\n    map_sheet.sheet_state = "veryHidden"\n    ws.sheet_state = "visible"\n    workbook.active = 0\n\n    try:\n        workbook.calculation.fullCalcOnLoad = False\n        workbook.calculation.forceFullCalc = False\n    except Exception:\n        pass\n\n    output = BytesIO()\n    workbook.save(output)\n    return output.getvalue()\n\n\n'


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig", errors="replace")


def write_text(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8")


def db_state() -> dict:
    con = sqlite3.connect(str(DB))
    try:
        result = {
            "integrity": str(
                con.execute("PRAGMA integrity_check").fetchone()[0]
            ),
            "fk_count": len(
                con.execute("PRAGMA foreign_key_check").fetchall()
            ),
        }
        for table in (
            "survey_forms",
            "survey_people",
            "survey_person_year_records",
        ):
            result[table] = int(
                con.execute(
                    f"SELECT COUNT(*) FROM {table}"
                ).fetchone()[0]
            )
        return result
    finally:
        con.close()


def backup_file(path: Path) -> None:
    target = BACKUP / path.relative_to(PROJECT)
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, target)


def restore_file(path: Path) -> None:
    source = BACKUP / path.relative_to(PROJECT)
    if source.exists():
        shutil.copy2(source, path)


def backup_db() -> None:
    src = sqlite3.connect(str(DB))
    dst = sqlite3.connect(str(DB_BACKUP))
    try:
        src.backup(dst)
        dst.commit()
    finally:
        dst.close()
        src.close()


def restore_db() -> None:
    if not DB_BACKUP.exists():
        return
    src = sqlite3.connect(str(DB_BACKUP))
    dst = sqlite3.connect(str(DB))
    try:
        src.backup(dst)
        dst.commit()
    finally:
        dst.close()
        src.close()


def get_function_node(tree: ast.Module, name: str):
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            if node.name == name:
                return node
    raise RuntimeError(
        f"Không xác định được hàm {name} bằng AST."
    )


def replace_function_ast(
    source: str,
    func_name: str,
    replacement: str,
) -> str:
    tree = ast.parse(source)
    node = get_function_node(tree, func_name)

    lines = source.splitlines(keepends=True)
    start = node.lineno - 1
    end = int(getattr(node, "end_lineno", node.lineno) or node.lineno)

    replacement_text = replacement
    if not replacement_text.endswith("\n"):
        replacement_text += "\n"

    new_lines = (
        lines[:start]
        + [replacement_text]
        + lines[end:]
    )
    result = "".join(new_lines)
    ast.parse(result)
    return result


def ensure_workbook_import(source: str) -> str:
    tree = ast.parse(source)

    for node in tree.body:
        if isinstance(node, ast.ImportFrom) and node.module == "openpyxl":
            names = {alias.name for alias in node.names}
            if "Workbook" in names:
                return source

    # Thêm import riêng ngay sau __future__ để không phụ thuộc format import cũ.
    lines = source.splitlines(keepends=True)
    insert_at = 0
    for index, line in enumerate(lines):
        if line.startswith("from __future__ import "):
            insert_at = index + 1
            break

    lines.insert(insert_at, "from openpyxl import Workbook\n")
    result = "".join(lines)
    ast.parse(result)
    return result


def patch_helper(source: str) -> str:
    if MARKER in source:
        return source

    required = (
        'FIELD_VERSION = "PCGDMN_FIELD_FORM_V240"',
        "def _select_rows",
        "def add_field_form_view_to_bytes",
        "def sync_field_form_to_tracking",
    )
    for token in required:
        if token not in source:
            raise RuntimeError(
                "Thiếu nền V2.4.1/V2.4.1.1A: " + token
            )

    source = ensure_workbook_import(source)
    source = replace_function_ast(
        source,
        "_select_rows",
        NEW_SELECT_ROWS,
    )
    source = replace_function_ast(
        source,
        "add_field_form_view_to_bytes",
        NEW_ADD_FUNCTION,
    )
    ast.parse(source)
    return source


def load_helper_module():
    spec = importlib.util.spec_from_file_location(
        "household_field_excel_v2412a_test",
        HELPER,
    )
    if spec is None or spec.loader is None:
        raise RuntimeError("Không nạp được helper V2.4.1.2A.")

    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def build_smoke_source() -> bytes:
    wb = Workbook()
    tracking = wb.active
    tracking.title = "Sổ theo dõi PCGDMN"

    tracking["A1"] = "PHIẾU ĐIỀU TRA PHỔ CẬP GIÁO DỤC"
    tracking["A2"] = (
        "Xóm/khối: Test-Xóm1; Xã/phường: Xã Test; "
        "Tỉnh: Nghệ An; Năm học: 2026-2027"
    )
    tracking["A3"] = "Họ và tên chủ hộ: TEST CHỦ HỘ"
    tracking["K3"] = "Địa chỉ: Test-Xóm1, số 01"

    for col, start_year in zip(range(12, 19), range(2020, 2027)):
        tracking.cell(7, col).value = start_year
        tracking.cell(8, col).value = start_year + 1

    for idx in range(4):
        row = 11 + idx
        tracking.cell(row, 3).value = f"TEST Người {idx + 1}"
        tracking.cell(row, 5).value = f"0{idx + 1}/01/2010"
        tracking.cell(row, 6).value = "Nam"
        tracking.cell(row, 7).value = "Test-Xóm1"
        tracking.cell(row, 8).value = "Xã Test"
        tracking.cell(row, 9).value = "Kinh"
        tracking.cell(row, 10).value = "Thành viên"
        tracking.cell(row, 18).value = str(idx + 1)
        tracking.cell(row, 19).value = "Trường Test"
        tracking.cell(row, 20).value = (
            "Cấp/nhóm: TEST; Nhập từ Excel IMP-TEST; "
            "Ghi chú nghiệp vụ"
        )

    meta = wb.create_sheet("_PCGDMN_META")
    values = [
        ("VERSION", "PCGDMN_HO_V1"),
        ("BATCH_ID", 1),
        ("BATCH_CODE", "DT-TEST"),
        ("FORM_ID", 1),
        ("FORM_NUMBER", "TEST-001"),
        ("HOUSEHOLD_ID", 1),
        ("HOUSEHOLD_CODE", "HO-TEST"),
        ("HOUSEHOLD_HEAD_NAME", "TEST CHỦ HỘ"),
        ("COMMUNE_NAME", "Xã Test"),
        ("SCHOOL_YEAR_ID", 2),
        ("SCHOOL_YEAR_CODE", "2026-2027"),
        ("CURRENT_YEAR_COLUMN", 18),
    ]
    for row, (key, value) in enumerate(values, start=1):
        meta.cell(row, 1).value = key
        meta.cell(row, 2).value = value

    for col, header in enumerate(
        ("visible_row", "person_id", "person_code"),
        start=1,
    ):
        meta.cell(20, col).value = header

    for idx in range(4):
        meta.cell(21 + idx, 1).value = 11 + idx
        meta.cell(21 + idx, 2).value = 100 + idx
        meta.cell(21 + idx, 3).value = f"DT-TEST-{idx + 1}"

    tracking["Z1"] = "=[External.xlsx]Sheet1!A1"

    raw = BytesIO()
    wb.save(raw)
    return raw.getvalue()


def smoke_test() -> None:
    module = load_helper_module()
    exported = module.add_field_form_view_to_bytes(
        build_smoke_source()
    )

    with zipfile.ZipFile(BytesIO(exported), "r") as zf:
        bad = zf.testzip()
        if bad is not None:
            raise RuntimeError(
                "Smoke test ZIP lỗi tại: " + str(bad)
            )

    wb = load_workbook(
        BytesIO(exported),
        data_only=False,
        keep_links=False,
    )

    required_sheets = (
        "Phiếu điều tra",
        "Sổ theo dõi PCGDMN",
        "_PCGDMN_META",
        "_PCGDMN_FIELD_MAP",
    )
    for name in required_sheets:
        if name not in wb.sheetnames:
            raise RuntimeError(
                "Smoke test thiếu sheet: " + name
            )

    if wb["Phiếu điều tra"].sheet_state != "visible":
        raise RuntimeError(
            "Phiếu điều tra không visible."
        )

    for name in required_sheets[1:]:
        if wb[name].sheet_state != "veryHidden":
            raise RuntimeError(
                name + " chưa veryHidden."
            )

    ws = wb["Phiếu điều tra"]
    all_text = "\n".join(
        str(cell.value or "")
        for row in ws.iter_rows()
        for cell in row
    )

    if "Trang 1/2" in all_text:
        raise RuntimeError(
            "Hộ 4 người vẫn bị tạo 2 trang."
        )

    if "Cấp/nhóm:" in all_text or "IMP-TEST" in all_text:
        raise RuntimeError(
            "Ghi chú kỹ thuật vẫn bị lộ."
        )

    if "Ghi chú nghiệp vụ" not in all_text:
        raise RuntimeError(
            "Ghi chú nghiệp vụ bị mất."
        )

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if (
                    isinstance(value, str)
                    and value.startswith("=")
                    and "[" in value
                    and "]" in value
                ):
                    raise RuntimeError(
                        "Vẫn còn công thức external."
                    )


def verify_source(source: str) -> None:
    for token in (
        MARKER,
        "workbook = Workbook()",
        "source_workbook = load_workbook(",
        "value_workbook = load_workbook(",
        "slot_count = max(7",
        "is_last_page",
    ):
        if token not in source:
            raise RuntimeError(
                "Verifier source thiếu: " + token
            )

    ast.parse(source)
    py_compile.compile(str(HELPER), doraise=True)


def main() -> int:
    print("=" * 124)
    print(
        "BÀI 13B-12 V2.4.1.2A - "
        "ỔN ĐỊNH FILE EXCEL / BỎ REPAIRED / "
        "KHÔNG TẠO TRANG TRẮNG THỪA"
    )
    print("=" * 124)

    for path in (DB, HELPER):
        if not path.exists():
            print("DỪNG AN TOÀN: thiếu", path)
            return 2

    before = db_state()
    print("integrity_check:", before["integrity"])
    print("foreign_key_check:", before["fk_count"], "lỗi")

    if before["integrity"].lower() != "ok" or before["fk_count"] != 0:
        print("DỪNG: database chưa đạt kiểm tra an toàn.")
        return 3

    source = read_text(HELPER)

    if MARKER in source:
        print("V2.4.1.2A đã cài. Không cài lặp.")
        return 0

    try:
        patched = patch_helper(source)
        ast.parse(patched)
    except Exception as exc:
        print()
        print("DỪNG AN TOÀN TRƯỚC KHI GHI SOURCE:")
        print(type(exc).__name__ + ":", exc)
        return 4

    EXPORTS.mkdir(parents=True, exist_ok=True)
    BACKUP.mkdir(parents=True, exist_ok=False)
    backup_file(HELPER)
    backup_db()
    print("Backup:", BACKUP)

    try:
        write_text(HELPER, patched)
        verify_source(read_text(HELPER))
        smoke_test()

        after = db_state()
        if after["integrity"].lower() != "ok":
            raise RuntimeError(
                "integrity_check không đạt sau cài."
            )
        if after["fk_count"] != 0:
            raise RuntimeError(
                "foreign_key_check có lỗi sau cài."
            )

        for table in (
            "survey_forms",
            "survey_people",
            "survey_person_year_records",
        ):
            if after[table] != before[table]:
                raise RuntimeError(
                    f"Số bản ghi {table} thay đổi ngoài dự kiến."
                )

        for cache in APP.rglob("__pycache__"):
            if cache.is_dir():
                shutil.rmtree(
                    cache,
                    ignore_errors=True,
                )

    except Exception as exc:
        print()
        print("CÓ LỖI - ĐANG ROLLBACK SOURCE + DATABASE...")
        print(type(exc).__name__ + ":", exc)
        restore_file(HELPER)
        restore_db()
        print("ĐÃ KHÔI PHỤC.")
        return 9

    REPORT.write_text(
        "\n".join(
            [
                "=" * 124,
                "BÁO CÁO CÀI BÀI 13B-12 V2.4.1.2A",
                "=" * 124,
                f"Thời gian: {datetime.now():%d/%m/%Y %H:%M:%S}",
                "",
                "ĐÃ SỬA:",
                "- Thay hàm bằng AST, không phụ thuộc cách xuống dòng của source.",
                "- Tạo workbook mới sạch, không mang external relationship cũ.",
                "- Hộ <= 7 người chỉ 1 trang; >7 người mới có trang tiếp.",
                "- Không tạo trang mới chỉ vì dòng dự phòng.",
                "- Vẫn giữ các slot trống trong trang hiện tại để nhập phát sinh.",
                "- Giữ cơ chế 2.2.3 xuất -> chính file đó 2.2.4 nhập.",
                "",
                "SMOKE TEST: ĐẠT",
                f"integrity_check: {after['integrity']}",
                f"foreign_key_check: {after['fk_count']} lỗi",
                f"Backup: {BACKUP}",
            ]
        ) + "\n",
        encoding="utf-8-sig",
    )

    print()
    print("=" * 124)
    print("CÀI ĐẶT THÀNH CÔNG BÀI 13B-12 V2.4.1.2A")
    print("=" * 124)
    print(" - Thay hàm bằng AST: ĐẠT.")
    print(" - Smoke test file Excel: ĐẠT.")
    print(" - Không sửa dữ liệu hiện có.")
    print("Báo cáo:", REPORT)
    print("Backup:", BACKUP)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
